#!/usr/bin/env python3
"""Генератор второго набора примеров — «Университет: приём, наука и результаты 2026».

Все данные вымышленные. Случайность — только `random.Random(SEED)`, поэтому файлы воспроизводятся один в один.
Запуск (из корня репозитория):

    PYTHONPATH="$PWD" .venv/bin/python examples/data-science-2026/source/build.py

Скрипт пишет 6 файлов в папку `examples/data-science-2026/` и печатает числа для SCENARIO.md. Никакой код анализа
(`verstka.data`) он не вызывает и от него не зависит: данные описаны так, как их выгрузила бы настоящая учебная часть.
"""

from __future__ import annotations

import json
import math
import random
import statistics
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SEED = 20261016
OUT = Path(__file__).resolve().parent.parent
NBSP = " "


def fmt_ru(x: float, digits: int = 1, nbsp: bool = True) -> str:
    """1234.5 → «1 234,5» (разделитель тысяч — неразрывный пробел, как в выгрузке из Excel)."""
    s = f"{x:,.{digits}f}".replace(",", NBSP if nbsp else " ").replace(".", ",")
    return s


def neg_paren(x: float, digits: int = 1) -> str:
    """Бухгалтерская запись: минус — в скобках, «(6,3)»."""
    return f"({fmt_ru(abs(x), digits)})" if x < 0 else fmt_ru(x, digits)


# ───────────────────────────── 1. Часы подготовки и баллы ЕГЭ ─────────────────────────────

PROFILES = {
    # профиль: (доля в выборке, балл при нуле часов, прибавка за час в неделю, разброс)
    "Технический": (0.46, 49.0, 2.15, 6.8),
    "Естественнонаучный": (0.30, 46.0, 2.45, 7.2),
    "Гуманитарный": (0.24, 55.0, 1.85, 6.4),
}


def build_ege(rng: random.Random) -> dict:
    n = 240
    names = list(PROFILES)
    weights = [PROFILES[k][0] for k in names]
    rows = []
    for i in range(n):
        prof = rng.choices(names, weights)[0]
        _, base, slope, sigma = PROFILES[prof]
        hours = min(max(rng.gauss(11.0, 5.2), 1.0), 26.0)
        hours = round(hours * 2) / 2
        score = base + slope * hours + rng.gauss(0, sigma)
        score = round(min(max(score, 38.0), 100.0), 1)
        rows.append([f"АБ-26{1000 + i * 3 + rng.randint(0, 2)}", prof, hours, score])
    # у семи абитуриентов часы в анкете не указаны («н/д»)
    gaps = set(rng.sample(range(n), 7))
    lines = ["Абитуриент,Профиль,Подготовка (ч в неделю),Средний балл ЕГЭ"]
    pairs = []
    for i, (aid, prof, h, s) in enumerate(rows):
        hs = "н/д" if i in gaps else f"{h:g}"
        lines.append(f"{aid},{prof},{hs},{s:.1f}")
        if i not in gaps:
            pairs.append((h, s))
    (OUT / "Часы подготовки и баллы ЕГЭ.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")

    xs = [p[0] for p in pairs]
    ys = [p[1] for p in pairs]
    mx, my = statistics.fmean(xs), statistics.fmean(ys)
    sxx = sum((x - mx) ** 2 for x in xs)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    syy = sum((y - my) ** 2 for y in ys)
    slope = sxy / sxx
    return {
        "rows": n, "with_hours": len(pairs), "r": sxy / math.sqrt(sxx * syy), "slope": slope,
        "intercept": my - slope * mx, "mean_hours": mx, "mean_score": statistics.fmean(s for *_, s in rows),
        "mean_score_pairs": my, "min_hours": min(xs), "max_hours": max(xs), "min_score": min(ys), "max_score": max(ys),
        "by_profile": {k: sum(1 for r in rows if r[1] == k) for k in names},
        "score_low": statistics.fmean(y for x, y in pairs if x < 8), "score_high": statistics.fmean(y for x, y in pairs if x >= 16),
        "n_low": sum(1 for x, _ in pairs if x < 8), "n_high": sum(1 for x, _ in pairs if x >= 16),
    }


# ───────────────────────────── 2. Приём по направлениям ─────────────────────────────

ADMISSION = [
    # направление, бюджет, платное, целевое (None — мест по целевому приёму нет, в ведомости «—»)
    ("Прикладная информатика", 120, 140, 25),
    ("Программная инженерия", 110, 150, 20),
    ("Информационная безопасность", 90, 60, 45),
    ("Электроэнергетика и электротехника", 85, 40, 70),
    ("Машиностроение", 100, 55, 60),
    ("Строительство", 80, 120, 15),
    ("Экономика", 45, 190, None),
    ("Менеджмент", 35, 165, 5),
    ("Юриспруденция", 40, 210, None),
    ("Биотехнологии", 50, 30, 20),
    ("Лингвистика", 25, 85, None),
    ("Педагогическое образование", 60, 45, 35),
]


def build_admission() -> dict:
    wb = Workbook()
    wb.properties.creator = "Приёмная комиссия"
    wb.properties.created = datetime(2026, 8, 5, 12, 0, 0)
    wb.properties.modified = datetime(2026, 8, 5, 12, 0, 0)
    ws = wb.active
    ws.title = "Приём 2026"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = "Приём на первый курс в 2026 году по направлениям подготовки, человек"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:E1")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    # строка 2 пустая, шапка — строки 3–4
    ws["A3"] = "Направление подготовки"
    ws.merge_cells("A3:A4")
    ws["B3"] = "Форма обучения"
    ws.merge_cells("B3:D3")
    ws["B4"], ws["C4"], ws["D4"] = "Бюджет", "Платное", "Целевое"
    ws["E3"] = "Всего"
    ws.merge_cells("E3:E4")
    for row in ws["A3:E4"]:
        for c in row:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = center
            c.border = box
    totals = [0, 0, 0, 0]
    for k, (name, b, p, t) in enumerate(ADMISSION):
        r = 5 + k
        ws.cell(r, 1, name).border = box
        ws.cell(r, 2, b)
        ws.cell(r, 3, p)
        ws.cell(r, 4, t if t is not None else "—")
        ws.cell(r, 5, b + p + (t or 0))
        for j in range(2, 6):
            ws.cell(r, j).border = box
            ws.cell(r, j).alignment = Alignment(horizontal="right")
        for j, v in enumerate((b, p, t or 0, b + p + (t or 0))):
            totals[j] += v
    ws.column_dimensions["A"].width = 38
    for col in "BCDE":
        ws.column_dimensions[col].width = 12
    wb.save(OUT / "Приём по направлениям.xlsx")
    top = max(ADMISSION, key=lambda r: r[1] + r[2] + (r[3] or 0))
    return {
        "rows": len(ADMISSION), "budget": totals[0], "paid": totals[1], "target": totals[2], "total": totals[3],
        "share_budget": totals[0] / totals[3], "share_paid": totals[1] / totals[3], "share_target": totals[2] / totals[3],
        "top": top[0], "top_total": top[1] + top[2] + (top[3] or 0),
        "no_target": [r[0] for r in ADMISSION if r[3] is None],
        "top_budget": max(ADMISSION, key=lambda r: r[1])[0],
        "top_paid_share": max(ADMISSION, key=lambda r: r[2] / (r[1] + r[2] + (r[3] or 0))),
    }


# ───────────────────────────── 3. Публикации по годам ─────────────────────────────

MISSING_YEAR = 2016  # за этот год выгрузку из Scopus не сохранили


def build_publications(rng: random.Random) -> dict:
    years, scopus, rinc = [], [], []
    for y in range(2005, 2026):
        if y == MISSING_YEAR:
            continue
        t = y - 2005
        s = 41 * math.exp(0.113 * t) * rng.uniform(0.93, 1.07)
        r = 260 + 740 * (1 - math.exp(-0.19 * t))
        if y > 2021:
            r *= 1 - 0.022 * (y - 2021)
        r *= rng.uniform(0.96, 1.04)
        years.append(y)
        scopus.append(round(s))
        rinc.append(round(r))
    obj = {"Год": years, "Scopus, статей": scopus, "РИНЦ, статей": rinc}
    (OUT / "Публикации по годам.json").write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    peak = max(range(len(rinc)), key=rinc.__getitem__)
    return {
        "rows": len(years), "first_year": years[0], "last_year": years[-1], "missing": MISSING_YEAR,
        "scopus_first": scopus[0], "scopus_last": scopus[-1], "scopus_x": scopus[-1] / scopus[0],
        "rinc_first": rinc[0], "rinc_last": rinc[-1], "rinc_x": rinc[-1] / rinc[0],
        "rinc_peak": rinc[peak], "rinc_peak_year": years[peak],
        "scopus_2024": scopus[-2], "scopus_growth_last": scopus[-1] / scopus[-2] - 1,
        "series": list(zip(years, scopus, rinc)),
    }


# ───────────────────────────── 4. Бюджет НИОКР ─────────────────────────────

RND = [
    # институт, кафедра, бюджет 2026 (млн ₽), изменение к 2025 (млн ₽)
    ("Институт информационных технологий", "Кафедра искусственного интеллекта", 142.6, 31.4),
    ("Институт информационных технологий", "Кафедра кибербезопасности", 96.2, 22.8),
    ("Институт информационных технологий", "Кафедра программной инженерии", 71.5, 4.1),
    ("Институт информационных технологий", "Кафедра сетевых технологий", 38.9, -6.3),
    ("Институт энергетики", "Кафедра электроэнергетики", 118.4, 9.7),
    ("Институт энергетики", "Кафедра теплоэнергетики", 64.0, -3.2),
    ("Институт энергетики", "Кафедра возобновляемых источников энергии", 87.3, 25.5),
    ("Институт машиностроения и материалов", "Кафедра технологий машиностроения", 79.8, -5.6),
    ("Институт машиностроения и материалов", "Кафедра материаловедения", 103.7, 12.2),
    ("Институт машиностроения и материалов", "Кафедра робототехники и мехатроники", 91.4, 18.9),
    ("Институт машиностроения и материалов", "Кафедра аддитивных технологий", 55.2, 14.0),
    ("Строительный институт", "Кафедра строительных конструкций", 52.1, -1.5),
    ("Строительный институт", "Кафедра геотехники", 33.6, 0.8),
    ("Строительный институт", "Кафедра архитектуры", 21.7, -2.4),
    ("Институт естественных наук", "Кафедра физики", 74.5, 3.3),
    ("Институт естественных наук", "Кафедра химии", 66.8, -0.9),
    ("Институт естественных наук", "Кафедра прикладной математики", 48.3, 7.6),
    ("Институт естественных наук", "Кафедра биотехнологии", 59.2, 11.1),
]


def build_rnd() -> dict:
    total = round(sum(r[2] for r in RND), 1)
    delta = round(sum(r[3] for r in RND), 1)
    lines = ["Институт;Кафедра;Бюджет 2026, млн руб.;Изменение к 2025, млн руб."]
    for inst, dep, v, d in RND:
        lines.append(f"{inst};{dep};{fmt_ru(v)};{neg_paren(d)}")
    lines.append(f"Итого;;{fmt_ru(total)};{neg_paren(delta)}")
    # Windows-1251, как выгружает 1С/Excel с русской локалью; CRLF. Знака ₽ в cp1251 нет, поэтому «млн руб.»
    (OUT / "Бюджет НИОКР.csv").write_bytes(("\r\n".join(lines) + "\r\n").encode("cp1251"))
    by_inst: dict[str, float] = {}
    by_inst_delta: dict[str, float] = {}
    for inst, _, v, d in RND:
        by_inst[inst] = by_inst.get(inst, 0.0) + v
        by_inst_delta[inst] = by_inst_delta.get(inst, 0.0) + d
    lead = max(by_inst, key=by_inst.get)
    dep_top = max(RND, key=lambda r: r[2])
    return {
        "rows": len(RND), "total": total, "delta": delta, "prev": total - delta, "growth": delta / (total - delta),
        "by_inst": {k: round(v, 1) for k, v in by_inst.items()},
        "by_inst_delta": {k: round(v, 1) for k, v in by_inst_delta.items()},
        "lead": lead, "lead_sum": by_inst[lead], "lead_share": by_inst[lead] / total,
        "dep_top": dep_top[1], "dep_top_sum": dep_top[2], "dep_top_share": dep_top[2] / total,
        "falling": [r[1] for r in RND if r[3] < 0], "top3_share": sum(sorted((r[2] for r in RND), reverse=True)[:3]) / total,
        "top_growth": max(RND, key=lambda r: r[3])[1:],
    }


# ───────────────────────────── 5. Стипендии ─────────────────────────────

ASSIGNED = [("01.09.2025", 0.52), ("01.02.2026", 0.30), ("01.09.2026", 0.18)]


def build_stipends(rng: random.Random) -> dict:
    n = 420
    kinds = ["base", "social", "raised", "named"]
    weights = [0.57, 0.15, 0.24, 0.04]
    amounts = []
    lines = ["Студент;Назначена с;Размер стипендии, руб."]
    ids = rng.sample(range(100, 9999), n)
    for i in range(n):
        k = rng.choices(kinds, weights)[0]
        if k == "base":
            a = 2590 + (rng.choice([500, 1000, 1500]) if rng.random() < 0.25 else 0)
        elif k == "social":
            a = 4120
        elif k == "raised":
            a = rng.choice([6500, 8000, 8000, 10000, 12500])
        else:
            a = rng.choice([20000, 25000, 30000, 40000, 55000])
        d = rng.choices([x for x, _ in ASSIGNED], [w for _, w in ASSIGNED])[0]
        amounts.append(a)
        lines.append(f"СТ-26-{ids[i]:04d};{d};{a}")
    (OUT / "Стипендии.csv").write_text("\n".join(lines) + "\n", encoding="utf-8-sig")
    srt = sorted(amounts)
    q1, q2, q3 = (statistics.quantiles(srt, n=4, method="inclusive"))
    fence = q3 + 1.5 * (q3 - q1)
    p90 = statistics.quantiles(srt, n=10, method="inclusive")[8]
    return {
        "rows": n, "mean": statistics.fmean(amounts), "median": q2, "q1": q1, "q3": q3, "p90": p90, "max": srt[-1],
        "min": srt[0], "fence": fence, "outliers": sum(1 for a in amounts if a > fence),
        "over_30k": sum(1 for a in amounts if a >= 30000), "share_below_3k": sum(1 for a in amounts if a < 3000) / n,
        "share_up_to_median": sum(1 for a in amounts if a <= q2) / n,
        "top10_share": sum(srt[-int(n * 0.1):]) / sum(amounts),
    }


# ───────────────────────────── 6. Итоги года ─────────────────────────────

def build_summary(ege: dict, adm: dict, rnd: dict) -> dict:
    admitted, admitted_prev = adm["total"], 2288
    score, score_prev = round(ege["mean_score"], 1), 72.9
    employed, employed_prev = 87.0, 84.0
    index, index_prev = 78, 74
    expelled, expelled_prev = 412, 450
    rnd_now, rnd_prev = rnd["total"], rnd["prev"]

    def pct(a: float, b: float) -> str:
        v = (a / b - 1) * 100
        return f"{v:+.0f}%".replace("-", "−")

    rows = [
        ("Принято на 1 курс, чел.", f"{admitted:,}".replace(",", " "), f"{admitted_prev:,}".replace(",", " "), pct(admitted, admitted_prev)),
        ("Средний балл ЕГЭ", fmt_ru(score), fmt_ru(score_prev), f"+{fmt_ru(score - score_prev)}"),
        ("Трудоустроено выпускников, %", f"{employed:.0f}%", f"{employed_prev:.0f}%", f"+{employed - employed_prev:.0f} п.п."),
        ("Индекс удовлетворённости студентов", f"{index}", f"{index_prev}", f"+{index - index_prev} п."),
        ("Объём НИОКР, млн ₽", fmt_ru(rnd_now, 0, nbsp=False), fmt_ru(rnd_prev, 0, nbsp=False), pct(rnd_now, rnd_prev)),
        ("Отчислено студентов, чел.", f"{expelled}", f"{expelled_prev}", pct(expelled, expelled_prev)),
    ]
    adm_s = f"{admitted:,}".replace(",", " ")
    text = (
        "Итоги 2026 года — университет, для учёного совета\n"
        "\n"
        "Подготовлено 14 октября 2026 года. Данные — по приказам о зачислении, анкете первокурсников, ведомостям "
        "научной части и опросу выпускников, сравнение с 2025 годом.\n"
        "\n"
        f"На первый курс принято {adm_s} человек — на {(admitted / admitted_prev - 1) * 100:.0f}% больше, чем годом ранее. "
        f"Средний балл ЕГЭ поступивших вырос до {fmt_ru(score)}: по анкете первокурсников, чем больше часов в неделю "
        "абитуриент готовился, тем выше его балл. Трудоустроенных выпускников стало "
        f"{employed:.0f}%, а индекс удовлетворённости студентов вырос до {index} пунктов. Объём НИОКР составил "
        f"{fmt_ru(rnd_now, 0, nbsp=False)} млн ₽, отчислений стало меньше на {(1 - expelled / expelled_prev) * 100:.0f}%.\n"
        "\n"
        "| Показатель | 2026 | 2025 | Изменение |\n"
        "|---|---|---|---|\n"
    )
    for r in rows:
        text += "| " + " | ".join(r) + " |\n"
    text += (
        "\n"
        "Публикации в Scopus за 2016 год в базу не попали: выгрузку не сохранили, поэтому ряд по годам идёт с пропуском. "
        "Стипендии назначены трём потокам студентов — с 1 сентября 2025 года, с 1 февраля и с 1 сентября 2026 года.\n"
        "\n"
        "Что дальше: сохранить рост приёма на бюджетные места ИТ-направлений и перераспределить НИОКР в пользу "
        "кафедр, где объём заметно падает.\n"
    )
    (OUT / "Итоги года.txt").write_text(text, encoding="utf-8")
    return {"rows": rows}


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)
    ege = build_ege(rng)
    adm = build_admission()
    pub = build_publications(rng)
    rnd = build_rnd()
    stip = build_stipends(rng)
    summ = build_summary(ege, adm, rnd)
    for name, d in (("ege", ege), ("admission", adm), ("publications", pub), ("rnd", rnd), ("stipends", stip), ("summary", summ)):
        print(f"--- {name}")
        for k, v in d.items():
            if k == "series":
                continue
            print(f"  {k}: {v}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
