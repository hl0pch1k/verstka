"""«Заменить изображение»: a new picture for the place of slide `n` of a finished variant — made on the person's wish
(«фото сверху, больше зелени») — or the free place given back («Убрать изображение»: their own photo goes there).

Like «Исправить слайд» (pipeline/revise.py fix_slide): only slide n changes; nothing in the variant's folder is written
before the new picture is in, the deck audited and its previews rendered; then the previous version is kept
(versions/vN, its images.json with it), the files are replaced and the change is logged — «верни как было» undoes it."""

from __future__ import annotations

import json
import logging
import random
import shutil
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional

from verstka.images.fill import (
    WHY_REMOVED,
    descr_of,
    fitted,
    made_fields,
    public_records,
    read_records,
    summary,
    write_records,
)
from verstka.images.places import places_on, put_picture, swap_picture, take_picture_out
from verstka.images.prompt import ImageAsk, load_config
from verstka.images.providers import nearest_aspect
from verstka.images.service import ImageMaker

log = logging.getLogger(__name__)


@dataclass
class PictureChange:
    applied: bool
    reply: str
    slide: int
    action: str  # generate | remove
    record: Optional[dict] = None
    version: Optional[int] = None
    at: Optional[float] = None
    score_before: Optional[float] = None
    score_after: Optional[float] = None

    def result(self, strategy: str) -> dict:
        return {
            "reply": self.reply, "changed": self.applied, "applied": self.applied, "kind": "picture", "action": self.action, "strategy": strategy,
            "slide": self.slide, "picture": public_records([self.record])[0] if self.record else None, "version": self.version, "at": self.at,
            "score_before": self.score_before, "new_score": self.score_after,
        }


def picture_label(n: int, action: str, wish: Optional[str]) -> str:
    """The change as the edit log and «верни как было» name it."""
    if action == "remove":
        return f"Слайд {n}: убрать изображение"
    return f"Слайд {n}: другое изображение" + (f" — «{wish.strip()[:160]}»" if wish and wish.strip() else "")


def slide_picture(records: list[dict], n: int) -> Optional[dict]:
    """The record of slide n's picture place: a generated picture first, then a free place."""
    mine = [r for r in records if r.get("slide") == n]
    return next((r for r in mine if r.get("status") == "generated"), None) or (mine[0] if mine else None)


def has_place(vdir: Path, n: int) -> bool:
    """Slide n of the variant has a picture or a free place for one (images.json, else the deck itself)."""
    if slide_picture(read_records(vdir), n) is not None:
        return True
    try:
        from pptx import Presentation

        prs = Presentation(str(Path(vdir) / "deck.pptx"))
        return 1 <= n <= len(prs.slides) and bool(places_on(prs.slides[n - 1], n))
    except Exception:  # noqa: BLE001
        return False


def _say(progress: Optional[Callable[..., None]], message: str, frac: float) -> None:
    if progress is not None:
        progress(message, frac)


def _agent(progress: Optional[Callable[..., None]], step: str, message: str, frac: float, slide: int, variant: str) -> None:
    """An event of the agent's timeline (the panel shows the job's steps like a slide fix's)."""
    if progress is not None:
        progress({"type": "agent", "step": step, "message": message, "slide": slide, "variant": variant}, frac)


def change_picture(
    vdir: Path,
    strategy: str,
    n: int,
    *,
    manifest,
    ws,
    provider: Any = None,
    action: str = "generate",
    wish: Optional[str] = None,
    skills: Any = None,
    providers: Any = None,
    exports: Optional[list[str]] = None,
    brief_text: Optional[str] = None,
    progress: Optional[Callable[..., None]] = None,
    cfg: Optional[dict] = None,
) -> PictureChange:
    from pptx import Presentation

    from verstka.api.remarks import carry_report_ids
    from verstka.audit.runner import run_audit
    from verstka.pipeline.generate import render_outputs
    from verstka.pipeline.revise import _previews, image_drift, log_edit, snapshot
    from verstka.schemas.audit import AuditReport
    from verstka.schemas.outline import DeckOutline

    vdir = Path(vdir)
    cfg = cfg if cfg is not None else load_config()
    wish = (wish or "").strip() or None
    outline = DeckOutline.model_validate_json((vdir / "outline.json").read_text(encoding="utf-8"))
    report0 = AuditReport.model_validate_json((vdir / "audit_report.json").read_text(encoding="utf-8"))
    records = read_records(vdir)
    oslide = outline.slides[n - 1] if 1 <= n <= len(outline.slides) else None
    scratch = vdir.parent / f".{strategy}.picture"
    shutil.rmtree(scratch, ignore_errors=True)
    scratch.mkdir(parents=True)
    try:
        work = scratch / "deck.pptx"
        shutil.copy2(vdir / "deck.pptx", work)
        prs = Presentation(str(work))
        if not 1 <= n <= len(prs.slides):
            raise ValueError(f"слайда {n} нет")
        slide = prs.slides[n - 1]
        rec = slide_picture(records, n)
        place = None
        if rec is None or rec.get("status") != "generated":
            places = places_on(slide, n, oslide)
            place = next((p for p in places if rec is not None and p.shape_id == rec.get("shape_id")), None) or (places[0] if places else None)
            if place is None:
                raise ValueError(f"на слайде {n} нет изображения и места под него")
            if rec is None:
                rec = {"id": f"s{n}-{place.shape_id}", "slide": n, "shape_id": place.shape_id, "outline_id": place.outline_id, "headline": place.headline,
                       "label": place.label, "hint": place.hint, "own": place.own, "box": list(place.box), "aspect": nearest_aspect(place.box[2], place.box[3]),
                       "status": "free", "why": None, "wish": None}
        rec = dict(rec)
        before = rec.get("status")

        if action == "remove":
            if before != "generated":
                raise ValueError(f"на слайде {n} нет сгенерированного изображения")
            _say(progress, "Убираю изображение", 0.2)
            if not take_picture_out(slide, str(rec.get("shape_id")), rec.get("kept") or {}):
                raise ValueError("место под изображение не удалось вернуть")
            rec.update({"status": "removed", "why": WHY_REMOVED, "at": time.time()})
            rec.pop("kept", None)
            reply = f"Слайд {n}: изображение убрано — место оставлено под ваше фото."
        else:
            if provider is None or getattr(provider, "off", False):
                raise ValueError("модель изображений не подключена")
            ask = ImageAsk(hint=str(rec.get("hint") or ""), headline=str(rec.get("headline") or (oslide.headline if oslide else "")), deck=outline.title or "",
                           label=str(rec.get("label") or ""), wish=wish)
            box = rec.get("box") or (list(place.box) if place is not None else None)
            aspect = str(rec.get("aspect") or (nearest_aspect(box[2], box[3]) if box else "4:3"))
            deadline = time.monotonic() + float(cfg.get("timeout_s", 90)) * 2 + 15
            maker = ImageMaker(provider, manifest, vdir.parent / "images", cfg=cfg, skills=skills, providers=providers, deadline=deadline, deck_title=outline.title or "")
            _say(progress, "Рисую изображение", 0.1)
            _agent(progress, "designer", f"Дизайнер: слайд {n} — новое изображение" + (f": «{wish[:120]}»" if wish else "") + ".", 0.15, n, strategy)
            try:
                made = maker.make(ask, aspect, seed=random.SystemRandom().randrange(1, 2_147_483_647))
            finally:
                maker.close()
            _say(progress, "Ставлю изображение на слайд", 0.55)
            if before == "generated":
                pic_box = rec.get("box") or box
                data = fitted(made.path, pic_box, cfg)
                rec.update(made_fields(made))
                rec["wish"] = wish
                if not swap_picture(slide, str(rec.get("shape_id")), data, descr=descr_of(rec)):
                    raise ValueError("изображение на слайде не найдено")
            else:
                data = fitted(made.path, place.box, cfg)
                rec.update(made_fields(made))
                rec["wish"] = wish
                kept = put_picture(slide, place, data, descr=descr_of(rec))
                if kept is None:
                    raise ValueError("место под изображение не найдено")
                rec.update({"kept": kept, "shape_id": place.shape_id, "box": list(place.box)})
            rec.update({"status": "generated", "why": None, "error": None, "at": time.time()})
            reply = f"Слайд {n}: новое изображение" + (f" — «{wish}»" if wish else "") + f". Нарисовала модель {made.model_label}."
        prs.save(str(work))

        # the deck audited again (the deterministic checks); the model's remarks and the fixes carry over
        _say(progress, "Проверяю слайд", 0.65)
        report = carry_report_ids(report0, run_audit(work, manifest, outline, ws, use_vlm=False, use_llm=False, render=False, strategy=strategy, brief_text=brief_text))
        report.issues.extend(i.model_copy(deep=True) for i in report0.issues if i.kind == "model")
        report.applied_fixes = list(report0.applied_fixes) + [{"action": "picture", "slide": n, "result": action}]
        report.iterations = report0.iterations
        report.recompute()

        # previews: slide n from the new render, every other slide keeps its image
        _say(progress, "Готовлю превью", 0.75)
        out = scratch / "out"
        out.mkdir()
        shutil.copy2(work, out / "deck.pptx")
        exports = [e for e in (exports or []) if e in ("pdf", "html")]
        _, images, warns, _ = render_outputs(out, manifest, outline.title, exports, images=True)
        for w in warns:
            log.warning("picture change: %s", w)
        merged_dir = scratch / "slides"
        merged_dir.mkdir()
        old_images = {int(p.stem.rsplit("-", 1)[-1]): p for p in _previews(vdir)}
        new_images = {k: p for k, p in enumerate(images, 1)}
        for k in range(1, len(outline.slides) + 1):
            src = new_images.get(k) if (k == n or k not in old_images) else old_images[k]
            src = src or old_images.get(k)
            if src is None:
                continue
            if k != n and k in new_images and k in old_images:
                d = image_drift(old_images[k], new_images[k])
                if d is not None and d > 3:
                    log.info("picture change: image drift on slide %d (%.1f) — the old preview is kept", k, d)
            shutil.copy2(src, merged_dir / f"slide-{k:03d}.jpg")

        # E. apply: the previous version kept, the files replaced, the change logged
        version = snapshot(vdir)
        (out / "deck.pptx").replace(vdir / "deck.pptx")
        for fmt in ("pdf", "html"):
            made_file = out / f"deck.{fmt}"
            if made_file.exists():
                made_file.replace(vdir / f"deck.{fmt}")
            elif (vdir / f"deck.{fmt}").exists():
                (vdir / f"deck.{fmt}").unlink()
        if any(merged_dir.iterdir()):
            shutil.rmtree(vdir / "slides", ignore_errors=True)
            merged_dir.replace(vdir / "slides")
        records = [rec if (r.get("slide") == n and r.get("id") == rec.get("id")) else r for r in records]
        if not any(r.get("id") == rec.get("id") for r in records):
            records.append(rec)
        write_records(vdir, records)
        report.deck = str(vdir / "deck.pptx")
        report.slide_images = {k: str(p) for k, p in enumerate(_previews(vdir), 1)}
        (vdir / "audit_report.json").write_text(report.model_dump_json(indent=2), encoding="utf-8")
        try:
            rm = json.loads((vdir / "run_manifest.json").read_text(encoding="utf-8"))
            rm["audit"] = report.summary.model_dump()
            rm["applied_fixes"] = report.applied_fixes
            rm["pictures"] = summary(records)
            (vdir / "run_manifest.json").write_text(json.dumps(rm, ensure_ascii=False, indent=2), encoding="utf-8")
        except (OSError, ValueError):
            pass
        at = time.time()
        score_before, score_after = report0.summary.score, report.summary.score
        tail = " Не понравится — «Вернуть как было»."
        log_edit(vdir, {
            "at": at, "version": version, "undo_of": None, "kind": "picture", "slides": [n], "request": picture_label(n, action, wish),
            "reply": reply, "wish": wish, "action": action, "score_before": score_before, "score_after": score_after,
        })
        return PictureChange(applied=True, reply=reply + tail, slide=n, action=action, record=rec, version=version, at=at, score_before=score_before, score_after=score_after)
    finally:
        shutil.rmtree(scratch, ignore_errors=True)
