"""Revising one variant of a finished generation — the chat agent's edits («на слайде 3 покажи расходы таблицей»,
«убери слайд 5», «поменяй местами 2 и 3», «верни как было»).

Every revision keeps the variant's previous version (versions/vN/: its plan, layout, audit, deck and run manifest), so
«верни как было» restores it exactly; the edited plan is rendered again with the template the variant was built on and
audited against the brief (the figures on the slides included), and the variant's edit log (edits.json) records what was
asked, what changed and the scores before and after — the versioning of the agent's work the ТЗ asks for.

A fix of one slide's remarks («Исправить слайд», fix_slide) never re-renders the variant: the fixed slide is built in a
scratch deck and spliced into the current deck (splice_slide), every other slide is checked by its fingerprint
(slide_fingerprints: the slide's XML with every related part hashed), and «верни как было» puts the kept files back
exactly (restore_version) instead of rendering the old plan again."""

from __future__ import annotations

import copy
import hashlib
import io
import json
import logging
import math
import re
import shutil
import threading
import time
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Optional

from verstka.schemas.outline import Brief, DeckOutline

log = logging.getLogger(__name__)

# what a version of a variant is: enough to show it again without re-rendering (the deck) and to rebuild it (the plan)
VERSION_FILES = ("outline.json", "layout_plan.json", "audit_report.json", "run_manifest.json", "deck.pptx", "images.json")
# kept with a version too, so «верни как было» shows exactly the old previews and files without LibreOffice (older
# versions lack them: their previews and exports are rendered again from the kept deck)
VERSION_EXTRA = ("deck.pdf", "deck.html")
# what a re-render replaces in the variant's folder (planner_raw.json, the model's own answer, stays)
RENDER_OUTPUTS = ("outline.json", "layout_plan.json", "audit_report.json", "run_manifest.json", "deck.pptx", "deck.pdf", "deck.html", "images.json")


def versions_dir(vdir: Path) -> Path:
    return vdir / "versions"


def version_numbers(vdir: Path) -> list[int]:
    root = versions_dir(vdir)
    if not root.is_dir():
        return []
    return sorted(int(p.name[1:]) for p in root.iterdir() if p.is_dir() and p.name[:1] == "v" and p.name[1:].isdigit())


def snapshot(vdir: Path) -> int:
    """Keep the variant as it is now (versions/vN/) before it changes; returns N."""
    n = (version_numbers(vdir) or [0])[-1] + 1
    dst = versions_dir(vdir) / f"v{n}"
    dst.mkdir(parents=True, exist_ok=True)
    for name in VERSION_FILES + VERSION_EXTRA:
        src = vdir / name
        if src.exists():
            shutil.copy2(src, dst / name)
    if (vdir / "slides").is_dir():
        shutil.rmtree(dst / "slides", ignore_errors=True)
        shutil.copytree(vdir / "slides", dst / "slides")
    return n


def last_version(vdir: Path) -> Optional[tuple[int, DeckOutline]]:
    """The newest kept version and its plan (None when the variant was never edited)."""
    nums = version_numbers(vdir)
    if not nums:
        return None
    path = versions_dir(vdir) / f"v{nums[-1]}" / "outline.json"
    if not path.exists():
        return None
    return nums[-1], DeckOutline.model_validate_json(path.read_text(encoding="utf-8"))


def drop_version(vdir: Path, n: int) -> None:
    shutil.rmtree(versions_dir(vdir) / f"v{n}", ignore_errors=True)


def read_edits(vdir: Path) -> list[dict]:
    path = vdir / "edits.json"
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return data if isinstance(data, list) else []


def log_edit(vdir: Path, entry: dict) -> None:
    edits = read_edits(vdir)
    edits.append({"at": time.time(), **entry})
    tmp = vdir / "edits.json.tmp"
    tmp.write_text(json.dumps(edits, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(vdir / "edits.json")


def _read_json(path: Path) -> Optional[dict]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    return data if isinstance(data, dict) else None


def _read_audit(path: Path):
    """The AuditReport at `path`, or None (missing or unreadable)."""
    from verstka.schemas.audit import AuditReport

    try:
        return AuditReport.model_validate_json(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


_PROVENANCE = ("planner", "providers", "usage", "skills", "inputs", "verstka_version", "git_commit", "created_at")


def add_usage(base: Optional[dict], extra: Optional[dict]) -> dict:
    """The generation's usage with an edit's model calls on top (N-E2): the totals and `by_skill` grow, and `edits`
    keeps what the edits alone cost."""
    out = copy.deepcopy(base) if isinstance(base, dict) else {}
    if not extra or not extra.get("calls"):
        return out
    for key in ("calls", "requests", "tokens_in", "tokens_out", "failed"):
        out[key] = int(out.get(key) or 0) + int(extra.get(key) or 0)
    if extra.get("cost_rub") is not None:
        out["cost_rub"] = round(float(out.get("cost_rub") or 0.0) + float(extra["cost_rub"]), 4)
    by = out.setdefault("by_skill", {})
    for skill, b in (extra.get("by_skill") or {}).items():
        mine = by.setdefault(skill, {"calls": 0, "requests": 0, "tokens_in": 0, "tokens_out": 0, "failed": 0, "seconds": 0.0, "cost_rub": None})
        for key in ("calls", "requests", "tokens_in", "tokens_out", "failed"):
            mine[key] = int(mine.get(key) or 0) + int(b.get(key) or 0)
        mine["seconds"] = round(float(mine.get("seconds") or 0.0) + float(b.get("seconds") or 0.0), 2)
        if b.get("cost_rub") is not None:
            mine["cost_rub"] = round(float(mine.get("cost_rub") or 0.0) + float(b["cost_rub"]), 4)
    edits = out.setdefault("edits", {"calls": 0, "requests": 0, "tokens_in": 0, "tokens_out": 0, "failed": 0, "cost_rub": None, "count": 0})
    for key in ("calls", "requests", "tokens_in", "tokens_out", "failed"):
        edits[key] = int(edits.get(key) or 0) + int(extra.get(key) or 0)
    if extra.get("cost_rub") is not None:
        edits["cost_rub"] = round(float(edits.get("cost_rub") or 0.0) + float(extra["cost_rub"]), 4)
    edits["count"] = int(edits.get("count") or 0) + 1
    out.setdefault("scope", "generation")
    return out


def carry_provenance(old: Optional[dict], new: Optional[dict], edit_usage: Optional[dict] = None) -> Optional[dict]:
    """The re-rendered variant's run manifest with the generation's provenance (N-E2): who planned it, the models and
    skills it used, what the calls cost (plus the edit's own calls) and the inputs; the re-render keeps its own audit,
    slides, timings and the new outline's hash."""
    if new is None:
        return None
    if old is None:
        if edit_usage:
            new["usage"] = add_usage(new.get("usage"), edit_usage)
        return new
    for key in _PROVENANCE:
        if old.get(key) is None:
            continue
        if key == "inputs":
            inputs = dict(old["inputs"]) if isinstance(old["inputs"], dict) else {}
            mine = new.get("inputs") if isinstance(new.get("inputs"), dict) else {}
            if mine.get("outline_sha256"):
                inputs["outline_sha256"] = mine["outline_sha256"]
            new["inputs"] = inputs
        else:
            new[key] = copy.deepcopy(old[key])
    if edit_usage:
        new["usage"] = add_usage(new.get("usage"), edit_usage)
    timings = new.get("timings_s") if isinstance(new.get("timings_s"), dict) else None
    if timings is not None and isinstance(old.get("timings_s"), dict):
        timings["generation"] = old["timings_s"].get("generation") or old["timings_s"]
    return new


def _slide_state(s) -> str:
    """A slide's plan as it renders (everything but the agent's words about it)."""
    return s.model_dump_json(exclude={"rationale", "alternatives"})


def rerender_variant(
    vdir: Path,
    strategy: str,
    outline: DeckOutline,
    template_source: Path,
    workspace_root: Path,
    *,
    brief: Optional[Brief] = None,
    exports: Optional[list[str]] = None,
    pictures: Optional[bool] = None,
):
    """Render the edited plan of one variant into a scratch folder, then put its files in place of the variant's (the
    slide previews too). No text model is asked anything here: the plan is final; the audit and its autofix run as for
    a new deck, against the brief. The pictures of the variant stay (its records: the same place keeps its picture, a
    picture taken out stays out); a new place the author described gets one when `pictures` allows (None: the config).
    Returns the pipeline's VariantResult."""
    from verstka.images.fill import read_records

    from verstka.pipeline.generate import generate_variants

    edit_usage: dict = {}
    try:
        from verstka.planning.slide_edit import take_edit_usage

        edit_usage = take_edit_usage(outline)
    except Exception:  # noqa: BLE001
        edit_usage = {}
    try:
        # the slides the edit did not touch keep their layout exactly (N-E1): only the edited ones are laid out again
        pinned = _rerender_pinned(Path(vdir), strategy, outline, Path(template_source), Path(workspace_root), brief=brief, exports=exports, edit_usage=edit_usage)
    except Exception:  # noqa: BLE001 - the variant is laid out again as a whole, as before
        log.warning("pinned re-render failed: the variant is laid out again", exc_info=True)
        pinned = None
    if pinned is not None:
        return pinned
    scratch = vdir.parent / f".{strategy}.revise"
    shutil.rmtree(scratch, ignore_errors=True)
    # who planned the deck (the agent, its model and calls) stays the variant's: an edit re-renders a supplied plan
    old_manifest = _read_json(vdir / "run_manifest.json")
    old_audit = _read_audit(vdir / "audit_report.json")
    try:
        res = generate_variants(
            template_source, brief=brief, outline=outline, strategies=[strategy], out_dir=scratch, workspace_root=workspace_root,
            providers=None, skills=None, use_llm=False, use_vlm=False, audit=True, autofix=True, exports=list(exports or []),
            pictures=pictures, picture_dir=vdir.parent / "images", picture_reuse={strategy: read_records(vdir)},
        )
        v = res.variants[0]
        made = scratch / strategy
        for name in RENDER_OUTPUTS:
            src = made / name
            if src.exists():
                src.replace(vdir / name)
            elif name in ("deck.pdf", "deck.html") and (vdir / name).exists() and name.split(".")[-1] not in (exports or []):
                (vdir / name).unlink()  # an export the variant no longer has would show the old deck
            elif name == "images.json" and (vdir / name).exists():
                (vdir / name).unlink()  # the edited plan has no picture places left
        if (made / "slides").is_dir():
            shutil.rmtree(vdir / "slides", ignore_errors=True)
            (made / "slides").replace(vdir / "slides")
        # the audit report and the run manifest name the previews and exports by path: they now live in the variant's
        # own folder
        for name in ("audit_report.json", "run_manifest.json"):
            path = vdir / name
            if path.exists():
                path.write_text(path.read_text(encoding="utf-8").replace(str(made), str(vdir)), encoding="utf-8")
        new_manifest = carry_provenance(old_manifest, _read_json(vdir / "run_manifest.json"), edit_usage)
        if new_manifest is not None:
            (vdir / "run_manifest.json").write_text(json.dumps(new_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        # the remarks that stayed keep their ids (the audit numbers them through the deck, so an edit of one slide would
        # rename the remarks of every slide after it)
        new_audit = _read_audit(vdir / "audit_report.json")
        if old_audit is not None and new_audit is not None:
            from verstka.api.remarks import carry_report_ids

            carry_report_ids(old_audit, new_audit)
            (vdir / "audit_report.json").write_text(new_audit.model_dump_json(indent=2), encoding="utf-8")
            v.audit = new_audit
        v.out_dir = vdir
        return v
    finally:
        shutil.rmtree(scratch, ignore_errors=True)


def _rerender_pinned(vdir: Path, strategy: str, outline: DeckOutline, template_source: Path, workspace_root: Path, *, brief: Optional[Brief], exports: Optional[list[str]], edit_usage: Optional[dict] = None):
    """The edited plan rendered with every untouched slide pinned (N-E1: a caption edit on slide 3 re-laid out slide 9):
    the edited slides are matched and rendered again in a scratch deck, fixed by the autofix on those slides only, and
    spliced into the current deck; every other slide keeps its XML (checked by fingerprints), its preview and its
    layout entry. None when that is not possible — another number or order of slides, nothing kept, missing files —
    and the variant is laid out again as a whole."""
    from verstka.analysis.manifest import analyze_template
    from verstka.api.remarks import carry_report_ids
    from verstka.audit.autofix import autofix_loop
    from verstka.audit.runner import run_audit
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.matching.matcher import match_outline
    from verstka.pipeline.generate import VariantResult, render_outputs
    from verstka.planning.strategies import load_strategies
    from verstka.rendering.renderer import RenderResult, render_deck, strip_outline_marks
    from verstka.schemas.audit import AuditReport, Issue
    from verstka.schemas.layout import LayoutPlan

    t0 = time.time()
    need = ("outline.json", "layout_plan.json", "audit_report.json", "deck.pptx")
    if not all((vdir / f).exists() for f in need):
        return None
    outline0 = DeckOutline.model_validate_json((vdir / "outline.json").read_text(encoding="utf-8"))
    if [s.id for s in outline0.slides] != [s.id for s in outline.slides]:
        return None
    changed = [k for k, (a, b) in enumerate(zip(outline0.slides, outline.slides), 1) if _slide_state(a) != _slide_state(b)]
    if not changed or len(changed) == len(outline.slides):
        return None
    plan0 = LayoutPlan.model_validate_json((vdir / "layout_plan.json").read_text(encoding="utf-8"))
    if any(plan0.for_outline(s.id) is None for s in outline.slides):
        return None
    report0 = AuditReport.model_validate_json((vdir / "audit_report.json").read_text(encoding="utf-8"))
    deck0 = vdir / "deck.pptx"
    manifest = analyze_template(template_source, workspace_root=workspace_root, providers=None, skills=None, use_llm=False, use_vlm=False)
    ws = TemplateWorkspace.create(template_source, workspace_root)
    st = load_strategies()
    strategy_obj = st.get(strategy) or next(iter(st.values()))
    fresh = match_outline(outline, manifest, strategy_obj)
    plan = plan0.model_copy(deep=True)
    for k in changed:
        oid = outline.slides[k - 1].id
        plan = _with_entry(plan, fresh.for_outline(oid), oid)
    brief_text = brief.text if brief is not None else None
    exports = [e for e in (exports or []) if e in ("pdf", "html")]
    scratch = vdir.parent / f".{strategy}.pinned"
    shutil.rmtree(scratch, ignore_errors=True)
    scratch.mkdir(parents=True)
    try:
        work = scratch / "work.pptx"
        o = outline.model_copy(deep=True)
        render_deck(o, plan, manifest, ws, work)  # the whole scratch deck; only the edited slides are used
        rep = run_audit(work, manifest, o, ws, use_vlm=False, use_llm=False, render=False, strategy=strategy, brief_text=brief_text)
        ids = {i.id for i in rep.issues if i.slide in changed and i.severity in ("error", "warn") and i.kind != "model" and i.autofix is not None and i.autofix.action != "none"}
        records: list[dict] = []
        if ids:
            start = len(rep.applied_fixes)
            rep, plan, o, _ = autofix_loop(work, rep, o, plan, manifest, ws, max_iterations=2, only_ids=ids, render=False, brief_text=brief_text)
            records = list(rep.applied_fixes[start:])
        final = scratch / "final.pptx"
        shutil.copy2(deck0, final)
        for k in changed:
            splice_slide(final, work, k, scratch / "spliced.pptx")
            (scratch / "spliced.pptx").replace(final)
        try:
            strip_outline_marks(final)
        except Exception:  # noqa: BLE001
            log.warning("outline markers left in the notes", exc_info=True)
        # a picture the edited slide had goes back into its place (no model call)
        pic_records = None
        for k in changed:
            got = _keep_slide_picture(vdir, final, o, k)
            if got is not None:
                pic_records = got
        report = carry_report_ids(report0, run_audit(final, manifest, o, ws, use_vlm=False, use_llm=False, render=False, strategy=strategy, brief_text=brief_text))
        report.issues.extend(i.model_copy(deep=True) for i in report0.issues if i.kind == "model" and i.slide not in changed)
        report.applied_fixes = list(report0.applied_fixes) + [{"action": "edit", "slides": changed}] + records
        report.iterations = report0.iterations
        report.recompute()
        fp0, fp1 = slide_fingerprints(deck0), slide_fingerprints(final)
        other = changed_slides(fp0, fp1, skip=set(changed))
        out = scratch / "out"
        out.mkdir()
        shutil.copy2(final, out / "deck.pptx")
        paths, images, warns, render_ok = render_outputs(out, manifest, o.title, exports, images=True)
        for w in warns:
            log.warning("pinned re-render: %s", w)
        if render_ok is False:
            report.issues.append(Issue(id="file_opens-0-render", slide=0, check_id="file_opens", severity="warn", kind="deterministic", message="LibreOffice не смог отрендерить файл"))
            report.recompute()
        merged = scratch / "slides"
        merged.mkdir()
        old_images = {int(p.stem.rsplit("-", 1)[-1]): p for p in _previews(vdir)}
        new_images = {k: p for k, p in enumerate(images, 1)}
        for k in range(1, len(o.slides) + 1):
            fresh_img = k in changed or k in other or k not in old_images
            src = (new_images.get(k) if fresh_img else old_images.get(k)) or old_images.get(k)
            if src is not None:
                shutil.copy2(src, merged / f"slide-{k:03d}.jpg")
        # the files in place (the caller kept the previous version)
        (out / "deck.pptx").replace(deck0)
        if pic_records is not None:
            from verstka.images.fill import write_records

            write_records(vdir, pic_records)
        for fmt in ("pdf", "html"):
            made = out / f"deck.{fmt}"
            if made.exists():
                made.replace(vdir / f"deck.{fmt}")
            elif (vdir / f"deck.{fmt}").exists():
                (vdir / f"deck.{fmt}").unlink()  # it would show the deck before the edit
        if any(merged.iterdir()):
            shutil.rmtree(vdir / "slides", ignore_errors=True)
            merged.replace(vdir / "slides")
        report.deck = str(deck0)
        report.slide_images = {k: str(p) for k, p in enumerate(_previews(vdir), 1)}
        (vdir / "outline.json").write_text(o.model_dump_json(indent=2), encoding="utf-8")
        (vdir / "layout_plan.json").write_text(plan.model_dump_json(indent=2), encoding="utf-8")
        (vdir / "audit_report.json").write_text(report.model_dump_json(indent=2), encoding="utf-8")
        rm = _read_json(vdir / "run_manifest.json")
        if rm is not None:
            rm["audit"] = report.summary.model_dump()
            rm["applied_fixes"] = report.applied_fixes
            rm.setdefault("inputs", {})["outline_sha256"] = hashlib.sha256(o.model_dump_json().encode("utf-8")).hexdigest()[:16]
            slides = rm.get("slides")
            for k in changed:
                e = plan.for_outline(o.slides[k - 1].id)
                if e is not None and isinstance(slides, list) and len(slides) >= k and isinstance(slides[k - 1], dict):
                    slides[k - 1].update({"mode": e.mode, "pattern_id": e.pattern_id, "composition": e.composition})
            if edit_usage:
                rm["usage"] = add_usage(rm.get("usage"), edit_usage)
            rm.setdefault("timings_s", {})["edit"] = round(time.time() - t0, 2)
            (vdir / "run_manifest.json").write_text(json.dumps(rm, ensure_ascii=False, indent=2), encoding="utf-8")
        v = VariantResult(strategy=strategy, out_dir=vdir, outline=o, plan=plan, render=RenderResult(pptx_path=deck0), images=_previews(vdir),
                          warnings=[], seconds=round(time.time() - t0, 2), audit=report,
                          exports={fmt: vdir / f"deck.{fmt}" for fmt in exports if (vdir / f"deck.{fmt}").exists()})
        v.planner = (rm or {}).get("planner") or {}
        v.other_slides = other  # type: ignore[attr-defined]  (the slides besides the edited ones that changed: none expected)
        v.changed_slides = changed  # type: ignore[attr-defined]
        if other:
            log.warning("pinned re-render: slides %s changed besides the edited %s", other, changed)
        return v
    finally:
        shutil.rmtree(scratch, ignore_errors=True)


# ---------------------------------------------------------------------------- one slide into a deck: splice, fingerprints


_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_R_ATTRS = tuple(f"{{{_R_NS}}}{a}" for a in ("embed", "link", "id", "pict", "dm", "lo", "qs", "cs"))


def _elements(el) -> Iterable:
    from lxml import etree

    return el.iter(tag=etree.Element)


def _remap_rids(el, rid_map: dict[str, str]) -> list[str]:
    """Rewrite every r:* reference by `rid_map` in one pass; returns the references it had no mapping for."""
    unknown = []
    for node in _elements(el):
        for attr in _R_ATTRS:
            v = node.get(attr)
            if v is None:
                continue
            if v in rid_map:
                node.set(attr, rid_map[v])
            else:
                unknown.append(v)
    return unknown


def _next_partname(pkg, template: str, reserved: set[str]):
    """A free part name for `template` ("/ppt/charts/chart%d.xml"): unused in the package and not handed out yet (a
    part cloned but not related yet is not in the package's graph)."""
    from pptx.opc.packuri import PackURI

    used = {str(p.partname) for p in pkg.iter_parts()} | reserved
    for i in range(1, 100000):
        cand = template % i
        if cand not in used:
            reserved.add(cand)
            return PackURI(cand)
    raise RuntimeError(f"no free part name for {template}")


def _clone_part(pkg, src, memo: dict[int, Any], reserved: set[str]):
    """A copy of `src` (and of the parts it relates to) inside `pkg`, with its references renamed."""
    from pptx.opc.package import PartFactory

    hit = memo.get(id(src))
    if hit is not None:
        return hit
    name = str(src.partname)
    template = re.sub(r"\d+(\.[^./]+)$", r"%d\1", name)
    if "%d" not in template:
        stem, dot, ext = name.rpartition(".")
        template = f"{stem}%d.{ext}" if dot else f"{name}%d"
    new = PartFactory(_next_partname(pkg, template, reserved), src.content_type, pkg, src.blob)
    memo[id(src)] = new
    rid_map: dict[str, str] = {}
    for rid, rel in list(src.rels.items()):
        if rel.is_external:
            rid_map[rid] = new.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rid_map[rid] = new.relate_to(_clone_part(pkg, rel.target_part, memo, reserved), rel.reltype)
    el = getattr(new, "_element", None)
    if el is not None and any(k != v for k, v in rid_map.items()):
        _remap_rids(el, rid_map)
    return new


def splice_slide(base: Path, donor: Path, index: int, out: Path) -> None:
    """`out` = `base` with slide `index` (1-based) taken from `donor` (the same template); every other part of `base`
    stays as it is. The slide's pictures, charts (with their workbooks) and links come along; its notes (the outline-id
    marker the audit reads) are replaced; the old slide's own parts that nothing uses any more are not written."""
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    bp, dp = Presentation(str(base)), Presentation(str(donor))
    if not (1 <= index <= len(bp.slides) and index <= len(dp.slides)):
        raise IndexError(f"slide {index} is not in both decks ({len(bp.slides)} and {len(dp.slides)} slides)")
    bs, ds = bp.slides[index - 1], dp.slides[index - 1]
    pkg = bp.part.package
    layouts = [lay for m in bp.slide_masters for lay in m.slide_layouts]
    want = ds.slide_layout
    layout = next((lay for lay in layouts if lay.part.partname == want.part.partname), None) or next((lay for lay in layouts if lay.name == want.name), None)
    if layout is None:
        raise ValueError(f"the base deck has no layout {want.part.partname} «{want.name}»")
    # the base slide keeps only its notes; the layout is related again (the donor's may be another one)
    for rid, rel in list(bs.part.rels.items()):
        if rel.reltype != RT.NOTES_SLIDE:
            bs.part.rels.pop(rid)
    bs.part.relate_to(layout.part, RT.SLIDE_LAYOUT)
    donor_slides = {id(s.part): i for i, s in enumerate(dp.slides)}
    memo: dict[int, Any] = {}
    reserved: set[str] = set()
    rid_map: dict[str, str] = {}
    used = {node.get(attr) for node in _elements(ds._element) for attr in _R_ATTRS if node.get(attr) is not None}
    for rid, rel in ds.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rid not in used:  # a part the slide no longer shows (a sample's chart the render replaced): left behind
            continue
        if rel.is_external:
            rid_map[rid] = bs.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif rel.reltype == RT.SLIDE:  # a link to another slide of the deck: the base deck's slide at that place
            j = donor_slides.get(id(rel.target_part))
            if j is not None and j < len(bp.slides):
                rid_map[rid] = bs.part.relate_to(bp.slides[j].part, RT.SLIDE)
        elif rel.reltype == RT.IMAGE:
            try:
                _, rid_map[rid] = bs.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))  # the same picture is shared
            except Exception:  # noqa: BLE001 - a format python-pptx cannot size (svg, wdp): copied as it is
                rid_map[rid] = bs.part.relate_to(_clone_part(pkg, rel.target_part, memo, reserved), rel.reltype)
        else:  # a chart with its workbook, media, anything else: copied with the parts it relates to
            rid_map[rid] = bs.part.relate_to(_clone_part(pkg, rel.target_part, memo, reserved), rel.reltype)
    new_el = copy.deepcopy(ds._element)
    unknown = _remap_rids(new_el, rid_map)
    if unknown:
        log.warning("splice: slide %d refers to %s that the donor slide does not relate", index, sorted(set(unknown)))
    old = bs._element
    for child in list(old):
        old.remove(child)
    old.attrib.clear()
    old.attrib.update(dict(new_el.attrib))
    for child in list(new_el):
        old.append(child)
    if ds.has_notes_slide:
        d_tree = ds.notes_slide._element.cSld.spTree
        b_tree = bs.notes_slide._element.cSld.spTree  # created from the notes master when the base slide had none
        b_tree.getparent().replace(b_tree, copy.deepcopy(d_tree))
    out.parent.mkdir(parents=True, exist_ok=True)
    bp.save(str(out))


_ZIP_VOLATILE = frozenset({"docProps/core.xml"})  # an embedded workbook's creation time: new on every render


def _blob_digest(blob: bytes) -> str:
    """A binary part's hash; an embedded package (a chart's workbook) by its entries, without its timestamps."""
    if blob[:2] == b"PK":
        import zipfile

        try:
            h = hashlib.sha1()
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                for name in sorted(z.namelist()):
                    if name in _ZIP_VOLATILE:
                        continue
                    h.update(name.encode("utf-8") + b"\0" + z.read(name) + b"\0")
            return h.hexdigest()
        except (zipfile.BadZipFile, OSError, ValueError):
            pass
    return hashlib.sha1(blob).hexdigest()


def _part_digest(part, memo: dict[int, str], depth: int) -> str:
    key = id(part)
    if key in memo:
        return memo[key]
    memo[key] = "cycle"
    el = getattr(part, "_element", None)
    if el is not None:
        from lxml import etree

        el = copy.deepcopy(el)
        _hash_refs(el, part.rels, memo, depth + 1)
        digest = hashlib.sha1(etree.tostring(el, method="c14n", exclusive=True)).hexdigest()
    else:
        digest = _blob_digest(part.blob)
    memo[key] = digest
    return digest


def _hash_refs(el, rels, memo: dict[int, str], depth: int) -> None:
    """Every r:* value of `el` replaced by what it points at (the target's content, the external url)."""
    for node in _elements(el):
        for attr in _R_ATTRS:
            v = node.get(attr)
            if v is None:
                continue
            rel = rels.get(v)
            if rel is None:
                node.set(attr, "missing")
            elif rel.is_external:
                node.set(attr, "ext:" + str(rel.target_ref))
            elif depth > 4:
                node.set(attr, str(rel.target_part.partname))
            else:
                node.set(attr, _part_digest(rel.target_part, memo, depth))


def slide_fingerprints(pptx: Path) -> list[str]:
    """Per slide: the SHA-1 of its XML (exclusive C14N) with every r:* reference replaced by a hash of its target (the
    part's C14N or blob, `ext:{url}`), plus its layout's part name. Equal fingerprints — the same slide."""
    from lxml import etree
    from pptx import Presentation

    prs = Presentation(str(pptx))
    memo: dict[int, str] = {}
    out = []
    for s in prs.slides:
        el = copy.deepcopy(s._element)
        _hash_refs(el, s.part.rels, memo, 1)
        data = etree.tostring(el, method="c14n", exclusive=True) + str(s.slide_layout.part.partname).encode()
        out.append(hashlib.sha1(data).hexdigest())
    return out


def changed_slides(a: Path | list[str], b: Path | list[str], skip: Iterable[int] = ()) -> list[int]:
    """The 1-based slides whose fingerprints differ between two decks (or fingerprint lists), every slide past the
    shorter deck included; `skip` is left out."""
    fa = a if isinstance(a, list) else slide_fingerprints(Path(a))
    fb = b if isinstance(b, list) else slide_fingerprints(Path(b))
    skip = set(skip)
    out = [i for i, (x, y) in enumerate(zip(fa, fb), 1) if x != y and i not in skip]
    out += [i for i in range(min(len(fa), len(fb)) + 1, max(len(fa), len(fb)) + 1) if i not in skip]
    return out


def image_drift(a: Path, b: Path) -> Optional[float]:
    """The mean absolute grey difference (0–255) of two slide previews at 64×36; None when one cannot be read."""
    try:
        from PIL import Image, ImageChops, ImageStat

        with Image.open(a) as ia, Image.open(b) as ib:
            ga, gb = ia.convert("L").resize((64, 36)), ib.convert("L").resize((64, 36))
            return float(ImageStat.Stat(ImageChops.difference(ga, gb)).mean[0])
    except Exception:  # noqa: BLE001
        return None


# ---------------------------------------------------------------------------- «верни как было»: the kept files back


def _previews(vdir: Path) -> list[Path]:
    root = vdir / "slides"
    return sorted(root.glob("slide-*.jpg")) if root.is_dir() else []


def restore_version(vdir: Path, n: int, manifest, title: Optional[str], exports: Optional[list[str]] = None):
    """Put version `n` back exactly: its deck, plan, layout, audit and run manifest; its previews and exports when it
    kept them, else they are rendered again from the kept deck (a failed render keeps the files and warns). Nothing is
    planned or rendered as a deck again. Returns the restored AuditReport (None when the version had none)."""
    from verstka.pipeline.generate import render_outputs
    from verstka.schemas.audit import AuditReport

    src = versions_dir(vdir) / f"v{n}"
    if not src.is_dir():
        raise FileNotFoundError(f"version {n} of {vdir} is gone")
    exports = [e for e in (exports or []) if e in ("pdf", "html")]
    for name in VERSION_FILES:
        if (src / name).exists():
            shutil.copy2(src / name, vdir / name)
    kept_slides = (src / "slides").is_dir()
    if kept_slides:
        tmp = vdir / ".slides.restore"
        shutil.rmtree(tmp, ignore_errors=True)
        shutil.copytree(src / "slides", tmp)
        shutil.rmtree(vdir / "slides", ignore_errors=True)
        tmp.replace(vdir / "slides")
    missing = []
    for fmt in ("pdf", "html"):
        kept = src / f"deck.{fmt}"
        if kept.exists():
            shutil.copy2(kept, vdir / f"deck.{fmt}")
        elif fmt in exports:
            missing.append(fmt)
        elif (vdir / f"deck.{fmt}").exists():
            (vdir / f"deck.{fmt}").unlink()  # an export the version did not have would show another deck
    if missing or not kept_slides:
        try:
            _, _, warns, _ = render_outputs(vdir, manifest, title, exports, images=not kept_slides)
            for w in warns:
                log.warning("restore v%d: %s", n, w)
        except Exception as e:  # noqa: BLE001 - the files are back; the previews may lag behind
            log.warning("restore v%d: render failed: %s", n, e)
    path = vdir / "audit_report.json"
    if not path.exists():
        return None
    report = AuditReport.model_validate_json(path.read_text(encoding="utf-8"))
    report.deck = str(vdir / "deck.pptx")
    images = _previews(vdir)
    if images:
        report.slide_images = {k: str(p) for k, p in enumerate(images, 1)}
    path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    return report


# ---------------------------------------------------------------------------- «Исправить слайд»


@dataclass
class SlideFix:
    """What a fix of one slide did (REMARKS_SPEC §2.6)."""

    applied: bool
    how: str  # "model" | "rules" | "autofix"
    what: str
    reply: str
    report: Any  # the variant's AuditReport after the fix (the old one when not applied)
    requested: list = field(default_factory=list)  # the Issues asked to fix, numbered like the UI
    fixed: list[str] = field(default_factory=list)
    remaining: list[dict] = field(default_factory=list)
    other: list[int] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    why: Optional[str] = None
    version: Optional[int] = None
    at: Optional[float] = None
    score_before: Optional[float] = None
    score_after: Optional[float] = None
    n_slides: int = 0
    # nothing to change after all: the slide is the same, the model sees none of the asked problems, the remarks are gone
    # from the report (the reply says «замечаний больше нет»)
    cleared: bool = False


_EVENT_PREFIX = re.compile(r"^(?:Аналитик|Архитектор|Дизайнер|Критик|Правка|Сборка|Вёрстка|Проверка)\s*:\s*")
_EVENT_SLIDE_LEAD = re.compile(r"^слайд\s*\d+\s*[—–:]\s*", re.I)


def _event_text(message: str) -> str:
    """An event's text as the timeline shows it under its step: no step name and no «слайд N —» in front (the event
    carries the slide), like the planning agent's events."""
    text = _EVENT_PREFIX.sub("", message.strip(), count=1)
    text = _EVENT_SLIDE_LEAD.sub("", text, count=1).strip()
    return text[:1].upper() + text[1:] if text else message


class _Voice:
    """The fix's progress: plain job messages and agent events {"type": "agent", step, message, slide, variant}."""

    def __init__(self, progress: Optional[Callable[..., None]], slide: int, variant: str) -> None:
        self.progress = progress
        self.slide = slide
        self.variant = variant
        self.lines: list[str] = []
        self.designer_events = 0

    def _send(self, *args: Any) -> None:
        if self.progress is None:
            return
        try:
            self.progress(*args)
        except Exception:  # noqa: BLE001 - a listener never stops the fix
            log.debug("slide fix: progress listener failed", exc_info=True)

    def plain(self, message: str, share: Optional[float] = None) -> None:
        self._send(message, share)

    def agent(self, step: str, message: str, share: Optional[float] = None) -> None:
        self.lines.append(message)
        ev = {"type": "agent", "step": step, "message": _event_text(message), "slide": self.slide, "variant": self.variant}
        if share is not None:
            ev["progress"] = share
        self._send(ev)

    def relay(self) -> Callable[..., None]:
        """The designer's progress callback (redesign_slide): its events get this slide, this variant and a share."""

        def progress(msg: Any = "", frac: Any = None, *args: Any, **kwargs: Any) -> None:
            ev = msg if isinstance(msg, dict) else kwargs.get("event")
            if isinstance(ev, dict) and ev.get("type") == "agent":
                ev = {**ev, "slide": self.slide, "variant": self.variant}
                if ev.get("step") in ("designer", "revise"):
                    ev["progress"] = min(0.15 + 0.12 * self.designer_events, 0.5)
                    self.designer_events += 1
                self._send(ev)
            elif msg:
                self._send(msg, frac)

        return progress


class _Creep:
    """A long step with no steps of its own (LibreOffice rendering the previews): while it runs, the job's bar creeps
    from `start` toward `end` without reaching it — the person sees the work go on, never a bar that stopped."""

    def __init__(self, voice: "_Voice", message: str, start: float, end: float, every: float = 1.5, tau: float = 12.0) -> None:
        self.voice, self.message, self.start, self.end, self.every, self.tau = voice, message, start, end, every, tau
        self._stop = threading.Event()
        self._thread: Optional[threading.Thread] = None

    def _run(self) -> None:
        t0 = time.monotonic()
        while not self._stop.wait(self.every):
            share = self.start + (self.end - self.start) * (1.0 - math.exp(-(time.monotonic() - t0) / self.tau))
            self.voice.plain(self.message, round(share, 3))

    def __enter__(self) -> "_Creep":
        if self.voice.progress is not None:
            self._thread = threading.Thread(target=self._run, daemon=True, name="slide-fix-creep")
            self._thread.start()
        return self

    def __exit__(self, *exc: Any) -> None:
        self._stop.set()
        if self._thread is not None:
            self._thread.join(timeout=5)


def _with_slide(base: DeckOutline, src: DeckOutline, n: int) -> DeckOutline:
    """`base` with slide n taken from `src`, and the registry entries `src` added (series, facts, tables)."""
    out = base.model_copy(deep=True)
    out.slides[n - 1] = src.slides[n - 1].model_copy(deep=True)
    have = {x.id for x in out.series}
    out.series.extend(x.model_copy(deep=True) for x in src.series if x.id not in have)
    have = {x.id for x in out.facts}
    out.facts.extend(x.model_copy(deep=True) for x in src.facts if x.id not in have)
    for t in src.tables:
        if not any(t.columns == u.columns and t.rows == u.rows for u in out.tables):
            out.tables.append(t.model_copy(deep=True))
    return out


def _with_entry(base, entry, oid: str):
    """The layout plan `base` with slide `oid`'s entry replaced by `entry`."""
    out = base.model_copy(deep=True)
    if entry is None:
        return out
    entry = entry.model_copy(deep=True)
    for k, s in enumerate(out.slides):
        if s.outline_id == oid:
            out.slides[k] = entry
            return out
    out.slides.append(entry)
    return out


_ACTION_RU = {"move_inside": "элемент возвращён на место", "shrink_text": "шрифт уменьшен, чтобы текст поместился", "recolor": "цвет приведён к палитре шаблона", "refont": "шрифт заменён на шрифт шаблона", "drop_element": "убран лишний элемент"}


def _what_of(records: list[dict]) -> str:
    """What the rules did to the slide (the edits that stayed), in plain Russian («элемент возвращён внутрь слайда; …»)."""
    out: list[str] = []
    for r in records:
        act, res = r.get("action"), str(r.get("result") or "")
        if act == "xml":
            text = res or _ACTION_RU.get(str(r.get("kind")), "")
        elif act in ("rematch", "synth") and res and res != "skip":
            text = "слайд собран на другом макете шаблона" if res.startswith("rematch") else "слайд собран заново по его содержанию"
        elif act == "condense_text" and res and res != "skip":
            text = "тексты сокращены"
        else:
            continue
        text = text[:1].lower() + text[1:]
        if text and text not in out:
            out.append(text)
    return "; ".join(out[:3])


def _frame_name(kind: str) -> tuple[str, str]:
    """(the slide in the designer's line, the note for wishes) of a cover, a divider or a closing slide."""
    if kind == "section":
        return "разделитель", "Разделители агент собирает из разделов текста — исправлено только оформление"
    if kind == "thanks":
        return "финальный слайд", "Финальный слайд агент собирает из названия текста — исправлено только оформление"
    return "обложку", "Обложку агент собирает из названия и разделов текста — исправлено только оформление"


def _keep_slide_picture(vdir: Path, deck: Path, outline: DeckOutline, n: int) -> Optional[list[dict]]:
    """Slide `n` built again comes with its free picture place: the picture it had goes back in (refitted to the new
    place, no model call). The variant's records with slide n's renewed, or None when nothing changed."""
    try:
        from verstka.images.fill import apply_pictures, merged, plan_pictures, read_records

        old = read_records(vdir)
        mine = [r for r in old if r.get("slide") == n]
        if not mine:
            return None
        pend = plan_pictures(deck, outline, None, enabled=False, reuse=mine, folder=vdir.parent / "images", slides={n})
        if not pend:
            return None  # fixed in place: the picture never left
        got = apply_pictures(deck, pend, wait_s=0)
        return merged(old, got.records, {n})
    except Exception:  # noqa: BLE001 - the slide keeps its free place rather than the fix being lost
        log.warning("slide fix: the slide's picture was not put back", exc_info=True)
        return None


_WORD_RE = re.compile(r"[\w\d][\w\d\-]{2,}", re.U)


def slide_words(pptx: Path, n: int) -> set[str]:
    """The words slide `n` of `pptx` shows (its shapes' text, groups and tables included; lower case, 3+ letters)."""
    from pptx import Presentation

    try:
        s = Presentation(str(pptx)).slides[n - 1]
    except (IndexError, OSError, ValueError):
        return set()
    text = " ".join(t.text or "" for t in s._element.cSld.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
    return {w.lower() for w in _WORD_RE.findall(text)}


def lost_words_share(before: set[str], after: set[str]) -> float:
    """The share of the slide's words a new layout of it no longer shows (a cover sample without a subtitle place)."""
    return len(before - after) / float(len(before)) if before else 0.0


def remaining_asked(fixed: list[str], remarks: list) -> bool:
    """Whether a remark the person asked to fix is not among the fixed ones."""
    done = set(fixed)
    return any(r.id not in done for r in remarks)


def fix_label(n: int, wishes: Optional[str]) -> str:
    """The history line of a fix («Ваши правки»)."""
    wishes = (wishes or "").strip()
    return f"Исправь слайд {n}: {wishes}" if wishes else f"Исправь замечания на слайде {n}"


def fix_slide(
    vdir: Path,
    strategy: str,
    n: int,
    *,
    manifest,
    ws,
    remarks: list,
    wishes: Optional[str] = None,
    brief: Optional[Brief] = None,
    use_models: bool = False,
    skills: Any = None,
    providers: Any = None,
    exports: Optional[list[str]] = None,
    progress: Optional[Callable[..., None]] = None,
) -> SlideFix:
    """Fix the remarks of slide `n` of one variant — only this slide (REMARKS_SPEC §2.2).

    A: the remarks' in-place edits (an element moved back, a smaller size); B: the slide designer redesigns the slide
    when its content is at stake (or the person wished something), the slide is rendered by the template's layouts and
    fixed in place again; C: the structural fixes of what is left (another layout, a composed one, shorter texts),
    rolled back when worse; D: the slide is spliced into the current deck, every other slide is checked by its
    fingerprint, the deck is audited; E: when the slide got better, the previous version is kept (versions/vN), the
    files are replaced, the previews rendered and the fix logged. Nothing in `vdir` is written before E, so a failure
    leaves the variant as it was."""
    from verstka.api import remarks as RM
    from verstka.audit import seen_fix as SF
    from verstka.audit.autofix import RERENDER_ACTIONS, autofix_loop, plan_fixes
    from verstka.audit.runner import run_audit
    from verstka.matching.matcher import match_outline
    from verstka.pipeline.generate import render_outputs
    from verstka.planning.slide_edit import FRAME_KINDS, describe_change, redesign_slide
    from verstka.planning.strategies import load_strategies
    from verstka.rendering.renderer import render_deck
    from verstka.schemas.audit import AuditReport, Issue
    from verstka.schemas.layout import LayoutPlan

    vdir = Path(vdir)
    voice = _Voice(progress, n, strategy)
    titles = RM.check_titles()
    outline0 = DeckOutline.model_validate_json((vdir / "outline.json").read_text(encoding="utf-8"))
    plan0 = LayoutPlan.model_validate_json((vdir / "layout_plan.json").read_text(encoding="utf-8"))
    report0 = AuditReport.model_validate_json((vdir / "audit_report.json").read_text(encoding="utf-8"))
    deck0 = vdir / "deck.pptx"
    brief_text = brief.text if brief is not None else None
    wishes = (wishes or "").strip() or None
    remarks = RM.order_remarks(remarks)
    before_stage = RM.stage_remarks(report0.issues, n)
    osl0 = outline0.slides[n - 1]
    oid = osl0.id
    frame = osl0.kind.value in FRAME_KINDS
    asked_checks = {r.check_id for r in remarks}
    notes: list[str] = []
    records: list[dict] = []  # every fix record of this fix, rolled-back tries included (the history)
    kept: list[dict] = []  # the edits that stayed (what the reply says was done)
    how, what = "autofix", ""
    redesigned = False
    spent: list[dict] = []  # what the designer's calls of this fix cost (run_manifest.usage)

    voice.plain("Критик читает замечания", 0.03)
    voice.agent("critic", RM.critic_line(n, remarks, titles), 0.05)

    def fixable(rep: AuditReport, xml_only: bool) -> set[str]:
        return {
            i.id for i in rep.issues
            if i.slide == n and i.severity in ("error", "warn") and i.kind != "model" and i.autofix is not None and i.autofix.action != "none"
            and (not xml_only or i.autofix.action not in RERENDER_ACTIONS)
        }

    def loop(rep: AuditReport, outline: DeckOutline, plan: LayoutPlan, ids: set[str], iterations: int):
        start = len(rep.applied_fixes)
        rep2, plan2, outline2, _ = autofix_loop(work, rep, outline, plan, manifest, ws, max_iterations=iterations, only_ids=ids, render=False, brief_text=brief_text)
        new = [dict(r, slide=n) for r in rep2.applied_fixes[start:]]
        rolled = {r.get("iteration") for r in new if r.get("action") == "rollback"}
        records.extend(new)
        kept.extend(r for r in new if r.get("action") != "rollback" and r.get("iteration") not in rolled)
        return rep2, outline2, plan2

    def _fix_seen(deck: Path, k: int, codes: set, mf, where: Path, shown: Optional[list] = None) -> SF.SeenFix:
        """The in-place fixes of the review's `codes` on slide k of `deck`; their records join the history."""
        pic = SF.slide_picture(deck, k, where) if codes & {"low_contrast", "text_cut"} else None
        again = iter(range(1, 10))
        got = SF.fix_seen(deck, k, codes, mf, pic, shown or (), render=lambda d: SF.slide_picture(d, k, where / f"r{next(again)}"))
        recs = [{"iteration": 0, "action": "xml", **r} for r in got.records]
        records.extend(recs)
        kept.extend(recs)
        return got

    def mark() -> tuple[int, int]:
        return len(records), len(kept)

    def back_to(m: tuple[int, int]) -> tuple[list[dict], list[dict]]:
        """Forget the records since mark `m`; returns them (to put them back)."""
        gone = (records[m[0]:], kept[m[1]:])
        del records[m[0]:]
        del kept[m[1]:]
        return gone

    def slide_score(rep: AuditReport) -> tuple[int, int]:
        own = [i for i in rep.issues if i.slide == n and i.kind != "model"]
        return sum(1 for i in own if i.severity == "error"), sum(1 for i in own if i.severity == "warn")

    def inplace(rep: AuditReport, outline: DeckOutline, plan: LayoutPlan, ids: set[str]):
        """The XML edits of `ids`. A move into the template's margins can land on another block (a footnote at the
        bottom edge moved up onto the text above it): the same edits with the move kept inside the slide only are
        tried too, and the version with fewer errors, then fewer warnings on this slide wins."""
        into_margins = {i.id for i in rep.issues if i.id in ids and i.autofix is not None and i.autofix.action == "move_inside" and i.autofix.params.get("safe")}
        if not into_margins:
            return loop(rep, outline, plan, ids, 1)
        start_deck, m = scratch / "inplace_start.pptx", mark()
        shutil.copy2(work, start_deck)
        first = loop(rep, outline, plan, ids, 1)
        if slide_score(first[0]) == (0, 0):
            return first
        first_records = back_to(m)
        shutil.copy2(work, scratch / "inplace_first.pptx")
        shutil.copy2(start_deck, work)
        soft = rep.model_copy(deep=True)
        for i in soft.issues:
            if i.id in into_margins and i.autofix is not None:
                i.autofix.params = {**i.autofix.params, "safe": False}
        second = loop(soft, outline, plan, ids, 1)
        if slide_score(second[0]) < slide_score(first[0]):
            log.info("slide fix: slide %d — moved inside the slide rather than into the margins (%s < %s)", n, slide_score(second[0]), slide_score(first[0]))
            return second
        back_to(m)
        records.extend(first_records[0])
        kept.extend(first_records[1])
        shutil.copy2(scratch / "inplace_first.pptx", work)
        return first

    scratch = vdir.parent / f".{strategy}.fix"
    shutil.rmtree(scratch, ignore_errors=True)
    scratch.mkdir(parents=True)
    try:
        work = scratch / "work.pptx"
        shutil.copy2(deck0, work)
        outline, plan = outline0.model_copy(deep=True), plan0.model_copy(deep=True)
        # autofix reads the slides' outline ids from report.deck; nothing below renders images or touches vdir
        report = report0.model_copy(update={"deck": str(work)}, deep=True)
        audit_kw = dict(render=False, strategy=strategy, brief_text=brief_text)
        rerendered = False

        # A. in place first: the remarks' own XML edits
        ids_xml = {r.id for r in remarks if r.autofix is not None and r.autofix.action not in RERENDER_ACTIONS and r.autofix.action != "none" and r.kind != "model"}
        rep = report
        if ids_xml:
            rep, outline, plan = inplace(report, outline, plan, ids_xml)
        # A'. what «Взгляд дизайнера» saw, fixed in place by measuring (audit/seen_fix.py): the text's colour or a plate
        # where it does not read on the ground the slide shows, a larger size, a box cut short of what it runs into, no
        # letter-spacing, a smaller size — on a cover too (the designer never lays a frame out again)
        seen_codes = {c for c in (SF.code_of(r) for r in remarks if r.kind == "model") if c}
        seen: Optional[SF.SeenFix] = None
        if seen_codes:
            voice.plain(f"Правлю то, что увидел дизайнер на слайде {n}", 0.07)
            shown = [p for p in (vdir / "slides" / f"slide-{n:03d}.jpg",) if p.exists()]  # the preview the person saw
            seen = _fix_seen(work, n, seen_codes, manifest, scratch / "seen", shown)
            if seen.records:
                rep = run_audit(work, manifest, outline, ws, **audit_kw)
        # what is left of the asked remarks (the model's own remarks stay until the designer has redesigned the slide,
        # or the rules fixed what they say in place)
        left = [i for i in rep.issues if i.slide == n and RM.is_stage(i) and i.kind != "model" and i.check_id in asked_checks]
        left += [r for r in remarks if r.kind == "model" and not (seen is not None and SF.code_of(r) in seen.cured)]

        # B. the designer, when content is at stake
        # (a remark the rules have no fix for — a figure hardly larger than its label — goes to the designer too: the
        # button never leaves the slide as it was while a remark is on it)
        need_designer = not frame and bool(wishes or any(i.check_id in RM.CONTENT_CHECKS or i.autofix is None or i.autofix.action == "none" for i in left))
        if need_designer and use_models:
            voice.plain(f"Дизайнер переделывает слайд {n}", 0.1)
            left_ids = {id(i) for i in left}
            source = left if left else remarks
            note_items = [i for i in source if i.check_id not in RM.TEMPLATE_CHECKS and (i.check_id not in RM.INPLACE_FIRST or id(i) in left_ids)]
            if note_items:
                start = "переделываю по замечаниям" + (f" и пожеланию «{wishes[:120]}»" if wishes else "")
            else:
                start = f"переделываю по пожеланию «{(wishes or '')[:120]}»"
            named = RM.titles_phrase(remarks, titles)
            rationale = (f"Исправлены замечания: {named}" + (f" · по вашей просьбе «{wishes}»" if wishes else "")) if named else f"По вашей просьбе «{wishes or ''}»"

            def design(base: DeckOutline, note: str, start_line: str):
                """One designer's version of the slide, rendered in the scratch deck, audited and fixed in place."""
                try:
                    r = redesign_slide(
                        base, n, brief, manifest, request=wishes or "", note=note, start_line=start_line, rationale=rationale,
                        log_line=f"Исправление: слайд {n} — {{what}}.", skills=skills, providers=providers, progress=voice.relay(),
                    )
                except ValueError as e:
                    return None, str(e)
                spent.append(r.usage)
                o = _with_slide(outline0, r.outline, n)
                strategy_obj = load_strategies().get(strategy) or next(iter(load_strategies().values()))
                pl = _with_entry(plan0, match_outline(o, manifest, strategy_obj).for_outline(oid), oid)
                voice.plain(f"Собираю слайд {n}", 0.52)
                render_deck(o, pl, manifest, ws, work)  # the whole scratch deck; only slide n is used
                rp = run_audit(work, manifest, o, ws, **audit_kw)
                ids = fixable(rp, xml_only=True)
                if ids:
                    rp, o, pl = inplace(rp, o, pl, ids)
                return (r, o, pl, rp), None

            got, reason = design(outline0, RM.remarks_note(note_items, titles, wishes), start)
            if got is None and reason and "не изменился" in reason:
                # the model gave the slide back as it was: once more, told so (a remark never leaves the slide as it was)
                got, reason = design(outline0, RM.remarks_note(note_items, titles, wishes) + "\n" + RM.UNCHANGED_LINE, "пробую ещё раз — прежний ответ слайд не менял")
            if got is not None:
                r, o, pl, rp = got
                # (any asked remark of the rules the new version still has: a figure still hardly larger than its label)
                still = [i for i in rp.issues if i.slide == n and RM.is_stage(i) and i.kind != "model" and i.check_id in asked_checks]
                if still and r.how == "model":
                    # one more try: the form the designer kept cannot show the slide's content
                    first_deck, m = scratch / "design_first.pptx", mark()
                    shutil.copy2(work, first_deck)
                    retry_note = RM.remarks_note(still, titles, wishes) + "\n" + RM.retry_line(o.slides[n - 1].kind.value)
                    again, _ = design(o, retry_note, "пробую другую форму — прежняя не вместила всё")
                    if again is not None and slide_score(again[3]) < slide_score(rp):
                        got = again
                    else:
                        back_to(m)
                        shutil.copy2(first_deck, work)
                r, o, pl, rp = got
                how, redesigned, rerendered = r.how, True, True
                what = describe_change(osl0, o.slides[n - 1])
                outline, plan, rep = o, pl, rp
            elif reason == "model_unavailable":
                if wishes:
                    notes.append("Пожелания не учтены: модель не ответила")
                voice.agent("designer", "Дизайнер: модель не ответила — правлю оформление по правилам.", 0.3)
            else:
                if wishes:
                    notes.append(f"Пожелания не учтены: {reason}")
                voice.agent("designer", f"Дизайнер: слайд {n} — {reason}; правлю оформление по правилам.", 0.3)
        elif need_designer:
            if wishes:
                notes.append("Пожелания не учтены: модель недоступна")
            voice.agent("designer", f"Дизайнер: слайд {n} — правлю оформление по правилам шаблона.", 0.15)
        elif frame:
            name, note = _frame_name(osl0.kind.value)
            if wishes:
                notes.append(note)
            voice.agent("designer", f"Дизайнер: {name} не переделываю — правлю только оформление.", 0.15)
        else:
            voice.agent("designer", f"Дизайнер: слайд {n} — правлю оформление по правилам шаблона.", 0.15)

        # C. structural fixes of what is left on the slide, rolled back when worse — none after a redesign for a wish
        # about the look: a shortened text or another sample would undo what was asked and cost the slide's conclusions
        from verstka.planning.slide_edit import style_only as _style_only

        ids = fixable(rep, xml_only=bool(redesigned and wishes and _style_only(wishes)))
        if ids:
            structural = any(a.action in RERENDER_ACTIONS for acts in plan_fixes(rep, plan, ids).values() for a in acts)
            if structural and not rerendered:
                # the other slides as a re-render shows them, so the loop's before/after compares this slide only
                try:
                    baseline = scratch / "baseline.pptx"
                    render_deck(outline.model_copy(deep=True), plan.model_copy(deep=True), manifest, ws, baseline)
                    splice_slide(baseline, work, n, scratch / "work_c.pptx")
                    (scratch / "work_c.pptx").replace(work)
                    rep = run_audit(work, manifest, outline, ws, **audit_kw)
                    ids = fixable(rep, xml_only=False)
                    rerendered = True
                except Exception:  # noqa: BLE001 - the loop still runs on the work deck as it is
                    log.warning("slide fix: no baseline for the structural pass", exc_info=True)
            if ids:
                # the autofix ladder on this slide (audit/autofix.py): the next sample, the one after it, then the slide
                # composed — a rung that made the slide worse rolled back, the next one tried (two rungs stopped at the
                # second sample, and a slide whose texts no sample holds was never composed)
                rep, outline, plan = loop(rep, outline, plan, ids, 3)
        # C'. a remark of the review the rules found nothing to change for (or could not fix in place) and the designer
        # did not cure — a cover, a divider, no model, the designer's slide the same: the slide on another sample of the
        # template (the eyes' ladder: the next sample, then the composer). A remark never leaves the slide as it was
        relaid = laddered = False
        pending = [r for r in remarks if r.kind == "model" and not (seen is not None and SF.code_of(r) in seen.cured)]
        asked_rules = Counter(r.check_id for r in remarks if r.kind != "model")

        def asked_open(rp: AuditReport) -> int:
            """The asked remarks of the rules still on the slide in report `rp` (by check, at most as many as asked)."""
            have = Counter(i.check_id for i in rp.issues if i.slide == n and RM.is_stage(i) and i.kind != "model")
            return sum(min(k, have[c]) for c, k in asked_rules.items())

        if (pending or asked_open(rep)) and not wishes:
            fp_now = slide_fingerprints(work)
            same = len(fp_now) < n or fp_now[n - 1] == slide_fingerprints(deck0)[n - 1]
            # (the designer's version that kept an asked remark of the rules goes up the ladder too)
            if not redesigned or same or asked_open(rep) >= asked_open(report):
                from verstka.audit.eyes import relayout

                voice.plain(f"Собираю слайд {n} на другом образце шаблона", 0.5)
                errors_now, warns_now = slide_score(rep)
                trial = scratch / "relayout.pptx"
                pl_try = plan.model_copy(deep=True)
                for _ in range(3):
                    laddered = True
                    o_try = outline.model_copy(deep=True)
                    step = relayout(pl_try, o_try, oid, manifest, set())
                    log.info("slide fix: slide %d — next layout: %s", n, step)
                    if step == "skip":
                        break
                    try:
                        render_deck(o_try, pl_try, manifest, ws, trial)
                        rp = run_audit(trial, manifest, o_try, ws, **audit_kw)
                    except Exception:  # noqa: BLE001 - the next rung
                        log.warning("slide fix: slide %d on another sample failed", n, exc_info=True)
                        pl_try = pl_try.model_copy(deep=True)
                        continue
                    fp_try = slide_fingerprints(trial)
                    e_try, w_try = slide_score(rp)
                    log.info("slide fix: slide %d on it — errors %d, warnings %d (were %d, %d), asked open %d (was %d)", n, e_try, w_try, errors_now, warns_now, asked_open(rp), asked_open(rep))
                    # a layout that brings a remark of its own (a line gone to the notes, text on the template's art) is
                    # no fix: the next rung (pl_try keeps what was tried)
                    if e_try > errors_now or w_try > warns_now or len(fp_try) < n or fp_try[n - 1] == slide_fingerprints(deck0)[n - 1]:
                        continue
                    lost = lost_words_share(slide_words(deck0, n), slide_words(trial, n))
                    if lost > 0.05:
                        log.info("slide fix: slide %d on it loses %.0f%% of its words — the next rung", n, lost * 100)
                        continue  # a layout that drops the slide's text (a subtitle with no place) is no fix
                    if not pending and asked_open(rp) >= asked_open(rep):
                        continue  # the asked remark of the rules is still there: the next rung
                    shutil.copy2(trial, work)
                    outline, plan, rep = o_try, pl_try, rp
                    relaid = rerendered = True
                    how = "relayout"
                    what = "слайд собран на другом образце шаблона" if step.startswith("rematch") and "synth" not in step else "слайд собран заново по его содержанию"
                    log.info("slide fix: slide %d laid out again for the review's remarks (%s)", n, step)
                    voice.agent("designer", f"Дизайнер: слайд {n} собран заново — {what[:1].lower() + what[1:]}.", 0.55)
                    break
        if seen is not None and seen.cured and (redesigned or rerendered):
            # the slide laid out again (the designer, another sample) came without the in-place fixes: made again where
            # they are still due (each fix measures first)
            _fix_seen(work, n, set(seen.cured), manifest, scratch / "seen_again")
        if not what:
            what = _what_of(kept)

        voice.plain(f"Собираю слайд {n}", 0.58)
        if rerendered:
            voice.agent("compile", f"Вёрстка: слайд {n} собран по макету шаблона, остальные слайды не трогаю.", 0.6)
        else:
            voice.agent("compile", f"Вёрстка: слайд {n} поправлен на месте, остальные слайды не трогаю.", 0.6)

        # D. splice + verify
        voice.plain("Проверяю слайд и остальные", 0.65)
        final = scratch / "final.pptx"
        try:
            splice_slide(deck0, work, n, final)
        except Exception:  # noqa: BLE001 - the verification below tells the truth about the other slides
            log.warning("slide fix: splice failed, the scratch deck is used as it is", exc_info=True)
            shutil.copy2(work, final)
        try:
            from verstka.rendering.renderer import strip_outline_marks

            strip_outline_marks(final)  # the spliced slide came with its machine marker
        except Exception:  # noqa: BLE001
            log.warning("outline markers left in the notes", exc_info=True)
        final_outline = _with_slide(outline0, outline, n)
        final_plan = _with_entry(plan0, plan.for_outline(oid), oid)
        pic_records = _keep_slide_picture(vdir, final, final_outline, n)
        # the audit numbers its remarks through the deck: the ids of the remarks that stayed (on every slide) are carried
        # over, so the person's remark keeps its id through the fix and a fixed remark's id never comes back
        report = RM.carry_report_ids(report0, run_audit(final, manifest, final_outline, ws, **audit_kw))
        # the model's remarks cannot be checked again without a model: the other slides keep theirs, and this slide
        # keeps its own unless the designer redesigned it
        cured_ids = {r.id for r in remarks if r.kind == "model" and seen is not None and SF.code_of(r) in seen.cured}
        report.issues.extend(i.model_copy(deep=True) for i in report0.issues if i.kind == "model" and (i.slide != n or not (redesigned or relaid or i.id in cured_ids)))
        fp0, fp1 = slide_fingerprints(deck0), slide_fingerprints(final)
        slide_changed = len(fp1) >= n and fp0[n - 1] != fp1[n - 1]
        # the model looks again at the slide as it is now: a remark it still sees stays (never called fixed on the
        # rules' word alone); without a vision model the rules' measure stands
        still_seen: set[str] = set()
        asked_seen = [r for r in remarks if r.kind == "model"]
        looked = False
        if use_models and asked_seen:
            voice.plain(f"Дизайнер смотрит на слайд {n} ещё раз", 0.7)
            pic = SF.slide_picture(final, n, scratch / "relook")
            got = SF.relook(pic, n, providers=providers, skills=skills, headline=final_outline.slides[n - 1].headline or "") if pic else None
            log.info("slide fix: slide %d looked at again — the model sees %s", n, got if got is not None else "nothing (no answer)")
            if got is not None:
                looked = True
                # the asked remarks are judged by this look alone: kept ones go, the ones it still sees come back
                asked_ids = {r.id for r in asked_seen}
                report.issues = [i for i in report.issues if i.id not in asked_ids]
                # by count: a problem the model still sees is first the remark of its kind the slide kept (one not
                # asked, one not fixed), then an asked one — the heading moved off the circle is fixed while the card
                # under the circle is still seen
                seen_now = Counter(got)
                held = Counter(SF.code_of(i) for i in report.issues if i.kind == "model" and i.slide == n)
                have = {i.id for i in report.issues}
                # a remark about the heading when the heading was what the rules changed is the last to stay
                titled = seen is not None and seen.titled
                for r in sorted(asked_seen, key=lambda r: titled and "заголов" in (r.message or "").lower()):
                    c = SF.code_of(r)
                    if r.id not in have and seen_now[c] > held[c]:
                        report.issues.append(r.model_copy(deep=True))
                        held[c] += 1
                        still_seen.add(c or "")
                if still_seen:
                    voice.agent("check", f"Взгляд дизайнера: на слайде всё ещё {', '.join(SF.seen_words(c) for c in sorted(still_seen))}.", 0.72)
                else:
                    voice.agent("check", "Взгляд дизайнера: посмотрел на исправленный слайд — замечание ушло.", 0.72)
        report.applied_fixes = list(report0.applied_fixes) + [{"action": "slide_fix", "slide": n, "outline_id": oid, "result": how}] + records
        report.iterations = report0.iterations
        report.recompute()
        other = changed_slides(fp0, fp1, skip={n})
        after_stage = RM.stage_remarks(report.issues, n)
        known = {i.id for i in before_stage}
        fixed, remaining = RM.match_fixed(remarks, after_stage, before=before_stage + [r for r in remarks if r.id not in known])
        errors_before = sum(1 for i in report0.issues if i.slide == n and i.severity == "error")
        errors_after = sum(1 for i in report.issues if i.slide == n and i.severity == "error")
        worse = errors_after > errors_before
        applied = not worse and slide_changed and (bool(fixed) or bool(wishes))
        # a wish about the look («цифры крупнее») never costs a conclusion — neither the designer's version nor the
        # autofix's shortening of it afterwards (qa4 NEW-E1, live: «…больше 120 мин, 3 — больше 10 часов» condensed away)
        from verstka.planning.slide_edit import lost_lines, style_only

        lost_words = bool(wishes) and style_only(wishes or "") and not frame and bool(lost_lines(osl0, final_outline.slides[n - 1]))
        if lost_words:
            applied = False
        why = None
        # nothing to change after all: the slide is the same, the model looked again and sees none of the asked problems,
        # and no asked remark of the rules is left — the remarks go, and the reply says so (not «остался прежним»)
        clear = not slide_changed and not worse and looked and not still_seen and not wishes and not remaining_asked(fixed, remarks)
        if clear:
            why = "Замечаний больше нет — дизайнер посмотрел на слайд ещё раз и проблем не видит"
        elif not applied:
            if lost_words:
                why = "Сделать крупнее, не потеряв выводы слайда, не получилось — оставил как было"
            elif worse:
                why = "Новая версия слайда вышла хуже — оставил как было"
            elif wishes:
                # (a wish was given: «попробуйте добавить пожелание» would ask for what the person already did)
                why = "Пожелание не удалось выполнить — слайд оставлен как был"
            elif still_seen:
                said = ", ".join(SF.seen_words(c) for c in sorted(still_seen))
                tried = "; на других образцах шаблона слайд выходит хуже" if laddered and not relaid else ""
                why = f"Дизайнер всё ещё видит: {said}{tried} — попробуйте другой вариант презентации или пожелание"
            elif seen is not None and seen.left:
                why = seen.why_for(frame, {"section": "разделителя", "thanks": "финального слайда"}.get(osl0.kind.value, "обложки") if frame else None)
            elif laddered and not relaid:
                why = "Замечания остались: ни правки на месте, ни другие образцы шаблона их не убрали — попробуйте пожелание или другой вариант презентации"
            else:
                why = "Замечания остались — попробуйте добавить пожелание"
        voice.agent("check", RM.check_line(fixed, remarks, remaining, applied, worse, titles, other=other), 0.8)
        if applied and other:
            voice.agent("check", f"Проверка: изменились и слайды {RM.slides_phrase(other)}.", 0.8)
        score_before, score_after = report0.summary.score, report.summary.score
        result = SlideFix(
            applied=applied, how=how, what=what, reply="", report=report if applied else report0, requested=remarks, fixed=fixed if applied else [],
            remaining=remaining if applied else [{"id": i.id, "check_id": i.check_id, "severity": i.severity, "message": i.message, "new": False} for i in before_stage],
            other=other if applied else [], notes=notes, why=why, score_before=score_before, score_after=score_after if applied else score_before,
            n_slides=len(final_outline.slides),
        )
        result.reply = RM.fix_reply(n, result.fixed, remarks, result.remaining, score_before, result.score_after, notes, applied, why, result.other, what, titles)
        if clear:
            # the report without the remarks the model no longer sees (the deck and its previews stay as they are)
            report.deck, report.slide_images = report0.deck, dict(report0.slide_images)
            (vdir / "audit_report.json").write_text(report.model_dump_json(indent=2), encoding="utf-8")
            result.report, result.fixed, result.remaining, result.cleared = report, fixed, remaining, True
            result.score_after = report.summary.score
            result.reply = f"Слайд {n}: замечаний больше нет — дизайнер посмотрел на него ещё раз и проблем не видит."
            log_edit(vdir, {
                "at": time.time(), "version": None, "undo_of": None, "kind": "fix", "slides": [n], "request": fix_label(n, wishes),
                "reply": result.reply, "score_before": score_before, "score_after": result.score_after, "issue_ids": [r.id for r in remarks],
                "fixed": fixed, "wishes": wishes, "how": "look", "other_slides": [],
            })
            return result
        if not applied:
            return result

        # E. apply: previews, the previous version kept, the files replaced, the fix logged. LibreOffice takes a while
        # (≈ 20 s): the check step says the previews are on their way and the bar keeps moving, so the timeline never
        # reads as finished while the job still works
        voice.plain("Готовлю превью", 0.85)
        voice.agent("check", RM.preview_line(result.fixed, remarks, result.remaining), 0.85)
        out = scratch / "out"
        out.mkdir()
        shutil.copy2(final, out / "deck.pptx")
        exports = [e for e in (exports or []) if e in ("pdf", "html")]
        with _Creep(voice, "Готовлю превью", 0.85, 0.95):
            _, images, warns, render_ok = render_outputs(out, manifest, final_outline.title, exports, images=True)
        for w in warns:
            log.warning("slide fix: %s", w)
        if render_ok is False:
            report.issues.append(Issue(id="file_opens-0-render", slide=0, check_id="file_opens", severity="warn", kind="deterministic", message="LibreOffice не смог отрендерить файл"))
            report.recompute()
            result.score_after = report.summary.score
        voice.plain("Готовлю превью", 0.95)
        # previews: the fixed slide (and any slide that changed) from the new render; every other slide keeps its image
        merged = scratch / "slides"
        merged.mkdir()
        old_images = {int(p.stem.rsplit("-", 1)[-1]): p for p in _previews(vdir)}
        new_images = {k: p for k, p in enumerate(images, 1)}
        for k in range(1, len(final_outline.slides) + 1):
            src = new_images.get(k) if (k == n or k in other or k not in old_images) else old_images[k]
            if src is None:
                src = old_images.get(k)
            if src is None:
                continue
            if k != n and k not in other and k in new_images and k in old_images:
                d = image_drift(old_images[k], new_images[k])
                if d is not None and d > 3:
                    log.info("slide fix: image drift on slide %d (%.1f) with an equal fingerprint — the old preview is kept", k, d)
            shutil.copy2(src, merged / f"slide-{k:03d}.jpg")
        version = snapshot(vdir)
        (out / "deck.pptx").replace(deck0)
        if pic_records is not None:
            from verstka.images.fill import write_records

            write_records(vdir, pic_records)
        for fmt in ("pdf", "html"):
            made = out / f"deck.{fmt}"
            if made.exists():
                made.replace(vdir / f"deck.{fmt}")
            elif (vdir / f"deck.{fmt}").exists():
                (vdir / f"deck.{fmt}").unlink()  # it would show the deck before the fix
                if fmt in exports:
                    log.warning("slide fix: deck.%s was not exported again", fmt)
        if any(merged.iterdir()):
            shutil.rmtree(vdir / "slides", ignore_errors=True)
            merged.replace(vdir / "slides")
        final_outline.agent_log = [*outline0.agent_log, f"Исправление: слайд {n} — {what or 'поправлено оформление'}."]
        report.deck = str(deck0)
        report.slide_images = {k: str(p) for k, p in enumerate(_previews(vdir), 1)}
        (vdir / "outline.json").write_text(final_outline.model_dump_json(indent=2), encoding="utf-8")
        (vdir / "layout_plan.json").write_text(final_plan.model_dump_json(indent=2), encoding="utf-8")
        (vdir / "audit_report.json").write_text(report.model_dump_json(indent=2), encoding="utf-8")
        rm = _read_json(vdir / "run_manifest.json")
        if rm is not None:
            rm["audit"] = report.summary.model_dump()
            rm["applied_fixes"] = report.applied_fixes
            rm.setdefault("inputs", {})["outline_sha256"] = hashlib.sha256(final_outline.model_dump_json().encode("utf-8")).hexdigest()[:16]
            for u in spent:
                if u:
                    rm["usage"] = add_usage(rm.get("usage"), u)
            entry = final_plan.for_outline(oid)
            slides = rm.get("slides")
            if entry is not None and isinstance(slides, list) and len(slides) >= n and isinstance(slides[n - 1], dict):
                slides[n - 1].update({"mode": entry.mode, "pattern_id": entry.pattern_id, "composition": entry.composition})
            (vdir / "run_manifest.json").write_text(json.dumps(rm, ensure_ascii=False, indent=2), encoding="utf-8")
        result.report, result.version, result.at = report, version, time.time()
        result.reply = RM.fix_reply(n, result.fixed, remarks, result.remaining, score_before, result.score_after, notes, applied, why, result.other, what, titles)
        log_edit(vdir, {
            "at": result.at, "version": version, "undo_of": None, "kind": "fix", "slides": [n], "request": fix_label(n, wishes),
            "reply": result.reply, "score_before": score_before, "score_after": result.score_after, "issue_ids": [r.id for r in remarks],
            "fixed": result.fixed, "wishes": wishes, "how": how, "other_slides": result.other,
        })
        return result
    finally:
        shutil.rmtree(scratch, ignore_errors=True)
