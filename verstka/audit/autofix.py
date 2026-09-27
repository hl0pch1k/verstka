"""Autofix: turn audit issues into plan/outline changes and XML edits, re-render, re-audit (≤ 2 iterations, rollback on regression)."""

from __future__ import annotations

import copy
import logging
import shutil
from pathlib import Path
from typing import Optional

from lxml import etree
from pptx import Presentation

from verstka.analysis.colors import nearest_palette_color
from verstka.analysis.xmlns import q
from verstka.audit.runner import run_audit
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.compat import composition_for
from verstka.planning.condense import trim_words
from verstka.providers.registry import ProviderRegistry
from verstka.rendering.deck import element_bbox, is_nested, set_element_pos, slide_shape_elements
from verstka.rendering.renderer import RenderResult, render_deck
from verstka.schemas.audit import AuditReport, FixAction, Issue
from verstka.schemas.common import Bbox
from verstka.schemas.layout import LayoutPlan
from verstka.schemas.outline import DeckOutline
from verstka.schemas.template import TemplateManifest
from verstka.skills_registry.registry import SkillsRegistry

log = logging.getLogger(__name__)

STRUCTURAL_ACTIONS = ("rematch", "synth")  # need a full re-render of the slide
RERENDER_ACTIONS = STRUCTURAL_ACTIONS + ("condense_text",)
MIN_FONT_PT = 8.0  # body text below this is unreadable on a projector; footnote sizes of the scale are not shrink targets


def plan_fixes(report: AuditReport, plan: LayoutPlan, only_ids: Optional[set[str]] = None, *, xml_only: bool = False) -> dict[str, list[FixAction]]:
    """outline_id → actions (at most one structural action per slide, plus text/colour actions).

    Info-level findings are advisory and never fixed unless the caller asked for them by issue id.
    With xml_only, only in-place XML edits are returned (used to re-derive fixes after a re-render)."""
    by_slide: dict[str, list[FixAction]] = {}
    for issue in report.issues:
        if only_ids is not None and issue.id not in only_ids:
            continue
        if not issue.autofix or issue.kind == "model" or issue.autofix.action == "none":
            continue
        if issue.severity == "info" and only_ids is None:
            continue
        if xml_only and issue.autofix.action in RERENDER_ACTIONS:
            continue
        oid = issue.outline_id or issue.autofix.params.get("outline_id")
        if not oid:
            continue
        by_slide.setdefault(oid, []).append(issue.autofix)
    # collapse: one structural action per slide; a re-render supersedes geometry/size edits of the same slide
    out: dict[str, list[FixAction]] = {}
    for oid, actions in by_slide.items():
        structural = [a for a in actions if a.action in STRUCTURAL_ACTIONS]
        others = [a for a in actions if a.action not in STRUCTURAL_ACTIONS]
        chosen: list[FixAction] = []
        if structural:
            chosen.append(structural[0])
            others = [a for a in others if a.action not in ("shrink_text", "move_inside")]
        chosen.extend(others)
        if chosen:
            out[oid] = chosen
    return out


def _empty_frame(pattern, osl) -> bool:
    """A sample that would show an empty photo frame or a device mock-up on this slide: it has a picture slot or a
    mock-up and the slide brings no picture of its own."""
    if pattern is None or getattr(osl.content, "image_hint", None):
        return False
    from verstka.schemas.common import SlotRole

    return bool(getattr(pattern, "mockup_boxes", None)) or any(sl.role == SlotRole.image for sl in pattern.slots)


def _unsay(outline: DeckOutline) -> int:
    """Before a re-render: the lines a previous render said aloud because its slide had no room for them (a cover's
    goal or small print, `clone._note_to_speaker_notes` / synth `_to_notes`) come off the speaker notes — the new
    render places them on the slide or says them again. Without this a cover re-rendered by another sample showed its
    goal on the slide and repeated it in the notes (LO Vivid long). Returns the number of lines removed."""
    from verstka.matching.scorer import cover_goal

    removed = 0
    for s in outline.slides:
        if not s.notes:
            continue
        said = {" ".join(t.split()) for t in (cover_goal(s), s.footnote) if t and t.strip()}
        if not said:
            continue
        lines = s.notes.split("\n")
        keep = [ln for ln in lines if " ".join(ln.split()) not in said]
        if len(keep) != len(lines):
            removed += len(lines) - len(keep)
            try:
                s.notes = "\n".join(keep).strip()
            except Exception:  # noqa: BLE001 - a frozen outline keeps its notes
                pass
    return removed


def _rematch(plan: LayoutPlan, outline: DeckOutline, oid: str, manifest: Optional[TemplateManifest] = None) -> str:
    ps = plan.for_outline(oid)
    osl = next((s for s in outline.slides if s.id == oid), None)
    if ps is None or osl is None:
        return "skip"
    if ps.mode == "synth" and ps.composition and ps.composition not in ("title", "section", "thanks"):
        return "skip"  # a composed slide is laid out from the content: a sample slide would only fit it worse
    tried = set(ps.fit.get("tried", []))
    if ps.pattern_id:
        tried.add(ps.pattern_id)
    bookend = getattr(osl.kind, "value", osl.kind) in ("title", "section", "thanks")
    patterns = {p.id: p for p in manifest.patterns} if manifest is not None else {}
    for pid, score in ps.alternatives:
        if bookend and _empty_frame(patterns.get(pid), osl):
            continue  # a cover never trades its heading for an empty photo frame or a phone mock-up
        if pid not in tried and score >= 0.3:
            ps.mode = "clone"
            ps.pattern_id = pid
            ps.fit["tried"] = sorted(tried)
            ps.reasons.append(f"автофикс: перевыбран паттерн {pid} ({score:.2f})")
            return f"rematch → {pid}"
    if ps.mode != "synth":
        ps.mode = "synth"
        ps.pattern_id = None
        ps.composition = composition_for(osl)
        ps.fit["tried"] = sorted(tried)
        ps.reasons.append(f"автофикс: синтез композиции {ps.composition}")
        return f"synth → {ps.composition}"
    return "skip"


def _condense(outline: DeckOutline, oid: str, skills: Optional[SkillsRegistry], providers: Optional[ProviderRegistry]) -> str:
    from verstka.planning.condense import condense_text

    osl = next((s for s in outline.slides if s.id == oid), None)
    if osl is None:
        return "skip"
    c = osl.content
    changed = 0
    new_bullets = []
    for b in c.bullets[:6]:
        nb = condense_text(b, 15, skills, providers, outline.language)
        changed += nb != b
        new_bullets.append(nb)
    changed += len(c.bullets) > 6
    c.bullets = new_bullets
    for it in c.items + c.columns:
        if it.text and len(it.text.split()) > 24:
            it.text = condense_text(it.text, 24, skills, providers, outline.language)
            changed += 1
        it.bullets = [condense_text(b, 15, skills, providers, outline.language) for b in it.bullets[:6]]
    for p_i, p in enumerate(c.paragraphs):
        if len(p.split()) > 60:
            c.paragraphs[p_i] = trim_words(p, 60)
            changed += 1
    return f"condensed {changed} texts" if changed else "skip"


# ----------------------------------------------------------------------------- XML edits


def _snap_size(size_pt: float, factor: float, scale: list[float]) -> float:
    """Next smaller step of the template's type scale that is ≤ size·factor, so the result stays in scale.

    With no step small enough, the smallest step below the current size is used (a partial shrink is still in scale);
    below the scale's floor the raw target is used only while it stays within 2 pt of a scale step (audit: info, not warn),
    otherwise the size is left alone — an out-of-scale 6 pt run is worse than the overflow it tries to hide."""
    target = size_pt * factor
    lower = sorted((s for s in scale if MIN_FONT_PT <= s < size_pt - 0.25), reverse=True)
    for s in lower:
        if s <= target + 0.25:
            return max(s, MIN_FONT_PT)
    if lower:
        return max(lower[-1], MIN_FONT_PT)
    raw = max(round(target * 2) / 2, MIN_FONT_PT)
    if not scale or any(abs(raw - s) <= 2.0 for s in scale):
        return raw
    return size_pt


def _force_run_color(rPr: etree._Element, hex_: str) -> int:
    """Give a run an explicit solid colour when its colour is inherited or a scheme colour. Returns 1 if changed."""
    fill = rPr.find(q("a:solidFill"))
    if fill is not None:
        if fill.find(q("a:srgbClr")) is not None:
            return 0  # handled by the srgbClr pass
        for child in list(fill):
            fill.remove(child)
        etree.SubElement(fill, q("a:srgbClr")).set("val", hex_)
        return 1
    if any(rPr.find(q(t)) is not None for t in ("a:noFill", "a:gradFill", "a:pattFill", "a:blipFill", "a:grpFill")):
        return 0
    fill = etree.Element(q("a:solidFill"))
    etree.SubElement(fill, q("a:srgbClr")).set("val", hex_)
    ln = rPr.find(q("a:ln"))
    rPr.insert(rPr.index(ln) + 1 if ln is not None else 0, fill)  # schema order: ln, fill, effects, …, latin
    return 1


def _recolor(el: etree._Element, scope: str, only_hex: Optional[str], target_hex: Optional[str], palette: list[str]) -> int:
    """Replace colours inside el. scope='text' touches p:txBody only (runs, bullets); 'all' also the shape fill/outline."""
    roots = list(el.iter(q("p:txBody"))) if scope == "text" else [el]
    changed = 0
    for root in roots:
        for clr in list(root.iter(q("a:srgbClr"))):
            val = (clr.get("val") or "").upper()
            if only_hex and val != only_hex:
                continue
            new_hex = target_hex or (nearest_palette_color(val, palette)[0] if palette and val else None)
            if new_hex and new_hex.upper() != val:
                clr.set("val", new_hex.upper())
                changed += 1
        if scope == "text" and target_hex and not only_hex:
            for rPr in root.iter(q("a:rPr")):  # inherited/scheme colours become explicit so the fix is visible
                changed += _force_run_color(rPr, target_hex.upper())
    return changed


def _move_inside(el: etree._Element, safe: Bbox, slide_w: int, slide_h: int, into_safe: bool) -> Optional[tuple[int, int, int, int]]:
    """New (x, y, w, h) that puts el inside the slide (or the safe area); oversized elements are shrunk to fit. None = no change."""
    if is_nested(el):
        return None  # group children use group coordinates
    box = element_bbox(el)
    if not box:
        return None
    x, y, w, h = box
    bx, by, bw, bh = (safe.x, safe.y, safe.w, safe.h) if into_safe else (0, 0, slide_w, slide_h)
    nw, nh, nx, ny = w, h, x, y
    if w > bw:
        nw = bw if into_safe else max(slide_w - 2 * safe.x, 1)
        nx = bx if into_safe else safe.x
    else:
        nx = min(max(x, bx), bx + bw - w)
    if h > bh:
        nh = bh if into_safe else max(slide_h - 2 * safe.y, 1)
        ny = by if into_safe else safe.y
    else:
        ny = min(max(y, by), by + bh - h)
    if (nx, ny, nw, nh) == (x, y, w, h):
        return None
    return nx, ny, nw, nh


# name families of composed blocks (compose.py names its shapes «Card 12», «Legend 45», «Step 7»…): a marker, its
# label and its card move together or not at all
_BLOCK_FAMILIES = (("Card", "Rule", "Column title", "Column"), ("Swatch", "Legend", "Share", "Legend title"), ("Step", "Index", "Badge"))
_NEW_OVERLAP = 0.02  # a move may not create an overlap larger than this share of the smaller box
_NAME_NUM_RE = __import__("re").compile(r"\s*\d+$")


def _name_of(el: etree._Element) -> str:
    for nv in el.iter(q("p:cNvPr")):
        return _NAME_NUM_RE.sub("", nv.get("name") or "")
    return ""


def _family_of(el: etree._Element) -> Optional[int]:
    name = _name_of(el)
    return next((i for i, fam in enumerate(_BLOCK_FAMILIES) if name in fam), None)


def _box(el: etree._Element) -> Optional[Bbox]:
    b = element_bbox(el)
    return Bbox(x=b[0], y=b[1], w=b[2], h=b[3]) if b else None


def _paints_or_reads(el: etree._Element) -> bool:
    from verstka.rendering.textfill import has_visible_style, shape_text

    tag = etree.QName(el).localname
    return tag in ("pic", "graphicFrame", "grpSp") or has_visible_style(el) or bool(shape_text(el).strip())


def _block_of(el: etree._Element, top: list[etree._Element], slide_w: int, slide_h: int) -> list[etree._Element]:
    """What moves with `el`: the smallest filled or outlined card holding ≥ 80 % of it (not a slide-sized panel) with
    everything standing inside that card, and the members of its name family (card / rule / column; swatch / legend
    / share; step / index / badge) in the same row or column."""
    from verstka.rendering.textfill import has_visible_style

    box = _box(el)
    if box is None:
        return [el]
    members = [el]
    boxes = {id(o): _box(o) for o in top}
    cards = [o for o in top if o is not el and boxes[id(o)] is not None and has_visible_style(o) and boxes[id(o)].area > box.area
             and boxes[id(o)].intersection(box) >= 0.8 * box.area and boxes[id(o)].area < 0.5 * slide_w * slide_h]
    if cards:
        card = min(cards, key=lambda o: boxes[id(o)].area)
        cb = boxes[id(card)]
        members.append(card)
        members += [o for o in top if o not in members and boxes[id(o)] is not None and boxes[id(o)].intersection(cb) >= 0.8 * max(boxes[id(o)].area, 1)]
    fam = _family_of(el)
    if fam is not None:
        cx, cy = box.x + box.w / 2, box.y + box.h / 2
        for o in top:
            b = boxes[id(o)]
            if o in members or b is None or _family_of(o) != fam:
                continue
            ox, oy = b.x + b.w / 2, b.y + b.h / 2
            if abs(oy - cy) <= max(box.h, b.h) / 2 or abs(ox - cx) <= max(box.w, b.w) / 2:
                members.append(o)
    return members


def _move_block(el: etree._Element, slide, safe: Bbox, slide_w: int, slide_h: int, into_safe: bool) -> Optional[list[tuple[etree._Element, tuple[int, int, int, int]]]]:
    """New boxes for `el` and its block (`_block_of`) inside the slide (or the safe area): one shift for the whole
    block; a lone element that is larger than the bounds is shrunk to them as before. None when nothing moves, when
    the block does not fit, or when the move would lay it over another element (> 2 % of the smaller box) — the issue
    then stays for a re-render rather than an XML move that breaks the layout."""
    if is_nested(el):
        return None
    tree = slide._element.cSld.find(q("p:spTree"))
    top = [o for o in tree if isinstance(o.tag, str) and etree.QName(o).localname in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp")]
    members = _block_of(el, top, slide_w, slide_h) if el in top else [el]
    if len(members) == 1:
        moved = _move_inside(el, safe, slide_w, slide_h, into_safe)
        plan = [(el, moved)] if moved else []
    else:
        boxes = [_box(m) for m in members]
        if any(b is None for b in boxes):
            return None
        u = boxes[0]
        for b in boxes[1:]:
            u = u.union(b)
        bx, by, bw, bh = (safe.x, safe.y, safe.w, safe.h) if into_safe else (0, 0, slide_w, slide_h)
        if u.w > bw or u.h > bh:
            return None  # a block is never squeezed: it is re-rendered instead
        dx = min(max(u.x, bx), bx + bw - u.w) - u.x
        dy = min(max(u.y, by), by + bh - u.h) - u.y
        if dx == 0 and dy == 0:
            return None
        plan = [(m, (b.x + dx, b.y + dy, b.w, b.h)) for m, b in zip(members, boxes)]
    if not plan:
        return None
    others = [o for o in top if o not in members and _paints_or_reads(o)]
    slide_area = slide_w * slide_h
    for m, (nx, ny, nw, nh) in plan:
        old = _box(m)
        new = Bbox(x=nx, y=ny, w=nw, h=nh)
        for o in others:
            ob = _box(o)
            if ob is None or ob.area >= 0.85 * slide_area:
                continue  # a background picture or panel is under everything
            limit = _NEW_OVERLAP * max(min(new.area, ob.area), 1)
            if new.intersection(ob) <= limit:
                continue
            if old is not None and old.intersection(ob) > limit:
                continue  # they overlapped before the move: not a new overlap
            return None
    return plan


_SCOPE_RU = {"text": "текста", "fill": "заливки", "line": "обводки", "all": "элемента"}


def _xml_fixes(pptx: Path, actions: list[tuple[int, FixAction]], manifest: TemplateManifest, slide_w: int, slide_h: int) -> list[dict]:
    """In-place edits that need no re-render: recolor, refont, drop_element, move_inside, shrink_text.

    Returns what changed, one record per edit: {"slide", "kind", "element_id", "result"} — the result in plain Russian
    for the history of fixes shown to people."""
    if not actions:
        return []
    prs = Presentation(str(pptx))
    applied: list[dict] = []

    def done(slide_index: int, kind: str, eid: object, text: str) -> None:
        applied.append({"slide": slide_index, "kind": kind, "element_id": str(eid), "result": text})

    palette = manifest.tokens.palette()
    primary_font = manifest.tokens.typography.primary_family
    text_primary = manifest.tokens.color_for("text.primary")
    typo = manifest.tokens.typography
    scale = sorted({float(s) for s in (typo.sizes_used or [st.size_pt for st in typo.scale])})
    safe = manifest.tokens.spacing.safe_area.to_emu(slide_w, slide_h)
    for slide_index, act in actions:
        if not (1 <= slide_index <= len(prs.slides)):
            continue
        by_id = slide_shape_elements(prs.slides[slide_index - 1])
        ids = act.params.get("element_ids") or ([act.params["element_id"]] if act.params.get("element_id") else [])
        for eid in ids:
            el = by_id.get(str(eid))
            if el is None:
                continue
            if act.action == "drop_element":
                el.getparent().remove(el)
                done(slide_index, "drop_element", eid, "Убран лишний элемент")
            elif act.action == "refont" and primary_font:
                n = 0
                for latin in el.iter(q("a:latin")):
                    if latin.get("typeface") != primary_font:
                        latin.set("typeface", primary_font)
                        n += 1
                for rPr in el.iter(q("a:rPr")):
                    if rPr.find(q("a:latin")) is None:
                        etree.SubElement(rPr, q("a:latin")).set("typeface", primary_font)
                        n += 1
                if n:
                    done(slide_index, "refont", eid, f"Шрифт заменён на {primary_font} — основной шрифт шаблона")
            elif act.action == "recolor":
                target_hex = (act.params.get("to") or "").upper() or None
                if not target_hex and act.params.get("target") == "text.primary" and text_primary:
                    target_hex = text_primary.upper()
                only_hex = (act.params.get("hex") or "").upper() or None
                scope = act.params.get("scope") or "text"
                n = _recolor(el, scope, only_hex, target_hex, palette)
                if n:
                    what = _SCOPE_RU.get(scope, "элемента")
                    done(slide_index, "recolor", eid, f"Цвет {what} заменён на #{target_hex}" if target_hex else f"Цвет {what} приведён к палитре шаблона")
            elif act.action == "move_inside":
                plan = _move_block(el, prs.slides[slide_index - 1], safe, slide_w, slide_h, bool(act.params.get("safe")))
                if plan:
                    resized = False
                    for m, (nx, ny, nw, nh) in plan:
                        x, y, w, h = element_bbox(m)  # type: ignore[misc]
                        set_element_pos(m, x=nx, y=ny, w=nw if nw != w else None, h=nh if nh != h else None)
                        resized = resized or (nw, nh) != (w, h)
                    where = "в поля шаблона" if act.params.get("safe") else "внутрь слайда"
                    what = "Элемент" if len(plan) == 1 else f"Блок из {len(plan)} элементов"
                    done(slide_index, "move_inside", eid, f"{what} возвращён {where}{' и уменьшен по размеру' if resized else ''}")
            elif act.action == "shrink_text":
                ratio = float(act.params.get("ratio", 1.2))
                factor = max(0.6, min(0.95, 1 / ratio))
                sizes: list[float] = []
                for rPr in el.iter(q("a:rPr")):
                    sz = rPr.get("sz")
                    if not sz:
                        continue
                    cur = int(sz) / 100.0
                    new = _snap_size(cur, factor, scale)
                    if new < cur:
                        rPr.set("sz", str(int(round(new * 100))))
                        sizes.append(new)
                if sizes:
                    done(slide_index, "shrink_text", eid, f"Шрифт уменьшен до {', '.join(f'{s:g}' for s in sorted(set(sizes), reverse=True))} пт, чтобы текст поместился")
    if applied:
        prs.save(str(pptx))
    return applied


# ----------------------------------------------------------------------------- loop


class _Snapshot:
    """deck.pptx + plan + outline + render result before an iteration, to undo it when the audit gets worse."""

    def __init__(self, pptx: Path, plan: LayoutPlan, outline: DeckOutline, render_result: RenderResult | None, iteration: int):
        self.bak = pptx.with_name(f"{pptx.stem}.iter{iteration}.bak")
        self.had_file = pptx.exists()
        if self.had_file:
            shutil.copy2(pptx, self.bak)
        self.plan = copy.deepcopy(plan)
        self.outline = copy.deepcopy(outline)
        self.render_result = render_result

    def restore(self, pptx: Path) -> tuple[LayoutPlan, DeckOutline, RenderResult | None]:
        if self.had_file and self.bak.exists():
            shutil.copy2(self.bak, pptx)
        return self.plan, self.outline, self.render_result

    def discard(self) -> None:
        try:
            self.bak.unlink(missing_ok=True)
        except OSError:
            pass


def _better(new: AuditReport, old: AuditReport) -> bool:
    """Fewer errors, then fewer warnings, then a higher score."""
    return (new.summary.errors, new.summary.warnings, -new.summary.score) < (old.summary.errors, old.summary.warnings, -old.summary.score)


def _issue_key(i: Issue) -> tuple:
    return (i.check_id, i.outline_id or i.slide)


def autofix_loop(
    pptx: Path,
    report: AuditReport,
    outline: DeckOutline,
    plan: LayoutPlan,
    manifest: TemplateManifest,
    ws: TemplateWorkspace,
    *,
    providers: Optional[ProviderRegistry] = None,
    skills: Optional[SkillsRegistry] = None,
    max_iterations: int = 2,
    use_models: bool = False,
    only_ids: Optional[set[str]] = None,
    images_dir: Optional[Path] = None,
    render: Optional[bool] = None,
    brief_text: Optional[str] = None,
) -> tuple[AuditReport, LayoutPlan, DeckOutline, RenderResult | None]:
    """Apply fixes while the audit improves; an iteration that makes it worse is rolled back (deck, plan, outline).

    Slide images are rendered for every re-audit only when the model checks need them (`render`, default: with
    models) — LibreOffice is the slowest step of the loop and the deterministic checks work on the XML.
    Returns the final report/plan/outline (and last render)."""
    pptx = Path(pptx)
    current = report
    render_result: RenderResult | None = None
    slide_w, slide_h = manifest.slide_size.w, manifest.slide_size.h
    render = use_models if render is None else render
    audit_kw = dict(providers=providers if use_models else None, skills=skills if use_models else None, use_vlm=use_models, use_llm=use_models, images_dir=images_dir, strategy=report.strategy, render=render, brief_text=brief_text)
    selected_keys = {_issue_key(i) for i in current.issues if i.id in only_ids} if only_ids is not None else None
    for it in range(1, max_iterations + 1):
        fixes = plan_fixes(current, plan, only_ids)
        if not fixes:
            break
        snapshot = _Snapshot(pptx, plan, outline, render_result, it)
        try:
            applied: list[dict] = []
            need_render = False
            xml_actions: list[tuple[int, FixAction]] = []
            slide_index_of = {s.outline_id: s.index for s in _ir_slides(current)}
            for oid, actions in fixes.items():
                for act in actions:
                    if act.action in STRUCTURAL_ACTIONS:
                        res = _rematch(plan, outline, oid, manifest)
                        applied.append({"iteration": it, "outline_id": oid, "action": act.action, "result": res})
                        need_render = need_render or res != "skip"
                    elif act.action == "condense_text":
                        res = _condense(outline, oid, skills if use_models else None, providers if use_models else None)
                        applied.append({"iteration": it, "outline_id": oid, "action": act.action, "result": res})
                        need_render = need_render or res != "skip"
                    else:
                        idx = slide_index_of.get(oid)
                        if idx:
                            xml_actions.append((idx, act))
            if need_render:
                _unsay(outline)
                render_result = render_deck(outline, plan, manifest, ws, pptx)
                # element ids change after a re-render: re-derive the XML fixes from a render-free audit of the new deck
                fresh = run_audit(pptx, manifest, outline, ws, render=False, images_dir=images_dir, strategy=report.strategy, brief_text=brief_text)
                fresh_ids = {i.id for i in fresh.issues if _issue_key(i) in selected_keys} if selected_keys is not None else None
                fresh_index_of = {s.outline_id: s.index for s in _ir_slides(fresh)}
                xml_actions = []
                for oid, actions in plan_fixes(fresh, plan, fresh_ids, xml_only=True).items():
                    idx = fresh_index_of.get(oid)
                    if idx:
                        xml_actions.extend((idx, a) for a in actions)
            if xml_actions:
                for fix in _xml_fixes(pptx, xml_actions, manifest, slide_w, slide_h):
                    applied.append({"iteration": it, "action": "xml", **fix})
            if not applied:
                break
            new_report = run_audit(pptx, manifest, outline, ws, **audit_kw)
            new_report.applied_fixes = current.applied_fixes + applied
            new_report.iterations = it
            if _better(new_report, current):
                current = new_report
                continue
            # the iteration made things worse: ship the previous deck, re-audit it so slide images match
            plan, outline, render_result = snapshot.restore(pptx)
            restored = run_audit(pptx, manifest, outline, ws, **audit_kw)
            old_s, new_s = current.summary, new_report.summary
            why = (
                f"ошибок стало {new_s.errors} вместо {old_s.errors}" if new_s.errors > old_s.errors
                else f"предупреждений стало {new_s.warnings} вместо {old_s.warnings}" if new_s.warnings > old_s.warnings
                else f"оценка стала {new_s.score:g} вместо {old_s.score:g}" if new_s.score < old_s.score
                else "правки ничего не улучшили"
            )
            rollback = {"iteration": it, "action": "rollback", "result": f"{why} — возвращена предыдущая версия"}
            log.info("autofix iteration %d rolled back: %s", it, rollback["result"])
            restored.applied_fixes = current.applied_fixes + applied + [rollback]
            restored.iterations = it
            current = restored
            break
        finally:
            snapshot.discard()
    return current, plan, outline, render_result


def _ir_slides(report: AuditReport):
    from verstka.audit.ir import build_deck_ir

    try:
        return build_deck_ir(report.deck, with_images=False).slides
    except Exception:  # noqa: BLE001
        return []
