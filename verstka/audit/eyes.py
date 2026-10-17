"""«Взгляд дизайнера» (Verstka 3.0): the vision model looks at every rendered slide of a variant like an art director
checking a proof before it goes to the client.

The deterministic audit reads the deck's XML: it measures overflow, fonts, colours and contrast, but it never sees
the slide. The model does (skill `slide_design_review`): a word broken mid-word in a narrow card, an empty row of a
ruled grid under a callout, text over a rule, a half-empty slide. A slide with a major problem is laid out again —
the autofix ladder's next rung: the slide's next template sample, then the composer (autofix._rematch) — and the
change has to pass three gates: the deterministic audit of the slide must not get worse, the new layout must really
differ from the old one (not a hairline moved), and the model, shown the two layouts side by side in BOTH orders (skill
`slide_design_compare`), must prefer the new one both times — a model that answers by the position of the image, or
mistakes one image for the other, never passes. Otherwise the slide goes back as it was. Problems no new layout cured stay in the audit report as the model's remarks (check `design_review`),
with the model's own words, so the person sees them and may ask the helper to fix them.

Bounded by the generation's time: it starts only with `min_seconds_left` of the model budget left, stops asking when
the budget runs out, lays out at most `max_fixes` slides per variant (configs/eyes.yaml). Without a vision model, a
skills registry or LibreOffice it does nothing.

Checked live on Qwen3.8-27B (2026-10-07): it found the four words broken mid-word on Marketing Report slide 8 and the
empty grid row of Классика slide 3, passed a clean VK Tech slide, and chose the layout without the broken words in
both orders of the pair (≈ 2–7 s and 0.03–0.11 ₽ a look).
"""

from __future__ import annotations

import logging
import re
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional

import yaml
from pydantic import BaseModel, Field, field_validator

from verstka.schemas.audit import AuditReport, CheckSpec, Issue

log = logging.getLogger(__name__)

_CONFIG = Path(__file__).resolve().parents[2] / "configs" / "eyes.yaml"
DEFAULTS: dict[str, Any] = {
    "enabled": True,
    "max_fixes": 4,  # slides laid out again per variant
    "min_seconds_left": 40.0,  # the model budget left to start the review
    "fix_seconds_left": 20.0,  # ... to lay out again and compare
    "compare": True,  # a new layout stays only when the model prefers it side by side, in both orders
    "min_change": 0.02,  # the share of the slide a new layout has to change (grey, 160 px wide) to be compared at all
    "dpi": 100,  # the slide images the model sees
    "skip_kinds": ["thanks", "agenda"],
    "workers": 4,
}

# the problems the review may name — and their words for the person (the timeline, the remarks)
CODES: dict[str, str] = {
    "text_cut": "текст обрезан или выходит за блок",
    "word_broken": "слова разорваны посередине",
    "overlap": "элементы наезжают друг на друга",
    "cramped": "текст набит слишком плотно",
    "too_small": "текст слишком мелкий",
    "low_contrast": "текст плохо читается на фоне",
    "unbalanced": "пустоты и перекос композиции",
    "leftover": "пустая рамка или лишний элемент",
    "chart_unreadable": "диаграмма или таблица не читается",
}
# what helps against each problem — the remark's «Как исправить» (the person may ask the helper to do it)
FIX_HINTS: dict[str, str] = {
    "text_cut": "сократите текст или дайте ему больше места: другая форма слайда, шире блок",
    "word_broken": "дайте словам место: меньше карточек в ряд или шире колонки",
    "overlap": "разведите элементы: другая форма слайда или меньше элементов",
    "cramped": "разгрузите слайд: меньше текста или часть — на следующий слайд",
    "too_small": "крупнее текст: меньше элементов на слайде",
    "low_contrast": "текст на спокойном фоне или в цвете шаблона, который читается",
    "unbalanced": "распределите блоки по слайду: другая форма или крупнее главное",
    "leftover": "уберите пустую рамку или лишний элемент",
    "chart_unreadable": "меньше рядов или категорий, крупнее подписи, или таблица вместо диаграммы",
}
_SYNONYMS = {"word_break": "word_broken", "broken_word": "word_broken", "cut": "text_cut", "overflow": "text_cut", "contrast": "low_contrast",
             "empty": "unbalanced", "imbalance": "unbalanced", "placeholder": "leftover", "small_text": "too_small", "crowded": "cramped"}

# why the review did not look at the variant, by its code (`EyesOutcome.skipped_code`) — the person's words: the build's
# timeline says «Взгляд дизайнера пропущен: …», «Подробнее» → «Агент» shows the step skipped with the reason
# (web/src/lib/designEye.ts `EYES_SKIP` has the same words for runs recorded before the timeline said it). `skipped`
# keeps the reason in English for the log.
SKIP_RU: dict[str, str] = {
    "off": "выключен в настройках",
    "no_model": "модель недоступна",
    "no_time": "не хватило времени на сборку",
    "no_images": "не получилось снять картинки слайдов",
    "no_answer": "модель не ответила",
    "failed": "не получилось посмотреть на слайды",
}
_SKIP_OF = (
    ("off in configs", "off"), ("no vision model", "no_model"), ("no time left", "no_time"), ("no slide images", "no_images"),
    ("the model did not answer", "no_answer"), ("failed", "failed"),
)


def skip_code(why: Optional[str]) -> Optional[str]:
    """The code of a skip reason (`EyesOutcome.skipped`, English): runs recorded before the code was kept are read by
    their words; None — the review ran, or a reason this module does not know."""
    if not why:
        return None
    low = str(why).lower()
    return next((code for lead, code in _SKIP_OF if low.startswith(lead)), None)


def skip_message(code: Optional[str]) -> str:
    """The timeline's line of a skipped review: «Взгляд дизайнера пропущен: не хватило времени на сборку.»"""
    said = SKIP_RU.get(code or "")
    return f"Взгляд дизайнера пропущен: {said}." if said else "Взгляд дизайнера пропущен."

DESIGN_REVIEW = CheckSpec(
    id="design_review", title="Взгляд дизайнера (VLM)", severity="warn", kind="model", category="layout",
    description=(
        "«Взгляд дизайнера» (Verstka 3.0, `audit/eyes.py`): модель смотрит на картинку каждого готового слайда, как "
        "арт-директор на вёрстку перед отправкой клиенту, — обрезанный текст, слова, разорванные посередине, наложения, "
        "тесный или мелкий текст, плохой контраст, перекос и пустоты, пустые рамки, нечитаемые диаграммы (скилл "
        "`slide_design_review`). Слайд с серьёзной проблемой собирается заново (следующий образец шаблона, композиция, "
        "стиль другого варианта); новая вёрстка остаётся, только если детерминированный аудит слайда не стал хуже, слайд "
        "заметно изменился и модель, сравнив обе вёрстки рядом, выбрала новую в обоих порядках пары (скилл "
        "`slide_design_compare`). Замечание в отчёте — то, что новая вёрстка не исправила, со словами модели. Идёт по "
        "умолчанию вместе с моделью текста."
    ),
)


class ReviewProblem(BaseModel):
    code: str = ""
    severity: str = "minor"
    where: str = ""
    evidence: str = ""

    @field_validator("code", mode="before")
    @classmethod
    def _code(cls, v: Any) -> str:
        c = str(v or "").strip().lower().replace("-", "_").replace(" ", "_")
        return _SYNONYMS.get(c, c)

    @field_validator("severity", mode="before")
    @classmethod
    def _severity(cls, v: Any) -> str:
        s = str(v or "").strip().lower()
        return "major" if s.startswith(("maj", "crit", "high", "серьёз", "серьез", "крит", "сильн")) else "minor"

    @field_validator("where", "evidence", mode="before")
    @classmethod
    def _text(cls, v: Any) -> str:
        return " ".join(str(v or "").split())[:300]


class DesignReview(BaseModel):
    """Output of the `slide_design_review` skill (read leniently: a null or a word where a flag or a text is due costs
    no second call)."""

    ok: bool = True
    problems: list[ReviewProblem] = Field(default_factory=list)
    verdict: str = ""

    @field_validator("ok", mode="before")
    @classmethod
    def _ok(cls, v: Any) -> bool:
        if isinstance(v, bool):
            return v
        return str(v or "").strip().lower() not in ("false", "0", "no", "нет", "не", "плохо")

    @field_validator("verdict", mode="before")
    @classmethod
    def _verdict(cls, v: Any) -> str:
        return " ".join(str(v or "").split())[:300]

    @field_validator("problems", mode="before")
    @classmethod
    def _problems(cls, v: Any) -> Any:
        return [p for p in v if isinstance(p, dict)] if isinstance(v, list) else []

    def known(self) -> list[ReviewProblem]:
        return [p for p in self.problems if p.code in CODES]

    def major(self) -> list[ReviewProblem]:
        return [p for p in self.known() if p.severity == "major"]


_CHOICE_RE = re.compile(r"^\W*(?:(?:layout|image|вариант|макет|вёрстка|верстка|изображение)\s+)?([ABАВБ])(?![A-Za-zА-Яа-яЁё])", re.I)


class DesignCompare(BaseModel):
    """Output of the `slide_design_compare` skill: what stands where on each layout (it makes the model look at both
    before it chooses), the choice and why."""

    a: str = ""
    b: str = ""
    better: str = "same"
    reason: str = ""

    @field_validator("a", "b", "reason", mode="before")
    @classmethod
    def _text(cls, v: Any) -> str:
        return " ".join(str(v or "").split())[:300]

    @field_validator("better", mode="before")
    @classmethod
    def _better(cls, v: Any) -> str:
        # «B», «Layout B», «Вариант Б», «B (новая)»; Cyrillic letters the model may answer with; «both» is no «B»
        m = _CHOICE_RE.match(str(v or ""))
        if not m:
            return "same"
        c = m.group(1).upper()
        return {"А": "A", "В": "B", "Б": "B"}.get(c, c)


def load_config(path: Optional[Path] = None) -> dict[str, Any]:
    p = Path(path) if path else _CONFIG
    try:
        data = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError):
        data = {}
    return {**DEFAULTS, **{k: v for k, v in data.items() if k in DEFAULTS}}


@dataclass
class EyesOutcome:
    """What the review did on one variant (run_manifest.json `eyes`, the timeline)."""

    looked: int = 0
    calls: int = 0
    flagged: list[dict] = field(default_factory=list)  # {"slide", "outline_id", "problems": [{code, severity, where, evidence}]}
    fixed: list[dict] = field(default_factory=list)  # {"slide", "outline_id", "result", "reason"}: the new layout stayed
    kept: list[dict] = field(default_factory=list)  # tried again, the old layout stayed: {"slide", "outline_id", "why"}
    skipped: Optional[str] = None
    skipped_code: Optional[str] = None  # SKIP_RU: off · no_model · no_time · no_images · no_answer · failed
    seconds: float = 0.0

    def summary(self) -> dict:
        return {
            "looked": self.looked, "calls": self.calls, "flagged": self.flagged, "fixed": self.fixed, "kept": self.kept,
            "skipped": self.skipped, "skipped_code": self.skipped_code, "seconds": round(self.seconds, 2),
        }


def slide_images(pptx: Path, out_dir: Path, dpi: int = 100, timeout_s: float = 240.0) -> dict[int, Path]:
    """Slide number → image of the deck as it is now (LibreOffice → PDF → JPEG)."""
    import shutil

    from verstka.ingest.render import pdf_to_images, pptx_to_pdf

    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.jpg"):
        old.unlink(missing_ok=True)
    made = pptx_to_pdf(Path(pptx), out_dir, timeout_s=timeout_s)
    pdf = out_dir / "deck.pdf"
    if made != pdf:
        shutil.move(str(made), str(pdf))
    try:
        imgs = pdf_to_images(pdf, out_dir, dpi=dpi)
    finally:
        pdf.unlink(missing_ok=True)
    return {i: p for i, p in enumerate(imgs, 1)}


def changed_share(a: Path, b: Path, width: int = 160, level: int = 24) -> float:
    """The share of the slide that looks different on two images of it (grey, `width` px wide, a pixel counts when it
    differs by more than `level` of 255): a new layout that moved nothing but a hairline is no new layout."""
    from PIL import Image, ImageChops

    with Image.open(a) as ia, Image.open(b) as ib:
        size = (width, max(1, round(width * ia.height / max(ia.width, 1))))
        diff = ImageChops.difference(ia.convert("L").resize(size, Image.BILINEAR), ib.convert("L").resize(size, Image.BILINEAR))
    hist = diff.histogram()
    return sum(hist[level + 1:]) / float(size[0] * size[1])


def slide_hashes(pptx: Path) -> dict[int, str]:
    """Slide number → a fingerprint of its XML: a slide set exactly as before has the same one."""
    import hashlib

    from lxml import etree
    from pptx import Presentation

    return {i: hashlib.sha1(etree.tostring(s._element)).hexdigest() for i, s in enumerate(Presentation(str(pptx)).slides, 1)}


def shown_as(render_result: Any, oid: str) -> Optional[str]:
    """What a render put on the slide of `oid`: «rematch → p7» (the sample shown — a fallback, when the one chosen read
    badly there) or «synth → cards»; None when the render does not say."""
    for rs in getattr(render_result, "slides", None) or []:
        if getattr(rs, "outline_id", None) == oid:
            if rs.mode == "clone" and rs.pattern_id:
                return f"rematch → {rs.pattern_id}"
            return f"synth → {rs.composition or 'composition'}"
    return None


def _left(deadline: Optional[float]) -> float:
    return float("inf") if deadline is None else deadline - time.monotonic()


def _form(osl: Any) -> str:
    if osl is None:
        return "слайд"
    try:
        from verstka.planning.agent import form_ru

        return form_ru(osl)
    except Exception:  # noqa: BLE001 - the form only helps the model read the slide
        return str(getattr(osl.kind, "value", osl.kind))


def _data_note(osl: Any) -> str:
    """What the data layer already checked on a slide made from a file (its categories, its gaps): the eyes judge how it
    reads, not whether a point sits at its label (qa3 U2: «52 … далеко от своей отметки» on a correct chart)."""
    d = getattr(osl, "data", None) if osl is not None else None
    ch = getattr(getattr(osl, "content", None), "chart", None) if osl is not None else None
    if d is None or ch is None:
        return ""
    cats = [str(c) for c in (ch.categories or [])]
    gaps = [cats[k] for s in (ch.series or [])[:1] for k, v in enumerate(s.values) if v is None and k < len(cats)]
    note = f"categories in order: {', '.join(cats[:24])}." if cats else ""
    if gaps:
        note += f" No data for {', '.join(gaps)}: the gap in the line is intended."
    return note.strip()


def _said(problems: list[ReviewProblem]) -> str:
    return "; ".join(f"{CODES.get(p.code, p.code)} ({p.where})" if p.where else CODES.get(p.code, p.code) for p in problems) or "—"


def look(
    vdir: Path,
    report: AuditReport,
    outline: Any,
    plan: Any,
    manifest: Any,
    ws: Any,
    *,
    providers: Any,
    skills: Any,
    deadline: Optional[float] = None,
    others: Optional[dict] = None,
    brief_text: Optional[str] = None,
    writer: Any = None,
    emit: Optional[Callable[..., None]] = None,
    cfg: Optional[dict] = None,
    images: Optional[Callable[[Path, Path, int], dict[int, Path]]] = None,
    others_live: Optional[Callable[[], dict]] = None,
    on_pick: Optional[Callable[[Any, str], None]] = None,
) -> tuple[AuditReport, Any, Any, Any, EyesOutcome]:
    """Review the variant in `vdir` (its deck.pptx, audit report, outline and plan) and lay out again the slides the
    model sees a major problem on. Returns (audit report, plan, outline, render result or None, outcome); the deck on
    disk is the one the report describes. `others`: what the variants before this one show on each slide (the render's
    `shown`); `others_live`: the same, read when a slide takes another sample — every other variant, also those laid
    out while this one is being looked at; `on_pick`: told each sample a slide takes (outline slide, sample id)."""
    from verstka.analysis.classify import image_bytes_for_vlm
    from verstka.audit.autofix import _ir_slides, _restore_slide, _slide_weight, _Snapshot, _unsay, _xml_fixes, plan_fixes
    from verstka.audit.runner import run_audit
    from verstka.rendering.renderer import render_deck

    cfg = {**DEFAULTS, **(cfg or {})}
    # LibreOffice gets the time the generation has left, not its own four minutes
    shoot = images or (lambda p, d, dpi: slide_images(p, d, dpi, timeout_s=max(15.0, min(240.0, _left(deadline)))))
    out = EyesOutcome()
    t0 = time.monotonic()
    say = emit or (lambda message, slide=None, fix=None: None)
    pptx = Path(vdir) / "deck.pptx"
    render_result = None

    def done(why: Optional[str] = None) -> tuple[AuditReport, Any, Any, Any, EyesOutcome]:
        out.skipped = why
        out.skipped_code = skip_code(why)
        out.seconds = time.monotonic() - t0
        return report, plan, outline, render_result, out

    def skip(why: str) -> tuple[AuditReport, Any, Any, Any, EyesOutcome]:
        """Not looked at all: the timeline says so (the step ends skipped, not «now» till the build's end)."""
        got = done(why)
        say(skip_message(out.skipped_code), None)
        return got

    if not cfg.get("enabled", True):
        return skip("off in configs/eyes.yaml")
    if providers is None or skills is None or not providers.has("vlm"):
        return skip("no vision model")
    if _left(deadline) < float(cfg["min_seconds_left"]):
        return skip("no time left in the generation's model budget")
    try:
        before = shoot(pptx, Path(vdir) / "eyes" / "before", int(cfg["dpi"]))
    except Exception as e:  # noqa: BLE001 - no LibreOffice, a deck it cannot open: the deck stays as it is
        log.warning("design review: slide images failed: %s", e)
        return skip(f"no slide images ({str(e)[:120]})")

    ir = _ir_slides(report, outline)
    oid_of = {s.index: s.outline_id for s in ir}
    by_id = {s.id: s for s in outline.slides}
    skip_kinds = set(cfg.get("skip_kinds") or ())
    todo = [i for i in sorted(before) if (by_id.get(oid_of.get(i) or "") is not None) and by_id[oid_of[i]].kind.value not in skip_kinds]
    title = getattr(outline, "title", "") or ""

    def review(i: int) -> tuple[int, Optional[DesignReview]]:
        if _left(deadline) < 8.0:
            return i, None
        osl = by_id.get(oid_of.get(i) or "")
        try:
            res = skills.run(
                "slide_design_review", providers,
                {"slide_index": i, "deck_title": title, "form": _form(osl), "headline": (osl.headline if osl else "") or "", "data_note": _data_note(osl)},
                images=[image_bytes_for_vlm(before[i])],
            )
            return i, res.parsed
        except Exception as e:  # noqa: BLE001 - one slide unseen never stops the deck
            log.info("design review of slide %s failed: %s", i, str(e)[:160])
            return i, None

    workers = max(1, min(int(cfg["workers"]), getattr(getattr(providers, "limits", None), "max_concurrency", 4) or 4, len(todo) or 1))
    with ThreadPoolExecutor(max_workers=workers) as ex:
        reviews = dict(ex.map(review, todo))
    out.calls += len(todo)
    out.looked = sum(1 for r in reviews.values() if r is not None)
    if todo and not out.looked:
        say("Взгляд дизайнера: модель не ответила — слайды остались как были.", None)
        return done("the model did not answer")
    seen: dict[int, list[ReviewProblem]] = {}
    for i, r in reviews.items():
        if r is None or not r.known():
            continue
        seen[i] = r.known()
        out.flagged.append({"slide": i, "outline_id": oid_of.get(i), "problems": [p.model_dump() for p in r.known()]})
    majors = sorted(((i, [p for p in ps if p.severity == "major"]) for i, ps in seen.items()), key=lambda x: (-len(x[1]), x[0]))
    majors = [(i, ps) for i, ps in majors if ps][: int(cfg["max_fixes"])]
    if not seen:
        say(f"Взгляд дизайнера: посмотрел на {ru_slides(out.looked)} — замечаний нет.", None)
    else:
        say(f"Взгляд дизайнера: посмотрел на {ru_slides(out.looked)} — замечания на {_slides_list(sorted(seen))}.", None)

    fixed_ids: set[str] = set()
    if majors and _left(deadline) >= float(cfg["fix_seconds_left"]):
        snapshot = _Snapshot(pptx, plan, outline, None, "eyes")
        tried: dict[str, int] = {}  # outline id → slide number: every slide sent up the ladder

        def rerender() -> tuple[Any, Optional[AuditReport]]:
            """The deck laid out again from the plan, and the fixes the autofix made in its XML (a recolour, a shrunk
            text, a box moved inside) made again on it: they live only in the deck file — a render from the plan alone
            would drop them on every slide. Returns the render and the deck's audit when no fix had to be made again."""
            _unsay(outline)
            rr = render_deck(outline, plan, manifest, ws, pptx, shown=others)
            fresh = run_audit(pptx, manifest, outline, ws, render=False, strategy=report.strategy, brief_text=brief_text, writer=writer)
            at = {s.outline_id: s.index for s in _ir_slides(fresh, outline)}
            acts = [(at[o], a) for o, aa in plan_fixes(fresh, plan, xml_only=True).items() if at.get(o) for a in aa]
            if acts:
                _xml_fixes(pptx, acts, manifest, manifest.slide_size.w, manifest.slide_size.h)
            return rr, (None if acts else fresh)

        def step(oid: str) -> str:
            res = relayout(plan, outline, oid, manifest, taken, others=others_live() if others_live else others)
            if res.startswith("rematch → ") and on_pick is not None:
                on_pick(by_id.get(oid), res.split("→", 1)[1].strip())  # the next variants know the sample is shown here
            return res

        try:
            moved: dict[str, tuple[int, str]] = {}
            taken: set = set()
            for i, _ in majors:
                oid = oid_of.get(i)
                if not oid:
                    continue
                res = step(oid)
                if res != "skip":
                    moved[oid] = (i, res)
            tried = {oid: i for oid, (i, _r) in moved.items()}
            audited: Optional[AuditReport] = None
            if moved:
                was = slide_hashes(pptx)
                render_result, audited = rerender()
                # a rung that sets the slide exactly as it was (a style the composer draws alike) is no new layout: the
                # next rung of the ladder, twice at most, before anything is judged
                for _ in range(2):
                    now = slide_hashes(pptx)
                    stepped = False
                    for oid, (i, _r) in list(moved.items()):
                        if now.get(i) == was.get(i):
                            res = step(oid)
                            if res != "skip":
                                moved[oid] = (i, res)
                                stepped = True
                    if not stepped:
                        break
                    render_result, audited = rerender()
                now = slide_hashes(pptx)
                for oid in [oid for oid, (i, _r) in moved.items() if now.get(i) == was.get(i)]:
                    i, _r = moved.pop(oid)
                    _restore_slide(plan, outline, snapshot, oid)  # the slide on disk is the one it had: no render needed
                    out.kept.append({"slide": i, "outline_id": oid, "why": "другой вёрстки для этого слайда не нашлось"})
            if moved:
                trial = audited or run_audit(pptx, manifest, outline, ws, render=False, strategy=report.strategy, brief_text=brief_text, writer=writer)
                worse = {oid for oid in moved if _slide_weight(trial, oid) > _slide_weight(report, oid)}
                for oid in worse:
                    i, res = moved.pop(oid)
                    _restore_slide(plan, outline, snapshot, oid)
                    out.kept.append({"slide": i, "outline_id": oid, "why": "автоматическая проверка нашла на новой вёрстке больше замечаний"})
                if worse and moved:
                    render_result, _a = rerender()
                reasons: dict[str, str] = {}
                if moved and cfg.get("compare", True):
                    # the model has to see the new layout: without the time to look, the old one stays
                    after = shoot(pptx, Path(vdir) / "eyes" / "after", int(cfg["dpi"])) if _left(deadline) >= 8.0 else {}

                    asks: list[tuple[str, int, bool]] = []
                    verdicts: list[tuple[str, str, str]] = []
                    for oid, (i, _) in moved.items():
                        if i not in before or i not in after or _left(deadline) < 5.0:
                            verdicts.append((oid, "same", "не хватило времени сравнить вёрстки"))
                            continue
                        try:
                            share = changed_share(before[i], after[i])
                        except Exception:  # noqa: BLE001 - an unreadable image: the model compares as before
                            share = 1.0
                        if share < float(cfg["min_change"]):
                            verdicts.append((oid, "same", "новая вёрстка почти не отличается от прежней"))
                            continue
                        asks.extend([(oid, i, True), (oid, i, False)])  # both orders: neither a side nor a guess decides

                    def compare(ask: tuple[str, int, bool]) -> tuple[str, bool, Optional[str], str]:
                        oid, i, new_first = ask
                        if _left(deadline) < 5.0:
                            return oid, new_first, None, "не хватило времени сравнить вёрстки"
                        old_img, new_img = image_bytes_for_vlm(before[i]), image_bytes_for_vlm(after[i])
                        osl = by_id.get(oid)
                        try:
                            res = skills.run(
                                "slide_design_compare", providers,
                                {"slide_index": i, "form": _form(osl), "headline": (osl.headline if osl else "") or "", "original": "B" if new_first else "A", "problems": _said([p for p in seen.get(i, []) if p.severity == "major"])},
                                images=[new_img, old_img] if new_first else [old_img, new_img],
                            )
                            better, reason = res.parsed.better, res.parsed.reason
                        except Exception as e:  # noqa: BLE001 - not compared: the old layout stays
                            log.info("design compare of slide %s failed: %s", i, str(e)[:160])
                            return oid, new_first, None, "сравнение не удалось"
                        if better == "same":
                            return oid, new_first, "same", reason
                        return oid, new_first, ("new" if better == ("A" if new_first else "B") else "old"), reason

                    if asks:
                        with ThreadPoolExecutor(max_workers=workers) as ex:
                            answers = list(ex.map(compare, asks))
                        out.calls += len(asks)
                        for oid in dict.fromkeys(a[0] for a in asks):
                            votes = [(v, why) for o, _, v, why in answers if o == oid]
                            said = [v for v, _ in votes]
                            if said == ["new", "new"]:
                                verdicts.append((oid, "new", next((why for v, why in votes if why), "")))
                            elif "new" in said:
                                # chosen when shown in one order only: the position of the picture chose, not the model
                                verdicts.append((oid, "same", "модель выбрала новую вёрстку только в одном порядке сравнения"))
                            elif said == ["old", "old"]:
                                why = next((why for v, why in votes if why), "")
                                verdicts.append((oid, "same", f"модель сочла её лучше — {_lower(why)}" if why else "модель сочла её лучше"))
                            elif None in said:
                                verdicts.append((oid, "same", next((why for v, why in votes if v is None and why), "сравнение не удалось")))
                            else:
                                # «same» in at least one order (its words may even praise the new layout): no clear gain
                                verdicts.append((oid, "same", "модель не увидела, что новая вёрстка лучше"))
                    back = [(oid, why) for oid, v, why in verdicts if v != "new"]
                    for oid, why in back:
                        i, res = moved.pop(oid)
                        _restore_slide(plan, outline, snapshot, oid)
                        out.kept.append({"slide": i, "outline_id": oid, "why": why or "модель не выбрала новую вёрстку"})
                    if back and moved:
                        render_result, _a = rerender()
                    reasons = {oid: why for oid, v, why in verdicts if v == "new"}
                for oid, (i, res) in moved.items():
                    fixed_ids.add(oid)
                    # what the render put on the slide: a sample that read badly there gives way to a fallback
                    out.fixed.append({"slide": i, "outline_id": oid, "result": shown_as(render_result, oid) or res, "reason": reasons.get(oid, "")})
            if tried and not out.fixed:
                # nothing changed in the end: the deck exactly as it was, byte for byte (the autofix's own fixes in it)
                plan, outline, _r = snapshot.restore(pptx)
                render_result = None
        except Exception as e:  # noqa: BLE001 - laying out again failed (LibreOffice, a render): the deck as it was
            log.warning("design review: laying the slides out again failed: %s", e, exc_info=True)
            plan, outline, _r = snapshot.restore(pptx)
            render_result = None
            fixed_ids = set()
            out.fixed = []
            kept = {k["outline_id"] for k in out.kept}
            out.kept += [{"slide": i, "outline_id": oid, "why": "пересобрать слайд не удалось"} for oid, i in tried.items() if oid not in kept]
        finally:
            snapshot.discard()
        for item in out.fixed:
            problems = _said([p for p in seen.get(item["slide"], []) if p.severity == "major"])
            fix = f"Собрал заново{_how(item['result'])}" + (f": {_lower(item['reason'])}." if item["reason"] else ".")
            say(f"Взгляд дизайнера: слайд {item['slide']} — {problems}.", item["slide"], fix)
        for item in out.kept:
            problems = _said([p for p in seen.get(item["slide"], []) if p.severity == "major"])
            say(f"Взгляд дизайнера: слайд {item['slide']} — {problems}.", item["slide"], f"Оставил прежнюю вёрстку: {_lower(item['why'] or 'модель не выбрала новую')}.")

    # the report of the deck as it is now: the deterministic checks again when a slide changed, the model's remarks of
    # the slides that kept their problems
    final = report
    if fixed_ids:
        final = run_audit(pptx, manifest, outline, ws, render=False, strategy=report.strategy, brief_text=brief_text, writer=writer)
        final.applied_fixes = list(report.applied_fixes)
        final.iterations = report.iterations
        final.issues.extend(i for i in report.issues if i.kind == "model" and (i.outline_id or "") not in fixed_ids)
    final.applied_fixes = list(final.applied_fixes) + [
        {"iteration": "eyes", "outline_id": f["outline_id"], "action": "rematch", "result": f["result"], "by": "design_review", "reason": f["reason"]} for f in out.fixed
    ] + [{"iteration": "eyes", "outline_id": k["outline_id"], "action": "rematch", "result": "kept", "rolled_back": True, "by": "design_review", "reason": k["why"]} for k in out.kept]
    # the serious problems no new layout cured; a designer's polish («minor») stays in run_manifest `eyes` only
    final.issues.extend(problem_issues({i: ps for i, ps in seen.items() if oid_of.get(i) not in fixed_ids}, oid_of))
    final.recompute()
    report = final
    return done(None)


# a closing or contents slide of a deck the review does not know the plan of, by its title: the review of a new deck
# skips them too (`skip_kinds`), so «было» and «стало» are judged on the same kind of slides
_SKIP_TITLE_RE = re.compile(r"(?i)^\W*(спасибо|благодар|вопросы|ваши вопросы|thank|q\s*&\s*a|questions|содержание|оглавление|agenda|план презентации)")


def look_only(pptx: Path, out_dir: Path, *, providers: Any, skills: Any, cfg: Optional[dict] = None, titles: Optional[dict[int, str]] = None,
              deadline: Optional[float] = None, images: Optional[Callable[[Path, Path, int], dict[int, Path]]] = None) -> dict[int, list[ReviewProblem]]:
    """The review alone, on any deck (an old deck uploaded to be improved): slide number → the problems the model sees
    on it; nothing is laid out again. {} when the review is off (configs/eyes.yaml), without a vision model, slide
    images or time. The images go to a folder of their own under `out_dir`, removed afterwards: two uploads of one
    file never share them."""
    import shutil
    import tempfile

    from verstka.analysis.classify import image_bytes_for_vlm

    cfg = {**load_config(), **(cfg or {})}
    if not cfg.get("enabled", True) or providers is None or skills is None or not providers.has("vlm") or _left(deadline) < 10.0:
        return {}
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="look-", dir=str(out_dir)))
    try:
        try:
            shoot = images or (lambda p, d, dpi: slide_images(p, d, dpi, timeout_s=max(15.0, min(240.0, _left(deadline)))))
            shots = shoot(Path(pptx), work, int(cfg["dpi"]))
        except Exception as e:  # noqa: BLE001 - no LibreOffice: nothing to look at
            log.warning("design review of the deck: slide images failed: %s", e)
            return {}
        names = titles or {}
        todo = [i for i in sorted(shots) if not _SKIP_TITLE_RE.match(names.get(i, "") or "")]

        def review(i: int) -> tuple[int, list[ReviewProblem]]:
            if _left(deadline) < 5.0:
                return i, []
            try:
                res = skills.run("slide_design_review", providers, {"slide_index": i, "deck_title": "", "form": "слайд", "headline": names.get(i, ""), "data_note": ""},
                                 images=[image_bytes_for_vlm(shots[i])])
                return i, res.parsed.known()
            except Exception as e:  # noqa: BLE001
                log.info("design review of slide %s failed: %s", i, str(e)[:160])
                return i, []

        limit = getattr(getattr(providers, "limits", None), "max_concurrency", 4) or 4
        workers = max(1, min(int(cfg["workers"]), int(limit), len(todo) or 1))
        with ThreadPoolExecutor(max_workers=workers) as ex:
            return {i: ps for i, ps in ex.map(review, todo) if ps}
    finally:
        shutil.rmtree(work, ignore_errors=True)


def problem_issues(problems: dict[int, list[ReviewProblem]], outline_ids: Optional[dict[int, Optional[str]]] = None) -> list[Issue]:
    """The serious problems the review saw, as the audit's model remarks (`design_review`)."""
    out: list[Issue] = []
    for i, ps in sorted(problems.items()):
        for k, p in enumerate(q for q in ps if q.severity == "major"):
            out.append(Issue(
                id=f"design_review-{i}-{k}", slide=i, check_id=DESIGN_REVIEW.id, severity="warn", kind="model",
                message=f"{CODES[p.code].capitalize()}" + (f" ({p.where})" if p.where else "") + (f": {p.evidence}" if p.evidence else ""),
                details={"code": p.code, "severity": p.severity}, outline_id=(outline_ids or {}).get(i),
                suggestion=FIX_HINTS.get(p.code, "переложите слайд: другая форма, шире блоки, меньше элементов в ряд"),
            ))
    return out


_STYLES = ("compact", "visual", "structured")


def relayout(plan: Any, outline: Any, oid: str, manifest: Any, taken: Optional[set] = None, others: Optional[dict] = None) -> str:
    """The slide's next layout: a slide on a template sample takes its next untried sample, then the composer (the
    autofix ladder, autofix._rematch); a composed slide is composed again in another variant's style (each style once),
    the composer's own knob — a sample left behind because it fit worse stays behind."""
    from verstka.audit.autofix import _rematch

    ps = plan.for_outline(oid)
    if ps is None:
        return "skip"
    if ps.mode == "synth" and ps.composition and ps.composition not in ("title", "section", "thanks"):
        tried = set(ps.fit.get("styles_tried", [])) | {ps.fit.get("style") or outline.strategy}
        for style in _STYLES:
            if style not in tried:
                ps.fit["styles_tried"] = sorted(tried | {style})
                ps.fit["style"] = style
                ps.reasons.append(f"взгляд дизайнера: композиция в стиле варианта «{style}»")
                return f"restyle → {style}"
        return "skip"
    # a sample another variant shows on this slide is never the new layout: the three variants stay apart (the ladder goes
    # on to the composer instead)
    return _rematch(plan, outline, oid, manifest, taken, others=others, strict=True)


def ru_slides(n: int) -> str:
    tail = n % 100
    word = "слайдов" if 11 <= tail <= 14 else {1: "слайд", 2: "слайда", 3: "слайда", 4: "слайда"}.get(n % 10, "слайдов")
    return f"{n} {word}"


def _slides_list(ids: list[int]) -> str:
    return ("слайде " if len(ids) == 1 else "слайдах ") + ", ".join(map(str, ids))


def _lower(text: str) -> str:
    """The model's sentence after a colon: its first letter small, its full stop off («Новая вёрстка…» → «новая вёрстка…»)."""
    t = " ".join(str(text or "").split()).rstrip(".")
    return t[:1].lower() + t[1:] if t[:2] != t[:2].upper() else t


def _how(result: str) -> str:
    if result.startswith("rematch"):
        return " на другом макете шаблона"
    if result.startswith("synth"):
        return " композицией из стилей шаблона"
    if result.startswith("restyle"):
        return " в другой композиции"
    return ""
