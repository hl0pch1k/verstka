"""Template compliance checks: fonts, sizes, colours, layouts, chrome, contrast."""

from __future__ import annotations

import re

from typing import Optional

from verstka.analysis.chrome import shape_signature
from verstka.audit.checks.common import composite_hex, fill_alpha, fix, ground_of, is_chrome_like, is_template_chrome, ru_count, template_grounds, text_elements, text_height_needed_pt
from verstka.audit.registry import AuditContext, check
from verstka.schemas.audit import CheckSpec, Issue
from verstka.schemas.common import EMU_PER_PT, Bbox, Color, contrast_ratio
from verstka.schemas.outline import PHOTO_PLACE_NAME

FONT_NOT_IN_TEMPLATE = CheckSpec(id="font_not_in_template", title="Шрифт не из шаблона или гарнитур больше двух", severity="error", category="template", description="Гарнитура рана отсутствует среди шрифтов шаблона (используемых, встроенных или шрифтов темы), либо на слайде больше двух гарнитур. Имена сравниваются без регистра и дефисов, начертание в имени — та же гарнитура («Bebas» и «Bebas Neue», «Montserrat-Regular» и «Montserrat»); ссылки на шрифты темы (+mj-lt, +mn-lt) раскрываются; номер слайда и колонтитулы шаблона не проверяются.")
SIZE_NOT_IN_SCALE = CheckSpec(id="size_not_in_scale", title="Размер шрифта не из шкалы шаблона", severity="warn", category="template", description="Размер шрифта отличается более чем на 0,75 пт от всех размеров, встречающихся в шаблоне. Крупные числа (значение с единицей, от наибольшего кегля шаблона или вдвое крупнее основного текста) — отдельная ступень шкалы: они не проверяются.")
COLOR_NOT_IN_PALETTE = CheckSpec(id="color_not_in_palette", title="Цвет не из палитры шаблона", severity="warn", category="template", description="Цвет текста или заливки отстоит от ближайшего цвета палитры шаблона больше чем на ΔE 6 и не является более светлым или тёмным вариантом цветного цвета палитры (тот же тон и насыщенность, как ряды оттенков темы в PowerPoint). Своими считаются и цвета, которые задают макеты и мастер шаблона, а также чёрный и белый; служебные элементы шаблона (номер слайда, колонтитулы) не проверяются.")
LAYOUT_NOT_FROM_TEMPLATE = CheckSpec(id="layout_not_from_template", title="Слайд собран не на макете из шаблона", severity="error", category="template", description="Слайд ссылается на макет, которого нет в пакете шаблона.")
CHROME_MOVED = CheckSpec(id="chrome_moved", title="Логотип или колонтитул сдвинуты с положенного места", severity="warn", category="template", description="Элемент хрома шаблона (логотип, колонтитул на слайдах) отсутствует или стоит в другом месте.")
TABLE_CONTRAST_LOW = CheckSpec(id="table_contrast_low", title="Контраст текста в таблице ниже нормы", severity="warn", category="template", description="Контраст по WCAG между текстом ячейки таблицы и её заливкой (или фоном слайда) ниже 4.5:1 для обычного текста и 3:1 для крупного (≥18 пт или ≥14 пт жирным) и для значков ✓ / —. Пара цветов шаблона (белый на фирменном синем) от 3:1 — только справка.")
CONTRAST_LOW = CheckSpec(id="contrast_low", title="Контраст текста к фону ниже 4.5:1", severity="warn", category="template", description="Контраст по WCAG между цветом текста и фоном ниже 4.5:1 для обычного текста и 3:1 для крупного (≥18 пт или ≥14 пт жирным). Фон — карточка слайда с учётом прозрачности заливки, иначе плашка, панель или картинка, которую рисуют макет и мастер шаблона, иначе фон слайда; градиент и картинка берутся в том месте, где стоит текст. На картинке или градиенте цвет фона — оценка: от 2:1 — сведение, ниже — ошибка. Строки, которые выходят за край своей плашки на другой фон (двухстрочный заголовок в однострочной полосе), проверяются и на том фоне; у текста на фигуре макета (треугольник, произвольная фигура) каждая строка проверяется по ширине своих букв — буквы, сошедшие с наклонного края на фон рядом, сравниваются с тем фоном.")

CONTRAST_CANDIDATE_ROLES = ("text.primary", "text.secondary")


def _norm_font(name: str) -> str:
    """«Bebas Neue» → «bebas neue», «Montserrat-Regular» → «montserrat regular»: lower case, separators as spaces."""
    return " ".join(re.sub(r"[-_]+", " ", (name or "").lower()).split())


def _theme_font(name: str, ctx: AuditContext) -> str:
    """A theme reference (+mj-lt, +mn-ea …) resolved through the deck's theme fonts."""
    tf = list(getattr(ctx.ir, "theme_fonts", []) or [])
    if name.startswith("+mj") and tf and tf[0]:
        return tf[0]
    if name.startswith("+mn") and len(tf) > 1 and tf[1]:
        return tf[1]
    return name


def same_family(a: str, b: str) -> bool:
    """Two names of one family: equal after normalisation, or one is the other plus more words and the shorter has
    at least 4 characters («Bebas» ≡ «Bebas Neue», «Open Sans» ≡ «Open Sans Light», «Montserrat» ≡ «Montserrat-Regular»)."""
    a, b = _norm_font(a), _norm_font(b)
    if not a or not b:
        return False
    if a == b:
        return True
    short, long_ = (a, b) if len(a) <= len(b) else (b, a)
    return len(short) >= 4 and long_.startswith(short + " ")


def _allowed_fonts(ctx: AuditContext) -> set[str]:
    typo = ctx.manifest.tokens.typography
    # anything the template itself uses, and the fonts of its theme (what its placeholders fall back to)
    fam = {f.family.lower() for f in typo.families if f.weight > 0 or getattr(f, "source", "text") == "theme"}
    fam |= {f.lower() for f in ctx.manifest.embedded_fonts}
    fam |= {f.lower() for f in ctx.ir.embedded_fonts}
    fam |= {f.lower() for f in (getattr(ctx.ir, "theme_fonts", []) or []) if f}
    # the faces the template's own slots are set in (its samples' and their layouts' placeholder styles): a template
    # whose sample slides are empty placeholders counts no text of them, yet its headings are designed in that face —
    # a heading the renderer measured and set in it, «refonted» into the theme's wider face, ran to a third line
    fam |= {s.style.font_family.lower() for p in ctx.manifest.patterns for s in p.slots if s.style.font_family and not s.style.font_family.startswith("+")}
    # the stand-ins the renderer sets Russian text in where a template family has no Cyrillic (rendering/cyrillic.py)
    if ctx.ws is not None:
        try:
            from verstka.rendering.cyrillic import cyrillic_substitutes

            fam |= {v.lower() for v in cyrillic_substitutes(ctx.ws.source, extra=[f.family for f in typo.families]).values()}
        except Exception:  # noqa: BLE001 - the template's own families then
            pass
    return fam


def _font_groups(fonts: list[str]) -> list[list[str]]:
    """Fonts grouped by family (`same_family`), first appearance order."""
    groups: list[list[str]] = []
    for f in fonts:
        for g in groups:
            if any(same_family(f, x) for x in g):
                g.append(f)
                break
        else:
            groups.append([f])
    return groups


@check(FONT_NOT_IN_TEMPLATE)
def font_not_in_template(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    allowed = _allowed_fonts(ctx)
    if not allowed:
        return out
    for s in ctx.ir.slides:
        used: dict[str, list[str]] = {}
        for e in text_elements(s):
            if is_template_chrome(e, ctx.manifest):
                continue  # the template's own page number / footer keeps the template's own font
            for p in e.paragraphs:
                for r in p.runs:
                    if r.font and r.text.strip():
                        used.setdefault(_theme_font(r.font, ctx), []).append(e.id)
        # a monospace face is the code's own (a code slide sets its lines in Consolas whatever the template's font): never
        # «refonted» into the template's proportional face — the code's indents would no longer line up
        from verstka.rendering.compose import MONO_RE

        bad = {f: ids for f, ids in used.items() if not f.startswith("+") and not any(same_family(f, a) for a in allowed) and not MONO_RE.search(f)}
        for f, ids in bad.items():
            out.append(ctx.new_issue(FONT_NOT_IN_TEMPLATE, s.index, f"шрифт «{f}» отсутствует в шаблоне", element_ids=sorted(set(ids)), bboxes=[s.by_id(i).bbox_frac for i in sorted(set(ids))[:3] if s.by_id(i)], autofix=fix("refont", "заменить на основной шрифт шаблона", element_ids=sorted(set(ids)))))
        # (the code's monospace face is no third face of the slide's design: heading, text and the code — as designed)
        groups = _font_groups([f for f in used if not MONO_RE.search(f)])
        if len(groups) > 2:
            out.append(ctx.new_issue(FONT_NOT_IN_TEMPLATE, s.index, f"на слайде {ru_count(len(groups), 'шрифт', 'шрифта', 'шрифтов')}: {', '.join(sorted(g[0] for g in groups))}", severity="warn"))
    return out


_FIGURE_RE = re.compile(r"^[\s+\-−–×x~≈«»]*\d[\d\s\u00a0\u202f.,]*(%|[a-zа-яё₽$€]{0,6}\.?(?:[\s\u00a0][₽$€])?)?(\s?[→\-–/]\s?[\d\s.,]+(%|[a-zа-яё₽]{0,6}\.?)?)?\s*$", re.I)


def is_figure(text: str) -> bool:
    """«6,5», «40%», «×4,8», «27 млн», «31% → 12%»: a value set as a figure, not running text."""
    t = text.strip()
    return 0 < len(t) <= 16 and bool(_FIGURE_RE.match(t))


@check(SIZE_NOT_IN_SCALE)
def size_not_in_scale(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    sizes = ctx.manifest.tokens.typography.sizes_used or [st.size_pt for st in ctx.manifest.tokens.typography.scale]
    if not sizes:
        return out
    # sizes the samples inherit from their layouts (a cover title set at 48 pt by the placeholder) are the template's own
    sizes = sorted(set(sizes) | {sl.style.size_pt for p in ctx.manifest.patterns for sl in p.slots if sl.style.size_pt})
    body = ctx.manifest.tokens.typography.size_for("body", 14.0)
    display_from = min(max(sizes), 2 * body)
    from verstka.rendering.compose import MONO_RE
    for s in ctx.ir.slides:
        bad: dict[float, list[str]] = {}
        for e in text_elements(s):
            if is_template_chrome(e, ctx.manifest):
                continue  # the template's own page number (a big «02») is set at the template's own size
            for p in e.paragraphs:
                for r in p.runs:
                    if r.size_pt and r.size_pt >= display_from and is_figure(r.text):
                        continue  # the display step of the scale: figures may be set larger than any text
                    if r.font and MONO_RE.search(r.font):
                        continue  # the user's code: sized by its longest line, as its monospace face is its own
                    if r.size_pt and r.text.strip() and all(abs(r.size_pt - t) > 0.75 for t in sizes):
                        bad.setdefault(r.size_pt, []).append(e.id)
        for sz, ids in bad.items():
            out.append(ctx.new_issue(SIZE_NOT_IN_SCALE, s.index, f"размер шрифта {sz:g} пт не из шкалы шаблона", element_ids=sorted(set(ids)), severity="info" if any(abs(sz - t) <= 2.0 for t in sizes) else "warn", details={"size": sz}))
    return out


SHADE_HUE = 4.0  # degrees
SHADE_SAT = 0.08
SHADE_MIN_SAT = 0.15


def is_shade_of(hex_: str, palette: list[str]) -> bool:
    """A lighter or darker variant of a chromatic palette colour — the same hue and saturation (HSL), only another
    lightness: what PowerPoint lists under each theme colour («Акцент 1, темнее 25%»). A greyed or re-hued colour is
    not one, and neither is a grey (every grey would be a «shade» of black)."""
    import colorsys

    def hsl(h: str) -> tuple[float, float, float]:
        r, g, b = (int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
        hh, ll, ss = colorsys.rgb_to_hls(r, g, b)
        return hh * 360.0, ss, ll

    try:
        h0, s0, l0 = hsl(hex_)
    except ValueError:
        return False
    if s0 < SHADE_MIN_SAT or not 0.08 <= l0 <= 0.92:
        return False
    for p in palette:
        try:
            h1, s1, _ = hsl(p)
        except ValueError:
            continue
        if s1 < SHADE_MIN_SAT:
            continue
        dh = abs(h0 - h1) % 360.0
        if min(dh, 360.0 - dh) <= SHADE_HUE and abs(s0 - s1) <= SHADE_SAT:
            return True
    return False


@check(COLOR_NOT_IN_PALETTE)
def color_not_in_palette(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    palette = [Color(hex=h) for h in ctx.manifest.tokens.palette()]
    if not palette:
        return out
    # the colours the template's layouts and masters set themselves (its page-number grey, its text styles), and
    # pure black and white, are the template's own
    palette += [Color(hex=h) for h in dict.fromkeys(list(getattr(ctx.ir, "template_colors", []) or []) + ["000000", "FFFFFF"])]

    def nearest(hex_: str) -> float:
        c = Color(hex=hex_)
        return min(Color.delta_e(c, p) for p in palette)

    shades_of = [p.hex for p in palette]

    for s in ctx.ir.slides:
        seen: dict[str, list[str]] = {}
        for e in s.elements:
            if e.type == "chart" or (e.name or "").startswith("Swatch "):
                continue  # chart colours (and a legend's swatches, which repeat them) come from the palette by construction; sample charts are removed
            if e.name == PHOTO_PLACE_NAME:
                continue  # the free place for the user's photo: a stand-in tint of the ground, gone once the photo is in
            if is_template_chrome(e, ctx.manifest):
                continue
            cands = []
            if e.fill_hex and e.type in ("shape", "text"):
                cands.append(e.fill_hex)
            for p in e.paragraphs:
                for r in p.runs:
                    if r.color_hex and r.text.strip():
                        cands.append(r.color_hex)
            for h in cands:
                try:
                    d = nearest(h)
                except ValueError:
                    continue
                if d > 6.0 and not is_shade_of(h, shades_of):
                    seen.setdefault(h, []).append(e.id)
        for h, ids in seen.items():
            out.append(ctx.new_issue(COLOR_NOT_IN_PALETTE, s.index, f"цвет #{h} не из палитры шаблона", element_ids=sorted(set(ids)), autofix=fix("recolor", "заменить ближайшим цветом палитры", hex=h, element_ids=sorted(set(ids)), scope="all")))
    return out


@check(LAYOUT_NOT_FROM_TEMPLATE)
def layout_not_from_template(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    known = set(ctx.ir.layout_parts)
    for s in ctx.ir.slides:
        if s.layout_part and s.layout_part not in known:
            out.append(ctx.new_issue(LAYOUT_NOT_FROM_TEMPLATE, s.index, f"макет {s.layout_part} отсутствует в шаблоне"))
    return out


def _chrome_name(c) -> str:
    """What a person calls a repeated template element: a logo, the slide number, a caption."""
    text = (c.text or "").strip()
    if c.kind == "pic":
        return "логотип или картинка шаблона"
    if text.isdigit():
        return "номер слайда"
    if text:
        return f"надпись шаблона «{text[:30]}»"
    return "элемент оформления шаблона"


@check(CHROME_MOVED)
def chrome_moved(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    slide_chrome = [c for c in ctx.manifest.tokens.chrome if c.source == "slide" and c.share >= 0.6]
    if not slide_chrome:
        return out
    for s in ctx.ir.slides:
        for c in slide_chrome:
            present = False
            for e in s.elements:
                if (c.kind == "pic") != (e.type == "picture"):
                    continue
                if e.bbox_frac.close_to(c.bbox, 0.02) and (not c.text or c.text.strip()[:20] == e.text.strip()[:20] or c.text.strip().isdigit()):
                    present = True
                    break
            if not present:
                out.append(ctx.new_issue(CHROME_MOVED, s.index, f"{_chrome_name(c)} отсутствует или стоит не на своём месте", bboxes=[c.bbox], severity="info"))
    return out


def _contrast_replacement(tokens, bg: str, need: float) -> Optional[str]:
    """First template text colour (then white, black) that reaches the required ratio against bg; None if nothing does."""
    cands: list[str] = []
    for role in CONTRAST_CANDIDATE_ROLES:
        hx = tokens.color_for(role)
        if hx:
            cands.append(hx.upper())
    cands.extend(("FFFFFF", "000000"))
    for c in dict.fromkeys(cands):
        try:
            if contrast_ratio(c, bg) >= need:
                return c
        except ValueError:
            continue
    return None


SPILL_LINE = 0.1  # share of a line height the text's lines may run past its band before it counts as spilling


def _line_band(e) -> Bbox:
    """The band the text's lines really fill at the box's anchor — not clipped to the box: an overflowing text runs
    past its box (down from a top anchor, both ways from a centre anchor, up from a bottom anchor)."""
    need, _ = text_height_needed_pt(e)
    h = max(int(need * EMU_PER_PT) + e.insets_emu[1] + e.insets_emu[3], 1)
    anchor = e.anchor or "t"
    y = e.bbox.y if anchor == "t" else (e.bbox.y2 - h if anchor == "b" else e.bbox.y + (e.bbox.h - h) // 2)
    return Bbox(x=e.bbox.x, y=y, w=e.bbox.w, h=h)


def _band_spill(s, e, bg_slide: str, color: str, size: float) -> Optional[tuple[float, str]]:
    """(contrast, ground) of the worst strip where the lines of a text that stands on a template band run past the
    band's top or bottom edge by more than SPILL_LINE of a line; None when they stay inside it."""
    grounds = template_grounds(s, e)
    if not grounds:
        return None
    band = min(grounds, key=lambda o: o.bbox.area)
    lines = e.bbox if e.autofit == "norm" else _line_band(e)  # a shrink-on-overflow box keeps its text inside it
    line_h = size * 1.2 * EMU_PER_PT
    strips = []
    if band.bbox.y - lines.y >= SPILL_LINE * line_h:
        strips.append(Bbox(x=lines.x, y=lines.y, w=lines.w, h=band.bbox.y - lines.y))
    if lines.y2 - band.bbox.y2 >= SPILL_LINE * line_h:
        strips.append(Bbox(x=lines.x, y=band.bbox.y2, w=lines.w, h=lines.y2 - band.bbox.y2))
    worst = None
    for strip in strips:
        probe = e.model_copy(update={"bbox": strip, "fill_hex": None})
        under = ground_of(s, probe, bg_slide)[0] or bg_slide
        try:
            cr = contrast_ratio(color, under)
        except ValueError:
            continue
        if worst is None or cr < worst[0]:
            worst = (cr, under)
    return worst


LINE_OFF = 0.08  # share of a line's letters off its shaped ground before the ground beside it counts


def _line_off_ground(s, e, bg_slide: str, color: str, need: float) -> Optional[tuple[float, str, str, float]]:
    """(contrast, ground, line) of the worst line of a text standing on a shaped template layer (a triangle, a
    freeform) whose letters leave that shape — measured line by line at the width each line really takes: a centred
    heading's second line may start on the white gap beside a slanted edge while the box as a whole stands on the
    triangle (LO Focus, gate 2 — C2). The ground there is the topmost other template layer painting it, else the
    slide's own ground. Returns (contrast, ground, line, share of the line's letters there) for the worst ground that
    lacks `need` under at least two samples of a line; None when every line keeps on its shape or on a ground it
    reads on."""
    from verstka.audit.checks.common import line_boxes
    from verstka.audit.checks.layout import _paints_at

    shaped = [o for o in template_grounds(s, e) if o.outline]
    if not shaped:
        return None
    band = shaped[-1]
    tpl = sorted((o for o in getattr(s, "template_elements", []) or [] if o is not band and o.fill_hex and o.type in ("shape", "text") and o.opaque), key=lambda o: (0 if o.source == "master" else 1, o.z))
    cache: dict = {}
    worst = None
    texts = [ln for p in e.paragraphs for ln in re.split(r"[\n\x0b]", p.text) if ln.strip()]
    for i, (box, _size) in enumerate(line_boxes(e)):
        nx = 48  # a letter is about 1/15 of a line: one at the end off the shape is seen
        pts = [(box.x + (k + 0.5) * box.w / nx, box.y + f * box.h) for k in range(nx) for f in (0.3, 0.5, 0.7)]
        off = [pt for pt in pts if not _paints_at(band, pt[0], pt[1])]
        if len(off) < 2:
            continue
        seen: dict[str, int] = {}
        for px, py in off:
            under = next((o.fill_hex.upper() for o in reversed(tpl) if _paints_at(o, px, py, None, cache)), None) or bg_slide
            seen[under] = seen.get(under, 0) + 1
        for under, n in seen.items():
            try:
                cr = contrast_ratio(color, under)
            except ValueError:
                continue
            if n < 2 or cr >= need:
                continue
            if worst is None or cr < worst[0]:
                worst = (cr, under, texts[i] if i < len(texts) else e.text, n / float(len(pts)))
    return worst


def _letter_probes(ctx: AuditContext, s, e) -> list[str]:
    """The colours under the letters of text `e` read on the deck itself (`layers.letter_grounds`); [] when the deck
    or the shape cannot be read."""
    try:
        if "_deck" not in ctx.__dict__:
            from pptx import Presentation

            ctx.__dict__["_deck"] = Presentation(str(ctx.ir.source))
        prs = ctx.__dict__["_deck"]
        if prs is None:
            return []
        from verstka.rendering.deck import slide_shape_elements
        from verstka.rendering.layers import letter_grounds

        slide = prs.slides[s.index - 1]
        el = slide_shape_elements(slide).get(str(e.id))
        return letter_grounds(slide, el, ctx.manifest.tokens.typography.primary_family) if el is not None else []
    except Exception:  # noqa: BLE001 - the box's estimate then
        ctx.__dict__.setdefault("_deck", None)
        return []


@check(CONTRAST_LOW)
def contrast_low(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    tokens = ctx.manifest.tokens
    for s in ctx.ir.slides:
        bg_slide = s.background_hex or tokens.color_for("background.dark" if s.family.value == "dark" else "background.light") or ("000000" if s.family.value == "dark" else "FFFFFF")
        for e in text_elements(s):
            if is_chrome_like(e, ctx.ir, ctx.manifest):
                continue
            color = e.dominant_color
            if not color:
                continue
            enclosing, ground_kind = ground_of(s, e, bg_slide)
            under = enclosing or bg_slide
            # a ground of unknown colour: report, never recolour
            uncertain = enclosing is None and not s.background_hex
            # a picture or gradient ground (a layout photo, a gradient panel or background): its colour under the text
            # is only estimated (the region's mean, or the whole ground's median) — below 2:1 it is still an error,
            # above it a note, and it is never recoloured
            estimated = ground_kind in ("picture", "gradient") or (enclosing is None and (s.background_kind or "") in ("image", "gradient", "render"))
            try:
                bg = composite_hex(e.fill_hex, fill_alpha(e), under) if e.fill_hex else under
                cr = contrast_ratio(color, bg)
            except ValueError:
                continue
            if estimated and not e.fill_hex and e.ph_type not in ("title", "ctrTitle"):
                # a picture or a gradient under the text: read where its letters stand, as the renderer reads it when it
                # chooses the colour (`layers.letter_grounds`) — the box's mean read a pink glow as the dark purple
                # around it. The second-worst probe: one probe on a speck is no line. (A heading keeps the template's
                # own colour on the template's own picture: its box's estimate, as before.)
                probes = _letter_probes(ctx, s, e)
                if probes:
                    try:
                        crs = sorted((contrast_ratio(color, g), g) for g in probes)
                    except ValueError:
                        crs = []
                    if crs:
                        cr, bg = crs[1] if len(crs) >= 3 else crs[0]
            size = e.dominant_size or 14.0
            bold = any(r.bold for p in e.paragraphs for r in p.runs if r.text.strip())
            need = 3.0 if size >= 18 or (bold and size >= 14) else 4.5  # WCAG: large text is 18 pt, or 14 pt bold
            spill = _band_spill(s, e, bg_slide, color, size) if ground_kind == "template" and not e.fill_hex else None
            if spill is not None and spill[0] < need:
                # the text is read on its band, but its lines run out of the band onto another ground (a two-line
                # heading in a one-line band: the second line white on white)
                cr2, under2 = spill
                out.append(ctx.new_issue(CONTRAST_LOW, s.index, f"строки «{e.text[:30]}» выходят за плашку шаблона: контраст {cr2:.1f}:1 (#{color} на #{under2})", bboxes=[e.bbox_frac], element_ids=[e.id], severity="error" if cr2 < 2.5 else "warn", details={"contrast": round(cr2, 2), "background": under2, "ground": "band_spill"}))
                continue
            off = _line_off_ground(s, e, bg_slide, color, need) if ground_kind == "template" and not e.fill_hex else None
            if off is not None:
                # the box stands on a triangle of the layout, but a line's letters run off its slanted edge onto the
                # ground beside it (white on the white gap): a letter or two at the edge — a warning, more — an error
                cr3, under3, line, share = off
                out.append(ctx.new_issue(CONTRAST_LOW, s.index, f"строка «{line[:30]}» выходит за фигуру шаблона на другой фон: контраст {cr3:.1f}:1 (#{color} на #{under3})", bboxes=[e.bbox_frac], element_ids=[e.id], severity="error" if cr3 < 2.5 and share >= LINE_OFF else "warn", details={"contrast": round(cr3, 2), "background": under3, "ground": "line_off_shape", "share": round(share, 3)}))
                continue
            if cr < need:
                accent_text = any(t.hex == color and any(r.startswith("accent.") for r in t.roles) for t in tokens.colors)
                severity = "error" if cr < 2.5 else ("info" if (accent_text and cr >= 3.0) else "warn")  # brand accent on dark is the template's own choice
                # the heading keeps the template's own colour (white on a pink pill is the design, not a slip)
                is_heading = e.ph_type in ("title", "ctrTitle")
                if is_heading and severity != "error":
                    severity = "info"
                to = None if uncertain or estimated or is_heading else _contrast_replacement(tokens, bg, need)
                autofix = fix("recolor", f"заменить цвет текста на #{to}", element_ids=[e.id], to=to, scope="text") if to else None
                if uncertain or (estimated and cr >= 2.0):
                    severity = "info"
                out.append(ctx.new_issue(CONTRAST_LOW, s.index, f"контраст {cr:.1f}:1 у «{e.text[:30]}» (#{color} на #{bg})", bboxes=[e.bbox_frac], element_ids=[e.id], severity=severity, details={"contrast": round(cr, 2), "background": bg, "ground": ground_kind or ("estimated" if estimated else "slide")}, autofix=autofix))
    return out


_TABLE_MARKS = {"✓", "✔", "✗", "✕", "×", "—", "–", "-", "−", "+"}


@check(TABLE_CONTRAST_LOW)
def table_contrast_low(ctx: AuditContext) -> list[Issue]:
    """Table text against its cell fill, or the slide ground under an unfilled cell. One issue per table (the worst
    cell). A pairing of two template colours (white on the brand blue) between 3:1 and 4.5:1 is the template's own
    choice: reported as info."""
    out: list[Issue] = []
    tokens = ctx.manifest.tokens
    palette = {t.hex.upper() for t in tokens.colors} | {"FFFFFF", "000000"}
    for s in ctx.ir.slides:
        bg_slide = s.background_hex
        for e in s.elements:
            if e.type != "table" or not e.table or not e.table.cells:
                continue
            # an unfilled cell shows what lies under the table: the panel it stands on (a white half of a purple
            # slide), else the slide's ground
            try:
                enclosing, _ = ground_of(s, e, bg_slide or "FFFFFF")
            except Exception:  # noqa: BLE001
                enclosing = None
            ground = enclosing or bg_slide
            worst = None
            for r_i, row in enumerate(e.table.cells):
                for c_i, cell in enumerate(row):
                    text = e.table.rows[r_i][c_i] if r_i < len(e.table.rows) and c_i < len(e.table.rows[r_i]) else ""
                    bg = cell.fill_hex or ground
                    if not text.strip() or not cell.color_hex or not bg:
                        continue
                    try:
                        cr = contrast_ratio(cell.color_hex, bg)
                    except ValueError:
                        continue
                    size = cell.size_pt or 14.0
                    mark = text.strip() in _TABLE_MARKS  # ✓ / — are marks, not text: the icon gate (3:1) applies
                    need = 3.0 if mark or size >= 18 or (cell.bold and size >= 14) else 4.5
                    if cr < need and (worst is None or cr / need < worst[0]):
                        worst = (cr / need, cr, need, text, cell.color_hex, bg, cell.fill_hex is None)
            if worst is None:
                continue
            _, cr, need, text, color, bg, _ = worst
            brand = color.upper() in palette and bg.upper() in palette and cr >= 3.0
            severity = "error" if cr < 2.5 else ("info" if brand else "warn")
            out.append(ctx.new_issue(TABLE_CONTRAST_LOW, s.index, f"контраст {cr:.1f}:1 у «{text[:30]}» в таблице (#{color} на #{bg}, нужно {need:g}:1)", bboxes=[e.bbox_frac], element_ids=[e.id], severity=severity, details={"contrast": round(cr, 2), "background": bg, "brand_pair": brand}))
    return out


# ── qa3 (N-U1): the brand logo on every content slide, in one corner ──────────────────────────────────────────────
BRAND_LOGO_PLACE = CheckSpec(
    id="brand_logo_place",
    title="Логотип бренда не на всех слайдах или в разных углах",
    severity="warn",
    category="template",
    description="В колоде по брендбуку логотип бренда стоит на каждом слайде с содержимым и в одном и том же углу "
                "(обложка, разделы и финальный слайд — по своим образцам).",
)


@check(BRAND_LOGO_PLACE)
def brand_logo_place(ctx: AuditContext) -> list[Issue]:
    from verstka.rendering.brand_logo import is_brand_logo_any as is_brand_logo

    try:
        from pptx import Presentation

        prs = Presentation(str(ctx.ir.source))
    except Exception:  # noqa: BLE001 - no file to read: nothing to say
        return []
    W, H = float(prs.slide_width), float(prs.slide_height)
    kinds = {o.id: (o.kind.value if hasattr(o.kind, "value") else str(o.kind)) for o in (ctx.outline.slides if ctx.outline is not None else [])}
    corners: dict[int, Optional[str]] = {}
    any_logo = False
    for k, (slide, s) in enumerate(zip(prs.slides, ctx.ir.slides), 1):
        # (a slide the deck's marks no longer name — a finished file audited again — is a cover or a closing slide when
        # it is the first or the last one)
        if kinds.get(s.outline_id or "") in ("title", "section", "thanks") or (not s.outline_id and k in (1, len(ctx.ir.slides))):
            continue
        corner = None
        own = [el for el in slide._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}pic") if is_brand_logo(el)]
        # (the slide's own logo, else the one its layout draws)
        drawn = own or [el for el in slide.slide_layout._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}pic") if is_brand_logo(el)]
        for el in drawn:
            from lxml import etree as _et

            from verstka.rendering.deck import element_bbox

            top = el  # (a logo inside a group sits in the group's coordinates: the group's own box says where it is)
            while top.getparent() is not None and _et.QName(top.getparent()).localname not in ("spTree",):
                top = top.getparent()
            b = element_bbox(top)
            if not b or b[2] <= 0 or b[3] <= 0:
                continue
            cx, cy = b[0] + b[2] / 2, b[1] + b[3] / 2
            corner = ("сверху" if cy < H / 2 else "снизу") + " " + ("слева" if cx < W / 3 else ("справа" if cx > 2 * W / 3 else "по центру"))
            any_logo = True
            break
        corners[s.index] = corner
    if not any_logo:
        return []
    out = []
    missing = [i for i, c in corners.items() if c is None]
    for i in missing:
        out.append(ctx.new_issue(BRAND_LOGO_PLACE, i, "На слайде нет логотипа бренда, хотя на других слайдах колоды он есть",
                                 suggestion="Оставьте под логотип угол колоды: содержимое его обходит."))
    seen = [c for c in corners.values() if c]
    if len(set(seen)) > 1:
        main = max(set(seen), key=seen.count)
        for i, c in corners.items():
            if c and c != main:
                out.append(ctx.new_issue(BRAND_LOGO_PLACE, i, f"Логотип бренда {c}, а на других слайдах — {main}",
                                         suggestion="Один угол логотипа на всю колоду."))
    return out


# ── qa4: the typographic floor of every slide, chips, logos (a score of 100 means what it says) ─────────────────────
FIGURE_LABEL_RATIO = CheckSpec(
    id="figure_label_ratio",
    title="Ключевая цифра почти не крупнее своей подписи",
    severity="warn",
    category="template",
    description="Ключевая цифра слайда (число плана с подписью, цифра карточки, цифра, с которой начинается строка "
                "под короткой меткой «Было / Стало») набрана меньше чем в 1,6 раза крупнее своей подписи: цифра и подпись "
                "читаются одним кеглем. Цифра и подпись ищутся по тексту плана, пара — ближайшие абзацы.",
)
HEADLINE_IN_CHIP = CheckSpec(
    id="headline_in_chip",
    title="Заголовок слайда стоит на плашке-метке",
    severity="error",
    category="template",
    description="Заголовок слайда с содержимым набран на плашке-метке (невысокая цветная плашка не шире половины слайда "
                "под началом текста) вместо метки раздела: заголовка под ней нет, метка «съела» вывод.",
)
BODY_BELOW_TEMPLATE = CheckSpec(
    id="body_below_template",
    title="Основной текст мельче основного кегля шаблона",
    severity="info",  # a note until calibrated on unseen templates (final review C1)
    category="template",
    description="Строка основного текста (от 25 знаков и четырёх слов: пункт, текст карточки, шаг, подпись цифры) "
                "набрана мельче основного кегля шаблона (больше чем на полпункта). Сноски, источники, подписи "
                "диаграмм и таблиц — мелкий текст по замыслу, их проверяет «Текст мельче читаемого».",
)
TEXT_TOO_SMALL = CheckSpec(
    id="text_too_small",
    title="Текст мельче читаемого",
    severity="error",
    category="template",
    description="Текст слайда с содержимым (и мелкий: сноска, источник, ячейка таблицы) набран мельче 1,7% высоты "
                "слайда (самый мелкий кегль плотной вёрстки) — 9,2 пт на слайде 7,5″, 7 пт на слайде 5,6″: такой текст с "
                "экрана не читается. Служебные "
                "элементы шаблона (номер слайда, колонтитулы) не проверяются.",
)
LOGO_OFF_SLIDE = CheckSpec(
    id="logo_off_slide",
    title="Логотип выходит за край слайда",
    severity="error",
    category="template",
    description="Видимая часть логотипа бренда (его непрозрачные пиксели, с учётом группы, в которой он стоит) выходит "
                "за край слайда: логотип обрезан («VK Clouc»).",
)

_FOOTNOTE_RE = re.compile(r"^\s*(?:источник|источники|данные|по данным|примечание|\*|source)\b", re.I)
_CONTENT_SKIP_KINDS = ("title", "section", "thanks")


def _content_slides(ctx: AuditContext):
    """(IR slide, outline slide or None) of the content slides: covers, dividers and closing slides set display type."""
    by_id = {o.id: o for o in (ctx.outline.slides if ctx.outline is not None else [])}
    n = len(ctx.ir.slides)
    for s in ctx.ir.slides:
        o = by_id.get(s.outline_id or "")
        kind = (o.kind.value if hasattr(o.kind, "value") else str(o.kind)) if o is not None else None
        if kind in _CONTENT_SKIP_KINDS or (o is None and s.index in (1, n)):
            continue
        yield s, o


def _records_of(s) -> list[tuple]:
    """The slide's paragraphs as `clone_check.match_figures` records (text, digits' size, largest size, x, y, element,
    paragraph, box)."""
    out = []
    for e in text_elements(s):
        b = e.bbox
        paras = [p for p in e.paragraphs]
        for j, p in enumerate(paras):
            sizes = [(r.text, r.size_pt) for r in p.runs if r.size_pt and r.text.strip()]
            if not sizes or not p.text.strip():
                continue
            dig = [z for t, z in sizes if any(ch.isdigit() for ch in t)]
            out.append((p.text, max(dig) if dig else 0.0, max(z for _, z in sizes), b.x + b.w / 2, b.y + (j + 0.5) * b.h / max(len(paras), 1), e, p, (b.x, b.y, b.w, b.h)))
    return out


@check(FIGURE_LABEL_RATIO)
def figure_label_ratio(ctx: AuditContext) -> list[Issue]:
    from verstka.rendering.clone_check import FIG_RATIO, _sq, box_gap, figure_pairs, match_figures

    out: list[Issue] = []
    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    for s, o in _content_slides(ctx):
        if o is None:
            continue
        pairs = figure_pairs(o)
        if not pairs:
            continue
        recs = _records_of(s)
        found = match_figures(pairs, recs)
        for fig, lab, fs, ls, fr, lr in found:
            if fs < FIG_RATIO * ls - 0.3:
                out.append(ctx.new_issue(FIGURE_LABEL_RATIO, s.index, f"цифра «{fig.strip()[:16]}» {fs:g} пт почти не крупнее подписи «{lab.strip()[:24]}» {ls:g} пт (нужно от {FIG_RATIO:g}×)",
                                         element_ids=[fr[5].id, lr[5].id], bboxes=[fr[5].bbox_frac, lr[5].bbox_frac], details={"figure_pt": fs, "label_pt": ls},
                                         suggestion="Цифра крупнее, подпись — основным кеглем шаблона; иначе другой образец или композиция."))
                break
            gx, gy = box_gap(fr[7], lr[7])
            if gx > 0.1 * W or gy > 0.12 * H:
                # qa4 NEW-L3: «до 5 часов» alone at the slide's top-right edge, its words in a list at the left
                out.append(ctx.new_issue(FIGURE_LABEL_RATIO, s.index, f"цифра «{fig.strip()[:16]}» стоит отдельно от своей подписи «{lab.strip()[:24]}»",
                                         element_ids=[fr[5].id, lr[5].id], bboxes=[fr[5].bbox_frac, lr[5].bbox_frac], details={"apart": True},
                                         suggestion="Цифра и её подпись — в одном блоке."))
                break
        else:
            kind = o.kind.value if hasattr(o.kind, "value") else str(o.kind)
            if kind in ("stat_row", "big_number"):
                shown = {_sq(f[0]) for f in found}
                lone = next((fig for fig, _, lead in pairs if not lead and _sq(fig) not in shown and any(r[1] and _sq(r[0]) == _sq(fig) for r in recs)), None)
                if lone:
                    out.append(ctx.new_issue(FIGURE_LABEL_RATIO, s.index, f"у цифры «{lone.strip()[:16]}» нет подписи рядом", details={"unlabelled": True},
                                             suggestion="У каждой ключевой цифры — подпись рядом, что она считает."))
    return out


def _plate_under(s, e, W: int, H: int):
    """A filled plate (the slide's own shape, or the template's) under the start of a text's first line: (element, box)."""
    from verstka.audit.checks.common import line_boxes

    try:
        rows = line_boxes(e)
    except Exception:  # noqa: BLE001
        rows = []
    first = rows[0][0] if rows else e.bbox
    px, py = first.x + min(first.w, int(0.02 * W)) // 2, first.y + first.h // 2
    best = None
    for o in list(s.elements) + list(getattr(s, "template_elements", []) or []):
        if o is e or o.type not in ("shape", "text") or not o.fill_hex or fill_alpha(o) < 0.6 or (o.type == "text" and o.has_text and o is not e):
            continue
        b = o.bbox
        if b.x <= px <= b.x2 and b.y <= py <= b.y2 and b.area < 0.25 * W * H and (best is None or b.area < best[1].area):
            best = (o, b)
    if e.fill_hex and fill_alpha(e) >= 0.6 and best is None:
        best = (e, e.bbox)
    return best


@check(HEADLINE_IN_CHIP)
def headline_in_chip(ctx: AuditContext) -> list[Issue]:
    from verstka.rendering.clone_check import _sq

    out: list[Issue] = []
    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    for s, o in _content_slides(ctx):
        head = _sq((o.headline if o is not None else "") or "")
        if len(head) < 6:
            continue
        for e in text_elements(s):
            if e.bbox_frac.y > 0.35:
                continue
            t = _sq(e.text)
            if not t or not (t == head or (len(t) >= 12 and head.startswith(t))):
                continue
            plate = _plate_under(s, e, W, H)
            if plate is not None and plate[1].h < 0.1 * H and plate[1].w <= 0.5 * W:
                out.append(ctx.new_issue(HEADLINE_IN_CHIP, s.index, f"заголовок «{e.text.strip()[:40]}» стоит на плашке-метке вместо названия раздела",
                                         element_ids=[e.id], bboxes=[e.bbox_frac], suggestion="Плашка называет раздел или вид слайда, заголовок — под ней."))
                break
    return out


_SMALL_NAME_RE = re.compile(r"^(?:Footnote|Source|Caption|Chart note|Chart title|Legend|Unit|Axis)")


def is_small_print(name: str, text: str) -> bool:
    """A box or a line that is small print by design (by the box's name or the line's opening): a source or a note
    line, a caption, a chart's note, legend or axis — set under the body size on purpose. Autofix shares it."""
    return bool(_SMALL_NAME_RE.match(name or "") or _FOOTNOTE_RE.match(text or ""))


def is_running_text(text: str) -> bool:
    """A line of running text (a bullet, a card's text, a step, a figure's label): 25 characters and four words."""
    t = (text or "").strip()
    return len(t) >= 25 and len(t.split()) >= 4


_SMALL_HERE_RE = re.compile(r"^(?:Photo caption|Label|Takeaway)")


def _small_print(e, p, o) -> bool:
    """A paragraph that is small print by design: a source or a note line, the slide's footnote, a caption — and, for
    this check, a picture place's description, a figure's label and the conclusion of a slide of data beside or under
    its chart (final review C1: «Вид теплицы сверху…» at 12 pt took two brief decks from 100 to 91)."""
    if is_small_print(e.name or "", p.text or ""):
        return True
    name = e.name or ""
    if _SMALL_HERE_RE.match(name):
        return True
    if name.startswith("Conclusion") and o is not None and getattr(o, "data", None) is not None and (o.content.chart is not None or o.content.table is not None):
        return True
    if o is not None and (o.footnote or "").strip() and (p.text or "").strip()[:30] == (o.footnote or "").strip()[:30]:
        return True
    return False


@check(BODY_BELOW_TEMPLATE)
def body_below_template(ctx: AuditContext) -> list[Issue]:
    from verstka.rendering.clone_check import body_floor

    out: list[Issue] = []
    floor = body_floor(ctx.manifest)
    for s, o in _content_slides(ctx):
        worst = None
        for e in text_elements(s):
            if is_chrome_like(e, ctx.ir, ctx.manifest):
                continue
            from verstka.rendering.compose import MONO_RE

            for p in e.paragraphs:
                txt = (p.text or "").strip()
                if not is_running_text(txt) or _small_print(e, p, o):
                    continue
                if any(r.font and MONO_RE.search(r.font) for r in p.runs):
                    continue
                sizes = [r.size_pt for r in p.runs if r.size_pt and r.text.strip()]
                if not sizes:
                    continue
                z = max(sizes)
                if z < floor - 0.5 and (worst is None or z < worst[0]):
                    worst = (z, e, txt)
        if worst is not None:
            z, e, txt = worst
            out.append(ctx.new_issue(BODY_BELOW_TEMPLATE, s.index, f"текст «{txt[:40]}» набран {z:g} пт — мельче основного кегля шаблона {floor:g} пт",
                                     element_ids=[e.id], bboxes=[e.bbox_frac], details={"size": z, "floor": floor},
                                     suggestion="Основной текст — не мельче основного кегля шаблона: меньше текста на слайде или другая раскладка."))
    return out


@check(TEXT_TOO_SMALL)
def text_too_small(ctx: AuditContext) -> list[Issue]:
    from verstka.rendering.clone_check import small_floor

    out: list[Issue] = []
    floor = small_floor(ctx.manifest)
    for s, o in _content_slides(ctx):
        worst = None
        for e in s.elements:
            if e.type == "table" and e.table is not None:
                for row, cells in zip(e.table.rows, e.table.cells):
                    for txt, c in zip(row, cells):
                        if txt.strip() and c.size_pt and c.size_pt < floor - 0.3 and (worst is None or c.size_pt < worst[0]):
                            worst = (c.size_pt, e, txt.strip())
                continue
            if e.type != "text" or not e.has_text or is_template_chrome(e, ctx.manifest):
                continue
            if is_chrome_like(e, ctx.ir, ctx.manifest) and not e.text.strip()[:1].isalpha():
                continue  # a page number, a mark in the corner
            for p in e.paragraphs:
                for r in p.runs:
                    if r.size_pt and r.text.strip() and r.size_pt < floor - 0.3 and (worst is None or r.size_pt < worst[0]):
                        worst = (r.size_pt, e, p.text.strip())
        if worst is not None:
            z, e, txt = worst
            out.append(ctx.new_issue(TEXT_TOO_SMALL, s.index, f"текст «{txt[:40]}» набран {z:g} пт — мельче читаемого ({floor:.1f} пт на этом слайде)",
                                     element_ids=[e.id], bboxes=[e.bbox_frac], details={"size": z, "floor": round(floor, 1)},
                                     suggestion="Меньше текста или строк таблицы, крупнее кегль; лишнее — в заметки."))
    return out


@check(LOGO_OFF_SLIDE)
def logo_off_slide(ctx: AuditContext) -> list[Issue]:
    from verstka.rendering.brand_logo import is_brand_logo_any as is_brand_logo
    from verstka.rendering.deck import abs_bbox
    from verstka.rendering.layers import opaque_box

    try:
        from pptx import Presentation

        prs = Presentation(str(ctx.ir.source))
    except Exception:  # noqa: BLE001 - no file to read: nothing to say
        return []
    W, H = int(prs.slide_width), int(prs.slide_height)
    tol = int(0.002 * W)
    out: list[Issue] = []
    for k, (slide, s) in enumerate(zip(prs.slides, ctx.ir.slides), 1):
        holders = [(slide.part, slide._element)]
        try:
            holders.append((slide.slide_layout.part, slide.slide_layout._element))
        except Exception:  # noqa: BLE001
            pass
        for part, root in holders:
            for el in root.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}pic"):
                if not is_brand_logo(el):
                    continue
                b = abs_bbox(el)
                if not b or b[2] <= 0 or b[3] <= 0:
                    continue
                try:
                    vis = opaque_box(part, el, Bbox(x=int(b[0]), y=int(b[1]), w=int(b[2]), h=int(b[3])))
                except Exception:  # noqa: BLE001
                    vis = Bbox(x=int(b[0]), y=int(b[1]), w=int(b[2]), h=int(b[3]))
                if vis.w <= 0 or vis.h <= 0:
                    continue
                if vis.x < -tol or vis.y < -tol or vis.x2 > W + tol or vis.y2 > H + tol:
                    cut = max(-vis.x, -vis.y, vis.x2 - W, vis.y2 - H) / max(vis.w if (vis.x < 0 or vis.x2 > W) else vis.h, 1)
                    out.append(ctx.new_issue(LOGO_OFF_SLIDE, s.index, f"логотип выходит за край слайда на {cut:.0%} своей ширины",
                                             details={"box": [vis.x / W, vis.y / H, vis.w / W, vis.h / H]},
                                             suggestion="Логотип целиком на слайде, с полем от края, как в брендбуке."))
                    break
            else:
                continue
            break
    return out
