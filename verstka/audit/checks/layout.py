"""Layout checks: bounds, overlaps, overflow, margins, stretched images, table cells."""

from __future__ import annotations

import re
from typing import Optional

from verstka.audit.checks.template import is_figure
from verstka.audit.checks.common import CONTENT_TYPES, at_template_position, contains, fill_alpha, fix, ground_of, is_chrome_like, is_template_chrome, ru_count, ru_times, text_elements, text_height_needed_pt, title_element, usable_height_pt
from verstka.audit.registry import AuditContext, check
from verstka.rendering.fonts import text_width_pt
from verstka.rendering.groups import PLATE_MARK
from verstka.schemas.audit import CheckSpec, Issue
from verstka.schemas.common import EMU_PER_PT, Bbox

OUT_OF_BOUNDS = CheckSpec(id="out_of_bounds", title="Элемент вышел за границы слайда", severity="error", category="layout", description="Текст, таблица или диаграмма выходит за край слайда более чем на 1%; картинка — более чем на 25% своей площади. Служебные элементы шаблона (номер слайда, дата, колонтитул) стоят там, где их поставил шаблон, и не проверяются.")
OVERLAP = CheckSpec(id="overlap", title="Два блока наложились друг на друга", severity="error", category="layout", description="Два текстовых блока пересекаются более чем на 30% меньшего из них (вложенность в карточку не считается). У рамки без заливки и обводки считается полоса, которую занимают её строки (по привязке к верху, центру или низу): пустая часть рамки ничего не закрывает.")
TEXT_OVERFLOW = CheckSpec(id="text_overflow", title="Текст не поместился в свою рамку", severity="error", category="layout", description="Оценка высоты текста по метрикам шрифта превышает высоту рамки более чем на 8% (сильное переполнение — ошибка, лёгкое — предупреждение). Текст, который шаблон набирает прописными (cap=all в стиле текста или плейсхолдера макета), меряется прописными. Поля номера слайда и даты шаблона не проверяются. Одна строка в рамке, подогнанной шаблоном под её буквы (выше рамки не больше чем в 1,2 раза, привязка к верху), — не переполнение: под рамку уходит только межстрочный просвет.")
TEXT_CLIPPED = CheckSpec(id="text_clipped", title="Текст обрезан краем слайда", severity="error", category="layout", description="Рамка с текстом частично за пределами слайда (кроме номера слайда, даты и колонтитула шаблона).")
MARGIN_VIOLATION = CheckSpec(id="margin_violation", title="Контент заходит в поля у краёв", severity="warn", category="layout", description="Текстовый блок начинается за пределами безопасной области шаблона (допуск 3% ширины). Служебные элементы шаблона не проверяются.")
GRID_ALIGNMENT = CheckSpec(id="grid_alignment", title="Блоки не выровнены по направляющим макета", severity="info", category="layout", description="Левый край текстового блока не совпадает ни с одной колонкой шаблона (допуск 1.5% ширины).")
IMAGE_STRETCHED = CheckSpec(id="image_stretched", title="Картинка растянута, пропорции нарушены", severity="warn", category="layout", description="Пропорции рамки картинки отличаются от пропорций исходного изображения (с учётом кадрирования) более чем на 12%. Подложка карточки шаблона, нарисованная картинкой (ровная плашка без рисунка), растягивается вместе со своей карточкой, как нарисованная фигура, — её вёрстка помечает в имени «· подложка», и такая картинка не считается искажённой.")
TEXT_OUTSIDE_CARD = CheckSpec(id="text_outside_card", title="Текст выходит за свою карточку", severity="error", category="layout", description="Строки текстового блока, который начинается внутри карточки (залитой или обведённой фигуры), заходят за её нижний или правый край более чем на 1,5% размера слайда (больше 3% — ошибка): текст висит под карточкой и наезжает на то, что ниже. Подложка, на которой текст тоже начинается (две перекрывающиеся панели обложки), — его фон, а не карточка, в которую он заходит.")
TABLE_CELL_WRAP = CheckSpec(id="table_cell_wrap", title="Слово в ячейке таблицы переносится по буквам", severity="warn", category="layout", description="Самое длинное слово в ячейке нативной таблицы шире своей колонки (ширина из a:gridCol минус поля ячейки; без сетки — ширина рамки / число колонок минус 2×7.2 пт) при кегле и начертании самой ячейки — PowerPoint рвёт его посреди слова.")

WORD_BREAK = CheckSpec(id="word_break", title="Слово шире своей рамки и рвётся посередине", severity="warn", category="layout", description="Слово абзаца (по метрикам шрифта, при кегле и начертании своего фрагмента, прописными — если шаблон набирает этот текст прописными) шире рамки текста за вычетом полей и отступа маркера — рендерер разорвёт его посреди слова («Корректировк / а»). Допуск 3%: рендерер набирает текст чуть шире, чем его меряют; неразрывный пробел держит слова вместе, дефис без словосоединителя — место переноса. Рамки, которые сами уменьшают шрифт (normAutofit), и служебные элементы шаблона (номер слайда, дата, колонтитул) не проверяются.")

TABLE_CELL_INSET_PT = 7.2  # default a:tcPr marL/marR
WORD_ROOM = 0.97  # a word wider than this share of its line is broken by the renderer
_HARD_SPACE_RE = re.compile(r"[ \t\r\n]+")  # NBSP stays inside a word, as PowerPoint keeps it on one line


def _template_bleed(f, manifest, ws=None) -> bool:
    """A picture standing exactly where a sample of the template sets one (a 3D render cut by the slide's edge on a
    template slide the deck stands on): the template's design, not a picture that slid off."""
    if manifest is None:
        return False
    boxes = []
    for p in manifest.patterns:
        boxes += list(p.decor_boxes) + [s.bbox for s in p.slots if s.role.value in ("image", "icon", "decoration")]
    if ws is not None:
        try:
            from verstka.analysis.pictures import sample_pictures

            boxes += [pic.box for pics in sample_pictures(ws, manifest).values() for pic in pics]
        except Exception:  # noqa: BLE001 - the manifest's boxes then
            pass
    for b in boxes:
        if abs(b.x - f.x) <= 0.006 and abs(b.y - f.y) <= 0.006 and abs(b.w - f.w) <= 0.02 * max(f.w, 1e-6) + 0.004 and abs(b.h - f.h) <= 0.02 * max(f.h, 1e-6) + 0.004:
            return True
    return False


@check(OUT_OF_BOUNDS)
def out_of_bounds(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    for s in ctx.ir.slides:
        for e in s.elements:
            if e.type not in CONTENT_TYPES or e.bbox.area <= 0:
                continue
            if is_template_chrome(e, ctx.manifest):
                continue  # the template's own page number / footer stands where the template put it
            f = e.bbox_frac
            if e.type == "picture":
                inside_w = max(0.0, min(f.x2, 1.0) - max(f.x, 0.0))
                inside_h = max(0.0, min(f.y2, 1.0) - max(f.y, 0.0))
                if inside_w * inside_h < 0.75 * f.area and not _template_bleed(f, ctx.manifest, ctx.ws):
                    out.append(ctx.new_issue(OUT_OF_BOUNDS, s.index, f"картинка «{e.name}» на {100 - int(100 * inside_w * inside_h / max(f.area, 1e-9))}% за краем слайда", bboxes=[f], element_ids=[e.id], severity="warn", autofix=fix("move_inside", "сдвинуть внутрь слайда", element_id=e.id)))
                continue
            if f.x < -0.01 or f.y < -0.01 or f.x2 > 1.01 or f.y2 > 1.01:
                out.append(ctx.new_issue(OUT_OF_BOUNDS, s.index, f"«{(e.text[:40] or e.name)}» выходит за край слайда", bboxes=[f], element_ids=[e.id], autofix=fix("move_inside", "сдвинуть внутрь слайда", element_id=e.id)))
    return out


@check(TEXT_CLIPPED)
def text_clipped(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    for s in ctx.ir.slides:
        for e in text_elements(s):
            if is_template_chrome(e, ctx.manifest):
                continue
            f = e.bbox_frac
            if (f.x2 > 1.005 or f.y2 > 1.005 or f.x < -0.005 or f.y < -0.005) and f.x < 1.0 and f.y < 1.0:
                out.append(ctx.new_issue(TEXT_CLIPPED, s.index, f"текст «{e.text[:40]}» обрезан краем слайда", bboxes=[f], element_ids=[e.id], autofix=fix("rematch", "перевыбрать макет", outline_id=s.outline_id)))
    return out


def _occupied(e, spacing: float) -> Bbox:
    """The band of a text box its lines fill: a box without fill or outline shows only its text, so its empty part
    (the ascender room over a figure set on the box's bottom) overlaps nothing."""
    if e.fill_hex or e.line_hex:
        return e.bbox
    need, lines = text_height_needed_pt(e, spacing)
    if lines == 1 and is_figure(e.text.strip()):
        # a figure shows its digits: the ascender room of a 160 pt line over them is empty
        from verstka.rendering.fonts import figure_metrics_em

        size = max((r.size_pt or 0.0) for p in e.paragraphs for r in p.runs) or need
        descent, digit_h = figure_metrics_em(any(r.bold for p in e.paragraphs for r in p.runs))
        need = min(need, (descent + digit_h) * size * 1.05)
    h = int(need * EMU_PER_PT) + e.insets_emu[1] + e.insets_emu[3]
    if h >= e.bbox.h:
        return e.bbox
    anchor = e.anchor or "t"
    y = e.bbox.y if anchor == "t" else (e.bbox.y2 - h if anchor == "b" else e.bbox.y + (e.bbox.h - h) // 2)
    return Bbox(x=e.bbox.x, y=y, w=e.bbox.w, h=h)


@check(OVERLAP)
def overlap(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    spacing = 1.2  # the IR carries the spacing each paragraph is shown with; one set nowhere is single-spaced
    for s in ctx.ir.slides:
        texts = [e for e in text_elements(s) if not is_chrome_like(e, ctx.ir, ctx.manifest)]
        seen: set[tuple[str, str]] = set()
        for i, a in enumerate(texts):
            for b in texts[i + 1 :]:
                if (a.id, b.id) in seen:
                    continue
                oa, ob = _occupied(a, spacing), _occupied(b, spacing)
                inter = oa.intersection(ob)
                if inter <= 0:
                    continue
                smaller = min(oa.area, ob.area)
                if smaller <= 0:
                    continue
                # a filled card holding a text box is intended nesting
                if (a.fill_hex and contains(a.bbox, b.bbox)) or (b.fill_hex and contains(b.bbox, a.bbox)):
                    continue
                ratio = inter / smaller
                if ratio > 0.3:
                    seen.add((a.id, b.id))
                    out.append(
                        ctx.new_issue(
                            OVERLAP,
                            s.index,
                            f"«{a.text[:30]}» и «{b.text[:30]}» пересекаются на {int(ratio * 100)}%",
                            bboxes=[a.bbox_frac, b.bbox_frac],
                            element_ids=[a.id, b.id],
                            severity="error" if ratio > 0.5 else "warn",
                            details={"ratio": round(ratio, 2)},
                            autofix=fix("rematch", "перевыбрать макет слайда", outline_id=s.outline_id),
                        )
                    )
    return out




def _visible_card(slide, o) -> bool:
    """A shape a reader sees as a card: an outline of its own, or a fill that stands out from the ground under it (a
    background-coloured ellipse of the template — a soft glow — is no card a text could run out of), not nearly clear."""
    from verstka.schemas.common import contrast_ratio

    if o.line_hex:
        return True
    if fill_alpha(o) < 0.3:
        return False
    ground = ground_of(slide, o)[0] or slide.background_hex
    if not ground or not o.fill_hex:
        return True
    try:
        return contrast_ratio(o.fill_hex, ground) >= 1.1
    except (ValueError, TypeError):
        return True


@check(TEXT_OUTSIDE_CARD)
def text_outside_card(ctx: AuditContext) -> list[Issue]:
    """A text that starts inside a card (the smallest filled or outlined shape holding its first line) and whose lines
    run past the card's bottom or right edge: the column of a two-column slide whose list hangs below its card."""
    out: list[Issue] = []
    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    spacing = 1.2  # the IR carries the spacing each paragraph is shown with; one set nowhere is single-spaced
    for s in ctx.ir.slides:
        cards = [o for o in s.elements if o.type in ("shape", "text") and (o.fill_hex or o.line_hex) and o.bbox.area > 0 and not o.has_text and _visible_card(s, o)]
        for e in text_elements(s):
            if is_chrome_like(e, ctx.ir, ctx.manifest) or e.fill_hex or e.line_hex:
                continue
            occ = _occupied(e, spacing)
            x0, y0 = occ.x + min(occ.w // 20, int(0.01 * W)), occ.y + min(occ.h // 20, int(0.01 * H))
            hosts = [c for c in cards if c.bbox.x <= x0 <= c.bbox.x2 and c.bbox.y <= y0 <= c.bbox.y2 and c.bbox.area > occ.area * 0.3]
            if not hosts:
                continue
            host = min(hosts, key=lambda c: c.bbox.area)
            below = occ.y2 - host.bbox.y2
            right = occ.x2 - host.bbox.x2
            over = max(below / max(H, 1), right / max(W, 1))
            # or its lines run into another card under it (a card that grew into the conclusion strip). A card that
            # also holds the text's first line is its ground too, not a card it runs into (two panels of a cover laid
            # over each other, neither holding the other), and only what hangs below the host counts: running past its
            # right edge is `right` above
            host_ids = {id(c) for c in hosts}
            for c in cards:
                if id(c) in host_ids or c.bbox.area <= 0 or contains(c.bbox, host.bbox) or contains(host.bbox, c.bbox):
                    continue
                inter_h = min(occ.y2, c.bbox.y2) - max(occ.y, c.bbox.y, host.bbox.y2)
                inter_w = min(occ.x2, c.bbox.x2) - max(occ.x, c.bbox.x)
                if inter_h > 0 and inter_w > 0.3 * occ.w:
                    over = max(over, inter_h / max(H, 1))
            if over <= 0.015:
                continue
            out.append(
                ctx.new_issue(
                    TEXT_OUTSIDE_CARD,
                    s.index,
                    f"текст «{e.text[:40]}» выходит за свою карточку на {int(round(over * 100))}% слайда",
                    bboxes=[e.bbox_frac, host.bbox_frac],
                    element_ids=[e.id, host.id],
                    severity="error" if over > 0.03 else "warn",
                    details={"over": round(over, 3)},
                    autofix=fix("rematch", "перевыбрать макет слайда", outline_id=s.outline_id),
                )
            )
    return out


@check(TEXT_OVERFLOW)
def text_overflow(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    spacing = 1.2  # the IR carries the spacing each paragraph is shown with; one set nowhere is single-spaced
    for s in ctx.ir.slides:
        title = title_element(s)
        for e in text_elements(s):
            if e.bbox.h <= 0 or e.ph_type in ("sldNum", "dt"):
                continue  # the template's own slide-number and date fields keep the template's own boxes
            if is_template_chrome(e, ctx.manifest):
                continue  # the template's own menu row or footer keeps the template's own boxes (a spAutoFit label grows)
            need, lines = text_height_needed_pt(e, spacing)
            have = usable_height_pt(e)
            if have <= 0:
                continue
            ratio = need / have
            if lines == 1 and ratio <= 1.2 and (e.anchor or "t") == "t":
                # one line in a box the template sized to its letters (Google Slides measures a line at about 1.1 em,
                # the check at 1.2): only the line's leading hangs under the box, its letters stand inside — what
                # stands right under it is the overlap check's (VK Tech's card titles: «1,1 times their box», fine)
                continue
            if ratio > 1.08:
                grows = e.autofit == "sp"  # PowerPoint enlarges such a box downwards
                shrinks = e.autofit == "norm" and ratio <= 2.5  # PowerPoint shrinks the text itself; nothing to fix
                grown_bottom = e.bbox.y + (need + (e.insets_emu[1] + e.insets_emu[3]) / 12700.0) * 12700
                is_title = e.ph_type in ("title", "ctrTitle") or (e is title)
                severe = ratio > (2.5 if is_title else 1.5)  # title boxes are generous and rarely collide
                if grows:
                    severity = "error" if grown_bottom > ctx.ir.slide_h * 1.0 else "warn"
                    msg = f"текст «{e.text[:40]}» вырастит рамку {ru_times(ratio)} ({ru_count(lines, 'строка', 'строки', 'строк')})"
                elif shrinks:
                    severity = "info"
                    msg = f"текст «{e.text[:40]}» выше рамки {ru_times(ratio)} ({ru_count(lines, 'строка', 'строки', 'строк')}), рамка сама уменьшит шрифт"
                else:
                    severity = "error" if severe else "warn"
                    msg = f"текст «{e.text[:40]}» выше рамки {ru_times(ratio)} ({ru_count(lines, 'строка', 'строки', 'строк')})"
                rematch = severe and not grows and not shrinks
                if ctx.verbatim(s) and not is_title and not grows and not shrinks:
                    # the author's text set at the least size its layout allows: one whole sentence less (it stays in the
                    # speaker notes) — never a smaller type than the template's, never a cut inside a sentence
                    out.append(ctx.new_issue(TEXT_OVERFLOW, s.index, msg, bboxes=[e.bbox_frac], element_ids=[e.id], severity=severity,
                                             details={"ratio": round(ratio, 2), "lines": lines, "autofit": e.autofit},
                                             autofix=fix("condense_text", "показать на одно предложение меньше (оно останется в заметках)", outline_id=s.outline_id, element_id=e.id)))
                    continue
                out.append(
                    ctx.new_issue(
                        TEXT_OVERFLOW,
                        s.index,
                        msg,
                        bboxes=[e.bbox_frac],
                        element_ids=[e.id],
                        severity=severity,
                        details={"ratio": round(ratio, 2), "lines": lines, "autofit": e.autofit},
                        autofix=None if shrinks else fix("rematch" if rematch else "shrink_text", "перевыбрать макет" if rematch else "уменьшить шрифт по шкале шаблона — не мельче основного кегля, иначе другая раскладка", outline_id=s.outline_id, element_id=e.id, ratio=round(ratio, 2)),
                    )
                )
    return out


@check(MARGIN_VIOLATION)
def margin_violation(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    safe = ctx.manifest.tokens.spacing.safe_area
    prs = None
    for s in ctx.ir.slides:
        home = None
        for e in text_elements(s):
            if is_chrome_like(e, ctx.ir, ctx.manifest) or e.nested or at_template_position(e, ctx.manifest):
                continue
            if e.ph_type in ("title", "ctrTitle"):
                # a heading inside the box its own layout gives it stands where the template sets headings — the
                # margins measured on other samples (a cover's) are not this layout's
                if home is None:
                    prs = prs if prs is not None else _deck(ctx)
                    try:
                        home = _title_home(prs.slides[s.index - 1], ctx.ir.slide_w, ctx.ir.slide_h) if prs is not None else None
                    except Exception:  # noqa: BLE001
                        home = None
                    home = home or False
                if home and contains(home, e.bbox, 0.0):
                    continue
            f = e.bbox_frac
            if f.x < safe.x - 0.03 or f.x2 > safe.x2 + 0.03 or f.y2 > safe.y2 + 0.05:
                if ctx.verbatim(s) and f.y2 > safe.y2 + 0.05 and f.x >= safe.x - 0.03 and f.x2 <= safe.x2 + 0.03:
                    # the author's text ran down past the foot (more of it than its column holds): one whole sentence
                    # less (it stays in the speaker notes) — a box moved up would cover the text above it
                    out.append(ctx.new_issue(MARGIN_VIOLATION, s.index, f"«{e.text[:40]}» заходит в поля шаблона внизу", bboxes=[f], element_ids=[e.id],
                                             autofix=fix("condense_text", "показать на одно предложение меньше (оно останется в заметках)", outline_id=s.outline_id, element_id=e.id)))
                    continue
                out.append(ctx.new_issue(MARGIN_VIOLATION, s.index, f"«{e.text[:40]}» заходит в поля шаблона", bboxes=[f], element_ids=[e.id], autofix=fix("move_inside", "сдвинуть в безопасную область", element_id=e.id, safe=True)))
    return out


@check(GRID_ALIGNMENT)
def grid_alignment(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    cols = ctx.manifest.tokens.spacing.columns
    if not cols:
        return out
    for s in ctx.ir.slides:
        misaligned = []
        for e in text_elements(s):
            if is_chrome_like(e, ctx.ir, ctx.manifest) or e.nested or e.bbox_frac.w < 0.2:
                continue
            if all(abs(e.bbox_frac.x - c) > 0.015 for c in cols):
                misaligned.append(e)
        if len(misaligned) >= 2:
            out.append(ctx.new_issue(GRID_ALIGNMENT, s.index, f"{ru_count(len(misaligned), 'блок не выровнен', 'блока не выровнены', 'блоков не выровнены')} по колонкам шаблона", bboxes=[e.bbox_frac for e in misaligned[:4]], element_ids=[e.id for e in misaligned]))
    return out


@check(IMAGE_STRETCHED)
def image_stretched(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    for s in ctx.ir.slides:
        for e in s.elements:
            if e.type != "picture" or not e.image_size or e.bbox.h <= 0:
                continue
            if e.name.endswith(PLATE_MARK):
                continue  # a card's plate stretched with its card (rendering/groups.py `mark_plate`)
            iw, ih = e.image_size
            if iw <= 0 or ih <= 0:
                continue
            l, t, r, b = e.crop or (0.0, 0.0, 0.0, 0.0)
            cw = iw * max(1 - l - r, 0.01)
            ch = ih * max(1 - t - b, 0.01)
            img_ar = cw / ch
            box_ar = e.bbox.w / e.bbox.h
            dev = abs(box_ar / img_ar - 1)
            if dev > 0.12:
                out.append(ctx.new_issue(IMAGE_STRETCHED, s.index, f"картинка «{e.name}» искажена на {int(dev * 100)}%", bboxes=[e.bbox_frac], element_ids=[e.id], details={"deviation": round(dev, 2)}))
    return out


@check(TABLE_CELL_WRAP)
def table_cell_wrap(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    family = ctx.manifest.tokens.typography.primary_family
    size = ctx.manifest.components.table_style.font_size_pt or 12.0
    for s in ctx.ir.slides:
        for e in s.elements:
            if e.type != "table" or not e.table or not e.table.rows or e.bbox.w <= 0:
                continue
            n_cols = max((len(r) for r in e.table.rows), default=0)
            if n_cols == 0:
                continue
            grid = e.table.col_widths_emu if len(e.table.col_widths_emu) == n_cols and all(w > 0 for w in e.table.col_widths_emu) else None
            uniform_pt = e.bbox.w / EMU_PER_PT / n_cols - 2 * TABLE_CELL_INSET_PT  # no grid: uniform columns
            worst: tuple[float, str, int, float, float] | None = None
            for r_i, row in enumerate(e.table.rows):
                for c_i, cell in enumerate(row):
                    info = e.table.cells[r_i][c_i] if r_i < len(e.table.cells) and c_i < len(e.table.cells[r_i]) else None
                    if grid and c_i < len(grid):  # the real column, less this cell's own margins, at its own size
                        col_pt = grid[c_i] / EMU_PER_PT - ((info.mar_l_pt + info.mar_r_pt) if info else 2 * TABLE_CELL_INSET_PT)
                    else:
                        col_pt = uniform_pt
                    cell_size = (info.size_pt if info and info.size_pt else size)
                    cell_bold = info.bold if info and info.size_pt else (r_i == 0)
                    for word in _HARD_SPACE_RE.split(cell or ""):
                        if not word:
                            continue
                        w_pt = text_width_pt(word, family, cell_size, bold=cell_bold)
                        if w_pt > col_pt + 0.5 and (worst is None or w_pt - col_pt > worst[0]):
                            worst = (w_pt - col_pt, word, c_i, col_pt, cell_size)
            if worst is not None:
                over, word, c_i, col_pt, w_size = worst
                out.append(
                    ctx.new_issue(
                        TABLE_CELL_WRAP,
                        s.index,
                        f"слово «{word}» ({col_pt + over:.0f} пт) шире колонки {c_i + 1} таблицы ({col_pt:.0f} пт при {w_size:g} пт) и рвётся по буквам",
                        bboxes=[e.bbox_frac],
                        element_ids=[e.id],
                        details={"word": word, "column": c_i, "word_pt": round(col_pt + over, 1), "column_pt": round(col_pt, 1), "cols": n_cols},
                        autofix=fix("rematch", "перевыбрать макет с более широкой областью таблицы", outline_id=s.outline_id),
                    )
                )
    return out


_WORD_SPLIT_RE = re.compile(r"[ \t\r\n]+|(?<=-)(?!\u2060)")


@check(WORD_BREAK)
def word_break(ctx: AuditContext) -> list[Issue]:
    """A word wider than the line it stands on: a column too narrow for it (six steps in a row, a card of a long
    word) — the renderer breaks it in the middle."""
    out: list[Issue] = []
    for s in ctx.ir.slides:
        for e in text_elements(s):
            if not e.wrap or e.autofit == "norm" or e.rotation or e.bbox.w <= 0:
                continue
            if e.ph_type in ("sldNum", "dt", "ftr") or is_chrome_like(e, ctx.ir, ctx.manifest):
                continue  # the template's own page number, date and footer boxes: its chrome, not the deck's text
            usable = (e.bbox.w - e.insets_emu[0] - e.insets_emu[2]) / EMU_PER_PT
            if usable <= 0:
                continue
            worst: tuple[float, str, float] | None = None
            for p in e.paragraphs:
                size_p = next((r.size_pt for r in p.runs if r.size_pt), None) or e.dominant_size or 14.0
                line = usable - (size_p * 1.1 if p.bullet else 0.0)
                for r in p.runs:
                    size = r.size_pt or size_p
                    for word in _WORD_SPLIT_RE.split((r.text or "").upper() if e.caps else (r.text or "")):
                        if not word.strip():
                            continue
                        w_pt = text_width_pt(word, r.font, size, r.bold)
                        if w_pt > line * WORD_ROOM and (worst is None or w_pt - line > worst[0] - worst[2]):
                            worst = (w_pt, word, line)
            if worst is not None:
                w_pt, word, line = worst
                out.append(
                    ctx.new_issue(
                        WORD_BREAK,
                        s.index,
                        f"слово «{word.strip()}» ({w_pt:.0f} пт) не помещается в строку рамки ({line:.0f} пт) и рвётся посередине",
                        bboxes=[e.bbox_frac],
                        element_ids=[e.id],
                        details={"word": word.strip(), "word_pt": round(w_pt, 1), "line_pt": round(line, 1)},
                        autofix=fix("shrink_text", "уменьшить шрифт по шкале шаблона, чтобы слово встало в строку, — не мельче основного кегля, иначе другая раскладка", outline_id=s.outline_id, element_id=e.id, ratio=round(w_pt / (line * WORD_ROOM), 2)),
                    )
                )
    return out


CONTENT_OVER_ART = CheckSpec(
    id="content_over_art",
    title="Контент лежит на рисунке шаблона",
    severity="warn",
    category="layout",
    description=(
        "Текст (по месту, которое занимают его строки), диаграмма, таблица или картинка слайда перекрывает рисунок макета или мастера "
        "(иллюстрацию, фотографию, фигуру оформления — по её настоящему контуру, а не по рамке) на 15–90% своей площади; текст, "
        "который пересекает край фотографии или фигурного рисунка (перекрытие 10–90%), — ошибка; горизонтальная линия шаблона "
        "через строки текста — предупреждение. Фоном, а не рисунком, считаются заливка или спокойная картинка во весь слайд и "
        "прямоугольная плашка, на которой блок стоит целиком. Текст проверяется и построчно: строка, которая пересекает край "
        "фотографии или фигурного рисунка, хотя блок целиком — нет, — ошибка; строка, которая со спокойной части фотографии "
        "во весь слайд заходит на её пёструю часть (12–90% букв строки на ячейках с разбросом яркости от 24), — предупреждение."
    ),
)

ART_MIN_AREA = 0.003  # share of the slide: smaller template marks are not art
ART_WARN = (0.15, 0.90)
ART_CROSS = (0.10, 0.90)


def _inside_poly(px: float, py: float, poly) -> bool:
    inside = False
    n = len(poly)
    j = n - 1
    for i in range(n):
        xi, yi = poly[i]
        xj, yj = poly[j]
        if (yi > py) != (yj > py) and px < (xj - xi) * (py - yi) / ((yj - yi) or 1e-9) + xi:
            inside = not inside
        j = i
    return inside


VEIL_ALPHA = 0.25  # a layer nowhere more opaque than this is a glow or a veil over the ground, not art
SOLID_VEIL_ALPHA = 0.35  # a see-through solid shape up to this opacity is a veil too (white bubbles at 20–30 % over LibreOffice «Lights»)
CELL_ALPHA = 0.5  # a picture cell paints when at least half of it is opaque …
CELL_STD = 18.0  # … and it is textured (luminance stdev) …
CELL_CONTRAST = 1.3  # … or differs from the slide's ground


def _painted_cells(o, ground: Optional[str], cache: dict) -> Optional[list[bool]]:
    key = (id(o), ground)
    if key not in cache:
        c = o.cells
        out = None
        if c is not None and c.w * c.h == len(c.alpha) == len(c.hex) == len(c.std):
            out = []
            for a, hx, sd in zip(c.alpha, c.hex, c.std):
                paint = a >= CELL_ALPHA and (sd >= CELL_STD or not ground or _contrast(hx, ground) >= CELL_CONTRAST)
                out.append(paint)
        cache[key] = out
    return cache[key]


def _contrast(a: str, b: str) -> float:
    from verstka.schemas.common import contrast_ratio

    try:
        return contrast_ratio(a, b)
    except ValueError:
        return 21.0


def _paints_at(o, px: float, py: float, ground: Optional[str] = None, cache: Optional[dict] = None) -> bool:
    """Whether the template layer `o` paints at (px, py): its outline for shapes that are not rectangles, its opaque,
    textured or ground-contrasting cells for pictures, else its box."""
    b = o.bbox
    if not (b.x <= px <= b.x2 and b.y <= py <= b.y2):
        return False
    if o.outline:
        hits = sum(1 for poly in o.outline if len(poly) > 2 and _inside_poly(px, py, poly))
        return hits % 2 == 1
    if o.cells is not None:
        grid = _painted_cells(o, ground, cache if cache is not None else {})
        if grid is not None:
            cx = min(int((px - b.x) / max(b.w, 1) * o.cells.w), o.cells.w - 1)
            cy = min(int((py - b.y) / max(b.h, 1) * o.cells.h), o.cells.h - 1)
            return grid[cy * o.cells.w + cx]
    return True


def _is_art(o, W: int, H: int, ground: Optional[str] = None) -> bool:
    """A template layer that is drawing, not ground: painted (fill, outline or picture), not a text box, not a thin
    rule, not a tiny mark, not a full-bleed ground, not a plain shape in the ground's own colour."""
    if not (o.fill_hex or o.line_hex or o.type == "picture"):
        return False
    if o.type != "picture" and not o.line_hex and o.fill_hex and ground and (o.paint_kind or "solid") == "solid":
        from verstka.audit.checks.common import composite_hex

        try:
            seen = composite_hex(o.fill_hex, fill_alpha(o), ground)
        except ValueError:
            seen = o.fill_hex
        if _contrast(seen, ground) < CELL_CONTRAST:
            return False  # a panel of (almost) the ground's own colour — a faint veil — draws nothing
    if o.type == "text" and not o.fill_hex:
        return False
    if o.type in ("connector",) or o.bbox.w <= 0 or o.bbox.h <= 0:
        return False
    if o.cells is not None and o.cells.alpha and max(o.cells.alpha) < VEIL_ALPHA and not o.line_hex:
        return False  # a glow or a veil (a gradient at 7 %, a PNG glow at 14 % — Office's «Ion»): the ground shows through
    if o.cells is None and o.fill_hex and fill_alpha(o) < SOLID_VEIL_ALPHA and not o.line_hex:
        return False  # a see-through solid shape (a bubble at 30 %): a veil as well

    f = o.bbox_frac
    cover = max(0.0, min(f.x2, 1.0) - max(f.x, 0.0)) * max(0.0, min(f.y2, 1.0) - max(f.y, 0.0))
    if cover < ART_MIN_AREA or f.h <= 0.02:
        return False
    if cover >= 0.85:
        return False  # the slide's ground, a photo included (text on a photo is contrast_low's question)
    return True


_WRITING_ROLES = ("title", "subtitle", "body", "bullet_list", "card_title", "card_body", "number", "number_label", "caption")


def _template_writes_on(o, s, manifest, W: int, H: int, ground: Optional[str], cache: dict) -> bool:
    """The template's own sample slides set text on this layer (a text slot of a sample on the same layout — any
    layout for a master layer — lies at least half on what the layer paints): it is a ground the template designed for
    text (a glow inside a panel, a band under the heading), not art to keep content off."""
    key = ("writes", id(o), s.layout_part)
    if key in cache:
        return cache[key]
    got = False
    for p in getattr(manifest, "patterns", []) or []:
        if o.source == "layout" and p.layout_part and s.layout_part and p.layout_part != s.layout_part:
            continue
        for sl in p.slots:
            role = getattr(sl.role, "value", sl.role)
            if role not in _WRITING_ROLES:
                continue
            b = Bbox(x=int(sl.bbox.x * W), y=int(sl.bbox.y * H), w=max(int(sl.bbox.w * W), 1), h=max(int(sl.bbox.h * H), 1))
            if b.intersection(o.bbox) <= 0:
                continue
            hit = sum(1 for i in range(8) for j in range(4) if _paints_at(o, b.x + (i + 0.5) * b.w / 8, b.y + (j + 0.5) * b.h / 4, ground, cache))
            if hit >= 16:
                got = True
                break
        if got:
            break
    cache[key] = got
    return got


PHOTO_CROSS = (0.12, 0.90)  # a line of text with this share of its letters on the busy part of a photo ground crosses onto it
PHOTO_CROSS_MIN_W = 0.02  # … and at least this share of the slide's width of it


def _busy_cell(o, px: float, py: float, y0: float, y1: float) -> bool:
    """The photo `o` is textured at (px, py) for a line spanning y0..y1 (EMU): its cell there is opaque, its luminance
    stdev at least layers.PHOTO_BUSY_STD and, on the fine grid, not flat at the line's height (`layers.busy_for_line`
    — the renderer keeps cover lines off such cells, round 4.1, C3-2)."""
    from verstka.rendering.layers import busy_for_line

    b, c = o.bbox, o.cells
    if c is None or not (b.x <= px <= b.x2 and b.y <= py <= b.y2):
        return False
    cx = min(int((px - b.x) / max(b.w, 1) * c.w), c.w - 1)
    cy = min(int((py - b.y) / max(b.h, 1) * c.h), c.h - 1)
    if c.alpha[cy * c.w + cx] < CELL_ALPHA:
        return False
    fine = c.fine_std if c.w == 24 and c.h == 14 else None
    return busy_for_line(c.std, fine, cx, cy, (y0 - b.y) / max(b.h, 1), (y1 - b.y) / max(b.h, 1))


def letter_band(row: Bbox) -> Bbox:
    """The part of a line's box (`common.line_boxes`: a line's height, 1.2 em) its letters fill: from the top of the
    capitals to the baseline — the half-leading over them and the descenders' room under the baseline aside."""
    top = int(row.h * 0.15)
    return Bbox(x=row.x, y=row.y + top, w=row.w, h=max(int(row.h * 0.65), 1))


def _photo_crossings(e, s, photos: list, tpl: list, W: int) -> Optional[tuple[float, object]]:
    """(share, photo) of the first line of a text that runs from the calm part of a photo ground onto its busy part
    (LO Candy's subtitle on the candies): each line's letters (common.line_boxes), sampled 24 × 3, on the busy cells of
    a full-slide photo of the template, with no painted layer of the template over the photo there. None when no line
    crosses."""
    from verstka.audit.checks.common import fill_alpha, line_boxes

    if e.fill_hex and fill_alpha(e) >= 0.6:
        return None  # a text on a card of its own stands on the card
    try:
        rows = line_boxes(e)
    except Exception:  # noqa: BLE001
        return None
    for box, _size in rows:
        if box.w <= 0 or box.h <= 0:
            continue
        cx, cy = box.x + box.w / 2.0, box.y + box.h / 2.0
        for o in photos:
            if not (o.bbox.x <= cx <= o.bbox.x2 and o.bbox.y <= cy <= o.bbox.y2):
                continue
            own = o in s.elements
            over = [t for t in tpl if not own and t.z > o.z] + [t for t in s.elements if t is not e and (not own or t.z > o.z)]
            if any(t is not o and t.type in ("shape", "text", "group") and ((t.fill_hex and fill_alpha(t) >= 0.6) or t.outline) and _paints_at(t, cx, cy) for t in over):
                continue  # a panel over the photo (the layout's or the slide's own card): the line stands on the panel
            nx, ny = 24, 3
            # (the band the letters fill — the line's leading and its descenders' room aside: a full-slide parchment
            # whose ink blot starts right under a source line was taken for the line's ground)
            g = letter_band(box)
            hit = sum(1 for i in range(nx) for j in range(ny) if _busy_cell(o, g.x + (i + 0.5) * g.w / nx, g.y + (j + 0.5) * g.h / ny, g.y, g.y2))
            share = hit / float(nx * ny)
            if PHOTO_CROSS[0] <= share <= PHOTO_CROSS[1] and share * box.w >= PHOTO_CROSS_MIN_W * W:
                return share, o
            break
    return None


def _line_crossings(e, under: list, ground: Optional[str], cache: dict) -> Optional[tuple[float, object]]:
    """(share, art) of the first line of a multi-line text that runs onto the edge of a photo or a shaped drawing of
    the template (a line's letters 10–90 % on it) when the text as a block does not: a long last line reaching a tree,
    a second line under a triangle's slanted edge."""
    from verstka.audit.checks.common import line_boxes

    try:
        rows = line_boxes(e)
    except Exception:  # noqa: BLE001
        return None
    if len(rows) < 2:
        return None
    for box, _size in rows:
        if box.w <= 0 or box.h <= 0:
            continue
        for o in under:
            if not (o.busy or o.outline or o.geometry == "custom") or o.bbox.intersection(box) <= 0:
                continue
            nx, ny = 16, 3
            hit = sum(1 for i in range(nx) for j in range(ny) if _paints_at(o, box.x + (i + 0.5) * box.w / nx, box.y + (j + 0.5) * box.h / ny, ground, cache))
            share = hit / float(nx * ny)
            if ART_CROSS[0] <= share <= ART_CROSS[1]:
                return share, o
    return None


def _rule_through(o, box: Bbox, H: int) -> bool:
    """A horizontal template line (a connector, a line shape or a thin filled bar) crossing the inside of a text's
    line band over at least half of its width."""
    f = o.bbox_frac
    if not (f.h <= 0.02 and (o.line_hex or o.fill_hex)):
        return False
    return _line_through(o.bbox, box)


def _line_through(rule: Bbox, box: Bbox) -> bool:
    y = rule.y + rule.h / 2
    margin = 0.15 * box.h
    if not (box.y + margin < y < box.y2 - margin):
        return False
    return min(rule.x2, box.x2) - max(rule.x, box.x) >= 0.5 * box.w


def _hairlines(source: str, W: int, H: int) -> dict[int, list[Bbox]]:
    """The horizontal hairlines the slides' layouts and masters draw, by slide number: a connector or a line has no area,
    so the template layers of the IR leave them out (Классика: the rules of its comparison grid, drawn by the layout,
    struck the seventh line of a column through and no check saw it). Read from the deck itself; {} when it cannot be."""
    out: dict[int, list[Bbox]] = {}
    try:
        from pptx import Presentation

        from verstka.rendering.layers import drawn_layers

        prs = Presentation(source)
    except Exception:  # noqa: BLE001 - no deck to read (an IR built in memory): no hairlines known
        return out
    for i, slide in enumerate(prs.slides, 1):
        try:
            layers = drawn_layers(slide)
        except Exception:  # noqa: BLE001
            continue
        for lay in layers:
            b = lay.box
            if lay.source in ("layout", "master") and not lay.placeholder and not lay.has_text and lay.paints and not lay.picture and b.h <= 0.006 * H and b.w >= 0.05 * W:
                out.setdefault(i, []).append(b)
    return out


def _deck(ctx: AuditContext):
    """The audited deck as python-pptx reads it (once per audit), or None."""
    if "_deck" not in ctx.__dict__:
        try:
            from pptx import Presentation

            ctx.__dict__["_deck"] = Presentation(str(ctx.ir.source))
        except Exception:  # noqa: BLE001 - an IR built in memory: no file to read
            ctx.__dict__["_deck"] = None
    return ctx.__dict__["_deck"]


def _picture_lines(ctx: AuditContext) -> dict[int, list]:
    """By slide number, the lines the pictures under each whole slide draw (`rendering.picture_lines`: a panel's edges,
    VK Tech's divider down the middle of «3_Содержание»), less the lines the template's own samples on that layout
    write over (its designed ground). {} without a deck file."""
    if "_picture_lines" in ctx.__dict__:
        return ctx.__dict__["_picture_lines"]
    out: dict[int, list] = {}
    prs = _deck(ctx)
    if prs is not None:
        from verstka.rendering.picture_lines import designed_lines, ground_lines

        W, H = int(prs.slide_width), int(prs.slide_height)
        for i, (slide, s) in enumerate(zip(prs.slides, ctx.ir.slides), 1):
            try:
                lines = ground_lines(slide)
            except Exception:  # noqa: BLE001
                continue
            lines = [ln for ln in lines if ln.strength >= 6.0 and (ln.end - ln.start) >= 0.3 * (H if ln.vertical else W)]
            if lines:
                out[i] = designed_lines(lines, ctx.manifest, s.layout_part, W, H)
    ctx.__dict__["_picture_lines"] = out
    return out


def _line_cuts_row(ln, row: Bbox) -> bool:
    """A ground line through a line of text: a vertical one through its letters (not at their ends) along its whole
    height, a horizontal one through its middle band over half its width."""
    if ln.vertical:
        return row.x + 0.04 * row.w < ln.pos < row.x2 - 0.04 * row.w and ln.start <= row.y + 0.2 * row.h and ln.end >= row.y2 - 0.2 * row.h
    return row.y + 0.25 * row.h < ln.pos < row.y2 - 0.25 * row.h and min(ln.end, row.x2) - max(ln.start, row.x) >= 0.5 * row.w


@check(CONTENT_OVER_ART)
def content_over_art(ctx: AuditContext) -> list[Issue]:
    from verstka.audit.checks.density import _bleeds, _ink

    out: list[Issue] = []
    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    cache: dict = {}
    hairlines = _hairlines(ctx.ir.source, W, H)
    plines = _picture_lines(ctx)
    for s in ctx.ir.slides:
        if plines.get(s.index):
            # qa4 NEW-L1: a line drawn inside the picture under the slide runs through a text or under a card
            out += _picture_line_issues(ctx, s, plines[s.index], W, H)
        tpl = list(getattr(s, "template_elements", []) or [])
        if not tpl and not hairlines.get(s.index) and not any(o.type == "picture" and o.cells is not None for o in s.elements):
            continue
        ground = s.background_hex
        art = [o for o in tpl if _is_art(o, W, H, ground) and not _template_writes_on(o, s, ctx.manifest, W, H, ground, cache)]
        rules = [o for o in tpl if o.bbox_frac.h <= 0.02 and o.bbox_frac.w >= 0.05]
        # a full-slide photo is the ground, not art — but a line of text may still run from its calm part onto its busy
        # part (the candies of LO Candy): checked line by line on the photo's cells
        photos = [o for o in tpl if o.type == "picture" and o.cells is not None and _is_full_slide(o)]
        photos += [o for o in s.elements if o.type == "picture" and o.cells is not None and _is_full_slide(o)]  # the slide's own
        for e in s.elements:
            if e.type not in CONTENT_TYPES or e.bbox.area <= 0 or (e.type == "text" and not e.has_text):
                continue
            if is_chrome_like(e, ctx.ir, ctx.manifest) or (e.type == "picture" and _bleeds(e)):
                continue
            box = _ink(e) if e.type == "text" else e.bbox
            if box.w <= 0 or box.h <= 0:
                continue
            n_before = len(out)
            under = []
            for o in art:
                if o.bbox.intersection(box) <= 0:
                    continue
                # a calm rectangular panel (or picture) the block stands on entirely is its ground, not art
                if not o.busy and not o.outline and (o.geometry in (None, "rect", "roundRect") or o.type == "picture") and contains(o.bbox, box, 0.0) and o.bbox.intersection(box) >= 0.9 * box.area:
                    continue
                under.append(o)
            if under:
                nx, ny = 16, 8
                hit = 0
                per = {id(o): 0 for o in under}
                for i in range(nx):
                    px = box.x + (i + 0.5) * box.w / nx
                    for j in range(ny):
                        py = box.y + (j + 0.5) * box.h / ny
                        any_hit = False
                        for o in under:
                            if _paints_at(o, px, py, ground, cache):
                                per[id(o)] += 1
                                any_hit = True
                        hit += any_hit
                share = hit / float(nx * ny)
                crossing = [o for o in under if (o.busy or o.outline or o.geometry == "custom") and ART_CROSS[0] <= per[id(o)] / float(nx * ny) <= ART_CROSS[1]]
                what = (e.text[:40] if e.type == "text" else e.name) or e.type
                if e.type == "text" and crossing:
                    o = crossing[0]
                    out.append(ctx.new_issue(CONTENT_OVER_ART, s.index, f"текст «{what}» заходит на {'фотографию' if o.busy else 'рисунок'} шаблона и пересекает его край", severity="error", bboxes=[e.bbox_frac, o.bbox_frac], element_ids=[e.id], details={"share": round(share, 2), "art": o.id}, autofix=fix("rematch", "перевыбрать макет, где контент не ложится на рисунок", outline_id=s.outline_id)))
                    continue
                if ART_WARN[0] <= share <= ART_WARN[1]:
                    o = max(under, key=lambda o: per[id(o)])
                    out.append(ctx.new_issue(CONTENT_OVER_ART, s.index, f"«{what}» лежит на рисунке шаблона ({int(round(share * 100))}% блока)", bboxes=[e.bbox_frac, o.bbox_frac], element_ids=[e.id], details={"share": round(share, 2), "art": o.id}, autofix=fix("rematch", "перевыбрать макет, где контент не ложится на рисунок", outline_id=s.outline_id)))
                    continue
            if e.type == "text":
                r = next((o for o in rules if _rule_through(o, box, H)), None)
                if r is not None:
                    out.append(ctx.new_issue(CONTENT_OVER_ART, s.index, f"линия шаблона проходит через текст «{e.text[:40]}»", bboxes=[e.bbox_frac, r.bbox_frac], element_ids=[e.id], details={"art": r.id, "rule": True}, autofix=fix("rematch", "перевыбрать макет", outline_id=s.outline_id)))
                else:
                    # (line by line: a hairline through a block's last line lies in the block's lower margin)
                    from verstka.audit.checks.common import line_boxes

                    try:
                        rows = [b for b, _ in line_boxes(e)] or [box]
                    except Exception:  # noqa: BLE001
                        rows = [box]
                    # (a line the measure puts outside the block's own box is the estimate's, not the slide's: the
                    # composer's figure «423 млн ₽» on its accent rule, bottom-anchored on one line, was measured as two
                    # and its «first line» crossed the layout's rule above it — qa3 U2)
                    rows = [b for b in rows if b.y >= e.bbox.y - int(0.005 * H)] or [box]
                    # (a plate of the slide itself under the text — the conclusion's strip, a card — hides the layout's
                    # hairline: the slide's shapes are drawn over its layout)
                    plates = [o.bbox for o in s.elements if o.type in ("shape", "text") and o.fill_hex and fill_alpha(o) >= 0.9 and o.bbox.area > 0]

                    def shows(rule: Bbox, row: Bbox) -> bool:
                        x0, x1 = max(rule.x, row.x), min(rule.x2, row.x2)
                        y = rule.y + rule.h // 2
                        return not any(pl.x <= x0 and pl.x2 >= x1 and pl.y <= y <= pl.y2 for pl in plates)

                    hl = next((b for b in hairlines.get(s.index, []) for row in rows if _line_through(b, row) and shows(b, row)), None)
                    if hl is not None:
                        out.append(ctx.new_issue(CONTENT_OVER_ART, s.index, f"линия шаблона проходит через текст «{e.text[:40]}»", bboxes=[e.bbox_frac, hl.to_frac(W, H)], element_ids=[e.id], details={"rule": True, "hairline": True}, autofix=fix("rematch", "перевыбрать макет", outline_id=s.outline_id)))
            if e.type != "text" or len(out) > n_before:
                continue
            # line by line (round 4.1): one line of a block may run onto the art or the photo although the block does not
            what = e.text[:40]
            got = _line_crossings(e, under, ground, cache) if under else None
            if got is not None:
                share, o = got
                out.append(ctx.new_issue(CONTENT_OVER_ART, s.index, f"строка текста «{what}» заходит на {'фотографию' if o.busy else 'рисунок'} шаблона и пересекает его край", severity="error", bboxes=[e.bbox_frac, o.bbox_frac], element_ids=[e.id], details={"share": round(share, 2), "art": o.id, "line": True}, autofix=fix("rematch", "перевыбрать макет, где контент не ложится на рисунок", outline_id=s.outline_id)))
                continue
            got = _photo_crossings(e, s, photos, tpl, W) if photos else None
            if got is not None:
                share, o = got
                out.append(ctx.new_issue(CONTENT_OVER_ART, s.index, f"строка текста «{what}» заходит на пёструю часть фотографии шаблона ({int(round(share * 100))}% строки)", bboxes=[e.bbox_frac, o.bbox_frac], element_ids=[e.id], details={"share": round(share, 2), "art": o.id, "photo": True}, autofix=fix("rematch", "перевыбрать макет, где текст стоит на спокойной части фотографии", outline_id=s.outline_id)))
    return out


def _picture_line_issues(ctx: AuditContext, s, lines: list, W: int, H: int) -> list[Issue]:
    """The texts a ground picture's line runs through (where no plate of the slide hides it) and the cards it runs
    under (a card across a panel's divider or edge reads as laid over the frame, not in it): one issue a slide."""
    from verstka.audit.checks.common import line_boxes

    plates = [o for o in s.elements if o.type in ("shape", "text") and o.fill_hex and fill_alpha(o) >= 0.9 and o.bbox.area > 0]
    for e in s.elements:
        if e.type != "text" or not e.has_text or is_chrome_like(e, ctx.ir, ctx.manifest):
            continue
        try:
            rows = [b for b, _ in line_boxes(e)]
        except Exception:  # noqa: BLE001
            rows = []
        for row in rows:
            for ln in lines:
                if not _line_cuts_row(ln, row):
                    continue
                y = row.y + row.h // 2 if ln.vertical else ln.pos
                x = ln.pos if ln.vertical else row.x + row.w // 2
                if any(pl is not e and pl.bbox.x <= x <= pl.bbox.x2 and pl.bbox.y <= y <= pl.bbox.y2 for pl in plates):
                    continue  # a card of the slide stands over the line there: the text is on the card
                return [ctx.new_issue(CONTENT_OVER_ART, s.index, f"линия рисунка шаблона проходит через текст «{e.text.strip()[:40]}»", bboxes=[e.bbox_frac, ln.box(int(0.002 * W)).to_frac(W, H)], element_ids=[e.id], details={"picture_line": True}, autofix=fix("rematch", "перевыбрать макет, где линии подложки не пересекают контент", outline_id=s.outline_id))]
    for e in plates:
        if e.type != "shape" or e.bbox_frac.area < 0.01 or is_chrome_like(e, ctx.ir, ctx.manifest):
            continue
        hit = next((ln for ln in lines if ln.crosses(e.bbox, inset=0.12, span=0.3)), None)
        if hit is not None:
            return [ctx.new_issue(CONTENT_OVER_ART, s.index, "карточка лежит поперёк линии рисунка шаблона (край или разделитель его панели)", bboxes=[e.bbox_frac, hit.box(int(0.002 * W)).to_frac(W, H)], element_ids=[e.id], details={"picture_line": True}, autofix=fix("rematch", "перевыбрать макет, где линии подложки не пересекают контент", outline_id=s.outline_id))]
    return []


TEXT_OVER_DECOR = CheckSpec(
    id="text_over_decor",
    title="Текст лежит на рисунке подложки (знак, крупная цифра)",
    severity="warn",
    category="layout",
    description="Строка текста слайда лежит на знаке, который рисует картинка под всем слайдом (бледные «01–04» в "
                "карточках VK Tech, значок): под строкой на спокойном листе подложки есть пятно другого цвета. Текст на "
                "фотографии проверяет «Контент лежит на рисунке шаблона».",
)


def _title_home(slide, W: int, H: int) -> Optional[Bbox]:
    """The box the slide's layout (or its master) gives its heading, a hair wider — the place the template draws its
    heading's ground for (a ribbon, a banner baked into the picture under the slide). None when it has none."""
    try:
        for owner in (slide.slide_layout, slide.slide_layout.slide_master):
            for ph in owner.placeholders:
                kind = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0]
                if kind in ("TITLE", "CENTER_TITLE") and None not in (ph.left, ph.top, ph.width, ph.height):
                    pad_x, pad_y = int(0.02 * W), int(0.02 * H)
                    return Bbox(x=int(ph.left) - pad_x, y=int(ph.top) - pad_y, w=int(ph.width) + 2 * pad_x, h=int(ph.height) + 2 * pad_y)
    except Exception:  # noqa: BLE001
        return None
    return None


MARK_IN_LETTERS = 0.04  # the share of a line's letter band a mark must cover to stand under the letters


def _marks_in(slide, row: Bbox, band: Bbox) -> bool:
    """The marks of the ground cover a part of the line's letter band, not only the bounding box of specks around it
    (an old paper's texture and a blot under the descenders united into one box as wide as the line)."""
    from verstka.rendering.picture_lines import mark_share

    try:
        share = mark_share(slide, row, band, sheet_share=0.55)
    except Exception:  # noqa: BLE001 - unread: the box test stands
        return True
    return share is None or share >= MARK_IN_LETTERS


@check(TEXT_OVER_DECOR)
def text_over_decor(ctx: AuditContext) -> list[Issue]:
    from verstka.audit.checks.common import line_boxes
    from verstka.rendering.picture_lines import _grounds, cell_marks, ground_is_photo

    out: list[Issue] = []
    prs = _deck(ctx)
    if prs is None:
        return out
    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    n = len(ctx.ir.slides)
    kinds = {o.id: (o.kind.value if hasattr(o.kind, "value") else str(o.kind)) for o in (ctx.outline.slides if ctx.outline is not None else [])}
    for slide, s in zip(prs.slides, ctx.ir.slides):
        if kinds.get(s.outline_id or "") in ("title", "section", "thanks") or (not s.outline_id and s.index in (1, n)):
            continue
        try:
            if not _grounds(slide) or ground_is_photo(slide):
                continue  # (a photograph under the slide: text over it is `content_over_art`'s — final review C1, lo_candy)
        except Exception:  # noqa: BLE001
            continue
        plates = [o.bbox for o in s.elements if o.type in ("shape", "text", "picture") and (o.fill_hex or o.type == "picture") and fill_alpha(o) >= 0.9 and o.bbox.area > 0]
        # (a mark the template's own sample sets its words over is still a mark: VK Tech's «01»…«04» under its card
        # texts read as noise behind them — qa4 NEW-L3)
        display = 2.0 * ctx.manifest.tokens.typography.size_for("body", 14.0)
        hit = None
        for e in s.elements:
            if e.type != "text" or not e.has_text or is_chrome_like(e, ctx.ir, ctx.manifest):
                continue
            if (e.dominant_size or 0) >= display:
                continue  # a display figure or a heading stands over the ground's art as a large shape of its own
            try:
                rows = [b for b, _ in line_boxes(e)]
            except Exception:  # noqa: BLE001
                continue
            home = _title_home(slide, W, H) if e.ph_type in ("title", "ctrTitle") else None
            boxes, own = [], []
            for r in rows:
                if home is not None and contains(home, r, 0.0):
                    continue  # the heading's letters where the layout sets its heading: the ribbon under them is its band
                if r.w <= 0 or r.h <= 0 or any(p.x <= r.x and p.x2 >= r.x2 and p.y <= r.y and p.y2 >= r.y2 for p in plates if p != e.bbox):
                    continue
                pad = int(0.35 * r.h)
                boxes.append(Bbox(x=r.x - pad, y=r.y - pad, w=r.w + 2 * pad, h=r.h + 2 * pad))
                own.append(r)
            if not boxes:
                continue
            try:
                marks = cell_marks(slide, boxes, sheet_share=0.55)
            except Exception:  # noqa: BLE001
                continue
            for r0, r, m in zip(own, boxes, marks):
                # (the mark under the letters themselves, not beside them: a sheet's edge next to a short figure is none,
                # nor a blot that starts under the line's descender room)
                g0 = letter_band(r0)
                if m is not None and m.w < 0.9 * r.w and m.intersection(g0) >= 0.15 * min(m.area, g0.area) and _marks_in(slide, r, g0):
                    hit = (e, m)
                    break
            if hit:
                break
        if hit:
            e, m = hit
            out.append(ctx.new_issue(TEXT_OVER_DECOR, s.index, f"текст «{e.text.strip()[:40]}» лежит на знаке, который рисует подложка шаблона", bboxes=[e.bbox_frac, m.to_frac(W, H)], element_ids=[e.id], autofix=fix("rematch", "перевыбрать образец без знаков под текстом", outline_id=s.outline_id)))
    return out


LEGEND_OVER_AXIS_TITLE = CheckSpec(
    id="legend_over_axis_title",
    title="Легенда диаграммы наезжает на подпись оси",
    severity="warn",
    category="layout",
    description="Легенда внизу диаграммы не помещается в одну строку ширины диаграммы (значки и названия рядов по "
                "метрикам шрифта) и переносится на вторую, которую под неё не оставили: она ложится на подпись оси X, "
                "поставленную вручную над ней. Считается по разметке самой диаграммы (c:legend, c:title оси, c:manualLayout).",
)


def _chart_legend_overlap(gf_el, part, W: int) -> Optional[str]:
    """Why the bottom legend of a chart frame lies over its X axis's title (its rows by the fonts' metrics against the
    room under the title's manual place), or None."""
    from lxml import etree

    from verstka.rendering.deck import element_bbox

    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    ref = next(iter(gf_el.iter(C + "chart")), None)
    rid = ref.get(R + "id") if ref is not None else None
    if not rid:
        return None
    try:
        cs = etree.fromstring(part.related_part(rid).blob)
    except Exception:  # noqa: BLE001
        return None
    leg = cs.find(f".//{C}legend")
    if leg is None:
        return None
    pos = leg.find(C + "legendPos")
    if (pos.get("val") if pos is not None else "r") != "b":
        return None
    b = element_bbox(gf_el)
    if not b or b[2] <= 0 or b[3] <= 0:
        return None
    fw, fh = b[2] / EMU_PER_PT, b[3] / EMU_PER_PT
    # the X axis's title (a scatter's X is a value axis at the bottom): its manual place
    title = None
    for ax in list(cs.iter(C + "catAx")) + list(cs.iter(C + "dateAx")) + list(cs.iter(C + "valAx")):
        axpos = ax.find(C + "axPos")
        if axpos is None or axpos.get("val") != "b":
            continue
        t = ax.find(C + "title")
        if t is not None:
            title = t
            break
    if title is None:
        return None
    ml = title.find(f"{C}layout/{C}manualLayout")
    ty = ml.find(C + "y") if ml is not None else None
    if ty is None:
        return None
    t_size = 10.0
    rp = next(iter(title.iter(A + "rPr")), None)
    if rp is not None and (rp.get("sz") or "").isdigit():
        t_size = int(rp.get("sz")) / 100
    title_bottom = float(ty.get("val")) * fh + t_size * 1.0  # (its letters: the line's leading under them is no ink)
    # the legend's entries: the series that keep their entry, a marker and a gap with each name
    deleted = {e.find(C + "idx").get("val") for e in leg.findall(C + "legendEntry") if e.find(C + "delete") is not None and e.find(C + "delete").get("val") == "1" and e.find(C + "idx") is not None}
    names = []
    for ser in cs.iter(C + "ser"):
        idx = ser.find(C + "idx")
        if idx is not None and idx.get("val") in deleted:
            continue
        v = next(iter(ser.iter(C + "v")), None)
        if v is not None and (v.text or "").strip():
            names.append(v.text.strip())
    if not names:
        return None
    l_size = 10.0
    drp = next(iter(leg.iter(A + "defRPr")), None)
    if drp is not None and (drp.get("sz") or "").isdigit():
        l_size = int(drp.get("sz")) / 100
    lat = next(iter(leg.iter(A + "latin")), None)
    fam = lat.get("typeface") if lat is not None else None
    lml = leg.find(f"{C}layout/{C}manualLayout")
    # the renderer lays a bottom legend out in columns as wide as its widest entry (its marker, the gap, the name),
    # as many as 80 % of the frame's width holds (qa4 NEW-H2: three entries of a scatter in a 476 pt frame took two
    # rows, «Гуманитарный» on the second, over «Подготовка, ч/нед.»)
    avail = (float(lml.find(C + "w").get("val")) * fw if lml is not None and lml.find(C + "w") is not None else fw * 0.8)
    col = max(text_width_pt(nm, fam, l_size) for nm in names) + l_size * 1.6
    per_row = max(1, int(avail // max(col, 1.0)))
    rows = -(-len(names) // per_row)
    leg_h = rows * l_size * 1.25 + l_size * 0.3
    leg_top = (float(lml.find(C + "y").get("val")) * fh) if lml is not None and lml.find(C + "y") is not None else fh - leg_h
    if leg_top < title_bottom - 1.0:
        return f"легенда в {rows} {'строку' if rows == 1 else ('строки' if rows < 5 else 'строк')} ({leg_h:.0f} пт) ложится на подпись оси X"
    return None


@check(LEGEND_OVER_AXIS_TITLE)
def legend_over_axis_title(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    prs = _deck(ctx)
    if prs is None:
        return out
    for slide, s in zip(prs.slides, ctx.ir.slides):
        for el in slide._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}graphicFrame"):
            try:
                why = _chart_legend_overlap(el, slide.part, int(prs.slide_width))
            except Exception:  # noqa: BLE001
                why = None
            if why:
                e = next((x for x in s.elements if x.type == "chart"), None)
                out.append(ctx.new_issue(LEGEND_OVER_AXIS_TITLE, s.index, why, bboxes=[e.bbox_frac] if e is not None else [], element_ids=[e.id] if e is not None else [], suggestion="Легенде — своя строка под подписью оси или место справа от диаграммы."))
                break
    return out


def _is_full_slide(o) -> bool:
    f = o.bbox_frac
    return max(0.0, min(f.x2, 1.0) - max(f.x, 0.0)) * max(0.0, min(f.y2, 1.0) - max(f.y, 0.0)) >= 0.85


BOTTOM_EMPTY = CheckSpec(
    id="bottom_empty",
    title="Низ слайда пуст: блок не заполняет свою область",
    severity="warn",
    category="layout",
    description="Содержимое слайда (текст по его строкам, карточки, таблица, диаграмма) стоит в верхней части свободной "
    "области под заголовком, а ниже пусто: под блоком остаётся больше 42% высоты области и пустота снизу больше пустоты "
    "сверху на четверть области. Свободная область — от низа заголовка до низа безопасной области шаблона или до верха "
    "рисунка шаблона, который тянется по ширине слайда у его низа (лента, полоса, орнамент). Не проверяются обложка, "
    "разделы, цитаты, финальный слайд и слайд с одним крупным числом: там воздух задуман.",
)
BOTTOM_EMPTY_EXEMPT = ("title", "section", "thanks", "quote", "big_number", "image_text", "mockup")
BOTTOM_EMPTY_SHARE = 0.42  # the empty band under the content, of the free area's height
BOTTOM_EMPTY_LOPSIDED = 0.25  # …and that much more than the air over it (a centred block is balanced, not empty-bottomed)


def _slide_kind(ctx: AuditContext, s) -> Optional[str]:
    if ctx.outline is None or not s.outline_id:
        return None
    o = next((x for x in ctx.outline.slides if x.id == s.outline_id), None)
    if o is None:
        return None
    return o.kind.value if hasattr(o.kind, "value") else str(o.kind)


def _touches_edge(f, tol: float = 0.012) -> bool:
    return f.x <= tol or f.y <= tol or f.x2 >= 1.0 - tol or f.y2 >= 1.0 - tol


def content_span(ctx: AuditContext, s) -> Optional[tuple[int, int, int, int]]:
    """(heading bottom, content top, content bottom, floor) in EMU of a content slide: the heading's foot, the ink of the
    slide's content (text by its lines, cards around it, tables, charts, the slide's own pictures that do not bleed off
    its edge), and the foot of the free area under it — the safe area's foot, or the top of the template's art that runs
    across the slide's foot (a band, a ribbon, an ornament). None when the slide has no content under a heading."""
    from verstka.audit.checks.density import _ink

    W, H = ctx.ir.slide_w, ctx.ir.slide_h
    safe = ctx.manifest.tokens.spacing.safe_area
    title = title_element(s)
    head_bottom = title.bbox.y2 if title is not None and title.bbox_frac.y < 0.4 else int(safe.y * H)
    if title is not None and title.bbox_frac.y < 0.4:
        # a heading's box may run far under its letters: its ink
        head_bottom = min(head_bottom, _ink(title).y2)
    els = [e for e in s.elements if e is not title and e.bbox.area > 0 and not is_chrome_like(e, ctx.ir, ctx.manifest)]
    content = []
    for e in els:
        if e.type == "text":
            if e.has_text and e.bbox.y >= head_bottom - int(0.01 * H):
                content.append(_ink(e))
        elif e.type in ("chart", "table") or e.formula:
            content.append(e.bbox)
        elif e.type == "picture" and not _touches_edge(e.bbox_frac) and e.bbox_frac.area < 0.6:
            content.append(e.bbox)
    if not content:
        return None
    texts = [b for b in content]
    # a card or a plate the content's text stands in counts with its text (the frame is part of the block)
    for e in els:
        if e.type == "shape" and (e.fill_hex or e.line_hex) and e.bbox_frac.area < 0.6 and not _touches_edge(e.bbox_frac):
            b = e.bbox
            if any(b.x <= t.x + t.w // 2 <= b.x2 and b.y <= t.y + t.h // 2 <= b.y2 for t in texts):
                content.append(b)
    top = min(b.y for b in content)
    bottom = max(b.y2 for b in content)
    floor = int(safe.y2 * H)
    # the marks at the slide's foot — a page number, a footer line, a logo: small things in its bottom fifth, as the
    # renderer keeps the content clear of them (`synth._content_bottom_limit`)
    for o in list(s.elements) + list(s.template_elements):
        f = o.bbox_frac
        if f.y >= 0.8 and f.y2 <= 1.01 and f.area < 0.2 and f.h <= 0.15 and o.bbox.y > bottom:
            floor = min(floor, o.bbox.y - int(0.035 * H))
    # the template's art across the slide's foot (its layers, or drawn on the slide itself) ends the free area
    for o in list(s.template_elements) + [e for e in els if e.type in ("picture", "shape") and e.bbox not in content]:
        f = o.bbox_frac
        if f.w < 0.4 or f.area >= 0.85 or o.type not in ("picture", "shape", "group", "other"):
            continue
        if bottom < o.bbox.y < floor and o.bbox.y > head_bottom + int(0.2 * H):
            floor = o.bbox.y
    if floor - head_bottom < int(0.2 * H):
        return None
    return head_bottom, top, bottom, floor


@check(BOTTOM_EMPTY)
def bottom_empty(ctx: AuditContext) -> list[Issue]:
    out: list[Issue] = []
    n = len(ctx.ir.slides)
    for s in ctx.ir.slides:
        kind = _slide_kind(ctx, s)
        if s.index == 1 or kind in BOTTOM_EMPTY_EXEMPT or (kind is None and s.index == n):
            continue
        try:
            span = content_span(ctx, s)
        except Exception:  # noqa: BLE001 - a measure never breaks the audit
            span = None
        if span is None:
            continue
        head, top, bottom, floor = span
        region = floor - head
        below, above = floor - bottom, max(top - head, 0)
        if below > BOTTOM_EMPTY_SHARE * region and below - above > BOTTOM_EMPTY_LOPSIDED * region:
            used = int(round(100 * (bottom - head) / region))
            out.append(ctx.new_issue(BOTTOM_EMPTY, s.index, f"содержимое занимает верхние {used}% свободной области, ниже пусто",
                                     details={"used": round((bottom - head) / region, 2), "empty_below": round(below / region, 2)},
                                     suggestion="Крупнее текст или другой макет: блок должен занимать 55–80% области под заголовком."))
    return out
