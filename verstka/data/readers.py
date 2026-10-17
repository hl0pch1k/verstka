"""Чтение файлов с данными в таблицы `Table` (описание решения §2.1).

Форматы: CSV/TSV (кодировка, разделитель, кавычки, десятичная запятая, пробелы-разделители тысяч), XLSX (все листы,
лист — таблица, даты Excel, проценты по формату ячейки, объединённые ячейки, формулы без сохранённых значений), JSON
(массив записей, `{columns, rows}`, «ключ → записи», «категория → число», колонки-массивы), текст .txt/.md (таблицы
Markdown, блоки через табуляцию, «;», запятую и выровненные пробелами «столбцы»; текст вокруг таблицы — в `notes`).

Разбор чисел, единиц, дат и периодов, пропусков и строки итога — общий с профилем (`verstka/data/parse.py`). Ячейки
типизируются по колонке, а не по одной ячейке: числа — int/float («1 234,5», «46 %», «−12», «1,2 млн ₽»), даты и
периоды — ISO-строки («2025-01-31», «2025-01», «2025-Q4», «2025», «2025-03-01T10:00»), да/нет — bool, пропуски
(«», «—», «н/д») — None. Единица уходит из заголовка в `Table.units` («Выручка, млн ₽» → «Выручка» + «млн ₽»);
одинаковая единица в ячейках («46 %») — тоже. Колонка, где у чисел разные единицы (строка показателей «1,24 млрд ₽»,
«3 480», «356 тыс. ₽»), остаётся строками как в файле: число и единицу каждой ячейки даёт `parse_quantity`.

Шапка — первая «широкая» строка (заголовки и пустые строки над ней уходят в `notes`); двухэтажная шапка (группа над
подзаголовками, строка единиц под названиями) склеивается. Пустые строки и колонки, безымянная колонка-нумерация,
повтор шапки и сноски под таблицей в данные не попадают; строка итога («Итого», «Всего:») остаётся и отмечается в
`Table.totals`. Широкая таблица (периоды в колонках: «Регион · 2023 · 2024 · 2025») разворачивается в длинную
(«Регион · Год · Значение»). Титульный лист книги — не таблица; «выгрузка в Excel», которая на деле текст через
табуляцию (.xls), читается как CSV.

`Table` собирается без повторной проверки pydantic (`model_construct`): типы ячеек гарантирует разбор, а 100 тысяч
строк читаются быстрее секунды.

Всё детерминированно и без модели; ошибки — `DataReadError` с человеческим русским текстом."""

from __future__ import annotations

import csv
import datetime as dt
import functools
import hashlib
import io
import json
import math
import re
import sys
import warnings
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Optional, Union

from verstka.data.parse import (
    BOOL,
    COUNT_UNITS,
    MISSING,
    MISSING_FIRST,
    TIME_HEADER,
    TOTAL_FIRST,
    Quantity,
    dayfirst_hint,
    decimal_hint,
    excel_serial_iso,
    forecast_period,
    header_unit,
    is_total_label,
    iso_of_datetime,
    iso_ord,
    parse_number,
    parse_period,
    parse_quantity,
    period_iso,
    period_str,
    quantity,
    word_unit_column,
    split_unit,
    unit_text,
)
from verstka.data.parse import _CUR_SYMBOLS
from verstka.schemas.data import Cell, SourceKind, Table

__all__ = [
    "DataReadError",
    "parse_number",
    "parse_period",
    "parse_quantity",
    "read_bytes",
    "read_file",
    "split_unit",
]

csv.field_size_limit(min(sys.maxsize, 1 << 30))

# короткие имена внутри модуля (разбор — общий с профилем, verstka/data/parse.py)
_quantity = quantity
_period = period_str
_header_unit = header_unit
_unit_text = unit_text
_iso_of_datetime = iso_of_datetime
_Q = Quantity
_BOOL = BOOL
_TIME_HEADER = TIME_HEADER


class DataReadError(Exception):
    """Файл не читается как таблица. Текст — для человека, по-русски: «Файл пустой», «Не нашёл таблицу в тексте»."""

    def __init__(self, message: str, *, file: Optional[str] = None):
        super().__init__(message)
        self.message = message
        self.file = file

    def __str__(self) -> str:
        return self.message


class _NoData(Exception):
    """Внутреннее: в этом блоке/листе нет строк с данными (сообщение — для единственной таблицы файла)."""


# ======================================================================================================== ячейки и колонки

_MISSING = MISSING
_MISSING_FIRST = MISSING_FIRST
_TOTAL_FIRST = TOTAL_FIRST
_FOOTNOTE = re.compile(r"^(?:\*|примечани|прим\.|источник|сноск|note\b|notes\b|source\b)", re.IGNORECASE)


def _clean_str(s: str) -> str:
    return " ".join(s.split()) if ("  " in s or "\n" in s or "\t" in s or "\r" in s or "\xa0" in s) else s


def _norm_raw(v: Any) -> Any:
    """Сырая ячейка → str (обрезанная) | int | float | bool | None; даты — ISO-строки."""
    if v.__class__ is str:
        s = v.strip()
        if not s or (s[0] in _MISSING_FIRST and len(s) <= 10 and s.lower() in _MISSING):
            return None
        return s
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return _iso_of_datetime(v)
    if isinstance(v, dt.timedelta):
        return str(v)
    return _norm_raw(str(v))


def _norm_row(r: list[Any]) -> list[Any]:
    """`_norm_raw` для строки целиком (быстрый путь для строк: так читается CSV в сотни тысяч строк)."""
    return [
        (s if (s := v.strip()) and not (s[0] in _MISSING_FIRST and len(s) <= 10 and s.lower() in _MISSING) else None)
        if v.__class__ is str else _norm_raw(v)
        for v in r
    ]


def _is_num(v: Any) -> bool:
    return v.__class__ is int or v.__class__ is float or (isinstance(v, (int, float)) and not isinstance(v, bool))


_SERIAL_STR = re.compile(r"\d{5}(?:[.,]\d+)?")
_ISO_DATE_CELL = re.compile(r"\d{4}-\d{2}-\d{2}(?:T\d{2}:\d{2}(?::\d{2})?)?")
# шапка колонки дат: с неё начинается имя («Дата», «Дата договора», «Date»), а не «Заказов за день»
_DATE_HEADER = re.compile(r"^(?:дата|date|день|day|период|period|месяц|month|недел[яи]|week|datetime|timestamp)\b",
                          re.IGNORECASE)


def _excel_serials(vals: list[Any], counts: Counter, nonstr_num: list, name: str,
                   epoch: Optional[dt.datetime]) -> Optional[list[Cell]]:
    """Даты Excel числами (45658 → «2025-01-01»): под шапкой дат («Дата», «Date») или в колонке, где рядом настоящие
    даты (книга Excel, ячейки с форматом «Общий»). Числа — от 20 000 до 60 000 (1954–2064)."""
    date_header = bool(_DATE_HEADER.search(name.strip()))
    if not nonstr_num and not date_header:
        return None
    serial_strs: list[str] = []
    has_dates = False
    for s in counts:
        if _SERIAL_STR.fullmatch(s):
            serial_strs.append(s)
        elif _ISO_DATE_CELL.fullmatch(s):
            has_dates = True
        else:
            return None  # в колонке есть что-то, кроме дат и чисел-дат
    if not nonstr_num and not serial_strs:
        return None
    if not date_header and not (has_dates and not serial_strs):
        return None
    smap = {s: excel_serial_iso(float(s.replace(",", ".")), epoch) for s in serial_strs}
    if any(v is None for v in smap.values()):
        return None
    nmap = {v: excel_serial_iso(v, epoch) for v in set(nonstr_num)}
    if any(v is None for v in nmap.values()):
        return None
    out: list[Cell] = []
    for v in vals:
        if v.__class__ is str:
            out.append(smap.get(v, v))
        elif _is_num(v):
            out.append(nmap[v])
        else:
            out.append(None)
    return _uniform_time(out)


def _uniform_time(cells: list[Cell]) -> list[Cell]:
    """В колонке, где есть время суток («2025-03-01T10:00»), полночь пишется так же («2025-03-01T00:00»)."""
    has_time = any(v.__class__ is str and len(v) > 10 and v[10:11] == "T" for v in cells)
    if not has_time:
        return cells
    return [v + "T00:00" if v.__class__ is str and len(v) == 10 and v[4:5] == "-" and v[7:8] == "-" else v
            for v in cells]


def _type_column(vals: list[Any], name: str, file_decimal: Optional[str], *, english: bool = False,
                 epoch: Optional[dt.datetime] = None, rowwise: bool = False) -> tuple[list[Cell], Optional[str]]:
    """Колонка нормализованных значений (`_norm_raw`) → типизированные ячейки и единица из ячеек (если у всех чисел
    она одна). Решение — по колонке: ≥ 80 % непустых ячеек — числа (с одной единицей) или даты/периоды. english —
    заголовки таблицы по-английски («01/02/2025» — 2 января), epoch — начало отсчёта дат книги Excel. rowwise —
    колонка таблицы показателей по строкам: числа с разными единицами (и без единицы) остаются как в файле."""
    counts: Counter = Counter(v for v in vals if v.__class__ is str)
    time_header = bool(_TIME_HEADER.search(name))
    nonstr_num = [v for v in vals if v.__class__ is int or v.__class__ is float]
    years_only = bool(nonstr_num) and all(float(v).is_integer() and 1900 <= v <= 2100 for v in nonstr_num)
    if nonstr_num or counts:
        serials = _excel_serials(vals, counts, nonstr_num, name, epoch)
        if serials is not None:
            return serials, None
    if not counts:
        # уже типизированные значения (XLSX, JSON): годы в колонке «Год» — периоды
        if time_header and years_only:
            return [str(int(v)) if _is_num(v) else v for v in vals], None
        return vals, None
    distinct = list(counts)
    total = sum(counts.values()) + len(nonstr_num)

    # да/нет
    if not nonstr_num and all(s.lower() in _BOOL for s in distinct):
        return [_BOOL[v.lower()] if v.__class__ is str else v for v in vals], None

    dec = decimal_hint(distinct, file_decimal)
    dayfirst = dayfirst_hint(distinct, english)
    sample = distinct[:300]
    s_total = sum(counts[s] for s in sample) or 1
    q_sample = {s: _quantity(s, dec) for s in sample}
    q_share = sum(counts[s] for s in sample if q_sample[s] is not None) / s_total
    p_sample = {s: _period(s, dayfirst, years=time_header) for s in sample}
    p_share = sum(counts[s] for s in sample if p_sample[s] is not None) / s_total
    period_only = any(p_sample[s] is not None and q_sample[s] is None for s in sample)

    if p_share >= 0.8 and (time_header or period_only) and (not nonstr_num or (time_header and years_only)):
        pmap = {s: (p_sample[s] if s in p_sample else _period(s, dayfirst, years=time_header)) for s in distinct}
        hits = sum(counts[s] for s in distinct if pmap[s] is not None) + len(nonstr_num)
        if hits >= 0.8 * total:
            # «2026 (прогноз)», «2027П» — не пропуск: подпись остаётся, строку-прогноз узнают дальше и не смешивают с фактом
            out = [(pmap[v] if pmap[v] is not None else (_clean_str(v) if forecast_period(v) is not None else None))
                   if v.__class__ is str else (str(int(v)) if _is_num(v) else None) for v in vals]
            return _uniform_time(out), None

    if q_share >= 0.8:
        typed = _numeric_column(vals, distinct, counts, q_sample, dec, total, rowwise)
        if typed is not None:
            return typed
    elif not nonstr_num and not rowwise and len(distinct) <= 5000:
        # «30 задач», «12 задач» — одно слово у всех ячеек: это единица (не «5 класс, 6 класс» — номера)
        wu = word_unit_column(distinct, dec)
        if wu is not None:
            qmap, unit = wu
            out = [(int(qmap[v].value) if qmap[v].is_int else qmap[v].value) if v.__class__ is str else v for v in vals]
            return out, unit
    return [_clean_str(v) if v.__class__ is str else v for v in vals], None


def _numeric_column(vals: list[Any], distinct: list[str], counts: Counter, q_sample: dict, dec: Optional[str],
                    total: int, rowwise: bool = False) -> Optional[tuple[list[Cell], Optional[str]]]:
    qmap: dict[str, Optional[_Q]] = {}
    for s in distinct:
        qmap[s] = q_sample[s] if s in q_sample else _quantity(s, dec)
    nonstr = total - sum(counts.values())
    hits = sum(counts[s] for s in distinct if qmap[s] is not None) + nonstr
    if hits < 0.8 * total:
        return None
    units: Counter = Counter()
    for s in distinct:
        q = qmap[s]
        if q is not None:
            units[(q.scale, q.factor, q.base)] += counts[s]
    plain = (None, 1.0, None)
    if rowwise and (len(units) > 1 or (units and nonstr and plain not in units)):
        return None  # «2 425», «74,1», «87%» — у каждой строки своя единица: как в файле
    keys = [k for k in units if k != plain]
    if len(keys) <= 1:
        if keys and keys[0][0] is not None and (units.get(plain) or nonstr):
            return None  # «1,2 тыс.» рядом с «800» — неясно, в чём «800»
        target = keys[0] if keys else plain
    else:
        bases = {k[2] for k in units}
        if len(bases) != 1 or (None in bases and units.get(plain)):
            return None  # разные единицы (строка показателей «1,24 млрд ₽», «3 480») — оставить как в файле
        # «1,2 млн ₽» и «800 тыс. ₽» — к самому частому множителю
        target = max(keys, key=lambda k: (units[k], k[1]))
    t_scale, t_factor, t_base = target

    def conv(q: _Q) -> Union[int, float]:
        if q.factor == t_factor:
            v = q.value
            return int(v) if q.is_int and abs(v) < 2**53 else v
        v = float(f"{q.value * q.factor / t_factor:.12g}")
        return int(v) if v.is_integer() and abs(v) < 2**53 else v

    cmap = {s: (conv(q) if q is not None else None) for s, q in qmap.items()}
    out: list[Cell] = [cmap[v] if v.__class__ is str else (None if v.__class__ is bool else v) for v in vals]
    return out, _unit_text(t_scale, t_base)


# ======================================================================================================== таблица из строк

_YEARISH = re.compile(r"(?:19|20)\d\d(?:\s*(?:г\.?|год))?")
# год с единицей в шапке: «2024, ГВт·ч», «2025 (тыс. чел.)» — имя колонки-года, а не число «2 024 ГВт·ч»
_YEAR_UNIT = re.compile(r"(?:19|20)\d\d(?:\s*(?:г\.?|год))?\s*(?:,\s*[^\d\s,][^,]{0,24}|\(\s*[^\d()][^()]{0,24}\))")


def _kind(v: Any) -> str:
    """Вид нормализованной ячейки для поиска шапки: empty | num | year | period | text."""
    if v is None:
        return "empty"
    if v.__class__ is bool:
        return "num"
    if _is_num(v):
        return "year" if float(v).is_integer() and 1900 <= v <= 2100 else "num"
    if _YEARISH.fullmatch(v) or _YEAR_UNIT.fullmatch(v.strip()):
        return "year"
    if _quantity(v, strict=True) is not None:
        return "num"
    if _period(v, True, years=False) is not None:
        return "period"
    return "text"


def _header_like(kinds: list[str]) -> bool:
    """Строка похожа на шапку: больше половины заполненных ячеек — названия (годы и периоды — если их ≥ 2: шапка
    «Регион · 2024 · 2025»)."""
    filled = [k for k in kinds if k != "empty"]
    if not filled:
        return False
    text = filled.count("text")
    timey = filled.count("year") + filled.count("period")
    score = text + (timey if timey >= 2 else 0)
    return (text >= 1 or timey >= 2) and score * 2 > len(filled)


def _numbered_header(row: list[Any]) -> bool:
    """Шапка из номеров: название в первой ячейке и номера подряд 1..k (k ≥ 3) — «Месяц · 1 · 2 · 3 · 4 · 5»."""
    filled = [v for v in row if v is not None]
    if len(filled) < 4 or not isinstance(filled[0], str) or _kind(filled[0]) != "text":
        return False
    nums: list[int] = []
    for v in filled[1:]:
        if v.__class__ is int:
            nums.append(v)
        elif isinstance(v, str) and v.isascii() and v.isdigit() and len(v) <= 3:
            nums.append(int(v))
        else:
            return False
    return nums[0] in (0, 1) and nums == list(range(nums[0], nums[0] + len(nums)))


def _filled(row: list[Any]) -> int:
    return sum(1 for v in row if v is not None)


def _row_text(row: list[Any]) -> str:
    return " ".join(str(v) for v in row if v is not None)


def _numbers_below(below: list[list[Any]], j: int) -> bool:
    kinds = [_kind(r[j]) for r in below[:20] if j < len(r) and r[j] is not None]
    return bool(kinds) and sum(k in ("num", "year") for k in kinds) * 2 > len(kinds)


_Header = tuple[list[str], list[Optional[str]], list[str]]  # имена колонок, группы, заголовки как в файле


def _merge_headers(h1: list[Any], h2: list[Any], below: list[list[Any]], width: int) -> Optional[_Header]:
    """Двухэтажная шапка → (имена, группы, заголовки как в файле); None — второй ряд не шапка. Строка единиц под
    названиями («Выручка» / «млн ₽») приклеивается к названиям. Подзаголовки под объединённой группой: названия
    («Форма обучения» над «Бюджет · Платное · Целевое») — имена колонок, группа — отдельно; периоды («Выручка» над
    «2024 · 2025») — «Выручка 2024», группа «Выручка» (так длинная таблица узнаёт ряды)."""
    k2 = [_kind(h2[j]) if j < len(h2) else "empty" for j in range(width)]
    filled2 = [j for j in range(width) if k2[j] != "empty"]
    if not filled2:
        return None
    s1 = [str(h1[j]) if j < len(h1) and h1[j] is not None else "" for j in range(width)]
    s2 = [str(h2[j]) if k2[j] != "empty" else "" for j in range(width)]
    # строка единиц под названиями: «млн ₽», «%», «чел.»
    if (
        all(s1[j] and _header_unit(s2[j]) for j in filled2)
        and any(_numbers_below(below, j) for j in filled2)
    ):
        units_row = [f"{s1[j]}, {s2[j]}" if s2[j] else s1[j] for j in range(width)]
        return units_row, [None] * width, units_row
    if not _header_like(k2):
        return None
    k1 = [_kind(h1[j]) if j < len(h1) else "empty" for j in range(width)]
    # объединённые ячейки группы, заполненные чтением («2025» над каждым кварталом), и годы над периодами — тоже группа
    spans = any(s1[j] and s1[j] == s1[j - 1] and s2[j] for j in range(1, width))
    years_over = any(k1[j] in ("year", "period") and s2[j] for j in range(width))
    if not any(not s1[j] and s2[j] for j in range(width)) and not spans and not years_over:
        return None
    # подзаголовок стоит над числами хотя бы в одной колонке
    if not any(k2[j] in ("text", "year", "period") and _numbers_below(below, j) for j in filled2):
        return None
    names: list[str] = []
    groups: list[Optional[str]] = []
    leaves: dict[int, str] = {}
    group = ""
    for j in range(width):
        if s1[j]:
            group = s1[j]
        elif not s2[j]:
            group = ""
        parent = s1[j] or (group if s2[j] else "")
        if parent and s2[j] and s2[j] == s1[j]:
            names.append(s1[j])  # объединённая по вертикали ячейка шапки («Регион» на два этажа)
            groups.append(None)
        elif parent and s2[j]:
            pn, pu = split_unit(parent)
            cn, cu = split_unit(s2[j])
            unit = cu or pu
            if re.fullmatch(r"(?:19|20)\d\d(?:\s*(?:г\.?|год))?", pn.strip()) and period_str(f"{cn} {pn}") is not None:
                # год над кварталом или месяцем: «I кв. 2025» — период, ось времени, а не «I кв. (2)»
                names.append(f"{cn} {pn.strip()}" + (f", {unit}" if unit else ""))
                groups.append(None)
                continue
            names.append(f"{pn} {cn}" + (f", {unit}" if unit else ""))
            groups.append(pn or None)
            if k2[j] == "text":
                leaves[j] = cn + (f", {unit}" if unit else "")
        else:
            names.append(parent or s2[j])
            groups.append(None)
    # строка единиц под группой: «Выручка» над «млн ₽ | %» — только если все подзаголовки группы — единицы, и не
    # слова-подписи («Мин | Макс», «Год | Месяц» — названия колонок, а не минуты и годы)
    for js in _group_spans(groups):
        units_here = [_leaf_unit(s2[j]) for j in js]
        if len(js) >= 2 and all(units_here):
            for j, u in zip(js, units_here):
                names[j] = f"{split_unit(groups[j])[0]}, {u}"
                leaves.pop(j, None)
    raw = list(names)
    # названия под группой — сами по себе имена колонок, если не повторяются («Бюджет», а не «Форма обучения Бюджет»)
    if leaves and len(leaves) == sum(1 for g in groups if g):
        leaf_names = [split_unit(v)[0] for v in leaves.values()]
        others = {split_unit(names[j])[0] for j in range(width) if j not in leaves}
        # «Цена | Мин · Макс» — «Цена Мин», а не «Мин» (минуты): слово-подпись под группой само именем не бывает
        labelish = any(x.strip().lower().rstrip(".") in _LABEL_WORDS or x.strip().lower() in _LABEL_WORDS
                       for x in leaf_names)
        if len(set(leaf_names)) == len(leaf_names) and not set(leaf_names) & others and not labelish:
            for j, v in leaves.items():
                names[j] = v
    return names, groups, raw


# слова, которые в подзаголовке — названия, а не единицы: «Мин | Макс», «Год | Месяц», «Ч» (часть), «Т» (тип)
_LABEL_WORDS = frozenset({"мин", "макс", "ч", "г", "т", "год", "года", "месяц", "день", "неделя", "лет", "раз", "ед",
                          "ед.", "с", "м", "л", "д", "п", "р", "п.", "р.", "мес", "мес.", "нед", "нед.", "дн", "дн."})


def _leaf_unit(text: str) -> Optional[str]:
    """Подзаголовок — единица: «млн ₽», «%», «чел.», «шт.»; слова-подписи («Мин», «Год») — нет, если не записаны
    символом или с валютой."""
    t = " ".join(text.split())
    unit = _header_unit(t)
    if not unit:
        return None
    if t.lower() in _LABEL_WORDS and not any(ch in t for ch in "%₽$€£¥₸"):
        return None
    return unit


def _group_spans(groups: list[Optional[str]]) -> list[list[int]]:
    """Колонки подряд под одной группой."""
    spans: list[list[int]] = []
    for j, g in enumerate(groups):
        if g and spans and groups[spans[-1][-1]] == g and spans[-1][-1] == j - 1:
            spans[-1].append(j)
        elif g:
            spans.append([j])
    return spans


def _group_row(row: list[Any]) -> bool:
    """Строка групп над шапкой: одни названия или годы («2025 · 2026» над кварталами), и не одно название в первой
    ячейке (это заголовок отчёта)."""
    pos = [j for j, v in enumerate(row) if v is not None]
    if not pos or not all(_kind(row[j]) in ("text", "year", "period") for j in pos):
        return False
    return len(pos) >= 2 or pos[0] >= 1


def _table_id(digest: str, index: int) -> str:
    return hashlib.sha1(f"{digest}:{index}".encode()).hexdigest()[:12]


def _build_table(
    rows: list[list[Any]],
    *,
    file: str,
    kind: SourceKind,
    sheet: Optional[str],
    tid: str,
    notes: str = "",
    header_row: Optional[int] = -1,
    decimal: Optional[str] = None,
    first_name: str = "",
    epoch: Optional[dt.datetime] = None,
    formula_cols: frozenset = frozenset(),
    extra_notes: tuple = (),
) -> Table:
    """Строки сырых ячеек → Table. header_row: -1 — найти шапку, None — шапки нет, k — шапка в строке k.
    first_name — имя безымянной первой колонки (ключи JSON-объекта), иначе «Колонка 1». epoch — начало отсчёта дат
    книги Excel; formula_cols — номера колонок листа с формулами без сохранённых значений (не выбрасываются)."""
    # «—» в ячейке: пропуск, но в таблице счётов с итогом — ноль; где стояли прочерки, помним (по строке-объекту)
    dash_at: dict[int, list[int]] = {}
    normed = []
    for r in rows:
        nr = _norm_row(r)
        if None in nr:
            ds = [j for j, v in enumerate(r) if v.__class__ is str and v.strip() in _DASHES]
            if ds:
                dash_at[id(nr)] = ds
        normed.append(nr)
    rows = normed
    for r in rows:
        while r and r[-1] is None:
            r.pop()
    while rows and not rows[-1]:
        rows.pop()
    lead = 0
    while lead < len(rows) and not rows[lead]:
        lead += 1
    rows = rows[lead:]
    if header_row is not None and header_row >= 0:
        header_row = max(0, header_row - lead)
    if not rows:
        raise _NoData("Файл пустой")

    note_lines: list[str] = []
    scan = rows[:60]
    # пустые ячейки слева от названий — безымянная колонка-индекс («,Выручка»): строка всё равно «широкая»
    leads = [next((k for k, v in enumerate(r) if v is not None), 0) for r in scan]
    edge = min((k for k, r in zip(leads, scan) if r), default=0)
    counts = [_filled(r) + (k - edge if r else 0) for k, r in zip(leads, scan)]
    wide_counts = Counter(c for c in counts if c >= 2)
    width_est = max(wide_counts.items(), key=lambda kv: (kv[1], kv[0]))[0] if wide_counts else 1
    min_fill = 1 if width_est < 2 else max(2, math.ceil(width_est * 0.6))

    header: Optional[list[Any]] = None
    merged_names: Optional[_Header] = None
    title_rows: list[list[Any]] = []
    if header_row == -1:
        # шапка — первая «широкая» строка; название отчёта и пустые строки над ней — заметки
        h = next((i for i, c in enumerate(counts) if c >= min_fill), 0)
        # «Источник:;Росстат» над настоящей шапкой
        while h + 1 < len(scan) and counts[h + 1] >= min_fill:
            first = next((v for v in rows[h] if v is not None), None)
            if isinstance(first, str) and first.endswith(":") and _header_like([_kind(v) for v in rows[h + 1]]):
                h += 1
            else:
                break
        if _header_like([_kind(v) for v in rows[h]]) or _numbered_header(rows[h]):
            header, data = rows[h], rows[h + 1:]
            # группа над шапкой: «;Выручка;;Клиенты» над «Регион;2024;2025;2024;2025»
            if h >= 1 and _group_row(rows[h - 1]):
                width = max([len(header)] + [len(r) for r in data[:200]])
                group = _merge_headers(rows[h - 1], header, data, width)
                if group is not None:
                    merged_names, h = group, h - 1
        else:
            header, data = None, rows[h:]
        title_rows = [r for r in rows[:h] if r]
        note_lines += [_row_text(r) for r in title_rows]
    elif header_row is None:
        header, data = None, rows
    else:
        title_rows = [r for r in rows[:header_row] if r]
        note_lines += [_row_text(r) for r in title_rows]
        header, data = rows[header_row], rows[header_row + 1:]

    width = max([len(header) if header else 0] + [len(r) for r in data])
    raw_names: list[str] = [""] * width
    raw_groups: list[Optional[str]] = [None] * width
    raw_full: Optional[list[str]] = None
    if merged_names is not None:
        raw_names = (merged_names[0] + raw_names)[:width]
        raw_groups = (merged_names[1] + raw_groups)[:width]
        raw_full = (merged_names[2] + [""] * width)[:width]
    elif header is not None:
        raw_names = ([str(v) if v is not None else "" for v in header] + raw_names)[:width]
        # вторая строка шапки: подзаголовки под группой или строка единиц (шапку искали сами — у JSON и
        # Markdown-таблицы с линией под шапкой следующая строка уже данные)
        if data and header_row == -1:
            merged = _merge_headers(header, data[0], data[1:], width)
            if merged is not None:
                data = data[1:]
                raw_names, raw_groups = merged[0], merged[1]
                raw_full = merged[2]

    # пустые строки и повтор шапки (склеенные выгрузки) — прочь; сноски в конце — не данные (строка «Итого» остаётся:
    # её отмечает Table.totals)
    head_cells = [v for v in (header or []) if v is not None]
    head_first = head_cells[0] if head_cells else None
    data = [r for r in data if r and not (
        head_first is not None and head_first in r and [v for v in r if v is not None] == head_cells)]
    tail_notes: list[str] = []
    while data:
        last = data[-1]
        filled = [v for v in last if v is not None]
        first = filled[0] if filled else None
        numbers = any(_is_num(v) or (isinstance(v, str) and _quantity(v) is not None) for v in filled[1:])
        if (
            width >= 2
            and isinstance(first, str)
            and not numbers  # «Источник питания;50», «Notebook;50» — данные, а не сноска
            and (
                _FOOTNOTE.match(first)
                or (len(filled) == 1 and len(first) >= 25 and _quantity(first) is None and _period(first) is None)
            )
        ):
            tail_notes.insert(0, _row_text(last))
            data.pop()
        else:
            break
    if not data:
        raise _NoData("В таблице только заголовок — строк с данными нет" if header is not None else "Файл пустой")

    # строки итога («Итого», «Всего:») — по строгому правилу; их подпись не участвует в выборе типа колонки
    total_labels: dict[int, int] = {}  # строка → колонка подписи
    if width >= 2:
        for i, r in enumerate(data):
            k = 0 if r[0] is not None else next((k for k, v in enumerate(r) if v is not None), None)
            if k is not None and r[k].__class__ is str and r[k][:1] in _TOTAL_FIRST and is_total_label(r[k]):
                total_labels[i] = k
        if len(total_labels) == len(data):
            total_labels = {}

    # колонки
    names_units = [split_unit(n) for n in raw_names]
    english = _english_header(raw_names)
    dash_rows = {i: dash_at[id(r)] for i, r in enumerate(data) if id(r) in dash_at}
    typed_cols: list[list[Cell]] = []
    units: list[Optional[str]] = []
    keep: list[int] = []
    padded = [r if len(r) == width else r + [None] * (width - len(r)) for r in data]
    columns = [list(c) for c in zip(*padded)]
    saved: dict[int, list[tuple[int, Any]]] = {}
    for i, k in total_labels.items():
        saved.setdefault(k, []).append((i, columns[k][i]))
        columns[k][i] = None
    # таблица показателей: каждая строка — свой показатель со своей единицей («Принято, чел.», «Трудоустроено, %»);
    # единица строки не становится единицей колонки
    label_j = _metric_label_column(columns, names_units)
    row_units_all: Optional[list[Optional[str]]] = None
    if label_j is not None:
        row_units_all = _row_units(columns, label_j, names_units)
        # «Принято, чел.» → «Принято» (единица — в row_units); «Молоко 3,2%», «Скидка 10%» — имена, их не трогаем
        columns[label_j] = [_label_unit(v)[0] if v.__class__ is str else v for v in columns[label_j]]
    totals_evident = bool(total_labels) or any(is_total_label(nm) for nm, _ in names_units)
    approx: list[str] = []
    for j in range(width):
        name, unit = names_units[j]
        typed, cell_unit = _type_column(columns[j], name, decimal, english=english, epoch=epoch,
                                        rowwise=label_j is not None and j != label_j)
        if unit is not None and (unit in COUNT_UNITS or _RATE_ONLY.match(unit)) and (
                not _numeric_values(typed) or _NOT_A_MEASURE.search(name)):
            # «Менеджер, клиент», «Номер, заказ» — колонка не числа (или номер): это имя целиком, а не единица
            name, unit = " ".join(raw_names[j].split()), None
            names_units[j] = (name, unit)
        for i, label in saved.get(j, ()):
            if typed is columns[j]:
                typed = list(typed)
            typed[i] = _clean_str(label)
        if _numeric_values(typed) and any(v.__class__ is str and v.lstrip()[:1] in "≈~" for v in columns[j]):
            approx.append(name or f"Колонка {j + 1}")
        if totals_evident and dash_rows and not _NOT_A_COUNT.search(name):
            typed = _dashes_as_zero(typed, j, dash_rows, set(total_labels))
        if all(v is None for v in typed) and not (j in formula_cols and name):
            continue  # пустая колонка (колонку формул без сохранённых значений оставляем — о ней заметка)
        unnamed = not name or name.lower().startswith("unnamed:")
        if unnamed and j == 0 and not first_name and _is_row_number(typed):
            continue  # безымянная нумерация строк (индекс pandas)
        keep.append(j)
        typed_cols.append(typed)
        units.append(unit or cell_unit)
    if not keep:
        raise _NoData("В таблице нет значений")

    # счёт матчей текстом («41-18» в колонке «Мячи»): две колонки чисел — «Забито» и «Пропущено», иначе сравнить
    # команды по забитым нечем
    renamed: dict[int, str] = {}
    for pos in range(len(keep) - 1, -1, -1):
        j = keep[pos]
        vals = [v for v in typed_cols[pos] if v is not None]
        hdr = names_units[j][0] or ""
        halves = [x.strip() for x in re.split(r"\s*/\s*", hdr)] if hdr.count("/") == 1 else []
        if len(vals) >= 2 and len(halves) == 2 and all(re.fullmatch(r"[А-ЯЁа-яёA-Za-z][\w\s.-]*", x) for x in halves) \
                and all(v.__class__ is str and _FRACTION.fullmatch(v) for v in vals):
            # дробь «22/25» под шапкой «Сдали / сдавали» — две колонки чисел: «Сдали» и «Сдавали» (вес среднего)
            pairs = [_FRACTION.fullmatch(v) if v.__class__ is str else None for v in typed_cols[pos]]
            typed_cols[pos] = [int(m_.group(1)) if m_ else None for m_ in pairs]
            typed_cols.insert(pos + 1, [int(m_.group(2)) if m_ else None for m_ in pairs])
            keep.insert(pos + 1, j)
            units.insert(pos + 1, None)
            renamed = {(k + 1 if k > pos else k): v for k, v in renamed.items()}
            renamed[pos], renamed[pos + 1] = halves[0][:1].upper() + halves[0][1:], halves[1][:1].upper() + halves[1][1:]
            continue
        if len(vals) >= 2 and re.search(r"(?i)мяч|гол|шайб", names_units[j][0] or "") and \
                all(v.__class__ is str and _SCORE.fullmatch(v) for v in vals):
            pairs = [_SCORE.fullmatch(v) if v.__class__ is str else None for v in typed_cols[pos]]
            typed_cols[pos] = [int(m_.group(1)) if m_ else None for m_ in pairs]
            typed_cols.insert(pos + 1, [int(m_.group(2)) if m_ else None for m_ in pairs])
            keep.insert(pos + 1, j)
            units.insert(pos + 1, None)
            renamed = {(k + 1 if k > pos else k): v for k, v in renamed.items()}
            renamed[pos], renamed[pos + 1] = "Забито", "Пропущено"

    names: list[str] = []
    for pos, j in enumerate(keep):
        name = renamed.get(pos) or names_units[j][0]
        if not name or name.lower().startswith("unnamed:"):
            name = (first_name if (j == 0 and first_name) else None) or _name_by_values(typed_cols[pos]) \
                or f"Колонка {pos + 1}"
        names.append(name)
    names = _dedupe(names, units)
    shown = raw_full if raw_full is not None else raw_names
    raw_header = [" ".join(shown[j].split()) if not shown[j].lower().startswith("unnamed:") else "" for j in keep]
    groups = [raw_groups[j] for j in keep]
    column_groups = groups if any(groups) else []
    # единица из названия таблицы («Приём … по направлениям, человек»): только колонкам с прямым признаком — слово этой
    # единицы в имени, колонки-периоды одной меры, части и их «Всего», доли с суммой 100 при «%». Номера, коды,
    # возраст, оценки и номера подряд — никогда. Нет признаков — единица остаётся в заметках
    title_unit = _title_unit(title_rows) if label_j is None else None
    if title_unit:
        data_idx = [i for i in range(len(data)) if i not in total_labels]
        cands = [pos for pos in range(len(keep)) if units[pos] is None and _numeric_values(typed_cols[pos])
                 and not _NOT_A_MEASURE.search(names[pos])
                 and not _is_row_number([typed_cols[pos][i] for i in data_idx if typed_cols[pos][i] is not None])]
        for pos in _title_evidence(cands, names, typed_cols, data_idx, title_unit):
            units[pos] = title_unit

    n = len(typed_cols)
    out_rows: list[list[Cell]] = []
    kept_rows: list[int] = []
    for i, r in enumerate(zip(*typed_cols)):
        if r.count(None) != n:
            out_rows.append(list(r))
            kept_rows.append(i)
    if not kept_rows:
        raise _NoData("В таблице нет значений")
    totals = [k for k, i in enumerate(kept_rows) if i in total_labels]
    row_units = [row_units_all[i] for i in kept_rows] if row_units_all is not None and label_j in keep else []
    melt_notes: list[str] = []
    melted = _melt_wide(names, units, raw_header, typed_cols, kept_rows, set(totals), row_units,
                        _title_measure(title_rows))
    if melted is None and column_groups:
        melted = _melt_groups(names, units, raw_header, typed_cols, kept_rows, set(totals), column_groups)
    if melted is not None:
        names, units, raw_header, out_rows, totals, melt_notes = melted
        row_units, column_groups = [], []
    if approx:
        listed = ", ".join(f"«{a}»" for a in approx)
        extra_notes = (*extra_notes, f"Приближённые значения («≈», «~») прочитаны как числа: {listed}")
    all_notes = "\n".join(p for p in [notes.strip(), *note_lines, *tail_notes, *melt_notes, *extra_notes] if p)
    return Table.model_construct(
        id=tid, file=file, kind=kind, sheet=sheet, columns=names, units=units, rows=out_rows,
        raw_header=raw_header, notes=all_notes, totals=totals, row_units=row_units, column_groups=column_groups,
    )


_DASHES = frozenset({"-", "--", "---", "—", "–", "−"})
# прочерк у кода, номера, оценки, возраста, года — «нет», а не ноль
_NOT_A_COUNT = re.compile(r"код|№|(?:^|[^а-яёa-z])(?:id|no|номер)(?:[^а-яёa-z]|$)|оценк|балл|возраст|(?:^|\s)год", re.IGNORECASE)
_RATE_ONLY = re.compile(r"^(?:в|за|на) ")


def _numeric_values(typed: list[Cell]) -> bool:
    """Колонка — числа: есть int/float, а строки — только подписи итога."""
    has = False
    for v in typed:
        if v is None:
            continue
        if v.__class__ is str:
            if not is_total_label(v):
                return False
        elif _is_num(v):
            has = True
        else:
            return False
    return has


# не мера, даже если числа: номер, код, возраст, год, рейтинг, оценка, место в списке
_NOT_A_MEASURE = re.compile(r"(?:^|[^а-яёa-z])(?:id|ид|no|номер|код|артикул|инн|возраст|год|года|рейтинг|оценк|балл|"
                            r"место|ранг|разряд|класс|курс|этаж)|№", re.IGNORECASE)
_PERIODISH = re.compile(r"^(?:q[1-4]|[1-4]\s*кв|[ivx]{1,3}\s*кв|[1-4]-?(?:й|ое)?\s*(?:квартал|полугодие)|"
                        r"янв|фев|мар|апр|ма[йя]|июн|июл|авг|сен|окт|ноя|дек|jan|feb|mar|apr|may|jun|jul|aug|sep|oct|"
                        r"nov|dec|h[12]\b)", re.IGNORECASE)


def _title_evidence(cands: list[int], names: list[str], cols: list[list[Cell]], data_idx: list[int],
                    unit: str) -> set:
    """Колонки, которым единица названия таблицы подходит по прямому признаку (не просто «не противоречит»)."""
    money = any(c in unit for c in "₽$€£¥₸")
    ok = [pos for pos in cands if _takes_title_unit(names[pos], unit)]
    ev = {pos for pos in ok if _names_the_unit(names[pos], unit) or _period_header(names[pos])}

    def vals(pos: int) -> list[float]:
        return [float(cols[pos][i]) if _is_num(cols[pos][i]) else 0.0 for i in data_idx]

    # части и их «Всего»: итог в каждой строке равен сумме остальных — у них одна единица
    for t in ok:
        if not (is_total_label(names[t]) or names[t].strip().lower().startswith(("всего", "итого"))):
            continue
        parts = [pos for pos in ok if pos != t]
        if len(parts) >= 2 and data_idx and all(
                abs(sum(vals(pos)[k] for pos in parts) - vals(t)[k]) <= max(0.01 * abs(vals(t)[k]), 0.5)
                for k in range(len(data_idx))):
            ev |= set(parts) | {t}
    # «%» в названии: колонки-доли с суммой 100
    if unit == "%":
        ev |= {pos for pos in ok if 95 <= sum(vals(pos)) <= 105 and all(0 <= v <= 100 for v in vals(pos))}
    return ev


def _period_header(name: str) -> bool:
    """Имя колонки — период («2024», «I кв. 2025», «Q1», «янв»): колонки-периоды — одна мера в разное время."""
    n = name.strip()
    return bool(_PERIODISH.match(n)) or period_str(n, True, years=True) is not None


def _dashes_as_zero(typed: list[Cell], j: int, dash_rows: dict[int, list[int]], totals: set[int]) -> list[Cell]:
    """В таблице счётов с итогом («Всего», строка «Итого») прочерк в колонке целых неотрицательных чисел — явный
    ноль («Целевое: —»), а не пропуск. «н/д», пусто и «нет данных» остаются пропусками."""
    vals = [v for i, v in enumerate(typed) if v is not None and i not in totals]
    if not vals or not all(v.__class__ is int and v >= 0 for v in vals):
        return typed
    hits = [i for i, cols in dash_rows.items() if j in cols and i < len(typed) and typed[i] is None]
    if not hits:
        return typed
    out = list(typed)
    for i in hits:
        out[i] = 0
    return out


# подписи показателей в первой колонке: «Показатель», «Метрика», «Статья»
_KPI_HEADER_ROWS = re.compile(r"показател|метрик|kpi|индикатор|параметр|стать[яи]|наименовани", re.IGNORECASE)
_CHANGE_COLUMN = re.compile(r"изменени|прирост|(?:^|\s)рост|динамик|отклонени|разниц|δ|г/г|yoy|\bvs\b|change|delta",
                            re.IGNORECASE)


_PAREN_LABEL = re.compile(r"^(?P<name>.*\S)\s*\((?P<unit>[^()]{1,30})\)$")
# имя колонки значений, которому единицы строк противоречат: деньги, счёт, доля
_CONTRADICTS_ROWS = re.compile(r"выручк|доход|прибыл|стоимост|цен[аыу]|сумм|руб|затрат|расход|зарплат|оклад|чек|"
                               r"кол-во|количеств|число|клиент|сделок|заказ|штук|дол[яиюей]|процент", re.IGNORECASE)


def _label_unit(label: str) -> tuple[str, Optional[str]]:
    """Подпись строки → (имя, единица), если единица — после запятой или в скобках, а имя не кончается цифрой:
    «Принято на 1 курс, чел.» → («Принято на 1 курс», «чел.»), «Объём (млн ₽)» → («Объём», «млн ₽»). «Молоко 3,2%»,
    «Скидка 10%», «НДС 20%» — имена: (подпись как есть, None)."""
    text = " ".join(label.split())
    m = _PAREN_LABEL.match(text)
    if m:
        name, unit = m.group("name").strip(), header_unit(m.group("unit"))
    elif "," in text:
        name, _, tail = text.rpartition(",")
        name, unit = name.strip(), header_unit(tail.strip())
    else:
        return label, None
    if not unit or not name or name[-1].isdigit():
        return label, None
    return name, unit


def _metric_label_column(columns: list[list[Any]], names_units: list[tuple[str, Optional[str]]]) -> Optional[int]:
    """Таблица показателей по строкам: первая заполненная колонка — разные подписи (не числа и не даты), у других
    числовых колонок нет своих единиц в шапке, и есть прямой признак единиц у строк — у двух и больше подписей единица
    после запятой или в скобках («Принято, чел.», «Объём НИОКР, млн ₽»), или в ячейках строки одна единица, а у разных
    строк — разные («1,24 млрд ₽» и «3 480»). Тогда единицы — у строк (`Table.row_units`). Слово «Показатель» в шапке
    само по себе не признак. None — обычная таблица."""
    if len(columns) < 2:
        return None
    n = len(columns[0])
    if not 2 <= n <= 30:
        return None
    for j, col in enumerate(columns):
        filled = [v for v in col if v is not None]
        if not filled:
            continue
        if len(filled) < 0.8 * n or not all(v.__class__ is str and _quantity(v) is None and _period(v) is None
                                            for v in filled):
            return None
        if len(set(filled)) != len(filled):
            return None
        values = [k for k in range(len(columns)) if k != j]
        if not any(_is_num(v) or (v.__class__ is str and _quantity(v) is not None) for k in values for v in columns[k]):
            return None
        if any(names_units[k][1] for k in values if any(v is not None for v in columns[k])):
            return None  # у колонок свои единицы («Выручка, тыс. ₽») — единицы не у строк
        if any(_CONTRADICTS_ROWS.search(names_units[k][0]) for k in values):
            return None  # «Цена», «Количество», «Доля» — мера колонки, единица строки («кг») ей противоречит
        with_units = sum(1 for v in filled if _label_unit(v)[1])
        return j if with_units >= 2 or _cells_differ_by_row(columns, j, names_units) else None
    return None


def _row_cell_units(columns: list[list[Any]], label_j: int, names_units: list[tuple[str, Optional[str]]],
                    i: int) -> Optional[set]:
    """Единицы значений строки i (колонки изменений не в счёт): множество `Quantity.unit` (None — без единицы);
    None — в строке есть не число."""
    skip = {j for j, (nm, _) in enumerate(names_units) if _CHANGE_COLUMN.search(nm)}
    seen: set = set()
    for j, col in enumerate(columns):
        v = col[i] if i < len(col) else None
        if j == label_j or j in skip or v is None:
            continue
        if v.__class__ is str:
            q = _quantity(v)
            if q is None:
                return None
            seen.add((q.unit, q.base))
        elif _is_num(v):
            seen.add((None, None))
    return seen


def _cells_differ_by_row(columns: list[list[Any]], label_j: int,
                         names_units: list[tuple[str, Optional[str]]]) -> bool:
    """В каждой строке значения в одной единице, а у строк единицы разные по сути (разные базы: «млрд ₽» и «без
    единицы»; «млн ₽» и «тыс. ₽» — одна база, это не признак)."""
    bases = set()
    for i in range(len(columns[label_j])):
        seen = _row_cell_units(columns, label_j, names_units, i)
        if seen is None or len(seen) > 1:
            return False
        if seen:
            bases.add(next(iter(seen))[1])
    return len(bases) >= 2 and any(b is not None for b in bases)


def _row_units(columns: list[list[Any]], label_j: int, names_units: list[tuple[str, Optional[str]]]) -> list[Optional[str]]:
    """Единица каждой строки-показателя: из подписи после запятой или в скобках («…, чел.» → «чел.»), иначе общая у её
    значений («1,24 млрд ₽» и «1,01 млрд ₽» → «млрд ₽»); колонки изменений («Изменение», «Рост») не в счёт."""
    out: list[Optional[str]] = []
    for i, label in enumerate(columns[label_j]):
        unit = _label_unit(label)[1] if label.__class__ is str else None
        if unit is None:
            seen = _row_cell_units(columns, label_j, names_units, i)
            unit = next(iter(seen))[0] if seen is not None and len(seen) == 1 else None
        out.append(unit)
    return out


# слова, по которым колонка — счёт людей и штук, деньги или доля/рост: единица названия таблицы им может противоречить
_COUNT_NAME = re.compile(r"кол-во|количеств|число|клиент|сделок|заказ|человек|чел\.|студент|сотрудник|штук|шт\.|мест|"
                         r"магазин|филиал|визит|лид|посещ|единиц|площад|объём|объем|вес|длин|расстоян",
                         re.IGNORECASE)
_MONEYISH_NAME = re.compile(r"выручк|доход|прибыл|стоимост|цен[аыу]|сумм|руб|затрат|расход|зарплат|оклад|чек|фонд|"
                            r"преми|стипенди|выплат|оборот|бюджет\s|бюджета|финанс", re.IGNORECASE)
_RATE_NAME = re.compile(r"дол[яиюей]|процент|(?:^|\s)рост|изменени|прирост|темп|конверси|%|индекс|балл|рейтинг|оценк|"
                        r"текучест|конкурс", re.IGNORECASE)
# слово самой единицы в имени колонки: «Численность» при «человек», «Фонд оплаты» при «тыс. руб.»
_PERSON_NAME = re.compile(r"численност|человек|чел\.|сотрудник|студент|работник|персонал|штат|учащ|ученик|участник|"
                          r"посетител|пациент|жител|населени|принят|зачислен|при[её]м|абитуриент|выпускник",
                          re.IGNORECASE)


def _title_unit(title_rows: list[list[Any]]) -> Optional[str]:
    """Единица из строки-названия над шапкой: «Приём на первый курс по направлениям, человек» → «чел.»,
    «Выручка по регионам (млн руб.)» → «млн ₽». Строка-название — одна заполненная ячейка с текстом."""
    for r in title_rows:
        filled = [v for v in r if v is not None]
        if len(filled) == 1 and isinstance(filled[0], str):
            # «Расходы на облако, тыс. ₽ (без НДС)»: уточнение в скобках в конце — не единица, единица — перед ним
            text = re.sub(r"\s*\([^()]*\)\s*$", "", filled[0]) if not split_unit(filled[0])[1] else filled[0]
            unit = split_unit(text)[1]
            if unit:
                return unit
    return None


def _takes_title_unit(name: str, unit: str) -> bool:
    """Единица названия таблицы не противоречит колонке: это не доля/рост/индекс, у денег — не счёт людей и штук,
    у счёта — не деньги."""
    if _RATE_NAME.search(name):
        return False
    money = any(c in unit for c in "₽$€£¥₸")
    return not (_COUNT_NAME.search(name) if money else _MONEYISH_NAME.search(name))


def _names_the_unit(name: str, unit: str) -> bool:
    """В имени колонки — слово самой единицы: деньги — «выручка», «фонд», «сумма»; «чел.» — «численность»,
    «сотрудники»; другие — сама единица («Статей в Scopus» при «статей»)."""
    if _RATE_NAME.search(name):
        return False
    if any(c in unit for c in "₽$€£¥₸"):
        return bool(_MONEYISH_NAME.search(name))
    if unit == "чел.":
        return bool(_PERSON_NAME.search(name))
    word = re.sub(r"[^а-яёa-z]", "", unit.lower())
    return len(word) >= 3 and word in name.lower()


def _dedupe(names: list[str], units: Optional[list[Optional[str]]] = None) -> list[str]:
    """Одинаковые имена колонок: с разными единицами — имя с единицей («Экспорт, т» и «Экспорт, $», «Брак, шт.» и
    «Брак, ppm»), а не «Экспорт (2)»; иначе — номер."""
    count = Counter(names)
    if units is not None:
        names = list(names)
        for nm in [n for n, c in count.items() if c > 1]:
            js = [j for j, n in enumerate(names) if n == nm]
            us = [units[j] if j < len(units) else None for j in js]
            if all(us) and len(set(us)) == len(us):
                for j, u in zip(js, us):
                    names[j] = f"{nm}, {_short_unit(u)}"
    seen: Counter = Counter()
    out = []
    for name in names:
        seen[name] += 1
        out.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return out


def _short_unit(u: Optional[str]) -> str:
    """Единица в имени колонки — без множителя: «тыс. т» → «т», «млн $» → «$»."""
    return re.sub(r"^(?:тыс\.|млн|млрд|трлн)\s*", "", (u or "").strip()) or (u or "")


_LATIN = re.compile(r"[A-Za-z]")
_CYRILLIC = re.compile(r"[А-Яа-яЁё]")


def _english_header(raw_names: list[str]) -> bool:
    """Заголовки таблицы по-английски (есть латиница, нет кириллицы): «01/02/2025» тогда — 2 января."""
    text = " ".join(raw_names)
    return bool(_LATIN.search(text)) and not _CYRILLIC.search(text)


# ── широкая таблица → длинная

_PERIOD_COLUMN = {"year": "Год", "month": "Месяц", "quarter": "Квартал", "week": "Неделя", "day": "Дата",
                  "hour": "Время"}
_KPI_HEADER = re.compile(r"показател|метрик|kpi|индикатор|параметр|стать[яи]|наименовани", re.IGNORECASE)
_DERIVED_COLUMN = re.compile(
    r"изменени|прирост|рост|динамик|отклонени|разниц|итог|всего|сумм|средн|total|change|growth|avg|average|δ|%",
    re.IGNORECASE,
)
_WIDE_NOTE = "Таблица развёрнута: периоды были в колонках"
_SCORE = re.compile(r"\s*(\d{1,4})\s*[-–:]\s*(\d{1,4})\s*")
_FRACTION = re.compile(r"\s*(\d{1,7})\s*/\s*(\d{1,7})\s*")
_ID_HEADER = re.compile(r"(?:^|[^а-яёa-z])(?:№|n|no|id|ид|код|номер|п/п|артикул)(?:[^а-яёa-z]|$)|№", re.IGNORECASE)


def _header_period(name: str) -> Optional[tuple[str, str, int]]:
    """Имя колонки — период («2024», «I кв. 2025», «янв 2025») или «Выручка 2024» → (префикс, зернистость, номер)."""
    if not name or len(name) > 60:
        return None
    words = name.split()
    for k in range(len(words), 0, -1):
        if k > 3 and k != len(words):
            continue
        iso = period_str(" ".join(words[-k:]), True, years=True)
        if iso is None:
            continue
        r = iso_ord(iso)
        if r is None:
            return None
        return " ".join(words[:-k]), r[0], r[1]
    return None


_TITLE_STOP = re.compile(r"\s+(?:по|за|в|во|на|для|от|у|с|и)\s+|[,(:;—–]|\s-\s", re.IGNORECASE)
_TITLE_NOT_A_MEASURE = {"данные", "таблица", "лист", "отчёт", "отчет", "итоги", "список", "выгрузка", "сводка",
                        "структура", "распределение", "состав", "динамика", "показатели", "результаты", "сведения"}


def _title_measure(title_rows: list[list[Any]]) -> Optional[str]:
    """Мера из строки-названия над шапкой: «Выручка по регионам, млн ₽» → «Выручка» — имя значения широкой таблицы
    вместо безликого «Значение». Только 1–3 слова без цифр и не «Отчёт», «Итоги», «Данные»."""
    for r in title_rows:
        filled = [v for v in r if v is not None]
        if len(filled) != 1 or not isinstance(filled[0], str):
            continue
        head = _TITLE_STOP.split(" ".join(filled[0].split()), maxsplit=1)[0].strip()
        words = head.split()
        if 1 <= len(words) <= 3 and not re.search(r"\d", head) and re.fullmatch(r"[А-Яа-яЁё\s-]+", head) \
                and words[0].lower() not in _TITLE_NOT_A_MEASURE:
            return head[0].upper() + head[1:]
    return None


def _melt_wide(names: list[str], units: list[Optional[str]], raw_header: list[str], cols: list[list[Cell]],
               kept_rows: list[int], totals: set[int], row_units: Optional[list[Optional[str]]] = None,
               value_name: Optional[str] = None) -> Optional[tuple]:
    """Периоды в колонках («Регион · 2023 · 2024 · 2025», «Показатель · янв 2025 · фев 2025 · мар 2025») → длинная
    таблица: колонки-подписи, колонка периода («Год», «Месяц», «Квартал»), колонка значения. Если подписи —
    показатели («Показатель»: «Выручка, млн ₽», «Клиенты») или строка одна, каждый показатель становится своей
    колонкой. Нужно ≥ 3 периодов одной зернистости с постоянным шагом и числа под ними; иначе None — таблица
    остаётся как есть."""
    per = [_header_period(n) for n in names]
    pcols = [j for j, p in enumerate(per) if p is not None]
    if len(pcols) < 3:
        return None
    grans = {per[j][1] for j in pcols}
    if len(grans) != 1:
        return None
    gran = grans.pop()
    ords = {j: per[j][2] for j in pcols}
    if gran == "day":  # даты-заголовки 1-го числа подряд по месяцам — месяцы
        days = [dt.date.fromordinal(ords[j]) for j in pcols]
        if all(d.day == 1 for d in days):
            gran = "month"
            ords = {j: d.year * 12 + d.month - 1 for j, d in zip(pcols, days)}
    groups: dict[str, list[int]] = {}
    for j in pcols:
        groups.setdefault(per[j][0], []).append(j)
    seqs = [[ords[j] for j in js] for js in groups.values()]
    first = seqs[0]
    if len(first) < 3 or any(s != first for s in seqs):
        return None
    steps = {b - a for a, b in zip(first, first[1:])}
    if len(steps) != 1 or 0 in steps:
        return None
    n_rows = len(kept_rows)
    texty = False  # числа строками («87%», «2 425»): таблица показателей с единицами по строкам
    for j in pcols:  # под периодами — числа
        for i in range(n_rows):
            v = cols[j][kept_rows[i]]
            if v is None or i in totals or _is_num(v):
                continue
            if v.__class__ is str and _quantity(v) is not None:
                texty = True
                continue
            return None
    ids: list[int] = []
    dropped: list[int] = []
    last_period = max(pcols)
    for j in range(len(names)):
        if per[j] is not None:
            continue
        vals = [cols[j][kept_rows[i]] for i in range(n_rows) if i not in totals and cols[j][kept_rows[i]] is not None]
        numeric = bool(vals) and sum(1 for v in vals if _is_num(v)) * 2 > len(vals)
        if numeric and j < min(pcols) and (_ID_HEADER.search(names[j]) or _is_row_number(vals)):
            ids.append(j)  # «№ п/п», «Код» слева от периодов — подпись строки, а не мера
            continue
        if numeric or is_total_label(names[j]):
            if j > last_period and _DERIVED_COLUMN.search(names[j]):
                dropped.append(j)  # «Итого», «Изменение, %» за периодами — считаются из них
                continue
            return None
        ids.append(j)
    data_rows = [i for i in range(n_rows) if i not in totals]
    if not ids and len(data_rows) != 1:
        return None
    pname = _PERIOD_COLUMN[gran]
    pvalues = [period_iso(gran, o) for o in first]
    notes = [_WIDE_NOTE]
    if dropped:
        listed = ", ".join(f"«{names[j]}»" for j in dropped)
        notes.append(f"Не вошли колонки, которые считаются из периодов: {listed}")
    prefixes = list(groups)

    def group_unit(js: list[int]) -> Optional[str]:
        us = {units[j] for j in js}
        return us.pop() if len(us) == 1 else None

    labels = [cols[ids[0]][kept_rows[i]] for i in data_rows] if len(ids) == 1 else []
    by_metric = (
        prefixes == [""] and len(ids) == 1 and not totals and 1 <= len(labels) <= 12
        and all(isinstance(v, str) for v in labels) and len(set(labels)) == len(labels)
        and bool(_KPI_HEADER.search(names[ids[0]]) or any(split_unit(v)[1] for v in labels) or len(labels) == 1)
    )
    if by_metric:
        common = group_unit(groups[""])
        out_names, out_units, out_raw = [pname], [None], [pname]
        new_cols: list[list[Cell]] = []
        for i, label in zip(data_rows, labels):
            nm, u = split_unit(label)
            vals = [cols[j][kept_rows[i]] for j in groups[""]]
            cell_unit = None
            if any(v.__class__ is str for v in vals):  # строка показателя — колонка: свои единицы в ячейках
                vals, cell_unit = _type_column(vals, nm or label, None)
            out_names.append(nm or label)
            out_units.append(u or (row_units[i] if row_units else None) or cell_unit or common)
            out_raw.append(" ".join(label.split()))
            new_cols.append(vals)
        rows = [[pv] + [c[k] for c in new_cols] for k, pv in enumerate(pvalues)]
        return _dedupe(out_names), out_units, out_raw, rows, [], notes
    if texty:
        return None  # разные единицы под периодами у обычной таблицы — оставить как есть
    vname = value_name if value_name and len(prefixes) == 1 and value_name not in names else "Значение"
    out_names = [names[j] for j in ids] + [pname] + [p or vname for p in prefixes]
    out_units = [units[j] for j in ids] + [None] + [group_unit(groups[p]) for p in prefixes]
    out_raw = [raw_header[j] for j in ids] + [pname] + [p or vname for p in prefixes]
    rows: list[list[Cell]] = []
    new_totals: list[int] = []
    for i in range(n_rows):
        r = kept_rows[i]
        head = [cols[j][r] for j in ids]
        for k, pv in enumerate(pvalues):
            if i in totals:
                new_totals.append(len(rows))
            rows.append(head + [pv] + [cols[groups[p][k]][r] for p in prefixes])
    return _dedupe(out_names), out_units, out_raw, rows, new_totals, notes


def _melt_groups(names: list[str], units: list[Optional[str]], raw_header: list[str], cols: list[list[Cell]],
                 kept_rows: list[int], totals: set[int], groups: list[Optional[str]]) -> Optional[tuple]:
    """Двухуровневая шапка с периодами сверху («2024: Штат · Вакансии | 2025: Штат · Вакансии») → длинная таблица:
    подписи, «Год» (период), по колонке на величину («Штат», «Вакансии»). Годы в шапке — ось времени, а не часть
    имени («2025 Штат»). Нужны ≥ 2 периодов одной зернистости с одинаковым набором величин; другие числовые колонки
    вне групп — таблица остаётся как есть."""
    per: dict[str, tuple[str, int]] = {}
    for g in {g for g in groups if g}:
        iso = period_str(str(g), True, years=True)
        r = iso_ord(iso) if iso else None
        if r is None:
            return None
        per[g] = r
    if len(per) < 2 or len({r[0] for r in per.values()}) != 1:
        return None
    gran = next(iter(per.values()))[0]
    order = sorted(per, key=lambda g: per[g][1])
    subs: dict[str, dict[str, int]] = {g: {} for g in order}
    for j, g in enumerate(groups):
        if not g:
            continue
        nm = names[j]
        sub = nm[len(g):].strip(" ,:-—") if nm.startswith(g) else re.sub(rf"\s*{re.escape(g)}\s*", " ", nm).strip()
        if not sub or sub in subs[g]:
            return None
        subs[g][sub] = j
    keys = list(subs[order[0]])
    if any(list(subs[g]) != keys for g in order):
        return None
    n_rows = len(kept_rows)
    ids = [j for j, g in enumerate(groups) if not g]
    for j in ids:
        vals = [cols[j][kept_rows[i]] for i in range(n_rows) if i not in totals and cols[j][kept_rows[i]] is not None]
        if vals and sum(1 for v in vals if _is_num(v)) * 2 > len(vals):
            return None  # числовая колонка вне групп — не подпись
    if not ids:
        return None
    pname = _PERIOD_COLUMN[gran]
    out_names = [names[j] for j in ids] + [pname] + keys
    out_units = [units[j] for j in ids] + [None] + [
        (lambda us: us.pop() if len(us) == 1 else None)({units[subs[g][k]] for g in order}) for k in keys]
    out_raw = [raw_header[j] for j in ids] + [pname] + keys
    rows: list[list[Cell]] = []
    new_totals: list[int] = []
    for i in range(n_rows):
        r = kept_rows[i]
        head = [cols[j][r] for j in ids]
        for g in order:
            if i in totals:
                new_totals.append(len(rows))
            rows.append(head + [period_iso(gran, per[g][1])] + [cols[subs[g][k]][r] for k in keys])
    return _dedupe(out_names), out_units, out_raw, rows, new_totals, [_WIDE_NOTE]


_ISO_DAY = re.compile(r"\d{4}-\d{2}-\d{2}(?:T.*)?")
_ISO_PERIOD = re.compile(r"\d{4}(?:-\d{2}|-Q[1-4]|-W\d{2})?")


def _name_by_values(vals: list[Cell]) -> Optional[str]:
    """Безымянная колонка дат — «Дата», периодов — «Период» (ось времени без подписи в файле)."""
    strs = [v for v in vals if v is not None]
    if not strs or not all(isinstance(v, str) for v in strs):
        return None
    if all(_ISO_DAY.fullmatch(v) for v in strs):
        return "Дата"
    if all(_ISO_PERIOD.fullmatch(v) for v in strs):
        return "Период"
    return None


def _is_row_number(vals: list[Cell]) -> bool:
    if len(vals) < 2 or not all(v.__class__ is int for v in vals):
        return False
    start = vals[0]
    return start in (0, 1) and all(v == start + i for i, v in enumerate(vals))


# ======================================================================================================== CSV и текст

_RU_WEIGHTS = {"о": 11, "е": 8, "а": 8, "и": 7, "н": 7, "т": 6, "с": 5, "р": 5, "в": 4, "л": 4, "к": 3, "м": 3, "д": 3,
               "п": 3, "у": 3}


def _decode(data: bytes) -> str:
    """Байты → текст: BOM, UTF-8, UTF-16 без BOM, иначе cp1251/KOI8-R/cp866 — по частым русским буквам."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    head = data[:4000]
    if len(head) >= 8:
        # UTF-16 без BOM: старшие байты — 0 (латиница, цифры) или 4 (кириллица)
        odd, even = head[1::2], head[0::2]
        hi_odd = sum(1 for b in odd if b in (0, 4)) > 0.9 * len(odd) and odd.count(0) > 0
        hi_even = sum(1 for b in even if b in (0, 4)) > 0.9 * len(even) and even.count(0) > 0
        if hi_odd != hi_even:
            try:
                return data[: len(data) // 2 * 2].decode("utf-16-le" if hi_odd else "utf-16-be")
            except UnicodeDecodeError:
                pass
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    loose = data.decode("utf-8", errors="replace")
    bad = loose.count("\ufffd")
    good = len(loose) - len(loose.encode("ascii", "ignore")) - bad  # настоящие не-ASCII буквы UTF-8
    if bad <= max(1, len(loose) // 2000) and good >= 4 * bad:
        return loose  # UTF-8 с парой битых байтов (рядом много настоящих русских букв)
    if good <= 0:
        return data.decode(_single_byte_encoding(data), errors="replace")
    # строки по отдельности: шапка в UTF-8, тело в cp1251 (склеенные выгрузки) — каждая строка в своей кодировке
    lines = data.split(b"\n")
    utf8_lines: dict[int, str] = {}
    for k, ln in enumerate(lines):
        if not ln.isascii():
            try:
                utf8_lines[k] = ln.decode("utf-8")
            except UnicodeDecodeError:
                pass
    enc = _single_byte_encoding(b"\n".join(ln for k, ln in enumerate(lines) if k not in utf8_lines))
    if not utf8_lines:
        return data.decode(enc, errors="replace")
    return "\n".join(utf8_lines[k] if k in utf8_lines else ln.decode(enc, errors="replace")
                     for k, ln in enumerate(lines))


def _single_byte_encoding(data: bytes) -> str:
    """Однобайтовая русская кодировка по частым буквам: cp1251, KOI8-R, cp866 или MacCyrillic."""
    best, best_score = "latin-1", -1
    for enc in ("cp1251", "koi8_r", "cp866", "mac_cyrillic"):
        try:
            text = data[:60000].decode(enc)
        except UnicodeDecodeError:
            continue
        score = sum(_RU_WEIGHTS.get(ch, 0) for ch in text[:30000])
        if score > best_score:
            best, best_score = enc, score
    return best


def _looks_binary(text: str) -> bool:
    sample = text[:20000]
    if not sample:
        return False
    bad = sum(1 for ch in sample if (ord(ch) < 32 and ch not in "\t\r\n\f\v") or ch == "\ufffd")
    return bad > len(sample) * 0.02


def _sniff_delimiter(lines: list[str]) -> Optional[str]:
    """Разделитель по первым строкам: «\\t», «;», «,», «|»; None — одна колонка. Кандидаты — те, что делят строки
    согласованно (≥ 80% строк на одно число полей) и без оборванных кавычек; из них — тот, при котором больше полей
    читаются как числа и меньше полей — текст с другим разделителем внутри: при неверном разделителе числа склеиваются
    с текстом («Товар 1;1» | «5;730» в «;»-файле с дробной запятой; «Статья 1,Иванов И.И.» в CSV через запятую с «;»
    в списке авторов). Дальше — доля согласованных строк и порядок (табуляция, «;», запятая, «|»)."""
    sample = [ln for ln in lines if ln.strip()][:200]
    if not sample:
        return None
    scored: list[tuple[float, int, int, str, bool]] = []
    for order, d in enumerate(("\t", ";", ",", "|")):
        if not any(d in ln for ln in sample):
            continue
        rows = list(csv.reader(sample, delimiter=d, skipinitialspace=d in (",", "|")))
        counts = [len(r) for r in rows]
        mode, freq = Counter(c for c in counts).most_common(1)[0]
        if mode < 2:
            continue
        share = freq / len(counts)
        stray = sum(1 for r in rows if any('"' in c for c in r)) > 0.1 * len(rows)
        others = [x for x in ("\t", ";", ",", "|") if x != d]
        numeric = glued = 0
        for r in rows:
            for c in r:
                c = c.strip()
                if not c:
                    continue
                if _quantity(c) is not None:
                    numeric += 1
                elif any(x in c for x in others):
                    glued += 1  # текст с чужим разделителем внутри: «Тариф 1;772», «Статья 1,Иванов И.И.»
        scored.append((share, numeric - glued, -order, d, stray))
    if not scored:
        return None
    steady = [x for x in scored if x[0] >= 0.8 and not x[4]]
    if steady:
        best = max(steady, key=lambda x: (x[1], round(x[0], 1), x[2]))
    else:
        best = max(scored, key=lambda x: (round(x[0], 1), x[2]))
    if best[0] < 0.5:
        return None
    return best[3]


def _csv_rows(name: str, text: str, delim: Optional[str]) -> list[list[str]]:
    if delim is None:
        return [[ln] for ln in text.splitlines()]
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delim, skipinitialspace=delim in (",", "|")))
    except csv.Error as e:
        raise DataReadError("Не получилось разобрать CSV: проверьте кавычки в строках", file=name) from e
    if delim == "|":
        rows = [_strip_pipes(r) for r in rows]
    return rows


def _strip_pipes(r: list[str]) -> list[str]:
    if r and not r[0].strip():
        r = r[1:]
    if r and not r[-1].strip():
        r = r[:-1]
    return r


def _read_delimited(name: str, text: str, kind: SourceKind, digest: str, delim: Optional[str] = "auto") -> list[Table]:
    text = text.replace("\x00", "")
    m = re.match(r"\s*sep=(.)\s*?\r?\n", text, flags=re.IGNORECASE)
    if m:  # «sep=;» — подсказка Excel в первой строке
        delim, text = m.group(1), text[m.end():]
    lines = text.splitlines()
    if delim == "auto":
        delim = "\t" if kind == "tsv" else _sniff_delimiter(lines)
    rows = _csv_rows(name, text, delim)
    file_dec = "." if delim == "," else None  # в CSV через запятую дробь пишут точкой
    chunks = _split_stacked(rows)
    tables: list[Table] = []
    first_error = None
    for c, chunk in enumerate(chunks):
        try:
            tables.append(_build_table(chunk, file=name, kind=kind, sheet=None,
                                       tid=_table_id(digest, 0 if c == 0 else 1000 + c), decimal=file_dec))
        except _NoData as e:
            first_error = first_error or str(e)
    if not tables:
        raise DataReadError(first_error or "Файл пустой", file=name) from None
    if len(tables) > 1:
        for k, t in enumerate(tables):
            t.sheet = f"Таблица {k + 1}"
    return tables


def _blank(r: list[Any]) -> bool:
    return not r or all(v is None or (v.__class__ is str and not v.strip()) for v in r)


def _split_stacked(rows: list[list[Any]]) -> list[list[list[Any]]]:
    """Две таблицы одна под другой через пустую строку (лист Excel, CSV): после пустой строки — строка из одних
    названий (≥ 2), под ней числа, а выше уже была таблица с числами и другой шапкой → вторая таблица. Повтор той же
    шапки (склеенные выгрузки) — не разрез."""
    blank = [_blank(r) for r in rows]
    if not any(blank):
        return [rows]
    chunks: list[list[list[Any]]] = []
    start = 0
    for i in range(1, len(rows)):
        if not blank[i - 1] or blank[i]:
            continue
        head = [v for v in rows[i] if not (v is None or (v.__class__ is str and not v.strip()))]
        if len(head) < 2 or not all(v.__class__ is str and _kind(" ".join(v.split())) == "text" for v in head):
            continue
        before = [r for r in rows[start:i] if not _blank(r)]
        if len(before) < 2 or not any(_numeric_cell(v) for r in before for v in r):
            continue
        after = [r for r in rows[i + 1:i + 4] if not _blank(r)]
        if not any(_numeric_cell(v) for r in after for v in r):
            continue
        first_head = next(([" ".join(str(v).split()) for v in r if v is not None and str(v).strip()]
                           for r in before if sum(1 for v in r if v is not None and str(v).strip()) >= 2), None)
        if first_head == [" ".join(v.split()) for v in head]:
            continue
        chunks.append(rows[start:i])
        start = i
    chunks.append(rows[start:])
    return chunks


def _numeric_cell(v: Any) -> bool:
    return _is_num(v) or (v.__class__ is str and bool(v.strip()) and _quantity(v.strip()) is not None)


# --- таблицы внутри текста

_BOX_TRANS = str.maketrans({"│": "|", "┃": "|", "║": "|"})
_BORDER_CHARS = set(" \t-=─━═┌┐└┘├┤┬┴┼╔╗╚╝╠╣╦╩╬╒╕╘╛╞╡╤╧╪+|:")
_PIPE_SPLIT = re.compile(r"(?<!\\)\|")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def _is_border(line: str) -> bool:
    """Линия таблицы: «|---|:--:|», «--- | ---», «+----+», «├───┼───┤»."""
    s = line.strip()
    return bool(s) and set(s) <= _BORDER_CHARS and sum(s.count(c) for c in "-=─━═") >= 3


def _pipe_cells(line: str) -> list[str]:
    s = line.strip()
    cells = _PIPE_SPLIT.split(s)
    if s.startswith("|"):
        cells = cells[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        cells = cells[:-1]
    return [re.sub(r"(\*\*|__|`)", "", c.replace("\\|", "|")).strip() for c in cells]


def _is_pipe_line(line: str) -> bool:
    s = line.strip()
    return bool(_PIPE_SPLIT.search(s)) and len(s.strip("| ")) > 0


def _pipe_block(lines: list[str], i: int) -> Optional[tuple[int, list[list[str]], Optional[int]]]:
    if not (_is_pipe_line(lines[i]) or _is_border(lines[i])):
        return None
    j = i
    rows: list[list[str]] = []
    header_row: Optional[int] = -1
    has_border = False
    while j < len(lines) and lines[j].strip() and (_is_pipe_line(lines[j]) or _is_border(lines[j])):
        if _is_border(lines[j]):
            has_border = True
            if len(rows) == 1 and header_row == -1:
                header_row = 0
        else:
            cells = _pipe_cells(lines[j])
            if len(cells) >= 2:
                rows.append(cells)
        j += 1
    single_pipe = all(len(_PIPE_SPLIT.findall(lines[k].strip())) <= 1 for k in range(i, j) if not _is_border(lines[k]))
    if len(rows) < 2 or (single_pipe and not has_border):
        return None
    return j, rows, header_row


def _tab_block(lines: list[str], i: int) -> Optional[tuple[int, list[list[str]], Optional[int]]]:
    if "\t" not in lines[i] or len([c for c in lines[i].split("\t") if c.strip()]) < 2:
        return None
    j = i
    while j < len(lines) and lines[j].strip() and "\t" in lines[j]:
        j += 1
    rows = [[c.strip() for c in ln.split("\t")] for ln in lines[i:j]]
    if len(rows) < 2:
        return None
    return j, rows, -1


def _semicolon_block(lines: list[str], i: int) -> Optional[tuple[int, list[list[str]], Optional[int]]]:
    first = lines[i]
    if ";" not in first:
        return None
    n = len(next(csv.reader([first], delimiter=";")))
    if n < 2:
        return None
    j = i
    rows = []
    while j < len(lines) and lines[j].strip():
        r = next(csv.reader([lines[j]], delimiter=";"))
        if len(r) != n or any(len(c) > 80 for c in r):
            break
        rows.append([c.strip() for c in r])
        j += 1
    if len(rows) < 2 or not any(_quantity(c.strip()) is not None for r in rows[1:] for c in r if c.strip()):
        return None
    return j, rows, -1


def _comma_block(lines: list[str], i: int) -> Optional[tuple[int, list[list[str]], Optional[int]]]:
    """CSV через запятую внутри письма или в блоке ```csv: одинаковое число полей (≥ 2) в строках подряд, короткие
    ячейки без выравнивания пробелами, под шапкой — числа."""
    first = lines[i]
    if "," not in first or "  " in first.strip() or "\t" in first or _BULLET.match(first):
        return None
    try:
        head = next(csv.reader([first], skipinitialspace=True))
    except csv.Error:
        return None
    n = len(head)
    if n < 2 or any(_prose_cell(c) for c in head):
        return None
    j = i
    rows = []
    while j < len(lines) and lines[j].strip():
        try:
            r = next(csv.reader([lines[j]], skipinitialspace=True))
        except csv.Error:
            break
        if len(r) != n or any(len(c) > 40 or "  " in c.strip() for c in r):
            break
        rows.append([c.strip() for c in r])
        j += 1
    if len(rows) < 2 or any(not c for c in rows[0]):
        return None
    # «Выручка 1,5 млрд» — дробная запятая в тексте, а не разделитель: «Выручка 1» + «5 млрд»
    raw_pairs = [re.split(r",(?=\S)", lines[k]) for k in range(i, j)]
    glued = sum(1 for parts in raw_pairs[1:] if any(
        _GLUED_TAIL.search(a.strip()) and b[:1].isdigit() and _period(a.strip()) is None
        for a, b in zip(parts, parts[1:])))
    if glued * 2 >= len(rows) - 1:
        return None
    numeric_rows = sum(1 for r in rows[1:] if any(_quantity(c) is not None or _period(c) is not None for c in r if c))
    if not any(_quantity(c) is not None for r in rows[1:] for c in r if c) or numeric_rows * 2 < len(rows) - 1:
        return None
    return j, rows, -1


_BULLET = re.compile(r"^\s*(?:[-–—•*·]\s|\d{1,2}[.)]\s)")
_GLUED_TAIL = re.compile(r"[^\W\d_][^,]*\s\d{1,3}$")  # «Выручка 1» перед «5 млрд»


def _prose_cell(cell: str) -> bool:
    """Ячейка шапки похожа на прозу, а не на название колонки: «коллеги!», «прошу посмотреть.», «Выручка по месяцам:
    100», четыре слова и больше со знаком предложения. «Среднее время ответа клиенту» — название."""
    c = cell.strip()
    if not c:
        return False
    if c[-1] in "!?" or ":" in c:
        return True
    words = c.split()
    if len(words) >= 4 and re.search(r"[.!?…]", c):  # длинное название колонки без знаков предложения — не проза
        return True
    return c.endswith(".") and len(words[-1].rstrip(".")) > 4  # «команда.», «посмотреть.» — не «руб.», «чел.»


_SPACED = re.compile(r"\s{2,}|\t")


def _spaced_cells(line: str) -> list[str]:
    s = line.rstrip()
    cells = _SPACED.split(s.strip())
    if s[:2] == "  " and cells:
        cells = [""] + cells  # пустая левая верхняя ячейка
    return [c.strip() for c in cells]


def _spaced_block(lines: list[str], i: int) -> Optional[tuple[int, list[list[str]], Optional[int]]]:
    first = _spaced_cells(lines[i])
    if len([c for c in first if c]) < 2 or len(lines[i]) > 300:
        return None
    n = len(first)
    j = i + 1
    rows = [first]
    while j < len(lines) and lines[j].strip() and len(lines[j]) <= 300:
        cells = _spaced_cells(lines[j])
        if len(cells) != n or any(len(c) > 60 for c in cells):
            break
        rows.append(cells)
        j += 1
    if len(rows) < 2 or any(len(c) > 60 for c in first):
        return None
    numeric_rows = sum(1 for r in rows[1:] if any(_quantity(c) is not None for c in r if c))
    if numeric_rows * 2 < len(rows) - 1:
        return None
    return j, rows, -1


def _text_blocks(lines: list[str]) -> list[tuple[int, int, list[list[str]], Optional[int], str]]:
    blocks = []
    i = 0
    while i < len(lines):
        if not lines[i].strip() or _MD_FENCE.match(lines[i]):
            i += 1
            continue
        for how, finder in (("pipe", _pipe_block), ("tab", _tab_block), ("semicolon", _semicolon_block),
                            ("comma", _comma_block), ("spaced", _spaced_block)):
            found = finder(lines, i)
            if found is not None:
                end, rows, header_row = found
                blocks.append((i, end, rows, header_row, how))
                i = end
                break
        else:
            i += 1
    return blocks


def _read_text(name: str, text: str, digest: str) -> list[Table]:
    text = text.replace("\x00", "").translate(_BOX_TRANS)
    lines = text.splitlines()
    blocks = _text_blocks(lines)
    nonempty = [k for k, ln in enumerate(lines) if ln.strip() and not _MD_FENCE.match(ln)]
    if not blocks:
        # весь файл — CSV через запятую с расширением .txt
        delim = _sniff_delimiter(lines)
        if delim in (",", ";", "\t") and len(nonempty) >= 2:
            parsed = list(csv.reader([lines[k] for k in nonempty], delimiter=delim, skipinitialspace=True))
            counts = [len(r) for r in parsed]
            mode, freq = Counter(counts).most_common(1)[0]
            # проза с запятыми («Цели: …, …», «Привет, коллеги, как дела?») — не таблица
            prose = delim == "," and any(_prose_cell(c) for c in parsed[0])
            if mode >= 2 and freq >= 0.9 * len(counts) and not prose:
                return _read_delimited(name, text, "tsv" if delim == "\t" else "csv", digest, delim)
        raise DataReadError("Не нашёл таблицу в тексте", file=name)
    # весь файл — одна таблица через табуляцию, «;» или запятую: это CSV/TSV, а не «таблица в тексте»
    if len(blocks) == 1 and blocks[0][4] in ("tab", "semicolon", "comma"):
        start, end = blocks[0][0], blocks[0][1]
        if all(start <= k < end for k in nonempty):
            how = blocks[0][4]
            body = "\n".join(ln for ln in lines if not _MD_FENCE.match(ln))
            return _read_delimited(name, body, "tsv" if how == "tab" else "csv", digest,
                                   {"tab": "\t", "semicolon": ";", "comma": ","}[how])
    covered = set()
    for start, end, *_ in blocks:
        covered.update(range(start, end))
    outside = [ln.rstrip() for k, ln in enumerate(lines) if k not in covered and not _MD_FENCE.match(ln)]
    notes = re.sub(r"\n{3,}", "\n\n", "\n".join(outside)).strip()
    tables: list[Table] = []
    titles: list[Optional[str]] = []
    for k, (_start, _end, rows, header_row, _how) in enumerate(blocks):
        # строка-название прямо над таблицей («Выручка по продуктам, млн ₽:») — её единица и имя, как у строки-названия
        # над шапкой в Excel; иначе единица теряется в заметках
        title = _text_title(lines, _start, covered)
        if title is not None:
            width = max((len(r) for r in rows), default=1)
            rows = [[title] + [None] * (width - 1)] + rows
            header_row = header_row + 1 if header_row is not None and header_row >= 0 else header_row
        try:
            tables.append(_build_table(
                rows, file=name, kind="text", sheet=None, tid=_table_id(digest, k), notes=notes,
                header_row=header_row, decimal=None,
            ))
            titles.append(title)
        except _NoData:
            continue  # блок без значений — не таблица
    if not tables:
        raise DataReadError("Не нашёл таблицу в тексте", file=name)
    if len(tables) > 1:
        for k, t in enumerate(tables):
            nm = split_unit(titles[k])[0].strip() if titles[k] else ""
            t.sheet = nm if nm and len(nm) <= 60 and nm not in [x.sheet for x in tables[:k]] else f"Таблица {k + 1}"
    return tables


def _text_title(lines: list[str], start: int, covered: set[int]) -> Optional[str]:
    """Название таблицы в тексте: короткая строка прямо над ней (через одну пустую), с двоеточием в конце или с
    единицей («Выручка по продуктам, млн ₽:», «Численность команды, чел.»). Абзац прозы — не название."""
    k = start - 1
    while k >= 0 and k not in covered and not lines[k].strip() and start - k <= 2:
        k -= 1
    if k < 0 or k in covered or start - k > 2:
        return None
    ln = " ".join(lines[k].strip().strip("#*_ ").split())
    if not ln or len(ln) > 160 or not re.search(r"[А-Яа-яЁёA-Za-z]", ln) or _MD_FENCE.match(lines[k]):
        return None
    body = ln.rstrip(":").strip()
    if (ln.endswith(":") or split_unit(body)[1]) and body.count(".") <= 2 and len(ln) <= 100:
        return body
    # фраза с единицей счёта: «Отчёт продуктового отдела, тыс. абонентов на конец квартала.» → «Абоненты, тыс.
    # абонентов»: имя величины и её единица, а не безликое «Значение»
    m = re.search(r"(?:^|[\s,(])((тыс\.|млн|млрд)\s+([а-яё]{4,}))", ln)
    if m:
        noun = phrase_nominative(m.group(3))
        return f"{noun}, {m.group(2)}"  # «Абоненты, тыс.»: множитель — единица, слово — имя величины
    return None


def phrase_nominative(gen_pl: str) -> str:
    """Именительный по родительному множественного счёта: «абонентов» → «Абоненты», «клиентов» → «Клиенты»,
    «пользователей» → «Пользователи»; иначе — как есть, с заглавной."""
    w = gen_pl.lower()
    for end, new in (("ов", "ы"), ("ев", "и"), ("ей", "и")):
        if w.endswith(end) and len(w) > len(end) + 2:
            nom = w[: -len(end)] + new
            if new == "ы" and nom[-2:-1] in "гкхжшчщ":
                nom = nom[:-1] + "и"
            return nom[:1].upper() + nom[1:]
    return w[:1].upper() + w[1:]


# ======================================================================================================== XLSX

@functools.lru_cache(maxsize=256)
def _format_unit(fmt: str) -> Optional[str]:
    """Числовой формат ячейки Excel → «%» (значение — доля), символ валюты или None («0.0%», «#,##0 ₽», «[$₽-419]»)."""
    f = re.sub(r"\[\$-[0-9A-Fa-f]+\]", "", fmt)  # «[$-409]» — язык, не валюта
    if "%" in f.replace('"%"', ""):
        return "%"
    cur = next((c for c in _CUR_SYMBOLS if c in f), None)
    if cur is None and ("р." in f or "руб" in f):
        cur = "₽"
    return cur


def _xlsx_cell(cell: Any) -> Any:
    v = getattr(cell, "value", None)
    if v is None:
        return None
    if v.__class__ is int or v.__class__ is float:
        fmt = getattr(cell, "number_format", None)
        if not fmt or fmt == "General":
            return v
        unit = _format_unit(fmt)
        if unit == "%":
            return f"{v * 100:.10g}%"
        if unit is not None:
            return f"{v:.15g} {unit}"
        return v
    if isinstance(v, str):
        if v.startswith("#") and v.lower() in _MISSING:
            return None
        return v
    return v


_MERGE_REF = re.compile(rb'<(?:\w+:)?mergeCell\b[^>]*?\bref="\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)"')
_CELL_REF = re.compile(rb'\br="([A-Z]{1,3})\d+"')
_CELL_VALUE = re.compile(rb"<(?:\w+:)?v>([^<]*)</(?:\w+:)?v>")
_FORMULA_TAG = re.compile(rb"<(?:\w+:)?f[ >/]")
FORMULA_NOTE = "В книге есть формулы без сохранённых значений — откройте и сохраните её в Excel"


def _col_index(letters: bytes) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ch - 64)
    return n - 1


_HIDDEN_COL = re.compile(rb'<(?:\w+:)?col\b[^>]*?\bmin="(\d+)"[^>]*?\bmax="(\d+)"[^>]*?\bhidden="(?:1|true)"')
_HIDDEN_COL2 = re.compile(rb'<(?:\w+:)?col\b[^>]*?\bhidden="(?:1|true)"[^>]*?\bmin="(\d+)"[^>]*?\bmax="(\d+)"')
_HIDDEN_ROW = re.compile(rb'<(?:\w+:)?row\b[^>]*?\br="(\d+)"[^>]*?\bhidden="(?:1|true)"')
_HIDDEN_ROW2 = re.compile(rb'<(?:\w+:)?row\b[^>]*?\bhidden="(?:1|true)"[^>]*?\br="(\d+)"')


def _drop_hidden(rows: list[list[Any]], hidden: dict) -> str:
    """Скрытые колонки и строки листа — как скрытые листы: не читаются, а в заметках сказано сколько."""
    cols, rws = hidden.get("cols", set()), hidden.get("rows", set())
    for i in rws:
        if i < len(rows):
            rows[i] = []
    if cols:
        for r in rows:
            for j in cols:
                if j < len(r):
                    r[j] = None
    if not cols and not rws:
        return ""
    parts = [f"колонок — {len(cols)}" if cols else "", f"строк — {len(rws)}" if rws else ""]
    return "Скрытые колонки и строки листа не прочитаны, как и скрытые листы: " + ", ".join(p for p in parts if p)


def _sheet_extras(wb: Any, ws: Any, hidden: Optional[dict] = None) -> tuple[list[tuple[int, int, int, int]], set[int]]:
    """Из XML листа: объединённые диапазоны (строка₁, колонка₁, строка₂, колонка₂ — с нуля) и номера колонок, где есть
    формулы без сохранённого значения (книгу записала программа, а не Excel: openpyxl, pandas)."""
    path = getattr(ws, "_worksheet_path", None)
    archive = getattr(wb, "_archive", None)
    if not path or archive is None:
        return [], set()
    try:
        xml = archive.read(path)
    except Exception:
        return [], set()
    merges = []
    if b"mergeCell" in xml:
        for m in _MERGE_REF.finditer(xml):
            merges.append((int(m.group(2)) - 1, _col_index(m.group(1)), int(m.group(4)) - 1, _col_index(m.group(3))))
    if hidden is not None and b'hidden="' in xml:
        cols: set[int] = set()
        for rx in (_HIDDEN_COL, _HIDDEN_COL2):
            for m in rx.finditer(xml):
                lo, hi = int(m.group(1)), int(m.group(2))
                if hi - lo < 1000:
                    cols.update(range(lo - 1, hi))
        rws = {int(m.group(1)) - 1 for rx in (_HIDDEN_ROW, _HIDDEN_ROW2) for m in rx.finditer(xml)}
        hidden["cols"], hidden["rows"] = cols, rws
    formulas: set[int] = set()
    for m in _FORMULA_TAG.finditer(xml):
        pos = m.start()
        lo = max(0, pos - 4000)  # начало ячейки — рядом с <f>: поиск назад не дальше 4 КБ (иначе квадратично)
        start = max(xml.rfind(b"<c ", lo, pos), xml.rfind(b":c ", lo, pos))
        hi = pos + 20000  # и вперёд — в пределах одной ячейки
        end = xml.find(b"</c>", pos, hi)
        if end < 0:
            end = xml.find(b"c>", pos, hi)
        if start < 0 or end < 0:
            continue
        cell = xml[start:end]
        head = cell[:cell.find(b">") + 1]
        if b't="str"' in head or b't="e"' in head:
            continue  # формула с текстом или ошибкой: значение сохранено (может быть пустой строкой)
        v = _CELL_VALUE.search(cell)
        if v is not None and v.group(1).strip():
            continue
        ref = _CELL_REF.search(head)
        if ref is not None:
            formulas.add(_col_index(ref.group(1)))
    return merges, formulas


def _fill_merged(rows: list[list[Any]], merges: list[tuple[int, int, int, int]]) -> None:
    """Объединённые по вертикали ячейки с текстом («Центр» на строки 2–4) → значение верхней ячейки в каждой строке
    диапазона (числа не размножаются).
    Объединения в одну строку (заголовок отчёта на всю ширину, группа «Выручка» над «2024 · 2025») не размножаются:
    шапку из них собирает разбор двухэтажной шапки, а заголовок отчёта уходит в заметки."""
    for r1, c1, r2, c2 in merges:
        if r2 <= r1 or c2 != c1 or r1 >= len(rows) or c1 >= len(rows[r1]):
            continue
        top = rows[r1][c1]
        if not isinstance(top, str) or not top.strip():
            continue  # числа не размножаем: «План региона» 100 на три города — не 100 у каждого
        for r in range(r1 + 1, min(r2, len(rows) - 1) + 1):
            row = rows[r]
            if len(row) <= c1:
                row.extend([None] * (c1 + 1 - len(row)))
            if row[c1] is None:
                row[c1] = top


# книга Excel: строк листа читается не больше этого (дальше — «прочитаны первые …», как у API), а листы в распакованном
# виде больше предела не открываются вовсе — сжатый XML может развернуться в гигабайты (I6)
XLSX_MAX_ROWS = 200_000
XLSX_MAX_XML_BYTES = 300 * 1024 * 1024


def _xlsx_too_big(data: bytes) -> bool:
    """Листы и общие строки книги в распакованном виде больше `XLSX_MAX_XML_BYTES` (по оглавлению ZIP, не распаковывая)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            total = sum(i.file_size for i in z.infolist()
                        if i.filename.startswith("xl/worksheets/") or i.filename == "xl/sharedStrings.xml")
    except (zipfile.BadZipFile, OSError, ValueError):
        return False
    return total > XLSX_MAX_XML_BYTES


def _read_xlsx(name: str, data: bytes, digest: str) -> list[Table]:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise DataReadError("Для чтения Excel нужен пакет openpyxl", file=name) from e
    if data[:8] == _OLE:
        raise DataReadError(
            "Книга защищена паролем или сохранена в старом формате .xls — сохраните её как .xlsx без пароля или как CSV",
            file=name,
        )
    if _xlsx_too_big(data):
        raise DataReadError(
            f"Книга Excel слишком большая: листы в распакованном виде больше {XLSX_MAX_XML_BYTES // (1024 * 1024)} МБ — "
            "сохраните нужный лист как CSV или уменьшите книгу", file=name)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise DataReadError("Не получилось открыть книгу Excel: файл повреждён или это не .xlsx", file=name) from e
    tables: list[Table] = []
    visible = 0
    first_error = None
    try:
        sheets = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
        for k, ws in enumerate(sheets):
            if not hasattr(ws, "iter_rows"):
                continue
            visible += 1
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            hidden: dict = {}
            merges, formula_cols = _sheet_extras(wb, ws, hidden)
            rows: list[list[Any]] = []
            empty_run = 0
            cut = False
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                for row in ws.iter_rows():
                    if len(rows) > XLSX_MAX_ROWS:
                        cut = True  # строка сверх предела прочитана — значит, лист длиннее: дальше не читаем
                        break
                    vals = [_xlsx_cell(c) for c in row]
                    while vals and (vals[-1] is None or (isinstance(vals[-1], str) and not vals[-1].strip())):
                        vals.pop()
                    if not vals:
                        empty_run += 1
                        if empty_run > 10000:
                            break
                    else:
                        empty_run = 0
                    rows.append(vals)
            if cut:
                rows = rows[:XLSX_MAX_ROWS]
            _fill_merged(rows, merges)
            notes_here = [FORMULA_NOTE] if formula_cols else []
            if cut:
                notes_here.append(f"Лист длиннее {XLSX_MAX_ROWS:,} строк — прочитаны первые {XLSX_MAX_ROWS:,}".replace(",", "\u00a0"))
            hidden_note = _drop_hidden(rows, hidden)
            if hidden_note:
                notes_here.append(hidden_note)
            epoch = getattr(wb, "epoch", None)
            chunks = _split_stacked(rows)
            for c, chunk in enumerate(chunks):
                try:
                    t = _build_table(
                        chunk, file=name, kind="xlsx", sheet=ws.title,
                        tid=_table_id(digest, k if c == 0 else k * 1000 + c),
                        epoch=epoch if isinstance(epoch, dt.datetime) else None, formula_cols=frozenset(formula_cols),
                        extra_notes=tuple(notes_here),
                    )
                    if len(chunks) > 1:
                        t.sheet = f"{ws.title}, таблица {c + 1}"
                    tables.append(t)
                except _NoData as e:
                    first_error = first_error or str(e)
                    if formula_cols and first_error in ("Файл пустой", "В таблице нет значений"):
                        first_error = FORMULA_NOTE
    finally:
        wb.close()
    if not tables:
        if visible <= 1:
            raise DataReadError("Лист без данных" if first_error in (None, "Файл пустой") else first_error, file=name)
        raise DataReadError("В книге нет листов с данными", file=name)
    # титульный лист и списки подписей (одна колонка текста) рядом с настоящими таблицами — не данные
    useful = [t for t in tables if not _is_prose(t)]
    tables = useful or tables
    if len(tables) == 1:
        tables[0].sheet = None
    return tables


def _is_prose(t: Table) -> bool:
    if len(t.columns) != 1:
        return False
    vals = [v for v in t.column_values(0) if v is not None]
    return all(isinstance(v, str) and not (_ISO_DAY.fullmatch(v) or _ISO_PERIOD.fullmatch(v)) for v in vals)


# ======================================================================================================== JSON

_COLS_KEYS = ("columns", "headers", "header", "fields", "cols", "колонки", "столбцы")
_ROWS_KEYS = ("rows", "data", "values", "records", "items", "строки", "данные")


def _json_load(name: str, data: bytes) -> Any:
    text = _decode(data).strip()
    if not text:
        raise DataReadError("Файл пустой", file=name)
    try:
        return json.loads(text)
    except json.JSONDecodeError as e:
        # JSON Lines: запись на строку
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if len(lines) > 1:
            try:
                return [json.loads(ln) for ln in lines]
            except json.JSONDecodeError:
                pass
        raise DataReadError(f"Не получилось прочитать JSON: ошибка в строке {e.lineno}, позиция {e.colno}",
                            file=name) from None


def _flat(v: Any) -> Any:
    if isinstance(v, list):
        if all(not isinstance(x, (dict, list)) for x in v):
            return ", ".join("" if x is None else str(x) for x in v)
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def _records_rows(records: list[Any]) -> list[list[Any]]:
    """Записи → [шапка] + строки; вложенные объекты — колонки «родитель.ключ»."""
    keys: dict[str, None] = {}
    flat_recs = []
    for rec in records:
        if not isinstance(rec, dict):
            continue
        flat: dict[str, Any] = {}
        for k, v in rec.items():
            if isinstance(v, dict) and v and all(not isinstance(x, (dict, list)) for x in v.values()):
                for k2, v2 in v.items():
                    flat[f"{k}.{k2}"] = v2
            else:
                flat[str(k)] = _flat(v)
        for k in flat:
            keys.setdefault(k, None)
        flat_recs.append(flat)
    header = list(keys)
    return [header] + [[r.get(k) for k in header] for r in flat_recs]


def _col_name(c: Any) -> str:
    if isinstance(c, dict):
        for k in ("name", "title", "label", "key", "id", "field", "название", "имя"):
            if k in c and c[k] is not None:
                base = str(c[k])
                unit = c.get("unit") or c.get("units") or c.get("единица")
                return f"{base}, {unit}" if unit else base
        return json.dumps(c, ensure_ascii=False)
    return "" if c is None else str(c)


def _is_records(v: Any) -> bool:
    return isinstance(v, list) and any(isinstance(x, dict) for x in v) and all(
        isinstance(x, dict) or x is None for x in v)


def _is_scalar(v: Any) -> bool:
    return v is None or isinstance(v, (str, int, float, bool))


def _is_numberish(v: Any) -> bool:
    return _is_num(v) or (isinstance(v, str) and parse_quantity(v) is not None)


def _is_tableish(v: Any) -> bool:
    """Значение ключа — таблица: записи, массив строк, колонки-массивы или «категория → число»."""
    if _is_records(v):
        return True
    if isinstance(v, list) and v and all(isinstance(x, list) for x in v):
        return True
    if isinstance(v, list) and len(v) >= 2 and all(_is_scalar(x) for x in v):
        return sum(_is_numberish(x) for x in v) >= 0.8 * len(v)
    return False


# (имя таблицы, строки, строка шапки, заметки, имя безымянной первой колонки)
_JsonTable = tuple[Optional[str], list[list[Any]], Optional[int], str, str]


def _json_tables(obj: Any, path: str = "", depth: int = 0) -> list[_JsonTable]:
    """JSON → таблицы: массив записей; {columns, rows|data}; pandas split/table; колонки-массивы
    {"Месяц": [...], "Выручка": [...]}; «категория → число»; объект объектов; «ключ → таблица» (каждый ключ — своя)."""
    name = path or None
    if isinstance(obj, list):
        if not obj:
            return []
        if _is_records(obj):
            return [(name, _records_rows(obj), 0, "", "")]
        if all(isinstance(x, list) for x in obj):
            return [(name, [[_flat(c) for c in r] for r in obj], -1, "", "")]
        if all(_is_scalar(x) for x in obj):
            return [(name, [["Значение"]] + [[x] for x in obj], 0, "", "")]
        return []
    if not isinstance(obj, dict) or not obj:
        return []
    low = {str(k).lower(): k for k in obj}
    cols_key = next((low[k] for k in _COLS_KEYS if k in low and isinstance(obj[low[k]], list)), None)
    rows_key = next((low[k] for k in _ROWS_KEYS if k in low and isinstance(obj[low[k]], list)), None)

    # pandas orient="table": {"schema": {"fields": [...]}, "data": [записи]}
    schema = obj.get(low.get("schema", "\0"))
    if cols_key is None and isinstance(schema, dict) and rows_key is not None and _is_records(obj[rows_key]):
        rows = _records_rows(obj[rows_key])
        if rows[0] and rows[0][0] == "index":
            rows = [r[1:] for r in rows]
        return [(name, rows, 0, _json_notes(obj, {low.get("schema"), rows_key}), "")]

    # {"columns": [...], "rows" | "data": [[...], ...]} (и pandas orient="split" с "index")
    if cols_key is not None and rows_key is not None:
        header = [_col_name(c) for c in obj[cols_key]]
        data = obj[rows_key]
        notes = _json_notes(obj, {cols_key, rows_key, low.get("index")})
        if _is_records(data):
            rows = _records_rows(data)
            order = [rows[0].index(h) for h in header if h in rows[0]]
            if len(order) == len(header):
                rows = [[r[i] for i in order] for r in rows]
            return [(name, rows, 0, notes, "")]
        body = [[_flat(c) for c in r] if isinstance(r, list) else [_flat(r)] for r in data]
        index = obj.get(low.get("index", "\0"))
        first = ""
        if isinstance(index, list) and len(index) == len(body) and index != list(range(len(body))):
            header, body, first = [""] + header, [[_flat(ix)] + r for ix, r in zip(index, body)], "Категория"
        return [(name, [header] + body, 0, notes, first)]

    values = list(obj.values())
    # колонки-массивы одинаковой длины: {"Месяц": [...], "Выручка": [...]} (pandas orient="list")
    if len(values) >= 2 and all(isinstance(v, list) and all(_is_scalar(x) for x in v) for v in values):
        if len({len(v) for v in values}) == 1 and values[0]:
            return [(name, [[str(k) for k in obj]] + [list(r) for r in zip(*values)], 0, "", "")]
    # «категория → число»: одна серия; одни простые значения без чисел — одна запись
    if all(_is_scalar(v) for v in values):
        if sum(_is_numberish(v) for v in values) >= max(2, 0.8 * len(values)):
            return [(name, [["Категория", "Значение"]] + [[str(k), v] for k, v in obj.items()], 0, "", "")]
        return [(name, [[str(k) for k in obj], values], 0, "", "")] if depth == 0 else []
    # объект объектов: {"Москва": {"Выручка": 1433}} (внешние ключи — строки) или pandas orient="columns"
    # {"Выручка": {"Москва": 1433}} (внешние ключи — колонки): строк обычно больше, чем колонок
    if all(isinstance(v, dict) and v and all(_is_scalar(x) for x in v.values()) for v in values):
        inner = list(dict.fromkeys(str(k) for v in values for k in v))
        if len(obj) > len(inner):
            header = [""] + inner
            body = [[str(k)] + [v.get(ik) for ik in inner] for k, v in obj.items()]
        else:
            header = [""] + [str(k) for k in obj]
            body = [[ik] + [v.get(ik) for v in values] for ik in inner]
        return [(name, [header] + body, 0, "", "Категория")]
    # «ключ → таблица»: каждый ключ — своя таблица; рядом с таблицами простые объекты и строки — заметки
    notes = _json_notes(obj)
    tableish = {k for k, v in obj.items() if _is_tableish(v)}
    out: list[_JsonTable] = []
    if depth < 4:
        for k, v in obj.items():
            if not isinstance(v, (list, dict)):
                continue
            if tableish and k not in tableish and not (isinstance(v, dict) and _has_table(v)):
                continue  # метаданные рядом с таблицами
            if isinstance(v, list) and k not in tableish:
                continue  # список подписей, а не таблица
            sub = _json_tables(v, str(k) if not path else f"{path}.{k}", depth + 1)
            out += [(n, r, h, "\n".join(p for p in (notes, nt) if p), f) for n, r, h, nt, f in sub]
    return out


def _has_table(obj: dict, depth: int = 0) -> bool:
    if depth > 3:
        return False
    return any(_is_tableish(v) or (isinstance(v, dict) and _has_table(v, depth + 1)) for v in obj.values())


def _json_notes(obj: dict, skip: Optional[set] = None) -> str:
    skip = skip or set()
    parts = []
    for k, v in obj.items():
        if k in skip:
            continue
        if isinstance(v, str) and v.strip():
            parts.append(f"{k}: {v.strip()}")
        elif _is_num(v):
            parts.append(f"{k}: {v}")
        elif isinstance(v, dict) and v and all(_is_scalar(x) for x in v.values()) and not _is_tableish(v):
            # «meta»: {total, generated} рядом с таблицей — в заметки, а не молча прочь
            parts.append(f"{k}: " + ", ".join(f"{k2} {x}" for k2, x in v.items() if x is not None and str(x).strip()))
    return "\n".join(parts)


def _read_json(name: str, data: bytes, digest: str) -> list[Table]:
    obj = _json_load(name, data)
    found = _json_tables(obj)
    if not found:
        raise DataReadError(
            "В JSON нет таблицы: нужен массив записей, объект {columns, rows} или «ключ → массив записей»", file=name
        )
    tables: list[Table] = []
    first_error = None
    for k, (sheet, rows, header_row, notes, first_name) in enumerate(found):
        try:
            tables.append(_build_table(rows, file=name, kind="json", sheet=sheet, tid=_table_id(digest, k),
                                       notes=notes, header_row=header_row, first_name=first_name))
        except _NoData as e:
            first_error = first_error or str(e)
    if not tables:
        raise DataReadError(
            "В JSON нет строк с данными" if first_error in (None, "Файл пустой") else first_error, file=name
        )
    if len(tables) == 1:
        tables[0].sheet = None
    return tables


# ======================================================================================================== вход

_EXT_KIND: dict[str, str] = {
    ".csv": "csv", ".tsv": "tsv", ".tab": "tsv",
    ".xlsx": "xlsx", ".xlsm": "xlsx", ".xltx": "xlsx", ".xltm": "xlsx",
    ".json": "json", ".jsonl": "json", ".ndjson": "json", ".geojson": "json",
    ".txt": "text", ".md": "text", ".markdown": "text", ".text": "text",
}
_UNSUPPORTED: dict[str, str] = {
    ".xls": "Старый формат Excel (.xls) не читается — сохраните файл как .xlsx или .csv",
    ".xlsb": "Двоичная книга Excel (.xlsb) не читается — сохраните файл как .xlsx или .csv",
    ".ods": "Таблица OpenDocument (.ods) не читается — сохраните её как .xlsx или .csv",
    ".numbers": "Файл Numbers не читается — экспортируйте таблицу в .xlsx или .csv",
}
_NOT_DATA = {".pdf", ".doc", ".docx", ".rtf", ".odt", ".ppt", ".pptx", ".key", ".png", ".jpg", ".jpeg", ".gif",
             ".webp", ".svg", ".zip", ".rar", ".7z", ".parquet", ".sqlite", ".db"}


def read_bytes(name: str, data: bytes) -> list[Table]:
    """Файл с данными (имя и содержимое) → таблицы: лист XLSX, ключ JSON, таблица в тексте — каждая своя `Table`.
    Не получилось — `DataReadError` с русским текстом («Файл пустой», «Не нашёл таблицу в тексте», «Лист без данных»)."""
    name = Path(str(name)).name or "данные"
    ext = Path(name).suffix.lower()
    if ext in _NOT_DATA:
        raise DataReadError(
            f"Файлы {ext} с данными не читаю — приложите таблицу как CSV, XLSX, JSON или текст с таблицей", file=name
        )
    if not data or not data.strip():
        raise DataReadError("Файл пустой", file=name)
    kind = _EXT_KIND.get(ext)
    if ext in _UNSUPPORTED or kind == "xlsx":
        # «выгрузка в Excel», которая на деле текст через табуляцию или «;» с расширением .xls/.xlsx
        if _text_disguised_as_excel(name, data):
            kind = "csv"
        elif ext in _UNSUPPORTED:
            raise DataReadError(_UNSUPPORTED[ext], file=name)
    digest = hashlib.sha1(name.encode() + b"\0" + data).hexdigest()
    if kind is None or (kind != "xlsx" and data[:4] == b"PK\x03\x04"):
        kind = _sniff_kind(data)
    if kind == "xlsx":
        return _read_xlsx(name, data, digest)
    if kind == "json":
        return _read_json(name, data, digest)
    text = _decode(data)
    if _looks_binary(text):
        raise DataReadError("Файл не похож на таблицу: внутри двоичные данные, а не текст", file=name)
    if not text.strip():
        raise DataReadError("Файл пустой", file=name)
    if kind == "text":
        return _read_text(name, text, digest)
    return _read_delimited(name, text, kind, digest)  # csv / tsv


_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _text_disguised_as_excel(name: str, data: bytes) -> bool:
    if data[:4] == b"PK\x03\x04" or data[:8] == _OLE:
        return False
    head = _decode(data[:8000])
    if _looks_binary(head):
        return False
    if head.lstrip()[:1] == "<":
        raise DataReadError(
            "Файл — страница HTML или XML, а не книга Excel: откройте его в Excel и сохраните как .xlsx или .csv",
            file=name,
        )
    return True


def _sniff_kind(data: bytes) -> str:
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if any(n.startswith("xl/") for n in z.namelist()):
                    return "xlsx"
        except zipfile.BadZipFile:
            pass
        raise DataReadError("Архив не похож на книгу Excel — приложите таблицу как CSV, XLSX, JSON или текст")
    head = data.lstrip()[:1]
    if head in (b"[", b"{"):
        return "json"
    return "text"


def read_file(path: Union[str, Path]) -> list[Table]:
    """Файл на диске → таблицы (см. `read_bytes`)."""
    p = Path(path)
    if not p.exists():
        raise DataReadError(f"Файл «{p.name}» не найден", file=p.name)
    if p.is_dir():
        raise DataReadError(f"«{p.name}» — папка, а не файл с данными", file=p.name)
    return read_bytes(p.name, p.read_bytes())
