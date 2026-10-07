"""«Улучшить презентацию» (Verstka 3.0): a finished deck read back as the author's text — the second way in the
organisers named («есть старая презентация — нужна улучшенная/обновлённая»).

Every visible slide becomes a «Слайд N. Заголовок» section of a brief — the first one, when it is the deck's cover, in
the brief's own cover format («Слайд 1. Титульный» / «Название: «…».» / «Подзаголовок: «…».»): its title, its text in
reading order (a list stays a list, the boxes of one card stay together, a KPI tile is one line «Средний чек — 300 ₽»),
its tables as markdown, its charts as a chart note with their data («[Столбчатая диаграмма: выручка, млн ₽ — 2023: 1,2;
2024: 2,5]»; a chart of several series as a note over a markdown table of all of them), its SmartArt as a list, its
pictures as a place for the picture («[Место под изображение: Изображение со слайда 3 исходной презентации · 1a2b3c]»)
whose file is kept, so the improved deck shows the original picture there (images/fill.py reuses it by its
description — the deck's code in it, so another deck's picture never takes the place — and no picture is generated for
it), its speaker notes as «[Заметки докладчика: …]» (the improved slide's notes, never its text: planning/
brief_structure.py takes them out of the slide's text). What every slide repeats — a footer, the company name, the page
number, a logo — is the template's, not the author's: left out; so are hidden slides, step badges («1», «2», «3») and a
legend drawn as text boxes beside its chart. The deck itself becomes the template, so the improved deck keeps its look;
the deterministic audit of the original tells «было → стало».
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

from verstka.analysis.xmlns import local_name, q

log = logging.getLogger(__name__)

SOURCE_IMAGES = "source_images.json"  # in the template's workspace: {hint: file name in source_images/}
SPEAKER_NOTES = "Заметки докладчика"  # the label of the note that carries a slide's speaker notes in the brief
_PICTURE_HINT = "Изображение{k} со слайда {n} исходной презентации · {deck}"
# a place for a picture of the original deck: its picture is the original's or none — never one a model draws from these
# words («Изображение со слайда 3 исходной презентации», with the deck's code since 3.0)
_SOURCE_HINT_RE = re.compile(r"^Изображение(?:\s+\d+)?\s+со\s+слайда\s+\d+\s+исходной\s+презентации\b", re.I)
_TINY_PICTURE = 0.015  # of the slide's area: an icon, a bullet, a mark — never the author's picture
_SMALL_PICTURE = 0.06  # a logo, a portrait: the author's unless the slides repeat it
_BACKGROUND = 0.92  # a picture over (almost) the whole slide is its background art
_CHART_RU = {
    "pie": "Круговая диаграмма", "doughnut": "Кольцевая диаграмма", "bar": "Горизонтальная диаграмма", "line": "Линейная диаграмма",
    "area": "Диаграмма с областями", "column": "Столбчатая диаграмма",
}
_PLOT_KIND = {
    "pieChart": "pie", "pie3DChart": "pie", "ofPieChart": "pie", "doughnutChart": "doughnut", "lineChart": "line", "line3DChart": "line",
    "stockChart": "line", "scatterChart": "line", "radarChart": "line", "areaChart": "area", "area3DChart": "area",
}
_SHAPE_TAGS = {"sp", "grpSp", "graphicFrame", "cxnSp", "pic"}
_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_ENDS = (".", "!", "?", ":", ";", "…")


def is_source_hint(hint: Optional[str]) -> bool:
    """The place names a picture of the original deck (its hint as deck_text wrote it)."""
    return bool(_SOURCE_HINT_RE.match((hint or "").strip()))


@dataclass
class DeckText:
    brief: str
    title: str = ""
    slides: int = 0  # the slides read into the brief (hidden and empty ones not counted)
    tables: int = 0
    charts: int = 0
    pictures: list[tuple[str, bytes, str]] = field(default_factory=list)  # (hint, image bytes, extension)

    def summary(self) -> dict:
        return {"title": self.title, "slides": self.slides, "tables": self.tables, "charts": self.charts, "pictures": len(self.pictures)}


@dataclass
class _Part:
    """One thing a slide shows, where it shows it (slide coordinates, EMU)."""

    text: str
    x: float
    y: float
    w: float
    h: float
    kind: str = "text"  # text | list | code | table | chart | picture | diagram
    size: float = 0.0  # the largest font size of its text, in points (inherited sizes resolved)
    lines: list[str] = field(default_factory=list)  # a text's paragraphs
    bulleted: bool = False  # a list the deck marks with bullets
    role: str = ""  # title | subtitle (its placeholder)
    names: list[str] = field(default_factory=list)  # a chart's series and category names (its legend repeats them)
    values: list[float] = field(default_factory=list)  # a chart's values as its labels may show them
    image: Optional[tuple[bytes, str]] = None  # a picture's bytes (as the slide shows them) and extension
    label: str = ""  # a KPI line read as one («Средний чек — 300 ₽»): its label and its figure
    figure: str = ""
    art: bool = False  # a picture the template analysis reads as art (a render, an illustration): one of a chart's bars?

    @property
    def right(self) -> float:
        return self.x + self.w

    @property
    def bottom(self) -> float:
        return self.y + self.h

    @property
    def line(self) -> bool:
        """One line of text: a title, a label or a figure may be it."""
        return self.kind == "text" and "\n" not in self.text


def _norm(text: str) -> str:
    return " ".join((text or "").split()).strip().lower()


_INVISIBLE_RE = re.compile("[\u200b\u200c\u200d\u2060\ufeff\u00ad]")


def _clean(text: str) -> str:
    return " ".join(_INVISIBLE_RE.sub("", (text or "").replace("\v", " ").replace("\u000b", " ")).split()).strip()


def _words(text: str) -> int:
    return len((text or "").split())


# ------------------------------------------------------------------ shapes


def _members(shapes: Any) -> list[Any]:
    """The shapes of a tree, each made on its own: an element python-pptx cannot make (ink, a part it does not know) is
    skipped, never the slide; the fallback of an alternate content (an equation, a 3D model) is read like any shape."""
    try:
        tree = shapes._spTree
    except Exception:  # noqa: BLE001
        return []
    els = []
    for el in tree.iterchildren():
        name = local_name(el) if isinstance(el.tag, str) else ""
        if name == "AlternateContent":
            fb = el.find(q("mc:Fallback"))
            els.extend(c for c in (fb if fb is not None else []) if isinstance(c.tag, str) and local_name(c) in _SHAPE_TAGS)
        elif name in _SHAPE_TAGS:
            els.append(el)
    out = []
    for el in els:
        try:
            out.append(shapes._shape_factory(el))
        except Exception:  # noqa: BLE001 - a shape python-pptx cannot make is left out
            log.debug("shape not read: %s", local_name(el), exc_info=True)
    return out


def _tag(sh: Any) -> str:
    try:
        return local_name(sh._element)
    except Exception:  # noqa: BLE001
        return ""


def _flat(shapes: Any, offset: tuple[float, float, float, float] = (0.0, 0.0, 1.0, 1.0), groups: Optional[list] = None) -> list[tuple[Any, float, float, float, float]]:
    """(shape, x, y, w, h) of every shape in slide coordinates, groups opened (`groups` collects each group's box)."""
    ox, oy, kx, ky = offset
    out = []
    for sh in _members(shapes):
        try:
            x, y, w, h = (float(v or 0) for v in (sh.left, sh.top, sh.width, sh.height))
        except Exception:  # noqa: BLE001 - a shape without a position (a broken one) is read where it lies
            x = y = w = h = 0.0
        ax, ay, aw, ah = ox + x * kx, oy + y * ky, w * kx, h * ky
        if _tag(sh) == "grpSp":
            if groups is not None:
                groups.append((ax, ay, aw, ah))
            try:
                xfrm = sh._element.grpSpPr.xfrm
                ch_off, ch_ext = xfrm.chOff, xfrm.chExt
                gkx = (aw / ch_ext.cx) if ch_ext is not None and ch_ext.cx else kx
                gky = (ah / ch_ext.cy) if ch_ext is not None and ch_ext.cy else ky
                gx = ax - (ch_off.x if ch_off is not None else 0) * gkx
                gy = ay - (ch_off.y if ch_off is not None else 0) * gky
                out.extend(_flat(sh.shapes, (gx, gy, gkx, gky), groups))
            except Exception:  # noqa: BLE001 - a group read flat
                out.extend(_flat(sh.shapes, offset, groups))
            continue
        out.append((sh, ax, ay, aw, ah))
    return out


def _placeholder_kind(sh: Any) -> Optional[str]:
    if not getattr(sh, "is_placeholder", False):
        return None
    try:
        t = sh.placeholder_format.type
    except Exception:  # noqa: BLE001
        return None
    name = str(t).split(".")[-1].split(" ")[0].upper()
    if name in ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE"):
        return "title"
    if name in ("SUBTITLE",):
        return "subtitle"
    if name in ("DATE", "FOOTER", "SLIDE_NUMBER", "HEADER"):
        return "chrome"
    return "body"


# ------------------------------------------------------------------ text: sizes, bullets, paragraphs


def _styles(sh: Any) -> list:
    """The list styles the box's paragraphs inherit from, nearest first: its own, its layout's and master's placeholder,
    the master's title or body style — or, for a text box, the presentation's default text style."""
    out = []
    try:
        body = sh._element.find(q("p:txBody"))
        if body is not None:
            out.append(body.find(q("a:lstStyle")))
        if getattr(sh, "is_placeholder", False):
            base = sh
            for _ in range(2):
                try:
                    base = base._base_placeholder
                except Exception:  # noqa: BLE001
                    base = None
                if base is None:
                    break
                body = base._element.find(q("p:txBody"))
                if body is not None:
                    out.append(body.find(q("a:lstStyle")))
            master = sh.part.slide.slide_layout.slide_master._element
            which = "p:titleStyle" if _placeholder_kind(sh) == "title" else "p:bodyStyle"
            out.append(master.find(q("p:txStyles") + "/" + q(which)))
        else:
            out.append(sh.part.package.presentation_part._element.find(q("p:defaultTextStyle")))
    except Exception:  # noqa: BLE001 - what is found so far
        pass
    return [x for x in out if x is not None]


def _level_pr(styles: list, lvl: int) -> list:
    return [el for el in (st.find(q(f"a:lvl{min(max(lvl, 1), 9)}pPr")) for st in styles) if el is not None]


def _sz(el: Any) -> Optional[float]:
    try:
        v = el.get("sz") if el is not None else None
        return int(v) / 100.0 if v else None
    except (TypeError, ValueError):
        return None


def _sizes(sh: Any, styles: Optional[list] = None) -> list[float]:
    """The font size of every run with text (points): its own, else its paragraph's or the inherited one; shrunk as
    the box's autofit shrinks it."""
    styles = _styles(sh) if styles is None else styles
    out: list[float] = []
    try:
        for p in sh.text_frame.paragraphs:
            lvl = int(p.level or 0) + 1
            ppr = p._p.pPr
            base = _sz(ppr.find(q("a:defRPr"))) if ppr is not None else None
            if base is None:
                base = next((s for s in (_sz(el.find(q("a:defRPr"))) for el in _level_pr(styles, lvl)) if s), None)
            for r in p.runs:
                if r.text.strip():
                    out.append(_sz(r._r.find(q("a:rPr"))) or base or 18.0)
        fit = sh._element.find(".//" + q("a:bodyPr") + "/" + q("a:normAutofit"))
        scale = int(fit.get("fontScale")) / 100000.0 if fit is not None and fit.get("fontScale") else 1.0
    except Exception:  # noqa: BLE001
        return out
    return [round(s * scale, 1) for s in out]


def _max_size(sh: Any) -> float:
    """The largest font size of the box's text, in points (inherited sizes resolved from the layout, the master and the
    deck's default text style)."""
    return max(_sizes(sh), default=18.0)


def _bullet_of(ppr: Any) -> Optional[bool]:
    if ppr is None:
        return None
    for ch in ppr:
        name = local_name(ch) if isinstance(ch.tag, str) else ""
        if name in ("buChar", "buAutoNum", "buBlip"):
            return True
        if name == "buNone":
            return False
    return None


def _paragraphs(sh: Any, styles: Optional[list] = None) -> list[tuple[int, str, bool]]:
    """The box's paragraphs as (level, text, bulleted). A paragraph of plain text that only continues the one before
    it (a label the layout wrapped: «Продукты, упаковка» / «и списания») is joined to it — never an item of a list."""
    styles = _styles(sh) if styles is None else styles
    out: list[tuple[int, str, bool]] = []
    for p in sh.text_frame.paragraphs:
        t = _clean(p.text)
        if not t:
            continue
        lvl = int(p.level or 0)
        bu = _bullet_of(p._p.pPr)
        if bu is None:
            bu = next((b for b in (_bullet_of(el) for el in _level_pr(styles, lvl + 1)) if b is not None), False)
        if out and not bu and not out[-1][2] and t[:1].islower() and not out[-1][1].endswith(_ENDS) and lvl == out[-1][0]:
            out[-1] = (out[-1][0], f"{out[-1][1]} {t}", False)
            continue
        out.append((lvl, t, bu))
    return out


_MONO_RE = re.compile(r"courier|consolas|menlo|monaco|mono|source code|fira code|jetbrains", re.I)
_BRACE_LINE_RE = re.compile(r"^[{}\[\]()]+[,;]?$")
# what only code writes: a JSON key, a call «name(…)», an assignment or comparison of an identifier, a shell or template
# variable, a tag, a comment marker, a statement's end after code
_CODE_SIGNAL_RE = re.compile(
    r"^\"[\w\-. ]+\"\s*:|[A-Za-z_][\w.]*\(|[A-Za-z_][\w.\]\[]*\s*(?:[+\-*/%]?=|==|!=|<=|>=|=>|->|:=)\s*\S|\$[A-Za-z_{(]|</?[a-z][\w-]*[^<>]*>|"
    r"^(?:#|//|--)\s?\S|[;{]$"
)
_CODE_WORD_RE = re.compile(
    r"^(?:def|class|return|import|from|if|elif|else|for|while|var|let|const|function|public|private|static|void|int|try|except|"
    r"catch|select|insert|update|delete|create|async|await)\b"
)


_STRING_RE = re.compile(r"\"[^\"]*\"|'[^']*'")


def _code_line(line: str, mono: bool) -> bool:
    """A line only code writes. Its strings are left out first (a message in Russian inside code is still code); a line
    of words in Russian («снизить затраты на 10%;») never is."""
    t = line.replace("\u00a0", " ").strip()
    if not t:
        return False
    if _BRACE_LINE_RE.match(t) or re.fullmatch(r"[)\]}]+[:;,]?", t) or re.match(r"^\"[\w\-. ]+\"\s*:", t):
        return True  # a bracket closing a block, a JSON key
    bare = _STRING_RE.sub('""', t)
    cyr = len(re.findall(r"[а-яё]", bare, re.I))
    lat = len(re.findall(r"[a-z]", bare, re.I))
    if cyr >= lat:
        return mono and bare.startswith(("#", "//")) and lat + cyr > 0  # a comment of the code in Russian
    if _CODE_SIGNAL_RE.search(bare) or _CODE_WORD_RE.match(bare):
        return True
    return mono and bool(re.match(r"^[ \t\u00a0]{2,}\S", line))  # an indented line of a monospace box


def _raw_lines(sh: Any) -> list[str]:
    return [ln for p in sh.text_frame.paragraphs for ln in p.text.replace("\v", "\n").split("\n") if ln.strip()]


def _is_code(sh: Any, lines: Optional[list[str]] = None) -> bool:
    """Code needs code's own marks on most of its lines (two or more lines; a monospace face lowers the share) — a
    list whose items end with «;» or a heading set in a monospace design face («01 / Проблема») is never code."""
    lines = _raw_lines(sh) if lines is None else lines
    if len(lines) < 2:
        return False
    try:
        fonts = {r.font.name for p in sh.text_frame.paragraphs for r in p.runs if r.font is not None and r.font.name}
    except Exception:  # noqa: BLE001
        fonts = set()
    mono = any(_MONO_RE.search(f) for f in fonts)
    hits = sum(1 for ln in lines if _code_line(ln, mono))
    return hits >= 2 and hits >= (0.4 if mono else 0.6) * len(lines)


def _code_text(sh: Any) -> str:
    raw = "\n".join(p.text.replace("\v", "\n") for p in sh.text_frame.paragraphs).strip("\n")
    return f"```\n{raw}\n```"


def _list_text(paras: list[tuple[int, str, bool]]) -> str:
    low = min(lv for lv, _, _ in paras)
    return "\n".join(("  " * (lv - low)) + "- " + t for lv, t, _ in paras)


def _text_block(paras: list[tuple[int, str, bool]]) -> tuple[str, str]:
    """(kind, text) of a box of several paragraphs: a list (bulleted, or lines that are not sentences — each stays an
    item of its own), the lead line ending with «:» over its items, paragraphs of prose — or a name in capitals the box
    broke over two lines («СВЕТЛАНА» / «АРТЕМЬЕВА»)."""
    if not any(bu for _, _, bu in paras) and len(paras) == 2 and all(t.upper() == t and re.search(r"[^\W\d_]", t) and _words(t) <= 2 for _, t, _ in paras):
        return "text", " ".join(t for _, t, _ in paras)
    if not any(bu for _, _, bu in paras):
        if all(t.endswith((".", "!", "?", "…")) for _, t, _ in paras):
            return "text", "\n\n".join(t for _, t, _ in paras)  # prose: paragraphs of whole sentences
        if len(paras) >= 3 and paras[0][1].endswith(":"):
            return "list", paras[0][1] + "\n" + _list_text(paras[1:])
    return "list", _list_text(paras)


# ------------------------------------------------------------------ figures and labels


_NUM_TXT = r"[+\-−±]?\d[\d   ]*(?:[.,]\d+)?"
_UNIT_TXT = (
    r"(?:%|‰|₽|\$|€|£|°C|°F|°|руб\.?|рубл(?:ей|я|ь)|тыс\.?|млн\.?|млрд\.?|трлн\.?|ч\.?|час(?:а|ов)?|мин\.?|минут[аы]?|сек\.?|секунд[аы]?|с\.?|"
    r"дн\.?|дн(?:ей|я)|день|нед\.?|мес\.?|месяц(?:а|ев)?|лет|год(?:а|ов)?|г\.|кг|т|км|м²|м³|м2|м3|м|см|мм|га|л|шт\.?|чел\.?|человек|"
    r"x|×|раз[а]?|п\.\s?п\.|pp|k|m|bn|b|гб|мб|тб|gb|mb|tb|квт(?:·ч)?|мвт|вт|ггц|мгц|fps|мс|ms)"
)
_ONE_FIG = rf"(?:(?:[≤≥<>~≈]|до|от|около|более|менее|свыше|почти)\s*)?(?:[x×]\s?)?[$€£₽]?\s?{_NUM_TXT}(?:\s?{_UNIT_TXT}(?![а-яёa-z])){{0,2}}"
# a figure alone: «300 ₽», «25,4 °C», «12 400», «+34%», «2,1 часа», «20 000 → 35 000 ₽», «35% → 33%», «6 × 4 м», «≤ 3 мин»,
# «113 с», «1 мин 18 с», «99 из 100», «24/7»
_FIGURE_RE = re.compile(
    rf"^{_ONE_FIG}(?:\s*(?:→|->|—|–|-|/|×|x|·|из|до|и)\s*{_ONE_FIG}){{0,3}}(?:\s+{_NUM_TXT}\s?{_UNIT_TXT}(?![а-яёa-z]))?\.?$", re.I
)
_OPERATOR_RE = re.compile(r"^[×x*+=→←↑↓–—\-/:]{1,2}$")
_BADGE_RE = re.compile(r"^(\d{1,2})[.)]?$")


def _is_figure(text: str) -> bool:
    return bool(_FIGURE_RE.match((text or "").strip()))


def _label_like(p: _Part, most: int = 8) -> bool:
    """A short label of a figure: one line of a few words, not a figure itself (a label may hold a digit: «Выручка за 30
    дней»)."""
    n = len(re.findall(r"\w+", p.text))
    return p.line and p.role != "title" and 1 <= n <= most and not _is_figure(p.text) and not p.text.startswith(("[", "|", "`"))


# ------------------------------------------------------------------ tables and charts


def _md_table(rows: list[list[str]]) -> str:
    rows = [[_clean(c).replace("|", "/") for c in r] for r in rows if any(_clean(c) for c in r)]
    if not rows:
        return ""
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * n]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _plain_fmt(fmt: str) -> str:
    """A number format without its quoted literals and bracketed codes."""
    return re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", (fmt or "").split(";")[0])


def _percent(fmt: str) -> bool:
    return "%" in _plain_fmt(fmt)


def _fmt_unit(fmt: str) -> str:
    """The unit a number format writes after or before the figure («#,##0 "₽"», «[$₽-419]#,##0», «0.0 "млн"»)."""
    head = (fmt or "").split(";")[0]
    lit = " ".join(x.strip() for x in re.findall(r'"([^"]*)"', head) if x.strip())
    cur = re.search(r"\[\$([^\-\]]+)", head)
    sym = re.search(r"[₽$€£]", _plain_fmt(head))
    unit = lit or (cur.group(1).strip() if cur else "") or (sym.group(0) if sym else "")
    return unit if re.search(r"[^\W\d_]|[₽$€£]", unit) else ""


def _fmt_scale(fmt: str) -> float:
    """«#,##0,,» shows millions: each comma after the last digit place divides by a thousand."""
    m = re.search(r"[0#](,+)(?![0#])", _plain_fmt(fmt))
    return 1000.0 ** len(m.group(1)) if m else 1.0


def _num_text(f: float) -> str:
    """A figure as a Russian slide writes it: thousands by a space, a decimal comma, up to four decimals (four significant
    digits of a value under one: «0,00105», never «0»)."""
    if not math.isfinite(f):
        return ""
    dec = 4
    if f and abs(f) < 1:
        dec = min(12, max(4, 3 - int(math.floor(math.log10(abs(f))))))  # four significant digits of a small value
    s = f"{f:,.{dec}f}"
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    if s in ("-0", ""):
        s = "0"
    return s.replace(",", " ").replace(".", ",")


def _figure(v: Any, fmt: str = "") -> str:
    """A chart value as the chart shows it: a percent format times a hundred with «%», a scaled format («#,##0,,»)
    scaled, its unit after the figure."""
    if v is None:
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return _clean(str(v))
    if _percent(fmt):
        return _num_text(f * 100) + "%"
    s = _num_text(f / _fmt_scale(fmt))
    unit = _fmt_unit(fmt)
    return f"{s} {unit}" if unit else s


_RU_MONTHS = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")


def _category(v: Optional[str], fmt: str) -> str:
    """A category label: a date a number format shows (Excel's serial day) as a date, a number as a slide writes it."""
    if v is None:
        return ""
    s = _clean(v)
    plain = _plain_fmt(fmt).lower()
    try:
        f = float(s)
    except ValueError:
        return s
    if plain and re.search(r"[dy]|m{3,}", plain) and not re.search(r"[0#]", plain):
        try:
            import datetime as dt

            d = dt.datetime(1899, 12, 30) + dt.timedelta(days=f)
            if "d" in plain:
                return d.strftime("%d.%m.%Y")
            return f"{_RU_MONTHS[d.month - 1]} {d.year}" if "mmm" in plain else d.strftime("%m.%Y")
        except (OverflowError, ValueError):
            return s
    return _num_text(f) if f != int(f) or abs(f) >= 10000 else str(int(f))


def _points(ref: Any) -> tuple[list[Optional[str]], str]:
    """The values a chart's reference (or literal) holds, by index, and their number format."""
    if ref is None:
        return [], ""
    cache = None
    for tag in ("c:strCache", "c:numCache", "c:strLit", "c:numLit"):
        cache = ref.find(".//" + q(tag))
        if cache is not None:
            break
    count_el = cache
    if cache is None:
        multi = ref.find(".//" + q("c:multiLvlStrCache"))
        if multi is None:
            return [], ""
        count_el, cache = multi, multi.find(q("c:lvl"))  # the innermost level: the labels by the axis
        if cache is None:
            return [], ""
    fmt_el = cache.find(q("c:formatCode"))
    fmt = (fmt_el.text or "") if fmt_el is not None else ""
    pc = count_el.find(q("c:ptCount"))
    n = int(pc.get("val")) if pc is not None and (pc.get("val") or "").isdigit() else 0
    pts = cache.findall(q("c:pt"))
    n = max([n] + [int(p.get("idx") or 0) + 1 for p in pts])
    out: list[Optional[str]] = [None] * min(n, 10000)
    for p in pts:
        i = int(p.get("idx") or 0)
        v = p.find(q("c:v"))
        if v is not None and 0 <= i < len(out):
            out[i] = v.text
    return out, fmt


def _rich_text(el: Any) -> str:
    if el is None:
        return ""
    rich = [_clean("".join(t.text or "" for t in p.iter(q("a:t")))) for p in el.iter(q("a:p"))]
    text = " ".join(x for x in rich if x)
    if text:
        return text
    v = el.find(".//" + q("c:v"))
    return _clean(v.text or "") if v is not None else ""


def _chart_data(space: Any) -> tuple[str, str, str, list[tuple[str, list[Optional[str]], str, list[Optional[float]], str]]]:
    """(kind, title, category axis title, series) of a chart part, read from its XML whatever its type (3D, pie of pie,
    stock, scatter): each series as (name, categories, their format, values, number format)."""
    chart = space.find(q("c:chart"))
    plot = chart.find(q("c:plotArea")) if chart is not None else None
    if plot is None:
        return "column", "", "", []
    kind = ""
    series = []
    ax_fmt = ""
    num_ax = plot.find(q("c:valAx") + "/" + q("c:numFmt"))
    if num_ax is not None:
        ax_fmt = num_ax.get("formatCode") or ""
    for pl in plot:
        name = local_name(pl) if isinstance(pl.tag, str) else ""
        if not name.endswith("Chart"):
            continue
        if name in ("barChart", "bar3DChart"):
            d = pl.find(q("c:barDir"))
            k = "bar" if d is not None and d.get("val") == "bar" else "column"
        else:
            k = _PLOT_KIND.get(name, "column")
        kind = kind or k
        for ser in pl.findall(q("c:ser")):
            sname = _rich_text(ser.find(q("c:tx")))
            cat_ref = ser.find(q("c:cat"))
            if cat_ref is None:
                cat_ref = ser.find(q("c:xVal"))
            cats, cat_fmt = _points(cat_ref)
            val_ref = ser.find(q("c:val"))
            if val_ref is None:
                val_ref = ser.find(q("c:yVal"))
            raw, fmt = _points(val_ref)
            label_fmt = ser.find(q("c:dLbls") + "/" + q("c:numFmt"))
            if label_fmt is not None and label_fmt.get("formatCode") and label_fmt.get("sourceLinked") != "1":
                fmt = label_fmt.get("formatCode") or fmt
            if not fmt or fmt.lower() == "general":
                fmt = ax_fmt if ax_fmt.lower() != "general" else ""
            vals: list[Optional[float]] = []
            for v in raw:
                try:
                    vals.append(float(v) if v is not None and str(v).strip() else None)
                except ValueError:
                    vals.append(None)
            series.append((sname, cats, cat_fmt, vals, fmt))
    title = ""
    t = chart.find(q("c:title"))
    deleted = chart.find(q("c:autoTitleDeleted"))
    if t is not None:
        title = _rich_text(t.find(q("c:tx")))
        if not title and len(series) == 1 and not (deleted is not None and deleted.get("val") in ("1", "true")):
            title = series[0][0]  # an automatic title: the series' name
    axis = plot.find(q("c:catAx"))
    if axis is None:
        axis = plot.find(q("c:dateAx"))
    cat_title = _rich_text(axis.find(q("c:title"))) if axis is not None else ""
    return kind or "column", title, cat_title, series


def _inline_ok(body: str, n: int) -> bool:
    """The note's data reads back whole (planning/faithful.py): every category an item with its one figure."""
    try:
        from verstka.planning.faithful import chart_items, note_covered

        items = chart_items(body)
        return len(items) == n and all(len(it.values) == 1 for it in items) and note_covered(body, items)
    except Exception:  # noqa: BLE001 - the table then
        return False


def _chart_note(space: Any) -> Optional[tuple[str, list[str], list[float]]]:
    """A chart as the brief writes it, with all its data — one series inline («[Круговая диаграмма: Доля, % — Онлайн:
    55; Офлайн: 35]»), several (or data an inline note would not read back whole) as a note over a markdown table:
    categories by rows, every series a column — and the names and values a legend beside it may repeat. None when the
    chart holds no data."""
    kind, title, cat_title, series = _chart_data(space)
    series = [s for s in series if any(v is not None for v in s[3])]
    if not series:
        return None
    cats_src = max(series, key=lambda s: len(s[1]))
    n = max(max(len(s[3]) for s in series), len(cats_src[1]))
    cats = [_category(cats_src[1][i], cats_src[2]) if i < len(cats_src[1]) else "" for i in range(n)]
    cats = [c or str(i + 1) for i, c in enumerate(cats)]
    names = [_clean(s[0]) or f"Ряд {k}" for k, s in enumerate(series, 1)]
    shown: list[float] = []
    for _, _, _, vals, fmt in series:
        for v in vals:
            if v is not None:
                shown.append(v * 100 if _percent(fmt) else v / _fmt_scale(fmt))
    legend = [*(_clean(s[0]) for s in series if _clean(s[0])), *cats]
    head = _CHART_RU.get(kind, _CHART_RU["column"])
    if len(series) == 1:
        _, _, _, vals, fmt = series[0]
        what = _clean(title) or names[0]
        items = [(c, _figure(v, fmt)) for c, v in zip(cats, vals) if v is not None]
        body = f"{what} — " + "; ".join(f"{c}: {v}" for c, v in items)
        if _inline_ok(body, len(items)):
            return f"[{head}: {body}]", legend, shown
    rows = [[cat_title or "", *names]]
    for i, c in enumerate(cats):
        rows.append([c, *(_figure(s[3][i], s[4]) if i < len(s[3]) else "" for s in series)])
    note = f"[{head}: {_clean(title)}]" if _clean(title) else f"[{head}]"
    return f"{note}\n\n{_md_table(rows)}", legend, shown


def _diagram_lines(sh: Any) -> list[tuple[int, str]]:
    """A SmartArt's text in its own order (the data part: its nodes under the document, by their order), each node with
    its depth."""
    from lxml import etree

    gd = sh._element.find(".//" + q("a:graphicData"))
    rel = gd.find(q("dgm:relIds")) if gd is not None else None
    rid = rel.get(q("r:dm")) if rel is not None else None
    if not rid:
        return []
    root = etree.fromstring(sh.part.related_part(rid).blob)
    pts: dict[str, tuple[str, str]] = {}
    order: list[str] = []
    for pt in root.iter(q("dgm:pt")):
        mid = pt.get("modelId") or ""
        body = pt.find(q("dgm:t"))
        text = _clean(" ".join("".join(t.text or "" for t in p.iter(q("a:t"))) for p in body.iter(q("a:p")))) if body is not None else ""
        pts[mid] = (pt.get("type") or "node", text)
        order.append(mid)
    kids: dict[str, list[tuple[int, str]]] = {}
    for cx in root.iter(q("dgm:cxn")):
        if (cx.get("type") or "parOf") == "parOf":
            kids.setdefault(cx.get("srcId") or "", []).append((int(cx.get("srcOrd") or 0), cx.get("destId") or ""))
    out: list[tuple[int, str]] = []
    seen: set[str] = set()

    def walk(mid: str, depth: int) -> None:
        for _, k in sorted(kids.get(mid, [])):
            if k in seen:
                continue
            seen.add(k)
            typ, text = pts.get(k, ("", ""))
            if typ != "node":
                continue
            if text:
                out.append((depth, text))
            walk(k, depth + 1 if text else depth)

    doc = next((m for m in order if pts[m][0] == "doc"), None)
    if doc is not None:
        walk(doc, 0)
    if not out:
        out = [(0, pts[m][1]) for m in order if pts[m][0] == "node" and pts[m][1]]
    return out


# ------------------------------------------------------------------ pictures


def _plain_panel(blob: bytes) -> Optional[bool]:
    """A picture that shows nothing: one flat colour (a card's ground, a rounded panel drawn as a PNG), nothing but
    transparency, or a frame (its opaque pixels only along the edges) — True: decoration of the template, never a
    picture the author put on the slide. False: line art on a transparent ground (a scheme, a signature: strokes inside
    it) — the author's. None: neither (the picture's kind decides, `_decoration`)."""
    try:
        from PIL import Image, ImageStat

        with Image.open(io.BytesIO(blob)) as im:
            im.thumbnail((128, 128))
            rgba = im.convert("RGBA")
        a = rgba.getchannel("A")
        if ImageStat.Stat(a).mean[0] < 40:
            w, h = a.size
            bx, by = max(1, w // 8), max(1, h // 8)
            inner = a.crop((bx, by, w - bx, h - by)) if w > 2 * bx and h > 2 * by else a
            if inner.getextrema()[1] < 100:
                return True  # nothing opaque inside the border: a frame, a glow or nothing at all
            solid = a.point(lambda v: 255 if v >= 100 else 0)
            box = solid.getbbox()
            if box is not None:
                filled = ImageStat.Stat(solid.crop(box)).mean[0] / 255.0
                if filled >= 0.75 and max(ImageStat.Stat(rgba.convert("RGB"), solid).stddev) < 24.0:
                    return True  # one plain shape on its shadow (a card drawn as a picture) — strokes are line art
            return False
        opaque = a.point(lambda v: 255 if v > 128 else 0)
        if max(ImageStat.Stat(rgba.convert("RGB"), opaque).stddev) < 6.0:
            return True  # one colour where it is opaque
        return None
    except Exception:  # noqa: BLE001 - a picture PIL cannot read: nothing here could show it either
        return True


def _decoration(blob: bytes, wf: float, hf: float) -> tuple[bool, bool]:
    """(the picture is the template's decoration, it is art) — read as the template analysis reads a sample's pictures
    (analysis/pictures.py): a panel, a line, an icon, or art cut out on transparency (a 3D render, a glow, a ring) is
    the template's; a photograph, a screenshot, a scheme with words, a picture on a ground of its own (a toned photo,
    an illustration) and line art on transparency are the author's. `wf`, `hf`: its size as fractions of the slide."""
    try:
        from verstka.analysis.pictures import _stats, classify_blob

        st = _stats(blob)
        kind = classify_blob(st, wf, hf, covers=False) if st is not None else "art"
    except Exception:  # noqa: BLE001 - the picture is kept
        st, kind = None, "photo"
    if kind in ("photo", "content"):
        return False, False  # (a chart with its labels on transparency is the author's whatever its ground)
    plain = _plain_panel(blob)
    if plain is not None:
        return plain, False
    if kind in ("panel", "line", "icon"):
        return True, False
    if kind == "art":
        return st is None or st.coverage < 0.85, True
    return False, False


def _bars_left_out(parts: list[_Part], W: float, H: float) -> list[_Part]:
    """The bars of a chart drawn as pictures (art of one thickness standing on one line, of different lengths — or two
    long ones of different lengths) are the deck's figures drawn, not pictures of the author's: no place of their own.
    Pictures of one size side by side (cards, a team) are no bars."""
    pics = [p for p in parts if p.kind == "picture" and p.art]
    drop: set[int] = set()
    for horizontal in (True, False):
        def length(p: _Part) -> float:
            return p.w if horizontal else p.h

        def thick(p: _Part) -> float:
            return p.h if horizontal else p.w

        long_ = [p for p in pics if length(p) >= 2.0 * thick(p)]
        rows: list[list[_Part]] = []
        for p in long_:
            # one baseline: the bars' left ends (a horizontal chart) or their feet (a column chart)
            base = p.x if horizontal else p.bottom
            row = next((r for r in rows if abs((r[0].x if horizontal else r[0].bottom) - base) <= 0.03 * (W if horizontal else H)
                        and abs(thick(r[0]) - thick(p)) <= 0.1 * thick(r[0])), None)
            if row is None:
                rows.append([p])
            else:
                row.append(p)
        for row in rows:
            if len(row) >= 3 and max(map(length, row)) >= 1.15 * min(map(length, row)):
                drop.update(id(p) for p in row)
        very = [p for p in long_ if length(p) >= 3.5 * thick(p)]
        if len(very) >= 2 and max(map(length, very)) >= 1.15 * min(map(length, very)):
            drop.update(id(p) for p in very)
    return [p for p in parts if id(p) not in drop]


def _repeats_left_out(parts: list[_Part]) -> list[_Part]:
    """One picture set several times on a slide (a sample's stand-in photo in every card) gets one place."""
    seen: set[str] = set()
    out = []
    for p in parts:
        if p.kind == "picture" and p.image is not None:
            d = hashlib.sha1(p.image[0]).hexdigest()
            if d in seen:
                continue
            seen.add(d)
        out.append(p)
    return out


def _cropped(blob: bytes, ext: str, rect: Any) -> tuple[bytes, str]:
    """The picture as the slide shows it: the crop of its frame («srcRect», thousandths of a percent) applied."""
    if rect is None:
        return blob, ext
    cut = [max(0, int(rect.get(k) or 0)) / 100000.0 for k in ("l", "t", "r", "b")]
    if not any(cut):
        return blob, ext
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as im:
            w, h = im.size
            box = (round(cut[0] * w), round(cut[1] * h), round(w * (1 - cut[2])), round(h * (1 - cut[3])))
            if box[2] - box[0] < 2 or box[3] - box[1] < 2:
                return blob, ext
            part = im.crop(box)
            buf = io.BytesIO()
            if ext.lower() in ("jpg", "jpeg") and part.mode in ("RGB", "L", "CMYK"):
                part.save(buf, "JPEG", quality=92)
                return buf.getvalue(), "jpg"
            if part.mode not in ("RGB", "RGBA", "L", "LA", "P"):
                part = part.convert("RGBA")
            part.save(buf, "PNG")
            return buf.getvalue(), "png"
    except Exception:  # noqa: BLE001 - the whole picture rather than none
        return blob, ext


def _is_vector(ext: str, ctype: str = "") -> bool:
    return (ext or "").lower() in ("emf", "wmf") or "emf" in (ctype or "") or "wmf" in (ctype or "")


def _convert_vectors(blobs: dict[str, tuple[bytes, str]], timeout_s: float = 60.0) -> dict[str, bytes]:
    """EMF / WMF pictures as PNG (LibreOffice, one run for the deck): the improved deck's picture tools decode neither.
    Whatever does not convert is missing from the answer."""
    if not blobs:
        return {}
    import os
    import subprocess
    import tempfile

    try:
        from verstka.ingest.render import find_soffice

        soffice = find_soffice()
    except Exception:  # noqa: BLE001
        soffice = None
    if not soffice:
        return {}
    out: dict[str, bytes] = {}
    try:
        with tempfile.TemporaryDirectory(prefix="verstka_vec_", ignore_cleanup_errors=True) as tmp:
            src, dst = Path(tmp) / "src", Path(tmp) / "png"
            src.mkdir()
            dst.mkdir()
            files = []
            for key, (blob, ext) in blobs.items():
                f = src / f"{key}.{ext.lower()}"
                f.write_bytes(blob)
                files.append(str(f))
            env = {**os.environ, "SAL_USE_VCLPLUGIN": "svp"}
            cmd = [soffice, f"-env:UserInstallation={(Path(tmp) / 'profile').as_uri()}", "--headless", "--norestore", "--convert-to", "png", "--outdir", str(dst), *files]
            subprocess.run(cmd, env=env, capture_output=True, timeout=timeout_s)
            for key in blobs:
                png = dst / f"{key}.png"
                if png.is_file() and png.stat().st_size > 0:
                    out[key] = png.read_bytes()
    except Exception:  # noqa: BLE001 - such a picture gets no place
        log.warning("EMF/WMF pictures not converted", exc_info=True)
    return out


def _picture_blob(sh: Any) -> Optional[tuple[bytes, str, str]]:
    """(bytes, extension, content type) of an embedded picture; None for a linked one, a video or a sound."""
    nv = sh._element.find(q("p:nvPicPr") + "/" + q("p:nvPr"))
    if nv is not None and any(local_name(c) in ("videoFile", "audioFile", "quickTimeFile", "wavAudioFile") for c in nv if isinstance(c.tag, str)):
        return None
    blip = sh._element.find(".//" + q("a:blip"))
    rid = blip.get(q("r:embed")) if blip is not None else None
    if not rid:
        return None  # a linked picture: its bytes are not in the deck
    try:
        part = sh.part.related_part(rid)
        ext = str(part.partname).rsplit(".", 1)[-1].lower() if "." in str(part.partname) else "png"
        return part.blob, ext, (part.content_type or "")
    except Exception:  # noqa: BLE001 - a relationship that is not there
        return None


# ------------------------------------------------------------------ what the deck repeats


class _Repeats:
    """What the deck's slides repeat — the template's chrome, not the author's text: a text at the same place (to a tenth
    of the slide) on three slides or more (and on at least three in ten), or in the top or bottom band on as many
    (wherever the layout moved it); a picture on as many slides, a small one (a logo) on two."""

    def __init__(self, slides: list, W: float, H: float) -> None:
        self.W, self.H = W, H
        self.need = max(3, int(0.3 * len(slides) + 0.5))
        self.at: Counter = Counter()
        self.band: Counter = Counter()
        self.pics: Counter = Counter()
        self.small: Counter = Counter()
        for s in slides:
            at, band, pics, small = set(), set(), set(), set()
            for sh, x, y, w, h in _flat(s.shapes):
                try:
                    tag = _tag(sh)
                    if tag == "sp" and sh.has_text_frame and _placeholder_kind(sh) not in ("title", "chrome"):
                        t = _norm(sh.text_frame.text)
                        if t and len(t) <= 120:
                            at.add(self._key(t, x, y, w, h))
                            if self._in_band(y, h):
                                band.add(t)
                    elif tag == "pic":
                        got = _picture_blob(sh)
                        if got is not None:
                            d = hashlib.sha1(got[0]).hexdigest()
                            pics.add(d)
                            if w * h < _SMALL_PICTURE * W * H:
                                small.add(d)
                except Exception:  # noqa: BLE001
                    continue
            self.at.update(at)
            self.band.update(band)
            self.pics.update(pics)
            self.small.update(small)

    def _key(self, t: str, x: float, y: float, w: float, h: float) -> tuple:
        return (t, round((x + w / 2) / self.W, 1), round((y + h / 2) / self.H, 1))

    def _in_band(self, y: float, h: float) -> bool:
        c = (y + h / 2) / self.H
        return c < 0.12 or c > 0.88

    def text(self, text: str, x: float, y: float, w: float, h: float, size: float) -> bool:
        t = _norm(text)
        band = self._in_band(y, h)
        if not t or not (self.at[self._key(t, x, y, w, h)] >= self.need or (band and self.band[t] >= self.need)):
            return False
        if not band and (_is_figure(text) or size >= 20):
            return False  # a figure, a heading in the slide's body: the author's, wherever it repeats
        return True

    def picture(self, digest: str, small: bool) -> bool:
        return self.pics[digest] >= self.need or (small and self.small[digest] >= 2)


# ------------------------------------------------------------------ a slide


def _contains(box: tuple, x: float, y: float, tol: float = 0.0) -> bool:
    return box[0] - tol <= x <= box[0] + box[2] + tol and box[1] - tol <= y <= box[1] + box[3] + tol


def _v_overlap(a: _Part, b: _Part) -> float:
    return max(0.0, min(a.bottom, b.bottom) - max(a.y, b.y))


def _h_overlap(a: _Part, b: _Part) -> float:
    return max(0.0, min(a.right, b.right) - max(a.x, b.x))


def _legend_left_out(parts: list[_Part], W: float, H: float) -> list[_Part]:
    """A chart's legend drawn as text boxes beside it repeats the chart note: its series and category names, a row of it
    read as one line («Кофе — 60%»: a name and the chart's own figure), and — with them — the chart's figures set as
    labels on their rows. Left out."""
    charts = [p for p in parts if p.kind == "chart" and p.names]
    if not charts:
        return parts
    drop: set[int] = set()
    for c in charts:
        names = {_norm(x) for x in c.names}
        mx, my = max(0.1 * W, 0.5 * c.w), max(0.1 * H, 0.3 * c.h)

        def near(p: _Part) -> bool:
            return p.x < c.right + mx and p.right > c.x - mx and p.y < c.bottom + my and p.bottom > c.y - my

        texts = [p for p in parts if p.kind in ("text", "list") and p.role not in ("title", "subtitle") and p.size < 20]
        named = [p for p in texts if near(p) and ((p.lines and all(_norm(ln.lstrip("- ")) in names for ln in p.lines))
                                                     or (p.label and _norm(p.label) in names and _chart_value(p.figure, c.values)))]
        drop.update(id(p) for p in named)
        for p in texts:
            if named and p.line and _chart_value(p.text, c.values) and (near(p) or any(_v_overlap(p, q) >= 0.5 * min(p.h, q.h) for q in named)):
                drop.add(id(p))
    return [p for p in parts if id(p) not in drop]


def _chart_value(text: str, values: list[float]) -> bool:
    """The text is a figure of the chart (one of its values, as a legend's label writes it: «60%», «315 000 ₽ · 40%»)."""
    t = (text or "").strip()
    if not t or not _is_figure(t):
        return False
    for m in re.finditer(r"\d[\d\u00a0\u202f ]*(?:[.,]\d+)?", t):
        try:
            v = float(re.sub(r"[\s\u00a0\u202f]", "", m.group(0)).replace(",", "."))
        except ValueError:
            continue
        if any(abs(v - x) <= 1e-6 * max(1.0, abs(x)) for x in values):
            return True
    return False


def _badges_left_out(parts: list[_Part], W: float, H: float) -> list[_Part]:
    """Step badges — boxes of a lone number that count 1, 2, 3 … (or 01, 02, 03 …) — number the cards beside them: left
    out (read alone they were paragraphs «1», «2», and «1» could become the slide's heading); so is the page number."""
    badges = [p for p in parts if p.line and _BADGE_RE.match(p.text) and p.w <= 0.15 * W and p.h <= 0.15 * H]
    by: dict[int, list[_Part]] = {}
    for p in badges:
        by.setdefault(int(_BADGE_RE.match(p.text).group(1)), []).append(p)
    run = []
    start = 0 if 0 in by else 1
    while start + len(run) in by and len(by[start + len(run)]) == 1:
        run.append(by[start + len(run)][0])
    gone = {id(p) for p in run} if len(run) >= 2 else set()
    return [p for p in parts if id(p) not in gone and p.role != "page"]


def _title_of(parts: list[_Part], W: float, H: float) -> tuple[Optional[_Part], list[_Part]]:
    """The slide's heading and the parts left: the title placeholder's text — unless a clearly larger line high on the
    slide is the real heading (or, for a kicker of one or two capitalised words, a larger line right under it); without
    one, the largest line high on the slide (a word or two only when it is the largest there). A figure is never the
    heading; a title that gives way stays in the body."""
    placed = next((p for p in parts if p.role == "title"), None)
    tops = [p for p in parts if p.line and p.role != "title" and p.y < 0.4 * H and _words(p.text) <= 16 and not _is_figure(p.text)
            and not p.text.startswith(("[", "|", "`"))]
    if placed is not None and _is_figure(placed.text):
        placed.role = ""  # («1 мин 18 с» set as a title is a figure of the slide)
        placed = None
    if placed is not None:
        kicker = _words(placed.text) <= 3 and placed.text.upper() == placed.text and any(ch.isalpha() for ch in placed.text)
        under = [p for p in tops if p.size > placed.size and p.y >= placed.y + 0.5 * placed.h and p.y - placed.bottom <= 0.12 * H
                 and (_h_overlap(p, placed) > 0 or abs(p.x - placed.x) <= 0.05 * W)]
        larger = [p for p in tops if p.size >= 1.3 * placed.size and _words(p.text) >= 3 and p.y < 0.35 * H]
        heading = min(under, key=lambda p: p.y) if kicker and under else (max(larger, key=lambda p: (p.size, -p.y)) if larger else None)
        if heading is None:
            return placed, [p for p in parts if p is not placed]
        placed.role = ""  # the kicker («ПРОБЛЕМА») stays on the slide, over its heading
        return heading, [p for p in parts if p is not heading]
    if not tops:
        # nothing high on the slide: a cover's or a divider's title set large in its middle («Marketing Report» at
        # 185 pt) — the largest line when it is clearly the largest text of the slide
        lines = [p for p in parts if p.line and _words(p.text) <= 16 and not _is_figure(p.text) and not p.text.startswith(("[", "|", "`"))]
        texts = [p for p in parts if p.kind in ("text", "list", "code")]
        big = max(lines, key=lambda p: (p.size, -p.y), default=None)
        if big is not None and all(big.size >= 1.2 * p.size for p in texts if p is not big):
            return big, [p for p in parts if p is not big]
        return None, parts
    top = max(p.size for p in tops)
    best = sorted(tops, key=lambda p: (-p.size, p.y, p.x))
    pick = best[0]
    if _words(pick.text) <= 2 and sum(1 for p in tops if p.size >= pick.size) > 1:
        # a word or two heads the slide only as its largest line («Команда» over «Шесть человек, три города»)
        longer = [p for p in best if _words(p.text) >= 3 and p.size >= top]
        pick = longer[0] if longer else min((p for p in tops if p.size >= top), key=lambda p: (p.y, p.x))
    return pick, [p for p in parts if p is not pick]


def _paired(parts: list[_Part], W: float, H: float, panel_of: Optional[dict] = None) -> list[_Part]:
    """A figure alone in its box with its label beside it — a KPI tile, a legend's row — is one line: «Средний чек —
    300 ₽»; read apart, the deck's figures lost what they measure. The label: on the figure's row to its left first
    («Насос и клапаны | 16 000 ₽»), else under it, else over it — close, in the same card; a pairing that is not clear
    by the layout is never made."""
    panel_of = panel_of or {}
    cands: list[tuple[tuple, int, int]] = []

    def between(a: _Part, b: _Part) -> bool:
        """Another box lies between the figure and the label in the figure's column: the pairing is not clear."""
        top, bottom = (a.bottom, b.y) if b.y >= a.y else (b.bottom, a.y)
        return any(c is not a and c is not b and c.kind in ("text", "list") and _h_overlap(c, a) > 0 and c.y + c.h / 2 > top and c.y + c.h / 2 < bottom
                   for c in parts)

    for i, a in enumerate(parts):
        if not (a.line and _is_figure(a.text)):
            continue
        for j, b in enumerate(parts):
            card = panel_of.get(id(a))
            if i == j or card != panel_of.get(id(b)) or not _label_like(b, 12 if card is not None else 8):
                continue
            if _v_overlap(a, b) >= 0.5 * min(a.h, b.h) and b.x < a.x and -0.02 * W <= a.x - b.right <= 0.08 * W:
                cands.append(((0, a.x - b.right), i, j))
                continue
            column = _h_overlap(a, b) >= 0.5 * min(a.w, b.w) or abs(a.x - b.x) <= 0.03 * W
            if not column:
                continue
            d_below, d_above = b.y - a.bottom, a.y - b.bottom
            if b.y >= a.y and -0.02 * H <= d_below <= 0.06 * H and not between(a, b):
                cands.append(((1, d_below), i, j))
            elif b.y < a.y and -0.02 * H <= d_above <= 0.06 * H and not between(a, b):
                cands.append(((2, d_above), i, j))
    used: set[int] = set()
    joined: list[_Part] = []
    for _, i, j in sorted(cands, key=lambda c: c[0]):
        if i in used or j in used:
            continue
        a, b = parts[i], parts[j]
        used.update((i, j))
        x, y = min(a.x, b.x), min(a.y, b.y)
        joined.append(_Part(text=f"{H_strip(b.text)} — {a.text}", x=x, y=y, w=max(a.right, b.right) - x, h=max(a.bottom, b.bottom) - y,
                            size=max(a.size, b.size), lines=[f"{H_strip(b.text)} — {a.text}"], label=H_strip(b.text), figure=a.text))
        if id(a) in panel_of:
            panel_of[id(joined[-1])] = panel_of[id(a)]
    return [p for k, p in enumerate(parts) if k not in used] + joined


def H_strip(text: str) -> str:
    return text.rstrip(" .:;,")


# ------------------------------------------------------------------ reading order


@dataclass
class _Node:
    box: tuple
    kids: list = field(default_factory=list)


def _box_of(item: Any) -> tuple:
    return item.box if isinstance(item, _Node) else (item.x, item.y, item.w, item.h)


def _split(items: list, axis: int, dim: float) -> tuple[list[list], float]:
    """The items in bands along one axis (0: x, 1: y), cut where nothing crosses — a small overlap of two boxes does not
    join them — and how wide its cuts are on average (a share of the slide; nothing where boxes overlap a little)."""
    def span(it: Any) -> tuple[float, float]:
        b = _box_of(it)
        return (b[axis], b[axis] + b[axis + 2])

    order = sorted(items, key=lambda it: span(it))
    groups: list[list] = []
    end = None
    last_len = 0.0
    gaps: list[float] = []
    for it in order:
        s0, s1 = span(it)
        ln = s1 - s0
        slack = min(0.02 * dim, max(0.003 * dim, 0.25 * min(ln, last_len) if last_len else 0.003 * dim))
        if end is None or s0 >= end - slack:
            if end is not None:
                gaps.append(max(0.0, s0 - end) / dim)
            groups.append([it])
            end, last_len = s1, ln
        else:
            groups[-1].append(it)
            if s1 > end:
                end, last_len = s1, ln
    return groups, (sum(gaps) / len(gaps) if gaps else 0.0)


def _xycut(items: list, W: float, H: float) -> list:
    """Reading order by cuts along the gaps of the slide — the wider cuts first (rows when both are as wide): the
    columns of a slide with a gutter, each person of a team (photo, name, role) whole — and so on down."""
    if len(items) <= 1:
        return items
    rows, row_gap = _split(items, 1, H)
    cols, col_gap = _split(items, 0, W)
    cuts = [(gap, k, g) for k, (g, gap) in enumerate(((rows, row_gap), (cols, col_gap))) if len(g) > 1]
    if cuts:
        _, _, groups = max(cuts, key=lambda c: (round(c[0], 3), -c[1]))
        return [x for g in groups for x in _xycut(g, W, H)]
    return sorted(items, key=lambda it: (_box_of(it)[1], _box_of(it)[0]))


def _ordered(parts: list[_Part], panels: list[tuple], W: float, H: float) -> tuple[list[_Part], dict]:
    """The parts in reading order: the boxes of one card (a panel shape or a group they lie in) together, the cards and
    the free boxes by the cuts of the slide (_xycut). Also each part's card (for the pairing of a KPI tile)."""
    area = W * H
    boxes = sorted({b for b in panels if 0.004 * area <= b[2] * b[3] <= 0.45 * area}, key=lambda b: b[2] * b[3])
    nodes = [_Node(box=b) for b in boxes]
    root = _Node(box=(0.0, 0.0, W, H))

    def holder(x: float, y: float, start: int = 0) -> _Node:
        return next((nd for nd in nodes[start:] if _contains(nd.box, x, y, 0.004 * W)), root)

    def inside(b: tuple, o: tuple) -> bool:
        ow = max(0.0, min(b[0] + b[2], o[0] + o[2]) - max(b[0], o[0]))
        oh = max(0.0, min(b[1] + b[3], o[1] + o[3]) - max(b[1], o[1]))
        return ow * oh >= 0.8 * b[2] * b[3]

    for k, nd in enumerate(nodes):
        parent = next((o for o in nodes[k + 1 :] if inside(nd.box, o.box)), root)
        parent.kids.append(nd)
    panel_of: dict = {}
    for p in parts:
        nd = holder(p.x + p.w / 2, p.y + p.h / 2)
        nd.kids.append(p)
        if nd is not root:
            panel_of[id(p)] = id(nd)

    def count(nd: _Node) -> int:
        return sum(count(k) if isinstance(k, _Node) else 1 for k in nd.kids)

    def flatten(nd: _Node) -> list[_Part]:
        items = [k for k in nd.kids if not isinstance(k, _Node) or count(k)]
        out: list[_Part] = []
        for it in _xycut(items, W, H):
            out.extend(flatten(it) if isinstance(it, _Node) else [it])
        return out

    return flatten(root), panel_of


def _panel_box(sh: Any, x: float, y: float, w: float, h: float) -> Optional[tuple]:
    """The box of a shape that may hold a card's boxes: an autoshape without text of its own (not a text box, not a
    line, not a placeholder) of some size."""
    try:
        el = sh._element
        if local_name(el) != "sp" or w <= 0 or h <= 0 or getattr(sh, "is_placeholder", False):
            return None
        nv = el.find(q("p:nvSpPr") + "/" + q("p:cNvSpPr"))
        if nv is not None and nv.get("txBox") in ("1", "true"):
            return None
        if "".join(t.text or "" for t in el.iter(q("a:t"))).strip():
            return None
        geom = el.find(q("p:spPr") + "/" + q("a:prstGeom"))
        if geom is not None and (geom.get("prst") or "").startswith(("line", "straightConnector")):
            return None
        return (x, y, w, h)
    except Exception:  # noqa: BLE001
        return None


# ------------------------------------------------------------------ the deck


def _deck_code(pptx: Path) -> str:
    try:
        return hashlib.sha1(Path(pptx).read_bytes()).hexdigest()[:6]
    except OSError:
        return hashlib.sha1(str(pptx).encode("utf-8")).hexdigest()[:6]


_MACHINE_MARK_RE = re.compile(r"\[verstka:[A-Za-z0-9_\-]+\]")


def _notes_of(slide: Any) -> list[str]:
    """The slide's speaker notes, paragraph by paragraph (a page number and Verstka's own markers left out)."""
    try:
        if not slide.has_notes_slide:
            return []
        tf = slide.notes_slide.notes_text_frame
        if tf is None:
            return []
        return [t for t in (_clean(_MACHINE_MARK_RE.sub(" ", p.text)) for p in tf.paragraphs) if t and not re.fullmatch(r"\d{1,4}", t)]
    except Exception:  # noqa: BLE001
        return []


def _new_notes(notes: list[str], shown: str) -> str:
    """The speaker notes the slide does not already say (a Verstka deck keeps the slide's whole text there): each
    paragraph whose every sentence is on the slide left out; brackets made round, so the note stays one note."""
    seen = _flat_words(shown)
    keep = []
    for para in notes:
        sents = [s for s in re.split(r"(?<=[.!?…])\s+", para) if s.strip()]
        if sents and all(_flat_words(s) and _flat_words(s) in seen for s in sents):
            continue
        keep.append(para)
    return " ".join(keep).replace("[", "(").replace("]", ")").strip()


def _flat_words(text: str) -> str:
    return " ".join(re.findall(r"[\w%₽$€]+", (text or "").lower().replace("ё", "е")))


def _quoted(text: str) -> str:
    t = H_strip(text.strip())
    return t if "«" in t or "»" in t else f"«{t}»"


def deck_text(pptx: Path | str) -> DeckText:
    """The deck as a «Слайд N» brief (see the module's doc)."""
    from pptx import Presentation

    prs = Presentation(str(pptx))
    W, H = float(prs.slide_width or 1), float(prs.slide_height or 1)
    code = _deck_code(Path(pptx))
    every = list(enumerate(prs.slides, 1))
    shown = [(n, s) for n, s in every if (s._element.get("show") or "1") not in ("0", "false")]  # hidden slides are not the deck's
    repeats = _Repeats([s for _, s in shown], W, H)
    vectors = _vector_pictures(shown, repeats, W, H)
    out = DeckText(brief="")
    sections: list[str] = []
    for n, s in shown:
        try:
            got = _read_slide(n, s, W, H, repeats, vectors, code, out, first=n == shown[0][0])
        except Exception:  # noqa: BLE001 - a slide that cannot be read is left out, never the deck
            log.warning("slide %d of %s not read", n, pptx, exc_info=True)
            continue
        if got:
            sections.append(got)
    out.slides = len(sections)
    out.brief = "\n\n".join(sections).strip() + "\n"
    return out


def _vector_pictures(slides: list, repeats: _Repeats, W: float, H: float) -> dict[str, bytes]:
    """The deck's EMF / WMF pictures that may get a place, converted to PNG in one run."""
    todo: dict[str, tuple[bytes, str]] = {}
    for _, s in slides:
        for sh, x, y, w, h in _flat(s.shapes):
            try:
                if _tag(sh) != "pic" or not (_TINY_PICTURE * W * H <= w * h <= _BACKGROUND * W * H):
                    continue
                got = _picture_blob(sh)
                if got is not None and _is_vector(got[1], got[2]):
                    d = hashlib.sha1(got[0]).hexdigest()
                    if not repeats.picture(d, w * h < _SMALL_PICTURE * W * H):
                        todo[d] = (got[0], got[1] or "emf")
            except Exception:  # noqa: BLE001
                continue
    return _convert_vectors(todo) if todo else {}


def _read_slide(n: int, s: Any, W: float, H: float, repeats: _Repeats, vectors: dict, code: str, out: DeckText, first: bool) -> str:
    """One slide's section of the brief ("" when it says nothing)."""
    parts: list[_Part] = []
    groups: list[tuple] = []
    panels: list[tuple] = []
    for sh, x, y, w, h in _flat(s.shapes, groups=groups):
        try:
            kind = _placeholder_kind(sh)
            if kind == "chrome":
                continue
            tag = _tag(sh)
            if tag == "sp":
                box = _panel_box(sh, x, y, w, h)
                if box is not None:
                    panels.append(box)
            if tag in ("sp",) and getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                part = _text_part(sh, kind, x, y, w, h, repeats, H)
                if part is not None:
                    parts.append(part)
            elif tag == "graphicFrame":
                part = _frame_part(sh, x, y, w, h, out)
                if part is not None:
                    parts.append(part)
            elif tag == "pic":
                part = _picture_part(sh, x, y, w, h, W, H, repeats, vectors)
                if part is not None:
                    parts.append(part)
        except Exception:  # noqa: BLE001 - a shape that cannot be read is left out, never the slide
            log.debug("shape on slide %d not read", n, exc_info=True)
    parts = _legend_left_out(parts, W, H)
    parts = _badges_left_out(parts, W, H)
    parts = _repeats_left_out(_bars_left_out(parts, W, H))
    heading, parts = _title_of(parts, W, H)
    title = heading.text if heading is not None else ""
    notes = _notes_of(s)
    if first and title and _cover_like(title, parts):
        section = _cover(n, title, heading, parts, panels + groups, W, H, code, out)
        out.title = title
    else:
        if first and title:
            out.title = title
        subtitle = [p for p in parts if p.role == "subtitle"]
        rest = [p for p in parts if p.role != "subtitle"]
        _, panel_of = _ordered(rest, panels + groups, W, H)
        rest = _legend_left_out(_paired(rest, W, H, panel_of), W, H)
        rest, _ = _ordered(rest, panels + groups, W, H)
        _placed(rest, n, code, out)
        body = [p.text for p in subtitle] + [p.text for p in rest]
        if not title and not body:
            return ""
        section = "\n\n".join([f"Слайд {n}. {title}".rstrip(". ").rstrip() if title else f"Слайд {n}.", *body])
    said = _new_notes(notes, section)
    if said:
        section += f"\n\n[{SPEAKER_NOTES}: {said}]"
    return section


_AGENDA_RE = re.compile(r"^(?:план|содержание|повестка|оглавление|agenda|contents)\b", re.I)


def _cover_like(title: str, parts: list[_Part]) -> bool:
    """The first slide is the deck's cover: a few short lines under its title (and pictures) — no table, chart, code,
    list of its own or paragraph of prose; never an agenda."""
    if _AGENDA_RE.match(title or "") or any(p.kind in ("table", "chart", "code", "diagram") for p in parts):
        return False
    texts = [p for p in parts if p.kind in ("text", "list")]
    if any(p.kind == "list" and len(p.lines) >= (3 if p.bulleted else 5) for p in texts):
        return False
    if any(len(re.findall(r"[.!?…](?:\s|$)", x)) >= 2 and _words(x) >= 12 for p in texts for x in (p.lines or [p.text])):
        return False  # a paragraph of prose: the slide's text, not a cover's lines («Москва. 2026.» is)
    return sum(_words(p.text) for p in texts) <= 50 and all(_words(x) <= 30 for p in texts for x in (p.lines or [p.text]))


def _cover(n: int, title: str, heading: _Part, parts: list[_Part], panels: list, W: float, H: float, code: str, out: DeckText) -> str:
    """The cover in the brief's own cover format: «Слайд 1. Титульный» / «Название: …» / «Подзаголовок: …», its other
    lines after them (planning/faithful.py reads the fields; a cover named «Итоги …» is no first content slide)."""
    texts = [p for p in parts if p.kind in ("text", "list")]
    ordered, _ = _ordered(parts, panels, W, H)
    _placed(ordered, n, code, out)
    sub = next((p for p in texts if p.role == "subtitle"), None)
    if sub is None:
        under = [p for p in texts if p.y >= heading.y and _words(" ".join(p.lines or [p.text])) <= 25]
        sub = min(under, key=lambda p: (p.y - heading.bottom if p.y >= heading.bottom - 0.02 * H else H, p.x)) if under else None
    lines = [f"Слайд {n}. Титульный", f"Название: {_quoted(title)}."]
    if sub is not None:
        lines.append(f"Подзаголовок: {_quoted(' · '.join(H_strip(x) for x in (sub.lines or [sub.text])))}.")
    rest = []
    for p in ordered:
        if p is sub:
            continue
        if p.kind in ("text", "list"):
            rest.extend(p.lines or [p.text])
        else:
            rest.append(p.text)
    return "\n".join(lines) + "".join(f"\n\n{x}" for x in rest)


def _text_part(sh: Any, kind: Optional[str], x: float, y: float, w: float, h: float, repeats: _Repeats, H: float) -> Optional[_Part]:
    styles = _styles(sh)
    paras = _paragraphs(sh, styles)
    if not paras:
        return None
    size = max(_sizes(sh, styles), default=18.0)
    joined = " ".join(t for _, t, _ in paras)
    if kind == "title":
        return _Part(text=joined, x=x, y=y, w=w, h=h, size=size, lines=[joined], role="title")
    glyph = len(joined) <= 3 and not re.search(r"\w", joined)
    if glyph or _OPERATOR_RE.match(joined) or repeats.text(sh.text_frame.text, x, y, w, h, size):
        return None  # a footer, the company's name, a mark («✓», the «×» of a formula's tiles)
    if re.fullmatch(r"\d{1,3}", joined) and h < 0.08 * H and (y < 0.1 * H or y + h > 0.88 * H) and size <= 20:
        # a page number (small, at an edge) — unless it is one of the slide's step badges (`_badges_left_out`)
        return _Part(text=joined, x=x, y=y, w=w, h=h, size=size, lines=[joined], role="page")
    lines = [t for _, t, _ in paras]
    if kind == "subtitle":
        return _Part(text=" · ".join(lines) if len(lines) > 1 else joined, x=x, y=y, w=w, h=h, size=size, lines=lines, role="subtitle")
    if _is_code(sh):
        return _Part(text=_code_text(sh), x=x, y=y, w=w, h=h, kind="code", size=size)
    if len(paras) == 1:
        return _Part(text=paras[0][1], x=x, y=y, w=w, h=h, size=size, lines=lines)
    k, text = _text_block(paras)
    return _Part(text=text, x=x, y=y, w=w, h=h, kind=k, size=size, lines=lines, bulleted=any(bu for _, _, bu in paras))


def _frame_part(sh: Any, x: float, y: float, w: float, h: float, out: DeckText) -> Optional[_Part]:
    """A table, a chart or a SmartArt."""
    gd = sh._element.find(".//" + q("a:graphicData"))
    uri = gd.get("uri") if gd is not None else ""
    if getattr(sh, "has_table", False) and sh.has_table:
        md = _md_table([[c.text for c in r.cells] for r in sh.table.rows])
        if md:
            out.tables += 1
            return _Part(text=md, x=x, y=y, w=w, h=h, kind="table")
        return None
    if getattr(sh, "has_chart", False) and sh.has_chart:
        try:
            space = sh.chart_part._element
        except Exception:  # noqa: BLE001 - a chart whose part is missing (a dangling relationship)
            return None
        got = _chart_note(space)
        if got is None:
            return None
        out.charts += 1
        note, names, values = got
        return _Part(text=note, x=x, y=y, w=w, h=h, kind="chart", names=names, values=values)
    if uri == _DIAGRAM_URI:
        items = _diagram_lines(sh)
        if not items:
            return None
        if len(items) == 1:
            return _Part(text=items[0][1], x=x, y=y, w=w, h=h, kind="text", lines=[items[0][1]])
        low = min(d for d, _ in items)
        return _Part(text="\n".join(("  " * (d - low)) + "- " + t for d, t in items), x=x, y=y, w=w, h=h, kind="diagram", lines=[t for _, t in items])
    return None  # an OLE object, a video frame: nothing of the author's text


def _picture_part(sh: Any, x: float, y: float, w: float, h: float, W: float, H: float, repeats: _Repeats, vectors: dict) -> Optional[_Part]:
    area = w * h
    if area < _TINY_PICTURE * W * H or area > _BACKGROUND * W * H:
        return None  # an icon, a mark — or the slide's background art
    got = _picture_blob(sh)
    if got is None:
        return None
    blob, ext, ctype = got
    digest = hashlib.sha1(blob).hexdigest()
    if repeats.picture(digest, area < _SMALL_PICTURE * W * H):
        return None  # a logo on every slide
    if _is_vector(ext, ctype):
        png = vectors.get(digest)
        if png is None:
            return None  # an EMF / WMF picture nothing here can show: no place for it
        blob, ext = png, "png"
    else:
        blob, ext = _cropped(blob, ext, sh._element.find(".//" + q("a:srcRect")))
    decor, art = _decoration(blob, w / W, h / H)
    if decor:
        return None  # a flat panel of the template (a card's ground drawn as a picture), a frame, a render cut out as art
    return _Part(text="", x=x, y=y, w=w, h=h, kind="picture", image=(blob, ext or "png"), art=art)


def _placed(parts: list[_Part], n: int, code: str, out: DeckText) -> None:
    """Each picture of the slide (in reading order) gets its description — the place's hint — and its file is kept."""
    k = 0
    for p in parts:
        if p.kind == "picture" and p.image is not None:
            k += 1
            hint = _PICTURE_HINT.format(k=f" {k}" if k > 1 else "", n=n, deck=code)
            p.text = f"[Место под изображение: {hint}]"
            out.pictures.append((hint, p.image[0], p.image[1]))


def keep_pictures(text: DeckText, folder: Path) -> dict[str, str]:
    """The deck's pictures saved in `folder` (the template's workspace) with SOURCE_IMAGES: {hint: file}."""
    folder.mkdir(parents=True, exist_ok=True)
    index: dict[str, str] = {}
    for k, (hint, blob, ext) in enumerate(text.pictures, 1):
        name = f"picture-{k:02d}.{ext.lower().lstrip('.') or 'png'}"
        (folder / name).write_bytes(blob)
        index[hint] = name
    (folder.parent / SOURCE_IMAGES).write_text(json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")
    return index


def source_records(ws_root: Path, pic_dir: Path) -> list[dict]:
    """Picture records for the generation (images/fill.py `reuse`): each picture of the original deck, copied into the
    generation's pictures folder, goes to the place its description names — never generated again. The description
    carries the deck's code: the pictures of another improved deck never take these places."""
    import shutil

    idx_file = Path(ws_root) / SOURCE_IMAGES
    if not idx_file.is_file():
        return []
    try:
        index = json.loads(idx_file.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    if not isinstance(index, dict):
        return []
    pic_dir.mkdir(parents=True, exist_ok=True)
    out = []
    for hint, name in index.items():
        src = Path(ws_root) / "source_images" / str(name)
        if not src.is_file():
            continue
        dst = pic_dir / f"source-{name}"
        if not dst.exists():
            shutil.copy2(src, dst)
        out.append({
            "status": "generated", "file": dst.name, "hint": hint, "source": "original", "model": "original", "model_label": "исходная презентация",
            "aspect": "4:3", "key": f"source:{name}", "prompt": "", "style": "", "subject_by": "author", "seed": 0, "palette_match": None, "cost_rub": None,
        })
    return out
