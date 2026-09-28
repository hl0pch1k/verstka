"""Compose: lay out a content slide from the template's design system.

A sample slide of the template is a good source of style but a poor mould for new content: its cards are as tall as
its designer needed, its text as small as a placeholder allows. Here the content decides the geometry. The slide's
free area (below the heading, clear of chrome) is divided on a grid; text is measured with the template's fonts; sizes
come from the template's own type scale with readable minimums; cards, badges, dividers and accents take the
template's colours — and its card shape itself when a sample has one.
"""

from __future__ import annotations

from contextlib import contextmanager

import copy
import math
import re
from dataclasses import dataclass
from typing import Optional

from lxml import etree
from pptx.slide import Slide

from verstka.analysis.xmlns import q
from verstka.planning.heuristics import ABBR_END_RE, label_beside
from verstka.rendering.charts import add_chart
from verstka.rendering.fonts import figure_metrics_em, left_bearing_em, text_width_pt, wrap_lines
from verstka.rendering.tables import add_table
from verstka.schemas.common import EMU_PER_PT, Bbox, PatternKind, contrast_ratio, relative_luminance
from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideItem, TableData
from verstka.schemas.template import TemplateManifest
from verstka.ru import typeset

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


# ---------------------------------------------------------------------------------------------- text model


@dataclass
class Run:
    text: str
    size: float
    color: str
    bold: bool = False
    font: Optional[str] = None


@dataclass
class Para:
    runs: list[Run]
    align: str = "l"
    space_after: float = 0.0  # pt
    marker: Optional[str] = None  # bullet character
    marker_color: Optional[str] = None

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)

    @property
    def size(self) -> float:
        return max((r.size for r in self.runs), default=12.0)

    @property
    def bold(self) -> bool:
        return any(r.bold for r in self.runs)


def para_lines(p: Para, width_pt: float) -> int:
    font = p.runs[0].font if p.runs else None
    indent = p.size * 1.1 if p.marker else 0.0
    return max(len(wrap_lines(p.text, font, p.size, p.bold, max(width_pt - indent, 1.0))), 1)


def block_height_pt(paras: list[Para], width_pt: float, line: float) -> float:
    h = 0.0
    for i, p in enumerate(paras):
        h += para_lines(p, width_pt) * p.size * line
        if i < len(paras) - 1:
            h += p.space_after
    return h


def widest_word_pt(text: str, font: Optional[str], size: float, bold: bool) -> float:
    return max((text_width_pt(w, font, size, bold) for w in text.split()), default=0.0)


# ---------------------------------------------------------------------------------------------- kit


def _mix(fg: str, bg: str, share: float) -> str:
    a = [int(fg[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(bg[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{int(round(x * share + y * (1 - share))):02X}" for x, y in zip(a, b))


def palette_accents(tokens) -> list[str]:
    """The template's accent colours; a monochrome template (no accent of its own — stock theme accents are dropped by
    the analysis) accents with its own colours, the most saturated first, then the text colour and the neutrals —
    never a colour the template does not use."""
    acc = list(tokens.accents())
    if acc:
        return acc

    def sat(h: str) -> float:
        r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
        mx, mn = max(r, g, b), min(r, g, b)
        return 0.0 if mx == 0 else (mx - mn) / mx

    grounds = ("background", "surface")
    cands = [c for c in tokens.colors if c.hex and len(c.hex) == 6 and not any((r or "").startswith(grounds) for r in (c.roles or [c.role]))]
    cands.sort(key=lambda c: (-round(sat(c.hex), 1), 0 if (c.role or "").startswith("text") else 1, -(c.weight or 0.0)))
    return [c.hex.upper() for c in cands] or ["0077FF"]


@dataclass
class Colors:
    """Colours that read on one ground (the slide background or a card)."""

    ground: str
    text: str
    muted: str
    accent: str  # for large figures and markers (≥ 3:1)
    accent_text: str  # for small text in the accent colour (≥ 4.5:1)
    divider: str
    heading: str = ""  # titles of blocks (card titles, step titles): the template's own colour for them when it reads
    figure: str = ""  # figures (KPI values, indices)

    def __post_init__(self) -> None:
        self.heading = self.heading or self.text
        self.figure = self.figure or self.accent


@dataclass
class CardStyle:
    proto: Optional[etree._Element]  # a sample card shape to copy (keeps gradients, shadows, corner shape)
    fill: Optional[str]
    line: Optional[str]
    radius_emu: int
    colors: Colors  # colours inside the card


class Kit:
    """The template's design system as the composer needs it: sizes, colours, card style, grid steps."""

    def __init__(self, manifest: TemplateManifest, W: int, H: int, ground: str, card_proto: Optional[etree._Element] = None, heading_color: Optional[str] = None) -> None:
        self.manifest = manifest
        t = manifest.tokens
        typo = t.typography
        self.W, self.H = W, H
        self.hpt = H / EMU_PER_PT
        self.font = typo.primary_family
        self.line = typo.line_height
        self.spacing_pct = typo.line_spacing if abs(typo.line_spacing - 1.2) > 1e-6 else 1.0
        self.sizes = sorted({round(s.size_pt, 2) for s in typo.scale if s.size_pt >= 7} | {round(float(x), 2) for x in (typo.sizes_used or []) if x >= 7})
        hp = self.hpt
        # a sparse template (its scale completed by a derived ladder): its placeholder body (24–32 pt on an 11″ slide)
        # is not the size of dense composed content — the body keeps to 2.7–3.1 % of the slide height, the small size
        # to about 2.2 %
        self.sparse = bool(getattr(typo, "derived_sizes", None))
        body_t = max(typo.size_for("body", 14.0), 0.027 * hp)
        if self.sparse:
            body_t = min(body_t, 0.031 * hp)
        self.body = self.snap(body_t, 0.025 * hp, 0.037 * hp)
        self.small = self.snap(max(typo.size_for("small", self.body * 0.85), 0.022 * hp), 0.019 * hp, self.body - 0.4)
        if self.sparse:
            self.small = min(self.small, self.snap(0.022 * hp, 0.019 * hp, 0.025 * hp))
        self.h3 = self.snap(max(self.body * 1.18, 0.032 * hp), self.body * 1.07, 0.044 * hp)
        self.lead = self.snap(0.037 * hp, self.body * 1.12, 0.05 * hp)
        self.h2 = self.snap(0.046 * hp, self.lead * 1.08, 0.06 * hp)
        self.statement = self.snap(0.068 * hp, self.h2 * 1.1, 0.085 * hp)
        self.display = self.snap(0.09 * hp, 0.062 * hp, 0.13 * hp)
        self.mega = self.snap(0.2 * hp, max(self.display * 1.3, 0.11 * hp), 0.26 * hp)
        # the template's weight habits: VK decks set headings and figures in regular weight, the LCT deck in bold
        fam = typo.families[0] if typo.families else None
        heads_bold = max((st.weight_bold_share for st in typo.scale if st.role in ("h1", "h2", "display")), default=0.0)
        self.bold = bool((fam and fam.bold_share >= 0.06) or heads_bold >= 0.3)
        nums = [sl.style.bold for pt in manifest.patterns if not pt.reference for sl in pt.slots if sl.role.value == "number"]
        self.figure_bold = (sum(nums) > len(nums) / 2) if nums else self.bold  # figures weigh what the samples' figures weigh
        # grid steps
        self.gap = int(0.024 * W)  # between cards
        self.vgap = int(0.03 * H)  # between rows / blocks
        # colours
        self.accents = palette_accents(t)
        self.palette = [c.hex for c in t.colors]
        self.heading_color = heading_color
        self.head_size: Optional[float] = None  # the deck's heading size (set by the caller): statements stay under it
        self.colors = self.colors_on(ground, prefer=heading_color)
        self.card = self._card_style(ground, card_proto)

    # ---- sizes -------------------------------------------------------------------------------------------------
    def snap(self, target: float, lo: float, hi: float) -> float:
        """The template size closest to `target` inside [lo, hi]; outside the window the nearest template size when it
        is within a quarter of the target, else the target itself (rounded to half a point)."""
        inside = [s for s in self.sizes if lo - 0.05 <= s <= hi + 0.05]
        if inside:
            return min(inside, key=lambda s: (abs(s - target), -s))
        # nothing of the template's own in the window: the nearest size of the template that keeps the order of the
        # roles (never above hi — a «small» size must stay smaller than the body), else the target itself
        near = [s for s in self.sizes if abs(s - target) <= 0.25 * target and lo * 0.9 <= s <= hi + 0.05]
        if near:
            return min(near, key=lambda s: abs(s - target))
        return round(min(max(target, lo), hi) * 2) / 2

    def steps_down(self, size: float, floor: float) -> list[float]:
        """size and every template size below it down to floor (largest first)."""
        out = [size] + [s for s in reversed(self.sizes) if floor - 0.05 <= s < size - 0.05]
        return out

    _NICE = (40.0, 44.0, 48.0, 54.0, 60.0, 66.0, 72.0, 80.0, 88.0, 96.0, 110.0, 120.0, 140.0, 160.0)

    def figure_sizes(self, cap: float, floor: float) -> list[float]:
        """Sizes for figures, largest first: the template's own sizes plus display steps above them (a figure may be
        set larger than any text of the template — the audit's display rule)."""
        top = max(self.sizes) if self.sizes else 36.0
        cands = {s for s in self.sizes if floor - 0.05 <= s <= cap + 0.05} | {s for s in self._NICE if top < s <= cap + 0.05 and s >= floor}
        return sorted(cands, reverse=True) or [max(floor, min(cap, top))]

    def figure_cap(self, n: int, in_tiles: bool, strategy: str) -> float:
        cap = {1: 0.30, 2: 0.20, 3: 0.15, 4: 0.11}.get(n, 0.08) * self.hpt
        if in_tiles:
            cap = min(cap, 0.11 * self.hpt)
        return cap * {"visual": 1.2, "compact": 0.85}.get(strategy, 1.0)

    def steps_up(self, size: float, cap: float) -> list[float]:
        return [s for s in self.sizes if size + 0.05 < s <= cap + 0.05]

    # ---- colours -----------------------------------------------------------------------------------------------
    def _slot_colors(self, *roles: str) -> list[str]:
        """Colours the template gives a role in its samples, most frequent first."""
        from collections import Counter

        cnt: Counter = Counter()
        for p in self.manifest.patterns:
            if p.reference:
                continue
            for sl in p.slots:
                if sl.role.value in roles and sl.style.color_hex:
                    cnt[sl.style.color_hex.upper()] += 1
        return [c for c, _ in cnt.most_common()]

    def colors_on(self, ground: str, prefer: Optional[str] = None, inside_card: bool = False) -> Colors:
        """Colours for one ground. Running text prefers a neutral (black, white, greys): body copy in a brand colour
        reads as decoration. Headings of blocks and figures take the colours the template's samples give them."""
        t = self.manifest.tokens
        dark_ground = relative_luminance(ground) < 0.4

        def sat(hex_: str) -> float:
            r, g, b = (int(hex_[i : i + 2], 16) / 255 for i in (0, 2, 4))
            mx, mn = max(r, g, b), min(r, g, b)
            return 0.0 if mx == 0 else (mx - mn) / mx

        cands = [t.color_for("text.primary"), t.color_for("text.secondary")] + [c.hex for c in t.colors] + ["FFFFFF", "000000"]
        cands = [c for c in dict.fromkeys(c for c in cands if c)]
        strong = [c for c in cands if contrast_ratio(c, ground) >= 7.0] or [c for c in cands if contrast_ratio(c, ground) >= 4.5]
        neutral = [c for c in strong if sat(c) < 0.25]
        text = (neutral or strong or ["FFFFFF" if dark_ground else "000000"])[0]
        if prefer and not inside_card and contrast_ratio(prefer, ground) >= 4.5 and sat(prefer) < 0.25:
            text = prefer  # the slide heading is neutral: the body follows it
        palette = {c.hex.upper() for c in t.colors}
        own_text = t.color_for("text.primary")
        fallback = text.upper() in ("000000", "FFFFFF") and text.upper() not in palette
        if fallback and own_text and contrast_ratio(own_text, ground) >= 4.5:
            # plain black/white the template never uses: its own text colour when it reads (a deep brand-tinted
            # near-black is the template's body text, not decoration)
            text = own_text
        muted_c = [c for c in [t.color_for("text.secondary")] + [x.hex for x in t.colors if x.role and x.role.startswith("neutral")] if c and c != text]
        if fallback:
            muted_c = [c for c in muted_c if sat(c) < 0.25]
        muted = next((c for c in muted_c if 4.5 <= contrast_ratio(c, ground) and contrast_ratio(c, text) >= 1.25), text)
        acc_big = next((a for a in self.accents if contrast_ratio(a, ground) >= 3.0), None)
        acc_small = next((a for a in self.accents if contrast_ratio(a, ground) >= 4.5), None)
        accent = acc_big or text
        accent_text = acc_small or (acc_big if acc_big and contrast_ratio(acc_big, ground) >= 3.8 else text)
        div_c = [x.hex for x in t.colors if x.hex not in (ground,)]
        divider = next((c for c in sorted(div_c, key=lambda c: contrast_ratio(c, ground)) if 1.25 <= contrast_ratio(c, ground) <= 2.6), None)
        if divider is None:
            divider = _mix(text, ground, 0.22)
        head_roles = ("card_title",) if inside_card else ("card_title", "title")
        heading = next((c for c in self._slot_colors(*head_roles) if contrast_ratio(c, ground) >= 4.5), None)
        if not inside_card and prefer and contrast_ratio(prefer, ground) >= 4.5:
            heading = prefer
        figure = next((c for c in self._slot_colors("number") if contrast_ratio(c, ground) >= 3.0 and sat(c) >= 0.25), None) or accent
        return Colors(ground=ground, text=text, muted=muted, accent=accent, accent_text=accent_text, divider=divider, heading=heading or text, figure=figure)

    def on_accent(self, fill: str, fallback: str) -> str:
        """Text on an accent-filled shape (a badge, a result panel): white when it reads (3:1, the digits are large),
        else the fallback when it reads, else the best of the template's dark/light colours and black/white."""
        if contrast_ratio("FFFFFF", fill) >= 3.0:
            return "FFFFFF"
        if contrast_ratio(fallback, fill) >= 4.5:
            return fallback
        t = self.manifest.tokens
        cands = [c for c in ("FFFFFF", t.color_for("background.dark"), t.color_for("text.primary"), t.color_for("background.light"), "000000", fallback) if c]
        return max(cands, key=lambda c: contrast_ratio(c, fill))

    def _card_style(self, ground: str, proto: Optional[etree._Element]) -> CardStyle:
        card = self.manifest.components.card
        fill = line = None
        radius = int(0.12 * 914400)
        if proto is not None:
            fill = _shape_fill_hex(proto)
            line = _shape_line_hex(proto)
        elif card is not None:
            fill = card.fill_hex
            line = card.line_hex
            if card.radius is not None and card.width_frac and card.height_frac:
                side = min(card.width_frac * self.W, card.height_frac * self.H)
                radius = int(max(0.0, min(card.radius, 0.5)) * side)
        surface = self.manifest.tokens.color_for("surface")
        if proto is None:
            # a card must stand apart from its ground: its own fill, else the surface colour, else an outline
            if not fill or contrast_ratio(fill, ground) < 1.04:
                fill = surface if surface and 1.04 <= contrast_ratio(surface, ground) <= 6 else None
            if fill and contrast_ratio(fill, ground) > 12:
                fill = surface if surface and 1.04 <= contrast_ratio(surface, ground) <= 6 else fill
            if not fill:
                line = line if line and contrast_ratio(line, ground) >= 1.5 else self.colors.divider
            else:
                line = None
        inner = self.colors_on(fill or ground, prefer=self.heading_color if fill is None else None, inside_card=fill is not None)
        return CardStyle(proto=proto, fill=fill, line=line, radius_emu=radius, colors=inner)


def _solid_hex(fill_parent: Optional[etree._Element]) -> Optional[str]:
    if fill_parent is None:
        return None
    sf = fill_parent.find(q("a:solidFill"))
    if sf is not None:
        c = sf.find(q("a:srgbClr"))
        if c is not None:
            return (c.get("val") or "").upper() or None
    gf = fill_parent.find(q("a:gradFill"))
    if gf is not None:
        stops = gf.findall(".//" + q("a:srgbClr"))
        if stops:
            return (stops[0].get("val") or "").upper() or None
    return None


def _shape_fill_hex(el: etree._Element) -> Optional[str]:
    return _solid_hex(el.find(q("p:spPr")))


def _shape_line_hex(el: etree._Element) -> Optional[str]:
    spPr = el.find(q("p:spPr"))
    ln = spPr.find(q("a:ln")) if spPr is not None else None
    if ln is None or ln.find(q("a:noFill")) is not None:
        return None
    return _solid_hex(ln)


# ---------------------------------------------------------------------------------------------- primitives


class Canvas:
    """Adds shapes to a slide with fresh ids."""

    def __init__(self, slide: Slide, kit: Optional[Kit] = None, *, font: Optional[str] = None, spacing_pct: Optional[float] = None) -> None:
        self.slide = slide
        self.kit = kit
        self.font = font if font is not None else (kit.font if kit else None)
        self.spacing_pct = spacing_pct if spacing_pct is not None else (kit.spacing_pct if kit else 1.0)
        self.tree = slide._element.cSld.find(q("p:spTree"))
        ids = [int(el.get("id")) for el in slide._element.iter(q("p:cNvPr")) if (el.get("id") or "").isdigit()]
        self.next_id = (max(ids) + 1) if ids else 2
        self.added: list[etree._Element] = []
        self.W, self.H = (kit.W, kit.H) if kit is not None else _slide_size(slide)
        self.clamps: list[tuple[str, str]] = []  # (shape id, what was clamped): a last guard, never silent
        self._pending_clamp: Optional[str] = None

    def _id(self) -> int:
        """A fresh shape id: above every id of the slide — the slide may have grown by elements the canvas did not
        add (a chart or a table added through python-pptx takes the next free id of the slide)."""
        ids = [int(el.get("id")) for el in self.slide._element.iter(q("p:cNvPr")) if (el.get("id") or "").isdigit()]
        i = max(self.next_id, (max(ids) + 1) if ids else 2)
        self.next_id = i + 1
        if self._pending_clamp:
            self.clamps.append((str(i), self._pending_clamp))
            self._pending_clamp = None
        return i

    def _clamp(self, box: Bbox, name: str) -> Bbox:
        """Every box inside the slide: a coordinate past the slide is a defect (a file PowerPoint may repair); the
        box is cut to the slide and the cut is reported to the composer's warnings."""
        W, H = self.W, self.H
        if not W or not H:
            return box
        x = min(max(int(box.x), 0), W)
        y = min(max(int(box.y), 0), H)
        x2 = min(max(int(box.x + box.w), x), W)
        y2 = min(max(int(box.y + box.h), y), H)
        if (x, y, x2 - x, y2 - y) == (int(box.x), int(box.y), int(box.w), int(box.h)):
            return box
        self._pending_clamp = f"{name} clamped to the slide (x {box.x / W:.2f}, y {box.y / H:.2f}, w {box.w / W:.2f}, h {box.h / H:.2f})"
        return Bbox(x=x, y=y, w=x2 - x, h=y2 - y)

    def clamp_warnings(self) -> list[str]:
        """The clamps of shapes still on the slide (a composition pass that was undone does not count)."""
        live = {el.get("id") for el in self.slide._element.iter(q("p:cNvPr"))}
        return [msg for sid, msg in self.clamps if sid in live]

    def _append(self, el: etree._Element) -> etree._Element:
        """Content goes on top of the canvas — as python-pptx's own element classes, so that `slide.shapes` can read
        the slide while it is being composed (a chart looks for the unit caption written above it)."""
        from pptx.oxml import parse_xml

        el = parse_xml(etree.tostring(el))
        self.tree.append(el)
        self.added.append(el)
        return el

    # ---- text --------------------------------------------------------------------------------------------------
    def text(self, box: Bbox, paras: list[Para], *, anchor: str = "t", name: str = "Text", insets: tuple[int, int, int, int] = (0, 0, 0, 0)) -> etree._Element:
        box = self._clamp(box, name)
        sp = etree.Element(q("p:sp"))
        nv = etree.SubElement(sp, q("p:nvSpPr"))
        c = etree.SubElement(nv, q("p:cNvPr"))
        c.set("id", str(self._id()))
        c.set("name", f"{name} {c.get('id')}")
        etree.SubElement(nv, q("p:cNvSpPr")).set("txBox", "1")
        etree.SubElement(nv, q("p:nvPr"))
        spPr = etree.SubElement(sp, q("p:spPr"))
        _xfrm(spPr, box)
        geom = etree.SubElement(spPr, q("a:prstGeom"))
        geom.set("prst", "rect")
        etree.SubElement(geom, q("a:avLst"))
        etree.SubElement(spPr, q("a:noFill"))
        self._txbody(sp, paras, anchor, insets)
        return self._append(sp)

    def _txbody(self, sp: etree._Element, paras: list[Para], anchor: str, insets: tuple[int, int, int, int]) -> None:
        tx = etree.SubElement(sp, q("p:txBody"))
        bp = etree.SubElement(tx, q("a:bodyPr"))
        bp.set("wrap", "square")
        for n, v in zip(("lIns", "tIns", "rIns", "bIns"), insets):
            bp.set(n, str(int(v)))
        bp.set("anchor", anchor)
        bp.set("rtlCol", "0")
        etree.SubElement(bp, q("a:noAutofit"))
        etree.SubElement(tx, q("a:lstStyle"))
        for para in paras or [Para([Run("", 12.0, "000000")])]:
            p = etree.SubElement(tx, q("a:p"))
            pPr = etree.SubElement(p, q("a:pPr"))
            pPr.set("algn", para.align)
            ln = etree.SubElement(pPr, q("a:lnSpc"))
            etree.SubElement(ln, q("a:spcPct")).set("val", str(int(round(self.spacing_pct * 100000))))
            if para.space_after:
                sa = etree.SubElement(pPr, q("a:spcAft"))
                etree.SubElement(sa, q("a:spcPts")).set("val", str(int(round(para.space_after * 100))))
            if para.marker:
                hang = int(para.size * 1.1 * EMU_PER_PT)
                pPr.set("marL", str(hang))
                pPr.set("indent", str(-hang))
                if para.marker_color:
                    bc = etree.SubElement(pPr, q("a:buClr"))
                    etree.SubElement(bc, q("a:srgbClr")).set("val", para.marker_color)
                bf = etree.SubElement(pPr, q("a:buFont"))
                bf.set("typeface", "Arial")
                etree.SubElement(pPr, q("a:buChar")).set("char", para.marker)
            else:
                pPr.set("marL", "0")
                pPr.set("indent", "0")
                etree.SubElement(pPr, q("a:buNone"))
            for r in para.runs:
                if not r.text:
                    continue
                re_ = etree.SubElement(p, q("a:r"))
                rPr = etree.SubElement(re_, q("a:rPr"))
                rPr.set("lang", "ru-RU")
                rPr.set("sz", str(int(round(r.size * 100))))
                rPr.set("b", "1" if r.bold else "0")
                # never the template's inherited capitals, tracking or baseline shift (measured text is what is set)
                rPr.set("cap", "none")
                rPr.set("spc", "0")
                rPr.set("baseline", "0")
                rPr.set("dirty", "0")
                sf = etree.SubElement(rPr, q("a:solidFill"))
                etree.SubElement(sf, q("a:srgbClr")).set("val", r.color.upper())
                font = r.font or self.font
                if font:
                    etree.SubElement(rPr, q("a:latin")).set("typeface", font)
                    etree.SubElement(rPr, q("a:cs")).set("typeface", font)
                t = etree.SubElement(re_, q("a:t"))
                t.text = r.text
            end = etree.SubElement(p, q("a:endParaRPr"))
            end.set("lang", "ru-RU")
            end.set("sz", str(int(round(para.size * 100))))

    # ---- shapes ------------------------------------------------------------------------------------------------
    def rect(self, box: Bbox, fill: Optional[str], line: Optional[str] = None, radius_emu: int = 0, line_w_pt: float = 0.75, name: str = "Shape", geom: Optional[str] = None, alpha: Optional[float] = None) -> etree._Element:
        box = self._clamp(box, name)
        sp = etree.Element(q("p:sp"))
        nv = etree.SubElement(sp, q("p:nvSpPr"))
        c = etree.SubElement(nv, q("p:cNvPr"))
        c.set("id", str(self._id()))
        c.set("name", f"{name} {c.get('id')}")
        etree.SubElement(nv, q("p:cNvSpPr"))
        etree.SubElement(nv, q("p:nvPr"))
        spPr = etree.SubElement(sp, q("p:spPr"))
        _xfrm(spPr, box)
        prst = geom or ("roundRect" if radius_emu > 0 else "rect")
        g = etree.SubElement(spPr, q("a:prstGeom"))
        g.set("prst", prst)
        av = etree.SubElement(g, q("a:avLst"))
        if prst == "roundRect":
            side = max(min(box.w, box.h), 1)
            gd = etree.SubElement(av, q("a:gd"))
            gd.set("name", "adj")
            gd.set("fmla", f"val {int(min(radius_emu / side, 0.5) * 100000)}")
        if fill:
            sf = etree.SubElement(spPr, q("a:solidFill"))
            clr = etree.SubElement(sf, q("a:srgbClr"))
            clr.set("val", fill.upper())
            if alpha is not None and alpha < 1.0:
                etree.SubElement(clr, q("a:alpha")).set("val", str(int(round(max(alpha, 0.0) * 100000))))
        else:
            etree.SubElement(spPr, q("a:noFill"))
        ln = etree.SubElement(spPr, q("a:ln"))
        if line:
            ln.set("w", str(int(line_w_pt * EMU_PER_PT)))
            sf = etree.SubElement(ln, q("a:solidFill"))
            etree.SubElement(sf, q("a:srgbClr")).set("val", line.upper())
        else:
            etree.SubElement(ln, q("a:noFill"))
        return self._append(sp)

    def card(self, box: Bbox, style: CardStyle, name: str = "Card") -> etree._Element:
        if style.proto is not None:
            box = self._clamp(box, name)
            el = copy.deepcopy(style.proto)
            for tag in ("p:txBody",):
                for old in el.findall(q(tag)):
                    el.remove(old)
            nv = el.find(".//" + q("p:cNvPr"))
            if nv is not None:
                nv.set("id", str(self._id()))
                nv.set("name", f"{name} {nv.get('id')}")
                for k in ("descr", "title"):
                    nv.attrib.pop(k, None)
                for child in list(nv):
                    nv.remove(child)  # no hyperlinks of the sample
            spPr = el.find(q("p:spPr"))
            old_box = _xfrm_box(spPr)
            _xfrm(spPr, box, replace=True)
            geom = spPr.find(q("a:prstGeom"))
            if geom is not None and geom.get("prst") in ("roundRect", "snip1Rect", "snip2SameRect", "round2SameRect", "round1Rect") and old_box is not None:
                # the corner keeps its absolute size: the sample's ratio is relative to the sample's shorter side
                av = geom.find(q("a:avLst"))
                gds = av.findall(q("a:gd")) if av is not None else []
                old_side = max(min(old_box.w, old_box.h), 1)
                new_side = max(min(box.w, box.h), 1)
                # (capped at a quarter of the shorter side: a small card never turns into a capsule)
                for gd in gds:
                    m = re.match(r"val (\d+)", gd.get("fmla") or "")
                    if m:
                        gd.set("fmla", f"val {min(int(int(m.group(1)) * old_side / new_side), 25000)}")
                if not gds and geom.get("prst") == "roundRect":
                    # default roundRect corner is 16.667% of the sample's shorter side
                    if av is None:
                        av = etree.SubElement(geom, q("a:avLst"))
                    gd = etree.SubElement(av, q("a:gd"))
                    gd.set("name", "adj")
                    gd.set("fmla", f"val {min(int(16667 * old_side / new_side), 25000)}")
            # a group-transformed or rotated sample must not carry its rotation/flip
            xfrm = spPr.find(q("a:xfrm"))
            for a in ("rot", "flipH", "flipV"):
                xfrm.attrib.pop(a, None)
            return self._append(el)
        return self.rect(box, style.fill, style.line, style.radius_emu, name=name)

    def ellipse(self, box: Bbox, fill: str, name: str = "Badge") -> etree._Element:
        return self.rect(box, fill, None, 0, name=name, geom="ellipse")

    def line(self, x1: int, y1: int, x2: int, y2: int, color: str, width_pt: float = 1.0, name: str = "Line") -> etree._Element:
        cx = etree.Element(q("p:cxnSp"))
        nv = etree.SubElement(cx, q("p:nvCxnSpPr"))
        c = etree.SubElement(nv, q("p:cNvPr"))
        c.set("id", str(self._id()))
        c.set("name", f"{name} {c.get('id')}")
        etree.SubElement(nv, q("p:cNvCxnSpPr"))
        etree.SubElement(nv, q("p:nvPr"))
        spPr = etree.SubElement(cx, q("p:spPr"))
        xfrm = etree.SubElement(spPr, q("a:xfrm"))
        if y2 < y1:
            xfrm.set("flipV", "1")
        off = etree.SubElement(xfrm, q("a:off"))
        off.set("x", str(int(min(x1, x2))))
        off.set("y", str(int(min(y1, y2))))
        ext = etree.SubElement(xfrm, q("a:ext"))
        ext.set("cx", str(int(abs(x2 - x1))))
        ext.set("cy", str(int(abs(y2 - y1))))
        g = etree.SubElement(spPr, q("a:prstGeom"))
        g.set("prst", "line")
        etree.SubElement(g, q("a:avLst"))
        ln = etree.SubElement(spPr, q("a:ln"))
        ln.set("w", str(int(width_pt * EMU_PER_PT)))
        sf = etree.SubElement(ln, q("a:solidFill"))
        etree.SubElement(sf, q("a:srgbClr")).set("val", color.upper())
        return self._append(cx)


def _slide_size(slide: Slide) -> tuple[int, int]:
    try:
        prs = slide.part.package.presentation_part.presentation
        return int(prs.slide_width), int(prs.slide_height)
    except Exception:  # noqa: BLE001
        return 0, 0


def _xfrm(spPr: etree._Element, box: Bbox, replace: bool = False) -> None:
    xfrm = spPr.find(q("a:xfrm"))
    if xfrm is None:
        xfrm = etree.Element(q("a:xfrm"))
        spPr.insert(0, xfrm)
    elif replace:
        for child in list(xfrm):
            xfrm.remove(child)
    off = xfrm.find(q("a:off"))
    if off is None:
        off = etree.SubElement(xfrm, q("a:off"))
    off.set("x", str(int(box.x)))
    off.set("y", str(int(box.y)))
    ext = xfrm.find(q("a:ext"))
    if ext is None:
        ext = etree.SubElement(xfrm, q("a:ext"))
    ext.set("cx", str(max(int(box.w), 0)))
    ext.set("cy", str(max(int(box.h), 0)))


def _xfrm_box(spPr: Optional[etree._Element]) -> Optional[Bbox]:
    xfrm = spPr.find(q("a:xfrm")) if spPr is not None else None
    if xfrm is None:
        return None
    off, ext = xfrm.find(q("a:off")), xfrm.find(q("a:ext"))
    if off is None or ext is None:
        return None
    return Bbox(x=int(off.get("x")), y=int(off.get("y")), w=int(ext.get("cx")), h=int(ext.get("cy")))


# ---------------------------------------------------------------------------------------------- helpers

_FIG_RE = re.compile(r"(?<![\w])([+\-−–]?\d[\d\s\u00a0\u202f]*(?:[.,]\d+)?\s?(?:%|₽|×|(?:млн|млрд|тыс)\.?(?![а-яё])|рубл(?:ь|я|ей)(?![а-яё])|руб\.?(?![а-яё])|час(?:а|ов)?(?![а-яё])|ч\.?(?![а-яё])|мин(?:ут[аы]?|\.)?(?![а-яё])|сек(?:унд[аы]?|\.)?(?![а-яё])|дн(?:ей|я|\.)?(?![а-яё])|мес(?:яц(?:а|ев)?|\.)?(?![а-яё])|раз[а]?(?![а-яё]))?(?:\s?₽)?)", re.I)  # a unit is a whole word: «рублей», never «руб|лей»


def highlight_runs(text: str, size: float, color: str, accent: str, bold: bool = False, accent_bold: bool = True, font: Optional[str] = None) -> list[Run]:
    """Figures inside a statement («27 млн ₽», «4 GPU») set in the accent colour: the eye finds the ask."""
    out: list[Run] = []
    pos = 0
    text = typeset(text)
    for m in _FIG_RE.finditer(text):
        s, e = m.span(1)
        val = text[s:e].strip()
        if not any(ch.isdigit() for ch in val) or (len(val) <= 2 and not re.search(r"%|₽", val) and not re.match(r"\d{2}", val)):
            continue
        if s > pos:
            out.append(Run(text[pos:s], size, color, bold, font))
        out.append(Run(text[s:e], size, accent, accent_bold or bold, font))
        pos = e
    if pos < len(text):
        out.append(Run(text[pos:], size, color, bold, font))
    return out or [Run(text, size, color, bold, font)]


_HERO_FIG_RE = re.compile(r"(?<![\w])([+\-−–]?\d[\d\s\u00a0]*(?:[.,]\d+)?\s?(?:%|×|раз[а]?\b|(?:тыс\.?|млн|млрд)?\s?(?:₽|руб\.?|\$|€)|млн|млрд))", re.I)


def single_figure(text: str) -> Optional[tuple[str, str]]:
    """(figure, label) when a short sentence carries exactly one money/percent/multiple figure — the label is the
    sentence without it («Экономия в год на одну площадку»)."""
    t = " ".join(text.split())
    if len(t) > 90:
        return None
    found = [m for m in _HERO_FIG_RE.finditer(t) if any(ch.isdigit() for ch in m.group(1))]
    digits = re.findall(r"\d+(?:[.,]\d+)?", t)
    if len(found) != 1 or len(digits) > len(re.findall(r"\d+(?:[.,]\d+)?", found[0].group(1))):
        return None
    m = found[0]
    value = m.group(1).strip()
    before = t[: m.start(1)].rstrip()
    last = before.split()[-1].lower() if before.split() else ""
    if last in ("на", "в", "до", "с", "от", "за", "по", "около", "почти", "более", "менее"):
        before = before[: -len(last)]  # «рост на 34%» → «рост …»: the preposition leaves with its figure
    rest = (before + " " + t[m.end(1) :]).replace("—", " ").replace("–", " ")
    rest = " ".join(rest.split()).strip(" ,.:;")
    if len(rest.split()) < 2:
        return None
    return value, rest[:1].upper() + rest[1:]


_ORDINAL_RE = re.compile(r"(?<![\w])\d+\s?-\s?(?:й|я|е|го|му|м|х|ый|ой|ий|ая|ое|ые)(?![а-яё])", re.I)
_MONEY_WORD_RE = re.compile(r"(\d)\s?(?:рубл(?:ей|я|ь)|руб\.?)(?![а-яё])", re.I)
_PLAIN_FIG_RE = re.compile(r"(?<![\w.,:])\d{1,3}(?: \d{3})+(?:[.,]\d+)?(?![\w:])|(?<![\w.,:])\d+(?:[.,]\d+)?(?![\w:%])")
_LABEL_TAIL_PREP = ("на", "в", "до", "с", "со", "от", "за", "по", "около", "почти", "более", "менее", "—", "–", "-", ":")


def kpi_callout(text: str) -> Optional[tuple[str, str]]:
    """(figure, label) of a short side line that carries one figure — the figure a visual slide sets large beside its
    chart: «Ежемесячная выручка — 900 000 рублей» → («900 000 ₽», «Ежемесячная выручка»), «Прогноз: 1 138 500 ₽ на
    6-й месяц» → («1 138 500 ₽», «Прогноз на 6-й месяц»), «100 покупок в день» → («100», «Покупок в день»). An ordinal
    («6-й месяц») is a word, not a figure; a line with two figures, a time («15:00»), a year or a figure without words
    around it is no callout (None) and stays a line of text."""
    t = " ".join((text or "").replace("\xa0", " ").split()).rstrip(".")
    if not t or len(t) > 100:
        return None
    t = _MONEY_WORD_RE.sub(r"\1 ₽", t)
    masked = _ORDINAL_RE.sub(lambda m: "#" * len(m.group(0)), t)  # same length: positions stay valid in `t`
    digits = re.findall(r"\d+(?:[.,]\d+)?", masked)
    hero = [m for m in _HERO_FIG_RE.finditer(masked) if any(ch.isdigit() for ch in m.group(1))]
    if len(hero) == 1:
        s, e = hero[0].span(1)
    elif not hero:
        plain = list(_PLAIN_FIG_RE.finditer(masked))
        if len(plain) != 1:
            return None
        s, e = plain[0].span(0)
        if re.fullmatch(r"(?:19|20)\d\d", t[s:e].strip()):
            return None  # a year is a date, not a figure to set large
        if s and not re.search(r"[—–:]\s*$", t[:s]):
            return None  # a bare number after a word is an index («Месяц 2», «Этап 3»), not a figure
    else:
        return None
    if len(digits) != len(re.findall(r"\d+(?:[.,]\d+)?", masked[s:e])):
        return None
    value = " ".join(t[s:e].split())
    words = t[:s].split()
    while words and words[-1].lower().strip(":") in _LABEL_TAIL_PREP:
        words.pop()  # «рост на 34%» → «рост»: the preposition (and a dash before the figure) leaves with its figure
    before = " ".join(words).rstrip(" :—–-,")
    after = t[e:].strip(" ,;:—–-")
    label = " ".join(x for x in (before, after) if x)
    if not re.search(r"[A-Za-zА-Яа-яЁё]{3,}", label):
        return None
    return value, label[:1].upper() + label[1:]


_FIG_UNIT_RE = re.compile(r"^([+\-−–~≈×]?\s?\d[\d\s\u00a0]*(?:[.,]\d+)?\s?%?)\s*([A-Za-zА-Яа-яЁё₽$€].{0,12})$")


# a hedge before a figure — «около 615 тыс. тонн», «более 10», «почти 40%» — (the words, then the figure)
_HEDGE_RE = re.compile(r"^((?:не\s)?(?:около|более|менее|почти|свыше|порядка|примерно|приблизительно|до|от|ок\.|меньше|больше)(?:\sчем)?)\s+(?=[+\-−–~≈]?\d)", re.I)


_SHORT_UNIT_RE = re.compile(r"(?:%|‰|×|₽|\$|€|£|¥|тыс\.?|млн\.?|млрд\.?|трлн\.?|руб\.?|р\.|шт\.?|x|х)", re.I)


def figure_split(value: str) -> Optional[tuple[str, str]]:
    """(figure, words) of a figure followed by the words it counts — «27 миллионов человек» → («27», «миллионов
    человек»), «40 стран» → («40», «стран»), «12 млн ₽ выручки» → («12 млн ₽», «выручки»). The short units (%, ₽,
    млн) stay with the number. None when no words follow, for a change («31% → 12%») and for «4,6 из 5»."""
    t = " ".join((value or "").replace("\u00a0", " ").split())
    if not t or re.search(r"→|->|⟶", t):
        return None
    m = re.match(r"^([+\-−–~≈<>≤≥]?\s?\d[\d ]*(?:[.,]\d+)?)(.*)$", t)
    if not m:
        return None
    toks = m.group(2).split()
    head = [m.group(1).strip()]
    while toks and _SHORT_UNIT_RE.fullmatch(toks[0]):
        head.append(toks.pop(0))
    if not toks or not any(len(w) >= 3 and w[0].isalpha() for w in toks):
        return None
    return " ".join(head), " ".join(toks)


def distinct_label(value: str, label: str, headline: str) -> str:
    """A figure's label that does not repeat the heading above it: when the label is the heading's own sentence
    («оператор тратит в среднем 6,5 минуты» under «Оператор тратит в среднем 6,5 минуты на одно обращение»), the
    words around the figure in the heading say what it measures («минуты на одно обращение»). One rule shared with the
    planner (heuristics.label_beside): the figure is found as whole words («3 мес» is «3 месяца», never a piece of
    «13 месяцев»), so the plan says what the slide shows."""
    return label_beside(value, label, headline)


def _pt(emu: int) -> float:
    return emu / EMU_PER_PT


def _emu(pt: float) -> int:
    return int(round(pt * EMU_PER_PT))


# ---------------------------------------------------------------------------------------------- composer


class Composer:
    def __init__(self, slide: Slide, kit: Kit, oslide: OutlineSlide, outline: DeckOutline, strategy: str, manifest: TemplateManifest) -> None:
        self.slide = slide
        self.kit = kit
        self.cv = Canvas(slide, kit)
        self.o = oslide
        self.outline = outline
        self.strategy = strategy
        self.manifest = manifest
        self.warnings: list[str] = []

    # ---- measurement helpers ----------------------------------------------------------------------------------
    def h(self, paras: list[Para], width_emu: int) -> int:
        return _emu(block_height_pt(paras, _pt(width_emu), self.kit.line)) + _emu(2)

    def P(self, text: str, size: float, color: str, bold: bool = False, align: str = "l", space_after: float = 0.0, marker: Optional[str] = None, marker_color: Optional[str] = None) -> Para:
        # the run names the template font: measuring and rendering see the same face (Montserrat is 20% wider than Play)
        return Para([Run(typeset(text), size, color, bold, self.kit.font)], align=align, space_after=space_after, marker=marker, marker_color=marker_color)

    WORD_ROOM = 0.95  # a word may fill this much of its column: the renderer sets it a little wider than measured

    def fits_width(self, texts: list[str], size: float, bold: bool, width_emu: int) -> bool:
        """Every word of the texts fits the width at this size: a line breaks only on a real space, so the words a
        no-break space binds as typeset («по результатам», «180 000 ₽») are measured as one. A word that fills the
        column to the last point is broken by the renderer («Корректировк / а»): it keeps 5% of the width free."""
        w = _pt(width_emu) * self.WORD_ROOM
        return all(text_width_pt(word, self.kit.font, size, bold) <= w for t in texts for word in typeset(t).split(" ") if word)

    def _dense_floor_2pct(self) -> float:
        return max(0.02 * self.kit.hpt, 7.0)

    def _least_text(self, n0: int) -> float:
        """The smallest text size (pt) of what was composed after the first `n0` elements, the conclusion aside."""
        least = math.inf
        for el in list(self.cv.tree)[n0:]:
            nv = el.find(".//" + q("p:cNvPr"))
            if nv is not None and (nv.get("name") or "").startswith("Conclusion"):
                continue
            for r in el.iter(q("a:rPr")):
                if r.get("sz"):
                    least = min(least, int(r.get("sz")) / 100)
        return least

    def _dense_floor(self) -> float:
        """The least text size of a dense setting: 1.7 % of the slide height; 2 % while the block is set again to keep
        its conclusion on the slide (B3-5: 8 pt on a 6.2″ slide does not read from the back of the room)."""
        return max((0.02 if getattr(self, "_floor_2pct", False) else 0.017) * self.kit.hpt, 7.0)

    def _size_pairs(self, top: Optional[list[tuple[float, float]]] = None) -> list[tuple[float, float]]:
        """(title, text) sizes a block of columns, cards or steps tries, largest first: the roles of the type scale
        (heading 2 over lead, lead over lead, lead over body…) with every size the template uses between them — a
        block that does not fit at the lead size takes the next size of the template (18, 17, 16 pt) rather than the
        body size four points below, and fills its area. A title stays a visible step (≥ 1.1×) over its text."""
        k = self.kit
        base = top if top is not None else [(k.h2, k.lead), (k.lead, k.lead)]
        if top is None and getattr(self, "_boost", 0) >= 1:
            # the fill pass: a short block may climb above the usual cap — text up to the h2 size, titles up to the
            # statement size, a clear step over their text
            texts = [x for x in k.sizes if k.lead - 0.05 <= x <= k.h2 + 0.05]
            heads = [x for x in k.sizes if k.h2 - 0.05 <= x <= k.statement + 0.05]
            ext = [(ts, bs) for bs in texts for ts in heads if ts >= bs * 1.15 - 0.05]
            ext.sort(key=lambda p: (-p[1], -p[0]))
            base = ext + base
        mids = [x for x in reversed(k.sizes) if k.small - 0.05 <= x < k.lead - 0.05]
        tail: list[tuple[float, float]] = []
        for bs in dict.fromkeys(mids + [k.body, k.small]):
            for ts in (k.lead, k.h3, k.body):
                if ts >= bs * 1.1 - 0.05:
                    tail.append((ts, bs))
        tail.sort(key=lambda p: (-p[1], -p[0]))
        out = list(dict.fromkeys(base + tail))
        if getattr(self, "_dense", False):
            # the dense pass: text below the small size, down to 1.7 % of the slide height (the template's own
            # steps there — a derived ladder on a sparse template); to keep a conclusion on the slide never under
            # 2 % (B3-5: 8 pt on a 6.2″ slide does not read from the back of the room)
            floor = self._dense_floor()
            dense = [x for x in reversed(k.sizes) if floor - 0.05 <= x < k.small - 0.05]
            if not dense and k.small * 0.85 >= floor:
                dense = [math.floor(k.small * 0.85 * 2) / 2]
            extra = [(ts, bs) for bs in dense for ts in (k.body, k.small, bs) if ts >= bs * 1.1 - 0.05 or ts == bs]
            extra.sort(key=lambda p: (-p[1], -p[0]))
            out = list(dict.fromkeys(out + extra))
        return out

    # ---- entry -------------------------------------------------------------------------------------------------
    # the fill pass: a block that takes less than FILL_LOW of its area's height is composed again, larger — first its
    # type may climb the template's scale above the usual cap, then the block itself grows to FILL_TARGET
    FILL_LOW = 0.75
    FILL_TARGET = 0.86
    _GROWABLE = ("cards", "process", "two_column", "comparison", "stat_row", "table", "bullets")

    def compose(self, comp: str, area: Bbox) -> None:
        # the slide's footnote and conclusion (Agent v2) keep their room at the foot of the area: the content is
        # composed above them, the conclusion then follows it
        self._comp = comp
        self._unsay_conclusion()
        area = self._reserve_notes(area)
        n0 = len(self.cv.tree)
        o0, w0 = self.o, list(self.warnings)
        notes0 = getattr(self.o, "notes", None)  # a pass may say lines aloud (speaker notes): undone with the pass
        self._boost = 0
        self._dense = False
        self._area = area
        self._compose_body(comp, area)
        if self._overflows(n0, area):
            # the block runs past the foot of its area at every size of the usual ladder: set again denser — sizes
            # below the small size (down to 1.7 % of the slide height), grids instead of long rows, the lines under
            # the block in two columns or in the speaker notes — never drawn past the area
            self._undo(n0)
            self.o, self.warnings = o0, list(w0)
            self._restore_notes(notes0)
            self._dense = True
            lead = bool(getattr(self, "_take_lead", False))
            self._compose_body(comp, area)
            if lead and self._least_text(n0) < self._dense_floor_2pct() - 0.05:
                # B3-5: the compact variant's conclusion line over a block set dense below 2 % of the slide height
                # (8 pt on a 6.2″ slide): the block at ≥ 2 % under the line; else the conclusion as the strip at the
                # foot (its size, then the body and the small size); else said aloud with the block at ≥ 2 %; the dense
                # setting under the line only when none of these fits
                notes_saved = self._notes
                k = self.kit

                def reset() -> None:
                    self._undo(n0)
                    self.o, self.warnings = o0, list(w0)
                    self._restore_notes(notes0)
                    self._notes = notes_saved
                    self._take_lead = lead

                done = None
                self._floor_2pct = True
                try:
                    reset()
                    self._compose_body(comp, area)
                    if not self._overflows(n0, area):
                        done = "the block set at 2 % under the conclusion line"
                    for size in ((None,) + tuple(dict.fromkeys((k.body, k.small)))) if done is None else ():
                        reset()
                        self._take_lead = False
                        strip_area = self._takeaway_to_strip(area, size)
                        self._compose_body(comp, strip_area)
                        if not self._overflows(n0, strip_area) and self._conclusion_spot(n0) is not None:
                            done = "the conclusion set as the strip under the block: no room over it"
                            break
                    if done is None:
                        reset()
                        self._take_lead = False
                        self._compose_body(comp, area)
                        if not self._overflows(n0, area):
                            said = self._said()
                            try:
                                if said and said not in (self.o.notes or ""):
                                    self.o.notes = (f"{said} " + (self.o.notes or "")).strip()
                            except Exception:  # noqa: BLE001
                                pass
                            done = "the conclusion moved to the speaker notes: no room under the content"
                finally:
                    self._floor_2pct = False
                if done is None:
                    reset()
                    self._compose_body(comp, area)
                elif not done.startswith("the block set"):
                    self.warnings.append(done)
            if comp == "process" or self.o.kind in (PatternKind.process, PatternKind.timeline):
                self._dense_steps_beside(comp, area, n0, (o0, w0, notes0))
            self.warnings.append("dense content set smaller to stay inside its area")
            if self._overflows(n0, area):
                self.warnings.append("content runs past its area")
        growable = comp in self._GROWABLE and not (comp == "stat_row" and len(self.o.content.numbers) <= 1) and not self._dense
        fill = self._fill(n0, area) if growable else None
        if fill is not None and fill < self.FILL_LOW:
            # a short block over an empty band, its text small for the room it has: set again, larger, while it
            # fits; the level that fills the area best (never past its foot) is kept
            best = (fill, 0)
            for boost in (1, 2):
                self._undo(n0)
                self.o, self.warnings = o0, list(w0)
                self._restore_notes(notes0)
                self._boost = boost
                self._compose_body(comp, area)
                got = self._fill(n0, area)
                if got is None or got > 1.0 + 1e-3:
                    break
                if comp == "table" and self._coverage(n0, area) > self.FILL_MAX:
                    break  # a table grown into a wall over the slide (the audit's 80 % of the safe area): it stays smaller
                if got > best[0] + 0.02:
                    best = (got, boost)
                if got >= self.FILL_LOW:
                    break
            if best[1] != self._boost:
                self._undo(n0)
                self.o, self.warnings = o0, list(w0)
                self._restore_notes(notes0)
                self._boost = best[1]
                self._compose_body(comp, area)
        if comp == "table" and not self._dense and self._coverage(n0, area) > self.FILL_MAX:
            # a long table fills the slide past the audit's wall: set again with its rows at their own height (no
            # growth), then — still over — as wide as its columns need at the same type size (never under 60 % of the
            # area), at last a size smaller; the conclusion and the heading keep their room around it
            for level in (1, 2, 3):
                self._undo(n0)
                self.o, self.warnings = o0, list(w0)
                self._restore_notes(notes0)
                self._tight = level
                try:
                    self._compose_body(comp, area)
                finally:
                    self._tight = 0
                if self._coverage(n0, area) <= self.FILL_MAX:
                    break
        boost_used = self._boost
        self._boost = 0
        if comp != "table" and self._notes and self._notes[1] is not None and self._conclusion_spot(n0) is None:
            self._rehouse_conclusion(comp, area, n0, (o0, w0, notes0), boost_used)
        self._draw_notes(n0)
        self.warnings.extend(self.cv.clamp_warnings())

    def _rehouse_conclusion(self, comp: str, area: Bbox, n0: int, state: tuple, boost_used: int) -> None:
        """The block took the conclusion's room and `_conclusion_spot` found none under it (B2). Before the conclusion
        is said aloud: (1) the block is set again steered above that room — the size ladder takes a step down, a row
        of cards goes to two rows, the lines under number tiles step down —; (2) the conclusion leads the content
        under the heading (the compact variant's bold line); (3) the conclusion a size or two smaller in its place at
        the foot; (4) the block at the dense steps — never under 2 % of the slide height (B3-5) — above the room, then
        above the smaller conclusion's. A pass
        counts when the block stays above the content's real floor and the conclusion finds its spot. None: the first
        setting stays and `_draw_notes` says the conclusion aloud (its smallest form tried first)."""
        o0, w0, notes0 = state
        foot_box, plan, full = self._notes
        floor0, area0, dense0 = self._floor_y, self._area, self._dense
        w_first = list(self.warnings)  # the first setting's own warnings (its dense pass), kept when it stays

        def again(body_area: Bbox, floor: int, lead: bool, dense: bool = False) -> bool:
            self._undo(n0)
            self.o, self.warnings = o0, list(w0)
            self._restore_notes(notes0)
            self._floor_y, self._area, self._dense, self._boost = floor, body_area, dense0 or dense, 0
            self._take_lead = lead
            self._rehousing = not lead
            self._floor_2pct = True
            try:
                self._compose_body(comp, body_area)
            finally:
                self._take_lead = False
                self._rehousing = False
                self._floor_2pct = False
                self._floor_y = floor0
            # the tighter floor only steers the sizes: the block may not pass the content's real floor, and the
            # conclusion must find its room under it (`_conclusion_spot`)
            return not self._overflows(n0, body_area)

        # the block kept above the conclusion's room: a step down the ladder, cards in two rows
        shorter = area.y2 < floor0 - int(0.01 * self.kit.H)
        if shorter and again(area, area.y2, False) and self._conclusion_spot(n0) is not None:
            self._floor_y, self._area = floor0, area0
            self.warnings.append("the block set a step smaller: the conclusion keeps its room under it")
            return
        # the conclusion over the content, under the heading (the footnote keeps its place at the foot)
        body = Bbox(x=area.x, y=area.y, w=area.w, h=max(floor0 - area.y, int(0.2 * self.kit.H)))
        self._notes = (foot_box, None, full)
        if (self.o.takeaway or "").strip() and again(body, floor0, True):
            self._floor_y, self._area = floor0, area0
            self.warnings.append("the conclusion leads the content under the heading: no room under it")
            return
        self._notes = (foot_box, plan, full)
        # the conclusion a size or two smaller (the body size, then the small one, without the large figure): its room
        # shrinks and the block keeps its readable steps above it — then the dense steps (never under 2 % of the
        # slide height, B3-5) above the conclusion's own room, then above the smaller one's
        k = self.kit
        take = " ".join((self.o.takeaway or "").split())
        bottom = plan["y_max"] + plan["h"]
        smaller = []
        for size in dict.fromkeys((k.body, k.small)):
            if take and shorter and size < plan["size"] - 0.05:
                p2 = self._takeaway_plan(take, full.w, figure=False, sizes=(size,))
                if p2["h"] < plan["h"]:
                    p2["y_max"] = bottom - p2["h"]
                    smaller.append((p2, Bbox(x=area.x, y=area.y, w=area.w, h=max(p2["y_max"] - int(k.vgap * 1.2) - area.y, int(area.h * 0.4)))))
        for p2, a2 in smaller:
            self._notes = (foot_box, p2, full)
            if again(a2, a2.y2, False) and self._conclusion_spot(n0) is not None:
                self._floor_y, self._area = floor0, area0
                self.warnings.append("the conclusion set a size smaller: the block keeps its readable size above it")
                return
        self._notes = (foot_box, plan, full)
        tries = ([(plan, area)] if shorter else []) + smaller
        for p2, a2 in tries:
            self._notes = (foot_box, p2, full)
            if not dense0 and again(a2, a2.y2, False, True) and self._conclusion_spot(n0) is not None:
                self._floor_y, self._area = floor0, area0
                self.warnings.append("the block set denser: the conclusion keeps its room under it" + ("" if p2 is plan else " (a size smaller)"))
                return
        # neither: the first setting, the conclusion said aloud
        self._notes = (foot_box, plan, full)
        self._undo(n0)
        self.o, self.warnings = o0, list(w0)
        self._restore_notes(notes0)
        self._floor_y, self._area, self._dense, self._boost = floor0, area0, dense0, boost_used
        self._compose_body(comp, area)
        self._boost = 0
        self.warnings = w_first

    FILL_MAX = 0.8  # the share of the safe area past which a slide reads as a wall of content (the audit's fill_ratio)

    def _coverage(self, n0: int, area: Bbox) -> float:
        try:
            return self._coverage_of(n0, area)
        except Exception:  # noqa: BLE001 — a measure never breaks a slide: unmeasured, the table keeps its setting
            return 0.0

    def _coverage_of(self, n0: int, area: Bbox) -> float:
        """The share of the template's safe area the slide's content covers, measured as the audit's fill_ratio does
        (96×54 grid; past 80 % by the blocks' boxes, the text counted by the band its lines fill): the heading (the
        canvas's text above the area), what was composed after the first `n0` elements (text, tables, charts — not
        the bare rules and cards) and the room kept for the conclusion and the footnote."""
        from verstka.rendering.deck import element_bbox

        k = self.kit
        safe = k.manifest.tokens.spacing.safe_area
        sx, sy = safe.x * k.W, safe.y * k.H
        sw, sh = max(safe.w * k.W, 1.0), max(safe.h * k.H, 1.0)
        gx, gy = 96, 54
        boxes: list[tuple] = []
        inks: list[tuple] = []
        for i, el in enumerate(list(self.cv.tree)):
            b = element_bbox(el)
            if not b or b[2] <= 0 or b[3] <= 0:
                continue
            tag = etree.QName(el).localname
            text = "".join(t.text or "" for t in el.iter(q("a:t"))).strip()
            if i < n0:
                if not text or b[1] + b[3] > area.y + int(0.01 * k.H) or b[2] * b[3] >= 0.9 * k.W * k.H:
                    continue  # the template's own chrome and grounds are not content
            elif tag == "sp" and not text:
                continue
            boxes.append(b)
            inks.append(self._ink_box(el, b, heading=i < n0) if tag == "sp" else b)
        notes = getattr(self, "_notes", None)
        if notes:
            foot_box, plan, _a = notes
            if foot_box:
                fb = foot_box[0]
                boxes.append((fb.x, fb.y, fb.w, fb.h))
                inks.append(boxes[-1])
            if plan is not None:
                # the conclusion's text (its strip and bar are bare shapes: not content), where _draw_notes sets it
                bottom = self._content_bottom(n0) or area.y
                y = bottom + int(k.vgap * 1.2) + plan.get("pad_y", 0)
                x = area.x + plan.get("pad_x", 0)
                fig = plan.get("fig")
                if fig:
                    boxes.append((x, y, fig[1], fig[2]))  # its key figure, a text of its own
                    inks.append(boxes[-1])
                x += (fig[1] + fig[3]) if fig else plan.get("bar_w", 0) + plan.get("bar_gap", 0)
                w = plan.get("text_w") or area.w
                para = plan.get("para")
                h = plan.get("text_h") or plan["h"]
                boxes.append((x, y, w, h))
                if para is not None and para.runs:
                    size, font, bold = para.runs[0].size, para.runs[0].font, para.runs[0].bold
                    lines = wrap_lines(para.text, font, size, bold, _pt(w)) or [""]
                    widest = max(text_width_pt(t_, font, size, bold) for t_ in lines)
                    inks.append((x, y, min(w, _emu(widest)), min(h, _emu(len(lines) * size * 1.2 * self.cv.spacing_pct))))
                else:
                    inks.append(boxes[-1])

        def covered(bs) -> float:
            cells = set()
            for x, y, w, h in bs:
                x0, x1 = max(0, int((x - sx) / sw * gx)), min(gx, int((x + w - sx) / sw * gx) + 1)
                y0, y1 = max(0, int((y - sy) / sh * gy)), min(gy, int((y + h - sy) / sh * gy) + 1)
                cells.update((xi, yi) for xi in range(x0, x1) for yi in range(y0, y1))
            return len(cells) / (gx * gy)

        ratio = covered(boxes)
        return max(covered(inks), 0.25) if ratio > 0.8 else ratio

    def _heading_look(self) -> tuple[Optional[str], str, bool]:
        """(font, alignment, bold) the template's headings inherit when their runs do not say (the title's)."""
        from collections import Counter

        slots = [s_ for p in self.kit.manifest.patterns for s_ in p.slots if s_.role.value == "title"]
        fonts = Counter(s_.style.font_family for s_ in slots if s_.style.font_family)
        aligns = Counter(s_.style.align for s_ in slots if s_.style.align)
        bold = sum(1 for s_ in slots if s_.style.bold) > len(slots) / 2
        return (fonts.most_common(1)[0][0] if fonts else self.kit.font), (aligns.most_common(1)[0][0] if aligns else "l"), bold

    def _ink_box(self, el, b, heading: bool = False) -> tuple:
        """What a text box really covers (the audit's ink): the band its lines fill at its anchor and, without a fill of
        its own, no wider than its longest line. A heading's runs inherit the title's font and alignment."""
        k = self.kit
        own_font, own_align, own_bold = self._heading_look() if heading else (k.font, "l", False)
        body = el.find(q("p:txBody"))
        if body is None:
            return b
        bp = body.find(q("a:bodyPr"))
        ins = [int(bp.get(a_, d_)) if bp is not None else d_ for a_, d_ in (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720))]
        usable = max((b[2] - ins[0] - ins[2]) / EMU_PER_PT, 1.0)
        spPr = el.find(q("p:spPr"))
        filled = spPr is not None and any(spPr.find(q(t)) is not None for t in ("a:solidFill", "a:gradFill", "a:blipFill"))
        height = widest = 0.0
        align = own_align
        for p in body.findall(q("a:p")):
            txt = "".join(t.text or "" for t in p.iter(q("a:t")))
            rpr = next((r for r in p.iter(q("a:rPr")) if r.get("sz")), None)
            size = int(rpr.get("sz")) / 100 if rpr is not None else (k.head_size or k.body)
            bold = rpr.get("b") == "1" if rpr is not None and rpr.get("b") is not None else own_bold
            lat = rpr.find(q("a:latin")) if rpr is not None else None
            font = lat.get("typeface") if lat is not None and not (lat.get("typeface") or "").startswith("+") else own_font
            lines = wrap_lines(txt, font, size, bold, usable) if txt.strip() else [""]
            spc = p.find(q("a:pPr") + "/" + q("a:lnSpc") + "/" + q("a:spcPct"))
            pct = int(spc.get("val")) / 100000 if spc is not None and (spc.get("val") or "").isdigit() else None
            height += max(len(lines), 1) * size * (1.2 * pct if pct else 1.2)
            if txt.strip():
                widest = max(widest, max(text_width_pt(t_, font, size, bold) for t_ in lines))
                ppr = p.find(q("a:pPr"))
                align = (ppr.get("algn") if ppr is not None and ppr.get("algn") else align)
        h = max(min(b[3], int(height * EMU_PER_PT) + ins[1] + ins[3]), 1)
        anchor = bp.get("anchor", "t") if bp is not None else "t"
        y = b[1] if anchor == "t" else (b[1] + b[3] - h if anchor == "b" else b[1] + (b[3] - h) // 2)
        x, w = b[0], b[2]
        if not filled:
            w = min(b[2], int(widest * EMU_PER_PT) + ins[0] + ins[2])
            if align == "r":
                x = b[0] + b[2] - w
            elif align == "ctr":
                x = b[0] + (b[2] - w) // 2
        return (x, y, w, h)

    def _restore_notes(self, notes) -> None:
        try:
            self.o.notes = notes
        except Exception:  # noqa: BLE001
            pass

    def _overflows(self, n0: int, area: Bbox) -> bool:
        """What was composed after the first `n0` elements runs past the area's foot (or the slide's edge)."""
        bottom = self._content_bottom(n0)
        if bottom is None:
            return False
        tol = int(0.01 * self.kit.H)
        return bottom > min(self._floor(area.y2), self.kit.H) + tol

    def _floor(self, y2: int) -> int:
        """The foot a block (or the lines under it) may reach: the area's own foot, or — for the composer's whole area
        — the content floor under the conclusion's room (above the footnote)."""
        own = getattr(self, "_area", None)
        floor = getattr(self, "_floor_y", None)
        if own is not None and floor is not None and abs(y2 - own.y2) <= int(0.01 * self.kit.H):
            return max(y2, floor)
        return y2

    def _fill(self, n0: int, area: Bbox) -> Optional[float]:
        """The share of the area's height the content composed after the first `n0` elements takes."""
        bottom = self._content_bottom(n0)
        if bottom is None or area.h <= 0:
            return None
        return (bottom - area.y) / area.h

    def _grow_to(self, area_h: int, rows: int, gap: int, content_h: int) -> int:
        """The height a block's row may take at the fill pass's second level: the rows together about FILL_TARGET of
        the area, a row never more than 1.6 × its content (a card is a frame for its text, not a half-empty slab)."""
        if getattr(self, "_boost", 0) < 2 or rows <= 0:
            return content_h
        want = int((area_h * self.FILL_TARGET - gap * (rows - 1)) / rows)
        return max(content_h, min(want, int(content_h * 1.6)))

    def _compose_body(self, comp: str, area: Bbox) -> None:
        c = self.o.content
        kind = self.o.kind
        intro = self._intro_text(comp)
        if intro and comp not in ("statement", "quote"):
            area = self._intro(intro, area, formula=intro == (c.formula or "").strip())
        if getattr(self, "_take_lead", False):
            area = self._lead_line(area, " ".join((self.o.takeaway or "").split()))
        if comp == "formula" and (c.formula or "").strip():
            self.formula(area, c.formula)
        elif comp == "chart_pair" and c.chart is not None and c.chart2 is not None:
            self.chart_pair(area)
        elif comp in ("table",) and c.table is not None:
            with self._lines_under(area, list(c.bullets)) as main:
                self.table(main, c.table)
        elif comp in ("chart_text", "chart") and c.chart is not None:
            self.chart(area)
        elif comp in ("stat_row", "big_number") and c.numbers:
            if len(c.numbers) == 1 or comp == "big_number":
                self.big_number(area, c.numbers[0], extra=self._running(skip_intro=bool(intro)), others=c.numbers[1:])
            else:
                self.kpis(area, c.numbers, extra=self._running(skip_intro=bool(intro)))
        elif comp in ("process",) or kind in (PatternKind.process, PatternKind.timeline):
            items = self._items()
            lines = list(c.bullets) if (c.items or c.columns) else []
            if items and getattr(self, "_beside", False) and lines and self._steps_beside(area, items, lines):
                pass
            elif items:
                with self._lines_under(area, lines) as main:
                    self.process(main, items)
            else:
                self.bullets(area, self._running())
        elif comp == "agenda" or kind == PatternKind.agenda:
            self.agenda(area, self._items() or [SlideItem(title=b) for b in c.bullets])
        elif comp in ("comparison", "two_column") or kind in (PatternKind.comparison, PatternKind.two_column):
            items = self._items()
            if len(items) >= 2:
                with self._lines_under(area, list(c.bullets) if (c.items or c.columns) else []) as main:
                    self.columns(main, items)
            else:
                self.bullets(area, self._running())
        elif comp in ("cards",) or kind in (PatternKind.cards, PatternKind.team):
            items = self._items()
            if items:
                with self._lines_under(area, list(c.bullets) if (c.items or c.columns) else []) as main:
                    self.cards(main, items)
            else:
                self.bullets(area, self._running())
        elif comp == "quote" and c.quote:
            self.quote(area, c.quote, c.quote_author)
        elif comp == "statement" or (not c.bullets and not c.items and len(c.paragraphs) == 1 and len(c.paragraphs[0]) <= 260):
            text = (c.paragraphs or [c.quote or ""])[0]
            hero = single_figure(text)
            if hero is not None:
                # «Экономия — 18 млн ₽ в год на одну площадку»: one figure carries the message — it is set as one
                value, label = hero
                self.big_number(area, NumberCallout(value=value, label=label), extra=[], others=[])
            else:
                self.statement(area, text)
        else:
            items = self._items()
            if c.bullets:
                self.bullets(area, list(c.bullets) + ([p for p in c.paragraphs] if c.paragraphs and not intro else []))
            elif items:
                self.cards(area, items)
            elif c.paragraphs:
                self.paragraphs(area, c.paragraphs)
            elif c.numbers:
                self.kpis(area, c.numbers, extra=[])
            else:
                self.warnings.append("empty slide content")

    @contextmanager
    def _lines_under(self, area: Bbox, lines: list[str]):
        """Short lines that go with a block of steps, cards or a table (a timeline and the budget of the launch):
        the block gets the area less their height, the lines follow right under what it drew, in two columns when
        there are more than three. No lines: the block gets the whole area."""
        lines = [t for t in lines if t and t.strip()]
        if not lines:
            yield area
            return
        plan = self._under_plan(area, lines)
        n0 = len(self.cv.tree)
        yield plan["main"]
        self._under_draw(plan, n0)

    def _under_plan(self, area: Bbox, lines: list[str], spread: bool = False) -> dict:
        """Where short lines go under a block: the block's box (the area less their height) and the lines' columns —
        two columns for more than three lines; `spread` (the visual variant's lines under its charts): one column per
        line up to three, so they read as captions of the picture rather than as a list."""
        k = self.kit
        n = len(lines)
        cols = (n if n <= 3 else (2 if n == 4 else 3)) if spread else (2 if n > 3 else 1)
        size = k.body
        dense = getattr(self, "_dense", False)
        if dense and not spread and n >= 2:
            cols, size = 2, k.small  # the dense pass: the lines in two columns at the small size

        def lay(cols: int, size: float):
            per = math.ceil(n / cols)
            gap = k.gap * 2
            cw = int((area.w - gap * (cols - 1)) / cols) if cols > 1 else int(area.w * 0.8)
            groups = [lines[i * per:(i + 1) * per] for i in range(cols)]
            paras = [[self.P(t, size, k.colors.text, space_after=size * 0.5, marker="•", marker_color=k.colors.accent) for t in g] for g in groups if g]
            eh = max(self.h(p, cw) for p in paras)
            return per, gap, cw, paras, eh

        per, gap, cw, _, _ = lay(cols, size)
        if (spread or getattr(self, "_boost", 0) >= 1) and self.fits_width(lines, k.lead, False, cw - _emu(k.lead * 1.1)) and sum(len(t) for t in lines) <= 40 * cols * per:
            size = k.lead  # captions of a few words each (or the lines of a short block, on the fill pass) at the lead size
        per, gap, cw, paras, eh = lay(cols, size)
        room = eh + int(k.vgap * 1.4)
        if spread and getattr(self, "_rehousing", False) and room > 0.3 * area.h:
            # G4-19 (Nature short visual s4): set again above the conclusion's room, the captions under the charts
            # step down (the body size, two columns for three lines, the small size — never under the 2 % floor)
            # before the conclusion leaves its strip at the foot for a line under the heading
            floor = self._dense_floor()
            steps = [(cols, x) for x in (k.body, k.small) if floor - 0.05 <= x < size - 0.05]
            if n == 3:
                steps = [(c2, x) for x in dict.fromkeys((k.body, k.small)) if floor - 0.05 <= x <= size + 0.05 for c2 in (cols, 2)]
            for c2, s2 in steps:
                got = lay(c2, s2)
                r2 = got[4] + int(k.vgap * 1.4)
                if r2 < room:
                    per, gap, cw, paras, eh = got
                    cols, size, room = c2, s2, r2
                if room <= 0.3 * area.h:
                    break
        if dense and not spread and n >= 2 and room > 0.3 * area.h:
            # the dense pass: the block keeps 70 % of the area — the lines step down to the dense floor (1.7 % of the
            # slide height) in two columns, then one, before they take more
            floor = self._dense_floor()
            best = None
            for s2 in [x for x in reversed(k.sizes) if floor - 0.05 <= x < size - 0.05]:
                for c2 in (2, 1):
                    got = lay(c2, s2)
                    r2 = got[4] + int(k.vgap * 1.4)
                    if best is None or r2 < best[0]:
                        best = (r2, c2, s2, got)
                    if r2 <= 0.3 * area.h:
                        break
                if best is not None and best[0] <= 0.3 * area.h:
                    break
            if best is not None and best[0] < room:
                room, cols, size = best[0], best[1], best[2]
                per, gap, cw, paras, eh = best[3]
        if room > 0.5 * area.h and not spread:
            # the lines would starve the block: two columns, then a step smaller — the block keeps at least 35 %
            for c2, s2 in ((2, size), (2, k.small)):
                if n < 2:
                    break
                got = lay(c2, s2)
                if got[4] + int(k.vgap * 1.4) < room:
                    per, gap, cw, paras, eh = got
                    cols, size = c2, s2
                    room = eh + int(k.vgap * 1.4)
                if room <= 0.5 * area.h:
                    break
            if room > 0.65 * area.h:
                # still more than the block keeps: the lines are said aloud (speaker notes), the block takes the area
                self._lines_to_notes(lines)
                return dict(main=area, paras=[], eh=0, cw=cw, gap=gap, x=area.x, y2=area.y2, lines=lines)
            main = Bbox(x=area.x, y=area.y, w=area.w, h=area.h - room)
            return dict(main=main, paras=paras, eh=eh, cw=cw, gap=gap, x=area.x, y2=area.y2, lines=lines)
        main = Bbox(x=area.x, y=area.y, w=area.w, h=max(area.h - room, int(area.h * 0.5)))
        return dict(main=main, paras=paras, eh=eh, cw=cw, gap=gap, x=area.x, y2=area.y2, lines=lines)

    def _lines_to_notes(self, lines: list[str]) -> None:
        """Lines with no room on the slide go to its speaker notes (said aloud), never under the slide's foot."""
        text = " ".join("• " + " ".join(t.split()) for t in lines if t and t.strip())
        if not text:
            return
        try:
            if text not in (self.o.notes or ""):
                self.o.notes = ((self.o.notes or "").strip() + "\n" + text).strip()
        except Exception:  # noqa: BLE001
            pass
        self.warnings.append("строки ушли в заметки: нет места на слайде")

    def _under_draw(self, plan: dict, n0: int) -> None:
        """The lines of `_under_plan`, right under what the block drew (a block that grew past its share pushes the
        lines down with it)."""
        k = self.kit
        if not plan["paras"]:
            return
        bottom = self._content_bottom(n0)
        y = plan["main"].y2 + int(k.vgap * 1.4) if bottom is None else bottom + int(k.vgap * 1.4)
        limit = min(self._floor(plan.get("y2") or k.H), k.H)
        if y + plan["eh"] > limit + int(0.005 * k.H):
            # the block grew past its share: the lines are set again into what is left (two columns, smaller), else
            # they go to the speaker notes — never past the area's foot
            refit = self._under_refit(plan, limit - y)
            if refit is None and not getattr(self, "_dense", False):
                pass  # drawn where they fall: the overflow sends the slide to the dense pass (`compose`), which re-plans
            elif refit is None:
                self._lines_to_notes(plan.get("lines") or [])
                return
            else:
                plan = refit
        for i, p in enumerate(plan["paras"]):
            self.cv.text(Bbox(x=plan["x"] + i * (plan["cw"] + plan["gap"]), y=y, w=plan["cw"], h=plan["eh"]), p, name="Note")

    def _under_refit(self, plan: dict, room: int) -> Optional[dict]:
        """The lines of a plan set into `room` (EMU): one or two columns, the body, small and dense sizes."""
        k = self.kit
        lines = plan.get("lines") or []
        if not lines or room <= 0:
            return None
        w = plan["main"].w
        floor = self._dense_floor()
        sizes = list(dict.fromkeys([k.body, k.small] + [x for x in reversed(k.sizes) if floor - 0.05 <= x < k.small - 0.05]))
        for size in sizes:
            for cols in ((1, 2) if len(lines) >= 2 else (1,)):
                per = math.ceil(len(lines) / cols)
                gap = k.gap * 2
                cw = int((w - gap * (cols - 1)) / cols) if cols > 1 else int(w * 0.8)
                groups = [lines[i * per:(i + 1) * per] for i in range(cols)]
                paras = [[self.P(t, size, k.colors.text, space_after=size * 0.5, marker="•", marker_color=k.colors.accent) for t in g] for g in groups if g]
                eh = max(self.h(p, cw) for p in paras)
                if eh <= room:
                    return dict(plan, paras=paras, eh=eh, cw=cw, gap=gap)
        return None

    # ---- content accessors -------------------------------------------------------------------------------------
    def _items(self) -> list[SlideItem]:
        c = self.o.content
        if c.items:
            return list(c.items)
        if c.columns:
            return list(c.columns)
        if c.numbers and self.o.kind not in (PatternKind.stat_row, PatternKind.big_number):
            return [SlideItem(title=n.value, text=n.label, number=n.value) for n in c.numbers]
        return [SlideItem(title=b) for b in c.bullets]

    def _intro_text(self, comp: str = "") -> Optional[str]:
        """A subtitle, or the one paragraph that stands next to structured content (cards, a table, a chart), or —
        on a slide whose chart or table takes the area — its formula as one line."""
        c = self.o.content
        if self.o.subtitle and self.o.kind not in (PatternKind.title, PatternKind.section, PatternKind.thanks):
            return self.o.subtitle
        if comp == "formula":
            return None  # the formula is the slide; its paragraphs explain it under the equation
        if (c.formula or "").strip():
            return c.formula.strip()  # the slide's chart, table or cards take the area: the formula is a line over them
        structured = c.items or c.columns or c.table is not None or c.chart is not None or c.numbers
        if structured and len(c.paragraphs) == 1 and len(c.paragraphs[0]) <= 200 and not c.bullets:
            return c.paragraphs[0]
        return None

    def _running(self, skip_intro: bool = False) -> list[str]:
        c = self.o.content
        out = list(c.bullets)
        if not skip_intro:
            out += list(c.paragraphs)
        return out

    # ---- blocks ------------------------------------------------------------------------------------------------
    def _intro(self, text: str, area: Bbox, formula: bool = False) -> Bbox:
        k = self.kit
        width = min(area.w, int(area.w * 0.78))
        size = k.lead if len(text) <= 120 else k.body
        para = [self.P(text, size, k.colors.muted if k.colors.muted != k.colors.text else k.colors.text)]
        if formula:  # «100 × 300 × 30 = 900 000 ₽» as a line: its figures in the accent
            para = [Para(highlight_runs(text, size, k.colors.text, k.colors.accent_text, bold=False, accent_bold=k.bold, font=k.font))]
        hh = self.h(para, width)
        self.cv.text(Bbox(x=area.x, y=area.y, w=width, h=hh), para, name="Intro")
        dy = hh + int(k.vgap * 1.1)
        return Bbox(x=area.x, y=area.y + dy, w=area.w, h=max(area.h - dy, int(area.h * 0.4)))

    def _place_v(self, area: Bbox, block_h: int, fill_top: bool = False) -> int:
        """Top of a block: always the content line under the heading. Flipping through the deck, content starts at
        the same height on every slide; the air a short block leaves stays at the bottom (the type grows first)."""
        return area.y

    # ---- cards --------------------------------------------------------------------------------------------------
    def _grid(self, n: int, area: Bbox, longest: int) -> tuple[int, int]:
        if n <= 1:
            return 1, 1
        wide = area.w / max(area.h, 1) >= 1.6
        if n <= 3:
            return n, 1
        if n == 4:
            return (4, 1) if wide and longest <= 140 else (2, 2)
        if n == 5:
            return (5, 1) if wide and longest <= 70 else (3, 2)
        if n == 6:
            return 3, 2
        if n <= 8:
            return 4, 2
        return math.ceil(n / 3), 3

    def cards(self, area: Bbox, items: list[SlideItem], badge: Optional[str] = None) -> None:
        k = self.kit
        st = k.card
        n = len(items)
        if all(not i.text and not i.bullets and not i.number for i in items) and max(len(i.title.split()) for i in items) > 6:
            # theses, not titled items: each card holds its sentence as running text (regular weight, text colour)
            items = [SlideItem(title="", text=i.title, icon_hint=i.icon_hint) for i in items]
        longest = max(len(i.title) + len(i.text) + sum(len(b) for b in i.bullets) for i in items)
        cols, rows = self._grid(n, area, longest)
        if getattr(self, "_dense", False) and rows == 1 and n >= 5:
            cols, rows = math.ceil(n / 2), 2  # the dense pass: five or six cards in two rows, not one row of slivers
        if badge is None:
            # circles with digits say «in this order» — only steps get them; theses and risks carry an index numeral
            sequence = self.o.kind in (PatternKind.process, PatternKind.timeline)
            badge = {"visual": "number" if sequence else "index", "structured": "index", "compact": None}.get(self.strategy, "index")
            if self.o.kind in (PatternKind.team,):
                badge = None
            if all(i.number for i in items):
                badge = "figure"
        gap = k.gap
        colors = st.colors
        title_color = colors.heading

        forced_of: dict = {}  # (cols, rows, tight) → no size pair fits (a word breaks or the row overflows)
        lone = n <= 2 and all(not (i.title and i.text) and not i.bullets and not i.number for i in items)

        def pick(cols: int, rows: int, tight: bool = False):
            cw = int((area.w - gap * (cols - 1)) / cols)
            pad = max(int(min(cw * 0.085, 0.05 * k.H)), int(0.026 * k.H))
            if tight:
                pad = max(int(pad * 0.6), int(0.022 * k.H))
            inner_w = cw - 2 * pad
            # sizes: the largest pair at which every card reads well (title ≤ 3 lines, text ≤ 7) and the row fits
            max_row_h = int((area.h - k.vgap * (rows - 1)) / rows)
            ladders = self._size_pairs()
            if n <= 3 and longest <= 70:
                ladders.insert(0, (k.statement, k.h2))  # three short items: the type grows before the card does
            if lone:
                # one or two sentences in cards (a written slide's two theses, EV compact s7: body size in two tall
                # cards, 14 % of the slide in lines): set like a statement — the callout size under the heading, down
                # to the lead size — before the usual ladder
                top = min(k.statement, self._under_heading()) if k.head_size else k.statement
                ladders = [(x, x) for x in k.steps_down(top, k.lead)] + ladders
            chosen = None
            fallback = None
            long_ok = None  # fits, but with more lines than we like: kept when the next step down is a big drop
            bodies = [b for i in items for b in ([i.text] if i.text else []) + list(i.bullets)]
            for ts, bs in ladders:
                # a word wider than the card breaks in the middle: its title and its text are measured
                if not self.fits_width([i.title for i in items], ts, k.bold, inner_w) or not self.fits_width(bodies, bs, False, inner_w - _emu(bs * 1.1)):
                    continue
                tp = self._title_pads(items, cols, inner_w, ts, self._hang_w(ts) if badge == "index_h" else 0)
                h = max(self._card_content_h(it, inner_w, ts, bs, badge, colors, i, tp[i]) for i, it in enumerate(items)) + 2 * pad
                if h > max_row_h:
                    continue
                t_lines = max(para_lines(self.P(i.title, ts, colors.text, bold=k.bold), _pt(inner_w)) for i in items)
                b_lines = max((sum(para_lines(self.P(b, bs, colors.text), _pt(inner_w)) for b in ([i.text] if i.text else []) + list(i.bullets)) for i in items), default=0)
                words = sum(len(t.split()) for i in items for t in [i.title, i.text] + list(i.bullets) if t)
                lines_all = sum(para_lines(self.P(i.title, ts, colors.text, bold=k.bold), _pt(inner_w)) for i in items if i.title) + sum(para_lines(self.P(b, bs, colors.text), _pt(inner_w)) for i in items for b in ([i.text] if i.text else []) + list(i.bullets))
                if lines_all and words / lines_all < 1.8 and (ts, bs) != ladders[-1]:
                    continue  # one or two words a line reads as a column of words, not as text
                fallback = fallback or (ts, bs, h)
                t_max = 3 if b_lines else 5  # a title that is the whole card may take more lines
                if t_lines <= t_max and b_lines <= 7:
                    chosen = (ts, bs, h)
                    break
                if t_lines <= t_max and b_lines <= 10 and long_ok is None:
                    long_ok = (ts, bs, h)
            if long_ok is not None and (chosen is None or chosen[1] < 0.8 * long_ok[1]):
                chosen = long_ok  # a template with sparse sizes (18 → 12) keeps 18 pt in nine lines rather than 12 pt
            chosen = chosen or fallback
            forced_of[(cols, rows, tight)] = chosen is None
            if chosen is None:
                ts, bs = k.body, k.small
                tp = self._title_pads(items, cols, inner_w, ts, self._hang_w(ts) if badge == "index_h" else 0)
                chosen = (ts, bs, max(self._card_content_h(it, inner_w, ts, bs, badge, colors, i, tp[i]) for i, it in enumerate(items)) + 2 * pad)

            return cw, pad, inner_w, chosen

        cw, pad, inner_w, chosen = pick(cols, rows)
        forced = forced_of[(cols, rows, False)]
        if chosen[1] < k.body - 0.05 or forced:
            # text in small type inside roomy cards: a tighter padding gives the text the room first
            cw2, pad2, inner2, chosen2 = pick(cols, rows, tight=True)
            if (chosen2[1] > chosen[1] + 0.05 and forced_of[(cols, rows, True)] <= forced) or (forced and not forced_of[(cols, rows, True)]):
                cw, pad, inner_w, chosen = cw2, pad2, inner2, chosen2
                forced = forced_of[(cols, rows, True)]
        if forced and rows == 1 and n >= 4:
            # a row of cards too narrow for their words at every size (a 4:3 slide, a column beside art): two rows
            c2, r2 = (2, 2) if n == 4 else (math.ceil(n / 2), 2)
            for tight in (False, True):
                got = pick(c2, r2, tight)
                if not forced_of[(c2, r2, tight)]:
                    cols, rows = c2, r2
                    cw, pad, inner_w, chosen = got
                    forced = False
                    self.warnings.append(f"{n} карточек → сетка {c2}×{r2}")
                    break
        if forced and badge == "index":
            # the index numerals take the room the words need: the same cards without them (one row, then two)
            grids = [(cols, rows)] + ([(2, 2) if n == 4 else (math.ceil(n / 2), 2)] if rows == 1 and n >= 4 else [])
            for badge in ("index_s", "index_h", None):  # G4-18: the numerals a step over the title, hanging, then none
                for c_, r_ in dict.fromkeys(grids):
                    for tight in (False, True):
                        got = pick(c_, r_, tight)
                        if not forced_of[(c_, r_, tight)]:
                            cols, rows = c_, r_
                            cw, pad, inner_w, chosen = got
                            forced = False
                            break
                    if not forced:
                        break
                if not forced:
                    break
            if forced:
                badge = "index"
        if forced and all(not i.bullets and not (i.title and i.text) and not i.number for i in items):
            # theses that fit no grid of cards at a readable size: a list (the words keep whole, the type stays
            # readable) — never five slivers of 9 pt text
            self.warnings.append("тезисы не помещаются в карточки: набраны списком")
            self.bullets(area, [i.title or i.text for i in items], as_list=True)
            return
        if forced and not getattr(self, "_dense", False):
            # still no fit: the dense sizes (below the small size, down to 1.7 % of the slide height)
            self._dense = True
            try:
                got = pick(cols, rows, True)
            finally:
                self._dense = False
            if not forced_of[(cols, rows, True)]:
                cw, pad, inner_w, chosen = got
                forced = False
        if n == 4 and cols == 4 and chosen[1] < k.body:
            # four cards that only fit in small type read better as a 2×2 block
            cw2, pad2, inner2, chosen2 = pick(2, 2)
            if chosen2[1] > chosen[1]:
                cols, rows, cw, pad, inner_w, chosen = 2, 2, cw2, pad2, inner2, chosen2
        has_body = any(i.text or i.bullets for i in items)

        def reading(ch) -> float:
            return ch[1] if has_body else ch[0]  # the size the cards' words are read at

        if not forced and reading(chosen) < k.body - 0.05 and not getattr(self, "_dense", False):
            # B3-2: the words of the cards in small type (12 pt under 18 pt titles, theses at 12 pt in four slivers):
            # the other grid (two rows), the cards without their index numerals, a tighter padding — the setting that
            # reads largest is taken when it gains a step (ties keep the numerals, the padding, the first grid)
            grids = [(cols, rows)] + ([(2, 2) if n == 4 else (math.ceil(n / 2), 2)] if rows == 1 and n >= 4 else [])
            badge0 = badge
            best, best_kept, best_own = None, None, None
            # G4-18: the index numerals are tried at their own size, then a step over the title only, then left out
            for gi, (c_, r_) in enumerate(dict.fromkeys(grids)):
                for bi, b_ in enumerate(dict.fromkeys([badge0] + (["index_s", "index_h", None] if badge0 in ("index", "index_s", "index_h") else []))):
                    for tight in (False, True):
                        badge = b_
                        got = pick(c_, r_, tight)
                        if forced_of[(c_, r_, tight)]:
                            continue
                        key = (reading(got[3]), -bi, -int(tight), -gi)
                        if best is None or key > best[0]:
                            best = (key, c_, r_, b_, got)
                        if (b_ is not None or badge0 is None) and (best_kept is None or key > best_kept[0]):
                            best_kept = (key, c_, r_, b_, got)
                        if bi == 0 and (best_own is None or key > best_own[0]):
                            best_own = (key, c_, r_, b_, got)
            badge = badge0
            if best_own is not None and best_own[0][0] >= k.body - 0.05:
                best_kept = best_own  # the variant's own numerals, where the words reach the body size with them (VK Tech)
            keep_at = max(k.small, 0.85 * best[0][0]) if best is not None else k.body
            if best_kept is not None and (best_kept[0][0] >= k.body - 0.05 or (n >= 3 and best_kept[0][0] >= keep_at - 0.05)):
                # the numerals stay where the words reach the body size with them — a row of three or more keeps them
                # (G4-18) while the words stay within a step of what they reach without them
                best = best_kept
            if best is not None and best[0][0] > reading(chosen) + 0.05:
                _, cols, rows, badge, (cw, pad, inner_w, chosen) = best
                self.warnings.append(f"карточки {cols}×{rows}{' без номеров' if badge is None and badge0 in ('index', 'index_s', 'index_h') else (' с малыми номерами' if badge == 'index_s' else (' с номерами сбоку' if badge == 'index_h' else ''))}: текст {reading(chosen):g} пт")
        ts, bs, content_h = chosen
        max_row_h = int((area.h - k.vgap * (rows - 1)) / rows)
        if content_h > max_row_h and rows == 1 and 2 <= n <= 3 and any(it.bullets for it in items):
            # cards with lists that do not fit their row at any size: columns set the same content denser (no badge,
            # the figure over the lines) — never a card running into the conclusion under it
            self.columns(area, items)
            return
        # cards hug their content (the tallest one sets the row): a card is never a half-empty slab — nor shorter than
        # its text (when even the smallest step does not fit the row, the card grows rather than the text hanging out)
        ch = self._grow_to(area.h, rows, k.vgap, content_h)
        block_h = rows * ch + (rows - 1) * k.vgap
        y0 = self._place_v(area, block_h)
        last_row_n = n - cols * (rows - 1)
        pads = self._title_pads(items, cols, inner_w, ts, self._hang_w(ts) if badge == "index_h" else 0)
        for i, it in enumerate(items):
            r, cidx = divmod(i, cols)
            x_off = 0
            if r == rows - 1 and last_row_n < cols:
                x_off = int((cols - last_row_n) * (cw + gap) / 2)  # an incomplete last row is centred
            box = Bbox(x=area.x + x_off + cidx * (cw + gap), y=y0 + r * (ch + k.vgap), w=cw, h=ch)
            if st.fill or st.line or st.proto is not None:
                self.cv.card(box, st)
            else:
                # no card style: an accent rule over each item keeps the columns apart
                self.cv.rect(Bbox(x=box.x, y=box.y, w=box.w, h=max(_emu(2.25), 1)), colors.accent, name="Rule")
            self._card_content(it, Bbox(x=box.x + pad, y=box.y + pad, w=inner_w, h=ch - 2 * pad), ts, bs, badge, colors, i, title_color, pads[i])

    def figure_para(self, value: str, size: float, color: str, muted: str) -> tuple[Para, float]:
        """A figure as one line: «31% → 12%» is two figures — the old one small and muted, the new one in the
        accent at full size. Returns the paragraph and its width in points."""
        k = self.kit
        b = k.figure_bold
        parts = re.split(r"\s*(?:→|->|⟶)\s*", value.strip())
        if len(parts) == 2 and all(any(ch.isdigit() for ch in x) for x in parts):
            small = k.snap(size * 0.55, k.h3, size * 0.7)
            runs = [Run(typeset(parts[0]), small, muted, b, k.font), Run("\u00a0→\u00a0", small, muted, False, k.font)]
            width = text_width_pt(parts[0], k.font, small, b) + text_width_pt(" → ", k.font, small, False)
            m = _FIG_UNIT_RE.match(parts[1])
            if m and m.group(2).strip() and size >= 1.6 * k.h2:
                # «47 → 29 минут»: the new figure leads, its word unit rides on the baseline at the old figure's size
                num, unit = m.group(1).strip(), m.group(2).strip()
                runs += [Run(typeset(num), size, color, b, k.font), Run("\u00a0" + typeset(unit), small, color, b, k.font)]
                return Para(runs), width + text_width_pt(num, k.font, size, b) + text_width_pt(" " + unit, k.font, small, b)
            runs.append(Run(typeset(parts[1]), size, color, b, k.font))
            return Para(runs), width + text_width_pt(parts[1], k.font, size, b)
        hm = _HEDGE_RE.match(value.strip())
        if hm and size >= 1.6 * k.h2:
            # «около 615 тыс. тонн», «более 10»: the hedge is a word of the figure, not the figure — it is set at the
            # unit's size on the figure's baseline, so the number keeps the hero's size in its tile (G4 Focus s3)
            hedge, rest = hm.group(1), value.strip()[hm.end():]
            usz = k.snap(size * 0.5, k.h3, size * 0.62)
            head = [Run(typeset(hedge) + " ", usz, color, b, k.font)]
            hw = text_width_pt(hedge + " ", k.font, usz, b)
            m = _FIG_UNIT_RE.match(rest)
            if m and m.group(2).strip():
                num, unit = m.group(1).strip(), m.group(2).strip()
                runs = head + [Run(typeset(num), size, color, b, k.font), Run(" " + typeset(unit), usz, color, b, k.font)]
                return Para(runs), hw + text_width_pt(num, k.font, size, b) + text_width_pt(" " + unit, k.font, usz, b)
            return Para(head + [Run(typeset(rest), size, color, b, k.font)]), hw + text_width_pt(typeset(rest), k.font, size, b)
        m = _FIG_UNIT_RE.match(value.strip())
        if m and m.group(2).strip() and size >= 1.6 * k.h2:
            # the unit is set smaller on the figure's baseline: «145 000 ₽», «2,1 ч», «4,6 из 5»
            num, unit = m.group(1).strip(), m.group(2).strip()
            usz = k.snap(size * 0.5, k.h3, size * 0.62)
            runs = [Run(typeset(num), size, color, b, k.font), Run("\u00a0" + typeset(unit), usz, color, b, k.font)]
            return Para(runs), text_width_pt(num, k.font, size, b) + text_width_pt(" " + unit, k.font, usz, b)
        v = typeset(value)
        return Para([Run(v, size, color, b, k.font)]), text_width_pt(v, k.font, size, b)

    def _under_heading(self) -> float:
        """The largest size a callout (a statement, a quote) may take under the slide's heading: a clear step below it."""
        k = self.kit
        return k.snap(k.head_size * 0.84, k.lead, max(k.lead, k.head_size * 0.9))

    def _figure_optics(self, para: Para, box_h: int) -> tuple[int, int]:
        """(left, top) EMU a large figure set bottom-anchored in a box `box_h` tall is moved by, so that its digits —
        not the empty side bearing and ascender room of the font — stand on the left edge and the top line of the
        content. The measuring font's metrics (Play); other families are moved a little less."""
        k = self.kit
        if not para.runs or not para.runs[0].text:
            return 0, 0
        first = para.runs[0]
        big = max(r.size for r in para.runs)
        bold = bool(first.bold)
        damp = 1.0 if not k.font or k.font.lower().startswith("play") else 0.7
        ref = 0.07 * (k.head_size or k.h2)  # the heading's own letters stand off its edge about this much
        left = max(0.0, left_bearing_em(first.text.lstrip()[:1], bold) * first.size - ref) * damp
        descent, digit_h = figure_metrics_em(bold)
        top = max(0.0, _pt(box_h) - (descent + digit_h) * big) * damp
        return _emu(left), _emu(top * 0.95)

    @staticmethod
    def _index_ref(it: SlideItem, ts: float, bs: float, badge: Optional[str]) -> float:
        """The size a card's numeral is set from: the title's — the smaller numeral of a thesis (a card of words
        without a title) goes with its words."""
        return bs if badge == "index_s" and not (it.title and it.title != it.number) else ts

    def _index_size(self, ts: float, small: bool = False) -> float:
        """The «01» over a card: the accent numeral, a step above the card title (never the small accent text) — in
        the visual variant a display numeral, twice the title («крупные цифры»). `small`: the numeral a step over the
        title only, when the display numeral would take the room the cards' words need (G4-18)."""
        k = self.kit
        if small:
            return k.snap(ts * 1.3, ts * 1.1, ts * 1.5)
        if self.strategy == "visual":
            return k.snap(max(ts * 2.0, k.statement), ts * 1.5, max(k.display, ts * 2.4))
        return k.snap(max(ts * 1.4, k.h2), ts * 1.2, max(k.statement, ts * 1.8))

    def _badge_size(self, ts: float) -> int:
        """One badge for the deck (cards, steps): 7.5% of the slide height (never under 36 pt, so its digit can be
        set as large text), whatever the text size of the slide."""
        return _emu(max(self.kit.hpt * 0.075, 36.0))

    def _badge_digit(self) -> float:
        k = self.kit
        d = max(k.hpt * 0.075, 36.0)
        return k.snap(max(18.0, d * 0.48), 18.0, d * 0.6)

    def _card_paras(self, it: SlideItem, ts: float, bs: float, colors: Colors, title_color: Optional[str] = None, title_pad: float = 0.0) -> list[Para]:
        """A card's title and its words; `title_pad` (pt) lengthens the gap under a title shorter than the row's
        tallest one, so the words of a row start on one line (G4-18)."""
        paras: list[Para] = []
        head = it.title if it.title and it.title != it.number else ""
        body_texts = ([it.text] if it.text else []) + list(it.bullets)
        if head:
            paras.append(self.P(head, ts, title_color or colors.heading, bold=self.kit.bold, space_after=(bs * 0.55 + title_pad) if body_texts else 0))
        for j, b in enumerate(body_texts):
            last = j == len(body_texts) - 1
            marker = "•" if (it.bullets and j >= (1 if it.text else 0)) else None
            paras.append(self.P(b, bs, colors.muted if head else colors.text, space_after=0 if last else bs * 0.4, marker=marker, marker_color=colors.accent))
        return paras

    def _title_pads(self, items: list[SlideItem], cols: int, inner_w: int, ts: float, hang: int = 0) -> list[float]:
        """Per card, the points its title block is shorter than the tallest one of its row (cards with a title and
        words only): «Свободные часы» on two lines no longer pushes its row's words out of line (G4-18). `hang`: the
        titles stand beside a hanging numeral (that much narrower)."""
        k = self.kit
        lines = []
        for it in items:
            head = it.title if it.title and it.title != it.number else ""
            has_body = bool(it.text or it.bullets)
            lines.append(para_lines(self.P(head, ts, "000000", bold=k.bold), _pt(inner_w - hang)) if head and has_body else 0)
        pads = [0.0] * len(items)
        for r0 in range(0, len(items), max(cols, 1)):
            row = range(r0, min(r0 + cols, len(items)))
            top = max(lines[i] for i in row)
            for i in row:
                if lines[i]:
                    pads[i] = (top - lines[i]) * ts * k.line
        return pads

    def _hang_w(self, size: float) -> int:
        """The room of a hanging numeral «05» beside a card's first line: its digits and a gap."""
        k = self.kit
        return _emu(text_width_pt("00", k.font, size, k.bold) * 1.05 + size * 0.6)

    def _hung(self, it: SlideItem, paras: list[Para], width: int, ts: float, bs: float) -> tuple[int, list[tuple[int, int, list[Para], int]]]:
        """A card with its numeral hanging beside the first line (G4-18: the numeral costs a column, not a line):
        (numeral room, [(x offset, y offset, paras, height)]) — the title beside the numeral and the words under both
        at the card's width; a thesis (no title) all beside it."""
        head = bool(it.title and it.title != it.number)
        nw = self._hang_w(ts if head else bs)
        if not head or len(paras) < 2:
            return nw, [(nw, 0, paras, self.h(paras, width - nw))]
        tp = paras[0]
        t0 = Para(tp.runs, tp.align, 0.0, tp.marker, tp.marker_color)
        th = self.h([t0], width - nw)
        return nw, [(nw, 0, [t0], th), (0, th + _emu(tp.space_after), paras[1:], self.h(paras[1:], width))]

    def _card_content_h(self, it: SlideItem, inner_w: int, ts: float, bs: float, badge: Optional[str], colors: Colors, i: int, title_pad: float = 0.0) -> int:
        paras = self._card_paras(it, ts, bs, colors, title_pad=title_pad)
        if badge == "index_h" and paras:
            _, parts = self._hung(it, paras, inner_w, ts, bs)
            return max(dy + hh for _, dy, _, hh in parts)
        h = self.h(paras, inner_w)
        if badge == "number":
            h += self._badge_size(ts) + int(self.kit.vgap * 0.7)
        elif badge in ("index", "index_s"):
            h += _emu(self._index_size(self._index_ref(it, ts, bs, badge), small=badge == "index_s") * self.kit.line) + int(self.kit.vgap * 0.35)
        elif badge == "figure" and it.number:
            fs = self._figure_size(it.number, inner_w)
            h += _emu(fs * max(1.1, self.kit.line)) + int(self.kit.vgap * 0.3)
        return h

    def _figure_size(self, value: str, width: int) -> float:
        k = self.kit
        for s in k.figure_sizes(k.figure_cap(4, True, self.strategy), k.h3):
            if self.figure_para(value, s, "000000", "000000")[1] <= _pt(width) * 0.88:
                return s
        return k.h3

    def _card_content(self, it: SlideItem, box: Bbox, ts: float, bs: float, badge: Optional[str], colors: Colors, i: int, title_color: str, title_pad: float = 0.0) -> None:
        k = self.kit
        y = box.y
        if badge == "number":
            d = self._badge_size(ts)
            fill = colors.accent
            num_color = self.kit.on_accent(fill, colors.text)
            el = self.cv.ellipse(Bbox(x=box.x, y=y, w=d, h=d), fill)
            self._label_in(el, str(i + 1), self._badge_digit(), num_color)
            y += d + int(k.vgap * 0.7)
        elif badge in ("index", "index_s"):
            idx = f"{i + 1:02d}"
            isz = self._index_size(self._index_ref(it, ts, bs, badge), small=badge == "index_s")
            hh = _emu(isz * k.line)
            self.cv.text(Bbox(x=box.x, y=y, w=box.w, h=hh), [self.P(idx, isz, colors.accent, bold=k.bold)], name="Index")
            y += hh + int(k.vgap * 0.35)
        elif badge == "figure" and it.number:
            fs = self._figure_size(it.number, box.w)
            hh = _emu(fs * max(1.1, k.line))
            self.cv.text(Bbox(x=box.x, y=y, w=box.w, h=hh), [self.figure_para(it.number, fs, colors.figure, colors.muted)[0]], name="Figure", anchor="b")
            y += hh + int(k.vgap * 0.3)
        paras = self._card_paras(it, ts, bs, colors, title_color, title_pad=title_pad)
        if badge == "index_h" and paras:
            # the numeral hangs beside the first line, at its size (one baseline); the title beside it, the words under
            ref = paras[0].size
            nw, parts = self._hung(it, paras, box.w, ts, bs)
            self.cv.text(Bbox(x=box.x, y=y, w=nw, h=_emu(ref * k.line) + _emu(2)), [self.P(f"{i + 1:02d}", ref, colors.accent, bold=k.bold)], name="Index")
            for j, (dx, dy, ps, hh) in enumerate(parts):
                last = j == len(parts) - 1
                self.cv.text(Bbox(x=box.x + dx, y=y + dy, w=box.w - dx, h=max(hh, box.y2 - y - dy) if last else hh), ps, name="Card text")
            return
        if paras:
            hh = self.h(paras, box.w)
            self.cv.text(Bbox(x=box.x, y=y, w=box.w, h=max(hh, box.y2 - y)), paras, name="Card text")

    def _label_in(self, el: etree._Element, text: str, size: float, color: str) -> None:
        """Centered text inside a drawn shape (a badge)."""
        tx = etree.SubElement(el, q("p:txBody"))
        bp = etree.SubElement(tx, q("a:bodyPr"))
        for n in ("lIns", "tIns", "rIns", "bIns"):
            bp.set(n, "0")
        bp.set("anchor", "ctr")
        bp.set("wrap", "none")
        etree.SubElement(tx, q("a:lstStyle"))
        p = etree.SubElement(tx, q("a:p"))
        pPr = etree.SubElement(p, q("a:pPr"))
        pPr.set("algn", "ctr")
        r = etree.SubElement(p, q("a:r"))
        rPr = etree.SubElement(r, q("a:rPr"))
        rPr.set("lang", "ru-RU")
        rPr.set("sz", str(int(round(size * 100))))
        rPr.set("b", "1" if self.kit.bold else "0")
        rPr.set("cap", "none")
        rPr.set("spc", "0")
        rPr.set("baseline", "0")
        sf = etree.SubElement(rPr, q("a:solidFill"))
        etree.SubElement(sf, q("a:srgbClr")).set("val", color.upper())
        if self.kit.font:
            etree.SubElement(rPr, q("a:latin")).set("typeface", self.kit.font)
            etree.SubElement(rPr, q("a:cs")).set("typeface", self.kit.font)
        etree.SubElement(r, q("a:t")).text = text

    # ---- KPI tiles ---------------------------------------------------------------------------------------------
    def kpis(self, area: Bbox, numbers: list[NumberCallout], extra: list[str]) -> None:
        k = self.kit
        n = len(numbers)
        extra_block = None
        if extra:
            # the running text goes under the tiles as a short note — in two columns when there are more than three
            # lines (a list of measures under two figures), so the figures keep their size
            ecols = 2 if len(extra) > 3 else 1
            per = math.ceil(len(extra) / ecols)
            ew = int((area.w - k.gap) / 2) if ecols == 2 else int(area.w * 0.8)
            groups = [extra[i * per:(i + 1) * per] for i in range(ecols)]
            esize = k.lead if getattr(self, "_boost", 0) >= 1 and self.fits_width(extra, k.lead, False, ew - _emu(k.lead * 1.1)) else k.body
            steps = [esize]
            if getattr(self, "_rehousing", False):
                # set again above the conclusion's room (B2): the lines under the tiles step down (the small size, on
                # the dense pass the dense floor) while they take more than 45 % of the area
                steps += [x for x in (k.small,) if x < esize - 0.05]
                if getattr(self, "_dense", False):
                    steps.append(max(self._dense_floor(), k.small * 0.85))
            for esize in steps:
                paras_g = [[self.P(t, esize, k.colors.text, space_after=esize * 0.5, marker="•" if len(extra) > 1 else None, marker_color=k.colors.accent) for t in g] for g in groups if g]
                eh = max(self.h(pg, ew) for pg in paras_g)
                if eh <= 0.45 * area.h:
                    break
            extra_block = (paras_g, eh, ew)
        avail = Bbox(x=area.x, y=area.y, w=area.w, h=area.h - ((extra_block[1] + k.vgap) if extra_block else 0))
        st = k.card
        use_cards = bool(st.fill or st.proto is not None) and self.strategy != "compact"
        colors = st.colors if use_cards else k.colors
        gap = k.gap
        hero_min = 1.5 * (k.head_size or k.h2)  # a figure smaller than this stops being the hero of the slide

        # «27 миллионов человек»: the figure with the words it counts is too long for a tile at a figure's size — the
        # number keeps the figure's size and its words go on the line under it (same text: the value stays whole)
        splits = [figure_split(x.value) for x in numbers]

        def unit_paras(fs_: float, lab_size: float) -> list[Optional[Para]]:
            us = k.snap(max(lab_size * 1.2, fs_ * 0.4), lab_size, max(lab_size, fs_ * 0.5))
            return [self.P(sp[1], us, colors.figure, bold=k.figure_bold) if sp else None for sp in splits]

        slim = [False]

        def tile_pad_y(cols: int, pad: int) -> int:
            # tiles in two rows keep a slimmer top and bottom padding: 0.05 of the slide's height over and under each
            # figure took 45 % of a 2×2 tile and left the figures at 18 pt under 20 pt labels (B3); so does a row whose
            # figures the full padding sets under the heading's size (G4-20: «300 → 330 ₽» at 23.5 pt under 33.5)
            return min(pad, int(0.03 * k.H)) if (math.ceil(n / cols) > 1 or slim[0]) else pad

        def layout(cols: int, stacked: bool = False):
            cw = int((avail.w - gap * (cols - 1)) / cols)
            pad = int(min(cw * 0.1, 0.05 * k.H)) if use_cards else 0
            pad_y = tile_pad_y(cols, pad)
            inner = cw - 2 * pad
            rows_ = math.ceil(n / cols)
            max_tile_ = int((avail.h - k.vgap * (rows_ - 1)) / rows_)
            lab = max(self.h([self.P(distinct_label(x.value, x.label, self.o.headline), k.body, colors.muted)], inner) for x in numbers)
            # one figure size for the row: the largest size under the cap for this many figures at which every
            # figure fits its tile on one line and the tiles fit the area
            fs = k.h3
            cap_ = k.figure_cap(max(cols, 2) if cols < n else n, use_cards, self.strategy) * (1.3 if getattr(self, "_boost", 0) >= 1 else 1.0)
            wide_ok = False
            for s_ in k.figure_sizes(cap_, k.h3):
                fs = s_
                # slack for the rendering font: a figure that breaks («12 40 / 0») is the worst thing a slide can show
                heads = [sp[0] if (stacked and sp) else x.value for x, sp in zip(numbers, splits)]
                wide_ok = all(self.figure_para(v, s_, colors.figure, colors.muted)[1] <= _pt(inner) * (0.88 if use_cards else 0.8) for v in heads)
                units_h = max((self.h([u], inner) for u in unit_paras(s_, k.body) if u is not None), default=0) if stacked else 0
                units_ok = not stacked or all(u is None or para_lines(u, _pt(inner)) <= 2 for u in unit_paras(s_, k.body))
                tall_ok = _emu(s_ * max(1.15, k.line)) + units_h + lab + int(k.vgap * 1.4) + 2 * pad_y <= max_tile_
                if wide_ok and tall_ok and units_ok:
                    break
            return cw, pad, inner, fs, wide_ok

        cols = n if n <= 4 else math.ceil(n / 2)
        cw, pad, inner, fs, one_line = layout(cols)
        if n == 4 and fs < hero_min and avail.h > 0.55 * area.h:
            # four long figures read as a 2×2 block, not as four small numbers in a row — when the block gives them
            # a larger size (short figures «70 млн», «62» keep their row at 48 pt rather than a grid at 18 pt, B3)
            grid = layout(2)
            if grid[3] > fs + 0.05:
                cols = 2
                cw, pad, inner, fs, one_line = grid
        stacked = False
        if any(splits) and fs < 1.6 * k.h2:
            # a row of figures set at a text size (smaller than their labels) reads as captions: stack the words
            cw2, pad2, inner2, fs2, ok2 = layout(cols, stacked=True)
            if fs2 >= fs * 1.3 or (ok2 and not one_line):
                stacked, cw, pad, inner, fs, one_line = True, cw2, pad2, inner2, fs2, ok2
        head_pt = k.head_size or k.h2
        if use_cards and fs < head_pt - 0.05 and math.ceil(n / cols) == 1 and tile_pad_y(cols, pad) > int(0.03 * k.H):
            slim[0] = True
            got = layout(cols, stacked=stacked)
            if got[3] > fs + 0.05 and (got[4] or not one_line):
                cw, pad, inner, fs, one_line = got
            else:
                slim[0] = False
        if not one_line:
            # a figure never wraps (G4, Focus s3: «около 615 тыс. тонн» broke over its tile's accent rule): no size of
            # the row sets every figure on one line in its tile — the figures stand one under the other, each with its
            # label beside it; else the row's figures step down until each one reads on one line
            plan = self._kpi_rows_plan(avail, numbers, fs)
            if plan is not None:
                block_h = plan["h"] + ((extra_block[1] + k.vgap * 1.4) if extra_block else 0)
                y0 = self._place_v(area, int(block_h))
                self._kpi_rows_draw(avail.x, y0, plan)
                if extra_block:
                    paras_g, eh, ew = extra_block
                    ey = y0 + plan["h"] + int(k.vgap * 1.4)
                    for gi, pg in enumerate(paras_g):
                        self.cv.text(Bbox(x=area.x + gi * (ew + k.gap), y=ey, w=ew, h=eh), pg, name="Note")
                self.warnings.append("figures too long for a row of tiles stand one under the other, labels beside")
                return
            heads = [sp[0] if (stacked and sp) else x.value for x, sp in zip(numbers, splits)]
            room = _pt(inner) * (0.88 if use_cards else 0.8)
            for s_ in [x for x in k.steps_down(fs, max(self._dense_floor(), k.small * 0.85)) if x < fs - 0.05] + [max(self._dense_floor(), k.small * 0.85)]:
                fs = s_
                if all(self.figure_para(v, s_, colors.figure, colors.muted)[1] <= room for v in heads):
                    break
        rows = math.ceil(n / cols)
        pad_y = tile_pad_y(cols, pad)
        label_size = k.body
        for ls in (k.lead, k.body):
            label_size = ls
            if ls > k.body + 0.05 and fs <= ls + 0.05:
                continue  # a figure never set under its own label's size (B3: 18 pt figures under 20 pt labels)
            if max(para_lines(self.P(x.label, ls, colors.muted), _pt(inner)) for x in numbers) <= 3:
                break
        fig_h = _emu(fs * max(1.15, k.line))
        rule = 0 if use_cards else _emu(3)
        lab_h = max(self.h([self.P(distinct_label(x.value, x.label, self.o.headline), label_size, colors.muted)], inner) for x in numbers)
        figs = [self.figure_para(sp[0] if (stacked and sp) else x.value, fs, colors.figure, colors.muted)[0] for x, sp in zip(numbers, splits)]
        units = unit_paras(fs, label_size) if stacked else [None] * n
        unit_h = max((self.h([u], inner) - _emu(2) for u in units if u is not None), default=0)
        # the ascender room over the digits is not part of the tile's padding: every figure of the row moves up by
        # the same amount (one size), each one left by its own first glyph's bearing
        tg = min(self._figure_optics(f, fig_h)[1] for f in figs)
        content_h = rule + (int(k.vgap * 0.6) if rule else 0) + fig_h + unit_h - tg + int(k.vgap * 0.35) + lab_h
        max_tile = int((avail.h - k.vgap * (rows - 1)) / rows)
        if content_h + 2 * pad_y > max_tile and label_size > k.body:
            label_size = k.body  # a step down for the labels before the tile outgrows its share
            lab_h = max(self.h([self.P(distinct_label(x.value, x.label, self.o.headline), label_size, colors.muted)], inner) for x in numbers)
            if stacked:
                units = unit_paras(fs, label_size)
                unit_h = max((self.h([u], inner) - _emu(2) for u in units if u is not None), default=0)
            content_h = rule + (int(k.vgap * 0.6) if rule else 0) + fig_h + unit_h - tg + int(k.vgap * 0.35) + lab_h
        # the tile holds its figure and label, never shorter than them (a label hanging below its card)
        tile_h = self._grow_to(avail.h, rows, k.vgap, content_h + 2 * pad_y) if use_cards else content_h + 2 * pad_y
        block_h = rows * tile_h + (rows - 1) * k.vgap + ((extra_block[1] + k.vgap * 1.4) if extra_block else 0)
        y0 = self._place_v(area, int(block_h))
        for i, num in enumerate(numbers):
            r, cidx = divmod(i, cols)
            box = Bbox(x=avail.x + cidx * (cw + gap), y=y0 + r * (tile_h + k.vgap), w=cw, h=tile_h)
            y = box.y + pad_y
            if use_cards:
                self.cv.card(box, st)
            else:
                self.cv.rect(Bbox(x=box.x, y=y, w=min(_emu(k.hpt * 0.09), inner), h=rule), colors.accent, name="Rule")
                y += rule + int(k.vgap * 0.6)
            lsb = self._figure_optics(figs[i], fig_h)[0]
            if units[i] is not None:
                # one text: the number over its words, bottom-anchored so the digits keep the row's top line
                own_h = self.h([units[i]], inner) - _emu(2)
                self.cv.text(Bbox(x=box.x + pad - lsb, y=y - tg, w=inner + lsb, h=fig_h + own_h), [figs[i], units[i]], anchor="b", name="Figure")
            else:
                self.cv.text(Bbox(x=box.x + pad - lsb, y=y - tg, w=inner + lsb, h=fig_h), [figs[i]], anchor="b", name="Figure")
            y += fig_h + unit_h - tg + int(k.vgap * 0.35)
            self.cv.text(Bbox(x=box.x + pad, y=y, w=inner, h=max(lab_h, box.y2 - pad_y - y)), [self.P(distinct_label(num.value, num.label, self.o.headline), label_size, colors.muted)], name="Label")
        if extra_block:
            paras_g, eh, ew = extra_block
            ey = y0 + rows * tile_h + (rows - 1) * k.vgap + int(k.vgap * 1.4)
            for gi, pg in enumerate(paras_g):
                self.cv.text(Bbox(x=area.x + gi * (ew + k.gap), y=ey, w=ew, h=eh), pg, name="Note")

    def _kpi_rows_plan(self, avail: Bbox, numbers: list[NumberCallout], fs_row: float) -> Optional[dict]:
        """Figures too long for a row of tiles, one under the other: each figure on one line at the left (one size for
        all, larger than the row could give them), its label beside it at the lead or body size (never larger than
        the figure, at most three lines), hairlines between the rows. None when no such size fits the area."""
        k = self.kit
        colors = k.colors
        n = len(numbers)
        gap = int(k.gap * 1.5)
        between = int(k.vgap * 1.6)
        labels = [distinct_label(x.value, x.label, self.o.headline) for x in numbers]
        for fs in k.figure_sizes(k.figure_cap(n, False, self.strategy), k.h3):
            if fs <= fs_row + 0.05:
                break
            figs = [self.figure_para(x.value, fs, colors.figure, colors.muted) for x in numbers]
            fw = _emu(max(w for _, w in figs) / 0.94)
            if fw > 0.6 * avail.w:
                continue
            lw = avail.w - fw - gap
            fig_h = _emu(fs * max(1.15, k.line))
            for ls in dict.fromkeys((k.lead, k.body)):
                if ls > fs + 0.05:
                    continue  # a label never over its figure's size
                labs = [[self.P(t, ls, colors.muted)] for t in labels]
                if max(para_lines(lp[0], _pt(lw)) for lp in labs) > 3:
                    continue
                rows_h = [max(fig_h, self.h(lp, lw)) for lp in labs]
                total = sum(rows_h) + (n - 1) * between
                if total <= avail.h:
                    return dict(fs=fs, fw=fw, lw=lw, gap=gap, between=between, figs=[p for p, _ in figs], labs=labs, rows_h=rows_h, fig_h=fig_h, w=avail.w, h=total)
        return None

    def _kpi_rows_draw(self, x: int, y0: int, plan: dict) -> None:
        k = self.kit
        y = y0
        fig_h = plan["fig_h"]
        for i, (fig, lp, rh) in enumerate(zip(plan["figs"], plan["labs"], plan["rows_h"])):
            lsb, tg = self._figure_optics(fig, fig_h)
            fy = y + (rh - fig_h) // 2
            self.cv.text(Bbox(x=x - lsb, y=fy, w=plan["fw"] + lsb, h=fig_h), [fig], anchor="b", name="Figure")
            lh = self.h(lp, plan["lw"])
            mid = fy + (tg + fig_h) // 2  # the middle of the digits, not of the ascender room over them
            ly = min(max(y, mid - lh // 2), y + rh - lh)
            self.cv.text(Bbox(x=x + plan["fw"] + plan["gap"], y=ly, w=plan["lw"], h=lh), lp, name="Label")
            if i < len(plan["figs"]) - 1:
                ry = y + rh + plan["between"] // 2
                self.cv.line(x, ry, x + plan["w"], ry, k.colors.divider, 0.75)
            y += rh + plan["between"]

    def big_number(self, area: Bbox, num: NumberCallout, extra: list[str], others: list[NumberCallout]) -> None:
        k = self.kit
        colors = k.colors
        has_side = bool(extra or others)
        left_w = int(area.w * (0.46 if has_side else 0.9))
        fs = k.display
        for s in k.figure_sizes(k.figure_cap(1, False, self.strategy), k.display):
            fs = s
            if self.figure_para(num.value, s, colors.figure, colors.muted)[1] <= _pt(left_w) * 0.88:
                break
        fig_h = _emu(fs * max(1.12, k.line))
        fig = self.figure_para(num.value, fs, colors.figure, colors.muted)[0]
        lsb, tg = self._figure_optics(fig, fig_h)
        label = [self.P(distinct_label(num.value, num.label, self.o.headline), k.lead, colors.text)]
        lab_h = self.h(label, int(left_w * 0.92))
        left_h = fig_h - tg + int(k.vgap * 0.4) + lab_h
        right_paras: list[Para] = []
        rw = area.w - left_w - k.gap * 2
        for o in others:
            right_paras.append(self.P(o.value, k.snap(k.h3 * 1.4, k.h3, k.display), colors.accent, bold=k.bold, space_after=2))
            right_paras.append(self.P(o.label, k.body, colors.muted, space_after=k.body * 1.1))
        size = k.lead if sum(len(t) for t in extra) <= 260 else k.body
        for t in extra:
            right_paras.append(self.P(t, size, colors.text, space_after=size * 0.6, marker="•" if len(extra) > 1 else None, marker_color=colors.accent))
        right_h = self.h(right_paras, rw) if right_paras else 0
        block_h = max(left_h, right_h)
        y0 = self._place_v(area, block_h)
        self.cv.text(Bbox(x=area.x - lsb, y=y0 - tg, w=left_w + lsb, h=fig_h), [fig], anchor="b", name="Figure")
        self.cv.text(Bbox(x=area.x, y=y0 + fig_h - tg + int(k.vgap * 0.4), w=int(left_w * 0.92), h=lab_h), label, name="Label")
        if right_paras:
            x = area.x + left_w + k.gap * 2
            # a hairline divides the figure from its explanation
            self.cv.line(x - k.gap, y0, x - k.gap, y0 + max(left_h, right_h), colors.divider, 1.0)
            self.cv.text(Bbox(x=x, y=y0 + (max(0, (left_h - right_h) // 2) if right_h < left_h else 0), w=rw, h=right_h), right_paras, name="Explanation")

    # ---- lists ---------------------------------------------------------------------------------------------------
    def bullets(self, area: Bbox, texts: list[str], as_list: bool = False) -> None:
        k = self.kit
        texts = [t for t in texts if t and t.strip()]
        if not texts:
            self.warnings.append("empty slide content")
            return
        n = len(texts)
        total = sum(len(t) for t in texts)
        if self.strategy == "visual" and not as_list and 2 <= n <= 6 and max(len(t) for t in texts) <= 140 and area.w >= 0.6 * k.W:
            # a visual variant turns theses into cards: each thesis its own block
            self.cards(area, [SlideItem(title=t) for t in texts], badge="index")
            return
        two_cols = n >= 5 and total > 360 or n >= 7
        col_gap = k.gap * 2
        cols = 2 if two_cols else 1
        col_w = int((area.w - col_gap * (cols - 1)) / cols)
        text_w = col_w if cols == 2 else min(col_w, int(area.w * 0.8))
        per_col = math.ceil(n / cols)
        # the largest size at which the list fits; rows separated by hairlines. Every size the template uses between the
        # roles is tried (a list 0.02 pt too tall for 20 pt takes 18 pt, not the body size 14 pt), and at each size the
        # rows close up (the hairline gaps down to 0.45 of a line) before the type steps down (B3-2)
        sizes = [k.h2, k.lead, k.body, k.small] if n <= 4 and total <= 260 else ([k.lead, k.body, k.small] if total <= 700 else [k.body, k.small])
        sizes = list(dict.fromkeys(sizes[:1] + [x for x in reversed(k.sizes) if sizes[-1] - 0.05 <= x < sizes[0] - 0.05] + sizes))
        sizes.sort(reverse=True)
        if getattr(self, "_boost", 0) >= 1:
            sizes = [x for x in k.sizes if sizes[0] + 0.05 < x <= k.h2 + 0.05][::-1] + sizes
        for size in sizes:
            marker_w = _emu(size * 1.6)
            heights = [self.h([self.P(t, size, k.colors.text)], text_w - marker_w) for t in texts]
            for pad_k in (0.75, 0.6, 0.45):
                row_pad = int(max(size * pad_k, 6) * EMU_PER_PT)
                col_h = max(sum(heights[c * per_col : (c + 1) * per_col]) + (min(per_col, n - c * per_col) - 1) * 2 * row_pad for c in range(cols))
                if col_h <= area.h:
                    break
            if col_h <= area.h:
                break
        # the rows breathe: the gap between them grows (up to 2.5 lines) until the list fills about 60% of the band
        # (on the fill pass, most of it)
        rows_n = max(min(per_col, n), 1)
        share = self.FILL_TARGET if getattr(self, "_boost", 0) >= 2 else 0.6
        if rows_n > 1 and col_h < share * area.h:
            extra = min(int((share * area.h - col_h) / (rows_n - 1) / 2), _emu(size * k.line * 1.25) - row_pad)
            if extra > 0:
                row_pad += extra
                col_h += extra * 2 * (rows_n - 1)
        block_h = col_h
        y0 = self._place_v(area, block_h, fill_top=True)
        for c in range(cols):
            x = area.x + c * (col_w + col_gap)
            y = y0
            chunk = list(enumerate(texts))[c * per_col : (c + 1) * per_col]
            for j, (i, t) in enumerate(chunk):
                hh = heights[i]
                mk = _emu(size * 0.38)
                # a round accent marker on the first line's x-height (one bullet glyph for the deck)
                self.cv.ellipse(Bbox(x=x, y=y + _emu(size * k.line * 0.5) - mk // 2, w=mk, h=mk), k.colors.accent, name="Marker")
                self.cv.text(Bbox(x=x + marker_w, y=y, w=text_w - marker_w, h=hh), [self.P(t, size, k.colors.text)], name="Item")
                y += hh
                if j < len(chunk) - 1:
                    self.cv.line(x + marker_w, y + row_pad, x + text_w, y + row_pad, k.colors.divider, 0.75)
                    y += 2 * row_pad

    def paragraphs(self, area: Bbox, texts: list[str]) -> None:
        k = self.kit
        total = sum(len(t) for t in texts)
        width = min(area.w, int(area.w * 0.78))
        for size in (k.lead, k.body, k.small) if total <= 500 else (k.body, k.small):
            paras = [self.P(t, size, k.colors.text, space_after=size * 0.8) for t in texts]
            hh = self.h(paras, width)
            if hh <= area.h:
                break
        self.cv.text(Bbox(x=area.x, y=self._place_v(area, hh, fill_top=True), w=width, h=hh), paras, name="Text")

    def statement(self, area: Bbox, text: str) -> None:
        """One message — an ask, a conclusion: large type, figures in the accent colour, an accent bar or panel."""
        k = self.kit
        panel = self.strategy == "visual" and bool(k.card.fill or k.card.proto is not None)
        colors = k.card.colors if panel else k.colors
        width = int(area.w * (0.9 if panel else (0.86 if area.w > 0.6 * k.W else 1.0)))
        pad = int(0.06 * k.H) if panel else 0
        bar = 0 if panel else _emu(max(k.hpt * 0.011, 4))
        inset = bar + int(k.gap * 1.2) if bar else 0
        top = k.statement
        if k.head_size:
            # one headline per slide: the statement is larger than running text and a step below the heading — a
            # callout set above the heading turns the hierarchy upside down
            top = min(top, self._under_heading())
        size = top
        for s in k.steps_down(top, k.body):
            size = s
            paras = [Para(highlight_runs(text, s, colors.text, colors.accent if s >= 18 else colors.accent_text, bold=False, accent_bold=k.bold, font=k.font), space_after=0)]
            hh = self.h(paras, width - inset - 2 * pad)
            if hh + 2 * pad <= area.h * 0.8 and (para_lines(paras[0], _pt(width - inset - 2 * pad)) <= 5 or s <= k.lead):
                break
        block_h = hh + 2 * pad
        y0 = self._place_v(area, block_h)
        if panel:
            self.cv.card(Bbox(x=area.x, y=y0, w=width, h=block_h), k.card)
        elif bar:
            self.cv.rect(Bbox(x=area.x, y=y0 + _emu(size * 0.15), w=bar, h=max(hh - _emu(size * 0.3), bar)), k.colors.accent, name="Accent bar")
        self.cv.text(Bbox(x=area.x + inset + pad, y=y0 + pad, w=width - inset - 2 * pad, h=hh), paras, name="Statement", anchor="ctr" if panel else "t")

    def quote(self, area: Bbox, text: str, author: Optional[str]) -> None:
        """A large accent quotation mark hanging to the left of the words (set at statement size, without their own
        quotes), the author under them."""
        k = self.kit
        mark_size = k.figure_sizes(0.16 * k.hpt, k.display)[0]
        mark_w = _emu(mark_size * 0.75)
        width = int(area.w * 0.82) - mark_w
        words = text.strip().strip("«»\"“”„")
        top = min(k.statement, self._under_heading()) if k.head_size else k.statement
        size = top
        for s_ in k.steps_down(top, k.body):
            size = s_
            body = [self.P(words, s_, k.colors.text)]
            hh = self.h(body, width)
            if hh <= area.h * 0.7 and para_lines(body[0], _pt(width)) <= 5:
                break
        auth = [self.P(author, k.lead, k.colors.muted)] if author else []
        ah = self.h(auth, width) if auth else 0
        block = hh + (int(k.vgap) + ah if auth else 0)
        y = self._place_v(area, block)
        # an opening quote that hangs at the cap height of the first line (a guillemet sits on the x-height and
        # reads as a small bracket at any size)
        self.cv.text(Bbox(x=area.x, y=y - _emu(mark_size * 0.08), w=mark_w, h=_emu(mark_size * k.line)), [self.P("“", mark_size, k.colors.accent, bold=k.bold)], name="Quote mark")
        x = area.x + mark_w
        self.cv.text(Bbox(x=x, y=y, w=width, h=hh), body, name="Quote")
        if auth:
            self.cv.rect(Bbox(x=x, y=y + hh + int(k.vgap * 0.9), w=_emu(k.hpt * 0.06), h=_emu(2)), k.colors.accent, name="Rule")
            self.cv.text(Bbox(x=x, y=y + hh + int(k.vgap * 1.3), w=width, h=ah), auth, name="Author")

    # ---- process / agenda / columns -------------------------------------------------------------------------------
    def process(self, area: Bbox, items: list[SlideItem]) -> None:
        k = self.kit
        n = len(items)
        if all(not i.text and not i.bullets for i in items) and max(len(i.title.split()) for i in items) > 6:
            items = [SlideItem(title="", text=i.title) for i in items]  # steps told as sentences read as text
        if n > 6:
            self.cards(area, items, badge="number")
            return
        # five or six steps in a row: a narrower gutter and padding give each step the width its longest word needs
        # at a readable size («Корректировка» at the body size)
        gap = k.gap if n <= 4 else int(k.gap * 0.6)
        cw = int((area.w - gap * (n - 1)) / n)
        colors = k.colors
        st = k.card
        in_cards = self.strategy == "visual" and bool(st.fill or st.proto is not None)
        pad = int(min(cw * (0.09 if n <= 4 else 0.07), 0.045 * k.H)) if in_cards else 0
        inner = cw - 2 * pad
        tcolors = st.colors if in_cards else colors
        bodies = [b for i in items for b in ([i.text] if i.text else []) + list(i.bullets)]
        fitted = False
        for ts, bs in self._size_pairs():
            # a word wider than the step's column breaks in the middle («Корректиров|ка»): a smaller size, or a grid
            if not self.fits_width([i.title for i in items], ts, k.bold, inner) or not self.fits_width(bodies, bs, False, inner):
                continue
            fitted = True
            text_h = max(self.h(self._card_paras(it, ts, bs, tcolors), inner) for it in items)
            d = self._badge_size(ts)
            total = d + int(k.vgap * 0.9) + text_h + 2 * pad
            t_lines = max(para_lines(self.P(i.title, ts, tcolors.text, bold=k.bold), _pt(inner)) for i in items)
            titled = any(i.text or i.bullets for i in items)
            if total <= area.h and t_lines <= (2 if titled else 4):
                break
        if n >= 5 and (not fitted or bs < k.body - 0.05) and self._steps_grid(area, items, better_than=bs if fitted else 0.0):
            return  # too many steps for one row at a readable size: a grid of steps, each its badge beside its text
        if not fitted:
            # too narrow for a row of steps (a word would break): numbered cards in a grid
            self.cards(area, items, badge="number")
            return
        block_h = self._grow_to(area.h, 1, 0, total) if in_cards else total
        y0 = self._place_v(area, block_h)
        # the axis runs through the badges' centres, behind them
        cy = y0 + d // 2
        first_cx = area.x + (pad if in_cards else 0) + d // 2
        last_cx = area.x + (n - 1) * (cw + gap) + (pad if in_cards else 0) + d // 2
        if n > 1 and not in_cards:
            self.cv.line(first_cx, cy, last_cx, cy, colors.accent, 1.5, name="Axis")
        for i, it in enumerate(items):
            x = area.x + i * (cw + gap)
            by = y0
            if in_cards:
                # the step is a card; its badge sits inside it, a padding from the corner (never on the edge)
                self.cv.card(Bbox(x=x, y=y0, w=cw, h=block_h), st)
                by = y0 + pad
            bx = x + pad
            el = self.cv.ellipse(Bbox(x=bx, y=by, w=d, h=d), colors.accent)
            self._label_in(el, str(i + 1), self._badge_digit(), k.on_accent(colors.accent, colors.text))
            ty = by + d + int(k.vgap * 0.9)
            paras = self._card_paras(it, ts, bs, tcolors)
            self.cv.text(Bbox(x=bx, y=ty, w=inner, h=max(text_h, y0 + block_h - pad - ty)), paras, name="Step")

    def _dense_steps_beside(self, comp: str, area: Bbox, n0: int, state: tuple) -> None:
        """G4-20 (long visual s8 on Sidebar43, Focus, Nature, Handdrawn, Marketing): six steps in a row with the budget
        lines under them came out under 2 % of the slide height in the dense pass (the row is bound by its longest
        word, «Корректировка», the grid of steps by its height). The steps then stand one under the other at the
        left, the lines beside them at the right, at ≥ 2 % — taken only when the dense setting was under it and the
        new one fits."""
        c = self.o.content
        if not (c.items or c.columns) or not c.bullets or self._least_text(n0) >= self._dense_floor_2pct() - 0.05:
            return
        o0, w0, notes0 = state
        keep = (list(self.warnings), self.o, getattr(self, "_notes", None))
        dense_els = list(self.cv.tree)[n0:]
        dense_added = [el for el in self.cv.added if any(el is d for d in dense_els)]
        self._beside_done = False
        self._undo(n0)
        self.o, self.warnings = o0, list(w0)
        self._restore_notes(notes0)
        self._beside, self._floor_2pct = True, True
        try:
            self._compose_body(comp, area)
        finally:
            self._beside, self._floor_2pct = False, False
        if self._beside_done and not self._overflows(n0, area) and self._least_text(n0) >= self._dense_floor_2pct() - 0.05:
            self.warnings.append("steps one under the other, their lines beside them: the row set them under 2 % of the slide")
            return
        # the dense setting stays
        self._undo(n0)
        for el in dense_els:
            self.cv.tree.append(el)
        self.cv.added.extend(dense_added)
        self.warnings, self.o, self._notes = keep[0], keep[1], keep[2]

    def _steps_beside(self, area: Bbox, items: list[SlideItem], lines: list[str]) -> bool:
        """Steps as a column — a small number badge on an axis, the step's title leading its line in bold («1-й месяц —
        учет показателей…») — and the slide's lines as a column beside them, for a slide whose row of steps would set
        its words too small. The largest size (never under the dense floor) at which both columns fit, the steps' column
        taking 58 %, 64 % or 70 % of the width."""
        self._beside_done = False
        k = self.kit
        n = len(items)
        colors = k.colors
        gap = k.gap * 2
        floor = self._dense_floor()
        between = int(k.vgap * 0.8)
        dgap = int(k.gap * 0.7)

        def step_paras(bs: float) -> list[list[Para]]:
            out = []
            for it in items:
                if not it.title:
                    out.append(self._card_paras(it, bs, bs, colors))
                    continue
                text = it.text or ""
                if text[:2] != text[:2].upper():
                    text = text[:1].lower() + text[1:]
                runs = [Run(typeset(it.title), bs, colors.heading, k.bold, k.font)] + ([Run(typeset(" — " + text), bs, colors.text, False, k.font)] if text else [])
                out.append([Para(runs)] + [self.P(b, bs, colors.text, marker="•", marker_color=colors.accent) for b in it.bullets])
            return out

        sizes = list(dict.fromkeys(bs for _, bs in self._size_pairs() if bs >= floor - 0.05))
        for bs in sizes:
            d = _emu(max(bs * 1.9, 16.0))
            paras = step_paras(bs)
            offs = [max(0, (d - _emu(ps[0].size * k.line)) // 2) if ps else 0 for ps in paras]  # the first line on the badge
            line_paras = [self.P(t, bs, colors.text, space_after=bs * 0.5, marker="•", marker_color=colors.accent) for t in lines]
            for share in (0.58, 0.64, 0.7):
                left_w = int((area.w - gap) * share)
                right_w = area.w - gap - left_w
                inner = left_w - d - dgap
                if not self.fits_width([f"{i.title} {i.text}" for i in items], bs, k.bold, inner) or not self.fits_width(lines, bs, False, right_w - _emu(bs * 1.1)):
                    continue
                hs = [max(d, off + self.h(ps, inner)) for ps, off in zip(paras, offs)]
                steps_h = sum(hs) + between * (n - 1)
                lines_h = self.h(line_paras, right_w)
                block_h = max(steps_h, lines_h)
                if block_h > area.h:
                    continue
                y0 = self._place_v(area, block_h)
                x = area.x
                if n > 1:
                    self.cv.line(x + d // 2, y0 + d // 2, x + d // 2, y0 + sum(hs[:-1]) + between * (n - 1) + d // 2, colors.accent, 1.25, name="Axis")
                y = y0
                for i, (ps, hh, off) in enumerate(zip(paras, hs, offs)):
                    el = self.cv.ellipse(Bbox(x=x, y=y, w=d, h=d), colors.accent)
                    self._label_in(el, str(i + 1), bs, k.on_accent(colors.accent, colors.text))  # the number at the steps' size
                    self.cv.text(Bbox(x=x + d + dgap, y=y + off, w=inner, h=hh - off), ps, name="Step")
                    y += hh + between
                rx = area.x + left_w + gap
                self.cv.line(rx - gap // 2, y0, rx - gap // 2, y0 + block_h, colors.divider, 1.0)
                self.cv.text(Bbox(x=rx, y=y0, w=right_w, h=lines_h), line_paras, name="Note")
                self._beside_done = True
                return True
        return False

    def _steps_grid(self, area: Bbox, items: list[SlideItem], better_than: float) -> bool:
        """Five to eight steps too many for one row at the body size: rows of three (four for seven or eight), each
        step its number badge at the left of its title and text — a row as tall as its text, not a badge stacked over
        it, so two rows hold the type larger than one row of narrow columns did. Drawn only when its text comes out
        larger than `better_than`; returns whether it was drawn."""
        k = self.kit
        n = len(items)
        cols = 3 if n <= 6 else 4
        rows = math.ceil(n / cols)
        gap = k.gap
        st = k.card
        boxed = bool(st.fill or st.proto is not None)
        cw = int((area.w - gap * (cols - 1)) / cols)
        pad = int(min(cw * 0.07, 0.035 * k.H)) if boxed else 0
        if getattr(self, "_dense", False) and boxed:
            pad = int(min(cw * 0.05, 0.018 * k.H))  # the dense pass: a slimmer card around each step, before the type steps down
        d = _emu(max(k.hpt * 0.06, 28.0))
        dgap = int(k.gap * 0.6)
        inner = cw - 2 * pad - d - dgap
        colors = st.colors if boxed else k.colors
        lead_off = 0 if boxed else int(k.vgap * 0.6)  # under the hairline that opens a step without a card
        bodies = [b for i in items for b in ([i.text] if i.text else []) + list(i.bullets)]
        titled = any(i.text or i.bullets for i in items)
        chosen = None
        for ts, bs in self._size_pairs():
            if bs <= better_than + 0.05:
                break
            if not self.fits_width([i.title for i in items], ts, k.bold, inner) or not self.fits_width(bodies, bs, False, inner - _emu(bs * 1.1)):
                continue
            first = _emu((ts if any(i.title for i in items) else bs) * k.line)
            text_h = max(self.h(self._card_paras(it, ts, bs, colors), inner) for it in items)
            shift = max(0, (d - first) // 2)  # the first line centred on the badge
            row_h = max(d, shift + text_h) + 2 * pad + lead_off
            t_lines = max(para_lines(self.P(i.title, ts, colors.text, bold=k.bold), _pt(inner)) for i in items if i.title) if any(i.title for i in items) else 0
            if rows * row_h + (rows - 1) * k.vgap <= area.h and t_lines <= (2 if titled else 4):
                chosen = (ts, bs, text_h, shift, row_h)
                break
        if chosen is None:
            return False
        ts, bs, text_h, shift, row_h = chosen
        row_h = self._grow_to(area.h, rows, k.vgap, row_h)
        y0 = self._place_v(area, rows * row_h + (rows - 1) * k.vgap)
        digit = k.snap(max(_pt(d) * 0.5, 12.0), 10.0, _pt(d) * 0.62)
        on_accent = k.on_accent(colors.accent, colors.text)
        last_n = n - cols * (rows - 1)
        for i, it in enumerate(items):
            r, c_ = divmod(i, cols)
            x_off = int((cols - last_n) * (cw + gap) / 2) if r == rows - 1 and last_n < cols else 0
            box = Bbox(x=area.x + x_off + c_ * (cw + gap), y=y0 + r * (row_h + k.vgap), w=cw, h=row_h)
            if boxed:
                self.cv.card(box, st)
            else:
                # no card style: a hairline over each step keeps the grid
                self.cv.rect(Bbox(x=box.x, y=box.y, w=box.w, h=max(_emu(1.5), 1)), colors.divider, name="Rule")
            top = box.y + pad + lead_off
            el = self.cv.ellipse(Bbox(x=box.x + pad, y=top, w=d, h=d), colors.accent)
            self._label_in(el, str(i + 1), digit, on_accent)
            tx = box.x + pad + d + dgap
            ty = top + shift
            self.cv.text(Bbox(x=tx, y=ty, w=inner, h=max(text_h, box.y2 - pad - ty)), self._card_paras(it, ts, bs, colors), name="Step")
        return True

    def agenda(self, area: Bbox, items: list[SlideItem]) -> None:
        k = self.kit
        n = len(items)
        cols = 1 if n <= 4 else 2
        per = math.ceil(n / cols)
        col_gap = k.gap * 3
        col_w = int((area.w - col_gap * (cols - 1)) / cols)
        for ts in (k.statement if n <= 3 else k.h2, k.h2, k.lead, k.h3, k.body):
            idx_w = _emu(ts * 2.4)
            heights = [self.h([self.P(it.title, ts, k.colors.text)], col_w - idx_w) for it in items]
            row_gap = _emu(ts * 0.9)
            col_h = max(sum(heights[c * per : (c + 1) * per]) + (min(per, n - c * per) - 1) * 2 * row_gap for c in range(cols))
            if col_h <= area.h:
                break
        y0 = self._place_v(area, col_h)
        for c in range(cols):
            x = area.x + c * (col_w + col_gap)
            y = y0
            chunk = list(range(c * per, min((c + 1) * per, n)))
            for j, i in enumerate(chunk):
                hh = heights[i]
                self.cv.text(Bbox(x=x, y=y, w=idx_w, h=hh), [self.P(f"{i + 1:02d}", ts, k.colors.accent, bold=k.bold)], name="Index")
                self.cv.text(Bbox(x=x + idx_w, y=y, w=col_w - idx_w, h=hh), [self.P(items[i].title, ts, k.colors.text)], name="Agenda item")
                y += hh
                if j < len(chunk) - 1:
                    self.cv.line(x, y + row_gap, x + col_w, y + row_gap, k.colors.divider, 0.75)
                    y += 2 * row_gap

    def columns(self, area: Bbox, items: list[SlideItem]) -> None:
        """Two or three columns side by side (options, before/after, pros/cons): the first column carries the accent."""
        k = self.kit
        n = min(len(items), 3)
        items = items[:n]
        gap = k.gap
        cw = int((area.w - gap * (n - 1)) / n)
        st = k.card
        rule_h = _emu(3) + int(k.vgap * 0.6)

        def parts(ts: float, bs: float, pad: int, ws: list[int]):
            inners = [w - 2 * pad for w in ws]
            # titles share one band (as tall as the longest title), so every body starts on one line across the columns;
            # a column's figure («22 770 ₽ экономии») is set large over its lines, never left out
            heads = [self._card_paras(SlideItem(title=it.title), ts, bs, st.colors) if it.title else [] for it in items]
            bodies = [self._column_figure(it, ts, st.colors) + self._card_paras(SlideItem(title="", text=it.text, bullets=it.bullets), ts, bs, st.colors) for it in items]
            head_h = max((self.h(hp, iw) for hp, iw in zip(heads, inners) if hp), default=0)
            head_gap = int(bs * 0.55 * EMU_PER_PT) if head_h else 0
            body_h = max((self.h(bp, iw) for bp, iw in zip(bodies, inners) if bp), default=0)
            lines = max(sum(para_lines(p, _pt(iw)) for p in bp) for bp, iw in zip(bodies, inners)) if bodies else 0
            return heads, bodies, head_h, head_gap, body_h, rule_h + head_h + head_gap + body_h + 2 * pad, lines

        # ≤ 11 lines a column — a list of six two-line points is twelve lines, not a wall (G4-20, Focus long compact s8:
        # the cap sent it to 9 pt with a third of each card empty)
        max_lines = max(11, 2 * max((len(it.bullets) + (1 if it.text else 0) for it in items), default=0))

        def choose(pad: int, ws: list[int]):
            # the largest pair at which the columns fit the area (a little air kept under them) in ≤ max_lines each
            best = None
            for ts, bs in self._size_pairs():
                got = parts(ts, bs, pad, ws)
                if best is None or got[5] < best[1][5]:
                    best = ((ts, bs), got)
                if got[5] <= area.h * 0.96 and got[6] <= max_lines:
                    return ((ts, bs), got), True
            return best, False

        ws = [cw] * n
        pad = int(min(cw * 0.07, 0.05 * k.H))
        best, ok = choose(pad, ws)
        if best[0][1] < k.body - 0.05 or not ok:
            # small type inside roomy columns: a tighter padding gives the text the room first
            tight = max(int(pad * 0.6), int(0.022 * k.H))
            alt, ok2 = choose(tight, ws)
            if (ok2 and not ok) or (ok2 == ok and alt[0][1] > best[0][1] + 0.05):
                best, pad = alt, tight
        if n == 2 and (best[0][1] < k.body - 0.05 or not ok):
            # still small: two columns of unequal length (five short amounts beside six two-line steps) share the width
            # by their words — the longer list takes up to 62 % and both read a step or more larger (G4-20)
            size = [len(it.title or "") + len(it.text or "") + sum(len(b) for b in it.bullets) for it in items]
            share = min(max(size[0] / max(sum(size), 1), 0.38), 0.62)
            if abs(share - 0.5) > 0.04:
                w0 = int((area.w - gap) * share)
                ws2 = [w0, area.w - gap - w0]
                alt, ok2 = choose(pad, ws2)
                if (ok2 and not ok) or (ok2 == ok and alt[0][1] > best[0][1] + 0.05):
                    best, ws, ok = alt, ws2, ok2
        (ts, bs), (heads, bodies, head_h, head_gap, body_h, need, _) = best
        # the columns are as tall as their content, never a tall empty frame — and never shorter than it: when even
        # the smallest step does not fit the area, the cards grow with their text (the lines under them follow the
        # cards' real bottom) rather than the text running out of its card
        ch = self._grow_to(area.h, 1, 0, need)
        y0 = self._place_v(area, ch, fill_top=True)
        x = area.x
        for i, it in enumerate(items):
            box = Bbox(x=x, y=y0, w=ws[i], h=ch)
            x += ws[i] + gap
            inner = ws[i] - 2 * pad
            if st.fill or st.line or st.proto is not None:
                self.cv.card(box, st)
            self.cv.rect(Bbox(x=box.x + pad, y=box.y + pad, w=_emu(k.hpt * 0.07), h=_emu(3)), st.colors.accent, name="Rule")
            y = box.y + pad + rule_h
            if heads[i]:
                self.cv.text(Bbox(x=box.x + pad, y=y, w=inner, h=head_h), heads[i], name="Column title")
            y += head_h + head_gap
            if bodies[i]:
                self.cv.text(Bbox(x=box.x + pad, y=y, w=inner, h=max(body_h, box.y2 - pad - y)), bodies[i], name="Column")

    def _column_figure(self, it: SlideItem, ts: float, colors: Colors) -> list[Para]:
        """A column's own figure (SlideItem.number) over its lines, when its title and lines do not say it already."""
        num = (it.number or "").strip()
        if not num or num in f"{it.title} {it.text} {' '.join(it.bullets)}":
            return []
        para, _ = self.figure_para(num, max(ts * 1.25, self.kit.h3), colors.figure, colors.muted)
        para.space_after = ts * 0.4
        return [para]

    # ---- the slide's conclusion and footnote (Agent v2) ----------------------------------------------------------
    def _footnote_size(self) -> float:
        """Small print that still reads from the back of the room: about 2% of the slide height, never under 8 pt,
        never above the deck's small size."""
        k = self.kit
        return k.snap(max(0.02 * k.hpt, 8.0), max(0.017 * k.hpt, 7.5), max(k.small, 8.0))

    def _takeaway_plan(self, text: str, width: int, figure: bool = True, sizes: Optional[tuple] = None) -> dict:
        """How the conclusion is set: one line at the lead size (a step down, then two lines, when it is longer), on a
        strip of the template's card colour with an accent bar at its left — or, when the template's cards have no
        fill of their own, as an accent bar and the line beside it. `figure=False`: never the visual variant's large
        key figure; `sizes`: those sizes only (the slimmer forms of `_takeaway_fit`)."""
        k = self.kit
        st = k.card
        strip = bool(st.fill) and contrast_ratio(st.fill, k.colors.ground) >= 1.04
        colors = st.colors if strip else k.colors
        bar_w = max(_emu(3), int(0.006 * k.H))
        best = None
        cand = None
        text = text[:1].upper() + text[1:] if text[:1].islower() else text
        if text.endswith(".") and not text.endswith("..") and not ABBR_END_RE.search(text):
            text = text[:-1]  # a conclusion line, like a headline, has no period («… на 10 п. п.» keeps it)
        deck = self._deck_takeaway_size(width, strip, bar_w) if sizes is None else None
        for size in (sizes if sizes is not None else [deck] if deck else dict.fromkeys((k.lead, k.body, k.small))):
            pad_x = _emu(size * 1.0) if strip else 0
            pad_y = _emu(size * 0.62) if strip else 0
            bar_gap = _emu(size * 0.8)
            text_w = width - 2 * pad_x - bar_w - bar_gap
            runs = highlight_runs(text, size, colors.heading, colors.accent if size >= 18 else colors.accent_text, bold=k.bold, accent_bold=k.bold, font=k.font)
            para = Para(runs)
            lines = para_lines(para, _pt(text_w))
            cand = dict(size=size, strip=strip, colors=colors, pad_x=pad_x, pad_y=pad_y, bar_w=bar_w, bar_gap=bar_gap, text_w=text_w, para=para, lines=lines)
            if lines == 1 and size >= k.body - 0.05:
                best = cand
                break
            if lines <= 2 and size <= k.body + 0.05:
                best = cand  # two lines at the body size rather than one line of small type
                break
            if deck:
                best = cand  # the deck's one size: one or two lines
                break
        if best is None:
            best = cand
        text_h = self.h([best["para"]], best["text_w"])
        best["text_h"] = text_h
        best["h"] = text_h + 2 * best["pad_y"]
        fig = _aside_figure(text) if self.strategy == "visual" and figure else None
        if fig and self._figure_on_slide(fig):
            fig = None  # the slide's tiles already show it large («80%»): the strip does not say it twice (B3)
        if fig:
            # the visual variant's strip leads with the conclusion's key figure, set large in the template's figure
            # colour where the other variants have the accent bar: the eye takes the number first, then the sentence
            size, colors = best["size"], best["colors"]
            fs = k.snap(max(size * 2.0, k.h2), k.h2, max(k.statement, k.h2))
            para_f, fw = self.figure_para(fig, fs, colors.figure, colors.muted)
            fig_w = _emu(fw * 1.08) + _emu(4)  # a figure never wraps: its box keeps the renderer's slack
            fig_gap = _emu(size * 1.2)
            text_w = width - 2 * best["pad_x"] - fig_w - fig_gap
            if text_w >= 0.45 * width:
                text = label_without_figure(text, fig)
                para = Para(highlight_runs(text, size, colors.heading, colors.accent if size >= 18 else colors.accent_text, bold=k.bold, accent_bold=k.bold, font=k.font))
                t_h = self.h([para], text_w)
                fig_h = _emu(fs * max(1.15, k.line))
                inner = max(t_h, fig_h)
                best.update(fig=(para_f, fig_w, fig_h, fig_gap), para=para, text_w=text_w, text_h=t_h, h=inner + 2 * best["pad_y"])
        return best

    def _figure_on_slide(self, fig: str) -> bool:
        """Whether a figure is one the slide already sets large: a value of its number tiles, cards or columns."""
        c = self.o.content
        want = re.sub(r"[^\d%]", "", fig or "")
        if not re.search(r"\d", want):
            return False
        shown = [x.value for x in c.numbers] + [getattr(it, "number", None) or "" for it in list(c.items) + list(c.columns)]
        return any(re.sub(r"[^\d%]", "", v or "") == want for v in shown)

    def _deck_takeaway_size(self, width: int, strip: bool, bar_w: int) -> Optional[float]:
        """One size for every conclusion strip of the deck: the largest step (lead, body) at which every takeaway of the
        deck fits in two lines — the strip does not jump between 20 and 14 pt from slide to slide."""
        k = self.kit
        takes = [" ".join((s.takeaway or "").split()) for s in self.outline.slides if (s.takeaway or "").strip() and s.kind not in (PatternKind.title, PatternKind.thanks, PatternKind.section)]
        if not takes:
            return None
        chosen = None
        for size in dict.fromkeys((k.lead, k.body)):
            pad_x = _emu(size * 1.0) if strip else 0
            text_w = width - 2 * pad_x - bar_w - _emu(size * 0.8)
            if all(para_lines(Para(highlight_runs(x, size, "000000", "000000", bold=k.bold, font=k.font)), _pt(text_w)) <= 2 for x in takes):
                chosen = size
                break
        return chosen

    def _reserve_notes(self, area: Bbox) -> Bbox:
        """Room for the footnote (small print at the foot of the area, just above the template's footer) and the
        conclusion over it; the content is composed in what is left."""
        self._notes = None
        self._aside = None
        self._aside = self._takeaway_aside()
        # the compact variant says its conclusion first: a bold line over the content (a chart's text column leads
        # with it instead), no strip at the foot
        self._take_lead = (
            self.strategy == "compact" and bool((self.o.takeaway or "").strip()) and not self._aside
            and getattr(self, "_comp", None) not in (None, "chart", "chart_text", "chart_pair")
        )
        foot = " ".join((self.o.footnote or "").split())
        foot = foot[:1].upper() + foot[1:] if foot[:1].islower() else foot
        take = "" if (self._aside or self._take_lead) else " ".join((self.o.takeaway or "").split())
        # the hard floor of the content: the area's foot, above the footnote when there is one (the conclusion's room
        # is not a floor — a block that needs it takes it and the conclusion is said aloud, `_draw_notes`)
        self._floor_y = area.y2
        if not (foot or take):
            return area
        k = self.kit
        bottom = area.y2
        foot_box = None
        if foot:
            size = self._footnote_size()
            para = [self.P(foot, size, k.colors.muted)]
            fh = self.h(para, area.w)
            foot_box = (Bbox(x=area.x, y=bottom - fh, w=area.w, h=fh), para)
            bottom -= fh + int(k.vgap * (0.8 if take else 1.2))
            self._floor_y = bottom
        plan = None
        if take:
            plan = self._takeaway_plan(take, area.w)
            plan["y_max"] = bottom - plan["h"]
            bottom = plan["y_max"] - int(k.vgap * 1.2)
        self._notes = (foot_box, plan, area)
        return Bbox(x=area.x, y=area.y, w=area.w, h=max(bottom - area.y, int(area.h * 0.4)))

    def _takeaway_aside(self) -> bool:
        """Whether the slide's conclusion stands in the chart's side column instead of a strip under the content. The
        structured variant keeps the strip. The visual variant sets a lone chart's conclusion beside it, its key figure
        large, when the column has room for it (at most two other lines). The compact variant leads its text column
        with it (a chart or a pair of charts): no strip, the chart keeps the height. Never next to a table or a
        formula, and only on a slide composed as a chart (a chart slide the plan set otherwise keeps its strip)."""
        cached = getattr(self, "_aside", None)
        if cached is not None:
            return cached  # decided once, when the notes' room was reserved (a pair that falls back to one chart keeps it)
        c = self.o.content
        if not (self.o.takeaway or "").strip() or getattr(self, "_comp", None) not in ("chart", "chart_text", "chart_pair"):
            return False
        if c.chart is None or c.table is not None or (c.formula or "").strip():
            return False
        pair = getattr(self, "_comp", None) == "chart_pair" and c.chart2 is not None
        side = [t for t in list(c.bullets) + list(c.paragraphs) if t.strip()]
        if self.strategy == "compact":
            return True
        if self.strategy == "visual":
            # beside a lone chart when the column has room; a pair with a pie (which needs the whole width) sets it
            # as the first callout of the row under the charts, in the strip's place
            return self._pair_has_pie() if pair else len(side) <= 2
        return False

    def _lead_line(self, area: Bbox, text: str) -> Bbox:
        """The compact variant's conclusion over the content: one or two bold lines in the heading colour, figures in
        the accent, an accent bar at their left (the strip's bar, without the strip). Returns the area under it."""
        k = self.kit
        colors = k.colors
        text = text[:1].upper() + text[1:]
        if text.endswith(".") and not text.endswith("..") and not ABBR_END_RE.search(text):
            text = text[:-1]
        bar_w = max(_emu(3), int(0.006 * k.H))
        width = int(area.w * 0.9)
        for size in dict.fromkeys((k.lead, k.body)):
            gap = _emu(size * 0.8)
            para = Para(highlight_runs(text, size, colors.heading, colors.accent if size >= 18 else colors.accent_text, bold=True, accent_bold=True, font=k.font))
            if para_lines(para, _pt(width - bar_w - gap)) <= 2:
                break
        text_w = width - bar_w - gap
        hh = self.h([para], text_w)
        self.cv.rect(Bbox(x=area.x, y=area.y + _emu(size * 0.12), w=bar_w, h=max(hh - _emu(size * 0.24), bar_w * 3)), colors.accent, name="Conclusion bar")
        self.cv.text(Bbox(x=area.x + bar_w + gap, y=area.y, w=text_w, h=hh), [para], name="Conclusion")
        dy = hh + int(k.vgap * 1.2)
        return Bbox(x=area.x, y=area.y + dy, w=area.w, h=max(area.h - dy, int(area.h * 0.5)))

    def _takeaway_to_strip(self, area: Bbox, size: Optional[float] = None) -> Bbox:
        """The chart the conclusion was to stand beside could not be drawn: the conclusion takes its strip under what
        the slide shows instead (the footnote keeps its place). Returns the area left above the strip."""
        take = " ".join((self.o.takeaway or "").split())
        notes = getattr(self, "_notes", None)
        if not take or (notes and notes[1] is not None):
            return area
        k = self.kit
        plan = self._takeaway_plan(take, area.w) if size is None else self._takeaway_plan(take, area.w, figure=False, sizes=(size,))
        plan["y_max"] = area.y2 - plan["h"]
        self._notes = (notes[0] if notes else None, plan, notes[2] if notes else area)
        return Bbox(x=area.x, y=area.y, w=area.w, h=max(plan["y_max"] - int(k.vgap * 1.2) - area.y, int(area.h * 0.4)))

    def _content_bottom(self, n0: int) -> Optional[int]:
        from verstka.rendering.deck import element_bbox

        bottoms = []
        for el in list(self.cv.tree)[n0:]:
            b = element_bbox(el)
            if b and b[2] > 0 and b[3] > 0:
                bottoms.append(b[1] + b[3])
        return max(bottoms) if bottoms else None

    def _takeaway_fit(self, text: str, width: int, room_h: int, smaller: float = 0.0) -> Optional[dict]:
        """The conclusion's plan at `width` that fits `room_h`: its own form, then without the visual variant's large
        figure; `smaller`: at that size instead (the body size, at last the small one: two or three lines) — None when
        none fits."""
        k = self.kit
        if room_h <= 0 or width <= 0:
            return None
        if smaller:
            plans = [lambda: self._takeaway_plan(text, width, figure=False, sizes=(smaller,))]
        else:
            plans = [lambda: self._takeaway_plan(text, width), lambda: self._takeaway_plan(text, width, figure=False)]
        for mk in plans:
            p_ = mk()
            if p_["h"] <= room_h and p_["text_w"] >= 0.3 * k.W:
                return p_
        return None

    def _conclusion_spot(self, n0: int, last: bool = False):
        """Where the conclusion goes under what was composed after the first `n0` elements: (plan, x, y, w, foot_box)
        — its reserved place, or right under the content; else the same words set slimmer (no large figure, a size
        down) in the room left under the content; else a narrower strip in the free room under an illustration the
        area was cut above (`foot_room`, B2), the footnote moved to that room's foot. None: no room on the slide."""
        notes = getattr(self, "_notes", None)
        if not notes or notes[1] is None:
            return None
        foot_box, plan, area = notes
        k = self.kit
        bottom = self._content_bottom(n0)
        room_below = area.y2 - (foot_box[0].h + int(k.vgap * 0.8) if foot_box else 0)
        if bottom is None or not (bottom + int(k.vgap * 0.6) > plan["y_max"] and bottom + int(k.vgap * 0.6) + plan["h"] > room_below):
            y = plan["y_max"] if bottom is None else min(plan["y_max"], bottom + int(k.vgap * 1.4))
            if bottom is not None and bottom + int(k.vgap * 1.4) > plan["y_max"] and bottom + int(k.vgap * 1.4) + plan["h"] <= area.y2 - (foot_box[0].h if foot_box else 0):
                y = bottom + int(k.vgap * 1.4)  # the content grew past its share: the conclusion follows it, never over it
            if bottom is not None and plan["y_max"] - (bottom + int(k.vgap * 1.4)) > 0.28 * area.h:
                # a short block leaves a large empty band: the conclusion stands at the foot of the area (its reserved
                # place), so the slide has a top and a bottom rather than everything in its upper half
                y = plan["y_max"]
            return plan, area.x, max(y, area.y), area.w, foot_box
        take = " ".join((self.o.takeaway or "").split())
        top = bottom + int(k.vgap * 1.0)
        # a narrower strip in the free room under the art the area was cut above (trees, a grass band at one side)
        fr = getattr(self, "foot_room", None)
        room = None
        if fr is not None and fr.y2 > room_below + int(0.02 * k.H) and fr.y <= area.y2 + int(0.02 * k.H):
            x, w = (area.x, fr.x2 - area.x) if abs(fr.x - area.x) <= int(0.02 * k.W) else (fr.x, fr.w)
            w = min(w, area.x2 - x)
            fb2, floor = None, fr.y2
            if foot_box is not None:
                para_f = foot_box[1]
                fh = self.h(para_f, w)
                fb2 = (Bbox(x=x, y=fr.y2 - fh, w=w, h=fh), para_f)
                floor = fr.y2 - fh - int(k.vgap * 0.8)
            if w >= 0.4 * area.w:
                room = (x, w, floor, fb2)
        # its own form (without the large figure: a strip with one is twice as tall) where the content left room, then
        # in the free room by the art; then a size down, in the same order
        for smaller in ((0.0, k.body) + ((k.small,) if last and k.small < k.body - 0.05 else ())):
            got = self._takeaway_fit(take, area.w, room_below - top, smaller)
            if got is not None:
                return got, area.x, min(bottom + int(k.vgap * 1.4), room_below - got["h"]), area.w, foot_box
            if room is not None:
                x, w, floor, fb2 = room
                got = self._takeaway_fit(take, w, floor - top, smaller)
                if got is not None:
                    return got, x, min(bottom + int(k.vgap * 1.4), floor - got["h"]), w, fb2
        return None

    def _said(self) -> str:
        """The line `_draw_notes` says aloud for a conclusion with no room on the slide."""
        take = " ".join((self.o.takeaway or "").split())
        return f"Вывод: {take.rstrip('.')}." if take else ""

    def _unsay_conclusion(self) -> None:
        """A plan set before (a saved outline rendered again, a second pass over the same slide) may carry the line an
        earlier setting said aloud: it goes before this setting decides again — `_draw_notes` says it once when the
        conclusion still has no room, never twice, and never while the slide shows it."""
        said = self._said()
        try:
            if said and said in (self.o.notes or ""):
                self.o.notes = (self.o.notes or "").replace(said + " ", "").replace(said, "").strip()
        except Exception:  # noqa: BLE001
            pass

    def _draw_notes(self, n0: int) -> None:
        """The conclusion right under the content (a gap below its last block, never lower than its reserved place),
        the footnote at the foot of the area."""
        notes = getattr(self, "_notes", None)
        if not notes:
            return
        foot_box, plan, area = notes
        k = self.kit
        if plan is not None:
            spot = self._conclusion_spot(n0, last=True)
            said = self._said()
            if spot is None:
                # the content grew past its share and leaves no room for the conclusion: it goes to the speaker notes
                # rather than over the content (a slide with every line of its lists, the takeaway said aloud) —
                # once: a slide set again (a second pass, a re-render of a saved plan) finds it there already
                try:
                    if said not in (self.o.notes or ""):
                        self.o.notes = (f"{said} " + (self.o.notes or "")).strip()
                except Exception:  # noqa: BLE001
                    pass
                self.warnings.append("the conclusion moved to the speaker notes: no room under the content")
                plan = None
            else:
                if spot[0] is not plan:
                    self.warnings.append("the conclusion set slimmer to stay on the slide" if spot[3] == area.w else "the conclusion set in the free room beside the template's art")
                plan, x, y, w, foot_box = spot
        if plan is not None:
            if plan["strip"]:
                self.cv.card(Bbox(x=x, y=y, w=w, h=plan["h"]), k.card, name="Conclusion strip")
            bx = x + plan["pad_x"]
            if plan.get("fig"):
                # the key figure and the sentence beside it, both centred on the strip's middle line
                para_f, fig_w, fig_h, fig_gap = plan["fig"]
                inner = plan["h"] - 2 * plan["pad_y"]
                lsb, _ = self._figure_optics(para_f, fig_h)
                self.cv.text(Bbox(x=bx - lsb, y=y + plan["pad_y"] + (inner - fig_h) // 2, w=fig_w + lsb, h=fig_h), [para_f], anchor="ctr", name="Conclusion figure")
                tx = bx + fig_w + fig_gap
                self.cv.text(Bbox(x=tx, y=y + plan["pad_y"] + (inner - plan["text_h"]) // 2, w=plan["text_w"], h=plan["text_h"]), [plan["para"]], name="Conclusion")
            else:
                by = y + plan["pad_y"] + _emu(plan["size"] * 0.12)
                self.cv.rect(Bbox(x=bx, y=by, w=plan["bar_w"], h=max(plan["text_h"] - _emu(plan["size"] * 0.24), plan["bar_w"] * 3)), plan["colors"].accent, name="Conclusion bar")
                tx = bx + plan["bar_w"] + plan["bar_gap"]
                self.cv.text(Bbox(x=tx, y=y + plan["pad_y"], w=plan["text_w"], h=plan["text_h"]), [plan["para"]], name="Conclusion")
        if foot_box is not None:
            box, para = foot_box
            self.cv.text(box, para, name="Footnote")

    def _undo(self, n0: int) -> None:
        """Remove what was added to the slide after the first `n0` elements (a chart that failed half-way: its part
        relationship goes too)."""
        tree = self.cv.tree
        for el in list(tree)[n0:]:
            for ref in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/chart}chart"):
                rid = ref.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if rid:
                    try:
                        self.slide.part.drop_rel(rid)
                    except KeyError:
                        pass
            tree.remove(el)
            if el in self.cv.added:
                self.cv.added.remove(el)

    # ---- formula ---------------------------------------------------------------------------------------------------
    def _quiet(self) -> str:
        """A quiet colour for large glyphs (operators, rules): the template's grey that reads as large text on the
        ground (3:1) and stands apart from the running text; the muted text colour when it has none."""
        k = self.kit
        g, text = k.colors.ground, k.colors.text

        def sat(h: str) -> float:
            r, gg, b = (int(h[i : i + 2], 16) / 255 for i in (0, 2, 4))
            mx, mn = max(r, gg, b), min(r, gg, b)
            return 0.0 if mx == 0 else (mx - mn) / mx

        for c in [k.colors.muted] + [x for x in k.palette]:
            if c and c.upper() != text.upper() and sat(c) < 0.25 and contrast_ratio(c, g) >= 3.0 and contrast_ratio(c, text) >= 1.6:
                return c
        return k.colors.muted
    def formula(self, area: Bbox, text: str) -> None:
        """An equation set large: every figure in the accent at one display size (its unit smaller on the baseline),
        the operators muted and smaller, the result on an accent panel; the words of a term («покупок», «дней») under
        its figure. Too wide for one line at a display size, the result moves to a second line. The slide's other
        text follows under the equation."""
        k = self.kit
        terms, ops = parse_formula(text)
        if len(terms) < 2 or not any(fig for fig, _ in terms):
            self.statement(area, text)
            return
        words = max((i for i, (fig, _) in enumerate(terms) if not fig), default=None)
        if words is not None and words + 1 < len(terms) and ops[words] in ("=", "≈") and len(terms) - words - 1 >= 2:
            # «Прибыль = Выручка − Расходы = 900 000 − 780 000 = 120 000 ₽»: the definition in words is a line over
            # the equation of figures, which is set large
            line = " ".join(x for i, (_, w) in enumerate(terms[: words + 1]) for x in ([w] + ([ops[i]] if i < words else [])))
            quiet_ = self._quiet()
            size = k.h2
            runs: list[Run] = []
            for tok in re.split(r"(\s[=≈×−+÷]\s)", typeset(line)):
                op = tok.strip() in ("=", "≈", "×", "−", "+", "÷")
                runs.append(Run(tok, size, quiet_ if op else k.colors.heading, k.bold and not op, k.font))
            para = [Para(runs)]
            hh = self.h(para, area.w)
            words_line = (para, hh, hh + int(k.vgap * 1.2))
            dy = words_line[2]
            area = Bbox(x=area.x, y=area.y + dy, w=area.w, h=max(area.h - dy, int(area.h * 0.5)))
            terms, ops = terms[words + 1 :], ops[words + 1 :]
        else:
            words_line = None
        c = self.o.content
        # a term without words of its own takes the label of the slide's figure it repeats («100» → «покупок в день»)
        used: set[int] = set()
        labeled: list[tuple[Optional[str], str]] = []
        for fig, lab in terms:
            if fig and not lab:
                d = re.sub(r"\D", "", fig)
                hit = next((i for i, n in enumerate(c.numbers) if i not in used and re.sub(r"\D", "", n.value) == d and n.label.strip()), None)
                if hit is not None:
                    used.add(hit)
                    lab = c.numbers[hit].label.strip()
                    value = " ".join(c.numbers[hit].value.split())
                    if re.search(r"[^\d\s.,]", value) and not re.search(r"[^\d\s.,]", fig):
                        fig = value  # «300» of the formula is the slide's «300 ₽»: the figure keeps its unit
            labeled.append((fig, lab))
        terms = labeled
        res = next((i + 1 for i in range(len(ops) - 1, -1, -1) if ops[i] in ("=", "≈")), None)
        colors = k.colors
        lab_size = k.body
        on_accent = k.on_accent(colors.accent, colors.text)
        quiet = self._quiet()
        descent, digit_h = figure_metrics_em(k.figure_bold)

        def layout(fs: float) -> dict:
            op_size = k.snap(fs * 0.55, k.h3, fs * 0.7)
            pad = _emu(fs * 0.28)
            blocks = []
            for i, (fig, lab) in enumerate(terms):
                if fig:
                    color = on_accent if i == res else colors.figure
                    para, w = self.figure_para(fig, fs, color, on_accent if i == res else colors.muted)
                else:
                    para = self.P(lab, k.snap(fs * 0.45, k.lead, fs * 0.6), colors.text)
                    w = text_width_pt(lab, k.font, para.size, False)
                    lab = ""
                lw = text_width_pt(lab, k.font, lab_size, False) if lab else 0.0
                bw = _emu(max(w, min(lw, max(w * 1.8, 150.0))) * 1.04)
                if i == res:
                    bw += 2 * pad
                blocks.append(dict(para=para, w=bw, fig_w=_emu(w), lab=lab))
            op_w = [_emu(text_width_pt(op, k.font, op_size, False) + fs * 0.5) for op in ops]
            return dict(fs=fs, op_size=op_size, pad=pad, blocks=blocks, op_w=op_w, total=sum(b["w"] for b in blocks) + sum(op_w))

        cap = k.figure_cap(2, False, self.strategy)
        floor = max(k.h2 * 1.25, k.h3)
        chosen = None
        for fs in k.figure_sizes(cap, floor):
            lay = layout(fs)
            if lay["total"] <= area.w:
                chosen, rows = lay, [list(range(len(terms)))]
                break
        rest_texts = [t for t in list(c.bullets) + list(c.paragraphs) if t.strip()]
        if chosen is not None and res is not None and len(rows[0]) >= 3 and not rest_texts and self.strategy != "compact":
            # the equation is the slide, and one line leaves most of the area empty under it: the operands on one line
            # and «= result» under them set much larger (the compact variant keeps its one dense line)
            split = res
            room_h = int(area.h * 0.9)
            for fs2 in k.figure_sizes(cap, chosen["fs"] * 1.25):
                lay2 = layout(fs2)
                first = sum(lay2["blocks"][i]["w"] for i in range(split)) + sum(lay2["op_w"][: split - 1])
                second = sum(lay2["blocks"][i]["w"] for i in range(split, len(terms))) + sum(lay2["op_w"][split - 1:])
                fig_h2 = _emu(fs2 * max(1.15, k.line))
                lab_h2 = max((self.h([self.P(b["lab"], lab_size, colors.muted)], b["w"]) for b in lay2["blocks"] if b["lab"]), default=0)
                row_h2 = fig_h2 + lay2["pad"] + (int(k.vgap * 0.15) + lab_h2 if lab_h2 else 0)
                block2 = 2 * row_h2 + int(k.vgap * 1.4) + lay2["pad"]
                if max(first, second) <= area.w and block2 <= room_h:
                    chosen, rows = lay2, [list(range(split)), list(range(split, len(terms)))]
                    break
        if chosen is None:
            # two lines: the operands, then «= result»
            split = res if res is not None else math.ceil(len(terms) / 2)
            for fs in k.figure_sizes(cap, k.h3):
                lay = layout(fs)
                first = sum(lay["blocks"][i]["w"] for i in range(split)) + sum(lay["op_w"][: split - 1])
                second = sum(lay["blocks"][i]["w"] for i in range(split, len(terms))) + sum(lay["op_w"][split - 1 :])
                chosen, rows = lay, [list(range(split)), list(range(split, len(terms)))]
                if max(first, second) <= area.w:
                    break
            else:
                # too long for two lines at a readable size: the formula is a statement, its figures in the accent
                self.statement(area, text)
                return
        lay = chosen
        fs = lay["fs"]
        fig_h = _emu(fs * max(1.15, k.line))
        pad = lay["pad"]
        # the digits of a figure set bottom-anchored in its box: baseline and centre, for the operators and the panel
        base_off = _emu(descent * fs)
        dig_h = _emu(digit_h * fs)
        lab_h = max((self.h([self.P(b["lab"], lab_size, colors.muted)], b["w"]) for b in lay["blocks"] if b["lab"]), default=0)
        lab_gap = int(k.vgap * 0.15)
        row_h = fig_h + (pad if res is not None else 0) + (lab_gap + lab_h if lab_h else 0)
        block_h = len(rows) * row_h + (len(rows) - 1) * int(k.vgap * 1.4) + (pad if res is not None else 0)
        rest = [t for t in list(c.bullets) + list(c.paragraphs) if t.strip()]
        rest += [f"{n.value} — {n.label}" for i, n in enumerate(c.numbers) if i not in used and n.label.strip()]
        y = area.y + (pad if res is not None else 0)
        if not rest:
            # the equation is the slide: it stands in the upper middle of the free area (with the conclusion that
            # follows it), not pressed under the heading
            notes = getattr(self, "_notes", None)
            plan = notes[1] if notes else None
            group_h = block_h + ((plan["h"] + int(k.vgap * 1.4)) if plan else 0) + (words_line[2] if words_line else 0)
            free = (plan["y_max"] + plan["h"] - area.y) + (words_line[2] if words_line else 0) if plan else area.h + (words_line[2] if words_line else 0)
            y += max(0, min(int((free - group_h) * 0.42), area.h - block_h))
        if words_line is not None:
            # the definition in words right over its equation
            self.cv.text(Bbox(x=area.x, y=y - (pad if res is not None else 0) - words_line[2], w=area.w, h=words_line[1]), words_line[0], name="Formula words")
        for r, row in enumerate(rows):
            x = area.x
            base = y + fig_h - base_off  # the digits' baseline on this row
            for n, i in enumerate(row):
                if n or r:
                    oi = i - 1  # the operator before term i
                    ow = lay["op_w"][oi]
                    osz = lay["op_size"]
                    oh = _emu(osz * 1.3)
                    cy = base - dig_h // 2
                    self.cv.text(Bbox(x=x, y=cy - oh // 2, w=ow, h=oh), [self.P(ops[oi], osz, quiet, align="ctr")], anchor="ctr", name="Operator")
                    x += ow
                b = lay["blocks"][i]
                bx = x
                if i == res:
                    self.cv.rect(Bbox(x=x, y=base - dig_h - pad, w=b["w"], h=dig_h + 2 * pad), colors.accent, radius_emu=_emu(fs * 0.18), name="Result")
                    bx = x + pad
                lsb, _ = self._figure_optics(b["para"], fig_h)
                self.cv.text(Bbox(x=bx - lsb, y=y, w=b["fig_w"] + lsb + _emu(4), h=fig_h), [b["para"]], anchor="b", name="Figure" if i != res else "Result figure")
                if b["lab"]:
                    ly = y + fig_h + (pad if res is not None and terms[res][1] else 0) + lab_gap
                    self.cv.text(Bbox(x=bx, y=ly, w=b["w"] - (2 * pad if i == res else 0), h=lab_h), [self.P(b["lab"], lab_size, colors.muted)], name="Term label")
                x += b["w"]
            y += row_h + int(k.vgap * 1.4)
        # what explains the equation, under it
        if rest:
            below = Bbox(x=area.x, y=y + int(k.vgap * 0.6), w=area.w, h=max(area.y2 - y - int(k.vgap * 0.6), int(0.12 * k.H)))
            if len(rest) == 1 and not c.bullets:
                self.paragraphs(below, rest)
            else:
                self.bullets(below, rest)

    # ---- data ----------------------------------------------------------------------------------------------------
    def table(self, area: Bbox, table: TableData) -> None:
        """A native table measured for its content (tables.measure_table): one size, columns by content, rows that
        grow to fill about half the area; a note or the caption under it."""
        from verstka.rendering.tables import emphasis_column, measure_table, table_style_for_ground, template_bold

        k = self.kit
        style = self.manifest.components.table_style
        n_cols = max(len(table.columns), 1)
        n_rows = len(table.rows) + 1
        width = area.w if n_cols >= 3 else int(area.w * 0.72)
        # what changed most, beside the table (visual) — unless the slide states its own conclusion under it
        deltas = table_deltas(table) if self.strategy == "visual" and area.w * 0.36 >= 0.3 * k.W and not (self.o.takeaway or "").strip() else []
        if deltas:
            width = int(area.w * 0.64)  # the table and, beside it, what changed most
        floor = 0.022 * k.hpt
        # a table reads at body size; dense tables step down to the small size. A small one (a few rows, a few
        # columns) set at body size is a thin strip over an empty slide: it may take the lead size when nothing wraps
        header_bold = template_bold(k.manifest.tokens.typography)

        boost = getattr(self, "_boost", 0)

        def sizes_for(t: TableData) -> list[float]:
            small_table = len(t.rows) + 1 <= 5 and len(t.columns) <= 4
            lead = [k.lead] if (small_table or boost >= 1) and k.lead > k.body and (not k.head_size or k.lead <= 0.7 * k.head_size) else []
            mids = [x for x in k.sizes if k.body + 0.05 < x < k.lead - 0.05] if boost >= 1 else []
            return [s for s in sorted({*lead, *mids, k.body, k.small}, reverse=True) if s >= floor] or [k.small]

        tight = getattr(self, "_tight", 0)

        def measure(t: TableData, sizes: list[float]):
            # a cell is set single-spaced (the renderer's 1.2 em), whatever the template's leading for running text;
            # on the fill pass the rows grow until the table takes most of its area
            return measure_table(
                t, width, k.font, sizes, 1.2,
                max_h_emu=area.h, fill_h_emu=0 if tight else int((0.9 if boost >= 2 else 0.62) * area.h), min_row_h_emu=int((0.052 if tight else 0.065) * k.H), max_row_h_emu=int((0.13 if boost >= 2 else 0.10) * k.H),
                header_bold=header_bold,
            )

        # one table size for the deck: every table is set at the size the densest one allows (two tables of one deck
        # at 15 and 11 pt read as two designers)
        cache = self.slide.part.package.__dict__.setdefault("_verstka_table_size", {})  # one deck, one build
        key = (width, area.h, boost)
        if key not in cache:
            tables = [s.content.table for s in self.outline.slides if s.content.table is not None and s.content.table.columns]
            cache[key] = min((measure(t, sizes_for(t))[0] for t in tables), default=None)
        cap = cache[key]
        sizes = sizes_for(table)
        if cap:
            sizes = [s for s in sizes if s <= cap + 0.05] or [min(sizes)]
        if tight >= 2:
            # as wide as the columns need at the size the full width sets, with air — never under 60 % of the area;
            # the type keeps that size (a narrower table is not a smaller one) — one step smaller only as the last
            # level, when the narrower table still covers the slide past the wall
            probe = measure(table, sizes)
            sizes = [probe[0]]
            if tight >= 3:
                sizes = [x for x in sizes_for(table) if x < probe[0] - 0.05][:1] or sizes
                probe = measure(table, sizes)
            width = min(width, max(int(probe.content_w * 1.12), int(0.6 * area.w)))
        size, widths, heights = measure(table, sizes)
        spec = table_style_for_ground(style, ground_hex=k.colors.ground, text_hex=k.colors.text, accent_hex=k.colors.accent, divider_hex=k.colors.divider)
        total_h = sum(heights)
        add_table(
            self.slide, Bbox(x=area.x, y=area.y, w=sum(widths), h=total_h), table, spec, k.manifest.tokens.typography,
            font_family=k.font, col_widths=widths, row_heights=heights, size_pt=size,
            accent_hex=k.colors.accent, muted_hex=k.colors.muted, ground_hex=k.colors.ground, header_bold=header_bold,
            accent_text_hex=k.colors.accent_text, highlight_col=emphasis_column(table),
        )
        if table.caption and total_h + k.vgap + _emu(k.small * 2) <= area.h:
            self.cv.text(Bbox(x=area.x, y=area.y + total_h + int(k.vgap * 0.6), w=width, h=_emu(k.small * 1.6)), [self.P(table.caption, k.small, k.colors.muted)], name="Caption")
        if deltas:
            x = area.x + width + k.gap * 2
            w = area.x + area.w - x
            paras: list[Para] = []
            for value, label, span in deltas[:2]:
                fs = k.display
                for s_ in k.figure_sizes(k.figure_cap(2, False, self.strategy), k.h3):
                    fs = s_
                    if self.figure_para(value, s_, k.colors.figure, k.colors.muted)[1] <= _pt(w) * 0.8:
                        break
                paras.append(Para(self.figure_para(value, fs, k.colors.figure, k.colors.muted)[0].runs, space_after=2))
                paras.append(self.P(label, k.body, k.colors.text, space_after=1))
                paras.append(self.P(span, k.small, k.colors.muted, space_after=k.body * 1.3))
            hh = self.h(paras, w)
            self.cv.line(x - k.gap, area.y, x - k.gap, area.y + total_h, k.colors.divider, 1.0)  # as tall as the table
            self.cv.text(Bbox(x=x, y=area.y, w=w, h=min(hh, area.h)), paras, name="Takeaway")

    # ---- chart slides: one content, three presentations ------------------------------------------------------------
    #
    # The brief often dictates every slide's form (a pie here, two columns there): the three variants would then be
    # the same deck three times. Each variant presents a chart slide its own way, inside the template's design system:
    #   structured — the chart at the left (64%), the lines in a column at its right behind a hairline, the conclusion
    #                in a strip under both;
    #   visual     — a wider chart; the lines that carry a figure become callouts beside it (the figure large in the
    #                accent, its words under it), the other lines stay short lines; the conclusion of a lone chart stands
    #                beside it with its key figure large. Lines without a figure go under the chart as captions;
    #   compact    — mirrored and denser: a text column at the left led by the conclusion in bold, the lines under it,
    #                the chart at the right taking the full height (no strip).

    def _side_split(self, texts: list[str]) -> tuple[list[tuple[str, str, bool]], list[str]]:
        """The visual variant's reading of a chart's side lines: up to three lines that carry one figure become
        callouts (figure, label, statement=False); the others stay lines."""
        kpis: list[tuple[str, str, bool]] = []
        lines: list[str] = []
        for t in texts:
            got = kpi_callout(t) if len(kpis) < 3 else None
            if got is not None:
                kpis.append((got[0], got[1], False))
            else:
                lines.append(t)
        return kpis, lines

    def _kpi_column(self, box: Bbox, kpis: list[tuple[str, str, bool]], lines: list[str], dry: bool = False) -> Optional[int]:
        """The visual variant's column beside a chart: each callout a short accent rule, its figure large in the
        template's figure colour (one size for the column: the largest at which every figure stays on one line and the
        column fits the height) and its words under it; a statement (the slide's conclusion) in the text colour at the
        lead size; the lines without a figure under the callouts. Returns the height it takes (None: it does not fit
        even at the smallest sizes — the caller sets the lines as a plain column). `dry`: measure only."""
        k = self.kit
        colors = k.colors
        figs = [v for v, _, _ in kpis if v]
        w_pt = _pt(box.w)
        cap = min(k.figure_cap(max(len(figs), 2), False, "visual"), 0.15 * k.hpt)
        floor = max(k.h2, k.h3)
        fig_sizes = [s for s in k.figure_sizes(cap, floor) if all(self.figure_para(v, s, colors.figure, colors.muted)[1] <= w_pt * 0.94 for v in figs)] if figs else [0.0]
        if not fig_sizes:
            return None  # a figure too wide for the column at a callout size: no callouts
        rule_h = _emu(3)
        for rules in (True, False):
            for lab_size in dict.fromkeys((k.lead, k.body)):
                for line_size in dict.fromkeys((k.body, k.small)):
                    for fs in fig_sizes:
                        layout = self._kpi_layout(box, kpis, lines, fs, lab_size, line_size, rules, rule_h)
                        if layout["h"] <= box.h:
                            if not dry:
                                self._kpi_draw(box, layout)
                            return layout["h"]
        return None

    def _kpi_layout(self, box: Bbox, kpis, lines, fs: float, lab_size: float, line_size: float, rules: bool, rule_h: int) -> dict:
        k = self.kit
        colors = k.colors
        blocks = []
        y = 0
        between = int(k.vgap * 1.15)
        for i, (value, label, statement) in enumerate(kpis):
            b: dict = {"y": y}
            if rules:
                b["rule"] = y
                y += rule_h + int(k.vgap * 0.9)
            if value:
                para = self.figure_para(value, fs, colors.figure, colors.muted)[0]
                fig_h = _emu(fs * max(1.15, k.line))
                lsb, tg = self._figure_optics(para, fig_h)
                b["fig"] = (para, y - tg, fig_h, lsb)
                y += fig_h - tg + int(k.vgap * 0.3)
            size = k.lead if statement else lab_size
            if statement:
                lp = [Para(highlight_runs(label, size, colors.text, colors.accent if size >= 18 else colors.accent_text, bold=False, accent_bold=k.bold, font=k.font) if not value else [Run(typeset(label), size, colors.text, False, k.font)])]
            else:
                lp = [self.P(label, size, colors.muted)]
            lh = self.h(lp, box.w)
            b["label"] = (lp, y, lh)
            y += lh
            blocks.append(b)
            if i < len(kpis) - 1:
                y += between
        lines_block = None
        if lines:
            if blocks:
                y += between
            paras = [self.P(t, line_size, colors.text, space_after=line_size * 0.55, marker="•" if len(lines) > 1 else None, marker_color=colors.accent) for t in lines]
            lh = self.h(paras, box.w)
            lines_block = (paras, y, lh)
            y += lh
        return dict(h=y, blocks=blocks, lines=lines_block, rule_h=rule_h)

    def _kpi_draw(self, box: Bbox, layout: dict) -> None:
        k = self.kit
        colors = k.colors
        for b in layout["blocks"]:
            if "rule" in b:
                self.cv.rect(Bbox(x=box.x, y=box.y + b["rule"], w=min(_emu(k.hpt * 0.07), box.w), h=layout["rule_h"]), colors.accent, name="Rule")
            if "fig" in b:
                para, fy, fh, lsb = b["fig"]
                self.cv.text(Bbox(x=box.x - lsb, y=box.y + fy, w=box.w + lsb, h=fh), [para], anchor="b", name="Figure")
            lp, ly, lh = b["label"]
            self.cv.text(Bbox(x=box.x, y=box.y + ly, w=box.w, h=lh), lp, name="Label")
        if layout["lines"] is not None:
            paras, ly, lh = layout["lines"]
            self.cv.text(Bbox(x=box.x, y=box.y + ly, w=box.w, h=lh), paras, name="Chart note")

    def _text_column(self, box: Bbox, lead: Optional[str], figure: Optional[tuple[str, str]], texts: list[str]) -> int:
        """The compact variant's text column beside a chart: the conclusion as its bold lead line (in the heading
        colour, its figures in the accent), a computed figure of the series under it, then the lines — set dense (a
        short gap between them) at the largest size at which the column fits the height. A hairline at its right
        edge sets it apart from the chart. Returns the height it took."""
        k = self.kit
        colors = k.colors
        lead = (lead[:1].upper() + lead[1:]) if lead else None
        if lead and lead.endswith(".") and not lead.endswith("..") and not ABBR_END_RE.search(lead):
            lead = lead[:-1]
        best = None
        for lead_size, size in ((k.lead, k.lead), (k.lead, k.body), (k.h3, k.body), (k.body, k.body), (k.body, k.small)):
            paras: list[Para] = []
            if lead:
                paras.append(Para(highlight_runs(lead, lead_size, colors.heading, colors.accent if lead_size >= 18 else colors.accent_text, bold=True, accent_bold=True, font=k.font), space_after=lead_size * (0.9 if (texts or figure) else 0)))
            if figure is not None:
                value, label = figure
                fs = k.h2
                paras.append(Para(self.figure_para(value, fs, colors.figure, colors.muted)[0].runs, space_after=2))
                paras.append(self.P(label, size, colors.muted, space_after=size * (0.9 if texts else 0)))
            for j, t in enumerate(texts):
                paras.append(self.P(t, size, colors.text, space_after=size * 0.4 if j < len(texts) - 1 else 0, marker="•" if len(texts) > 1 else None, marker_color=colors.accent))
            hh = self.h(paras, box.w)
            best = (paras, hh)
            if hh <= box.h and self.fits_width(texts + ([lead] if lead else []), size, True, box.w - _emu(size * 1.1)):
                break
        paras, hh = best
        self.cv.text(Bbox(x=box.x, y=box.y, w=box.w, h=hh), paras, name="Chart note")
        return hh

    def _chart_split(self, area: Bbox) -> float:
        """The share of the width the visual variant's callout column takes beside a chart."""
        return 0.28 if area.w >= 0.6 * self.kit.W else 0.32

    def chart(self, area: Bbox) -> None:
        from verstka.rendering.charts import effective_chart_type

        k = self.kit
        c = self.o.content
        spec = c.chart
        side_texts = [t for t in (list(c.bullets) + list(c.paragraphs)) if t.strip()]
        kind = effective_chart_type(spec, self.outline)
        pie = kind in ("pie", "doughnut")
        # the figure a reader takes from a series (its growth) — never for the parts of a whole, and not when the
        # slide states its own conclusion under the chart
        takeaway = None if (pie or (self.o.takeaway or "").strip()) else _series_takeaway(spec, self.outline, self.o.headline)
        aside = self._takeaway_aside()
        take = " ".join((self.o.takeaway or "").split())
        take = take[:1].upper() + take[1:]
        n0 = len(self.cv.tree)
        tall_h = int(area.h * 0.92)
        gap = k.gap * 2
        if self.strategy == "compact" and (side_texts or aside or takeaway is not None):
            # mirrored: the text column at the left, the chart at the right over the full height
            col_w = int(area.w * (0.34 if sum(len(t) for t in side_texts + [take if aside else ""]) <= 220 else 0.38))
            hh = self._text_column(Bbox(x=area.x, y=area.y, w=col_w, h=area.h), take if aside else None, takeaway, side_texts)
            self.cv.line(area.x + col_w + gap // 2, area.y, area.x + col_w + gap // 2, area.y + min(hh, area.h), k.colors.divider, 1.0)
            err = self._chart_block(Bbox(x=area.x + col_w + gap, y=area.y, w=area.w - col_w - gap, h=tall_h), spec, prominent=False, center=False)
            if err is not None:
                self._chart_failed(n0, area, err)
            return
        if self.strategy == "visual":
            kpis, lines = self._side_split(side_texts)
            if aside:
                fig = _aside_figure(take) or ""
                kpis.insert(0, (fig, label_without_figure(take, fig), True))
            elif takeaway is not None:
                kpis.insert(0, (takeaway[0], takeaway[1], False))
            elif pie and not side_texts and kind == "pie":
                # a lone pie of the visual variant is a doughnut: the whole, when the deck states it, large in its hole
                spec = spec.model_copy(update={"type": "doughnut"})
                self._hole_large = True
            side_w = int(area.w * self._chart_split(area))
            side = Bbox(x=area.x + area.w - side_w, y=area.y, w=side_w, h=area.h)
            if kpis and self._kpi_column(side, kpis, lines, dry=True) is not None:
                err = self._chart_block(Bbox(x=area.x, y=area.y, w=area.w - side_w - gap, h=tall_h), spec, prominent=False, center=False)
                if err is not None:
                    self._chart_failed(n0, area, err)
                    return
                self._kpi_column(side, kpis, lines)
                return
            if lines and not kpis:
                # lines without a figure: captions under a chart as wide as the slide
                plan = self._under_plan(area, lines, spread=True)
                main = plan["main"]
                err = self._chart_block(Bbox(x=main.x, y=main.y, w=main.w, h=main.h), spec, prominent=False, center=True)
                if err is not None:
                    self._chart_failed(n0, area, err)
                    return
                self._under_draw(plan, n0)
                return
        if aside:
            takeaway = (_aside_figure(take) or "", take)  # a column of callouts did not fit: the conclusion leads the lines
        show_side = bool(side_texts) or (takeaway is not None and self.strategy != "structured")
        chart_w = int(area.w * (0.64 if show_side else 1.0))
        tall = show_side or kind == "bar" or pie
        chart_h = tall_h if tall else min(int(area.h * 0.82), int(area.w * 0.42))
        self._pie_mirror = self.strategy == "compact" and pie and not show_side
        err = self._chart_block(Bbox(x=area.x, y=area.y, w=chart_w, h=chart_h), spec, prominent=False, center=not show_side)
        self._pie_mirror = False
        self._hole_large = False
        if err is not None:
            self._chart_failed(n0, area, err)
            return
        if not show_side:
            return
        x = area.x + chart_w + gap
        w = area.x + area.w - x
        paras: list[Para] = []
        if takeaway is not None:
            value, label = takeaway
            if value:
                fs = k.display
                for s in k.steps_down(k.display, k.h3):
                    fs = s
                    if text_width_pt(value, k.font, s, k.bold) <= _pt(w) * 0.95:
                        break
                paras.append(self.P(value, fs, k.colors.accent, bold=k.bold, space_after=4))
            # the slide's own conclusion reads as a statement (lead, text colour); a computed figure's label is muted
            paras.append(self.P(label, k.lead if aside else k.body, k.colors.text if aside else k.colors.muted, space_after=k.body * 1.4))
        for t in side_texts:
            paras.append(self.P(t, k.body if len(t) > 90 else k.lead, k.colors.text, space_after=k.body * 0.7, marker="•" if len(side_texts) > 1 else None, marker_color=k.colors.accent))
        hh = self.h(paras, w)
        y = area.y + max(0, int((area.h - hh) * 0.35))
        # a hairline as tall as what it sets apart
        self.cv.line(x - k.gap, y, x - k.gap, y + min(hh, area.h), k.colors.divider, 1.0)
        self.cv.text(Bbox(x=x, y=y, w=w, h=min(hh, area.h)), paras, name="Takeaway")

    def _chart_failed(self, n0: int, area: Bbox, err: str) -> None:
        """No chart, no caption over an empty frame: a half-built chart goes (and the column set beside it), the slide
        shows what else it has — the conclusion the column was to hold goes back to its strip."""
        self._undo(n0)
        if self._takeaway_aside():
            area = self._takeaway_to_strip(area)
        self.warnings.append(f"chart failed: {err[:120]}; shown as {self._chart_fallback(area)}")

    def _pair_share(self) -> float:
        """The first chart's share of a pair's width: a trend over time (a line, an area, bars of five periods or more)
        is read along its axis and takes 60% — its labels stay on one line; the other one (a structure, a before/after
        pair) 40%: a pie sets its legend's names on two lines rather than its trend partner breaking its labels."""
        from verstka.rendering.charts import effective_chart_type, resolve_series

        def trend(spec) -> bool:
            kind = effective_chart_type(spec, self.outline)
            if kind in ("pie", "doughnut"):
                return False
            if kind in ("line", "area"):
                return True
            try:
                series = resolve_series(spec, self.outline)
            except Exception:  # noqa: BLE001
                return False
            return bool(series) and len(series[0].categories) >= 5

        c = self.o.content
        a, b = trend(c.chart), trend(c.chart2)
        return 0.6 if a and not b else (0.4 if b and not a else 0.5)

    def _pair_has_pie(self) -> bool:
        from verstka.rendering.charts import effective_chart_type

        c = self.o.content
        return any(spec is not None and effective_chart_type(spec, self.outline) in ("pie", "doughnut") for spec in (c.chart, c.chart2))

    def _text_band(self, area: Bbox, lead: Optional[str], texts: list[str]) -> int:
        """The compact variant's text over a pair of charts that needs the whole width: the conclusion as a bold lead
        line across the area, the lines under it side by side (up to three columns). Returns the band's height."""
        k = self.kit
        colors = k.colors
        y = area.y
        if lead:
            lead = lead[:1].upper() + lead[1:]
            if lead.endswith(".") and not lead.endswith("..") and not ABBR_END_RE.search(lead):
                lead = lead[:-1]
            for size in dict.fromkeys((k.lead, k.body)):
                para = [Para(highlight_runs(lead, size, colors.heading, colors.accent if size >= 18 else colors.accent_text, bold=True, accent_bold=True, font=k.font))]
                hh = self.h(para, area.w)
                if para_lines(para[0], _pt(area.w)) <= 2:
                    break
            self.cv.text(Bbox(x=area.x, y=y, w=area.w, h=hh), para, name="Chart note")
            y += hh + (int(k.vgap * 0.5) if texts else 0)
        if texts:
            n = len(texts)
            cols = n if n <= 3 else 2
            per = math.ceil(n / cols)
            gap = k.gap * 2
            cw = int((area.w - gap * (cols - 1)) / cols)
            size = k.body
            groups = [texts[i * per:(i + 1) * per] for i in range(cols)]
            paras = [[self.P(t, size, colors.text, space_after=size * 0.4, marker="•", marker_color=colors.accent) for t in g] for g in groups if g]
            eh = max(self.h(p, cw) for p in paras)
            for i, p in enumerate(paras):
                self.cv.text(Bbox(x=area.x + i * (cw + gap), y=y, w=cw, h=eh), p, name="Chart note")
            y += eh
        return y - area.y

    def _kpi_row_plan(self, area: Bbox, kpis: list[tuple[str, str, bool]], lines: list[str]) -> Optional[dict]:
        """The visual variant's callouts under a pair of charts that needs the whole width: each callout a column (an
        accent rule, the figure large, its words under it), the lines without a figure a last column. One figure size
        for the row: the largest at which every figure stays on one line in its column and the row keeps under 40% of
        the area. None when even the smallest size does not fit."""
        k = self.kit
        colors = k.colors
        cols = len(kpis) + (1 if lines else 0)
        gap = k.gap * 2
        cw = int((area.w - gap * (cols - 1)) / cols)
        budget = int(area.h * 0.36)
        rule_h = _emu(3)
        # the charts stay the picture: the row's figures stop at the display size (the column beside a lone chart
        # may go larger)
        for fs in k.figure_sizes(min(k.figure_cap(max(len(kpis), 2), False, "visual"), k.display), k.h3):
            if not all(self.figure_para(v, fs, colors.figure, colors.muted)[1] <= _pt(cw) * 0.94 for v, _, _ in kpis if v):
                continue
            blocks = []
            band = 0
            fig_h = _emu(fs * max(1.15, k.line))
            for value, label, statement in kpis:
                lp = [self.P(label, k.lead if statement and not value else k.body, colors.text if statement else colors.muted)]
                lh = self.h(lp, cw)
                y_fig = rule_h + int(k.vgap * 0.9)
                if value:
                    para = self.figure_para(value, fs, colors.figure, colors.muted)[0]
                    lsb, tg = self._figure_optics(para, fig_h)
                    y_lab = y_fig + fig_h - tg + int(k.vgap * 0.3)
                    blocks.append(dict(para=para, fig=(y_fig - tg, fig_h, lsb), label=(lp, y_lab, lh)))
                else:
                    y_lab = y_fig
                    blocks.append(dict(para=None, fig=None, label=(lp, y_lab, lh)))
                band = max(band, y_lab + lh)
            line_paras = [self.P(t, k.body, colors.text, space_after=k.body * 0.5, marker="•" if len(lines) > 1 else None, marker_color=colors.accent) for t in lines]
            # the lines without a figure stand level with the callouts' words, under the rule of their own column
            lines_y = blocks[0]["label"][1] if blocks else rule_h + int(k.vgap * 0.9)
            if line_paras:
                band = max(band, lines_y + self.h(line_paras, cw))
            if band <= budget:
                room = band + int(k.vgap * 1.4)
                main = Bbox(x=area.x, y=area.y, w=area.w, h=max(area.h - room, int(area.h * 0.5)))
                return dict(main=main, blocks=blocks, lines=line_paras, lines_y=lines_y, cw=cw, gap=gap, band=band, rule_h=rule_h, x=area.x, y2=area.y2)
        return None

    def _kpi_row_draw(self, row: dict, n0: int) -> None:
        k = self.kit
        colors = k.colors
        bottom = self._content_bottom(n0)
        y = row["main"].y2 + int(k.vgap * 1.4) if bottom is None else bottom + int(k.vgap * 1.4)
        limit = min(self._floor(row.get("y2") or k.H), k.H)
        if y + row["band"] > limit:
            # the charts ran past their share: the row keeps inside the area (a hair closer to the charts)
            y = max(limit - row["band"], row["main"].y2 + int(k.vgap * 0.6))
        cw, gap = row["cw"], row["gap"]
        for i, b in enumerate(row["blocks"]):
            x = row["x"] + i * (cw + gap)
            self.cv.rect(Bbox(x=x, y=y, w=min(_emu(k.hpt * 0.07), cw), h=row["rule_h"]), colors.accent, name="Rule")
            if b["fig"] is not None:
                fy, fh, lsb = b["fig"]
                self.cv.text(Bbox(x=x - lsb, y=y + fy, w=cw + lsb, h=fh), [b["para"]], anchor="b", name="Figure")
            lp, ly, lh = b["label"]
            self.cv.text(Bbox(x=x, y=y + ly, w=cw, h=lh), lp, name="Label")
        if row["lines"]:
            x = row["x"] + len(row["blocks"]) * (cw + gap)
            ly = y + row["lines_y"]
            self.cv.text(Bbox(x=x, y=ly, w=cw, h=self.h(row["lines"], cw)), row["lines"], name="Chart note")

    def chart_pair(self, area: Bbox) -> None:
        """Two charts side by side («до и после», a structure next to a trend), each under its own title in the heading
        colour — a trend takes the larger share. The slide's lines as the variant sets them: a column at the right
        (structured), callouts at the right or captions under the charts (visual), a text column at the left led by
        the conclusion (compact)."""
        k = self.kit
        c = self.o.content
        side_texts = [t for t in (list(c.bullets) + list(c.paragraphs)) if t.strip()]
        aside = self._takeaway_aside()
        take = " ".join((self.o.takeaway or "").split())
        gap = int(k.gap * 2.5)
        n0 = len(self.cv.tree)
        region = area
        after = None  # what is drawn once the charts stand
        # a pie carries its legend beside its circle and needs half the width: next to it the lines leave the side
        # column for a band over (compact) or under (structured, visual) the charts
        pie = self._pair_has_pie()
        if self.strategy == "compact" and (side_texts or aside):
            if pie:
                hh = self._text_band(area, take if aside else None, side_texts)
                dy = hh + int(k.vgap * 1.3)
                region = Bbox(x=area.x, y=area.y + dy, w=area.w, h=max(area.h - dy, int(area.h * 0.5)))
            else:
                col_w = int(area.w * 0.3)
                hh = self._text_column(Bbox(x=area.x, y=area.y, w=col_w, h=area.h), take if aside else None, None, side_texts)
                self.cv.line(area.x + col_w + gap // 2, area.y, area.x + col_w + gap // 2, area.y + min(hh, area.h), k.colors.divider, 1.0)
                region = Bbox(x=area.x + col_w + gap, y=area.y, w=area.w - col_w - gap, h=area.h)
        elif self.strategy == "visual" and (side_texts or aside):
            kpis, lines = self._side_split(side_texts)
            side_w = int(area.w * 0.25)
            side = Bbox(x=area.x + area.w - side_w, y=area.y, w=side_w, h=area.h)
            if pie:
                # the callouts (the conclusion first) in a row under the charts, in the strip's place
                if aside:
                    fig = _aside_figure(take) or ""
                    kpis.insert(0, (fig, label_without_figure(take[:1].upper() + take[1:], fig), True))
                row = self._kpi_row_plan(area, kpis, lines) if kpis else None
                if row is not None and row["main"].h >= 0.6 * area.h:
                    region = row["main"]
                    after = lambda: self._kpi_row_draw(row, n0)  # noqa: E731
                else:
                    if aside:
                        area = self._takeaway_to_strip(area)  # no room for the row: the conclusion takes its strip
                    if side_texts:
                        plan = self._under_plan(area, side_texts, spread=True)
                        region = plan["main"]
                        after = lambda: self._under_draw(plan, n0)  # noqa: E731
                    else:
                        region = area
            elif kpis and self._kpi_column(side, kpis, lines, dry=True) is not None:
                region = Bbox(x=area.x, y=area.y, w=area.w - side_w - gap, h=area.h)
                after = lambda: self._kpi_column(side, kpis, lines)  # noqa: E731
            else:
                plan = self._under_plan(area, side_texts, spread=True)
                region = plan["main"]
                after = lambda: self._under_draw(plan, n0)  # noqa: E731
        elif side_texts and pie:
            plan = self._under_plan(area, side_texts)
            region = plan["main"]
            after = lambda: self._under_draw(plan, n0)  # noqa: E731
        elif side_texts:
            text_w = int(area.w * 0.27)
            region = Bbox(x=area.x, y=area.y, w=area.w - text_w - gap, h=area.h)

            def after() -> None:
                x = area.x + area.w - text_w
                paras = [self.P(t, k.body if len(t) > 90 else k.lead, k.colors.text, space_after=k.body * 0.7, marker="•" if len(side_texts) > 1 else None, marker_color=k.colors.accent) for t in side_texts]
                hh = min(self.h(paras, text_w), area.h)
                self.cv.line(x - gap // 2, area.y, x - gap // 2, area.y + hh, k.colors.divider, 1.0)
                self.cv.text(Bbox(x=x, y=area.y, w=text_w, h=hh), paras, name="Chart note")
        from verstka.rendering.charts import effective_chart_type

        share = self._pair_share()
        # side by side first (two bar charts of a few bars read well even in a narrow region); a pie whose legend does
        # not fit its share takes half the width, then — in a region narrower than 1.25 slide heights (a 4:3 slide, a
        # column beside text) and tall enough — the two charts are set one over the other
        plans = [(share, False)]
        pies = [i for i, sp in enumerate((c.chart, c.chart2)) if sp is not None and effective_chart_type(sp, self.outline) in ("pie", "doughnut")]
        if len(pies) == 1 and (share if pies[0] == 0 else 1 - share) < 0.5:
            plans.append((0.5, False))
            # B5: a low region (a row of callouts under the charts, triangles over it) leaves a half-width pie a mark
            # (LO Focus s5: 0.088 H) — the pie takes 55 % of the width, its legend beside the circle; the trend reads
            # in the rest while its plot keeps its own width (`_trend_squeezed`: the value axis never cut to «1 200 0…»)
            plans.append((0.55 if pies[0] == 0 else 0.45, False))
        if pies and region.w < 1.25 * k.H and region.h >= 0.45 * k.H:
            plans.append((0.5, True))
        n1 = len(self.cv.tree)
        w_mark = len(self.warnings)

        def run(share: float, stacked: bool) -> Optional[bool]:
            """Draw the pair by one plan; None when a chart failed (the other one took the slide), else whether every
            pie of the pair set its legend cleanly (inside its box)."""
            del self.warnings[w_mark:]
            if stacked:
                vg = int(k.vgap * 1.2)
                h1 = int((region.h - vg) / 2)
                boxes = [Bbox(x=region.x, y=region.y, w=region.w, h=h1), Bbox(x=region.x, y=region.y + h1 + vg, w=region.w, h=region.h - h1 - vg)]
            else:
                w1 = int((region.w - gap) * share)
                boxes = [Bbox(x=region.x, y=region.y, w=w1, h=region.h), Bbox(x=region.x + w1 + gap, y=region.y, w=region.w - gap - w1, h=region.h)]
            self._legend_used: list[float] = []
            self._pie_ds: list[int] = []
            self._pie_unclean = False
            self._pie_small_any = False
            self._pie_shares_any = False
            for attempt in range(2):
                self._turned = False
                self._trend_ph = None
                self._trend_squeezed = False
                for i, spec in enumerate((c.chart, c.chart2)):
                    self._pie_clean = True
                    self._pie_small = False
                    err = self._chart_block(boxes[i], spec, prominent=True, center=stacked)
                    if err is not None:
                        # one chart of the pair could not be drawn: the other one takes the slide
                        self._undo(n0)
                        self._legend_force = None
                        self.warnings.append(f"chart {i + 1} of 2 failed: {err[:120]}; the other one is shown alone")
                        keep = c.chart2 if i == 0 else c.chart
                        self.o = self.o.model_copy(update={"content": c.model_copy(update={"chart": keep, "chart2": None})})
                        self.chart(area)
                        return None
                    if not getattr(self, "_pie_clean", True):
                        self._pie_unclean = True
                    if getattr(self, "_pie_small", False):
                        self._pie_small_any = True
                    if getattr(self, "_pie_shares_only", False):
                        self._pie_shares_any = True
                if attempt or len(self._legend_used) < 2 or (len(set(self._legend_used)) <= 1 and max(self._pie_ds) - min(self._pie_ds) <= _emu(2)):
                    break
                # two pies of one slide: one legend size and one diameter for both (the smaller ones), set again
                self._undo(n1)
                self._legend_force = min(self._legend_used)
                self._pie_force_d = min(self._pie_ds)
                self._legend_used, self._pie_ds = [], []
            self._legend_force = None
            self._pie_force_d = None
            # a plan that turns the trend chart's labels (45°, or every other one blank) where the first plan set them
            # straight is not taken for a larger pie (the dataset's LCT: straight months at full width)
            self._pair_squeezed = bool(sq0) and self._turned and not sq0[0]
            return not self._pie_unclean and not self._pie_small_any and not self._pair_squeezed

        # (the partner chart's labels as the first plan set them, legend inside its box, smallest pie diameter, plan
        # index) of the plans not taken at once: a legend past its box, a circle under the pie's least size, or — past
        # the first plan — a trend chart whose narrower slot turns its category labels
        # B3-3: another plan than the first is worth it only when its pie is ≥ 15 % larger, or reaches the pie's least
        # size, or its legend fits (or keeps the brief's amounts) where the first one's did not — and, one chart over
        # the other, when the trend chart keeps a plot ≥ 0.28 of the slide height (a stacked plan flattened the growth
        # line to 0.20 H for a pie 4 % smaller); side by side the trend keeps its height
        tried: list[tuple[bool, bool, bool, int, int]] = []
        chosen = None
        sq0: list[bool] = []
        d0, unclean0, shares0, small0 = 0, False, False, False
        for pi, (share, stacked) in enumerate(plans):
            ok = run(share, stacked)
            if ok is None:
                return
            d_now = min(self._pie_ds) if self._pie_ds else 0
            ph = getattr(self, "_trend_ph", None)
            gain = d_now >= 1.15 * d0 or (not self._pie_unclean and unclean0) or (not self._pie_shares_any and shares0) or (not self._pie_small_any and small0)
            worth = pi == 0 or (gain and not self._trend_squeezed and (not stacked or ph is None or ph >= self.PAIR_TREND_MIN_H * k.H))
            if pi == 0:
                sq0.append(self._turned)
                d0, unclean0, shares0, small0 = d_now, self._pie_unclean, self._pie_shares_any, self._pie_small_any
            if ok and worth:
                chosen = pi
                break
            tried.append((worth, not self._pair_squeezed, not self._pie_unclean, d_now, pi))
            if pi < len(plans) - 1:
                self._undo(n1)
        if chosen is None:
            # no plan sets every legend inside its box with a circle of its least size: a plan worth taking (B3-3), the
            # partner's labels whole, a clean legend, then the plan with the largest circle (the last one drawn when it
            # is that plan)
            best = max(tried)[4] if tried else len(plans) - 1
            if best != len(plans) - 1:
                self._undo(n1)
                if run(*plans[best]) is None:
                    return
            chosen = best
        if chosen:
            self.warnings.append("два графика одним над другим" if plans[chosen][1] else "круговой диаграмме — половина ширины")
        if after is not None:
            after()

    PAIR_TREND_MIN_H = 0.28  # a pair's other plan keeps the trend chart's plot at least this share of the slide height

    def _chart_caption(self, spec, named: bool = False) -> Optional[str]:
        """What stands over a chart, once: its title with the word unit («Прогноз выручки, тыс. ₽»), else the word
        unit alone (a one-glyph unit rides on the labels, a pie shows shares). A chart without a title of its own is
        named by its one series («Средний чек»): two charts side by side must say which is which."""
        from verstka.rendering.charts import resolve_series, unit_caption

        unit = unit_caption(spec, self.outline)
        title = " ".join((spec.title or "").split())
        if not title:
            series = resolve_series(spec, self.outline)
            name = " ".join((series[0].name or "").split()) if len(series) == 1 else ""
            generic = re.fullmatch(r"(?i)(series|ряд|серия|значени[яе]|values?)\s*\d*", name or "")
            if name and not generic and (named or not _same_words(name, self.o.headline)):
                title = name[:1].upper() + name[1:]
        if title and unit and unit.lower() not in title.lower():
            return f"{title.rstrip(' ,.:;')}, {unit}"
        return title or unit

    def _chart_block(self, box: Bbox, spec, *, prominent: bool, center: bool) -> Optional[str]:
        """A chart under its caption inside `box` (a pie with a legend of its own beside it). `prominent`: the caption
        is the chart's title in the heading colour (a pair of charts), else a muted caption line. Returns an error
        message when the chart could not be drawn (the caller removes what was added)."""
        from verstka.rendering.charts import effective_chart_type

        k = self.kit
        y = box.y
        cap = self._chart_caption(spec, named=prominent)
        pie = effective_chart_type(spec, self.outline) in ("pie", "doughnut")
        if cap and pie and not prominent:
            cap, legend_head = None, cap  # a lone pie's caption heads its legend
        else:
            legend_head = None
        if cap:
            size = k.lead if prominent else k.small
            para = [self.P(cap, size, k.colors.heading if prominent else k.colors.muted, bold=k.bold and prominent)]
            ch = self.h(para, box.w)
            # drawn before the chart: the chart finds its word unit written above it and prints plain numbers
            self.cv.text(Bbox(x=box.x, y=y, w=box.w, h=ch), para, name="Chart title" if prominent else "Unit")
            y += ch + int(k.vgap * (0.6 if prominent else 0.3))
        cbox = Bbox(x=box.x, y=y, w=box.w, h=max(box.y2 - y, int(0.2 * k.H)))
        plain = spec.model_copy(update={"title": None, "highlight_index": _after_index(spec, self.outline)})
        try:
            if pie:
                self._pie(cbox, plain, center=center, head=legend_head)
            else:
                gf = add_chart(self.slide, cbox, plain, self.outline, self._chart_style(), k.manifest.tokens.typography, text_hex=k.colors.text, neutral_hex=k.colors.divider, ground_hex=k.colors.ground)
                self._turned = bool(getattr(self, "_turned", False) or getattr(gf, "verstka_turned", False))
                from verstka.rendering.charts import chart_plot_bbox

                pb = chart_plot_bbox(gf)
                if pb is not None:
                    # the trend's plot height: a pair's other plan must not flatten it (B3-3) — nor press the plot to
                    # the chart's least width (its value axis and end label past the frame, cut by the renderer)
                    self._trend_ph = min(getattr(self, "_trend_ph", None) or pb.h, pb.h)
                if getattr(gf, "verstka_cramped", False):
                    self._trend_squeezed = True
        except Exception as e:  # noqa: BLE001
            return str(e) or type(e).__name__
        return None

    def _chart_style(self):
        from verstka.rendering.charts import chart_text_capped

        k = self.kit
        style = self.manifest.components.chart_style
        size = max(style.font_size_pt or 0, k.small)
        if chart_text_capped(k.manifest.tokens.typography, size, k.H):
            size = min(size, math.floor(0.026 * k.hpt * 2) / 2)  # chart text relative to the slide (≤ 2.6 % of its height)
        return style.model_copy(update={"font_size_pt": size, "font_family": k.font or style.font_family})

    def _pie(self, box: Bbox, spec, center: bool, head: Optional[str] = None) -> None:
        """A pie or a doughnut with its legend set as text beside it (below it in a narrow box): a swatch, the
        category and its share — the one place every share is read, the thin slices included. Shares are of the
        total, whatever the unit of the values (money, people)."""
        from verstka.rendering.charts import is_other_category, pie_values_are_percents, resolve_series
        from verstka.schemas.outline import InlineSeries

        k = self.kit
        series = resolve_series(spec, self.outline)
        if not series:
            raise ValueError("chart has no series data")
        s = series[0]
        cats0 = [str(x) for x in s.categories]
        vals0 = [max(0.0, float(v or 0.0)) for v in list(s.values)[: len(cats0)]]
        # the slices from the largest down, a remainder («Прочие», «Резерв») last: the legend's tints step down with
        # the sizes (the brief's order put «Резерв 17%» next to «Программа лояльности 19%» in nearly one tint)
        rest = lambda c: is_other_category(c) or bool(re.match(r"(?i)резерв|остал|прочи|друг", c))  # noqa: E731
        order = sorted(range(len(cats0)), key=lambda j: (rest(cats0[j]), -vals0[j]))
        cats = [cats0[j] for j in order]
        vals = [vals0[j] for j in order]
        if order != list(range(len(cats0))):
            spec = spec.model_copy(update={"series_ids": [], "categories": cats, "series": [InlineSeries(name=s.name, values=vals)], "unit": spec.unit or s.unit})
        total = sum(vals)
        if total <= 0:
            raise ValueError("pie without positive values")
        unit_txt = (spec.unit or s.unit or "").strip()
        # G4-11: a pie of the brief's own percents (unit «%», or no unit and the parts add up to 100 ± 1) shows one
        # value a row, as written — never «42 · 42%»
        is_pct = pie_values_are_percents(unit_txt, total)
        shares = [v if is_pct else v / total * 100 for v in vals]
        if is_pct:
            pct = [f"{round(sh, 1 if sh >= 1 else 2):g}".replace(".", ",") + "%" for sh in shares]
        else:
            pct = [(f"{sh:.1f}".replace(".", ",") if 0 < sh < 1 else f"{sh:.0f}") + "%" for sh in shares]
        # a pie of amounts: each part's amount (the brief's figure) and its share of the parts' total, the legend
        # headed as such — the share is not a figure of the brief («Продукты 40%» of the costs next to «35% от
        # выручки» read as a contradiction when unlabelled)
        amounts = not is_pct
        shares_only, head_plain = list(pct), head
        if amounts:
            amt = [typeset(_amount_text(v, unit_txt)) for v in vals]
            pct = [f"{a}  ·  {p_}" for a, p_ in zip(amt, pct)]
            share_head = "сумма · доля от общей суммы"
            head = f"{head}\n{share_head}" if head else share_head[:1].upper() + share_head[1:]
        rows = [j for j in range(len(cats)) if vals[j] > 0]
        colors = k.colors
        gap = k.gap
        big = max(rows, key=lambda j: vals[j])

        stack = False  # the amount · share line under each name (a narrow legend keeps the brief's amounts)

        def legend(size: float, name_cap: int, cols: int = 1) -> dict:
            sw = _emu(size * 0.72)
            pct_nat = _emu(max(text_width_pt(pct[j], k.font, size, k.figure_bold) for j in rows) + size * 0.5)
            pct_w = 0 if stack else pct_nat
            name_nat = _emu(max(text_width_pt(cats[j], k.font, size, False) for j in rows) + size * 0.6)
            if stack:
                name_nat = max(name_nat, pct_nat)
            name_w = max(min(name_nat, name_cap), _emu(size * 5), pct_nat if stack else 0)
            row_gap = _emu(size * 0.5)
            col_gap = _emu(size * 1.6)
            hs = [self.h([self.P(cats[j], size, colors.text)] + ([self.P(pct[j], size, colors.text, bold=k.figure_bold)] if stack else []), name_w) for j in rows]
            per = math.ceil(len(rows) / cols)
            chunks = [list(range(c_ * per, min((c_ + 1) * per, len(rows)))) for c_ in range(cols)]
            col_w = sw + int(sw * 0.8) + name_w + pct_w
            w = cols * col_w + (cols - 1) * col_gap
            head_paras = [self.P(line_, size, colors.muted) for line_ in head.split("\n")] if head else []
            head_h = self.h(head_paras, w) + _emu(size * 0.7) if head else 0
            body_h = max(sum(hs[i] for i in ch) + row_gap * (len(ch) - 1) for ch in chunks if ch)
            return dict(size=size, sw=sw, pct_w=pct_w, name_w=name_w, row_gap=row_gap, hs=hs, w=w, col_w=col_w, col_gap=col_gap,
                        chunks=chunks, h=head_h + body_h, head=head_paras, head_h=head_h, stack=stack)

        got: dict[str, tuple] = {}
        for legend_mode in (("amounts", "stacked", "shares") if amounts else ("shares",)):
            # the amounts beside their names; a narrower legend with each amount on the line under its name; the
            # shares only
            stack = legend_mode == "stacked"
            if legend_mode == "shares" and amounts:
                # a narrow box: the amounts stay on the slices, the legend gives the shares only, headed as such
                pct = shares_only
                head = f"{head_plain}\nдоля от общей суммы" if head_plain else "Доля от общей суммы"
            lg, d, below, sizes = self._pie_legend_fit(box, legend, rows, cats)
            if lg is not None:
                got[legend_mode] = (lg, d, below, list(pct), head)
            if legend_mode == "amounts" and lg is not None and ((d >= 0.4 * box.h and not below) or (below and d >= 0.35 * box.h)) and d >= int(self.PIE_SMALL * self._pie_min_d(box)):
                break  # the amounts beside a circle of 0.4 of the box's height (under it: 0.35) — the settled look
        pick = got.get("amounts") if len(got) == 1 and "amounts" in got else None
        if pick is None and got:
            # the brief's amounts never cost the circle its size (B1: the stacked legend shrank the dataset's pie by
            # half): an amounts form is kept when its circle is as large as the shares' one, or still of the pie's
            # least size (`_pie_min_d`); else the shares only
            amt = [got[m] for m in ("amounts", "stacked") if m in got]
            best_amt = max(amt, key=lambda v: v[1]) if amt else None
            shr = got.get("shares")
            if best_amt is not None and (shr is None or best_amt[1] >= min(int(shr[1] * 0.98), self._pie_min_d(box))):
                pick = best_amt
            else:
                pick = shr if shr is not None else best_amt
        self._pie_clean = pick is not None
        # a pie of amounts whose legend could only give the shares (the brief's amounts lost): a pair's other plan that
        # keeps them is worth taking (B3-3)
        self._pie_shares_only = bool(amounts and pick is not None and pick is got.get("shares"))
        if pick is not None:
            lg, d, below, pct, head = pick
        else:
            # nothing fits beside or under the circle in either mode: one column of shares under the circle, as wide as
            # the box, the names wrapped, a size down to the dense floor — the circle keeps 45 % of the height
            lg, d, below = self._pie_legend_last(box, legend, rows)
            self._pie_clean = bool(getattr(self, "_pie_last_ok", False))  # a dense legend inside its box is clean
        force_d = getattr(self, "_pie_force_d", None)
        if force_d and force_d < d:
            d = force_d  # two pies of one slide: one diameter
        # a circle well under the pie's least size (B5): a pair of charts tries its other plans (half the width, one
        # chart over the other) and keeps the one with the larger circle
        self._pie_small = d < int(self.PIE_SMALL * self._pie_min_d(box))
        if isinstance(getattr(self, "_legend_used", None), list):
            self._legend_used.append(lg["size"])
            self.__dict__.setdefault("_pie_ds", []).append(d)
        self._pie_draw(box, spec, center, lg, d, below, rows, cats, vals, pct, big, colors, amounts, total, unit_txt)

    PIE_MIN_H = 0.22  # the least circle of a pie, of the slide's height (B5): smaller, it reads as a mark, not a chart
    PIE_SMALL = 0.85  # a circle under this share of the least size sends a pair of charts to its other plans

    def _pie_min_d(self, box: Bbox) -> int:
        """The least diameter a pie keeps in `box`: 0.22 of the slide's height or 0.4 of the box's, whichever is
        larger — never more than the box holds (0.9 of its height, 0.6 of its width: the legend keeps its room)."""
        k = self.kit
        want = max(int(self.PIE_MIN_H * k.H), int(0.4 * box.h))
        return min(want, int(0.9 * box.h), int(0.6 * box.w))

    def _pie_legend_last(self, box: Bbox, legend, rows: list[int]):
        """The last resort of a pie's legend: one column under the circle, as wide as the box (names wrapped), the
        largest size down to the dense floor (1.7 % of the slide height) at which the circle keeps 45 % of the box's
        height. (legend, diameter, below=True)."""
        k = self.kit
        self._pie_last_ok = False
        min_d = int(0.45 * box.h)
        floor = max(0.017 * k.hpt, 7.0)
        sizes = [x for x in reversed(k.sizes) if floor - 0.05 <= x <= k.small + 0.05] or [k.small]
        # a wide, low box (the band left over a conclusion): beside the circle at the dense sizes first — the circle as
        # tall as the box, the legend in the width it leaves
        for size in sizes:
            best = None
            for cols in ((1, 2) if len(rows) >= 4 else (1,)):
                cand = legend(size, int(box.w * (0.45 if cols == 1 else 0.28)), cols)
                if cand["h"] > box.h:
                    continue
                dd = min(box.h, box.w - cand["w"] - k.gap * 2)
                if dd >= min_d and (best is None or dd > best[1]):
                    best = (cand, dd)
            if best is not None:
                self.warnings.append("pie legend set beside the circle at a dense size")
                self._pie_last_ok = True
                return best[0], best[1], False
        lg = None
        for size in sizes:
            sw = _emu(size * 0.72)
            probe = legend(size, _emu(size * 5))
            cap = box.w - sw - int(sw * 0.8) - probe["pct_w"]
            if cap < _emu(size * 5):
                continue
            lg = legend(size, cap, 1)
            d = min(int(box.w * 0.85), box.h - lg["h"] - k.vgap)
            if lg["w"] <= box.w and d >= min_d:
                self.warnings.append("pie legend set under the circle (one column)")
                self._pie_last_ok = True
                return lg, d, True
        self._pie_last_ok = False
        lg = lg or legend(sizes[-1], max(box.w - _emu(sizes[-1] * 4), _emu(sizes[-1] * 5)), 1)
        self.warnings.append("pie legend does not fit its box")
        return lg, max(min(int(box.w * 0.85), box.h - lg["h"] - k.vgap), int(0.3 * box.h)), True

    def _pie_legend_fit(self, box: Bbox, legend, rows: list[int], cats: list[str]):
        """(legend, diameter, legend below?, sizes) of a pie's legend set beside (or under) its circle; legend None
        when no legend fits the box beside or under a circle of at least a third of its height (a legend past the
        box — wider than it, or leaving the circle a dot — is no candidate)."""
        k = self.kit
        gap = k.gap
        min_d = int(0.33 * box.h)
        # the legend beside the circle: the largest size (and the narrowest name column, then two columns of rows) at
        # which the circle keeps most of the height (0.8 of it) — from the lead size (the body size for more than seven
        # parts) down to the body size; below it, the small sizes only when the circle would otherwise keep less than
        # 0.55 of the height. Below the circle only in a tall, narrow box
        top = k.lead if len(rows) <= 7 else k.body
        sizes = [x for x in reversed(k.sizes) if k.small - 0.05 <= x <= top + 0.05] or [k.body, k.small]
        sizes = list(dict.fromkeys(sizes))
        force = getattr(self, "_legend_force", None)
        if force:
            sizes = [s_ for s_ in sizes if s_ <= force + 0.05] or [force]
        best = None
        # two columns of names beside the circle only while it keeps the pie's least size (B5: a low wide box gave a
        # 0.11 H circle next to a two-column legend) — else the legend goes under the circle
        min_2col = min(int(self.PIE_MIN_H * k.H), int(0.8 * box.h))
        for size in sizes:
            if size < k.body - 0.05 and best is not None and best[1] >= 0.55 * box.h:
                break  # the body size keeps a circle of more than half the height: the legend never goes smaller
            natural = _emu(max(text_width_pt(cats[j], k.font, size, False) for j in rows) + size * 0.6)
            # the names on one line first (a legend of even rows), narrower columns when the circle needs the room
            tries = [(cap, 1) for cap in (min(natural, int(box.w * 0.55)), int(box.w * 0.45), int(box.w * 0.36))]
            if len(rows) >= 4:
                tries += [(min(natural, int(box.w * 0.3)), 2)]
            for cap, cols in tries:
                cand = legend(size, cap, cols)
                if cand["h"] > box.h:
                    continue
                dd = min(box.h, box.w - cand["w"] - gap * 2)
                if dd < min_d or (cols == 2 and dd < min_2col):
                    continue  # a legend that leaves the circle a sliver (or runs past the box) is no candidate
                if best is None or dd > best[1] + _emu(2):
                    best = (cand, dd)
                if dd >= 0.8 * box.h:
                    best = (cand, dd)
                    break
            if best is not None and best[1] >= 0.8 * box.h:
                break  # a legend that reads (the larger type) next to a circle of most of the height
        lg, d = best if best is not None else (None, 0)
        below = False
        if lg is None or d < 0.5 * box.h or box.h > 0.7 * box.w:
            # a narrow box (a third of the slide): the legend under the circle when that gives a clearly larger circle
            for size in sizes:
                alts = [legend(size, box.w - _emu(size * 2.2) - _emu(size * 3))]
                if len(rows) >= 4 and d < self._pie_min_d(box):
                    # a circle short of its least size beside the legend: the names in two columns under it (half the
                    # rows' height) before the legend steps down a size
                    alts.append(legend(size, int((box.w - _emu(size * 1.6)) / 2) - _emu(size * 2.2) - _emu(size * 3), 2))
                hit = None
                for alt in alts:
                    d_alt = min(int(box.w * 0.85), box.h - alt["h"] - k.vgap)
                    if alt["w"] > box.w or d_alt < min_d:
                        continue  # under the circle only as wide as the box, the circle still a third of the height
                    if d_alt > d * 1.15 and d_alt >= 0.35 * box.h and (hit is None or d_alt > hit[1]):
                        hit = (alt, d_alt)
                if hit is not None:
                    lg, d, below = hit[0], hit[1], True
                    break
        return lg, d, below, sizes

    def _pie_draw(self, box: Bbox, spec, center: bool, lg: dict, d: int, below: bool, rows: list[int], cats: list[str], vals: list[float], pct: list[str], big: int, colors, amounts: bool, total: float, unit_txt: str) -> None:
        k = self.kit
        gap = k.gap
        group_w = d if below else d + gap * 2 + lg["w"]
        x0 = box.x
        if center:
            x0 = box.x + max(0, (box.w - (max(d, lg["w"]) if below else group_w)) // 2)
        # the compact variant mirrors a lone pie as it mirrors its chart slides: the legend (the text) at the left
        mirror = bool(getattr(self, "_pie_mirror", False)) and not below
        cx = x0 + lg["w"] + gap * 2 if mirror else x0
        gf = add_chart(self.slide, Bbox(x=cx, y=box.y, w=d, h=d), spec, self.outline, self._chart_style(), k.manifest.tokens.typography, text_hex=colors.text, neutral_hex=colors.divider, ground_hex=colors.ground, legend=False, amounts=amounts)
        if amounts:
            self._doughnut_total(gf, spec, total, unit_txt)
        shades = list(getattr(gf, "verstka_shades", None) or []) or [(colors.accent, 1.0)] * len(cats)
        if below:
            lx, ly = x0, box.y + d + k.vgap
        elif mirror:
            lx, ly = x0, box.y + max(0, (d - lg["h"]) // 2)
        else:
            lx, ly = x0 + d + gap * 2, box.y + max(0, (d - lg["h"]) // 2)
        size, sw = lg["size"], lg["sw"]
        if lg["head"]:
            self.cv.text(Bbox(x=lx, y=ly, w=lg["w"], h=lg["head_h"]), lg["head"], name="Legend title")
            ly += lg["head_h"]
        line_h = _emu(size * k.line)
        for ci, chunk in enumerate(lg["chunks"]):
            y = ly
            cx = lx + ci * (lg["col_w"] + lg["col_gap"])
            for n in chunk:
                j = rows[n]
                nh = lg["hs"][n]
                base, share = shades[j % len(shades)]
                self.cv.rect(Bbox(x=cx, y=y + (line_h - sw) // 2, w=sw, h=sw), base, radius_emu=sw // 4, name="Swatch", alpha=share if share < 1.0 else None)
                nx = cx + sw + int(sw * 0.8)
                if lg.get("stack"):
                    # the amount · share on the line under the name: a narrow legend keeps the brief's amounts
                    name_p = self._set_lines(cats[j], size, colors.text, lg["name_w"])
                    name_h = self.h(name_p, lg["name_w"]) - _emu(2)
                    self.cv.text(Bbox(x=nx, y=y, w=lg["name_w"], h=name_h + _emu(2)), name_p, name="Legend")
                    self.cv.text(Bbox(x=nx, y=y + name_h, w=lg["name_w"], h=max(nh - name_h, line_h + _emu(2))), [self.P(pct[j], size, colors.accent_text if j == big else colors.text, bold=k.figure_bold)], name="Share")
                else:
                    self.cv.text(Bbox(x=nx, y=y, w=lg["name_w"], h=nh), self._set_lines(cats[j], size, colors.text, lg["name_w"]), name="Legend")
                    self.cv.text(Bbox(x=nx + lg["name_w"], y=y, w=lg["pct_w"], h=line_h + _emu(2)), [self.P(pct[j], size, colors.accent_text if j == big else colors.text, bold=k.figure_bold, align="r")], name="Share")
                y += nh + lg["row_gap"]

    def _set_lines(self, text: str, size: float, color: str, width: int) -> list[Para]:
        """A legend name broken into its lines where it was measured (one paragraph a line): the row's height counted
        the measured lines, and a renderer's face a little narrower than the estimate (Open Sans set in Noto Sans) no
        longer joins them into one line under a blank one (B3-4)."""
        para = self.P(text, size, color)
        lines = wrap_lines(para.text, self.kit.font, size, False, max(_pt(width), 1.0))
        if len(lines) <= 1:
            return [para]
        return [Para([Run(line_, size, color, False, self.kit.font)], align=para.align) for line_ in lines]

    def _doughnut_total(self, gf, spec, total: float, unit: str) -> None:
        """The whole in a doughnut's hole («780 000 ₽»), when the deck states that total (its facts or a slide's text):
        the parts' sum is not written on a slide as a figure the brief does not give."""
        from verstka.rendering.charts import doughnut_hole_bbox, effective_chart_type

        if effective_chart_type(spec, self.outline) != "doughnut" or total <= 0:
            return
        stated = False
        for f in self.outline.facts:
            v = _num(f.value or "")
            if v is not None and abs(v - total) <= 0.5:
                stated = True
                break
        if not stated:
            texts = [x for sl in self.outline.slides for x in (sl.headline, sl.subtitle or "", sl.takeaway or "", sl.notes or "", *sl.content.bullets)]
            want = _amount_text(total, "").replace(" ", "").replace("\u00a0", "")
            stated = any(want in re.sub(r"[\s\u00a0\u2060]", "", x) for x in texts)
        box = doughnut_hole_bbox(gf) if stated else None
        if box is None:
            return
        k = self.kit
        text = typeset(_amount_text(total, unit))
        size = k.h3
        if getattr(self, "_hole_large", False):
            # the visual variant's doughnut: the whole as its figure, as large as the hole takes it — one line through
            # the centre may run a little wider than the inscribed square (the hole's chord there is 1.38 × its side;
            # the slices' own labels keep the rest)
            wide = int(box.w * 1.15)
            box = Bbox(x=box.x - (wide - box.w) // 2, y=box.y, w=wide, h=box.h)
            size = next((s_ for s_ in k.figure_sizes(k.display, k.h3) if text_width_pt(text, k.font, s_, k.figure_bold) <= _pt(box.w) * 0.88 and s_ * 1.3 <= _pt(box.h) * 0.45), k.h3)
        while size > k.small and text_width_pt(text, k.font, size, k.figure_bold) > _pt(box.w) * 0.95:
            size = k.snap(size * 0.9, k.small, size - 0.5)
        hh = _emu(size * 1.3)
        color = k.colors.figure if getattr(self, "_hole_large", False) else k.colors.heading
        self.cv.text(Bbox(x=box.x, y=box.y + (box.h - hh) // 2, w=box.w, h=hh), [self.P(text, size, color, bold=k.figure_bold, align="c")], name="Total", anchor="ctr")

    def _chart_fallback(self, area: Bbox) -> str:
        """A chart that could not be drawn gives its area to the slide's figures (its numbers, else the facts the chart
        names) as tiles or one big number, else to its text or its cards. Returns what was set, for the warning."""
        from verstka.rendering.fallbacks import fact_numbers

        c = self.o.content
        numbers = list(c.numbers) or fact_numbers(self.o, self.outline, chart_only=True)
        running = [t for t in self._running() if t.strip()]
        items = [it for it in (c.items or c.columns)]
        if len(numbers) == 1:
            self.big_number(area, numbers[0], extra=running, others=[])
            return "one big number"
        if numbers:
            self.kpis(area, numbers[:4], extra=running)
            return f"{len(numbers[:4])} KPI tiles"
        if running:
            self.bullets(area, running)
            return "bullets"
        if items:
            self.cards(area, items)
            return "cards"
        self.warnings.append("empty slide content")
        return "nothing"


# ---------------------------------------------------------------------------------------------- data helpers


_FORMULA_OP_RE = re.compile(r"\s*(=|≈|×|\*|·|÷|\+|(?<=\s)[−–-](?=\s)|(?<=[\d\s])[xх](?=[\s\d])|(?<=\s)/(?=\s))\s*")
_OP_NORM = {"*": "×", "·": "×", "x": "×", "х": "×", "-": "−", "–": "−", "/": "÷"}
_TERM_NUM_RE = re.compile(
    r"(?<![\w])([+\-−]?\d[\d\s\u00a0]*(?:[.,]\d+)?)"
    r"((?:\s*(?:тыс\.?|млн\.?|млрд\.?))?(?:\s*(?:%|₽|\$|€|руб\.?(?![а-я])|рубл[а-я]*))?)",
    re.I,
)


_BEFORE_CAT_RE = re.compile(r"(?i)^\s*(сейчас|было|до\b|текущ|факт|исходн|сегодня|now|before|current)")
_AFTER_CAT_RE = re.compile(r"(?i)(цел[ьи]|план|прогноз|после|стало|через|будет|к\s+\d|target|goal|after|forecast)")


def _amount_text(v: float, unit: str) -> str:
    """315000 → «315 000 ₽», 13.3 → «13,3»: a part's amount as the brief writes it."""
    n = f"{int(round(v)):,}".replace(",", " ") if abs(v - round(v)) < 1e-9 else f"{v:g}".replace(".", ",")
    u = (unit or "").strip()
    if u.lower() in ("рублей", "руб.", "руб", "рубли", "rub"):
        u = "₽"
    return f"{n} {u}".strip() if u else n


def _after_index(spec, outline: DeckOutline) -> Optional[int]:
    """The bar to highlight: the chart's own choice, else — a before/after chart of one series («Сейчас» / «Через 6
    месяцев», «Было» / «Стало») — the «after» bar."""
    if spec.highlight_index is not None or spec.type not in ("column", "bar"):
        return spec.highlight_index
    from verstka.rendering.charts import resolve_series

    try:
        series = resolve_series(spec, outline)
    except Exception:  # noqa: BLE001
        return None
    if len(series) != 1 or len(series[0].categories) != 2:
        return None
    a, b = (str(x) for x in series[0].categories)
    return 1 if _BEFORE_CAT_RE.search(a) and _AFTER_CAT_RE.search(b) else None


def _same_words(a: str, b: str) -> bool:
    """`a` says nothing `b` does not (every word of it is in b): a chart named like its slide's heading."""
    wa = set(re.findall(r"\w+", (a or "").lower().replace("ё", "е")))
    wb = set(re.findall(r"\w+", (b or "").lower().replace("ё", "е")))
    return bool(wa) and wa <= wb


def parse_formula(text: str) -> tuple[list[tuple[Optional[str], str]], list[str]]:
    """«100 покупок × 300 ₽ × 30 дней = 900 000 рублей» → terms [(figure, words)] and the operators between them:
    [("100", "покупок"), ("300 ₽", ""), ("30", "дней"), ("900 000 ₽", "")], ["×", "×", "="]. A figure keeps its
    scale and currency («1,14 млн ₽»; рублей → ₽); the other words of a term are its label. A term without a figure
    («Выручка = покупки × чек») is (None, words)."""
    parts = _FORMULA_OP_RE.split(" ".join((text or "").replace("\u00a0", " ").split()))
    terms_raw, ops = parts[0::2], [_OP_NORM.get(o, o) for o in parts[1::2]]
    while len(terms_raw) > 1 and not terms_raw[0].strip() and ops and ops[0] in ("+", "−"):
        # a leading sign belongs to its figure («+15% × …»), it is no operator
        terms_raw = [(("−" if ops[0] == "−" else "+") + terms_raw[1].strip())] + terms_raw[2:]
        ops = ops[1:]
    terms: list[tuple[Optional[str], str]] = []
    for t in terms_raw:
        t = t.strip()
        m = next((m for m in _TERM_NUM_RE.finditer(t) if any(ch.isdigit() for ch in m.group(1))), None)
        if m is None:
            terms.append((None, t))
            continue
        num = " ".join(m.group(1).split())
        unit = " ".join((m.group(2) or "").split())
        unit = re.sub(r"(?i)рубл[а-я]*|руб\.?", "₽", unit)
        fig = f"{num} {unit}".strip() if unit and unit != "%" else num + unit
        label = " ".join((t[: m.start()] + " " + t[m.end() :]).split()).strip(" ,.:;—–-()")
        terms.append((fig, label))
    return terms, ops


def _num(cell: str) -> Optional[float]:
    m = re.match(r"^\s*[+\-−]?\d[\d\s\u00a0]*(?:[.,]\d+)?", cell or "")
    if not m:
        return None
    try:
        return float(m.group(0).replace("\u00a0", "").replace(" ", "").replace(",", ".").replace("−", "-"))
    except ValueError:
        return None


def table_deltas(table: TableData) -> list[tuple[str, str, str]]:
    """What changed most in a before/after table (a label column and exactly two numeric columns): the change of
    each row, largest first — in percentage points for percentages, as a relative change otherwise. Computed from
    the cells, never invented."""
    if len(table.columns) != 3 or not 2 <= len(table.rows) <= 5:
        return []
    out = []
    for row in table.rows:
        if len(row) < 3:
            return []
        a, b = _num(row[1]), _num(row[2])
        if a is None or b is None:
            return []
        if "%" in row[1] and "%" in row[2]:
            d = b - a
            if abs(d) < 0.5:
                continue
            value = (f"+{d:g}" if d > 0 else f"−{abs(d):g}").replace(".", ",") + " п.п."
            size = abs(d) / 100
        elif a:
            r = (b - a) / abs(a)
            if abs(r) < 0.05:
                continue
            value = (f"+{r * 100:.0f}%" if r > 0 else f"−{abs(r) * 100:.0f}%")
            size = abs(r)
        else:
            continue
        out.append((size, value, row[0].strip(), f"{row[1].strip()} → {row[2].strip()}"))
    out.sort(key=lambda t: -t[0])
    return [(v, lab, span) for _, v, lab, span in out]


_MONTHS = ("январ", "феврал", "март", "апрел", "май", "мая", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр")


def _in_sentence(cat: str) -> str:
    """A category inside a sentence: month names are lower case in Russian («рост за период июль → сентябрь»)."""
    c = cat.strip()
    return c[:1].lower() + c[1:] if c.lower().startswith(_MONTHS) else c


_MONEY_UNIT_RE = re.compile(r"\s*(?:рубл(?:ь|я|ей)|руб\.?)$", re.I)


_TAIL_LINK_RE = re.compile(r"(?:\s+(?:до|на|в|с|—|–|-|:|→|->)|\s+(?:составит|составляет|составят|составляют|равна|равен|равно|—\s*это))+\s*$", re.I)


def label_without_figure(text: str, fig: Optional[str]) -> str:
    """The words under a figure set large, without the figure said again: «Рентабельность увеличится до 22,4%» under
    «22,4%» reads «Рентабельность увеличится», «Месячная выручка составит 1 138 500 ₽» → «Месячная выручка», a figure
    in brackets goes with its brackets. A figure in the middle of the sentence stays (cutting it would break the
    sentence), as does the sentence when too little of it would be left."""
    text = " ".join((text or "").split())
    if not fig:
        return text
    want = re.sub(r"\D", "", fig)
    for m in _FIG_RE.finditer(text):
        if re.sub(r"\D", "", m.group(1)) != want:
            continue
        a, b = m.start(1), m.end(1)
        before, after = text[:a], text[b:]
        if before.rstrip().endswith("(") and after.lstrip().startswith(")"):
            out = (before.rstrip()[:-1].rstrip() + " " + after.lstrip()[1:].lstrip()).strip()
        elif re.fullmatch(r"\s*(?:в\s+месяц|в\s+год|в\s+день)?\s*[.!]?\s*", after):
            out = _TAIL_LINK_RE.sub("", before.rstrip()).strip(" ,;:—–-")
        elif re.match(r"\s*(?:в\s+месяц\s*)?,\s*или\s+\S", after):
            # «Операционная прибыль составляет 120 000 рублей, или 13,3% выручки» → «Операционная прибыль — 13,3% выручки»
            subject = _TAIL_LINK_RE.sub("", before.rstrip()).strip(" ,;:—–-")
            rest = re.sub(r"^\s*(?:в\s+месяц\s*)?,\s*или\s+", "", after).strip().rstrip(".")
            out = f"{subject} — {rest}" if subject else ""
        else:
            return text
        if not out or (len(out.split()) < 2 and len(out) < 4):
            return text
        return out[:1].upper() + out[1:]
    return text


_YEAR_RE = re.compile(r"(?:1[89]|20)\d{2}")
_YEAR_WORD_RE = re.compile(r"\s*(?:год|г\.|гг\.?|-?[мйхго]\b)", re.I)
_YEAR_BEFORE_RE = re.compile(r"(?:^|[\s(«])(?:в|во|с|со|к|до|по|от|из|квартал[еау]?|полугоди[еия]|январ[ья]|феврал[ья]|марта?|апрел[ья]|ма[йя]|июн[ья]|июл[ья]|августа?|сентябр[ья]|октябр[ья]|ноябр[ья]|декабр[ья])\s*$", re.I)


def _is_year(m: "re.Match", text: str) -> bool:
    """A figure match that is a year, not a quantity: four digits 1800–2099 with no unit, followed by «год/года/г.»,
    or standing without a counted word after it («в первом квартале 2024», «с 2019 по 2024»). «1854 электромобиля» is
    a count and stays a figure."""
    fig = m.group(1).strip()
    if not _YEAR_RE.fullmatch(fig):
        return False
    after = text[m.end(1):]
    if _YEAR_WORD_RE.match(after):
        return True
    if not re.match(r"\s*[а-яёa-z]{3,}", after, re.I):
        return True  # «… квартале 2024», «с 2019 по 2024»: no counted word after it
    return bool(_YEAR_BEFORE_RE.search(text[: m.start(1)]))  # «в 2024 выручка выросла»


_RANGE_TO_RE = re.compile(r"\s*(?:→|->|⟶|—>|до)\s*$", re.I)
_RANGE_FROM_RE = re.compile(r"(?:^|[\s(«])(?:с|со|от)\s*$", re.I)


def _range_start(text: str, m: "re.Match", nxt: Optional["re.Match"]) -> bool:
    """The figure opens a range whose end is the next figure: «с X до Y», «от X до Y», «X → Y»."""
    if nxt is None:
        return False
    between = text[m.end(1): nxt.start(1)]
    if not _RANGE_TO_RE.fullmatch(between) and not re.fullmatch(r"\s*(?:→|->|⟶|—>)\s*", between):
        return False
    if re.search(r"→|->|⟶|—>", between):
        return True
    return bool(_RANGE_FROM_RE.search(text[: m.start(1)]))


def _aside_figure(text: str) -> Optional[str]:
    """The key figure of a conclusion, as the large figure beside the chart: the first sum or share
    («120 000 рублей» → «120 000 ₽», «112,3%»), else the first figure of three digits or more; None without one."""
    text = text or ""
    ms = [m for m in _FIG_RE.finditer(text) if not _is_year(m, text)]
    # a change «с 13,3% до 22,4%», «от X до Y», «X → Y»: the figure is where it goes, the target (B6) — the first
    # figure beside a growth message read as the old value
    found = [m.group(1).strip() for i, m in enumerate(ms) if not _range_start(text, m, ms[i + 1] if i + 1 < len(ms) else None)]
    unit = [f for f in found if re.search(r"%|₽|руб|млн|млрд|тыс", f, re.I)]
    big = [f for f in found if len(re.sub(r"\D", "", f)) >= 3]
    pick = (unit or big or [None])[0]
    if pick is None:
        return None
    return typeset(_MONEY_UNIT_RE.sub(" ₽", pick).strip())


def _series_takeaway(spec, outline: DeckOutline, headline: str = "") -> Optional[tuple[str, str]]:
    """The figure a reader takes from a series: its growth (last/first) or its last value. Computed, never invented.
    When the heading already names the multiple («рост в 5 раз»), the panel gives the last value instead of a
    second, differently rounded multiple."""
    from verstka.rendering.charts import resolve_series

    try:
        series = resolve_series(spec, outline)
    except Exception:  # noqa: BLE001
        return None
    if not series or len(series) > 1:
        return None
    s = series[0]
    vals = list(s.values)
    if len(vals) < 2 or any(v is None for v in vals):
        return None
    first, last = vals[0], vals[-1]
    # the heading already states the change («вырастет на 112,3%», «в 2 раза»): a second, differently computed
    # measure of it beside the chart («×2,1») says nothing new and is not the brief's figure
    if re.search(r"\d\s?%|\bв\s+\d+([.,]\d+)?\s*раз|[×x]\s?\d|вдвое|втрое", headline or "", re.I):
        return None
    periods = all(re.search(r"(?i)месяц|квартал|год|недел|янв|фев|мар|апр|ма[йя]|июн|июл|авг|сен|окт|ноя|дек|\b\d{4}\b", c or "") for c in s.categories[1:]) if s.categories else False
    # a change is read only along an order: periods, or a before/after pair («Сейчас → Цель», «До → После»). Parts of a
    # set (countries, products, regions) have no first and last: «−89%: СССР → Япония» would be a made-up measure
    before_after = len(s.categories) == 2 and bool(re.search(r"(?i)сейчас|текущ|было|до\b|старт|начал|база|факт", s.categories[0] or "")) and bool(re.search(r"(?i)цель|прогноз|план|стало|после|итог|будет", s.categories[1] or ""))
    if not (periods or before_after):
        return None
    span = f"{_in_sentence(s.categories[0])} → {_in_sentence(s.categories[-1])}" if s.categories else ""
    span = f"за период {span}" if periods and span else (f"{span}" if span else "")
    says_multiple = False
    if first > 0 and last > 0 and not says_multiple:
        ratio = last / first
        if ratio >= 1.5:
            txt = f"×{ratio:.1f}".replace(".", ",").replace(",0", "")
            return txt, f"рост: {span}".strip(": ") if span and not span.startswith("за период") else f"рост {span}".strip()
        if ratio <= 0.67:
            pct = int(round((1 - ratio) * 100))
            return f"−{pct}%", f"снижение: {span}".strip(": ") if span and not span.startswith("за период") else f"снижение {span}".strip()
    # the change in the series' own unit — never the last value, which the highlighted bar already shows
    delta = last - first
    if not delta:
        return None
    unit = f" {s.unit}" if s.unit and s.unit not in ("%",) else (s.unit or "")
    mag = abs(delta)
    v = f"{mag:,.0f}".replace(",", " ") if mag >= 100 else f"{mag:g}".replace(".", ",")
    if re.sub(r"\s", "", v) in re.sub(r"\s", "", headline or ""):
        return None  # the heading already says it
    word = "прирост" if delta > 0 else "снижение"
    return f"{'+' if delta > 0 else '−'}{v}{unit}", (f"{word}: {span}" if span and not span.startswith("за период") else f"{word} {span}").strip(": ")
