"""Clone renderer: copy a sample slide and rewrite its slots with outline content."""

from __future__ import annotations

import copy
import logging
import math
import re
import weakref
from dataclasses import dataclass, field
from typing import Optional

from lxml import etree
from pptx.slide import Slide

from verstka.analysis.patterns import container_inset
from verstka.analysis.shapes import looks_like_placeholder, looks_like_sample_value
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.scorer import awkward_breaks, balanced_lines, bind_short_words, bookend_max_lines, cover_goal, display_fit, display_lines, is_placeholder_text, same_words, split_display_title
from verstka.rendering.assets_pick import pick_asset, pick_icon
from verstka.rendering.charts import add_chart, add_ring
from verstka.rendering.deck import DeckBuilder, element_bbox, is_nested, materialize_xfrm, remove_element, renumber_ids, set_element_pos, shift_element, slide_shape_elements
from verstka.rendering.fonts import text_width_pt, wrap_lines
from verstka.rendering.fit import fit_size, grow_size
from verstka.rendering.groups import adjust_group, cell_bbox, cell_riders, cells_elements, reading_order_key, transform_element
from verstka.rendering.images import replace_picture
from verstka.rendering.layers import ground_under, heading_band, opaque_box
from verstka.rendering.tables import add_table
from verstka.ru import typeset_figures
from verstka.rendering.textfill import ParagraphSpec, clear_text, effective_insets, ensure_txbody, fill_text, has_visible_style, set_text_size, shape_text, style_runs
from verstka.schemas.common import EMU_PER_PT, Bbox, PatternKind, SlotRole, contrast_ratio, relative_luminance
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideItem
from verstka.schemas.template import Pattern, RepeatGroup, Slot, TemplateManifest

log = logging.getLogger(__name__)

TEXT_ROLES = {SlotRole.title, SlotRole.subtitle, SlotRole.body, SlotRole.bullet_list, SlotRole.card_title, SlotRole.card_body, SlotRole.number, SlotRole.number_label, SlotRole.caption}
CELL_TEXT_ROLES = {SlotRole.card_title, SlotRole.card_body, SlotRole.number, SlotRole.number_label, SlotRole.bullet_list, SlotRole.body}
_BODY_ROLES = (SlotRole.card_body, SlotRole.bullet_list, SlotRole.body)
_ITEM_KINDS = {PatternKind.cards, PatternKind.process, PatternKind.timeline, PatternKind.team, PatternKind.comparison, PatternKind.agenda, PatternKind.two_column}
_DEFAULT_INSETS = (91440, 45720, 91440, 45720)
_ICON_MAX_W = 0.06  # pictures narrower than this share of the slide are icons: part of the card design


@dataclass
class RenderedSlide:
    outline_id: str
    index: int
    mode: str
    pattern_id: Optional[str] = None
    composition: Optional[str] = None
    warnings: list[str] = field(default_factory=list)


class _SlideCtx:
    def __init__(self, builder: DeckBuilder, slide: Slide, pattern: Pattern, manifest: TemplateManifest, ws: TemplateWorkspace, outline: DeckOutline) -> None:
        self.builder = builder
        self.slide = slide
        self.pattern = pattern
        self.manifest = manifest
        self.ws = ws
        self.outline = outline
        self.els = slide_shape_elements(slide)
        self.filled: set[str] = set()
        self.removed: set[str] = set()
        self.cell_member_ids: set[str] = set()  # every shape of a repeat-group cell that received an item
        self.warnings: list[str] = []
        self.next_id = builder.next_shape_id(slide)
        self.typo = manifest.tokens.typography
        self.scale = [s.size_pt for s in self.typo.scale]
        # text grows along every size the template really uses (VK Tech: 12 → 13.22, 14, 15, 16), not only the role
        # scale, whose steps can be far apart (12 → 24)
        self.grow_scale = sorted(set(self.scale) | {float(x) for x in (self.typo.sizes_used or []) if x >= 8})
        self.W, self.H = manifest.slide_size.w, manifest.slide_size.h
        self.role_of = {s.shape_id: s.role for s in pattern.slots}
        self.slot_of = {s.shape_id: s for s in pattern.slots}
        self.used_assets: set[str] = set()
        self.origin: dict[str, str] = {}  # id of a duplicated shape → id of the sample shape it was copied from
        # whether the slide keeps the sample's content pictures (it has an image of its own); when it does not, the
        # big pictures in cards are dropped at the end and must not hold text back while it is being placed
        self.keeps_pictures = True
        self.band_limited = False  # the heading being filled stands in a layout/master band (T12)
        self.goal_to_place: Optional[str] = None  # a cover goal that left its place under the subtitle (C3-1)

    # ---- helpers ----------------------------------------------------------------
    def slots(self, *roles: SlotRole, unfilled: bool = True) -> list[Slot]:
        out = [s for s in self.pattern.slots if s.role in roles and s.shape_id in self.els]
        if unfilled:
            out = [s for s in out if s.shape_id not in self.filled and s.shape_id not in self.removed]
        return out

    def id_of(self, el: etree._Element) -> Optional[str]:
        nv = el.find(".//" + q("p:cNvPr"))
        return nv.get("id") if nv is not None else None

    def fill_el(self, el: etree._Element, slot: Optional[Slot], paragraphs: list[ParagraphSpec], *, size_hint: Optional[float] = None, min_ratio: float = 0.6, grow_to: Optional[float] = None, widen: bool = False) -> None:
        if not paragraphs:
            return
        box = element_bbox(el)
        style = slot.style if slot else None
        slot_size = style.size_pt if style and style.size_pt else self.typo.size_for("body")
        size = size_hint or slot_size
        bold = bool(style and style.bold) if style else False
        family = style.font_family if style and style.font_family else self.typo.primary_family
        target_size = size
        container = slot is not None and slot.container and etree.QName(el).localname == "sp"
        if container and box and box[2] > 0 and box[3] > 0:
            pad = container_inset(box[2], box[3])
            ensure_txbody(el, (pad, pad, pad, pad))
        insets = _body_insets(el)
        if box and box[2] > 0 and box[3] > 0:
            _wrap_inside(el, [p.text for p in paragraphs], family, size, bold, box, insets)
        if box and box[2] > 0 and box[3] > 0 and slot is not None and slot.role in _GROWABLE and not container:
            box = self._room_in_holder(el, box)
        self.band_limited = False
        if box and box[2] > 0 and box[3] > 0 and slot is not None and (widen or slot.role in (SlotRole.title, SlotRole.subtitle)):
            box = self._widen_on_backing(el, box, [p.text for p in paragraphs], family, size, bold, insets)
        if self.band_limited:
            min_ratio = min(min_ratio, 0.55)  # a heading in a band steps down rather than leaving the band
        if box and box[2] > 0 and box[3] > 0:
            res = fit_size([p.text for p in paragraphs], Bbox(x=box[0], y=box[1], w=box[2], h=box[3]), family, size, bold, self.grow_scale, insets_emu=insets, line_spacing=self.typo.line_height, min_ratio=min_ratio)
            target_size = res.size_pt
            if grow_to and res.fits and target_size >= size:
                target_size = grow_size([p.text for p in paragraphs], Bbox(x=box[0], y=box[1], w=box[2], h=box[3]), family, target_size, bold, self.grow_scale, grow_to, insets_emu=insets, line_spacing=self.typo.line_height)
            if not res.fits:
                if slot is not None and slot.role == SlotRole.number and len(paragraphs) == 1 and not is_nested(el):
                    # a single figure never wraps well: widen the box instead of clipping
                    need_w = int(text_width_pt(paragraphs[0].text, family, target_size, True) * EMU_PER_PT * 1.15) + insets[0] + insets[2]
                    if need_w > box[2] and box[0] + need_w <= self.W * 0.97:
                        set_element_pos(el, w=need_w)
                        self.warnings.append(f"widened number slot {slot.id}")
                    elif need_w > box[2]:
                        shrink = fit_size([p.text for p in paragraphs], Bbox(x=box[0], y=box[1], w=box[2], h=box[3]), family, size, bold, self.grow_scale, insets_emu=insets, line_spacing=self.typo.line_height, min_ratio=0.2)
                        target_size = shrink.size_pt
                else:
                    self.warnings.append(f"text may overflow in {slot.id if slot else 'shape'} ({res.lines} lines)")
        if box and box[2] > 0 and box[3] > 0 and any(p.size_pt for p in paragraphs):
            # an emphasised line (a figure) may be larger than the rest only as far as the box height allows
            ls = self.typo.line_height
            usable_w = max((box[2] - insets[0] - insets[2]) / EMU_PER_PT, 1.0)
            usable_h = max((box[3] - insets[1] - insets[3]) / EMU_PER_PT, 1.0)
            rest = sum(max(len(wrap_lines(p.text, family, target_size, bold, usable_w)), 1) for p in paragraphs if not p.size_pt) * target_size * ls
            paragraphs = [ParagraphSpec(p.text, bullet=p.bullet, level=p.level, bold=p.bold, size_pt=p.size_pt, color_hex=p.color_hex) for p in paragraphs]
            for p in paragraphs:
                if p.size_pt and p.size_pt > target_size:
                    size = p.size_pt
                    if len(p.text) <= 16:
                        # a figure is read at a glance: one line («4,6 из 5», not «4,6 из» over «5»)
                        while size > target_size and text_width_pt(p.text, family, size, True) > usable_w * 0.98:
                            size = round(size * 0.94, 1)
                    lines = max(len(wrap_lines(p.text, family, size, True, usable_w)), 1)
                    room = (usable_h - rest) / (ls * lines)
                    p.size_pt = _snap_down(max(target_size, min(size, room)), self.grow_scale)
        explicit = size_hint is not None or target_size != slot_size or container
        fill_text(el, paragraphs, size_pt=target_size if explicit else None)
        if container:
            style_runs(el, family, style.color_hex if style else None)
        sid = self.id_of(el)
        if sid:
            self.filled.add(sid)

    def _room_in_holder(self, el: etree._Element, box: tuple[int, int, int, int]) -> tuple[int, int, int, int]:
        """A text box drawn on a card may take the card's empty height below it — down to the next text or picture
        in the card, or to the card's inner edge (the sample's one-line box would otherwise force tiny type)."""
        if is_nested(el):
            return box
        bodyPr = el.find(q("p:txBody") + "/" + q("a:bodyPr"))
        if bodyPr is not None and bodyPr.get("anchor") == "b":
            return box
        bx = Bbox(x=box[0], y=box[1], w=box[2], h=box[3])
        tol = int(0.01 * self.W)
        tree = el.getparent()
        holder = None
        for other in tree:
            if other is el or etree.QName(other).localname != "sp" or shape_text(other).strip() or not _has_fill(other):
                continue
            b = element_bbox(other)
            if not b or b[2] * b[3] > 0.5 * self.W * self.H or b[2] * b[3] < 1.5 * bx.area:
                continue
            ob = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
            if ob.x - tol <= bx.x and ob.y - tol <= bx.y and ob.x2 + tol >= bx.x2 and ob.y2 + tol >= bx.y2 and (holder is None or ob.area < holder.area):
                holder = ob
        if holder is None:
            return box
        pad = max(min(bx.x - holder.x, int(0.04 * self.H)), int(0.012 * self.H))
        limit = holder.y2 - pad
        for other in tree:
            if other is el or etree.QName(other).localname not in ("sp", "pic", "graphicFrame", "grpSp"):
                continue
            b = element_bbox(other)
            if not b or (b[0], b[1], b[2], b[3]) == (holder.x, holder.y, holder.w, holder.h):
                continue
            inside = holder.x <= b[0] + b[2] / 2 <= holder.x2 and holder.y <= b[1] + b[3] / 2 <= holder.y2
            overlaps_x = b[0] < bx.x2 and b[0] + b[2] > bx.x
            if etree.QName(other).localname == "pic" and not self.keeps_pictures and b[2] > _ICON_MAX_W * self.W:
                continue  # a sample screenshot in the card: it leaves at the end, the text may take its place
            is_content = etree.QName(other).localname != "sp" or shape_text(other).strip() or (self.id_of(other) or "") in self.slot_of
            if inside and overlaps_x and b[1] >= bx.y2 - tol and is_content:
                limit = min(limit, b[1] - int(0.012 * self.H))
        if limit <= bx.y2 + int(0.02 * self.H):
            return box
        set_element_pos(el, h=limit - bx.y)
        return (bx.x, bx.y, bx.w, limit - bx.y)

    def _widen_on_backing(self, el: etree._Element, box: tuple[int, int, int, int], texts: list[str], family: Optional[str], size: float, bold: bool, insets: tuple[int, int, int, int]) -> tuple[int, int, int, int]:
        """Headings keep to their own ground.

        On a label (a pill or plate painted under the start of the heading box — the LCT titles sit on short pills
        inside a wide box) the text is held to the label's width, and the label grows to the right with the text up
        to the next obstacle (logos of the layout included) or the safe area; a heading that still needs two lines
        makes both one line taller. Without a label the box is cut before logos and other chrome in its band, so the
        heading wraps instead of running under them.
        """
        if is_nested(el):
            return box
        bx = Bbox(x=box[0], y=box[1], w=box[2], h=box[3])
        tol = int(0.01 * self.W)
        backing = None
        for other in self.els.values():
            if other is el or is_nested(other) or other.getparent() is None or etree.QName(other).localname != "sp" or shape_text(other).strip() or not _has_fill(other):
                continue
            b = element_bbox(other)
            if not b or b[2] * b[3] > 0.25 * self.W * self.H:
                continue
            ob = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
            under_start = ob.x - tol <= bx.x <= ob.x2 - int(0.04 * self.W) and ob.y - tol <= bx.y and ob.y2 + tol >= bx.y2
            if under_start and (backing is None or ob.area < backing[1].area):
                backing = (other, ob)
        limit = int(self.manifest.tokens.spacing.safe_area.x2 * self.W)
        obstacles = [element_bbox(o) for o in self.els.values() if o is not el and (backing is None or o is not backing[0]) and not is_nested(o) and o.getparent() is not None]
        obstacles += [(lambda b: (b.x, b.y, b.w, b.h))(c.bbox.to_emu(self.W, self.H)) for c in self.manifest.tokens.chrome]
        band = backing[1] if backing is not None else bx
        for b in obstacles:
            if not b or b[2] * b[3] >= 0.6 * self.W * self.H:
                continue
            if b[1] < band.y2 and b[1] + b[3] > band.y and b[0] >= bx.x + int(0.1 * self.W):
                limit = min(limit, b[0] - int(0.015 * self.W))
        need_w = int(max(text_width_pt(t, family, size, bold) for t in texts) * EMU_PER_PT * 1.05) + insets[0] + insets[2]
        if backing is None:
            banded = self._hold_in_band(el, bx, need_w, limit)
            if banded is not None:
                return banded
            if bx.x2 > limit and bx.x + need_w > limit and limit - bx.x >= int(0.25 * self.W):
                set_element_pos(el, w=limit - bx.x)
                self.warnings.append("heading kept clear of the logos")
                return (bx.x, bx.y, limit - bx.x, bx.h)
            return box
        back_el, ob = backing
        pad_r = max(ob.x2 - (bx.x + need_w), 0) if bx.x + need_w <= ob.x2 else int(0.012 * self.W)
        inner = ob.x2 - bx.x - int(0.008 * self.W)  # the label's own width is the text's to use
        width = min(max(bx.w, min(need_w, inner)), max(inner, int(0.05 * self.W)))
        grow = 0
        if need_w > width:
            target_x2 = min(limit, bx.x + need_w + pad_r)
            grow = max(0, target_x2 - ob.x2)
            width = min(max(width, target_x2 - bx.x), max(bx.w, target_x2 - bx.x))
        height_grow = 0
        res = fit_size(texts, Bbox(x=bx.x, y=bx.y, w=width, h=bx.h), family, size, bold, self.grow_scale, insets_emu=insets, line_spacing=self.typo.line_height, min_ratio=0.8)
        if not res.fits:
            # two lines at a size down to 0.7 of the heading's: the label grows by one line, the heading stays legible
            h2 = int(2 * size * self.typo.line_height * EMU_PER_PT) + insets[1] + insets[3]
            two = fit_size(texts, Bbox(x=bx.x, y=bx.y, w=width, h=h2), family, size, bold, self.grow_scale, insets_emu=insets, line_spacing=self.typo.line_height, min_ratio=0.7)
            if two.fits and two.lines <= 2:
                height_grow = max(0, int(two.height_pt * EMU_PER_PT) + insets[1] + insets[3] - bx.h)
        set_element_pos(el, w=width, h=bx.h + height_grow)
        if grow or height_grow:
            set_element_pos(back_el, w=ob.w + grow, h=ob.h + height_grow)
            self.warnings.append(f"label widened by {grow * 100 // self.W}% of the slide for the heading")
        return (bx.x, bx.y, width, bx.h + height_grow)

    def _hold_in_band(self, el: etree._Element, bx: Bbox, need_w: int, limit: int) -> Optional[tuple[int, int, int, int]]:
        """A heading standing in a band its layout or master paints (T12: the band is a hard limit): the box keeps to
        the band's height — a one-line band holds one line, set smaller rather than spilling its second line onto the
        slide's ground in the band's text colour — and may widen along the band up to the next obstacle. None when no
        such band holds the heading (the slide's own pills are the label logic's)."""
        try:
            band = heading_band(self.slide, bx)
        except Exception:  # noqa: BLE001 - layers are advice
            band = None
        if band is None or band.source == "slide":
            return None
        b = band.box
        tol = int(0.01 * self.H)
        y0, y1 = max(bx.y, b.y), min(bx.y2, b.y2)
        if y1 - y0 < 0.5 * min(bx.h, b.h):
            return None  # the heading box mostly lies off the band: the band is not its ground
        x2 = min(limit, b.x2 - int(0.015 * self.W))
        width = max(bx.w, min(x2 - bx.x, need_w)) if x2 > bx.x else bx.w
        if bx.y >= b.y - tol and bx.y2 <= b.y2 + tol and width == bx.w:
            self.band_limited = True
            return (bx.x, bx.y, bx.w, bx.h)
        set_element_pos(el, x=bx.x, y=y0, w=width, h=y1 - y0)
        self.band_limited = True
        self.warnings.append("heading held inside the layout's band")
        return (bx.x, y0, width, y1 - y0)

    def fill_slot(self, slot: Slot, paragraphs: list[ParagraphSpec], **kw) -> bool:
        el = self.els.get(slot.shape_id)
        if el is None:
            return False
        self.fill_el(el, slot, paragraphs, **kw)
        return True

    def remove_slot_shape(self, slot: Slot) -> None:
        el = self.els.get(slot.shape_id)
        if el is not None and not is_nested(el):
            remove_element(el)
            self.removed.add(slot.shape_id)

    def remove_el(self, el: etree._Element) -> None:
        sid = self.id_of(el)
        remove_element(el)
        if sid:
            self.removed.add(sid)

    def largest(self, slots: list[Slot]) -> Optional[Slot]:
        return max(slots, key=lambda s: s.bbox.area) if slots else None

    def smallest_font(self, slots: list[Slot]) -> Optional[Slot]:
        return min(slots, key=lambda s: (s.style.size_pt or 99.0)) if slots else None



_GROWABLE = {SlotRole.body, SlotRole.bullet_list, SlotRole.card_body, SlotRole.card_title, SlotRole.number_label, SlotRole.caption}


def clear_width(manifest: TemplateManifest, W: int, H: int, box: Bbox) -> int:
    """Width a heading box may use without running under logos and other chrome standing in its band."""
    limit = int(manifest.tokens.spacing.safe_area.x2 * W)
    for c in manifest.tokens.chrome:
        b = c.bbox.to_emu(W, H)
        if b.w * b.h >= 0.6 * W * H:
            continue
        if b.y < box.y2 and b.y2 > box.y and b.x >= box.x + int(0.1 * W):
            limit = min(limit, b.x - int(0.015 * W))
    return box.w if limit - box.x < int(0.25 * W) else min(box.w, limit - box.x)


def _wrap_inside(el: etree._Element, texts: list[str], family: Optional[str], size: float, bold: bool, box: tuple[int, int, int, int], insets: tuple[int, int, int, int]) -> None:
    """A box set to «no wrap» for a one-word sample must wrap a longer text inside its own width, not run across
    the neighbouring column."""
    bodyPr = el.find(q("p:txBody") + "/" + q("a:bodyPr"))
    if bodyPr is None or bodyPr.get("wrap") != "none":
        return
    widest = max((text_width_pt(t, family, size, bold) for t in texts), default=0.0) * EMU_PER_PT
    if widest > box[2] - insets[0] - insets[2]:
        bodyPr.set("wrap", "square")


def _has_fill(el: etree._Element) -> bool:
    """A shape that paints an area: explicit solid/gradient/picture fill or a theme fill reference."""
    spPr = el.find(q("p:spPr"))
    if spPr is not None:
        if spPr.find(q("a:noFill")) is not None:
            return False
        if any(spPr.find(q(t)) is not None for t in ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")):
            return True
    ref = el.find(q("p:style") + "/" + q("a:fillRef"))
    return ref is not None and (ref.get("idx") or "0") != "0"


def _body_insets(el: etree._Element) -> tuple[int, int, int, int]:
    """Text insets of a shape (the templates set explicit zeros; PowerPoint defaults apply when absent)."""
    txBody = el.find(q("p:txBody"))
    bodyPr = txBody.find(q("a:bodyPr")) if txBody is not None else None
    if bodyPr is None:
        return _DEFAULT_INSETS
    out = []
    for name, default in zip(("lIns", "tIns", "rIns", "bIns"), _DEFAULT_INSETS):
        v = bodyPr.get(name)
        try:
            out.append(int(v) if v is not None else default)
        except ValueError:
            out.append(default)
    return out[0], out[1], out[2], out[3]


# ---------------------------------------------------------------------------- content helpers


def _items_for(oslide: OutlineSlide) -> list[SlideItem]:
    c = oslide.content
    if c.items:
        return list(c.items)
    if c.columns:
        return list(c.columns)
    if c.numbers:
        return [SlideItem(title=n.value, text=n.label, number=n.value) for n in c.numbers]
    if oslide.kind in _ITEM_KINDS and c.bullets:
        return [SlideItem(title=b) for b in c.bullets]
    return []


def _item_body_paragraphs(item: SlideItem) -> list[ParagraphSpec]:
    out = [ParagraphSpec(item.text, bullet=False)] if item.text else []
    out += [ParagraphSpec(b, bullet=True) for b in item.bullets]
    return out


def _label_paragraphs(item: SlideItem) -> list[ParagraphSpec]:
    """Everything of an item except its figure: the title (bold when a text follows) and the text/bullets."""
    body = _item_body_paragraphs(item)
    head = item.title if item.title and item.title != item.number else ""
    if head and body:
        return [ParagraphSpec(head, bullet=False, bold=True)] + body
    if head:
        return [ParagraphSpec(head, bullet=False)]
    if body:
        return body
    return [ParagraphSpec(item.title, bullet=False)] if item.title else []


def _item_single_paragraphs(item: SlideItem, include_number: bool = True) -> list[ParagraphSpec]:
    """Everything of an item in one text shape (cells that have a single text slot)."""
    if item.number:
        if not include_number:
            return _label_paragraphs(item)
        return [ParagraphSpec(item.number, bullet=False, bold=True)] + _label_paragraphs(item)
    return _label_paragraphs(item)


def _ordinal(idx: int, sample: Optional[str]) -> str:
    """01 / 1 style sequence number matching the template's sample."""
    sm = (sample or "").strip()
    if len(sm) == 2 and sm.isdigit() and sm.startswith("0"):
        return f"{idx + 1:02d}"
    return str(idx + 1)


def _merge_items(items: list[SlideItem], n: int) -> list[SlideItem]:
    """Reduce items to n cells: plain theses are spread evenly as bullet lists, richer items fold into the last cell."""
    if len(items) <= n or n <= 0:
        return items
    if all(not x.text and not x.bullets for x in items):
        # every cell gets a share: 5 theses into 4 cells is 2+1+1+1, never 2+2+1 and an empty fourth card
        base, extra = divmod(len(items), n)
        out, k = [], 0
        for j in range(n):
            take = base + (1 if j < extra else 0)
            chunk = items[k : k + take]
            k += take
            if len(chunk) == 1:
                out.append(chunk[0])
            elif chunk:
                out.append(SlideItem(title="", bullets=[x.title for x in chunk], icon_hint=chunk[0].icon_hint, number=chunk[0].number))
        return out
    head = items[: n - 1]
    tail = items[n - 1 :]
    texts = [tail[0].text]
    bullets = list(tail[0].bullets)
    for x in tail[1:]:
        if x.text:
            texts.append(f"{x.title}: {x.text}")
        bullets += x.bullets
        if not x.text and not x.bullets:
            bullets.append(x.title)
    merged = SlideItem(title=tail[0].title, text="; ".join(t for t in texts if t), icon_hint=tail[0].icon_hint, number=tail[0].number, bullets=bullets)
    return head + [merged]


def _text_groups(ctx: _SlideCtx, exclude_number_only: bool = False) -> list[RepeatGroup]:
    out = []
    for g in ctx.pattern.repeat_groups:
        ids = {sid for cell in g.member_shape_ids for sid in cell}
        roles = {ctx.role_of.get(sid) for sid in ids}
        if roles & CELL_TEXT_ROLES and len(g.member_shape_ids) >= 2:
            if exclude_number_only and not (roles & {SlotRole.card_title, SlotRole.card_body, SlotRole.bullet_list, SlotRole.body}):
                continue
            out.append(g)
    return out


def _slot_chars(ctx: _SlideCtx, slot: Slot) -> int:
    if slot.capacity.max_chars:
        return slot.capacity.max_chars
    size = slot.style.size_pt or ctx.typo.size_for("body", 14.0)
    w_pt = slot.bbox.w * ctx.W / EMU_PER_PT
    h_pt = slot.bbox.h * ctx.H / EMU_PER_PT
    return int(max(w_pt / (0.55 * size), 0) * max(h_pt / (1.2 * size), 0))


def _cells_can_hold_bullets(ctx: _SlideCtx, bullets: list[str]) -> bool:
    """Bullets become cell items only when the biggest text group can actually show them (not legend chips)."""
    groups = _text_groups(ctx, exclude_number_only=True)
    if not groups:
        return False
    group = max(groups, key=lambda g: len(g.member_shape_ids))
    slots = [ctx.slot_of[sid] for cell in group.member_shape_ids for sid in cell if sid in ctx.slot_of and ctx.slot_of[sid].role in CELL_TEXT_ROLES]
    if not slots:
        return False
    if max(s.bbox.w for s in slots) < 0.10:
        return False
    need = sum(len(b) for b in bullets)
    return sum(_slot_chars(ctx, s) for s in slots) >= 0.5 * need


# ---------------------------------------------------------------------------- repeat-group cells


def _cell_sort_key(cells: list[list[etree._Element]], slide_h: int):
    return reading_order_key(cells, slide_h)


def _companions(ctx: _SlideCtx, group: RepeatGroup, before: list[Optional[Bbox]]) -> list[tuple[RepeatGroup, list[tuple[int, list[etree._Element]]]]]:
    """Other groups on the same axis whose cells sit in this group's cell positions: [(group, [(cell index, elements)])]."""
    boxes = [b for b in before if b is not None]
    if group.axis not in ("row", "column") or len(boxes) != len(before) or not boxes:
        return []
    row = group.axis == "row"
    starts = [b.x if row else b.y for b in boxes]
    size0 = boxes[0].w if row else boxes[0].h
    pitch = (starts[1] - starts[0]) if len(starts) > 1 else size0 + int(group.gap * (ctx.W if row else ctx.H))
    if pitch <= 0:
        return []
    band_lo = min((b.y if row else b.x) for b in boxes)
    band_hi = max((b.y2 if row else b.x2) for b in boxes)
    band_size = max((b.h if row else b.w) for b in boxes)
    out = []
    for g in ctx.pattern.repeat_groups:
        if g.id == group.id or g.axis != group.axis:
            continue
        mapped: list[tuple[int, list[etree._Element]]] = []
        ok = True
        for c in cells_elements(ctx.slide, g):
            b = cell_bbox(c)
            if b is None:
                ok = False
                break
            center = (b.x + b.w / 2) if row else (b.y + b.h / 2)
            i = int((center - starts[0]) // pitch)
            lo, hi = (b.y, b.y2) if row else (b.x, b.x2)
            gap = max(0, max(lo, band_lo) - min(hi, band_hi))
            # a companion row may sit far below its labels (VK Tech stat row: titles at 25% height, figures at 58%):
            # alignment along the axis is what matters, the cross-axis distance only has to stay within the content area
            if i < 0 or i >= len(starts) or gap > max(1.2 * max(band_size, hi - lo), 0.4 * (ctx.H if row else ctx.W)):
                ok = False
                break
            mapped.append((i, c))
        idx = [i for i, _ in mapped]
        if ok and mapped and len(set(idx)) == len(idx):
            out.append((g, mapped))
    return out


def _cell_entries(ctx: _SlideCtx, group: RepeatGroup, cell: list[etree._Element], extra: list[etree._Element]) -> list[tuple[etree._Element, Optional[Slot], Optional[SlotRole]]]:
    """(element, slot, role) for every member of a cell; a duplicated shape takes the role of the shape it was copied
    from (never a role by position: the cells of a sample list their shapes in different orders)."""
    out = []
    for e in cell:
        sid = ctx.id_of(e) or ""
        src = ctx.origin.get(sid, sid)
        if src in ctx.role_of:
            out.append((e, ctx.slot_of.get(src), ctx.role_of[src]))
        else:
            out.append((e, None, None))
    for e in extra:
        sid = ctx.id_of(e) or ""
        src = ctx.origin.get(sid, sid)
        out.append((e, ctx.slot_of.get(src), ctx.role_of.get(src)))
    return out


def _reading_order(entries, under=None):
    """Entries by rows then columns; with `under` (a title element) the slots lined up below it come first."""
    tb = element_bbox(under) if under is not None else None

    def key(x):
        b = element_bbox(x[0])
        if not b:
            return (0, 0, 0)
        aligned = 0
        if tb:
            cx = b[0] + b[2] / 2
            aligned = 0 if tb[0] - b[2] * 0.2 <= cx <= tb[0] + tb[2] + b[2] * 0.2 else 1
        return (aligned, round(b[1] / 45720), b[0])

    return sorted(entries, key=key)


def _figure_color(ctx: "_SlideCtx", el: etree._Element, slot: Optional[Slot]) -> Optional[str]:
    """A figure set large reads as the point of the card: when the card's own text colour is a muted grey (VK Tech
    captions) or weak on its ground, it takes the first accent of the template that stands out there."""
    from verstka.schemas.common import contrast_ratio, hex_to_rgb

    b = element_bbox(el)
    ground = _ground_hex(ctx, Bbox(x=b[0], y=b[1], w=b[2], h=b[3])) if b else None
    own = slot.style.color_hex if slot is not None and slot.style and slot.style.color_hex else None
    if own and ground:
        r, g, bl = hex_to_rgb(own)
        grey = max(r, g, bl) - min(r, g, bl) < 40 and 70 < (r + g + bl) / 3 < 200
        if not grey and contrast_ratio(own, ground) >= 4.5:
            return None  # the template's own colour for such text is strong enough
    for a in ctx.manifest.tokens.accents()[:3]:
        if ground is None or contrast_ratio(a, ground) >= 3.0:
            return a
    return None


def _figure_size(ctx: "_SlideCtx", base: float) -> float:
    """A figure leading a card is set well above its label — about 2.4× (12 pt → 29 pt), never past the template's
    display size; the box it lands in may still cut it down to fit one line."""
    return min(max(base * 2.4, base + 12), ctx.typo.size_for("display", base * 3))


def _holds_line(el: etree._Element, size_pt: float) -> bool:
    """Whether a text box is tall enough for one line set at `size_pt`."""
    b = element_bbox(el)
    if not b or b[3] <= 0:
        return False
    ins = _body_insets(el)
    return (b[3] - ins[1] - ins[3]) / EMU_PER_PT >= size_pt * 1.05


def _fill_one_cell(ctx: _SlideCtx, entries, chunk: list[SlideItem], cell_idx: int, *, include_number: bool, use_ordinals: bool, per_cell: int) -> None:
    """Write one item (or `per_cell` consecutive figures) into the text slots of a cell."""
    text = _reading_order([x for x in entries if x[2] in CELL_TEXT_ROLES and etree.QName(x[0]).localname == "sp"])
    nums = [x for x in text if x[2] == SlotRole.number]
    labels = [x for x in text if x[2] == SlotRole.number_label]
    titles = [x for x in text if x[2] == SlotRole.card_title]
    bodies = [x for x in text if x[2] in _BODY_ROLES]
    if not nums:
        bodies = _reading_order(labels + bodies)
        labels = []
    if not titles and len(bodies) >= 2:
        # two text boxes of different size in one cell: the bigger one is the card title (WorkSpace p13/p16 rows)
        sizes = [(s.style.size_pt if s is not None and s.style.size_pt else 0.0) for _, s, _ in bodies]
        big = max(range(len(bodies)), key=lambda i: sizes[i])
        if sizes[big] > 1.15 * min(sizes):
            titles = [bodies.pop(big)]
    item = chunk[0] if chunk else None

    def write(slot_list, paras, **kw):
        for k, (e, s, _) in enumerate(slot_list):
            if k == 0 and paras:
                ctx.fill_el(e, s, paras, **kw)
            else:
                clear_text(e)

    if item is None:
        write(text, [])
        return
    if per_cell > 1:
        # a cell with several figures (Education p46): consecutive items into consecutive number/label slots
        label_slots = labels or bodies
        for k, (e, s, _) in enumerate(nums):
            if k < len(chunk):
                ctx.fill_el(e, s, [ParagraphSpec(chunk[k].number or chunk[k].title)], min_ratio=0.35)
            else:
                clear_text(e)
        for k, (e, s, _) in enumerate(label_slots):
            paras = _label_paragraphs(chunk[k]) if k < len(chunk) else []
            if paras:
                ctx.fill_el(e, s, paras)
            else:
                clear_text(e)
        for k, it in enumerate(chunk):
            if k >= len(label_slots) and (it.text or it.bullets):
                ctx.warnings.append(f"label of «{it.number or it.title}» had no slot in cell {cell_idx + 1}")
        write(titles, [])
        if labels:
            write(bodies, [])
        return
    if len(text) == 1:
        e, s, _ = text[0]
        paras = _item_single_paragraphs(item, include_number=include_number)
        if item.number and include_number and paras and paras[0].text == item.number:
            base = s.style.size_pt if s is not None and s.style.size_pt else ctx.typo.size_for("body", 14.0)
            paras[0] = ParagraphSpec(item.number, bullet=False, bold=True, size_pt=_figure_size(ctx, base))
        ctx.fill_el(e, s, paras)
        return
    if item.number and not include_number:
        # the figure goes to a standalone number slot: the cell shows the label only
        item = SlideItem(title=item.text or item.title, icon_hint=item.icon_hint, bullets=item.bullets)
    head: Optional[ParagraphSpec] = None
    if item.number:
        if nums:
            write(nums, [ParagraphSpec(item.number)], min_ratio=0.35)
        else:
            # no figure slot in the card: the figure leads it, set a few steps above the card's text
            host = titles[0][1] if titles else (bodies[0][1] if bodies else None)
            host_el = titles[0][0] if titles else (bodies[0][0] if bodies else None)
            base = (host.style.size_pt if host is not None and host.style.size_pt else ctx.typo.size_for("body", 14.0))
            big = _figure_size(ctx, base)
            head = ParagraphSpec(item.number, bullet=False, bold=True, size_pt=big, color_hex=_figure_color(ctx, host_el, host) if host_el is not None else None)
        title = item.title if item.title and item.title != item.number else ""
    else:
        if nums:
            write(nums, [ParagraphSpec(_ordinal(cell_idx, nums[0][1].sample_text if nums[0][1] else None))] if use_ordinals else [], min_ratio=0.35)
        title = item.title
    body_paras = _item_body_paragraphs(item)
    body_slots = _reading_order(labels + bodies, under=titles[0][0] if titles else None)
    if head is not None:
        if titles and (_holds_line(titles[0][0], head.size_pt or 0) or not body_slots):
            write(titles, [head])
            body = ([ParagraphSpec(title, bullet=False, bold=True)] if title else []) + body_paras
        else:
            # the card's title line is too low for a figure (VK Tech cards: an 11 pt caption over a tall text box):
            # the figure leads the text box at its size and the caption line leaves the card
            for e, _, _ in titles:
                if not is_nested(e) and e.getparent() is not None:
                    ctx.remove_el(e)
                else:
                    clear_text(e)
            titles = []
            if body_slots:
                head.color_hex = _figure_color(ctx, body_slots[0][0], body_slots[0][1])
            body = [head] + ([ParagraphSpec(title, bullet=False, bold=True)] if title else []) + body_paras
    elif titles and title:
        if not body_paras:
            # a thesis card: the empty description slot leaves, the heading takes the card's height instead of
            # shrinking into its one-line box
            for e, _, _ in body_slots:
                if not is_nested(e) and e.getparent() is not None:
                    ctx.remove_el(e)
            body_slots = []
        write(titles, [ParagraphSpec(title)])
        body = body_paras
    elif title and body_paras:
        write(titles, [])
        body = [ParagraphSpec(title, bullet=False, bold=True)] + body_paras
    else:
        write(titles, [])
        body = body_paras or ([ParagraphSpec(title)] if title else [])
    if body and not body_slots:
        if titles and not head and title and not (item.text or item.bullets):
            pass  # the title slot already shows the whole item
        else:
            lost = "; ".join(p.text for p in body if p.text)
            ctx.warnings.append(f"text of «{item.number or item.title or lost[:20]}» had no slot in cell {cell_idx + 1}")
    write(body_slots, body)


def _fill_cells(ctx: _SlideCtx, items: list[SlideItem], use_ordinals: bool = True, text_only: bool = False) -> bool:
    """Write items into the pattern's biggest text-bearing repeat group (and the groups aligned with it). Returns True when handled."""
    groups = _text_groups(ctx, exclude_number_only=text_only)
    if not groups or not items:
        return False
    group = max(groups, key=lambda g: len(g.member_shape_ids))
    orig = cells_elements(ctx.slide, group)
    if not orig:
        return False
    orig.sort(key=_cell_sort_key(orig, ctx.H))
    before = [cell_bbox(c) for c in orig]
    companions = _companions(ctx, group, before)
    companion_ids = {g.id for g, _ in companions}
    src_roles = [ctx.role_of.get(sid) for sid in group.member_shape_ids[0]]
    comp_roles = {ctx.role_of.get(sid) for g, _ in companions for cell in g.member_shape_ids for sid in cell}
    numbers_in_cells = SlotRole.number in src_roles or SlotRole.number in comp_roles
    numeric_items = any(i.number for i in items)
    n_numeric = sum(1 for i in items if i.number)
    per_cell = max(1, src_roles.count(SlotRole.number)) if numeric_items else 1
    n_needed = -(-len(items) // per_cell)
    # titles, subtitles and companion cells are never riders: companions are mapped to their primary cell below
    protected = {sl.shape_id for sl in ctx.pattern.slots if sl.role in (SlotRole.title, SlotRole.subtitle)}
    protected |= {sid for g, _ in companions for cell in g.member_shape_ids for sid in cell}
    cells, ctx.next_id = adjust_group(ctx.slide, group, n_needed, ctx.W, ctx.H, ctx.next_id, protected_ids=protected, origin=ctx.origin)
    if not cells:
        return False
    cells.sort(key=_cell_sort_key(cells, ctx.H))  # same order as `before`: items, shifts and companions go by index
    after = [cell_bbox(c) for c in cells]
    # companion cells follow their primary cell: moved and resized with it, removed with it, copied for new cells
    extra: list[list[etree._Element]] = [[] for _ in cells]
    for g, mapped in companions:
        for i, c in mapped:
            if i < len(cells):
                b0, b1 = before[i], after[i]
                if b0 is not None and b1 is not None and (b0.x, b0.y, b0.w, b0.h) != (b1.x, b1.y, b1.w, b1.h):
                    for e in c:
                        if not is_nested(e):
                            transform_element(e, b0, b1, group.axis)
                extra[i].extend(c)
            else:
                for e in c:
                    ctx.remove_el(e)
        src_i = max((i for i, _ in mapped if i < len(before)), default=None)
        if src_i is None or len(cells) <= len(before) or src_i >= len(after):
            continue
        src = dict(mapped)[src_i]
        if any(is_nested(e) for e in src):
            continue
        for i in range(len(before), len(cells)):
            if after[i] is None or after[src_i] is None:
                continue
            anchor = src[-1]
            for e in src:
                ne = copy.deepcopy(e)
                ctx.next_id = renumber_ids(ne, ctx.next_id)
                anchor.addnext(ne)
                anchor = ne
                transform_element(ne, after[src_i], after[i], group.axis)
                for a, b in zip(e.iter(q("p:cNvPr")), ne.iter(q("p:cNvPr"))):
                    ctx.origin[b.get("id")] = ctx.origin.get(a.get("id"), a.get("id"))
                new_id = ctx.id_of(ne)
                if new_id:
                    ctx.els[new_id] = ne
                extra[i].append(ne)
    capacity = len(cells) * per_cell
    overflow: list[SlideItem] = []
    if numeric_items:
        overflow = items[capacity:]
        items = items[:capacity]
    else:
        items = _merge_items(items, len(cells))
        if len(items) > len(cells):
            ctx.warnings.append(f"{len(items) - len(cells)} items dropped: group holds {len(cells)}")
            items = items[: len(cells)]
    standalone_numbers = sorted([s for s in ctx.slots(SlotRole.number) if s.group_id != group.id and s.group_id not in companion_ids], key=lambda s: -(s.style.size_pt or 0.0))
    include_number = numbers_in_cells or not standalone_numbers or len(standalone_numbers) < n_numeric
    seen: set[str] = set()
    cell_bodies: list[tuple[etree._Element, Optional[Slot]]] = []
    cell_numbers: list[tuple[etree._Element, Optional[Slot]]] = []
    cell_titles: list[tuple[etree._Element, Optional[Slot]]] = []
    cell_heads: list[etree._Element] = []
    for cell_idx, cell in enumerate(cells):
        chunk = items[cell_idx * per_cell : (cell_idx + 1) * per_cell]
        entries = _cell_entries(ctx, group, cell, extra[cell_idx])
        # a shape listed in two cells (overlapping detection) belongs to the first one
        entries = [x for x in entries if (ctx.id_of(x[0]) or "") not in seen]
        seen.update(ctx.id_of(x[0]) or "" for x in entries)
        _fill_one_cell(ctx, entries, chunk, cell_idx, include_number=include_number, use_ordinals=use_ordinals, per_cell=per_cell)
        item = chunk[0] if chunk else None
        for e, slot, role in entries:
            if role == SlotRole.icon and item is not None and item.icon_hint and etree.QName(e).localname == "pic":
                picked = pick_icon(ctx.manifest, ctx.ws, item.icon_hint, exclude=ctx.used_assets)
                if picked:
                    aid, path = picked
                    if replace_picture(ctx.slide, e, path):
                        ctx.used_assets.add(aid)
                        sid = ctx.id_of(e)
                        if sid:
                            ctx.filled.add(sid)
        for e, _, _ in entries:
            for nv in e.iter(q("p:cNvPr")):
                ctx.cell_member_ids.add(nv.get("id"))
        for e, slot, role in entries:
            if (ctx.id_of(e) or "") in ctx.filled and etree.QName(e).localname == "sp" and e.getparent() is not None:
                if role == SlotRole.number:
                    cell_numbers.append((e, slot))
                elif role == SlotRole.card_title:
                    if item is not None and item.number and shape_text(e).strip() == item.number:
                        cell_heads.append(e)  # an emphasised figure: sized with the other figures, not with titles
                    else:
                        cell_titles.append((e, slot))
                elif role in _BODY_ROLES or role == SlotRole.number_label:
                    cell_bodies.append((e, slot))
                    if item is not None and item.number and shape_text(e).startswith(item.number):
                        cell_heads.append(e)  # the figure leads the text box: one size with the other figures
    # one size for the same role in every card: the smallest that fits all of them, grown along the scale when all have room
    _harmonize(ctx, cell_bodies, cap=_text_cap(ctx))
    # thesis cards (a heading and nothing under it) are read like running text: they may grow the same way
    _harmonize(ctx, cell_titles, cap=None if cell_bodies else _text_cap(ctx))
    _keep_title_above_body(ctx, cell_titles, cell_bodies)
    _harmonize(ctx, cell_numbers, cap=None, min_ratio=0.35)
    _harmonize_first_lines(cell_heads)
    if numeric_items and not include_number:
        # figures live outside the group (big_number patterns): the largest standalone slot takes the first one
        for slot, item in zip(standalone_numbers, [i for i in items if i.number]):
            ctx.fill_slot(slot, [ParagraphSpec(item.number)], min_ratio=0.35)
    if overflow:
        _fill_standalone_items(ctx, overflow)
    return True


def _running_size(e: etree._Element) -> Optional[float]:
    """The size carrying most of a shape's characters."""
    from collections import Counter

    chars: Counter = Counter()
    for r in e.iter(q("a:r")):
        rpr, t = r.find(q("a:rPr")), r.find(q("a:t"))
        if rpr is not None and rpr.get("sz"):
            chars[int(rpr.get("sz")) / 100] += len(t.text or "") if t is not None else 0
    return chars.most_common(1)[0][0] if chars else None


def _keep_title_above_body(ctx: _SlideCtx, titles: list[tuple[etree._Element, Optional[Slot]]], bodies: list[tuple[etree._Element, Optional[Slot]]]) -> None:
    """Card text that grew into its card must not outgrow the card's heading («Неделя 1» at 11 pt over a 16 pt
    text): the headings follow at the template's own ratio of heading to text, as far as their boxes allow."""
    if not titles or not bodies:
        return
    t_sizes = [s for s in (_running_size(e) for e, _ in titles) if s]
    b_sizes = [s for s in (_running_size(e) for e, _ in bodies) if s]
    if not t_sizes or not b_sizes:
        return
    t, b = min(t_sizes), min(b_sizes)
    ts = next((s.style.size_pt for _, s in titles if s is not None and s.style.size_pt), None)
    bs = next((s.style.size_pt for _, s in bodies if s is not None and s.style.size_pt), None)
    ratio = max(1.0, ts / bs) if ts and bs else 1.15
    want = _snap_down(b * ratio, ctx.grow_scale)
    if want <= t:
        return
    fitted = []
    for e, slot in titles:
        box = element_bbox(e)
        if not box or box[2] <= 0 or box[3] <= 0:
            return
        style = slot.style if slot is not None else None
        family = style.font_family if style and style.font_family else ctx.typo.primary_family
        texts = shape_text(e).split("\n")
        res = fit_size(texts, Bbox(x=box[0], y=box[1], w=box[2], h=box[3]), family, want, bool(style and style.bold), ctx.grow_scale, insets_emu=_body_insets(e), line_spacing=ctx.typo.line_height, min_ratio=t / want)
        fitted.append(res.size_pt if res.fits else t)
    common = min(fitted)
    if common > t:
        for e, _ in titles:
            _scale_text(e, common, ctx.grow_scale)


def _text_cap(ctx: _SlideCtx) -> float:
    """How far running text may grow when it has room: to a comfortable reading size on a slide (16 pt, or 1.4× the
    template's body when that is larger), always clearly below the slide heading. VK Tech's body is 9 pt: three
    bullets in a panel two thirds of the slide high at 9–12 pt read as a mistake."""
    body = ctx.typo.size_for("body", 14.0)
    h1 = ctx.typo.size_for("h1", body * 2)
    return max(body, min(max(body * 1.4, 16.0), h1 * 0.8))


def _harmonize(ctx: _SlideCtx, pairs: list[tuple[etree._Element, Optional[Slot]]], cap: Optional[float], min_ratio: float = 0.6) -> None:
    if len(pairs) < 2 and cap is None:
        return
    sizes: list[float] = []
    for e, slot in pairs:
        b = element_bbox(e)
        texts = [t for t in shape_text(e).split("\n")]
        if not b or b[2] <= 0 or b[3] <= 0 or not any(t.strip() for t in texts):
            continue
        style = slot.style if slot is not None else None
        base = style.size_pt if style and style.size_pt else ctx.typo.size_for("body", 14.0)
        family = style.font_family if style and style.font_family else ctx.typo.primary_family
        bold = bool(style and style.bold)
        box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
        ins = _body_insets(e)
        res = fit_size(texts, box, family, base, bold, ctx.grow_scale, insets_emu=ins, line_spacing=ctx.typo.line_height, min_ratio=min_ratio)
        size = res.size_pt
        if cap is not None and res.fits:
            size = grow_size(texts, box, family, size, bold, ctx.grow_scale, cap, insets_emu=ins, line_spacing=ctx.typo.line_height)
        sizes.append(size)
    if sizes:
        common = min(sizes)
        for e, _ in pairs:
            _scale_text(e, common, ctx.grow_scale)


def _snap_down(size: float, scale: list[float]) -> float:
    """The largest size the template uses that is not above `size` (an emphasised figure stays on the template's
    scale: 43.2 pt → 40 pt), or `size` itself below the scale."""
    below = [s for s in scale if s <= size + 0.05]
    return max(below) if below else round(size, 1)


def _scale_text(e: etree._Element, size: float, snap: Optional[list[float]] = None) -> None:
    """Set the running text of a shape to `size`; emphasised lines (a figure set larger) keep their proportion."""
    runs = [r for r in e.iter(q("a:rPr")) if r.get("sz")]
    if not runs:
        set_text_size(e, size)
        return
    from collections import Counter

    # the running text is the size that carries the most characters — a figure of four characters over a label of
    # thirty is the emphasis, not the text (counting runs would pick the figure and flatten it to the label size)
    chars: Counter = Counter()
    for r in e.iter(q("a:r")):
        rpr, t = r.find(q("a:rPr")), r.find(q("a:t"))
        if rpr is not None and rpr.get("sz"):
            chars[int(rpr.get("sz"))] += len(t.text or "") if t is not None else 0
    modal = chars.most_common(1)[0][0] if chars else Counter(int(r.get("sz")) for r in runs).most_common(1)[0][0]
    for r in list(e.iter(q("a:rPr"))) + list(e.iter(q("a:endParaRPr"))):
        old = int(r.get("sz")) if r.get("sz") else modal
        if old > modal:
            # an emphasised figure shrinks with its text but never grows past the size that fitted its box
            new = min(old, size * 100 * old / modal)
            if snap and new < old:
                new = _snap_down(new / 100, snap) * 100
            r.set("sz", str(int(round(new))))
        else:
            r.set("sz", str(int(round(size * 100))))


def _harmonize_first_lines(els: list[etree._Element]) -> None:
    """Figures leading their cards share one size: the smallest one that fitted."""
    firsts = []
    for e in els:
        p0 = e.find(q("p:txBody") + "/" + q("a:p"))
        sizes = [int(r.get("sz")) for r in p0.iter(q("a:rPr")) if r.get("sz")] if p0 is not None else []
        if sizes:
            firsts.append((p0, min(sizes)))
    if len(firsts) < 2:
        return
    common = min(sz for _, sz in firsts)
    for p0, _ in firsts:
        for r in list(p0.iter(q("a:rPr"))) + list(p0.iter(q("a:endParaRPr"))):
            r.set("sz", str(common))


def _nearest_label(slot: Slot, labels: list[Slot]) -> Optional[Slot]:
    if not labels:
        return None
    cx, cy = slot.bbox.x + slot.bbox.w / 2, slot.bbox.y + slot.bbox.h / 2

    def dist(lab: Slot) -> float:
        lx, ly = lab.bbox.x + lab.bbox.w / 2, lab.bbox.y + lab.bbox.h / 2
        return abs(lx - cx) + abs(ly - cy) + (0.5 if ly < cy - slot.bbox.h / 2 else 0.0)  # labels sit below or beside a figure

    return min(labels, key=dist)


def _fill_standalone_items(ctx: _SlideCtx, items: list[SlideItem]) -> None:
    """Items into ungrouped slots: numbers/labels first, then card titles/bodies, rest into a list slot."""
    numbers = sorted([s for s in ctx.slots(SlotRole.number) if s.shape_id not in ctx.cell_member_ids], key=lambda s: -(s.style.size_pt or 0.0))
    labels = [s for s in ctx.slots(SlotRole.number_label) if s.shape_id not in ctx.cell_member_ids]
    if numbers and any(i.number for i in items):
        num_items = [i for i in items if i.number]
        for slot, item in zip(numbers, num_items):
            ctx.fill_slot(slot, [ParagraphSpec(item.number)], min_ratio=0.35)
            lab = _nearest_label(slot, labels)
            paras = _label_paragraphs(item)
            if lab is not None and paras:
                ctx.fill_slot(lab, paras)
                labels.remove(lab)
            elif paras:
                ctx.warnings.append(f"label of «{item.number}» had no slot")
        placed = min(len(numbers), len(num_items))
        if len(num_items) > placed:
            ctx.warnings.append(f"{len(num_items) - placed} numbers had no slot: {', '.join(i.number or '' for i in num_items[placed:])}")
        items = [i for i in items if not i.number] + num_items[placed:]
    elif numbers and len(numbers) >= len(items) >= 2:
        # ordinal markers next to standalone card slots (agenda / process patterns)
        for idx, (slot, item) in enumerate(zip(sorted(numbers, key=lambda s: (s.bbox.y, s.bbox.x)), items)):
            ctx.fill_slot(slot, [ParagraphSpec(_ordinal(idx, slot.sample_text))], min_ratio=0.35)
    titles = ctx.slots(SlotRole.card_title)
    bodies = ctx.slots(SlotRole.card_body)
    n = max(len(titles), len(bodies))
    if n and items:
        items2 = _merge_items(items, n)
        for slot, item in zip(titles, items2):
            ctx.fill_slot(slot, [ParagraphSpec(item.title or item.number or "")] if (item.title or item.number) else [])
        for k, (slot, item) in enumerate(zip(bodies, items2)):
            if k < len(titles):
                ctx.fill_slot(slot, _item_body_paragraphs(item) or ([ParagraphSpec(item.title)] if item.title else []))
            else:
                ctx.fill_slot(slot, _item_single_paragraphs(item))  # no title slot for this body: the title goes bold on top
        if len(titles) > len(bodies):
            for item in items2[len(bodies) : len(titles)]:
                if item.text or item.bullets:
                    ctx.warnings.append(f"text of «{item.title}» had no slot")
        items = items2[n:]
    if items:
        target = ctx.slots(SlotRole.bullet_list, SlotRole.body)
        if target:
            big = ctx.largest(target)
            paras: list[ParagraphSpec] = []
            for i in items:
                head = i.number or i.title
                if i.bullets:
                    if head:
                        paras.append(ParagraphSpec(head, bullet=False, bold=True))
                    if i.text:
                        paras.append(ParagraphSpec(i.text, bullet=False))
                    paras += [ParagraphSpec(b, bullet=True) for b in i.bullets]
                else:
                    paras.append(ParagraphSpec(f"{head}: {i.text}" if i.text and head else (i.text or head), bullet=True))
            ctx.fill_slot(big, paras)
        else:
            ctx.warnings.append(f"{len(items)} items had no slot: {', '.join((i.number or i.title)[:30] for i in items)}")


# ---------------------------------------------------------------------------- native objects


def _clamp_to_safe(box: Bbox, ctx: _SlideCtx) -> Bbox:
    safe = ctx.manifest.tokens.spacing.safe_area
    x = max(box.x, int(safe.x * ctx.W))
    y = max(box.y, int(safe.y * ctx.H))
    x2 = min(box.x2, int(safe.x2 * ctx.W))
    y2 = min(box.y2, int(safe.y2 * ctx.H))
    if x2 - x >= 0.2 * ctx.W and y2 - y >= 0.15 * ctx.H:
        return Bbox(x=x, y=y, w=x2 - x, h=y2 - y)
    return box


_SCHEME_BASE = {"bg1": "FFFFFF", "lt1": "FFFFFF", "bg2": "EEEEEE", "lt2": "EEEEEE", "tx1": "000000", "dk1": "000000", "tx2": "333333", "dk2": "333333"}


def _fill_hex(el: etree._Element) -> Optional[str]:
    """The solid fill of a shape: an sRGB value, or the light/dark base colours of the theme scheme (bg1, tx1…)."""
    spPr = el.find(q("p:spPr"))
    sf = spPr.find(q("a:solidFill")) if spPr is not None else None
    if sf is None:
        if spPr is not None and spPr.find(q("a:noFill")) is not None:
            return None
        ref = el.find(q("p:style") + "/" + q("a:fillRef"))  # the fill of the shape style («Прямоугольник» by default)
        if ref is None or (ref.get("idx") or "0") == "0":
            return None
        sf = ref
    rgb = sf.find(q("a:srgbClr"))
    if rgb is not None and rgb.get("val"):
        return rgb.get("val").upper()
    sch = sf.find(q("a:schemeClr"))
    return _SCHEME_BASE.get(sch.get("val")) if sch is not None else None


def _ground_hex(ctx: _SlideCtx, box: Bbox) -> Optional[str]:
    """What lies under a box: the smallest painted shape of the slide holding it, else what the layout or master
    paints there (a band, a panel, a gradient picture — measured where the box is) when it differs from the sample
    slide's ground, else the sample slide's own ground."""
    best = None
    for el in ctx.slide._element.cSld.find(q("p:spTree")):
        if etree.QName(el).localname != "sp" or not _has_fill(el):
            continue
        b = element_bbox(el)
        if not b:
            continue
        ob = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
        if ob.intersection(box) >= 0.8 * box.area and (best is None or ob.area < best[0]):
            hx = _fill_hex(el)
            if hx:
                best = (ob.area, hx)
    if best:
        return best[1]
    bg = next((b.hex for b in ctx.manifest.tokens.backgrounds if ctx.pattern.source_slide in b.slides and b.hex), None)
    bg = bg or ctx.manifest.tokens.color_for("background.dark" if ctx.pattern.family.value == "dark" else "background.light")
    try:
        under = ground_under(ctx.slide, box)
    except Exception:  # noqa: BLE001 - layers are advice
        under = None
    if under is not None and (not bg or contrast_ratio(under[0], bg) >= 1.2):
        return under[0]
    return bg


def _mixed_ground(ctx: _SlideCtx, box: Bbox) -> bool:
    """The box straddles an edge of what the layout or master paints: no single ground colour to judge it on."""
    from verstka.rendering.layers import ground_is_mixed

    try:
        return ground_is_mixed(ctx.slide, box)
    except Exception:  # noqa: BLE001 - layers are advice: unsure means mixed
        return True


def _ground_panel(ctx: _SlideCtx, box: Bbox) -> Optional[Bbox]:
    """The painted rectangular panel (the slide's own, or its layout's or master's) a heading box stands on when it is
    not the whole slide: the topmost opaque rectangle or picture that holds most of the box, covers 15–85 % of the
    slide and differs from the slide's ground. None when the heading stands on the slide's own ground."""
    from verstka.rendering.layers import drawn_layers

    try:
        layers = drawn_layers(ctx.slide)
    except Exception:  # noqa: BLE001 - layers are advice
        return None
    W, H = ctx.W, ctx.H
    bg = next((b.hex for b in ctx.manifest.tokens.backgrounds if ctx.pattern.source_slide in b.slides and b.hex), None)
    for lay in reversed(layers):
        if lay.placeholder or not lay.paints or lay.kind not in ("sp", "pic") or lay.custom_geom or lay.has_text:
            continue
        if lay.kind == "sp":
            prst = lay.el.find(q("p:spPr") + "/" + q("a:prstGeom"))
            if prst is None or (prst.get("prst") or "rect") not in ("rect", "snip1Rect", "round1Rect", "round2SameRect"):
                continue
        if lay.box.intersection(box) < 0.8 * max(box.area, 1):
            continue
        if not 0.15 <= lay.cover < 0.85:
            return None  # the heading's ground is the slide (a full-bleed layer) or a pill the label logic handles
        if not lay.fill_hex or (bg and contrast_ratio(lay.fill_hex, bg) < 1.2):
            return None
        return Bbox(x=max(lay.box.x, 0), y=max(lay.box.y, 0), w=min(lay.box.x2, W) - max(lay.box.x, 0), h=min(lay.box.y2, H) - max(lay.box.y, 0))
    return None


def _readable_on(ground: Optional[str], preferred: list[Optional[str]], manifest: TemplateManifest) -> Optional[str]:
    if not ground:
        return next((c for c in preferred if c), None)
    from verstka.schemas.common import contrast_ratio

    cands = [c for c in preferred if c] + [manifest.tokens.color_for(r) for r in ("text.secondary", "background.light", "background.dark")] + ["FFFFFF", "000000"]
    cands = [c for c in cands if c]
    for c in cands:
        if contrast_ratio(c, ground) >= 4.5:
            return c
    return max(cands, key=lambda c: contrast_ratio(c, ground))


def _grow_into_free_area(box: Bbox, ctx: _SlideCtx, oslide: Optional[OutlineSlide] = None) -> Bbox:
    """A sample table/chart area far smaller than the free content area (Education keeps a 25%-wide sample table)
    grows inside the safe area below its top edge — over text slots that will stay empty, never over anything else."""
    safe = ctx.manifest.tokens.spacing.safe_area
    if box.area >= 0.3 * ctx.W * ctx.H:
        return box
    keeps_text = bool(oslide and (oslide.content.paragraphs or oslide.content.bullets))
    tree = ctx.slide._element.cSld.find(q("p:spTree"))
    blockers = []
    for el in tree:
        if etree.QName(el).localname not in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            continue
        b = element_bbox(el)
        sid = ctx.id_of(el) or ""
        if not b or b[2] * b[3] >= 0.6 * ctx.W * ctx.H or sid in ctx.pattern.chrome_shape_ids:
            continue
        slot = ctx.slot_of.get(sid)
        if slot is not None and slot.role in (SlotRole.title, SlotRole.subtitle):
            if b[1] + b[3] <= box.y + int(0.01 * ctx.H):
                continue  # the heading above
        elif slot is not None and slot.role in TEXT_ROLES and sid not in ctx.filled and not has_visible_style(el) and not keeps_text:
            continue  # an empty text slot: the cleanup removes it
        blockers.append(Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    x1, x2 = int(safe.x * ctx.W), int(safe.x2 * ctx.W)
    y2 = int(safe.y2 * ctx.H)
    gap_x, gap_y = int(0.02 * ctx.W), int(0.02 * ctx.H)
    for o in blockers:
        if o.y2 <= box.y or o.y >= y2:
            continue
        if o.x2 <= box.x + 1:
            x1 = max(x1, o.x2 + gap_x)
        elif o.x >= box.x2 - 1:
            x2 = min(x2, o.x - gap_x)
    for o in blockers:
        if o.x < x2 and o.x2 > x1 and o.y >= box.y2 - 1:
            y2 = min(y2, o.y - gap_y)
    grown = Bbox(x=min(box.x, x1), y=box.y, w=max(box.x2, x2) - min(box.x, x1), h=max(box.h, y2 - box.y))
    return grown if grown.area > box.area else box


def _place_native_object(ctx: _SlideCtx, oslide: OutlineSlide) -> None:
    c = oslide.content
    manifest = ctx.manifest
    # sample charts/tables/images on the slide are replaced by the native object
    candidates = ctx.slots(SlotRole.image)
    frames = [ctx.slot_of[sid] for sid, el in ctx.els.items() if etree.QName(el).localname == "graphicFrame" and sid in ctx.slot_of and sid not in ctx.removed]
    candidates = list({s.shape_id: s for s in candidates + frames}.values())
    box: Optional[Bbox] = None
    holders = [s for s in ctx.slots(SlotRole.body, SlotRole.card_body) if s.container and s.bbox.area >= 0.08 and not s.group_id]
    if holders and not any(s.bbox.area >= holders[0].bbox.area for s in candidates if s.bbox.area >= 0.08):
        # the template's empty panel is the card of the chart/table: keep it, draw the object inside its padding
        frame = max(holders, key=lambda s: s.bbox.area)
        el = ctx.els.get(frame.shape_id)
        b = element_bbox(el) if el is not None else None
        if b and b[2] > 0 and b[3] > 0:
            pad = container_inset(b[2], b[3])
            box = Bbox(x=b[0] + pad, y=b[1] + pad, w=b[2] - 2 * pad, h=b[3] - 2 * pad)
            ctx.filled.add(frame.shape_id)
            candidates = [s for s in candidates if s.bbox.area < 0.08]
            for s in candidates:
                ctx.remove_slot_shape(s)
            candidates = []
    big_candidates = [s for s in candidates if s.bbox.area >= 0.08]
    if box is None and ctx.pattern.kind in (PatternKind.chart, PatternKind.table) and not big_candidates:
        # the sample chart/table is drawn with shapes: clear everything except title/subtitle and use that area
        keep = {s.shape_id for s in ctx.pattern.slots if s.role in (SlotRole.title, SlotRole.subtitle)}
        union: Optional[Bbox] = None
        for sid, el in list(ctx.els.items()):
            if sid in keep or sid in ctx.removed or is_nested(el):
                continue
            b = element_bbox(el)
            if b and b[2] > 0 and b[3] > 0 and b[1] > 0.12 * ctx.H:
                bb = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
                union = bb if union is None else union.union(bb)
            remove_element(el)
            ctx.removed.add(sid)
        candidates = []
        if union is not None and union.area >= 0.1 * ctx.W * ctx.H:
            box = union
    if candidates:
        best = max(candidates, key=lambda s: s.bbox.area)
        el = ctx.els.get(best.shape_id)
        b = element_bbox(el) if el is not None else None
        box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3]) if b and b[2] > 0 else best.bbox.to_emu(ctx.W, ctx.H)
        tbl = el.find(".//" + q("a:tbl")) if el is not None and etree.QName(el).localname == "graphicFrame" else None
        if tbl is not None:
            # Google-Slides exports keep a dummy frame extent; the grid columns and rows carry the real size
            grid_w = sum(int(g.get("w") or 0) for g in tbl.iter(q("a:gridCol")))
            rows_h = sum(int(r.get("h") or 0) for r in tbl.findall(q("a:tr")))
            box = Bbox(x=box.x, y=box.y, w=max(box.w, grid_w), h=max(box.h, rows_h))
        for s in candidates:
            ctx.remove_slot_shape(s)
    if box is None or box.area < 0.08 * ctx.W * ctx.H:
        text_slots = ctx.slots(SlotRole.body, SlotRole.bullet_list)
        if text_slots:
            big = ctx.largest(text_slots)
            el = ctx.els.get(big.shape_id)
            b = element_bbox(el) if el is not None else None
            if b and b[2] * b[3] > (box.area if box else 0):
                box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
                ctx.remove_slot_shape(big)
    if box is None:
        safe = manifest.tokens.spacing.safe_area
        title_bottom = max((s.bbox.y2 for s in ctx.pattern.slots if s.role == SlotRole.title), default=safe.y)
        box = Bbox(x=int(safe.x * ctx.W), y=int((title_bottom + 0.03) * ctx.H), w=int(safe.w * ctx.W), h=int((safe.y2 - title_bottom - 0.05) * ctx.H))
    box = _clamp_to_safe(box, ctx) if not ctx.filled & {s.shape_id for s in holders} else box
    if not ctx.filled & {s.shape_id for s in holders}:
        box = _grow_into_free_area(box, ctx, oslide)
    # pictures the object would cover (sample screenshots, chart images marked as decoration) are removed
    for sid, el in list(ctx.els.items()):
        if sid in ctx.removed or etree.QName(el).localname != "pic" or is_nested(el):
            continue
        b = element_bbox(el)
        if not b or b[2] <= 0 or b[3] <= 0:
            continue
        pic_box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
        if box.intersection(pic_box) > 0.3 * pic_box.area:
            remove_element(el)
            ctx.removed.add(sid)
    ground = _ground_hex(ctx, box)
    text_hex = _readable_on(ground, [manifest.tokens.color_for("text.primary"), manifest.components.table_style.body_text_hex], manifest)
    try:
        if c.chart is not None:
            add_chart(ctx.slide, box, c.chart, ctx.outline, manifest.components.chart_style, ctx.typo, text_hex=text_hex, neutral_hex=next((t.hex for t in manifest.tokens.colors if t.role and t.role.startswith("neutral")), None))
        elif c.table is not None:
            style = manifest.components.table_style
            if text_hex and text_hex != style.body_text_hex:
                # the template's table text is dark; on a dark ground (LCT purple) it must turn light, bands go
                style = style.model_copy(update={"body_text_hex": text_hex, "band_fill_hex": None})
            add_table(ctx.slide, box, c.table, style, ctx.typo)
    except Exception as e:  # noqa: BLE001
        if c.chart is not None:
            # the sample's chart area left empty is worse than a composed slide: the renderer rolls this clone back
            # and composes the slide, which shows its figures or its text instead
            raise RuntimeError(f"native chart failed: {str(e)[:120]}") from e
        ctx.warnings.append(f"native object failed: {str(e)[:120]}")


# ---------------------------------------------------------------------------- cleanup helpers


def _remove_empty_cells(ctx: _SlideCtx) -> None:
    """A repeat-group cell with a text slot that received nothing loses all its shapes — and its riders (the icon
    chip next to a list line): no orphan avatars, chips or card backgrounds."""
    protected = {sl.shape_id for sl in ctx.pattern.slots if sl.role in (SlotRole.title, SlotRole.subtitle)} | set(ctx.pattern.chrome_shape_ids)
    for g in ctx.pattern.repeat_groups:
        live = [c for c in cells_elements(ctx.slide, g) if c]
        riders_of: dict[int, list[etree._Element]] = {}
        if live and not any(is_nested(e) for c in live for e in c):
            for c, r in zip(live, cell_riders(ctx.slide, g, live, protected)):
                for e in c:
                    riders_of[id(e)] = r
        for cell in g.member_shape_ids:
            if not any(ctx.role_of.get(sid) in CELL_TEXT_ROLES for sid in cell):
                continue
            if any(sid in ctx.filled for sid in cell):
                continue
            for sid in cell:
                el = ctx.els.get(sid)
                for rider in riders_of.get(id(el), []) if el is not None else []:
                    if rider.getparent() is not None and (ctx.id_of(rider) or "") not in ctx.filled:
                        ctx.remove_el(rider)
            for sid in cell:
                el = ctx.els.get(sid)
                if el is None or sid in ctx.removed or el.getparent() is None:
                    continue
                if is_nested(el):
                    if etree.QName(el).localname == "sp":
                        clear_text(el)
                    continue
                remove_element(el)
                ctx.removed.add(sid)


def _fill_label(ctx: _SlideCtx, title: Slot, oslide: OutlineSlide) -> None:
    """A short tag printed above the heading («Кейс», «Проблема» in a pill) carries the section name; without a
    section it is emptied (its pill then goes as an orphan) so running text never lands in a 1-line tag."""
    labels = [
        s for s in ctx.slots(SlotRole.body, SlotRole.caption, SlotRole.subtitle)
        if s.bbox.y2 <= title.bbox.y + 0.02 and s.bbox.h <= 0.1 and len(s.sample_text or "") <= 25 and s.shape_id != title.shape_id
    ]
    if not labels:
        return
    label = min(labels, key=lambda s: abs(s.bbox.y2 - title.bbox.y))
    tag = (oslide.section or "").strip()
    if tag and tag.lower() != oslide.headline.strip().lower() and len(tag) <= 40:
        ctx.fill_slot(label, [ParagraphSpec(tag)], min_ratio=0.8, widen=True)
    else:
        el = ctx.els.get(label.shape_id)
        if el is not None and has_visible_style(el) and not is_nested(el):
            ctx.remove_el(el)  # the tag is its own pill: an empty pill must not stay
        elif el is not None:
            clear_text(el)
        ctx.removed.add(label.shape_id)


_PCT_RE = re.compile(r"^\s*(\d{1,3}(?:[.,]\d+)?)\s*%\s*$")


def _true_rings(ctx: _SlideCtx) -> None:
    """Ring infographics drawn as pictures (VK Tech 41, 42, 45: a ring filled to 75% under the sample «10%») cannot
    show our figure — «18%» over a three-quarter ring is a lie. A percentage gets a native doughnut in the ring's place
    showing that share; any other figure («5 ч», «31% → 12%») stands without the ring."""
    accent = next(iter(ctx.manifest.tokens.accents()), None) or "0077FF"

    def pieces_of(ring: Bbox) -> list[tuple[str, etree._Element]]:
        """The ring and the pictures it is drawn with (a gradient arc over a grey track, a highlight)."""
        out = []
        for pid, pic in list(ctx.els.items()):
            if etree.QName(pic).localname != "pic" or pid in ctx.filled or pid in ctx.removed or pic.getparent() is None or is_nested(pic):
                continue
            pb = element_bbox(pic)
            if pb and pb[2] * pb[3] > 0 and Bbox(x=pb[0], y=pb[1], w=pb[2], h=pb[3]).intersection(ring) >= 0.6 * pb[2] * pb[3]:
                out.append((pid, pic))
        return out

    def ring_at(x: float, y: float, min_area: float) -> Optional[tuple[str, etree._Element, Bbox]]:
        for pid, pic in list(ctx.els.items()):
            if etree.QName(pic).localname != "pic" or pid in ctx.filled or pid in ctx.removed or pic.getparent() is None or is_nested(pic):
                continue
            pb = element_bbox(pic)
            if not pb or pb[3] <= 0 or not (pb[0] <= x <= pb[0] + pb[2] and pb[1] <= y <= pb[1] + pb[3]):
                continue
            if 0.8 <= pb[2] / pb[3] <= 1.25 and pb[2] * pb[3] >= min_area:
                return pid, pic, Bbox(x=pb[0], y=pb[1], w=pb[2], h=pb[3])
        return None

    # rings drawn for sample figures that are no longer inside them — the cell left, or the row reflowed and moved
    # the figure away (two figures on a three-ring slide): such a ring leaves with all its pieces
    for slot in ctx.pattern.slots:
        if slot.role != SlotRole.number:
            continue
        b = slot.bbox.to_emu(ctx.W, ctx.H)
        hit = ring_at(b.x + b.w / 2, b.y + b.h / 2, 2 * b.w * b.h)
        if hit is None:
            continue
        ring = hit[2]
        el = ctx.els.get(slot.shape_id)
        cur = element_bbox(el) if el is not None and el.getparent() is not None and slot.shape_id in ctx.filled else None
        if cur and ring.x <= cur[0] + cur[2] / 2 <= ring.x2 and ring.y <= cur[1] + cur[3] / 2 <= ring.y2:
            continue  # its figure is still there: handled below
        for pid, pic in pieces_of(ring):
            remove_element(pic)
            ctx.removed.add(pid)
    for sid in list(ctx.filled):
        el = ctx.els.get(sid)
        slot = ctx.slot_of.get(ctx.origin.get(sid, sid))
        if el is None or el.getparent() is None or slot is None or slot.role != SlotRole.number:
            continue
        nb = element_bbox(el)
        text = shape_text(el).strip()
        if not nb or not text:
            continue
        cx, cy = nb[0] + nb[2] / 2, nb[1] + nb[3] / 2
        hit = ring_at(cx, cy, 2 * nb[2] * nb[3])  # a square picture around the figure: a ring, not a photo
        if hit is None:
            continue
        _, pic, box = hit
        m = _PCT_RE.match(text)
        if m and float(m.group(1).replace(",", ".")) <= 100:
            gf = add_ring(ctx.slide, box, float(m.group(1).replace(",", ".")), accent, _ground_hex(ctx, box))
            pic.addprevious(gf._element)  # the ring's place in the z-order: under the figure
            # the figure sits in the hole: centred on the ring
            set_element_pos(el, x=int(box.x + box.w / 2 - nb[2] / 2), y=int(box.y + box.h / 2 - nb[3] / 2))
            for p in el.iter(q("a:p")):
                ppr = p.find(q("a:pPr"))
                if ppr is None:
                    ppr = etree.Element(q("a:pPr"))
                    p.insert(0, ppr)
                ppr.set("algn", "ctr")
            ctx.warnings.append(f"ring picture under «{text}» replaced by a doughnut of the real value")
        else:
            ctx.warnings.append(f"ring picture under «{text}» removed: it cannot show this figure")
        for pid, piece in pieces_of(box):
            remove_element(piece)
            ctx.removed.add(pid)


def _remove_qr_codes(ctx: _SlideCtx) -> None:
    """QR codes of the template (a contact card on the closing slide) point to someone else's page."""
    qr_parts = {a.media_part for a in ctx.manifest.assets if a.kind == "qr" and a.media_part}
    if not qr_parts:
        return
    for pic in list(ctx.slide._element.cSld.iter(q("p:pic"))):
        blip = pic.find(".//" + q("a:blip"))
        rid = blip.get(q("r:embed")) if blip is not None else None
        try:
            part = ctx.slide.part.related_part(rid) if rid else None
        except KeyError:
            part = None
        if part is not None and str(part.partname).lstrip("/") in qr_parts:
            ctx.remove_el(pic)


def _remove_orphan_holders(ctx: _SlideCtx) -> None:
    """A plate whose text slots all stayed empty (a label pill without its tag, a note box without its note) goes,
    together with the small icons drawn on it. Plates that hold anything written, a picture or a chart stay."""
    W, H = ctx.W, ctx.H
    chrome = set(ctx.pattern.chrome_shape_ids)
    text_slots = [s for s in ctx.pattern.slots if s.role in TEXT_ROLES]
    tree = ctx.slide._element.cSld.find(q("p:spTree"))
    for sid, el in list(ctx.els.items()):
        if sid in ctx.removed or sid in ctx.filled or sid in chrome or el.getparent() is not tree or etree.QName(el).localname not in ("sp", "pic"):
            continue
        if etree.QName(el).localname == "sp" and (shape_text(el).strip() or not (_has_fill(el) or el.find(q("p:spPr") + "/" + q("a:ln")) is not None)):
            continue
        b = element_bbox(el)
        if not b or b[2] <= 0 or b[3] <= 0 or b[2] * b[3] >= 0.5 * W * H:
            continue
        box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])

        def inside(bx: Bbox) -> bool:
            cx, cy = bx.x + bx.w / 2, bx.y + bx.h / 2
            return box.x <= cx <= box.x2 and box.y <= cy <= box.y2

        hosted = [t for t in text_slots if t.shape_id != sid and inside(t.bbox.to_emu(W, H))]
        if not hosted or any(t.shape_id in ctx.filled for t in hosted):
            continue
        occupants = []
        busy = False
        for other in list(tree):
            if other is el or etree.QName(other).localname not in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
                continue
            ob = element_bbox(other)
            if not ob or not inside(Bbox(x=ob[0], y=ob[1], w=ob[2], h=ob[3])):
                continue
            osid = ctx.id_of(other) or ""
            if etree.QName(other).localname == "graphicFrame" or osid in ctx.filled or shape_text(other).strip():
                busy = True
                break
            if ob[2] * ob[3] <= 0.2 * box.area:
                occupants.append(other)
        if busy:
            continue
        for o in [el] + occupants:
            ctx.remove_el(o)


def _remove_placeholder_boxes(ctx: _SlideCtx) -> None:
    """«Вставить фото / QR» boxes: the text shape plus the empty box drawn behind it (and any picture inside that box)."""
    for slot in ctx.pattern.slots:
        if slot.shape_id in ctx.filled or slot.shape_id in ctx.removed or not slot.sample_text:
            continue
        if not re.search(r"вставить|insert|qr|фото|photo|логотип|logo", slot.sample_text, re.I):
            continue
        el = ctx.els.get(slot.shape_id)
        if el is None or is_nested(el):
            continue
        tb = element_bbox(el)
        remove_element(el)
        ctx.removed.add(slot.shape_id)
        if not tb:
            continue
        text_box = Bbox(x=tb[0], y=tb[1], w=tb[2], h=tb[3])
        tol = int(0.005 * ctx.W)
        holders = []
        for sid, other in ctx.els.items():
            if sid in ctx.removed or sid in ctx.filled or is_nested(other) or etree.QName(other).localname not in ("sp", "pic"):
                continue
            if etree.QName(other).localname == "sp" and shape_text(other).strip():
                continue
            b = element_bbox(other)
            if not b or b[2] <= 0 or b[3] <= 0:
                continue
            ob = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
            if ob.area > 0.6 * ctx.W * ctx.H:
                continue  # a slide-sized background is not a photo box
            if ob.x - tol <= text_box.x and ob.y - tol <= text_box.y and ob.x2 + tol >= text_box.x2 and ob.y2 + tol >= text_box.y2:
                holders.append((ob, sid, other))
        if not holders:
            continue
        holders.sort(key=lambda h: h[0].area)
        box_b, box_sid, box_el = holders[0]
        remove_element(box_el)
        ctx.removed.add(box_sid)
        for ob, sid, other in holders[1:]:
            if etree.QName(other).localname == "pic" and box_b.intersection(ob) >= 0.8 * ob.area:
                remove_element(other)
                ctx.removed.add(sid)


_PAGE_NUMBER_RE = re.compile(r"^\s*(\d{1,3})\s*$")


def renumber_page_chrome(els: dict[str, etree._Element], chrome_ids: list[str], source_slide: int, index: int) -> int:
    """Hand-typed page numbers are chrome copied with the sample: write the slide's own position into them.

    A chrome text whose whole content is a 1–3 digit number close to the sample's position (designers sometimes do
    not count the title) is a page number; the zero padding of the sample is kept, native slide-number fields are
    left alone. Returns how many shapes were rewritten.
    """
    done = 0
    for sid in chrome_ids:
        el = els.get(sid)
        if el is None or etree.QName(el).localname != "sp" or el.find(".//" + q("a:fld")) is not None:
            continue
        m = _PAGE_NUMBER_RE.match(shape_text(el))
        if not m or abs(int(m.group(1)) - source_slide) > 2:
            continue
        ts = list(el.iter(q("a:t")))
        if not ts:
            continue
        ts[0].text = f"{index:0{len(m.group(1))}d}"
        for t in ts[1:]:
            t.text = ""
        done += 1
    return done


# ---------------------------------------------------------------------------- covers, dividers, closing slides

_BOOKENDS = (PatternKind.title, PatternKind.section, PatternKind.thanks)
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _bookend_lines(oslide: OutlineSlide) -> list[str]:
    """The short text lines a cover or closing slide may carry beyond its heading (a contact, the audience) — a
    cover's goal line (`cover_goal`) is set on its own, under the subtitle."""
    c = oslide.content
    paras = list(c.paragraphs)[1:] if cover_goal(oslide) else list(c.paragraphs)
    return [t.strip() for t in paras + list(c.bullets) if t and t.strip()]


def _is_bookend(oslide: OutlineSlide) -> bool:
    """A cover, a section divider or a closing slide with nothing but its heading and a line or two under it (a
    subtitle, a contact, a cover's goal): these are set as a heading and a subtitle on the sample's ground."""
    if oslide.kind not in _BOOKENDS:
        return False
    c = oslide.content
    lines = _bookend_lines(oslide)
    only_text = not (c.items or c.numbers or c.table or c.chart or c.quote or c.columns)
    return only_text and len(lines) <= 2 and sum(len(t) for t in lines) <= 140 and len(cover_goal(oslide) or "") <= 220


def _cap_first(text: str) -> str:
    """«продуктовый комитет» → «Продуктовый комитет»; an address or a link («team@vk.com») stays as written."""
    first = text.split(" ", 1)[0] if text else ""
    if not text or not text[:1].islower() or re.search(r"[@/]|\w\.\w", first):
        return text
    return text[:1].upper() + text[1:]


def _ph(el: etree._Element) -> Optional[etree._Element]:
    nv = next((c for c in el if etree.QName(c).localname.startswith("nv")), None)
    nvpr = nv.find(q("p:nvPr")) if nv is not None else None
    return nvpr.find(q("p:ph")) if nvpr is not None else None


def _layout_placeholder(ctx: _SlideCtx, el: etree._Element) -> Optional[etree._Element]:
    """The layout shape a slide placeholder inherits from (matched by idx, else by type)."""
    ph = _ph(el)
    if ph is None:
        return None
    try:
        layout = ctx.slide.slide_layout
    except Exception:  # noqa: BLE001
        return None
    idx, typ = ph.get("idx"), ph.get("type") or "body"
    titles = ("title", "ctrTitle")
    best = None
    for sh in layout.placeholders:
        lph = _ph(sh._element)
        if lph is None:
            continue
        if idx is not None and lph.get("idx") == idx:
            return sh._element
        ltyp = lph.get("type") or "body"
        if ltyp == typ or (typ in titles and ltyp in titles):
            best = best or sh._element
    return best


def _style_levels(ctx: _SlideCtx, el: etree._Element) -> list[etree._Element]:
    """Paragraph property levels a text shape inherits, nearest first: its own first paragraph and list style, the
    layout placeholder's, the master's title/body style."""
    out = []
    txb = el.find(q("p:txBody"))
    if txb is not None:
        p0 = txb.find(q("a:p"))
        if p0 is not None and p0.find(q("a:pPr")) is not None:
            out.append(p0.find(q("a:pPr")))
        lvl = txb.find(q("a:lstStyle") + "/" + q("a:lvl1pPr"))
        if lvl is not None:
            out.append(lvl)
    lay = _layout_placeholder(ctx, el)
    if lay is not None:
        lvl = lay.find(q("p:txBody") + "/" + q("a:lstStyle") + "/" + q("a:lvl1pPr"))
        if lvl is not None:
            out.append(lvl)
    ph = _ph(el)
    if ph is not None:
        try:
            master = ctx.slide.slide_layout.slide_master._element
        except Exception:  # noqa: BLE001
            master = None
        if master is not None:
            style = "p:titleStyle" if (ph.get("type") or "") in ("title", "ctrTitle") else "p:bodyStyle"
            lvl = master.find(q("p:txStyles") + "/" + q(style) + "/" + q("a:lvl1pPr"))
            if lvl is not None:
                out.append(lvl)
    return out


def _eff_rpr(ctx: _SlideCtx, el: etree._Element, attr: str) -> Optional[str]:
    """A run attribute (b, sz, cap) as the text will show it: the first run, the paragraph end, then the inherited
    default run properties."""
    for r in el.iter(q("a:rPr")):
        if r.get(attr) is not None:
            return r.get(attr)
        break
    for r in el.iter(q("a:endParaRPr")):
        if r.get(attr) is not None:
            return r.get(attr)
        break
    for lvl in _style_levels(ctx, el):
        d = lvl.find(q("a:defRPr"))
        if d is not None and d.get(attr) is not None:
            return d.get(attr)
    return None


def _eff_line_spacing(ctx: _SlideCtx, el: etree._Element) -> float:
    for lvl in _style_levels(ctx, el):
        pct = lvl.find(q("a:lnSpc") + "/" + q("a:spcPct"))
        if pct is not None and pct.get("val"):
            try:
                return max(0.7, min(1.6, int(pct.get("val")) / 100000))
            except ValueError:
                pass
    return 1.0


def _eff_anchor(ctx: _SlideCtx, el: etree._Element) -> str:
    bp = el.find(q("p:txBody") + "/" + q("a:bodyPr"))
    if bp is not None and bp.get("anchor"):
        return bp.get("anchor")
    lay = _layout_placeholder(ctx, el)
    lbp = lay.find(q("p:txBody") + "/" + q("a:bodyPr")) if lay is not None else None
    return (lbp.get("anchor") if lbp is not None and lbp.get("anchor") else "t")


def _first_line_offset(ctx: _SlideCtx, el: etree._Element) -> int:
    """Where the first line starts inside the text box (marL + indent, inherited): every line of a heading and its
    subtitle start there."""
    marl = indent = None
    for lvl in _style_levels(ctx, el):
        if marl is None and lvl.get("marL") is not None:
            marl = int(lvl.get("marL"))
        if indent is None and lvl.get("indent") is not None:
            indent = int(lvl.get("indent"))
    return max(0, (marl or 0) + (indent or 0))


def _ensure_xfrm(ctx: _SlideCtx, el: etree._Element) -> Optional[tuple[int, int, int, int]]:
    """The geometry of a placeholder that inherits its position from the layout, written onto the slide (a thin
    alias of `deck.materialize_xfrm`; DeckBuilder already materializes every placeholder of a cloned slide)."""
    b = element_bbox(el)
    if b:
        return b
    sid = ctx.id_of(el)
    sh = next((s for s in ctx.slide.shapes if str(s.shape_id) == sid), None)
    if sh is None:
        return None
    materialize_xfrm(sh)
    return element_bbox(el)


def _opaque_part(part, pic: etree._Element, box: Bbox) -> Bbox:
    """The painted part of a picture (a thin alias of `layers.opaque_box`)."""
    return opaque_box(part, pic, box)


def _turned_box(el: etree._Element, box: Bbox) -> Bbox:
    """The box a rotated shape really covers: a sidebar's vertical brand line is a wide box turned by 270°, standing
    in the sidebar — not a band across the slide's middle (Synth Sidebar, gate 2 — C3)."""
    tag = etree.QName(el).localname
    if tag == "grpSp":
        xfrm = el.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
    elif tag == "graphicFrame":
        xfrm = el.find(q("p:xfrm"))
    else:
        sp_pr = el.find(q("p:spPr"))
        xfrm = sp_pr.find(q("a:xfrm")) if sp_pr is not None else None
    try:
        rot = (int(xfrm.get("rot") or 0) / 60000.0) % 180.0 if xfrm is not None else 0.0
    except ValueError:
        rot = 0.0
    if rot < 1.0 or rot > 179.0:
        return box
    a = math.radians(rot)
    c, s_ = abs(math.cos(a)), abs(math.sin(a))
    w, h = int(box.w * c + box.h * s_), int(box.w * s_ + box.h * c)
    return Bbox(x=box.x + box.w // 2 - w // 2, y=box.y + box.h // 2 - h // 2, w=w, h=h)


VEIL_ALPHA = 0.35  # a see-through solid shape up to this opacity is a veil over the ground (the audit's SOLID_VEIL_ALPHA)


def _veil(el: etree._Element) -> bool:
    """A see-through shape without text or outline — a white bubble at 20–30 % over the ground of LibreOffice «Lights»:
    decoration a heading may run over, as the template's own heading does (the ground shows through)."""
    if etree.QName(el).localname != "sp" or shape_text(el).strip():
        return False
    sppr = el.find(q("p:spPr"))
    if sppr is None:
        return False
    ln = sppr.find(q("a:ln"))
    if ln is not None and ln.find(q("a:noFill")) is None and (ln.find(q("a:solidFill")) is not None or ln.find(q("a:gradFill")) is not None):
        return False
    sf = sppr.find(q("a:solidFill"))
    if sf is None or not len(sf) or sppr.find(q("a:gradFill")) is not None or sppr.find(q("a:blipFill")) is not None:
        return False
    a = sf[0].find(q("a:alpha"))
    try:
        return a is not None and int(a.get("val") or 100000) / 100000.0 < VEIL_ALPHA
    except ValueError:
        return False


def _bookend_photo(ctx: "_SlideCtx", slot, el: etree._Element) -> bool:
    """A picture of a cover, divider or closing sample that is part of its design: a real picture (not an empty
    «insert your photo» frame) of at least 3 % of the slide, standing clear of the sample's title (a photo beside the
    title card — Google Slides «Marketing Campaign», «McKinsey»)."""
    if not getattr(ctx, "keep_photos", True) or slot.bbox.area < 0.03 or el.find(".//" + q("a:blip")) is None:
        return False
    for t in ctx.pattern.slots:
        if t.role == SlotRole.title and t.bbox.area > 0:
            inter = max(0.0, min(slot.bbox.x2, t.bbox.x2) - max(slot.bbox.x, t.bbox.x)) * max(0.0, min(slot.bbox.y2, t.bbox.y2) - max(slot.bbox.y, t.bbox.y))
            if inter > 0.1 * t.bbox.area:
                return False
    return True


def _veil_boxes(ctx: _SlideCtx) -> list[Bbox]:
    """The boxes of the see-through shapes (`_veil`) the slide, its layout and its master draw."""
    out: list[Bbox] = []
    holders = [ctx.slide]
    try:
        holders += [ctx.slide.slide_layout, ctx.slide.slide_layout.slide_master]
    except Exception:  # noqa: BLE001
        pass
    for h in holders:
        tree = h._element.find(q("p:cSld") + "/" + q("p:spTree"))
        if tree is None:
            continue
        for el in tree:
            if etree.QName(el).localname == "sp" and _veil(el):
                b = element_bbox(el)
                if b and b[2] > 0 and b[3] > 0:
                    out.append(Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    return out


def _drawn_boxes(ctx: _SlideCtx, exclude: set[int]) -> list[Bbox]:
    """Everything painted on the slide that a heading must not run under: the slide's own shapes and pictures, the
    pictures and logos of its layout and master (placeholders and slide-sized backgrounds aside), pictures measured by
    their opaque pixels."""
    W, H = ctx.W, ctx.H
    out: list[Bbox] = []

    def add(el: etree._Element, part, box: Bbox) -> None:
        box = _turned_box(el, box)
        if box.w <= 0 or box.h <= 0 or box.w * box.h >= 0.6 * W * H:
            return
        if etree.QName(el).localname == "pic":
            box = _opaque_part(part, el, box)
            if box.w <= 0 or box.h <= 0:
                return
        out.append(box)

    tree = ctx.slide._element.cSld.find(q("p:spTree"))
    for el in tree:
        if id(el) in exclude or etree.QName(el).localname not in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp"):
            continue
        if etree.QName(el).localname == "sp" and not shape_text(el).strip() and not has_visible_style(el):
            continue
        if _veil(el):
            continue
        b = element_bbox(el)
        if b:
            add(el, ctx.slide.part, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    try:
        layout = ctx.slide.slide_layout
        parts = [(layout.part, layout.shapes)]
        if layout._element.get("showMasterSp") != "0" and ctx.slide._element.get("showMasterSp") != "0":
            parts.append((layout.slide_master.part, layout.slide_master.shapes))
    except Exception:  # noqa: BLE001
        parts = []
    for part, shapes in parts:
        for sh in shapes:
            el = sh._element
            if _ph(el) is not None or sh.width is None:
                continue
            if etree.QName(el).localname == "sp" and not shape_text(el).strip() and not has_visible_style(el):
                continue
            if _veil(el):
                continue
            add(el, part, Bbox(x=int(sh.left or 0), y=int(sh.top or 0), w=int(sh.width or 0), h=int(sh.height or 0)))
    for c in ctx.manifest.tokens.chrome:
        if c.source == "background":  # logos drawn into a background picture, found on the renders
            b = c.bbox.to_emu(W, H)
            if 0 < b.w * b.h < 0.6 * W * H:
                out.append(b)
    return out


def _detail_map(ctx: _SlideCtx, hidden: list[Bbox]):
    """(detail, kx, ky) of the sample's own render at 300 px wide: the difference between the render and its blur,
    with what the renderer removed or rewrote (`hidden`) masked out. None when the render is missing."""
    try:
        import numpy as np
        from PIL import Image, ImageFilter

        path = ctx.ws.slide_image(ctx.pattern.source_slide)
        if not path.exists():
            return None
        im = Image.open(path).convert("RGB")
    except Exception:  # noqa: BLE001
        return None
    w = 300
    h = max(int(w * im.height / im.width), 1)
    small = im.resize((w, h), Image.BILINEAR)
    act = np.abs(np.asarray(small, dtype=float) - np.asarray(small.filter(ImageFilter.GaussianBlur(3)), dtype=float)).sum(axis=2)
    kx, ky = w / ctx.W, h / ctx.H
    for b in hidden:
        act[max(int(b.y * ky) - 1, 0):int(b.y2 * ky) + 2, max(int(b.x * kx) - 1, 0):int(b.x2 * kx) + 2] = 0.0
    return act, kx, ky


def _art_edge(ctx: _SlideCtx, band: tuple[int, int], x_from: int, hidden: list[Bbox], designed_to: int = 0) -> Optional[int]:
    """Where the sample's art begins to the right of `x_from` inside a horizontal band, read off the sample's own
    render: art painted into a background (the LCT city, the VK Tech glass cube and its rings) has no shape to
    measure. Up to `designed_to` (the right edge of the sample's own heading box) only strong art counts — the
    template designer set the heading over its faint glow and thin line work there. None when the band is calm to
    the right edge or the render is missing."""
    import numpy as np

    dm = _detail_map(ctx, hidden)
    if dm is None:
        return None
    act, kx, ky = dm
    h, w = act.shape
    y0, y1 = max(int(band[0] * ky), 0), min(int(band[1] * ky) + 1, h)
    if y1 - y0 < 2:
        return None
    col = np.percentile(act[y0:y1], 90, axis=0)
    sm = np.convolve(col, np.ones(5) / 5, mode="same")
    run = max(int(0.025 * w), 3)
    strong_to = int(designed_to * kx)
    for x in range(max(int(x_from * kx), 0), w - run):
        if all(sm[x + i] > (55 if x + i < strong_to else 10) for i in range(run)):
            return int(x / kx)
    return None


def _art_above(ctx: _SlideCtx, x0: int, x1: int, y_from: int, hidden: list[Bbox]) -> Optional[int]:
    """The lower edge of the art painted above `y_from` over the columns x0…x1 (a dot pattern or a stripe drawn into
    the background, which no shape measures), read off the sample's render; None when the column is calm."""
    import numpy as np

    dm = _detail_map(ctx, hidden)
    if dm is None:
        return None
    act, kx, ky = dm
    h, w = act.shape
    c0, c1 = max(int(x0 * kx), 0), min(int(x1 * kx) + 1, w)
    if c1 - c0 < 3:
        return None
    rows = np.percentile(act[:, c0:c1], 95, axis=1)
    run = max(int(0.012 * h), 2)
    for y in range(min(int(y_from * ky), h) - 1, run - 1, -1):
        if all(rows[y - i] > 12 for i in range(run)):
            return int((y + 1) / ky)
    return None


def _art_below(ctx: _SlideCtx, x0: int, x1: int, y_from: int, hidden: list[Bbox]) -> Optional[int]:
    """The upper edge of the art painted under `y_from` over the columns x0…x1 (the VK Tech glass cube under a centred
    cover heading: art drawn into the background, which no shape measures), read off the sample's render with the
    sample's own texts masked out; None when the column is calm to the foot of the slide."""
    import numpy as np

    dm = _detail_map(ctx, hidden)
    if dm is None:
        return None
    act, kx, ky = dm
    h, w = act.shape
    c0, c1 = max(int(x0 * kx), 0), min(int(x1 * kx) + 1, w)
    if c1 - c0 < 3:
        return None
    rows = np.percentile(act[:, c0:c1], 75, axis=1)  # most of the column busy: art, not a stray line
    run = max(int(0.02 * h), 3)
    for y in range(max(int(y_from * ky), 0), h - run):
        if all(rows[y + i] > 25 for i in range(run)):
            return int(y / ky)
    return None


def _sample_text_boxes(ctx: _SlideCtx, slots: list[Slot]) -> list[Bbox]:
    """How far the sample's own texts reach on its render (the box of a placeholder is often far wider than its text)."""
    out = []
    Wpt = ctx.W / EMU_PER_PT
    for s in slots:
        if not s.sample_text:
            continue
        size = s.style.size_pt or ctx.typo.size_for("body", 14.0)
        # generous: the render was made with whatever face stood in for the template font, often a wider one
        tw = (1.4 * max(text_width_pt(t, s.style.font_family, size, s.style.bold) for t in s.sample_text.split("\n")) + 24) / Wpt
        b = s.bbox
        x0 = b.x + b.w / 2 - tw / 2 if (s.style.align or "l") == "ctr" else b.x
        tw = min(tw, b.w) if (s.style.align or "l") != "ctr" else tw
        out.append(Bbox(x=int(x0 * ctx.W), y=int(b.y * ctx.H), w=int(tw * ctx.W), h=int(b.h * ctx.H)))
    return out


def _subtitle_carrier(ctx: _SlideCtx, title: Slot) -> Optional[Slot]:
    """The sample's subtitle box, else the first free text box under the heading (roles.py cannot always name it)."""
    subs = ctx.slots(SlotRole.subtitle)
    if subs:
        return min(subs, key=lambda s: abs(s.bbox.y - title.bbox.y2))
    below = [s for s in ctx.slots(SlotRole.body, SlotRole.bullet_list, SlotRole.caption) if s.bbox.y >= title.bbox.y2 - 0.02 and not s.group_id]
    return min(below, key=lambda s: s.bbox.y) if below else None


def _strip_bookend(ctx: _SlideCtx, keep: set[str], oslide: OutlineSlide) -> tuple[list[Bbox], list[etree._Element]]:
    """A cover or closing slide keeps its heading, subtitle, chrome and art. Every other text box of the sample
    (speaker name and role, «Вставить фото», a QR label) leaves with what was drawn for it — the avatar circle, the
    photo frame, a small picture pointing at a removed QR tile — and so do empty placeholders. Outlines that frame
    nothing (the LCT corner frame) are returned to the caller, which fills one with a kicker or removes them.
    Returns (boxes of what left, frames)."""
    W, H = ctx.W, ctx.H
    chrome = set(ctx.pattern.chrome_shape_ids)
    tree = ctx.slide._element.cSld.find(q("p:spTree"))
    gone: list[Bbox] = []
    cleared: list[Bbox] = []  # where the sample showed something that is gone now (its render must not count it as art)
    tiles: list[Bbox] = []  # painted boxes that left (a QR tile, a button, a photo frame): pictures pointing at them go too

    def drop(el: etree._Element) -> None:
        b = element_bbox(el)
        if b and b[2] > 0 and b[3] > 0:
            box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
            cleared.append(box)
            if not is_nested(el) and (etree.QName(el).localname == "pic" or has_visible_style(el)):
                tiles.append(box)
        ctx.remove_el(el)

    for slot in ctx.pattern.slots:
        if slot.shape_id in keep or slot.shape_id in ctx.removed:
            continue
        el = ctx.els.get(slot.shape_id)
        if el is None or el.getparent() is None:
            continue
        if slot.role == SlotRole.image and oslide.content.image_hint:
            continue
        if slot.role == SlotRole.image and _bookend_photo(ctx, slot, el):
            if KEPT_PHOTO not in ctx.warnings:
                ctx.warnings.append(KEPT_PHOTO)
            continue  # the photo beside the title card is the cover's design: without it half the slide stands empty
        b = element_bbox(el)
        if is_nested(el):
            if etree.QName(el).localname == "sp":
                clear_text(el)
            continue
        drop(el)
        if b:
            gone.append(Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    # a device mock-up (a phone frame around an empty screen) or an empty photo frame shows nothing without the
    # content's own picture: the slide's copy of it leaves (one the layout draws stays — the matcher avoids those)
    if not oslide.content.image_hint and ctx.pattern.kind == PatternKind.title:
        for mb in getattr(ctx.pattern, "mockup_boxes", None) or []:
            m = mb.to_emu(W, H)
            if m.area <= 0:
                continue
            for sid, el in list(ctx.els.items()):
                if sid in keep or sid in chrome or sid in ctx.removed or el.getparent() is not tree:
                    continue
                b = element_bbox(el)
                if not b or b[2] <= 0 or b[3] <= 0:
                    continue
                eb = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
                inside = eb.intersection(m) >= 0.8 * eb.area
                # the picture that draws the device (a phone on its decorative circles) holds the mock-up's box
                device = etree.QName(el).localname in ("pic", "grpSp") and eb.intersection(m) >= 0.8 * m.area and eb.area < 0.6 * W * H
                if inside or device:
                    drop(el)
                    gone.append(eb)
    # the subtitle moves under the heading: whatever was drawn beside it where it stood leaves too
    for sid in keep:
        slot = ctx.slot_of.get(sid)
        if slot is not None and slot.role != SlotRole.title and is_placeholder_text(slot.sample_text):
            gone.append(slot.bbox.to_emu(W, H))
    for sid, el in list(ctx.els.items()):
        if sid in keep or sid in chrome or sid in ctx.removed or el.getparent() is not tree or etree.QName(el).localname not in ("sp", "pic"):
            continue
        if etree.QName(el).localname == "sp" and (shape_text(el).strip() or _is_dot(el, W)):
            continue  # pager dots are handled on their own
        b = element_bbox(el)
        if not b or b[2] <= 0 or b[3] <= 0:
            continue
        box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
        if etree.QName(el).localname == "pic":
            box = _opaque_part(ctx.slide.part, el, box)  # a glass cursor drawn on a slide-wide transparent picture
            if box.w <= 0 or box.h <= 0:
                continue
        if box.area > 0.06 * W * H:
            continue
        for g in gone:
            over = box.intersection(g)
            beside = (0 <= g.x - box.x2 <= int(0.03 * W) or 0 <= box.x - g.x2 <= int(0.03 * W)) and min(box.y2, g.y2) - max(box.y, g.y) >= 0.5 * min(box.h, g.h)
            if over >= 0.3 * min(box.area, g.area) or beside:
                drop(el)
                break
        else:
            if etree.QName(el).localname == "pic":
                # a rider: a small picture drawn onto a removed tile (the paper plane pointing at a QR code)
                near = Bbox(x=box.x - int(0.02 * W), y=box.y - int(0.02 * H), w=box.w + int(0.04 * W), h=box.h + int(0.04 * H))
                if any(near.intersection(t) >= 0.1 * box.area for t in tiles):
                    drop(el)
    # placeholders nobody filled show their prompt («Логотип», «Иконка задачи») the moment the deck is opened
    for sid, el in list(ctx.els.items()):
        if sid in keep or sid in ctx.removed or el.getparent() is None or is_nested(el) or etree.QName(el).localname != "sp":
            continue
        ph = _ph(el)
        if ph is None or (ph.get("type") or "") in ("sldNum", "dt", "ftr") or shape_text(el).strip():
            continue
        spPr = el.find(q("p:spPr"))
        if spPr is not None and spPr.find(q("a:blipFill")) is not None:
            continue
        drop(el)
    # an outline that frames nothing (the LCT corner frame around a logo placeholder) is an unfinished corner
    kept_centres = []
    for sid, el in ctx.els.items():
        if sid in ctx.removed or el.getparent() is None:
            continue
        name = etree.QName(el).localname
        if sid in keep or name in ("pic", "graphicFrame") or (name == "sp" and shape_text(el).strip()):
            b = element_bbox(el)
            if b:
                kept_centres.append((b[0] + b[2] / 2, b[1] + b[3] / 2))
    frames = []
    for sid, el in list(ctx.els.items()):
        if sid in keep or sid in chrome or sid in ctx.removed or el.getparent() is not tree or etree.QName(el).localname != "sp" or shape_text(el).strip():
            continue
        spPr = el.find(q("p:spPr"))
        ln = spPr.find(q("a:ln")) if spPr is not None else None
        filled = _has_fill(el)
        outlined = ln is not None and ln.find(q("a:noFill")) is None and (ln.find(q("a:solidFill")) is not None or ln.find(q("a:gradFill")) is not None)
        if filled or not outlined:
            continue
        b = element_bbox(el)
        if not b or b[2] * b[3] >= 0.6 * W * H:
            continue
        if not any(b[0] <= cx <= b[0] + b[2] and b[1] <= cy <= b[1] + b[3] for cx, cy in kept_centres):
            frames.append(el)
    _remove_placeholder_boxes(ctx)
    return cleared, frames


def _drop_frames(ctx: _SlideCtx, frames: list[etree._Element], cleared: list[Bbox]) -> None:
    for el in frames:
        b = element_bbox(el)
        if b and b[2] > 0 and b[3] > 0:
            cleared.append(Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
        ctx.remove_el(el)


# ---- pagination dots


def _abs_scale(el: etree._Element) -> float:
    """How many slide EMU one child-space EMU of a nested shape is (groups scale their children)."""
    k = 1.0
    parent = el.getparent()
    while parent is not None and etree.QName(parent).localname == "grpSp":
        xfrm = parent.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
        ext, chext = (xfrm.find(q("a:ext")), xfrm.find(q("a:chExt"))) if xfrm is not None else (None, None)
        if ext is not None and chext is not None and int(chext.get("cx") or 0) > 0:
            k *= int(ext.get("cx")) / int(chext.get("cx"))
        parent = parent.getparent()
    return k


def _is_dot(el: etree._Element, W: int) -> bool:
    if etree.QName(el).localname != "sp" or shape_text(el).strip():
        return False
    geom = el.find(q("p:spPr") + "/" + q("a:prstGeom"))
    if geom is None or geom.get("prst") not in ("ellipse", "roundRect", "flowChartConnector"):
        return False
    b = element_bbox(el)
    if not b or b[2] <= 0 or b[3] <= 0 or not 0.8 <= b[2] / b[3] <= 1.25:
        return False
    return b[2] * _abs_scale(el) <= 0.03 * W and _has_fill(el)


def _dot_rows(ctx: _SlideCtx) -> list[list[etree._Element]]:
    """Pager dots: three or more equal small circles in one row at an even pitch, drawn loose or as a group."""
    tree = ctx.slide._element.cSld.find(q("p:spTree"))
    rows = []
    for parent in [tree] + list(tree.iter(q("p:grpSp"))):
        dots = [c for c in parent if _is_dot(c, ctx.W)]
        used: set[int] = set()
        for d in dots:
            if id(d) in used:
                continue
            b = element_bbox(d)
            row = [e for e in dots if id(e) not in used and abs(element_bbox(e)[2] - b[2]) <= 0.12 * b[2] and abs(element_bbox(e)[1] + element_bbox(e)[3] / 2 - b[1] - b[3] / 2) <= 0.5 * b[3]]
            row.sort(key=lambda e: element_bbox(e)[0])
            if len(row) < 3:
                continue
            xs = [element_bbox(e)[0] for e in row]
            steps = [x2 - x1 for x1, x2 in zip(xs, xs[1:])]
            med = sorted(steps)[len(steps) // 2]
            if med <= 0 or med > 5 * b[2] or any(abs(s - med) > 0.2 * med for s in steps):
                continue
            used.update(id(e) for e in row)
            rows.append(row)
    return rows


def _fill_node(el: etree._Element) -> Optional[etree._Element]:
    spPr = el.find(q("p:spPr"))
    if spPr is None:
        return None
    return next((c for c in spPr if etree.QName(c).localname in ("solidFill", "gradFill", "noFill", "pattFill")), None)


def _set_fill_node(el: etree._Element, node: etree._Element) -> None:
    spPr = el.find(q("p:spPr"))
    old = _fill_node(el)
    new = copy.deepcopy(node)
    if old is not None:
        old.addprevious(new)
        spPr.remove(old)
    else:
        anchor = spPr.find(q("a:prstGeom")) if spPr.find(q("a:prstGeom")) is not None else spPr.find(q("a:custGeom"))
        if anchor is not None:
            anchor.addnext(new)
        else:
            spPr.append(new)


def _abs_bbox(el: etree._Element) -> Optional[Bbox]:
    """A shape's box in slide coordinates, through the transforms of the groups holding it."""
    b = element_bbox(el)
    if not b:
        return None
    x, y, w, h = (float(v) for v in b)
    parent = el.getparent()
    while parent is not None and etree.QName(parent).localname == "grpSp":
        xfrm = parent.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
        if xfrm is None:
            break
        off, ext, choff, chext = (xfrm.find(q(t)) for t in ("a:off", "a:ext", "a:chOff", "a:chExt"))
        if None in (off, ext, choff, chext) or not int(chext.get("cx") or 0) or not int(chext.get("cy") or 0):
            break
        kx, ky = int(ext.get("cx")) / int(chext.get("cx")), int(ext.get("cy")) / int(chext.get("cy"))
        x = int(off.get("x")) + (x - int(choff.get("x"))) * kx
        y = int(off.get("y")) + (y - int(choff.get("y"))) * ky
        w, h = w * kx, h * ky
        parent = parent.getparent()
    return Bbox(x=int(x), y=int(y), w=int(w), h=int(h))


def _row_visible(ctx: _SlideCtx, row: list[etree._Element]) -> bool:
    """Whether a row of dots shows on the sample's render: a second row painted in the ground colour (Education keeps
    one under its visible dots) is invisible clutter in the editable deck."""
    try:
        import numpy as np
        from PIL import Image

        path = ctx.ws.slide_image(ctx.pattern.source_slide)
        if not path.exists():
            return True
        im = np.asarray(Image.open(path).convert("RGB"), dtype=float)
    except Exception:  # noqa: BLE001
        return True
    ih, iw = im.shape[:2]
    kx, ky = iw / ctx.W, ih / ctx.H
    for d in row:
        b = _abs_bbox(d)
        if b is None or b.w <= 0:
            return True
        cx, cy, r = (b.x + b.w / 2) * kx, (b.y + b.h / 2) * ky, max(b.w * kx / 2, 2.0)
        x0, x1, y0, y1 = int(cx - 2 * r), int(cx + 2 * r) + 1, int(cy - 2 * r), int(cy + 2 * r) + 1
        if x0 < 0 or y0 < 0 or x1 > iw or y1 > ih:
            return True
        patch = im[y0:y1, x0:x1]
        yy, xx = np.mgrid[y0:y1, x0:x1]
        dist = np.hypot(xx + 0.5 - cx, yy + 0.5 - cy)
        inner, ring = patch[dist <= 0.5 * r], patch[(dist >= 1.4 * r) & (dist <= 1.9 * r)]
        if not len(inner) or not len(ring):
            return True
        if np.abs(inner.mean(axis=0) - ring.mean(axis=0)).max() >= 14:
            return True
    return False


def _pagination(ctx: _SlideCtx, oslide: OutlineSlide) -> list[list[etree._Element]]:
    """Pager dots count the deck's sections: on a divider they are redrawn to that count with the current section
    lit; on a cover or closing slide (or with fewer than two sections) they are an empty promise and leave, and so
    does a row nobody can see. Returns the rows that stay."""
    rows = _dot_rows(ctx)
    if not rows:
        return []
    sections = [s.id for s in ctx.outline.slides if s.kind == PatternKind.section]
    cur = sections.index(oslide.id) if oslide.id in sections else None
    n = len(sections)
    kept = []
    for row in rows:
        parent = row[0].getparent()
        if oslide.kind != PatternKind.section or n < 2 or cur is None or not _row_visible(ctx, row):
            for d in row:
                ctx.remove_el(d)
            if etree.QName(parent).localname == "grpSp" and not any(etree.QName(c).localname in ("sp", "pic", "grpSp", "cxnSp", "graphicFrame") for c in parent):
                remove_element(parent)
            continue
        keys = [etree.tostring(_fill_node(d)) if _fill_node(d) is not None else b"" for d in row]
        common = max(set(keys), key=keys.count)
        inactive = _fill_node(row[keys.index(common)])
        odd = [i for i, k in enumerate(keys) if k != common]
        active = _fill_node(row[odd[0]]) if odd else None
        if active is None:
            accent = next(iter(ctx.manifest.tokens.accents()), None) or "0077FF"
            active = etree.Element(q("a:solidFill"))
            etree.SubElement(active, q("a:srgbClr")).set("val", accent)
        xs = [element_bbox(d)[0] for d in row]
        pitch = sorted(x2 - x1 for x1, x2 in zip(xs, xs[1:]))[(len(xs) - 1) // 2]
        while len(row) > n:
            ctx.remove_el(row.pop())
        while len(row) < n:
            nd = copy.deepcopy(row[-1])
            ctx.next_id = renumber_ids(nd, ctx.next_id)
            row[-1].addnext(nd)
            b = element_bbox(row[-1])
            set_element_pos(nd, x=b[0] + pitch)
            row.append(nd)
        if etree.QName(parent).localname == "grpSp":
            # the group's frame follows its children, so the dots keep their size and pitch
            xfrm = parent.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
            ext, chext, choff = xfrm.find(q("a:ext")), xfrm.find(q("a:chExt")), xfrm.find(q("a:chOff"))
            if ext is not None and chext is not None and choff is not None and int(chext.get("cx") or 0) > 0:
                k = int(ext.get("cx")) / int(chext.get("cx"))
                right = max(element_bbox(d)[0] + element_bbox(d)[2] for d in row)
                new_cx = right - int(choff.get("x"))
                chext.set("cx", str(new_cx))
                ext.set("cx", str(int(new_cx * k)))
        for i, d in enumerate(row):
            _set_fill_node(d, active if i == cur else inactive)
        kept.append(row)
    return kept


def _move_rows(rows: list[list[etree._Element]], dy: int) -> None:
    """Shift pager dot rows vertically (the group that holds a row, or each loose dot)."""
    for row in rows:
        parent = row[0].getparent()
        if etree.QName(parent).localname == "grpSp" and all(d.getparent() is parent for d in row):
            shift_element(parent, 0, dy)
        else:
            for d in row:
                shift_element(d, 0, dy)


# ---- heading and subtitle


def _inherits_caps(ctx: _SlideCtx, el: etree._Element) -> bool:
    """The text of this shape shows in capitals whatever case it is typed in: `cap="all"` (or small capitals) on its
    first run, its list style, its layout's and master's placeholder or the master's title/body/other text style."""
    if (_eff_rpr(ctx, el, "cap") or "") in ("all", "small"):
        return True
    from verstka.rendering.textfill import inherited_caps_spc

    shp = next((s for s in ctx.slide.shapes if s._element is el), None)
    try:
        return bool(inherited_caps_spc(shp if shp is not None else el)[0])
    except Exception:  # noqa: BLE001 - unknown: measured as typed
        return False


def _caps_convention(ctx: _SlideCtx, el: etree._Element, slot: Slot) -> bool:
    """Whether the template sets this heading in capitals: its layout prompt or sample text is written so."""
    if (_eff_rpr(ctx, el, "cap") or "") == "all":
        return False  # already set in capitals by the template: nothing to add
    texts = [slot.sample_text or ""]
    lay = _layout_placeholder(ctx, el)
    if lay is not None:
        texts.append(shape_text(lay))
    for t in texts:
        letters = [c for c in t if c.isalpha()]
        if len(letters) >= 4 and all(c.isupper() for c in letters):
            return True
    return False


def _new_text(ctx: _SlideCtx, src: etree._Element, name: str, family: Optional[str], color: Optional[str]) -> etree._Element:
    """A text box drawn like `src` (a subtitle, a kicker, a footer line on a sample that has none): a copy of the
    sample's own text box without its placeholder link, so it keeps the template's face and insets."""
    el = copy.deepcopy(src)
    ctx.next_id = renumber_ids(el, ctx.next_id)
    ph = _ph(el)
    if ph is not None:
        ph.getparent().remove(ph)
    nv = el.find(q("p:nvSpPr") + "/" + q("p:cNvPr"))
    if nv is not None:
        nv.set("name", name)
    spPr = el.find(q("p:spPr"))
    if spPr is not None and spPr.find(q("a:prstGeom")) is None and spPr.find(q("a:custGeom")) is None:
        # a placeholder takes its geometry from the layout; cut loose from it, the box needs its own (a plain
        # rectangle with no fill), or editors cannot tell what kind of shape it is
        geom = etree.Element(q("a:prstGeom"))
        geom.set("prst", "rect")
        etree.SubElement(geom, q("a:avLst"))
        xfrm = spPr.find(q("a:xfrm"))
        if xfrm is not None:
            xfrm.addnext(geom)
        else:
            spPr.insert(0, geom)
        if not any(spPr.find(q(t)) is not None for t in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")):
            geom.addnext(etree.Element(q("a:noFill")))
    for r in list(el.iter(q("a:rPr"))) + list(el.iter(q("a:endParaRPr"))):
        # the heading's capitals are its own; cut loose from its placeholder, the box would otherwise take the master's
        # «other text» style — capitals and letter-spacing (spc 500) that spread a sentence wider than it was measured
        r.set("cap", "none")
        if r.get("spc") is None:
            r.set("spc", "0")
    for para in el.iter(q("a:p")):
        if all(r.find(q("a:rPr")) is None for r in para.findall(q("a:r"))) and para.find(q("a:endParaRPr")) is None:
            para.append(etree.Element(q("a:endParaRPr"), {"lang": "ru-RU", "cap": "none", "spc": "0"}))
    # the new box stands on the slide itself, never inside the group its source sits in (the title of a sample may be
    # grouped with a decorative bar: a subtitle placed in the group's child coordinates would land off the slide)
    anchor = src
    while anchor.getparent() is not None and etree.QName(anchor.getparent()).localname == "grpSp":
        anchor = anchor.getparent()
    anchor.addnext(el)
    style_runs(el, family or ctx.typo.primary_family, color, align=None)
    sid = ctx.id_of(el)
    if sid:
        ctx.els[sid] = el
    return el


def _set_paragraph_box(el: etree._Element, marl: int, align: Optional[str], caps: bool = False, bold: Optional[bool] = None, line_spacing: Optional[float] = None) -> None:
    """Every paragraph starts at `marl` with no hanging indent; capitals, weight and line spacing are written out, so
    that every renderer — and the audit — sets the text the way it was measured."""
    for p in el.findall(q("p:txBody") + "/" + q("a:p")):
        pPr = p.find(q("a:pPr"))
        if pPr is None:
            pPr = etree.Element(q("a:pPr"))
            p.insert(0, pPr)
        pPr.set("marL", str(int(marl)))
        pPr.set("indent", "0")
        if align:
            pPr.set("algn", align)
        if line_spacing:
            for old in pPr.findall(q("a:lnSpc")):
                pPr.remove(old)
            ln = etree.Element(q("a:lnSpc"))
            etree.SubElement(ln, q("a:spcPct")).set("val", str(int(round(line_spacing * 100000))))
            pPr.insert(0, ln)
        for r in list(p.iter(q("a:rPr"))) + list(p.iter(q("a:endParaRPr"))):
            if caps:
                r.set("cap", "all")
            if bold is not None:
                r.set("b", "1" if bold else "0")


def _bookend_subtitle_size(ctx: _SlideCtx) -> float:
    """The template's own subtitle size on covers and dividers (the modal one), else its h2."""
    from collections import Counter

    sizes = Counter(round(s.style.size_pt, 1) for p in ctx.manifest.patterns if p.kind in _BOOKENDS for s in p.slots if s.role == SlotRole.subtitle and s.style.size_pt)
    return sizes.most_common(1)[0][0] if sizes else ctx.typo.size_for("h2", 16.0)


def _split_like(text: str, measured: list[str]) -> list[str]:
    """`text` cut where its measured twin (its capitals) was cut. The space at a break is dropped: LibreOffice sets
    a trailing space that does not fit on a line of its own, an empty line inside the heading."""
    out, pos = [], 0
    for line in measured:
        out.append(text[pos : pos + len(line)])
        pos += len(line)
        while pos < len(text) and text[pos] == " ":
            pos += 1
    if pos < len(text) and out:
        out[-1] += " " + text[pos:]
    return out


def _write_lines(el: etree._Element, lines: list[str]) -> None:
    """The heading's lines as runs joined by line breaks: every renderer breaks it where it was balanced."""
    p = el.find(q("p:txBody") + "/" + q("a:p"))
    runs = p.findall(q("a:r")) if p is not None else []
    if not runs or len(lines) < 2:
        return
    r0 = runs[0]
    rpr = r0.find(q("a:rPr"))
    for c in list(p):
        if etree.QName(c).localname in ("r", "br"):
            p.remove(c)
    end = p.find(q("a:endParaRPr"))
    for i, line in enumerate(lines):
        if i:
            br = etree.Element(q("a:br"))
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            if end is not None:
                end.addprevious(br)
            else:
                p.append(br)
        r = copy.deepcopy(r0)
        t = r.find(q("a:t"))
        t.text = typeset_figures(line)
        if line != line.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


def _no_autofit(el: etree._Element, anchor: Optional[str] = None) -> None:
    bp = el.find(q("p:txBody") + "/" + q("a:bodyPr"))
    if bp is None:
        return
    for c in list(bp):
        if etree.QName(c).localname in ("normAutofit", "spAutoFit", "noAutofit"):
            bp.remove(c)
    etree.SubElement(bp, q("a:noAutofit"))
    bp.set("wrap", "square")
    if anchor:
        bp.set("anchor", anchor)


@dataclass
class _Bookends:
    """What a deck decides once for all its covers, dividers and closing slides, the way a designer sets them as a
    family: the cover is the loudest heading, every divider shares one size and one top line, the closing slide
    answers the cover, and all of them speak in the cover's weight, case and subtitle colour."""

    cover_size: Optional[float] = None
    cover_bold: Optional[bool] = None
    caps: Optional[bool] = None
    sub_color: Optional[str] = None
    sub_dark: Optional[bool] = None
    sub_size: Optional[float] = None
    divider_size: Optional[float] = None
    divider_pattern: Optional[str] = None
    divider_lines: int = 1


_DECK_BOOKENDS: "weakref.WeakKeyDictionary[DeckBuilder, _Bookends]" = weakref.WeakKeyDictionary()


def _bookends(ctx: _SlideCtx) -> _Bookends:
    """The bookend decisions of the deck being built (one per DeckBuilder: an autofix re-render starts afresh)."""
    try:
        st = _DECK_BOOKENDS.get(ctx.builder)
        if st is None:
            st = _Bookends()
            _DECK_BOOKENDS[ctx.builder] = st
        return st
    except TypeError:
        return _Bookends()


_MONTHS = {
    "ru": "Январь Февраль Март Апрель Май Июнь Июль Август Сентябрь Октябрь Ноябрь Декабрь".split(),
    "en": "January February March April May June July August September October November December".split(),
}


def _deck_date(language: Optional[str]) -> str:
    """«Сентябрь 2026»: the month the deck was made, the footer line of a cover."""
    import datetime

    d = datetime.date.today()
    names = _MONTHS["ru" if (language or "ru").lower().startswith("ru") else "en"]
    return f"{names[d.month - 1]} {d.year}"


_SECTION_NO_RE = re.compile(r"(?i)(раздел|часть|глава|блок|этап|section|part|chapter)\s*(№\s*)?[0-9IVX]{1,4}\.?")


def _bookend_texts(ctx: _SlideCtx, oslide: OutlineSlide) -> tuple[str, Optional[str], Optional[str], Optional[str]]:
    """(heading, subtitle, kicker, footer) of a cover, divider or closing slide.

    A cover heading of two phrases («…: итоги пилота и план запуска») is set as a title and its subtitle; the audience
    line it displaces becomes the kicker over the title; the footer carries the date. A divider's kicker is its
    number in the deck («02»). A closing slide with nothing to add signs off with the deck title."""
    head = " ".join(oslide.headline.split())
    lines = [" ".join(t.split()) for t in [oslide.subtitle or ""] + _bookend_lines(oslide) if t and t.strip()]
    kicker = footer = None
    if oslide.kind == PatternKind.title:
        head, tail = split_display_title(head)
        if tail:
            if lines:
                kicker = _cap_first(lines.pop(0))
            lines.insert(0, tail)
        footer = _deck_date(ctx.outline.language)
    elif oslide.kind == PatternKind.section:
        secs = [s.id for s in ctx.outline.slides if s.kind == PatternKind.section]
        if len(secs) >= 2 and oslide.id in secs:
            kicker = f"{secs.index(oslide.id) + 1:02d}"
        # «Раздел 2» under «02» says the number twice: a divider keeps only a subtitle that says something
        lines = [t for t in lines if not _SECTION_NO_RE.fullmatch(t)]
    elif not lines and ctx.outline.title:
        deck = split_display_title(ctx.outline.title)[0]
        if deck and len(deck) <= 60 and not same_words(deck, head):
            lines = [deck]
    sub = _cap_first(" · ".join(lines)) if lines else None
    return head, sub, kicker, footer


def _pick_frame(ctx: _SlideCtx, frames: list[etree._Element]) -> Optional[etree._Element]:
    """An empty rounded frame big enough to hold a line of text (the LCT corner frame): it takes the kicker."""
    best = None
    for el in frames:
        b = element_bbox(el)
        if not b:
            continue
        x0, y0 = max(b[0], 0), max(b[1], 0)
        x1, y1 = min(b[0] + b[2], ctx.W), min(b[1] + b[3], ctx.H)
        w, h = x1 - x0, y1 - y0
        if w >= 0.2 * ctx.W and h >= 0.1 * ctx.H and (best is None or w * h > best[0]):
            best = (w * h, el)
    return best[1] if best else None


def _straddles(ctx: _SlideCtx, el: Optional[etree._Element]) -> bool:
    """The element's box straddles an edge of what the layout or master paints (half on a triangle, half off it).
    Unsure (no geometry, layers unreadable) is False: nothing moves on a guess."""
    b = element_bbox(el) if el is not None else None
    if not b or b[2] <= 0 or b[3] <= 0:
        return False
    from verstka.rendering.layers import ground_is_mixed

    try:
        return ground_is_mixed(ctx.slide, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    except Exception:  # noqa: BLE001 - layers are advice
        return False


def _calm_at(ctx: _SlideCtx, box: Bbox) -> bool:
    """Nothing the layout or master paints has an edge inside the box: the box lies wholly on one painted layer (or on
    none of them). Stricter than `ground_is_mixed`: the tip of a triangle reaching 5 % into the box already counts."""
    from verstka.rendering.layers import _ground_layers

    try:
        for lay, share, _got in _ground_layers(ctx.slide, box) or ():
            if lay is None or share >= 0.9:
                return True
            if share > 0.03:
                return False
    except Exception:  # noqa: BLE001 - unsure is not calm
        return False
    return True


def _recolor_for_ground(ctx: _SlideCtx, el: etree._Element, color: Optional[str], size: float, bold: bool, prefs: list[Optional[str]]) -> Optional[str]:
    """The line's colour where it really stands (a note at the foot on another triangle than the heading's): kept when
    it reads there (4.5:1, large text 3:1), else a colour of the template that does. Returns the colour set."""
    b = element_bbox(el)
    if not b or not color:
        return color
    g = _ground_hex(ctx, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    if not g:
        return color
    need = 3.0 if size >= 18 or (bold and size >= 14) else 4.5
    if contrast_ratio(color, g) >= need:
        return color
    # the first colour of the template that reads at this size (white on a brand band holds 3:1 for a large line —
    # it need not turn black), else the most readable one
    toks = ctx.manifest.tokens
    cands = [c for c in prefs if c] + [toks.color_for(r) for r in ("text.secondary", "background.light", "background.dark")] + ["FFFFFF", "000000"]
    better = next((c for c in cands if c and contrast_ratio(c, g) >= need), None) or _readable_on(g, [c for c in prefs if c] + [color], ctx.manifest)
    if better and better != color and contrast_ratio(better, g) > contrast_ratio(color, g):
        style_runs(el, None, better, align=None)
        return better
    return color


def _lines_ink(lines: list[str], family: Optional[str], size: float, bold: bool, align: str, x: int, y: int, w: int, ins: tuple[int, int, int, int], marl: int, ls: float) -> Optional[Bbox]:
    """The box a heading's letters cover: each line measured at its alignment inside the box, the lines stacked from
    the top inset. None without lines."""
    pitch = int(1.2 * ls * size * EMU_PER_PT)
    xs: list[tuple[int, int]] = []
    for ln in lines:
        if not ln.strip():
            continue
        wi = int(text_width_pt(ln, family, size, bold) * 1.04 * EMU_PER_PT)
        if align == "ctr":
            c = x + ins[0] + marl + (w - ins[0] - ins[2] - marl) // 2
            xs.append((c - wi // 2, c + wi // 2))
        elif align == "r":
            xs.append((x + w - ins[2] - wi, x + w - ins[2]))
        else:
            xs.append((x + ins[0] + marl, x + ins[0] + marl + wi))
    if not xs:
        return None
    a, b = min(v[0] for v in xs), max(v[1] for v in xs)
    return Bbox(x=a, y=y + ins[1], w=max(b - a, 1), h=max(pitch * len(lines), 1))


def _heading_shift_on_ground(ctx: _SlideCtx, lines: list[str], family: Optional[str], size: float, bold: bool, align: str, x: int, y: int, w: int, ins: tuple[int, int, int, int], marl: int, ls: float) -> Optional[int]:
    """How far a cover heading's box must move sideways so that every line's letters lie inside the run of the
    triangle or freeform of the layout it stands on, with 1 % of the slide's width to spare, measured over that line's
    own height (gate 2, C2): 0 when they already do (or no line stands on a slanted ground), None when no shift of up
    to 8 % of the slide's width does."""
    from verstka.rendering.layers import _outline_of, ground_span

    W = ctx.W
    pad = int(0.01 * W)
    pitch = int(1.2 * ls * size * EMU_PER_PT)
    lo, hi = -int(0.08 * W), int(0.08 * W)
    checked = False
    for i, ln in enumerate(lines):
        wi = int(text_width_pt(ln, family, size, bold) * 1.04 * EMU_PER_PT)
        if align == "ctr":
            c = x + ins[0] + marl + (w - ins[0] - ins[2] - marl) // 2
            a, b = c - wi // 2, c + wi // 2
        elif align == "r":
            b = x + w - ins[2]
            a = b - wi
        else:
            a = x + ins[0] + marl
            b = a + wi
        band = Bbox(x=a, y=y + ins[1] + i * pitch, w=max(b - a, 1), h=max(pitch, 1))
        try:
            got = ground_span(ctx.slide, band)
        except Exception:  # noqa: BLE001 - layers are advice
            got = None
        if got is None or _outline_of(got[2]) is None or got[2].cover >= 0.85:
            continue  # a slide ground, a rectangle or a picture: the heading's panel logic rules there
        checked = True
        lo = max(lo, got[0] + pad - a, pad - a)
        hi = min(hi, got[1] - pad - b, W - pad - b)
    if not checked or lo <= 0 <= hi:
        return 0
    if lo > hi:
        return None
    return int(lo) if lo > 0 else int(hi)


def _lines_on_one_ground(ctx: _SlideCtx, oslide: OutlineSlide, title_box: Bbox, s_el: Optional[etree._Element], s_box0, carrier, g_el: Optional[etree._Element], goal: Optional[str], g_size: float, sub_size: float, family: Optional[str], own_sub: Optional[str], foot: int) -> tuple[bool, Optional[etree._Element], bool]:
    """A cover's lines under the heading each stand on one ground (T15): a subtitle that straddles the edge of what
    the layout paints (half on a triangle, half off it) goes back to where the sample set its subtitle when that place
    is calm and clear of the heading; a goal line that straddles an edge follows the subtitle there when it fits under
    it, else it is said aloud (speaker notes). Returns (subtitle moved, goal element or None, goal went to the notes)."""
    W, H = ctx.W, ctx.H
    s_moved = goal_gone = False
    home: Optional[Bbox] = None
    s_color = own_sub
    if s_el is not None and s_box0 and carrier is not None and _straddles(ctx, s_el):
        hb = Bbox(x=s_box0[0], y=s_box0[1], w=s_box0[2], h=s_box0[3])
        ins = _body_insets(s_el)
        text = shape_text(s_el).strip()
        lines = max(1, len(display_lines(text, family, sub_size, False, max((hb.w - ins[0] - ins[2]) / EMU_PER_PT * 0.92, 10.0))))
        hh = int((lines * 1.2 + 0.25) * sub_size * EMU_PER_PT) + ins[1] + ins[3]
        # the lines in the middle of the sample's box (the calm place the template designed for them), else at its
        # top or its foot
        cand = None
        for top in (hb.y + (hb.h - hh) // 2, hb.y, hb.y2 - hh):
            c = Bbox(x=hb.x, y=max(top, 0), w=hb.w, h=hh)
            if c.y2 <= foot and c.intersection(title_box) <= 0 and _calm_at(ctx, c):
                cand = c
                break
        if cand is not None:
            set_element_pos(s_el, x=cand.x, y=cand.y, w=cand.w, h=cand.h)
            own_align = carrier.style.align if carrier.style.align in ("l", "ctr", "r") else None
            if own_align:
                for p in s_el.findall(q("p:txBody") + "/" + q("a:p")):
                    pPr = p.find(q("a:pPr"))
                    if pPr is not None:
                        pPr.set("algn", own_align)
            color = carrier.style.color_hex or own_sub
            if color:
                style_runs(s_el, None, color, align=None)
            s_color = _recolor_for_ground(ctx, s_el, color, sub_size, False, [carrier.style.color_hex, own_sub])
            ctx.warnings.append("подзаголовок обложки поставлен на место подзаголовка образца: под заголовком он пересекал край рисунка шаблона")
            s_moved, home = True, cand
    if g_el is not None and (s_moved or _straddles(ctx, g_el)):
        placed = False
        if s_moved and home is not None:
            gb = element_bbox(g_el)
            fit = _goal_after_subtitle(ctx, oslide, title_box, s_el, s_box0, home, g_el, goal or shape_text(g_el), g_size, sub_size, family, foot)
            if gb and fit is not None:
                s_box, cand, s_new, g_new = fit
                if s_new != sub_size:
                    # the subtitle a step smaller on one line, so that the goal keeps its place on the slide (C5)
                    set_text_size(s_el, s_new)
                    st = _bookends(ctx)
                    if st.sub_size and st.sub_size > s_new:
                        st.sub_size = s_new
                    ctx.warnings.append(f"подзаголовок обложки набран {s_new:g} пт вместо {sub_size:g}: под ним встала цель")
                set_element_pos(s_el, x=s_box.x, y=s_box.y, w=s_box.w, h=s_box.h)
                if g_new != g_size:
                    set_text_size(g_el, g_new)
                    g_size = g_new
                set_element_pos(g_el, x=cand.x, y=cand.y, w=cand.w, h=cand.h)
                for p in g_el.findall(q("p:txBody") + "/" + q("a:p")):
                    pPr = p.find(q("a:pPr"))
                    if pPr is not None and s_el is not None:
                        sp = s_el.find(q("p:txBody") + "/" + q("a:p") + "/" + q("a:pPr"))
                        if sp is not None and sp.get("algn"):
                            pPr.set("algn", sp.get("algn"))
                if s_color:
                    style_runs(g_el, None, s_color, align=None)
                _recolor_for_ground(ctx, g_el, s_color, g_size, False, [carrier.style.color_hex if carrier is not None else None, own_sub])
                placed = True
        if not placed:
            # a goal that neither stands calm under the heading nor follows the subtitle to the sample's place (a goal
            # left alone under the heading without its subtitle reads as a second subtitle): it leaves the stack and
            # looks for a calm panel of the layout once every other line of the cover stands (`_goal_on_panel`), else
            # it is said aloud (speaker notes)
            ctx.goal_to_place = goal or shape_text(g_el)
            sid = ctx.id_of(g_el)
            remove_element(g_el)
            if sid:
                ctx.filled.discard(sid)
            g_el, goal_gone = None, True
    return s_moved, g_el, goal_gone


def _goal_after_subtitle(ctx: _SlideCtx, oslide: OutlineSlide, title_box: Bbox, s_el: etree._Element, s_box0, home: Bbox, g_el: etree._Element, goal: str, g_size: float, sub_size: float, family: Optional[str], foot: int) -> Optional[tuple[Bbox, Bbox, float, float]]:
    """The band fit of a subtitle that went back to the sample's own place and the goal following it (gate 2, C5 — LO
    Vivid long): the goal at its size under the subtitle; else a smaller goal on one line (down to a step over the
    small print); else the subtitle a step smaller (down to 0.8 of it, one line) with the goal under it — the room of
    the cover's small print under the goal kept free when the slide has one. The block stands in the middle of the
    sample's subtitle box, else at its top, else on the foot; every line calm and clear of the heading. Returns
    (subtitle box, goal box, subtitle size, goal size) or None when the band really lacks the room."""
    H = ctx.H
    hb = Bbox(x=s_box0[0], y=s_box0[1], w=s_box0[2], h=s_box0[3])
    s_ins, g_ins = _body_insets(s_el), _body_insets(g_el)
    s_text = shape_text(s_el).strip()
    scale = sorted({float(x) for x in ctx.grow_scale} | {float(x) for x in (ctx.typo.sizes_used or [])})
    has_note = bool(" ".join((oslide.footnote or "").split()))
    n_size = _small_print_size(sub_size, H, ctx.grow_scale, scale, ctx.typo.size_for("caption", 10.0)) if has_note else 0.0
    note_room = int((1.2 + 0.25) * n_size * EMU_PER_PT) + int(0.015 * H) + g_ins[1] + g_ins[3] if has_note else 0
    g_floor = max(n_size + 0.5, 0.018 * H / EMU_PER_PT) if has_note else 0.018 * H / EMU_PER_PT
    g_opts = [g_size] + sorted((x for x in scale if g_floor - 0.05 <= x < g_size - 0.05), reverse=True)
    s_opts = [sub_size] + sorted((x for x in scale if 0.8 * sub_size - 0.05 <= x < sub_size - 0.05), reverse=True)
    s_room = max((hb.w - s_ins[0] - s_ins[2]) / EMU_PER_PT * 0.92, 10.0)
    g_room = max((hb.w - g_ins[0] - g_ins[2]) / EMU_PER_PT * 0.92, 10.0)
    for si, ss in enumerate(s_opts):
        s_lines = max(1, len(display_lines(s_text, family, ss, False, s_room)))
        if si and s_lines > 1:
            continue  # a subtitle set smaller is set to stand on one line
        hh = int((s_lines * 1.2 + 0.25) * ss * EMU_PER_PT) + s_ins[1] + s_ins[3]
        for gi, gs in enumerate(g_opts):
            g_lines = max(1, len(display_lines(goal, family, gs, False, g_room)))
            if gi and g_lines > 1:
                continue  # a goal set smaller is set to stand on one line
            gh = int((g_lines * 1.2 + 0.25) * gs * EMU_PER_PT) + g_ins[1] + g_ins[3]
            gap = int(max(0.5 * gs, 0.012 * H / EMU_PER_PT) * EMU_PER_PT)
            total = hh + gap + gh + note_room
            tops = [home.y] if si == 0 else []
            tops += [hb.y + (hb.h - total) // 2, hb.y, foot - total]
            for top in tops:
                sb = Bbox(x=home.x, y=max(top, 0), w=home.w, h=hh)
                gb = Bbox(x=home.x, y=sb.y2 + gap, w=home.w, h=gh)
                if gb.y2 + note_room > foot or sb.intersection(title_box) > 0 or gb.intersection(title_box) > 0:
                    continue
                if sb.y < title_box.y2 and sb.y2 > title_box.y:
                    continue
                if _calm_at(ctx, sb) and _calm_at(ctx, gb):
                    return sb, gb, ss, gs
    return None


def _wrap_to_widths(words: list[str], family: Optional[str], size: float, bold: bool, widths_pt: list[float]) -> Optional[list[str]]:
    """The words set greedily into lines of their own widths (line i at most widths_pt[i]); None when they need more
    lines than there are widths, or a word is wider than its line."""
    lines: list[str] = []
    cur = ""
    for w in words:
        if len(lines) >= len(widths_pt):
            return None
        cand = f"{cur} {w}" if cur else w
        if text_width_pt(cand, family, size, bold) <= widths_pt[len(lines)]:
            cur = cand
            continue
        if not cur:
            return None
        lines.append(cur)
        cur = w
        if len(lines) >= len(widths_pt) or text_width_pt(cur, family, size, bold) > widths_pt[len(lines)]:
            return None
    if cur:
        lines.append(cur)
    return lines


def _box_gap(a: Bbox, b: Bbox) -> float:
    dx = max(0, max(a.x, b.x) - min(a.x2, b.x2))
    dy = max(0, max(a.y, b.y) - min(a.y2, b.y2))
    return math.hypot(dx, dy)


def _goal_on_panel(ctx: _SlideCtx, goal: str, family: Optional[str], sizes: list[float], prefs: list[Optional[str]], near: list[Bbox], hidden: list[Bbox], src: etree._Element) -> Optional[etree._Element]:
    """A cover's goal line that has no calm place under the heading or beside its subtitle stands on a painted panel
    of the layout (a triangle, a band) near the stack rather than going to the speaker notes (round 4.1, C3-1 — LO
    Focus long: the purple triangle under the date). Every line is measured against the panel's run at its own height
    with 1 % of the slide's width to spare (a slanted edge moves with the line; the lines follow it), inside the safe
    area, clear of every shape and picture of the slide, on one calm ground and quiet on the sample's own render, and
    set in the first of `prefs` (the heading's colour first) that reads there as small text. Sizes are tried from the
    largest: the first size that fits anywhere wins, and the place nearest the stack (`near`). Returns the new text
    element, or None when no panel holds it."""
    from verstka.rendering.layers import drawn_layers, ground_under, panel_run

    W, H = ctx.W, ctx.H
    safe = ctx.manifest.tokens.spacing.safe_area
    sx0, sx1, sy0, sy1 = int(safe.x * W), int(safe.x2 * W), int(safe.y * H), int(safe.y2 * H)
    pad, clear, step = int(0.01 * W), int(0.012 * H), max(int(0.01 * H), 1)
    try:
        layers = drawn_layers(ctx.slide)
    except Exception:  # noqa: BLE001 - layers are advice
        return None
    panels = [lay for lay in layers if lay.source in ("layout", "master") and not lay.placeholder and lay.paints and lay.kind == "sp"
              and not lay.picture and not lay.busy and lay.fill_hex and not lay.has_text and 0.03 <= lay.cover < 0.85]
    if not panels:
        return None
    obstacles: list[Bbox] = []
    for el in ctx.slide._element.cSld.find(q("p:spTree")):
        tag = etree.QName(el).localname
        if tag not in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp"):
            continue
        if tag == "sp" and not shape_text(el).strip() and not has_visible_style(el):
            continue
        b = element_bbox(el)
        if b:
            obstacles.append(_turned_box(el, Bbox(x=b[0], y=b[1], w=b[2], h=b[3])))
    for lay in layers:
        if lay.source != "slide" and not lay.placeholder and lay.paints and (lay.picture or lay.has_text or lay.kind in ("pic", "graphicFrame", "grpSp")) and lay.cover < 0.6:
            obstacles.append(lay.box)
    for c in ctx.manifest.tokens.chrome:
        if c.source == "background":
            b = c.bbox.to_emu(W, H)
            if 0 < b.w * b.h < 0.6 * W * H:
                obstacles.append(b)
    words = [w for w in bind_short_words(goal).split(" ") if w]
    ins = _body_insets(src)
    for size in sizes:
        pitch = int(1.2 * size * EMU_PER_PT)
        need = 3.0 if size >= 18 else 4.5
        best: Optional[tuple] = None
        for lay in panels:
            color = next((c for c in prefs if c and contrast_ratio(c, lay.fill_hex) >= need), None)
            if color is None:
                continue
            y_lo, y_hi = max(lay.box.y, sy0 + ins[1]), min(lay.box.y2, sy1 - ins[3])
            for n in (1, 2, 3, 4):
                y = y_lo
                while y + n * pitch <= y_hi:
                    rows = []
                    for i in range(n):
                        r = panel_run(lay, y + i * pitch, y + (i + 1) * pitch)
                        if not r:
                            break
                        rows.append(max(r, key=lambda ab: ab[1] - ab[0]))
                    if len(rows) == n:
                        for align in ("l", "r"):
                            if align == "l":
                                x0 = max(max(a for a, _ in rows) + pad, sx0 + ins[0])
                                edges = [(x0, min(b - pad, sx1 - ins[2])) for _, b in rows]
                            else:
                                x1 = min(min(b for _, b in rows) - pad, sx1 - ins[2])
                                edges = [(max(a + pad, sx0 + ins[0]), x1) for a, _ in rows]
                            widths = [(b - a) / EMU_PER_PT * 0.95 for a, b in edges]
                            if min(widths) < 4 * size:
                                continue
                            lines = _wrap_to_widths(words, family, size, False, widths)
                            if lines is None or len(lines) != n:
                                continue
                            boxes = []
                            for (a, b), ln in zip(edges, lines):
                                w_ln = int(text_width_pt(ln, family, size, False) * EMU_PER_PT)
                                bx = a if align == "l" else b - w_ln
                                boxes.append(Bbox(x=bx, y=y + len(boxes) * pitch, w=max(w_ln, 1), h=pitch))
                            bx0, bx1 = min(b.x for b in boxes), max(b.x2 for b in boxes)
                            block = Bbox(x=bx0, y=y, w=bx1 - bx0, h=n * pitch)
                            room = Bbox(x=block.x - pad, y=block.y - clear, w=block.w + 2 * pad, h=block.h + 2 * clear)
                            if any(ob.intersection(room) > 0 for ob in obstacles):
                                continue
                            dist = min((_box_gap(block, nb) for nb in near), default=0.0)
                            if best is not None and dist >= best[0]:
                                continue
                            if not all(_calm_at(ctx, b) for b in boxes):
                                continue
                            try:
                                under = [ground_under(ctx.slide, b) for b in boxes]
                            except Exception:  # noqa: BLE001
                                continue
                            if any(u is None or u[1].el is not lay.el for u in under) or not _quiet_on_render(ctx, block, hidden):
                                continue
                            best = (dist, n, y, align, lines, edges, color, block)
                    y += step
        if best is None:
            continue
        _dist, n, y, align, lines, edges, color, block = best
        el = _new_text(ctx, src, "Цель", family, color)
        fill_text(el, [ParagraphSpec(" ".join(lines), bullet=False)], size_pt=size)
        _set_paragraph_box(el, 0, align, bold=False, line_spacing=1.0)
        _write_lines(el, lines)
        style_runs(el, None, color, align=None)
        _no_autofit(el, "t")
        x0 = min(a for a, _ in edges) if align == "l" else block.x
        x1 = block.x2 if align == "l" else max(b for _, b in edges)
        slack = int(0.5 * size * EMU_PER_PT)  # a face a little wider than measured keeps its lines
        if align == "l":
            set_element_pos(el, x=x0 - ins[0], y=y - ins[1], w=x1 - x0 + ins[0] + ins[2] + slack, h=int((n * 1.2 + 0.25) * size * EMU_PER_PT) + ins[1] + ins[3])
        else:
            set_element_pos(el, x=x0 - ins[0] - slack, y=y - ins[1], w=x1 - x0 + ins[0] + ins[2] + slack, h=int((n * 1.2 + 0.25) * size * EMU_PER_PT) + ins[1] + ins[3])
        sid = ctx.id_of(el)
        if sid:
            ctx.filled.add(sid)
        return el
    return None


def _calm_photo_lines(ctx: _SlideCtx, text: str, family: Optional[str], size: float, bold: bool, align: str, x: int, y: int, w: int, ins: tuple[int, int, int, int], marl: int = 0, max_lines: int = 3) -> Optional[tuple]:
    """A cover line set on a photo of the template keeps to the photo's calm part (round 4.1, C3-2 — LO Candy: the
    subtitle ran from the blurred table onto the candies): each line of the text as set in the box (x, y, w) is
    measured against the calm stretch of the photo at its own height (`layers.photo_calm_run`, the audit's cells).
    None when the text is not on a photo or already keeps to its calm part; else (lines, box width) of the fewest lines
    (up to `max_lines`) whose every line keeps to the calm stretch at its own height, with 1 % of the slide's width to
    spare — or () when no such wrap exists (the caller may try a smaller size)."""
    from verstka.rendering.layers import photo_calm_run

    W = ctx.W
    pad, tol = int(0.01 * W), int(0.012 * W)
    pitch = int(1.2 * size * EMU_PER_PT)
    inner_w = max(w - ins[0] - ins[2] - marl, 1)
    x0 = x + ins[0] + marl  # where the letters start (left) / the text column

    def ink(ln: str, i: int, col_w: int) -> Bbox:
        lw = int(text_width_pt(ln, family, size, bold) * EMU_PER_PT)
        if align == "ctr":
            a = x0 + (inner_w - lw) // 2
        elif align == "r":
            a = x0 + inner_w - lw
        else:
            a = x0
        return Bbox(x=a, y=y + ins[1] + i * pitch, w=max(lw, 1), h=pitch)

    def anchor(b: Bbox) -> int:
        return b.x + b.w // 2 if align == "ctr" else (b.x2 - 1 if align == "r" else b.x)

    def runs(boxes: list[Bbox]) -> Optional[list[Optional[tuple[int, int]]]]:
        try:
            return [photo_calm_run(ctx.slide, b, anchor(b)) for b in boxes]
        except Exception:  # noqa: BLE001 - layers are advice
            return None

    def keeps(boxes: list[Bbox], got: list[Optional[tuple[int, int]]]) -> bool:
        return all(r is None or (b.x >= r[0] - tol and b.x2 <= r[1] + tol) for b, r in zip(boxes, got))

    lines = display_lines(text, family, size, bold, inner_w / EMU_PER_PT * 0.92) or [text]
    boxes = [ink(ln, i, inner_w) for i, ln in enumerate(lines)]
    got = runs(boxes)
    if got is None or all(r is None for r in got) or keeps(boxes, got):
        return None
    words = [wd for wd in text.split(" ") if wd]
    for n in range(max(len(lines), 2), max_lines + 1):
        # the width the calm stretch leaves on each of the n rows, from where the lines start (or around their centre):
        # each line takes the room of its own row, as the photo's calm part narrows or widens with the height
        widths: list[int] = []
        for i in range(n):
            rb = Bbox(x=x0, y=y + ins[1] + i * pitch, w=inner_w, h=pitch)
            try:
                r = photo_calm_run(ctx.slide, rb, x0 + inner_w // 2 if align == "ctr" else (x0 + inner_w - 1 if align == "r" else x0))
            except Exception:  # noqa: BLE001
                r = None
            lim = inner_w
            if r is not None:
                if align == "ctr":
                    c = x0 + inner_w // 2
                    lim = min(lim, 2 * min(c - r[0], r[1] - c) - 2 * pad)
                elif align == "r":
                    lim = min(lim, x0 + inner_w - r[0] - pad)
                else:
                    lim = min(lim, r[1] - x0 - pad)
            widths.append(lim)
        if min(widths[:1]) < 0.12 * W:
            return ()
        cand = _wrap_to_widths(words, family, size, bold, [max(wd, 1) / EMU_PER_PT * 0.95 for wd in widths])
        if not cand:
            continue
        new_w = max(widths[: len(cand)]) + ins[0] + ins[2] + marl
        c_boxes = []
        for i, ln in enumerate(cand):
            lw = int(text_width_pt(ln, family, size, bold) * EMU_PER_PT)
            a = x0 + (inner_w - lw) // 2 if align == "ctr" else (x0 + inner_w - lw if align == "r" else x0)
            c_boxes.append(Bbox(x=a, y=y + ins[1] + i * pitch, w=max(lw, 1), h=pitch))
        c_got = runs(c_boxes)
        if c_got is not None and keeps(c_boxes, c_got):
            return cand, int(new_w)
    return ()


def _keep_off_photo(ctx: _SlideCtx, el: etree._Element, text: str, family: Optional[str], size: float, align: str, scale: list[float], what: str, marl: int = 0) -> Optional[int]:
    """`_calm_photo_lines` applied to a placed cover line (its box, insets and alignment as set): the line is set again
    on the lines that keep to the photo's calm part — at its size, else a step smaller (down to 0.85 of it). Returns the
    box's new bottom (EMU), or None when nothing changed."""
    b = element_bbox(el)
    if not b:
        return None
    ins = _body_insets(el)
    if _calm_photo_lines(ctx, text, family, size, False, align, b[0], b[1], b[2], ins, marl, max_lines=1) is None:
        return None  # not on a photo, or already on its calm part
    # two lines at its size, else two lines a step smaller (down to 0.85), else three lines: a subtitle of three
    # short lines over a photo reads worse than one a step smaller on two
    sizes = [size] + sorted((z for z in scale if 0.85 * size - 0.05 <= z < size - 0.05), reverse=True)[:2]
    tries = [(sz, 2) for sz in sizes] + [(sz, 3) for sz in sizes]
    for sz, n_max in tries:
        got = _calm_photo_lines(ctx, text, family, sz, False, align, b[0], b[1], b[2], ins, marl, max_lines=n_max)
        if not got:
            continue
        lines, new_w = got
        if sz != size:
            set_text_size(el, sz)
        _write_lines(el, _split_like(text, lines))
        h = int((len(lines) * 1.2 + 0.25) * sz * EMU_PER_PT) + ins[1] + ins[3]
        x = b[0] + (b[2] - new_w) // 2 if align == "ctr" else (b[0] + b[2] - new_w if align == "r" else b[0])
        set_element_pos(el, x=x, w=new_w, h=h)
        ctx.warnings.append(f"{what} обложки набран{'а' if what == 'цель' else ''} в {len(lines)} строки{'' if sz == size else f' ({sz:g} пт)'}: одной строкой {'она заходила' if what == 'цель' else 'он заходил'} на пёструю часть фотографии шаблона")
        return b[1] + h
    return None


def _note_to_speaker_notes(oslide: OutlineSlide, text: str) -> None:
    """A line the slide has no room for is said aloud: it joins the slide's speaker notes (written into the notes page
    after the render), once — a re-render of the same slide does not repeat it."""
    text = " ".join((text or "").split())
    if not text:
        return
    notes = (oslide.notes or "").strip()
    if text in notes:
        return
    try:
        oslide.notes = (notes + "\n" + text).strip()
    except Exception:  # noqa: BLE001 - a frozen outline keeps the line off the slide rather than failing the render
        pass


def _quiet_on_render(ctx: _SlideCtx, box: Bbox, hidden: list[Bbox]) -> bool:
    """Nothing is drawn inside the box on the sample's own render (its texts masked): no stripe, dot or line work a
    shape does not measure. True when the render is missing (the layers decide alone)."""
    import numpy as np

    dm = _detail_map(ctx, hidden)
    if dm is None:
        return True
    act, kx, ky = dm
    y0, y1 = max(int(box.y * ky), 0), min(int(box.y2 * ky) + 1, act.shape[0])
    x0, x1 = max(int(box.x * kx), 0), min(int(box.x2 * kx) + 1, act.shape[1])
    if y1 <= y0 or x1 <= x0:
        return True
    return float(np.percentile(act[y0:y1, x0:x1], 98)) <= 25.0


def _note_beside_stack(ctx: _SlideCtx, stack: list[Optional[etree._Element]], note: str, family: Optional[str], size: float, bold: bool, ins: tuple[int, int, int, int], hidden: list[Bbox], foot: int) -> Optional[tuple[int, int, int, list[str], int, str]]:
    """A cover's small print with no room under its stack stands in the column of a line set beside the stack (the
    template's date right of a rule, LO Grey Elegant — gate 2, C1): under that line, on its text edge and in its width,
    when the note fits above the foot in at most three lines, on one calm ground, clear of everything drawn and of the
    sample's own art. Returns (text x, top, width, lines, height, alignment) or None."""
    W, H = ctx.W, ctx.H
    safe = ctx.manifest.tokens.spacing.safe_area
    boxes = [b for b in (element_bbox(e) for e in stack if e is not None) if b]
    if not boxes:
        return None
    sx0, sx1 = min(b[0] for b in boxes), max(b[0] + b[2] for b in boxes)
    sy0, sy1 = min(b[1] for b in boxes), max(b[1] + b[3] for b in boxes)
    own = {id(e) for e in stack if e is not None}
    tree = ctx.slide._element.cSld.find(q("p:spTree"))
    cands: list[tuple[etree._Element, Bbox]] = []
    for el in tree:
        if id(el) in own or etree.QName(el).localname != "sp" or not shape_text(el).strip():
            continue
        ph = _ph(el)
        if ph is not None and ph.get("type") == "sldNum":
            continue  # a page number is no column
        b = element_bbox(el)
        if not b:
            continue
        bx = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
        if bx.w < 0.12 * W or bx.y2 <= sy0 or bx.y >= sy1:
            continue
        if bx.x >= sx1 - int(0.005 * W) or bx.x2 <= sx0 + int(0.005 * W):
            cands.append((el, bx))
    if not cands:
        return None
    drawn = _drawn_boxes(ctx, exclude=set())
    for el, bx in sorted(cands, key=lambda c: c[1].y):
        sh = next((x for x in ctx.slide.shapes if x._element is el), None)
        e_ins = effective_insets(sh) if sh is not None else _body_insets(el)  # the layout's zero insets count
        x = bx.x + e_ins[0]
        w = min(bx.x2 - e_ins[2], int(safe.x2 * W)) - x
        if w < 0.1 * W:
            continue
        lines = display_lines(bind_short_words(note), family, size, bold, w / EMU_PER_PT * 0.92)
        if not lines or len(lines) > 3:
            continue
        h = int((1.2 * len(lines) + 0.25) * size * EMU_PER_PT) + ins[1] + ins[3]
        top = bx.y2 + int(0.01 * H)
        nb = Bbox(x=x - ins[0], y=top, w=w + ins[0] + ins[2], h=h)
        if nb.y2 - ins[3] > foot or any(o.intersection(nb) > 0 for o in drawn):
            continue  # the letters keep above the foot (the box's own bottom inset may reach past it)
        # the line's own letters on the sample's render are no art (their blur reaches a few pixels under its box)
        own_text = Bbox(x=bx.x, y=bx.y, w=bx.w, h=bx.h + int(0.012 * H))
        if not _calm_at(ctx, nb) or not _quiet_on_render(ctx, nb, hidden + [own_text]):
            continue
        algn = next((lv.get("algn") for lv in _style_levels(ctx, el) if lv.get("algn")), None)
        return x, top, w, lines, h, algn if algn in ("ctr", "r") else "l"
    return None


def _note_under_goal(ctx: _SlideCtx, oslide: OutlineSlide, g_el: etree._Element, goal: str, g_size: float, family: Optional[str], align: str, t_el: etree._Element, s_el: Optional[etree._Element], hidden: list[Bbox], obstacles: list[Bbox], note: tuple[int, int, int, float], floor: int, drop_goal: bool = True) -> tuple[Optional[int], Optional[etree._Element]]:
    """Room for a cover's small print under its goal line when there is none at the foot (G1-01): (1) a goal that
    wrapped needlessly («…со 120 000 до 255 000 / рублей») is set on one line in the width free at its own height (up to
    the art or what stands to its right), and the note goes under it; (2) else (with `drop_goal`) the goal is said
    aloud (speaker notes) and the note takes its place under the subtitle — the brief asks for the small print on the
    slide, the goal is also on the notes page. Returns (the note's top or None, the goal element or None when it
    left)."""
    W, H = ctx.W, ctx.H
    safe = ctx.manifest.tokens.spacing.safe_area
    n_x, n_w, n_h, n_size = note
    gb = element_bbox(g_el)
    if not gb:
        return None, g_el
    gap = max(int(0.04 * H), int(1.2 * n_size * EMU_PER_PT))

    def fits_at(top: int) -> bool:
        nb = Bbox(x=n_x, y=top, w=n_w, h=n_h)
        return nb.y2 <= floor and not any(o.intersection(nb) > 0 for o in obstacles)

    g_ins = _body_insets(g_el)
    one_h = int((1.2 + 0.25) * g_size * EMU_PER_PT) + g_ins[1] + g_ins[3]
    if align == "l" and gb[3] > one_h * 1.3:
        room = int(safe.x2 * W)
        for ob in obstacles:
            if ob.y < gb[1] + one_h and ob.y2 > gb[1] and ob.x > gb[0] + int(0.1 * W):
                room = min(room, ob.x - int(0.02 * W))
        art = _art_edge(ctx, (gb[1], gb[1] + one_h), gb[0] + int(0.05 * W), hidden)
        if art is not None:
            room = min(room, art - int(0.02 * W))
        room_pt = (room - gb[0] - g_ins[0] - g_ins[2]) / EMU_PER_PT
        top = gb[1] + one_h + gap
        if room_pt > 0 and len(display_lines(goal, family, g_size, False, room_pt * 0.92)) == 1 and fits_at(top):
            need = int(text_width_pt(goal, family, g_size, False) * 1.1 * EMU_PER_PT) + g_ins[0] + g_ins[2]
            set_element_pos(g_el, w=min(room - gb[0], max(gb[2], need)), h=one_h)
            ctx.warnings.append("цель обложки набрана в одну строку: под ней встала сноска")
            return top, g_el
    if not drop_goal:
        return None, g_el
    # the note takes the goal's own place under the subtitle (the stack's gap is already there); never closer to the
    # heading and the subtitle standing over it than a line's distance
    prev = 0
    for el in (t_el, s_el):
        b = element_bbox(el) if el is not None else None
        if b and b[1] < gb[1]:
            prev = max(prev, b[1] + b[3])
    top = max(gb[1], prev + int(0.5 * n_size * EMU_PER_PT)) if prev else gb[1]
    if not fits_at(top):
        return None, g_el
    ctx.goal_to_place = goal  # a calm panel of the layout, once the cover's lines stand; else the speaker notes
    gid = ctx.id_of(g_el)
    remove_element(g_el)
    if gid:
        ctx.filled.discard(gid)
    ctx.warnings.append("цель обложки уступила место сноске под подзаголовком: ниже для сноски нет места")
    return top, None


def _small_print_color(ctx: _SlideCtx, box: Bbox, color: Optional[str]) -> Optional[str]:
    """The colour of a cover's small print where it stands (G1-18): its line colour when that reads at the small size
    (4.5:1); on a light ground where it does not (orange on white), the template's dark text colour — text.primary, else
    the darkest colour of the palette that reads (black counts: it is the template's own). None: keep the colour."""
    g = _ground_hex(ctx, box)
    if not g or not color or contrast_ratio(color, g) >= 4.5 or relative_luminance(g) < 0.5:
        return None
    toks = ctx.manifest.tokens
    primary = toks.color_for("text.primary")
    cands = [primary, toks.color_for("text.secondary")] + [c.hex for c in toks.colors] + ["000000"]
    good = [c for c in cands if c and contrast_ratio(c, g) >= 4.5]
    if not good:
        return None
    return primary if primary in good else min(good, key=relative_luminance)


def _small_print_size(ref: float, H: int, grow_scale: list[float], sizes: list[float], caption: float) -> float:
    """The size of a cover's small print («мелким текстом», gate 2 — C3): a step under the subtitle `ref` as before
    (0.62 of it, never under the caption size or 1.8 % of the slide height, on the template's scale), but never above
    the scale step below the subtitle, 0.7 of the subtitle, nor 2.6 % of the slide height — a template whose caption
    is 20 pt (Synth Sidebar) or a slide twice the usual size (26.67″) no longer sets it at the subtitle's size. On the
    template's scale where a size of it lies between the floor and that cap."""
    h_pt = H / EMU_PER_PT
    floor = 0.018 * h_pt
    size = _snap_up(max(_snap_down(0.62 * ref, grow_scale), caption, floor), sizes)
    below = [x for x in sizes if x < ref - 0.05]
    cap = min(0.7 * ref, 0.026 * h_pt, max(below) if below else ref)
    if size <= cap + 0.05:
        return size
    on_scale = [x for x in sizes if floor - 0.05 <= x <= cap + 0.05]
    if on_scale:
        return max(on_scale)
    return max(round(cap * 2) / 2, round(floor * 2 + 0.49) / 2)


def _snap_up(size: float, sizes: list[float]) -> float:
    """The smallest size of the template's scale at or above `size` (a floor derived from the slide height lands on
    the scale, not between its steps); `size` itself when the scale has nothing that large."""
    up = [s for s in sizes if s >= size - 0.25]
    return min(up) if up else size


def _snap_display(size: float, sizes: list[float]) -> float:
    """A display size of the template's own scale when one is close below (38.4 pt → 36 pt), else an even point
    size (45.6 pt → 44 pt): no heading is set at 45.6 pt."""
    on_scale = _snap_down(size, sizes)
    if on_scale >= 0.88 * size:
        return on_scale
    if size >= 24:
        return float(int(size) // 2 * 2)  # 45.6 → 44, 57 → 56
    return float(int(size)) if size >= 12 else round(size, 1)


def _absolute_bbox(el: etree._Element) -> Optional[tuple[int, int, int, int]]:
    """The box of an element in slide coordinates, through every group it sits in (their chOff/chExt scaling)."""
    b = element_bbox(el)
    if b is None:
        return None
    x, y, w, h = (float(v) for v in b)
    parent = el.getparent()
    while parent is not None and etree.QName(parent).localname == "grpSp":
        xfrm = parent.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
        if xfrm is None:
            break
        off, ext = xfrm.find(q("a:off")), xfrm.find(q("a:ext"))
        choff, chext = xfrm.find(q("a:chOff")), xfrm.find(q("a:chExt"))
        if off is None or ext is None:
            break
        gx, gy, gw, gh = (float(off.get("x") or 0), float(off.get("y") or 0), float(ext.get("cx") or 0), float(ext.get("cy") or 0))
        cx, cy = (float(choff.get("x") or 0), float(choff.get("y") or 0)) if choff is not None else (gx, gy)
        cw, ch = (float(chext.get("cx") or 0), float(chext.get("cy") or 0)) if chext is not None else (gw, gh)
        sx, sy = (gw / cw if cw else 1.0), (gh / ch if ch else 1.0)
        x, y, w, h = gx + (x - cx) * sx, gy + (y - cy) * sy, w * sx, h * sy
        parent = parent.getparent()
    return int(x), int(y), int(w), int(h)


def _lift_out_of_group(el: etree._Element) -> None:
    """Move an element out of the groups it sits in onto the slide itself, at the same place (its box in slide
    coordinates): what is written and placed in it afterwards is placed on the slide."""
    parent = el.getparent()
    if parent is None or etree.QName(parent).localname != "grpSp":
        return
    b = _absolute_bbox(el)
    top = parent
    while top.getparent() is not None and etree.QName(top.getparent()).localname == "grpSp":
        top = top.getparent()
    parent.remove(el)
    top.addnext(el)
    if b is not None:
        set_element_pos(el, x=b[0], y=b[1], w=b[2], h=b[3])


def _text_over_pictures(ctx: _SlideCtx) -> None:
    """A cover whose sample sets its title behind a picture (a word behind a can of soda, a product shot over a big
    word — a design of the sample's own short title): the deck's longer words would stand hidden behind it. Every text
    the slide carries is brought in front of the pictures and drawings that cover a quarter of it."""
    tree = next(iter(ctx.slide._element.iter(q("p:spTree"))), None)
    if tree is None:
        return
    kids = [c for c in tree if isinstance(c.tag, str) and etree.QName(c).localname in ("sp", "pic", "grpSp", "graphicFrame")]
    texts = [c for c in kids if etree.QName(c).localname == "sp" and "".join(t.text or "" for t in c.iter(q("a:t"))).strip()]
    moved = 0
    for el in texts:
        b = element_bbox(el)
        if not b or b[2] <= 0 or b[3] <= 0:
            continue
        area = b[2] * b[3]
        after = kids[kids.index(el) + 1:] if el in kids else []
        for o in after:
            if etree.QName(o).localname not in ("pic", "grpSp") or o.find(".//" + q("a:blip")) is None:
                continue
            ob = element_bbox(o)
            if not ob:
                continue
            ox = max(0, min(b[0] + b[2], ob[0] + ob[2]) - max(b[0], ob[0]))
            oy = max(0, min(b[1] + b[3], ob[1] + ob[3]) - max(b[1], ob[1]))
            if ox * oy >= 0.25 * area:
                tree.remove(el)
                tree.append(el)
                moved += 1
                break
    if moved:
        ctx.warnings.append(f"{moved} text block(s) brought in front of the pictures that covered them")


def _render_bookend(ctx: _SlideCtx, oslide: OutlineSlide) -> tuple[float, int]:
    """Cover, divider, closing slide: the heading set large on the sample's own ground and art — up to three
    balanced lines at no less than 0.85 of the sample size where the room allows, never under a logo or a picture;
    the subtitle right under it on the heading's left edge; a kicker over it (the audience on a cover, the section
    number on a divider) and the date as a cover's footer; speaker and QR blocks, empty frames and pager dots that
    have nothing to show leave. Sizes are decided for the deck: the cover is the loudest heading, dividers share
    one size and one top line, the closing slide is never louder than the cover. Returns (size, lines)."""
    W, H = ctx.W, ctx.H
    typo = ctx.typo
    st = _bookends(ctx)
    kind = oslide.kind
    title = max(ctx.slots(SlotRole.title), key=lambda s: s.bbox.area)
    t_el = ctx.els[title.shape_id]
    head, sub_text, kicker_text, footer_text = _bookend_texts(ctx, oslide)
    goal = bind_short_words(cover_goal(oslide) or "") or None
    carrier = _subtitle_carrier(ctx, title) if sub_text else None
    s_el = ctx.els.get(carrier.shape_id) if carrier is not None else None
    if s_el is not None and has_visible_style(s_el) and len(sub_text) > max(20, 2 * len((carrier.sample_text or "").strip())):
        # a painted button («Ссылка» on the WorkSpace call-to-action) is no place for a sentence: it leaves with the
        # sample's other boxes and the subtitle is set as a plain line under the heading
        carrier, s_el = None, None
    for el_ in (t_el, s_el):
        if el_ is not None:
            _lift_out_of_group(el_)  # a heading or a subtitle grouped with a bar is placed in slide coordinates
    t_box0 = _ensure_xfrm(ctx, t_el)
    s_box0 = _ensure_xfrm(ctx, s_el) if s_el is not None else None
    if not t_box0:
        ctx.fill_slot(title, [ParagraphSpec(oslide.headline)], min_ratio=0.7)
        return 0.0, 0
    keep = {title.shape_id} | ({carrier.shape_id} if carrier is not None else set())
    cleared, frames = _strip_bookend(ctx, keep, oslide)
    # an empty frame of the sample holds the kicker (on a cover: the audience and the date), else it leaves
    frame_text = " · ".join(t for t in (kicker_text, footer_text) if t) if kind == PatternKind.title else kicker_text
    frame_el = _pick_frame(ctx, frames) if frame_text else None
    _drop_frames(ctx, [f for f in frames if f is not frame_el], cleared)
    if frame_el is not None:
        kicker_text, footer_text = None, None
    dot_rows = _pagination(ctx, oslide)
    tb = Bbox(x=t_box0[0], y=t_box0[1], w=t_box0[2], h=t_box0[3])
    size0 = title.style.size_pt or (int(_eff_rpr(ctx, t_el, "sz") or 0) / 100) or typo.size_for("display", 40.0)
    family = title.style.font_family or typo.primary_family
    bold = (_eff_rpr(ctx, t_el, "b") or "0") in ("1", "true") or bool(title.style.bold)
    if kind != PatternKind.title and st.cover_bold is not None:
        bold = st.cover_bold  # the deck's bookends speak in the cover's weight
    align = title.style.align if title.style.align in ("l", "ctr", "r") else "l"
    max_lines = bookend_max_lines(kind, title)
    text = bind_short_words(head)
    marl = _first_line_offset(ctx, t_el)
    ls = _eff_line_spacing(ctx, t_el)
    anchor = _eff_anchor(ctx, t_el)
    # the heading's own letter-spacing (a badge's title tracked 8 pt a letter — Office «Badge») counts in every width its
    # lines are measured at: measured without it, «ТОЧКА КОФЕ» broke inside a word
    try:
        t_spc = max(int(_eff_rpr(ctx, t_el, "spc") or 0) / 100.0, 0.0)
    except ValueError:
        t_spc = 0.0
    safe = ctx.manifest.tokens.spacing.safe_area
    h1 = typo.size_for("h1", 0.0)
    obstacles = _drawn_boxes(ctx, exclude={id(t_el)} | ({id(s_el)} if s_el is not None else set()))

    # ---- the size this heading aims at, decided for the deck
    display_sizes = sorted({s.style.size_pt for p in ctx.manifest.patterns if p.kind in _BOOKENDS for s in p.slots
                            if s.role == SlotRole.title and s.style.size_pt and (not h1 or s.style.size_pt <= 2.5 * h1)})
    snap_sizes = sorted(set(ctx.grow_scale) | {float(x) for x in (typo.sizes_used or [])} | set(display_sizes))
    target = size0
    if kind == PatternKind.section and st.cover_size and h1 and size0 <= 1.3 * h1:
        # a divider sample set barely above the slide headings (29 pt against 24) on an empty page reads unfinished:
        # it takes a display size a clear step below the cover
        up = [s for s in snap_sizes if size0 < s <= 0.8 * st.cover_size + 0.05]
        if up:
            target = max(up)
    if kind != PatternKind.title and st.cover_size:
        target = min(target, st.cover_size)

    s_size = 0.0
    if sub_text:
        # the sample's own subtitle size; a speaker line standing in for it is set at the template's subtitle size
        own = carrier is not None and not is_placeholder_text(carrier.sample_text)
        s_size = (carrier.style.size_pt if own and carrier.style.size_pt else None) or (int(_eff_rpr(ctx, s_el, "sz") or 0) / 100 if own and s_el is not None else 0.0) or _bookend_subtitle_size(ctx)
        # a step below the heading, never below the template's h2: a cover subtitle is read from across the room
        s_size = min(max(s_size, typo.size_for("h2", s_size)), _snap_down(target * 0.5, ctx.grow_scale))
        if st.sub_size:
            s_size = min(st.sub_size, max(_snap_down(target * 0.5, ctx.grow_scale), typo.size_for("h2", 16.0)))  # one subtitle size for the deck
    def kicker_size(title_size: float) -> float:
        if kind == PatternKind.section:
            return max(_snap_down(0.75 * title_size, snap_sizes), typo.size_for("h2", 16.0))  # a display numeral
        return st.sub_size or s_size or _bookend_subtitle_size(ctx)

    k_size = kicker_size(target) if (kicker_text or frame_text) else 0.0

    # the vertical room: under the logos above the heading, over what stands below it (pager dots, footer logos)
    col_x2 = tb.x + max(tb.w, int(0.3 * W))
    if align == "l":
        # a heading set from the left reaches as far as its words on one line at the target size: a mark above the far
        # end of a wide sample box (a square in the top right corner of LibreOffice «Inspiration») does not hold a
        # short heading down under it
        reach = tb.x + int(text_width_pt(text, family, target, bold) * 1.05 * EMU_PER_PT) + int(0.03 * W)
        col_x2 = min(col_x2, max(reach, tb.x + int(0.3 * W)))
    top_lim, bot_lim = int(safe.y * H), int(safe.y2 * H)
    mid = tb.y + tb.h // 2
    below: list[Bbox] = []
    for ob in obstacles:
        if ob.x2 <= tb.x or ob.x >= col_x2:
            continue
        if ob.y2 <= mid:
            top_lim = max(top_lim, ob.y2 + int(max(0.05 * H, 0.6 * target * EMU_PER_PT)))
        elif ob.y >= mid:
            bot_lim = min(bot_lim, ob.y - int(0.03 * H))
            below.append(ob)
    # the panel the heading stands on (a teal band over a navy ground): the whole stack keeps to it — a subtitle set
    # past its edge lands on another ground in a colour chosen for this one
    panel = _ground_panel(ctx, tb)
    if panel is not None:
        pad = int(0.03 * H)
        if panel.y2 < int(safe.y2 * H):
            bot_lim = min(bot_lim, panel.y2 - pad)
            below.append(Bbox(x=0, y=panel.y2 - pad + int(0.03 * H), w=W, h=max(H - panel.y2, 1)))
        if panel.y > int(safe.y * H):
            top_lim = max(top_lim, panel.y + pad)

    def floor_under(x_right: int) -> int:
        """How low text reaching `x_right` may come: over what stands below it (art to the right of the text's
        actual end — the WorkSpace stripes — does not hold it up)."""
        lim = int(safe.y2 * H)
        for ob in below:
            if ob.x < x_right and ob.x2 > tb.x:
                lim = min(lim, ob.y - int(0.03 * H))
        return lim
    # where the heading may come to stand: grown from the sample's anchor by up to `max_lines` lines
    size_est = min(target, 2.5 * h1) if h1 else target  # a 144 pt «Q&A» will not stay 144 pt
    grow = int(max_lines * size_est * 1.2 * ls * EMU_PER_PT)
    extra = int(2.5 * s_size * EMU_PER_PT) if sub_text else 0
    if anchor == "b":
        band = (tb.y2 - grow, tb.y2 + extra)
    elif anchor == "ctr":
        band = (mid - grow // 2, mid + grow // 2 + extra)
    else:
        band = (tb.y, tb.y + grow + extra)
    band = (max(band[0], top_lim), min(band[1], bot_lim))
    # horizontal room: from the heading's left edge to the first picture, logo or piece of art standing to its right
    limit = int(safe.x2 * W)
    for ob in obstacles:
        if ob.y < band[1] and ob.y2 > band[0] and ob.x >= tb.x + int(0.12 * W):
            limit = min(limit, ob.x - int(0.02 * W))
    hidden = cleared + _sample_text_boxes(ctx, [title] + ([carrier] if carrier is not None else []))
    if s_box0:
        hidden.append(Bbox(x=s_box0[0], y=s_box0[1], w=s_box0[2], h=s_box0[3]))
    # a heading printed on a band stays in the band: the band (its edges, its shadow) is the heading's ground, never art
    # the heading must keep clear of — read off the render, the light green band under the title of LibreOffice
    # «Inspiration» pushed the heading down across its lower edge and cut its width to a fifth of the slide
    try:
        from verstka.rendering.layers import heading_band

        hb = heading_band(ctx.slide, tb)
    except Exception:  # noqa: BLE001 - layers are advice
        hb = None
    if hb is not None:
        pad = int(0.01 * H)
        hidden.append(Bbox(x=hb.box.x - pad, y=hb.box.y - pad, w=hb.box.w + 2 * pad, h=hb.box.h + 2 * pad))
    hidden.extend(_veil_boxes(ctx))  # see-through bubbles the heading may run over, as the template's own heading does
    art = _art_edge(ctx, band, tb.x + int(0.12 * W), hidden, designed_to=tb.x2)
    if art is not None and kind == PatternKind.title and art - int(0.02 * W) - tb.x < 0.6 * tb.w:
        # the art over the whole band the heading might grow into cuts the sample's own heading box (a triangle
        # narrowing towards its tip, LO Focus): the heading stands at the wide top — the edge over the two lines the
        # sample's box holds at the target size, where the heading really stands (G1-17)
        ins2 = _body_insets(t_el)
        two = (max(band[0], tb.y), min(band[1], tb.y + max(tb.h, int(2 * 1.2 * ls * size_est * EMU_PER_PT) + ins2[1] + ins2[3])))
        art2 = _art_edge(ctx, two, tb.x + int(0.12 * W), hidden, designed_to=tb.x2) if two[1] > two[0] else art
        if art2 is None or art2 > art:
            art = art2
    if art is not None:
        limit = min(limit, art - int(0.02 * W))  # a gutter from the art, not a hairline
    art_low = None
    if kind == PatternKind.title:
        # art painted under the heading (a centred cover over its picture): the whole stack keeps above it — the
        # heading steps down a size rather than running over the art
        art_low = _art_below(ctx, tb.x + tb.w // 10, tb.x2 - tb.w // 10, tb.y + tb.h // 2, hidden)
        if art_low is not None:
            # the sample's own texts reach over the edge of its art (a subtitle on the cube's rings): the stack may go
            # as low as they do, a little more, never onto the body of the art
            own = max((sl.bbox.to_emu(W, H).y2 for sl in ctx.pattern.slots if sl.role in (SlotRole.title, SlotRole.subtitle)), default=0)
            art_low = max(art_low, own + int(0.05 * H))
            bot_lim = min(bot_lim, art_low)
            below.append(Bbox(x=tb.x, y=art_low + int(0.03 * H), w=tb.w, h=max(int(safe.y2 * H) - art_low, 1)))
    w_sample = max(min(tb.w, limit - tb.x), int(0.2 * W))
    w_max = max(w_sample, min(limit - tb.x, max(tb.w, int(0.6 * W) - tb.x)))
    col_x0 = tb.x
    if kind == PatternKind.title and align == "ctr":
        # a centred cover's column grows on both sides of the sample's centre, as far as the nearer side allows (a
        # subtitle of 50 characters on one line under «История VK», not two lines running onto the cube — G1-17/G1-20)
        c_mid = tb.x + tb.w // 2
        half = min(c_mid - int(safe.x * W), limit - c_mid)
        w_ctr = min(2 * half, max(tb.w, int(0.72 * W)))
        if w_ctr > w_max:
            w_max = w_ctr
            col_x0 = c_mid - w_max // 2
    top_lim_col = top_lim  # the room above without the art a line as wide as `w_max` would meet
    above = None if hb is not None else _art_above(ctx, col_x0, col_x0 + w_max, tb.y + min(tb.h // 2, int(0.08 * H)), hidden)
    if above is not None:
        top_lim = max(top_lim, above + int(max(0.05 * H, 0.6 * target * EMU_PER_PT)))
    t_ins = _body_insets(t_el)
    word_room = 0.75 if bold else 0.8

    def inner(w: int) -> float:
        # 5 % of the line is kept free: the face that renders the deck (the template font, or a substitute where it is
        # not installed) is never exactly the one measured
        return max((w - t_ins[0] - t_ins[2] - marl) / EMU_PER_PT * 0.95, 10.0)

    def floor_of(ref: float) -> float:
        # a word keeps a fifth of the line free: a face wider than the measured one may wrap the heading once more,
        # but never breaks a word in two. A giant sample size (a 144 pt «Q&A») may step down as far as the h1 size.
        return max(0.2, min(0.45, typo.size_for("h1", 0.45 * ref) / ref))

    def fit_lines(txt: str, n: int, ref: float) -> tuple[float, int, int]:
        s1, n1 = display_fit(txt, family, ref, bold, inner(w_sample), n, ctx.grow_scale, floor=floor_of(ref), word_room=word_room, spc_pt=t_spc)
        if s1 >= 0.85 * ref - 0.05 or w_max <= w_sample:
            return s1, n1, w_sample
        s2, n2 = display_fit(txt, family, ref, bold, inner(w_max), n, ctx.grow_scale, floor=floor_of(ref), word_room=word_room, spc_pt=t_spc)
        return (s2, n2, w_max) if s2 > s1 else (s1, n1, w_sample)

    def breaks(txt: str, sz: float, w: int, n_min: int = 1) -> list[str]:
        n = max(len(display_lines(txt, family, sz, bold, inner(w), spc_pt=t_spc)), n_min)
        return (balanced_lines(txt, family, sz, bold, inner(w), n, spc_pt=t_spc) if n > 1 else None) or display_lines(txt, family, sz, bold, inner(w), spc_pt=t_spc)

    def clean(txt: str, got: tuple[float, int, int]) -> tuple[float, int, int]:
        """A heading torn between an adjective and its noun («Умные / напоминания») is set a step smaller on fewer
        lines when that breaks it cleanly — down to 0.85 of the size it had."""
        sz, n, w = got
        if n < 2 or not awkward_breaks(breaks(txt, sz, w)):
            return got
        cands = sorted({s for s in snap_sizes if 0.85 * sz - 0.05 <= s < sz} | {round(sz * k, 1) for k in (0.95, 0.9, 0.85)}, reverse=True)
        for c in cands:
            for ww in (w, w_max):
                lines_c = breaks(txt, c, ww)
                if len(lines_c) <= n and not awkward_breaks(lines_c) and not any(text_width_pt(x, family, c, bold) > inner(ww) * word_room for x in re.split(r"[\s\u00a0]+", txt) if x):
                    return c, len(lines_c), ww
        return got

    def fit(txt: str, ref: float) -> tuple[float, int, int]:
        best = fit_lines(txt, max_lines, ref)
        if kind == PatternKind.title and h1 and best[0] < 1.25 * h1:
            # a very long title in a narrow column: a fourth line keeps it a display heading, three lines would set it
            # at the size of a slide heading
            four = fit_lines(txt, max_lines + 1, ref)
            if four[0] >= 1.12 * best[0]:
                return clean(txt, four)
        return clean(txt, best)

    def heads_of(k: PatternKind) -> list[str]:
        out = []
        for s in ctx.outline.slides:
            if s.kind == k and s.headline:
                h = split_display_title(s.headline)[0] if k == PatternKind.title else " ".join(s.headline.split())
                out.append(bind_short_words(h))
        return out

    # capitals are a deck-wide decision: every cover, divider and closing heading of the deck must keep its display
    # size in them, or none is set in capitals (one long title would otherwise mix the two)
    caps = False
    if _caps_convention(ctx, t_el, title) and len(text.upper()) == len(text):
        if st.caps is None:
            heads = set(heads_of(PatternKind.title) + heads_of(PatternKind.section) + heads_of(PatternKind.thanks)) | {text}
            st.caps = True
            for h in heads:
                c_size, c_lines, _ = fit_lines(h.upper(), 3, target)
                if c_size < 0.85 * target - 0.05 or c_lines > 3:
                    st.caps = False
                    break
        caps = st.caps
    # a heading the template already sets in capitals (cap="all" in its own run, list or title style) renders wider
    # than its lower-case letters: it is measured the way it shows, or «ЧАШКИ» wraps onto the subtitle
    shown_caps = caps or _inherits_caps(ctx, t_el)
    measured = text.upper() if shown_caps else text

    size, lines, w_used = fit(measured, target)
    if kind == PatternKind.title:
        # a short cover title grows to the largest display size of the template's covers and dividers, on at most two
        # lines inside the room the sample's three lines would take
        for cand in sorted((s for s in display_sizes if s > target + 0.05), reverse=True):
            s_c, n_c, w_c = fit_lines(measured, 2, cand)
            if s_c >= cand - 0.05 and n_c * cand <= 0.9 * max_lines * target and not awkward_breaks(breaks(measured, s_c, w_c)):
                size, lines, w_used = s_c, n_c, w_c
                break
    elif kind == PatternKind.section:
        if st.divider_size is None or st.divider_pattern != ctx.pattern.id:
            # every divider of the deck: one size (the one the longest heading allows) and one top line
            fits = [fit(h.upper() if caps else h, target) for h in heads_of(PatternKind.section) or [text]]
            shared = _snap_display(min(f[0] for f in fits), snap_sizes)
            log.debug("dividers: %s → %g", fits, shared)
            n_max = 1
            for h in heads_of(PatternKind.section) or [text]:
                m = h.upper() if caps else h
                n_s = len(display_lines(m, family, shared, bold, inner(w_sample), spc_pt=t_spc))
                if n_s > max_lines:
                    n_s = len(display_lines(m, family, shared, bold, inner(w_max), spc_pt=t_spc))
                n_max = max(n_max, n_s)
            st.divider_size, st.divider_pattern, st.divider_lines = shared, ctx.pattern.id, n_max
        size = min(size, st.divider_size) if st.divider_pattern == ctx.pattern.id else size
        if len(display_lines(measured, family, size, bold, inner(w_sample), spc_pt=t_spc)) <= max_lines:
            w_used = w_sample
        else:
            w_used = max(w_used, w_max)
    size = _snap_display(size, snap_sizes)
    if kind != PatternKind.title and st.cover_size:
        size = min(size, st.cover_size)

    # the sample's own line break is part of its design («Спасибо / за внимание»): a heading of about its length keeps it
    sample = (title.sample_text or "").replace("\x0b", "\n")
    sample_lines = sample.count("\n") + 1 if sample.strip() else 1
    min_lines = sample_lines if kind != PatternKind.title and sample_lines > 1 and len(head) <= 1.4 * len(sample) and len(head.split()) >= sample_lines else 1

    # ---- colours of the lines around the heading
    ground = _ground_hex(ctx, tb)
    dark = relative_luminance(ground) < 0.4 if ground else ctx.pattern.family.value == "dark"
    own_sub = (carrier.style.color_hex if carrier is not None else None) or title.style.color_hex or _readable_on(ground, [], ctx.manifest)
    sub_color = st.sub_color if st.sub_color and st.sub_dark == dark else own_sub
    if sub_color and ground and contrast_ratio(sub_color, ground) < 3.0:
        sub_color = _readable_on(ground, [own_sub, title.style.color_hex], ctx.manifest)
    if st.sub_color is None and sub_color and (sub_text or kicker_text or footer_text):
        st.sub_color, st.sub_dark = sub_color, dark
    if kind == PatternKind.section:
        accents = [a for a in ctx.manifest.tokens.accents() if not ground or contrast_ratio(a, ground) >= 3.0]
        k_color = accents[0] if accents else (title.style.color_hex or sub_color)
    else:
        k_color = sub_color

    def legible(sz: float, color: Optional[str], bold: bool = False) -> float:
        """A line in a colour that holds only 3:1 against the ground (white on a brand blue) is set as large text:
        18 pt at least, or 14 pt in bold (WCAG); small text needs 4.5:1."""
        big = 14.0 if bold else 18.0
        if color and ground and contrast_ratio(color, ground) < 4.5 and sz < big:
            return big
        return sz

    def low_contrast(color: Optional[str]) -> bool:
        return bool(color and ground and contrast_ratio(color, ground) < 4.5)

    # ---- lines, heights, the vertical stack: kicker, heading, subtitle
    sample_gap = (s_box0[1] - (tb.y + tb.h)) / EMU_PER_PT if s_box0 else None
    s_ins = _body_insets(s_el) if s_el is not None else t_ins
    k_ins = s_ins if kind != PatternKind.section else t_ins  # the kicker is drawn like the subtitle, a section number like the heading
    k_lines_room = inner(w_max)
    kicker_above = bool(kicker_text) and text_width_pt(kicker_text, family, k_size, bold and kind == PatternKind.section) <= k_lines_room
    if kicker_text and not kicker_above and kind == PatternKind.title:
        footer_text = " · ".join(t for t in (kicker_text, footer_text) if t)
        kicker_text = None
    y = tb.y
    h_title = h_sub = gap = h_k = gap_k = h_goal = gap_goal = 0
    sub_size = g_size = 0.0
    split: list[str] = []
    for _ in range(14):
        n = len(display_lines(measured, family, size, bold, inner(w_used), spc_pt=t_spc))
        if n < min_lines:
            n = min_lines
        split = (balanced_lines(measured, family, size, bold, inner(w_used), n, spc_pt=t_spc) if n > 1 else None) or display_lines(measured, family, size, bold, inner(w_used), spc_pt=t_spc)
        lines = len(split)
        h_title = int((lines * 1.2 * ls + 0.15) * size * EMU_PER_PT) + t_ins[1] + t_ins[3]
        if sub_text:
            sub_size = legible(max(min(s_size, _snap_down(size * 0.6, ctx.grow_scale)), min(s_size, typo.size_for("caption", 10.0))), sub_color)
            s_lines = max(1, len(display_lines(bind_short_words(sub_text), family, sub_size, False, inner(w_max) * 0.92)))
            h_sub = int((s_lines * 1.2 + 0.25) * sub_size * EMU_PER_PT) + s_ins[1] + s_ins[3]
            want = sample_gap if sample_gap is not None and carrier is not None and not is_placeholder_text(carrier.sample_text) else 0.45 * size
            gap = int(max(0.35 * size, min(want, 0.6 * size)) * EMU_PER_PT)
        if goal:
            # the goal: a short line a step under the subtitle, in its colour, never under the caption size
            ref = sub_size or min(s_size or _bookend_subtitle_size(ctx), _snap_down(size * 0.6, ctx.grow_scale))
            # (a template whose caption is set at 20 pt, or a slide twice the usual size, does not lift it to the
            # subtitle's own size: the caption floor reaches 0.8 of the subtitle at most — gate 2, C3)
            g_size = legible(_snap_up(max(_snap_down(0.8 * ref, ctx.grow_scale), min(typo.size_for("caption", 10.0), 0.8 * ref), 0.022 * H / EMU_PER_PT), snap_sizes), sub_color)
            g_lines = max(1, len(display_lines(goal, family, g_size, False, inner(w_max) * 0.92)))
            h_goal = int((g_lines * 1.2 + 0.25) * g_size * EMU_PER_PT) + s_ins[1] + s_ins[3]
            gap_goal = int(max(0.5 * g_size, 0.012 * H / EMU_PER_PT) * EMU_PER_PT) if sub_text else int(0.45 * size * EMU_PER_PT)
        k_size = legible(kicker_size(size), k_color) if (kicker_text or frame_text) else 0.0
        if kicker_above:
            h_k = int((1.2 + 0.25) * k_size * EMU_PER_PT) + k_ins[1] + k_ins[3]
            gap_k = max(int(max(0.3 * size, 0.5 * k_size) * EMU_PER_PT) - k_ins[3] - t_ins[1], 0)
        k_room = h_k + gap_k if kicker_above else 0
        body = h_title + gap + h_sub + gap_goal + h_goal
        reach = max(text_width_pt(t, family, size, bold) for t in split)
        if sub_text:
            reach = max(reach, min(text_width_pt(sub_text, family, sub_size, False), inner(w_max)))
        if goal:
            reach = max(reach, min(text_width_pt(goal, family, g_size, False), inner(w_max)))
        bot_room = floor_under(tb.x + t_ins[0] + marl + int(reach * EMU_PER_PT) + int(0.02 * W) if align == "l" else tb.x + w_used)
        ref_h = h_title
        if kind == PatternKind.section and st.divider_pattern == ctx.pattern.id and st.divider_lines > lines:
            ref_h = int((st.divider_lines * 1.2 * ls + 0.15) * size * EMU_PER_PT) + t_ins[1] + t_ins[3]
        if anchor == "b":
            y = tb.y2 - ref_h
        elif anchor == "ctr":
            y = mid - ref_h // 2
        else:
            y = tb.y
        y = max(min(y, bot_room - body), top_lim + k_room)
        if y + body <= bot_room + int(0.01 * H):
            break
        if kicker_above:
            # the kicker gives way before the heading shrinks: on a cover it joins the footer, a divider drops its number
            kicker_above = False
            if kind == PatternKind.title:
                footer_text = " · ".join(t for t in (kicker_text, footer_text) if t)
            kicker_text = None
            continue
        if size <= 0.5 * size0:
            break
        smaller = [s for s in snap_sizes if 0.85 * size <= s < size - 0.05]
        size = max(smaller) if smaller else _snap_display(size * 0.92, snap_sizes)
    if kind == PatternKind.title:
        st.cover_size, st.cover_bold = size, bold
    if sub_text and not st.sub_size:
        st.sub_size = sub_size
    x = tb.x if align == "l" else (tb.x + tb.w // 2 - w_used // 2 if align == "ctr" else tb.x2 - w_used)
    if kind == PatternKind.title:
        # every line of the heading on its own piece of the ground: a triangle's slanted edge moves with the height, so
        # a centred second line may start on the white gap beside it (LO Focus — gate 2, C2): the box shifts, else
        # the heading steps down a size or two
        dx = _heading_shift_on_ground(ctx, split, family, size, bold, align, x, y, w_used, t_ins, marl, ls)
        if dx is None:
            for smaller in sorted((s_ for s_ in snap_sizes if 0.8 * size - 0.05 <= s_ < size - 0.05), reverse=True)[:3]:
                n_s = max(len(display_lines(measured, family, smaller, bold, inner(w_used), spc_pt=t_spc)), min_lines)
                split_s = (balanced_lines(measured, family, smaller, bold, inner(w_used), n_s, spc_pt=t_spc) if n_s > 1 else None) or display_lines(measured, family, smaller, bold, inner(w_used), spc_pt=t_spc)
                dx_s = _heading_shift_on_ground(ctx, split_s, family, smaller, bold, align, x, y, w_used, t_ins, marl, ls)
                if dx_s is not None:
                    ctx.warnings.append(f"заголовок обложки набран {smaller:g} пт вместо {size:g}: строки не помещались на своей части рисунка шаблона")
                    size, split, lines, dx = smaller, split_s, len(split_s), dx_s
                    h_title = int((lines * 1.2 * ls + 0.15) * size * EMU_PER_PT) + t_ins[1] + t_ins[3]
                    st.cover_size = size
                    break
        if dx:
            x += dx
    fill_text(t_el, [ParagraphSpec(text, bullet=False)], size_pt=size)
    _write_lines(t_el, _split_like(text, split))
    _set_paragraph_box(t_el, marl, align, caps=caps, bold=bold, line_spacing=ls)
    t_color = title.style.color_hex
    # judged where the letters stand: a title box wider than the band it is printed on (a light green band of half the
    # width under a short white title — LibreOffice «Inspiration») straddles the band's end, the letters do not
    ink = _lines_ink(split, family, size, bold, align, x, y, w_used, t_ins, marl, ls) or Bbox(x=x, y=y, w=w_used, h=h_title)
    g_ink = _ground_hex(ctx, ink) or ground
    if t_color and g_ink and contrast_ratio(t_color, g_ink) < 2.5 and not _mixed_ground(ctx, ink):
        # the sample's heading colour is unreadable on the ground it stands on here (white on a light teal band): a
        # colour of the template that reads on it — the heading's own colour is kept from 2.5:1 up, as the template's
        # design
        better = _readable_on(g_ink, [sub_color, own_sub], ctx.manifest)
        if better and contrast_ratio(better, g_ink) >= 3.0:
            style_runs(t_el, None, better, align=None)
            ctx.warnings.append(f"цвет заголовка #{t_color} на #{g_ink} ({contrast_ratio(t_color, g_ink):.1f}:1) заменён на #{better}")
    _no_autofit(t_el, "t")
    set_element_pos(t_el, x=x, y=y, w=w_used, h=h_title)
    ctx.filled.add(title.shape_id)
    if size < 0.85 * size0 - 0.05:
        ctx.warnings.append(f"заголовок набран {size:g} пт ({lines} стр.) при {size0:g} пт в образце")
    log.debug("bookend %s p%d: sample %g target %g → %g pt × %d, y %.3f, room y %.3f–%.3f x→%.3f (w %.3f/%.3f)", kind.value, ctx.pattern.source_slide, size0, target, size, lines, y / H, top_lim / H, bot_lim / H, limit / W, w_sample / W, w_max / W)
    text_x = x + t_ins[0] + marl  # where the heading's letters start: every other line starts there too
    block_bottom = y + h_title
    if s_el is None and sub_text:
        s_el = _new_text(ctx, t_el, "Подзаголовок", family, sub_color)
    if s_el is not None and sub_text:
        fill_text(s_el, [ParagraphSpec(bind_short_words(sub_text), bullet=False)], size_pt=sub_size)
        _set_paragraph_box(s_el, marl, align, bold=False, line_spacing=1.0)
        if sub_color:
            style_runs(s_el, None, sub_color, align=None)
        _no_autofit(s_el, "t")
        s_w = w_max - (t_ins[0] - s_ins[0])
        s_x = tb.x + t_ins[0] - s_ins[0] if align == "l" else (x + w_used // 2 - s_w // 2 if align == "ctr" else x + w_used - s_w)
        set_element_pos(s_el, x=s_x, y=y + h_title + gap, w=s_w, h=h_sub)
        if kind == PatternKind.title and sub_color:
            # the subtitle on another ground than the heading (under a heading on white, on the layout's blue band —
            # LO Vivid): its colour is checked where it stands
            _recolor_for_ground(ctx, s_el, sub_color, sub_size, False, [own_sub, title.style.color_hex])
        block_bottom = y + h_title + gap + h_sub
        # on a photo of the template the subtitle keeps to its calm part: more lines, not a run onto the candies (C3-2)
        s_bottom = _keep_off_photo(ctx, s_el, bind_short_words(sub_text), family, sub_size, align, ctx.grow_scale, "подзаголовок", marl)
        if s_bottom is not None:
            h_sub = s_bottom - (y + h_title + gap)
            block_bottom = s_bottom
        sid = ctx.id_of(s_el)
        if sid:
            ctx.filled.add(sid)
    g_el = k_el = None
    if goal:
        # the cover's goal under the subtitle (on the subtitle's edge and in its colour), drawn like it
        g_src = s_el if s_el is not None else t_el
        g_el = _new_text(ctx, g_src, "Цель", family, sub_color)
        fill_text(g_el, [ParagraphSpec(goal, bullet=False)], size_pt=g_size)
        _set_paragraph_box(g_el, marl, align, bold=False, line_spacing=1.0)
        if sub_color:
            style_runs(g_el, None, sub_color, align=None)
        _no_autofit(g_el, "t")
        g_ins = _body_insets(g_el)
        g_w = w_max - (t_ins[0] - g_ins[0])
        g_x = tb.x + t_ins[0] - g_ins[0] if align == "l" else (x + w_used // 2 - g_w // 2 if align == "ctr" else x + w_used - g_w)
        g_y = block_bottom + gap_goal
        g_lines = max(1, len(display_lines(goal, family, g_size, False, inner(w_max) * 0.92)))
        set_element_pos(g_el, x=g_x, y=g_y, w=g_w, h=int((g_lines * 1.2 + 0.25) * g_size * EMU_PER_PT) + g_ins[1] + g_ins[3])
        if kind == PatternKind.title and sub_color:
            _recolor_for_ground(ctx, g_el, sub_color, g_size, False, [own_sub, title.style.color_hex])
        block_bottom = g_y + int((g_lines * 1.2 + 0.25) * g_size * EMU_PER_PT) + g_ins[1] + g_ins[3]
        g_bottom = _keep_off_photo(ctx, g_el, goal, family, g_size, align, ctx.grow_scale, "цель", marl)
        if g_bottom is not None:
            block_bottom = g_bottom
        gid = ctx.id_of(g_el)
        if gid:
            ctx.filled.add(gid)
    moved_sub: Optional[Bbox] = None
    if kind == PatternKind.title and ((s_el is not None and sub_text) or g_el is not None):
        s_moved, g_el, goal_gone = _lines_on_one_ground(ctx, oslide, Bbox(x=x, y=y, w=w_used, h=h_title), s_el if sub_text else None, s_box0, carrier, g_el, goal, g_size, sub_size, family, own_sub, int(safe.y2 * H))
        mb = element_bbox(s_el) if s_moved and s_el is not None else None
        if mb:
            moved_sub = Bbox(x=mb[0], y=mb[1], w=mb[2], h=mb[3])  # the date and the small print keep off it
            gbx = element_bbox(g_el) if g_el is not None else None
            if gbx:
                # the goal followed the subtitle: the small print goes under both
                x0, y0 = min(moved_sub.x, gbx[0]), min(moved_sub.y, gbx[1])
                moved_sub = Bbox(x=x0, y=y0, w=max(moved_sub.x2, gbx[0] + gbx[2]) - x0, h=max(moved_sub.y2, gbx[1] + gbx[3]) - y0)
        if s_moved or goal_gone:
            # the heading's column now ends where what stayed in it ends
            block_bottom = y + h_title
            for el in (None if s_moved else (s_el if sub_text else None), g_el):
                b = element_bbox(el) if el is not None else None
                if b and b[0] < x + w_used and b[0] + b[2] > x:
                    block_bottom = max(block_bottom, b[1] + b[3])

    from verstka.rendering.fonts import is_pinned

    def roomy(tight: int, room: int) -> int:
        """A left-aligned line box of a face the deck was not tuned with takes the room it may use, not the width
        measured with 10 % slack: the renderer may set the face heavier than it was measured (Futura regular shown as
        its medium cut) and must not wrap a one-line note onto a second line past the slide's foot."""
        return max(tight, room) if align == "l" and not is_pinned(family) else tight

    def line_box(txt: str, sz: float, name: str, color: Optional[str], weight: bool, box: tuple[int, int, int], anchor_to: str = "t", al: Optional[str] = None) -> etree._Element:
        sz = legible(sz, color, weight)
        src = s_el if s_el is not None and kind != PatternKind.section else t_el
        el = _new_text(ctx, src, name, family, color)
        rows = txt.split("\n")
        fill_text(el, [ParagraphSpec(t, bullet=False) for t in rows], size_pt=sz)
        _set_paragraph_box(el, 0, al or ("l" if align == "l" else align), bold=weight, line_spacing=1.0)
        style_runs(el, None, color, align=None)
        _no_autofit(el, anchor_to)
        ins = _body_insets(el)
        bx, by, bw = box
        set_element_pos(el, x=bx - ins[0], y=by, w=bw + ins[0] + ins[2], h=int((1.2 * len(rows) + 0.25) * sz * EMU_PER_PT) + ins[1] + ins[3])
        if kind == PatternKind.title:
            # a line at the foot may stand on another ground than the heading (a note on the green triangle under a
            # heading on the red one): its colour is checked where it stands
            _recolor_for_ground(ctx, el, color, sz, weight, [own_sub, title.style.color_hex])
        sid = ctx.id_of(el)
        if sid:
            ctx.filled.add(sid)
        return el

    # the kicker over the heading
    if kicker_above and kicker_text:
        k_w = int(text_width_pt(kicker_text, family, k_size, bold and kind == PatternKind.section) * 1.1 * EMU_PER_PT) + int(0.02 * W)
        k_el = line_box(kicker_text, k_size, "Номер раздела" if kind == PatternKind.section else "Надзаголовок", k_color, bold and kind == PatternKind.section, (text_x, y - gap_k - h_k, roomy(min(k_w, w_max), w_max - (text_x - x))))
    # a frame of the sample holds the kicker at its inner bottom edge, on the heading's letter line
    if frame_el is not None and frame_text:
        fb = element_bbox(frame_el)
        f_size = k_size
        heavy = bold and kind == PatternKind.section
        pad = int(max(0.035 * H, 0.6 * f_size * EMU_PER_PT))
        fx2 = min(fb[0] + fb[2], W) - pad
        f_lines = [frame_text]
        if text_x + int(text_width_pt(frame_text, family, f_size, heavy) * 1.1 * EMU_PER_PT) > fx2 and " · " in frame_text:
            f_lines = frame_text.split(" · ")  # the audience over the date
        f_w = max(int(text_width_pt(t, family, f_size, heavy) * 1.1 * EMU_PER_PT) for t in f_lines)
        f_h = int((1.2 * len(f_lines) + 0.25) * f_size * EMU_PER_PT)
        fy = min(fb[1] + fb[3], H) - pad - f_h
        if text_x + f_w <= fx2 and fy >= max(fb[1], 0) + pad // 2:
            line_box("\n".join(f_lines), f_size, "Номер раздела" if kind == PatternKind.section else "Надзаголовок", k_color, heavy, (text_x, fy, fx2 - text_x))
        else:
            _drop_frames(ctx, [frame_el], cleared)
            if kind == PatternKind.title:
                footer_text = frame_text
    # the slide's footnote (Agent v2: a cover's disclaimer «Все цифры условные»): small print at the very foot of the
    # heading's column, under the date; with no room there, a line under the subtitle
    note = " ".join((oslide.footnote or "").split())
    note_top: Optional[int] = None
    beside: Optional[tuple] = None  # the small print in the column beside the stack (C1)
    beside_al: Optional[str] = None
    if note:
        ref = sub_size or s_size or _bookend_subtitle_size(ctx)
        n_size = _small_print_size(ref, H, ctx.grow_scale, snap_sizes, typo.size_for("caption", 10.0))
        # «мелким текстом»: on a dark or brand ground where the subtitle's colour lacks contrast for small text (white
        # on VK blue) the small print is set at 14 pt bold — large text by WCAG — never at the subtitle's or the date's
        # 18 pt; on a light ground it keeps the small size and takes the template's dark text colour where it stands
        # (an orange heading on white: its notes are dark — G1-18, `_small_print_color` below)
        n_color = sub_color
        n_bold = low_contrast(sub_color) and n_size < 14.0 and not relative_luminance(ground) >= 0.5
        if n_bold:
            n_size = 14.0
        # the column under a heading that stands over art is no place for small print: it goes to the slide's foot
        # at the left margin, clear of the art
        n_x = int(safe.x * W) if art_low is not None else text_x
        n_room = int(safe.x2 * W)
        for ob in obstacles:
            if ob.y2 > int(0.7 * H) and ob.x > n_x + int(0.1 * W):
                n_room = min(n_room, ob.x - int(0.02 * W))
        n_art = _art_edge(ctx, (int(0.8 * H), int(safe.y2 * H)), n_x + int(0.05 * W), hidden)
        if n_art is not None:
            n_room = min(n_room, n_art - int(0.02 * W))
        n_width = max(min(w_max, n_room - n_x), int(0.2 * W))
        n_lines = display_lines(bind_short_words(note), family, n_size, n_bold, n_width / EMU_PER_PT * 0.92) or [note]
        n_w = int(max(text_width_pt(t, family, n_size, n_bold) for t in n_lines) * 1.1 * EMU_PER_PT)
        l_ins = s_ins if s_el is not None and kind != PatternKind.section else t_ins  # the insets line_box gives its box
        n_h = int((1.2 * len(n_lines) + 0.25) * n_size * EMU_PER_PT) + l_ins[1] + l_ins[3]
        if align == "ctr":
            # a centred stack keeps its small print centred under it, not at the left edge of its column; over art
            # (a band of waves at the foot) when the foot above the art has the room, else at the left margin
            cx = max(int(safe.x * W), x + w_used // 2 - min(n_w, n_width) // 2)
            if art_low is None or floor_under(cx + n_w + int(0.02 * W)) - n_h - block_bottom >= max(int(0.05 * H), int(1.2 * n_size * EMU_PER_PT)):
                n_x = cx
        ny = floor_under(n_x + n_w + int(0.02 * W)) - n_h
        # at the foot beside art the note still never runs into the heading block (the goal line) it shares columns with
        crosses_block = art_low is not None and n_x < x + w_used and n_x + n_w > x and ny < block_bottom + int(0.01 * H)
        if (ny - block_bottom < max(int(0.05 * H), int(1.2 * n_size * EMU_PER_PT)) and art_low is None) or crosses_block:
            ny = block_bottom + max(int(0.04 * H), int(1.2 * n_size * EMU_PER_PT))  # no foot room: under the subtitle
            # the small print under the block must not fall off the foot of the slide nor onto what stands under the
            # column (a logo): the whole stack rises for it as far as the room above the heading allows
            floor_n = int(safe.y2 * H)
            ob_floor = H - int(0.01 * H)  # what the note may never cross: the slide's foot, what stands under it
            for ob in obstacles:
                if ob.x < n_x + min(n_w, n_width) and ob.x2 > n_x and ob.y >= block_bottom:
                    floor_n = min(floor_n, ob.y - int(0.015 * H))
                    ob_floor = min(ob_floor, ob.y)
            lift = ny + n_h - floor_n
            k_room_now = h_k + gap_k if k_el is not None else 0
            room_up = y - (top_lim + k_room_now)
            if lift > max(room_up, 0):
                # the art above was looked for over the widest line the heading might have taken: over what the
                # stack's lines really span there may be more room (a moon to the right of a two-line heading)
                parts = [(t_el, max(text_width_pt(t, family, size, bold) for t in split), int(max(0.05 * H, 0.6 * size * EMU_PER_PT)))]
                if sub_text and s_el is not None:
                    parts.append((s_el, text_width_pt(bind_short_words(sub_text), family, sub_size, False), int(0.03 * H)))
                if g_el is not None:
                    parts.append((g_el, text_width_pt(goal or "", family, g_size, False), int(0.03 * H)))
                if k_el is not None:
                    parts.append((k_el, text_width_pt(kicker_text or "", family, k_size, False), int(0.03 * H)))
                room2 = None
                for el_p, w_pt, pad_p in parts:
                    b = element_bbox(el_p)
                    if not b:
                        continue
                    x2 = min(b[0] + b[2], text_x + int(w_pt * EMU_PER_PT) + int(0.02 * W))
                    a = _art_above(ctx, tb.x, x2, b[1], hidden)
                    lim = max(top_lim_col, a + pad_p) if a is not None else top_lim_col
                    room2 = b[1] - lim if room2 is None else min(room2, b[1] - lim)
                if room2 is not None:
                    room_up = max(room_up, room2 - k_room_now)
            if lift > 0 and room_up > 0:
                lift = min(lift, room_up)
                for el in (t_el, s_el if sub_text else None, g_el, k_el):
                    if el is not None:
                        b = element_bbox(el)
                        if b:
                            set_element_pos(el, y=b[1] - lift)
                ny -= lift
            if ny + n_h > ob_floor:
                # something of the template stands under the block (a short rule under the subtitle): the small print
                # goes under it when the slide has room there; else (a heading block at the foot under the template's
                # art) it is said aloud — the speaker notes, never off the slide or onto what stands below
                n_foot = min(int(safe.y2 * H) + int(0.01 * H), H - int(0.01 * H))
                spot = None
                for ob in sorted((o for o in obstacles if o.x < n_x + min(n_w, n_width) and o.x2 > n_x and o.y >= block_bottom), key=lambda o: o.y):
                    cand_n = Bbox(x=n_x, y=ob.y2 + int(0.015 * H), w=min(n_w, n_width), h=n_h)
                    if cand_n.y2 <= n_foot and not any(o.intersection(cand_n) > 0 for o in obstacles):
                        spot = cand_n.y
                        break
                if spot is None and g_el is not None and goal:
                    # the brief's small print («Все данные условные» on the first slide) is not given up for a goal
                    # line: first the goal takes one line where it wrapped needlessly («…255 000 / рублей»)
                    spot, g_el = _note_under_goal(ctx, oslide, g_el, goal, g_size, family, align, t_el, s_el if sub_text else None, hidden, obstacles, (n_x, min(n_w, n_width), n_h, n_size), ob_floor, drop_goal=False)
                if spot is None:
                    # then the column of a line standing beside the stack (the template's date right of its rule)
                    beside = _note_beside_stack(ctx, [t_el, s_el if sub_text else None, g_el, k_el], note, family, n_size, n_bold, l_ins, hidden, n_foot)
                    if beside is not None:
                        n_x, spot, n_width, n_lines, n_h, beside_al = beside
                        n_w = n_width
                        ctx.warnings.append("сноска обложки поставлена в колонку рядом с заголовком: под ним для неё нет места")
                if spot is None and beside is None and g_el is not None and goal:
                    # then the goal is said aloud and the small print takes its place under the subtitle
                    spot, g_el = _note_under_goal(ctx, oslide, g_el, goal, g_size, family, align, t_el, s_el if sub_text else None, hidden, obstacles, (n_x, min(n_w, n_width), n_h, n_size), ob_floor)
                if spot is None:
                    _note_to_speaker_notes(oslide, note)
                    ctx.warnings.append("сноска обложки перенесена в заметки докладчика: на слайде для неё нет места")
                    note = ""
                else:
                    ny = spot
                    if beside is None:
                        block_bottom = ny + n_h
            else:
                block_bottom = ny + n_h
        else:
            note_top = ny  # at the foot: the date stands over it
    if note and moved_sub is not None and moved_sub.intersection(Bbox(x=n_x, y=ny, w=min(n_w, n_width), h=n_h)) > 0:
        # the subtitle went back to the sample's own place at the foot (LO Vivid's band): the small print goes under
        # it when the slide has room there. Else it stays (the brief wants it on the slide): the audit sees the
        # overlap and the cover is rematched, never quietly emptied into the speaker notes
        under = moved_sub.y2 + int(0.015 * H)
        nb = Bbox(x=n_x, y=under, w=min(n_w, n_width), h=n_h)
        if nb.y2 <= min(int(safe.y2 * H) + int(0.01 * H), H - int(0.01 * H)) and not any(ob.intersection(nb) > 0 for ob in obstacles):
            ny = under
    if note:
        # small print never stands on the template's drawing (a tree at the foot of the column): it steps right of it
        for _ in range(4):
            nb = Bbox(x=n_x, y=ny, w=min(n_w, n_width), h=n_h)
            hit = next((ob for ob in obstacles if ob.intersection(nb) > 0.02 * nb.area), None)
            if hit is None:
                break
            nx2 = hit.x2 + int(0.02 * W)
            if nx2 + min(n_w, n_width) > int(safe.x2 * W):
                break
            n_x = nx2
        if not n_bold:
            n_color = _small_print_color(ctx, Bbox(x=n_x, y=ny, w=min(n_w, n_width), h=n_h), sub_color) or n_color
        line_box("\n".join(n_lines), n_size, "Сноска", n_color, n_bold, (n_x, ny, n_width if beside is not None else roomy(min(n_w, n_width), min(n_width, n_room - n_x))), al=beside_al)
    # the footer: the date (and whatever had no room above the heading) at the foot of the heading's column
    if footer_text:
        f_step = _snap_down(0.85 * (s_size or _bookend_subtitle_size(ctx)), ctx.grow_scale)
        f_size = legible(max(f_step, min(typo.size_for("caption", 10.0), f_step)), sub_color)  # a 20 pt caption does not set the date at the subtitle's size (C3)
        f_h = int((1.2 + 0.25) * f_size * EMU_PER_PT)
        f_w = int(text_width_pt(footer_text, family, f_size, False) * 1.1 * EMU_PER_PT)
        # a centred stack keeps its date centred under it, never at the edge of a slide-wide heading box
        f_x = max(int(safe.x * W), x + w_used // 2 - f_w // 2) if align == "ctr" else text_x
        fy = floor_under(f_x + f_w + int(0.02 * W)) - f_h
        if note_top is not None:
            f_ins = s_ins if s_el is not None and kind != PatternKind.section else t_ins
            fy = min(fy, note_top - int(0.6 * f_size * EMU_PER_PT) - f_h - f_ins[1] - f_ins[3])  # the date stands over the small print
        f_room = int(safe.x2 * W)
        for ob in obstacles:
            if ob.y < fy + f_h and ob.y2 > fy and ob.x > f_x:
                f_room = min(f_room, ob.x - int(0.02 * W))
        f_art = _art_edge(ctx, (fy, fy + f_h), f_x + int(0.05 * W), hidden)
        if f_art is not None:
            f_room = min(f_room, f_art - int(0.02 * W))
        clear = fy - block_bottom >= max(int(0.06 * H), int(1.5 * (sub_size or f_size) * EMU_PER_PT))
        if note_top is not None and not clear:
            # the small print took the foot: the date keeps a plain line's distance from the subtitle
            clear = fy - block_bottom >= max(int(0.035 * H), int(0.9 * (sub_size or f_size) * EMU_PER_PT))
        if moved_sub is not None and moved_sub.intersection(Bbox(x=f_x, y=fy, w=f_w, h=f_h)) > 0:
            clear = False  # never on the subtitle that went back to the sample's place
        if clear and f_x + f_w <= f_room:
            line_box(footer_text, f_size, "Дата", sub_color, False, (f_x, fy, roomy(f_w, min(f_room - f_x, w_max))))
        elif s_el is not None and sub_text and kind == PatternKind.title and footer_text != _deck_date(ctx.outline.language):
            # an audience line that found no place of its own stays with the subtitle rather than being lost
            longer = bind_short_words(sub_text + " · " + footer_text.rsplit(" · ", 1)[0])
            fill_text(s_el, [ParagraphSpec(longer, bullet=False)], size_pt=sub_size)
            _set_paragraph_box(s_el, marl, align, bold=False, line_spacing=1.0)
            if sub_color:
                style_runs(s_el, None, sub_color, align=None)
            s_lines = max(1, len(display_lines(longer, family, sub_size, False, inner(w_max) * 0.92)))
            set_element_pos(s_el, h=int((s_lines * 1.2 + 0.25) * sub_size * EMU_PER_PT) + s_ins[1] + s_ins[3])
    if ctx.goal_to_place and kind == PatternKind.title:
        # the goal that found no calm place under the heading or beside its subtitle: on a calm painted panel of the
        # layout near the stack (the purple triangle under the date on LO Focus), in the heading's colour where it
        # reads, never under 1.8 % of the slide's height nor at the small print's size (round 4.1, C3-1)
        homeless, ctx.goal_to_place = ctx.goal_to_place, None
        floor_pt = max(0.018 * H / EMU_PER_PT, (n_size + 0.5) if note else 0.0)
        top_pt = max(g_size, floor_pt)
        g_sizes = [top_pt] + sorted((z for z in snap_sizes if floor_pt - 0.05 <= z < top_pt - 0.05), reverse=True)
        if not any(abs(z - floor_pt) < 0.3 for z in g_sizes):
            g_sizes.append(round(floor_pt, 1))
        near = [Bbox(x=b[0], y=b[1], w=b[2], h=b[3]) for b in (element_bbox(el) for el in (s_el if sub_text else None, t_el) if el is not None) if b][:1]
        src_el = s_el if s_el is not None and sub_text else t_el
        toks = ctx.manifest.tokens
        placed_el = None
        for prefs in ([title.style.color_hex], [sub_color, own_sub], [toks.color_for("background.light"), toks.color_for("background.dark"), toks.color_for("text.primary"), "FFFFFF", "000000"]):
            if not any(prefs):
                continue
            placed_el = _goal_on_panel(ctx, homeless, family, g_sizes, prefs, near, hidden, src_el)
            if placed_el is not None:
                break
        if placed_el is not None:
            ctx.warnings.append("цель обложки поставлена на спокойную часть рисунка шаблона рядом с заголовком: под ним для неё нет места")
        else:
            _note_to_speaker_notes(oslide, homeless)
            ctx.warnings.append("цель с обложки перенесена в заметки докладчика: на слайде для неё нет спокойного места")
    # pager dots follow the text: they stand under it at the sample's own distance, not where a subtitle used to be
    if dot_rows and kind == PatternKind.section:
        tops = [min(_abs_bbox(d).y for d in row if _abs_bbox(d)) for row in dot_rows]
        row_top = min(tops)
        if row_top > block_bottom:
            texts_above = [s.bbox.to_emu(W, H) for s in ctx.pattern.slots if s.role in TEXT_ROLES and s.bbox.y * H < row_top]
            ref = max((b.y2 for b in texts_above if b.y2 <= row_top + int(0.005 * H)), default=None)
            if ref is not None and row_top - ref <= int(0.12 * H):  # dots that belong to the text block, not a page pager
                gap_d = max(row_top - ref, int(0.03 * H), int(0.5 * size * EMU_PER_PT))
                if block_bottom + gap_d < row_top:
                    _move_rows(dot_rows, block_bottom + gap_d - row_top)
    return size, lines


# ---------------------------------------------------------------------------- main


KEPT_PHOTO = "фото обложки шаблона оставлено"  # the warning a clone gives when it keeps a bookend's photo


def render_clone(builder: DeckBuilder, plan_slide: LayoutSlide, oslide: OutlineSlide, pattern: Pattern, manifest: TemplateManifest, ws: TemplateWorkspace, outline: DeckOutline, keep_photos: bool = True) -> tuple[Slide, list[str]]:
    """`keep_photos`: a cover's, divider's or closing slide's photo beside its title stays (`_bookend_photo`); the
    renderer sets the cover again without it when the kept photo leaves the title no room."""
    slide = builder.clone_slide(pattern.source_slide)
    ctx = _SlideCtx(builder, slide, pattern, manifest, ws, outline)
    ctx.keep_photos = keep_photos
    c = oslide.content
    ctx.keeps_pictures = bool(c.image_hint or c.chart is not None or c.table is not None)
    renumber_page_chrome(ctx.els, pattern.chrome_shape_ids, pattern.source_slide, len(builder.created))
    # a header's or a footer's sample values of a free template («JOHN DOE», «NEW YORK», «2023», the template site)
    for sid in pattern.chrome_shape_ids:
        el = ctx.els.get(sid)
        if el is None or el.getparent() is None or etree.QName(el).localname != "sp":
            continue
        txt = shape_text(el)
        if txt.strip() and looks_like_sample_value(txt):
            remove_element(el)
            ctx.removed.add(sid)
    if _is_bookend(oslide) and ctx.slots(SlotRole.title):
        size, lines = _render_bookend(ctx, oslide)
        _text_over_pictures(ctx)
        if size:
            # the plan explains the slide as it was set, not as the scorer estimated it
            plan_slide.fit["title_fit"] = f"{size:g} пт, строк {lines}"
            plan_slide.reasons = [re.sub(r"^обложка: заголовок [\d.,]+ пт", f"обложка: заголовок {size:g} пт", r) for r in plan_slide.reasons]
        for sid, el in list(ctx.els.items()):
            if sid in ctx.filled or sid in ctx.removed or el.getparent() is None or etree.QName(el).localname != "sp":
                continue
            txt = shape_text(el)
            if txt.strip() and looks_like_placeholder(txt):
                clear_text(el)
        _remove_qr_codes(ctx)
        if oslide.notes:
            try:
                slide.notes_slide.notes_text_frame.text = oslide.notes
            except Exception:  # noqa: BLE001
                pass
        return slide, ctx.warnings

    # title / subtitle
    titles = ctx.slots(SlotRole.title)
    if titles:
        ctx.fill_slot(titles[0], [ParagraphSpec(oslide.headline)], min_ratio=0.7)
        _fill_label(ctx, titles[0], oslide)
    elif oslide.headline and oslide.kind not in (PatternKind.thanks,):
        ctx.warnings.append("pattern has no title slot")
    subtitle_text = oslide.subtitle or (oslide.section if oslide.kind not in (PatternKind.title, PatternKind.thanks, PatternKind.section) else None)
    subs = ctx.slots(SlotRole.subtitle)
    if not subs and oslide.subtitle and oslide.kind in (PatternKind.title, PatternKind.section, PatternKind.thanks) and titles:
        # title samples often keep an empty text box under the title that roles.py cannot name: it is the subtitle
        top = titles[0].bbox
        below = [s for s in ctx.slots(SlotRole.body, SlotRole.bullet_list, SlotRole.caption) if s.bbox.y >= top.y2 - 0.02]
        if below:
            subs = [min(below, key=lambda s: s.bbox.y)]
    if subs and subtitle_text:
        ctx.fill_slot(subs[0], [ParagraphSpec(subtitle_text, bullet=False)])
    elif oslide.subtitle:
        ctx.warnings.append("subtitle had no slot")

    # repeated items (cards, KPIs, steps, columns) — bullets too when the pattern is a list of cells
    items = _items_for(oslide)
    bullets_as_items = False
    if not items and c.bullets and _cells_can_hold_bullets(ctx, c.bullets):
        items = [SlideItem(title=b) for b in c.bullets]
        bullets_as_items = True
    use_ordinals = oslide.kind in (PatternKind.agenda, PatternKind.process, PatternKind.timeline)
    if items:
        handled = _fill_cells(ctx, items, use_ordinals=use_ordinals, text_only=bullets_as_items)
        if not handled:
            if bullets_as_items:
                items = []
                bullets_as_items = False
            else:
                _fill_standalone_items(ctx, items)

    # bullets / paragraphs / quote
    body_cap = ctx.typo.size_for("h2", ctx.typo.size_for("body", 14.0) * 1.3)
    text_cap = _text_cap(ctx)

    def _running_size(slot: Slot) -> Optional[float]:
        sz = slot.style.size_pt or 0
        return body_cap if sz > body_cap * 1.2 else None  # a KPI-sized slot must not carry running text at 60 pt

    if c.bullets and not bullets_as_items and oslide.kind not in _ITEM_KINDS:
        # the biggest text area of the slide, whatever the sample called it (VK Tech p9 has a one-line «body» next to
        # a panel-sized «bullet list»: the role order alone would put the text into the one-liner)
        target = ctx.slots(SlotRole.bullet_list, SlotRole.body) or ctx.slots(SlotRole.card_body)
        if target:
            big = ctx.largest(target)
            ctx.fill_slot(big, [ParagraphSpec(b, bullet=True) for b in c.bullets], size_hint=_running_size(big), grow_to=text_cap)
        else:
            ctx.warnings.append("no slot for bullets")
    if c.paragraphs:
        target = ctx.slots(SlotRole.body, SlotRole.bullet_list) or ctx.slots(SlotRole.card_body) or ctx.slots(SlotRole.caption)
        if target:
            big = ctx.largest(target)
            lone = len(c.paragraphs) == 1 and len(c.paragraphs[0].split()) <= 25 and not c.bullets and not items
            cap = max(body_cap, ctx.typo.size_for("h1", body_cap) * 0.75) if lone else text_cap
            ctx.fill_slot(big, [ParagraphSpec(p, bullet=False) for p in c.paragraphs], size_hint=_running_size(big), grow_to=cap)
        elif oslide.kind not in (PatternKind.chart, PatternKind.table):
            ctx.warnings.append("no slot for paragraphs")
    if c.quote:
        target = ctx.slots(SlotRole.body, SlotRole.bullet_list, SlotRole.card_body, SlotRole.subtitle)
        if target:
            big = ctx.largest(target)
            ctx.fill_slot(big, [ParagraphSpec("«" + c.quote.strip("«»\"") + "»", bullet=False)], grow_to=body_cap)
            if c.quote_author:
                small_slots = ctx.slots(SlotRole.caption, SlotRole.card_body, SlotRole.body, SlotRole.subtitle, SlotRole.number_label)
                author_slot = ctx.smallest_font(small_slots)
                if author_slot is not None:
                    cap = ctx.typo.size_for("small", ctx.typo.size_for("body", 14.0))
                    size = min(author_slot.style.size_pt or cap, cap)
                    ctx.fill_slot(author_slot, [ParagraphSpec(c.quote_author)], size_hint=size)
        else:
            ctx.warnings.append("no slot for quote")

    # native chart / table
    if c.chart is not None or c.table is not None:
        _place_native_object(ctx, oslide)

    # illustration by hint
    if c.image_hint:
        imgs = ctx.slots(SlotRole.image)
        if imgs:
            path = pick_asset(manifest, ws, c.image_hint, exclude=ctx.used_assets)
            el = ctx.els.get(imgs[0].shape_id)
            if path and el is not None and etree.QName(el).localname == "pic":
                replace_picture(slide, el, path)
                ctx.filled.add(imgs[0].shape_id)

    # cleanup: cells that got nothing lose their anchors and chips, «Вставить фото / QR» boxes go with their text,
    # then unfilled text slots and stray placeholder text
    _remove_empty_cells(ctx)
    _remove_placeholder_boxes(ctx)
    for slot in list(ctx.pattern.slots):
        if slot.role not in TEXT_ROLES or slot.shape_id in ctx.filled or slot.shape_id in ctx.removed:
            continue
        el = ctx.els.get(slot.shape_id)
        if el is None or el.getparent() is None:
            continue
        if slot.container and not is_nested(el):
            remove_element(el)
            ctx.removed.add(slot.shape_id)
        elif has_visible_style(el) and not is_nested(el) and slot.bbox.area <= 0.08 and len(slot.sample_text or "") <= 40:
            remove_element(el)  # a painted label («Примечание» in a pill) without its text is an empty lozenge
            ctx.removed.add(slot.shape_id)
        elif has_visible_style(el) or is_nested(el):
            clear_text(el)
        else:
            remove_element(el)
            ctx.removed.add(slot.shape_id)
    for sid, el in list(ctx.els.items()):
        if sid in ctx.filled or sid in ctx.removed or etree.QName(el).localname != "sp":
            continue
        txt = shape_text(el)
        if txt.strip() and looks_like_placeholder(txt):
            clear_text(el)
    _remove_orphan_holders(ctx)
    _true_rings(ctx)
    # sample content pictures (photos, screenshots, chart images) that nothing replaced are stale: drop them
    has_visual = bool(c.image_hint or c.chart is not None or c.table is not None)
    drop_icons = oslide.kind in (PatternKind.thanks, PatternKind.title, PatternKind.section, PatternKind.quote)
    for slot in ctx.slots(SlotRole.image, SlotRole.icon):
        el = ctx.els.get(slot.shape_id)
        if el is None or is_nested(el) or el.getparent() is None:
            continue
        if slot.shape_id in ctx.cell_member_ids:
            # inside a card that received an item: a big picture is a stale sample chart/photo, a small one is the card's icon
            if slot.bbox.w > _ICON_MAX_W:
                remove_element(el)
                ctx.removed.add(slot.shape_id)
            continue
        if slot.role == SlotRole.icon and not drop_icons:
            continue
        if drop_icons and slot.role == SlotRole.image and _bookend_photo(ctx, slot, el):
            if KEPT_PHOTO not in ctx.warnings:
                ctx.warnings.append(KEPT_PHOTO)
            continue  # a cover's photo beside its title card is the cover's design: without it half the slide is empty
        if not has_visual or slot.bbox.area >= 0.05:
            remove_element(el)
            ctx.removed.add(slot.shape_id)
    _remove_qr_codes(ctx)
    if oslide.notes:
        try:
            slide.notes_slide.notes_text_frame.text = oslide.notes
        except Exception:  # noqa: BLE001
            pass
    return slide, ctx.warnings
