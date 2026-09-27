"""DeckBuilder: build a new deck inside a copy of the template package by cloning sample slides."""

from __future__ import annotations

import copy
import logging
from pathlib import Path
from typing import Optional

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.slide import Slide

from verstka.analysis.xmlns import NS, q

log = logging.getLogger(__name__)

_R_ATTRS = (q("r:embed"), q("r:link"), q("r:id"), q("r:pict"), q("r:dm"), q("r:lo"), q("r:qs"), q("r:cs"))
_SKIP_RELTYPES = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.SLIDE}


class DeckBuilder:
    def __init__(self, template: Path | str) -> None:
        self.template = Path(template)
        self.prs = Presentation(str(self.template))
        self.n_original = len(self.prs.slides)
        self._original_ids = [sld.get("id") for sld in self.prs.slides._sldIdLst]
        self.slide_w = int(self.prs.slide_width)
        self.slide_h = int(self.prs.slide_height)
        self.created: list[Slide] = []

    # ---- access -----------------------------------------------------------------
    def source_slide(self, index_1based: int) -> Slide:
        if not (1 <= index_1based <= self.n_original):
            raise IndexError(f"template has {self.n_original} slides, requested {index_1based}")
        return self.prs.slides[index_1based - 1]

    def layouts(self):
        """Every layout of every master (python-pptx's `slide_layouts` covers the first master only)."""
        return [layout for master in self.prs.slide_masters for layout in master.slide_layouts]

    # ---- cloning ----------------------------------------------------------------
    def clone_slide(self, index_1based: int) -> Slide:
        src = self.source_slide(index_1based)
        new = self.prs.slides.add_slide(src.slide_layout)
        # drop placeholders python-pptx created from the layout
        for shp in list(new.shapes):
            shp._element.getparent().remove(shp._element)
        src_cSld = src._element.cSld
        new_cSld = new._element.cSld
        bg = src_cSld.find(q("p:bg"))
        if bg is not None:
            new_cSld.insert(0, copy.deepcopy(bg))
        new_tree = new_cSld.find(q("p:spTree"))
        for child in src_cSld.find(q("p:spTree")):
            tag = etree.QName(child).localname
            if tag in ("nvGrpSpPr", "grpSpPr"):
                continue
            new_tree.append(copy.deepcopy(child))
        # relationships referenced from the copied XML
        rid_map: dict[str, str] = {}
        for rid, rel in src.part.rels.items():
            if rel.reltype in _SKIP_RELTYPES:
                continue
            if rel.is_external:
                new_rid = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            else:
                new_rid = new.part.relate_to(rel.target_part, rel.reltype)
            rid_map[rid] = new_rid
        # the whole cSld: a picture background (<p:bg>) references the part rels as well
        for el in new_cSld.iter():
            for attr in _R_ATTRS:
                v = el.get(attr)
                if v is not None and v in rid_map:
                    el.set(attr, rid_map[v])
        # timing/transitions from the source (optional, harmless)
        # placeholders that inherit their geometry from the layout get it written explicitly: every later partial
        # write (a lone `.width =`) and every bbox reader then sees the real box, not x=0 / None
        materialize_placeholders(new)
        self.created.append(new)
        return new

    def add_blank_slide(self, layout_index: int = 0) -> Slide:
        layout = self.prs.slide_layouts[layout_index]
        return self.add_layout_slide(layout)

    def add_layout_slide(self, layout) -> Slide:
        """A new slide made from a layout (its placeholders with their inherited geometry written explicitly)."""
        new = self.prs.slides.add_slide(layout)
        materialize_placeholders(new)
        self.created.append(new)
        return new

    # ---- bookkeeping --------------------------------------------------------------
    def delete_original_slides(self) -> None:
        sld_lst = self.prs.slides._sldIdLst
        for sld in list(sld_lst):
            if sld.get("id") in self._original_ids:
                rid = sld.rId
                sld_lst.remove(sld)
                self.prs.part.drop_rel(rid)
        self.n_original = 0

    def next_shape_id(self, slide: Slide) -> int:
        ids = [int(el.get("id")) for el in slide._element.iter() if etree.QName(el).localname == "cNvPr" and (el.get("id") or "").isdigit()]
        return (max(ids) + 1) if ids else 2

    def save(self, path: Path | str) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


_PH_ALIASES = {"ctrTitle": ("ctrTitle", "title"), "title": ("title", "ctrTitle"), "subTitle": ("subTitle", "body"), "obj": ("obj", "body"), "body": ("body", "obj")}


def _ph_of(el: etree._Element) -> Optional[etree._Element]:
    for path in ("p:nvSpPr/p:nvPr/p:ph", "p:nvPicPr/p:nvPr/p:ph", "p:nvGraphicFramePr/p:nvPr/p:ph"):
        ph = el.find("/".join(q(t) for t in path.split("/")))
        if ph is not None:
            return ph
    return None


def _match_ph(root: etree._Element, ph_type: str, ph_idx: Optional[str]) -> Optional[etree._Element]:
    """The placeholder of a layout/master a placeholder inherits from — the rule the template analysis uses
    (`analysis.shapes._match_placeholder`): an explicit idx with a compatible type, then the idx alone, then the type
    (python-pptx matches idx only, so an untyped `<p:ph/>` body would take the title's box)."""
    phs = []
    for sp in root.iter(q("p:sp")):
        ph = _ph_of(sp)
        if ph is not None:
            phs.append((ph.get("type") or "body", ph.get("idx"), sp))
    if ph_idx is not None:
        for t, i, sp in phs:
            if i == ph_idx and (t == ph_type or (ph_type in ("body", "obj") and t in ("body", "obj"))):
                return sp
        for t, i, sp in phs:
            if i == ph_idx:
                return sp
    for cand in _PH_ALIASES.get(ph_type, (ph_type,)):
        for t, i, sp in phs:
            if t == cand:
                return sp
    return None


def _xfrm_box(sp: etree._Element) -> Optional[tuple[int, int, int, int]]:
    spPr = sp.find(q("p:spPr"))
    xfrm = spPr.find(q("a:xfrm")) if spPr is not None else None
    off = xfrm.find(q("a:off")) if xfrm is not None else None
    ext = xfrm.find(q("a:ext")) if xfrm is not None else None
    if off is None or ext is None:
        return None
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def placeholder_chain(shape) -> list[etree._Element]:
    """The placeholders a slide (or layout) placeholder inherits from, nearest first: the layout's, then the master's
    (matched like the template analysis does). Empty for a shape that is not a placeholder."""
    from pptx.parts.slide import SlideLayoutPart, SlidePart

    el = getattr(shape, "_element", None)
    ph = _ph_of(el) if el is not None else None
    part = getattr(shape, "part", None)
    if ph is None or part is None:
        return []
    ph_type, ph_idx = ph.get("type") or "body", ph.get("idx")
    chain: list[tuple[etree._Element, Optional[str]]] = []
    try:
        if isinstance(part, SlidePart):
            layout = part.slide.slide_layout
            chain = [(layout._element, ph_idx), (layout.slide_master._element, None)]
        elif isinstance(part, SlideLayoutPart):
            chain = [(part.slide_layout.slide_master._element, None)]
    except Exception:  # noqa: BLE001
        return []
    out = []
    for root, idx in chain:
        base = _match_ph(root, ph_type, idx)
        if base is not None:
            out.append(base)
    return out


def inherited_box(shape) -> Optional[tuple[int, int, int, int]]:
    """The geometry a placeholder inherits (layout placeholder, then master placeholder), ignoring its own xfrm."""
    for base in placeholder_chain(shape):
        box = _xfrm_box(base)
        if box is not None:
            return box
    return None


def materialize_xfrm(shape) -> bool:
    """Write the inherited geometry of a placeholder (resolved through the layout and master placeholders) as an
    explicit `a:xfrm` — the first child of `p:spPr`. No-op (False) when the shape already has a complete xfrm, is not
    a `p:sp`, or its geometry is unknown. A partial xfrm (a lone python-pptx `.width =` writes `a:ext` with cy=0 and
    no `a:off`) keeps what was deliberately written and takes the rest from the placeholder chain."""
    el = getattr(shape, "_element", None)
    if el is None or etree.QName(el).localname != "sp":
        return False
    spPr = el.find(q("p:spPr"))
    if spPr is None:
        return False
    xfrm = spPr.find(q("a:xfrm"))
    off = xfrm.find(q("a:off")) if xfrm is not None else None
    ext = xfrm.find(q("a:ext")) if xfrm is not None else None
    if off is not None and ext is not None:
        return False
    inh = inherited_box(shape)
    if inh is None:
        return False
    x, y, w, h = inh
    if off is not None:
        x, y = int(off.get("x", x)), int(off.get("y", y))
    if ext is not None:
        w = int(ext.get("cx") or 0) or w
        h = int(ext.get("cy") or 0) or h
    if xfrm is not None:
        for child in list(xfrm):
            xfrm.remove(child)
    else:
        xfrm = etree.Element(q("a:xfrm"))
        spPr.insert(0, xfrm)
    etree.SubElement(xfrm, q("a:off")).attrib.update({"x": str(int(x)), "y": str(int(y))})
    etree.SubElement(xfrm, q("a:ext")).attrib.update({"cx": str(int(w)), "cy": str(int(h))})
    return True


def materialize_placeholders(slide: Slide) -> int:
    """`materialize_xfrm` on every top-level placeholder of the slide; returns how many were written."""
    n = 0
    for shp in slide.shapes:
        if getattr(shp, "is_placeholder", False):
            try:
                n += bool(materialize_xfrm(shp))
            except Exception:  # noqa: BLE001 - never break a render over bookkeeping
                log.debug("materialize_xfrm failed on %s", getattr(shp, "shape_id", "?"), exc_info=True)
    return n


def slide_shape_elements(slide: Slide) -> dict[str, etree._Element]:
    """cNvPr id → drawable element (sp/pic/grpSp/graphicFrame/cxnSp), any depth."""
    out: dict[str, etree._Element] = {}
    tree = slide._element.cSld.find(q("p:spTree"))
    for nv in tree.iter(q("p:cNvPr")):
        sid = nv.get("id")
        drawable = nv.getparent().getparent()  # nvXxPr → sp/pic/...
        if sid and drawable is not None and drawable is not tree and etree.QName(drawable).localname != "spTree":
            out[sid] = drawable
    return out


def element_bbox(el: etree._Element) -> Optional[tuple[int, int, int, int]]:
    tag = etree.QName(el).localname
    if tag == "graphicFrame":
        xfrm = el.find(q("p:xfrm"))
    elif tag == "grpSp":
        xfrm = el.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
    else:
        spPr = el.find(q("p:spPr"))
        xfrm = spPr.find(q("a:xfrm")) if spPr is not None else None
    if xfrm is None:
        return None
    off = xfrm.find(q("a:off"))
    ext = xfrm.find(q("a:ext"))
    if off is None or ext is None:
        return None
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def set_element_pos(el: etree._Element, x: Optional[int] = None, y: Optional[int] = None, w: Optional[int] = None, h: Optional[int] = None) -> None:
    tag = etree.QName(el).localname
    if tag == "graphicFrame":
        xfrm = el.find(q("p:xfrm"))
    elif tag == "grpSp":
        xfrm = el.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
    else:
        spPr = el.find(q("p:spPr"))
        xfrm = spPr.find(q("a:xfrm")) if spPr is not None else None
    if xfrm is None:
        return
    off = xfrm.find(q("a:off"))
    ext = xfrm.find(q("a:ext"))
    if off is not None:
        if x is not None:
            off.set("x", str(int(x)))
        if y is not None:
            off.set("y", str(int(y)))
    if ext is not None:
        if w is not None:
            ext.set("cx", str(int(w)))
        if h is not None:
            ext.set("cy", str(int(h)))


def shift_element(el: etree._Element, dx: int, dy: int) -> None:
    box = element_bbox(el)
    if box is None:
        return
    set_element_pos(el, x=box[0] + dx, y=box[1] + dy)


def is_nested(el: etree._Element) -> bool:
    parent = el.getparent()
    return parent is not None and etree.QName(parent).localname == "grpSp"


def remove_element(el: etree._Element) -> None:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def renumber_ids(el: etree._Element, start: int) -> int:
    """Assign fresh cNvPr ids to every id inside el; returns the next free id."""
    nxt = start
    for nv in el.iter(q("p:cNvPr")):
        nv.set("id", str(nxt))
        nxt += 1
    return nxt
