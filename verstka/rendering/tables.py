"""Native PowerPoint tables styled from the template.

Two entry points:

* :func:`measure_table` — pure geometry. It picks one type size, the column widths and the row heights for a table set
  in a given width (Play metrics × the family width factor, with 8 % slack) and says whether the table overflows.
* :func:`add_table` — draws the native, editable table: explicit fills, rules, margins and typefaces on every cell, no
  built-in PowerPoint table style. The look comes from the template: from a native sample table when the template has
  one (the header fill XML, gradients included; the header weight; the rule colour and width; a body cell fill; the
  accent-row fill), else from the template's table tokens, adjusted so that everything reads on the actual ground.

How the composer calls them (H = slide height, CB = content box, h1 = the deck's heading size)::

    sizes = table_sizes(k.lead, k.body, k.small, h1=h1, slide_h_emu=H, n_body_rows=len(table.rows),
                        compact=strategy == "compact")        # below h1 / 1.25; B floor above 6 body rows
    fit = measure_table(
        table, width, k.font, sizes, k.line,
        max_h_emu=area.h,                      # never taller than the free area
        fill_h_emu=int(0.55 * CB.h),           # short tables grow their rows until they fill 55 % of the CB …
        min_row_h_emu=int(0.065 * H),          # … from rows of at least 0.065 H …
        max_row_h_emu=int(0.10 * H),           # … to rows of at most 0.10 H
        header_bold=template_bold(typography))
    size, widths, heights = fit                # fit.content_w: the natural width (≤ 4 rows: shrink it, add a takeaway)
                                               # fit.overflow / fit.too_dense: split the table or move rows on
    top = y0 + int(0.4 * (CB.h - sum(heights)))
    add_table(slide, Bbox(x=x0, y=top, w=sum(widths), h=sum(heights)), table, style, typography,
              font_family=k.font, col_widths=widths, row_heights=heights, size_pt=size,
              accent_hex=k.colors.accent, accent_text_hex=k.colors.accent_text, muted_hex=k.colors.muted,
              ground_hex=k.colors.ground)      # omit on a gradient or picture ground: the tint then stays translucent

``add_table`` adapts ``style`` to the ground itself (:func:`table_style_for_ground`), so the manifest's
``components.table_style`` may be passed as is. ``band_every`` may be left at ``None``: tables of six body rows or more
are banded every second row, shorter ones are separated by hairlines only, and banding that would give fewer than two
stripes is never drawn. Without the geometry kwargs ``add_table`` measures the table itself inside ``bbox`` (the clone
and synth renderers rely on that); without ``ground_hex`` it reads the ground from the slide background.
"""

from __future__ import annotations

import copy
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Optional, Union

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.slide import Slide
from pptx.util import Emu, Pt

from verstka.analysis.theme import _PRESET_COLORS, _apply_modifiers
from verstka.analysis.xmlns import q
from verstka.rendering.fonts import text_width_pt
from verstka.schemas.common import EMU_PER_PT, Bbox, contrast_ratio, relative_luminance
from verstka.schemas.outline import TableData
from verstka.schemas.template import TableStyleSpec, Typography

# built-in "No Style, No Grid": nothing from a table style can leak into the cells we style explicitly
NO_STYLE_TABLE_ID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
TABLE_NAME = "Table"
MEASURE_SLACK = 1.08  # measured text is 8 % wider than Play × factor says: the rendering font may be wider
RULE_PT = 0.75  # hairline between body rows
TOTAL_RULE_PT = 1.25  # accent rule above a total row
HEADER_RULE_PT = 1.5  # accent rule under an unfilled header
MARK_RULE_PT = 3.0  # accent bar at the left of a recommended row whose label cannot be set in the accent
SEPARATOR_PT = 2.0  # a ground-coloured gap between a filled header and an accent row of the same colour
GUTTER_EM = 1.0  # extra left margin of a left-aligned column right after a right-aligned one

_S = r"[\s  ]"
_N = r"\d[\d\s  .,]*"
_QUAL = r"(?:до|от|около|более|менее|свыше|почти|≈|~|<|>|≤|≥)"
_NUM_RE = re.compile(
    rf"^(?:{_QUAL}{_S}?)?"  # a qualifier: «до 5 мин», «≈ 40 %»
    rf"[+\-−–×xх]?{_S}?{_N}"  # value (with thousands groups), «×2», «x3»
    rf"(?:{_S}?[–—\-]{_S}?[+\-−]?{_N})?"  # a range: «3–5 дней»
    rf"(?:{_S}?[%‰×xх])?"  # percent, times: «86%», «2x»
    rf"(?:{_S}из{_S}{_N})?"  # a grade: «4,7 из 5»
    rf"(?:{_S}[^\s\d]{{1,12}}\.?){{0,2}}"  # up to two unit words: «млн ₽», «тыс.», «месяцев», «мин»
    rf"(?:{_S}?[^\s\d]{{1,2}})?$",  # a glued unit: «12ч», «5₽»
    re.I,
)
_PERIOD_RE = re.compile(r"^(?:[QH][1-4]|[1-4]\s?(?:кв|пг|п/г)\.?|[IV]{1,3}\s?кв\.?)\s?(?:\d{4}|\d{2})(?:\s?г\.?)?$", re.I)
_EMPTY_CELLS = {"", "—", "–", "-", "−", "н/д", "n/a", "нет данных"}
_DASHES = {"—", "–", "-", "−"}
_YES = {"да", "есть", "yes", "true", "+", "✓", "✔", "✔️", "✅", "☑"}
_NO = {"нет", "no", "false", "✗", "✕", "✖", "❌", "×"}
_PARTIAL = {"частично", "опц.", "опционально", "по запросу", "partial", "±"}
_FLAG_WORDS = _YES | _NO | _DASHES | _PARTIAL
_TOTAL_RE = re.compile(r"^\s*(итого|итог|всего|total|сумма|в\s+сумме|суммарно)\b", re.I)
_RECOMMENDED_RE = re.compile(r"^\s*(наш|наша|наше|наши|рекомендуем\w*|рекомендуемый|предлагаем\w*|our|recommended)\b", re.I)
_SPLIT_WORDS = re.compile(r"[ \t\r\n]+")  # NBSP keeps a number with its unit: it is one «word»
CHECK = "✓"
DASH = "—"

# units a numeric column may carry in every cell; they move into the header («Стоимость в год, млн ₽»)
_UNIT_TOKENS = {
    "%": "%", "‰": "‰", "₽": "₽", "$": "$", "€": "€", "руб": "₽", "руб.": "₽", "рубль": "₽", "рубля": "₽", "рублей": "₽",
    "млн": "млн", "млн.": "млн", "млрд": "млрд", "млрд.": "млрд", "трлн": "трлн", "тыс": "тыс.", "тыс.": "тыс.",
    "мин": "мин", "мин.": "мин", "минута": "мин", "минуты": "мин", "минут": "мин",
    "ч": "ч", "ч.": "ч", "час": "ч", "часа": "ч", "часов": "ч", "сек": "сек", "сек.": "сек", "секунд": "сек", "секунды": "сек",
    "мс": "мс", "день": "дн.", "дня": "дн.", "дней": "дн.", "дн": "дн.", "дн.": "дн.",
    "неделя": "нед.", "недели": "нед.", "недель": "нед.", "нед": "нед.", "нед.": "нед.",
    "месяц": "мес.", "месяца": "мес.", "месяцев": "мес.", "мес": "мес.", "мес.": "мес.", "год": "лет", "года": "лет", "лет": "лет",
    "чел": "чел.", "чел.": "чел.", "человек": "чел.", "шт": "шт.", "шт.": "шт.", "п.п.": "п.п.", "пп": "п.п.",
    "гб": "ГБ", "тб": "ТБ", "мб": "МБ", "кг": "кг", "км": "км",
}
_UNIT_SPLIT_RE = re.compile(
    rf"^(?P<num>(?:{_QUAL}{_S})?[+\-−–]?\d[\d\s  ]*(?:[.,]\d+)?(?:{_S}?[–—\-]{_S}?\d[\d\s  ]*(?:[.,]\d+)?)?)(?P<unit>[^\d]*)$",
    re.I,
)


# ------------------------------------------------------------------------------------------------ classification


def _is_numeric(text: str) -> bool:
    t = text.strip()
    return bool(t) and bool(_NUM_RE.match(t) or _PERIOD_RE.match(t)) and any(ch.isdigit() for ch in t)


def _body_cells(rows: list[list[str]], j: int) -> list[str]:
    return [row[j].strip() for row in rows if j < len(row) and row[j] and row[j].strip()]


def _column_is_numeric(rows: list[list[str]], j: int) -> bool:
    """Every non-empty body cell is a number (dashes and «н/д» for missing values do not count against it)."""
    cells = [c for c in _body_cells(rows, j) if c.lower() not in _EMPTY_CELLS]
    return bool(cells) and all(_is_numeric(c) for c in cells)


def _column_is_flag(rows: list[list[str]], j: int) -> bool:
    """Short categorical values (да/нет, ✓, —, grades of ≤ 3 characters): centred."""
    cells = _body_cells(rows, j)
    return bool(cells) and all(c.lower() in _FLAG_WORDS or len(c) <= 3 for c in cells)


def _column_is_boolean(rows: list[list[str]], j: int) -> bool:
    """Yes/no values in any spelling (да/нет, ✓/—, +/−, Да/-, есть/—, a «частично» among them) with at least one yes."""
    cells = [c.lower() for c in _body_cells(rows, j)]
    return bool(cells) and all(c in _YES or c in _NO or c in _DASHES or c in _PARTIAL for c in cells) and any(c in _YES for c in cells)


def column_kinds(table: TableData) -> list[str]:
    """One kind per column: ``label`` (the first column of a multi-column table names the row), ``numeric``
    (right-aligned, equal widths), ``flag`` (short categorical values, centred) or ``text``."""
    n = len(table.columns)
    kinds = []
    for j in range(n):
        if j == 0 and n > 1:
            kinds.append("label")
        elif _column_is_numeric(table.rows, j):
            kinds.append("numeric")
        elif _column_is_flag(table.rows, j) or _column_is_boolean(table.rows, j):
            kinds.append("flag")
        else:
            kinds.append("text")
    return kinds


def is_total_row(row: list[str]) -> bool:
    return bool(row) and bool(_TOTAL_RE.match(row[0] or ""))


_TARGET_COL_RE = re.compile(r"(?<![\w])(цел[ьи]|план\w*|прогноз\w*|после|стало|будет|через\s+\d|к\s+\d|target|goal|plan|forecast|after)", re.I)
_PAST_COL_RE = re.compile(r"(?<![\w])(сейчас|было|до|текущ\w*|факт\w*|исходн\w*|now|before|current|actual)(?![\w])", re.I)
_YEAR_COL_RE = re.compile(r"(?<!\d)(19\d\d|20\d\d)(?!\d)")


def emphasis_column(table: TableData) -> Optional[int]:
    """The column a comparison table argues for — the target («Цель», «План», «Прогноз», «После», «Стало»,
    «Через 6 месяцев») when exactly one header names it, else the last column of a table whose other value columns
    are the past («Сейчас», «Было», earlier years). None for a table of options, where no column is the answer."""
    cols = [c or "" for c in table.columns]
    if len(cols) < 3 or not table.rows:
        return None
    hits = [j for j, c in enumerate(cols) if j > 0 and _TARGET_COL_RE.search(c)]
    if len(hits) == 1:
        return hits[0]
    last = len(cols) - 1
    years = {j: int(m.group(1)) for j, c in enumerate(cols) if j > 0 and (m := _YEAR_COL_RE.search(c))}
    if len(years) >= 2 and max(years, key=lambda j: years[j]) == last:
        return last
    if any(_PAST_COL_RE.search(c) for c in cols[1:last]) and not _PAST_COL_RE.search(cols[last]):
        return last
    return None


def recommended_row(table: TableData) -> Optional[int]:
    """The body row the table argues for («Наш подсказчик», «Рекомендуем…»): exactly one such row among ≥ 2."""
    if len(table.rows) < 2 or len(table.columns) < 2:
        return None
    hits = [i for i, row in enumerate(table.rows) if row and _RECOMMENDED_RE.match(row[0] or "") and not is_total_row(row)]
    return hits[0] if len(hits) == 1 else None


NBSP = " "
_THOUSANDS_RE = re.compile(r"(?<=\d) (?=\d{3}(?!\d))")
_BEFORE_SYMBOL_RE = re.compile(r" (?=[%‰₽$€£¥—–](?:\s|$))")
_SHORT_WORD_RE = re.compile(r"(?<![\w\-.,])(\w{1,2}) (?=\S)")
_UNIT_RE = re.compile(r"(?<=\d) (?=(?:млн|млрд|трлн|тыс|руб|мин|сек|ч|шт|мес|дн|нед|кг|км|см|мм|гб|тб|мб|pp|п\.п\.|из)(?:\.|\b))", re.I)


def typeset(text: str) -> str:
    """Russian typesetting for a cell: NBSP in thousands groups, between a value and its unit, before a unit symbol
    or a dash, and after words of 1–2 letters (so «в», «из», «8» never end a line)."""
    t = (text or "").strip()
    if " " not in t:
        return t
    t = _THOUSANDS_RE.sub(NBSP, t)
    t = _UNIT_RE.sub(NBSP, t)
    t = _BEFORE_SYMBOL_RE.sub(NBSP, t)
    t = _SHORT_WORD_RE.sub(lambda m: m.group(1) + NBSP, t)
    return t


def _split_unit(cell: str) -> Optional[tuple[str, str]]:
    """(value, normalised unit) of a numeric cell («9 млн ₽» → «9», «млн ₽»; «2 месяца» → «2», «мес.»), or None when the
    text after the number is not a unit."""
    m = _UNIT_SPLIT_RE.match(cell.strip())
    if not m:
        return None
    num, raw = m.group("num").strip(), m.group("unit").strip()
    if not raw:
        return num, ""
    units: list[str] = []
    for tok in raw.lower().replace(NBSP, " ").replace(" ", " ").split():
        u = _UNIT_TOKENS.get(tok)
        if u is None:
            return None
        if u not in units:
            units.append(u)
    return num, " ".join(units)


def _shared_unit(rows: list[list[str]], j: int) -> Optional[str]:
    """The unit every value of column j carries (at least two values), when they all carry the same one."""
    cells = [c for c in _body_cells(rows, j) if c.lower() not in _EMPTY_CELLS]
    if len(cells) < 2:
        return None
    parts = [_split_unit(c) for c in cells]
    if any(p is None for p in parts):
        return None
    units = {p[1] for p in parts if p}
    return units.pop() if len(units) == 1 and "" not in units else None


def _header_with_unit(head: str, unit: str) -> str:
    """«Стоимость в год» + «млн ₽» → «Стоимость в год, млн ₽»; a unit the header already names is replaced, not doubled."""
    h = (head or "").strip()
    if not h:
        return unit
    if "," in h:
        body, tail = h.rsplit(",", 1)
        toks = tail.strip().lower().replace(NBSP, " ").split()
        if toks and all(t in _UNIT_TOKENS or t in _UNIT_TOKENS.values() for t in toks):
            return f"{body.rstrip()}, {unit}"
    m = re.search(r"\(([^()]*)\)\s*$", h)
    if m:
        toks = m.group(1).strip().lower().replace(NBSP, " ").split()
        if toks and all(t in _UNIT_TOKENS or t in _UNIT_TOKENS.values() for t in toks):
            return h  # «Доля (%)»: the header names its unit already, in its own words
    words = {w for w in re.split(r"[^\w%‰₽$€.]+", h.lower()) if w}
    forms = {t for t, u in _UNIT_TOKENS.items() if u in unit.split()} | set(unit.lower().split())
    if words & forms:
        return h  # «Доля в %»: the header already names the unit
    return f"{h}, {unit}"


def _no_widow(text: str) -> str:
    """Bind the last two words of a sentence-long cell so a wrapped cell never ends on a lone word."""
    words = text.split(" ")
    if len(words) >= 4 and len(words[-1]) <= 12 and len(words[-2]) <= 14:
        return " ".join(words[:-2]) + " " + words[-2] + NBSP + words[-1]
    return text


@dataclass
class _Prep:
    """The table as it is set: kinds, header texts, body texts and the boolean (✓ / —) columns."""

    kinds: list[str]
    heads: list[str]
    rows: list[list[str]]
    booleans: list[bool]


def _prepare(table: TableData, boolean_marks: bool = True) -> _Prep:
    n = len(table.columns)
    kinds = column_kinds(table) if n else []
    raw = [[((row[j] if j < len(row) else "") or "").strip() for j in range(n)] for row in table.rows]
    heads = [(table.columns[j] or "").strip() for j in range(n)]
    for j in range(n):
        if kinds[j] != "numeric":
            continue
        unit = _shared_unit(raw, j)
        if unit:  # the unit goes into the header once, the cells keep the bare values (they align on the digits)
            heads[j] = _header_with_unit(heads[j], unit)
            for row in raw:
                sp = _split_unit(row[j]) if row[j] and row[j].lower() not in _EMPTY_CELLS else None
                if sp:
                    row[j] = sp[0]
    booleans = [boolean_marks and n > 1 and j > 0 and _column_is_boolean(table.rows, j) for j in range(n)]
    rows = [[typeset(c) for c in row] for row in raw]
    for row in rows:
        for j in range(n):
            if kinds[j] == "text":
                row[j] = _no_widow(row[j])
            elif booleans[j]:
                v = row[j].strip().lower()
                if v in _YES:
                    row[j] = CHECK
                elif v in _NO or v in _DASHES:
                    row[j] = DASH
    return _Prep(kinds=kinds, heads=[typeset(h) for h in heads], rows=rows, booleans=booleans)


def header_cells(table: TableData) -> list[str]:
    """Header texts as they are set (typeset; a column's shared unit appended)."""
    return _prepare(table).heads


def display_rows(table: TableData, boolean_marks: bool = True) -> list[list[str]]:
    """The body as it is set: rows padded to the column count, typeset (NBSP; sentence-long text cells without a
    widow), a unit shared by a whole numeric column moved into its header, boolean columns as ✓ / —."""
    return _prepare(table, boolean_marks).rows


def _align_for(kind: str) -> PP_ALIGN:
    return {"numeric": PP_ALIGN.RIGHT, "flag": PP_ALIGN.CENTER}.get(kind, PP_ALIGN.LEFT)


def _gutters(kinds: list[str], size: float) -> list[float]:
    """Extra left margin (pt) of a left-aligned column right after a right-aligned one: «2 недели | Команда данных»
    must not read as one phrase."""
    al = [_align_for(k) for k in kinds]
    return [GUTTER_EM * size if j > 0 and al[j - 1] == PP_ALIGN.RIGHT and al[j] == PP_ALIGN.LEFT else 0.0 for j in range(len(kinds))]


def template_bold(typography: Typography) -> bool:
    """Whether the template sets its headings in bold (weight_bold_share ≥ 0.3 on h1 or h2): a table header follows
    the template's own emphasis, Play decks stay regular, Montserrat decks go bold."""
    return any(st.role in ("h1", "h2") and st.weight_bold_share >= 0.3 for st in typography.scale)


# ------------------------------------------------------------------------------------------------ geometry


def cell_padding_pt(size_pt: float) -> tuple[float, float]:
    """Cell margins (horizontal, vertical) proportional to the type size: 0.7 em and 0.3 em."""
    return min(max(0.7 * size_pt, 5.0), 14.0), min(max(0.3 * size_pt, 2.0), 6.0)


def _w(text: str, font: Optional[str], size: float, bold: bool) -> float:
    if not text:
        return 0.0
    # ✓ is not in the template fonts: the fallback glyph is about as wide as an «M»
    return text_width_pt(text.replace(CHECK, "M"), font, size, bold) * MEASURE_SLACK


def _words(text: str) -> list[str]:
    return [w for w in _SPLIT_WORDS.split(text.strip()) if w]


def _longest_word(text: str, font: Optional[str], size: float, bold: bool) -> float:
    return max((_w(w, font, size, bold) for w in _words(text)), default=0.0)


def _lines(text: str, font: Optional[str], size: float, bold: bool, width_pt: float) -> int:
    """Greedy word wrap on ordinary spaces (an NBSP group never breaks); a word wider than the line breaks by letters."""
    words = _words(text)
    if not words:
        return 1
    width_pt = max(width_pt, 1.0)
    space = _w(" ", font, size, bold)
    lines, cur = 1, 0.0
    for word in words:
        ww = _w(word, font, size, bold)
        if cur and cur + space + ww <= width_pt + 1e-6:
            cur += space + ww
            continue
        if cur:
            lines += 1
        if ww > width_pt:
            extra = math.ceil(ww / width_pt) - 1
            lines += extra
            cur = ww - extra * width_pt
        else:
            cur = ww
    return lines


def _level(values: list[float], total: float, cap: float = math.inf) -> list[float]:
    """Raise the smallest values to a common level L ≤ cap so that sum(max(v, L)) == total (water-filling)."""
    if not values or sum(values) >= total:
        return list(values)
    lo, hi = 0.0, min(total, cap)
    if sum(max(v, hi) for v in values) <= total:
        return [max(v, hi) for v in values]
    for _ in range(60):
        mid = (lo + hi) / 2
        if sum(max(v, mid) for v in values) > total:
            hi = mid
        else:
            lo = mid
    return [max(v, lo) for v in values]


def _level_capped(values: list[float], caps: list[float], total: float) -> list[float]:
    """Water-filling with a cap per value: sum(min(max(v, L), cap)) == total, as far as the caps allow."""
    caps = [max(c, v) for v, c in zip(values, caps)]
    if total <= sum(values):
        return list(values)
    if total >= sum(caps):
        return list(caps)
    lo, hi = min(values), max(caps)
    for _ in range(60):
        mid = (lo + hi) / 2
        if sum(min(max(v, mid), c) for v, c in zip(values, caps)) > total:
            hi = mid
        else:
            lo = mid
    return [min(max(v, lo), c) for v, c in zip(values, caps)]


def _column_needs(prep: _Prep, total_rows: set[int], width_pt: float, font: Optional[str], size: float, header_bold: bool) -> tuple[list[float], list[float], list[float]]:
    """(minimum, body-natural, natural) width of every column in pt, padding and gutter included: the longest word; the
    widest body cell on one line (the header may wrap); every cell including the header on one line."""
    heads, rows, kinds = prep.heads, prep.rows, prep.kinds
    n = len(heads)
    px, _ = cell_padding_pt(size)
    gut = _gutters(kinds, size)
    mins, bnats, nats = [], [], []
    for j in range(n):
        head = heads[j]
        body = [(r[j], header_bold and i in total_rows) for i, r in enumerate(rows)]
        head_word = _longest_word(head, font, size, header_bold)
        if kinds[j] == "numeric":  # a number never wraps: its column is at least as wide as the widest value
            word = max([head_word] + [_w(t, font, size, b) for t, b in body])
        else:
            word = max([head_word] + [_longest_word(t, font, size, b) for t, b in body])
        bnat = max([word] + [_w(t, font, size, b) for t, b in body])
        nat = max(bnat, _w(head, font, size, header_bold))
        pad = 2 * px + gut[j]
        mins.append(word + pad)
        bnats.append(bnat + pad)
        nats.append(nat + pad)
    if n > 1 and kinds[0] == "label":  # the label column: natural width, ≤ 0.4 of the table
        bnats[0] = max(mins[0], min(bnats[0], 0.4 * width_pt))
        nats[0] = max(bnats[0], min(nats[0], 0.4 * width_pt))
    for kind in ("numeric", "flag"):  # numeric columns are equal, and so are the ✓/— columns of a feature matrix
        same = [j for j in range(n) if kinds[j] == kind]
        if len(same) > 1:
            for arr in (mins, bnats, nats):
                m = max(arr[j] for j in same)
                for j in same:
                    arr[j] = m
    return mins, bnats, nats


def _is_long(kind: str, nat: float, width_pt: float) -> bool:
    return kind == "text" and nat > 0.25 * width_pt


def _share(lo: list[float], hi: list[float], total: float) -> list[float]:
    """lo + a share of (total − Σlo) proportional to what each column still wants (hi − lo)."""
    spare = total - sum(lo)
    want = [max(h - l, 0.0) for l, h in zip(lo, hi)]
    tw = sum(want)
    if tw <= 0:
        return _level(lo, total)
    return [l + spare * w / tw for l, w in zip(lo, want)]


def _allot_spare(nats: list[float], kinds: list[str], width_pt: float, size: float) -> list[float]:
    """Everything fits on one line: share the spare width. Numbers and flags get a little air (up to
    max(1.5 × natural, natural + 3 em)), never a river between a label and its values; the rest goes to the label and
    text columns (a label keeps ≤ 0.4 of the table while a text column can take the rest)."""
    n = len(nats)
    growers = [j for j in range(n) if kinds[j] in ("label", "text")]
    fixed = [j for j in range(n) if j not in growers]
    if not growers or not fixed:
        return _level(nats, width_pt)
    widths = list(nats)
    spare = width_pt - sum(nats)
    grown = _level_capped([nats[j] for j in fixed], [max(1.5 * nats[j], nats[j] + 3 * size) for j in fixed], sum(nats[j] for j in fixed) + spare)
    for j, w in zip(fixed, grown):
        widths[j] = w
    rest = width_pt - sum(widths)
    texts = [j for j in growers if kinds[j] == "text"]
    if rest <= 1e-6:
        return widths
    if not texts:
        for j in growers:
            widths[j] += rest / len(growers)
        return widths
    tw = sum(nats[j] for j in growers) or 1.0
    for j in growers:
        widths[j] += rest * nats[j] / tw
    for j in growers:
        if kinds[j] == "label" and widths[j] > max(nats[j], 0.4 * width_pt):
            excess = widths[j] - max(nats[j], 0.4 * width_pt)
            widths[j] -= excess
            tt = sum(nats[t] for t in texts) or 1.0
            for t in texts:
                widths[t] += excess * nats[t] / tt
    return widths


def _allot(mins: list[float], bnats: list[float], nats: list[float], kinds: list[str], width_pt: float, size: float = 12.0) -> list[float]:
    """Share the table width between columns.

    Everything fits on one line → :func:`_allot_spare`. Otherwise the short columns (label, numbers, flags, short text)
    first get their body on one line, then the rest is shared by what each column still wants: long text columns wrap
    less, short headers go on one line. When not even the words fit, the columns get the width in proportion to their
    minimum.
    """
    n = len(mins)
    if sum(nats) <= width_pt:
        return _allot_spare(nats, kinds, width_pt, size)
    if sum(mins) > width_pt:
        tm = sum(mins) or 1.0
        return [width_pt * m / tm for m in mins]
    first = [mins[j] if _is_long(kinds[j], nats[j], width_pt) else bnats[j] for j in range(n)]
    if sum(first) > width_pt:
        return _share(mins, first, width_pt)
    return _share(first, nats, width_pt)


def _cell_lines(prep: _Prep, total_rows: set[int], widths_pt: list[float], font: Optional[str], size: float, header_bold: bool) -> list[list[int]]:
    px, _ = cell_padding_pt(size)
    gut = _gutters(prep.kinds, size)
    out = []
    for r, row in enumerate([prep.heads] + prep.rows):
        bold = header_bold if r == 0 else (header_bold and (r - 1) in total_rows)
        out.append([_lines(t, font, size, bold, widths_pt[j] - 2 * px - gut[j]) for j, t in enumerate(row[: len(widths_pt)])])
    return out


def _wrap_score(lines: list[list[int]], long_cols: set[int]) -> tuple[float, float]:
    """How ragged a table reads at one size, as (total, body): a wrapped body cell in a short column costs 1, a
    long-text cell beyond 3 lines 1, a header on two lines ¼ (a header «Стоимость в год, / млн ₽» reads fine), a header
    beyond two lines 1."""
    head = body = 0.0
    for r, ls in enumerate(lines):
        for j, k in enumerate(ls):
            if r == 0:
                head += 0.25 if k == 2 else (1.0 if k > 2 else 0.0)
            elif j in long_cols:
                body += 1.0 if k > 3 else 0.0
            else:
                body += 1.0 if k > 1 else 0.0
    return head + body, body


class TableFit(tuple):
    """What :func:`measure_table` returns: unpacks as ``(size_pt, col_widths, row_heights)`` and also carries

    * ``content_w`` — the natural width (EMU) of the table at that size: every cell on one line, capped at the width
      it was measured for. A composer may shrink a short table to it and put a takeaway beside it;
    * ``overflow`` — no size fits: words break inside their column or the rows do not fit ``max_h_emu``;
    * ``too_dense`` — more than 8 body rows or 5 columns, or rows below the minimum row height: split the table."""

    size_pt: float
    col_widths: list[int]
    row_heights: list[int]
    content_w: int
    overflow: bool
    too_dense: bool

    def __new__(cls, size_pt: float, col_widths: list[int], row_heights: list[int], *, content_w: int = 0, overflow: bool = False, too_dense: bool = False) -> "TableFit":
        obj = super().__new__(cls, (size_pt, col_widths, row_heights))
        obj.size_pt, obj.col_widths, obj.row_heights = size_pt, col_widths, row_heights
        obj.content_w, obj.overflow, obj.too_dense = content_w, overflow, too_dense
        return obj


def _grow_rows(heights: list[float], fill: float, cap: float) -> list[float]:
    """Grow the body rows evenly (shortest first, each ≤ cap) until the table is ``fill`` tall; the header follows at
    ¾ of the body level, never below its own need."""
    h0, body = heights[0], heights[1:]
    if not body:
        return _level(heights, fill, cap)

    def total(level: float) -> float:
        return max(h0, 0.75 * level) + sum(max(b, min(level, max(cap, b))) for b in body)

    lo, hi = min(body), max(max(body), cap if math.isfinite(cap) else fill)
    if total(lo) >= fill:
        return list(heights)
    if total(hi) <= fill:
        level = hi
    else:
        for _ in range(60):
            mid = (lo + hi) / 2
            if total(mid) > fill:
                hi = mid
            else:
                lo = mid
        level = lo
    body = [max(b, min(level, max(cap, b))) for b in body]
    return [max(h0, 0.75 * level)] + body


def table_sizes(lead: float, body: float, small: float, *, h1: float, slide_h_emu: int, n_body_rows: int, compact: bool = False) -> list[float]:
    """Candidate table sizes, largest first: the deck's lead, body and small sizes that stay below h1 / 1.25 (a table
    never competes with the heading) and above the dense-table floor 0.022 H; no lead above 6 body rows; the compact
    variant starts one step lower."""
    floor = 0.022 * slide_h_emu / EMU_PER_PT
    top = h1 / 1.25 if h1 else math.inf
    out = [s for s in dict.fromkeys((lead, body, small)) if s and floor - 0.05 <= s <= top + 1e-6]
    if n_body_rows > 6:
        out = [s for s in out if s <= body + 1e-6] or out
    if compact and len(out) > 1:
        out = out[1:]
    return out or [max(min(body, top), round(floor * 2) / 2)]


def measure_table(
    table: TableData,
    width_emu: int,
    font: Optional[str],
    sizes_desc: list[float],
    line_height: float,
    max_h_emu: int,
    fill_h_emu: int,
    min_row_h_emu: int,
    max_row_h_emu: int,
    *,
    header_bold: bool = True,
    max_size_pt: Optional[float] = None,
) -> TableFit:
    """Type size, column widths and row heights (EMU) of a table set ``width_emu`` wide.

    * Size: the largest of ``sizes_desc`` (≤ ``max_size_pt``) at which no word breaks inside its column (column minimum
      = longest word + 2 × cell padding; a number never wraps) and the rows fit ``max_h_emu``. Among those, the size
      that reads least ragged (short columns and headers on one line, long text in ≤ 3 lines), where every step down
      the scale costs as much as one wrapped cell and the first (lead) size is taken only when nothing wraps. When no
      size fits, the last one is used and ``overflow`` is set.
    * Columns: the label column takes its natural (one-line) width, capped at 0.4 of the table; numeric columns share
      one width, so do flag (✓/—) columns; spare width gives numbers and flags a little air and goes on to the label
      and text columns; long text columns absorb a shortfall. They sum to ``width_emu``.
    * Rows: text + 2 × padding, at least ``min_row_h_emu`` (as far as ``max_h_emu`` allows). Body rows share one height
      when that costs little, then grow evenly (shortest first, each ≤ ``max_row_h_emu``) until the table is
      ``fill_h_emu`` tall; the header is at least ¾ of a body row.

    Measured as :func:`add_table` sets the table: typeset with NBSP, units moved into the header, header (and total
    rows) bold unless ``header_bold=False``, boolean columns as ✓ / —.
    """
    n = max(len(table.columns), 1)
    width_pt = width_emu / EMU_PER_PT
    sizes = [s for s in sizes_desc if s and s > 0] or [12.0]
    if max_size_pt:
        sizes = [s for s in sizes if s <= max_size_pt + 1e-6] or [min(sizes)]
    if not table.columns:
        h = min_row_h_emu or int(sizes[-1] * 2 * EMU_PER_PT)
        return TableFit(sizes[-1], [width_emu], [h], content_w=width_emu)
    prep = _prepare(table)
    kinds = prep.kinds
    total_rows = {i for i, r in enumerate(table.rows) if is_total_row(r)}
    lh = line_height or 1.2
    min_row = (min_row_h_emu or 0) / EMU_PER_PT
    max_h = max_h_emu / EMU_PER_PT if max_h_emu and max_h_emu > 0 else math.inf
    candidates = []
    for size in sizes:
        mins, bnats, nats = _column_needs(prep, total_rows, width_pt, font, size, header_bold)
        widths = _allot(mins, bnats, nats, kinds, width_pt, size)
        fits = sum(mins) <= width_pt + 1e-6
        lines = _cell_lines(prep, total_rows, widths, font, size, header_bold)
        _, py = cell_padding_pt(size)
        natural = [max(ls) * size * lh + 2 * py for ls in lines]
        long_cols = {j for j in range(n) if _is_long(kinds[j], nats[j], width_pt)}
        candidates.append((size, widths, natural, fits and sum(natural) <= max_h + 1e-6, _wrap_score(lines, long_cols), min(sum(nats), width_pt)))
    ok = [(i, c) for i, c in enumerate(candidates) if c[3]]
    overflow = not ok
    if ok:
        # each step down the scale costs as much as one wrapped cell; the first (lead) size must be clean — a lead-size
        # header on two lines outweighs the body (a heavy band over light rows)
        def cost(ic):
            i, c = ic
            total, _ = c[4]
            return total + i + (1.0 if i == 0 and total > 0 and len(candidates) > 1 else 0.0), i

        size, widths, natural, _, _, content = min(ok, key=cost)[1]
    else:
        size, widths, natural, _, _, content = candidates[-1]
    heights = [max(h, min_row) for h in natural]
    if sum(heights) > max_h:
        # the minimum row height gives way before the type does: level the short rows up within the budget
        heights = _level(natural, max_h) if sum(natural) < max_h else list(natural)
    body = heights[1:]
    if body:
        # one body row height when that costs little: a steady rhythm reads cleaner than rows of mixed heights
        uniform = max(body)
        if uniform <= 1.8 * min(body) and heights[0] + uniform * len(body) <= max_h + 1e-6:
            heights = [heights[0]] + [uniform] * len(body)
    fill = min(fill_h_emu / EMU_PER_PT if fill_h_emu else 0.0, max_h)
    cap = max(max_row_h_emu / EMU_PER_PT, min_row) if max_row_h_emu else math.inf
    if sum(heights) < fill:
        heights = _grow_rows(heights, fill, cap)
    elif len(heights) > 1:
        # no growth: the header still reads as a band of the same rhythm (≥ ¾ of a body row) when there is room
        body_med = sorted(heights[1:])[len(heights[1:]) // 2]
        want = min(0.75 * body_med, 0.75 * cap)
        if want > heights[0] and sum(heights) - heights[0] + want <= max_h + 1e-6:
            heights[0] = want
    overflow = overflow or sum(heights) > max_h + 1.0
    too_dense = len(table.rows) > 8 or len(table.columns) > 5 or (min_row > 0 and min(heights) < min_row - 1.0)
    col_emu = [int(round(w * EMU_PER_PT)) for w in widths[:n]]
    col_emu[-1] += width_emu - sum(col_emu)
    row_emu = [int(round(h * EMU_PER_PT)) for h in heights]
    return TableFit(size, col_emu, row_emu, content_w=int(round(content * EMU_PER_PT)), overflow=overflow, too_dense=too_dense)


def neutral_run(rPr: etree._Element) -> None:
    """An engine-written run does not inherit the template's capitals, tracking or baseline shift (a master body style
    with cap="all" spc="500" would set «П О К А З А Т Е Л Ь»): capitals the design wants are typed as capitals."""
    rPr.set("cap", "none")
    rPr.set("spc", "0")
    rPr.set("baseline", "0")


def default_sizes(typography: Typography, slide_h_emu: int, style_size: Optional[float] = None) -> list[float]:
    """Candidate table sizes, largest first, when no composer chose them: body B = max(scale body, 0.026 H) down to
    the dense-table floor 0.022 H, through the template's own sizes in between."""
    h_pt = slide_h_emu / EMU_PER_PT
    body = typography.size_for("body", style_size or 12.0)
    b = round(max(body, 0.026 * h_pt) * 2) / 2
    floor = round(0.022 * h_pt * 2) / 2
    sparse = bool(getattr(typography, "derived_sizes", None))
    if sparse or (style_size or 0.0) > 0.026 * h_pt:
        # a bullet placeholder's 24 pt is not a table size on an 11″ slide: capped relative to the slide
        b = min(b, math.floor(0.026 * h_pt * 2) / 2)
    cands = {b, floor}
    cands.update(s for s in list(typography.sizes_used) + [st.size_pt for st in typography.scale] + [style_size or 0] if floor <= s <= b)
    return sorted(cands, reverse=True)


# ------------------------------------------------------------------------------------------------ colours


def _mix(fg: str, bg: str, share: float) -> str:
    a = [int(fg[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(bg[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{int(round(x * share + y * (1 - share))):02X}" for x, y in zip(a, b))


def _saturation(hex_: str) -> float:
    r, g, b = (int(hex_[i : i + 2], 16) / 255 for i in (0, 2, 4))
    mx, mn = max(r, g, b), min(r, g, b)
    return 0.0 if mx == 0 else (mx - mn) / mx


def _is_dark(hex_: str) -> bool:
    return relative_luminance(hex_) < 0.18


def _is_vivid(hex_: str) -> bool:
    """A saturated mid-tone ground (brand blue): neither the dark nor the light scheme reads well on it."""
    return _saturation(hex_) >= 0.5 and 0.05 <= relative_luminance(hex_) <= 0.5 and not _is_dark(hex_)


def _hue_gap(a: str, b: str) -> float:
    """Distance between two hues in degrees (0–180)."""
    import colorsys

    ha = colorsys.rgb_to_hsv(*(int(a[i : i + 2], 16) / 255 for i in (0, 2, 4)))[0] * 360
    hb = colorsys.rgb_to_hsv(*(int(b[i : i + 2], 16) / 255 for i in (0, 2, 4)))[0] * 360
    d = abs(ha - hb) % 360
    return min(d, 360 - d)


def _need(size_pt: float, bold: bool) -> float:
    """WCAG contrast for text: 3:1 for large text (≥ 18.67 pt, or ≥ 14 pt bold), else 4.5:1."""
    return 3.0 if size_pt >= 18.6 or (bold and size_pt >= 14.0) else 4.5


def _mix_until(fg: str, bg: str, target: float, start: float = 0.04, stop: float = 0.6, step: float = 0.02) -> str:
    """fg mixed into bg, the smallest share from ``start`` on whose contrast against bg reaches ``target``."""
    s = start
    while s <= stop + 1e-9:
        m = _mix(fg, bg, s)
        if contrast_ratio(m, bg) >= target:
            return m
        s += step
    return _mix(fg, bg, stop)


def _shade_until(fill: str, text: str, need: float) -> str:
    """The fill deepened a touch (towards black, same hue) until ``text`` reads on it; at most 30 % darker."""
    for i in range(1, 16):
        f = _mix("000000", fill, 0.02 * i)
        if contrast_ratio(text, f) >= need:
            return f
    return _mix("000000", fill, 0.3)


def _readable(color: str, bg: str, need: float, toward: str) -> Optional[str]:
    """``color`` when it reads on bg at ``need``, else it mixed towards ``toward`` (the text colour) until it does."""
    if contrast_ratio(color, bg) >= need:
        return color
    for s in (0.15, 0.3, 0.45, 0.6, 0.75):
        m = _mix(toward, color, s)
        if contrast_ratio(m, bg) >= need:
            return m
    return None


def _close(a: str, b: str, tol: int = 24) -> bool:
    return all(abs(int(a[i : i + 2], 16) - int(b[i : i + 2], 16)) <= tol for i in (0, 2, 4))


def _brand_pair(fill: str, text: str, pairs: Optional[set]) -> bool:
    """The template itself sets these two colours on each other (either way round, within a small tolerance): a
    pairing people already see in the brand, accepted at 3:1."""
    for f, t in pairs or ():
        if (_close(f, fill) and _close(t, text)) or (_close(t, fill) and _close(f, text)):
            return True
    return False


def header_text_on(fill_hex: str, requested: Optional[str], size_pt: Optional[float] = None, bold: bool = False, brand_pairs: Optional[set] = None) -> str:
    """Header text on a filled header row. A saturated brand fill carries white whenever white reads (the templates'
    own convention: VK blue #0077FF with black text is the wrong way round even though black has more contrast);
    otherwise the requested colour when it reads, else the better of white and black. «Reads» is 3:1 without a size,
    else 4.5:1 below 18.67 pt (14 pt bold) — or 3:1 for a pairing the template itself sets (``brand_pairs``)."""
    fill = fill_hex.upper()
    need = 3.0 if size_pt is None else _need(size_pt, bold)
    pairs = brand_pairs or set()

    def reads(c: str) -> bool:
        cr = contrast_ratio(c, fill)
        return cr >= need or (cr >= 3.0 and _brand_pair(fill, c, pairs))

    if _saturation(fill) >= 0.5 and 0.05 <= relative_luminance(fill) <= 0.45 and (not requested or relative_luminance(requested) < 0.2) and reads("FFFFFF"):
        return "FFFFFF"
    if requested and reads(requested.upper()):
        return requested.upper()
    return "FFFFFF" if contrast_ratio("FFFFFF", fill) >= contrast_ratio("000000", fill) else "000000"


@dataclass
class NativeTableStyle:
    """What a native sample table of the template says about tables (every field optional)."""

    header_fill_xml: Optional[etree._Element] = None  # a:solidFill / a:gradFill of the header cells, colours resolved
    header_fill_hex: Optional[str] = None  # its representative colour
    header_text_hex: Optional[str] = None
    header_bold: Optional[bool] = None
    header_rule: Optional[tuple[str, float]] = None  # (hex, pt) under an unfilled header
    rule: Optional[tuple[str, float]] = None  # (hex, pt) between body rows
    rule_votes: int = 0
    body_fill_hex: Optional[str] = None  # a fill every body cell carries (WS: dark rows on a black ground)
    accent_fill_hex: Optional[str] = None  # a highlighted row or column (Edu «Акцент»)
    accent_text_hex: Optional[str] = None

    def empty(self) -> bool:
        return not any((self.header_fill_xml is not None, self.header_bold is not None, self.header_rule, self.rule, self.body_fill_hex, self.accent_fill_hex))


@dataclass
class _Look:
    """Every colour of one table on one ground."""

    ground: str
    flat: bool  # the ground is one known colour (opaque tints leave no seams); else tints stay translucent
    row_bg: str  # what body text sits on: the ground, or the template's body cell fill
    body_fill: Optional[str]
    header_fill: Optional[str]
    header_fill_xml: Optional[etree._Element]
    header_text: str
    header_rule: tuple[str, float]
    text: str
    muted: str
    accent: str
    rule: tuple[str, float]
    band: Optional[str]
    tint: str
    tint_alpha: Optional[float]
    tint_mark: bool  # the tint is a neutral step (the accent would turn muddy on these rows): an accent bar marks the row
    hl_fill: Optional[str]
    hl_text: Optional[str]


def _pick_rule(cands: list[Optional[str]], row_bg: str, dark: bool) -> Optional[str]:
    for c in cands:
        if not c:
            continue
        c = c.upper()
        cr = contrast_ratio(c, row_bg)
        if _saturation(c) >= 0.25:
            continue  # a coloured hairline reads as decoration
        if dark and 1.2 <= cr <= 2.0 and relative_luminance(c) > relative_luminance(row_bg):
            return c
        if not dark and 1.12 <= cr <= 2.2:
            return c
    return None


def _resolve_look(
    style: TableStyleSpec,
    *,
    ground: str,
    text: Optional[str] = None,
    accent: Optional[str] = None,
    muted: Optional[str] = None,
    divider: Optional[str] = None,
    native: Optional[NativeTableStyle] = None,
    pairs: Optional[set] = None,
    flat: bool = True,
) -> _Look:
    ground = ground.upper()
    pairs = pairs or set()
    dark, vivid = _is_dark(ground), _is_vivid(ground)
    text = (text or style.body_text_hex or ("FFFFFF" if dark else "000000")).upper()
    if contrast_ratio(text, ground) < 3.0:
        text = "FFFFFF" if contrast_ratio("FFFFFF", ground) >= contrast_ratio("000000", ground) else "000000"
    if vivid and relative_luminance(text) < 0.2 and _brand_pair(ground, "FFFFFF", pairs) and contrast_ratio("FFFFFF", ground) >= 3.0:
        text = "FFFFFF"  # the template sets white on its brand ground: so does the table
    nat = native if native is not None and not native.empty() else None

    # body cell fill (a sample table whose rows are dark cards on a black ground)
    body_fill = None
    if nat and nat.body_fill_hex and 1.05 <= contrast_ratio(nat.body_fill_hex, ground) <= 1.8 and contrast_ratio(text, nat.body_fill_hex) >= 4.5:
        body_fill = nat.body_fill_hex.upper()
    row_bg = body_fill or ground
    row_dark = _is_dark(row_bg)

    # header: the sample's fill XML, else the template's header colour, else the accent — when it stands off the
    # ground; on a dark or brand ground with none of them, a light band carrying the ground colour
    header_fill = header_xml = None
    header_text = None
    if nat and nat.header_fill_xml is not None and nat.header_fill_hex and contrast_ratio(nat.header_fill_hex, ground) >= 1.3:
        header_fill, header_xml = nat.header_fill_hex.upper(), nat.header_fill_xml
        # the sample's own text colour on its own fill: a gradient's mean colour understates how it reads, so the
        # pairing is trusted down to 2.5:1 on that mean
        requested = nat.header_text_hex or style.header_text_hex
        ok = requested and contrast_ratio(requested, header_fill) >= (2.5 if nat.header_text_hex else 3.0)
        header_text = requested.upper() if ok else header_text_on(header_fill, requested, brand_pairs=pairs)
    elif style.header_fill_hex:  # a template without a header fill keeps its header unfilled (muted over a rule)
        cands = [style.header_fill_hex, accent]
        for c in cands:
            if not c:
                continue
            c = c.upper()
            if contrast_ratio(c, ground) < 1.3:
                continue
            if (vivid or dark) and _saturation(c) < 0.25 and relative_luminance(c) < 0.5:
                continue  # black on a brand blue, near-black on purple: never
            header_fill = c
            break
        if header_fill is None and (vivid or dark):
            header_fill = "FFFFFF"
        if header_fill:
            if relative_luminance(header_fill) > 0.8 and _saturation(header_fill) < 0.1 and (vivid or dark):
                header_text = ground if contrast_ratio(ground, header_fill) >= 3.0 else ("000000" if contrast_ratio("000000", header_fill) >= 4.5 else text)
            else:
                header_text = header_text_on(header_fill, style.header_text_hex, brand_pairs=pairs)

    # ✓ marks and rules: the accent when it reads on the rows; on a brand ground the text colour
    acc = (accent or style.header_fill_hex or text).upper()
    if vivid or contrast_ratio(acc, row_bg) < 3.0:
        alt = header_fill if header_fill and header_xml is None and _saturation(header_fill) >= 0.25 and contrast_ratio(header_fill, row_bg) >= 3.0 else None
        acc = alt or text
    mut = (muted or "").upper()
    if not mut or contrast_ratio(mut, row_bg) < 3.0 or mut == text or (relative_luminance(mut) < 0.4) != (relative_luminance(text) < 0.4):
        mut = _mix(text, row_bg, 0.62)
    if not header_fill:
        header_text = mut if contrast_ratio(mut, row_bg) >= 4.5 else text
    header_rule = (acc, HEADER_RULE_PT)
    if nat and nat.header_rule and contrast_ratio(nat.header_rule[0], ground) >= 1.5:
        header_rule = (nat.header_rule[0].upper(), max(nat.header_rule[1], 1.0))

    # hairlines between rows: the sample's rule, else a quiet neutral step from the rows' ground
    rule_hex = None
    rule_pt = RULE_PT
    if nat and nat.rule and 1.1 <= contrast_ratio(nat.rule[0], row_bg) <= 3.0:
        rule_hex, rule_pt = nat.rule[0].upper(), min(max(nat.rule[1], 0.5), 1.5)
    if rule_hex is None and vivid:
        rule_hex = _mix(text, row_bg, 0.35)
    if rule_hex is None:
        rule_hex = _pick_rule([style.border_hex, divider], row_bg, row_dark)
    if rule_hex is None:
        rule_hex = _mix_until(text, row_bg, 1.3 if row_dark else 1.5, start=0.08)

    # bands: the template's band colour when it is a subtle step the text still reads on, else a step of the text
    band = style.band_fill_hex.upper() if style.band_fill_hex else None
    band_lo = 1.1 if row_dark else 1.04  # a step of 1.05 on black reads as noise, not as a stripe
    if not band or vivid or not (band_lo <= contrast_ratio(band, row_bg) <= 1.6) or contrast_ratio(text, band) < 4.5 or (row_dark and relative_luminance(band) < relative_luminance(row_bg)):
        band = _mix_until(text, row_bg, 1.12 if row_dark else 1.06, start=0.03)

    # the recommended row: the sample's accent-row fill, else an accent tint that visibly stands off the rows
    hl_fill = hl_text = None
    if nat and nat.accent_fill_hex and nat.accent_text_hex and contrast_ratio(nat.accent_fill_hex, ground) >= 1.3 and contrast_ratio(nat.accent_text_hex, nat.accent_fill_hex) >= 3.0:
        hl_fill, hl_text = nat.accent_fill_hex.upper(), nat.accent_text_hex.upper()
    tint_src = acc if _saturation(acc) >= 0.25 else (style.header_fill_hex or acc).upper()
    tint_mark = False
    if contrast_ratio(tint_src, row_bg) < 1.5:  # the brand colour is the ground itself: lighten it with the text colour
        tint_src = text
    elif _saturation(row_bg) >= 0.3 and relative_luminance(row_bg) < 0.5 and _hue_gap(tint_src, row_bg) > 60:
        tint_src, tint_mark = text, True  # orange into forest green is olive: step the rows' own colour instead
    target = 1.3 if row_dark or vivid else 1.12
    tint = _mix_until(tint_src, row_bg, target, start=0.10, stop=0.45)
    tint_alpha = None
    if not flat and body_fill is None:
        share = next((s / 100 for s in range(10, 46, 2) if contrast_ratio(_mix(tint_src, row_bg, s / 100), row_bg) >= target), 0.45)
        tint, tint_alpha = tint_src, share
    return _Look(
        ground=ground, flat=flat, row_bg=row_bg, body_fill=body_fill, header_fill=header_fill, header_fill_xml=header_xml,
        header_text=(header_text or text).upper(), header_rule=header_rule, text=text, muted=mut, accent=acc,
        rule=(rule_hex.upper(), rule_pt), band=band, tint=tint, tint_alpha=tint_alpha, tint_mark=tint_mark, hl_fill=hl_fill, hl_text=hl_text,
    )


def table_style_for_ground(
    style: TableStyleSpec,
    *,
    ground_hex: str,
    text_hex: str,
    accent_hex: Optional[str] = None,
    divider_hex: Optional[str] = None,
    muted_hex: Optional[str] = None,
    native: Optional[NativeTableStyle] = None,
    brand_pairs: Optional[set] = None,
) -> TableStyleSpec:
    """The template's table style made to read on one ground.

    * Header: the sample table's fill (``native``) or the template's header colour or the accent — the first that
      stands off the ground (≥ 1.3:1) and is not a dark neutral on a dark or brand ground; on such a ground with none of
      them, a white band set in the ground colour. Text per :func:`header_text_on`.
    * Body text: the ground's text colour (white on a brand ground the template itself sets white on).
    * Rules: the sample's rule, else the template border or the divider when it is a quiet neutral step from the ground
      (1.12–2.2:1 on light grounds; 1.2–2.0:1 and lighter than the ground on dark ones), else the text colour stepped
      into the ground until it reads as a hairline (1.5:1 light, 1.3:1 dark); on a brand ground the text colour 35 %.
    * Band: the template band when it is a subtle step the text reads on, else the text colour stepped into the ground
      (1.06:1 light, 1.12:1 dark)."""
    look = _resolve_look(style, ground=ground_hex, text=text_hex, accent=accent_hex, muted=muted_hex, divider=divider_hex, native=native, pairs=brand_pairs)
    return style.model_copy(
        update={
            "header_fill_hex": look.header_fill,
            "header_text_hex": look.header_text,
            "body_text_hex": look.text,
            "band_fill_hex": look.band,
            "border_hex": look.rule[0],
        }
    )


# ------------------------------------------------------------------------------------------------ the template's own tables


def _local(el: etree._Element) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def _theme_scheme(slide: Slide) -> dict[str, str]:
    """clrScheme of the slide's master theme (dk1 … accent6) plus the master's colour map (bg1 → lt1 …)."""
    out: dict[str, str] = {}
    try:
        master = slide.slide_layout.slide_master
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    except Exception:  # noqa: BLE001
        return out
    scheme = theme.find(".//" + q("a:clrScheme"))
    if scheme is not None:
        for child in scheme:
            if not isinstance(child.tag, str) or not len(child):
                continue
            clr = child[0]
            val = clr.get("val") if _local(clr) == "srgbClr" else clr.get("lastClr")
            if val and len(val) == 6:
                out[_local(child)] = val.upper()
    cmap = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
    try:
        cm = master._element.find(q("p:clrMap"))
        if cm is not None:
            cmap.update(dict(cm.attrib))
    except Exception:  # noqa: BLE001
        pass
    for k, v in cmap.items():
        if v in out and k not in out:
            out[k] = out[v]
    return out


def _color_of(el: Optional[etree._Element], scheme: dict[str, str]) -> Optional[tuple[str, float]]:
    """(hex, alpha) of a colour element (srgbClr, schemeClr, sysClr, prstClr) with its modifiers applied."""
    if el is None:
        return None
    tag = _local(el)
    base = None
    if tag == "srgbClr":
        base = el.get("val")
    elif tag == "schemeClr":
        base = scheme.get(el.get("val") or "")
    elif tag == "sysClr":
        base = el.get("lastClr") or {"windowText": "000000", "window": "FFFFFF"}.get(el.get("val") or "")
    elif tag == "prstClr":
        base = _PRESET_COLORS.get(el.get("val") or "")
    if not base or len(base) != 6:
        return None
    try:
        hex_ = _apply_modifiers(base.upper(), el).upper()
    except Exception:  # noqa: BLE001
        hex_ = base.upper()
    a = el.find(q("a:alpha"))
    alpha = int(a.get("val", "100000")) / 100000 if a is not None else 1.0
    return hex_, alpha


def _fill_color(fill_el: Optional[etree._Element], scheme: dict[str, str]) -> Optional[tuple[str, float]]:
    """Representative (hex, alpha) of an a:solidFill / a:gradFill (a gradient: the mean of its stops)."""
    if fill_el is None:
        return None
    tag = _local(fill_el)
    if tag == "solidFill" and len(fill_el):
        return _color_of(fill_el[0], scheme)
    if tag == "gradFill":
        stops = [_color_of(gs[0], scheme) for gs in fill_el.iter(q("a:gs")) if len(gs)]
        stops = [s for s in stops if s]
        if not stops:
            return None
        rgb = [sum(int(s[0][i : i + 2], 16) for s in stops) / len(stops) for i in (0, 2, 4)]
        return "".join(f"{int(round(v)):02X}" for v in rgb), sum(s[1] for s in stops) / len(stops)
    return None


def _to_srgb(fill_el: etree._Element, scheme: dict[str, str]) -> etree._Element:
    """A copy of a fill whose theme colours are resolved to srgbClr (alpha kept): it reads the same under any master."""
    el = copy.deepcopy(fill_el)
    for c in list(el.iter()):
        if _local(c) in ("schemeClr", "sysClr", "prstClr"):
            res = _color_of(c, scheme)
            if res is None:
                continue
            new = etree.Element(q("a:srgbClr"))
            new.set("val", res[0])
            a = c.find(q("a:alpha"))
            if a is not None:
                new.append(copy.deepcopy(a))
            c.getparent().replace(c, new)
    return el


def _cell_fill_el(tc: etree._Element) -> Optional[etree._Element]:
    tcPr = tc.find(q("a:tcPr"))
    if tcPr is None:
        return None
    for ch in tcPr:
        if _local(ch) in ("solidFill", "gradFill"):
            return ch
        if _local(ch) == "noFill":
            return None
    return None


def _cell_text(tc: etree._Element) -> str:
    return "".join(t.text or "" for t in tc.iter(q("a:t"))).strip()


def _cell_run(tc: etree._Element, scheme: dict[str, str]) -> tuple[Optional[bool], Optional[str]]:
    """(bold, colour) of the first run with text in a cell; bold is None when the run does not say."""
    for r in tc.iter(q("a:r")):
        t = r.find(q("a:t"))
        if t is None or not (t.text or "").strip():
            continue
        rPr = r.find(q("a:rPr"))
        if rPr is None:
            return None, None
        b = rPr.get("b")
        sf = rPr.find(q("a:solidFill"))
        col = _color_of(sf[0], scheme) if sf is not None and len(sf) else None
        return (None if b is None else b in ("1", "true")), (col[0] if col else None)
    return None, None


def _cell_line(tc: etree._Element, side: str, scheme: dict[str, str]) -> Optional[tuple[str, float]]:
    tcPr = tc.find(q("a:tcPr"))
    ln = tcPr.find(q("a:" + side)) if tcPr is not None else None
    if ln is None:
        return None
    w = int(ln.get("w", "12700") or 0)
    sf = ln.find(q("a:solidFill"))
    col = _color_of(sf[0], scheme) if sf is not None and len(sf) else None
    if not col or col[1] < 0.3 or w <= 0:
        return None
    return col[0], round(w / EMU_PER_PT, 2)


def native_table_style(tbl: etree._Element, scheme: dict[str, str]) -> Optional[NativeTableStyle]:
    """What one native table (a:tbl) of the template says: header fill XML and text, header weight, the rule under an
    unfilled header, the rule between body rows, a fill every body cell carries, a highlighted row or column."""
    trs = tbl.findall(q("a:tr"))
    if len(trs) < 2:
        return None

    def cells(tr):
        return [tc for tc in tr.findall(q("a:tc")) if tc.get("hMerge") not in ("1", "true") and tc.get("vMerge") not in ("1", "true")]

    out = NativeTableStyle()
    head = cells(trs[0])
    fills = [_cell_fill_el(tc) for tc in head]
    visible = [(tc, f) for tc, f in zip(head, fills) if f is not None and (_fill_color(f, scheme) or ("", 0.0))[1] >= 0.3]
    if head and len(visible) >= 0.6 * len(head):
        tag = Counter(_local(f) for _, f in visible).most_common(1)[0][0]
        tc0, src = next(((tc, f) for tc, f in visible if _local(f) == tag and _cell_text(tc)), visible[0])
        rep = _fill_color(src, scheme)
        out.header_fill_xml = _to_srgb(src, scheme)
        out.header_fill_hex = rep[0] if rep else None
        bold, col = _cell_run(tc0, scheme)
        out.header_bold, out.header_text_hex = bold, col
    else:
        rules = Counter(r for r in (_cell_line(tc, "lnB", scheme) for tc in head) if r)
        if rules:
            (hex_, pt), _ = rules.most_common(1)[0]
            if pt >= 1.0:
                out.header_rule = (hex_, pt)
        bolds = [b for b in (_cell_run(tc, scheme)[0] for tc in head if _cell_text(tc)) if b is not None]
        if bolds:
            out.header_bold = Counter(bolds).most_common(1)[0][0]
    body = [cells(tr) for tr in trs[1:]]
    ruled = body[:-1] if len(body) > 1 else body
    n_ruled = sum(len(r) for r in ruled)
    rules = Counter(r for row in ruled for r in (_cell_line(tc, "lnB", scheme) for tc in row) if r)
    if rules:
        (hex_, pt), votes = rules.most_common(1)[0]
        if votes >= 0.5 * max(n_ruled, 1):
            out.rule, out.rule_votes = (hex_, pt), votes
    solid = [(_fill_color(_cell_fill_el(tc), scheme) or (None, 0.0)) for row in body for tc in row]
    counts = Counter(h for h, a in solid if h and a >= 0.5)
    n_body = len(solid)
    if counts and n_body:
        h, c = counts.most_common(1)[0]
        if c >= 0.7 * n_body:
            out.body_fill_hex = h

    def accent_of(group: list[etree._Element]) -> Optional[tuple[str, str]]:
        cols = [_fill_color(_cell_fill_el(tc), scheme) for tc in group]
        if len(group) < 2 or not all(c and c[1] >= 0.5 for c in cols):
            return None
        hexes = {c[0] for c in cols}
        if len(hexes) != 1:
            return None
        fill = hexes.pop()
        if _saturation(fill) < 0.45 or fill == out.body_fill_hex:
            return None
        texts = Counter(c for c in (_cell_run(tc, scheme)[1] for tc in group if _cell_text(tc)) if c)
        if not texts:
            return None
        txt = texts.most_common(1)[0][0]
        return (fill, txt) if contrast_ratio(txt, fill) >= 3.0 else None

    hit = next((a for a in (accent_of(row) for row in body) if a), None)
    if hit is None and body:
        n_cols = min(len(r) for r in body)
        hit = next((a for a in (accent_of([row[j] for row in body]) for j in range(1, n_cols)) if a), None)
    if hit:
        out.accent_fill_hex, out.accent_text_hex = hit
    return None if out.empty() else out


def _merge_native(found: list[NativeTableStyle]) -> Optional[NativeTableStyle]:
    if not found:
        return None
    out = NativeTableStyle()
    src = next((n for n in found if n.header_fill_xml is not None), None)
    if src is not None:
        out.header_fill_xml, out.header_fill_hex, out.header_text_hex, out.header_bold = src.header_fill_xml, src.header_fill_hex, src.header_text_hex, src.header_bold
    if out.header_bold is None:
        bolds = [n.header_bold for n in found if n.header_bold is not None]
        out.header_bold = Counter(bolds).most_common(1)[0][0] if bolds else None
    out.header_rule = next((n.header_rule for n in found if n.header_rule), None)
    votes: Counter = Counter()
    for n in found:
        if n.rule:
            votes[n.rule] += n.rule_votes
    if votes:
        out.rule, out.rule_votes = votes.most_common(1)[0]
    out.body_fill_hex = next((n.body_fill_hex for n in found if n.body_fill_hex), None)
    acc = next((n for n in found if n.accent_fill_hex), None)
    if acc is not None:
        out.accent_fill_hex, out.accent_text_hex = acc.accent_fill_hex, acc.accent_text_hex
    return None if out.empty() else out


@dataclass
class _TemplateCtx:
    native: Optional[NativeTableStyle] = None
    pairs: set = field(default_factory=set)  # (fill, text) colour pairs the template sets (pills, cards, table cells)


def _is_ours(frame: etree._Element, tbl: etree._Element) -> bool:
    nv = frame.find(".//" + q("p:cNvPr"))
    sid = tbl.find(q("a:tblPr") + "/" + q("a:tableStyleId"))
    return nv is not None and nv.get("name") == TABLE_NAME and sid is not None and sid.text == NO_STYLE_TABLE_ID


def _scan_presentation(prs) -> _TemplateCtx:
    found: list[NativeTableStyle] = []
    pairs: set = set()
    schemes: dict[str, dict[str, str]] = {}
    for s in prs.slides:
        try:
            key = str(s.slide_layout.slide_master.part.partname)
        except Exception:  # noqa: BLE001
            key = ""
        if key not in schemes:
            schemes[key] = _theme_scheme(s)
        scheme = schemes[key]
        root = s._element
        for frame in root.iter(q("p:graphicFrame")):
            tbl = frame.find(".//" + q("a:tbl"))
            if tbl is None or _is_ours(frame, tbl):
                continue
            nt = native_table_style(tbl, scheme)
            if nt is not None:
                found.append(nt)
            for tc in tbl.iter(q("a:tc")):
                fc = _fill_color(_cell_fill_el(tc), scheme)
                txt = _cell_run(tc, scheme)[1]
                if fc and fc[1] >= 0.5 and txt:
                    pairs.add((fc[0], txt))
        for sp in root.iter(q("p:sp")):
            spPr = sp.find(q("p:spPr"))
            sf = spPr.find(q("a:solidFill")) if spPr is not None else None
            fc = _fill_color(sf, scheme) if sf is not None else None
            if not fc or fc[1] < 0.5:
                continue
            for r in sp.iter(q("a:r")):
                t = r.find(q("a:t"))
                rPr = r.find(q("a:rPr"))
                if t is None or not (t.text or "").strip() or rPr is None:
                    continue
                rf = rPr.find(q("a:solidFill"))
                col = _color_of(rf[0], scheme) if rf is not None and len(rf) else None
                if col:
                    pairs.add((fc[0], col[0]))
    return _TemplateCtx(native=_merge_native(found), pairs=pairs)


def template_table_context(slide: Slide) -> _TemplateCtx:
    """The template's native table style and its colour pairings, read once per package from the slides it holds (a
    deck is built inside a copy of the template, so the sample slides are there while the deck is rendered)."""
    try:
        package = slide.part.package
    except Exception:  # noqa: BLE001
        return _TemplateCtx()
    cached = getattr(package, "_verstka_table_ctx", None)
    if isinstance(cached, _TemplateCtx):
        return cached
    try:
        ctx = _scan_presentation(package.presentation_part.presentation)
    except Exception:  # noqa: BLE001
        ctx = _TemplateCtx()
    try:
        setattr(package, "_verstka_table_ctx", ctx)
    except Exception:  # noqa: BLE001
        pass
    return ctx


def _slide_ground(slide: Slide) -> Optional[tuple[str, bool]]:
    """(colour, flat) of the slide background — the slide's own, else its layout's, else its master's; flat is False
    for a gradient (its mean colour is returned); None for a picture or pattern background."""
    try:
        holders = [slide, slide.slide_layout, slide.slide_layout.slide_master]
    except Exception:  # noqa: BLE001
        holders = [slide]
    scheme = _theme_scheme(slide)
    for h in holders:
        try:
            cSld = h._element.find(q("p:cSld"))
        except Exception:  # noqa: BLE001
            continue
        bg = cSld.find(q("p:bg")) if cSld is not None else None
        if bg is None:
            continue
        bgPr = bg.find(q("p:bgPr"))
        if bgPr is not None:
            for ch in bgPr:
                tag = _local(ch)
                if tag in ("solidFill", "gradFill"):
                    c = _fill_color(ch, scheme)
                    return (c[0], tag == "solidFill") if c else None
                if tag in ("blipFill", "pattFill"):
                    return None
            return None
        ref = bg.find(q("p:bgRef"))
        if ref is not None and len(ref):
            c = _color_of(ref[0], scheme)
            return (c[0], (ref.get("idx") or "1001") == "1001") if c else None
    return None


# ------------------------------------------------------------------------------------------------ XML helpers


def _solid(parent: etree._Element, hex_: str, alpha: Optional[float] = None) -> None:
    sf = etree.SubElement(parent, q("a:solidFill"))
    c = etree.SubElement(sf, q("a:srgbClr"))
    c.set("val", hex_.upper())
    if alpha is not None and alpha < 1.0:
        etree.SubElement(c, q("a:alpha")).set("val", str(int(round(alpha * 100000))))


_LN_TAGS = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")
_FILL_TAGS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")
Fill = Union[None, tuple, etree._Element]


def _style_tcpr(cell, lines: dict[str, Optional[tuple[str, float]]], fill: Fill) -> None:
    """Rewrite a cell's borders and fill in schema order (lnL, lnR, lnT, lnB, fill). ``lines`` maps L/R/T/B to
    (hex, width pt) or None (no line); ``fill`` is (hex, alpha), a fill element to copy, or None (transparent)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in _LN_TAGS + ("a:lnTlToBr", "a:lnBlToTr") + _FILL_TAGS:
        for old in tcPr.findall(q(tag)):
            tcPr.remove(old)
    new: list[etree._Element] = []
    for side, tag in zip("LRTB", _LN_TAGS):
        spec = lines.get(side)
        ln = etree.Element(q(tag))
        if spec:
            hex_, w_pt = spec
            ln.set("w", str(int(round(w_pt * EMU_PER_PT))))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")
            _solid(ln, hex_)
            etree.SubElement(ln, q("a:prstDash")).set("val", "solid")
        else:
            ln.set("w", "0")
            etree.SubElement(ln, q("a:noFill"))
        new.append(ln)
    if isinstance(fill, etree._Element):
        new.append(copy.deepcopy(fill))
    else:
        holder = etree.Element("holder")
        if fill:
            _solid(holder, fill[0], fill[1])
        else:
            etree.SubElement(holder, q("a:noFill"))
        new.extend(list(holder))
    for i, el in enumerate(new):
        tcPr.insert(i, el)


def _set_typeface(rPr: etree._Element, family: str) -> None:
    """latin + ea + cs typefaces in schema order (after the fill), theme references mapped per script."""
    for tag in ("a:latin", "a:ea", "a:cs"):
        for old in rPr.findall(q(tag)):
            rPr.remove(old)
    if family.startswith("+mj") or family.startswith("+mn"):
        base = family[:3]
        faces = {"a:latin": base + "-lt", "a:ea": base + "-ea", "a:cs": base + "-cs"}
    else:
        faces = {"a:latin": family, "a:ea": family, "a:cs": family}
    # schema: ln, fill, effect, highlight, uLnTx/uLn, uFillTx/uFill, latin, ea, cs, sym, hlink…
    after = None
    for tag in ("a:uFill", "a:uFillTx", "a:uLn", "a:uLnTx", "a:highlight", "a:effectDag", "a:effectLst", "a:solidFill", "a:noFill", "a:gradFill", "a:ln"):
        after = rPr.find(q(tag))
        if after is not None:
            break
    idx = list(rPr).index(after) + 1 if after is not None else 0
    for k, (tag, face) in enumerate(faces.items()):
        el = etree.Element(q(tag))
        el.set("typeface", face)
        rPr.insert(idx + k, el)


def _gradient_slice(grad: etree._Element, a: float, b: float) -> etree._Element:
    """The part of a header-row gradient that falls on one cell (row fractions a…b): a mostly horizontal gradient runs
    once across the whole header instead of restarting in every cell. A vertical one is kept as is (it reads the same
    in every cell of one row)."""
    lin = grad.find(q("a:lin"))
    if lin is None or b <= a:
        return grad
    ang = math.radians(int(lin.get("ang", "0") or 0) / 60000)
    if abs(math.cos(ang)) < 0.35:
        return grad
    stops = []
    for gs in grad.iter(q("a:gs")):
        clr = gs[0] if len(gs) else None
        if clr is None or _local(clr) != "srgbClr" or not clr.get("val"):
            return grad
        al = clr.find(q("a:alpha"))
        stops.append((int(gs.get("pos", "0")) / 100000, clr.get("val").upper(), int(al.get("val")) / 100000 if al is not None else 1.0))
    if len(stops) < 2:
        return grad
    stops.sort()
    fwd = math.cos(ang) >= 0

    def at(t: float) -> tuple[str, float]:
        p = t if fwd else 1 - t
        if p <= stops[0][0]:
            return stops[0][1], stops[0][2]
        for (p0, c0, a0), (p1, c1, a1) in zip(stops, stops[1:]):
            if p0 <= p <= p1:
                f = (p - p0) / (p1 - p0) if p1 > p0 else 0.0
                return _mix(c1, c0, f), a0 + (a1 - a0) * f
        return stops[-1][1], stops[-1][2]

    inner = sorted(t for t in ((p if fwd else 1 - p) for p, _, _ in stops) if a + 1e-6 < t < b - 1e-6)
    new = etree.Element(q("a:gradFill"))
    for k, v in grad.attrib.items():
        new.set(k, v)
    lst = etree.SubElement(new, q("a:gsLst"))
    for t in [a] + inner + [b]:
        c, al = at(t)
        gs = etree.SubElement(lst, q("a:gs"))
        gs.set("pos", str(int(round((t - a) / (b - a) * 100000))))
        clr = etree.SubElement(gs, q("a:srgbClr"))
        clr.set("val", c)
        if al < 0.999:
            etree.SubElement(clr, q("a:alpha")).set("val", str(int(round(al * 100000))))
    lin_new = etree.SubElement(new, q("a:lin"))
    lin_new.set("ang", "0")
    lin_new.set("scaled", "0")
    return new


# ------------------------------------------------------------------------------------------------ drawing


_AUTO = object()


def _header_fits_bold(heads: list[str], widths: list[int], gut: list[float], family: Optional[str], size: float) -> bool:
    """A header measured regular may still be set bold when no word breaks and no cell gains a line."""
    px, _ = cell_padding_pt(size)
    for j, h in enumerate(heads):
        w = widths[j] / EMU_PER_PT - 2 * px - gut[j]
        if _longest_word(h, family, size, True) > w + 0.01 or _lines(h, family, size, True, w) > _lines(h, family, size, False, w):
            return False
    return True


def add_table(
    slide: Slide,
    bbox: Bbox,
    table: TableData,
    style: TableStyleSpec,
    typography: Typography,
    font_family: Optional[str] = None,
    *,
    col_widths: Optional[list[int]] = None,
    row_heights: Optional[list[int]] = None,
    size_pt: Optional[float] = None,
    band_every: Optional[int] = None,
    accent_hex: Optional[str] = None,
    accent_text_hex: Optional[str] = None,
    muted_hex: Optional[str] = None,
    header_bold: Optional[bool] = None,
    boolean_marks: bool = True,
    highlight_row: Optional[int] = None,
    auto_highlight: bool = True,
    ground_hex: Optional[str] = None,
    native: object = _AUTO,
    highlight_col: Optional[int] = None,
):
    """A native table. The composer passes the geometry from :func:`measure_table` (column widths, row heights, one
    type size); without it the table measures itself inside ``bbox`` (rows capped at 0.10 H, never taller than
    needed to fill the box).

    Look (:func:`_resolve_look`): the template's sample table when it has one (``native``, read from the package by
    default), else ``style`` — both adjusted to the ground (``ground_hex``, else the slide background, else the ground
    the body text implies). The header row is filled (the sample's fill XML, gradients included) or, with no fill, set
    in the muted colour over an accent rule; vertically centred and aligned like its column; bold when the template
    sets its headings in bold (:func:`template_bold`), ``header_bold=True`` or the sample header is bold and still fits.
    Header text on a brand fill reads at 4.5:1 below 18.67 pt (3:1 for a pairing the template itself sets), else the
    fill is deepened a touch. Numeric columns are right-aligned (their shared unit moved into the header), short
    categorical columns centred and boolean ones set as an accent ✓ / muted —. Body rows are separated by hairlines;
    ``band_every`` rows get the band colour instead (default: every 2nd row from 6 body rows; never fewer than two
    stripes; never in a shorter table with a recommended row). A total row (Итого/Всего/Total/Сумма) sits under a thin
    accent rule, bold only in a bold template. The recommended row («Наш …») takes the sample's accent-row fill, else a
    tint that visibly stands off the ground, with a label in ``accent_text_hex``/the accent at a readable contrast (or
    an accent bar when no accent reads). ``highlight_col`` (the target column, see :func:`emphasis_column`): its body
    cells take the same tint and their values the accent at a readable contrast, an unfilled header the accent too.
    Margins scale with the size; every run names its latin/ea/cs typeface.
    """
    n_rows = len(table.rows) + 1
    n_cols = max(len(table.columns), 1)
    family = font_family or typography.primary_family
    hb = template_bold(typography) if header_bold is None else bool(header_bold)
    slide_h = _slide_height(slide) or int(bbox.h * 1.6) or 6858000
    if size_pt is None or not col_widths or not row_heights:
        sizes = [size_pt] if size_pt else default_sizes(typography, slide_h, style.font_size_pt)
        m_size, m_cols, m_rows = measure_table(
            table,
            bbox.w,
            family,
            sizes,
            typography.line_height,
            max_h_emu=bbox.h,
            fill_h_emu=bbox.h,
            min_row_h_emu=int(0.065 * slide_h),
            max_row_h_emu=int(0.10 * slide_h),
            header_bold=hb,
        )
        size_pt = size_pt or m_size
        col_widths = col_widths if col_widths and len(col_widths) == n_cols else m_cols
        row_heights = row_heights if row_heights and len(row_heights) == n_rows else m_rows
    size = float(size_pt)
    widths = [int(w) for w in col_widths] if len(col_widths) == n_cols else [int(bbox.w / n_cols)] * n_cols
    heights = [int(h) for h in row_heights] if len(row_heights) == n_rows else [int(bbox.h / n_rows)] * n_rows

    # ---- the look
    ctx = template_table_context(slide) if native is _AUTO else _TemplateCtx(native=native if isinstance(native, NativeTableStyle) else None)
    nat = ctx.native
    body_req = (style.body_text_hex or "").upper() or None
    flat = True
    ground = ground_hex.upper() if ground_hex else None
    if ground is None:
        found = _slide_ground(slide)
        if found and (not body_req or contrast_ratio(body_req, found[0]) >= 3.0):
            ground, flat = found
        else:  # the table sits on something else (a card, a picture): the ground its text implies
            ground = "FFFFFF" if not body_req or relative_luminance(body_req) < 0.4 else "000000"
            flat = False
    look = _resolve_look(style, ground=ground, text=body_req, accent=accent_hex, muted=muted_hex, native=nat, pairs=ctx.pairs, flat=flat)

    prep = _prepare(table, boolean_marks) if table.columns else _Prep(["text"] * n_cols, [""] * n_cols, [[""] * n_cols for _ in table.rows], [False] * n_cols)
    kinds = prep.kinds + ["text"] * (n_cols - len(prep.kinds))
    heads = prep.heads + [""] * (n_cols - len(prep.heads))
    rows = [r + [""] * (n_cols - len(r)) for r in prep.rows]
    booleans = prep.booleans + [False] * (n_cols - len(prep.booleans))
    gut = _gutters(kinds, size)
    aligns = [_align_for(k) for k in kinds]

    total_bold = hb  # measured so: a total row follows the template's weight, never the sample header's
    # the sample header's weight, when it only takes weight away or still fits
    if nat is not None and nat.header_bold is not None and nat.header_bold != hb:
        hb = nat.header_bold if not nat.header_bold else _header_fits_bold(heads, widths, gut, family, size)
    header_fill = look.header_fill
    header_text = look.header_text
    if header_fill and look.header_fill_xml is None:
        need = _need(size, hb)
        cr = contrast_ratio(header_text, header_fill)
        if cr < need and not (cr >= 3.0 and _brand_pair(header_fill, header_text, ctx.pairs)):
            if _saturation(header_fill) >= 0.25 and relative_luminance(header_text) > 0.5:
                header_fill = _shade_until(header_fill, header_text, need)  # white on brand blue: deepen the blue a touch
            else:
                header_text = header_text_on(header_fill, header_text, size, hb, ctx.pairs)

    gf = slide.shapes.add_table(n_rows, n_cols, Emu(bbox.x), Emu(bbox.y), Emu(sum(widths)), Emu(sum(heights)))
    gf.name = TABLE_NAME
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    style_id = tblPr.find(q("a:tableStyleId"))
    if style_id is None:
        style_id = etree.SubElement(tblPr, q("a:tableStyleId"))
    style_id.text = NO_STYLE_TABLE_ID
    for r, h in zip(tbl.rows, heights):
        r.height = Emu(h)
    for c, w in zip(tbl.columns, widths):
        c.width = Emu(w)

    px, py = cell_padding_pt(size)
    highlight = highlight_row if highlight_row is not None else (recommended_row(table) if auto_highlight else None)
    n_body = len(rows)
    every = band_every if band_every is not None else (2 if n_body >= 6 else 0)
    # bands replace the hairlines (a closing rule ends the table); fewer than two stripes read as a highlight, and in a
    # short table the recommended row's fill is the one fill that means something
    banding = bool(every) and every > 0 and look.band is not None and n_body // every >= 2 and n_body >= 4 and not (highlight is not None and n_body < 6)

    # the recommended row: the sample's accent fill, else the tint; a label that reads on it
    hl_full = highlight is not None and look.hl_fill is not None
    solid_head = look.header_fill_xml is None or _local(look.header_fill_xml) != "gradFill"
    if hl_full and header_fill and solid_head and contrast_ratio(header_fill, look.hl_fill) < 1.3:
        hl_full = False  # the sample's accent row is the solid header's colour: under it, it reads as a second header
    tint_bg = look.tint if look.tint_alpha is None else _mix(look.tint, look.row_bg, look.tint_alpha)
    label_hex, mark_bar = None, False
    if highlight is not None and not hl_full:
        need = _need(size, hb)
        # the tint's own hue first (made darker/lighter until it reads), the deck's accent-for-text colour next
        for cand in (look.accent, accent_text_hex):
            if cand and _saturation(cand) >= 0.2:
                label_hex = _readable(cand.upper(), tint_bg, need, look.text)
                if label_hex:
                    break
        if label_hex is None or _saturation(label_hex) < 0.2:
            label_hex, mark_bar = look.text, True
        mark_bar = mark_bar or look.tint_mark
    check_on_tint = _readable(look.accent, tint_bg, 3.0, look.text) or look.text
    hcol = highlight_col if highlight_col is not None and 0 < highlight_col < n_cols else None
    col_text = None
    if hcol is not None:
        need_c = _need(size, False)
        for cand in (look.accent, accent_text_hex):
            if cand and _saturation(cand) >= 0.2:
                col_text = _readable(cand.upper(), tint_bg, need_c, look.text)
                if col_text and _saturation(col_text) >= 0.2:
                    break
                col_text = None
        col_text = col_text or look.text
    separator = None
    if hl_full and highlight == 0 and header_fill and contrast_ratio(header_fill, look.hl_fill) < 1.3:
        separator = (look.ground, SEPARATOR_PT)  # a gap between two bands of one colour

    def write(cell, text: str, *, bold: bool, color_hex: str, align: PP_ALIGN, extra_left: float = 0.0) -> None:
        cell.text = ""
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = False
        run.font.color.rgb = RGBColor.from_string(color_hex)
        if family:
            _set_typeface(run._r.get_or_add_rPr(), family)
        neutral_run(run._r.get_or_add_rPr())
        cell.margin_left = Emu(int((px + extra_left) * EMU_PER_PT))
        cell.margin_right = Emu(int(px * EMU_PER_PT))
        cell.margin_top = cell.margin_bottom = Emu(int(py * EMU_PER_PT))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE  # tcPr@anchor: the cell property PowerPoint reads, not bodyPr

    # header
    edges = [sum(widths[:j]) / max(sum(widths), 1) for j in range(n_cols + 1)]
    for j in range(n_cols):
        cell = tbl.cell(0, j)
        h_color = header_text if header_fill else look.header_text
        if j == hcol and not header_fill:
            h_color = col_text or h_color  # an unfilled header names the target column in the accent
        write(cell, heads[j], bold=hb, color_hex=h_color, align=aligns[j], extra_left=gut[j])
        if header_fill:
            xml = look.header_fill_xml
            if xml is not None and _local(xml) == "gradFill":
                xml = _gradient_slice(xml, edges[j], edges[j + 1])
            fill: Fill = xml if xml is not None else (header_fill, None)
            _style_tcpr(cell, {"B": separator} if separator else {}, fill)
        else:
            _style_tcpr(cell, {"B": look.header_rule}, None)

    total_idx = {i for i, row in enumerate(table.rows) if is_total_row(row)}
    rule = look.rule
    for i, row in enumerate(rows):
        r = i + 1
        is_total = i in total_idx
        is_hl = highlight == i and not is_total
        banded = banding and (i + 1) % every == 0
        if is_hl and hl_full:
            fill = (look.hl_fill, None)
        elif is_hl:
            fill = (look.tint, look.tint_alpha)
        elif banded:
            fill = (look.band, None)
        elif look.body_fill:
            fill = (look.body_fill, None)
        else:
            fill = None
        for j in range(n_cols):
            text = row[j]
            color = look.text
            bold = is_total and total_bold
            cell_fill = fill
            if j == hcol and not (is_hl and hl_full):
                cell_fill = (look.tint, look.tint_alpha)
                if not (booleans[j] and text in (CHECK, DASH)):
                    color = col_text or look.text
            if is_hl and hl_full:
                color = look.hl_text
                bold = bold or (j == 0 and hb)
            elif booleans[j] and text in (CHECK, DASH):
                color = (check_on_tint if is_hl else look.accent) if text == CHECK else look.muted
            elif is_hl and j == 0:
                color, bold = label_hex or look.text, hb
            cell = tbl.cell(r, j)
            write(cell, text, bold=bold, color_hex=color, align=aligns[j], extra_left=gut[j])
            lines: dict[str, Optional[tuple[str, float]]] = {}
            if (i + 1) in total_idx:  # the rule above the next total row is drawn by both cells sharing the edge
                lines["B"] = (look.accent, TOTAL_RULE_PT)
            elif not banding or i == n_body - 1:
                lines["B"] = rule
            if is_total:
                lines["T"] = (look.accent, TOTAL_RULE_PT)
            elif is_hl and separator:
                lines["T"] = separator
            if is_hl and mark_bar and j == 0:
                lines["L"] = (look.accent, MARK_RULE_PT)
            _style_tcpr(cell, lines, cell_fill)
    return gf


def _slide_height(slide: Slide) -> Optional[int]:
    try:
        return int(slide.part.package.presentation_part.presentation.slide_height)
    except Exception:  # noqa: BLE001
        return None
