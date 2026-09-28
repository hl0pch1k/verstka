"""Synth renderer: compose a slide from template tokens and components when no sample pattern fits."""

from __future__ import annotations

import copy
import dataclasses
import math
from collections import Counter
from dataclasses import dataclass
from typing import Optional

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Emu, Pt

from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.assets_pick import pick_asset
from verstka.rendering.charts import add_chart
from verstka.rendering.clone import _fill_hex, _SlideCtx, clear_width, renumber_page_chrome
from verstka.rendering.fonts import text_width_pt
from verstka.rendering.deck import DeckBuilder, element_bbox, is_nested, remove_element, slide_shape_elements
from verstka.rendering.fit import fit_size
from verstka.rendering.images import insert_picture
from verstka.rendering.tables import add_table
from verstka.rendering.textfill import ParagraphSpec, fill_text
from verstka.schemas.common import EMU_PER_INCH, EMU_PER_PT, Bbox, Family, PatternKind, SlotRole, contrast_ratio
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideItem
from verstka.schemas.template import Pattern, TemplateManifest

_DEFAULT_RADIUS_EMU = int(0.12 * EMU_PER_INCH)
_CONTENT_KINDS = (PatternKind.bullets, PatternKind.cards, PatternKind.freeform, PatternKind.two_column, PatternKind.stat_row)


def card_adj(radius_emu: int, w: int, h: int) -> float:
    """roundRect adjustment for an absolute corner radius: relative to the shorter side, capped so cards never turn into capsules."""
    side = min(w, h)
    if radius_emu <= 0 or side <= 0:
        return 0.0
    return min(radius_emu / side, 0.25)


def _card_radius_emu(manifest: TemplateManifest) -> int:
    """Corner radius of the template's cards as an absolute size.

    A sample's `adj` is relative to the sample's shorter side, so it is converted with the sample card size;
    when only a ratio is known (e.g. from pill-shaped chips) a conventional 0.12 inch is used instead.
    """
    card = manifest.components.card
    if card and card.radius is not None and card.width_frac and card.height_frac:
        side = min(card.width_frac * manifest.slide_size.w, card.height_frac * manifest.slide_size.h)
        return int(max(0.0, min(card.radius, 0.5)) * side)
    return _DEFAULT_RADIUS_EMU


class _Palette:
    def __init__(self, manifest: TemplateManifest, family: Family) -> None:
        t = manifest.tokens
        self.family = family
        self.bg = t.color_for("background.dark" if family == Family.dark else "background.light") or ("000000" if family == Family.dark else "FFFFFF")
        text = t.color_for("text.primary") or ("FFFFFF" if family == Family.dark else "000000")
        if contrast_ratio(text, self.bg) < 3.0:
            text = "FFFFFF" if family == Family.dark else "000000"
        self.text = text
        sec = t.color_for("text.secondary")
        self.text2 = sec if sec and contrast_ratio(sec, self.bg) >= 3.0 else self.text
        from verstka.rendering.compose import palette_accents

        self.accent = palette_accents(t)[0]
        surf = t.color_for("surface")
        self.surface = surf if surf and 1.02 < contrast_ratio(surf, self.bg) < 2.5 else None
        card = manifest.components.card
        self.card_fill = (card.fill_hex if card and card.fill_hex and 1.0 < contrast_ratio(card.fill_hex, self.bg) < 3.0 else self.surface)
        self.card_line = card.line_hex if card else None
        self.radius_emu = _card_radius_emu(manifest)


def _family_for(oslide: OutlineSlide, manifest: TemplateManifest) -> Family:
    fams = Counter(p.family for p in manifest.patterns if p.kind == oslide.kind)
    if fams:
        return fams.most_common(1)[0][0]
    fams = Counter(p.family for p in manifest.patterns)
    return fams.most_common(1)[0][0] if fams else Family.light


def _layout_family(manifest: TemplateManifest, part: str, default: Family) -> Family:
    fams = Counter(p.family for p in manifest.patterns if p.layout_part == part)
    return fams.most_common(1)[0][0] if fams else default


def _layout_for(builder: DeckBuilder, manifest: TemplateManifest, family: Family, kind: PatternKind):
    """(layout, family): the layout used by the most sample slides of this family, searched over every master.

    Preference: same kind → content kinds → any kind of the family → any family (then the returned family is the
    one of the samples on that layout, so the palette matches its background). When no sample resolves to a
    layout of the package, a title-only-like layout is chosen — never blindly the last one.
    """
    by_partname = {str(l.part.partname).lstrip("/"): l for l in builder.layouts()}
    preferred = builder.__dict__.get("_prefer_layout", {}).get(family)
    if preferred is not None and kind not in _COVER_KINDS:
        # every sample of the family carries art over the content band: the package's cleanest layout
        return preferred, _layout_family(manifest, str(preferred.part.partname).lstrip("/"), family)
    known = [p for p in manifest.patterns if p.layout_part and p.layout_part in by_partname]
    tiers = [
        [p for p in known if p.family == family and p.kind == kind],
        [p for p in known if p.family == family and p.kind in _CONTENT_KINDS],
        [p for p in known if p.family == family],
        known,
    ]
    for prefs in tiers:
        if prefs:
            part = Counter(p.layout_part for p in prefs).most_common(1)[0][0]
            return by_partname[part], _layout_family(manifest, part, family)

    def rank(layout):
        phs = list(layout.placeholders)
        titled = any(str(ph.placeholder_format.type).split(".")[-1].split(" ")[0] in ("TITLE", "CENTER_TITLE") for ph in phs)
        return (0 if titled else 1, len(phs))

    layouts = builder.layouts()
    return (min(layouts, key=rank) if layouts else builder.prs.slide_layouts[0]), family


def _add_layout_slide(builder: DeckBuilder, layout) -> Slide:
    """A new slide on a layout, its placeholders' inherited geometry written out (contract C1:
    DeckBuilder.add_layout_slide when it exists)."""
    add = getattr(builder, "add_layout_slide", None)
    if add is not None:
        return add(layout)
    slide = builder.prs.slides.add_slide(layout)
    builder.created.append(slide)
    for shp in slide.shapes:
        if shp.is_placeholder:
            _materialize(shp)
    return slide


_CANVAS_KINDS = {
    "title": (PatternKind.title, PatternKind.section, PatternKind.thanks),
    "section": (PatternKind.section,),
    "thanks": (PatternKind.thanks, PatternKind.title, PatternKind.section),
}
_NOT_CONTENT_CANVAS = {PatternKind.title, PatternKind.section, PatternKind.thanks, PatternKind.quote}


def _top_title(p: Pattern) -> Optional[object]:
    titles = [s for s in p.slots if s.role == SlotRole.title]
    return min(titles, key=lambda s: s.bbox.y) if titles else None


def _grid_x(manifest: TemplateManifest) -> float:
    """The deck's heading grid: the left edge most content samples start their heading at (a share of the width).
    Centred and narrow headings (a title set in a column, a label) do not vote; no vote → the safe area's edge."""
    votes: Counter = Counter()
    for p in _content_patterns(manifest):
        t = _top_title(p)
        if t is None or (t.style.align or "l") == "ctr" or t.bbox.w < 0.45:
            continue
        votes[round(t.bbox.x, 2)] += 1
    return float(votes.most_common(1)[0][0]) if votes else float(manifest.tokens.spacing.safe_area.x)


def _on_grid(manifest: TemplateManifest, title) -> bool:
    """A content canvas whose heading stands on the deck's grid: at its left edge (± 4 % of the width), not a narrow
    centred box (a composed slide takes its content column from the heading)."""
    gx = _grid_x(manifest)
    return abs(title.bbox.x - gx) <= 0.04 and not ((title.style.align or "l") == "ctr" and title.bbox.w < 0.45)


def _canvas_for(manifest: TemplateManifest, family: Family, comp: str) -> Optional[Pattern]:
    """The sample slide a synthesized slide is drawn on: same family, a writable title at the top, as little content
    as possible. Cloning it keeps what a layout alone would lose — backgrounds and chrome that a designer drew on
    the slides themselves (hand-made decks, picture backgrounds) and the look of the template's headings. A content
    canvas keeps the deck's heading grid (its heading at the grid's left edge) unless no such canvas exists."""
    kinds = _CANVAS_KINDS.get(comp)
    if kinds is None:
        on = _canvas_among(manifest, family, None, grid=True)
        return on if on is not None else _canvas_among(manifest, family, None, grid=False)
    return _canvas_among(manifest, family, kinds, grid=False)


def _canvas_among(manifest: TemplateManifest, family: Family, kinds, grid: bool) -> Optional[Pattern]:
    best: Optional[tuple] = None
    for p in manifest.patterns:
        if p.family != family or p.quality < 0.5 or p.reference:
            continue
        title = _top_title(p)
        if title is None:
            continue
        if kinds is not None:
            if p.kind not in kinds:
                continue
            rank = kinds.index(p.kind)
        else:
            if p.kind in _NOT_CONTENT_CANVAS or title.bbox.y > 0.3 or title.bbox.w < 0.2:
                continue
            if grid and not _on_grid(manifest, title):
                continue
            rank = 0
        big_pictures = sum(1 for s in p.slots if s.role == SlotRole.image and s.bbox.area >= 0.05) + sum(1 for b in p.decor_boxes if b.area >= 0.05)
        clutter = len(p.slots) + sum(len(g.member_shape_ids) for g in p.repeat_groups) + len(p.decor_assets) + 5 * big_pictures
        key = (rank, clutter, -p.quality, p.source_slide)
        if best is None or key < best[0]:
            best = (key, p)
    return best[1] if best else None


def _readable(ground: str, preferred: list) -> str:
    cands = [c for c in preferred if c] + ["000000", "FFFFFF"]
    return next((c for c in cands if contrast_ratio(c, ground) >= 4.5), max(cands, key=lambda c: contrast_ratio(c, ground)))


def _paints(el) -> bool:
    spPr = el.find(q("p:spPr"))
    if spPr is not None:
        if spPr.find(q("a:noFill")) is not None:
            return False
        if any(spPr.find(q(t)) is not None for t in ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")):
            return True
    ref = el.find(q("p:style") + "/" + q("a:fillRef"))
    return ref is not None and (ref.get("idx") or "0") != "0"


def _veil(el) -> bool:
    """A see-through solid shape without an outline (a white bubble at 20–30 %): the ground shows through it."""
    from verstka.rendering.clone import _veil as veil

    try:
        return veil(el)
    except Exception:  # noqa: BLE001
        return False


def _cover_art_ids(els: dict, title_box: Optional[tuple], W: int, H: int) -> set[str]:
    """The pictures of a cover sample that are its design, not its content: a photo or an illustration of at least 3 %
    of the slide that stands clear of the title (a photo beside a title card — the template's own cover), never an
    empty picture placeholder. A composed cover keeps them; the words are the deck's."""
    keep: set[str] = set()
    for sid, el in els.items():
        tag = etree.QName(el).localname
        if tag not in ("pic", "grpSp") or is_nested(el):
            continue
        if tag == "grpSp" and el.find(".//" + q("p:pic")) is None:
            continue
        if el.find(".//" + q("a:blip")) is None:
            continue  # an empty picture placeholder: the template's «insert your photo»
        b = element_bbox(el)
        if not b or b[2] * b[3] < 0.03 * W * H:
            continue
        if "".join(t.text or "" for t in el.iter(q("a:t"))).strip():
            continue
        if title_box is not None:
            bx = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
            tb = Bbox(x=title_box[0], y=title_box[1], w=title_box[2], h=title_box[3])
            if bx.intersection(tb) > 0.1 * max(tb.area, 1):
                continue
        keep.add(sid)
    return keep


def _slide_on_canvas(builder: DeckBuilder, canvas: Pattern, W: int, H: int, keep_extra: Optional[set[str]] = None, keep_art: bool = False):
    """Clone the canvas sample and strip it to background + chrome + the title slot (+ the art to keep; `keep_art`: a
    cover keeps its own photos and illustrations clear of the title). Returns (slide, title shape, heading ground)."""
    from verstka.analysis.shapes import looks_like_sample_value

    slide = builder.clone_slide(canvas.source_slide)
    els = slide_shape_elements(slide)
    title = _top_title(canvas)
    # chrome that is the template's sample copy («DEMO SLIDE», «Title Text Demo») or a sample value of its header or
    # footer («JOHN DOE», «NEW YORK», «2023», «SLIDESCARNIVAL.COM») is not kept on a composed slide
    chrome = {sid for sid in canvas.chrome_shape_ids if not (sid in els and looks_like_sample_value("".join(t.text or "" for t in els[sid].iter(q("a:t")))))}
    keep = chrome | ({title.shape_id} if title else set()) | set(keep_extra or ())
    # the label the heading is printed on (a pill) is part of the heading's style: its colour was chosen for it
    tb = element_bbox(els[title.shape_id]) if title and title.shape_id in els else None
    if keep_art:
        keep |= _cover_art_ids(els, tb, W, H)
    ground: Optional[tuple[Bbox, Optional[str]]] = None  # a half-slide panel the heading stands on
    if tb:
        tol = int(0.01 * W)
        for sid, el in els.items():
            b = element_bbox(el)
            if sid in keep or not b or is_nested(el) or etree.QName(el).localname != "sp":
                continue  # the heading's ground may be a small pill or a half-slide panel — it stays either way
            if "".join(t.text or "" for t in el.iter(q("a:t"))).strip() or not _paints(el):
                continue
            if b[0] - tol <= tb[0] <= b[0] + b[2] - int(0.04 * W) and b[1] - tol <= tb[1] and b[1] + b[3] + tol >= tb[1] + tb[3]:
                keep.add(sid)
                if b[2] * b[3] > 0.25 * W * H and (ground is None or b[2] * b[3] > ground[0].area):
                    ground = (Bbox(x=b[0], y=b[1], w=b[2], h=b[3]), _fill_hex(el))
    for sid, el in els.items():
        if sid in keep or is_nested(el) or el.getparent() is None:
            continue
        b = element_bbox(el)
        full_bleed = b is not None and b[2] * b[3] >= 0.9 * W * H and not "".join(t.text or "" for t in el.iter(q("a:t"))).strip()
        if full_bleed:
            continue  # a background drawn as a slide-sized rectangle or picture
        remove_element(el)
    renumber_page_chrome(slide_shape_elements(slide), canvas.chrome_shape_ids, canvas.source_slide, len(builder.created))
    title_shape = next((sh for sh in slide.shapes if title is not None and str(sh.shape_id) == title.shape_id), None)
    if title_shape is not None and is_nested(title_shape._element):
        title_shape = None
    return slide, title_shape, ground


def _adopt_canvas_colors(pal: "_Palette", canvas: Pattern, manifest: TemplateManifest) -> None:
    """On a cloned sample the ground is the sample's own (often a picture): its heading colour is known to read on it,
    so running text takes that colour too; a solid ground of the sample replaces the family's generic one."""
    bg = next((b.hex for b in manifest.tokens.backgrounds if canvas.source_slide in b.slides and b.hex), None)
    if bg:
        pal.bg = bg
    title = _top_title(canvas)
    color = title.style.color_hex if title is not None else None
    if color and contrast_ratio(color, pal.bg) >= 3.0 or (color and not bg):
        pal.text = pal.text2 = color


def _title_style(manifest: TemplateManifest, family: Family, pal: _Palette, content_only: bool = True) -> tuple[float, str, bool, Optional[str]]:
    """(size, colour, bold, font) of the template's headings on this ground. The size of a content heading is the
    median of the content samples' titles (a cover's display title or a one-off size does not set the deck's
    headings); covers and dividers (`content_only=False`) keep the most frequent title size of every sample."""
    slots = [s for p in manifest.patterns if p.family == family for s in p.slots if s.role == SlotRole.title and s.style.size_pt]
    size = manifest.tokens.typography.size_for("h1", 28.0)
    color = pal.text
    bold = False
    font = manifest.tokens.typography.primary_family
    if slots:
        sizes = Counter(round(s.style.size_pt or size) for s in slots)
        size = float(sizes.most_common(1)[0][0])
        if content_only:
            content = sorted(round(s.style.size_pt) for p in manifest.patterns if p.family == family and p.kind not in _NOT_CONTENT_CANVAS for s in p.slots if s.role == SlotRole.title and s.style.size_pt)
            if content:
                size = float(content[len(content) // 2])  # the median (of two middle sizes, the larger)
        colors = Counter(s.style.color_hex for s in slots if s.style.color_hex)
        if colors:
            cand = colors.most_common(1)[0][0]
            if contrast_ratio(cand, pal.bg) >= 3.0:
                color = cand
        bold = sum(1 for s in slots if s.style.bold) > len(slots) / 2
        fonts = Counter(s.style.font_family for s in slots if s.style.font_family)
        if fonts:
            font = fonts.most_common(1)[0][0]
    return size, color, bold, font


def _clamp_box(slide: Slide, box: Bbox, warnings: Optional[list[str]] = None, name: str = "text") -> Bbox:
    """A box cut to the slide (a last guard: a coordinate past the slide is a defect), reported when it happens."""
    from verstka.rendering.compose import _slide_size

    W, H = _slide_size(slide)
    if not W or not H:
        return box
    x, y = min(max(int(box.x), 0), W), min(max(int(box.y), 0), H)
    x2, y2 = min(max(int(box.x2), x), W), min(max(int(box.y2), y), H)
    if (x, y, x2 - x, y2 - y) == (int(box.x), int(box.y), int(box.w), int(box.h)):
        return box
    if warnings is not None:
        warnings.append(f"{name} clamped to the slide (y {box.y / H:.2f}, h {box.h / H:.2f})")
    return Bbox(x=x, y=y, w=x2 - x, h=y2 - y)


def _textbox(slide: Slide, box: Bbox, paragraphs: list[ParagraphSpec], *, size: float, color: str, font: Optional[str], bold: bool = False, align: str = "l", anchor: str = "t", scale: Optional[list[float]] = None, line_spacing: float = 1.2, fit: bool = True, warnings: Optional[list[str]] = None):
    box = _clamp_box(slide, box, warnings)
    tb = slide.shapes.add_textbox(Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.margin_top = tf.margin_bottom = Emu(22860)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    # a size asked as a multiple of a role («display × 0.8») lands on the next size the template really uses
    on_scale = [s for s in (scale or []) if s <= size + 0.05]
    if on_scale:
        size = max(on_scale)
    target = size
    if fit:
        res = fit_size([p.text for p in paragraphs], box, font, size, bold, scale or [], insets_emu=(45720, 22860, 45720, 22860), line_spacing=line_spacing, min_ratio=0.55)
        target = res.size_pt
    p0 = tf.paragraphs[0]
    p0.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    r = p0.add_run()
    r.text = ""
    r.font.size = Pt(target)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    if font:
        r.font.name = font
    specs = [ParagraphSpec(p.text, bullet=p.bullet, level=p.level, bold=p.bold if p.bold is not None else None, size_pt=p.size_pt, color_hex=p.color_hex) for p in paragraphs]
    fill_text(tb._element, specs, size_pt=target)
    from verstka.rendering.tables import neutral_run

    for rPr in tb._element.iter(q("a:rPr")):
        neutral_run(rPr)  # an engine textbox does not inherit the master's capitals or tracking
    if any(p.bullet for p in paragraphs):
        for pPr in tb._element.iter(q("a:pPr")):
            if pPr.find(q("a:buChar")) is not None:
                pPr.set("marL", "285750")
                pPr.set("indent", "-285750")
    for p in tb.text_frame.paragraphs:
        p.alignment = p0.alignment
        p.space_after = Pt(max(target * 0.35, 3))
    return tb


def _rect(slide: Slide, box: Bbox, fill: Optional[str], line: Optional[str], radius_emu: int):
    adj = card_adj(radius_emu, box.w, box.h)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if adj > 0 else MSO_SHAPE.RECTANGLE, Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    if adj > 0:
        shape.adjustments[0] = adj
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def _circle(slide: Slide, box: Bbox, fill: str, text: str, text_color: str, size: float, font: Optional[str]):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(text_color)
    if font:
        r.font.name = font
    return shape


def _items(oslide: OutlineSlide) -> list[SlideItem]:
    c = oslide.content
    if c.items:
        return list(c.items)
    if c.columns:
        return list(c.columns)
    if c.numbers:
        return [SlideItem(title=n.value, text=n.label, number=n.value) for n in c.numbers]
    return [SlideItem(title=b) for b in c.bullets]


_COVER_COMPS = ("title", "section", "thanks")
# compositions that need the whole width of the slide (a half-slide panel of the canvas is too narrow for them)
_WIDE_COMPS = ("cards", "process", "agenda", "comparison", "two_column", "table", "chart_text", "stat_row", "chart_pair", "formula")


def _needs_width(oslide: OutlineSlide, comp: str) -> bool:
    c = oslide.content
    n = len(c.items) or len(c.columns) or len(c.numbers) or len(c.bullets)
    if comp in ("table", "chart_text", "chart_pair", "formula") or c.table is not None or c.chart is not None or c.chart2 is not None or (c.formula or "").strip():
        return True
    if comp in _WIDE_COMPS and n >= 3:
        return True
    if comp == "bullets" and (len(c.bullets) >= 5 or sum(len(b) for b in c.bullets) > 320):
        return True
    return False


def _canvas_ground(builder: DeckBuilder, canvas: Pattern) -> Optional[Bbox]:
    """The half-slide panel a canvas's heading stands on (LCT: a white half next to a photo), read from the sample."""
    W, H = builder.slide_w, builder.slide_h
    try:
        els = slide_shape_elements(builder.source_slide(canvas.source_slide))
    except IndexError:
        return None
    title = _top_title(canvas)
    if title is None or title.shape_id not in els:
        return None
    tb = element_bbox(els[title.shape_id])
    if not tb:
        return None
    tol = int(0.01 * W)
    best: Optional[Bbox] = None
    for sid, el in els.items():
        b = element_bbox(el)
        if sid == title.shape_id or not b or is_nested(el) or etree.QName(el).localname != "sp":
            continue
        if "".join(t.text or "" for t in el.iter(q("a:t"))).strip() or not _paints(el):
            continue
        if b[0] - tol <= tb[0] <= b[0] + b[2] - int(0.04 * W) and b[1] - tol <= tb[1] and b[1] + b[3] + tol >= tb[1] + tb[3] and b[2] * b[3] > 0.25 * W * H:
            bb = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
            if best is None or bb.area > best.area:
                best = bb
    return best


def _layout_pictures(builder: DeckBuilder, pattern: Pattern) -> float:
    """Share of the slide covered by what the sample's layout draws besides its background — a photo collage or a
    dark code panel of the layout comes with every slide cloned from that sample."""
    W, H = builder.slide_w, builder.slide_h
    try:
        layout = builder.source_slide(pattern.source_slide).slide_layout
    except IndexError:
        return 0.0
    total = 0
    for el in layout._element.cSld.find(q("p:spTree")):
        tag = etree.QName(el).localname
        if tag not in ("sp", "pic", "grpSp", "graphicFrame") or el.find(".//" + q("p:ph")) is not None:
            continue  # placeholders of the layout are not drawn on a slide that does not use them
        if tag == "sp" and not _paints(el) and el.find(".//" + q("a:blip")) is None:
            continue
        b = element_bbox(el)
        if b and 0.03 * W * H <= b[2] * b[3] < 0.9 * W * H:
            total += b[2] * b[3]
    return total / (W * H)


def _has_pill(builder: DeckBuilder, canvas: Pattern) -> bool:
    t = _top_title(canvas)
    try:
        src = builder.source_slide(canvas.source_slide)
    except IndexError:
        return False
    el = slide_shape_elements(src).get(t.shape_id) if t else None
    return el is not None and _backing_of(src, el, builder.slide_w, builder.slide_h) is not None


def _bottom_chrome(builder: DeckBuilder, pattern: Pattern) -> int:
    """Brand marks a canvas carries at the foot of the slide (logos, footers) — its own chrome and its layout's."""
    cache = builder.__dict__.setdefault("_bottom_chrome", {})
    if pattern.id in cache:
        return cache[pattern.id]
    W, H = builder.slide_w, builder.slide_h
    n = 0
    try:
        src = builder.source_slide(pattern.source_slide)
        els = slide_shape_elements(src)
        items = [els[i] for i in pattern.chrome_shape_ids if i in els]
        items += [el for el in src.slide_layout._element.cSld.find(q("p:spTree")) if el.find(".//" + q("p:ph")) is None]
        for el in items:
            b = element_bbox(el)
            if b and b[1] >= 0.85 * H and b[2] * b[3] < 0.2 * W * H:
                n += 1
    except Exception:  # noqa: BLE001
        pass
    cache[pattern.id] = n
    return n


def _chrome_floor(builder: DeckBuilder, manifest: TemplateManifest, family: Family) -> int:
    """The foot chrome most content samples of the family have: a canvas without it would drop the brand marks."""
    counts = sorted(_bottom_chrome(builder, p) for p in _content_patterns(manifest) if p.family == family)
    return counts[len(counts) // 2] if counts else 0


GROUND_STRUCTURE_MAX = 7.0  # a ground a composed slide may stand on: a gradient, a paper texture, a starry sky stay under it


def _picture_structure(part) -> float:
    """Object-scale structure of a picture as the slide shows it (over white where it is transparent): the mean
    luminance step between neighbouring cells of a 24×14 thumbnail. Calm grounds — gradients, paper, a night sky —
    stay under GROUND_STRUCTURE_MAX; a photograph of things (a vase on a table, a beach, a cup) is above it."""
    import io

    import numpy as np
    from PIL import Image

    im = Image.open(io.BytesIO(part.blob))
    im = im.convert("RGBA")
    a = np.asarray(im.resize((24, 14)), dtype=np.float32)
    alpha = a[..., 3] / 255.0
    lum = (0.299 * a[..., 0] + 0.587 * a[..., 1] + 0.114 * a[..., 2]) * alpha + 255.0 * (1.0 - alpha)
    return float((np.abs(np.diff(lum, axis=1)).mean() + np.abs(np.diff(lum, axis=0)).mean()) / 2.0)


def _ground_structure(builder: DeckBuilder, canvas: Pattern) -> float:
    """The structure of the slide-sized pictures the sample itself lays behind its content (the largest step of
    them; 0 without) — a photo of the sample's own story that a composed slide on it would keep. The ground the
    layout or the master draws under every slide (Office «Celestial»'s night sky) is the template's designed ground:
    every canvas shares it, so it never tells one canvas from another."""
    cache = builder.__dict__.setdefault("_ground_structure", {})
    if canvas.id in cache:
        return cache[canvas.id]
    W, H = builder.slide_w, builder.slide_h
    worst = 0.0
    try:
        src = builder.source_slide(canvas.source_slide)
        holders = [(src, el) for el in slide_shape_elements(src).values()]
        for holder, el in holders:
            if not isinstance(el.tag, str) or etree.QName(el).localname not in ("pic", "sp", "grpSp"):
                continue
            if el.find(".//" + q("p:ph")) is not None and holder is not src:
                continue
            b = element_bbox(el)
            if not b or b[2] * b[3] < 0.85 * W * H:
                continue
            for part in _pic_parts(holder, el):
                try:
                    worst = max(worst, _picture_structure(part))
                except Exception:  # noqa: BLE001 - an unreadable picture (EMF, SVG) is not judged
                    continue
    except Exception:  # noqa: BLE001
        worst = 0.0
    cache[canvas.id] = worst
    return worst


def _title_turned(builder: DeckBuilder, canvas: Pattern) -> bool:
    """The canvas's heading is set turned or vertically («Agenda» running up the side of a photo): a composed
    heading written into it would run up the slide too."""
    t = _top_title(canvas)
    if t is None:
        return False
    try:
        el = slide_shape_elements(builder.source_slide(canvas.source_slide)).get(t.shape_id)
    except Exception:  # noqa: BLE001
        return False
    if el is None:
        return False
    xfrm = el.find(".//" + q("a:xfrm"))
    rot = int(xfrm.get("rot") or 0) if xfrm is not None else 0
    turned = abs(((rot / 60000.0) + 180.0) % 360.0 - 180.0) > 3.0
    body = el.find(".//" + q("a:bodyPr"))
    vert = body is not None and (body.get("vert") or "horz") != "horz"
    return turned or vert


def _pick_canvas(builder: DeckBuilder, manifest: TemplateManifest, family: Family, comp: str, wide: bool, need_pill: bool = False) -> Optional[Pattern]:
    """The cleanest canvas for a composition: no pictures of its layout in the content area, and — for a composition
    that needs the whole width — no half-slide panel under the heading."""
    tried: set[str] = set()
    m = manifest
    first = canvas = _canvas_for(manifest, family, comp)
    fallback = None
    seen: list[Pattern] = []
    while canvas is not None:
        seen.append(canvas)
        floor = min(_chrome_floor(builder, manifest, family), 2)
        ok = _layout_pictures(builder, canvas) < 0.12 and (not need_pill or _has_pill(builder, canvas)) and _bg_art(builder, canvas.source_slide) is None and _bottom_chrome(builder, canvas) >= floor and not _canvas_photo(builder, canvas)
        # a heading set turned or vertically, or a photograph of things behind the content: never a composed slide's canvas
        ok = ok and not _title_turned(builder, canvas) and _ground_structure(builder, canvas) < GROUND_STRUCTURE_MAX
        if ok and wide:
            g = _canvas_ground(builder, canvas)
            if g is not None and g.w < 0.7 * builder.slide_w:
                fallback = fallback or canvas
                ok = False
        if ok:
            return canvas
        tried.add(canvas.id)
        m = m.model_copy(update={"patterns": [p for p in m.patterns if p.id not in tried]})
        canvas = _canvas_for(m, family, comp)
    if fallback is not None or not seen:
        return fallback or first
    # no candidate passes: the one whose content band is freest of the template's art (not the most cluttered one),
    # on the heading grid, without a photo behind the content; a wide composition needs 60 % of the width free —
    # or a clean layout of the package when every sample carries art over the content band
    best = min(seen, key=lambda p: _canvas_rank(builder, manifest, p, wide))
    rank = _canvas_rank(builder, manifest, best, wide)
    lay = _clean_layout(builder, manifest)
    if lay is not None and (rank[0] or rank[1] or lay[1] >= -rank[3] + 0.2) and lay[1] >= 0.7:
        builder.__dict__.setdefault("_prefer_layout", {})[family] = lay[0]
        return None
    return best


def _clean_layout(builder: DeckBuilder, manifest: TemplateManifest):
    """(layout, free share of its content band): the layout of the package with a title placeholder whose content band
    is freest of drawn art (its master's art included) — for a template whose every sample carries art over the
    content. None when no layout is measurable."""
    cache = builder.__dict__.setdefault("_clean_layout", {})
    if "best" in cache:
        return cache["best"]
    best = None
    try:
        from verstka.rendering.layers import drawn_layers, largest_free
    except ImportError:
        cache["best"] = None
        return None
    W, H = builder.slide_w, builder.slide_h
    safe = manifest.tokens.spacing.safe_area
    for layout in builder.layouts():
        title = next((ph for ph in layout.placeholders if str(ph.placeholder_format.type).split(".")[-1].split(" ")[0] in ("TITLE",)), None)
        if title is None or title.top is None or title.height is None:
            continue
        if int(title.top) > 0.2 * H:
            continue  # a cover's or a divider's layout (its title stands mid-slide) is no content page
        try:
            y0 = max(int(safe.y * H), int(title.top) + int(title.height) + int(0.04 * H))
            band = Bbox(x=int(safe.x * W), y=y0, w=int(safe.w * W), h=max(int(safe.y2 * H) - y0, 1))
            if band.h < 0.3 * H:
                continue
            obstacles = _layout_art(drawn_layers(layout), band, W, H)
            if obstacles is None:
                continue
            pad = int(0.015 * W)
            free = largest_free(band, [Bbox(x=b.x - pad, y=b.y - pad, w=b.w + 2 * pad, h=b.h + 2 * pad) for b in obstacles])
            share = free.area / max(band.area, 1)
        except Exception:  # noqa: BLE001
            continue
        if best is None or share > best[1] + 1e-6:
            best = (layout, share)
    cache["best"] = best
    return best


def _layout_art(layers, band: Bbox, W: int, H: int) -> Optional[list[Bbox]]:
    """What a bare layout draws over a content band, strictly: every painted layer of the layout and its master that
    is not a placeholder, a rule or a speck — only a calm full-bleed background and a plain panel holding the whole band
    are grounds (a mosaic, a group or a picture holding the title is art, not a ground: a composed slide would be laid
    over it). None when a photo covers 30 % of the slide."""
    out: list[Bbox] = []
    for lay in layers:
        if lay.placeholder or not lay.paints or lay.kind in ("graphicFrame",):
            continue
        if lay.busy and lay.cover >= 0.3:
            return None
        if lay.box.h <= 0.02 * H or lay.cover < 0.003:
            continue
        if lay.cover >= 0.85 and not lay.busy:
            continue
        if lay.kind == "sp" and not lay.picture and not lay.custom_geom and lay.box.intersection(band) >= 0.95 * band.area:
            continue  # a plain panel under the whole content band: its ground
        out.append(lay.box)
    return out


def _canvas_rank(builder: DeckBuilder, manifest: TemplateManifest, p: Pattern, wide: bool) -> tuple:
    cache = builder.__dict__.setdefault("_canvas_rank", {})
    key = (p.id, wide)
    if key not in cache:
        share, busy, free_w = _canvas_free(builder, manifest, p)
        t = _top_title(p)
        on_grid = t is not None and _on_grid(manifest, t)
        clutter = len(p.slots) + sum(len(g.member_shape_ids) for g in p.repeat_groups) + len(p.decor_assets)
        rough = _title_turned(builder, p) or _ground_structure(builder, p) >= GROUND_STRUCTURE_MAX
        cache[key] = (rough, busy or _canvas_photo(builder, p), bool(wide and free_w < 0.6), -round(share, 1), not on_grid, clutter, -p.quality, p.source_slide)
    return cache[key]


def _canvas_free(builder: DeckBuilder, manifest: TemplateManifest, p: Pattern) -> tuple[float, bool, float]:
    """(free share of the content band, a photo ≥ 30 % of the slide behind the content, free width share) of a
    canvas, measured on its sample with only what a cloned canvas keeps (chrome, title, full-bleed grounds)."""
    W, H = builder.slide_w, builder.slide_h
    fs = getattr(p, "free_share", None)
    try:
        from verstka.rendering.layers import art_layers, free_rect

        src = builder.source_slide(p.source_slide)
        els = slide_shape_elements(src)
        t = _top_title(p)
        keep = set(p.chrome_shape_ids) | ({t.shape_id} if t else set())
        skip = [sid for sid in els if sid not in keep]
        safe = manifest.tokens.spacing.safe_area
        y0 = max(safe.y, (t.bbox.y2 if t else safe.y) + 0.04)
        band = Bbox(x=int(safe.x * W), y=int(y0 * H), w=int(safe.w * W), h=max(int((safe.y2 - y0) * H), 1))
        free = free_rect(src, band, skip=skip)
        busy = any(l.busy and l.cover >= 0.3 for l in art_layers(src, skip=skip))
        share = free.area / max(band.area, 1) if fs is None else float(fs)
        return share, busy, free.w / W
    except Exception:  # noqa: BLE001
        return (float(fs) if fs is not None else 0.5), False, 1.0


def _scheme(builder: DeckBuilder) -> dict[str, str]:
    """Theme colours of the first master (bg1/tx1/accent1… → hex) for sample shapes painted with scheme colours."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    out: dict[str, str] = {}
    try:
        master = builder.prs.slide_masters[0]
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        cs = theme.find(".//" + q("a:clrScheme"))
        for entry in cs if cs is not None else []:
            if len(entry) == 0:
                continue
            c = entry[0]
            tag = etree.QName(c).localname
            name = etree.QName(entry).localname
            if tag == "srgbClr":
                out[name] = (c.get("val") or "000000").upper()
            elif tag == "sysClr":
                out[name] = (c.get("lastClr") or ("000000" if c.get("val") == "windowText" else "FFFFFF")).upper()
        cmap = master._element.find(q("p:clrMap"))
        for k, v in (cmap.attrib.items() if cmap is not None else [("bg1", "lt1"), ("tx1", "dk1"), ("bg2", "lt2"), ("tx2", "dk2")]):
            if v in out:
                out[k] = out[v]
    except Exception:  # noqa: BLE001
        pass
    return out


def _resolved_fill(el, scheme: dict[str, str], line: bool = False) -> Optional[str]:
    from verstka.analysis.theme import _apply_modifiers

    spPr = el.find(q("p:spPr"))
    if spPr is None:
        return None
    holder = spPr.find(q("a:ln")) if line else spPr
    if holder is None or holder.find(q("a:noFill")) is not None:
        return None
    for tag in ("a:solidFill", "a:gradFill"):
        f = holder.find(q(tag))
        if f is None:
            continue
        clr = next((c for c in f.iter() if etree.QName(c).localname in ("srgbClr", "schemeClr")), None)
        if clr is None:
            return None
        base = (clr.get("val") or "").upper() if etree.QName(clr).localname == "srgbClr" else scheme.get(clr.get("val") or "")
        return _apply_modifiers(base, clr) if base and len(base) == 6 else None
    if not line:
        ref = el.find(q("p:style") + "/" + q("a:fillRef"))
        if ref is not None and (ref.get("idx") or "0") != "0":
            clr = next((c for c in ref if etree.QName(c).localname in ("srgbClr", "schemeClr")), None)
            if clr is not None:
                base = (clr.get("val") or "").upper() if etree.QName(clr).localname == "srgbClr" else scheme.get(clr.get("val") or "")
                return _apply_modifiers(base, clr) if base and len(base) == 6 else None
    return None


_CARD_GEOMS = ("rect", "roundRect", "round2SameRect", "round1Rect", "snip1Rect", "snip2SameRect")


def _card_proto(builder: DeckBuilder, manifest: TemplateManifest, family: Family, ground: str):
    """The template's own card: the painted rectangle that holds the text of a repeat-group cell on a sample slide of
    this family. The most frequent style wins (by fill/line/geometry), so one odd sample does not define the look.
    Returned as a detached copy with its colours resolved: (element, fill, line) or None. Found once per deck."""
    cache = builder.__dict__.setdefault("_card_protos", {})
    key = (family, ground)
    if key not in cache:
        cache[key] = _find_card_proto(builder, manifest, family, ground)
    found = cache[key]
    return (copy.deepcopy(found[0]), found[1], found[2]) if found else None


def _find_card_proto(builder: DeckBuilder, manifest: TemplateManifest, family: Family, ground: str):
    W, H = builder.slide_w, builder.slide_h
    scheme = _scheme(builder)
    votes: dict[tuple, list] = {}
    for p in manifest.patterns:
        if p.reference or p.family != family or not p.repeat_groups or p.quality < 0.4:
            continue
        try:
            els = slide_shape_elements(builder.source_slide(p.source_slide))
        except IndexError:
            continue
        slot_of = {s.shape_id: s for s in p.slots}
        for g in p.repeat_groups:
            for cell in g.member_shape_ids[:1]:
                members = [(sid, els[sid]) for sid in cell if sid in els]
                text_boxes = [element_bbox(e) for sid, e in members if sid in slot_of]
                for sid, el in members:
                    if etree.QName(el).localname != "sp" or is_nested(el):
                        continue
                    geom = el.find(q("p:spPr") + "/" + q("a:prstGeom"))
                    if geom is None or geom.get("prst") not in _CARD_GEOMS:
                        continue
                    b = element_bbox(el)
                    if not b or not (0.012 <= b[2] * b[3] / (W * H) <= 0.4):
                        continue
                    fill = _resolved_fill(el, scheme)
                    line = _resolved_fill(el, scheme, line=True)
                    if not fill and not line:
                        continue
                    if not fill and not line:
                        continue
                    holds = sid in slot_of or any(tb and b[0] <= tb[0] + tb[2] / 2 <= b[0] + b[2] and b[1] <= tb[1] + tb[3] / 2 <= b[1] + b[3] for tb in text_boxes if tb)
                    if not holds:
                        continue
                    key = (fill, line, geom.get("prst"))
                    votes.setdefault(key, []).append((b[2] * b[3], el))
    if not votes:
        return None
    # the plain card wins over a highlighted one (an accent outline marks one card of a row, not every card)
    key = max(votes, key=lambda k: (len(votes[k]), k[1] is None, max(a for a, _ in votes[k])))
    el = copy.deepcopy(max(votes[key], key=lambda t: t[0])[1])
    fill, line = key[0], key[1]
    spPr = el.find(q("p:spPr"))
    if fill and not line and contrast_ratio(fill, ground) < 1.12 and spPr is not None and spPr.find(q("a:effectLst")) is None:
        # a white card on an almost white ground (VK Tech) stands apart by a soft shadow, as in the template
        eff = etree.SubElement(spPr, q("a:effectLst"))
        sh = etree.SubElement(eff, q("a:outerShdw"))
        sh.set("blurRad", str(int(0.03 * H)))
        sh.set("dist", str(int(0.004 * H)))
        sh.set("dir", "5400000")
        sh.set("algn", "t")
        sh.set("rotWithShape", "0")
        clr = etree.SubElement(sh, q("a:srgbClr"))
        clr.set("val", "0B1F44")
        etree.SubElement(clr, q("a:alpha")).set("val", "9000")
    # the copy is painted with resolved colours: a scheme colour of another master must not change on our slide
    return el, fill, line


def _title_anchor(title_shape) -> str:
    """Vertical anchor of a title shape: its own bodyPr, else the placeholder chain (layout, master)."""
    el = title_shape._element
    node = el
    for _ in range(3):
        bp = node.find(".//" + q("a:bodyPr")) if node is not None else None
        if bp is not None and bp.get("anchor"):
            return bp.get("anchor")
        try:
            base = title_shape._base_placeholder if node is el else None
        except Exception:  # noqa: BLE001
            base = None
        if base is None:
            break
        title_shape = base
        node = base._element
    return "t"


def _backing_of(slide: Slide, title_el, W: int, H: int):
    """A small label (pill) painted under the start of the heading box: (element, bbox) or None."""
    tb = element_bbox(title_el)
    if not tb:
        return None
    tol = int(0.01 * W)
    best = None
    for el in slide._element.cSld.find(q("p:spTree")):
        if el is title_el or etree.QName(el).localname != "sp":
            continue
        if "".join(t.text or "" for t in el.iter(q("a:t"))).strip() or not _paints(el):
            continue
        b = element_bbox(el)
        if not b or b[2] * b[3] > 0.1 * W * H or b[3] > 0.16 * H:
            continue
        if b[0] - tol <= tb[0] <= b[0] + b[2] - int(0.04 * W) and b[1] - tol <= tb[1] + tb[3] // 2 <= b[1] + b[3] + tol:
            if best is None or b[2] * b[3] < best[1].area:
                best = (el, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
    return best


def render_synth(builder: DeckBuilder, plan_slide: LayoutSlide, oslide: OutlineSlide, manifest: TemplateManifest, ws: TemplateWorkspace, outline: DeckOutline) -> tuple[Slide, list[str]]:
    comp = plan_slide.composition or "bullets"
    if comp in _COVER_COMPS:
        return _render_cover(builder, plan_slide, oslide, manifest, ws, outline)
    return _render_content(builder, comp, oslide, manifest, ws, outline)


@dataclass
class Art:
    """A sample whose decorative picture stands on one side: the picture stays, the content takes the other side."""

    pattern: Pattern
    keep_ids: set[str]
    box: Bbox
    side: str  # "left" | "right"
    photo: bool = False  # the art is (or holds) a photograph — the template's sample story, not brand decoration


def _is_photo(builder: DeckBuilder, part) -> bool:
    """A photograph (people, a place, a product shot, a screenshot full of photos) rather than drawn art: an opaque
    raster with many colours and fine mid-contrast texture. A 3D render or an illustration is cut out (alpha) or
    smooth; a flat pattern has few colours. Measured on the picture, 128 px thumbnail."""
    cache = builder.__dict__.setdefault("_is_photo", {})
    key = str(getattr(part, "partname", id(part)))
    if key in cache:
        return cache[key]
    out = False
    try:
        import io

        import numpy as np
        from PIL import Image

        im = Image.open(io.BytesIO(part.blob)).convert("RGBA")
        im.thumbnail((128, 128))
        a = np.asarray(im)
        op = a[..., 3] > 200
        if op.mean() >= 0.5 and op.sum() >= 400:
            rgb = a[..., :3][op].astype(np.int32) >> 3
            colours = len(np.unique(rgb[:, 0] * 1024 + rgb[:, 1] * 32 + rgb[:, 2]))
            g = np.asarray(im.convert("L"), dtype=np.float32)
            d = np.abs(np.diff(g, axis=1))
            both = op[:, 1:] & op[:, :-1]
            texture = float(((d > 4) & (d < 40))[both].mean()) if both.any() else 0.0
            out = colours >= 500 and texture >= 0.2
    except Exception:  # noqa: BLE001
        out = False
    cache[key] = out
    return out


def _pic_parts(holder, el) -> list:
    """The raster parts an element (a picture, a group) draws."""
    out = []
    for bl in el.iter(q("a:blip")):
        rid = bl.get(q("r:embed"))
        if not rid:
            continue
        try:
            out.append(holder.part.related_part(rid))
        except KeyError:
            continue
    return out


def _canvas_photo(builder: DeckBuilder, canvas: Pattern) -> bool:
    """A composed slide cloned from this canvas would carry a photograph of the template's own story: its layout or
    master draws one (≥ 8 % of the slide), or the sample keeps one as chrome or as a full-bleed ground."""
    cache = builder.__dict__.setdefault("_canvas_photo", {})
    if canvas.id in cache:
        return cache[canvas.id]
    W, H = builder.slide_w, builder.slide_h
    out = False
    try:
        src = builder.source_slide(canvas.source_slide)
        holders = [(src.slide_layout, el) for el in src.slide_layout._element.cSld.find(q("p:spTree"))]
        holders += [(src.slide_layout.slide_master, el) for el in src.slide_layout.slide_master._element.cSld.find(q("p:spTree"))]
        els = slide_shape_elements(src)
        for sid, el in els.items():
            b = element_bbox(el)
            if sid in set(canvas.chrome_shape_ids) or (b is not None and b[2] * b[3] >= 0.9 * W * H):
                holders.append((src, el))
        for holder, el in holders:
            if el.find(".//" + q("p:ph")) is not None and holder is not src:
                continue  # a layout's empty picture placeholder draws nothing on a slide that does not fill it
            b = element_bbox(el)
            if not b or b[2] * b[3] < 0.08 * W * H:
                continue
            if any(_is_photo(builder, part) for part in _pic_parts(holder, el)):
                out = True
                break
    except Exception:  # noqa: BLE001
        out = False
    cache[canvas.id] = out
    return out


_ART_KINDS_EXCLUDED = (PatternKind.title, PatternKind.section, PatternKind.thanks, PatternKind.mockup, PatternKind.image_text, PatternKind.code, PatternKind.team, PatternKind.quote)


def _bg_art(builder: DeckBuilder, source_slide: int) -> Optional[tuple[str, float]]:
    """Art baked into the slide's background picture (the LCT city render): (side, inner edge as a share of the width)
    when one half of the picture is busy and the other calm. Measured on the picture itself — edges per column."""
    cache = builder.__dict__.setdefault("_bg_art", {})
    if source_slide in cache:
        return cache[source_slide]
    out = None
    try:
        import io

        import numpy as np
        from PIL import Image, ImageFilter

        s = builder.source_slide(source_slide)
        blip, holder = None, None
        for h in (s, s.slide_layout, s.slide_layout.slide_master):
            bg = h._element.cSld.find(q("p:bg"))
            blip = bg.find(".//" + q("a:blip")) if bg is not None else None
            if blip is not None:
                holder = h
                break
        if blip is not None:
            blob = holder.part.related_part(blip.get(q("r:embed"))).blob
            im = Image.open(io.BytesIO(blob)).convert("L").resize((320, 180))
            e = np.asarray(im.filter(ImageFilter.FIND_EDGES), dtype=float)[5:-5, 5:-5].mean(axis=0)
            bins = [float(e[j * 31 : (j + 1) * 31].mean()) for j in range(10)]
            left, right = sum(bins[:5]) / 5, sum(bins[5:]) / 5
            top = max(bins)
            if right >= 6 and right >= 4 * max(left, 0.5):
                first = next(j for j in range(10) if bins[j] >= 0.35 * top)
                out = ("right", first / 10)
            elif left >= 6 and left >= 4 * max(right, 0.5):
                last = max(j for j in range(10) if bins[j] >= 0.35 * top)
                out = ("left", (last + 1) / 10)
    except Exception:  # noqa: BLE001
        out = None
    cache[source_slide] = out
    return out


def _bg_photo(builder: DeckBuilder, source_slide: int) -> bool:
    """The background picture of a slide (its own, its layout's or its master's) is a photograph."""
    try:
        s = builder.source_slide(source_slide)
        for h in (s, s.slide_layout, s.slide_layout.slide_master):
            bg = h._element.cSld.find(q("p:bg"))
            blip = bg.find(".//" + q("a:blip")) if bg is not None else None
            if blip is not None:
                return _is_photo(builder, h.part.related_part(blip.get(q("r:embed"))))
    except Exception:  # noqa: BLE001
        return False
    return False


def _art_of(builder: DeckBuilder, p: Pattern, manifest: Optional[TemplateManifest] = None) -> Optional[Art]:
    """Brand art is repeated: a picture the template draws on several slides (or on a layout) is decoration; a picture
    drawn once — a screenshot in a phone, a chart snapshot — belongs to its sample's story."""
    cache = builder.__dict__.setdefault("_art_of", {})
    if (p.id, manifest is not None) not in cache:
        cache[(p.id, manifest is not None)] = _find_art(builder, p, manifest)
    return cache[(p.id, manifest is not None)]


def _find_art(builder: DeckBuilder, p: Pattern, manifest: Optional[TemplateManifest] = None) -> Optional[Art]:
    W, H = builder.slide_w, builder.slide_h
    if p.reference or p.kind in _ART_KINDS_EXCLUDED or p.quality < 0.5 or len(p.slots) > 16:
        return None
    if p.kind in (PatternKind.big_number, PatternKind.stat_row, PatternKind.chart, PatternKind.table):
        return None
    uses = {a.media_part.lstrip("/"): len(a.used_on_slides) for a in (manifest.assets if manifest else [])}
    t = _top_title(p)
    if t is None or t.bbox.y > 0.3:
        return None
    try:
        src = builder.source_slide(p.source_slide)
    except IndexError:
        return None
    if any(s.role == SlotRole.image and s.bbox.area >= 0.04 for s in p.slots):
        return None  # a picture slot is content (a screenshot to replace), not decoration
    tree = src._element.cSld.find(q("p:spTree"))
    keep: set[str] = set()
    union: Optional[Bbox] = None
    photo = False
    for el in tree:
        tag = etree.QName(el).localname
        if tag not in ("pic", "grpSp"):
            continue
        if tag == "grpSp" and (el.find(".//" + q("a:blip")) is None or "".join(x.text or "" for x in el.iter(q("a:t"))).strip()):
            continue
        b = element_bbox(el)
        if not b or not (0.05 * W * H <= b[2] * b[3] < 0.9 * W * H):
            continue
        if manifest is not None:
            blip = el.find(".//" + q("a:blip"))
            rid = blip.get(q("r:embed")) if blip is not None else None
            try:
                part = str(src.part.related_part(rid).partname).lstrip("/") if rid else ""
            except KeyError:
                part = ""
            if uses.get(part, 0) < 2:
                return None
        nv = el.find(".//" + q("p:cNvPr"))
        if nv is not None:
            keep.add(nv.get("id"))
        photo = photo or any(_is_photo(builder, part) for part in _pic_parts(src, el))
        bb = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
        union = bb if union is None else union.union(bb)
    lay: list[Bbox] = []
    lay_photo = False
    try:
        layout_part = str(src.slide_layout.part.partname).lstrip("/")
        usage = builder.__dict__.get("_layout_usage")
        if usage is None:
            usage = Counter(str(builder.source_slide(i).slide_layout.part.partname).lstrip("/") for i in range(1, builder.n_original + 1))
            builder.__dict__["_layout_usage"] = usage
        shared = usage.get(layout_part, 0) >= 2
        for el in src.slide_layout._element.cSld.find(q("p:spTree")):
            if el.find(".//" + q("p:ph")) is not None or el.find(".//" + q("a:blip")) is None:
                continue
            b = element_bbox(el)
            if b and 0.05 * W * H <= b[2] * b[3] < 0.9 * W * H:
                lay.append(Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
                lay_photo = lay_photo or any(_is_photo(builder, part) for part in _pic_parts(src.slide_layout, el))
        if lay and not shared:
            # a layout of its own: only a large, broad picture reads as art (a mood collage); a narrow one is a
            # device mockup waiting for its screenshot
            u = lay[0]
            for bb in lay[1:]:
                u = u.union(bb)
            if u.area < 0.3 * W * H or any(bb.w < 0.6 * bb.h for bb in lay):
                lay = []
        for bb in lay:
            union = bb if union is None else union.union(bb)
    except Exception:  # noqa: BLE001
        pass
    if union is None:
        bga = _bg_art(builder, p.source_slide)
        if bga is not None:
            side, edge = bga
            union = Bbox(x=int(edge * W), y=0, w=W - int(edge * W), h=H) if side == "right" else Bbox(x=0, y=0, w=int(edge * W), h=H)
            photo = photo or _bg_photo(builder, p.source_slide)
    if union is None or union.area < 0.15 * W * H:
        return None
    photo = photo or lay_photo  # the layout draws its pictures on every slide cloned from the sample
    pieces = [element_bbox(el) for el in tree if (nv := el.find(".//" + q("p:cNvPr"))) is not None and nv.get("id") in keep]
    biggest = max([b[2] * b[3] for b in pieces if b] + [bb.area for bb in lay] + ([union.area] if not keep and not lay else [0]))
    if biggest < 0.15 * W * H:
        return None  # thin strips and corner patterns are chrome, not art: they cannot carry a slide
    if union.x >= 0.38 * W and union.x2 >= 0.88 * W:
        side, free = "right", union.x
    elif union.x2 <= 0.62 * W and union.x <= 0.12 * W:
        side, free = "left", W - union.x2
    else:
        return None
    if free < 0.38 * W:
        return None
    return Art(pattern=p, keep_ids=keep, box=union, side=side, photo=photo)


def _choose_art(builder: DeckBuilder, manifest: TemplateManifest, ds: "DeckStyle", comp: str, oslide: OutlineSlide, outline: DeckOutline) -> Optional[Art]:
    """An art canvas for a slide that carries one message (a statement, a hero figure, a few theses) — in the visual
    variant, and never two in a row: at most one slide in three stands next to the template's art."""
    strategy = outline.strategy or "structured"
    c = oslide.content
    # the column beside the art also holds the lede, the conclusion strip and the footnote
    extra = sum(len(x or "") for x in (oslide.subtitle, oslide.takeaway, oslide.footnote))
    few = len(c.bullets) <= 4 and sum(len(b) for b in c.bullets) + extra <= 300 and not c.items and not c.table and not c.chart and not c.chart2 and not (c.formula or "").strip()
    eligible = comp in ("statement", "big_number", "quote") or (comp == "bullets" and few) or (comp == "stat_row" and len(c.numbers) == 1)
    if comp == "bullets" and not c.bullets and len(c.paragraphs) == 1:
        eligible = True  # a single paragraph is set as a statement
    if not eligible or strategy == "compact":
        return None
    if strategy == "structured":
        # the structured variant keeps its grid: only its closing message (the ask, the conclusion) stands next to art
        content = [s for s in outline.slides if s.kind not in _COVER_KINDS]
        last = bool(content) and content[-1].id == oslide.id
        if not (last and comp in ("bullets", "statement") and not c.bullets and len(c.paragraphs) == 1):
            return None
    history = builder.__dict__.setdefault("_art_slides", [])
    index = len(builder.created)
    if history and index - history[-1] < 4:
        return None
    hero = comp in ("statement", "big_number", "quote") or (comp == "bullets" and not c.bullets)
    from verstka.rendering.fonts import wrap_lines
    from verstka.ru import typeset

    W = builder.slide_w
    safe = manifest.tokens.spacing.safe_area
    head = typeset(oslide.headline)
    font = manifest.tokens.typography.primary_family

    def column_fits(a: "Art") -> bool:
        """The free column beside the art holds the headline in three lines at the deck's heading size — a long
        headline squeezed into five narrow lines next to a picture reads worse than the plain grid."""
        g = int(0.04 * W)
        col = (a.box.x - g - int(safe.x * W)) if a.side == "right" else (int(safe.x2 * W) - a.box.x2 - g)
        if col < 0.3 * W:
            return False
        return len(wrap_lines(head, font, ds.head_size, ds.head_bold, col / EMU_PER_PT - 14)) <= 3

    cands = []
    for p in manifest.patterns:
        if p.family != ds.family and not hero:
            continue  # a hero slide (one statement, one figure) may stand on the other ground as an accent
        a = _art_of(builder, p, manifest)
        if a is None or (ds.pill and not _has_pill(builder, p)) or p.id in builder.__dict__.get("_art_used", []):
            continue  # one appearance per deck: a repeated picture reads as a placeholder
        if a.photo and not c.image_hint:
            continue  # a photograph is the template's sample story (a smiling student on a war slide), not decoration
        if not column_fits(a):
            continue
        clutter = len(p.slots) + sum(len(g.member_shape_ids) for g in p.repeat_groups)
        cands.append((int(p.family != ds.family), -min(a.box.area / (builder.slide_w * builder.slide_h), 0.5), clutter, p.source_slide, a))
    if not cands:
        return None
    cands.sort(key=lambda t: t[:4])
    used = builder.__dict__.setdefault("_art_used", [])
    # the least used of the best art canvases: two statements do not repeat the same picture when another exists
    best = min(cands[:3], key=lambda t: used.count(t[4].pattern.id))[4]
    used.append(best.pattern.id)
    history.append(index)
    return best


def _render_content(builder: DeckBuilder, comp: str, oslide: OutlineSlide, manifest: TemplateManifest, ws: TemplateWorkspace, outline: DeckOutline) -> tuple[Slide, list[str]]:
    from verstka.rendering.compose import Composer, Kit

    warnings: list[str] = []
    ds = deck_style(builder, manifest, outline)
    family = ds.family
    if comp == "image_text" and not (oslide.content.image_hint and pick_asset(manifest, ws, oslide.content.image_hint, require_match=True)):
        comp = "bullets"  # no picture of the template says what the slide says: the text alone (an art canvas may frame it)
    wide = _needs_width(oslide, comp)
    art = _choose_art(builder, manifest, ds, comp, oslide, outline) if not wide else None
    canvas = art.pattern if art is not None else _pick_canvas(builder, manifest, family, comp, wide, need_pill=ds.pill)
    ground = None
    if canvas is not None:
        slide, title_ph, ground = _slide_on_canvas(builder, canvas, builder.slide_w, builder.slide_h, keep_extra=art.keep_ids if art else None)
        family = canvas.family
    else:
        layout, family = _layout_for(builder, manifest, family, oslide.kind)
        slide = _add_layout_slide(builder, layout)
        title_ph = None
        for shp in list(slide.shapes):
            if shp.is_placeholder and shp.placeholder_format.type is not None and str(shp.placeholder_format.type).split(".")[-1].split(" ")[0] in ("TITLE", "CENTER_TITLE"):
                title_ph = shp
            else:
                shp._element.getparent().remove(shp._element)
    if art is not None and title_ph is not None:
        _snap_header(builder, manifest, ds, slide, title_ph, art)
    pal = _Palette(manifest, family)
    if canvas is not None:
        _adopt_canvas_colors(pal, canvas, manifest)
    t = manifest.tokens
    typo = t.typography
    W, H = builder.slide_w, builder.slide_h
    safe = t.spacing.safe_area
    sx, sy, sw, sh = int(safe.x * W), int(safe.y * H), int(safe.w * W), int(safe.h * H)
    if ground is not None:
        gb, ghex = ground
        pad = int(0.03 * W)
        x1, x2 = max(sx, gb.x + pad), min(sx + sw, gb.x2 - pad)
        y2 = min(sy + sh, gb.y2 - pad)
        if x2 - x1 > 0.25 * W:
            sx, sw, sh = x1, x2 - x1, y2 - sy
        if ghex:
            pal.bg = ghex
            pal.text = pal.text2 = _readable(ghex, [pal.text, t.color_for("text.primary")])
    scale = sorted({s.size_pt for s in typo.scale} | {float(x) for x in (typo.sizes_used or []) if x >= 8})
    h1, title_color, title_bold, title_font = _title_style(manifest, family, pal)
    if contrast_ratio(title_color, pal.bg) < 3.0:
        title_color = pal.text
    if contrast_ratio(title_color, pal.bg) < 4.5:
        title_color = _heading_on(pal.bg, manifest, pal)
    if title_ph is not None:
        _heading_on_slide(title_ph, (sx, sy, sw, sh), W, H, warnings)
    if title_ph is not None and art is None and ground is None:
        _reclaim_left(slide, title_ph, (sx, sy, sw, sh), manifest, W, H, warnings)
    saved_title = copy.deepcopy(title_ph._element) if title_ph is not None else None
    n_warn = len(warnings)
    heading_bottom, content_left, lede_used = _place_heading(builder, slide, title_ph, canvas, oslide, manifest, ws, outline, pal, (sx, sy, sw, sh), h1, title_color, title_bold, title_font, scale, warnings, ds)
    h1 = _set_size(title_ph, h1)  # the gap under the heading follows the heading as set (a canvas's 150 pt display word was fitted to 26 pt)
    below = _content_top(slide, heading_bottom, h1, W, H)
    if saved_title is not None and _content_bottom_limit(slide, sy + sh, W, H) - below < int(0.30 * H) and title_ph is not None and _backing_of(slide, title_ph._element, W, H) is None:
        # the heading left less than 30 % of the slide for the content: it is set again a step (at most two) smaller
        for size in [x for x in sorted(scale, reverse=True) if x < ds.head_size - 0.05][:2]:
            title_ph = _restore_shape(slide, title_ph, saved_title)
            del warnings[n_warn:]
            heading_bottom, content_left, lede_used = _place_heading(builder, slide, title_ph, canvas, oslide, manifest, ws, outline, pal, (sx, sy, sw, sh), h1, title_color, title_bold, title_font, scale, warnings, dataclasses.replace(ds, head_size=size))
            h1 = _set_size(title_ph, h1)
            below = _content_top(slide, heading_bottom, h1, W, H)
            if _content_bottom_limit(slide, sy + sh, W, H) - below >= int(0.30 * H):
                break
        warnings.append("heading set smaller to leave room for the content")
    if title_ph is not None:
        turned = _keep_turned_heading(title_ph, manifest, W, H, warnings)
        if turned is not None and turned > heading_bottom:
            heading_bottom = turned
            below = _content_top(slide, heading_bottom, h1, W, H)
    if lede_used:
        oslide = _without_lede(oslide)
    page_left = sx
    if content_left is not None and not (abs(content_left - int(_grid_x(manifest) * W)) <= 0.03 * W or content_left <= page_left + int(0.02 * W)):
        content_left = None  # an indented or centred heading does not set the content column: the page grid does
    if content_left is not None and sx <= content_left < sx + sw // 2:
        sw = sx + sw - content_left
        sx = content_left
    # the right margin mirrors the left one — the page's, where the column starts past the art; a page whose own left
    # margin is wider than its right one (art down the left edge: a helix, a sidebar) keeps its own right edge
    # (a column that starts far into the slide — past a quarter of it — never mirrors itself: its right edge would
    # close in on its left one; the page's own margin is mirrored then)
    mirror = page_left if (art is not None and art.side == "left") or sx - page_left > int(0.15 * W) else sx
    lopsided = page_left - (W - int(safe.x2 * W)) > int(0.04 * W)
    if ground is None and sx + sw > W - mirror and not lopsided:
        sw = W - mirror - sx
    if art is not None:
        # the content keeps to the free side of the art, a gutter away from it
        g = int(0.04 * W)
        if art.side == "right":
            sw = max(min(sw, art.box.x - g - sx), int(0.3 * W))
        else:
            x1 = max(sx, art.box.x2 + g)
            sw, sx = max(sx + sw - x1, int(0.3 * W)), x1
    top = _content_top(slide, heading_bottom, h1, W, H)
    bottom = _content_bottom_limit(slide, sy + sh, W, H)
    if bottom - top < int(0.25 * H):
        # never an area past the slide's foot: what does not fit is restructured by the composer, not pushed off
        warnings.append("little room under the heading")
    if art is None and ground is None and sw < int(0.4 * W):
        # never a sliver: a column narrower than 40 % of the slide with nothing drawn beside it (a sample's column next
        # to a picture the composed slide does not keep) sets letters one under the other — the page's width instead
        wide_x = min(page_left, sx)
        wide_w = max(int(safe.x2 * W), sx + sw) - wide_x
        if wide_w > sw:
            warnings.append(f"content column widened from {sw / W:.0%} to {wide_w / W:.0%} of the slide")
            sx, sw = wide_x, wide_w
    area = Bbox(x=sx, y=top, w=sw, h=max(min(bottom, H) - top, int(0.1 * H)))
    if area.y2 > H:
        area = Bbox(x=area.x, y=max(H - area.h, 0), w=area.w, h=min(area.h, H))
    foot_room = None
    if art is None:
        full = area
        area = _clear_of_art(slide, area, title_ph, W, H, warnings)
        foot_room = _foot_room(slide, full, area, title_ph, W, H)
    _true_ground(slide, area, pal, manifest)
    proto = _card_proto(builder, manifest, family, pal.bg)
    kit = Kit(manifest, W, H, pal.bg, card_proto=proto[0] if proto else None, heading_color=title_color)
    kit.head_size = ds.head_size
    if proto:
        # colours of the copied card as the analysis sees them (scheme colours resolved)
        kit.card.fill, kit.card.line = proto[1], proto[2]
        kit.card.colors = kit.colors_on(proto[1] or pal.bg, prefer=None if proto[1] else title_color, inside_card=bool(proto[1]))
    composer = Composer(slide, kit, oslide, outline, outline.strategy or "structured", manifest)
    composer.foot_room = foot_room
    if comp == "image_text" and oslide.content.image_hint:
        # the picture and its text above the slide's conclusion and footnote (compose() keeps the same room)
        n0 = len(composer.cv.tree)
        composer._unsay_conclusion()
        comp = _image_text(builder, slide, oslide, manifest, ws, kit, composer, composer._reserve_notes(area))
        if comp is None:
            composer._draw_notes(n0)
    if comp:
        composer.compose(comp, area)
    _ink_pass(slide, composer.cv.added, pal, manifest, warnings)
    warnings.extend(composer.warnings)
    return slide, warnings


def _inks(pal: "_Palette", manifest: TemplateManifest) -> list[str]:
    """The template's own text colours for a ground they were not chosen for, in order: the palette's text colours,
    the template's text colour, its dark and its light ground (a navy heading on the light top of a navy template
    reads as the template's; plain black does not) — black and white last."""
    t = manifest.tokens
    out = [c for c in (pal.text, pal.text2, t.color_for("text.primary"), t.color_for("background.dark"), t.color_for("background.light")) if c]
    return list(dict.fromkeys(out + ["000000", "FFFFFF"]))


UNREADABLE = 2.0  # under this contrast a text is not read at all (the audit's error line on a picture or a gradient)


def unreadable_text_pass(slide: Slide, manifest: TemplateManifest) -> int:
    """The last guard of every slide, cloned or composed: a text whose colour all but vanishes into what lies under it
    (under 2:1 — grey type on the grey frame past the edge of a white page, white on a light band) takes a colour of
    the template that reads there: its text colours, its dark and light grounds, then black or white. Only colours
    written on the runs are judged. Returns the number of runs recoloured."""
    from verstka.rendering.charts import slide_ground

    t = manifest.tokens
    inks = list(dict.fromkeys([c for c in (t.color_for("text.primary"), t.color_for("text.secondary"), t.color_for("background.dark"), t.color_for("background.light")) if c] + ["000000", "FFFFFF"]))
    fixed = 0
    tree = slide._element.cSld.find(q("p:spTree"))
    for el in list(tree):
        if etree.QName(el).localname != "sp":
            continue
        runs = [r for r in el.iter(q("a:r")) if (r.findtext(q("a:t")) or "").strip()]
        b = element_bbox(el) if runs else None
        if not b or b[2] <= 0 or b[3] <= 0:
            continue
        try:
            g = slide_ground(slide, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
        except Exception:  # noqa: BLE001
            g = None
        if not g or len(g) != 6:
            continue
        for r in runs:
            rpr = r.find(q("a:rPr"))
            clr = rpr.find(q("a:solidFill") + "/" + q("a:srgbClr")) if rpr is not None else None
            if clr is None or not clr.get("val") or contrast_ratio(clr.get("val"), g) >= UNREADABLE:
                continue
            clr.set("val", next((c for c in inks if contrast_ratio(c, g) >= 4.5), max(inks, key=lambda c: contrast_ratio(c, g))))
            fixed += 1
    return fixed


def _ink_pass(slide: Slide, added: list, pal: "_Palette", manifest: TemplateManifest, warnings: list[str]) -> None:
    """Text the composer set in the palette's colours for the area's ground, standing where the ground is another (a
    background running from light at the top to navy at the bottom — Office «Circuit»: a legend in the light top half
    was set white for the dark half): such a text takes a colour that reads on what lies under it — the template's
    own inks first (`_inks`). Text on a card or a badge of its own keeps its colours (the card is its ground); text
    whose ground is the one the palette was chosen for is never touched."""
    from verstka.rendering.charts import delta_e, slide_ground

    inks = _inks(pal, manifest)
    changed = 0
    for el in added:
        if etree.QName(el).localname != "sp" or el.find(q("p:txBody")) is None or _paints(el):
            continue
        runs = [r for r in el.iter(q("a:r")) if (r.findtext(q("a:t")) or "").strip()]
        b = element_bbox(el) if runs else None
        if not b or b[2] <= 0 or b[3] <= 0:
            continue
        try:
            g = slide_ground(slide, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]))
        except Exception:  # noqa: BLE001
            g = None
        if not g or len(g) != 6:
            continue
        # on the ground the palette was chosen for only what fails outright (under 3:1: an orange figure on the orange
        # gradient of Office «Berlin») — its deliberate choices stand; elsewhere the full line
        own = delta_e(g, pal.bg) <= 10
        for r in runs:
            rpr = r.find(q("a:rPr"))
            clr = rpr.find(q("a:solidFill") + "/" + q("a:srgbClr")) if rpr is not None else None
            if clr is None or not clr.get("val"):
                continue
            size = int(rpr.get("sz") or 1800) / 100.0
            need = 3.0 if own or size >= 18 or (rpr.get("b") in ("1", "true") and size >= 14) else 4.5
            if contrast_ratio(clr.get("val"), g) >= need:
                continue
            clr.set("val", next((c for c in inks if contrast_ratio(c, g) >= need), max(inks, key=lambda c: contrast_ratio(c, g))))
            changed += 1
    if changed:
        warnings.append(f"{changed} text run(s) recoloured for the ground under them")


def _keep_turned_heading(title_ph, manifest: TemplateManifest, W: int, H: int, warnings: list[str]) -> Optional[int]:
    """A heading turned with its band (LibreOffice «Progress»: titles tilted 8° along a diagonal band): its box narrows
    to its letters and moves in until its turned corners stand on the slide — a turned box as wide as the slide lifted
    the end of a long heading over the top edge. Returns the lowest point the turned heading reaches; None when the
    heading is not turned."""
    try:
        rot = float(title_ph.rotation or 0.0)
    except Exception:  # noqa: BLE001
        return None
    a = ((rot + 180.0) % 360.0) - 180.0
    if abs(a) < 1.5 or None in (title_ph.left, title_ph.top, title_ph.width, title_ph.height):
        return None
    box = Bbox(x=int(title_ph.left), y=int(title_ph.top), w=int(title_ph.width), h=int(title_ph.height))
    insets = _heading_insets(title_ph)
    ink = _heading_ink(title_ph, box, manifest)
    x, w = box.x, box.w
    need = ink.w + insets[0] + insets[2] + int(0.01 * W)
    if need < w:
        x, w = max(ink.x - insets[0], box.x), need
    t = math.radians(a)
    cx, cy = x + w / 2.0, box.y + box.h / 2.0
    pts = [(cx + dx * math.cos(t) - dy * math.sin(t), cy + dx * math.sin(t) + dy * math.cos(t)) for dx in (-w / 2.0, w / 2.0) for dy in (-box.h / 2.0, box.h / 2.0)]
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    mx, my = 0.02 * W, 0.02 * H
    dx = (mx - min(xs)) if min(xs) < mx else ((W - mx) - max(xs) if max(xs) > W - mx else 0.0)
    dy = (my - min(ys)) if min(ys) < my else ((H - my) - max(ys) if max(ys) > H - my else 0.0)
    title_ph.left, title_ph.top, title_ph.width = Emu(int(x + dx)), Emu(int(box.y + dy)), Emu(int(w))
    if dx or dy or w != box.w:
        warnings.append("turned heading narrowed to its letters and kept on the slide")
    return int(max(ys) + dy)


def _set_size(title_ph, default: float) -> float:
    """The size the heading was set at: its largest run (the canvas's nominal title size when it has no runs)."""
    if title_ph is None:
        return default
    sizes = [int(r.get("sz")) / 100.0 for r in title_ph._element.iter(q("a:rPr")) if (r.get("sz") or "").isdigit()]
    return max(sizes) if sizes else default


def _heading_on_slide(title_ph, safe_box: tuple[int, int, int, int], W: int, H: int, warnings: list[str]) -> None:
    """A canvas's heading box that starts off the slide or runs past it (a display word set wider than the slide,
    «Our Services» from x = −2 %): the composed heading keeps to the page — from the safe area's left edge at the
    least, to its right edge at the most."""
    if title_ph.left is None or title_ph.width is None:
        return
    x, w = int(title_ph.left), int(title_ph.width)
    sx, sy, sw, sh = safe_box
    left = max(x, min(sx, int(0.03 * W)) if x < 0 else x)
    right = min(x + w, W - max(W - (sx + sw), int(0.03 * W)) if x + w > W else x + w)
    if left == x and right == x + w:
        return
    title_ph.left = Emu(left)
    title_ph.width = Emu(max(right - left, int(0.2 * W)))
    warnings.append("heading box kept on the slide (the sample's box ran past its edge)")


def _reclaim_left(slide: Slide, title_ph, safe_box: tuple[int, int, int, int], manifest: TemplateManifest, W: int, H: int, warnings: list[str]) -> None:
    """A canvas whose heading stands in a column far into the slide — a sample with a picture of its own on the left
    («Planet One is Mercury» beside a planet, a report page beside a photo): the composed slide keeps none of that
    picture, so the heading would open a narrow column beside an empty half. When nothing the slide still draws stands
    left of the heading (no art, no panel, no chrome), the heading moves to the deck's heading grid (or the page
    margin) and keeps its right edge."""
    if title_ph.left is None or title_ph.width is None or title_ph.top is None:
        return
    x = int(title_ph.left)
    sx, sy, sw, sh = safe_box
    gx = int(_grid_x(manifest) * W)
    target = gx if sx - int(0.01 * W) <= gx < int(0.25 * W) else sx
    if x - target < int(0.2 * W):
        return
    top = int(title_ph.top)
    band = Bbox(x=target, y=top, w=x - target, h=max(sy + sh - top, int(0.2 * H)))
    try:
        from verstka.rendering.layers import art_boxes, art_share

        skip = (title_ph._element,)
        if art_share(slide, band, skip=skip) > 0.02 or any(b.intersection(band) > 0.004 * W * H for b in art_boxes(slide, skip=skip)):
            return
    except Exception:  # noqa: BLE001 - the layers are advice: the canvas stays as it is
        return
    # the slide's own shapes that stay (chrome kept from the canvas: a logo, a page number) must not stand there either
    for shp in slide.shapes:
        if shp._element is title_ph._element or shp.left is None or shp.width is None or shp.top is None or shp.height is None:
            continue
        b = Bbox(x=int(shp.left), y=int(shp.top), w=int(shp.width), h=int(shp.height))
        if b.w * b.h < 0.9 * W * H and b.intersection(band) > 0.3 * max(b.area, 1):
            return
    right = x + int(title_ph.width)
    title_ph.left = Emu(target)
    title_ph.width = Emu(max(right - target, int(title_ph.width)))
    warnings.append("heading moved to the page grid: its sample's column stood beside a picture the composed slide does not keep")


def _clear_of_art(slide: Slide, area: Bbox, title_ph, W: int, H: int, warnings: list[str]) -> Bbox:
    """The content area clear of the template's art (illustrations, triangles, trees, photos the layout or the master
    draws — contract C2 layers.free_rect): when art covers 3 % of the area or more, the content takes the largest free
    rectangle inside it (below that, small marks are left to the foot and logo rules)."""
    try:
        from verstka.rendering.layers import art_boxes, art_share, free_rect
    except ImportError:
        return area
    skip = (title_ph._element,) if title_ph is not None else ()
    try:
        share = art_share(slide, area, skip=skip)
        # an illustration standing mostly inside the area (a tree in its corner) is art to keep clear of even when it
        # covers little of the area; specks and small logos are not (the foot and logo rules take them)
        intrudes = any(b.w * b.h >= 0.008 * W * H and b.intersection(area) >= 0.4 * b.area for b in art_boxes(slide, skip=skip))
        if share < 0.03 and not intrudes:
            return area
        free = free_rect(slide, area, skip=skip)
        # a composed block reads across the page: a rectangle as wide as the area, cut below the art at its top (a
        # cloud, a mark under the heading) or above the art at its foot, beats a taller half-width column
        pad = int(0.015 * W)
        score = lambda b: b.area * (b.w / max(area.w, 1))  # noqa: E731
        for b in art_boxes(slide, skip=skip):
            if b.x2 <= area.x or b.x >= area.x2 or b.y2 <= area.y or b.y >= area.y2:
                continue
            cuts = []
            if b.y2 + pad < area.y + 0.4 * area.h:
                cuts.append(Bbox(x=area.x, y=b.y2 + pad, w=area.w, h=area.y2 - b.y2 - pad))
            if b.y - pad > area.y + 0.6 * area.h:
                cuts.append(Bbox(x=area.x, y=area.y, w=area.w, h=b.y - pad - area.y))
            for cut in cuts:
                alt = free_rect(slide, cut, skip=skip)
                if alt.h >= 0.6 * area.h and score(alt) > score(free) * 1.1:
                    free = alt
    except Exception:  # noqa: BLE001
        return area
    if free.w < 0.3 * W or free.h < 0.2 * H:
        warnings.append(f"template art covers {share:.0%} of the content area: no free rectangle large enough")
        return area
    warnings.append(f"content kept clear of the template's art ({share:.0%} of the area)")
    return free


def _foot_room(slide: Slide, full: Bbox, area: Bbox, title_ph, W: int, H: int) -> Optional[Bbox]:
    """The free room under a content area that was cut above the template's art (trees at one side of the foot, a
    grass band): the largest rectangle clear of the art between the area's foot and the full area's, at least 0.4 of
    the area wide and 0.06 of the slide high — where a conclusion the content left no room for can stand (B2)."""
    if area.y2 >= full.y2 - int(0.06 * H):
        return None
    try:
        from verstka.rendering.layers import free_rect

        skip = (title_ph._element,) if title_ph is not None else ()
        band = Bbox(x=full.x, y=area.y2, w=full.w, h=full.y2 - area.y2)
        room = free_rect(slide, band, skip=skip)
    except Exception:  # noqa: BLE001
        return None
    if room.w < 0.4 * area.w or room.h < 0.06 * H or room.y > area.y2 + int(0.02 * H):
        return None
    return room


def _true_ground(slide: Slide, area: Bbox, pal: "_Palette", manifest: TemplateManifest) -> None:
    """The ground the content really stands on (what the slide, its layout and master paint under the area — a
    picture ground, a panel): when it is not the manifest's ground (ΔE > 10) and the palette's text does not read on
    it, the palette takes that ground and readable text colours (contract C4, T07)."""
    try:
        from verstka.rendering.charts import delta_e, slide_ground

        g = slide_ground(slide, area)
    except Exception:  # noqa: BLE001
        return
    if not g or len(g) != 6 or delta_e(g, pal.bg) <= 10:
        return
    # the ground is recorded even when the text already reads on it: the charts, cards and figures choose their
    # colours for it (a dark teal accent on a purple picture ground is invisible, whatever the text does)
    pal.bg = g
    if contrast_ratio(pal.text, g) >= 4.5:
        return
    pal.text = _readable(g, [pal.text, manifest.tokens.color_for("text.primary")])
    pal.text2 = pal.text if contrast_ratio(pal.text2, g) < 4.5 else pal.text2


def _content_top(slide: Slide, heading_bottom: int, h1: float, W: int, H: int) -> int:
    """The content line: a gap under the heading, below a rule the layout draws under it."""
    gap = max(int(0.05 * H), int(h1 * 0.8 * EMU_PER_PT))
    return _below_rules(slide, heading_bottom, heading_bottom + gap, W, H, gap)


def _content_bottom_limit(slide: Slide, safe_bottom: int, W: int, H: int) -> int:
    """The foot of the content area: the safe area's, clear of the logos and the footer line."""
    bottom = safe_bottom
    foot = _foot_top(slide, W, H)
    if foot is not None:
        bottom = min(bottom, foot - int(0.035 * H))  # the content keeps clear of the logos and the footer line
    return bottom


def _restore_shape(slide: Slide, shape, saved) -> object:
    """Put a saved copy of a shape back in place of the shape (to set it again) and return its new proxy."""
    el = shape._element
    new = copy.deepcopy(saved)
    el.addprevious(new)
    el.getparent().remove(el)
    sid = new.find(".//" + q("p:cNvPr")).get("id")
    return next((sh for sh in slide.shapes if str(sh.shape_id) == sid), shape)


def _snap_header(builder: DeckBuilder, manifest: TemplateManifest, ds: "DeckStyle", slide: Slide, title_ph, art: "Art") -> None:
    """An art canvas keeps the deck's header grid: its heading (and the label under it) moves to where the heading
    of every other content slide stands — or into the free column when that spot is taken by the art."""
    W, H = builder.slide_w, builder.slide_h
    std = _pick_canvas(builder, manifest, ds.family, "bullets", False, need_pill=ds.pill)
    t = _top_title(std) if std is not None else None
    if t is None:
        return
    try:
        el_std = slide_shape_elements(builder.source_slide(std.source_slide)).get(t.shape_id)
    except IndexError:
        el_std = None
    sb = element_bbox(el_std) if el_std is not None else None
    cur = element_bbox(title_ph._element)
    if not sb or not cur:
        return
    backing = _backing_of(slide, title_ph._element, W, H)
    target_x, target_y = sb[0], sb[1]
    g = int(0.04 * W)
    if art.side == "left" and target_x < art.box.x2:
        target_x = art.box.x2 + g  # the heading moves off the art into the free column
    dx, dy = target_x - cur[0], target_y - cur[1]
    if abs(dx) < int(0.004 * W) and abs(dy) < int(0.004 * H):
        return
    from verstka.rendering.deck import shift_element

    from verstka.rendering.deck import set_element_pos

    shift_element(title_ph._element, dx, dy)
    if backing is not None:
        shift_element(backing[0], dx, dy)
    nb = element_bbox(title_ph._element)
    if nb:
        # the heading keeps inside its column: up to the art on the right, up to the right margin otherwise
        right = art.box.x - g if art.side == "right" else int(manifest.tokens.spacing.safe_area.x2 * W)
        if nb[0] + nb[2] > right:
            set_element_pos(title_ph._element, w=max(right - nb[0], int(0.25 * W)))


def _image_text(builder, slide, oslide, manifest, ws, kit, composer, area: Bbox) -> Optional[str]:
    """A picture of the template next to the text: the picture takes 42% of the width, the text the rest."""
    path = pick_asset(manifest, ws, oslide.content.image_hint, require_match=True)
    if not path:
        return "bullets"
    img_w = int(area.w * 0.42)
    insert_picture(slide, Bbox(x=area.x + area.w - img_w, y=area.y, w=img_w, h=area.h), path)
    text_area = Bbox(x=area.x, y=area.y, w=area.w - img_w - kit.gap * 2, h=area.h)
    c = oslide.content
    if c.bullets:
        composer.bullets(text_area, list(c.bullets))
    elif c.paragraphs:
        composer.paragraphs(text_area, list(c.paragraphs))
    return None


_KIND_LABEL = {
    PatternKind.stat_row: "Ключевые цифры", PatternKind.big_number: "Главная цифра", PatternKind.chart: "Динамика",
    PatternKind.table: "Сравнение", PatternKind.comparison: "Сравнение", PatternKind.two_column: "Сравнение",
    PatternKind.process: "План", PatternKind.timeline: "Этапы", PatternKind.cards: "Главное", PatternKind.bullets: "Тезисы",
    PatternKind.agenda: "Содержание", PatternKind.quote: "Цитата", PatternKind.image_text: "Пример", PatternKind.team: "Команда",
}
_CHART_LABEL = {"pie": "Структура", "doughnut": "Структура", "line": "Динамика", "area": "Динамика", "column": "Сравнение", "bar": "Сравнение"}
_COVER_KINDS = (PatternKind.title, PatternKind.section, PatternKind.thanks)


def _kind_label(oslide: OutlineSlide) -> str:
    """The label that names the kind of slide: a chart by what it shows (the parts of a whole are not a trend)."""
    c = oslide.content
    if oslide.kind == PatternKind.chart and c.chart is not None:
        return _CHART_LABEL.get(str(c.chart.type or "").lower(), "Данные")
    if oslide.kind == PatternKind.bullets and (c.formula or "").strip():
        return "Расчёт"
    return _KIND_LABEL.get(oslide.kind, "")


@dataclass
class DeckStyle:
    """Decisions taken once per deck, so every composed slide wears the same heading and ground."""

    family: Family
    pill: bool  # headings of the template sit on a label (a pill)
    kicker: bool  # the label carries a short kicker, the headline is set below it
    head_size: float  # size of plain headings (the whole deck)
    upper: bool  # labels in capitals, as in the samples
    head_bold: bool = False  # headlines set under a label are bold when the template's labels are
    head_cap_reason: str = ""  # why the heading size is not the template's own content heading size (for the log)
    caps: bool = False  # plain headings in capitals: the template's own content headings are all typed so


def _is_sparse(manifest: TemplateManifest) -> bool:
    """A template whose own type scale is too thin for composed content (the analysis derived a ladder of sizes from
    the slide height for it)."""
    return bool(getattr(manifest.tokens.typography, "derived_sizes", None))


def _scale_of(manifest: TemplateManifest) -> list[float]:
    typo = manifest.tokens.typography
    return sorted({s.size_pt for s in typo.scale} | {float(x) for x in (typo.sizes_used or []) if x >= 8})


def _capped_head(builder: DeckBuilder, manifest: TemplateManifest, outline: DeckOutline, family: Family, h1c: float, font: Optional[str], bold: bool, scale: list[float]) -> tuple[float, str]:
    """The deck's heading size on a template whose heading size does not suit composed content — a sparse scale or a
    display-sized heading (> 7.5 % of the slide height): the largest size of the scale, at most the template's own
    content heading and 6.2 % of the slide height, at which 80 % of the deck's headlines take at most two lines at
    the heading width; never below the plain rule's minimum (4.6 % of the slide height)."""
    from verstka.rendering.fonts import wrap_lines

    W, H = builder.slide_w, builder.slide_h
    hpt = H / EMU_PER_PT
    floor = _snap_heading(0.046 * hpt, scale)
    cap = min(h1c, 0.062 * hpt)
    safe = manifest.tokens.spacing.safe_area
    std = _canvas_for(manifest, family, "bullets")
    t = _top_title(std) if std is not None else None
    x = t.bbox.x if t is not None else safe.x
    width = min(0.86 * W, max(safe.x2 - x, 0.5) * W) - 2 * 91440
    heads = [s.headline for s in outline.slides if s.kind not in _COVER_KINDS and s.headline]
    cands = sorted({x_ for x_ in scale if floor - 0.05 <= x_ <= cap + 0.05}, reverse=True)
    if not heads:
        return (cands[0] if cands else floor), "sparse/display heading: capped"
    from verstka.ru import typeset

    for size in cands:
        ok = sum(1 for h in heads if len(wrap_lines(typeset(h), font, size, bold, width / EMU_PER_PT)) <= 2)
        if ok >= 0.8 * len(heads):
            return size, f"heading {h1c:g} pt capped to {size:g} pt (two lines for {ok}/{len(heads)} headlines)"
    return floor, f"heading {h1c:g} pt capped to the floor {floor:g} pt"




def _content_patterns(manifest: TemplateManifest) -> list[Pattern]:
    return [p for p in manifest.patterns if not p.reference and p.quality >= 0.5 and p.kind not in _NOT_CONTENT_CANVAS and (t := _top_title(p)) is not None and t.bbox.y <= 0.3]


def deck_style(builder: DeckBuilder, manifest: TemplateManifest, outline: DeckOutline) -> DeckStyle:
    cache = builder.__dict__.setdefault("_deck_styles", {})  # one deck per builder: the style lives with it
    if id(outline) in cache:
        return cache[id(outline)]
    W, H = builder.slide_w, builder.slide_h
    pats = _content_patterns(manifest)
    fams = Counter(p.family for p in pats) or Counter(p.family for p in manifest.patterns)
    family = fams.most_common(1)[0][0] if fams else Family.light
    same = [p for p in pats if p.family == family]
    pilled, sizes, uppers, texted, bolds = 0, Counter(), 0, 0, 0
    title_box: Optional[Bbox] = None
    for p in same:
        try:
            src = builder.source_slide(p.source_slide)
        except IndexError:
            continue
        els = slide_shape_elements(src)
        t = _top_title(p)
        el = els.get(t.shape_id) if t else None
        if el is None:
            continue
        if _backing_of(src, el, W, H) is not None:
            pilled += 1
            if t.style.size_pt:
                sizes[round(t.style.size_pt)] += 1
            if t.sample_text and t.sample_text.strip():
                texted += 1
                uppers += int(t.sample_text.strip().isupper())
            bolds += int(bool(t.style.bold))
            b = element_bbox(el)
            if b and title_box is None:
                title_box = Bbox(x=b[0], y=b[1], w=b[2], h=b[3])
    pill = bool(same) and pilled >= 0.5 * len(same)
    typo = manifest.tokens.typography
    scale = sorted({s.size_pt for s in typo.scale} | {float(x) for x in (typo.sizes_used or []) if x >= 8})
    pal = _Palette(manifest, family)
    h1, _, bold, font = _title_style(manifest, family, pal)
    heads = [s.headline for s in outline.slides if s.kind not in _COVER_KINDS and s.headline]
    kicker = False
    if pill and heads:
        size = float(sizes.most_common(1)[0][0]) if sizes else h1
        room = clear_width(manifest, W, H, title_box) if title_box else int(0.5 * W)
        fails = sum(1 for h in heads if len(h.split()) > 4 or text_width_pt(h, font, size, bold) * EMU_PER_PT + int(0.06 * W) > room)
        kicker = fails > len(heads) / 3
    # one heading size for every content slide of every deck: the template's own (a long headline takes a third
    # line rather than a smaller size) — unless it is a display size or the scale is too thin to trust (then capped)
    reason = ""
    if _is_sparse(manifest) or h1 > 0.075 * H / EMU_PER_PT:
        head, reason = _capped_head(builder, manifest, outline, family, h1, font, bold, scale)
    else:
        head = max(h1, _snap_heading(0.046 * H / EMU_PER_PT, scale))
    # a template whose content headings are all typed in capitals («OUR SERVICES», SlidesCarnival and Google Slides
    # decks) keeps that voice on the composed headings (not in pill mode: there the label carries it)
    typed = [(_top_title(p).sample_text or "").strip() for p in same if _top_title(p) is not None]
    typed = [t for t in typed if any(ch.isalpha() for ch in t)]
    caps = not pill and len(typed) >= 3 and sum(1 for t in typed if t.isupper()) >= 0.8 * len(typed)
    ds = DeckStyle(family=family, pill=pill, kicker=kicker, head_size=head, upper=pill and texted > 0 and uppers >= 0.6 * texted, head_bold=pill and bolds > 0, head_cap_reason=reason, caps=caps)
    cache[id(outline)] = ds
    return ds


def _place_heading(builder, slide, title_ph, canvas, oslide, manifest, ws, outline, pal, safe_box, h1, title_color, title_bold, title_font, scale, warnings, ds: Optional[DeckStyle] = None) -> tuple[int, Optional[int], bool]:
    """Write the headline; returns (bottom of the heading, left edge the content aligns to, whether the slide's
    lede was set as the headline)."""
    from verstka.rendering.compose import Canvas, Para, Run, block_height_pt

    typo = manifest.tokens.typography
    W, H = builder.slide_w, builder.slide_h
    sx, sy, sw, sh = safe_box
    insets = (91440, 45720, 91440, 45720)
    if title_ph is None:
        box = Bbox(x=sx, y=sy, w=sw, h=int(H * 0.16))
        tb = _textbox(slide, box, [ParagraphSpec(oslide.headline)], size=h1, color=title_color, font=title_font, bold=title_bold, scale=scale, anchor="t")
        res = fit_size([oslide.headline], box, title_font, h1, title_bold, scale, insets_emu=(45720, 22860, 45720, 22860), line_spacing=typo.line_height)
        th = int(res.height_pt * EMU_PER_PT) + 45720
        tb.height = Emu(th)
        return box.y + th, sx + 45720, False
    _materialize(title_ph)  # an inherited placeholder gets its own geometry before any write (heading at x=0)
    if title_ph.top is not None and title_ph.left is not None and title_ph.width is not None and title_ph.height is not None:
        box = Bbox(x=int(title_ph.left), y=int(title_ph.top), w=int(title_ph.width), h=int(title_ph.height))
    else:
        box = Bbox(x=sx, y=sy, w=sw, h=int(H * 0.16))
    backing = _backing_of(slide, title_ph._element, W, H) if canvas is not None else None
    if backing is not None:
        back_el, bb = backing
        kicker_mode = ds.kicker if ds is not None else (len(oslide.headline) > 34)
        if kicker_mode:
            # a pill is a label: it carries the section (a kicker), the assertion is set as a heading below it
            upper = ds.upper if ds is not None else False
            head_size = ds.head_size if ds is not None else _snap_heading(max(h1 * 1.25, 0.046 * H / EMU_PER_PT), scale)
            head_bold = ds.head_bold if ds is not None else title_bold
            headline = oslide.headline.strip()
            short = len(headline.split()) <= 5
            from verstka.ru import typeset

            display = typeset(headline)
            tag = (oslide.section or "").strip() if short else _section_of(oslide, outline)
            only_label = False
            if not tag or _same_words(tag, headline):
                if len(headline.split()) <= 5:
                    tag, only_label = headline, True  # the headline is itself a section name: it goes on the label
                else:
                    tag = _kind_label(oslide)
            short_tag = _short_kicker(tag)
            if only_label and short_tag.lower().strip(" .:") != headline.lower().strip(" .:"):
                # a label would cut the headline («Результаты пилота: до и после» → «Результаты пилота»): the whole
                # headline is set under a label that names the kind of slide
                only_label = False
                short_tag = _short_kicker(_kind_label(oslide) or short_tag)
            tag = short_tag
            label = tag.upper() if upper else tag
            hctx = _SlideCtx(builder, slide, canvas, manifest, ws, outline)
            nb = hctx._widen_on_backing(title_ph._element, (box.x, box.y, box.w, box.h), [label], title_font, h1, title_bold, insets)
            # the label keeps the template's own size on every slide; a long kicker was shortened, never shrunk
            _fill_keep(title_ph._element, [ParagraphSpec(label)], size_pt=h1)
            _hug_label(back_el, title_ph._element, label, title_font, h1, title_bold, insets, W)
            b2 = element_bbox(back_el)
            below = max(bb.y2, (b2[1] + b2[3]) if b2 else bb.y2)
            y = below + int(0.022 * H)
            one_line = int(head_size * typo.line_height * EMU_PER_PT)
            if only_label:
                lede = _lede_of(oslide)
                if not lede:
                    # the content starts where it starts under a one-line headline: one rhythm for the whole deck
                    return y + one_line, bb.x, False
                display, lede_used = typeset(lede), True  # the slide's lede is its message: it takes the heading's place
            else:
                lede_used = False
            ground = pal.bg
            color = _heading_on(ground, manifest, pal)
            right = min(int(manifest.tokens.spacing.safe_area.x2 * W), W - bb.x)
            width = min(_free_width(slide, Bbox(x=bb.x, y=y, w=right - bb.x, h=int(0.1 * H)), W, H, right, skip=(title_ph._element, back_el)), int(0.86 * W))
            size = head_size
            from verstka.rendering.fonts import wrap_lines

            n = len(wrap_lines(display, title_font, size, head_bold, width / EMU_PER_PT))
            if n > 3:
                for s_ in [x for x in reversed(scale) if 0.8 * head_size <= x < head_size]:
                    size = s_
                    n = len(wrap_lines(display, title_font, size, head_bold, width / EMU_PER_PT))
                    if n <= 3:
                        break
            if 2 <= n <= 3:
                # balanced lines: the narrowest width that keeps the same number of lines
                lo, hi = int(width * 0.55), width
                for _ in range(12):
                    mid = (lo + hi) // 2
                    if len(wrap_lines(display, title_font, size, head_bold, mid / EMU_PER_PT)) <= n:
                        hi = mid
                    else:
                        lo = mid
                width = min(width, int(hi * 1.03))
            hpt = block_height_pt([Para([Run(display, size, color, head_bold, title_font)])], width / EMU_PER_PT, typo.line_height)
            hh = int(hpt * EMU_PER_PT) + int(2 * EMU_PER_PT)
            cv = Canvas(slide, font=title_font, spacing_pct=typo.line_spacing if abs(typo.line_spacing - 1.2) > 1e-6 else 1.0)
            cv.text(Bbox(x=bb.x, y=y, w=width, h=hh), [Para([Run(display, size, color, head_bold, title_font)])], name="Title")
            warnings.append("heading set under its label")
            return y + max(hh, one_line), bb.x, lede_used
        hctx = _SlideCtx(builder, slide, canvas, manifest, ws, outline)
        nb = hctx._widen_on_backing(title_ph._element, (box.x, box.y, box.w, box.h), [oslide.headline], title_font, h1, title_bold, insets)
        box = Bbox(x=nb[0], y=nb[1], w=nb[2], h=nb[3])
        warnings.extend(hctx.warnings)
        res = fit_size([oslide.headline], box, title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=typo.line_height)
        _fill_keep(title_ph._element, [ParagraphSpec(oslide.headline)], size_pt=res.size_pt)
        b2 = element_bbox(back_el)
        return max(box.y2, (b2[1] + b2[3]) if b2 else box.y2), bb.x, False
    return _place_plain_heading(builder, slide, title_ph, box, oslide, manifest, pal, h1, title_color, title_bold, title_font, scale, warnings, ds, canvas)


def _place_plain_heading(builder, slide, title_ph, box: Bbox, oslide, manifest, pal, h1, title_color, title_bold, title_font, scale, warnings, ds, canvas) -> tuple[int, Optional[int], bool]:
    """A heading written into the canvas's title shape (no label under it). It takes the free width of its band (up
    to the logos and the right margin), not the sample's box; a centred heading widens about its centre; a heading
    printed on a band of the layout stays inside the band or is set under it; a rule under the heading is a limit."""
    import functools

    from verstka.rendering.fit import fit_size as _fit_size
    from verstka.rendering.fonts import wrap_lines as _wrap_lines
    from verstka.ru import typeset

    typo = manifest.tokens.typography
    W, H = builder.slide_w, builder.slide_h
    safe = manifest.tokens.spacing.safe_area
    sparse = _is_sparse(manifest)
    insets = _heading_insets(title_ph)
    lh = typo.line_height
    inner = lambda w: max((w - insets[0] - insets[2]) / EMU_PER_PT, 1.0)  # noqa: E731
    # the heading keeps the template's look — capitals and letter-spacing it inherits are measured as set (C5, T14)
    caps, spc = _inherited_caps_spc(title_ph)
    caps = caps or bool(ds is not None and getattr(ds, "caps", False))
    look = {"caps": caps, "spc_pt": spc} if (caps or spc) else {}
    fit_size = functools.partial(_fit_size, **look) if look else _fit_size  # noqa: N806
    wrap_lines = functools.partial(_wrap_lines, **look) if look else _wrap_lines  # noqa: N806
    # (the right margin mirrors the heading's left edge — the page margin's when the heading stands in a column past art)
    safe_x = int(safe.x * W)
    right = min(int(safe.x2 * W), W - (box.x if box.x < 0.25 * W else safe_x))
    align_ctr = _title_align(title_ph) == "ctr"
    centred = align_ctr and box.x > safe_x + int(0.04 * W)
    w_free = _free_width(slide, box, W, H, right, skip=(title_ph._element,))
    # marks baked into the background picture (a bookmark ribbon, a stamp) stand in the band too: the analysis found
    # them as visual chrome — the heading stops before them
    for c in manifest.tokens.chrome:
        if c.source != "background":
            continue
        b = c.bbox.to_emu(W, H)
        if b.w * b.h < 0.002 * W * H or b.w * b.h >= 0.6 * W * H:
            continue
        if b.y < box.y2 and b.y2 > box.y and b.x >= box.x + int(0.3 * W):
            w_free = min(w_free, b.x - int(0.015 * W) - box.x)
    w_full = min(w_free, int(0.86 * W))
    content_left: Optional[int] = None
    if centred:
        # a centred heading grows about its centre, as far as the free room on both sides allows — never only to the
        # right (a centred line would end off-centre)
        cx = box.x + box.w // 2
        lim_r = box.x + _free_width(slide, box, W, H, int(safe.x2 * W), skip=(title_ph._element,))
        lim_l = _free_left(slide, box, W, H, safe_x, skip=(title_ph._element,))
        half = min(cx - lim_l, lim_r - cx, int(0.43 * W))
        if 2 * half > box.w:
            box = Bbox(x=cx - half, y=box.y, w=2 * half, h=box.h)
        w_full = box.w
    elif w_full != box.w and w_full > 0.3 * W:
        box = Bbox(x=box.x, y=box.y, w=w_full, h=box.h)
    # one heading size for the whole deck (never autofitted per slide unless the headline cannot fit at all)
    h1 = ds.head_size if ds is not None else max(h1, _snap_heading(0.046 * H / EMU_PER_PT, scale))
    head_text = typeset(oslide.headline)
    # something the canvas keeps stands over the left part of the heading's box (a big page number «09» beside the
    # title): where the heading's letters would start under it, the heading starts after it (a centred line that
    # clears it, or an indent the template gives its heading to clear it, stays as it is)
    past = _left_clear(slide, box, W, H, skip=(title_ph._element,))
    keep_indent = False
    if past > box.x:
        indent = _inherited_indent(title_ph)
        text_left = box.x + insets[0] + indent
        if align_ctr:
            iw = inner(box.w) - indent / EMU_PER_PT
            widest = max((text_width_pt(ln, title_font, h1, title_bold, **look) for ln in wrap_lines(head_text, title_font, h1, title_bold, iw)), default=0.0)
            text_left += int(max(0.0, (iw - widest) / 2) * EMU_PER_PT)
        if text_left >= past - int(0.015 * W):
            keep_indent = indent > 0  # the template indents its heading to clear the mark: the indent stays
        elif box.x2 - past >= int(0.5 * W):
            # the box starts after the mark and keeps the free width up to the page's right margin (a line that
            # filled the old box still fits on one line)
            w2 = min(_free_width(slide, Bbox(x=past, y=box.y, w=box.x2 - past, h=box.h), W, H, right, skip=(title_ph._element,)), int(0.86 * W))
            box = Bbox(x=past, y=box.y, w=max(box.x2 - past, w2), h=box.h)
            w_full = box.w
            warnings.append("heading moved past a mark at its left")
    color: Optional[str] = None  # None: the template's own heading colour (inherited)
    band = _heading_band(slide, box, W, H)
    if band is not None:
        # a heading printed on a band of the layout keeps inside the band: it takes the band's whole free width and
        # steps down the scale (to 55 %) before it spills out of the band — out of its end too: a short tab at the
        # left edge (LibreOffice «Freshes»: a teal tab of a third of the width under white type) holds the heading's
        # letters only as far as it reaches; past its end they would stand white on the white page
        wide = box
        bx, bx2 = max(box.x, band.x), min(box.x2, band.x2 - int(0.01 * W))
        if bx2 < box.x2 - int(0.01 * W) and bx2 - bx >= int(0.15 * W):
            box = Bbox(x=bx, y=box.y, w=bx2 - bx, h=box.h)
        inner_h = band.y2 - box.y - int(0.004 * H)
        fitted = fit_size([head_text], Bbox(x=box.x, y=box.y, w=box.w, h=max(inner_h, 1)), title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=lh, min_ratio=0.55, ladder_only=sparse) if inner_h > 0 else None
        if fitted is not None and fitted.fits:
            text_h = int(fitted.height_pt * EMU_PER_PT) + insets[1] + insets[3]
            h_box = max(min(box.h, inner_h), text_h)
            _fill_heading(title_ph, head_text, fitted.size_pt, None, insets, W, keep_indent=keep_indent)
            _caps_heading(title_ph, ds)
            _set_box(title_ph, Bbox(x=box.x, y=box.y, w=box.w, h=h_box))
            _check_heading_color(slide, title_ph, canvas, band, pal, manifest)
            return max(band.y2, box.y + h_box), _heading_left(title_ph, box, insets), False
        # the headline does not fit the band at any readable size: it is set under the band, on the slide's ground —
        # from the page's margin across the free width (a tab at the slide's edge is no column)
        ux = max(wide.x, safe_x)
        box = Bbox(x=ux, y=band.y2 + int(0.02 * H), w=max(wide.x2 - ux, int(0.3 * W)), h=int(0.16 * H))
        w_full = max(w_full, box.w)
        color = _heading_on(pal.bg, manifest, pal)
        warnings.append("heading set under its band")
    n = len(wrap_lines(head_text, title_font, h1, title_bold, inner(box.w)))
    # the deck's heading size holds: a long headline takes up to three lines (the box grows downwards and the
    # content moves with it); only past three lines does the size give way
    rows = min(max(n, 2), 3)
    need = int((rows + 0.1) * h1 * lh * EMU_PER_PT) + insets[1] + insets[3]
    rule = _rule_under(slide, box, W, H)
    if rule is not None:
        room = rule - box.y - int(0.01 * H)  # the heading keeps above the layout's rule
        one = int(1.1 * h1 * lh * EMU_PER_PT) + insets[1] + insets[3]
        if box.h > room:
            box = Bbox(x=box.x, y=box.y, w=box.w, h=max(room, min(one, box.h)))  # the sample's box itself crosses the rule
        need = min(need, max(room, box.h))
    if box.h < need:
        box = Bbox(x=box.x, y=box.y, w=box.w, h=need)
    min_ratio = 0.6 if rule is not None else 0.8
    res = fit_size([head_text], box, title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=lh, min_ratio=min_ratio, ladder_only=sparse)
    if res.broken_word:
        # a word wider than the box: the box takes the whole free width first, then the size may give way to 55 %
        if not centred and w_full > box.w:
            box = Bbox(x=box.x, y=box.y, w=w_full, h=box.h)
            res = fit_size([head_text], box, title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=lh, min_ratio=min_ratio, ladder_only=sparse)
        if res.broken_word:
            res = fit_size([head_text], box, title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=lh, min_ratio=0.55, ladder_only=sparse)
        if res.broken_word:
            warnings.append("heading word broken")
    limit = int(0.42 * H)
    if box.y + int(res.height_pt * EMU_PER_PT) + insets[1] + insets[3] > limit and box.y < limit - int(0.06 * H):
        # the heading would take the upper half of the slide: once more at the full free width, a size or two smaller
        wide = Bbox(x=box.x, y=box.y, w=max(box.w, w_full), h=limit - box.y)
        again = fit_size([head_text], wide, title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=lh, min_ratio=0.6, ladder_only=sparse)
        if again.fits and not again.broken_word:
            box, res = Bbox(x=wide.x, y=wide.y, w=wide.w, h=max(box.h if box.h <= wide.h else wide.h, int(again.height_pt * EMU_PER_PT) + insets[1] + insets[3])), again
            warnings.append("heading refitted to keep the upper half free")
    size = res.size_pt
    lines = res.lines
    grid_box = box  # the heading's column (a centred heading's narrowed box does not move the content column)
    if 2 <= lines <= 3 and not res.broken_word:
        # balanced lines at the final size: the narrowest box that keeps the same number of lines (no one-word last line)
        lo, hi = int(box.w * 0.55), box.w
        for _ in range(12):
            mid = (lo + hi) // 2
            if len(wrap_lines(head_text, title_font, size, title_bold, inner(mid))) <= lines:
                hi = mid
            else:
                lo = mid
        balanced = min(box.w, int(hi * _balance_slack(title_font)))
        if balanced < box.w:
            if align_ctr:
                # centred lines keep their centre when the box narrows (a left-anchored narrowing moves them left)
                box = Bbox(x=box.x + (box.w - balanced) // 2, y=box.y, w=balanced, h=box.h)
            else:
                box = Bbox(x=box.x, y=box.y, w=balanced, h=box.h)
    _fill_heading(title_ph, head_text, size, color, insets, W, keep_indent=keep_indent)
    _caps_heading(title_ph, ds)
    text_h = int(res.height_pt * EMU_PER_PT) + insets[1] + insets[3]
    if text_h > box.h:
        _set_box(title_ph, Bbox(x=box.x, y=box.y, w=box.w, h=text_h))
        bottom = box.y + text_h
    else:
        # the box shrinks to its text where the text is drawn (top / middle / bottom of the sample's box)
        anchor = _title_anchor(title_ph)
        y = box.y if anchor == "t" else (box.y2 - text_h if anchor == "b" else box.y + (box.h - text_h) // 2)
        _set_box(title_ph, Bbox(x=box.x, y=y, w=box.w, h=text_h))
        bottom = y + text_h
    if color is None:
        _check_heading_color(slide, title_ph, canvas, None, pal, manifest)
    content_left = None if centred else _heading_left(title_ph, grid_box if align_ctr else box, insets)
    return bottom, content_left, False


def _inherited_caps_spc(shape) -> tuple[bool, float]:
    """(capitals, letter-spacing pt) a heading inherits (contract C5 textfill.inherited_caps_spc); none when unknown."""
    try:
        from verstka.rendering.textfill import inherited_caps_spc

        caps, spc = inherited_caps_spc(shape)
        return bool(caps), float(spc or 0.0)
    except Exception:  # noqa: BLE001
        return False, 0.0


def _heading_left(title_ph, box: Bbox, insets: tuple[int, int, int, int]) -> int:
    """Where the heading's letters start: the box plus its left inset (its own, else the effective one)."""
    left = box.x + insets[0]
    lins = title_ph._element.find(".//" + q("a:bodyPr"))
    if lins is not None and lins.get("lIns") is not None:
        try:
            left = box.x + int(lins.get("lIns"))
        except ValueError:
            pass
    return left


def _balance_slack(font: Optional[str]) -> float:
    """Slack of the balanced heading width: a face measured with its own metrics needs 3 %, a face measured as Play ×
    a factor needs 8 % (its line may be wider than estimated)."""
    try:
        from verstka.rendering.fonts import is_measured
    except ImportError:
        return 1.03
    try:
        return 1.03 if is_measured(font) else 1.08
    except Exception:  # noqa: BLE001
        return 1.03


def _set_box(shape, box: Bbox) -> None:
    """Every geometry write is a four-value write: a lone width or top written on a placeholder that inherits its
    geometry makes python-pptx write an offset of 0 (the heading jumps to the slide's left edge)."""
    shape.left, shape.top, shape.width, shape.height = Emu(int(box.x)), Emu(int(box.y)), Emu(max(int(box.w), 0)), Emu(max(int(box.h), 0))


def _materialize(shape) -> None:
    """The shape's inherited geometry written as its own xfrm (contract C1: deck.materialize_xfrm when it exists)."""
    try:
        from verstka.rendering.deck import materialize_xfrm
    except ImportError:
        materialize_xfrm = None
    if materialize_xfrm is not None:
        try:
            materialize_xfrm(shape)
            return
        except Exception:  # noqa: BLE001
            pass
    spPr = shape._element.find(q("p:spPr"))
    if spPr is None or spPr.find(q("a:xfrm")) is not None:
        return
    try:
        x, y, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:  # noqa: BLE001
        return
    if None in (x, y, w, h):
        return
    xf = etree.Element(q("a:xfrm"))
    off = etree.SubElement(xf, q("a:off"))
    off.set("x", str(int(x)))
    off.set("y", str(int(y)))
    ext = etree.SubElement(xf, q("a:ext"))
    ext.set("cx", str(int(w)))
    ext.set("cy", str(int(h)))
    spPr.insert(0, xf)


def _ph_chain(shape) -> list:
    """The shape's element, then its layout and master placeholders (the inheritance chain of a placeholder)."""
    out = [shape._element]
    cur = shape
    for _ in range(2):
        try:
            cur = cur._base_placeholder
        except Exception:  # noqa: BLE001
            cur = None
        if cur is None:
            break
        out.append(cur._element)
    return out


def _master_title_lvl1(shape):
    try:
        master = shape.part.slide.slide_layout.slide_master
    except Exception:  # noqa: BLE001
        try:
            master = shape.part.slide_layout.slide_master
        except Exception:  # noqa: BLE001
            return None
    ts = master._element.find(q("p:txStyles") + "/" + q("p:titleStyle"))
    return ts.find(q("a:lvl1pPr")) if ts is not None else None


def _heading_insets(shape) -> tuple[int, int, int, int]:
    """Insets the heading is measured with: the effective ones of the shape (its own bodyPr, then its layout and
    master placeholders — contract C5 textfill.effective_insets when it exists), never less than PowerPoint's
    defaults (a wordmark placeholder may inherit a large inset that the default measure does not see)."""
    default = (91440, 45720, 91440, 45720)
    eff = None
    try:
        from verstka.rendering.textfill import effective_insets

        eff = effective_insets(shape)
    except Exception:  # noqa: BLE001
        eff = None
    if eff is None:
        vals: dict[str, int] = {}
        for el in _ph_chain(shape):
            bp = el.find(".//" + q("a:bodyPr"))
            if bp is None:
                continue
            for k in ("lIns", "tIns", "rIns", "bIns"):
                if k not in vals and bp.get(k) is not None:
                    try:
                        vals[k] = int(bp.get(k))
                    except ValueError:
                        pass
        eff = tuple(vals.get(k, d) for k, d in zip(("lIns", "tIns", "rIns", "bIns"), default))
    return tuple(max(d, int(e)) for d, e in zip(default, eff))  # type: ignore[return-value]


def _title_align(title_shape) -> str:
    """Effective horizontal alignment of a heading's first paragraph: its own pPr / lstStyle, then the layout and
    master placeholders, then the master's title style."""
    for el in _ph_chain(title_shape):
        tx = el.find(q("p:txBody"))
        if tx is None:
            continue
        p0 = tx.find(q("a:p"))
        ppr = p0.find(q("a:pPr")) if p0 is not None else None
        if ppr is not None and ppr.get("algn"):
            return ppr.get("algn")
        lvl = tx.find(q("a:lstStyle") + "/" + q("a:lvl1pPr"))
        if lvl is not None and lvl.get("algn"):
            return lvl.get("algn")
    lvl = _master_title_lvl1(title_shape)
    return (lvl.get("algn") if lvl is not None else None) or "l"


def _inherited_indent(title_shape) -> int:
    """The left indent (marL + indent) a heading's first level inherits, in EMU."""
    for el in _ph_chain(title_shape):
        tx = el.find(q("p:txBody"))
        if tx is None:
            continue
        for node in [n for n in [tx.find(q("a:p") + "/" + q("a:pPr")), tx.find(q("a:lstStyle") + "/" + q("a:lvl1pPr"))] if n is not None]:
            if node.get("marL") is not None:
                try:
                    return int(node.get("marL")) + int(node.get("indent") or 0)
                except ValueError:
                    return 0
    lvl = _master_title_lvl1(title_shape)
    if lvl is not None and lvl.get("marL") is not None:
        try:
            return int(lvl.get("marL")) + int(lvl.get("indent") or 0)
        except ValueError:
            return 0
    return 0


_LOOK_ATTRS = ("cap", "spc", "baseline")


def _fill_keep(el, specs, **kw) -> None:
    """fill_text that keeps the heading's own run look: capitals / letter-spacing / baseline written on the sample's
    run (cap="none" spc="0" over a master that tracks its text) stay on the new runs — never dropped to inheritance."""
    look: dict = {}
    for tag in ("a:rPr", "a:endParaRPr"):  # the sample run's own attributes first, then its paragraph end's
        src = next(iter(el.iter(q(tag))), None)
        for a in _LOOK_ATTRS:
            if src is not None and src.get(a) is not None and a not in look:
                look[a] = src.get(a)
    fill_text(el, specs, **kw)
    if look:
        for r in el.iter(q("a:rPr")):
            for a, v in look.items():
                if r.get(a) is None:
                    r.set(a, v)


def _fill_heading(title_ph, text: str, size: float, color: Optional[str], insets, W: int, keep_indent: bool = False) -> None:
    """Write the headline into the title shape at `size` (and `color` when given). A heading that inherits a large
    indent (a wordmark placeholder: 1.27″) is set flush with its box — the content aligns to its letters."""
    reset = not keep_indent and _inherited_indent(title_ph) > 0.01 * W
    spec = ParagraphSpec(text, color_hex=color) if color else ParagraphSpec(text)
    done = False
    if reset:
        try:
            _fill_keep(title_ph._element, [spec], size_pt=size, reset_indent=True)
            done = True
        except TypeError:
            done = False
    if not done:
        _fill_keep(title_ph._element, [spec], size_pt=size)
        if reset:
            for p_ in title_ph._element.iter(q("a:p")):
                ppr = p_.find(q("a:pPr"))
                if ppr is None:
                    ppr = etree.Element(q("a:pPr"))
                    p_.insert(0, ppr)
                ppr.set("marL", "0")
                ppr.set("indent", "0")


def _caps_heading(title_ph, ds) -> None:
    """Capitals on the heading's runs when the deck sets its headings so (`DeckStyle.caps`): the words stay as written
    in the file, the capitals are the template's voice."""
    if ds is None or not getattr(ds, "caps", False):
        return
    for r in title_ph._element.iter(q("a:rPr")):
        r.set("cap", "all")


def _heading_band(slide: Slide, box: Bbox, W: int, H: int) -> Optional[Bbox]:
    """A band the layout (or master, or the slide) paints under the heading (contract C2: layers.heading_band)."""
    try:
        from verstka.rendering.layers import heading_band
    except ImportError:
        return None
    try:
        got = heading_band(slide, box)
    except Exception:  # noqa: BLE001
        return None
    if got is None:
        return None
    b = getattr(got, "box", got)
    return b if isinstance(b, Bbox) else None


def _heading_ink(title_ph, box: Bbox, manifest: TemplateManifest) -> Bbox:
    """Where the heading's letters stand inside its box: its lines as set (the size of its runs, the box's width less
    the insets), anchored and aligned as the box says. What lies under the rest of the box (a vapour trail across the
    top of Office «Vapor Trail») is not under the heading."""
    from verstka.rendering.fonts import text_width_pt, wrap_lines

    text = "".join(t.text or "" for t in title_ph._element.iter(q("a:t"))).strip()
    if not text:
        return box
    size = _set_size(title_ph, 28.0)
    rpr = next(iter(title_ph._element.iter(q("a:rPr"))), None)
    bold = rpr is not None and rpr.get("b") in ("1", "true")
    lat = rpr.find(q("a:latin")) if rpr is not None else None
    font = (lat.get("typeface") if lat is not None else None) or manifest.tokens.typography.primary_family
    if font and font.startswith("+"):
        font = manifest.tokens.typography.primary_family
    insets = _heading_insets(title_ph)
    inner_w = max(box.w - insets[0] - insets[2], 1)
    lines = wrap_lines(text, font, size, bold, inner_w / EMU_PER_PT) or [text]
    lh = int(size * manifest.tokens.typography.line_height * EMU_PER_PT)
    text_h = min(len(lines) * lh, box.h)
    body = title_ph._element.find(".//" + q("a:bodyPr"))
    anchor = (body.get("anchor") if body is not None else None) or "t"
    if anchor == "ctr":
        y = box.y + (box.h - text_h) // 2
    elif anchor == "b":
        y = box.y2 - insets[3] - text_h
    else:
        y = box.y + insets[1]
    widest = min(int(max(text_width_pt(ln, font, size, bold) for ln in lines) * EMU_PER_PT), inner_w)
    align = _title_align(title_ph)
    if align == "r":
        x = box.x2 - insets[2] - widest
    elif align == "ctr":
        x = box.x + (box.w - widest) // 2
    else:
        x = box.x + insets[0]
    return Bbox(x=max(x, box.x), y=max(y, box.y), w=max(widest, 1), h=max(text_h, 1))


def _check_heading_color(slide: Slide, title_ph, canvas: Optional[Pattern], band: Optional[Bbox], pal: "_Palette", manifest: TemplateManifest) -> None:
    """The template's heading colour must read on what really stands under the heading (a band or panel painted by
    the layout, a picture ground): when it does not (< 3:1), the heading takes a readable colour."""
    t = _top_title(canvas) if canvas is not None else None
    col = t.style.color_hex if t is not None else None
    if not col:
        # no sample heading to learn the colour from (a template whose slides are empty, a slide without a canvas):
        # the colour the heading inherits from its layout and master
        from verstka.rendering.textfill import inherited_color

        col = inherited_color(title_ph)
    if not col:
        return
    try:
        from verstka.rendering.charts import slide_ground

        b = element_bbox(title_ph._element)
        where = band if band is not None else (_heading_ink(title_ph, Bbox(x=b[0], y=b[1], w=b[2], h=b[3]), manifest) if b else None)
        ground = slide_ground(slide, where) if where is not None else None
    except Exception:  # noqa: BLE001
        ground = None
    if not ground or contrast_ratio(col, ground) >= 3.0:
        return
    new = _readable(ground, [col] + _inks(pal, manifest))
    for r in title_ph._element.iter(q("a:rPr")):
        for old in [c for c in r if etree.QName(c).localname in ("solidFill", "gradFill", "noFill", "pattFill", "blipFill", "grpFill")]:
            r.remove(old)
        sf = etree.Element(q("a:solidFill"))
        etree.SubElement(sf, q("a:srgbClr")).set("val", new)
        ln = r.find(q("a:ln"))
        if ln is not None:
            ln.addnext(sf)
        else:
            r.insert(0, sf)


def _left_clear(slide: Slide, box: Bbox, W: int, H: int, skip: tuple = ()) -> int:
    """The x a heading box must start at to clear a mark that stands over its own left part, in its band, starting
    before box.x + 10 % of the width and reaching past box.x: the text of a page number «09» kept beside the title
    (letters over letters), or a small ornament of the layout — a cluster of dots, a small picture, 0.3–3 % of the
    slide (B4: Grey Elegant's dots under the first letters). Larger art (triangles, illustrations) is the canvas's
    and the free-width rules' concern; a heading on a plate of its own (a filled title box) covers the art behind it.
    box.x when nothing."""
    x = box.x
    plate = any(etree.QName(el).localname == "sp" and _paints(el) for el in skip)
    for el in _drawn(slide):
        tag = etree.QName(el).localname
        if el in skip or tag not in ("sp", "pic", "grpSp"):
            continue
        has_text = bool("".join(t.text or "" for t in el.iter(q("a:t"))).strip())
        if plate and not has_text:
            continue
        if not has_text and _veil(el):
            continue  # a see-through bubble (LibreOffice «Lights») is no mark the heading must clear
        b = element_bbox(el)
        if not b or b[2] * b[3] >= 0.25 * W * H or b[2] <= 0 or b[3] <= 0:
            continue
        if not has_text:
            if tag == "sp" and not _paints(el):
                continue  # an empty frame is no mark
            if not (0.003 * W * H <= b[2] * b[3] <= 0.03 * W * H):
                continue
        top, bot = max(b[1], box.y), min(b[1] + b[3], box.y2)
        if bot - top < 0.5 * min(b[3], box.h):
            continue  # not in the heading's band
        if b[0] < box.x + int(0.1 * W) and b[0] + b[2] > box.x and b[0] + b[2] < box.x + int(0.3 * W):
            # the caller keeps a heading whose letters start 0.015 W before this x (a text mark's own margin): an
            # ornament keeps 0.012 W of air between its edge and the first letter
            right = _ink_right(el, b) if has_text else b[0] + b[2]
            x = max(x, right + int((0.015 if has_text else 0.027) * W))
    return x


def _ink_right(el, b) -> int:
    """Where a shape really ends on the right: a left-aligned unfilled text box («09») ends where its text does, not at
    its frame; anything else at its frame."""
    right = b[0] + b[2]
    if etree.QName(el).localname != "sp" or _paints(el):
        return right
    tx = el.find(q("p:txBody"))
    if tx is None:
        return right
    p0 = tx.find(q("a:p"))
    ppr = p0.find(q("a:pPr")) if p0 is not None else None
    if ppr is not None and (ppr.get("algn") or "l") not in ("l", "just"):
        return right
    text = "".join(t.text or "" for t in tx.iter(q("a:t"))).strip()
    rpr = next((r for r in tx.iter(q("a:rPr")) if r.get("sz")), None)
    if not text or rpr is None or "\n" in text:
        return right
    bp = tx.find(q("a:bodyPr"))
    lins = int(bp.get("lIns")) if bp is not None and (bp.get("lIns") or "").isdigit() else 91440
    lat = rpr.find(q("a:latin"))
    w = text_width_pt(text, lat.get("typeface") if lat is not None else None, int(rpr.get("sz")) / 100, rpr.get("b") == "1")
    return min(right, b[0] + lins + int(w * 1.08 * EMU_PER_PT))


def _free_left(slide: Slide, box: Bbox, W: int, H: int, left: int, skip: tuple = ()) -> int:
    """The leftmost x a heading may reach from its box leftwards: the page margin, cut after whatever stands in its
    band on the left (a picture, a panel of the layout)."""
    limit = left
    for el in _drawn(slide):
        if el in skip or etree.QName(el).localname not in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp"):
            continue
        b = element_bbox(el)
        if not b or b[2] * b[3] >= 0.6 * W * H:
            continue
        if b[1] < box.y2 and b[1] + b[3] > box.y and b[0] + b[2] <= box.x2 - int(0.1 * W) and b[0] + b[2] <= box.x + box.w // 2:
            limit = max(limit, b[0] + b[2] + int(0.015 * W))
    return min(limit, box.x)


def _free_width(slide: Slide, box: Bbox, W: int, H: int, right: int, skip: tuple = ()) -> int:
    """Width a heading may take from box.x: up to `right`, cut before whatever really stands in its band on this
    slide — shapes of the cleaned canvas and of its layout (logos, patterns, panels), not the chrome of other layouts."""
    limit = right
    items = _drawn(slide)
    for el in items:
        if el in skip or etree.QName(el).localname not in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp"):
            continue
        b = element_bbox(el)
        if not b or b[2] * b[3] >= 0.6 * W * H:
            continue
        if b[1] < box.y2 and b[1] + b[3] > box.y and b[0] >= box.x + int(0.1 * W):
            limit = min(limit, b[0] - int(0.015 * W))
    return max(limit - box.x, int(0.25 * W))


def _foot_top(slide: Slide, W: int, H: int) -> Optional[int]:
    """The top of the brand marks drawn at the foot of this slide (logos, a footer line, a page number): small
    things in the bottom fifth, as the slide, its layout and its master draw them."""
    tops = []
    for el in _drawn(slide):
        if etree.QName(el).localname not in ("sp", "pic", "grpSp", "cxnSp"):
            continue
        b = element_bbox(el)
        if not b or b[2] <= 0 or b[3] <= 0 or b[2] * b[3] >= 0.2 * W * H or b[3] > 0.15 * H:
            continue
        if b[1] >= 0.8 * H and b[1] + b[3] <= H * 1.01:
            tops.append(b[1])
    return min(tops) if tops else None


def _drawn(slide: Slide) -> list:
    """What is drawn on a slide besides its own shapes: the decoration of its layout and master (placeholders of
    the layout and master are not drawn on a slide that does not use them)."""
    items = list(slide._element.cSld.find(q("p:spTree")))
    try:
        for holder in (slide.slide_layout, slide.slide_layout.slide_master):
            items += [el for el in holder._element.cSld.find(q("p:spTree")) if el.find(".//" + q("p:ph")) is None]
    except Exception:  # noqa: BLE001
        pass
    return items


def _rule_under(slide: Slide, box: Bbox, W: int, H: int) -> Optional[int]:
    """The top of a rule drawn across the slide just under a heading box (the classic «title + line» layout)."""
    best = None
    for el in _drawn(slide):
        if etree.QName(el).localname not in ("sp", "cxnSp"):
            continue
        b = element_bbox(el)
        if not b or b[3] > 0.02 * H or b[2] < 0.3 * W:
            continue
        if box.y + int(0.03 * H) <= b[1] <= box.y2 + int(0.12 * H) and (best is None or b[1] < best):
            best = b[1]
    return best


def _below_rules(slide: Slide, heading_bottom: int, top: int, W: int, H: int, gap: int) -> int:
    """The content starts below a rule the layout draws under the heading (a line across the slide), never on it."""
    items = _drawn(slide)
    out = top
    for el in items:
        if etree.QName(el).localname not in ("sp", "cxnSp", "pic"):
            continue
        b = element_bbox(el)
        if not b or b[3] > 0.02 * H or b[2] < 0.3 * W:
            continue
        if heading_bottom - int(0.03 * H) <= b[1] <= top + int(0.1 * H):
            out = max(out, b[1] + b[3] + int(gap * 0.7))
    return out


def _short_kicker(tag: str, limit: int = 26) -> str:
    """A label is a few words: cut at the first colon or comma, then at a word boundary."""
    t = tag.strip()
    if len(t) <= limit:
        return t
    for sep in (":", ",", " —", " –", "("):
        if sep in t and 3 <= len(t.split(sep)[0].strip()) <= limit:
            return t.split(sep)[0].strip()
    out = ""
    for w in t.split():
        if len((out + " " + w).strip()) > limit:
            break
        out = (out + " " + w).strip()
    return out or t[:limit]


def _lede_of(oslide: OutlineSlide) -> Optional[str]:
    """The one sentence that introduces structured content (the composer's intro line)."""
    c = oslide.content
    if oslide.subtitle:
        return oslide.subtitle
    structured = c.items or c.columns or c.table is not None or c.chart is not None or c.numbers
    if structured and len(c.paragraphs) == 1 and len(c.paragraphs[0]) <= 200 and not c.bullets:
        return c.paragraphs[0]
    return None


def _without_lede(oslide: OutlineSlide) -> OutlineSlide:
    o = oslide.model_copy(deep=True)
    if o.subtitle:
        o.subtitle = None
    elif len(o.content.paragraphs) == 1:
        o.content.paragraphs = []
    return o


def _section_of(oslide: OutlineSlide, outline: DeckOutline) -> str:
    """The slide's section, else the section of its neighbours (the first slide of a section often carries none)."""
    own = (oslide.section or "").strip()
    if own:
        return own
    ids = [s.id for s in outline.slides]
    if oslide.id not in ids:
        return ""
    i = ids.index(oslide.id)
    for j in (i - 1, i + 1):
        if 0 <= j < len(outline.slides) and outline.slides[j].kind not in _COVER_KINDS and (outline.slides[j].section or "").strip():
            return outline.slides[j].section.strip()
    return ""


def _same_words(tag: str, headline: str) -> bool:
    a, b = tag.lower().strip(" .:"), headline.lower().strip(" .:")
    return a == b or b.startswith(a) or a.startswith(b)


def _hug_label(back_el, title_el, label: str, font: Optional[str], size: float, bold: bool, insets: tuple[int, int, int, int], W: int) -> None:
    """The label hugs its text: its width is the text plus the sample's side padding, never wider than it was."""
    b = element_bbox(back_el)
    tb = element_bbox(title_el)
    if not b or not tb:
        return
    pad = max(tb[0] - b[0], 0) + insets[0]
    need = int(text_width_pt(label, font, size, bold) * EMU_PER_PT * 1.06) + pad * 2
    need = max(need, int(0.1 * W))
    if need < b[2]:
        from verstka.rendering.deck import set_element_pos

        set_element_pos(back_el, w=need)
        set_element_pos(title_el, w=max(need - (tb[0] - b[0]), int(0.05 * W)))


def _heading_on(ground: str, manifest: TemplateManifest, pal: "_Palette") -> str:
    """A heading colour that reads on the ground: the brand's dark or light colour before plain black/white."""
    t = manifest.tokens
    cands = [t.color_for("background.dark"), t.color_for("text.primary"), pal.text, t.color_for("background.light"), "000000", "FFFFFF"]
    return next((c for c in cands if c and contrast_ratio(c, ground) >= 4.5), pal.text)


def _snap_heading(target: float, scale: list[float]) -> float:
    near = [s for s in scale if abs(s - target) <= 0.22 * target]
    return min(near, key=lambda s: abs(s - target)) if near else round(target)


def _art_painted(slide: Slide, box: Bbox, skip: tuple = ()) -> float:
    """Share of `box` the template's art really paints (a triangle by its outline, not its bounding box)."""
    if box.w <= 0 or box.h <= 0:
        return 0.0
    try:
        from verstka.rendering.layers import art_layers, painted_share

        return min(1.0, sum(painted_share(lay, box) for lay in art_layers(slide, skip=skip)))
    except Exception:  # noqa: BLE001
        return 0.0


def _clear_rect(slide: Slide, area: Bbox, skip: tuple = (), n: int = 32) -> Optional[Bbox]:
    """The largest rectangle inside `area` that the template's art does not paint — measured by the art's real
    outline (a triangle's empty corner is free), on an n×n grid of the area, a cell's margin kept. None when unknown."""
    if area.w <= 0 or area.h <= 0:
        return None
    try:
        from verstka.rendering import layers as L

        arts = []
        for lay in L.art_layers(slide, skip=skip):
            if lay.box.intersection(area) <= 0:
                continue
            arts.append((lay.box, L._outline_of(lay)))
    except Exception:  # noqa: BLE001
        return None
    cw, ch = area.w / n, area.h / n

    def painted(px: float, py: float) -> bool:
        for box, polys in arts:
            if not (box.x <= px <= box.x2 and box.y <= py <= box.y2):
                continue
            if polys is None or sum(1 for poly in polys if len(poly) > 2 and L._inside(px, py, poly)) % 2 == 1:
                return True
        return False

    # a cell is busy when its centre or a corner is painted (the rectangle keeps a hair off the art's edge)
    busy = [[any(painted(area.x + (i + dx) * cw, area.y + (j + dy) * ch) for dx, dy in ((0.5, 0.5), (0, 0), (1, 0), (0, 1), (1, 1))) for i in range(n)] for j in range(n)]
    return _largest_empty(busy, area, n)


def _panel_rect(slide: Slide, area: Bbox, skip: tuple = (), n: int = 32) -> Optional[Bbox]:
    """The largest rectangle inside `area` that stands wholly on one shape of the template's art (a coloured triangle
    of a mosaic: text set on it reads as on a panel), a cell's margin kept from its edge. None when there is none."""
    if area.w <= 0 or area.h <= 0:
        return None
    try:
        from verstka.rendering import layers as L

        arts = [(lay.box, L._outline_of(lay)) for lay in L.art_layers(slide, skip=skip) if lay.box.intersection(area) > 0 and not lay.picture]
    except Exception:  # noqa: BLE001
        return None
    cw, ch = area.w / n, area.h / n
    best: Optional[Bbox] = None
    for k_, (box, polys) in enumerate(arts):
        def inside(px: float, py: float, box=box, polys=polys) -> bool:
            if not (box.x <= px <= box.x2 and box.y <= py <= box.y2):
                return False
            return polys is None or sum(1 for poly in polys if len(poly) > 2 and L._inside(px, py, poly)) % 2 == 1

        def other(px: float, py: float) -> bool:  # another art shape drawn over this one there
            for b2, p2 in arts[k_ + 1:]:
                if b2.x <= px <= b2.x2 and b2.y <= py <= b2.y2 and (p2 is None or sum(1 for poly in p2 if len(poly) > 2 and L._inside(px, py, poly)) % 2 == 1):
                    return True
            return False

        pts = ((0.5, 0.5), (0, 0), (1, 0), (0, 1), (1, 1))
        busy = [[not all(inside(area.x + (i + dx) * cw, area.y + (j + dy) * ch) and not other(area.x + (i + dx) * cw, area.y + (j + dy) * ch) for dx, dy in pts) for i in range(n)] for j in range(n)]
        got = _largest_empty(busy, area, n)
        if got is not None and (best is None or got.area > best.area):
            best = got
    return best


def _largest_empty(busy: list, area: Bbox, n: int) -> Optional[Bbox]:
    """The largest all-free rectangle of an n×n cell grid over `area` (histogram method), in EMU."""
    cw, ch = area.w / n, area.h / n
    best = (0, None)
    heights = [0] * n
    for j in range(n):
        for i in range(n):
            heights[i] = 0 if busy[j][i] else heights[i] + 1
        stack: list[int] = []
        for i in range(n + 1):
            hgt = heights[i] if i < n else 0
            while stack and heights[stack[-1]] >= hgt:
                top_i = stack.pop()
                left = stack[-1] + 1 if stack else 0
                a = heights[top_i] * (i - left)
                if a > best[0]:
                    best = (a, (left, j - heights[top_i] + 1, i - left, heights[top_i]))
            stack.append(i)
    if best[1] is None:
        return None
    x0, y0, wc, hc = best[1]
    return Bbox(x=int(area.x + x0 * cw), y=int(area.y + y0 * ch), w=int(wc * cw), h=int(hc * ch))


def _cover_lines_region(slide: Slide, canvas: Optional[Pattern], title_ph, under: Bbox, safe: Bbox, W: int, H: int) -> Optional[tuple[Bbox, str]]:
    """Where a cover's subtitle, goal line and footnote go when the column under its title stands on the template's
    art (a mosaic of triangles, an illustration): the clear part of the sample's own subtitle box, else the largest
    clear rectangle under the title (by the art's real outline). (region, alignment); None when the column is clear
    (kept as it is) or no clear room is large enough."""
    skip = (title_ph._element,) if title_ph is not None else ()
    if under.h < 0.08 * H or _art_painted(slide, under, skip) < 0.05:
        return None
    top = under.y
    sub = next((s_ for s_ in canvas.slots if s_.role == SlotRole.subtitle), None) if canvas is not None else None
    if sub is not None:
        b = sub.bbox.to_emu(W, H)
        y0 = max(b.y, top)
        box = Bbox(x=b.x, y=y0, w=b.w, h=max(b.y2 - y0, 0))
        clear = _clear_rect(slide, box, skip) if box.h > 0 else None
        if clear is not None and clear.w >= 0.18 * W and clear.h >= 0.15 * H:
            return clear, (sub.style.align or "l")
    below = Bbox(x=safe.x, y=top, w=safe.w, h=max(safe.y2 - top, 1))
    clear = _clear_rect(slide, below, skip)
    if clear is not None and clear.w >= 0.2 * W and clear.h >= 0.15 * H:
        return clear, "l"
    # no clear room: the lines stand wholly on one shape of the art (white on a coloured triangle, as the template's
    # own title does) — never across the edge of a shape
    panel = _panel_rect(slide, below, skip)
    if panel is not None and panel.w >= 0.25 * W and panel.h >= 0.15 * H:
        return panel, "l"
    return None


def _cover_note(slide: Slide, oslide: OutlineSlide, col: Bbox, small: float, pal: "_Palette", font: Optional[str], scale: list[float], prefs: tuple = (), align: str = "l", own: Optional[str] = None) -> Optional[int]:
    """The footnote of a cover, divider or closing slide drawn without a sample (a disclaimer «Все цифры условные»):
    small print at the foot of the heading's column."""
    from verstka.ru import typeset

    note = typeset(" ".join((oslide.footnote or "").split()))  # «и прогнозы», «120 000» never torn apart (G1-16)
    if not note:
        return
    size = max(min(small, 12.0), 9.0)
    try:
        hpt = slide.part.package.presentation_part.presentation.slide_height / EMU_PER_PT
    except Exception:  # noqa: BLE001
        hpt = 0.0
    if size < 0.017 * hpt:
        # a large slide (15″ high): 12 pt would be a speck — small print of the slide's own scale (≥ 1.7 % of its height)
        size = min([x for x in scale if x >= 0.017 * hpt - 0.05] or [round(0.017 * hpt * 2) / 2])
    h = int(size * 2.8 * EMU_PER_PT)
    nb = Bbox(x=col.x, y=col.y2 - h, w=col.w, h=h)
    _textbox(slide, nb, [ParagraphSpec(note, bullet=False)], size=size, color=_ink(slide, nb, [own, pal.text2, *prefs, pal.text]) or pal.text2, font=font, scale=scale, anchor="b", align=align)
    from verstka.rendering.fonts import wrap_lines

    lines = min(max(1, len(wrap_lines(note, font, size, False, col.w / EMU_PER_PT - 14.4))), 2)
    return col.y2 - int((lines * 1.2 + 0.3) * size * EMU_PER_PT)  # where its text starts (the box is bottom-anchored)


def _cover_goal_line(slide: Slide, oslide: OutlineSlide, box: Bbox, y: int, size: float, pal: "_Palette", font: Optional[str], scale: list[float], H: int, align: str = "l", small: Optional[float] = None, limit: Optional[int] = None, prefs: tuple = (), own: Optional[str] = None) -> None:
    """A cover's goal (the first paragraph of a title slide, Agent v2: «Цель: …») as a short line under the subtitle,
    a step smaller and in the subtitle's colour, inside the heading's column (never over the art), above the cover's
    footnote (`limit`: where the footnote starts)."""
    from verstka.matching.scorer import cover_goal
    from verstka.rendering.fonts import wrap_lines

    from verstka.ru import typeset

    goal = cover_goal(oslide)
    if not goal:
        return
    shown = typeset(goal)  # «со 120 000 до 255 000 рублей»: a figure keeps its thousands and its unit (G1-16)
    floor = min(size, small) if small else size
    sizes = [size] + [s_ for s_ in sorted(scale, reverse=True) if floor - 0.05 <= s_ < size - 0.05]
    foot = int(H * 0.95) if limit is None else min(int(H * 0.95), limit - int(0.015 * H))
    for s_ in sizes:
        lines = max(1, len(wrap_lines(shown, font, s_, False, box.w / EMU_PER_PT - 7.2)))
        h = int((lines * 1.2 + 0.3) * s_ * EMU_PER_PT)
        if y + h <= foot or s_ == sizes[-1]:
            break
    if y + h > foot:
        # no room under the title and the subtitle even at the small size: the goal goes to the speaker notes —
        # never over the title
        _to_notes(oslide, goal)
        return
    line = Bbox(x=box.x, y=y, w=box.w, h=h)
    _textbox(slide, line, [ParagraphSpec(shown, bullet=False)], size=s_, color=_ink(slide, line, [own, pal.text2, *prefs, pal.text]) or pal.text2, font=font, scale=scale, align=align)


def _ink(slide: Slide, box: Bbox, prefs: list) -> Optional[str]:
    """The first colour of `prefs` that reads (4.5:1) on what the slide, its layout and its master really paint under
    `box` (a triangle, a band, a photo), else the most readable of them and black/white; None when the ground is
    unknown (the caller's colour stands)."""
    try:
        from verstka.rendering.charts import slide_ground

        g = slide_ground(slide, box)
    except Exception:  # noqa: BLE001
        return None
    cands = [c for c in prefs if c]
    if not g or len(g) != 6 or not cands:
        return None
    for c in cands:
        if contrast_ratio(c, g) >= 4.5:
            return c
    return max(cands + ["FFFFFF", "000000"], key=lambda c: contrast_ratio(c, g))


def _has_goal(oslide: OutlineSlide) -> bool:
    try:
        from verstka.matching.scorer import cover_goal

        return bool(cover_goal(oslide))
    except Exception:  # noqa: BLE001
        return False


def _to_notes(oslide: OutlineSlide, text: str) -> None:
    """A line the slide has no room for is said aloud: it goes to the slide's speaker notes (the renderer writes
    `oslide.notes` into the notes page)."""
    try:
        oslide.notes = ((oslide.notes or "").strip() + "\n" + text).strip()
    except Exception:  # noqa: BLE001
        pass


def _unbreak(text: str, box: Bbox, font: Optional[str], size: float, bold: bool, scale: list[float], insets, line_spacing: float, warnings: list[str]):
    """A title with a word wider than its box: the size may give way down to 55 % before the word is broken."""
    res = fit_size([text], box, font, size, bold, scale, insets_emu=insets, line_spacing=line_spacing, min_ratio=0.55)
    if res.broken_word:
        warnings.append("heading word broken")
    return res


def _sub_height(text: str, font: Optional[str], size: float, w: int) -> int:
    from verstka.rendering.fonts import wrap_lines

    return int((len(wrap_lines(text, font, size, False, w / EMU_PER_PT - 7.2)) * 1.2 + 0.2) * size * EMU_PER_PT)


def _render_cover(builder: DeckBuilder, plan_slide: LayoutSlide, oslide: OutlineSlide, manifest: TemplateManifest, ws: TemplateWorkspace, outline: DeckOutline) -> tuple[Slide, list[str]]:
    warnings: list[str] = []
    comp = plan_slide.composition or "bullets"
    family = _family_for(oslide, manifest)
    canvas = _canvas_for(manifest, family, comp)
    ground = None
    if canvas is not None:
        slide, title_ph, ground = _slide_on_canvas(builder, canvas, builder.slide_w, builder.slide_h, keep_art=comp in ("title", "thanks"))
    else:
        layout, family = _layout_for(builder, manifest, family, oslide.kind)
        slide = _add_layout_slide(builder, layout)
        title_ph = None
    pal = _Palette(manifest, family)  # palette of the real family: text must contrast with its background
    if canvas is not None:
        _adopt_canvas_colors(pal, canvas, manifest)
    t = manifest.tokens
    typo = t.typography
    W, H = builder.slide_w, builder.slide_h
    safe = t.spacing.safe_area
    sx, sy, sw, sh = int(safe.x * W), int(safe.y * H), int(safe.w * W), int(safe.h * H)
    if ground is not None:
        # the heading stands on a panel (LCT: white half, photo half): the composition lives inside the panel
        gb, ghex = ground
        pad = int(0.03 * W)
        x1, x2 = max(sx, gb.x + pad), min(sx + sw, gb.x2 - pad)
        y2 = min(sy + sh, gb.y2 - pad)
        if x2 - x1 > 0.25 * W:
            sx, sw, sh = x1, x2 - x1, y2 - sy
        if ghex:
            pal.bg = ghex
            pal.text = pal.text2 = _readable(ghex, [pal.text, t.color_for("text.primary")])
    # every size the template uses (not only the role scale): text fits and grows along the template's own steps
    scale = sorted({s.size_pt for s in typo.scale} | {float(x) for x in (typo.sizes_used or []) if x >= 8})
    font = typo.primary_family
    h1, title_color, title_bold, title_font = _title_style(manifest, family, pal, content_only=comp not in _COVER_COMPS)
    if comp in _COVER_COMPS:
        # a cover or a divider sets its title as the template's own covers do: the canvas's title size, else the
        # largest title of the template's cover/divider samples — never the deck's most frequent title (on a deck
        # whose headings are small pill labels, a 20 pt cover title)
        own = _top_title(canvas) if canvas is not None else None
        covers = [s_.style.size_pt for p in manifest.patterns if p.kind in _COVER_KINDS and not p.reference for s_ in p.slots if s_.role == SlotRole.title and s_.style.size_pt]
        cover_size = (own.style.size_pt if own is not None and own.style.size_pt else None) or (max(covers) if covers else None)
        if cover_size and cover_size > h1:
            h1 = float(cover_size)
    body = typo.size_for("body", 14.0)
    h2 = typo.size_for("h2", body * 1.25)
    small = typo.size_for("small", body * 0.85)
    display = typo.size_for("display", h1 * 1.8)
    if _is_sparse(manifest):
        # a sparse scale's roles are placeholder defaults (a 32 pt body, a 44 pt h2 on an 11″ slide): a cover's
        # subtitle and its goal line keep to the slide's proportions (≤ 5.5 % / 3.1 % of its height)
        hpt = H / EMU_PER_PT
        at_most = lambda cap, v: max([x for x in scale if x <= cap + 0.05] or [round(cap * 2) / 2]) if v > cap else v  # noqa: E731
        body = at_most(0.031 * hpt, body)
        h2 = at_most(0.055 * hpt, h2)
        small = min(small, body)

    # keep only the title (fill it) — everything else is drawn from tokens
    if canvas is None:
        for shp in list(slide.shapes):
            if shp.is_placeholder and shp.placeholder_format.type is not None and str(shp.placeholder_format.type).split(".")[-1].split(" ")[0] in ("TITLE", "CENTER_TITLE"):
                title_ph = shp
            else:
                shp._element.getparent().remove(shp._element)
    c = oslide.content
    if title_ph is not None:
        _materialize(title_ph)  # an inherited placeholder gets its own geometry before any write
    title_h = int(H * (0.16 if comp not in ("title", "section", "thanks") else 0.3))
    if title_ph is not None and comp not in ("title", "section", "thanks"):
        # measure the headline against the placeholder box: the content starts below the fitted text, not below
        # the placeholder's nominal height (which is one line tall on most layouts)
        if title_ph.top is not None:
            box = Bbox(x=int(title_ph.left), y=int(title_ph.top), w=int(title_ph.width), h=int(title_ph.height))
        else:
            box = Bbox(x=sx, y=sy, w=sw, h=title_h)
        insets = (91440, 45720, 91440, 45720)
        if canvas is not None:
            # the canvas keeps the template's heading label (a pill): the same rules as on a cloned slide
            hctx = _SlideCtx(builder, slide, canvas, manifest, ws, outline)
            nb = hctx._widen_on_backing(title_ph._element, (box.x, box.y, box.w, box.h), [oslide.headline], title_font, h1, title_bold, insets)
            box = Bbox(x=nb[0], y=nb[1], w=nb[2], h=nb[3])
            warnings.extend(hctx.warnings)
        else:
            w_clear = clear_width(manifest, W, H, box)
            if w_clear < box.w:  # never under the logos: the heading wraps instead
                box = Bbox(x=box.x, y=box.y, w=w_clear, h=box.h)
                _set_box(title_ph, box)
        res = fit_size([oslide.headline], box, title_font, h1, title_bold, scale, insets_emu=insets, line_spacing=typo.line_height)
        if res.broken_word:
            res = _unbreak(oslide.headline, box, title_font, h1, title_bold, scale, insets, typo.line_height, warnings)
        _fill_keep(title_ph._element, [ParagraphSpec(oslide.headline)], size_pt=res.size_pt)
        text_h = int(res.height_pt * EMU_PER_PT) + insets[1] + insets[3]
        if text_h > box.h or title_ph.top is None:
            # keep the headline inside its own box
            _set_box(title_ph, Bbox(x=box.x, y=box.y, w=box.w, h=max(box.h, text_h)))
        y_after = box.y + max(box.h, text_h)
    elif comp not in ("title", "section", "thanks"):
        if title_ph is not None:
            title_ph._element.getparent().remove(title_ph._element)
        _textbox(slide, Bbox(x=sx, y=sy, w=sw, h=title_h), [ParagraphSpec(oslide.headline)], size=h1, color=title_color, font=title_font, bold=title_bold, scale=scale, anchor="t")
        y_after = sy + title_h
    elif canvas is not None and title_ph is not None:
        # title / section / thanks on a sample of that kind: its own title box, the subtitle right under it
        box = Bbox(x=int(title_ph.left), y=int(title_ph.top), w=int(title_ph.width), h=int(title_ph.height))
        insets = _heading_insets(title_ph)
        from verstka.matching.scorer import split_display_title
        from verstka.ru import typeset

        raw = oslide.headline
        if comp == "title" and not oslide.subtitle:
            # a cover heading of two phrases («…: план увеличения прибыли») is a title and its subtitle, as on a
            # cloned cover
            head, tail = split_display_title(" ".join(raw.split()))
            if tail:
                oslide = oslide.model_copy(update={"headline": head, "subtitle": tail})
                raw = head
        headline = typeset(raw)  # «с каждой» never split after the preposition
        if box.w < int(0.45 * W):
            # a sample's title box made for one short word (a quarter of the slide): the heading takes the free room
            # of its band, up to the page's right margin
            right = int(safe.x2 * W)
            free = _free_width(slide, box, W, H, right, skip=(title_ph._element,))
            if free > box.w:
                box = Bbox(x=box.x, y=box.y, w=free, h=box.h)
                _set_box(title_ph, box)
                warnings.append("cover title box widened into the free room of its band")
        # a cover heading is read from across the room, but never louder than a sixth of the slide per line (a sample's
        # 168 pt display word set for «SODA»); the subtitle and the note keep room under it
        cap = max(0.16 * H / EMU_PER_PT, min(scale) if scale else 0.0)
        h1_cover = min(h1, _snap_heading(cap, scale) if scale else cap)
        tall = box.h > int(0.5 * H)  # a box set for one giant word: the heading is placed, not stretched, in it
        if tall:
            box = Bbox(x=box.x, y=box.y, w=box.w, h=int(0.5 * H))
        res = fit_size([headline], box, title_font, h1_cover, title_bold, scale, insets_emu=insets, line_spacing=typo.line_height, min_ratio=0.45)
        if res.broken_word:
            res = _unbreak(headline, box, title_font, h1_cover, title_bold, scale, insets, typo.line_height, warnings)
        _fill_keep(title_ph._element, [ParagraphSpec(headline)], size_pt=res.size_pt)
        text_h = int(res.height_pt * EMU_PER_PT) + insets[1] + insets[3]
        if text_h > box.h:
            _set_box(title_ph, Bbox(x=box.x, y=box.y, w=box.w, h=max(min(text_h, int(H * 0.95) - box.y), 0)))
        elif text_h < box.h - int(0.04 * H):
            # the box hugs its lines where the sample's box was set for a taller word: the subtitle follows the
            # letters, not the sample's empty box (a bottom- or centre-anchored box keeps its lines where they stood)
            anchor = (title_ph._element.find(".//" + q("a:bodyPr")).get("anchor") if title_ph._element.find(".//" + q("a:bodyPr")) is not None else None) or "t"
            y_new = box.y2 - text_h if anchor == "b" else box.y + (box.h - text_h) // 2 if anchor == "ctr" else box.y
            if tall:
                # the heading and its subtitle stand a little above the optical centre of the slide
                y_new = max(box.y, int(0.40 * H) - text_h // 2)
            box = Bbox(x=box.x, y=y_new, w=box.w, h=text_h)
            _set_box(title_ph, box)
        if box.x < int(0.035 * W) and _backing_of(slide, title_ph._element, W, H) is None:
            # the sample's title stood inside a panel this cover does not keep (a sign, an arrow): off the slide's edge
            m = max(int(safe.x * W), int(0.06 * W))
            box = Bbox(x=m, y=box.y, w=max(box.w - (m - box.x), int(0.4 * W)), h=box.h)
            _set_box(title_ph, box)
        # the heading's colour is read where it now stands (a light title of a dropped blue sign on a light page)
        _check_heading_color(slide, title_ph, canvas, None, pal, manifest)
        # the lines under the title start where its letters start (its own left inset), never at the slide's edge
        shift = max(insets[0] - 91440, 0)
        col = Bbox(x=box.x + shift, y=box.y, w=max(box.w - shift, int(0.2 * W)), h=box.h)
        sub = oslide.subtitle or (oslide.section if comp == "section" else None)
        sub = typeset(sub) if sub else sub  # «на 6 месяцев» never split after the preposition (G1-16)
        y_sub = box.y + max(box.h, text_h) + int(H * 0.02)
        note_col = Bbox(x=col.x, y=sy, w=min(col.w, sw), h=sh)
        align = "ctr" if _title_align(title_ph) == "ctr" else "l"  # the lines under a centred heading are centred too
        region = _cover_lines_region(slide, canvas, title_ph, Bbox(x=col.x, y=y_sub, w=col.w, h=max(sy + sh - y_sub, 0)), Bbox(x=sx, y=sy, w=sw, h=sh), W, H) if (sub or oslide.footnote or _has_goal(oslide)) else None
        # the sample's own subtitle: its colour, and its place when it stands lower under the title in the same
        # column (a band the template prints its subtitle on — never half on the band's edge)
        sub_slot = next((s_ for s_ in canvas.slots if s_.role == SlotRole.subtitle), None)
        slot_color = sub_slot.style.color_hex if sub_slot is not None else None
        if region is None and sub_slot is not None:
            sbb = sub_slot.bbox.to_emu(W, H)
            overlap = max(0, min(sbb.x2, col.x2) - max(sbb.x, col.x))
            if y_sub <= sbb.y <= sy + sh - int(0.12 * H) and overlap >= 0.5 * min(sbb.w, col.w):
                y_sub = sbb.y
        if region is not None:
            # the column under the title stands on the template's art: the lines go where the art leaves room
            col, align = region
            y_sub = col.y
            note_col = col
            warnings.append("cover lines set clear of the template's art")
        if sub:
            sb = Bbox(x=col.x, y=y_sub, w=col.w, h=min(int(H * 0.16), int(H * 0.92) - y_sub))
            _textbox(slide, sb, [ParagraphSpec(sub)], size=h2, color=_ink(slide, sb, [slot_color, pal.text2, title_color, pal.text]) or pal.text2, font=font, scale=scale, align=align)
            y_sub += min(_sub_height(sub, font, h2, col.w), int(H * 0.16)) + int(H * 0.015)
        note_top = _cover_note(slide, oslide, note_col, small, pal, font, scale, prefs=(title_color,), align=align, own=slot_color)
        _cover_goal_line(slide, oslide, col, y_sub, max(min(body, h2 * 0.8), small), pal, font, scale, H, align=align, small=small, limit=note_top, prefs=(title_color,), own=slot_color)
        return slide, warnings
    else:
        if title_ph is not None:
            title_ph._element.getparent().remove(title_ph._element)
        y_after = sy
    if title_ph is not None and title_ph.left is not None and comp not in ("title", "section", "thanks") and 0 < int(title_ph.left) < sx + sw // 2:
        # content is aligned to the heading's left edge, as the template's own slides are
        sw = sx + sw - int(title_ph.left)
        sx = int(title_ph.left)
    top = y_after + int(H * 0.02)
    avail_h = sy + sh - top
    gap = int((t.spacing.gutter or 0.03) * W)
    card_gap = min(gap, int(0.03 * W))  # the template gutter is a column gutter; card grids need a narrower one

    if comp in ("title", "section", "thanks"):
        # a page of its own (no sample of the kind): a comfortable margin, the heading's colour read on what really
        # stands under it, a heading of two phrases set as a title and its subtitle
        mx = max(sx, int(0.06 * W))
        sw = max(sx + sw - mx, int(0.5 * W)) if mx > sx else sw
        sx = mx
        if comp == "title" and not oslide.subtitle:
            from verstka.matching.scorer import split_display_title

            head, tail = split_display_title(" ".join(oslide.headline.split()))
            if tail:
                oslide = oslide.model_copy(update={"headline": head, "subtitle": tail})
        box = Bbox(x=sx, y=int(H * 0.30), w=int(sw * 0.8), h=int(H * 0.3))
        try:
            from verstka.rendering.charts import slide_ground

            g = slide_ground(slide, box)
        except Exception:  # noqa: BLE001
            g = None
        if g and contrast_ratio(title_color, g) < 3.0:
            title_color = _readable(g, [title_color, pal.text, manifest.tokens.color_for("text.primary")])
            pal.bg = g
            if contrast_ratio(pal.text, g) < 4.5:
                pal.text = _readable(g, [pal.text, title_color])
            if contrast_ratio(pal.text2, g) < 4.5:
                pal.text2 = pal.text
        _textbox(slide, box, [ParagraphSpec(oslide.headline)], size=max(h1, display * 0.8) if comp == "title" else h1, color=title_color, font=title_font, bold=title_bold, scale=scale, anchor="b")
        sub = oslide.subtitle or (oslide.section if comp == "section" else None)
        if sub:
            from verstka.ru import typeset

            sub = typeset(sub)  # «на 6 месяцев» never split after the preposition (G1-16)
        y_goal = int(H * 0.62)
        col = Bbox(x=sx, y=y_goal, w=int(sw * 0.8), h=max(sy + sh - y_goal, 0))
        note_col = Bbox(x=sx, y=sy, w=int(sw * 0.8), h=sh)
        align = "l"
        region = _cover_lines_region(slide, None, None, col, Bbox(x=sx, y=sy, w=sw, h=sh), W, H) if (sub or oslide.footnote or _has_goal(oslide)) else None
        if region is not None:
            col, align = region
            y_goal = col.y
            note_col = col
            warnings.append("cover lines set clear of the template's art")
        if sub:
            sb = Bbox(x=col.x, y=y_goal, w=col.w, h=int(H * 0.16))
            _textbox(slide, sb, [ParagraphSpec(sub)], size=h2, color=_ink(slide, sb, [pal.text2, title_color, pal.text]) or pal.text2, font=font, scale=scale, align=align)
            y_goal += min(_sub_height(sub, font, h2, col.w), int(H * 0.16)) + int(H * 0.015)
        note_top = _cover_note(slide, oslide, note_col, small, pal, font, scale, prefs=(title_color,), align=align)
        _cover_goal_line(slide, oslide, Bbox(x=col.x, y=y_goal, w=col.w, h=int(H * 0.1)), y_goal, max(min(body, h2 * 0.8), small), pal, font, scale, H, align=align, small=small, limit=note_top, prefs=(title_color,))
        return slide, warnings

    if comp == "bullets" or (comp in ("image_text",) and not c.image_hint):
        paras = [ParagraphSpec(b, bullet=True) for b in c.bullets] or [ParagraphSpec(p, bullet=False) for p in c.paragraphs] or [ParagraphSpec(i.title, bullet=True) for i in _items(oslide)]
        if oslide.subtitle:
            paras = [ParagraphSpec(oslide.subtitle, bullet=False, bold=True)] + paras
        _textbox(slide, Bbox(x=sx, y=top, w=int(sw * 0.9), h=avail_h), paras, size=body, color=pal.text, font=font, scale=scale, line_spacing=typo.line_height)
        return slide, warnings

    if comp == "image_text":
        path = pick_asset(manifest, ws, c.image_hint)
        img_w = int(sw * 0.42)
        if path:
            insert_picture(slide, Bbox(x=sx, y=top, w=img_w, h=avail_h), path)
            x_text = sx + img_w + gap
        else:
            x_text = sx
        paras = [ParagraphSpec(b, bullet=True) for b in c.bullets] or [ParagraphSpec(p) for p in c.paragraphs]
        _textbox(slide, Bbox(x=x_text, y=top, w=sx + sw - x_text, h=avail_h), paras, size=body, color=pal.text, font=font, scale=scale)
        return slide, warnings

    if comp in ("cards", "process", "agenda", "comparison", "two_column"):
        items = _items(oslide)
        n = max(len(items), 1)
        cols = n if n <= 4 else math.ceil(n / 2)
        rows = math.ceil(n / cols)
        if comp in ("two_column", "comparison"):
            cols, rows = min(n, 3), 1
        cw = int((sw - card_gap * (cols - 1)) / cols)
        ch = int((avail_h - card_gap * (rows - 1)) / rows)
        for i, item in enumerate(items):
            r_, c_ = divmod(i, cols)
            box = Bbox(x=sx + c_ * (cw + card_gap), y=top + r_ * (ch + card_gap), w=cw, h=ch)
            if comp in ("cards", "comparison") and (pal.card_fill or pal.card_line):
                _rect(slide, box, pal.card_fill, pal.card_line if not pal.card_fill else None, pal.radius_emu)
            inset = int(cw * 0.07)
            head_h = int(ch * 0.3)
            y_head = box.y + inset
            if comp in ("process", "agenda"):
                d = int(min(cw, ch) * 0.28)
                _circle(slide, Bbox(x=box.x, y=box.y, w=d, h=d), pal.accent, str(i + 1), "FFFFFF" if contrast_ratio("FFFFFF", pal.accent) >= 3 else pal.text, max(small, d / 12700 * 0.45), font)
                y_head = box.y + d + int(H * 0.015)
                inset = 0
                head_h = int(ch * 0.22)
            _textbox(slide, Bbox(x=box.x + inset, y=y_head, w=cw - 2 * inset, h=head_h), [ParagraphSpec(item.title)], size=h2 if comp != "agenda" else h2, color=pal.accent if comp in ("cards", "comparison") else pal.text, font=font, bold=True, scale=scale, anchor="t")
            body_paras = [ParagraphSpec(b, bullet=True) for b in item.bullets] or ([ParagraphSpec(item.text)] if item.text else [])
            if body_paras:
                _textbox(slide, Bbox(x=box.x + inset, y=y_head + head_h, w=cw - 2 * inset, h=box.y2 - (y_head + head_h) - inset), body_paras, size=body, color=pal.text, font=font, scale=scale, line_spacing=typo.line_height)
        return slide, warnings

    if comp in ("stat_row", "big_number"):
        nums = c.numbers
        n = max(len(nums), 1)
        if comp == "big_number" or n == 1:
            num = nums[0] if nums else None
            if num:
                _textbox(slide, Bbox(x=sx, y=top, w=int(sw * 0.5), h=int(avail_h * 0.55)), [ParagraphSpec(num.value)], size=display, color=pal.accent, font=font, bold=True, anchor="b", scale=scale, fit=False)
                _textbox(slide, Bbox(x=sx, y=top + int(avail_h * 0.57), w=int(sw * 0.5), h=int(avail_h * 0.3)), [ParagraphSpec(num.label)], size=h2, color=pal.text2, font=font, scale=scale)
            if c.paragraphs or c.bullets:
                paras = [ParagraphSpec(b, bullet=True) for b in c.bullets] or [ParagraphSpec(p) for p in c.paragraphs]
                _textbox(slide, Bbox(x=sx + int(sw * 0.55), y=top, w=int(sw * 0.45), h=avail_h), paras, size=body, color=pal.text, font=font, scale=scale)
            return slide, warnings
        cw = int((sw - card_gap * (n - 1)) / n)
        for i, num in enumerate(nums):
            x = sx + i * (cw + card_gap)
            _textbox(slide, Bbox(x=x, y=top, w=cw, h=int(avail_h * 0.45)), [ParagraphSpec(num.value)], size=display * 0.8, color=pal.accent, font=font, bold=True, anchor="b", scale=scale)
            _textbox(slide, Bbox(x=x, y=top + int(avail_h * 0.47), w=cw, h=int(avail_h * 0.35)), [ParagraphSpec(num.label)], size=body, color=pal.text2, font=font, scale=scale)
        return slide, warnings

    if comp == "chart_text" and c.chart is not None:
        has_text = bool(c.paragraphs or c.bullets)
        chart_w = int(sw * (0.62 if has_text else 1.0))
        try:
            add_chart(slide, Bbox(x=sx, y=top, w=chart_w, h=avail_h), c.chart, outline, manifest.components.chart_style, typo, text_hex=pal.text, neutral_hex=t.color_for("neutral.1"))
        except Exception as e:  # noqa: BLE001
            warnings.append(f"chart failed: {str(e)[:120]}")
        if has_text:
            paras = [ParagraphSpec(b, bullet=True) for b in c.bullets] or [ParagraphSpec(p) for p in c.paragraphs]
            _textbox(slide, Bbox(x=sx + chart_w + gap, y=top, w=sw - chart_w - gap, h=avail_h), paras, size=body, color=pal.text, font=font, scale=scale)
        return slide, warnings

    if comp == "table" and c.table is not None:
        rows = len(c.table.rows) + 1
        tbl_h = min(avail_h, int(rows * H * 0.075))
        style = manifest.components.table_style
        if not style.body_text_hex or contrast_ratio(style.body_text_hex, pal.bg) < 4.5:
            style = style.model_copy(update={"body_text_hex": pal.text, "band_fill_hex": None})  # dark template text on a dark ground
        try:
            add_table(slide, Bbox(x=sx, y=top, w=sw, h=tbl_h), c.table, style, typo)
        except Exception as e:  # noqa: BLE001
            warnings.append(f"table failed: {str(e)[:120]}")
        if c.paragraphs and tbl_h < avail_h - int(H * 0.08):
            _textbox(slide, Bbox(x=sx, y=top + tbl_h + int(H * 0.02), w=sw, h=avail_h - tbl_h - int(H * 0.02)), [ParagraphSpec(p) for p in c.paragraphs], size=small, color=pal.text2, font=font, scale=scale)
        return slide, warnings

    if comp == "quote" and c.quote:
        _textbox(slide, Bbox(x=sx, y=top, w=int(sw * 0.85), h=int(avail_h * 0.65)), [ParagraphSpec("«" + c.quote.strip("«»\"") + "»")], size=h2 * 1.15, color=pal.text, font=font, scale=scale, anchor="ctr")
        if c.quote_author:
            _textbox(slide, Bbox(x=sx, y=top + int(avail_h * 0.68), w=int(sw * 0.85), h=int(avail_h * 0.2)), [ParagraphSpec(c.quote_author)], size=small, color=pal.accent, font=font, scale=scale)
        return slide, warnings

    # fallback: bullets
    paras = [ParagraphSpec(b, bullet=True) for b in c.bullets] or [ParagraphSpec(p) for p in c.paragraphs] or [ParagraphSpec(i.title, bullet=True) for i in _items(oslide)]
    if paras:
        _textbox(slide, Bbox(x=sx, y=top, w=int(sw * 0.9), h=avail_h), paras, size=body, color=pal.text, font=font, scale=scale)
    else:
        warnings.append("empty slide content")
    return slide, warnings
