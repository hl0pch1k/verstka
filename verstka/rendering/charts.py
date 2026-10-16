"""Native PowerPoint charts styled from template tokens.

A chart is drawn the way a designer would finish it, whatever the template:

* colours come from the template: accent.1 first, then the series colours of the template's own chart parts
  (schemeClr resolved through the theme), the manifest series colours and the theme accents — every colour visible
  on the ground, never a text or background colour, and ≥ ΔE 20 from every other, so no two series look alike;
* one series with a highlight: the highlight in accent.1, the other bars accent.1 mixed toward the ground (a tint of
  the brand colour, never a neutral grey); several series that tell a story (Было/Стало, 2024/2025, Факт/План): the
  latest in the accent, earlier ones in its tint, a plan as an outline or a dashed line;
* a word unit is written once, in a muted caption at the top left (or not at all when the caller already put it
  above the chart), and every label is a plain number of that one scale; a one-glyph unit (%, ₽) rides on every
  label; zero labels are hidden; the value axis and gridlines go while labels carry the values;
* type from the template: sizes snapped to its type scale, bold only where the template speaks in bold (elsewhere
  the highlight steps up a size in the accent);
* long category labels (> 12 characters) or many categories turn a column chart into horizontal bars read top-down;
* the plot area has a manual layout that reserves room for every label, so nothing is cut and the bars fill the
  frame exactly; chart and plot area are transparent, so the chart sits on the slide's own ground;
* lines are 2.25 pt with ringed markers and labels placed off the line; an area is a soft fill under a 2.25 pt edge;
  pies and doughnuts show percentages inside the slices that fit them, no leader lines;
* a missing value is a break in the line (dispBlanksAs gap, when `spec.gaps` says so), never a zero (without the flag a
  line or an area runs through it): the end label, its marker and the default highlight go to the last point that has
  a value; `highlight_index` = −1 asks for no highlight at all;
* value axes of top-down bars stand at their bottom edge; a range label («1 250–1 500») breaks after its dash, never
  inside a number, else the labels turn;
* a histogram's bins touch (no gap, a hairline of the ground between them) and are read like columns; stacked bars
  are parts of one whole: the template's categorical colours (a remainder quiet), a hairline of the ground between
  the segments, each value inside its segment where it fits, else a quiet value axis; grouped bars are the
  several-series bars above; several series with a highlighted category ease the other categories toward the ground;
* a scatter shows points only: translucent ringed markers on two quiet value axes on round ticks (the Y axis stands
  at the left edge even when X spans zero; the zero only where the data are near it); in a dense cloud the highlighted
  point's value takes the nearest free ground, on a chip of the ground where it must touch a marker;
* axis titles (`x_title` the categories or X, `y_title` the values or Y) are set in the template's face at the label
  size in the muted text colour, along their axis; the value axis's title carries the word unit, so no caption repeats
  it; an axis whose ticks the labels replace stays only to carry its title;
* a source line («Источник: …») in small muted print under the chart, inside its frame, when the caller does not
  print it on the slide (`source_said`).
"""

from __future__ import annotations

import copy
import io
import math
import re
import weakref
from dataclasses import dataclass, field
from typing import Iterable, Optional, Sequence

from lxml import etree
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.slide import Slide
from pptx.util import Emu, Pt

from verstka.rendering.fonts import break_pieces, text_width_pt, wrap_lines
from verstka.schemas.common import Bbox, contrast_ratio, relative_luminance, rgb_to_lab
from verstka.schemas.outline import ChartSpec, DeckOutline, Series
from verstka.schemas.template import ChartStyleSpec, Typography

_TYPE_MAP = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "histogram": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}
_GROUPED = {"grouped_column": "column", "grouped_bar": "bar"}  # several series side by side: the clustered bars
_HORIZONTAL = ("bar", "stacked_bar")
_STACKED = ("stacked_column", "stacked_bar")
_BARS = ("column", "bar", "histogram", "stacked_column", "stacked_bar")
NO_HIGHLIGHT = -1  # highlight_index: no point stands out (a line's last point is not accented either)

_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
EMU_PER_PT = 12700

MIN_FILL_CONTRAST = 1.4  # a bar, slice or line must stand off its ground at least this much
LEAD_MIN_CONTRAST = 2.2  # the lead colour of a chart (its largest slice, its highlighted bar) reads clearly on the ground
DARK_MUTED_DE = 34.0  # on a dark ground the stepped-back bars are as bright as this distance from the accent allows
MIN_SERIES_DE = 20.0  # CIE76 distance between two series colours that a reader can tell apart
SAME_HUE_DE = 32.0  # … and between two shades of the same hue
MUTED_TOWARD_GROUND = 0.58  # the bars that step back: accent.1 moved this share of the way to the ground
LONG_LABEL_CHARS = 12  # a category label longer than this reads better as a horizontal bar
PAIR_LABEL_CHARS = 20  # … unless the chart has only two categories: two wide columns take labels this long
MAX_COLUMNS = 7  # more categories than this: horizontal bars
HOLE_SIZE = 56  # doughnut hole, % of the diameter
BOLD_VOICE_SHARE = 0.3  # the template sets its display / h1 in bold at least this often: emphasis may be bold
LABEL_SEP = ": "  # «Было: 48» — a series named on its label
SEAM_PT = 0.75  # the hairline of the ground between a histogram's bins and between stacked segments
STACK_GAP = (40, 180)  # gapWidth range of stacked bars: one bar per category, wider than a cluster's
SCATTER_ALPHA = (0.85, 0.65)  # marker fill opacity: a few points, many points (overlaps stay visible)
TITLE_LINE = 1.25  # an axis title's line height over its size
# a value printed on a filled segment or slice (stacked bars, pie, doughnut): the template's own ink on its accent fill is
# accepted from 3:1 (white on the brand blue 0077FF is 4.1:1, on the pink FF3985 3.3:1) — the same convention as the
# audit (audit/checks/template.py, TABLE_CONTRAST_LOW: a pair of template colours from 3:1 is the template's own choice)
STACK_LABEL_CONTRAST = 3.0


def _c(tag: str) -> str:
    return f"{{{_C}}}{tag}"


def _a(tag: str) -> str:
    return f"{{{_A}}}{tag}"


def _local(el) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


# ---------------------------------------------------------------------------------------------- colour helpers


def _norm(hex_: Optional[str]) -> Optional[str]:
    h = (hex_ or "").strip().lstrip("#").upper()
    return h if len(h) == 6 and all(ch in "0123456789ABCDEF" for ch in h) else None


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.upper())


def _rgb_t(hex_: str) -> tuple[int, int, int]:
    return int(hex_[0:2], 16), int(hex_[2:4], 16), int(hex_[4:6], 16)


def _mix(fg: str, bg: str, share: float) -> str:
    """`share` of fg over bg."""
    a = [int(fg[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(bg[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{max(0, min(255, int(round(x * share + y * (1 - share))))):02X}" for x, y in zip(a, b))


def delta_e(a: str, b: str) -> float:
    """CIE76 distance between two hex colours."""
    la, lb = rgb_to_lab(*_rgb_t(a)), rgb_to_lab(*_rgb_t(b))
    return math.sqrt(sum((p - q) ** 2 for p, q in zip(la, lb)))


def _chroma(hex_: str) -> float:
    _, a, b = rgb_to_lab(*_rgb_t(hex_))
    return math.hypot(a, b)


def _lightness(hex_: str) -> float:
    return rgb_to_lab(*_rgb_t(hex_))[0]


def _hue_gap(a: str, b: str) -> float:
    """Angle between two colours' Lab hues (degrees); a near-grey counts as a different hue from anything."""
    _, a1, b1 = rgb_to_lab(*_rgb_t(a))
    _, a2, b2 = rgb_to_lab(*_rgb_t(b))
    if math.hypot(a1, b1) < 12 or math.hypot(a2, b2) < 12:
        return 180.0
    d = abs(math.degrees(math.atan2(b1, a1)) - math.degrees(math.atan2(b2, a2))) % 360
    return min(d, 360 - d)


def _is_dark(hex_: str) -> bool:
    return relative_luminance(hex_) < 0.18


def shade_to(hex_: str, ground: str, target: float) -> Optional[str]:
    """The colour's own hue and saturation made darker (on a light ground) or lighter (on a dark one), in steps of 5 % of
    its lightness, until it reads at `target` against the ground; None when no step does."""
    import colorsys

    r, g, b = (x / 255 for x in _rgb_t(hex_))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    light_ground = relative_luminance(ground) >= 0.4
    for i in range(1, 20):
        l2 = l * (1 - i / 20) if light_ground else l + (1 - l) * i / 20
        rr, gg, bb = colorsys.hls_to_rgb(h, l2, s)
        c = "%02X%02X%02X" % tuple(int(round(min(max(v, 0.0), 1.0) * 255)) for v in (rr, gg, bb))
        if contrast_ratio(c, ground) >= target:
            return c
    return None


def _is_data_color(hex_: str) -> bool:
    """False for a text or background colour (a near-black or near-white grey): a slice or a series never takes one."""
    return not (_chroma(hex_) < 12 and not (22 <= _lightness(hex_) <= 92))


def muted_tint(accent: str, ground: str, toward: float = MUTED_TOWARD_GROUND) -> str:
    """accent.1 moved `toward` of the way to the ground: the colour of the bars that step back behind the highlight.
    It is a tint (or shade) of the brand colour, never a grey, it stands off the ground (≥ 1.4:1; on a dark ground as
    bright as a clear step from the accent allows, since a faint shade sinks into it) and it stays ≥ ΔE 20 away from
    the accent itself."""
    accent, ground = _norm(accent) or "0077FF", _norm(ground) or "FFFFFF"
    share = 1.0 - toward
    if _is_dark(ground) and toward == MUTED_TOWARD_GROUND:
        # on a dark ground a faint shade sinks into it: the brightest shade that is still a clear step from the accent
        for i in range(20):
            s = 0.8 - 0.02 * i
            if s <= share or delta_e(_mix(accent, ground, s), accent) >= DARK_MUTED_DE:
                share = max(share, s)
                break
    out = _mix(accent, ground, share)
    while contrast_ratio(out, ground) < MIN_FILL_CONTRAST and share < 0.86:
        share += 0.02
        out = _mix(accent, ground, share)
    # too close to the accent now: step back toward the ground as long as it stays visible
    while delta_e(out, accent) < MIN_SERIES_DE and share > 0.2 and contrast_ratio(_mix(accent, ground, share - 0.04), ground) >= MIN_FILL_CONTRAST:
        share -= 0.04
        out = _mix(accent, ground, share)
    if delta_e(out, accent) < MIN_SERIES_DE:
        # an accent that barely differs from its ground (pale accent on white): step toward the text side instead
        out = _mix(accent, "000000" if not _is_dark(ground) else "FFFFFF", 0.6)
    return out


def eased(fill: str, ground: str, share: float = 0.45) -> str:
    """A series colour stepped back toward the ground — a category that is not the highlighted one: `share` of the
    colour over the ground (the same for every series, so a group keeps its order of strength), more of it only where
    the step would sink into the ground (every bar ≥ 1.4:1 off it). Never another hue, never past the colour itself."""
    fill, ground = _norm(fill) or "0077FF", _norm(ground) or "FFFFFF"
    for s in (share, 0.55, 0.65, 0.75, 0.85):
        c = _mix(fill, ground, s)
        if contrast_ratio(c, ground) >= MIN_FILL_CONTRAST:
            return c
    return fill


def _readable_text(text_hex: Optional[str], ground: str) -> str:
    t = _norm(text_hex)
    if t and contrast_ratio(t, ground) >= 3.0:
        return t
    return "FFFFFF" if contrast_ratio("FFFFFF", ground) >= contrast_ratio("000000", ground) else "1A1A1A"


def _muted_text(text: str, ground: str) -> str:
    """The secondary text colour for category labels and legends: text eased toward the ground, still ≥ 4.5:1."""
    for share in (0.66, 0.72, 0.78, 0.85, 0.92):
        c = _mix(text, ground, share)
        if contrast_ratio(c, ground) >= 4.5:
            return c
    return text


def _rule_color(neutral_hex: Optional[str], text: str, ground: str) -> str:
    """The baseline hairline: the template's divider when it reads as a quiet line on this ground (visible, not loud,
    and on the text's side of the ground — lighter than a dark ground), else text eased far toward the ground."""
    n = _norm(neutral_hex)
    floor = 1.45 if _is_dark(ground) else 1.2  # the same ratio reads fainter near black
    if n and floor <= contrast_ratio(n, ground) <= 4.0:
        lg = relative_luminance(ground)
        if (relative_luminance(n) > lg) == (relative_luminance(text) > lg):
            return n
    for share in (0.25, 0.3, 0.4):
        c = _mix(text, ground, share)
        if contrast_ratio(c, ground) >= floor:
            return c
    return _mix(text, ground, 0.5)


def label_on(fill: str, text: str, ground: str, floor: float = STACK_LABEL_CONTRAST) -> str:
    """Ink for a label printed on a filled segment or slice — the one rule of the PPTX charts and of their HTML/SVG twin
    (verstka.export.svg_charts): white, the template's text, the ground, near-black — the first that reads on `fill` at
    `floor` (3:1: the template's own ink on its accent, white on 0077FF and on FF3985), else the one of white and
    pure black that reads best (a light fill such as 8BBFFA takes the dark ink; black reaches further than near-black)."""
    for c in ("FFFFFF", text, ground, "1A1A1A"):
        if contrast_ratio(c, fill) >= floor:
            return c
    return "FFFFFF" if contrast_ratio("FFFFFF", fill) >= contrast_ratio("000000", fill) else "000000"


def _quiet_fill(text: str, ground: str) -> str:
    """The colour of an «Прочее» slice: text eased toward the ground — present, deliberately without a hue."""
    for share in (0.22, 0.28, 0.36, 0.46):
        c = _mix(text, ground, share)
        if contrast_ratio(c, ground) >= MIN_FILL_CONTRAST:
            return c
    return _mix(text, ground, 0.55)


def chart_palette(base: Sequence[str], ground: Optional[str], n: int, extra: Iterable[str] = ()) -> list[str]:
    """`n` series colours: `base[0]` (accent.1) first, then candidates from `base` and `extra` in order that are data
    colours (not a near-black or near-white grey), visible on the ground (≥ 1.4:1) and ≥ ΔE 20 from every colour
    already taken. When the template runs out of such colours, tints and shades of the taken ones fill in — still
    ΔE ≥ 20 from all — and only then a nearer colour is accepted."""
    n = max(1, n)
    cands: list[str] = []
    for h in [*base, *extra]:
        h = _norm(h)
        if h and h not in cands:
            cands.append(h)
    if not cands:
        cands = ["0077FF"]
    g = _norm(ground)
    ground_ = g or "FFFFFF"
    first = cands[0]
    visible = [h for h in cands[1:] if _is_data_color(h) and contrast_ratio(h, ground_) >= MIN_FILL_CONTRAST]
    rest = sorted(visible, key=lambda h: (_chroma(h) < 12, visible.index(h)))
    out = [first]

    def far_from_all(h: str, de: float, strict: bool = True) -> bool:
        # two shades of one hue need a bigger step than two hues to read as different series
        return all(delta_e(h, o) >= de and (not strict or _hue_gap(h, o) >= 30 or delta_e(h, o) >= max(de, SAME_HUE_DE)) for o in out)

    # strong, clearly different hues first (another hue family, well visible), then merely distinct colours
    passes = (
        lambda h: _chroma(h) >= 12 and contrast_ratio(h, ground_) >= 2.0 and far_from_all(h, MIN_SERIES_DE * 1.5) and all(_hue_gap(h, o) >= 40 for o in out),
        lambda h: far_from_all(h, MIN_SERIES_DE * 1.5),
        lambda h: far_from_all(h, MIN_SERIES_DE),
    )
    for ok in passes:
        for h in rest:
            if len(out) >= n:
                return out
            if h not in out and ok(h):
                out.append(h)
    far = "FFFFFF" if _is_dark(ground_) else "000000"

    def synth() -> list[str]:
        s = [muted_tint(first, ground_), _mix(first, far, 0.55), _mix(first, ground_, 0.72), _mix(first, far, 0.3)]
        for c in out[1:]:
            s += [muted_tint(c, ground_), _mix(c, far, 0.55)]
        return [h for h in s if contrast_ratio(h, ground_) >= MIN_FILL_CONTRAST]

    for de, strict in ((MIN_SERIES_DE, True), (MIN_SERIES_DE, False), (MIN_SERIES_DE * 0.6, False)):
        for h in synth() + rest:
            if len(out) >= n:
                return out
            if h not in out and far_from_all(h, de, strict):
                out.append(h)
    # a narrow palette (Классика's beiges and browns: four colours for eight slices, cycling — the legend matched two
    # slices per swatch): every taken colour stepped toward the ground and toward the far end in finer steps, each
    # still a visible step (ΔE ≥ 8) from all taken, before any colour repeats
    for de in (MIN_SERIES_DE * 0.45, 8.0):
        cands_fine = [_mix(c, tgt, t) for c in list(out) for tgt in (far, ground_) for t in (0.18, 0.32, 0.46, 0.6, 0.74)]
        for h in cands_fine:
            if len(out) >= n:
                return out
            if h not in out and contrast_ratio(h, ground_) >= MIN_FILL_CONTRAST and far_from_all(h, de, False):
                out.append(h)
    i = 0
    while len(out) < n:  # a very long series list: repeat, never fail
        out.append(out[i % len(out)])
        i += 1
    return out


def pie_shades(values: Sequence[Optional[float]], accent: str, ground: Optional[str], quiet_idx: Iterable[int] = (), text: Optional[str] = None) -> list[tuple[str, float]]:
    """Slice colours of a pie or a doughnut as (base colour, share over the ground): the largest slice in the accent
    (share 1), the others in tints of it that step toward the ground with the slice's rank (the bigger the part, the
    stronger the colour), every tint still visible on the ground; a remainder («Прочее») the text colour eased
    toward the ground. One hue reads as one whole; the eye goes to the largest part first. A tint is the accent laid
    over the ground at that share — the solid colour of a slice, or the accent at that opacity on a legend swatch."""
    accent = _norm(accent) or "0077FF"
    ground_ = _norm(ground) or "FFFFFF"
    text_ = _readable_text(text, ground_)
    n = len(values)
    quiet_set = set(quiet_idx)
    quiet_share = next((sh for sh in (0.22, 0.28, 0.36, 0.46) if contrast_ratio(_mix(text_, ground_, sh), ground_) >= MIN_FILL_CONTRAST), 0.55)
    out: list[tuple[str, float]] = [(text_, quiet_share)] * n
    order = sorted((j for j in range(n) if j not in quiet_set), key=lambda j: (-(values[j] or 0.0), j))
    if not order:
        return out
    out[order[0]] = (accent, 1.0)
    m = len(order) - 1
    if m:
        # tints evenly spaced by how different they look (CIE ΔE along the way from the accent to the ground): two
        # neighbours up to ΔE 32 apart, closer only when many slices share the room; the palest still clearly off the
        # ground (a shade on a dark ground keeps 1.6:1)
        floor = 1.6 if _is_dark(ground_) else MIN_FILL_CONTRAST
        if _is_dark(ground_):
            # on a dark ground the parts step toward white: a shade toward the ground sinks into it (VK WorkSpace: a
            # navy slice on black at 1.6:1 read as a hole in the pie); each tint ≥ 2:1 off the ground — unless the
            # accent is already pale (Классика's beige on brown): its way to white is too short for distinct parts
            # (five slices of near-white), so they step toward the ground, each still ≥ 2:1 off it
            alt = _pole_shades(accent, ground_, DARK_SLICE_CONTRAST, m, poles=("FFFFFF",))
            alt_way = delta_e(accent, alt[-1]) if alt else 0.0
            if alt is not None and alt_way / m < 10.0:
                floor = DARK_SLICE_CONTRAST
                way, prev, share = 0.0, accent, 1.0
                while share > 0.1:
                    share = round(share - 0.02, 3)
                    c = _mix(accent, ground_, share)
                    if contrast_ratio(c, ground_) < floor:
                        break
                    way += delta_e(c, prev)
                    prev = c
                if way > 1.5 * alt_way:
                    alt = None  # the tints toward the ground below
            if alt is not None:
                for r, j in enumerate(order[1:]):
                    out[j] = (alt[r], 1.0)
                return out
        path: list[tuple[float, float]] = [(1.0, 0.0)]
        share = 1.0
        while share > 0.1:
            share = round(share - 0.02, 3)
            c = _mix(accent, ground_, share)
            if contrast_ratio(c, ground_) < floor:
                break
            path.append((share, path[-1][1] + delta_e(c, _mix(accent, ground_, share + 0.02))))
        total = path[-1][1]
        if total < POLE_SWITCH_DE and m >= 1:
            # the accent and the ground are near kin (a coral accent on a magenta slide): tints toward the ground all
            # look alike — the slices step toward white or black instead, whichever keeps them visible on the ground
            alt = _pole_shades(accent, ground_, floor, m)
            if alt is not None:
                for r, j in enumerate(order[1:]):
                    out[j] = (alt[r], 1.0)
                return out
        step = min(32.0, total / m) if total > 0 else 0.0
        for r, j in enumerate(order[1:]):
            target = step * (r + 1)
            sh = next((a for a, d in path if d >= target - 1e-6), path[-1][0])
            out[j] = (accent, sh)
    return out


POLE_SWITCH_DE = 40.0  # the way from the accent to the ground this short (ΔE) leaves no room for distinct tints
DARK_SLICE_CONTRAST = 2.0  # every part of a pie on a dark ground stands at least this far off it


def _pole_shades(accent: str, ground: str, floor: float, m: int, poles: tuple[str, ...] = ("FFFFFF", "000000")) -> Optional[list[str]]:
    """`m` solid colours stepping from the accent toward white or black (the pole with the longer visible way), each
    ≥ `floor`:1 off the ground and evenly spaced by ΔE (neighbours up to 32 apart); None when neither pole helps."""
    best: Optional[tuple[float, list[tuple[float, float]], str]] = None
    for pole in poles:
        path = [(0.0, 0.0)]
        t = 0.0
        while t < 0.9:
            t = round(t + 0.02, 3)
            c = _mix(accent, pole, 1.0 - t)  # the accent moved t of the way to the pole
            if contrast_ratio(c, ground) < floor:
                break
            path.append((t, path[-1][1] + delta_e(c, _mix(accent, pole, 1.0 - (t - 0.02)))))
        if len(path) > 1 and (best is None or path[-1][1] > best[0]):
            best = (path[-1][1], path, pole)
    if best is None or best[0] < 12.0:
        return None
    total, path, pole = best
    step = min(32.0, total / m)
    out = []
    for r in range(m):
        target = step * (r + 1)
        t = next((a for a, d in path if d >= target - 1e-6), path[-1][0])
        out.append(_mix(accent, pole, 1.0 - t))
    return out


def pie_palette(values: Sequence[Optional[float]], accent: str, ground: Optional[str], quiet_idx: Iterable[int] = (), text: Optional[str] = None) -> list[str]:
    """The solid colours of `pie_shades`: each base colour laid over the ground at its share."""
    ground_ = _norm(ground) or "FFFFFF"
    return [base if share >= 1.0 else _mix(base, ground_, share) for base, share in pie_shades(values, accent, ground_, quiet_idx, text)]


# ---------------------------------------------------------------------------------------------- theme & ground


def _theme_scheme(slide: Slide) -> dict[str, str]:
    """clrScheme of the slide's master theme, keyed dk1/lt1/…/accent6, plus the master's colour map (bg1→lt1 …)."""
    out: dict[str, str] = {}
    try:
        master = slide.slide_layout.slide_master
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    except Exception:  # noqa: BLE001
        return out
    scheme = theme.find(f".//{_a('clrScheme')}")
    if scheme is not None:
        for child in scheme:
            if not isinstance(child.tag, str) or not len(child):
                continue
            clr = child[0]
            val = clr.get("val") if _local(clr) == "srgbClr" else clr.get("lastClr")
            if _norm(val):
                out[_local(child)] = _norm(val)
    cmap = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
    try:
        cm = master._element.find(f"{{{_P}}}clrMap")
        if cm is not None:
            cmap.update({k: v for k, v in cm.attrib.items()})
    except Exception:  # noqa: BLE001
        pass
    for k, v in cmap.items():
        if v in out and k not in out:
            out[k] = out[v]
    return out


def _hsl_mod(hex_: str, lum_mod: float, lum_off: float, sat_mod: float = 1.0, hue_mod: float = 1.0) -> str:
    import colorsys

    r, g, b = (x / 255 for x in _rgb_t(hex_))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = max(0.0, min(1.0, l * lum_mod + lum_off))
    s = max(0.0, min(1.0, s * sat_mod))
    h = (h * hue_mod) % 1.0
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return f"{int(round(r * 255)):02X}{int(round(g * 255)):02X}{int(round(b * 255)):02X}"


_PRESET = {"white": "FFFFFF", "black": "000000", "red": "FF0000", "blue": "0000FF", "green": "008000", "gray": "808080", "grey": "808080"}


def _color_choice(el, scheme: dict[str, str]) -> Optional[tuple[str, float]]:
    """(hex, alpha) of an a:srgbClr / a:schemeClr / a:sysClr / a:prstClr element with its lumMod/lumOff/tint/shade."""
    if el is None:
        return None
    kind = _local(el)
    if kind == "srgbClr":
        hex_ = _norm(el.get("val"))
    elif kind == "schemeClr":
        hex_ = scheme.get(el.get("val") or "")
    elif kind == "sysClr":
        hex_ = _norm(el.get("lastClr")) or ("FFFFFF" if el.get("val") == "window" else "000000")
    elif kind == "prstClr":
        hex_ = _PRESET.get((el.get("val") or "").lower())
    else:
        hex_ = None
    if not hex_:
        return None
    alpha = 1.0
    lum_mod, lum_off, sat_mod, hue_mod = 1.0, 0.0, 1.0, 1.0
    for m in el:
        name, v = _local(m), m.get("val")
        try:
            f = int(v) / 100000 if v is not None else None
        except ValueError:
            f = None
        if f is None:
            continue
        if name == "alpha":
            alpha = f
        elif name == "lumMod":
            lum_mod = f
        elif name == "lumOff":
            lum_off = f
        elif name == "satMod":
            sat_mod = f
        elif name == "hueMod":
            hue_mod = f
        elif name == "tint":
            hex_ = _mix(hex_, "FFFFFF", f)
        elif name == "shade":
            hex_ = _mix(hex_, "000000", f)
    if lum_mod != 1.0 or lum_off != 0.0 or sat_mod != 1.0 or hue_mod != 1.0:
        hex_ = _hsl_mod(hex_, lum_mod, lum_off, sat_mod, hue_mod)
    return hex_, alpha


def _grad_paint(grad, scheme: dict[str, str], region=None, size=None) -> Optional[tuple[str, float]]:
    """(colour, opacity) a gradient fill paints over `region` = (x0, y0, x1, y1) fractions of its box, sampled on a
    4 × 4 grid there — by the geometry the audit measures with (`layers.gradient_position`): a background running from
    light cyan at the top to navy at the bottom is light under a heading at its top (Office «Circuit»), a white glow
    fading out from the bottom is nearly clear there (LibreOffice «Lights»). `size` = (w, h) of the filled box: the
    aspect a linear gradient's angle is drawn at. Without a region: the mean of its stops."""
    from verstka.rendering.layers import gradient_at, gradient_position

    stops = []
    for gs in grad.iter(_a("gs")):
        if not len(gs):
            continue
        cc = _color_choice(gs[0], scheme)
        if not cc:
            continue
        try:
            pos = int(gs.get("pos") or 0) / 100000.0
        except ValueError:
            pos = 0.0
        stops.append((pos, tuple(float(x) for x in _rgb_t(cc[0])), float(cc[1])))
    if not stops:
        return None
    stops.sort(key=lambda s: s[0])
    if region is None:
        rgb = [sum(s[1][i] for s in stops) / len(stops) for i in range(3)]
        return "".join(f"{int(round(x)):02X}" for x in rgb), sum(s[2] for s in stops) / len(stops)
    x0, y0, x1, y1 = (max(0.0, min(1.0, float(v))) for v in region)
    w, h = size if size and size[0] > 0 and size[1] > 0 else (1.0, 1.0)
    n = 4
    acc = [0.0, 0.0, 0.0]
    tot = 0.0
    for i in range(n):
        for j in range(n):
            u = x0 + (i + 0.5) * (x1 - x0) / n
            v = y0 + (j + 0.5) * (y1 - y0) / n
            rgb, a = gradient_at(stops, gradient_position(grad, u, v, w, h))
            for k in range(3):
                acc[k] += rgb[k] * a
            tot += a
    if tot <= 0:
        return None
    return "".join(f"{int(round(min(max(c / tot, 0.0), 255.0))):02X}" for c in acc), tot / (n * n)


def _fill_color(holder, scheme: dict[str, str], part=None, region=None, size=None) -> Optional[str]:
    """Visible colour of the fill inside `holder` (spPr / bgPr): solid, gradient (its colours where `region` lies,
    `_grad_paint`) or picture (mean of the region it shows). None for no fill, a fill mostly transparent there or an
    unreadable picture."""
    if holder is None:
        return None
    for child in holder:
        kind = _local(child)
        if kind == "noFill":
            return None
        if kind == "solidFill" and len(child):
            got = _color_choice(child[0], scheme)
            return got[0] if got and got[1] >= 0.5 else None
        if kind == "gradFill":
            got = _grad_paint(child, scheme, region, size)
            return got[0] if got is not None and got[1] >= 0.5 else None
        if kind == "pattFill":
            # a pattern reads as the blend of its two colours (half and half: the density is not modelled)
            fg, bg = child.find(_a("fgClr")), child.find(_a("bgClr"))
            cf = _color_choice(fg[0], scheme) if fg is not None and len(fg) else None
            cb = _color_choice(bg[0], scheme) if bg is not None and len(bg) else None
            if cf and cb:
                return _mix(cf[0], cb[0], 0.5)
            return (cf or cb or (None,))[0]
        if kind == "blipFill" and part is not None:
            blip = child.find(_a("blip"))
            rid = blip.get(f"{{{_R}}}embed") if blip is not None else None
            if rid:
                # a picture fill that is mostly transparent there (a grain or noise texture over a gradient) is no
                # ground: what lies under it is
                got = _picture_cover(part, rid, region)
                return got[0] if got is not None and got[1] >= 0.5 else None
    return None


_COVER_MEMO: dict[tuple, Optional[tuple[str, float]]] = {}


def _picture_cover(part, rid: str, region=None) -> Optional[tuple[str, float]]:
    """(mean colour of the picture's opaque pixels, their coverage 0–1) inside `region` = (x0, y0, x1, y1) fractions."""
    # (one picture is asked about the same region again and again — every text over a picture ground, at every render
    # of the deck: a template's background decoded once per question cost LCT a quarter of its render)
    try:
        blob = part.related_part(rid).blob
        reg = tuple(region) if region else None
        key = (len(blob), hash(blob), reg)
    except Exception:  # noqa: BLE001
        return None
    if key not in _COVER_MEMO:
        if len(_COVER_MEMO) >= 4096:
            _COVER_MEMO.clear()
        _COVER_MEMO[key] = _picture_cover_of(blob, reg)
    return _COVER_MEMO[key]


_THUMB_MEMO: dict[tuple[int, int], object] = {}
_THUMB_SIDE = 512


def _thumb(blob: bytes):
    """The picture decoded once, at most `_THUMB_SIDE` px a side, RGBA (None for an unreadable one): every text over a
    picture ground asks about its own region of the same picture — a 4K background decoded per question cost LCT ten
    seconds a generation."""
    key = (len(blob), hash(blob))
    got = _THUMB_MEMO.get(key, _THUMB_MEMO)
    if got is not _THUMB_MEMO:
        return got
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(blob))
        im.draft("RGB", (_THUMB_SIDE, _THUMB_SIDE))
        im = im.convert("RGBA")
        im.thumbnail((_THUMB_SIDE, _THUMB_SIDE))
    except Exception:  # noqa: BLE001
        im = None
    if len(_THUMB_MEMO) >= 48:
        _THUMB_MEMO.clear()
    _THUMB_MEMO[key] = im
    return im


def _picture_cover_of(blob: bytes, region=None) -> Optional[tuple[str, float]]:
    try:
        im = _thumb(blob)
        if im is None:
            return None
        w, h = im.size
        if region:
            x0, y0, x1, y1 = (max(0.0, min(1.0, v)) for v in region)
            if x1 - x0 > 0.01 and y1 - y0 > 0.01:
                im = im.crop((int(x0 * w), int(y0 * h), max(int(x0 * w) + 1, int(x1 * w)), max(int(y0 * h) + 1, int(y1 * h))))
        px = list(im.resize((24, 24)).getdata())
        wsum = sum(p[3] for p in px)
        cover = wsum / (255.0 * len(px))
        if wsum <= 0:
            return None
        rgb = [sum(p[i] * p[3] for p in px) / wsum for i in range(3)]
        return "".join(f"{int(round(x)):02X}" for x in rgb), cover
    except Exception:  # noqa: BLE001
        return None


def _picture_ground(part, rid: Optional[str], region, base) -> Optional[str]:
    """The colour a picture paints as a ground: None when it is mostly transparent there (a vapour trail, a glow — the
    ground is what lies under it), its opaque colour laid over `base()` (the slide's background) where partly see-through."""
    if not rid:
        return None
    got = _picture_cover(part, rid, region)
    if got is None:
        return None
    hx, cover = got
    if cover < 0.5:
        return None
    if cover >= 0.97:
        return hx
    under = base() or "FFFFFF"
    return _mix(hx, under, cover)


def _picture_mean(part, rid: str, region=None) -> Optional[str]:
    """Mean colour of a picture (over white where transparent), restricted to `region` = (x0, y0, x1, y1) fractions."""
    try:
        from PIL import Image

        im = _thumb(part.related_part(rid).blob)
        if im is None:
            return None
        w, h = im.size
        if region:
            x0, y0, x1, y1 = (max(0.0, min(1.0, v)) for v in region)
            if x1 - x0 > 0.01 and y1 - y0 > 0.01:
                im = im.crop((int(x0 * w), int(y0 * h), max(int(x0 * w) + 1, int(x1 * w)), max(int(y0 * h) + 1, int(y1 * h))))
        im = im.resize((24, 24))
        bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
        bg.alpha_composite(im)
        px = list(bg.convert("RGB").getdata())
        rgb = [sum(p[i] for p in px) / len(px) for i in range(3)]
        return "".join(f"{int(round(x)):02X}" for x in rgb)
    except Exception:  # noqa: BLE001
        return None


def _slide_size(slide: Slide) -> tuple[int, int]:
    try:
        prs = slide.part.package.presentation_part.presentation
        return int(prs.slide_width), int(prs.slide_height)
    except Exception:  # noqa: BLE001
        return 12192000, 6858000


def slide_ground(slide: Slide, bbox: Bbox) -> Optional[str]:
    """The colour a chart at `bbox` sits on: the topmost filled shape or picture under it that covers most of it,
    else the slide / layout / master background (solid, gradient or picture, resolved through the theme)."""
    scheme = _theme_scheme(slide)
    sw, sh = _slide_size(slide)
    cx, cy = bbox.x + bbox.w / 2, bbox.y + bbox.h / 2
    area = max(bbox.w * bbox.h, 1)
    try:
        shapes = list(slide.shapes)
    except Exception:  # noqa: BLE001
        shapes = []
    for shp in reversed(shapes):
        el = shp._element
        kind = _local(el)
        if kind == "grpSp":
            # a picture wrapped in a group (Canva and Google Slides exports group their backgrounds): its box in slide
            # coordinates through the group's scaling
            got = _group_picture_ground(slide, el, bbox, scheme)
            if got:
                return got
            continue
        if kind not in ("sp", "pic"):
            continue
        x, y, w, h = shp.left, shp.top, shp.width, shp.height
        if None in (x, y, w, h) or w <= 0 or h <= 0 or not (x <= cx <= x + w and y <= cy <= y + h):
            continue
        ox = max(0, min(x + w, bbox.x + bbox.w) - max(x, bbox.x))
        oy = max(0, min(y + h, bbox.y + bbox.h) - max(y, bbox.y))
        if ox * oy < 0.5 * area:
            continue
        region = ((bbox.x - x) / w, (bbox.y - y) / h, (bbox.x + bbox.w - x) / w, (bbox.y + bbox.h - y) / h)
        sppr = el.find(f"{{{_P}}}spPr")
        if kind == "pic":
            bf = el.find(f"{{{_P}}}blipFill")
            blip = bf.find(_a("blip")) if bf is not None else None
            rid = blip.get(f"{{{_R}}}embed") if blip is not None else None
            got = _picture_ground(slide.part, rid, region, lambda: _template_ground(slide, bbox, scheme) or _bg_color(slide, bbox, scheme))
        else:
            got = _fill_color(sppr, scheme, slide.part, region, (w, h))
            if got is None and sppr is not None and not any(_local(c) in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill") for c in sppr):
                ref = el.find(f"{{{_P}}}style/{_a('fillRef')}")
                if ref is not None and ref.get("idx") not in (None, "0") and len(ref):
                    cc = _color_choice(ref[0], scheme)
                    got = cc[0] if cc and cc[1] >= 0.5 else None
        if got:
            return got
    # what the layout and the master paint under the box (a band, a panel, a full-bleed picture), topmost first
    got = _template_ground(slide, bbox, scheme)
    if got:
        return got
    return _bg_color(slide, bbox, scheme)


def _group_picture_ground(slide: Slide, grp, bbox: Bbox, scheme: dict[str, str]) -> Optional[str]:
    """The ground a group paints under `bbox`: its topmost picture (at any depth) holding the box's centre and
    covering half of it, in slide coordinates."""
    cx, cy = bbox.x + bbox.w / 2, bbox.y + bbox.h / 2
    area = max(bbox.w * bbox.h, 1)

    def to_slide(el, box):
        x, y, w, h = box
        parent = el.getparent()
        while parent is not None and _local(parent) == "grpSp":
            xf = parent.find(f"{{{_P}}}grpSpPr/{_a('xfrm')}")
            if xf is None:
                return None
            off, ext, choff, chext = xf.find(_a("off")), xf.find(_a("ext")), xf.find(_a("chOff")), xf.find(_a("chExt"))
            if off is None or ext is None:
                return None
            gx, gy, gw, gh = float(off.get("x") or 0), float(off.get("y") or 0), float(ext.get("cx") or 0), float(ext.get("cy") or 0)
            ox, oy = (float(choff.get("x") or 0), float(choff.get("y") or 0)) if choff is not None else (gx, gy)
            ow, oh = (float(chext.get("cx") or 0), float(chext.get("cy") or 0)) if chext is not None else (gw, gh)
            sx, sy = (gw / ow if ow else 1.0), (gh / oh if oh else 1.0)
            x, y, w, h = gx + (x - ox) * sx, gy + (y - oy) * sy, w * sx, h * sy
            parent = parent.getparent()
        return x, y, w, h

    pics = [e for e in grp.iter(f"{{{_P}}}pic")]
    for el in reversed(pics):
        xf = el.find(f"{{{_P}}}spPr/{_a('xfrm')}")
        off = xf.find(_a("off")) if xf is not None else None
        ext = xf.find(_a("ext")) if xf is not None else None
        if off is None or ext is None:
            continue
        box = to_slide(el, (float(off.get("x") or 0), float(off.get("y") or 0), float(ext.get("cx") or 0), float(ext.get("cy") or 0)))
        if box is None:
            continue
        x, y, w, h = box
        if w <= 0 or h <= 0 or not (x <= cx <= x + w and y <= cy <= y + h):
            continue
        ox = max(0.0, min(x + w, bbox.x + bbox.w) - max(x, bbox.x))
        oy = max(0.0, min(y + h, bbox.y + bbox.h) - max(y, bbox.y))
        if ox * oy < 0.5 * area:
            continue
        blip = el.find(f"{{{_P}}}blipFill/{_a('blip')}")
        rid = blip.get(f"{{{_R}}}embed") if blip is not None else None
        region = ((bbox.x - x) / w, (bbox.y - y) / h, (bbox.x + bbox.w - x) / w, (bbox.y + bbox.h - y) / h)
        got = _picture_ground(slide.part, rid, region, lambda: _template_ground(slide, bbox, scheme) or _bg_color(slide, bbox, scheme))
        if got:
            return got
    return None


def _bg_color(slide: Slide, bbox: Bbox, scheme: dict[str, str]) -> Optional[str]:
    """The background the slide, its layout or its master sets (solid, gradient or picture, through the theme)."""
    sw, sh = _slide_size(slide)
    region = (bbox.x / sw, bbox.y / sh, (bbox.x + bbox.w) / sw, (bbox.y + bbox.h) / sh)
    holders = []
    try:
        holders = [(slide._element, slide.part), (slide.slide_layout._element, slide.slide_layout.part), (slide.slide_layout.slide_master._element, slide.slide_layout.slide_master.part)]
    except Exception:  # noqa: BLE001
        holders = [(slide._element, slide.part)]
    for el, part in holders:
        bg = el.find(f"{{{_P}}}cSld/{{{_P}}}bg")
        if bg is None:
            continue
        pr = bg.find(f"{{{_P}}}bgPr")
        if pr is not None:
            got = _fill_color(pr, scheme, part, region, (sw, sh))
            if got:
                return got
            continue
        ref = bg.find(f"{{{_P}}}bgRef")
        if ref is not None and len(ref):
            cc = _color_choice(ref[0], scheme)
            if cc:
                return _bg_style_fill(slide, ref, cc[0], scheme, part, region, (sw, sh)) or cc[0]
    return scheme.get("bg1") or None


def _bg_style_fill(slide: Slide, ref, ph_hex: str, scheme: dict[str, str], part, region, size) -> Optional[str]:
    """What a `p:bgRef idx="1001…"` paints: the theme's background fill style of that index (a:bgFillStyleLst), its
    placeholder colour (phClr) the ref's own colour — Office themes paint their grounds so (Berlin: an orange-to-red
    gradient built from bg2, measured where the box is). None for an index outside the list."""
    try:
        idx = int(ref.get("idx") or 0)
    except ValueError:
        return None
    if idx < 1001:
        return None
    try:
        theme = etree.fromstring(slide.slide_layout.slide_master.part.part_related_by(RT.THEME).blob)
    except Exception:  # noqa: BLE001
        return None
    lst = theme.find(f".//{_a('bgFillStyleLst')}")
    fills = [c for c in lst if isinstance(c.tag, str)] if lst is not None else []
    if idx - 1001 >= len(fills):
        return None
    holder = etree.Element("holder")
    holder.append(copy.deepcopy(fills[idx - 1001]))
    return _fill_color(holder, {**scheme, "phClr": ph_hex}, part, region, size)


_OWN_PARTS: "weakref.WeakSet" = weakref.WeakSet()
_TEMPLATE_COLORS: "weakref.WeakKeyDictionary" = weakref.WeakKeyDictionary()


def _template_ground(slide: Slide, bbox: Bbox, scheme: dict[str, str]) -> Optional[str]:
    """The colour the slide's layout (then its master, unless a showMasterSp="0" hides it) paints under `bbox`: the
    topmost non-placeholder shape or picture that holds the box's centre and covers half of it (contract C4)."""
    cx, cy = bbox.x + bbox.w / 2, bbox.y + bbox.h / 2
    area = max(bbox.w * bbox.h, 1)
    try:
        layout = slide.slide_layout
        holders = [layout]
        hide = slide._element.get("showMasterSp") == "0" or layout._element.get("showMasterSp") == "0"
        if not hide:
            holders.append(layout.slide_master)
    except Exception:  # noqa: BLE001
        return None
    for holder in holders:
        tree = holder._element.find(f"{{{_P}}}cSld/{{{_P}}}spTree")
        if tree is None:
            continue
        for el in reversed(list(tree)):
            kind = _local(el)
            if kind not in ("sp", "pic") or el.find(f".//{{{_P}}}ph") is not None:
                continue
            sppr = el.find(f"{{{_P}}}spPr")
            xfrm = sppr.find(_a("xfrm")) if sppr is not None else None
            off = xfrm.find(_a("off")) if xfrm is not None else None
            ext = xfrm.find(_a("ext")) if xfrm is not None else None
            if off is None or ext is None:
                continue
            try:
                x, y, w, h = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
            except (TypeError, ValueError):
                continue
            if w <= 0 or h <= 0 or not (x <= cx <= x + w and y <= cy <= y + h):
                continue
            ox = max(0, min(x + w, bbox.x + bbox.w) - max(x, bbox.x))
            oy = max(0, min(y + h, bbox.y + bbox.h) - max(y, bbox.y))
            if ox * oy < 0.5 * area:
                continue
            region = ((bbox.x - x) / w, (bbox.y - y) / h, (bbox.x + bbox.w - x) / w, (bbox.y + bbox.h - y) / h)
            if kind == "pic":
                bf = el.find(f"{{{_P}}}blipFill")
                blip = bf.find(_a("blip")) if bf is not None else None
                rid = blip.get(f"{{{_R}}}embed") if blip is not None else None
                got = _picture_ground(holder.part, rid, region, lambda: _bg_color(slide, bbox, scheme))
            else:
                got = _fill_color(sppr, scheme, holder.part, region, (w, h))
                if got is None and sppr is not None and not any(_local(c) in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill") for c in sppr):
                    ref = el.find(f"{{{_P}}}style/{_a('fillRef')}")
                    if ref is not None and ref.get("idx") not in (None, "0") and len(ref):
                        cc = _color_choice(ref[0], scheme)
                        got = cc[0] if cc and cc[1] >= 0.5 else None
            if got:
                return got
    return None


def template_chart_colors(slide: Slide) -> list[str]:
    """Series colours of the template's own chart parts, in order: fills (or line colours) given as schemeClr and
    resolved through the theme, or as srgbClr. Charts this module made are skipped."""
    try:
        pkg = slide.part.package
    except Exception:  # noqa: BLE001
        return []
    cached = _TEMPLATE_COLORS.get(pkg)
    if cached is not None:
        return list(cached)
    scheme = _theme_scheme(slide)
    out: list[str] = []
    try:
        parts = list(pkg.iter_parts())
    except Exception:  # noqa: BLE001
        parts = []
    for part in parts:
        if part.content_type != CT.DML_CHART or part in _OWN_PARTS:
            continue
        try:
            root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
        except Exception:  # noqa: BLE001
            continue
        for ser in root.iter(_c("ser")):
            sp = ser.find(_c("spPr"))
            if sp is None:
                continue
            fill = sp.find(_a("solidFill"))
            if fill is None:
                fill = sp.find(f"{_a('ln')}/{_a('solidFill')}")
            if fill is None or not len(fill):
                continue
            got = _color_choice(fill[0], scheme)
            if got and got[0] not in out:
                out.append(got[0])
    _TEMPLATE_COLORS[pkg] = list(out)
    return out


def _theme_accents(slide: Slide) -> list[str]:
    scheme = _theme_scheme(slide)
    return [scheme[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in scheme]


def _brand_theme_accents(slide: Slide) -> list[str]:
    """The theme's accents less the stock ones every office suite writes into a new file (LibreOffice green, Office
    blue): in a template that does not draw them they are no colour of its brand."""
    from verstka.analysis.colors import is_stock_theme_color

    return [c for c in _theme_accents(slide) if not is_stock_theme_color(_norm(c) or "")]


# ---------------------------------------------------------------------------------------------- numbers & units


def number_format(unit: Optional[str]) -> str:
    u = (unit or "").strip()
    if u == "%":
        return '0"%"'
    if u in ("₽", "руб", "руб.", "RUB"):
        return '#,##0" ₽"'
    if u in ("$", "USD"):
        return '"$"#,##0'
    if u:
        return f'#,##0" {u}"'
    return "#,##0"


_UNIT_STOP = {"в", "на", "за", "per", "/", "в/", "за/", "от", "по"}
_GLYPH_UNITS = {"%", "₽", "$", "€", "руб", "руб.", "RUB", "USD", "EUR"}


def short_unit(unit: Optional[str], limit: int = 7) -> str:
    """The unit as a label suffix: «тыс. обращений» → «тыс.», «млн ₽ в год» → «млн ₽». The full unit belongs in
    one caption (see `unit_caption`), never on every label."""
    u = " ".join((unit or "").split())
    if len(u) <= limit:
        return u
    toks = u.split(" ")
    out = toks[0]
    for tok in toks[1:]:
        if tok.lower() in _UNIT_STOP or len(out) + 1 + len(tok) > limit:
            break
        out = f"{out} {tok}"
    return out


def glyph_unit(unit: Optional[str]) -> bool:
    """A one-glyph unit (%, ₽, $, €) rides on every label; a word unit is written once."""
    u = " ".join((unit or "").split())
    return u in _GLYPH_UNITS or short_unit(u) in _GLYPH_UNITS


def bare_unit_on_labels(spec: ChartSpec, outline: DeckOutline, unit: Optional[str]) -> bool:
    """A chart with no title of its own and a short noun for a unit (5–12 letters: «посещений», «клиентов») over a
    few categories (≤ 6) carries its unit on its labels — «2 700 посещений» — rather than under a caption of
    the unit alone, which reads as a stray word over the bars."""
    u = " ".join((unit or "").split())
    # (a noun in full — «посещений», «клиентов» — reads as a stray word over the bars; an abbreviation «чел.», «тыс.»
    # or a phrase «млн ₽ в год» is an axis unit and stays the caption above)
    if not u or glyph_unit(u) or not re.fullmatch(r"[А-Яа-яЁёA-Za-z]{5,12}", u) or spec.type in ("pie", "doughnut", "line"):
        return False
    series = resolve_series(spec, outline)
    return bool(series) and len(series) == 1 and 0 < len(series[0].values) <= 6


def unit_caption(spec: ChartSpec, outline: DeckOutline) -> Optional[str]:
    """A word unit («тыс. обращений», «чел.») for a caption placed once above the chart, so that every label is a
    plain number of one scale (a unit on one label only reads as «12 against 58 thousand»). One-glyph units (%, ₽)
    ride on the labels and need no caption. A chart whose caller puts this text right above it (or in its title)
    writes no unit caption of its own."""
    if spec.type in ("pie", "doughnut"):
        return None
    series = resolve_series(spec, outline)
    if not series:
        return None  # no chart, no caption: a unit over an empty frame is a stray word
    unit = spec.unit or series[0].unit
    u = " ".join((unit or "").split())
    return u if u and not glyph_unit(u) else None


def pie_values_are_percents(unit: Optional[str], total: float) -> bool:
    """A pie whose values are percents already — its unit is «%», or it has no unit and its values add up to
    100 ± 1 (a writer's «Доходы по направлениям, %» with the unit lost): each part is read as its value, never as
    «42 · 42%» (G4-11)."""
    u = short_unit(unit)
    return u == "%" or (not (unit or "").strip() and abs(total - 100.0) <= 1.0)


def _decimals(values: Iterable[Optional[float]]) -> int:
    vals = [abs(float(v)) for v in values if v is not None]
    if not vals or all(abs(v - round(v)) < 1e-6 for v in vals):
        return 0
    mx = max(vals)
    if mx >= 100:
        return 0
    if mx < 10 and any(abs(v * 10 - round(v * 10)) > 1e-6 for v in vals):
        return 2
    return 1


def _unit_affix(unit: Optional[str]) -> tuple[str, str]:
    u = short_unit(unit)
    if not u:
        return "", ""
    if u == "%":
        return "", "%"
    if u in ("₽", "руб", "руб.", "RUB"):
        return "", " ₽"
    if u in ("$", "USD"):
        return "$", ""
    if u in ("€", "EUR"):
        return "", " €"
    return "", f" {u}"


def label_format(unit: Optional[str], values: Iterable[Optional[float]] = ()) -> str:
    """Excel number format for data labels: thousands grouped, decimals only when the data has them, the short unit
    as a suffix (non-breaking, so a label never wraps), and an empty zero section so a zero prints nothing."""
    d = _decimals(values)
    u = short_unit(unit)
    num = ("0" if u == "%" else "#,##0") + ("." + "0" * d if d else "")
    pre, suf = _unit_affix(unit)
    q = lambda s: f'"{s}"' if s else ""  # noqa: E731
    pos = f"{q(pre)}{num}{q(suf)}"
    return f'{pos};"−"{pos};'  # a true minus sign, and an empty zero section


def _tick_format(unit: Optional[str], decimals: int) -> str:
    """Value-axis ticks: plain numbers; only a one-glyph unit (%, ₽, $) rides along — words never repeat per tick."""
    num = "#,##0" + ("." + "0" * decimals if decimals else "")
    u = short_unit(unit)
    if u == "%":
        return num + '"%"'
    if u in ("₽", "руб", "руб.", "RUB"):
        return num + '" ₽"'
    if u in ("$", "USD"):
        return '"$"' + num
    return num


def tick_decimals(step: Optional[float], decimals: int) -> int:
    """The decimals a value axis prints: as many as its step has, not the data's — steps of 10 over 41,2…63,9 read
    «40 · 50 · 60», never «40,0 · 50,0 · 60,0» (the data's decimals stay on the data labels). Without a step (the
    renderer picks the ticks): the data's."""
    if not step:
        return decimals
    frac = f"{abs(step):.6f}".rstrip("0").split(".")[1]
    return min(len(frac), 3)


def _tick_text(v: Optional[float], unit: Optional[str], decimals: int) -> str:
    """What a value-axis tick prints (`_tick_format`): «1 200 000 ₽»."""
    s = f"{abs(v or 0.0):,.{decimals}f}".replace(",", " ").replace(".", ",")
    s = ("−" if (v or 0.0) < 0 else "") + s
    u = short_unit(unit)
    if u == "%":
        return s + "%"
    if u in ("₽", "руб", "руб.", "RUB"):
        return s + " ₽"
    if u in ("$", "USD"):
        return "$" + s
    return s


def _axis_room(values: Iterable[float], unit: Optional[str], decimals: int, family: Optional[str], fs: float, gap: float) -> float:
    """The room left of a plot for its value-axis ticks: the widest tick as printed and 0.45 of a size; a family
    without «₽» has its whole axis set by the renderer in the face that has it (Trebuchet MS → Noto Sans, 4 % wider)
    and LibreOffice cuts a tick once it runs within one average character of the room («1 200 0…» at 82 pt for a
    76 pt tick): that face's width, a tenth more and 0.3 of a size. Never less than the plain number and `gap` sizes,
    the room it always kept."""
    from verstka.rendering.fonts import glyph_standin

    vals = list(values) or [0.0]
    texts = [_tick_text(v, unit, decimals) for v in vals]
    face = None
    if any("₽" in t for t in texts):
        try:
            face = glyph_standin(family, "₽")
        except Exception:  # noqa: BLE001
            face = None
    if face:
        need = max(text_width_pt(t, face, fs) for t in texts) * 1.1 + fs * 0.3
    else:
        need = max(text_width_pt(t, family, fs) for t in texts) + fs * 0.45
    plain = max(text_width_pt(_fmt_value(v, None, decimals) or "0", family, fs) for v in vals)
    return max(need, plain + fs * gap)


def _fmt_value(v: Optional[float], unit: Optional[str], decimals: int) -> str:
    """What a label prints (for measuring): 12 400 чел., 4,8 %."""
    if v is None or v == 0:
        return ""
    s = f"{abs(v):,.{decimals}f}".replace(",", " ").replace(".", ",")
    pre, suf = _unit_affix(unit)
    sign = "−" if v < 0 else ""
    return f"{sign}{pre}{s}{suf}"


def _nice_scale(lo: float, hi: float) -> tuple[float, float, float]:
    """(min, max, major unit) of a quiet value axis: a round step giving 4–5 intervals, the ends snapped to it."""
    if hi <= lo:
        hi = lo + (abs(lo) or 1.0)
    best = None
    mag = 10 ** math.floor(math.log10((hi - lo) / 4.5))
    for m in (0.1, 1, 10):
        for base in (1, 2, 2.5, 5):
            step = base * mag * m
            a = math.floor(lo / step + 1e-9) * step
            b = math.ceil(hi / step - 1e-9) * step
            k = round((b - a) / step)
            if k < 3 or k > 7:
                continue
            cost = abs(k - 4.5) + (0.3 if base == 2.5 else 0.0)
            if best is None or cost < best[0]:
                best = (cost, a, b, step)
    if best is None:
        return lo, hi, (hi - lo) / 4
    _, a, b, step = best
    return a, b, step


# ---------------------------------------------------------------------------------------------- decisions


def inline_series(spec: ChartSpec) -> list[Series]:
    """The data the slide designer wrote into the chart itself (Agent v2: `categories` + `series`), as registry-like
    Series: values cut to the categories, a missing tail padded with None; series without a value are left out."""
    cats = [str(c) for c in (spec.categories or []) if str(c).strip()]
    if not cats:
        return []
    out: list[Series] = []
    for i, s in enumerate(spec.series or []):
        vals = [(float(v) if v is not None else None) for v in list(s.values or [])[: len(cats)]]
        if not any(v is not None for v in vals):
            continue
        name = (s.name or spec.title or "").strip()
        if len(vals) == len(cats) and all(v is not None for v in vals):
            out.append(Series(id=f"inline_{i + 1}", name=name, categories=cats, values=vals, unit=spec.unit))
        else:  # a short series: its missing tail is a gap in the chart, not a zero
            vals += [None] * (len(cats) - len(vals))
            out.append(Series.model_construct(id=f"inline_{i + 1}", name=name, categories=cats, values=vals, unit=spec.unit, source_span=None))
    return out


def resolve_series(spec: ChartSpec, outline: DeckOutline) -> list[Series]:
    """The series a chart draws: the registry series its ids name; else the data written into the chart itself
    (`categories` + `series`); else — ids that name nothing and no data of its own — the deck's first series, unless
    the ids name facts: a chart «of f1, f2» asks for those figures, and another series of the deck would be someone
    else's data."""
    if spec.type == "scatter":
        return []  # a scatter draws its (x, y) points (`xy_series`), never a category series of the deck
    out = [s for sid in spec.series_ids if (s := outline.series_by_id(sid)) is not None]
    if not out:
        out = inline_series(spec)
    if not out and outline.series and not (spec.categories or spec.series) and not any(outline.fact_by_id(sid) is not None for sid in spec.series_ids):
        out = outline.series[:1]
    return out


def xy_series(spec: ChartSpec) -> list[tuple[str, list[float], list[float]]]:
    """The points of a scatter as (name, xs, ys): pairs where both coordinates are numbers, series without one left out."""
    out = []
    for s in spec.xy or []:
        pairs = [(float(x), float(y)) for x, y in zip(s.x or [], s.y or []) if x is not None and y is not None and math.isfinite(float(x)) and math.isfinite(float(y))]
        if pairs:
            out.append(((s.name or spec.title or "").strip(), [p[0] for p in pairs], [p[1] for p in pairs]))
    return out


def last_value_index(values: Sequence[Optional[float]]) -> Optional[int]:
    """The last point of a series that has a value: a line with a missing tail ends (and is labelled) there."""
    return next((j for j in range(len(values) - 1, -1, -1) if values[j] is not None), None)


_SOURCE_HEAD_RE = re.compile(r"(?i)^\s*(источник|источники|данные|по данным|source)\b")


def source_text(source: Optional[str]) -> Optional[str]:
    """The source line under a chart: «Источник: Выручка по месяцам.csv» (a line that names itself as a source stays
    as written)."""
    s = " ".join((source or "").split())
    if not s:
        return None
    return s if _SOURCE_HEAD_RE.match(s) else f"Источник: {s}"


def _source_key(text: str) -> str:
    t = _SOURCE_HEAD_RE.sub("", " ".join((text or "").replace(" ", " ").split()).lower()).strip(" :—-–.")
    return t.replace("ё", "е")


def source_said(source: Optional[str], *texts: Optional[str]) -> bool:
    """The caller already prints the chart's source on the slide (its footnote «Источник: …» names it): the chart
    draws no second source line."""
    key = _source_key(source or "")
    if not key:
        return True
    return any(key in _source_key(t or "") for t in texts if t)


def chart_data_ok(spec: Optional[ChartSpec], outline: DeckOutline) -> bool:
    """True when `add_chart` has something to draw: a series with categories and at least one value (a scatter: one
    (x, y) point)."""
    if spec is None:
        return False
    if spec.type == "scatter":
        return bool(xy_series(spec))
    series = resolve_series(spec, outline)
    if not series or not series[0].categories:
        return False
    n = len(series[0].categories)
    return any(v is not None for s in series for v in list(s.values)[:n])


def prefer_bar(spec: ChartSpec, outline: DeckOutline) -> bool:
    """A bar/column chart (grouped or stacked too) reads better as horizontal bars: a category label longer than 12
    characters, or more than 7 categories (short labels such as years or quarters fit columns up to 12). A histogram's
    bins always stand as columns."""
    t = _GROUPED.get(spec.type, spec.type)
    if t not in ("column", "bar", "stacked_column", "stacked_bar"):
        return False
    if t in _HORIZONTAL:
        return True
    if getattr(spec, "kind_locked", False):
        return False  # chosen on purpose (the data's option, the person's «Пересобрать слайд так»): columns stay columns
    series = resolve_series(spec, outline)
    if not series:
        return False
    cats = [str(c) for c in series[0].categories]
    longest = max((len(c.strip()) for c in cats), default=0)
    # two columns (before / after) are wide: a label of a few words wraps under its column in two lines
    if longest > (PAIR_LABEL_CHARS if len(cats) <= 2 else LONG_LABEL_CHARS):
        return True
    return len(cats) > MAX_COLUMNS and (longest > 5 or len(cats) > 12)


def effective_chart_type(spec: ChartSpec, outline: DeckOutline) -> str:
    """The kind drawn: grouped bars are the clustered bars of several series; columns (and stacked columns) with long
    or many labels turn horizontal."""
    t = _GROUPED.get(spec.type, spec.type)
    if t in _STACKED:
        return "stacked_bar" if prefer_bar(spec, outline) else "stacked_column"
    if t in ("column", "bar"):
        return "bar" if prefer_bar(spec, outline) else "column"
    return t if t in _TYPE_MAP else "column"


_OTHER_RE = re.compile(r"^\s*(проч|други|другое|остальн|ины[ея]|иное|other|rest\b|misc)", re.I)
_PLAN_RE = re.compile(r"(план|цел[ьи]\b|прогноз|ожидан|таргет|target|plan\b|goal|forecast|budget|бюджет)", re.I)
_NOW_RE = re.compile(r"(стало|после|факт|сейчас|текущ|\bнов|\bwith\b|after|actual|\bnow\b|current)", re.I)
_PAST_RE = re.compile(r"(было|\bдо\b|\bбез\b|прошл|стар|without|before|previous|prior|baseline|базов)", re.I)
_YEAR_RE = re.compile(r"(?<!\d)(19\d\d|20\d\d)(?!\d)")


_MONTH_SHORT = {
    "январ": "янв", "феврал": "фев", "март": "мар", "апрел": "апр", "июн": "июн", "июл": "июл",
    "август": "авг", "сентябр": "сен", "октябр": "окт", "ноябр": "ноя", "декабр": "дек",
}
_PERIOD_WORDS = (
    (re.compile(r"(?i)\bмесяц(?:а|ев|ы)?\b"), "мес."),
    (re.compile(r"(?i)\bквартал(?:а|ов|ы)?\b"), "кв."),
    (re.compile(r"(?i)\bнедел(?:я|и|ь)\b"), "нед."),
    (re.compile(r"(?i)\bполугоди(?:е|я)\b"), "полуг."),
)
_MONTH_WORD_RE = re.compile(r"(?i)\b(январ|феврал|март|апрел|ма|июн|июл|август|сентябр|октябр|ноябр|декабр)[а-яё]*\b")


def short_categories(cats: Sequence[str], joiner: bool = True) -> list[str]:
    """Category labels of a time axis set short, for a chart whose slots are too narrow for them: «1-й месяц» →
    «1-й мес.», «Месяц 3» → «Мес. 3», «Январь 2026» → «Янв 2026», «2 квартал» → «2 кв.». Only the period words are
    shortened (the figures and the other words stay), and only in the chart: the plan keeps its words. `joiner=False`:
    no word joiner after a hyphen (U+2060 is not rendered by every reader: a box or a break instead of nothing)."""
    out = []
    for c in cats:
        t = " ".join(str(c).split())
        for rx, rep in _PERIOD_WORDS:
            t = rx.sub(lambda m, rep=rep: rep if m.group(0)[:1].islower() else rep[:1].upper() + rep[1:], t)

        def month(m: re.Match) -> str:
            word = m.group(0)
            if m.group(1).lower() == "ма":
                return word  # «май», «мая» are short already («март» is caught above; «макет» is no month)
            short = _MONTH_SHORT.get(m.group(1).lower(), word)
            return short[:1].upper() + short[1:] if word[:1].isupper() else short

        t = _MONTH_WORD_RE.sub(month, t)
        # a shortened label is one word for the renderer: «1-й мес.» never breaks after «1-» or «1-й»
        if t != " ".join(str(c).split()):
            t = t.replace(" ", "\u00a0")
            if joiner:
                t = t.replace("-", "-\u2060")
        out.append(t)
    return out


_ORDINAL_PERIOD_RE = re.compile(r"(?i)(\d+)\s*[-‑–]\s*(?:й|я|е|ой|ый|ий|ая|ое|го|м)\s+(?=(?:мес|кв|нед|полуг|год|г)\.?(?:\s|$))")

# LibreOffice sets a category label (VCartesianAxis): straight labels wrap at 95 % of the tick distance (the plot's
# width over the categories) — a word wider than that is broken inside («1-й ме / с.»), whatever tickLblSkip or blank
# neighbours say; a turned label never wraps but is cut to «…» when its turned height and one average character pass
# the room under the axis (the frame's bottom). Measured in bfix6/lo (e1–e3): 0.95 fits, 0.97 wraps.
LO_LABEL_ROOM = 0.95
CAT_LINE = 1.25  # a category label's line height over its size
TURN_SIN = 0.7071  # 45°


_DIGIT_GROUP_RE = re.compile(r"(?<=\d)[ \u00a0\u202f](?=\d{3}(?!\d))")


def glue_numbers(label: str) -> str:
    """A category label with the digit groups of its numbers bound by no-break spaces («1 250–1 500» is two numbers,
    never four words): a renderer breaks such a label after the dash or before a word, never inside a number."""
    return _DIGIT_GROUP_RE.sub("\u00a0", label)


def label_pieces(label: str) -> list[str]:
    """The pieces a renderer may break a category label between (ordinary spaces, and after the dash of a range;
    never at a no-break space, so never inside a number): each is as wide as the label's narrowest possible line."""
    return [p for p, _ in break_pieces(label, dashes=True)]


def period_levels(cats: Sequence[str]) -> list[list[str]]:
    """A time axis's labels at each shortening level: as written; the period word shortened («1-й месяц» → «1-й мес.»,
    «Январь 2026» → «Янв 2026», «2 квартал» → «2 кв.»); and the ordinal dropped as well («1-й мес.» → «1 мес.»). Plain
    spaces and no joiners: a label that must wrap after all breaks between its words, never inside one. Levels that
    change nothing are left out (the first one is always there)."""
    base = [glue_numbers(" ".join(str(c).replace("⁠", "").replace(" ", " ").split())) for c in cats]
    one = [glue_numbers(c.replace(" ", " ")) for c in short_categories(base, joiner=False)]
    two = [_ORDINAL_PERIOD_RE.sub(r"\1 ", c) for c in one]
    out = [base]
    for lv in (one, two):
        if lv != out[-1]:
            out.append(lv)
    return out


@dataclass
class CatFit:
    """How a category axis sets its labels (`fit_categories`)."""

    labels: list  # the labels written into the chart («» for one left out)
    mode: str  # "flat" (one line each), "wrap" (whole words on ≤ 2 lines), "turn" (45°)
    size: float  # their size
    lines: int = 1
    keep: int = 1  # every k-th turned label is shown (turned labels closer than their line height)
    below: float = 0.0  # pt the labels take under the axis
    longest: float = 0.0  # the widest label shown (unturned), pt
    level: int = 0  # shortening level of `period_levels`


def _cat_room(family: Optional[str]) -> float:
    """Share of the tick distance a label may take: LibreOffice's 95 %, less half a percent for a face measured with its
    own metrics, less 8 % for one estimated with a stand-in."""
    from verstka.rendering.fonts import is_measured

    return LO_LABEL_ROOM * (0.995 if is_measured(family) else 0.92)


# a long single word must keep a margin inside its room: it cannot wrap, it breaks inside («Автоматически» measured
# 0.4 pt inside a 63 pt slot still broke as «Автоматическ / и» in the render; «Сейчас» at 94 % of its slot did not —
# the renderer's own measure drifts with the word's length)
WORD_SAFETY = 0.965
LONG_WORD = 9


def _word_room(word: str, lim: float) -> float:
    return lim * WORD_SAFETY if len(word) >= LONG_WORD else lim


def fit_categories(cats: Sequence[str], family: Optional[str], sizes: Sequence[float], slot: float, below_max: float, *, turn_first: bool = True) -> CatFit:
    """Category labels measured one by one at the axis font against their slot (`slot` = the plot's width over the
    categories, pt; `below_max` = the most room under the axis the labels may take). In order, at each size of
    `sizes` (the axis size first, smaller ones after): every label on one line as written, then shortened
    (`period_levels`); then turned by 45° (shortest level; every k-th label when turned neighbours would touch) or
    wrapped between whole words on two lines — `turn_first` picks which is tried first. Nothing is ever left for the
    renderer to break inside a word: the last resort is the shortest labels turned at the smallest size, as many shown
    as keep apart."""
    levels = period_levels(cats)
    n = max(1, len(cats))
    room = _cat_room(family)
    last = levels[-1]

    def width(t: str, size: float) -> float:
        return text_width_pt(t, family, size) if t else 0.0

    def turned(size: float, force: bool = False) -> Optional[CatFit]:
        widths = [width(t, size) for t in last]
        # LibreOffice cuts a turned label once its turned height + one average character pass the room under the axis
        need = max((TURN_SIN * (w + CAT_LINE * size) + w / max(1, len(t)) for t, w in zip(last, widths) if t), default=0.0) + 0.35 * size
        if need > below_max and not force:
            return None
        keep = max(1, math.ceil(CAT_LINE * size / TURN_SIN / max(slot, 1.0) - 1e-6))
        labels = [t if (n - 1 - j) % keep == 0 else "" for j, t in enumerate(last)] if keep > 1 else list(last)
        return CatFit(labels, "turn", size, 1, keep, need, max(widths, default=0.0), len(levels) - 1)

    def wrapped(size: float) -> Optional[CatFit]:
        lim = slot * room
        for lv, labs in enumerate(levels):
            if any(width(w, size) > _word_room(w, lim) for t in labs for w in label_pieces(t)):
                continue  # a word (a number, a range's half) wider than its slot would be broken inside
            lines = max((len(wrap_lines(t, family, size, False, lim, dashes=True)) for t in labs if t), default=1)
            below = lines * size * CAT_LINE + size * 0.6
            if lines <= 2 and below <= below_max:
                return CatFit(list(labs), "wrap", size, lines, 1, below, max((width(t, size) for t in labs), default=0.0), lv)
        return None

    for size in sizes:
        lim = slot * room
        for lv, labs in enumerate(levels):
            longest = max((width(t, size) for t in labs), default=0.0)
            # (a long single word keeps a margin: it cannot wrap, a renderer breaks it inside — `_word_room`)
            if all(width(t, size) <= (_word_room(t, lim) if len(label_pieces(t)) == 1 else lim) for t in labs if t):
                return CatFit(list(labs), "flat", size, 1, 1, size * CAT_LINE + size * 0.6, longest, lv)
        for step in ((turned, wrapped) if turn_first else (wrapped, turned)):
            got = step(size)
            if got is not None:
                return got
    return turned(min(sizes) if sizes else 10.0, force=True)


def is_other_category(name: str) -> bool:
    """«Прочее», «Другие», «Other»: a remainder slice, coloured quietly rather than as a category of its own."""
    return bool(_OTHER_RE.match(name or ""))


def series_roles(names: Sequence[str]) -> Optional[list[str]]:
    """The story several series tell, from their names: «lead» (the latest / actual: Стало, После, Факт, the last
    year), «past» (Было, До, earlier years) and «plan» (План, Цель, Прогноз). None when the names tell no such story
    (Web / Mobile): those series keep distinct categorical colours."""
    k = len(names)
    if k < 2:
        return None
    roles: list[Optional[str]] = [None] * k
    for i, nm in enumerate(names):
        if _PLAN_RE.search(nm or ""):
            roles[i] = "plan"
    years = {i: int(m.group(1)) for i, nm in enumerate(names) if roles[i] is None and (m := _YEAR_RE.search(nm or ""))}
    if len(set(years.values())) >= 2:
        latest = max(years, key=lambda i: (years[i], i))
        for i in years:
            roles[i] = "lead" if i == latest else "past"
    for i, nm in enumerate(names):
        if roles[i] is None:
            if _NOW_RE.search(nm or ""):
                roles[i] = "lead"
            elif _PAST_RE.search(nm or ""):
                roles[i] = "past"
    if not any(r in ("past", "plan") for r in roles):
        return None
    leads = [i for i, r in enumerate(roles) if r == "lead"]
    free = [i for i, r in enumerate(roles) if r is None]
    if not leads and len(free) == 1:
        roles[free[0]] = "lead"  # «Выручка» next to «План»
        leads, free = free, []
    if len(leads) != 1 or free:
        return None
    return roles  # type: ignore[return-value]


_AX_TAGS = ("axId", "crossAx")
_AX_ID_BASE = 510_000_000  # well inside int32: LibreOffice reads axis ids as signed 32-bit and loses larger ones


def fix_axis_ids(chart_space) -> None:
    """python-pptx writes negative axis ids (-2068027336); the schema wants xs:unsignedInt, and LibreOffice reads them
    as signed 32-bit (an id ≥ 2^31 loses the axis pairing: no category labels, gridlines on the wrong axis). Every id
    is mapped to a small positive one consistently, so c:axId / c:crossAx pairs keep pointing at each other."""
    mapping: dict[str, str] = {}
    for el in chart_space.iter(*(_c(t) for t in _AX_TAGS)):
        v = el.get("val")
        if v is None:
            continue
        if v not in mapping:
            mapping[v] = str(_AX_ID_BASE + len(mapping) + 1)
        el.set("val", mapping[v])


# ---------------------------------------------------------------------------------------------- XML builders


def _insert_before(parent, new, successors: Sequence[str]) -> None:
    """Insert `new` before the first child whose local name is in `successors` (the schema order), else append."""
    for child in parent:
        if _local(child) in successors:
            child.addprevious(new)
            return
    parent.append(new)


def _replace(parent, tag: str, new, successors: Sequence[str]) -> None:
    old = parent.find(_c(tag))
    if old is not None:
        old.addprevious(new)
        parent.remove(old)
    else:
        _insert_before(parent, new, successors)


def _txpr(size: float, color: Optional[str] = None, bold: Optional[bool] = None, family: Optional[str] = None, horizontal: bool = False, nowrap: bool = False):
    tx = etree.Element(_c("txPr"))
    bp = etree.SubElement(tx, _a("bodyPr"))
    if nowrap:
        bp.set("wrap", "none")
    if horizontal:
        bp.set("rot", "0")
        bp.set("vert", "horz")
    etree.SubElement(tx, _a("lstStyle"))
    p = etree.SubElement(tx, _a("p"))
    ppr = etree.SubElement(p, _a("pPr"))
    d = etree.SubElement(ppr, _a("defRPr"))
    d.set("sz", str(int(round(size * 100))))
    if bold is not None:
        d.set("b", "1" if bold else "0")
    if color:
        sf = etree.SubElement(d, _a("solidFill"))
        etree.SubElement(sf, _a("srgbClr")).set("val", color)
    if family:
        for tag in ("latin", "ea", "cs"):
            etree.SubElement(d, _a(tag)).set("typeface", family)
    etree.SubElement(p, _a("endParaRPr")).set("lang", "ru-RU")
    return tx


def _fonts_everywhere(chart_space, family: Optional[str]) -> None:
    """Every run property that names a latin typeface names the same ea and cs typefaces (schema order latin, ea, cs),
    so no script falls back to the theme font."""
    if not family:
        return
    for latin in list(chart_space.iter(_a("latin"))):
        parent = latin.getparent()
        prev = latin
        for tag in ("ea", "cs"):
            el = parent.find(_a(tag))
            if el is None:
                el = etree.Element(_a(tag))
                el.set("typeface", latin.get("typeface") or family)
                prev.addnext(el)
            prev = el


def _sppr(fill: Optional[str] = None, alpha: Optional[float] = None, line: Optional[str] = None, line_w: Optional[float] = None, no_line: bool = True, dash: Optional[str] = None):
    sp = etree.Element(_c("spPr"))
    if fill:
        sf = etree.SubElement(sp, _a("solidFill"))
        clr = etree.SubElement(sf, _a("srgbClr"))
        clr.set("val", fill)
        if alpha is not None and alpha < 1:
            etree.SubElement(clr, _a("alpha")).set("val", str(int(alpha * 100000)))
    else:
        etree.SubElement(sp, _a("noFill"))
    if line:
        ln = etree.SubElement(sp, _a("ln"))
        ln.set("w", str(int((line_w or 0.75) * EMU_PER_PT)))
        ln.set("cap", "rnd")
        sf = etree.SubElement(ln, _a("solidFill"))
        etree.SubElement(sf, _a("srgbClr")).set("val", line)
        if dash:
            etree.SubElement(ln, _a("prstDash")).set("val", dash)
        etree.SubElement(ln, _a("round"))
    elif no_line:
        ln = etree.SubElement(sp, _a("ln"))
        etree.SubElement(ln, _a("noFill"))
    return sp


_DLBL_TAIL = ("showLegendKey", "showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize")
_SER_AFTER_DLBLS = ("trendline", "errBars", "cat", "val", "xVal", "yVal", "smooth", "shape", "extLst")


def _flags(parent, show_val: bool = True, leader_lines: Optional[bool] = None, ser_name: bool = False, sep: Optional[str] = None) -> None:
    for tag in _DLBL_TAIL:
        on = (tag == "showVal" and show_val) or (tag == "showSerName" and ser_name)
        etree.SubElement(parent, _c(tag)).set("val", "1" if on else "0")
    if sep is not None:
        etree.SubElement(parent, _c("separator")).text = sep
    if leader_lines is not None:
        etree.SubElement(parent, _c("showLeaderLines")).set("val", "1" if leader_lines else "0")


def _dlbls(fmt: str, size: float, color: str, family: Optional[str], pos: Optional[str], points: Optional[dict] = None, leader_lines: Optional[bool] = False, bold: bool = False, show_val: bool = True):
    """A complete c:dLbls: per-point overrides (`points[idx]` = {"delete": True} or any of "fmt", "size", "color",
    "bold", "pos", "dx"/"dy" (a nudge, as a fraction of the chart frame), "box" (a chip colour behind the text),
    "sername" (series name + value)), then the
    shared numFmt, transparent box, text, position and flags. Leader lines are off: a nudged label is never tied to
    its point by a stray line."""
    d = etree.Element(_c("dLbls"))
    for idx in sorted(points or {}):
        o = points[idx]
        lbl = etree.SubElement(d, _c("dLbl"))
        etree.SubElement(lbl, _c("idx")).set("val", str(idx))
        if o.get("delete"):
            etree.SubElement(lbl, _c("delete")).set("val", "1")
            continue
        if o.get("dx") or o.get("dy"):
            lay = etree.SubElement(lbl, _c("layout"))
            ml = etree.SubElement(lay, _c("manualLayout"))
            etree.SubElement(ml, _c("x")).set("val", f"{o.get('dx', 0.0):.4f}")
            etree.SubElement(ml, _c("y")).set("val", f"{o.get('dy', 0.0):.4f}")
        nf = etree.SubElement(lbl, _c("numFmt"))
        nf.set("formatCode", o.get("fmt", fmt))
        nf.set("sourceLinked", "0")
        lbl.append(_sppr(o["box"], 0.92) if o.get("box") else _sppr())  # a «box» colour: a chip behind the text
        lbl.append(_txpr(o.get("size", size), o.get("color", color), o.get("bold", bold), family, nowrap=True))
        p = o.get("pos", pos)
        if p:
            etree.SubElement(lbl, _c("dLblPos")).set("val", p)
        _flags(lbl, ser_name=bool(o.get("sername")), sep=LABEL_SEP if o.get("sername") else None)
    nf = etree.SubElement(d, _c("numFmt"))
    nf.set("formatCode", fmt)
    nf.set("sourceLinked", "0")
    d.append(_sppr())
    d.append(_txpr(size, color, bold, family, nowrap=True))
    if pos:
        etree.SubElement(d, _c("dLblPos")).set("val", pos)
    _flags(d, show_val=show_val, leader_lines=leader_lines)
    return d


def _set_ser_dlbls(ser_el, dlbls) -> None:
    old = ser_el.find(_c("dLbls"))
    if old is not None:
        ser_el.remove(old)
    _insert_before(ser_el, dlbls, _SER_AFTER_DLBLS)


def _set_plot_dlbls(chart_el, dlbls) -> None:
    """The chart-type level c:dLbls (read by the audit and by editors as the series default)."""
    old = chart_el.find(_c("dLbls"))
    if old is not None:
        old.addprevious(dlbls)
        chart_el.remove(old)
        return
    _insert_before(chart_el, dlbls, ("gapWidth", "overlap", "serLines", "dropLines", "hiLowLines", "upDownBars", "marker", "smooth", "firstSliceAng", "holeSize", "axId", "extLst"))


def _manual_layout(holder, x: float, y: float, w: Optional[float], h: Optional[float], inner: bool = True) -> None:
    layout = holder.find(_c("layout"))
    if layout is None:
        layout = etree.Element(_c("layout"))
        holder.insert(0, layout) if _local(holder) == "plotArea" else _insert_before(holder, layout, ("overlay", "spPr", "txPr", "extLst"))
    for child in list(layout):
        layout.remove(child)
    ml = etree.SubElement(layout, _c("manualLayout"))
    items = ([("layoutTarget", "inner")] if inner else []) + [("xMode", "edge"), ("yMode", "edge")]
    x, y = max(0.0, min(0.95, x)), max(0.0, min(0.95, y))
    items += [("x", f"{x:.4f}"), ("y", f"{y:.4f}")]
    if w is not None and h is not None:
        w, h = max(0.05, min(1.0 - x, w)), max(0.05, min(1.0 - y, h))
        items += [("w", f"{w:.4f}"), ("h", f"{h:.4f}")]
    for tag, val in items:
        etree.SubElement(ml, _c(tag)).set("val", val)


def _transparent(chart) -> None:
    """No white box behind the chart: chart space and plot area without fill or border, square corners."""
    cs = chart._chartSpace
    for holder in (cs, cs.find(f".//{_c('plotArea')}")):
        if holder is None:
            continue
        sp = holder.find(_c("spPr"))
        if sp is None:
            sp = etree.Element(_c("spPr"))
            if holder is cs:
                chart_el = holder.find(_c("chart"))
                if chart_el is not None:
                    chart_el.addnext(sp)
                else:
                    holder.append(sp)
            else:
                _insert_before(holder, sp, ("extLst",))
        for child in list(sp):
            sp.remove(child)
        etree.SubElement(sp, _a("noFill"))
        ln = etree.SubElement(sp, _a("ln"))
        etree.SubElement(ln, _a("noFill"))
    if cs.find(_c("roundedCorners")) is None:
        rc = etree.Element(_c("roundedCorners"))
        rc.set("val", "0")
        _insert_before(cs, rc, ("AlternateContent", "clrMapOvr", "pivotSource", "protection", "chart"))


def _axis_txpr(axis, size: float, color: str, family: Optional[str]) -> None:
    el = axis._element
    _replace(el, "txPr", _txpr(size, color, False, family, horizontal=True), ("crossAx", "crosses", "crossesAt", "crossBetween", "auto", "lblAlgn", "lblOffset", "extLst"))


def _set_child_val(parent, tag: str, val: str, successors: Sequence[str]) -> None:
    el = parent.find(_c(tag))
    if el is None:
        el = etree.Element(_c(tag))
        _insert_before(parent, el, successors)
    el.set("val", val)


def _snap(v: float) -> float:
    return round(v * 2) / 2


def _set_disp_blanks(chart_space, how: str = "gap") -> None:
    """c:chart/c:dispBlanksAs: how a missing value is drawn («gap» — a break in the line, never a zero)."""
    ch = chart_space.find(_c("chart"))
    if ch is None:
        return
    el = ch.find(_c("dispBlanksAs"))
    if el is None:
        el = etree.Element(_c("dispBlanksAs"))
        _insert_before(ch, el, ("showDLblsOverMax", "extLst"))
    el.set("val", how)


def _source_ink(muted_hex: Optional[str], text: str, muted: str, ground: str) -> str:
    """The source line's colour: the template's secondary text colour where it reads as small print on the ground
    (≥ 4.5:1), else the template's text colour; the chart's own eased text only for a caller that names neither."""
    m = _norm(muted_hex)
    if m and contrast_ratio(m, ground) >= 4.5:
        return m
    return text if muted_hex is not None else muted


def _source_box(slide: Slide, box: Bbox, text: str, size: float, color: str, family: Optional[str]):
    """The source line under a chart: a native text box in the template's face, small muted print, no autofit."""
    tb = slide.shapes.add_textbox(Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    tb.name = "Chart source"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.color.rgb = _rgb(color)
    if family:
        run.font.name = family
    run._r.get_or_add_rPr().set("lang", "ru-RU")
    _fonts_everywhere(tb._element, family)
    return tb


# ---------------------------------------------------------------------------------------------- ring


def add_ring(slide: Slide, bbox: Bbox, percent: float, color_hex: str, ground_hex: Optional[str] = None):
    """A native doughnut showing `percent` — the true value in place of a ring drawn as a picture. The track is a
    faint tint of the colour over the ground; no legend, no labels (the figure sits in the hole), transparent back."""
    percent = max(0.0, min(100.0, percent))
    data = CategoryChartData()
    data.categories = ["value", "rest"]
    data.add_series("share", (percent, 100.0 - percent))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Emu(bbox.x), Emu(bbox.y), Emu(bbox.w), Emu(bbox.h), data)
    _OWN_PARTS.add(gf.chart_part)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.has_data_labels = False
    track = _mix(color_hex, ground_hex or "FFFFFF", 0.14)
    for j, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(color_hex if j == 0 else track)
        pt.format.line.fill.background()
    dn = chart._chartSpace.find(f".//{_c('doughnutChart')}")
    if dn is not None:
        for tag, val in (("firstSliceAng", "0"), ("holeSize", "78")):
            el = dn.find(_c(tag))
            if el is not None:
                el.set("val", val)
    # the ring fills its frame, centred: the figure placed in the middle of the frame sits in the hole
    plot_area = chart._chartSpace.find(f".//{_c('plotArea')}")
    if plot_area is not None:
        _manual_layout(plot_area, 0.02, 0.02, 0.96, 0.96)
    _transparent(chart)
    return gf


# ---------------------------------------------------------------------------------------------- type from the template


def _font_floor(slide: Slide) -> float:
    """The smallest chart text, relative to the slide: 2 % of its height (≈ 11 pt on a 16:9 slide of 540 pt, 8 pt on
    one of 405 pt, where the template's 9 pt body is the same size on screen)."""
    _, sh = _slide_size(slide)
    return max(8.0, 0.02 * sh / EMU_PER_PT)


def _type_sizes(typography: Optional[Typography], used: bool = False) -> list[float]:
    """The template's type scale (role sizes); with `used`, also every other whole or half size it sets."""
    if typography is None:
        return []
    sizes = {_snap(s.size_pt) for s in (typography.scale or []) if s.size_pt and s.size_pt > 0}
    if used:
        sizes |= {_snap(v) for v in (typography.sizes_used or []) if v and abs(v * 2 - round(v * 2)) < 0.05}
    return sorted(sizes)


def _size_at_least(size: float, typography: Optional[Typography], cap: float = 1.4) -> float:
    """The smallest template size ≥ `size` (and not above `cap` × it): a role size of the type scale when one is
    near, else another size the template uses, else `size` itself rounded to half a point."""
    for pool in (_type_sizes(typography), _type_sizes(typography, used=True)):
        near = [s for s in pool if size - 0.05 <= s <= size * cap]
        if near:
            return min(near)
    return _snap(size)


def _size_at_most(size: float, typography: Optional[Typography], floor: float = 0.8) -> float:
    """The largest template size ≤ `size` (and not under `floor` × it), else `size` itself rounded down to half a
    point."""
    pool = [s for s in _type_sizes(typography, used=True) if size * floor - 0.05 <= s <= size + 0.05]
    return max(pool) if pool else math.floor(size * 2) / 2


def chart_text_capped(typography: Optional[Typography], font_size_pt: Optional[float], slide_h_emu: int) -> bool:
    """Chart and table text is capped relative to the slide when the template's own chart text is larger than 2.6 %
    of the slide height (a bullet placeholder's 24 pt taken for chart text) or its type scale is too thin to trust
    (a derived ladder). The dataset templates set 7–12 pt: never capped."""
    hpt = slide_h_emu / EMU_PER_PT if slide_h_emu else 0.0
    sparse = bool(getattr(typography, "derived_sizes", None)) if typography is not None else False
    return sparse or (hpt > 0 and (font_size_pt or 0.0) > 0.026 * hpt)


def _bold_voice(typography: Optional[Typography]) -> bool:
    """Does the template emphasise with weight? Only when its display / h1 are set in bold often enough."""
    if typography is None:
        return False
    return max((s.weight_bold_share for s in (typography.scale or []) if s.role in ("display", "h1")), default=0.0) >= BOLD_VOICE_SHARE


def _unit_shown_above(slide: Slide, bbox: Bbox, unit_full: str) -> bool:
    """The caller already wrote the full unit right above the chart (a caption or a heading ending in it)."""
    u = " ".join(unit_full.split()).lower()
    if not u:
        return False
    _, sh = _slide_size(slide)
    try:
        shapes = list(slide.shapes)
    except Exception:  # noqa: BLE001
        return False
    for shp in shapes:
        if not getattr(shp, "has_text_frame", False) or shp.top is None or shp.height is None:
            continue
        t = " ".join(shp.text_frame.text.replace(" ", " ").split()).lower()
        if not t:
            continue
        if u.replace("\u00a0", " ") not in t.replace("\u00a0", " "):
            # (the caption named the series in another form of the unit's word — «Проверки» over «проверок»: the
            # composer dropped the unit as a repeat, the chart must not print it as its title either)
            from verstka.rendering.compose import _same_stem

            if len(t.split()) > 6 or not _same_stem(unit_full, t):
                continue  # (a caption, not a heading that happens to share the word)
        bottom = shp.top + shp.height
        overlap = min(shp.left + shp.width, bbox.x + bbox.w) - max(shp.left, bbox.x)
        # (a caption of a row of charts stands on the row's line: a caption shorter than its neighbours ends higher)
        reach = 0.3 if (shp.name or "").startswith(("Chart title", "Unit")) else 0.14
        if overlap > 0 and bbox.y - reach * sh <= bottom <= bbox.y + 0.03 * sh:
            return True
    return False


# ---------------------------------------------------------------------------------------------- charts


@dataclass
class _Kit:
    """Everything the chart-kind stylers share."""

    chart: object
    plot: object
    chart_el: object
    kind: str
    series: list
    values: list
    cats: list
    spec: ChartSpec
    palette: list
    accent: str
    ground: str
    text: str
    muted: str
    rule: str
    fs: float  # category labels, legend, caption
    fs_value: float  # value labels
    fs_hl: float  # the emphasised value label (fs_value when the template emphasises with bold)
    bold: bool  # the template's emphasis is bold
    family: Optional[str]
    fmt_unit: str  # label format with the (short) unit
    fmt_plain: str  # the plain number (keeps a one-glyph unit)
    fmt_emph: str  # the one label that carries the unit (plain when a caption / title / the caller carries it)
    unit: Optional[str]
    unit_plain: Optional[str]  # the unit as the plain labels print it (glyph units only)
    unit_emph: Optional[str]
    decimals: int
    W: float
    H: float
    sw: int
    sh: int
    title_h: float
    legend_h: float = 0.0
    legend_w: float = 0.0
    roles: Optional[list] = None
    styles: list = field(default_factory=list)  # per series: {"fill", "line", "dash", "label"}
    outside: bool = True  # a pie may set a label outside a thin slice (False: the caller's legend carries it)
    slice_labels: bool = True  # a pie of amounts with the caller's legend: no computed shares on the slices
    pct_given: bool = False  # a pie of the brief's own percents: each slice says its value as written («18,6%», «42%»)
    capped: bool = False  # chart text capped relative to the slide (chart_text_capped): category labels measured strictly
    squeeze: float = 0.0  # a trend chart's longest category label over its slot (> 1: the renderer breaks it inside a word)
    turned: bool = False  # its category labels were turned by 45° or thinned out (every other one blank) to stay whole
    cramped: bool = False  # the plot kept its least width (0.3 of the frame) past the room its axis and end labels leave:
    # the renderer cuts the value axis's labels («1 200 0…»)
    x_title: Optional[str] = None  # the category axis's title (a scatter: the X axis's)
    y_title: Optional[str] = None  # the value axis's title (a scatter: the Y axis's), with the word unit
    inset_l: float = 0.0  # pt kept at the left of the frame for an axis title set along the left axis
    inset_b: float = 0.0  # pt kept above the legend for an axis title set under the bottom axis
    xs: list = field(default_factory=list)  # a scatter: the x of every point, per series

    @property
    def foot(self) -> float:
        """The room under the plot's labels: the legend and a bottom axis title."""
        return self.legend_h + self.inset_b

    def plain(self, v) -> str:
        return _fmt_value(v, self.unit_plain, self.decimals)

    def emph(self, v) -> str:
        return _fmt_value(v, self.unit_emph, self.decimals)

    @property
    def hl_color(self) -> str:
        return self.accent if contrast_ratio(self.accent, self.ground) >= 3.0 else self.text

    def hl_label(self, size_ratio: float = 1.0) -> dict:
        return {"color": self.hl_color, "bold": self.bold, "size": _snap(self.fs_hl * size_ratio), "fmt": self.fmt_emph}


def _legend(k: _Kit, names: Sequence[str], side: bool) -> None:
    """A bottom (or, for a pie in a wide frame, right-hand) legend in the muted text colour; sets k.legend_h / _w."""
    chart = k.chart
    chart.has_legend = True
    lg = chart.legend
    lg.include_in_layout = False
    lg.position = XL_LEGEND_POSITION.RIGHT if side else XL_LEGEND_POSITION.BOTTOM
    lg.font.size = Pt(k.fs)
    lg.font.color.rgb = _rgb(k.muted)
    if k.family:
        lg.font.name = k.family
    widest = max((text_width_pt(t, k.family, k.fs) for t in names), default=40.0)
    if side:
        k.legend_w = min(widest + k.fs * 2.4, k.W * 0.45)
    else:
        per_row = max(1, int(k.W // (widest + k.fs * 2.6)))
        rows = math.ceil(len(names) / per_row)
        k.legend_h = rows * k.fs * 1.3 + k.fs * 0.9


def _series_styles(k: _Kit) -> list[dict]:
    """Per-series paint. With a story (roles): the latest in the accent, earlier ones in its tints (older = fainter),
    a plan hollow (bars) or dashed (lines); labels of the stepped-back series in the muted text colour. Without: the
    categorical palette."""
    n = len(k.series)
    if not k.roles:
        return [{"fill": k.palette[i % len(k.palette)], "label": k.text, "role": None} for i in range(n)]
    out: list[dict] = [{} for _ in range(n)]
    pasts = [i for i, r in enumerate(k.roles) if r == "past"]
    # the most recent past is the one nearest the lead (years: the latest; words: the later one in the list)
    years = {i: int(m.group(1)) for i in pasts if (m := _YEAR_RE.search(k.series[i].name or ""))}
    pasts.sort(key=lambda i: (years.get(i, 0), i), reverse=True)
    taken = [k.accent]
    for rank, i in enumerate(pasts):
        c = muted_tint(k.accent, k.ground, MUTED_TOWARD_GROUND if rank == 0 else 0.8)
        if any(delta_e(c, t) < MIN_SERIES_DE * 0.75 for t in taken):
            c = next((p for p in k.palette[1:] if all(delta_e(p, t) >= MIN_SERIES_DE for t in taken)), c)
        taken.append(c)
        out[i] = {"fill": c, "label": k.muted, "role": "past"}
    for i, r in enumerate(k.roles):
        if r == "lead":
            out[i] = {"fill": k.accent, "label": k.text, "role": "lead"}
        elif r == "plan":
            out[i] = {"fill": None, "line": k.accent, "label": k.muted, "role": "plan"}
    return out


def add_chart(
    slide: Slide,
    bbox: Bbox,
    spec: ChartSpec,
    outline: DeckOutline,
    style: ChartStyleSpec,
    typography: Typography,
    text_hex: Optional[str] = None,
    neutral_hex: Optional[str] = None,
    *,
    ground_hex: Optional[str] = None,
    accent_hex: Optional[str] = None,
    legend: bool = True,
    amounts: bool = False,
    muted_hex: Optional[str] = None,
):
    """A native chart at `bbox`, finished to the rules in the module docstring. `amounts` (a pie or a doughnut of
    amounts, not percents, drawn with the caller's legend): the chart keeps the amounts as its data and labels its
    slices with them (the brief's figures) — no computed share is written on a slice as if it were a figure of the
    brief; the caller's legend gives each part's amount and its share, headed as a share. `ground_hex` is the colour
    under the chart (read from the slide when omitted); `accent_hex` overrides accent.1 (the highlight). `legend=False`: a pie
    or a doughnut without its own legend (the caller sets one beside it) — the circle takes the whole frame and a
    slice too thin for its label inside carries none (the caller's legend gives every share). `spec.x_title` /
    `spec.y_title`: axis titles; `spec.source`: a source line under the chart inside `bbox` (a caller that prints the
    source on the slide clears it — `source_said`); `spec.xy`: the points of a scatter. `muted_hex`: the template's own
    secondary text colour (the composer's footnote colour) for the source line, which stands on the slide outside the
    chart — a colour of the template, as the audit wants every text of the slide. The returned graphic frame carries
    `verstka_source` (the source line's text box, or None) and `verstka_source_dropped` (the line of a frame too small to
    hold it, which the caller then prints in its own note; None when nothing was left out)."""
    kind = effective_chart_type(spec, outline)
    xs_all: list[list[float]] = []
    if kind == "scatter":
        pts = xy_series(spec)
        if not pts:
            raise ValueError("chart has no series data")
        series = [Series.model_construct(id=f"xy_{i + 1}", name=nm, categories=[], values=list(ys), unit=spec.unit, source_span=None) for i, (nm, _, ys) in enumerate(pts)]
        xs_all = [list(xs) for _, xs, _ in pts]
    else:
        series = resolve_series(spec, outline)
        if not series:
            raise ValueError("chart has no series data")
    if kind in ("pie", "doughnut"):
        series = series[:1]
    cats = [str(c) for c in series[0].categories]
    if kind in _BARS:
        cats = [glue_numbers(c) for c in cats]  # a bin «1 250–1 500» is broken after its dash, never inside a number
    n = len(cats)
    multi = len(series) > 1
    if kind == "scatter":
        values = [[float(v) for v in s.values] for s in series]
    else:
        values = [[(float(v) if v is not None else None) for v in list(s.values)[:n]] for s in series]
    unit = spec.unit or series[0].unit

    # --- ground, text and colours
    ground = _norm(ground_hex) or slide_ground(slide, bbox)
    if not ground:
        ground = "000000" if (_norm(text_hex) and not _is_dark(_norm(text_hex)) and relative_luminance(_norm(text_hex)) > 0.5) else "FFFFFF"
    text = _readable_text(text_hex, ground)
    muted = _muted_text(text, ground)
    if _norm(muted_hex) and contrast_ratio(_norm(muted_hex), ground) >= 4.5:
        muted = _norm(muted_hex)  # one grey: the template's own secondary text colour (ticks, legend, axis titles, source)
    rule = _rule_color(neutral_hex, text, ground)
    # accent.1 leads; then the template's own chart colours, the theme accents (what Office charts use) and the
    # manifest's ranked fill colours. On a ground the accent cannot stand on (an accent-coloured panel) the text
    # colour takes the lead, as white bars on a brand-blue slide.
    manifest = [c for c in list(style.series_colors or []) if _norm(c)]
    lead = _norm(accent_hex) or (manifest[0] if manifest else None) or "0077FF"
    if contrast_ratio(lead, ground) < LEAD_MIN_CONTRAST:
        # the chart's lead colour must read on its ground (a dark teal on a dark purple picture is lost): the next
        # colour of the template that does; else the brand colour itself deepened in its own hue (a honey yellow on
        # white becomes gold — LibreOffice «Beehive» — not the theme's stock green); else a theme accent; else the
        # text colour (white bars on a brand-blue slide)
        cands = [_norm(c) for c in manifest[1:] + template_chart_colors(slide) if _norm(c)]
        got = next((c for c in cands if contrast_ratio(c, ground) >= LEAD_MIN_CONTRAST and _is_data_color(c)), None)
        if got is None and _is_data_color(lead) and delta_e(lead, ground) >= 10:
            # (on a panel of the accent's own colour the text colour leads instead: white bars on brand blue)
            got = shade_to(lead, ground, LEAD_MIN_CONTRAST + 0.4)
        if got is None:
            got = next((c for c in (_norm(x) for x in _brand_theme_accents(slide)) if c and contrast_ratio(c, ground) >= LEAD_MIN_CONTRAST and _is_data_color(c)), None)
        lead = got or text
    extra = template_chart_colors(slide) + _brand_theme_accents(slide) + manifest[1:] + ([manifest[0]] if manifest else [])
    others = [j for j, c in enumerate(cats) if is_other_category(c)] if kind in ("pie", "doughnut") and n > 2 else []
    if kind in ("pie", "doughnut"):
        # the largest slice in the accent, the others in its tints (a remainder quiet): one whole, one hue
        shades = pie_shades(values[0], lead, ground, others, text)
        ranked = sorted((j for j, v in enumerate(values[0]) if v and j not in others), key=lambda j: (-(values[0][j] or 0.0), j))
        if len(ranked) >= 5:
            # five parts and more: tints of one hue run too close to tell apart (eight tints of one pink: neighbours at
            # 1.0–1.3:1, the legend matched no slice) — the template's own distinct colours, by size, a remainder quiet
            distinct = chart_palette([lead], ground, len(ranked), extra=extra)
            for r, j in enumerate(ranked):
                shades[j] = (distinct[r % len(distinct)], 1.0)
        palette = [base if share >= 1.0 else _mix(base, ground, share) for base, share in shades]
        accent = lead
    else:
        palette = chart_palette([lead], ground, max(1, len(series)), extra=extra)
        accent = palette[0]

    # --- type: sizes on the template scale, weight from its voice
    W, H = max(bbox.w, EMU_PER_PT) / EMU_PER_PT, max(bbox.h, EMU_PER_PT) / EMU_PER_PT
    sw, sh = _slide_size(slide)
    family = style.font_family or typography.primary_family
    fs = _size_at_least(max(style.font_size_pt or 0.0, _font_floor(slide)), typography, cap=1.35)
    try:
        body = typography.size_for("body", fs)
    except Exception:  # noqa: BLE001
        body = fs
    # the values speak a step above the category labels, at the template's body size where that is such a step
    fs_value = _size_at_least(max(fs * 1.1, min(body or fs, fs * 1.5)), typography, cap=1.45)
    bold = _bold_voice(typography)
    fs_hl = fs_value if bold else _size_at_least(fs_value * 1.12, typography, cap=1.4)
    capped = chart_text_capped(typography, style.font_size_pt, sh)
    if capped:
        # chart text relative to the slide (≤ 2.6 % of its height) and to its frame (≤ a tenth of it): a bullet
        # placeholder's 24 pt on an 11″ slide is not a chart label size
        hpt = sh / EMU_PER_PT
        cap = max(min(0.026 * hpt, H / 10.0), 8.0)
        if fs > cap:
            fs = _size_at_most(cap, typography)
        fs_value = max(fs, _size_at_most(min(fs_value, 1.25 * fs, max(0.03 * hpt, fs)), typography))
        fs_hl = max(fs_value, _size_at_most(min(fs_hl, 1.12 * fs_value), typography)) if not bold else fs_value

    # --- the source line: small muted print under the chart, inside its frame (the chart takes what is left)
    src = source_text(spec.source)
    src_dropped: Optional[str] = None
    src_fs = src_h = 0.0
    chart_box = bbox
    if src:
        src_fs = min(fs, max(_font_floor(slide), _size_at_most(fs * 0.9, typography)))
        lines = max(1, min(2, len(wrap_lines(src, family, src_fs, False, W * 0.98))))
        src_h = lines * src_fs * 1.3
        room = src_h + src_fs * 0.6
        # (the line is already at the chart's floor size) the chart keeps at least 45 % of its frame
        if H - room >= 0.45 * H:
            H -= room
            chart_box = Bbox(x=bbox.x, y=bbox.y, w=bbox.w, h=max(1, int(round(H * EMU_PER_PT))))
        else:
            src_dropped, src = src, None  # a frame too small for a chart and its line: `verstka_source_dropped` hands it back

    # --- data
    data = XyChartData() if kind == "scatter" else CategoryChartData()
    pie_pct = False
    if kind == "scatter":
        for s, xs, ys in zip(series, xs_all, values):
            ser_data = data.add_series(s.name or "")
            for x, y in zip(xs, ys):
                ser_data.add_data_point(x, y)
    elif kind in ("pie", "doughnut"):
        data.categories = cats
        vals = [max(0.0, v or 0.0) for v in values[0]]
        total = sum(vals)
        is_pct = pie_values_are_percents(unit, total)
        keep = amounts and not legend and not is_pct and total > 0
        pie_pct = is_pct and total > 0
        shares = vals if is_pct or total <= 0 or keep else [round(v / total * 100, 1) for v in vals]
        values = [shares]
        unit = unit if keep else "%"
        data.add_series(series[0].name, shares)
    else:
        data.categories = cats
        for s, vals in zip(series, values):
            data.add_series(s.name, vals)
    all_vals = [v for vs in values for v in vs if v is not None]

    # --- the unit, once: the value axis's title, else a caption above the chart, else the labels
    unit_full = " ".join((unit or "").split())
    glyph = glyph_unit(unit_full)
    pie = kind in ("pie", "doughnut")
    x_title = None if pie else (" ".join((spec.x_title or "").split()) or None)
    y_title = None if pie else (" ".join((spec.y_title or "").split()) or None)
    axis_unit = False  # the value axis's title carries the word unit
    if y_title and unit_full and not glyph:
        low = y_title.lower()
        if unit_full.lower() in low or short_unit(unit_full).lower() in low:
            axis_unit = True
        elif not _unit_shown_above(slide, bbox, unit_full):
            y_title = f"{y_title.rstrip(' ,.:;')}, {unit_full}"
            axis_unit = True
    head = spec.title or None
    head_is_title = bool(head)
    unit_in_head = False
    _on_labels = not head and not axis_unit and bare_unit_on_labels(spec, outline, unit_full)
    if not pie and unit_full and not glyph:
        if axis_unit:
            unit_in_head = True  # the value axis's title says the unit: the labels are plain numbers
        elif head:
            low = head.lower()
            if unit_full.lower() not in low and short_unit(unit_full).lower() not in low:
                head = f"{head.rstrip(' ,.:;')}, {unit_full}"
            unit_in_head = True
        elif _unit_shown_above(slide, bbox, unit_full):
            unit_in_head = True
        elif _on_labels:
            unit_in_head = False  # (a short word unit of a few bars rides on the labels, `bare_unit_on_labels`)
        else:
            head = unit_full  # a word unit is written once, above the chart: the labels are plain numbers
            unit_in_head = True
    if pie and unit != "%":
        decimals = _decimals(all_vals)  # a pie of the brief's amounts: its slices say the amounts
        fmt_unit = fmt_plain = fmt_emph = label_format(unit, all_vals)
        unit_plain = unit_emph = unit
    elif pie:
        decimals = 0
        fmt_unit = fmt_plain = fmt_emph = label_format("%", [])
        unit_plain = unit_emph = "%"
    else:
        decimals = _decimals(all_vals)
        fmt_unit = label_format(unit, all_vals)
        fmt_plain = fmt_unit if glyph or _on_labels else label_format(None, all_vals)
        fmt_emph = fmt_plain if unit_in_head else fmt_unit
        unit_plain = unit if glyph or _on_labels else None
        unit_emph = unit_plain if unit_in_head else unit

    gf = slide.shapes.add_chart(_TYPE_MAP[kind], Emu(chart_box.x), Emu(chart_box.y), Emu(chart_box.w), Emu(chart_box.h), data)
    _OWN_PARTS.add(gf.chart_part)
    chart = gf.chart
    cs = chart._chartSpace
    fix_axis_ids(cs)
    if kind != "scatter":
        # `spec.gaps`: a missing value is a break in the line, never a zero (bars have nothing to draw there either
        # way); without the flag a line or an area is drawn through it (connected), as a series that has no break to show
        _set_disp_blanks(cs, "gap" if (spec.gaps or kind not in ("line", "area")) else "span")
    chart.font.size = Pt(fs)
    if family:
        chart.font.name = family
    chart.font.color.rgb = _rgb(text)
    plot = chart.plots[0]
    chart_el = next((c for c in cs.find(f".//{_c('plotArea')}") if _local(c).endswith("Chart")), None)

    # --- title / unit caption: one line at the top left, above the plot
    title_h = 0.0
    chart.has_title = bool(head)
    if head:
        tf = chart.chart_title.text_frame
        tf.text = head
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(fs)
        run.font.bold = bool(head_is_title and bold)
        run.font.color.rgb = _rgb(text if head_is_title else muted)
        if family:
            run.font.name = family
        title_el = cs.find(f"{_c('chart')}/{_c('title')}")
        if title_el is not None:
            _manual_layout(title_el, 0.0, 0.0, None, None, inner=False)
            ov = title_el.find(_c("overlay"))
            if ov is None:
                ov = etree.Element(_c("overlay"))
                _insert_before(title_el, ov, ("spPr", "txPr", "extLst"))
            ov.set("val", "0")
        title_h = fs * 1.5 + 4

    k = _Kit(
        chart=chart, plot=plot, chart_el=chart_el, kind=kind, series=series, values=values, cats=cats, spec=spec,
        palette=palette, accent=accent, ground=ground, text=text, muted=muted, rule=rule,
        fs=fs, fs_value=fs_value, fs_hl=fs_hl, bold=bold, family=family,
        fmt_unit=fmt_unit, fmt_plain=fmt_plain, fmt_emph=fmt_emph, unit=unit, unit_plain=unit_plain, unit_emph=unit_emph,
        decimals=decimals, W=W, H=H, sw=sw, sh=sh, title_h=title_h,
    )
    k.slice_labels = True
    k.capped = capped
    k.pct_given = pie_pct
    k.x_title, k.y_title, k.xs = x_title, y_title, xs_all
    # axis titles take a line of room along their axis: the left one at the frame's left edge, the bottom one above
    # the legend
    left_t, bottom_t = (x_title, y_title) if kind in _HORIZONTAL else (y_title, x_title)
    if left_t:
        k.inset_l = fs * TITLE_LINE + fs * 0.45
    if bottom_t:
        k.inset_b = fs * TITLE_LINE + fs * 0.35
    if multi and not pie and kind != "scatter" and kind not in _STACKED:
        k.roles = series_roles([s.name for s in series])  # (stacked series are parts of a whole, not a story)
    k.styles = _series_styles(k)
    if kind in _STACKED and multi:
        quiet = _quiet_fill(text, ground)
        for i, s in enumerate(series):
            if is_other_category(s.name or ""):
                k.styles[i] = {"fill": quiet, "label": muted, "role": None}  # «Прочее» is present, without a hue

    # --- legend: pies always; bars with several series (horizontal bars name their series on the first row
    # instead); lines and areas name each line at its end
    chart.has_legend = False
    k.outside = legend
    if pie and legend:
        _legend(k, cats, side=W / max(H, 1) >= 1.45)
    elif multi and kind in ("column", "stacked_column", "histogram", "scatter"):
        _legend(k, [s.name for s in series], side=False)

    if kind in _BARS:
        _style_bars(k)
    elif kind in ("line", "area"):
        _style_line(k)
    elif kind == "scatter":
        _style_scatter(k)
    else:
        _style_pie(k)
    _transparent(chart)
    _fonts_everywhere(cs, family)
    gf.verstka_source_dropped = src_dropped  # the source line a frame too small for it left out: the caller prints it
    gf.verstka_source = _source_box(slide, Bbox(x=bbox.x, y=bbox.y + bbox.h - int(round(src_h * EMU_PER_PT)), w=bbox.w, h=int(round(src_h * EMU_PER_PT))), src, src_fs, _source_ink(muted_hex, text, muted, ground), family) if src else None
    # for a caller's own legend: the colour of every slice (series), and for a pie its tint as (base, share over the
    # ground) — a swatch drawn in the base colour at that opacity looks the same and stays a colour of the template
    gf.verstka_colors = list(palette) if kind in ("pie", "doughnut") else [st.get("fill") or st.get("line") or accent for st in k.styles]
    gf.verstka_shades = list(shades) if kind in ("pie", "doughnut") else None
    # a caller trying narrower slots for the chart compares how its labels fit: their width over the slot, and whether
    # they had to be turned or thinned out
    gf.verstka_squeeze = k.squeeze
    gf.verstka_turned = k.turned
    gf.verstka_cramped = k.cramped
    return gf


# ---------------------------------------------------------------------------------------------- axes


def _cat_axis(k: _Kit, fs: float, reverse: bool = False, line: bool = True, skip: int = 1, rotate: bool = False) -> None:
    ca = k.chart.category_axis
    ca.visible = True
    ca.has_major_gridlines = False
    ca.has_minor_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if line:
        ca.format.line.color.rgb = _rgb(k.rule)
        ca.format.line.width = Pt(0.75)
    else:
        ca.format.line.fill.background()
    _axis_txpr(ca, fs, k.muted, k.family)
    el = ca._element
    if rotate:
        bp = el.find(_c("txPr") + "/" + _a("bodyPr"))
        if bp is not None:
            bp.set("rot", "-2700000")
            bp.attrib.pop("vert", None)
    if reverse:
        scaling = el.find(_c("scaling"))
        orient = scaling.find(_c("orientation")) if scaling is not None else None
        if orient is not None:
            orient.set("val", "maxMin")
    _set_child_val(el, "tickLblSkip", str(max(1, skip)), ("tickMarkSkip", "noMultiLvlLbl", "extLst"))


def _value_axis(k: _Kit, show: bool, fs: float, lo: Optional[float], hi: Optional[float], step: Optional[float] = None) -> None:
    va = k.chart.value_axis
    va.has_minor_gridlines = False
    if lo is not None:
        va.minimum_scale = lo
    if hi is not None:
        va.maximum_scale = hi
    if not show and k.y_title:
        # the labels carry the values: the axis stays only to carry its title — no ticks, no line, no gridlines
        va.visible = True
        va.has_major_gridlines = False
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.minor_tick_mark = XL_TICK_MARK.NONE
        va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        va.format.line.fill.background()
        return
    if not show:
        va.has_major_gridlines = False
        va.visible = False
        d = va._element.find(_c("delete"))
        if d is not None:
            d.set("val", "1")  # explicit: some readers take a bare <c:delete/> for «shown»
        return
    va.visible = True
    va.has_major_gridlines = True
    if step:
        va.major_unit = step
    gl = va.major_gridlines.format.line
    gl.color.rgb = _rgb(_mix(k.rule, k.ground, 0.7))
    gl.width = Pt(0.5)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()
    va.tick_labels.number_format = _tick_format(k.unit, tick_decimals(step, k.decimals))
    va.tick_labels.number_format_is_linked = False
    _axis_txpr(va, fs, k.muted, k.family)


def _axis_crosses(axis_el, where: str) -> None:
    """Where this axis crosses the other one: c:crosses «autoZero» | «min» | «max» (the other axis's value), whichever
    of c:crosses / c:crossesAt it had. A value axis of reversed categories («maxMin», top-down bars) crosses them at
    «max» to stand at the bottom, not at the top edge; a scatter's Y axis crosses X at its «min» to stand at the left,
    not in the middle of the points when X spans zero."""
    for tag in ("crossesAt", "crosses"):
        for old in axis_el.findall(_c(tag)):
            axis_el.remove(old)
    el = etree.Element(_c("crosses"))
    el.set("val", where)
    _insert_before(axis_el, el, ("crossBetween", "auto", "lblAlgn", "lblOffset", "majorUnit", "minorUnit", "dispUnits", "tickLblSkip", "tickMarkSkip", "noMultiLvlLbl", "extLst"))


def _nice_floor(v: float) -> float:
    if v <= 0:
        return 0.0
    mag = 10 ** math.floor(math.log10(v))
    return math.floor(v / mag) * mag


def _plot_layout(k: _Kit, left: float, top: float, pw: float, ph: float) -> None:
    pa = k.chart._chartSpace.find(f".//{_c('plotArea')}")
    if pa is not None:
        _manual_layout(pa, left / k.W, top / k.H, pw / k.W, ph / k.H)
    _axis_titles(k, left, top, pw, ph)


def _axis_title(k: _Kit, axis, text: str, vertical: bool, size: float):
    """A native axis title in the template's face at `size`, the muted text colour, not bold; turned to read bottom-up
    along a vertical axis. Returns its c:title."""
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.color.rgb = _rgb(k.muted)
    if k.family:
        run.font.name = k.family
    run._r.get_or_add_rPr().set("lang", "ru-RU")
    title_el = axis._element.find(_c("title"))
    bp = title_el.find(f"{_c('tx')}/{_c('rich')}/{_a('bodyPr')}")
    if bp is not None:
        bp.set("rot", "-5400000" if vertical else "0")
        bp.set("vert", "horz")
    ov = title_el.find(_c("overlay"))
    if ov is None:
        ov = etree.Element(_c("overlay"))
        _insert_before(title_el, ov, ("spPr", "txPr", "extLst"))
    ov.set("val", "0")
    return title_el


def _axis_titles(k: _Kit, left: float, top: float, pw: float, ph: float) -> None:
    """The axis titles set along their axes (`k.x_title` on the category or X axis, `k.y_title` on the value or Y
    axis): the left one turned, centred on the plot's height at the frame's left edge (the room `k.inset_l`); the
    bottom one centred under the plot, just above the legend (the room `k.inset_b`). A title longer than its side
    steps down the template's sizes to the chart's floor."""
    if not (k.x_title or k.y_title) or k.kind in ("pie", "doughnut"):
        return
    horizontal = k.kind in _HORIZONTAL
    floor = max(8.0, 0.02 * k.sh / EMU_PER_PT)
    for text, axis, is_x in ((k.x_title, k.chart.category_axis, True), (k.y_title, k.chart.value_axis, False)):
        if not text:
            continue
        vertical = (is_x and horizontal) or (not is_x and not horizontal)
        side = (ph if vertical else pw) * 0.98
        size = k.fs
        for s in (k.fs, _snap(k.fs * 0.9), _snap(k.fs * 0.82)):
            size = max(min(floor, k.fs), s)
            if text_width_pt(text, k.family, size) * 1.04 + 4 <= side:
                break
        tw = min(text_width_pt(text, k.family, size) * 1.04 + 4, (k.H if vertical else k.W) * 0.98)
        title_el = _axis_title(k, axis, text, vertical, size)
        if vertical:
            x = 0.0
            y = min(max(0.0, top + ph / 2 - tw / 2), max(0.0, k.H - tw))
        else:
            x = min(max(0.0, left + pw / 2 - tw / 2), max(0.0, k.W - tw))
            y = max(0.0, k.H - k.legend_h - size * TITLE_LINE)
        _manual_layout(title_el, x / k.W, y / k.H, None, None, inner=False)


# ---------------------------------------------------------------------------------------------- bars


def _style_bars(k: _Kit) -> None:
    """Columns and horizontal bars — one series, several side by side, or stacked parts of a whole — and a histogram's
    touching bins."""
    horizontal = k.kind in _HORIZONTAL
    stacked = k.kind in _STACKED
    hist = k.kind == "histogram"
    cats, values, series, W, H, fs, family = k.cats, k.values, k.series, k.W, k.H, k.fs, k.family
    n = max(1, len(cats))
    ns = len(series)
    multi = ns > 1
    flat = [v for vs in values for v in vs if v is not None]
    has_neg = any(v < 0 for v in flat)
    want = k.spec.highlight_index
    hl = want if (want is not None and not multi and 0 <= want < n) else None
    # several series of the data layer's grouped or stacked bars: the highlighted category stands out, the other
    # categories ease toward the ground (the planner's several-series bars keep ignoring a highlight, as before)
    data_kind = k.spec.type in ("grouped_column", "grouped_bar", "stacked_column", "stacked_bar")
    cat_hl = want if (want is not None and multi and data_kind and 0 <= want < n) else None
    # the one label that carries the unit: the highlight, else the top bar (horizontal) or the tallest column
    emph = hl
    if emph is None and not multi and values[0]:
        vals0 = [(v if v is not None else -math.inf) for v in values[0]]
        emph = 0 if horizontal else max(range(len(vals0)), key=lambda j: (vals0[j], j))
    # a stacked bar reaches the sum of its parts (positives up, negatives down)
    tops = [sum(max(0.0, vs[j] or 0.0) for vs in values if j < len(vs)) for j in range(n)] if stacked else []
    bots = [sum(min(0.0, vs[j] or 0.0) for vs in values if j < len(vs)) for j in range(n)] if stacked else []
    ext = flat + tops + bots  # what the value axis spans
    # horizontal bars with several series name them on the first row («Было: 48») instead of in a legend; stacked
    # parts always have their legend
    widest_name = max((text_width_pt(s.name, family, fs) for s in series), default=0.0)
    direct = horizontal and multi and not stacked and ns <= 4 and widest_name <= W * 0.3
    if horizontal and multi and not direct:
        _legend(k, [s.name for s in series], side=False)

    fs_val0 = k.fs_value if (not multi and n <= 8) else fs
    ratio = k.fs_hl / k.fs_value if k.fs_value else 1.0

    def labels(fsv: float) -> list[tuple[str, float, bool]]:
        out = []
        for i, vs in enumerate(values):
            for j, v in enumerate(vs):
                if v is None or v == 0:
                    if horizontal and v == 0:
                        out.append(("0", fsv, False))
                    continue
                if not multi and j == hl:
                    t, size, b = k.emph(v), _snap(fsv * ratio), k.bold
                elif not multi and j == emph:
                    t, size, b = k.emph(v), fsv, False
                else:
                    t, size, b = k.plain(v), fsv, False
                if direct and j == 0:
                    t = f"{series[i].name}{LABEL_SEP}{t}"
                out.append((t, size, b))
        return out

    def lab_w(fsv: float) -> float:
        return max((text_width_pt(t, family, s, b) for t, s, b in labels(fsv)), default=0.0)

    def lab_top(fsv: float) -> float:
        return max((s for _, s, _ in labels(fsv)), default=fsv)

    # category label geometry: horizontal bars carry their labels on the left, wrapped inside 42 % of the frame and
    # stepped down when the wrapped lines would not fit their row
    cat_fs = fs
    cat_w = 0.0
    # never under the chart's own floor (2 % of the slide's height): a label that does not fit at it wraps or turns
    # (below) — Ateliers' 7.5 pt labels under 12 pt values read as footnotes
    floor = max(8.0, 0.02 * k.sh / EMU_PER_PT)
    if horizontal:
        row = max(H - k.title_h - k.foot - fs, H * 0.3) / n
        for size in sorted({fs, max(floor, _snap(fs * 0.9)), max(floor, _snap(fs * 0.82))}, reverse=True):
            cat_fs = size
            cat_w = min(max((text_width_pt(c, family, size) for c in cats), default=20.0), W * 0.42)
            lines = max((len(wrap_lines(c, family, size, False, cat_w + 0.5, dashes=True)) for c in cats), default=1)
            if lines * size * 1.2 <= row * 0.92:
                break
    else:
        # single-word labels cannot wrap: one step down when they crowd their column (headroom for a wider font)
        pitch = W / n
        longest = max((max((text_width_pt(w, family, fs) for w in label_pieces(c)), default=0.0) for c in cats), default=0.0)
        if longest > pitch * 0.75:
            cat_fs = min(fs, max(_snap(fs * 0.85), _snap(fs * pitch * 0.75 / max(longest, 1)), floor))

    turned_below = 0.0  # the room turned column labels take under the axis (B3-1: a word never broken inside)

    def box(labels_on: bool, axis_on: bool, fsv: float):
        outside = labels_on and not stacked  # (stacked labels sit inside their segments)
        if horizontal:
            left = k.inset_l + cat_w + cat_fs * 0.8
            right = (lab_w(fsv) + fsv * 0.7) if outside else fs
            top = k.title_h + (fs * 0.3)
            bottom = k.foot + (fs * 1.8 if axis_on else fs * 0.3)
        else:
            pitch0 = W / n
            lines = max((len(wrap_lines(c, family, cat_fs, False, pitch0 * 0.92, dashes=True)) for c in cats), default=1)
            lines = min(lines, 3)
            left = k.inset_l + (_axis_room(ext, k.unit, k.decimals, family, fs, 0.8) if axis_on else 1.0)
            right = 1.0
            top = k.title_h + ((lab_top(fsv) * 1.3 + 5) if outside else fs * 0.7)
            bottom = k.foot + (turned_below if turned_below else lines * cat_fs * 1.22 + cat_fs * 0.6)
        return left, top, max(W - left - right, W * 0.3), max(H - top - bottom, H * 0.3)

    labels_on, axis_on = True, False  # values on the bars replace the value axis
    fs_val = fs_val0
    left, top, pw, ph = box(labels_on, axis_on, fs_val)
    slot = (ph if horizontal else pw) / n
    if hist:
        # a histogram's bins touch: the bars are the intervals, a hairline of the ground between them
        gap = 0
        group = slot
    else:
        # bar thickness: gap 60–80 %, but a few bars in a wide frame keep a designed width instead of turning into slabs
        cap = (max(fs_val * 2.6, 0.075 * k.sh / EMU_PER_PT) if horizontal else 0.11 * k.sw / EMU_PER_PT) * (1.0 if (not multi or stacked) else min(ns, 3) * 0.9)
        group = min(slot / (1.5 if stacked else 1.7), cap)
        lo_gap, hi_gap = STACK_GAP if stacked else (60, 300)
        gap = int(max(lo_gap, min(hi_gap, round((slot / group - 1) * 100))))
        group = slot / (1 + gap / 100)
    bar = group / ns if (multi and not stacked) else group
    if not horizontal and not stacked:
        # every value label must fit over its bar (plus a little of the gap)
        room = bar + (slot - group) * 0.5 if not multi else bar * 1.12
        sizes = (fs_val, _snap(fs * 1.15), fs, _snap(fs * 0.9)) if not multi else (fs, _snap(fs * 0.9), max(8.0, _snap(fs * 0.82)))
        for size in sizes:
            if all(text_width_pt(t, family, s, b) <= room * 0.98 for t, s, b in labels(size)):
                fs_val = size
                break
        else:
            if multi:
                labels_on, axis_on = False, True
            else:
                fs_val = _snap(fs * 0.9)
        left, top, pw, ph = box(labels_on, axis_on, fs_val)
    if horizontal and multi and not stacked and (ph / n / (1 + gap / 100) / ns) < fs * 0.9:
        labels_on, axis_on = False, True
    if multi and not labels_on and direct:
        direct = False
        _legend(k, [s.name for s in series], side=False)
    left, top, pw, ph = box(labels_on, axis_on, fs_val)
    rotate_cols = False
    if not horizontal and cats:
        # B3-1: LibreOffice breaks a column's label inside a word wider than 95 % of its slot — every word measured at
        # the axis font against the real slot: a size down (never under the floor), else the labels turn by 45°
        lim = pw / n * _cat_room(family)
        floor = max(8.0, 0.02 * k.sh / EMU_PER_PT)
        # (each word against its own room: a long word keeps a margin, `_word_room`)
        # (a label breaks at its ordinary spaces and after the dash of a range, never at a no-break space or inside a
        # number: its pieces, `label_pieces` — «1 250–1 500» is «1 250–» over «1 500»)
        over = lambda size: any(text_width_pt(w, family, size) > _word_room(w, lim) for c in cats for w in label_pieces(c))  # noqa: E731
        if over(cat_fs):
            for size in [x for x in (_snap(cat_fs * 0.9), _snap(cat_fs * 0.82)) if x >= min(floor, cat_fs) - 0.05]:
                if not over(size):
                    cat_fs = size
                    break
            else:
                room_h = H - k.title_h - k.foot
                fit = fit_categories(cats, family, [cat_fs], pw / n, max(0.0, room_h - max(H * 0.3, 0.42 * room_h)))
                if fit.mode == "turn":
                    rotate_cols = k.turned = True
                    turned_below = fit.below
                    if fit.labels != list(cats):
                        cats = k.cats = list(fit.labels)
                        _set_categories(k, cats)
                    left, top, pw, ph = box(labels_on, axis_on, fs_val)

    # stacked: each value inside its segment where it fits (its width across the bar, its line along it); a value
    # that does not fit leaves its segment bare and brings a quiet value axis to read it by
    inside: dict[tuple[int, int], bool] = {}

    def fit_inside(pw_: float, ph_: float, lo_v: float, hi_v: float) -> dict[tuple[int, int], bool]:
        span = (hi_v - lo_v) or 1.0
        along = pw_ if horizontal else ph_
        thick = (ph_ if horizontal else pw_) / n / (1 + gap / 100)
        out: dict[tuple[int, int], bool] = {}
        for i, vs in enumerate(values):
            for j, v in enumerate(vs):
                if v is None or v == 0:
                    continue
                length = abs(v) / span * along
                tw = text_width_pt(k.plain(v), family, fs_val) * 1.08 + 4
                th = fs_val * 1.2
                out[(i, j)] = (tw + fs_val * 0.4 <= length and th <= thick * 0.92) if horizontal else (th + 2 <= length and tw <= thick)
        return out

    if stacked and labels_on and flat:
        fs_val = fs if multi else fs_val
        inside = fit_inside(pw, ph, min(0.0, min(bots)), max(tops) or 1.0)
        if not all(inside.values()):
            axis_on = True
            left, top, pw, ph = box(labels_on, axis_on, fs_val)
            lo_n, hi_n, _ = _nice_scale(min(0.0, min(ext)), max(0.0, max(ext)))
            inside = fit_inside(pw, ph, lo_n, hi_n)
        if not any(inside.values()):
            labels_on = False

    plot = k.plot
    plot.gap_width = gap
    plot.overlap = 100 if stacked else (-8 if multi else 0)
    plot.vary_by_categories = False
    if hist and k.chart_el is not None:
        _set_child_val(k.chart_el, "overlap", "0", ("serLines", "axId", "extLst"))  # (written out: the bins never overlap)

    def seam(fmt) -> None:
        # a histogram's bins and stacked segments are set apart by a hairline of the ground; other bars have no edge
        if hist or stacked:
            fmt.line.color.rgb = _rgb(k.ground)
            fmt.line.width = Pt(SEAM_PT)
        else:
            fmt.line.fill.background()

    # colours: the story's paint for several series; one series steps back behind its highlight
    muted_fill = muted_tint(k.accent, k.ground)
    for i, ser in enumerate(plot.series):
        ser.invert_if_negative = False
        st = k.styles[i] if multi else {"fill": muted_fill if hl is not None else k.accent}
        if st.get("fill"):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _rgb(st["fill"])
            seam(ser.format)
        else:  # a plan: the outline of the bar it aims for
            ser.format.fill.background()
            ser.format.line.color.rgb = _rgb(st.get("line") or k.accent)
            ser.format.line.width = Pt(1.5)
        if not multi and hl is not None:
            pt = ser.points[hl]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(k.accent)
            seam(pt.format)
        if cat_hl is not None and st.get("fill"):
            back = eased(st["fill"], k.ground)
            for j, v in enumerate(values[i]):
                if j != cat_hl and v is not None:
                    pt = ser.points[j]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = _rgb(back)
                    seam(pt.format)

    # labels
    plot.has_data_labels = labels_on
    if labels_on and k.chart_el is not None:
        pos = "ctr" if stacked else "outEnd"
        _set_plot_dlbls(k.chart_el, _dlbls(k.fmt_unit, fs_val, k.text, family, pos))
        for i, ser in enumerate(plot.series):
            st = k.styles[i] if multi else {"label": k.text}
            lab = st.get("label", k.text)
            pts: dict[int, dict] = {}
            if stacked:
                fill = st.get("fill") or k.accent
                for j, v in enumerate(values[i]):
                    if v is None or v == 0 or not inside.get((i, j)):
                        pts[j] = {"delete": True}
                    elif cat_hl is not None and j != cat_hl:
                        pts[j] = {"color": label_on(eased(fill, k.ground), k.text, k.ground)}
                _set_ser_dlbls(ser._element, _dlbls(k.fmt_plain, fs_val, label_on(fill, k.text, k.ground), family, pos, pts))
                continue
            for j, v in enumerate(values[i]):
                if v is None:
                    pts[j] = {"delete": True}
                elif v == 0:
                    # a real zero: a column simply has none; a bar row says «0» at the axis instead of looking empty —
                    # unless another series' bar fills the row (a «0» beside «38» of the other series is noise)
                    row_shown = multi and any((values[o][j] or 0) != 0 for o in range(len(values)) if o != i and j < len(values[o]))
                    pts[j] = {"fmt": "0", "color": k.muted} if horizontal and not row_shown else {"delete": True}
                elif cat_hl is not None and j != cat_hl:
                    pts[j] = {"color": k.muted}
            if not multi and hl is not None and hl not in pts:
                pts[hl] = k.hl_label(fs_val / k.fs_value if k.fs_value else 1.0)
            elif not multi and emph is not None and emph not in pts:
                pts[emph] = {"fmt": k.fmt_emph}
            if direct and (0 not in pts or (cat_hl is not None and set(pts[0]) == {"color"})):
                fill = st.get("fill") or st.get("line") or lab
                color = fill if contrast_ratio(fill, k.ground) >= 3.0 else lab
                if cat_hl is not None and cat_hl != 0:
                    color = k.muted  # (the first row is not the highlighted category: its names step back with it)
                pts[0] = {"sername": True, "color": color}
            _set_ser_dlbls(ser._element, _dlbls(k.fmt_plain, fs_val, lab, family, pos, pts))

    # axes: columns stand on one hairline baseline; horizontal bars start at a common edge and need none (unless < 0)
    _cat_axis(k, cat_fs, reverse=horizontal, line=not horizontal or has_neg, rotate=rotate_cols)
    lo = hi = step = None
    if flat and not has_neg:
        lo = 0.0
        if labels_on and not axis_on:
            hi = (max(tops) if stacked else max(flat)) or None
    if axis_on and flat:
        lo, hi, step = _nice_scale(min(0.0, min(ext)), max(0.0, max(ext)))
    _value_axis(k, axis_on, fs, lo, hi, step)
    if horizontal:
        _axis_crosses(k.chart.value_axis._element, "max")  # (the categories run top-down: «max» is the bottom edge)
    _plot_layout(k, left, top, pw, ph)


# ---------------------------------------------------------------------------------------------- lines & areas


def _seg_in_box(p, q, box) -> float:
    """Length of the segment pq inside the box (x0, y0, x1, y1) — Liang–Barsky clipping."""
    (x0, y0), (x1, y1) = p, q
    dx, dy = x1 - x0, y1 - y0
    t0, t1 = 0.0, 1.0
    for pk, qk in ((-dx, x0 - box[0]), (dx, box[2] - x0), (-dy, y0 - box[1]), (dy, box[3] - y0)):
        if abs(pk) < 1e-12:
            if qk < 0:
                return 0.0
            continue
        t = qk / pk
        if pk < 0:
            t0 = max(t0, t)
        else:
            t1 = min(t1, t)
    return (t1 - t0) * math.hypot(dx, dy) if t1 > t0 else 0.0


def _box_overlap(a, b) -> float:
    return max(0.0, min(a[2], b[2]) - max(a[0], b[0])) * max(0.0, min(a[3], b[3]) - max(a[1], b[1]))


_CANDIDATES = (("t", 0, 0.0), ("t", -1, 0.35), ("t", 1, 0.35), ("b", 0, 0.6), ("b", -1, 0.8), ("b", 1, 0.8), ("r", 0, 1.0), ("l", 0, 1.0))


def place_point_labels(points, sizes, bounds, gap: float, lines=()) -> list[Optional[tuple[str, float]]]:
    """Where each point's label goes: (dLblPos, horizontal nudge in pt). Above the point is preferred; a label that
    the line would strike (a steady rise or fall runs through a label centred above its point) moves to the free
    side — above-left on a rise, above-right on a fall, below at a dip — keeping clear of every line segment, the
    other labels and markers, and the `bounds` (x0, y0, x1, y1). `points` in pt, y downward; None = no label."""
    segs = []
    for pts in (list(lines) or [points]):
        segs += [(pts[i], pts[i + 1]) for i in range(len(pts) - 1) if pts[i] is not None and pts[i + 1] is not None]
    markers = [p for pts in (list(lines) or [points]) for p in pts if p is not None]
    placed: list[tuple] = []
    out: list[Optional[tuple[str, float]]] = []
    for j, (p, sz) in enumerate(zip(points, sizes)):
        if p is None or sz is None:
            out.append(None)
            continue
        x, y = p
        w, h = sz
        best = None
        for pos, nudge, pref in _CANDIDATES:
            dx = nudge * (w / 2 + gap * 0.6)
            if pos == "t":
                bx = (x + dx - w / 2, y - gap - h, x + dx + w / 2, y - gap)
            elif pos == "b":
                bx = (x + dx - w / 2, y + gap, x + dx + w / 2, y + gap + h)
            elif pos == "r":
                bx = (x + gap, y - h / 2, x + gap + w, y + h / 2)
            else:
                bx = (x - gap - w, y - h / 2, x - gap, y + h / 2)
            pad = (bx[0] - 1.5, bx[1] - 1.5, bx[2] + 1.5, bx[3] + 1.5)
            cost = pref
            cost += 3.0 * sum(_seg_in_box(a, b, pad) for a, b in segs) / max(h, 1.0)
            cost += 4.0 * sum(_box_overlap(pad, o) for o in placed) / max(w * h, 1.0)
            cost += 3.0 * sum(1 for q in markers if q != p and pad[0] <= q[0] <= pad[2] and pad[1] <= q[1] <= pad[3])
            out_of = max(0.0, bounds[0] - bx[0]) + max(0.0, bx[2] - bounds[2]) + max(0.0, bounds[1] - bx[1]) + max(0.0, bx[3] - bounds[3])
            cost += 4.0 * out_of / max(h, 1.0)
            if best is None or cost < best[0] - 1e-9:
                best = (cost, pos, dx, bx)
        placed.append(best[3])
        out.append((best[1], best[2]))
    return out


def _repel(ys: dict, h: float, lo: float, hi: float) -> dict:
    """1-D label repel: move the end labels (centres `ys`) apart until no two are closer than `h`, within [lo, hi]."""
    order = sorted(ys, key=lambda i: ys[i])
    pos = {i: ys[i] for i in order}
    for _ in range(40):
        moved = False
        for a, b in zip(order, order[1:]):
            d = pos[b] - pos[a]
            if d < h:
                push = (h - d) / 2 + 0.01
                pos[a] -= push
                pos[b] += push
                moved = True
        for i in order:
            pos[i] = max(lo + h / 2, min(hi - h / 2, pos[i]))
        if not moved:
            break
    return pos


def _area_overlay(k: _Kit, area_el, colors: list[str], lasts: Sequence[Optional[int]] = ()) -> list:
    """The top edge of each area as a 2.25 pt line series over it (a c:lineChart on the same axes): the area itself
    has no outline, so its sides and bottom are not stroked, and the edge carries the end label."""
    sers = area_el.findall(_c("ser"))
    lc = etree.Element(_c("lineChart"))
    etree.SubElement(lc, _c("grouping")).set("val", "standard")
    etree.SubElement(lc, _c("varyColors")).set("val", "0")
    out = []
    n = len(k.cats)
    for i, aser in enumerate(sers):
        ls = etree.SubElement(lc, _c("ser"))
        etree.SubElement(ls, _c("idx")).set("val", str(len(sers) + i))
        etree.SubElement(ls, _c("order")).set("val", str(len(sers) + i))
        tx = aser.find(_c("tx"))
        if tx is not None:
            ls.append(copy.deepcopy(tx))
        color = colors[i % len(colors)]
        ls.append(_sppr(None, line=color, line_w=2.25, no_line=False))
        mk = etree.SubElement(ls, _c("marker"))
        etree.SubElement(mk, _c("symbol")).set("val", "none")
        # the last point with a value: a ringed marker where the end label sits
        end = lasts[i] if i < len(lasts) and lasts[i] is not None else n - 1
        dpt = etree.SubElement(ls, _c("dPt"))
        etree.SubElement(dpt, _c("idx")).set("val", str(max(0, end)))
        m2 = etree.SubElement(dpt, _c("marker"))
        etree.SubElement(m2, _c("symbol")).set("val", "circle")
        etree.SubElement(m2, _c("size")).set("val", "7")
        m2.append(_sppr(color, line=k.ground, line_w=1.5, no_line=False))
        etree.SubElement(dpt, _c("bubble3D")).set("val", "0")
        for tag in ("cat", "val"):
            el = aser.find(_c(tag))
            if el is not None:
                ls.append(copy.deepcopy(el))
        etree.SubElement(ls, _c("smooth")).set("val", "0")
        out.append(ls)
    etree.SubElement(lc, _c("marker")).set("val", "1")
    for ax in area_el.findall(_c("axId")):
        lc.append(copy.deepcopy(ax))
    area_el.addnext(lc)
    return out


def _set_categories(k: _Kit, cats: list[str]) -> None:
    """The chart's categories rewritten (its data sheet too): the shortened labels of a crowded time axis."""
    data = CategoryChartData()
    data.categories = cats
    for ser, vals in zip(k.series, k.values):
        data.add_series(ser.name, vals)
    try:
        k.chart.replace_data(data)
    except Exception:  # noqa: BLE001
        pass


def _style_line(k: _Kit) -> None:
    """One series (≤ 8 points): every point labelled off the line, the highlight (or the last point) in the accent on a
    larger marker, no value axis. Several series or many points: a quiet value axis on round ticks and each line named
    once at its end, in its own colour, the end labels kept apart. An area: a soft fill under a 2.25 pt edge."""
    cats, values, series, W, H, fs, family = k.cats, k.values, k.series, k.W, k.H, k.fs, k.family
    kind = k.kind
    n = max(1, len(cats))
    ns = len(series)
    multi = ns > 1
    flat = [v for vs in values for v in vs if v is not None]
    has_neg = any(v < 0 for v in flat)
    # a line ends (and is labelled) at its last point with a value: a missing tail is a break, not a zero
    lasts = [last_value_index(vs) for vs in values]
    hl = None
    want = k.spec.highlight_index
    if not multi and kind == "line" and want != NO_HIGHLIGHT:
        hl = want if (want is not None and 0 <= want < n and values[0][want] is not None) else lasts[0]
    slot = W / n
    fs_val = k.fs_value if (not multi and n <= 6) else fs
    ratio = k.fs_hl / k.fs_value if k.fs_value else 1.0

    def point_text(i: int, j: int) -> str:
        v = values[i][j]
        return k.emph(v) if j == hl else k.plain(v)

    def widest(size: float) -> float:
        return max((text_width_pt(point_text(0, j), family, _snap(size * ratio) if j == hl else size, k.bold and j == hl) for j in range(n) if values[0][j] not in (None, 0)), default=0.0)

    # a label on every point while they stay few and apart; an area reads as a mass (one end label, a quiet axis)
    every_point = not multi and kind == "line" and n <= 8
    if every_point and widest(fs_val) > slot * 0.95:
        fs_val = fs
        if widest(fs_val) > slot * 1.05:
            every_point = False
    axis_on = multi or not every_point
    ends = not every_point
    # end labels name their line when there are several (no legend); a long series name falls back to a legend
    names = [s.name for s in series]
    named = multi and max((text_width_pt(nm, family, fs) for nm in names), default=0.0) <= W * 0.22
    if multi and not named:
        _legend(k, names, side=False)

    def end_text(i: int) -> str:
        last = values[i][lasts[i]] if lasts[i] is not None else None
        t = k.emph(last) if not multi else k.plain(last)
        return f"{names[i]}{LABEL_SEP}{t}" if named else t

    # one line: its end label is the chart's emphasis (accent, a size up where the template does not use bold)
    end_size = fs_val if multi else _snap(fs_val * ratio)
    end_w = max((text_width_pt(end_text(i), family, end_size, k.bold) for i in range(ns)), default=0.0) if ends else 0.0

    # category labels at the chart's text size, measured one by one against the real slot — the plot's width over the
    # points (the value axis and the end label take their share of the frame): straight on one line as written, else
    # shortened («1-й месяц» → «1-й мес.» → «1 мес.»), else turned by 45° with the room under the axis for their turned
    # height, else wrapped between whole words; a size smaller only when none of these fits (`fit_categories`)
    cat_fs = fs

    # value range
    lo = hi = step = None
    if flat:
        mn, mx = min(flat), max(flat)
        if axis_on:
            if has_neg:
                lo_raw = mn
            elif kind == "area" or mn <= 0.45 * mx:
                lo_raw = 0.0
            else:
                lo_raw = max(0.0, mn - (mx - mn) * 0.2)
            lo, hi, step = _nice_scale(lo_raw, max(mx, 0.0))
            if not has_neg:
                lo = max(0.0, lo)
        else:
            # no axis, so no false zero: the lowest point sits a fifth of the height above the category labels
            lo = 0.0 if (mn >= 0 and mn <= 0.45 * mx) else (mn - (mx - mn) * 0.25)
            if mn >= 0:
                lo = max(0.0, lo)
            hi = mx if mx > lo else None

    marker = 7 if multi else 8
    lab_h = fs_val * 1.2
    top = k.title_h + ((max(fs_val, _snap(fs_val * ratio)) * 1.25 + marker / 2 + 6) if every_point else (fs * 0.9 if ends else fs * 0.9))
    if axis_on:
        left = k.inset_l + _axis_room([lo or 0.0, hi or 0.0] + flat, k.unit, k.decimals, family, fs, 1.2)
    else:
        left = k.inset_l + max(1.0, widest(fs_val) / 2 - slot / 2 + 4)
    right = max(1.0, end_w + fs * 0.9 - slot / 2 + 6) if ends else max(1.0, widest(fs_val) / 2 - slot / 2 + 4)
    pw = max(W - left - right, W * 0.3)
    floor = max(8.0, 0.02 * k.sh / EMU_PER_PT)
    sizes = sorted({cat_fs, max(min(floor, cat_fs), _snap(cat_fs * 0.9)), max(min(floor, cat_fs), _snap(cat_fs * 0.82))}, reverse=True)
    room_h = H - top - k.foot
    below_max = max(0.0, room_h - max(H * 0.3, 0.42 * room_h))
    fit = fit_categories(cats, family, sizes, pw / n, below_max, turn_first=n >= 5)
    if fit.mode == "turn" and fit.labels and fit.labels[0]:
        # the first label turned hangs left of its point: the plot starts clear of it (and of a left axis title)
        hang = k.inset_l + TURN_SIN * (text_width_pt(fit.labels[0], family, fit.size) + CAT_LINE * fit.size) - pw / n / 2 + 2
        if hang > left:
            left = hang
            pw = max(W - left - right, W * 0.3)
    cat_fs = fit.size
    if fit.labels != list(cats):
        cats = k.cats = list(fit.labels)
        _set_categories(k, cats)
    bottom = k.foot + fit.below
    ph = max(H - top - bottom, H * 0.3)
    real_slot = pw / n
    k.cramped = W - left - right < pw - 0.5
    rotate = fit.mode == "turn"
    skip = 1
    k.squeeze = fit.longest / max(real_slot, 1.0)
    k.turned = rotate or fit.keep > 1  # turned or thinned out: every label shown is whole
    k.cat_fit = fit

    # paint
    area_lines = []
    colors = []
    for i, ser in enumerate(k.plot.series):
        st = k.styles[i]
        color = st.get("fill") or st.get("line") or k.accent
        dashed = st.get("role") == "plan"
        if dashed and k.roles:
            color = k.muted
        colors.append(color)
        ser.smooth = False
        if kind == "line":
            ln = ser.format.line
            ln.color.rgb = _rgb(color)
            ln.width = Pt(1.5 if dashed else 2.25)
            if dashed:
                ln.dash_style = MSO_LINE_DASH_STYLE.DASH
                ser.marker.style = XL_MARKER_STYLE.NONE
            else:
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = marker
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = _rgb(color)
                ser.marker.format.line.color.rgb = _rgb(k.ground)
                ser.marker.format.line.width = Pt(1.5)
        else:
            sp = _sppr(color, 0.28 if not _is_dark(k.ground) else 0.38)
            _replace(ser._element, "spPr", sp, ("invertIfNegative", "pictureOptions", "marker", "dPt", "dLbls", "trendline", "errBars", "cat", "val", "smooth", "extLst"))
    if hl is not None:
        mk = k.plot.series[0].points[hl].marker
        mk.style = XL_MARKER_STYLE.CIRCLE
        mk.size = 11
        mk.format.fill.solid()
        mk.format.fill.fore_color.rgb = _rgb(k.accent)
        mk.format.line.color.rgb = _rgb(k.ground)
        mk.format.line.width = Pt(2)
    if kind == "area" and k.chart_el is not None:
        area_lines = _area_overlay(k, k.chart_el, colors, lasts)
        k.plot.has_data_labels = False

    # geometry of the points in the chart frame (pt, y down): points sit in the middle of their category slots
    span = ((hi if hi is not None else (max(flat) if flat else 1.0)) - (lo if lo is not None else (min(flat) if flat else 0.0))) or 1.0
    base_lo = lo if lo is not None else (min(flat) if flat else 0.0)

    def xy(j: int, v: Optional[float]):
        if v is None:
            return None
        return left + (j + 0.5) * pw / n, top + ph * (1 - (v - base_lo) / span)

    lines_xy = [[xy(j, v) for j, v in enumerate(vs)] for vs in values]
    gap = marker / 2 + 3

    def drawn(i: int) -> list:
        # the line as drawn: broken at a missing value (`spec.gaps`), else running through it
        return lines_xy[i] if k.spec.gaps else [q for q in lines_xy[i] if q is not None]

    if k.chart_el is not None:
        if every_point:
            k.plot.has_data_labels = True
            sizes = []
            for j in range(n):
                v = values[0][j]
                if v in (None, 0):
                    sizes.append(None)
                    continue
                size = _snap(fs_val * ratio) if j == hl else fs_val
                sizes.append((text_width_pt(point_text(0, j), family, size, k.bold and j == hl) * 1.08 + 2, size * 1.2))
            bounds = (0.0, k.title_h, W, H - bottom + cat_fs * 0.4)
            placed = place_point_labels(lines_xy[0], sizes, bounds, gap, lines=[drawn(0)])
            _set_plot_dlbls(k.chart_el, _dlbls(k.fmt_unit, fs_val, k.text, family, "t"))
            pts: dict[int, dict] = {}
            for j, v in enumerate(values[0]):
                if v in (None, 0) or placed[j] is None:
                    pts[j] = {"delete": True}
                    continue
                pos, dx = placed[j]
                o = k.hl_label(fs_val / k.fs_value if k.fs_value else 1.0) if j == hl else {}
                o = {**o, "pos": pos}
                if dx:
                    o["dx"] = dx / W
                pts[j] = o
            _set_ser_dlbls(k.plot.series[0]._element, _dlbls(k.fmt_plain, fs_val, k.text, family, "t", pts))
        else:
            if kind == "line":
                k.plot.has_data_labels = True
                _set_plot_dlbls(k.chart_el, _dlbls(k.fmt_unit, fs_val, k.text, family, "r"))
            targets = area_lines if kind == "area" else [s._element for s in k.plot.series]
            # end labels right of the last point with a value, pushed apart vertically where two lines end close together
            ends_y = {i: lines_xy[i][lasts[i]][1] for i in range(ns) if lasts[i] is not None and values[i][lasts[i]] != 0}
            moved = _repel(ends_y, lab_h * 1.05, k.title_h, H - bottom) if ends_y else {}
            for i, ser_el in enumerate(targets):
                st = k.styles[i]
                color = colors[i]
                lab_color = color if contrast_ratio(color, k.ground) >= 3.0 else (k.text if st.get("role") != "past" else k.muted)
                pts = {}
                if i in ends_y:
                    o = {"color": lab_color, "bold": k.bold, "pos": "r", "fmt": k.fmt_emph if not multi else k.fmt_plain, "size": end_size}
                    if named:
                        o["sername"] = True
                    dy = moved[i] - ends_y[i]
                    if abs(dy) > 0.5:
                        o["dy"] = dy / H
                    pts[lasts[i]] = o
                if hl is not None and hl != lasts[i] and values[i][hl] not in (None, 0):
                    p = lines_xy[i][hl]
                    size = _snap(fs_val * ratio)
                    w = text_width_pt(k.emph(values[i][hl]), family, size, k.bold) * 1.08 + 2
                    (pos, dx), = place_point_labels([p], [(w, size * 1.2)], (0.0, k.title_h, W, H - bottom), gap, lines=[drawn(i)])
                    o = {**k.hl_label(fs_val / k.fs_value if k.fs_value else 1.0), "pos": pos}
                    if dx:
                        o["dx"] = dx / W
                    pts[hl] = o
                _set_ser_dlbls(ser_el, _dlbls(k.fmt_plain, fs_val, lab_color, family, None, pts, show_val=False))

    # the baseline implies zero: without a value axis over a raised minimum there is none
    _cat_axis(k, cat_fs, skip=skip, line=axis_on or not (lo and lo > 0), rotate=rotate)
    _value_axis(k, axis_on, fs, lo, hi, step)
    ax = k.chart.category_axis._element
    _set_child_val(ax, "lblOffset", "100", ("tickLblSkip", "tickMarkSkip", "noMultiLvlLbl", "extLst"))
    va = k.chart.value_axis._element
    _set_child_val(va, "crossBetween", "between", ("majorUnit", "minorUnit", "dispUnits", "extLst"))
    _plot_layout(k, left, top, pw, ph)


# ---------------------------------------------------------------------------------------------- scatter


def _scatter_range(vals: Sequence[float], max_steps: int = 9, anchor: float = 0.15) -> tuple[float, float, float]:
    """(min, max, step) of a scatter axis on round ticks. A scatter has no baseline whose length a bar's would carry, so
    the axis starts at zero only when the data sit that near it (its lowest point within `anchor` of its highest), else
    at a round value under the lowest point; the step is the round one that wastes the least of the axis (up to
    `max_steps` intervals, as many as the frame's labels have room for): 137…590 reads 100…600, not 0…800."""
    mn, mx = min(vals), max(vals)
    span = mx - mn
    if mn >= 0 and mn <= anchor * mx:
        lo_raw = 0.0
    else:
        lo_raw = mn - span * 0.05
    hi_raw = mx + span * 0.02  # (a marker's radius of clearance at the high edge)
    if hi_raw <= lo_raw:
        hi_raw = lo_raw + (abs(lo_raw) or 1.0)
    width = hi_raw - lo_raw
    best = None
    mag = 10 ** math.floor(math.log10(width / 4.5))
    for m in (0.1, 1, 10):
        for base in (1, 2, 2.5, 5):
            step = base * mag * m
            a = math.floor(lo_raw / step + 1e-9) * step
            b = math.ceil(hi_raw / step - 1e-9) * step
            n = round((b - a) / step)
            if n < 3 or n > max(3, max_steps):
                continue
            cost = abs(n - 5.0) * 0.18 + 3.0 * ((b - a) - width) / (b - a) + (0.3 if base == 2.5 else 0.0)
            if best is None or cost < best[0]:
                best = (cost, a, b, step)
    if best is None:
        lo, hi, step = _nice_scale(lo_raw, hi_raw)
    else:
        _, lo, hi, step = best
    if mn >= 0:
        lo = max(0.0, lo)
    return lo, hi, step


def _label_spot(me, wh, pts, marker_r: float, own_r: float, bounds, gap: float):
    """Where the label of the highlighted point stands in a dense cloud: the box (w, h) nearest to the point whose
    area no other marker touches, within a few lines' height of it and inside `bounds` — else the box that touches
    the fewest. Returns (centre x, centre y, markers touched). `pts`: every marker centre (pt, y down), `marker_r` their
    radius, `own_r` the highlighted marker's (the label keeps `gap` clear of it)."""
    w, h = wh
    mx, my = me
    step = max(h / 2.0, 3.0)
    reach_x, reach_y = w + 5 * h, 5 * h
    near = [q for q in pts if q != me and abs(q[0] - mx) <= reach_x + w + 2 * marker_r and abs(q[1] - my) <= reach_y + h + 2 * marker_r]
    best = None
    nx, ny = int(reach_x / step), int(reach_y / step)
    for iy in range(-ny, ny + 1):
        for ix in range(-nx, nx + 1):
            cx, cy = mx + ix * step, my + iy * step
            bx = (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2)
            # clear of the highlighted marker itself
            ex = max(bx[0] - mx, 0.0, mx - bx[2])
            ey = max(bx[1] - my, 0.0, my - bx[3])
            if math.hypot(ex, ey) < own_r + gap * 0.5:
                continue
            out_of = max(0.0, bounds[0] - bx[0]) + max(0.0, bx[2] - bounds[2]) + max(0.0, bounds[1] - bx[1]) + max(0.0, bx[3] - bounds[3])
            if out_of > 0.5:
                continue
            touched = sum(1 for q in near if bx[0] - marker_r <= q[0] <= bx[2] + marker_r and bx[1] - marker_r <= q[1] <= bx[3] + marker_r)
            dist = math.hypot(ex, ey)
            cost = touched * 100.0 + dist
            if best is None or cost < best[0]:
                best = (cost, cx, cy, touched)
    return best[1:] if best else None


def _style_scatter(k: _Kit) -> None:
    """Points only, no lines: translucent ringed markers (the accent; several series the template's palette), two
    quiet value axes on round ticks — the Y axis with gridlines and the unit on its ticks or in its title, the X axis
    on a hairline baseline. The highlighted point is larger, in the accent, with its value beside it; the others step
    back into the accent's tint. No other labels: the axes carry the values."""
    W, H, fs, family = k.W, k.H, k.fs, k.family
    ns = len(k.series)
    multi = ns > 1
    xs_flat = [x for xs in k.xs for x in xs] or [0.0]
    ys_flat = [y for ys in k.values for y in ys] or [0.0]
    npts = len(ys_flat)
    # as many round steps as the frame's labels have room for: a tick label wants ~5.5 sizes of width, ~2.6 of height
    x_lo, x_hi, x_step = _scatter_range(xs_flat, max(4, min(9, int(W / (fs * 5.5)))))
    y_lo, y_hi, y_step = _scatter_range(ys_flat, max(4, min(9, int(H / (fs * 2.6)))))
    want = k.spec.highlight_index
    hl = want if (want is not None and not multi and 0 <= want < len(k.values[0])) else None
    size = 8 if npts <= 30 else (7 if npts <= 120 else 5)
    alpha = SCATTER_ALPHA[0] if npts <= 30 else SCATTER_ALPHA[1]

    x_dec = tick_decimals(x_step, _decimals(xs_flat))
    x_ticks = []
    v = x_lo
    while v <= x_hi + x_step * 1e-6 and len(x_ticks) < 20:
        x_ticks.append(_tick_text(v, None, x_dec))
        v += x_step
    widest_x = max((text_width_pt(t, family, fs) for t in x_ticks), default=fs)
    left = k.inset_l + _axis_room([y_lo, y_hi] + ys_flat, k.unit, k.decimals, family, fs, 0.8)
    right = max(2.0, widest_x / 2 + 2)
    top = k.title_h + fs * 0.9
    bottom = k.foot + fs * CAT_LINE + fs * 0.6
    pw = max(W - left - right, W * 0.3)
    ph = max(H - top - bottom, H * 0.3)

    muted_fill = muted_tint(k.accent, k.ground)
    for i, ser in enumerate(k.plot.series):
        st = k.styles[i] if multi else {"fill": muted_fill if hl is not None else k.accent}
        color = st.get("fill") or st.get("line") or k.accent
        el = ser._element
        _replace(el, "spPr", _sppr(None, no_line=True), ("marker", "dPt", "dLbls", "trendline", "errBars", "xVal", "yVal", "smooth", "extLst"))
        mk = el.find(_c("marker"))
        if mk is None:
            mk = etree.Element(_c("marker"))
            _insert_before(el, mk, ("dPt", "dLbls", "trendline", "errBars", "xVal", "yVal", "smooth", "extLst"))
        for child in list(mk):
            mk.remove(child)
        etree.SubElement(mk, _c("symbol")).set("val", "circle")
        etree.SubElement(mk, _c("size")).set("val", str(size))
        mk.append(_sppr(color, alpha, line=k.ground, line_w=0.75, no_line=False))
        _set_child_val(el, "smooth", "0", ("extLst",))
        if hl is not None and i == 0:
            dpt = etree.Element(_c("dPt"))
            etree.SubElement(dpt, _c("idx")).set("val", str(hl))
            m2 = etree.SubElement(dpt, _c("marker"))
            etree.SubElement(m2, _c("symbol")).set("val", "circle")
            etree.SubElement(m2, _c("size")).set("val", str(size + 4))
            m2.append(_sppr(k.accent, None, line=k.ground, line_w=1.5, no_line=False))
            etree.SubElement(dpt, _c("bubble3D")).set("val", "0")
            _insert_before(el, dpt, ("dLbls", "trendline", "errBars", "xVal", "yVal", "smooth", "extLst"))
            if k.chart_el is not None:
                # its value beside it, on the side clear of the other points (`place_point_labels`, markers only)
                def at(x: float, y: float) -> tuple[float, float]:
                    return left + (x - x_lo) / ((x_hi - x_lo) or 1.0) * pw, top + ph * (1 - (y - y_lo) / ((y_hi - y_lo) or 1.0))

                o = k.hl_label()
                pts = [at(x, y) for xs, ys in zip(k.xs, k.values) for x, y in zip(xs, ys)]
                me = at(k.xs[0][hl], k.values[0][hl])
                w = text_width_pt(k.emph(k.values[0][hl]), family, o["size"], k.bold) * 1.08 + 2
                lh = o["size"] * 1.2
                bounds = (left, top, left + pw, top + ph)
                spot = _label_spot(me, (w, lh), pts, size / 2.0 + 0.5, (size + 4) / 2.0 + 0.75, bounds, (size + 4) / 2 + 3) if npts > 30 else None
                if spot is None:
                    (pos, dx), = place_point_labels([me], [(w, lh)], (0.0, k.title_h, W, H - bottom), (size + 4) / 2 + 3, lines=[[q] for q in pts])
                    o = {**o, "pos": pos}
                    if dx:
                        o["dx"] = dx / W
                else:
                    # a dense cloud: the label stands on the nearest free ground (clear of the cloud where the room
                    # allows) and on a chip of the slide's ground where it must touch a marker — it stays legible
                    cx, cy, touched = spot
                    g = (size + 4) / 2 + 3
                    if abs(cx - me[0]) / (w / 2 + g) >= abs(cy - me[1]) / (lh / 2 + g):
                        pos = "r" if cx >= me[0] else "l"
                        bx_, by_ = me[0] + (g + w / 2) * (1 if cx >= me[0] else -1), me[1]
                    else:
                        pos = "b" if cy >= me[1] else "t"
                        bx_, by_ = me[0], me[1] + (g + lh / 2) * (1 if cy >= me[1] else -1)
                    o = {**o, "pos": pos}
                    if abs(cx - bx_) > 0.5:
                        o["dx"] = (cx - bx_) / W
                    if abs(cy - by_) > 0.5:
                        o["dy"] = (cy - by_) / H
                    if touched:
                        o["box"] = k.ground
                _set_ser_dlbls(el, _dlbls(k.fmt_plain, k.fs_value, k.hl_color, family, None, {hl: o}, show_val=False))
    if k.chart_el is not None:
        for d in k.chart_el.findall(_c("dLbls")):
            k.chart_el.remove(d)  # (the highlighted point's label is its series' own; no point else is labelled)

    # the X axis: ticks under a hairline baseline, no gridlines; the Y axis: quiet gridlines, the unit on its ticks
    xa = k.chart.category_axis
    xa.visible = True
    xa.has_major_gridlines = False
    xa.has_minor_gridlines = False
    xa.major_tick_mark = XL_TICK_MARK.NONE
    xa.minor_tick_mark = XL_TICK_MARK.NONE
    xa.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = x_lo, x_hi, x_step
    xa.format.line.color.rgb = _rgb(k.rule)
    xa.format.line.width = Pt(0.75)
    xa.tick_labels.number_format = _tick_format(None, x_dec)
    xa.tick_labels.number_format_is_linked = False
    _axis_txpr(xa, fs, k.muted, family)
    _value_axis(k, True, fs, y_lo, y_hi, y_step)
    k.chart.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    va = k.chart.value_axis._element
    # the Y axis crosses X at X's minimum: its ticks stand at the left edge, not in the middle of the points when X
    # spans zero (the X axis keeps crossing Y at zero: a negative Y has its hairline there)
    _axis_crosses(va, "min")
    _set_child_val(va, "crossBetween", "midCat", ("majorUnit", "minorUnit", "dispUnits", "extLst"))
    _plot_layout(k, left, top, pw, ph)


# ---------------------------------------------------------------------------------------------- pies


def _style_pie(k: _Kit) -> None:
    shares, cats, W, H, fs, family = k.values[0], k.cats, k.W, k.H, k.fs, k.family
    n = max(1, len(cats))
    ser = k.plot.series[0]
    total = sum(v for v in shares if v) or 1.0
    ring = k.kind == "doughnut"
    palette = k.palette
    for j in range(n):
        pt = ser.points[j]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(palette[j % len(palette)])
        pt.format.line.color.rgb = _rgb(k.ground)
        pt.format.line.width = Pt(1.5 if not ring else 2.0)
    fine = label_format("%", [0.5])  # one decimal, for a slice under 1 % only

    amounts = (k.unit or "%") != "%"

    def label(v: float) -> str:
        if amounts:
            return _fmt_value(v, k.unit, k.decimals)
        return _fmt_value(v, "%", max(1 if v < 1 else 0, _decimals([v]) if k.pct_given else 0))

    def geometry(outside: bool):
        m = fs * (2.6 if outside else 0.6)
        avail_w = W - k.legend_w - 2 * m
        avail_h = H - k.title_h - k.legend_h - 2 * m
        side = max(min(avail_w, avail_h), min(W, H) * 0.3)
        return m, avail_h, side

    # which labels fit inside their slice: the label against the chord of the slice where it sits (with headroom for
    # a substituted, wider font); a pie puts the others outside, a doughnut leaves them to the legend
    def decide(side: float, size: float) -> dict[int, str]:
        r = side / 2
        out: dict[int, str] = {}
        before = 0.0  # the share of the slices drawn before this one, clockwise from twelve o'clock
        for j, v in enumerate(shares):
            if not v:
                continue
            s = v / total
            w = text_width_pt(label(v), family, size, k.bold) * 1.15 + 6
            h = size * 1.2 + 2
            chord = lambda rho: 2 * rho * math.sin(min(math.pi * s, math.pi / 2)) if s < 0.5 else 2 * rho  # noqa: E731
            # the wedge narrows toward the centre: the chord is taken at the label's inner edge
            if ring:
                # a label is set level at any angle: at three and nine o'clock its width lies across the ring, where
                # «110 000 ₽» runs over the edge into the total in the hole (the width already carries the headroom
                # for a wider font: the ring's own tenth is the margin)
                a = 2 * math.pi * (before + s / 2)
                across = w * abs(math.sin(a)) + h * abs(math.cos(a))
                rho, band = r * (1 + HOLE_SIZE / 100) / 2, r * (1 - HOLE_SIZE / 100)
                out[j] = "in" if (chord(rho - h / 2) >= w and band * 0.8 >= h and band * 0.9 >= across) else "none"
            elif s >= 0.25 and chord(r * 0.5 - h / 2) >= w:
                out[j] = "ctr"
            elif chord(max(r - max(w, h) / 2 - 4, r * 0.4) - h / 2) >= w:
                out[j] = "inEnd"
            else:
                out[j] = "outEnd" if k.outside else "none"
            before += s
        return out

    size = k.fs_value
    m, avail_h, side = geometry(False)
    if k.capped:
        size = max(min(size, _snap(side / 8.0)), min(fs, 8.0))  # a slice label never larger than an eighth of the pie
    where = decide(side, size)
    if any(v in ("outEnd", "none") for v in where.values()) and size > fs:
        smaller = decide(side, fs)
        # a step down when it brings a label inside (without a legend of its own, a thin slice stays unlabelled
        # whatever the size: the rest keep the larger type)
        if k.outside or sum(v not in ("outEnd", "none") for v in smaller.values()) > sum(v not in ("outEnd", "none") for v in where.values()):
            size, where = fs, smaller
    outside = any(v == "outEnd" for v in where.values())
    if outside:
        m, avail_h, side = geometry(True)
        where = decide(side, size)
    labelled = sum(1 for v in where.values() if v not in ("outEnd", "none"))
    if not k.outside and 0 < labelled < len(where):
        # the caller's legend gives every part's value: figures on some slices and not on others (one of eight, the four
        # wide enough of eight) read as highlights nobody meant — the slices carry all their figures or none
        where = {j: "none" for j in where}

    pts: dict[int, dict] = {}
    for j, v in enumerate(shares):
        if not v or where.get(j) == "none" or not k.slice_labels:
            pts[j] = {"delete": True}
            continue
        fill = palette[j % len(palette)]
        if where[j] == "outEnd":
            pts[j] = {"color": k.text, "pos": "outEnd"}
        else:
            pts[j] = {"color": label_on(fill, k.text, k.ground)}
            if not ring:
                pts[j]["pos"] = where[j]
        if v < 1 and not amounts:
            pts[j]["fmt"] = fine
        elif k.pct_given and not amounts and _decimals([v]):
            pts[j]["fmt"] = label_format("%", [v])  # G4-11: «18,6%» as the brief writes it, «42%» beside it
    pos = None if ring else "ctr"
    k.plot.has_data_labels = k.slice_labels
    if k.chart_el is not None:
        if k.slice_labels:
            _set_plot_dlbls(k.chart_el, _dlbls(k.fmt_unit, size, k.text, family, pos, leader_lines=False, bold=k.bold))
            _set_ser_dlbls(ser._element, _dlbls(k.fmt_unit, size, k.text, family, pos, pts, leader_lines=False, bold=k.bold))
        for tag, val in (("firstSliceAng", "0"), ("holeSize", str(HOLE_SIZE))):
            el = k.chart_el.find(_c(tag))
            if el is not None:
                el.set("val", val)
            elif tag == "firstSliceAng":
                _set_child_val(k.chart_el, tag, val, ("holeSize", "extLst"))
    # the circle and its legend as one group: left-aligned in a column, centred when the chart has the slide to itself
    group = side + (fs * 2.0 + k.legend_w if k.legend_w else 0.0)
    x = (W - group) / 2 if (not k.legend_w or W > 0.7 * k.sw / EMU_PER_PT) else m
    y = k.title_h + m + max(0.0, (avail_h - side) / 2)
    _plot_layout(k, x, y, side, side)
    if k.legend_w and k.chart.has_legend:
        lg = k.chart._chartSpace.find(f".//{_c('legend')}")
        if lg is not None:
            lx = (x + side + fs * 2.0) / W
            _manual_layout(lg, lx, 0.5 - min(0.45, (n * fs * 1.5) / H / 2), min(1.0 - lx, k.legend_w / W + 0.02), min(0.9, (n * fs * 1.5 + fs) / H), inner=False)


# ---------------------------------------------------------------------------------------------- geometry for callers


def chart_plot_bbox(graphic_frame) -> Optional[Bbox]:
    """The plot area of a chart made here, in slide EMU (from its manual layout) — e.g. to align a caption with it."""
    try:
        cs = graphic_frame.chart._chartSpace
        ml = cs.find(f".//{_c('plotArea')}/{_c('layout')}/{_c('manualLayout')}")
        vals = {_local(e): float(e.get("val")) for e in ml if e.get("val") not in (None, "inner", "outer", "edge", "factor")}
        gx, gy, gw, gh = graphic_frame.left, graphic_frame.top, graphic_frame.width, graphic_frame.height
        return Bbox(x=int(gx + vals["x"] * gw), y=int(gy + vals["y"] * gh), w=int(vals["w"] * gw), h=int(vals["h"] * gh))
    except Exception:  # noqa: BLE001
        return None


def doughnut_hole_bbox(graphic_frame) -> Optional[Bbox]:
    """The square inscribed in a doughnut's hole, in slide EMU: where a composer can set the total or the lead share
    as a native text box."""
    pb = chart_plot_bbox(graphic_frame)
    if pb is None:
        return None
    d = min(pb.w, pb.h)
    inner = d * HOLE_SIZE / 100 / math.sqrt(2)
    cx, cy = pb.x + pb.w / 2, pb.y + pb.h / 2
    return Bbox(x=int(cx - inner / 2), y=int(cy - inner / 2), w=int(inner), h=int(inner))
