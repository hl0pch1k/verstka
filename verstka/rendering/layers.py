"""What a slide really paints: master → layout → slide layers (contract C2 of TEMPLATE_FIX_PLAN, T06/T12).

The engine used to see only a slide's own shapes; bands, panels, illustrations and photo grounds drawn by the layout
or the master were invisible, so content was laid over art and headings spilled out of their bands. This module
lists every drawn layer in z-order and answers three questions about them:

* `is_ground(layer, title_box, body_box)` — a layer text may stand on (a calm full-bleed ground, or a painted panel
  that holds the title or the body);
* `art_boxes(slide, skip=…)` / `free_rect(slide, area, skip=…, pad=…)` — the template art content must avoid, and the
  largest empty rectangle inside an area;
* `heading_band(slide, title_box)` — the painted band a heading sits in (a hard limit for the heading).

EMU `Bbox` in, `Bbox` out; pure lxml/python-pptx; layout and master layers are cached per part.
"""

from __future__ import annotations

import io
import logging
import weakref
from dataclasses import dataclass, replace
from typing import Iterable, Optional

from lxml import etree

from verstka.analysis.xmlns import q
from verstka.schemas.common import Bbox

log = logging.getLogger(__name__)

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_DRAWN = ("sp", "pic", "grpSp", "graphicFrame", "cxnSp")
_FILLS = ("solidFill", "gradFill", "pattFill", "blipFill")

FULL_BLEED = 0.85  # share of the slide a layer covers to count as full-bleed
BUSY_STDEV = 40.0  # luminance stdev (0–255) of a 32×18 thumbnail above which a picture is a photo / busy art
PANEL_MIN = 0.25  # share of the slide a painted panel covers to be a ground when it holds the title or the body
RULE_H = 0.02  # a layer at most this share of the slide height is a rule, not art
ART_MIN = 0.003  # art smaller than this share of the slide is ignored (dots, tiny marks)
BAND_MAX_H = 0.30
BAND_MIN_W = 0.30


@dataclass
class Layer:
    el: etree._Element
    box: Bbox
    source: str  # "master" | "layout" | "slide"
    kind: str  # "sp" | "pic" | "grpSp" | "graphicFrame" | "cxnSp"
    placeholder: bool = False
    paints: bool = False  # a visible fill, outline or picture
    picture: bool = False  # a picture or a picture-filled shape (or a group holding one)
    busy: bool = False  # a photo-like picture (luminance stdev of the thumbnail > BUSY_STDEV)
    has_text: bool = False
    full_bleed: bool = False
    custom_geom: bool = False
    fill_hex: Optional[str] = None  # the resolved fill colour (solid, gradient mean, pattern blend, picture mean)
    shape_id: Optional[str] = None
    ph_type: Optional[str] = None
    slide_w: int = 0
    slide_h: int = 0

    @property
    def cover(self) -> float:
        """Share of the slide the layer covers."""
        return _cover(self.box, self.slide_w, self.slide_h)


@dataclass
class Band:
    box: Bbox
    el: etree._Element
    source: str
    fill_hex: Optional[str] = None


# ---------------------------------------------------------------------------------------------------- geometry


def _cover(box: Bbox, W: int, H: int) -> float:
    if W <= 0 or H <= 0:
        return 0.0
    ix = max(0, min(box.x2, W) - max(box.x, 0))
    iy = max(0, min(box.y2, H) - max(box.y, 0))
    return ix * iy / float(W * H)


def _contains(outer: Bbox, inner: Bbox, share: float = 0.9) -> bool:
    return inner.area > 0 and outer.intersection(inner) >= share * inner.area


def _clip(box: Bbox, W: int, H: int) -> Optional[Bbox]:
    x, y = max(box.x, 0), max(box.y, 0)
    x2, y2 = min(box.x2, W), min(box.y2, H)
    if x2 <= x or y2 <= y:
        return None
    return Bbox(x=x, y=y, w=x2 - x, h=y2 - y)


def _xfrm_box(el: etree._Element) -> Optional[Bbox]:
    tag = etree.QName(el).localname
    if tag == "graphicFrame":
        xfrm = el.find(_P + "xfrm")
    elif tag == "grpSp":
        xfrm = el.find(_P + "grpSpPr/" + _A + "xfrm")
    else:
        xfrm = el.find(_P + "spPr/" + _A + "xfrm")
    off = xfrm.find(_A + "off") if xfrm is not None else None
    ext = xfrm.find(_A + "ext") if xfrm is not None else None
    if off is None or ext is None:
        return None
    try:
        return Bbox(x=int(off.get("x")), y=int(off.get("y")), w=int(ext.get("cx")), h=int(ext.get("cy")))
    except (TypeError, ValueError):
        return None


def _slide_size(obj) -> tuple[int, int]:
    try:
        prs = obj.part.package.presentation_part.presentation
        return int(prs.slide_width), int(prs.slide_height)
    except Exception:  # noqa: BLE001
        return 12192000, 6858000


# ---------------------------------------------------------------------------------------------------- paint


def _ph_of(el: etree._Element) -> Optional[etree._Element]:
    for nv in ("nvSpPr", "nvPicPr", "nvGraphicFramePr", "nvCxnSpPr"):
        ph = el.find(_P + nv + "/" + _P + "nvPr/" + _P + "ph")
        if ph is not None:
            return ph
    return None


def _shape_id(el: etree._Element) -> Optional[str]:
    for nv in ("nvSpPr", "nvPicPr", "nvGraphicFramePr", "nvCxnSpPr", "nvGrpSpPr"):
        c = el.find(_P + nv + "/" + _P + "cNvPr")
        if c is not None:
            return c.get("id")
    return None


def _has_text(el: etree._Element) -> bool:
    return any((t.text or "").strip() for t in el.iter(_A + "t"))


def _style_fill(el: etree._Element) -> Optional[etree._Element]:
    ref = el.find(_P + "style/" + _A + "fillRef")
    if ref is None or (ref.get("idx") or "0") == "0" or not len(ref):
        return None
    return ref


def _line_paints(sppr: Optional[etree._Element], el: etree._Element) -> bool:
    ln = sppr.find(_A + "ln") if sppr is not None else None
    if ln is not None:
        if ln.find(_A + "noFill") is not None:
            return False
        if ln.find(_A + "solidFill") is not None or ln.find(_A + "gradFill") is not None:
            return True
    ref = el.find(_P + "style/" + _A + "lnRef")
    return ref is not None and (ref.get("idx") or "0") != "0" and (ln is None or ln.find(_A + "noFill") is None)


class _Paint:
    """Resolves fills through the deck's theme (the helpers of `charts`, which already handle schemeClr + mods)."""

    def __init__(self, holder) -> None:
        from verstka.rendering import charts as _ch

        self._ch = _ch
        self.scheme: dict[str, str] = {}
        try:
            master = holder if _kind_of(holder) == "master" else (holder.slide_master if _kind_of(holder) == "layout" else holder.slide_layout.slide_master)
            self.scheme = _theme_of(master)
        except Exception:  # noqa: BLE001
            self.scheme = {}

    def color(self, clr: Optional[etree._Element]) -> Optional[tuple[str, float]]:
        return self._ch._color_choice(clr, self.scheme) if clr is not None else None

    def fill(self, sppr: Optional[etree._Element], el: etree._Element, part) -> tuple[bool, bool, Optional[str]]:
        """(paints, picture, hex) of a shape's own fill."""
        if sppr is not None:
            for child in sppr:
                kind = etree.QName(child).localname
                if kind == "noFill":
                    return False, False, None
                if kind == "pattFill":
                    fgc, bgc = child.find(_A + "fgClr"), child.find(_A + "bgClr")
                    fg = self.color(fgc[0]) if fgc is not None and len(fgc) else None
                    bg = self.color(bgc[0]) if bgc is not None and len(bgc) else None
                    cols = [c for c in (fg, bg) if c and c[1] >= 0.3]
                    if not cols:
                        return False, False, None
                    if len(cols) == 1:
                        return True, False, cols[0][0]
                    return True, False, self._ch._mix(cols[0][0], cols[1][0], 0.5)
                if kind in ("solidFill", "gradFill"):
                    hx = self._ch._fill_color(sppr, self.scheme, part)
                    return (hx is not None), False, hx
                if kind == "blipFill":
                    blip = child.find(_A + "blip")
                    rid = blip.get(_R_EMBED) if blip is not None else None
                    hx = self._ch._picture_mean(part, rid) if rid else None
                    return True, True, hx
                if kind == "grpFill":
                    return False, False, None
        ref = _style_fill(el)
        if ref is not None:
            cc = self.color(ref[0])
            if cc and cc[1] >= 0.5:
                return True, False, cc[0]
        return False, False, None


_THEME_CACHE: "weakref.WeakKeyDictionary" = weakref.WeakKeyDictionary()


def _theme_of(master) -> dict[str, str]:
    """clrScheme of the master's theme (dk1/lt1/…/accent6) plus its colour map (bg1 → lt1 …), cached per master."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    part = master.part
    got = _THEME_CACHE.get(part)
    if got is not None:
        return got
    out: dict[str, str] = {}
    try:
        theme = etree.fromstring(part.part_related_by(RT.THEME).blob)
        scheme = theme.find(".//" + _A + "clrScheme")
        for child in scheme if scheme is not None else []:
            if not isinstance(child.tag, str) or not len(child):
                continue
            clr = child[0]
            val = clr.get("val") if etree.QName(clr).localname == "srgbClr" else clr.get("lastClr")
            if val and len(val) == 6:
                out[etree.QName(child).localname] = val.upper()
        cmap = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
        cm = master._element.find(_P + "clrMap")
        if cm is not None:
            cmap.update(dict(cm.attrib))
        for k, v in cmap.items():
            if v in out and k not in out:
                out[k] = out[v]
    except Exception:  # noqa: BLE001
        pass
    _THEME_CACHE[part] = out
    return out


def _kind_of(obj) -> str:
    from pptx.slide import Slide, SlideLayout, SlideMaster

    if isinstance(obj, Slide):
        return "slide"
    if isinstance(obj, SlideLayout):
        return "layout"
    if isinstance(obj, SlideMaster):
        return "master"
    return "?"


_BUSY_CACHE: dict[str, bool] = {}
_OPAQUE_CACHE: dict[tuple, Optional[tuple[float, float, float, float]]] = {}


def _image_key(img_part) -> str:
    return getattr(img_part, "sha1", None) or str(getattr(img_part, "partname", id(img_part)))


def picture_busy(part, rid: Optional[str]) -> bool:
    """A photo / busy art: the luminance stdev of a 32×18 thumbnail of its opaque pixels (over white) is above
    BUSY_STDEV (cached per image)."""
    if not rid or part is None:
        return False
    try:
        img_part = part.related_part(rid)
    except Exception:  # noqa: BLE001
        return False
    key = _image_key(img_part)
    if key in _BUSY_CACHE:
        return _BUSY_CACHE[key]
    busy = False
    try:
        from PIL import Image, ImageStat

        im = Image.open(io.BytesIO(img_part.blob))
        im.draft("RGB", (256, 144))
        im = im.convert("RGBA")
        bb = im.getchannel("A").point(lambda v: 255 if v > 40 else 0).getbbox()
        if bb:
            im = im.crop(bb)
        bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
        bg.alpha_composite(im)
        g = bg.convert("L").resize((32, 18))
        busy = ImageStat.Stat(g).stddev[0] > BUSY_STDEV
    except Exception:  # noqa: BLE001 - an unreadable picture (EMF/WMF/SVG) is not judged busy
        busy = False
    _BUSY_CACHE[key] = busy
    return busy


def opaque_box(part, pic: etree._Element, box: Bbox) -> Bbox:
    """The painted part of a picture: a transparent PNG laid over the whole slide (a glass cube, a supergraphic)
    only occupies its opaque pixels (crop `srcRect` respected; cached per image and crop)."""
    blip = pic.find(".//" + _A + "blip")
    rid = blip.get(_R_EMBED) if blip is not None else None
    if not rid or part is None:
        return box
    try:
        img_part = part.related_part(rid)
    except Exception:  # noqa: BLE001
        return box
    src = pic.find(".//" + _A + "srcRect")
    crop = tuple(int(src.get(k) or 0) / 100000 for k in ("l", "t", "r", "b")) if src is not None else (0.0, 0.0, 0.0, 0.0)
    key = (str(img_part.partname), crop)
    if key not in _OPAQUE_CACHE:
        frac = None
        try:
            from PIL import Image

            im = Image.open(io.BytesIO(img_part.blob))
            if "A" in im.getbands():
                iw, ih = im.size
                l, t, r, b = crop
                vis = (int(l * iw), int(t * ih), max(int((1 - r) * iw), int(l * iw) + 1), max(int((1 - b) * ih), int(t * ih) + 1))
                a = im.getchannel("A").crop(vis)
                bb = a.point(lambda v: 255 if v > 40 else 0).getbbox()
                vw, vh = vis[2] - vis[0], vis[3] - vis[1]
                frac = (bb[0] / vw, bb[1] / vh, (bb[2] - bb[0]) / vw, (bb[3] - bb[1]) / vh) if bb else (0.0, 0.0, 0.0, 0.0)
        except Exception:  # noqa: BLE001
            frac = None
        _OPAQUE_CACHE[key] = frac
    frac = _OPAQUE_CACHE[key]
    if frac is None:
        return box
    return Bbox(x=int(box.x + frac[0] * box.w), y=int(box.y + frac[1] * box.h), w=int(frac[2] * box.w), h=int(frac[3] * box.h))


# ---------------------------------------------------------------------------------------------------- layers


def _layer_of(el: etree._Element, source: str, part, paint: _Paint, W: int, H: int, owner=None) -> Optional[Layer]:
    tag = etree.QName(el).localname
    if tag not in _DRAWN:
        return None
    box = _xfrm_box(el)
    ph = _ph_of(el)
    if box is None and ph is not None and owner is not None:
        from verstka.rendering.deck import inherited_box

        shp = next((s for s in owner.shapes if s._element is el), None)
        got = inherited_box(shp) if shp is not None else None
        box = Bbox(x=got[0], y=got[1], w=got[2], h=got[3]) if got else None
    if box is None or box.w <= 0 and box.h <= 0:
        return None
    layer = Layer(el=el, box=box, source=source, kind=tag, placeholder=ph is not None, shape_id=_shape_id(el),
                  ph_type=((ph.get("type") or "body") if ph is not None else None), slide_w=W, slide_h=H)
    layer.has_text = _has_text(el)
    if tag == "pic":
        blip = el.find(_P + "blipFill/" + _A + "blip")
        rid = blip.get(_R_EMBED) if blip is not None else None
        if rid:
            layer.box = opaque_box(part, el, box)
        layer.picture = layer.paints = rid is not None or ph is None
        layer.busy = picture_busy(part, rid)
        layer.fill_hex = paint._ch._picture_mean(part, rid) if rid else None
    elif tag == "grpSp":
        subs = [_layer_of(c, source, part, paint, W, H) for c in el if etree.QName(c).localname in _DRAWN]
        subs = [s for s in subs if s is not None]
        layer.paints = any(s.paints for s in subs)
        layer.picture = any(s.picture for s in subs)
        layer.busy = any(s.busy for s in subs)
        layer.has_text = layer.has_text or any(s.has_text for s in subs)
        layer.custom_geom = any(s.custom_geom for s in subs)
    elif tag == "graphicFrame":
        layer.paints = True  # a table / chart / diagram draws
    else:
        sppr = el.find(_P + "spPr")
        paints, pic, hx = paint.fill(sppr, el, part)
        layer.picture = pic
        layer.fill_hex = hx
        layer.paints = paints or _line_paints(sppr, el)
        if pic:
            blip = sppr.find(_A + "blipFill/" + _A + "blip") if sppr is not None else None
            layer.busy = picture_busy(part, blip.get(_R_EMBED) if blip is not None else None)
        layer.custom_geom = sppr is not None and sppr.find(_A + "custGeom") is not None
    layer.full_bleed = layer.cover >= FULL_BLEED
    return layer


def _own_layers(holder, source: str, paint: _Paint, W: int, H: int, *, placeholders: bool) -> list[Layer]:
    tree = holder._element.find(_P + "cSld/" + _P + "spTree")
    out: list[Layer] = []
    if tree is None:
        return out
    for el in tree:
        if not isinstance(el.tag, str):
            continue
        if not placeholders and _ph_of(el) is not None:
            continue  # layout/master placeholders are prompts: a slide never shows them
        try:
            lay = _layer_of(el, source, holder.part, paint, W, H, owner=holder)
        except Exception:  # noqa: BLE001
            log.debug("layer of %s failed", _shape_id(el), exc_info=True)
            lay = None
        if lay is not None:
            out.append(lay)
    return out


_PART_CACHE: "weakref.WeakKeyDictionary" = weakref.WeakKeyDictionary()


def _cached_layers(holder, source: str, W: int, H: int) -> list[Layer]:
    part = holder.part
    got = _PART_CACHE.get(part)
    if got is None or got[0] != (W, H):
        got = ((W, H), _own_layers(holder, source, _Paint(holder), W, H, placeholders=False))
        _PART_CACHE[part] = got
    return list(got[1])


def _shows_master(el) -> bool:
    return (el.get("showMasterSp") or "1") not in ("0", "false")


def drawn_layers(slide_or_layout) -> list[Layer]:
    """Every drawn layer under and on a slide (or a layout), bottom first: the master's non-placeholder shapes (unless
    hidden by `showMasterSp="0"`), the layout's non-placeholder shapes, then the object's own shapes (placeholders
    included, flagged). A slide with `showMasterSp="0"` shows neither the layout's nor the master's shapes."""
    kind = _kind_of(slide_or_layout)
    W, H = _slide_size(slide_or_layout)
    out: list[Layer] = []
    try:
        if kind == "slide":
            layout = slide_or_layout.slide_layout
            if _shows_master(slide_or_layout._element):
                if _shows_master(layout._element):
                    out += _cached_layers(layout.slide_master, "master", W, H)
                out += _cached_layers(layout, "layout", W, H)
            out += _own_layers(slide_or_layout, "slide", _Paint(slide_or_layout), W, H, placeholders=True)
        elif kind == "layout":
            if _shows_master(slide_or_layout._element):
                out += _cached_layers(slide_or_layout.slide_master, "master", W, H)
            out += _own_layers(slide_or_layout, "layout", _Paint(slide_or_layout), W, H, placeholders=True)
        elif kind == "master":
            out += _own_layers(slide_or_layout, "master", _Paint(slide_or_layout), W, H, placeholders=True)
    except Exception:  # noqa: BLE001 - layers are advice: never break a render over them
        log.debug("drawn_layers failed", exc_info=True)
    return out


# ---------------------------------------------------------------------------------------------------- questions


def is_ground(layer: Layer, title_box: Optional[Bbox] = None, body_box: Optional[Bbox] = None) -> bool:
    """A layer text may stand on: a calm (not busy) full-bleed layer, or a painted panel ≥ PANEL_MIN of the slide that
    holds the title or the body box (without either box: a panel ≥ 0.6 of the slide)."""
    if layer.placeholder or not layer.paints or layer.kind in ("graphicFrame", "cxnSp"):
        return False
    cover = layer.cover
    if cover >= FULL_BLEED:
        return not layer.busy
    if layer.busy or cover < PANEL_MIN:
        return False
    boxes = [b for b in (title_box, body_box) if b is not None]
    if not boxes:
        return cover >= 0.6
    return any(_contains(layer.box, b, 0.9) for b in boxes)


def _title_body(layers: list[Layer], slide) -> tuple[Optional[Bbox], Optional[Bbox]]:
    """The title and body boxes of a slide/layout: its own title placeholder (else the layout's), its largest
    body/object placeholder (else the layout's)."""
    title = body = None
    own = [l for l in layers if l.placeholder]
    for l in own:
        if l.ph_type in ("title", "ctrTitle") and title is None:
            title = l.box
    bodies = [l for l in own if l.ph_type in ("body", "obj", "subTitle")]
    if bodies:
        body = max(bodies, key=lambda l: l.box.area).box
    if (title is None or body is None) and _kind_of(slide) == "slide":
        try:
            layout = slide.slide_layout
            W, H = _slide_size(slide)
            lay = _own_layers(layout, "layout", _Paint(layout), W, H, placeholders=True)
            if title is None:
                title = next((l.box for l in lay if l.placeholder and l.ph_type in ("title", "ctrTitle")), None)
            if body is None:
                lb = [l for l in lay if l.placeholder and l.ph_type in ("body", "obj")]
                body = max(lb, key=lambda l: l.box.area).box if lb else None
        except Exception:  # noqa: BLE001
            pass
    return title, body


def _skipped(layer: Layer, skip: set) -> bool:
    return layer.source == "slide" and (layer.shape_id in skip or layer.el in skip)


def _norm_skip(skip: Iterable) -> set:
    out = set()
    for s in skip or ():
        if isinstance(s, (str, int)):
            out.add(str(s))
        else:
            out.add(getattr(s, "_element", s))
    return out


def art_layers(slide, skip: Iterable = (), *, title_box: Optional[Bbox] = None, body_box: Optional[Bbox] = None) -> list[Layer]:
    """The template art of a slide: painted or picture layers that are not placeholders, not grounds, not rules
    (height ≤ RULE_H of the slide) and not tiny (< ART_MIN of the slide). `skip`: ids (or elements) of the slide's own
    shapes to ignore."""
    layers = drawn_layers(slide)
    if not layers:
        return []
    W, H = layers[0].slide_w, layers[0].slide_h
    skip_set = _norm_skip(skip)
    if title_box is None or body_box is None:
        t, b = _title_body(layers, slide)
        title_box = title_box or t
        body_box = body_box or b
    out = []
    for lay in layers:
        if lay.placeholder or _skipped(lay, skip_set) or not lay.paints:
            continue
        if lay.kind == "graphicFrame" and lay.source == "slide":
            continue  # the slide's own tables/charts are content, not art
        if lay.box.h <= RULE_H * H:
            continue
        if lay.cover < ART_MIN:
            continue
        if is_ground(lay, title_box, body_box):
            continue
        out.append(lay)
    return out


def art_boxes(slide, skip: Iterable = ()) -> list[Bbox]:
    """Boxes (clipped to the slide) of `art_layers`."""
    out = []
    for lay in art_layers(slide, skip):
        b = _clip(lay.box, lay.slide_w, lay.slide_h)
        if b is not None:
            out.append(b)
    return out


def largest_free(area: Bbox, obstacles: Iterable[Bbox]) -> Bbox:
    """The largest axis-aligned rectangle inside `area` that crosses none of `obstacles` (exact, on the compressed
    grid of the obstacle edges)."""
    obs = []
    for o in obstacles:
        ix = max(0, min(o.x2, area.x2) - max(o.x, area.x))
        iy = max(0, min(o.y2, area.y2) - max(o.y, area.y))
        if ix > 0 and iy > 0:
            obs.append(Bbox(x=max(o.x, area.x), y=max(o.y, area.y), w=ix, h=iy))
    if not obs:
        return area
    xs = sorted({area.x, area.x2} | {o.x for o in obs} | {o.x2 for o in obs})
    ys = sorted({area.y, area.y2} | {o.y for o in obs} | {o.y2 for o in obs})
    nx, ny = len(xs) - 1, len(ys) - 1
    free = [[True] * ny for _ in range(nx)]
    for i in range(nx):
        cx = (xs[i] + xs[i + 1]) / 2
        for j in range(ny):
            cy = (ys[j] + ys[j + 1]) / 2
            if any(o.x <= cx <= o.x2 and o.y <= cy <= o.y2 for o in obs):
                free[i][j] = False
    best = (0, None)
    for i1 in range(nx):
        rows = [True] * ny
        for i2 in range(i1, nx):
            col = free[i2]
            rows = [r and c for r, c in zip(rows, col)]
            width = xs[i2 + 1] - xs[i1]
            j = 0
            while j < ny:
                if not rows[j]:
                    j += 1
                    continue
                j0 = j
                while j < ny and rows[j]:
                    j += 1
                a = width * (ys[j] - ys[j0])
                if a > best[0]:
                    best = (a, (xs[i1], ys[j0], xs[i2 + 1], ys[j]))
    if best[1] is None:
        return Bbox(x=area.x, y=area.y, w=0, h=0)
    x0, y0, x1, y1 = best[1]
    return Bbox(x=int(x0), y=int(y0), w=int(x1 - x0), h=int(y1 - y0))


def free_rect(slide, area: Bbox, skip: Iterable = (), pad: Optional[int] = None) -> Bbox:
    """The largest empty axis-aligned rectangle inside `area` that keeps `pad` (default 0.015 of the slide width) away
    from every art box of the slide. `area` itself when no art touches it."""
    layers_art = art_layers(slide, skip)
    if not layers_art:
        return area
    W = layers_art[0].slide_w
    pad = int(0.015 * W) if pad is None else int(pad)
    obs = []
    for lay in layers_art:
        b = _clip(lay.box, lay.slide_w, lay.slide_h)
        if b is None:
            continue
        obs.append(Bbox(x=b.x - pad, y=b.y - pad, w=b.w + 2 * pad, h=b.h + 2 * pad))
    return largest_free(area, obs)


def art_share(slide, area: Bbox, skip: Iterable = ()) -> float:
    """Share of `area` covered by the slide's art (union estimated on a coarse grid)."""
    boxes = art_boxes(slide, skip)
    if not boxes or area.area <= 0:
        return 0.0
    n = 48
    hit = 0
    for i in range(n):
        cx = area.x + (i + 0.5) * area.w / n
        for j in range(n):
            cy = area.y + (j + 0.5) * area.h / n
            if any(b.x <= cx <= b.x2 and b.y <= cy <= b.y2 for b in boxes):
                hit += 1
    return hit / float(n * n)


def heading_band(slide, title_box: Bbox) -> Optional[Band]:
    """The painted band a heading stands in: the topmost painted non-placeholder layer (any source) that holds the
    title box's text start and vertical centre, at most BAND_MAX_H of the slide high (and more than a rule), at least
    BAND_MIN_W wide. Small pills stay with the renderer's own backing logic."""
    layers = drawn_layers(slide)
    if not layers or title_box is None:
        return None
    W, H = layers[0].slide_w, layers[0].slide_h
    px = title_box.x + min(title_box.w // 20, W // 100)
    py = title_box.y + title_box.h / 2
    for lay in reversed(layers):
        if lay.placeholder or not lay.paints or lay.kind in ("graphicFrame", "cxnSp") or lay.has_text:
            continue
        b = lay.box
        if b.h > BAND_MAX_H * H or b.h <= RULE_H * H or b.w < BAND_MIN_W * W:
            continue
        if b.x <= px <= b.x2 and b.y <= py <= b.y2:
            return Band(box=b, el=lay.el, source=lay.source, fill_hex=lay.fill_hex)
    return None


def with_box(layer: Layer, box: Bbox) -> Layer:
    return replace(layer, box=box)


# ---------------------------------------------------------------------------------------------------- gradients


def gradient_at(stops: list[tuple[float, tuple[float, float, float], float]], t: float) -> tuple[tuple[float, float, float], float]:
    """(rgb, alpha) of a gradient at position t ∈ [0, 1] given sorted (pos, rgb, alpha) stops; the ends beyond the
    first and last stop keep their colour."""
    t = max(0.0, min(1.0, t))
    if t <= stops[0][0]:
        return stops[0][1], stops[0][2]
    for (p0, c0, a0), (p1, c1, a1) in zip(stops, stops[1:]):
        if t <= p1:
            k = (t - p0) / (p1 - p0) if p1 > p0 else 0.0
            return (c0[0] + (c1[0] - c0[0]) * k, c0[1] + (c1[1] - c0[1]) * k, c0[2] + (c1[2] - c0[2]) * k), a0 + (a1 - a0) * k
    return stops[-1][1], stops[-1][2]


def gradient_position(grad: etree._Element, u: float, v: float, w: float, h: float) -> float:
    """Where a point (u, v) — fractions of the shape's box — falls along a DrawingML gradient (0 = first stop).

    Linear (`a:lin ang`): the projection on the gradient's direction over the box's extent along it (in the unit
    square when `scaled`). Path (`a:path circle|rect|shape`): the distance from the focus rectangle (`a:fillToRect`)
    scaled to the box — an ellipse for «circle», a rectangle otherwise — as LibreOffice paints them."""
    import math

    path = grad.find(_A + "path")
    if path is not None:
        ftr = path.find(_A + "fillToRect")

        def get(k: str) -> float:
            try:
                return int(ftr.get(k) or 0) / 100000.0 if ftr is not None else 0.5
            except ValueError:
                return 0.5

        x0, x1 = get("l"), 1.0 - get("r")
        y0, y1 = get("t"), 1.0 - get("b")
        if x0 > x1:
            x0 = x1 = (x0 + x1) / 2
        if y0 > y1:
            y0 = y1 = (y0 + y1) / 2
        dx = max(x0 - u, 0.0, u - x1)
        dy = max(y0 - v, 0.0, v - y1)
        rx = max(x0, 1.0 - x1, 1e-6)
        ry = max(y0, 1.0 - y1, 1e-6)
        if (path.get("path") or "circle") == "circle":
            return min(1.0, math.hypot(dx / rx, dy / ry))
        return min(1.0, max(dx / rx, dy / ry))
    lin = grad.find(_A + "lin")
    try:
        ang = math.radians((int(lin.get("ang") or 0) if lin is not None else 0) / 60000.0)
    except ValueError:
        ang = 0.0
    scaled = lin is not None and lin.get("scaled") in ("1", "true")
    sx, sy = (1.0, 1.0) if scaled or w <= 0 or h <= 0 else (w / max(w, h), h / max(w, h))
    dx, dy = math.cos(ang), math.sin(ang)
    extent = abs(sx * dx) + abs(sy * dy)
    if extent <= 1e-9:
        return 0.5
    return max(0.0, min(1.0, 0.5 + ((u - 0.5) * sx * dx + (v - 0.5) * sy * dy) / extent))


def _hex_rgb(hx: str) -> tuple[float, float, float]:
    return float(int(hx[0:2], 16)), float(int(hx[2:4], 16)), float(int(hx[4:6], 16))


def _grad_stops(grad: etree._Element, paint: "_Paint") -> list[tuple[float, tuple[float, float, float], float]]:
    stops = []
    for gs in grad.iter(_A + "gs"):
        if not len(gs):
            continue
        got = paint.color(gs[0])
        if not got:
            continue
        try:
            pos = int(gs.get("pos") or 0) / 100000.0
        except ValueError:
            pos = 0.0
        stops.append((pos, _hex_rgb(got[0]), float(got[1])))
    stops.sort(key=lambda s: s[0])
    return stops


def _grad_region(grad: etree._Element, paint: "_Paint", box: Bbox, region: Bbox, n: int = 6) -> Optional[tuple[str, float]]:
    """(mean colour, mean opacity) of a gradient-filled box over `region` (sampled on an n×n grid)."""
    stops = _grad_stops(grad, paint)
    if not stops or box.w <= 0 or box.h <= 0:
        return None
    acc = [0.0, 0.0, 0.0]
    tot_a = 0.0
    for i in range(n):
        for j in range(n):
            u = (region.x + (i + 0.5) * region.w / n - box.x) / box.w
            v = (region.y + (j + 0.5) * region.h / n - box.y) / box.h
            rgb, a = gradient_at(stops, gradient_position(grad, u, v, box.w, box.h))
            for k in range(3):
                acc[k] += rgb[k] * a
            tot_a += a
    if tot_a <= 0:
        return None
    return "%02X%02X%02X" % tuple(int(round(min(max(c / tot_a, 0), 255))) for c in acc), tot_a / (n * n)


_REGION_CACHE: dict[tuple, Optional[tuple[str, float]]] = {}


def _picture_region(part, rid: Optional[str], crop: tuple[float, float, float, float], frac: tuple[float, float, float, float]) -> Optional[tuple[str, float]]:
    """(mean colour of the opaque pixels, opaque share) of a picture over `frac` = (x0, y0, x1, y1) of its visible
    part (crop applied); cached per image, crop and region."""
    if not rid or part is None:
        return None
    try:
        img_part = part.related_part(rid)
    except Exception:  # noqa: BLE001
        return None
    key = (_image_key(img_part), crop, tuple(round(v, 3) for v in frac))
    if key in _REGION_CACHE:
        return _REGION_CACHE[key]
    got = None
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(img_part.blob))
        im.draft("RGB", (512, 512))
        im = im.convert("RGBA")
        iw, ih = im.size
        l, t, r, b = crop
        vx0, vy0 = l * iw, t * ih
        vw, vh = max((1 - l - r) * iw, 1), max((1 - t - b) * ih, 1)
        x0, y0, x1, y1 = (max(0.0, min(1.0, v)) for v in frac)
        box = (int(vx0 + x0 * vw), int(vy0 + y0 * vh), max(int(vx0 + x1 * vw), int(vx0 + x0 * vw) + 1), max(int(vy0 + y1 * vh), int(vy0 + y0 * vh) + 1))
        small = im.crop(box).resize((16, 16))
        px = list(small.get_flattened_data() if hasattr(small, "get_flattened_data") else small.getdata())
        wsum = sum(p[3] for p in px)
        if wsum > 0:
            rgb = [sum(p[k] * p[3] for p in px) / wsum for k in range(3)]
            got = ("%02X%02X%02X" % tuple(int(round(c)) for c in rgb), wsum / (255.0 * len(px)))
    except Exception:  # noqa: BLE001 - an unreadable picture (EMF/WMF/SVG): no colour
        got = None
    _REGION_CACHE[key] = got
    return got


def layer_paint_at(layer: Layer, region: Bbox, holder=None) -> Optional[tuple[str, float]]:
    """(colour, opacity) a layer paints over `region`: a picture's opaque pixels there, a gradient's colours there,
    a solid or pattern fill's colour. None when it paints nothing we can read."""
    el = layer.el
    part = getattr(holder, "part", None)
    tag = etree.QName(el).localname
    if tag == "pic" or (tag == "sp" and layer.picture):
        blip = el.find(".//" + _A + "blip")
        rid = blip.get(_R_EMBED) if blip is not None else None
        src = el.find(".//" + _A + "srcRect")
        crop = tuple(int(src.get(k) or 0) / 100000 for k in ("l", "t", "r", "b")) if src is not None else (0.0, 0.0, 0.0, 0.0)
        b = _xfrm_box(el) or layer.box
        if b.w <= 0 or b.h <= 0:
            return None
        frac = ((region.x - b.x) / b.w, (region.y - b.y) / b.h, (region.x2 - b.x) / b.w, (region.y2 - b.y) / b.h)
        return _picture_region(part, rid, crop, frac)
    if tag != "sp":
        return None
    sppr = el.find(_P + "spPr")
    grad = sppr.find(_A + "gradFill") if sppr is not None else None
    if grad is not None:
        return _grad_region(grad, _Paint(holder) if holder is not None else _Paint(None), layer.box, region)
    if not layer.fill_hex:
        return None
    alpha = 1.0
    sf = sppr.find(_A + "solidFill") if sppr is not None else None
    if sf is not None and len(sf):
        a = sf[0].find(_A + "alpha")
        try:
            alpha = int(a.get("val") or 100000) / 100000.0 if a is not None else 1.0
        except ValueError:
            alpha = 1.0
    return layer.fill_hex, alpha


_RECT_PRESETS = ("rect", "roundRect", "snip1Rect", "snip2SameRect", "round1Rect", "round2SameRect", "flowChartProcess")


def _outline_of(layer: Layer) -> Optional[list]:
    """Polygons (slide EMU) a non-rectangular shape really paints (an ellipse, a triangle, a freeform); None for a
    rectangle-like shape (its box is what it paints) or an outline that cannot be read."""
    if layer.kind != "sp":
        return None
    prst = layer.el.find(_P + "spPr/" + _A + "prstGeom")
    if prst is not None and (prst.get("prst") or "rect") in _RECT_PRESETS:
        return None
    try:
        from verstka.audit.ir import shape_outline

        xfrm = layer.el.find(_P + "spPr/" + _A + "xfrm")
        rot = int(xfrm.get("rot") or 0) / 60000.0 if xfrm is not None else 0.0
        box = _xfrm_box(layer.el) or layer.box
        return shape_outline(layer.el, box, rot) or None
    except Exception:  # noqa: BLE001
        return None


def _inside(px: float, py: float, poly) -> bool:
    inside = False
    j = len(poly) - 1
    for i in range(len(poly)):
        xi, yi = poly[i]
        xj, yj = poly[j]
        if (yi > py) != (yj > py) and px < (xj - xi) * (py - yi) / ((yj - yi) or 1e-9) + xi:
            inside = not inside
        j = i
    return inside


def painted_share(layer: Layer, box: Bbox) -> float:
    """Share of `box` a layer paints: its box's overlap for rectangles and pictures, the points of a 12×6 grid inside
    its outline for an ellipse, a triangle or a freeform (a box only partly covered by a triangle's box may still lie
    wholly on the triangle, or wholly off it)."""
    if box.w <= 0 or box.h <= 0:
        return 0.0
    over = layer.box.intersection(box) / float(box.area)
    if over <= 0:
        return 0.0
    polys = _outline_of(layer)
    if polys is None:
        if layer.kind == "sp" and layer.custom_geom:
            return 0.0  # a freeform we cannot read: never a ground
        return over
    hit = 0
    for i in range(12):
        px = box.x + (i + 0.5) * box.w / 12
        for j in range(6):
            py = box.y + (j + 0.5) * box.h / 6
            if sum(1 for poly in polys if len(poly) > 2 and _inside(px, py, poly)) % 2 == 1:
                hit += 1
    return hit / 72.0


def _ground_layers(slide, box: Bbox):
    """(layer, share painted, (hex, opacity)) of every layout/master layer that paints under `box`, topmost first;
    stops at a picture of the slide itself holding the box (it hides the layout) — then the last item is
    (None, share, None)."""
    layers = drawn_layers(slide)
    if not layers or box.w <= 0 or box.h <= 0:
        return
    try:
        layout = slide.slide_layout
        holders = {"layout": layout, "master": layout.slide_master, "slide": slide}
    except Exception:  # noqa: BLE001
        return
    for lay in reversed(layers):
        if lay.placeholder or not lay.paints or lay.kind in ("graphicFrame", "cxnSp", "grpSp"):
            continue
        if lay.source == "slide":
            if lay.kind == "pic" and lay.box.intersection(box) >= 0.8 * box.area:
                yield None, 1.0, None
                return
            continue
        share = painted_share(lay, box)
        if share <= 0.02:
            continue
        got = layer_paint_at(lay, box, holders.get(lay.source))
        if got is None or got[1] < 0.6:
            continue
        yield lay, share, got


def ground_under(slide, box: Bbox, *, min_share: float = 0.8) -> Optional[tuple[str, Layer]]:
    """The colour the slide's layout or master paints under `box` (a band, a panel, a triangle, a gradient picture):
    the topmost non-placeholder layout/master layer that paints at least `min_share` of the box (by its real outline)
    and is ≥ 60 % opaque there — its colour over the box (a gradient or a picture measured where the box is). None
    when a picture of the slide itself covers the box (it hides the layout), or when nothing of the layout or master
    lies under it."""
    for lay, share, got in _ground_layers(slide, box) or ():
        if lay is None:
            return None
        if share >= min_share:
            return got[0], lay
    return None


def ground_is_mixed(slide, box: Bbox) -> bool:
    """True when the box straddles an edge of what the layout or master paints (a heading half on a triangle, half
    off it): its ground has no single colour."""
    for lay, share, got in _ground_layers(slide, box) or ():
        if lay is None:
            return False
        if share >= 0.8:
            return False
        if share >= 0.2:
            return True
    return False


def _row_runs(polys, py: float) -> list[tuple[float, float]]:
    """The x-intervals the polygons paint on the horizontal line y = py (even-odd over all their edges)."""
    xs: list[float] = []
    for poly in polys:
        n = len(poly)
        if n < 3:
            continue
        for i in range(n):
            x1, y1 = poly[i]
            x2, y2 = poly[(i + 1) % n]
            if (y1 > py) != (y2 > py):
                xs.append(x1 + (py - y1) * (x2 - x1) / ((y2 - y1) or 1e-9))
    xs.sort()
    return [(xs[i], xs[i + 1]) for i in range(0, len(xs) - 1, 2)]


def ground_span(slide, box: Bbox, rows: int = 5) -> Optional[tuple[int, int, Layer]]:
    """The run of the ground a line of text stands on, over the line's whole height (gate 2, C2): the topmost layout
    or master layer that paints the box's centre (by its real outline, ≥ 60 % opaque), and the x-interval (EMU) it
    paints on every sampled row of the box around the centre — a triangle's slanted edge moves with the height, so
    the interval is the narrowest of the rows. A rectangle's is its box. None when nothing of the layout paints the
    centre, a picture of the slide itself hides the layout there, or the outline cannot be read."""
    if box.w <= 0 or box.h <= 0:
        return None
    layers = drawn_layers(slide)
    try:
        layout = slide.slide_layout
        holders = {"layout": layout, "master": layout.slide_master, "slide": slide}
    except Exception:  # noqa: BLE001
        return None
    cx, cy = box.x + box.w / 2.0, box.y + box.h / 2.0
    for lay in reversed(layers):
        if lay.placeholder or not lay.paints or lay.kind in ("graphicFrame", "cxnSp", "grpSp"):
            continue
        if lay.source == "slide":
            if lay.kind == "pic" and lay.box.x <= cx <= lay.box.x2 and lay.box.y <= cy <= lay.box.y2:
                return None
            continue
        if not (lay.box.x <= cx <= lay.box.x2 and lay.box.y <= cy <= lay.box.y2):
            continue
        polys = _outline_of(lay)
        if polys is None:
            if lay.kind == "sp" and lay.custom_geom:
                continue  # a freeform we cannot read
            got = layer_paint_at(lay, box, holders.get(lay.source))
            if got is None or got[1] < 0.6:
                continue
            return lay.box.x, lay.box.x2, lay
        if sum(1 for poly in polys if len(poly) > 2 and _inside(cx, cy, poly)) % 2 != 1:
            continue
        got = layer_paint_at(lay, box, holders.get(lay.source))
        if got is None or got[1] < 0.6:
            continue
        lo, hi = float("-inf"), float("inf")
        for j in range(rows):
            py = box.y + box.h * (0.02 + 0.96 * j / max(rows - 1, 1))
            run = next(((a, b) for a, b in _row_runs(polys, py) if a <= cx <= b), None)
            if run is None:
                return int(cx), int(cx), lay  # the centre itself leaves the layer on some row
            lo, hi = max(lo, run[0]), min(hi, run[1])
        return int(lo), int(hi), lay
    return None
