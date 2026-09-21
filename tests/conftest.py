"""Shared fixtures: synthetic decks built with python-pptx so unit tests need no external files."""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

W, H = 12192000, 6858000  # 16:9


def _footer(slide, w: int = W, h: int = H) -> None:
    tb = slide.shapes.add_textbox(Emu(int(w * 0.05)), Emu(int(h * 0.92)), Emu(int(w * 0.2)), Emu(int(h * 0.05)))
    tb.text_frame.text = "ACME"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    tb.name = "Footer ACME"


def build_simple_deck(out: Path, img_dir: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Emu(W)
    prs.slide_height = Emu(H)

    # slide 1: title slide
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Quarterly results"
    s1.placeholders[1].text = "Q3 2026 overview"
    _footer(s1)

    # slide 2: title only + 3 cards
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Three pillars"
    for i in range(3):
        x, y = int(W * (0.06 + i * 0.30)), int(H * 0.35)
        w, h = int(W * 0.26), int(H * 0.40)
        rect = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(0xED, 0xF3, 0xFC)
        rect.line.fill.background()
        rect.name = f"Card {i + 1}"
        t = s2.shapes.add_textbox(Emu(x + int(w * 0.08)), Emu(y + int(h * 0.10)), Emu(int(w * 0.84)), Emu(int(h * 0.20)))
        t.text_frame.text = f"Pillar {i + 1}"
        r = t.text_frame.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.size = Pt(20)
        r.font.color.rgb = RGBColor(0x00, 0x77, 0xFF)
        b = s2.shapes.add_textbox(Emu(x + int(w * 0.08)), Emu(y + int(h * 0.35)), Emu(int(w * 0.84)), Emu(int(h * 0.50)))
        b.text_frame.word_wrap = True
        b.text_frame.text = "Description of the pillar with details"
        b.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    _footer(s2)

    # slide 3: title + picture + caption
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Product photo"
    img = img_dir / "pic.png"
    Image.new("RGB", (400, 200), (30, 120, 255)).save(img)
    s3.shapes.add_picture(str(img), Emu(int(W * 0.06)), Emu(int(H * 0.30)), width=Emu(int(W * 0.50)))
    c = s3.shapes.add_textbox(Emu(int(W * 0.60)), Emu(int(H * 0.35)), Emu(int(W * 0.34)), Emu(int(H * 0.30)))
    c.text_frame.word_wrap = True
    c.text_frame.text = "The device shown in the office environment"
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    _footer(s3)

    prs.save(out)
    return out


@pytest.fixture
def simple_deck(tmp_path: Path) -> Path:
    return build_simple_deck(tmp_path / "simple.pptx", tmp_path)


@pytest.fixture
def fixtures_dir() -> Path | None:
    """Directory with the three VK templates (integration tests only)."""
    env = os.environ.get("VERSTKA_FIXTURES_DIR")
    candidates = [Path(env)] if env else []
    candidates.append(Path(__file__).resolve().parents[2] / "Датасет")
    for c in candidates:
        if c.is_dir() and any(c.glob("*.pptx")):
            return c
    return None
