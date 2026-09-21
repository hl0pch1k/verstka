"""Render the demo outline on every VK template with every strategy (skipped when the dataset is absent)."""

from pathlib import Path

import pytest
from pptx import Presentation

from verstka.analysis.shapes import looks_like_placeholder
from verstka.pipeline.generate import generate_variants
from verstka.schemas.outline import DeckOutline

pytestmark = pytest.mark.integration
FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json"


@pytest.mark.parametrize("needle", ["VK Tech", "WorkSpace", "Education"])
def test_demo_outline_renders_on_vk_templates(fixtures_dir, tmp_path, needle):
    if fixtures_dir is None:
        pytest.skip("VK templates not available")
    pptx = next((p for p in fixtures_dir.glob("*.pptx") if needle.lower() in p.name.lower()), None)
    if pptx is None:
        pytest.skip(needle)
    outline = DeckOutline.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    res = generate_variants(pptx, outline=outline, out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False)
    assert len(res.variants) == 3
    for v in res.variants:
        prs = Presentation(str(v.out_dir / "deck.pptx"))
        assert len(prs.slides) == 12
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
        leftovers = [t for t in texts if t.strip() and looks_like_placeholder(t)]
        assert not leftovers, leftovers[:5]
        assert v.seconds < 60
        failed = [w for w in v.warnings if "failed" in w]
        assert not failed, failed[:5]
