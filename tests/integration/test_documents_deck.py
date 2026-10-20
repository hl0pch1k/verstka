"""DOCX и PDF с заголовками и таблицами → бриф и данные → колода (без модели, шаблон VK Tech; пропускается без датасета).

Текст документа (заголовки, абзацы, списки, таблица слов) — бриф; таблица с цифрами — файл данных второй ступени;
колода — история данных по брифу, как у CSV."""

from pathlib import Path

import pytest
from pptx import Presentation

from verstka.data.story import build_data_story
from verstka.ingest.documents import document_tables, read_document
from verstka.pipeline.generate import generate_variants
from verstka.planning.brief import parse_brief_text

pytestmark = pytest.mark.integration
DOCS = Path(__file__).resolve().parents[1] / "fixtures" / "documents"


@pytest.mark.parametrize("name", ["report.docx", "report.pdf"])
def test_a_document_becomes_a_brief_and_data_and_a_deck(fixtures_dir, tmp_path, name):
    if fixtures_dir is None:
        pytest.skip("VK templates not available")
    pptx = next((p for p in fixtures_dir.glob("*.pptx") if "vk tech" in p.name.lower()), None)
    if pptx is None:
        pytest.skip("VK Tech")
    data = (DOCS / name).read_bytes()
    doc = read_document(name, data)
    brief = parse_brief_text(doc.markdown(data_tables=False))
    tables = document_tables(name, data)
    assert len(tables) == 1 and tables[0].file == name and len(tables[0].rows) == 3
    story = build_data_story(brief, tables, use_llm=False)
    res = generate_variants(pptx, outline=story.outline, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws",
                            use_llm=False, use_vlm=False)
    v = res.variants[0]
    prs = Presentation(str(v.out_dir / "deck.pptx"))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes), "the table of figures is a chart"
    assert "Итоги третьего квартала 2026" in text.replace("\xa0", " ").replace("\x0b", " ")
    assert v.audit is None or v.audit.summary.errors == 0


def test_a_short_request_and_a_knowledge_base_without_a_model_give_a_deck_with_sources(fixtures_dir, tmp_path):
    """Без модели база знаний всё равно даёт колоде текст: предложения самих документов, у слайдов — источники."""
    if fixtures_dir is None:
        pytest.skip("VK templates not available")
    pptx = next((p for p in fixtures_dir.glob("*.pptx") if "vk tech" in p.name.lower()), None)
    if pptx is None:
        pytest.skip("VK Tech")
    from verstka.knowledge.store import KnowledgeBase

    ws = tmp_path / "ws"
    kb = KnowledgeBase.create(ws, "Квартал")
    for f in sorted((DOCS / "kb").iterdir()):
        kb.add_file(f.name, f.read_bytes())
    brief = parse_brief_text("Итоги квартала\n\n1. Выручка по кварталам — покажи на графике\n2. Что думают клиенты\n3. Команда и найм\n4. Риски\n")
    res = generate_variants(pptx, brief=brief, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=ws, use_llm=False, use_vlm=False,
                            knowledge_id=kb.id)
    assert res.writer is not None and res.writer.written and res.writer.basis == "kb"
    prs = Presentation(str(res.variants[0].out_dir / "deck.pptx"))
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("по базе знаний «Квартал»" in n for n in notes)
    assert sum("Источник" in n for n in notes[1:]) >= 3
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame).replace("\xa0", " ")
    assert "46 человек" in text or "46" in text


def test_a_plan_data_files_and_a_knowledge_base_give_one_deck_without_a_model(fixtures_dir, tmp_path):
    """План + файлы с данными + база знаний: графики истории данных и слайды по базе (с источниками) — в одной колоде."""
    if fixtures_dir is None:
        pytest.skip("VK templates not available")
    pptx = next((p for p in fixtures_dir.glob("*.pptx") if "vk tech" in p.name.lower()), None)
    if pptx is None:
        pytest.skip("VK Tech")
    from verstka.data.readers import read_file
    from verstka.knowledge.reference import KbReference
    from verstka.knowledge.store import KnowledgeBase
    from verstka.knowledge.with_data import add_knowledge

    ws = tmp_path / "ws"
    kb = KnowledgeBase.create(ws, "Квартал")
    for n in ("strategiya-2027.pdf", "clients.md", "team.txt"):
        kb.add_file(n, (DOCS / "kb" / n).read_bytes())
    plan = ("Закрытие третьего квартала: отчёт правлению\n\n1. Как росла выручка по месяцам — покажи на графике\n"
            "2. Какие сегменты клиентов растут быстрее\n3. Что думают клиенты\n4. Команда и найм\n5. Что просим у правления\n")
    tables = [t for f in sorted((DOCS / "data").iterdir()) for t in read_file(f)]
    brief = parse_brief_text(plan)
    story = build_data_story(brief, tables, use_llm=False)
    part = add_knowledge(story.outline, brief, KbReference(kb))
    assert part.items and "Что думают клиенты" in part.items and "Команда и найм" in part.items
    res = generate_variants(pptx, outline=part.outline, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=ws, use_llm=False, use_vlm=False)
    prs = Presentation(str(res.variants[0].out_dir / "deck.pptx"))
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart) >= 2, "both data files are charts"
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert sum("Источник" in n for n in notes) >= 2
    assert "по базе знаний «Квартал»" in notes[0]
