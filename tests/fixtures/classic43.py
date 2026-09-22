"""A synthetic CLASSIC template: 4:3, light, everything in placeholders of the standard layouts.

The opposite of the hand-drawn fixture: no free shapes to speak of, the structure lives in the layouts (Title Slide,
Title and Content, Two Content, Comparison, Section Header) and every slide fills their placeholders. Tahoma, a
teal accent, a thin rule under the heading drawn on the master. Used by tests/integration/test_unseen_templates.py.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

W, H = 9144000, 6858000  # 4:3
FONT = "Tahoma"
TEXT = RGBColor(0x22, 0x2B, 0x33)
ACCENT = RGBColor(0x00, 0x8C, 0x8C)
ACCENT_HEX = "008C8C"


def _style(tf, size: int, bold: bool = False, color: RGBColor = TEXT) -> None:
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def _title(slide, text: str, size: int = 30) -> None:
    slide.shapes.title.text = text
    _style(slide.shapes.title.text_frame, size, True, TEXT)


def _body(ph, lines: list[str], size: int = 18) -> None:
    tf = ph.text_frame
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    _style(tf, size)


def build_classic_deck(path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    master = prs.slide_masters[0]
    L = {l.name: l for l in prs.slide_layouts}
    # a thin accent rule under the heading lives on the master (python-pptx draws on slides only: move it over)
    tmp = prs.slides.add_slide(L["Blank"])
    rule = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(457200), Emu(1400000), Emu(8229600), Emu(38100))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()
    master.shapes._spTree.append(rule._element)
    sld_lst = prs.slides._sldIdLst
    rid = sld_lst[-1].rId
    sld_lst.remove(sld_lst[-1])
    prs.part.drop_rel(rid)

    s = prs.slides.add_slide(L["Title Slide"])
    _title(s, "Годовой отчёт отдела аналитики", 40)
    _body(s.placeholders[1], ["Итоги 2026 года и планы"], 22)

    s = prs.slides.add_slide(L["Title and Content"])
    _title(s, "Главное за год")
    _body(s.placeholders[1], ["Запустили единое хранилище данных", "Сократили время отчётов вдвое", "Обучили 120 сотрудников"])

    s = prs.slides.add_slide(L["Two Content"])
    _title(s, "Было и стало")
    _body(s.placeholders[1], ["Отчёты вручную", "Данные в 14 системах", "Ответ за неделю"])
    _body(s.placeholders[2], ["Отчёты по расписанию", "Единое хранилище", "Ответ за день"])

    s = prs.slides.add_slide(L["Comparison"])
    _title(s, "Два пути развития")
    _body(s.placeholders[1], ["Облако"], 20)
    _body(s.placeholders[2], ["Быстрый старт", "Оплата по потреблению"])
    _body(s.placeholders[3], ["Свой контур"], 20)
    _body(s.placeholders[4], ["Полный контроль", "Капитальные затраты"])

    s = prs.slides.add_slide(L["Section Header"])
    _title(s, "Результаты", 36)
    _body(s.placeholders[1], ["Цифры и динамика"], 20)

    s = prs.slides.add_slide(L["Title Only"])
    _title(s, "Число отчётов по кварталам")
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3", "Q4"]
    cd.add_series("Отчёты", (40, 55, 71, 90))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(457200), Emu(1600000), Emu(8229600), Emu(4600000), cd)
    gf.chart.plots[0].series[0].format.fill.solid()
    gf.chart.plots[0].series[0].format.fill.fore_color.rgb = ACCENT

    s = prs.slides.add_slide(L["Title Only"])
    _title(s, "Показатели команды")
    data = [["Показатель", "2025", "2026"], ["Отчётов", "180", "256"], ["Источников", "14", "3"], ["Время ответа, дн", "7", "1"]]
    tbl = s.shapes.add_table(4, 3, Emu(457200), Emu(1700000), Emu(8229600), Emu(2400000)).table
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val
            _style(tbl.cell(r, c).text_frame, 16, r == 0, TEXT)

    s = prs.slides.add_slide(L["Title and Content"])
    _title(s, "Планы на 2027 год")
    _body(s.placeholders[1], ["Самообслуживание для бизнеса", "Прогнозные модели продаж", "Каталог данных"])

    s = prs.slides.add_slide(L["Title Slide"])
    _title(s, "Спасибо!", 40)
    _body(s.placeholders[1], ["analytics@example.com"], 20)

    prs.save(str(path))
    return path


if __name__ == "__main__":  # pragma: no cover
    import sys

    print(build_classic_deck(sys.argv[1] if len(sys.argv) > 1 else "classic43.pptx"))
