"""A synthetic HAND-DRAWN deck: every slide sits on the BLANK layout, there is not a single placeholder.

It imitates what a designer hands over when they "just drew" a presentation: backgrounds are either a full-bleed
rectangle or a solid slide fill, all text lives in free text boxes, cards are rectangles with text boxes on top,
the timeline is made of connectors and ovals.  Fonts and colours are deliberately unlike the VK templates:
Georgia headings, Verdana body, dark green ground, orange accent.

Used by tests/integration/test_unseen_templates.py to check that the pipeline generalises to unseen templates.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

W, H = 12192000, 6858000  # 16:9

HEADING_FONT = "Georgia"
BODY_FONT = "Verdana"

BG = RGBColor(0x0B, 0x3D, 0x2E)  # dark green ground
SURFACE = RGBColor(0x14, 0x52, 0x3F)  # card surface
ACCENT = RGBColor(0xFF, 0x7A, 0x00)  # orange accent
TEXT = RGBColor(0xF4, 0xF1, 0xE8)  # cream text
MUTED = RGBColor(0xA9, 0xC4, 0xB8)  # secondary text

ACCENT_HEX = "FF7A00"
BG_HEX = "0B3D2E"


def _x(f: float) -> Emu:
    return Emu(int(W * f))


def _y(f: float) -> Emu:
    return Emu(int(H * f))


def _text(slide, x, y, w, h, text, *, font=BODY_FONT, size=16, bold=False, color=TEXT, align=None, anchor=None, name=None):
    tb = slide.shapes.add_textbox(_x(x), _y(y), _x(w), _y(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    if name:
        tb.name = name
    return tb


def _rect(slide, x, y, w, h, fill, *, shape=MSO_SHAPE.RECTANGLE, name=None):
    sh = slide.shapes.add_shape(shape, _x(x), _y(y), _x(w), _y(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    return sh


def _solid_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _rect_bg(slide) -> None:
    _rect(slide, 0, 0, 1, 1, BG, name="Background")


def _chrome(slide, n: int) -> None:
    """Things that repeat on every content slide: an orange brand tick, a footer and a page number."""
    _rect(slide, 0.05, 0.06, 0.035, 0.012, ACCENT, name="Brand tick")
    _text(slide, 0.05, 0.925, 0.30, 0.045, "Северный лес · 2026", size=10, color=MUTED, name="Footer")
    _text(slide, 0.90, 0.925, 0.05, 0.045, f"{n:02d}", size=10, color=MUTED, align=PP_ALIGN.RIGHT, name="Page number")


def _heading(slide, text: str) -> None:
    _text(slide, 0.05, 0.09, 0.90, 0.13, text, font=HEADING_FONT, size=32, bold=True, name="Heading")


def _picture(path: Path) -> Path:
    img = Image.new("RGB", (1200, 900), (0x1E, 0x6B, 0x52))
    d = ImageDraw.Draw(img)
    for i in range(0, 1200, 120):
        d.polygon([(i, 900), (i + 60, 300 + (i % 360)), (i + 120, 900)], fill=(0x0B, 0x3D, 0x2E))
    d.ellipse((880, 90, 1060, 270), fill=(0xFF, 0x7A, 0x00))
    img.save(path)
    return path


def build_handdrawn_deck(path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Emu(W)
    prs.slide_height = Emu(H)
    blank = prs.slide_layouts[6]
    assert blank.name == "Blank"

    # 1. title
    s = prs.slides.add_slide(blank)
    _rect_bg(s)
    _rect(s, 0.05, 0.30, 0.012, 0.30, ACCENT, name="Title bar")
    _text(s, 0.08, 0.28, 0.80, 0.24, "Стратегия развития питомника на 2027 год", font=HEADING_FONT, size=48, bold=True, name="Title")
    _text(s, 0.08, 0.54, 0.70, 0.08, "Отчёт для совета директоров", size=20, color=MUTED, name="Subtitle")
    _text(s, 0.08, 0.86, 0.40, 0.05, "Северный лес · октябрь 2026", size=12, color=MUTED, name="Date")

    # 2. agenda list
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 2)
    _heading(s, "Содержание")
    items = ["Итоги сезона", "Рынок и спрос", "Производственный план", "Инвестиции", "Риски и решения"]
    for i, it in enumerate(items):
        y = 0.28 + i * 0.12
        _text(s, 0.05, y, 0.07, 0.09, f"{i + 1:02d}", font=HEADING_FONT, size=28, bold=True, color=ACCENT, name=f"Agenda num {i + 1}")
        _text(s, 0.13, y + 0.012, 0.70, 0.07, it, size=20, name=f"Agenda item {i + 1}")

    # 3. three cards
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 3)
    _heading(s, "Три направления роста")
    cards = [
        ("Хвойные", "Расширяем посадки сосны и ели на 40 гектаров"),
        ("Лиственные", "Запускаем линию саженцев дуба и липы для городов"),
        ("Сервис", "Доставка и посадка под ключ для частных клиентов"),
    ]
    for i, (t, b) in enumerate(cards):
        x = 0.05 + i * 0.31
        _rect(s, x, 0.28, 0.28, 0.50, SURFACE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Card {i + 1}")
        _rect(s, x + 0.02, 0.32, 0.04, 0.008, ACCENT, name=f"Card tick {i + 1}")
        _text(s, x + 0.02, 0.35, 0.24, 0.09, t, font=HEADING_FONT, size=22, bold=True, name=f"Card title {i + 1}")
        _text(s, x + 0.02, 0.46, 0.24, 0.28, b, size=14, name=f"Card body {i + 1}")

    # 4. four KPI numbers
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 4)
    _heading(s, "Сезон в цифрах")
    kpis = [("1,2 млн", "саженцев продано"), ("+27%", "выручка к прошлому году"), ("94%", "приживаемость"), ("38", "регионов доставки")]
    for i, (v, lab) in enumerate(kpis):
        x = 0.05 + i * 0.23
        _text(s, x, 0.34, 0.21, 0.16, v, font=HEADING_FONT, size=44, bold=True, color=ACCENT, name=f"KPI value {i + 1}")
        _text(s, x, 0.52, 0.21, 0.12, lab, size=14, color=MUTED, name=f"KPI label {i + 1}")

    # 5. two-column text
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 5)
    _heading(s, "Рынок: что изменилось")
    cols = [
        ("Спрос", ["Города увеличили бюджеты на озеленение", "Частные клиенты покупают крупномеры", "Растёт интерес к местным породам"]),
        ("Предложение", ["Импорт саженцев сократился вдвое", "Мелкие питомники уходят с рынка", "Логистика подорожала на 15%"]),
    ]
    for i, (t, bullets) in enumerate(cols):
        x = 0.05 + i * 0.46
        _text(s, x, 0.27, 0.42, 0.08, t, font=HEADING_FONT, size=24, bold=True, color=ACCENT, name=f"Column title {i + 1}")
        _text(s, x, 0.37, 0.42, 0.45, [f"• {b}" for b in bullets], size=16, name=f"Column body {i + 1}")

    # 6. picture + text
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 6)
    _heading(s, "Новый участок под Вологдой")
    pic = _picture(path.parent / f"{path.stem}_photo.png")
    s.shapes.add_picture(str(pic), _x(0.05), _y(0.27), width=_x(0.44), height=_y(0.58))
    _text(s, 0.53, 0.27, 0.42, 0.58, [
        "Сто двадцать гектаров бывших сельхозземель переведены под питомник.",
        "Первая посадка — весной 2027 года, выход на полную мощность — через три сезона.",
    ], size=16, name="Picture text")

    # 7. native table
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 7)
    _heading(s, "Производственный план")
    data = [["Порода", "2026", "2027", "2028"], ["Сосна", "420", "510", "600"], ["Ель", "380", "400", "450"], ["Дуб", "90", "160", "240"], ["Липа", "60", "110", "180"]]
    gf = s.shapes.add_table(len(data), 4, _x(0.05), _y(0.28), _x(0.90), _y(0.50))
    gf.name = "Plan table"
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = gf.table.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if r == 0 else SURFACE
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.name = BODY_FONT
            run.font.size = Pt(14)
            run.font.bold = r == 0
            run.font.color.rgb = BG if r == 0 else TEXT

    # 8. native chart
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 8)
    _heading(s, "Выручка растёт четвёртый год подряд")
    cd = CategoryChartData()
    cd.categories = ["2023", "2024", "2025", "2026"]
    cd.add_series("Выручка, млн ₽", (310, 365, 440, 560))
    gfc = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, _x(0.05), _y(0.27), _x(0.58), _y(0.58), cd)
    gfc.name = "Revenue chart"
    ch = gfc.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.name = BODY_FONT
    ch.legend.font.size = Pt(11)
    ch.legend.font.color.rgb = TEXT
    ch.plots[0].series[0].format.fill.solid()
    ch.plots[0].series[0].format.fill.fore_color.rgb = ACCENT
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.name = BODY_FONT
        ax.tick_labels.font.size = Pt(11)
        ax.tick_labels.font.color.rgb = TEXT
    _text(s, 0.67, 0.30, 0.28, 0.50, "Рост обеспечили крупномеры и городские контракты: их доля в выручке достигла 46%.", size=16, name="Chart note")

    # 9. timeline drawn with lines and circles
    s = prs.slides.add_slide(blank)
    _solid_bg(s)
    _chrome(s, 9)
    _heading(s, "Дорожная карта")
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _x(0.08), _y(0.45), _x(0.92), _y(0.45))
    line.line.color.rgb = ACCENT
    line.line.width = Pt(2.5)
    line.name = "Timeline axis"
    steps = [("Q1", "Подготовка участка"), ("Q2", "Первая посадка"), ("Q3", "Запуск доставки"), ("Q4", "Городские тендеры")]
    for i, (q, t) in enumerate(steps):
        cx = 0.12 + i * 0.25
        d = 0.022
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, _x(cx - d / 2), Emu(int(H * 0.45 - W * d / 2)), _x(d), _x(d))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()
        dot.name = f"Timeline dot {i + 1}"
        _text(s, cx - 0.09, 0.30, 0.18, 0.09, q, font=HEADING_FONT, size=24, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, name=f"Timeline label {i + 1}")
        _text(s, cx - 0.10, 0.51, 0.20, 0.16, t, size=14, align=PP_ALIGN.CENTER, name=f"Timeline text {i + 1}")

    # 10. closing
    s = prs.slides.add_slide(blank)
    _rect_bg(s)
    _rect(s, 0.44, 0.30, 0.12, 0.012, ACCENT, name="Closing tick")
    _text(s, 0.10, 0.36, 0.80, 0.18, "Спасибо за внимание", font=HEADING_FONT, size=48, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Closing title")
    _text(s, 0.10, 0.56, 0.80, 0.08, "sever-les.example · +7 000 000-00-00", size=18, color=MUTED, align=PP_ALIGN.CENTER, name="Contacts")

    prs.save(str(path))
    return path


if __name__ == "__main__":  # pragma: no cover
    import sys

    print(build_handdrawn_deck(sys.argv[1] if len(sys.argv) > 1 else "handdrawn.pptx"))
