"""Render-time guards (verstka/rendering/fallbacks.py): whatever the plan says, no slide reaches the deck as a heading
over an empty frame, and no template stub reaches the page. The cases are the ones a live model plan produced
(workspace/runs/20260925-155033-1006b5, brief «47 → 29 минут, NPS 64, 14,5 млн ₽»): charts whose `series_ids` name
facts, a comparison of bare titles, a «freeform» slide, a closing slide with «[email@company.com]»."""

from __future__ import annotations

import pytest
from pptx import Presentation

from verstka.analysis.manifest import analyze_template
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering import charts as charts_mod
from verstka.rendering import compose as compose_mod
from verstka.rendering.charts import chart_data_ok, resolve_series, unit_caption
from verstka.rendering.fallbacks import fact_numbers, prepare_slide, scrub_stubs
from verstka.rendering.renderer import render_deck
from verstka.schemas.common import PatternKind as K
from verstka.schemas.layout import LayoutPlan, LayoutSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, Fact, NumberCallout, OutlineSlide, Series, SlideContent, SlideItem

FACTS = [
    Fact(id="f1", value="47", unit="минут", label="время на чтение чатов", source_span="Проблема: сотрудники тратят 47 минут в день на чтение чатов."),
    Fact(id="f2", value="29", unit="минут", label="время на чтение чатов", source_span="Результаты: время сократилось до 29 минут, NPS 64."),
    Fact(id="f3", value="64", unit="%", label="NPS после пилота", source_span="Результаты: время сократилось до 29 минут, NPS 64."),
    Fact(id="f4", value="14,5", unit="млн ₽", label="бюджет на масштабирование", source_span="Просим: бюджет 14,5 млн ₽ на масштабирование."),
]


def _deck(*slides: OutlineSlide, series: list[Series] | None = None, planned_by: str = "model") -> DeckOutline:
    return DeckOutline(title="Итоги пилота «Умные сводки»", strategy="structured", slides=list(slides), facts=list(FACTS), series=list(series or []), planned_by=planned_by)


def _chart_slide(ids: list[str], *, sid: str = "sl5", unit: str = "минут", **content) -> OutlineSlide:
    return OutlineSlide(id=sid, kind=K.chart, headline="Динамика времени на чтение чатов", content=SlideContent(chart=ChartSpec(type="bar", series_ids=ids, unit=unit), **content), fact_refs=list(ids))


def _plan(slide: OutlineSlide, comp: str = "chart_text") -> LayoutSlide:
    return LayoutSlide(outline_id=slide.id, mode="synth", composition=comp, reasons=["композиция из дизайн-системы"])


# ---------------------------------------------------------------------------------------------- template stubs


@pytest.mark.parametrize(
    "text, bookend, expected",
    [
        ("Вопросы? Контакты: [email@company.com]", True, "Вопросы?"),  # the live closing slide (structured / visual)
        ("Контакты: [email] | [телефон]", True, ""),  # the live closing slide (compact)
        ("Email: [email] | Тел.: [телефон]", True, ""),
        ("Вопросы — пишите на ivan@company.com", True, "Вопросы"),
        ("Сайт: www.example.com", True, ""),
        ("Тел. +7 (XXX) XXX-XX-XX", True, ""),
        ("Для связи: <email>, {телефон}", True, ""),
        ("Команда продукта [Имя Фамилия]", True, "Команда продукта"),
        ("Вопросы?\nКонтакты: [email]", True, "Вопросы?"),
        ("Пишите: team@vk.com", True, "Пишите: team@vk.com"),  # a real address stays
        ("Бюджет 14,5 млн ₽ [уточнить]", False, "Бюджет 14,5 млн ₽"),
        ("Источник [1]: опрос 2024", False, "Источник [1]: опрос 2024"),  # a reference is not a stub
        ("Итоги пилота [за квартал]", False, "Итоги пилота [за квартал]"),
        ("[...]", False, ""),
        # comparison signs are no brackets: the claim stays whole, on a cover as anywhere
        ("Время ответа <5 минут, у конкурентов >10 минут", True, "Время ответа <5 минут, у конкурентов >10 минут"),
        ("Ответ <5 минут против >10 минут у конкурентов", True, "Ответ <5 минут против >10 минут у конкурентов"),
        ("Время ответа <5 минут, у конкурентов >10 минут", False, "Время ответа <5 минут, у конкурентов >10 минут"),
        ("Команда продукта <Имя Фамилия>", True, "Команда продукта"),
    ],
)
def test_template_stubs_are_scrubbed_with_the_label_they_leave(text, bookend, expected):
    assert scrub_stubs(text, all_brackets=bookend) == expected


def test_closing_slide_subtitle_loses_its_stubs():
    thanks = OutlineSlide(id="sl12", kind=K.thanks, headline="Спасибо за внимание", subtitle="Вопросы? Контакты: [email@company.com]")
    s, ps, notes = prepare_slide(thanks, LayoutSlide(outline_id="sl12", mode="clone", pattern_id="p4"), _deck(thanks))
    assert s.subtitle == "Вопросы?" and s.headline == "Спасибо за внимание"
    assert ps.mode == "clone" and ps.pattern_id == "p4"  # only the text changed: the sample stays
    assert any("stubs" in n for n in notes)
    empty = OutlineSlide(id="sl12", kind=K.thanks, headline="Благодарим за внимание", subtitle="Контакты: [email] | [телефон]")
    s2, _, _ = prepare_slide(empty, LayoutSlide(outline_id="sl12", mode="clone", pattern_id="p4"), _deck(empty))
    assert s2.subtitle is None


def test_an_untouched_slide_comes_back_as_the_same_objects():
    s = OutlineSlide(id="b", kind=K.bullets, headline="Причины запроса бюджета", content=SlideContent(bullets=["Распространение на все отделы", "Интеграция с системами"]))
    ps = _plan(s, "bullets")
    s2, ps2, notes = prepare_slide(s, ps, _deck(s))
    assert s2 is s and ps2 is ps and notes == []


# ---------------------------------------------------------------------------------------------- charts without data


def test_fact_ids_never_borrow_another_series():
    other = Series(id="s1", name="Активные пользователи", categories=["Май", "Июнь"], values=[1200, 3400])
    deck = _deck(series=[other])
    assert resolve_series(ChartSpec(series_ids=["f1", "f2"]), deck) == []  # facts, not series: no one else's data
    assert resolve_series(ChartSpec(series_ids=["typo"]), deck) == [other]  # an unknown id keeps the lenient fallback
    assert not chart_data_ok(ChartSpec(series_ids=["f1", "f2"], unit="минут"), deck)
    assert unit_caption(ChartSpec(series_ids=["f1", "f2"], unit="минут"), deck) is None  # no chart, no caption
    assert chart_data_ok(ChartSpec(series_ids=["s1"]), deck)


def test_fact_numbers_join_a_change_and_keep_units():
    s = _chart_slide(["f1", "f2", "f3"])
    nums = fact_numbers(s, _deck(s))
    assert [(n.value, n.label) for n in nums] == [("47 → 29 минут", "время на чтение чатов"), ("64", "NPS после пилота")]  # NPS is no share
    budget = OutlineSlide(id="x", kind=K.bullets, headline="h", fact_refs=["f4"])
    assert [n.value for n in fact_numbers(budget, _deck(budget))] == ["14,5 млн ₽"]


def test_chart_of_two_facts_becomes_one_big_number():
    s = _chart_slide(["f1", "f2"])
    out, ps, notes = prepare_slide(s, _plan(s), _deck(s))
    assert out.kind == K.big_number and out.content.chart is None
    assert [n.value for n in out.content.numbers] == ["47 → 29 минут"]
    assert ps.mode == "synth" and ps.composition == "big_number" and any("при вёрстке" in r for r in ps.reasons)
    assert any("no series data" in n and "big number" in n for n in notes)


def test_chart_of_three_facts_becomes_kpi_tiles():
    s = _chart_slide(["f1", "f2", "f3"], unit="минуты / %")
    out, ps, _ = prepare_slide(s, LayoutSlide(outline_id=s.id, mode="clone", pattern_id="p44"), _deck(s))
    assert out.kind == K.stat_row and len(out.content.numbers) == 2
    assert ps.mode == "synth" and ps.composition == "stat_row" and ps.pattern_id is None


def test_chart_without_data_keeps_its_text():
    s = OutlineSlide(id="c", kind=K.chart, headline="Динамика", content=SlideContent(chart=ChartSpec(series_ids=["nope"]), bullets=["Время сократилось", "NPS вырос"]))
    out, ps, _ = prepare_slide(s, _plan(s), _deck(s))
    assert out.kind == K.bullets and out.content.bullets == ["Время сократилось", "NPS вырос"] and ps.composition == "bullets"


def test_chart_with_nothing_at_all_is_a_statement_of_its_headline():
    s = OutlineSlide(id="c", kind=K.chart, headline="Ожидаемые показатели после масштабирования", content=SlideContent(chart=ChartSpec(series_ids=["nope"], unit="минут")))
    out, ps, notes = prepare_slide(s, _plan(s), _deck(s))
    assert out.kind == K.section and out.headline == s.headline and out.content.is_empty
    assert ps.composition == "section" and any("statement" in n for n in notes)


def test_chart_with_real_series_is_untouched():
    series = Series(id="s1", name="Активные", categories=["Май", "Июнь", "Июль"], values=[1200, 3400, 6100])
    s = OutlineSlide(id="c", kind=K.chart, headline="Рост", content=SlideContent(chart=ChartSpec(series_ids=["s1"])))
    ps = _plan(s)
    out, ps2, notes = prepare_slide(s, ps, _deck(s, series=[series]))
    assert out is s and ps2 is ps and notes == []


def test_a_broken_chart_on_a_text_slide_is_dropped_the_text_stays():
    s = OutlineSlide(id="b", kind=K.bullets, headline="Итоги", content=SlideContent(bullets=["Время сократилось"], chart=ChartSpec(series_ids=["f1"])))
    out, ps, notes = prepare_slide(s, _plan(s, "bullets"), _deck(s))
    assert out.kind == K.bullets and out.content.chart is None and out.content.bullets == ["Время сократилось"]
    assert any("chart dropped" in n for n in notes)


# ---------------------------------------------------------------------------------------------- empty columns and cards


def test_comparison_of_bare_titles_shows_the_facts_it_refers_to():
    s = OutlineSlide(
        id="sl6", kind=K.comparison, headline="Сравнение до и после пилота", subtitle="Основные показатели до и после внедрения",
        content=SlideContent(items=[SlideItem(title="Время на чтение чатов"), SlideItem(title="NPS")]), fact_refs=["f1", "f2", "f3"],
    )
    out, ps, _ = prepare_slide(s, _plan(s, "comparison"), _deck(s))
    assert out.kind == K.stat_row and [n.value for n in out.content.numbers] == ["47 → 29 минут", "64"]
    assert out.subtitle == s.subtitle and ps.composition == "stat_row"


def test_comparison_drops_a_bare_column_when_two_full_ones_remain():
    cols = [SlideItem(title="До", bullets=["47 минут в день"]), SlideItem(title="После", text="29 минут в день"), SlideItem(title="NPS")]
    s = OutlineSlide(id="c", kind=K.two_column, headline="До и после", content=SlideContent(columns=cols))
    ps = _plan(s, "two_column")
    out, ps2, notes = prepare_slide(s, ps, _deck(s))
    assert out.kind == K.two_column and [c.title for c in out.content.columns] == ["До", "После"] and ps2 is ps
    assert any("empty column" in n for n in notes)


@pytest.mark.parametrize("planned_by", ["model", "shared:visual"])
def test_model_cards_of_bare_titles_without_facts_become_a_list(planned_by):
    # a model is asked for a text under every card title: titles alone are cards that lost their text
    s = OutlineSlide(id="c", kind=K.cards, headline="Преимущества", content=SlideContent(items=[SlideItem(title="Экономия времени"), SlideItem(title="Простота"), SlideItem(title="Скорость")]))
    out, ps, _ = prepare_slide(s, _plan(s, "cards"), _deck(s, planned_by=planned_by))
    assert out.kind == K.bullets and out.content.bullets == ["Экономия времени", "Простота", "Скорость"] and ps.composition == "bullets"


@pytest.mark.parametrize("planned_by", ["rules", "skeleton"])
def test_rules_cards_of_bare_titles_stay_cards(planned_by):
    # the rules planner sets a list of short lines as title-only cards on purpose (edu_program «Годовая программа из
    # 4 модулей», verstka_pitch «Verstka читает любой PPTX как набор правил»)
    modules = ["Основы SQL и Python", "Потоковая обработка", "Хранилища и витрины", "Инженерия ML-пайплайнов"]
    s = OutlineSlide(id="sl4", kind=K.cards, headline="Годовая программа из 4 модулей", content=SlideContent(items=[SlideItem(title=t) for t in modules]))
    ps = _plan(s, "cards")
    out, ps2, notes = prepare_slide(s, ps, _deck(s, planned_by=planned_by))
    assert out is s and ps2 is ps and notes == []
    hinted = OutlineSlide(id="sl10", kind=K.cards, headline="Qwen3.8-27B под лицензией Apache 2.0", content=SlideContent(items=[SlideItem(title="Смена на инференс VK", icon_hint="Смена на"), SlideItem(title="Без модели сервис работает полностью", icon_hint="Без модели")]))
    out2, _, notes2 = prepare_slide(hinted, _plan(hinted, "cards"), _deck(hinted, planned_by=planned_by))
    assert out2 is hinted and notes2 == []


def test_rules_comparison_of_bare_titles_still_shows_its_facts():
    # an empty column is an empty frame whoever planned it
    s = OutlineSlide(id="cmp", kind=K.comparison, headline="До и после", content=SlideContent(items=[SlideItem(title="Время на чтение чатов"), SlideItem(title="NPS")]), fact_refs=["f1", "f2", "f3"])
    out, _, _ = prepare_slide(s, _plan(s, "comparison"), _deck(s, planned_by="rules"))
    assert out.kind == K.stat_row and [n.value for n in out.content.numbers] == ["47 → 29 минут", "64"]


def test_a_cover_keeps_a_comparison_in_its_text():
    title = OutlineSlide(id="sl1", kind=K.title, headline="Время ответа <5 минут, у конкурентов >10 минут", subtitle="Итоги пилота")
    ps = LayoutSlide(outline_id="sl1", mode="clone", pattern_id="p1")
    out, ps2, notes = prepare_slide(title, ps, _deck(title))
    assert out is title and ps2 is ps and notes == []


def test_cards_with_text_are_untouched():
    s = OutlineSlide(id="c", kind=K.cards, headline="План", content=SlideContent(items=[SlideItem(title="Фаза 1", text="Внедрение"), SlideItem(title="Фаза 2", text="Интеграция")]))
    out, _, notes = prepare_slide(s, _plan(s, "cards"), _deck(s))
    assert out is s and notes == []


def test_figure_slide_without_figures_takes_its_facts():
    s = OutlineSlide(id="n", kind=K.big_number, headline="Бюджет на масштабирование", fact_refs=["f4"])
    out, _, _ = prepare_slide(s, _plan(s, "big_number"), _deck(s))
    assert out.kind == K.big_number and [n.value for n in out.content.numbers] == ["14,5 млн ₽"]


def test_unknown_kind_is_set_as_bullets_with_its_figures():
    s = OutlineSlide(
        id="sl7", kind=K.freeform, headline="Масштабирование — следующий шаг",
        content=SlideContent(paragraphs=["Пилот показал стабильный положительный эффект."], numbers=[NumberCallout(value="14,5", label="бюджет на масштабирование", fact_id="f4")]),
    )
    out, ps, notes = prepare_slide(s, LayoutSlide(outline_id="sl7", mode="clone", pattern_id="p9"), _deck(s))
    assert out.kind == K.bullets and out.content.bullets == ["Пилот показал стабильный положительный эффект.", "14,5 млн ₽ — бюджет на масштабирование"]
    assert not out.content.numbers and ps.mode == "synth" and ps.composition == "bullets"
    assert any("freeform" in n for n in notes)


# ---------------------------------------------------------------------------------------------- rendered decks


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return manifest, ws


def _texts(slide) -> list[str]:
    return [sh.text_frame.text.replace(" ", " ").replace("⁠", "") for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def test_render_deck_never_leaves_a_chart_slide_empty(env, tmp_path):
    manifest, ws = env
    chart = _chart_slide(["f1", "f2"])
    lone = OutlineSlide(id="lone", kind=K.chart, headline="Ожидаемые показатели", content=SlideContent(chart=ChartSpec(series_ids=["nope"], unit="минут")))
    comparison = OutlineSlide(id="cmp", kind=K.comparison, headline="Сравнение до и после пилота", content=SlideContent(items=[SlideItem(title="Время на чтение чатов"), SlideItem(title="NPS")]), fact_refs=["f1", "f2", "f3"])
    free = OutlineSlide(id="free", kind=K.freeform, headline="Следующий шаг", content=SlideContent(paragraphs=["Нужен бюджет."], numbers=[NumberCallout(value="14,5", label="бюджет на масштабирование", fact_id="f4")]))
    thanks = OutlineSlide(id="end", kind=K.thanks, headline="Спасибо за внимание", subtitle="Вопросы? Контакты: [email@company.com]")
    outline = _deck(chart, lone, comparison, free, thanks)
    plan = LayoutPlan(strategy="structured", template_id=manifest.template_id, slides=[_plan(chart), _plan(lone), _plan(comparison, "comparison"), LayoutSlide(outline_id="free", mode="synth", composition="bullets"), LayoutSlide(outline_id="end", mode="synth", composition="thanks")])
    res = render_deck(outline, plan, manifest, ws, tmp_path / "deck.pptx")
    prs = Presentation(str(res.pptx_path))
    assert len(prs.slides) == 5
    s_chart, s_lone, s_cmp, s_free, s_end = (_texts(s) for s in prs.slides)
    # the chart slide shows its figures, never the bare unit caption
    assert any("47" in t and "29" in t for t in s_chart) and "минут" not in [t.strip() for t in s_chart]
    assert not any(sh.name.startswith("Unit") for sh in prs.slides[0].shapes)
    # nothing to show: the headline alone, as a statement — no stray caption
    assert any("Ожидаемые показатели" in t for t in s_lone) and not any(t.strip() == "минут" for t in s_lone)
    assert any("64" in t for t in s_cmp) and any("29" in t for t in s_cmp)
    assert any("14,5 млн ₽ — бюджет на масштабирование" in t for t in s_free)
    assert not any("[" in t or "company.com" in t for t in s_end) and any("Вопросы?" in t for t in s_end)
    # the outline and the plan say what the deck shows (audit, outline.json, layout_plan.json)
    kinds = {s.id: s.kind for s in outline.slides}
    assert kinds == {"sl5": K.big_number, "lone": K.section, "cmp": K.stat_row, "free": K.bullets, "end": K.thanks}
    assert plan.for_outline("sl5").composition == "big_number" and plan.for_outline("lone").composition == "section"
    assert [r.composition for r in res.slides][:3] == ["big_number", "section", "stat_row"]
    warnings = " | ".join(res.warnings)
    assert "chart failed" not in warnings and "no series data" in warnings and "freeform" in warnings


def test_a_chart_that_fails_while_drawing_gives_its_area_to_the_figures(env, tmp_path, monkeypatch):
    manifest, ws = env
    series = Series(id="s1", name="Время", categories=["До", "После"], values=[47, 29], unit="минут")
    s = OutlineSlide(id="c", kind=K.chart, headline="Время на чтение чатов", content=SlideContent(chart=ChartSpec(series_ids=["s1"], unit="минут"), numbers=[NumberCallout(value="47 → 29 минут", label="время на чтение чатов")]))

    def boom(*a, **k):
        raise ValueError("python-pptx refused the chart")

    monkeypatch.setattr(compose_mod, "add_chart", boom)
    outline = _deck(s, series=[series])
    plan = LayoutPlan(strategy="structured", template_id=manifest.template_id, slides=[_plan(s)])
    res = render_deck(outline, plan, manifest, ws, tmp_path / "deck.pptx")
    slide = Presentation(str(res.pptx_path)).slides[0]
    assert any("47" in t and "29" in t for t in _texts(slide))
    assert not any(sh.name.startswith("Unit") for sh in slide.shapes)  # no caption over a chart that is not there
    assert any("chart failed" in w and "big number" in w for w in res.slides[0].warnings)


def test_a_half_built_chart_is_removed_before_the_fallback(env, tmp_path, monkeypatch):
    manifest, ws = env
    series = Series(id="s1", name="Время", categories=["До", "После"], values=[47, 29], unit="минут")
    s = OutlineSlide(id="c", kind=K.chart, headline="Время на чтение чатов", content=SlideContent(chart=ChartSpec(series_ids=["s1"], unit="минут"), bullets=["Время сократилось на 18 минут"]))
    real = compose_mod.add_chart

    def half(*a, **k):
        real(*a, **k)
        raise RuntimeError("styling failed after the frame was added")

    monkeypatch.setattr(compose_mod, "add_chart", half)
    outline = _deck(s, series=[series])
    plan = LayoutPlan(strategy="structured", template_id=manifest.template_id, slides=[_plan(s)])
    res = render_deck(outline, plan, manifest, ws, tmp_path / "deck.pptx")
    slide = Presentation(str(res.pptx_path)).slides[0]
    assert not any(sh.has_chart for sh in slide.shapes)
    assert any("Время сократилось на 18 минут" in t for t in _texts(slide))


def test_a_cloned_chart_that_fails_is_rolled_back_to_a_composed_slide(env, tmp_path, monkeypatch):
    manifest, ws = env
    pattern = next(p for p in manifest.patterns if p.kind not in (K.title, K.section, K.thanks))
    series = Series(id="s1", name="Время", categories=["До", "После"], values=[47, 29], unit="минут")
    s = OutlineSlide(id="c", kind=K.chart, headline="Время на чтение чатов", content=SlideContent(chart=ChartSpec(series_ids=["s1"], unit="минут"), bullets=["Время сократилось на 18 минут"]))

    def boom(*a, **k):
        raise ValueError("python-pptx refused the chart")

    from verstka.rendering import clone as clone_mod

    monkeypatch.setattr(clone_mod, "add_chart", boom)
    monkeypatch.setattr(compose_mod, "add_chart", boom)
    outline = _deck(s, series=[series])
    plan = LayoutPlan(strategy="structured", template_id=manifest.template_id, slides=[LayoutSlide(outline_id="c", mode="clone", pattern_id=pattern.id)])
    res = render_deck(outline, plan, manifest, ws, tmp_path / "deck.pptx")
    prs = Presentation(str(res.pptx_path))
    assert len(prs.slides) == 1 and res.slides[0].mode == "synth" and res.slides[0].composition == "chart_text"
    assert any("native chart failed" in w for w in res.slides[0].warnings)
    assert any("Время сократилось на 18 минут" in t for t in _texts(prs.slides[0]))


def test_add_chart_still_refuses_a_chart_without_data():
    # the renderer guards against it; the chart module itself stays strict
    s = _chart_slide(["f1"])
    with pytest.raises(ValueError):
        charts_mod.add_chart(None, None, s.content.chart, _deck(s), None, None)
