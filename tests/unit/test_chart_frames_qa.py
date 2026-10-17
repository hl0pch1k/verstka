"""Диаграммы по живой проверке второй ступени (H1–H5): подписи времени тонятся, а не поворачиваются; график не сжимается в
полоску в низкой рамке; строки полос не налезают друг на друга; ось значений всегда с единицей; ограниченная мера
(балл из 100, проценты) кончается на 100; цвета рядов читаются на фоне (3:1); название группы — перед легендой."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from verstka.rendering.charts import (
    MIN_PLOT,
    SERIES_CONTRAST,
    _nice_scale,
    _scatter_range,
    add_chart,
    bounded_cap,
    chart_palette,
    fit_categories,
    is_time_axis,
    legible_color,
    min_frame_height,
    period_levels,
    short_years,
    thin_picks,
)
from verstka.rendering.fonts import text_width_pt
from verstka.schemas.common import Bbox, contrast_ratio
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, XYSeries
from verstka.schemas.template import ChartStyleSpec, TypeStep, Typography

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
W, H = 12192000, 6858000
STYLE = ChartStyleSpec(series_colors=["0077FF", "00D3E6", "005EFF", "FF6C6C", "0563C1"], font_family="Play", font_size_pt=7.0)
TYPO = Typography(scale=[TypeStep(role="display", size_pt=47), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=12), TypeStep(role="body", size_pt=9), TypeStep(role="small", size_pt=7)])
OUTLINE = DeckOutline(title="t")
MONTHS = [f"{m} {y}" for y in (2025, 2026) for m in ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")][:21]
REV = [296, 301, 312, 309, 318, 326, 331, 329, 340, 347, 352, 361, 360, 366, 372, 379, 384, 391, 395, 402, 412]
REG = ["Прикладная информатика", "Программная инженерия", "Юриспруденция", "Экономика", "Машиностроение", "Строительство", "Менеджмент", "Информационная безопасность", "Электроэнергетика и электротехника", "Педагогическое образование", "Лингвистика", "Биотехнологии"]


def _frame(w_share: float, h_pt: float) -> Bbox:
    return Bbox(x=int(W * 0.06), y=int(H * 0.22), w=int(W * w_share), h=int(h_pt * 12700))


def _slide(bg: str = "FFFFFF"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return prs, s


def _chart(spec: ChartSpec, box: Bbox, bg: str = "FFFFFF", style: ChartStyleSpec = STYLE, **kw):
    prs, s = _slide(bg)
    gf = add_chart(s, box, spec, OUTLINE, style, TYPO, text_hex="000000" if bg == "FFFFFF" else "FFFFFF", **kw)
    return s, gf, gf.chart._chartSpace


def _line(cats=MONTHS, vals=REV, **kw) -> ChartSpec:
    return ChartSpec(type="line", categories=list(cats), series=[InlineSeries(name="Выручка", values=list(vals))], unit="млн ₽", y_title="Выручка", gaps=True, **kw)


def _labels(cs) -> list:
    return [(v.text or "") for v in cs.findall(f".//{C}cat//{C}pt/{C}v")] if cs.find(f".//{C}cat//{C}pt/{C}v") is not None else []


def _cat_cache(cs) -> list[str]:
    cache = cs.find(f".//{C}cat/{C}strRef/{C}strCache")
    n = int(cache.find(f"{C}ptCount").get("val"))
    out = [""] * n
    for p in cache.findall(f"{C}pt"):
        out[int(p.get("idx"))] = p.find(f"{C}v").text or ""
    return out


def _rot(cs) -> str:
    return cs.find(f".//{C}catAx/{C}txPr/{A}bodyPr").get("rot") or "0"


# ---------------------------------------------------------------------------------------------- H1: thin before turning


def test_a_time_axis_is_recognised_and_thinned_with_its_ends():
    assert is_time_axis(MONTHS) and is_time_axis([f"{y}" for y in range(2005, 2026)]) and is_time_axis(["I кв. 2025", "II кв. 2025", "III кв. 2025", "IV кв. 2025", "I кв. 2026"])
    assert is_time_axis([str(i) for i in range(1, 15)]) and is_time_axis(["2026-01", "2026-02", "2026-03", "2026-04", "2026-05"])
    assert not is_time_axis(["Москва", "Казань", "Самара", "Уфа", "Омск", "Тверь"]) and not is_time_axis(["янв", "фев"])
    for n, k in ((21, 2), (21, 3), (21, 4), (12, 3), (8, 2), (7, 4)):
        picks = thin_picks(n, k)
        assert picks[0] == 0 and picks[-1] == n - 1 and picks == sorted(set(picks))
        assert min(b - a for a, b in zip(picks, picks[1:])) >= 0.7 * k - 1e-9
    assert thin_picks(21, 2) == list(range(0, 21, 2))


def test_crowded_time_labels_are_thinned_not_turned():
    for w_share in (0.88, 0.6):
        s, gf, cs = _chart(_line(), _frame(w_share, 170))
        assert _rot(cs) == "0", w_share  # straight, whatever the crowding
        labels = _cat_cache(cs)
        shown = [i for i, t in enumerate(labels) if t]
        assert shown[0] == 0 and shown[-1] == len(MONTHS) - 1 and len(shown) < len(MONTHS)  # every k-th, the first and the last kept
        assert all(labels[i] in (MONTHS[i], short_years(period_levels(MONTHS)[-1])[i], period_levels(MONTHS)[-1][i]) for i in shown)  # (the labels, perhaps with the year cut)
        # every kept label's words stand in the slot: a renderer never breaks one inside («202 / 5»)
        ml = cs.find(f".//{C}plotArea/{C}layout/{C}manualLayout")
        slot = float(ml.find(f"{C}w").get("val")) * gf.width / 12700 / len(MONTHS)
        size = int(cs.find(f".//{C}catAx/{C}txPr/{A}p/{A}pPr/{A}defRPr").get("sz")) / 100
        assert all(text_width_pt(w, "Play", size) <= slot * 0.95 for i in shown for w in labels[i].split())


def test_the_year_is_cut_to_two_digits_before_a_label_turns():
    cats = ["янв 2025", "III кв. 2025", "2005", "Январь 2026"]
    assert short_years(period_levels(cats)[-1]) == ["янв ’25", "III кв. ’25", "’05", "Янв ’26"] and period_levels(cats)[0][0] == "янв 2025"
    assert period_levels(["2 квартал", "Январь 2026"])[-1] == ["2 кв.", "Янв 2026"]  # (the other callers' levels are as they were)
    # a very crowded axis: straight labels with the short year, not 45° ones
    s, gf, cs = _chart(_line(), _frame(0.6, 170))
    assert _rot(cs) == "0" and any("’" in t for t in _cat_cache(cs))


def test_fit_categories_prefers_a_smaller_size_then_thinning_then_turning_on_a_time_axis():
    cats = MONTHS
    slot = 24.0
    # one size is enough once smaller: flat
    got = fit_categories(cats, "Play", [10.8, 9.7, 8.9], 40.0, 40.0, ordered=True)
    assert got.mode in ("flat", "thin") and got.fits
    # no size fits flat in the slot: thinned (two short lines a label at worst), not turned
    got = fit_categories(cats, "Play", [10.8, 9.7, 8.9], slot, 40.0, ordered=True)
    assert got.mode == "thin" and got.keep >= 2 and got.fits and got.lines <= 2
    # a categorical axis is never thinned: it turns or wraps as before
    got = fit_categories(["Автоматизация обращений", "Маршрутизация звонков", "Обучение новых", "Контроль качества", "Прочее", "Другое"], "Play", [10.8], 60.0, 60.0)
    assert got.mode != "thin"
    # nothing fits the room under the axis: a time axis still keeps straight two-line labels (fits False), the frame is short
    got = fit_categories(cats, "Play", [10.8], slot, 5.0, ordered=True)
    assert got.mode == "thin" and not got.fits


# ---------------------------------------------------------------------------------------------- H1: the plot keeps its share


@pytest.mark.parametrize("kind", ["line21", "line21_titled", "column6", "stacked_legend", "bars12", "scatter"])
def test_the_plot_keeps_55_percent_of_a_frame_of_the_minimum_height(kind):
    fs = 12.0  # the chart's label size: the floor of a 540 pt slide (10.8 pt) snapped up this template's type scale
    if kind == "line21":
        spec, args = _line(), ("line", 21, fs, dict(label_lines=2))
    elif kind == "line21_titled":
        spec, args = _line(x_title="Месяц"), ("line", 21, fs, dict(label_lines=2, x_title=True))
    elif kind == "column6":
        spec = ChartSpec(type="column", categories=["I кв. 2025", "II кв. 2025", "III кв. 2025", "IV кв. 2025", "I кв. 2026", "II кв. 2026"], series=[InlineSeries(name="Выручка", values=[120, 134, 142, 151, 160, 171])], unit="млн ₽")
        args = ("column", 6, fs, dict(label_lines=2))
    elif kind == "stacked_legend":
        spec = ChartSpec(type="stacked_column", categories=["I кв.", "II кв.", "III кв.", "IV кв."], series=[InlineSeries(name="Прямые продажи", values=[120, 134, 142, 151]), InlineSeries(name="Партнёры", values=[71, 80, 86, 92])], unit="млн ₽", y_title="Выручка")
        args = ("stacked_column", 4, fs, dict(label_lines=1, legend_rows=1))
    elif kind == "bars12":
        spec = ChartSpec(type="stacked_bar", categories=REG, series=[InlineSeries(name="Бюджет", values=[120, 110, 40, 45, 100, 80, 35, 90, 85, 60, 25, 50]), InlineSeries(name="Платное", values=[140, 150, 210, 170, 55, 120, 165, 60, 40, 45, 85, 30])], unit="чел.", y_title="Принято, чел.")
        args = ("stacked_bar", 12, fs, dict(legend_rows=1, value_title=True))
    else:
        xs = [float(i) for i in range(1, 31)]
        spec = ChartSpec(type="scatter", xy=[XYSeries(name="a", x=xs, y=[x * 2.0 + (i % 5) for i, x in enumerate(xs)])], x_title="Подготовка, ч/нед.", y_title="Балл")
        args = ("scatter", 30, fs, dict(x_title=True))
    assert _chart(spec, _frame(0.88, 400))[1].verstka_font_pt == fs
    h_min = min_frame_height(args[0], args[1], args[2], **args[3])
    assert 60 < h_min < 300
    s, gf, cs = _chart(spec, _frame(0.88, h_min + 1))
    assert gf.verstka_plot_share >= MIN_PLOT - 0.01, (kind, round(h_min), gf.verstka_plot_share)
    ml = cs.find(f".//{C}plotArea/{C}layout/{C}manualLayout")
    assert float(ml.find(f"{C}h").get("val")) >= MIN_PLOT - 0.03  # (the written layout says the same)


def test_a_short_frame_keeps_both_axis_titles_and_squeezes_only_the_legends_air():
    # the audit asks every data chart for both titles: a short frame never drops one
    for h in (125, 260):
        cs = _chart(_line(x_title="Месяц"), _frame(0.88, h))[2]
        assert cs.find(f".//{C}catAx/{C}title") is not None and cs.find(f".//{C}valAx/{C}title") is not None
    spec = ChartSpec(type="stacked_column", categories=["I кв.", "II кв.", "III кв."], series=[InlineSeries(name="Прямые продажи", values=[3, 4, 5]), InlineSeries(name="Партнёры", values=[4, 5, 6])], unit="млн ₽", x_title="Квартал", y_title="Выручка")
    tall = _chart(spec, _frame(0.88, 300))[1]
    short = _chart(spec, _frame(0.88, 150))[1]
    assert short.verstka_plot_share > 0.45 and tall.verstka_plot_share > short.verstka_plot_share


def test_min_frame_height_grows_with_type_rows_legend_and_titles():
    base = min_frame_height("line", 21, 10.8, label_lines=2)
    assert min_frame_height("line", 21, 14.0, label_lines=2) > base and min_frame_height("line", 21, 10.8, label_lines=2, x_title=True, legend_rows=1) > base + 20
    assert min_frame_height("stacked_bar", 12, 10.8) > min_frame_height("stacked_bar", 6, 10.8) > min_frame_height("stacked_bar", 3, 10.8)
    assert 90 <= base <= 115  # ≈ 9½ lines of type


# ---------------------------------------------------------------------------------------------- H2: rows never touch


def _bars12(box: Bbox, labels_on: bool = True):
    spec = ChartSpec(type="stacked_bar", categories=REG, series=[InlineSeries(name="Бюджет", values=[120, 110, 40, 45, 100, 80, 35, 90, 85, 60, 25, 50]), InlineSeries(name="Платное", values=[140, 150, 210, 170, 55, 120, 165, 60, 40, 45, 85, 30])],
                     unit="чел.", x_title="Направление подготовки", y_title="Принято, чел.")
    return _chart(spec, box)


@pytest.mark.parametrize("h_pt", [120, 150, 190, 260, 340])
def test_twelve_horizontal_rows_never_have_text_over_text(h_pt):
    s, gf, cs = _bars12(_frame(0.88, h_pt))
    ml = cs.find(f".//{C}plotArea/{C}layout/{C}manualLayout")
    plot_h = float(ml.find(f"{C}h").get("val")) * gf.height / 12700
    pitch = plot_h / len(REG)
    labels = _cat_cache(cs)
    size = int(cs.find(f".//{C}catAx/{C}txPr/{A}p/{A}pPr/{A}defRPr").get("sz")) / 100
    shown = [i for i, t in enumerate(labels) if t]
    gaps = [(b - a) * pitch for a, b in zip(shown, shown[1:])]
    # the labels shown stand a line of type apart (a label of two lines two): none over another
    assert shown[0] == 0 and all(g >= size * 1.15 for g in gaps), (h_pt, size, round(pitch, 1), gaps)
    if h_pt >= 260:
        assert len(shown) == len(REG)  # room enough: every row named


def test_a_tall_enough_frame_names_every_row_and_keeps_the_bars_readable():
    s, gf, cs = _bars12(_frame(0.88, min_frame_height("stacked_bar", 12, 10.8, legend_rows=1, value_title=True) + 4))
    assert all(_cat_cache(cs))


# ---------------------------------------------------------------------------------------------- H3: the unit stays on the axis


def _headline(slide, text: str, name: str = "Title"):
    tb = slide.shapes.add_textbox(Emu(int(W * 0.06)), Emu(int(H * 0.08)), Emu(int(W * 0.88)), Emu(int(H * 0.1)))
    tb.name = name
    tb.text_frame.text = text
    return tb


def _axis_title_text(cs) -> str:
    return "".join(t.text or "" for t in cs.find(f".//{C}valAx/{C}title").iter(f"{A}t"))


def test_the_value_axis_carries_its_unit_whatever_the_headline_says():
    prs, s = _slide()
    _headline(s, "Выручка достигла 423 млн ₽ в сентябре, рост +23% год к году")
    cs = add_chart(s, _frame(0.88, 200), _line(), OUTLINE, STYLE, TYPO, text_hex="000000").chart._chartSpace
    assert _axis_title_text(cs) == "Выручка, млн ₽"
    assert cs.find(f"{C}chart/{C}title") is None  # (no second caption: the axis says it)
    # an explicit unit caption the composer set for this chart still counts
    prs, s = _slide()
    cap = _headline(s, "млн ₽", name="Unit")
    cap.top, cap.height = Emu(int(H * 0.22) - int(H * 0.05)), Emu(int(H * 0.03))
    cs = add_chart(s, _frame(0.88, 200), _line(), OUTLINE, STYLE, TYPO, text_hex="000000").chart._chartSpace
    assert _axis_title_text(cs) == "Выручка"
    # a title that names the unit already is kept as written
    cs = _chart(_line().model_copy(update={"y_title": "Выручка, млн ₽"}), _frame(0.88, 200))[2]
    assert _axis_title_text(cs) == "Выручка, млн ₽"


# ---------------------------------------------------------------------------------------------- H5: bounded measures, colours


def test_a_bounded_measure_ends_its_axis_on_its_natural_maximum():
    assert bounded_cap(None, "Средний балл ЕГЭ", [38.0, 74.0, 100.0]) == 100.0 and bounded_cap("%", None, [12, 87]) == 100.0
    assert bounded_cap("млн ₽", "Выручка", [10, 90]) is None and bounded_cap("%", None, [10, 140]) is None and bounded_cap(None, "Балл", [10, 140]) is None
    lo, hi, step = _nice_scale(0.0, 98.0, 100.0)
    assert (lo, hi) == (0.0, 100.0) and 100.0 % step == 0
    lo, hi, step = _nice_scale(0.0, 87.0)
    assert hi <= 100.0
    # a scatter of scores: its Y axis stops at 100, not 120
    lo, hi, step = _scatter_range([38.0, 52.0, 80.0, 100.0], 9, cap=100.0)
    assert hi == 100.0 and 100.0 % step == 0
    assert _scatter_range([38.0, 52.0, 80.0, 100.0], 9)[1] > 100.0  # (without the cap it would run past)
    xs = [float(i) for i in range(1, 31)]
    spec = ChartSpec(type="scatter", xy=[XYSeries(name="Технический", x=xs, y=[40.0 + x * 1.9 for x in xs])], x_title="Подготовка, ч/нед.", y_title="Средний балл ЕГЭ")
    cs = _chart(spec, _frame(0.88, 260))[2]
    ymax = float(next(a for a in cs.findall(f".//{C}valAx") if a.find(f"{C}axPos").get("val") == "l").find(f"{C}scaling/{C}max").get("val"))
    assert ymax == 100.0
    # a percent line likewise
    pct = ChartSpec(type="line", categories=MONTHS[:12], series=[InlineSeries(name="Своя", values=[61.0, 64, 70, 72, 79, 85, 91, 98, 97, 96, 99, 98]), InlineSeries(name="Чужая", values=[40.0, 42, 45, 47, 50, 52, 55, 57, 60, 61, 63, 66])], unit="%", gaps=True)
    va = _chart(pct, _frame(0.88, 260))[2].find(f".//{C}valAx")
    assert float(va.find(f"{C}scaling/{C}max").get("val")) == 100.0


def test_series_colours_stand_off_their_ground_at_three_to_one():
    light_green = ChartStyleSpec(series_colors=["A8D5A2", "C8E6C0", "FFE08A", "9CCC65"], font_family="Play", font_size_pt=7.0)
    for bg in ("FFFFFF", "0B1220"):
        pal = chart_palette(["A8D5A2"], bg, 3, extra=["C8E6C0", "FFE08A", "9CCC65"], floor=SERIES_CONTRAST)
        assert len(set(pal)) == 3 and all(contrast_ratio(c, bg) >= SERIES_CONTRAST for c in pal), (bg, pal)
    assert legible_color("A8D5A2", "FFFFFF") != "A8D5A2" and contrast_ratio(legible_color("A8D5A2", "FFFFFF"), "FFFFFF") >= 3.0
    assert legible_color("0077FF", "FFFFFF") == "0077FF"  # (a colour that stands off already is the template's own)
    # the same on the charts: lines, bars and markers of a light template
    two = ChartSpec(type="line", categories=MONTHS[:8], series=[InlineSeries(name="Scopus", values=[42, 60, 80, 100, 120, 160, 200, 381]), InlineSeries(name="РИНЦ", values=[255, 380, 500, 590, 650, 740, 770, 867])], gaps=True)
    cs = _chart(two, _frame(0.88, 260), style=light_green)[2]
    for ser in cs.iter(f"{C}ser"):
        col = ser.find(f"{C}spPr/{A}ln/{A}solidFill/{A}srgbClr").get("val")
        assert contrast_ratio(col, "FFFFFF") >= 3.0, col
    bars = ChartSpec(type="stacked_column", categories=["I", "II"], series=[InlineSeries(name="a", values=[3, 4]), InlineSeries(name="b", values=[2, 2]), InlineSeries(name="c", values=[1, 2])])
    cs = _chart(bars, _frame(0.6, 260), style=light_green)[2]
    for ser in cs.iter(f"{C}ser"):
        assert contrast_ratio(ser.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val"), "FFFFFF") >= 3.0


# ---------------------------------------------------------------------------------------------- H4: the group named before the legend


def test_the_group_name_stands_in_front_of_the_legend():
    spec = ChartSpec(type="stacked_column", categories=["I кв.", "II кв.", "III кв."], series=[InlineSeries(name="Бюджет", values=[3, 4, 5]), InlineSeries(name="Платное", values=[4, 5, 6])], unit="чел.", y_title="Принято, чел.")
    s, gf, cs = _chart(spec, _frame(0.88, 260), legend_title="Форма обучения")
    tb = gf.verstka_legend_title
    assert tb is not None and tb.name == "Chart legend title" and tb.text_frame.text == "Форма обучения"
    lg = cs.find(f".//{C}legend")
    ml = lg.find(f"{C}layout/{C}manualLayout")
    assert ml is not None and float(ml.find(f"{C}x").get("val")) > 0.2  # the legend stands after the title, on one row
    # the value axis is named for the measure: «Принято, чел.»
    assert _axis_title_text(cs) == "Принято, чел."
    # the title's box ends where the legend begins, inside the frame
    assert gf.left <= tb.left and tb.left + tb.width <= gf.left + gf.width and tb.top > gf.top + gf.height * 0.8
    assert tb.left + tb.width <= gf.left + int(float(ml.find(f"{C}x").get("val")) * gf.width) + 12700 * 4
    # no title asked: no box, the legend is the legend as before
    s, gf, cs = _chart(spec, _frame(0.88, 260))
    assert gf.verstka_legend_title is None and cs.find(f".//{C}legend/{C}layout") is None
