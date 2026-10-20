"""Поломки записи демо 20.10: числа базы знаний — источник для проверки чисел, число слайдов держится и с базой знаний."""

from pptx import Presentation
from pptx.util import Inches, Pt

from verstka.audit.checks.facts import figure_not_in_brief, figure_sources
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import AuditContext
from verstka.knowledge.with_data import _fit_count
from verstka.schemas.outline import Brief, OutlineSlide, PatternKind

BRIEF = "Закрытие Q3 2026 — VK Cloud, для правления. Госсектор против цели года."
KB = "Целевое число клиентов госсектора — не менее 150 к 31 декабря 2026 года. Документ утверждён 20 января 2026 года."


def _deck(tmp_path, lines):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[6])  # the cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, line in enumerate(lines):
        tb = s.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(10), Inches(0.8))
        tb.text_frame.text = line
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return build_deck_ir(path)


def test_the_knowledge_base_is_a_source_of_the_figures(tmp_path):
    ir = _deck(tmp_path, ["Цель — не менее 150 клиентов к 31 декабря 2026 года", "Утверждено 20 января 2026 года"])
    alone = figure_not_in_brief(AuditContext(ir=ir, manifest=None, outline=None, brief_text=BRIEF))
    assert alone  # without the base its figures are not the brief's
    with figure_sources([KB]):
        assert figure_not_in_brief(AuditContext(ir=ir, manifest=None, outline=None, brief_text=BRIEF)) == []
    # outside the block the check is as before
    assert figure_not_in_brief(AuditContext(ir=ir, manifest=None, outline=None, brief_text=BRIEF))


def test_a_share_of_two_given_figures_of_the_slide_is_derived(tmp_path):
    ir = _deck(tmp_path, ["Клиентов сейчас: 131", "Цель: 150", "87 %"])
    with figure_sources([KB + " Сейчас клиентов 131."]):
        assert figure_not_in_brief(AuditContext(ir=ir, manifest=None, outline=None, brief_text=BRIEF)) == []
    # a percent no two figures of the slide give stays a remark
    ir2 = _deck(tmp_path, ["Клиентов сейчас: 131", "Цель: 150", "64 %"])
    with figure_sources([KB + " Сейчас клиентов 131."]):
        assert figure_not_in_brief(AuditContext(ir=ir2, manifest=None, outline=None, brief_text=BRIEF))


def _s(sid, kind=PatternKind.bullets):
    return OutlineSlide(id=sid, kind=kind, headline=sid)


def test_the_slide_count_holds_with_the_knowledge_slides():
    slides = [_s("title", PatternKind.title), _s("kb-1"), _s("data-1"), _s("data-3"), _s("kb-2"), _s("kb-3"), _s("summary"), _s("data-plan-5")]
    kb = {id(s) for s in slides if s.id.startswith("kb-")}
    gone = _fit_count(slides, Brief(text="x", slide_count=6), kb)
    assert [s.id for s in slides] == ["title", "kb-1", "data-1", "data-3", "kb-2", "data-plan-5"]
    assert gone == ["summary", "kb-3"]
    # no count asked — the slides stay
    more = [_s("title", PatternKind.title), _s("kb-1"), _s("data-1")]
    assert _fit_count(more, Brief(text="x"), {id(more[1])}) == [] and len(more) == 3


def test_the_data_decks_proposed_steps_go_when_over_the_count():
    from types import SimpleNamespace

    from verstka.data.story import fit_count
    from verstka.schemas.outline import DeckOutline

    o = DeckOutline(title="Итоги", slides=[_s("title", PatternKind.title), _s("data-2"), _s("data-plan-1"), _s("data-5"), _s("data-3"),
                                           _s("data-4"), _s("data-1"), _s("summary"), _s("data-plan-6"), _s("next-steps")])
    said = fit_count(SimpleNamespace(outline=o), [], None, 8)
    assert [s.id for s in o.slides] == ["title", "data-2", "data-plan-1", "data-5", "data-3", "data-4", "data-1", "data-plan-6"]
    assert said and "не больше 8" in said[0]


def test_an_edits_grounding_keeps_the_knowledge_bases_figures():
    from verstka.planning.grounding import BriefIndex, figures

    idx = BriefIndex(BRIEF)
    f = next(iter(figures("не менее 150 клиентов")))
    assert idx.verdict(f) != "ok"
    with figure_sources([KB]):
        assert BriefIndex(BRIEF).verdict(f) == "ok"
        made_up = next(iter(figures("не менее 170 клиентов")))
        assert BriefIndex(BRIEF).verdict(made_up) != "ok"
    assert BriefIndex(BRIEF).verdict(f) != "ok"
