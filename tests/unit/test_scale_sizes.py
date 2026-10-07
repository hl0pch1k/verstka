"""Sizes on the template's scale: tile labels never break inside a word (one size for the row), the user's code takes
the template's steps before half points, a display heading keeps its weight on a scale with a gap."""
from __future__ import annotations

import pytest
from lxml import etree

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.compose import Composer, Kit
from verstka.rendering.deck import DeckBuilder, element_bbox
from verstka.rendering.fit import fit_size
from verstka.rendering.fonts import text_width_pt
from verstka.schemas.common import EMU_PER_PT, Bbox, PatternKind
from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent

from tests.unit.test_tpl_compose import BIG_H, BIG_W, build_sparse_deck

CLASSIC = [12.0, 14.0, 16.0, 18.0, 20.0, 24.0, 40.0, 54.0, 66.0, 72.0]  # a scale with a gap: 24 → 40


@pytest.mark.parametrize("min_ratio", [0.6, 0.7, 0.8])
def test_fit_size_never_returns_a_size_between_two_scale_steps(min_ratio):
    """Классика s3/s4/s9: a 40 pt heading that does not fit was set at 36.8 / 31.2 pt — the template sets nothing between
    24 and 40. The fit takes a step of the scale (or stays at the given size), never an invented one."""
    text = ["Редкие проверки пропускают изменения условий в теплице"]
    box = Bbox(x=0, y=0, w=int(864 * EMU_PER_PT), h=int(76 * EMU_PER_PT))
    res = fit_size(text, box, "Open Sans", 40.0, False, CLASSIC, min_ratio=min_ratio)
    assert res.fits
    assert any(abs(res.size_pt - s) <= 0.05 for s in CLASSIC), res.size_pt
    # under the scale's smallest step the template has no size: an in-between size is still allowed there
    tiny = fit_size(["Слово"] * 12, Bbox(x=0, y=0, w=int(60 * EMU_PER_PT), h=int(60 * EMU_PER_PT)), "Open Sans", 12.0, False, CLASSIC)
    assert tiny.size_pt <= 12.0


@pytest.fixture
def big_env(tmp_path):
    deck = build_sparse_deck(tmp_path / "big.pptx", BIG_W, BIG_H, title_pt=100, body_pt=25)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return deck, manifest, ws


def test_tile_labels_keep_their_words_whole_in_a_narrow_column(big_env):
    """Marketing s8: four tiles beside the photo place, three lines under them — labels at the body size broke inside
    their words («Температ / ура»). Every label word stands whole in its tile, at a size of the template's scale."""
    deck, manifest, _ = big_env
    nums = [NumberCallout(value=v, label=l) for v, l in [("25,4 °C", "Температура воздуха"), ("67%", "Влажность воздуха"), ("72%", "Заполненность бака"), ("40 сек", "Последнее обновление")]]
    extra = ["Графики за сутки рядом с текущими значениями", "Предупреждения объясняют причину и требуемое действие", "Отметки включения полива и вентиляции на графике"]
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Панель показывает текущие показания и историю за 24 часа", content=SlideContent(numbers=nums, bullets=extra, photo_slot="Фото"))
    slide = DeckBuilder(deck).add_blank_slide(6)
    kit = Kit(manifest, BIG_W, BIG_H, "FFFFFF")
    comp = Composer(slide, kit, o, DeckOutline(title="T", slides=[o], strategy="structured"), "structured", manifest)
    area = Bbox(x=int(269 * EMU_PER_PT), y=int(429 * EMU_PER_PT), w=int(869 * EMU_PER_PT), h=int(475 * EMU_PER_PT))  # the column beside the photo
    n0 = len(comp.cv.tree)
    comp.kpis(area, nums, extra)
    labels = [el for el in list(comp.cv.tree)[n0:] if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Label")]
    assert len(labels) == 4
    for el in labels:
        x, y, w, h = element_bbox(el)
        for r in el.iter(q("a:r")):
            rpr = r.find(q("a:rPr"))
            size = int(rpr.get("sz")) / 100
            assert size in kit.sizes, size
            for word in (r.find(q("a:t")).text or "").split(" "):
                assert text_width_pt(word, kit.font, size, rpr.get("b") == "1") <= w / EMU_PER_PT * 0.97, (word, size, w / EMU_PER_PT)


def test_tile_labels_of_a_row_share_one_size(big_env):
    """The autofix shrank each broken label on its own (Marketing compact s8: 25 / 30.5 / 22 / 27.5 pt in one row):
    the composer sets the row's labels at one size where every word stands whole."""
    deck, manifest, _ = big_env
    nums = [NumberCallout(value=v, label=l) for v, l in [("25,4 °C", "Температура воздуха"), ("67%", "Влажность воздуха"), ("72%", "Заполненность бака"), ("40 сек", "Последнее обновление")]]
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Панель показывает текущие показания", content=SlideContent(numbers=nums))
    slide = DeckBuilder(deck).add_blank_slide(6)
    comp = Composer(slide, Kit(manifest, BIG_W, BIG_H, "FFFFFF"), o, DeckOutline(title="T", slides=[o], strategy="compact"), "compact", manifest)
    n0 = len(comp.cv.tree)
    comp.kpis(Bbox(x=int(960 * EMU_PER_PT), y=int(578 * EMU_PER_PT), w=int(790 * EMU_PER_PT), h=int(420 * EMU_PER_PT)), nums, [])
    sizes = {int(r.get("sz")) / 100 for el in list(comp.cv.tree)[n0:] if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Label") for r in el.iter(q("a:rPr")) if r.get("sz")}
    assert len(sizes) == 1, sizes


def test_display_headings_keep_their_weight_on_a_scale_with_a_gap():
    """Классика s14: «Спасибо за внимание!» does not fit at 40 pt. The template sets nothing between 24 and 40: the
    heading takes an even 36 pt rather than the far 24 pt step (a closing at 24 pt over an empty half of the slide reads
    unfinished); a step close below is taken as it is."""
    from verstka.rendering.clone import _snap_display

    assert _snap_display(36.8, CLASSIC) == 36.0
    assert _snap_display(49.68, CLASSIC) == 40.0  # a step close below stays as it was
    lct_with_slots = [10.5, 12.0, 13.0, 14.0, 15.0, 16.0, 18.67, 20.0, 24.0, 28.0, 36.0, 44.0, 48.0]
    assert _snap_display(45.6, lct_with_slots) == 44.0  # LCT's cover keeps 44 pt (a size its samples set)


def test_code_keeps_its_own_size_but_running_text_keeps_the_scale(tmp_path):
    """Education s7: 25 lines of the user's code fit the template's code box only at 9 pt (its scale starts at 12 pt).
    The code's size is its own — set by its longest line, as its monospace face is its own — while running text at
    9 pt is still flagged."""
    from pptx import Presentation
    from pptx.util import Emu, Pt

    from verstka.audit.checks.template import size_not_in_scale
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext

    deck = build_sparse_deck(tmp_path / "t.pptx", title_pt=40, body_pt=24)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    # a template whose own scale starts at 12 pt (no derived ladder under it)
    manifest.tokens.typography.sizes_used = [12.0, 14.0, 16.0, 18.0, 20.0, 24.0, 40.0]
    manifest.tokens.typography.derived_sizes = []
    prs = Presentation(str(deck))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, face in enumerate(("Consolas", "Arial")):
        tf = s.shapes.add_textbox(Emu(914400), Emu(914400 * (1 + 2 * i)), Emu(4 * 914400), Emu(914400)).text_frame
        tf.text = "def irrigation_decision(soil_moisture):" if face == "Consolas" else "Обычный текст абзаца"
        r = tf.paragraphs[0].runs[0]
        r.font.size, r.font.name = Pt(9), face
    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    issues = [i for i in size_not_in_scale(AuditContext(ir=build_deck_ir(out), manifest=manifest)) if i.slide == len(prs.slides)]
    assert len(issues) == 1 and issues[0].details["size"] == 9.0
    ir = build_deck_ir(out)
    flagged = {e.id for e in ir.slides[-1].elements for i in issues if e.id in i.element_ids}
    assert all(r.font == "Arial" for e in ir.slides[-1].elements if e.id in flagged for p in e.paragraphs for r in p.runs)


def test_code_size_takes_the_templates_step_before_a_half_point():
    from verstka.rendering.compose import code_size

    lines = ["x = 1"] * 10
    # 13 pt would fit; the template sets 12 and 14: the code takes 12, never 13
    size, ok = code_size(lines, "Consolas", 400.0, 10 * 13.2 * 1.25, 16.0, scale=[12.0, 14.0, 16.0, 24.0])
    assert ok and size == 12.0
    # nothing of the scale fits: half points under its smallest step, as before
    size, ok = code_size(lines, "Consolas", 400.0, 10 * 9.1 * 1.25, 16.0, scale=[12.0, 14.0, 16.0])
    assert ok and size == 9.0
