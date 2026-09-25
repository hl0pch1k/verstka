"""The three variants present one content three ways, the type fills its area, no word is broken, a time axis keeps
its labels on one line, a cover carries its goal line.

Synthetic decks only (the conftest deck and a hand-drawn cover): the rules hold for any template."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.audit.checks.layout import word_break
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import AuditContext, all_checks
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.compat import composition_for
from verstka.matching.scorer import cover_goal
from verstka.rendering.charts import short_categories
from verstka.rendering.clone import render_clone
from verstka.rendering.compose import kpi_callout
from verstka.rendering.deck import DeckBuilder
from verstka.rendering.fallbacks import prepare_slide
from verstka.rendering.fonts import text_width_pt
from verstka.rendering.synth import render_synth
from verstka.schemas.common import BboxFrac, Family, PatternKind, SlotRole
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, OutlineSlide, SlideContent, SlideItem
from verstka.schemas.template import Capacity, FontUsage, Pattern, SlideSize, Slot, SlotStyle, Spacing, TemplateManifest, Tokens, TypeStep, Typography

W, H = 12192000, 6858000

PIE = ChartSpec(type="pie", title="Доля выручки", unit="%", categories=["Кофе", "Десерты и выпечка", "Чай и другие напитки"], series=[InlineSeries(name="Доля", values=[60, 25, 15])])
PROFIT = ChartSpec(type="column", title="Операционная прибыль", unit="₽", categories=["Текущая", "Прогноз"], series=[InlineSeries(name="Прибыль", values=[120000, 254795])])
MONTHS = ChartSpec(
    type="line", title="Выручка", unit="₽",
    categories=["Сейчас", "1-й месяц", "2-й месяц", "3-й месяц", "4-й месяц", "5-й месяц", "6-й месяц"],
    series=[InlineSeries(name="Выручка", values=[900000, 930000, 970000, 1015000, 1060000, 1100000, 1138500])],
)
SPEND = ChartSpec(type="column", title="Вложения по статьям", unit="₽", categories=["Витрина", "Лояльность"], series=[InlineSeries(name="Вложения", values=[70000, 35000])])
SIDE = ["Ежемесячная выручка — 900 000 рублей", "Средний чек — 300 рублей", "Результат не гарантирован"]
TAKE = "Операционная прибыль достигнет 254 795 рублей в месяц"


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _render(env, oslide: OutlineSlide, strategy: str):
    deck, manifest, ws = env
    outline = DeckOutline(title="T", strategy=strategy, slides=[oslide])
    oslide, ps, _ = prepare_slide(oslide, LayoutSlide(outline_id=oslide.id, mode="synth", composition=composition_for(oslide)), outline)
    slide, warnings = render_synth(DeckBuilder(deck), ps, oslide, manifest, ws, outline)
    return slide, warnings


def _named(slide, name: str):
    return [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] == name]


def _text(sh) -> str:
    if not getattr(sh, "has_text_frame", False):
        return ""
    return sh.text_frame.text.replace("\u00a0", " ").replace("\u2060", "")


def _sizes(sh) -> set[float]:
    return {r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()}


def _chart_slide(strategy: str, env, chart=PROFIT, side=SIDE, take=TAKE):
    s = OutlineSlide(id="c", kind=PatternKind.chart, headline="Прибыль вырастет более чем в 2 раза", content=SlideContent(chart=chart, bullets=list(side)), takeaway=take)
    return _render(env, s, strategy)


# ---------------------------------------------------------------------------------------------- callouts


@pytest.mark.parametrize("line, want", [
    ("Ежемесячная выручка — 900 000 рублей", ("900 000 ₽", "Ежемесячная выручка")),
    ("Рост операционной прибыли на 112,3%", ("112,3%", "Рост операционной прибыли")),
    ("Прогноз: 1 138 500 ₽ на 6-й месяц", ("1 138 500 ₽", "Прогноз на 6-й месяц")),
    ("100 покупок в день", ("100", "Покупок в день")),
    ("Результат не гарантирован", None),
    ("Дневные предложения с 15:00 до 18:00", None),
    ("Добавки к напиткам за 40–60 рублей", None),
    ("Месяц 2 — запуск комбо", None),
    ("Запуск в 2026 году", None),
])
def test_a_line_with_one_figure_is_a_callout(line, want):
    assert kpi_callout(line) == want


# ---------------------------------------------------------------------------------------------- one chart, three looks


def test_structured_chart_keeps_its_column_and_its_strip(env):
    slide, warnings = _chart_slide("structured", env)
    assert not [w for w in warnings if "fail" in w], warnings
    chart = next(sh for sh in slide.shapes if sh.has_chart)
    side = _named(slide, "Takeaway")
    assert side and side[0].left >= chart.left + chart.width, "the lines stand in a column right of the chart"
    assert _named(slide, "Conclusion strip") or _named(slide, "Conclusion"), "the conclusion under the chart"
    assert not _named(slide, "Figure")


def test_visual_chart_sets_the_figures_of_its_lines_as_callouts(env):
    slide, warnings = _chart_slide("visual", env)
    assert not [w for w in warnings if "fail" in w], warnings
    chart = next(sh for sh in slide.shapes if sh.has_chart)
    figs = _named(slide, "Figure")
    texts = [_text(f) for f in figs]
    assert any("900 000" in t for t in texts) and any("300" in t for t in texts), texts
    labels = [_text(sh) for sh in _named(slide, "Label")]
    assert "Ежемесячная выручка" in labels and "Средний чек" in labels
    body = max(_sizes(next(sh for sh in _named(slide, "Label") if _text(sh) == "Средний чек")))
    for f in figs:
        assert f.left >= chart.left + chart.width, "the callouts stand beside the chart"
        assert max(_sizes(f)) >= 1.6 * body, "a callout's figure is set large"
        assert f.width >= text_width_pt(_text(f).split("₽")[0].strip(), "Play", max(_sizes(f))) * 12700, "a figure on one line"
    # a line without a figure stays a line, under the callouts
    note = next(sh for sh in slide.shapes if "Результат не гарантирован" in _text(sh))
    assert note.top > max(f.top for f in figs)
    # the visual chart is wider than the structured one
    s_slide, _ = _chart_slide("structured", env)
    s_chart = next(sh for sh in s_slide.shapes if sh.has_chart)
    assert chart.width > s_chart.width


def test_visual_chart_with_lines_without_figures_sets_them_under_a_full_width_chart(env):
    slide, _ = _chart_slide("visual", env, side=["Комбо с напитками и выпечкой", "Программа лояльности", "Предложения для офисов"], take=None)
    chart = next(sh for sh in slide.shapes if sh.has_chart)
    notes = _named(slide, "Note")
    assert len(notes) == 3, [sh.name for sh in slide.shapes]
    assert all(n.top >= chart.top + chart.height for n in notes) and len({n.top for n in notes}) == 1


def test_compact_chart_is_mirrored_with_the_conclusion_leading_its_text(env):
    slide, warnings = _chart_slide("compact", env)
    assert not [w for w in warnings if "fail" in w], warnings
    chart = next(sh for sh in slide.shapes if sh.has_chart)
    col = _named(slide, "Chart note")
    assert col, [sh.name for sh in slide.shapes]
    col = col[0]
    assert col.left + col.width <= chart.left, "the text column stands left of the chart"
    first = col.text_frame.paragraphs[0]
    assert "254 795" in first.text.replace(" ", " ") and all(r.font.bold for r in first.runs if r.text.strip()), "the conclusion is the bold lead line"
    assert not _named(slide, "Conclusion strip") and not _named(slide, "Conclusion"), "no strip under the chart"
    assert all(b in _text(col) for b in ("Средний чек", "Результат не гарантирован"))


def test_the_three_variants_differ_on_a_dictated_chart_slide(env):
    def layout(strategy: str):
        slide, _ = _chart_slide(strategy, env, chart=PIE)
        return sorted((sh.name.rsplit(" ", 1)[0], round(sh.left / W, 2), round(sh.width / W, 2)) for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Chart", "Figure", "Takeaway", "Chart note", "Conclusion"))

    got = {s: layout(s) for s in ("structured", "visual", "compact")}
    assert got["structured"] != got["visual"] != got["compact"] != got["structured"]


# ---------------------------------------------------------------------------------------------- a pair of charts


def test_a_time_series_takes_the_larger_share_of_a_pair(env):
    s = OutlineSlide(id="p", kind=PatternKind.chart, headline="Выручка вырастет на 26,5% за 6 месяцев", content=SlideContent(chart=SPEND, chart2=MONTHS))
    slide, warnings = _render(env, s, "structured")
    assert not [w for w in warnings if "fail" in w], warnings
    charts = sorted((sh for sh in slide.shapes if sh.has_chart), key=lambda sh: sh.left)
    assert len(charts) == 2
    assert charts[1].width > 1.3 * charts[0].width, "the line (second) is wider than the columns"


def test_crowded_time_labels_are_shortened_in_the_chart_only():
    assert short_categories(["Сейчас", "1-й месяц", "Месяц 3", "Январь 2026", "2 квартал"]) == ["Сейчас", "1-⁠й мес.", "Мес. 3", "Янв 2026", "2 кв."]
    assert short_categories(["Q1", "Q2"]) == ["Q1", "Q2"]


# ---------------------------------------------------------------------------------------------- the area is filled, no word is broken


def test_cards_take_a_readable_size_before_small_type(env):
    items = [
        SlideItem(title="Контролируйте показатели", bullets=["Количество покупок в день", "Средний чек", "Доля чеков с едой", "Доля переменных расходов в выручке", "Операционная прибыль за месяц"]),
        SlideItem(title="Риски и меры", bullets=["Слабый отклик на предложения", "Рост закупочных цен", "Перегрузка сотрудников", "Тестировать акции небольшими запусками", "Сравнивать поставщиков", "Корректировать графики смен"]),
    ]
    s = OutlineSlide(id="k", kind=PatternKind.cards, headline="Три фактора определяют рост прибыли", content=SlideContent(items=items), takeaway="Рост прибыли зависит от трех измеримых изменений: больше покупок, выше средний чек и меньше потерь")
    deck, manifest, ws = env
    slide, _ = _render(env, s, "structured")
    from verstka.rendering.compose import Kit

    kit = Kit(manifest, W, H, "FFFFFF")
    texts = [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Card text", "Column")]
    assert texts, [sh.name for sh in slide.shapes]
    assert min(min(_sizes(t)) for t in texts) >= kit.body - 0.05, [_sizes(t) for t in texts]


def test_no_word_is_wider_than_its_step(env):
    steps = [SlideItem(title="", text="Учет показателей и обновление меню"), SlideItem(title="Месяц 2", text="Запуск комбо и обучение сотрудников"),
             SlideItem(title="", text="Программа лояльности и партнерства"), SlideItem(title="Месяц 4", text="Продвижение дневных предложений"),
             SlideItem(title="", text="Корректировка предложений"), SlideItem(title="Месяц 6", text="Оценка результатов")]
    s = OutlineSlide(id="t", kind=PatternKind.timeline, headline="План на шесть месяцев", content=SlideContent(items=steps))
    for strategy in ("structured", "visual", "compact"):
        slide, _ = _render(env, s, strategy)
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.text_frame.word_wrap is False:
                continue
            bp = sh.text_frame._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
            if bp is not None and bp.get("wrap") == "none":
                continue
            usable = sh.width / 12700 - int(bp.get("lIns", 91440) if bp is not None else 91440) / 12700 - int(bp.get("rIns", 91440) if bp is not None else 91440) / 12700
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    for word in re.split(r"[ \t\r\n]+", r.text):
                        if word:
                            assert text_width_pt(word, r.font.name, r.font.size.pt, bool(r.font.bold)) <= usable, (strategy, word, usable)


def test_the_audit_flags_a_word_wider_than_its_box(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for x, width, text in ((0.1, 0.08, "Корректировка предложений"), (0.5, 0.3, "Корректировка предложений")):
        tb = s.shapes.add_textbox(Emu(int(x * W)), Emu(int(0.3 * H)), Emu(int(width * W)), Emu(int(0.3 * H)))
        tb.text_frame.word_wrap = True
        r = tb.text_frame.paragraphs[0].add_run()
        r.text, r.font.size, r.font.name = text, Pt(18), "Play"
    pptx = tmp_path / "words.pptx"
    prs.save(pptx)
    manifest = analyze_template(pptx, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    issues = word_break(AuditContext(ir=build_deck_ir(pptx), manifest=manifest))
    assert len(issues) == 1 and issues[0].check_id == "word_break" and "Корректировка" in issues[0].message
    assert issues[0].autofix is not None and issues[0].autofix.action == "shrink_text"
    assert "word_break" in {spec.id for spec, _ in all_checks()}


# ---------------------------------------------------------------------------------------------- the cover's goal line


def _cover(tmp: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Emu(int(0.06 * W)), Emu(int(0.28 * H)), Emu(int(0.55 * W)), Emu(int(0.16 * H)))
    t.text_frame.word_wrap = True
    r = t.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.name = "Название", Pt(48), "Play"
    pptx = tmp / "cover.pptx"
    prs.save(pptx)
    typo = Typography(
        families=[FontUsage(family="Play", weight=1.0)],
        scale=[TypeStep(role="display", size_pt=54), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=16), TypeStep(role="body", size_pt=12), TypeStep(role="caption", size_pt=10)],
        sizes_used=[54, 48, 36, 24, 20, 16, 14, 12, 10],
    )
    slot = Slot(id="title_t", role=SlotRole.title, shape_id=str(t.shape_id), bbox=BboxFrac(x=0.06, y=0.28, w=0.55, h=0.16), style=SlotStyle(font_family="Play", size_pt=48), capacity=Capacity(max_chars=20, max_lines=1), sample_text="Название")
    p = Pattern(id="pc", source_slide=1, kind=PatternKind.title, family=Family.light, slots=[slot])
    m = TemplateManifest(template_id="synthetic", source_file="s.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(typography=typo, spacing=Spacing(safe_area=BboxFrac(x=0.04, y=0.06, w=0.92, h=0.86))), patterns=[p], n_slides=1)
    return pptx, p, m


def _render_cover(tmp_path, o: OutlineSlide):
    pptx, p, m = _cover(tmp_path)
    ws = TemplateWorkspace.create(pptx, tmp_path / "ws")
    slide, _ = render_clone(DeckBuilder(pptx), LayoutSlide(outline_id=o.id, mode="clone", pattern_id="pc"), o, p, m, ws, DeckOutline(title="x", slides=[o]))
    return slide


GOAL = "Цель: увеличить ежемесячную операционную прибыль со 120 000 до 255 000 рублей"


def test_the_cover_sets_its_goal_under_the_subtitle(tmp_path):
    o = OutlineSlide(id="t", kind=PatternKind.title, headline="Кофейня «Точка кофе»", subtitle="План увеличения прибыли за 6 месяцев", content=SlideContent(paragraphs=[GOAL]), footnote="Все цифры условные")
    assert cover_goal(o) == GOAL
    slide = _render_cover(tmp_path, o)
    goal = next(sh for sh in slide.shapes if sh.has_text_frame and "Цель:" in sh.text_frame.text)
    sub = next(sh for sh in slide.shapes if sh.has_text_frame and "План увеличения" in sh.text_frame.text)
    note = next(sh for sh in slide.shapes if sh.has_text_frame and "условные" in sh.text_frame.text)
    assert "·" not in _text(sub), "the goal is not glued to the subtitle"
    assert goal.top >= sub.top + sub.height - Emu(Pt(2)) and abs(goal.left - sub.left) <= int(0.01 * W)
    assert max(_sizes(goal)) < max(_sizes(sub)) and max(_sizes(goal)) >= 10
    assert note.top >= goal.top + goal.height - Emu(Pt(2)), "the small print keeps its place under everything"
    assert goal.top + goal.height <= int(0.95 * H)


def test_a_cover_without_a_goal_is_unchanged(tmp_path):
    o = OutlineSlide(id="t", kind=PatternKind.title, headline="Кофейня «Точка кофе»", subtitle="План увеличения прибыли за 6 месяцев")
    assert cover_goal(o) is None
    slide = _render_cover(tmp_path, o)
    assert not any(sh.name.startswith("Цель") for sh in slide.shapes)


# ---------------------------------------------------------------------------------------------- the conclusion, per variant


COLS = [SlideItem(title="Цель", text="Снизить долю расходов с 35% до 33% выручки"), SlideItem(title="Меры", text="Ежедневный учет остатков и пересмотр цен"), SlideItem(title="Списания", text="Сократить списания с 27 000 до 15 000 рублей")]
SAVING = "Экономия 22 770 рублей при целевой выручке 1 138 500 рублей"


def test_compact_says_its_conclusion_first(env):
    s = OutlineSlide(id="k", kind=PatternKind.cards, headline="Снижение доли расходов даст экономию", content=SlideContent(items=COLS), takeaway=SAVING)
    slide, _ = _render(env, s, "compact")
    take = _named(slide, "Conclusion")
    assert take and not _named(slide, "Conclusion strip"), [sh.name for sh in slide.shapes]
    blocks = [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Card", "Column", "Card text")]
    assert blocks and take[0].top + take[0].height <= min(b.top for b in blocks), "the conclusion stands over the content"
    assert all(r.font.bold for p in take[0].text_frame.paragraphs for r in p.runs if r.text.strip())


def test_visual_strip_leads_with_its_key_figure(env):
    s = OutlineSlide(id="k", kind=PatternKind.cards, headline="Снижение доли расходов даст экономию", content=SlideContent(items=COLS), takeaway=SAVING)
    slide, _ = _render(env, s, "visual")
    fig, take = _named(slide, "Conclusion figure"), _named(slide, "Conclusion")
    assert fig and take and "22 770" in _text(fig[0]) and fig[0].left + fig[0].width <= take[0].left
    assert max(_sizes(fig[0])) > max(_sizes(take[0]))
    # the structured strip keeps its accent bar
    s_slide, _ = _render(env, s, "structured")
    assert _named(s_slide, "Conclusion bar") and not _named(s_slide, "Conclusion figure")


# ---------------------------------------------------------------------------------------------- the fill pass


SHORT_CARDS = [SlideItem(title="Средний чек", text="Только 20% чеков содержат еду"), SlideItem(title="Свободные часы", text="65% покупок — утром"), SlideItem(title="Потери", text="27 000 рублей в месяц")]


def _card_block(slide):
    cards = [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Card", "Column") and not (sh.has_text_frame and sh.text_frame.text.strip())]
    texts = [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Card text", "Column")]
    return cards, texts


def _compose(env, oslide: OutlineSlide, comp: str, strategy: str, area):
    from verstka.rendering.compose import Composer, Kit

    deck, manifest, _ = env
    b = DeckBuilder(deck)
    slide = b.prs.slides.add_slide(b.prs.slide_layouts[6])
    kit = Kit(manifest, W, H, "FFFFFF")
    kit.head_size = 36.0
    outline = DeckOutline(title="T", strategy=strategy, slides=[oslide])
    cp = Composer(slide, kit, oslide, outline, strategy, manifest)
    cp.compose(comp, area)
    return slide


@pytest.mark.parametrize("strategy", ["structured", "visual", "compact"])
def test_a_short_block_grows_to_fill_its_area(env, strategy, monkeypatch):
    from verstka.rendering.compose import Composer
    from verstka.schemas.common import Bbox

    area = Bbox(x=int(0.06 * W), y=int(0.1 * H), w=int(0.88 * W), h=int(0.8 * H))
    s = OutlineSlide(id="k", kind=PatternKind.cards, headline="Средний чек и пустые часы — основные потери", content=SlideContent(items=SHORT_CARDS))
    grown = _compose(env, s, "cards", strategy, area)
    monkeypatch.setattr(Composer, "FILL_LOW", 0.0)  # the fill pass off: the block as it was composed first
    plain = _compose(env, s, "cards", strategy, area)

    def bottom(slide):
        return max(sh.top + sh.height for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Card", "Card text", "Column", "Index", "Badge"))

    def size(slide):
        return max(max(_sizes(sh)) for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] in ("Card text", "Column"))

    p_fill = (bottom(plain) - area.y) / area.h
    g_fill = (bottom(grown) - area.y) / area.h
    assert p_fill < 0.75, "the block leaves the area half empty when composed at the usual cap"
    assert g_fill > p_fill + 0.1 and size(grown) >= size(plain), (p_fill, g_fill)
    assert g_fill <= 1.0 + 1e-3, "never past the foot of its area"


def test_a_table_is_set_at_the_body_size_when_its_rows_fit(env):
    from verstka.rendering.compose import Kit
    from verstka.schemas.outline import TableData

    t = TableData(columns=["Показатель", "Сейчас", "Цель"], rows=[["Покупки в день", "100", "115"], ["Средний чек", "300 рублей", "330 рублей"], ["Месячная выручка", "900 000 рублей", "1 138 500 рублей"], ["Операционная прибыль", "120 000 рублей", "254 795 рублей"]])
    s = OutlineSlide(id="t", kind=PatternKind.table, headline="Операционная прибыль вырастет", content=SlideContent(table=t))
    slide, _ = _render(env, s, "structured")
    _, manifest, _ = env
    kit = Kit(manifest, W, H, "FFFFFF")
    table = next(sh for sh in slide.shapes if sh.has_table)
    sizes = {r.font.size.pt for row in table.table.rows for c in row.cells for p in c.text_frame.paragraphs for r in p.runs if r.font.size}
    assert len(sizes) == 1 and min(sizes) >= kit.body - 0.05, sizes  # one size in every cell, never under the body size


def test_a_lone_pie_keeps_its_legend_at_the_body_size_or_larger(env):
    from verstka.rendering.compose import Kit

    expenses = ChartSpec(type="pie", title="Расходы", unit="₽", categories=["Продукты, упаковка и списания", "Зарплаты и связанные начисления", "Аренда", "Коммунальные услуги", "Маркетинг", "Другие расходы"],
                         series=[InlineSeries(name="Расходы", values=[315000, 270000, 120000, 25000, 20000, 30000])])
    s = OutlineSlide(id="p", kind=PatternKind.chart, headline="Общие расходы — 780 000 рублей в месяц", content=SlideContent(chart=expenses))
    _, manifest, _ = env
    kit = Kit(manifest, W, H, "FFFFFF")
    for strategy in ("structured", "visual", "compact"):
        slide, _ = _render(env, s, strategy)
        legend = [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] == "Legend"]
        assert legend and min(min(_sizes(sh)) for sh in legend) >= kit.body - 0.05, strategy
        chart = next(sh for sh in slide.shapes if sh.has_chart)
        assert chart.height >= 0.3 * H, strategy  # the circle keeps most of the area's height
    # the visual variant sets the lone pie as a doughnut with the whole in its hole
    slide, _ = _render(env, s, "visual")
    total = [sh for sh in slide.shapes if sh.name.rsplit(" ", 1)[0] == "Total"]
    assert total and "780" in _text(total[0])


def test_a_large_figure_is_not_said_again_in_its_label():
    from verstka.rendering.compose import label_without_figure

    assert label_without_figure("Рентабельность увеличится до 22,4%", "22,4%") == "Рентабельность увеличится"
    assert label_without_figure("Месячная выручка составит 1 138 500 ₽", "1 138 500 ₽") == "Месячная выручка"
    assert label_without_figure("Рост на 112,3%", "112,3%") == "Рост"
    assert label_without_figure("Низкая рентабельность (13,3%) требует оптимизации", "13,3%") == "Низкая рентабельность требует оптимизации"
    mid = "Экономия 22 770 рублей при целевой выручке 1 138 500 рублей"
    assert label_without_figure(mid, "22 770 ₽") == mid
    assert label_without_figure("Операционная прибыль составляет 120 000 рублей, или 13,3% выручки", "120 000 ₽") == "Операционная прибыль — 13,3% выручки"
