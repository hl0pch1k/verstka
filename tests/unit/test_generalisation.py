"""Generalisation defects found on unseen templates (the hand-drawn fixture and the LCT pitch template).

Each test builds the smallest deck that reproduces one defect and names the code it guards.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from verstka.analysis.xmlns import q
from verstka.rendering.deck import DeckBuilder

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _png(path: Path, rgb=(40, 10, 90)) -> Path:
    Image.new("RGB", (64, 36), rgb).save(path)
    return path


def _picture_background_deck(tmp_path: Path) -> Path:
    """One slide whose background is a picture fill; its notes slide takes a lower rId than the image."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "notes first, so the image rId is shifted"
    pic = slide.shapes.add_picture(str(_png(tmp_path / "bg.png")), Emu(0), Emu(0), Emu(914400), Emu(514350))
    rid = pic._element.find(".//" + q("a:blip")).get(f"{{{NS_R}}}embed")
    pic._element.getparent().remove(pic._element)
    bg = etree.fromstring(
        f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="{NS_R}">'
        f'<p:bgPr><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'
    )
    slide._element.cSld.insert(0, bg)
    path = tmp_path / "picture_bg.pptx"
    prs.save(str(path))
    return path


def test_clone_remaps_picture_background(tmp_path):
    """DeckBuilder.clone_slide: r:embed inside <p:bg> must point at the image in the new slide's rels."""
    b = DeckBuilder(_picture_background_deck(tmp_path))
    new = b.clone_slide(1)
    new.notes_slide.notes_text_frame.text = "the renderer adds notes after cloning"
    b.delete_original_slides()
    out = b.save(tmp_path / "out.pptx")
    slide = Presentation(str(out)).slides[0]
    blips = list(slide._element.cSld.find(q("p:bg")).iter(q("a:blip")))
    assert blips
    for blip in blips:
        rel = slide.part.rels.get(blip.get(f"{{{NS_R}}}embed"))
        assert rel is not None and rel.reltype.endswith("/image"), rel.reltype if rel is not None else None


def test_card_surface_is_not_the_first_accent():
    """colors.assign_color_roles: a saturated card fill that barely differs from the ground (1.3:1) is a surface;
    the brand accent is the colour that stands out, even when it covers less area."""
    from verstka.analysis.colors import assign_color_roles
    from verstka.schemas.common import Family
    from verstka.schemas.template import ColorToken, Tokens

    tokens = [
        ColorToken(hex="0B3D2E", weight=900.0, context_weight={"background": 900.0}),
        ColorToken(hex="14523F", weight=60.0, context_weight={"fill": 60.0}),
        ColorToken(hex="F4F1E8", weight=40.0, context_weight={"text": 40.0}),
        ColorToken(hex="FF7A00", weight=6.0, context_weight={"fill": 2.0, "text": 3.0, "line": 1.0}),
    ]
    assign_color_roles(tokens, Family.dark)
    t = Tokens(colors=tokens)
    assert t.accents()[0] == "FF7A00", [(c.hex, c.roles) for c in tokens]
    assert t.color_for("surface") == "14523F"


def test_template_chrome_is_not_a_margin_violation():
    """audit.checks.common.is_chrome_like: a wide footer the template itself places in the margin is chrome,
    not content — the audit must not flag it and the autofix must not move it (hand-drawn deck, 30% wide footer)."""
    from verstka.audit.checks.layout import margin_violation
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import Bbox
    from verstka.schemas.deck_ir import DeckIR, IRElement, IRParagraph, IRRun, IRSlide
    from verstka.schemas.template import ChromeElement, SlideSize, Spacing, TemplateManifest, Tokens

    W, H = 12192000, 6858000
    footer = Bbox(x=int(W * 0.05), y=int(H * 0.925), w=int(W * 0.30), h=int(H * 0.045))
    chrome = ChromeElement(signature="sp|footer", bbox=footer.to_frac(W, H), share=0.8, kind="text", sample_slide=2, text="Северный лес · 2026")
    spacing = Spacing(safe_area=Bbox(x=int(W * 0.05), y=int(H * 0.09), w=int(W * 0.9), h=int(H * 0.74)).to_frac(W, H))
    m = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(chrome=[chrome], spacing=spacing))
    el = IRElement(id="3", type="text", bbox=footer, bbox_frac=footer.to_frac(W, H), paragraphs=[IRParagraph(text="Северный лес · 2026", runs=[IRRun(text="Северный лес · 2026", size_pt=10)])])
    ir = DeckIR(source="x.pptx", slide_w=W, slide_h=H, slides=[IRSlide(index=2, elements=[el])])
    assert margin_violation(AuditContext(ir=ir, manifest=m)) == []


def test_hand_typed_page_numbers_follow_the_new_order():
    """clone.renumber_page_chrome: «03» typed on sample 3 becomes the position of the new slide; other chrome stays."""
    from verstka.rendering.clone import renumber_page_chrome
    from verstka.rendering.textfill import shape_text

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    page = s.shapes.add_textbox(Emu(0), Emu(0), Emu(100000), Emu(100000))
    page.text_frame.text = "03"
    footer = s.shapes.add_textbox(Emu(0), Emu(200000), Emu(100000), Emu(100000))
    footer.text_frame.text = "2026"
    els = {str(page.shape_id): page._element, str(footer.shape_id): footer._element}
    n = renumber_page_chrome(els, list(els), source_slide=3, index=7)
    assert n == 1
    assert shape_text(page._element) == "07" and shape_text(footer._element) == "2026"


def test_title_pattern_without_a_title_slot_is_rejected():
    """matching.scorer: the LCT title sample has its heading baked into the layout picture — nothing to write into."""
    from verstka.matching.scorer import score_pattern
    from verstka.planning.strategies import get_strategy
    from verstka.schemas.common import BboxFrac, Family, PatternKind, SlotRole
    from verstka.schemas.outline import OutlineSlide, SlideContent
    from verstka.schemas.template import Pattern, SlideSize, Slot, TemplateManifest, Tokens

    m = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=12192000, h=6858000), tokens=Tokens())
    baked = Pattern(id="p1", source_slide=1, kind=PatternKind.title, family=Family.dark, slots=[Slot(id="caption_1", role=SlotRole.caption, shape_id="5", bbox=BboxFrac(x=0.05, y=0.85, w=0.3, h=0.05))])
    ok = Pattern(id="p7", source_slide=7, kind=PatternKind.section, family=Family.dark, slots=[Slot(id="title_1", role=SlotRole.title, shape_id="2", bbox=BboxFrac(x=0.05, y=0.4, w=0.6, h=0.15))])
    slide = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Умные напоминания в VK WorkSpace", content=SlideContent())
    st = get_strategy("structured")
    assert score_pattern(slide, baked, m, st).score == 0.0
    assert score_pattern(slide, ok, m, st).score > 0.3


def test_slide_number_placeholder_is_chrome_not_body():
    """analysis.roles: a lone slide-number placeholder («1» at the bottom right) is never a content slot."""
    from verstka.analysis.roles import heuristic_roles
    from verstka.analysis.shapes import ShapeInfo
    from verstka.schemas.common import Bbox, ShapeKind, SlotRole
    from verstka.schemas.template import Typography

    s = ShapeInfo(id="3", name="Номер слайда 2", kind=ShapeKind.sp, bbox=Bbox(x=11460000, y=6380000, w=700000, h=340000), z=1, is_placeholder=True, ph_type="sldNum")
    roles = heuristic_roles([s], [], set(), Typography(), 12192000, 6858000)
    assert roles.get("3") in (None, SlotRole.chrome), roles


def test_icon_sheets_and_palettes_are_reference_slides():
    """analysis.patterns.reference_reason: the LCT «ИКОНКИ» sheets and a palette of HEX swatches are documentation."""
    from verstka.analysis.patterns import reference_reason
    from verstka.analysis.shapes import ShapeInfo, TextInfo
    from verstka.schemas.common import Bbox, ShapeKind, SlotRole

    def pic(i):
        return ShapeInfo(id=str(i), name=f"icon {i}", kind=ShapeKind.pic, bbox=Bbox(x=i * 300000, y=2000000, w=250000, h=250000), z=i)

    icons = [pic(i) for i in range(10, 50)]
    roles = {s.id: SlotRole.icon for s in icons}
    assert reference_reason(icons, roles)
    cards = [pic(i) for i in range(10, 14)]
    assert reference_reason(cards, {s.id: SlotRole.icon for s in cards}) is None


def test_new_bullets_follow_the_schema_order_and_hang():
    """textfill._set_bullet: buChar goes before tabLst/defRPr (CT_TextParagraphProperties order) and a paragraph
    that had no bullet gets a hanging indent, otherwise «•Задачи» is glued to its marker (hand-drawn body text)."""
    from verstka.rendering.textfill import ParagraphSpec, fill_text

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Emu(0), Emu(0), Emu(3000000), Emu(1000000))
    p = tb.text_frame.paragraphs[0]
    p.text = "plain"
    pPr = p._p.get_or_add_pPr()
    etree.SubElement(pPr, q("a:defRPr")).set("sz", "1400")
    fill_text(tb._element, [ParagraphSpec("Задачи живут в 4 инструментах", bullet=True)], size_pt=14)
    pPr = tb._element.find(".//" + q("a:pPr"))
    names = [etree.QName(c).localname for c in pPr]
    assert names.index("buChar") < names.index("defRPr"), names
    assert int(pPr.get("marL", "0")) > 0 and int(pPr.get("indent", "0")) < 0


def test_render_copy_blanks_layout_prompts_only(tmp_path):
    """ingest.render.render_copy: LibreOffice gets layouts without prompt text; slides and the original stay intact."""
    import zipfile

    from verstka.ingest.render import render_copy

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Настоящий заголовок"
    src = tmp_path / "deck.pptx"
    prs.save(str(src))
    before = zipfile.ZipFile(src).read("ppt/slideLayouts/slideLayout2.xml").decode()
    assert "<a:t>Click to edit" in before
    out = render_copy(src, tmp_path / "copy")
    assert out.name == src.name
    z = zipfile.ZipFile(out)
    assert "<a:t>Click to edit" not in z.read("ppt/slideLayouts/slideLayout2.xml").decode()
    assert "Настоящий заголовок" in z.read("ppt/slides/slide1.xml").decode()
    assert zipfile.ZipFile(src).read("ppt/slideLayouts/slideLayout2.xml").decode() == before


def _frames_deck(tmp_path: Path) -> Path:
    """A slide drawn like the LCT template: a heading and three empty white cards on a dark ground."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0x52, 0x09, 0x77)
    head = s.shapes.add_textbox(Emu(600000), Emu(400000), Emu(8000000), Emu(700000))
    head.text_frame.text = "Заголовок"
    head.text_frame.paragraphs[0].runs[0].font.size = Emu(32 * 12700)
    head.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i in range(3):
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(600000 + i * 3700000), Emu(1700000), Emu(3400000), Emu(4200000))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        card.line.fill.background()
    path = tmp_path / "frames.pptx"
    prs.save(str(path))
    return path


def test_empty_frames_become_container_slots_and_take_text(tmp_path):
    """analysis.roles.is_empty_frame + clone container slots: empty white cards receive the items, dark text on white."""
    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import render_clone
    from verstka.rendering.textfill import shape_text
    from verstka.schemas.common import PatternKind
    from verstka.schemas.layout import LayoutSlide
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent, SlideItem

    m = analyze_template(_frames_deck(tmp_path), workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    p = m.patterns[0]
    containers = [s for s in p.slots if s.container]
    assert len(containers) == 3 and p.repeat_groups, [(s.role.value, s.container) for s in p.slots]
    assert all(s.style.color_hex and s.style.color_hex != "FFFFFF" for s in containers)
    oslide = OutlineSlide(id="s1", kind=PatternKind.cards, headline="Три шага", content=SlideContent(items=[SlideItem(title=f"Шаг {i}", text=f"Описание шага {i}") for i in (1, 2)]))
    ws = TemplateWorkspace.open(m.template_id, tmp_path / "ws")
    b = DeckBuilder(ws.source)
    slide, warnings = render_clone(b, LayoutSlide(outline_id="s1", mode="clone", pattern_id=p.id), oslide, p, m, ws, DeckOutline(title="t", slides=[oslide]))
    texts = [shape_text(el) for el in slide.shapes._spTree.iter(q("p:sp"))]
    assert any("Шаг 1" in t and "Описание шага 1" in t for t in texts), texts
    assert any("Шаг 2" in t for t in texts)
    cards = [sh for sh in slide.shapes if sh.shape_type == 1 and sh.width > 2000000]
    assert len(cards) == 2, "the third, empty card is removed"
    assert cards[1].left > cards[0].left + cards[0].width and cards[1].left + cards[1].width > 7400000, "two cards reflow over the span"


def test_contrast_on_a_picture_ground_is_never_recoloured():
    """audit contrast_low: with a picture background the colour under the text is unknown — info, no autofix
    (the LCT template writes white on purple photos; the autofix used to turn it black)."""
    from verstka.audit.checks.template import contrast_low
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import Bbox, Family
    from verstka.schemas.deck_ir import DeckIR, IRElement, IRParagraph, IRRun, IRSlide
    from verstka.schemas.template import ColorToken, SlideSize, TemplateManifest, Tokens

    W, H = 12192000, 6858000
    m = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(colors=[ColorToken(hex="FFFFFF", roles=["background.light"])]))
    box = Bbox(x=int(W * 0.1), y=int(H * 0.4), w=int(W * 0.5), h=int(H * 0.1))
    el = IRElement(id="5", type="text", bbox=box, bbox_frac=box.to_frac(W, H), paragraphs=[IRParagraph(text="Белый текст", runs=[IRRun(text="Белый текст", size_pt=16, color_hex="FFFFFF")])])
    ir = DeckIR(source="x.pptx", slide_w=W, slide_h=H, slides=[IRSlide(index=1, family=Family.light, background_hex=None, elements=[el])])
    issues = contrast_low(AuditContext(ir=ir, manifest=m))
    assert all(i.severity == "info" and i.autofix is None for i in issues), issues


def test_merge_items_fills_every_cell():
    """clone._merge_items: title-only items are spread over all cells (no empty card at the end)."""
    from verstka.rendering.clone import _merge_items
    from verstka.schemas.outline import SlideItem

    items = [SlideItem(title=f"Пункт {i}") for i in range(5)]
    merged = _merge_items(items, 4)
    assert len(merged) == 4
    flat = [b for m in merged for b in (m.bullets or [m.title])]
    assert flat == [f"Пункт {i}" for i in range(5)]


def test_a_panel_hosting_empty_placeholders_is_not_an_empty_card():
    """roles.is_empty_frame: the white half-slide panel of LCT slide 12 holds empty title/object placeholders — it is
    the ground of those slots and must survive rendering."""
    from verstka.analysis.roles import is_empty_frame
    from verstka.analysis.shapes import ShapeInfo, TextInfo
    from verstka.schemas.common import Bbox, ShapeKind

    W, H = 12192000, 6858000
    panel = ShapeInfo(id="30", name="Прямоугольник 29", kind=ShapeKind.sp, bbox=Bbox(x=0, y=int(H * 0.1), w=int(W * 0.5), h=int(H * 0.8)), z=1, fill_hex="FFFFFF", geometry="rect")
    body = ShapeInfo(id="29", name="Объект 28", kind=ShapeKind.sp, bbox=Bbox(x=int(W * 0.04), y=int(H * 0.2), w=int(W * 0.4), h=int(H * 0.6)), z=2, is_placeholder=True, ph_type="obj", text=TextInfo())
    assert not is_empty_frame(panel, [panel, body], W, H)
    lone = ShapeInfo(id="31", name="card", kind=ShapeKind.sp, bbox=Bbox(x=int(W * 0.55), y=int(H * 0.2), w=int(W * 0.3), h=int(H * 0.5)), z=3, fill_hex="FFFFFF", geometry="roundRect")
    assert is_empty_frame(lone, [panel, body, lone], W, H)


def test_a_label_pill_without_its_text_is_removed(tmp_path):
    """clone._remove_orphan_holders: a coloured pill whose tag slot stays empty must not remain as a blank lozenge."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import render_clone
    from verstka.schemas.common import PatternKind
    from verstka.schemas.layout import LayoutSlide
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(600000), Emu(300000), Emu(1800000), Emu(400000))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(0x00, 0x77, 0xFF)
    tag = s.shapes.add_textbox(Emu(650000), Emu(330000), Emu(1700000), Emu(340000))
    tag.text_frame.text = "Проблема"
    tag.text_frame.paragraphs[0].runs[0].font.size = Emu(12 * 12700)
    head = s.shapes.add_textbox(Emu(600000), Emu(900000), Emu(9000000), Emu(900000))
    head.text_frame.text = "Большой заголовок слайда"
    head.text_frame.paragraphs[0].runs[0].font.size = Emu(36 * 12700)
    body = s.shapes.add_textbox(Emu(600000), Emu(2200000), Emu(9000000), Emu(3000000))
    body.text_frame.text = "Первый тезис\\nВторой тезис\\nТретий тезис"
    path = tmp_path / "pill.pptx"
    prs.save(str(path))
    m = analyze_template(path, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    p = m.patterns[0]
    oslide = OutlineSlide(id="s1", kind=PatternKind.bullets, headline="Новый заголовок", content=SlideContent(bullets=["Один", "Два"]))
    ws = TemplateWorkspace.open(m.template_id, tmp_path / "ws")
    slide, _ = render_clone(DeckBuilder(ws.source), LayoutSlide(outline_id="s1", mode="clone", pattern_id=p.id), oslide, p, m, ws, DeckOutline(title="t", slides=[oslide]))
    names = [sh.shape_id for sh in slide.shapes]
    assert pill.shape_id not in names, "the empty pill stayed"
