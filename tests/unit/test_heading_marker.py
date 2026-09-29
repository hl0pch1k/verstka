"""A headline is never a list item: Google Slides templates write their sample titles as numbered paragraphs
(«1. Идея проекта»), and a heading filled into that shape must not keep the «1.» or its hanging indent."""
from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from verstka.analysis.xmlns import q
from verstka.rendering.synth import _fill_heading


def _numbered_box(in_list_style: bool = False):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Emu(600000), Emu(400000), Emu(8000000), Emu(900000))
    tb.text_frame.text = "Образец заголовка"
    tx = tb._element.find(q("p:txBody"))
    if in_list_style:  # the marker comes from the shape's list style, not from the paragraph
        lst = tx.find(q("a:lstStyle"))
        lvl = etree.SubElement(lst, q("a:lvl2pPr"), marL="1209041", indent="-604520")
        etree.SubElement(lvl, q("a:buAutoNum"), type="arabicPeriod")
        ppr = tx.find(q("a:p")).get_or_add_pPr() if hasattr(tx.find(q("a:p")), "get_or_add_pPr") else None
        p = tx.find(q("a:p"))
        ppr = p.find(q("a:pPr"))
        if ppr is None:
            ppr = etree.SubElement(p, q("a:pPr"))
            p.insert(0, ppr)
        ppr.set("lvl", "1")
    else:
        p = tx.find(q("a:p"))
        ppr = etree.Element(q("a:pPr"), lvl="1", marL="1209041", indent="-604520")
        etree.SubElement(ppr, q("a:buAutoNum"), type="arabicPeriod")
        p.insert(0, ppr)
    return tb, prs


def _markers(tb) -> list[str]:
    return [etree.QName(el).localname for el in tb._element.iter() if etree.QName(el).localname in ("buAutoNum", "buChar", "buNone")]


def test_a_numbered_sample_title_loses_its_number_and_hanging_indent():
    tb, prs = _numbered_box()
    _fill_heading(tb, "Оформление съедает больше времени, чем содержание", 28.0, None, (0, 0, 0, 0), prs.slide_width)
    ppr = tb._element.find(".//" + q("a:p")).find(q("a:pPr"))
    assert "buAutoNum" not in _markers(tb) and "buNone" in _markers(tb)
    assert ppr.get("marL") in (None, "0") and ppr.get("indent") in (None, "0")
    assert tb.text_frame.text == "Оформление съедает больше времени, чем содержание"


def test_a_number_from_the_shapes_list_style_is_switched_off_too():
    tb, prs = _numbered_box(in_list_style=True)
    _fill_heading(tb, "Три варианта за 71 секунду", 28.0, None, (0, 0, 0, 0), prs.slide_width)
    ppr = tb._element.find(".//" + q("a:p")).find(q("a:pPr"))
    assert ppr.find(q("a:buNone")) is not None
    assert ppr.get("marL") == "0" and ppr.get("indent") == "0"
