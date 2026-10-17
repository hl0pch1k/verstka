"""Beauty round after the independent critique (qa5, 2026-10-17): figures and words in the middle of a grown tile or
card; the three variants show the same figures in their own forms; tiles whose label words would break become rows;
the template's small decor is covered where the slide's text stands on it; a sample whose text runs past the slide's
edge, or that draws quotation marks for text that is no quote, is turned down; our text is never justified; a heading
on a slot the template's picture draws keeps inside it; a label chip names what the slide is, not its layout."""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.compose import Composer, Kit
from verstka.rendering.deck import DeckBuilder, element_bbox
from verstka.rendering.synth import _kind_label, render_synth
from verstka.schemas.common import Bbox, PatternKind
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _render(env, oslide: OutlineSlide, comp: str, strategy: str = "structured"):
    deck, manifest, ws = env
    outline = DeckOutline(title="T", strategy=strategy, slides=[oslide])
    return render_synth(DeckBuilder(deck), LayoutSlide(outline_id=oslide.id, mode="synth", composition=comp), oslide, manifest, ws, outline)


def _named(slide, prefix: str) -> list:
    return [sh for sh in slide.shapes if sh.name.startswith(prefix)]


FIGS = [NumberCallout(value="1 840 млн ₽", label="Выручка за квартал"), NumberCallout(value="105%", label="Выполнение плана"),
        NumberCallout(value="412 млн ₽", label="EBITDA")]


def test_the_variants_show_the_same_figures_in_their_own_forms(env):
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Выручка 1 840 млн ₽, план выполнен на 105%", content=SlideContent(numbers=list(FIGS)))
    structured, _ = _render(env, o.model_copy(deep=True), "stat_row", "structured")
    visual, wv = _render(env, o.model_copy(deep=True), "stat_row", "visual")
    compact, wc = _render(env, o.model_copy(deep=True), "stat_row", "compact")
    # every variant shows every figure, each over or beside its label
    for slide in (structured, visual, compact):
        assert len(_named(slide, "Figure")) == 3 and len(_named(slide, "Label")) == 3
    # visual: the figure the heading names leads, the others beside it, smaller
    vs = sorted((max(r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), sh.text_frame.text) for sh in _named(visual, "Figure"))
    assert "1 840" in vs[-1][1].replace(" ", " ") and vs[-1][0] >= 1.4 * vs[0][0]
    assert any("visual variant's form" in w for w in wv)
    # compact: rows — the label beside its figure, not under it
    f0, l0 = _named(compact, "Figure")[0], _named(compact, "Label")[0]
    assert l0.left > f0.left + f0.width // 2 and abs(l0.top - f0.top) < f0.height
    assert any("compact variant's form" in w for w in wc)
    # the three slides are not the same picture
    def boxes(slide):
        return sorted((sh.left, sh.top, sh.width) for sh in _named(slide, "Figure"))
    assert len({tuple(boxes(s)) for s in (structured, visual, compact)}) == 3


def test_a_grown_tile_holds_its_figure_and_label_in_its_middle(env):
    deck, manifest, _ws = env
    prs = Presentation(str(deck))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = manifest.slide_size.w, manifest.slide_size.h
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Три цифры", content=SlideContent(numbers=list(FIGS)))
    c = Composer(slide, Kit(manifest, W, H, "FFFFFF"), o, DeckOutline(title="T", strategy="structured", slides=[o]), "structured", manifest)
    c._boost = 2  # the fill pass's level that grows the tiles
    area = Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.88 * W), h=int(0.7 * H))
    n0 = len(c.cv.tree)
    c._kpis_once(area, list(FIGS), [])
    els = list(c.cv.tree)[n0:]
    cards = [element_bbox(e) for e in els if (e.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Card")]
    figs = [element_bbox(e) for e in els if (e.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Figure")]
    labs = [element_bbox(e) for e in els if (e.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Label")]
    if not cards:
        pytest.skip("the template sets its figures without cards")
    card, fig, lab = cards[0], figs[0], labs[0]
    above = fig[1] - card[1]
    # (the label box runs to the card's foot less the same air: measure its text by the figure's line)
    assert above > 0.2 * card[3], "the figure is no longer pressed to the top of a tall tile"


def test_tile_labels_that_would_break_inside_a_word_become_rows(env):
    nums = [NumberCallout(value="100 м²", label="Тренировочный зал"), NumberCallout(value="35 м²", label="Раздевалки и душевые"),
            NumberCallout(value="25 м²", label="Ресепшен и ожидание"), NumberCallout(value="20 м²", label="Подсобные помещения")]
    deck, manifest, _ws = env
    prs = Presentation(str(deck))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = manifest.slide_size.w, manifest.slide_size.h
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Площади", content=SlideContent(numbers=nums))
    c = Composer(slide, Kit(manifest, W, H, "FFFFFF"), o, DeckOutline(title="T", strategy="structured", slides=[o]), "structured", manifest)
    area = Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.24 * W), h=int(0.7 * H))  # a column beside a photo place
    n0 = len(c.cv.tree)
    c._kpis_once(area, nums, [])
    from verstka.rendering.fonts import text_width_pt

    for el in list(c.cv.tree)[n0:]:
        if not (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Label"):
            continue
        b = element_bbox(el)
        size = max(int(r.get("sz")) / 100 for r in el.iter(q("a:rPr")) if r.get("sz"))
        words = "".join(t.text or "" for t in el.iter(q("a:t"))).replace("\u00a0", " ").split()
        assert all(text_width_pt(w, c.kit.font, size, False) <= b[2] / 12700 for w in words), words


# ------------------------------------------------------------------------------------------------- small decor


def _decor_deck(path: Path, bg: str = "FFFFFF") -> Path:
    """A layout with a small cross at the middle of the slide, the slide a text over it."""
    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    layout = prs.slide_layouts[6]
    scratch = prs.slides.add_slide(layout)
    mark = scratch.shapes.add_shape(MSO_SHAPE.CROSS, Emu(int(0.4 * W)), Emu(int(0.45 * H)), Emu(int(0.03 * W)), Emu(int(0.04 * H)))
    mark.fill.solid()
    mark.fill.fore_color.rgb = RGBColor(0xF2, 0xC2, 0x00)
    el = mark._element
    el.getparent().remove(el)
    layout.shapes._spTree.insert(2, el)
    prs.save(str(path))
    return path


def test_a_small_mark_of_the_layout_under_the_text_is_covered_in_the_ground(tmp_path):
    from verstka.rendering.decor import clear_small_decor, marks_under_text

    prs = Presentation(str(_decor_deck(tmp_path / "d.pptx")))
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Emu(int(0.1 * W)), Emu(int(0.43 * H)), Emu(int(0.8 * W)), Emu(int(0.1 * H)))
    tb.text_frame.text = "Клиентов в 15 минутах пешком"
    assert marks_under_text(s)
    assert clear_small_decor(s) == 1
    patch = next(sh for sh in s.shapes if sh.name.startswith("Decor cover"))
    assert patch.fill.fore_color.rgb == RGBColor(0xFF, 0xFF, 0xFF)
    tree = s._element.cSld.find(q("p:spTree"))
    assert tree.index(patch._element) < tree.index(tb._element), "the patch lies under the text"
    # a slide whose text stands clear of the mark is left alone
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Emu(int(0.1 * W)), Emu(int(0.1 * H)), Emu(int(0.8 * W)), Emu(int(0.1 * H))).text_frame.text = "Заголовок"
    assert clear_small_decor(s2) == 0


# ------------------------------------------------------------------------------------------------- clone gates


def test_text_past_the_edge_and_quote_marks_turn_a_sample_down(tmp_path):
    from verstka.rendering.clone_check import _quote_marks_drawn, text_past_the_edge
    from verstka.schemas.common import BboxFrac, Family
    from verstka.schemas.template import Pattern, SlideSize, TemplateManifest, Tokens

    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Emu(int(0.3 * W)), Emu(int(0.02 * H)), Emu(int(0.8 * W)), Emu(int(0.08 * H)))
    tb.text_frame.text = "Откат не сработал из-за несовместимости форматов кэша"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    tb.rotation = -8.0
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=int(W), h=int(H)), tokens=Tokens(), patterns=[], n_slides=1)
    pat = Pattern(id="p1", source_slide=1, kind=PatternKind.bullets, family=Family.light, slots=[])
    assert "past the slide's edge" in (text_past_the_edge(s, pat, man) or "")
    tb.left, tb.width, tb.rotation = Emu(int(0.1 * W)), Emu(int(0.7 * W)), 0.0
    tb.top = Emu(int(0.2 * H))
    assert text_past_the_edge(s, pat, man) is None
    q_box = s.shapes.add_textbox(Emu(int(0.05 * W)), Emu(int(0.5 * H)), Emu(int(0.1 * W)), Emu(int(0.1 * H)))
    q_box.text_frame.text = "“"
    assert _quote_marks_drawn(s)


def test_our_text_is_never_justified(env):
    from types import SimpleNamespace

    from verstka.rendering.clone import _no_justify

    deck, _m, _ws = env
    prs = Presentation(str(deck))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(0, 0, Emu(3000000), Emu(800000))
    tb.text_frame.text = "Организационная причина: инструкция по ручному исправлению устарела"
    ppr = tb.text_frame.paragraphs[0]._p.get_or_add_pPr()
    ppr.set("algn", "just")
    nv = tb._element.find(".//" + q("p:cNvPr"))
    ctx = SimpleNamespace(filled={nv.get("id")}, els={nv.get("id"): tb._element}, slide=s)
    from verstka.rendering import clone as cl

    orig = cl._style_levels
    cl._style_levels = lambda ctx_, el: []
    try:
        _no_justify(ctx)
    finally:
        cl._style_levels = orig
    assert tb.text_frame.paragraphs[0]._p.find(q("a:pPr")).get("algn") == "l"


# --------------------------------------------------------------------------------------------------- label chips


def test_a_label_chip_names_what_the_slide_is():
    ask = OutlineSlide(id="a", kind=PatternKind.stat_row, headline="Утвердить бюджет 96 млн ₽ и программу рекомендаций", content=SlideContent(numbers=list(FIGS[:2])))
    assert _kind_label(ask) == "Решение"
    funnel = OutlineSlide(id="f", kind=PatternKind.chart, headline="1 000 обращений приведут к 175 покупателям",
                          content=SlideContent(chart=ChartSpec(type="bar", categories=["Обращения", "Записи", "Посещения", "Покупатели"], series=[{"name": "шт.", "values": [1000, 500, 350, 175]}])))
    assert _kind_label(funnel) == "Данные"
    months = OutlineSlide(id="m", kind=PatternKind.chart, headline="Выручка растёт",
                          content=SlideContent(chart=ChartSpec(type="column", categories=["2023", "2024", "2025", "2026"], series=[{"name": "млн ₽", "values": [1, 2, 3, 4]}])))
    assert _kind_label(months) == "Динамика"
    hist = OutlineSlide(id="h", kind=PatternKind.chart, headline="Распределение", content=SlideContent(chart=ChartSpec(type="histogram", categories=["0–20", "20–40"], series=[{"name": "n", "values": [3, 4]}])))
    assert _kind_label(hist) == "Распределение"


def test_a_list_moved_to_the_middle_keeps_its_rules_between_its_rows(env):
    o = OutlineSlide(id="b", kind=PatternKind.bullets, headline="Найм", content=SlideContent(bullets=["Аналитик закрывается за 41 день", "Продавец — за 23 дня"]))
    slide, warnings = _render(env, o, "bullets", "structured")
    assert any("optical middle" in w for w in warnings)
    items = sorted(sh.top for sh in _named(slide, "Item"))
    rules = [sh.top for sh in _named(slide, "Line")]
    assert rules and all(items[0] < r < items[-1] for r in rules), (items, rules)


def _composer(env, o: OutlineSlide, strategy: str = "structured"):
    deck, manifest, _ws = env
    prs = Presentation(str(deck))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = manifest.slide_size.w, manifest.slide_size.h
    return Composer(slide, Kit(manifest, W, H, "FFFFFF"), o, DeckOutline(title="T", strategy=strategy, slides=[o]), strategy, manifest), W, H


def test_a_pie_legend_of_last_resort_keeps_the_readable_size_and_the_least_circle(env):
    """C, qa5: the last resort set the legend at 1.7 % of the slide height (7.5 pt under a 95 pt circle) and shrank the
    circle to 0.3 of its frame, past the pie's least size."""
    from verstka.rendering.charts import readable_floor

    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Доли", content=SlideContent())
    c, W, H = _composer(env, o)
    floor = float(readable_floor(c.slide, c.kit.manifest.tokens.typography))
    box = Bbox(x=int(0.1 * W), y=int(0.3 * H), w=int(0.35 * W), h=int(0.35 * H))
    seen = []

    def never(size, cap, cols=1):  # a legend that fits nowhere
        seen.append(size)
        return dict(size=size, h=box.h * 2, w=box.w * 2, pct_w=0)

    lg, d, below = c._pie_legend_last(box, never, list(range(9)))
    assert min(seen) >= 0.85 * floor - 0.05 and lg["size"] >= 0.85 * floor - 0.05  # a step within 15 % under it reads
    assert d >= min(c._pie_min_d(box), int(0.85 * box.w))

    def narrow(size, cap, cols=1):  # a legend that fits beside a circle of the least size only
        return dict(size=size, h=int(box.h * 0.9), w=min(cap + 1, box.w - c._pie_min_d(box) - 2 * c.kit.gap), pct_w=0)

    lg, d, below = c._pie_legend_last(box, narrow, list(range(9)))
    assert not below and d >= c._pie_min_d(box) and lg["size"] >= 0.85 * floor - 0.05


def test_a_table_that_reads_larger_at_full_height_takes_its_conclusion_beside_it(env, monkeypatch):
    """C, qa5: a ten-row table squeezed over its conclusion's strip and footnote (219 pt of 446) — the table keeps the
    area's whole height, its conclusion stands beside it in a column behind an accent bar."""
    import verstka.rendering.tables as tables

    rows = [[f"Направление {i}", str(100 + i), str(50 + i), str(150 + 2 * i)] for i in range(10)]
    from verstka.schemas.outline import TableData

    t = TableData(columns=["Направление", "Бюджет", "Платное", "Всего"], rows=rows)
    take = "Платный приём растёт быстрее бюджетного: две трети прироста дали платные места"
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Приём по направлениям", takeaway=take, footnote="Источник: приём.xlsx",
                     content=SlideContent(table=t))
    c, W, H = _composer(env, o)
    c._comp = "table"
    area = Bbox(x=int(0.06 * W), y=int(0.22 * H), w=int(0.88 * W), h=int(0.7 * H))
    # at the body size the table needs more than the strip leaves, less than the whole height at 0.68 of the width
    monkeypatch.setattr(tables, "table_need_h", lambda tb, w, font, size, **kw: int(area.h * (0.8 if size >= c.kit.body - 0.05 else 0.5)))
    assert c._table_takes_side(area, take, True)
    monkeypatch.setattr(tables, "table_need_h", lambda tb, w, font, size, **kw: int(area.h * 0.3))
    assert not c._table_takes_side(area, take, True)  # it holds over the strip: the strip stays
    c.strategy = "compact"
    assert not c._table_takes_side(area, take, True)  # the compact variant leads with its conclusion instead
    # drawn: the conclusion in the column beside the table, no strip under it
    monkeypatch.undo()
    monkeypatch.setattr(Composer, "_table_takes_side", lambda self, area, take, foot: getattr(self, "_comp", None) == "table")
    o2 = o.model_copy(deep=True)
    slide, warnings = _render(env, o2, "table", "structured")
    tbl = next(sh for sh in slide.shapes if sh.has_table)
    concl = _named(slide, "Conclusion")
    assert concl and all(sh.left >= tbl.left + tbl.width for sh in concl)
    assert not _named(slide, "Conclusion strip")
    assert any("beside the table" in w for w in warnings)


def test_a_pie_legend_of_amounts_is_headed_by_its_unit_not_a_service_line(env):
    """qa5: «сумма · доля от общей суммы» over the legend read as a note to the designer — the line names the unit of
    the amounts and the share of the total."""
    from verstka.schemas.outline import InlineSeries

    pie = ChartSpec(type="pie", title="Расходы за месяц", unit="₽", categories=["Аренда", "Зарплаты", "Реклама", "Прочее"],
                    series=[InlineSeries(name="Расходы", values=[150000, 260000, 40000, 20000])])
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Зарплаты — больше половины расходов", content=SlideContent(chart=pie))
    slide, _ = _render(env, o, "chart")
    text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    assert "сумма ·" not in text and "общей суммы" not in text
    head = _named(slide, "Legend title")
    assert head and "₽ и % от итога" in head[0].text_frame.text.replace("\u00a0", " ")


def test_a_conclusion_strip_never_lies_across_a_row_of_cards(tmp_path):
    """qa6 (VK + база знаний, сл. 6): the free stretch between two rows of white cards took the conclusion strip
    across the second row — a card drawn by its shadow alone is content all the same."""
    from verstka.rendering.clone_notes import _lies_on_content

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = prs.slide_width, prs.slide_height
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(0.05 * W)), Emu(int(0.56 * H)), Emu(int(0.4 * W)), Emu(int(0.3 * H)))
    card.fill.background()
    card.line.fill.background()
    across = Bbox(x=int(0.05 * W), y=int(0.53 * H), w=int(0.9 * W), h=int(0.1 * H))
    under = Bbox(x=int(0.05 * W), y=int(0.88 * H), w=int(0.9 * W), h=int(0.08 * H))
    assert _lies_on_content(slide, across, W, H)
    assert not _lies_on_content(slide, under, W, H)
