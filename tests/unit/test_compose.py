"""The composer lays content slides out from the template's design system: one grid, one type scale, the template's
own card; figures large; Russian typesetting; one heading treatment per deck."""

from __future__ import annotations

import pytest
from pptx.util import Pt

from verstka.analysis.manifest import analyze_template
from verstka.audit.checks.template import is_figure
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.compose import Kit, _series_takeaway, distinct_label, highlight_runs
from verstka.rendering.deck import DeckBuilder
from verstka.rendering.fonts import text_width_pt, wrap_lines
from verstka.rendering.synth import deck_style, render_synth
from verstka.ru import NBSP, WJ, typeset
from verstka.schemas.common import EMU_PER_PT, PatternKind
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, NumberCallout, OutlineSlide, Series, SlideContent, SlideItem

W, H = 12192000, 6858000


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _render(env, oslide: OutlineSlide, comp: str, strategy: str = "structured"):
    deck, manifest, ws = env
    b = DeckBuilder(deck)
    outline = DeckOutline(title="T", strategy=strategy, slides=[oslide])
    slide, warnings = render_synth(b, LayoutSlide(outline_id=oslide.id, mode="synth", composition=comp), oslide, manifest, ws, outline)
    return slide, warnings


def _named(slide, prefix: str):
    return [sh for sh in slide.shapes if sh.name.startswith(prefix)]


def test_kit_sizes_come_from_the_template_scale_with_readable_minimums(env):
    _, manifest, _ = env
    kit = Kit(manifest, W, H, "FFFFFF")
    sizes = set(kit.sizes)
    hpt = H / EMU_PER_PT
    assert kit.body >= 0.024 * hpt and kit.body in sizes | {round(kit.body * 2) / 2}
    assert kit.small < kit.body < kit.h3 <= kit.h2 and kit.lead > kit.body
    # figures may go past the largest text size of the template (the display step), never below the heading step
    figs = kit.figure_sizes(0.2 * hpt, kit.h3)
    # (a sparse template's derived ladder already reaches the display step: past the template's own sizes)
    own = sizes - {round(float(x), 2) for x in (manifest.tokens.typography.derived_sizes or [])}
    assert figs == sorted(figs, reverse=True) and figs[0] > max(own) and figs[-1] >= kit.h3
    assert kit.figure_cap(1, False, "structured") > kit.figure_cap(4, False, "structured") > kit.figure_cap(4, True, "compact") * 0.99


def test_cards_share_one_grid(env):
    items = [SlideItem(title=f"Сценарий {i}", text="Короткое описание сценария в одну-две строки") for i in range(1, 5)]
    slide, warnings = _render(env, OutlineSlide(id="c", kind=PatternKind.cards, headline="Четыре сценария", content=SlideContent(items=items)), "cards")
    cards = sorted(_named(slide, "Card ") , key=lambda s: s.left)
    cards = [c for c in cards if not c.has_text_frame or not c.text_frame.text.strip()]
    assert len(cards) == 4, [c.name for c in slide.shapes]
    widths = {c.width for c in cards}
    heights = {c.height for c in cards}
    tops = sorted({c.top for c in cards})
    assert max(widths) - min(widths) <= 2 and len(heights) == 1 and len(tops) in (1, 2)  # a row of four or a 2×2 block
    row = sorted((c for c in cards if c.top == tops[0]), key=lambda c: c.left)
    gaps = {row[i + 1].left - (row[i].left + row[i].width) for i in range(len(row) - 1)}
    assert max(gaps) - min(gaps) <= 2  # one gutter
    if len(tops) == 2:
        assert tops[1] - (tops[0] + cards[0].height) > 0  # rows do not touch
    # the grid ends where the template's safe area ends, never past the mirror of the left margin
    _, manifest, _ = env
    left, right_edge = row[0].left, row[-1].left + row[-1].width
    assert right_edge == pytest.approx(min(int(manifest.tokens.spacing.safe_area.x2 * W), W - left), abs=0.005 * W)
    title = next(sh for sh in slide.shapes if sh.is_placeholder)
    assert min(c.top for c in cards) >= title.top + title.height


def test_kpi_figures_share_one_size_and_never_wrap(env):
    nums = [NumberCallout(value=v, label=l) for v, l in [("12 400", "участников пилота"), ("+34%", "задач завершены в срок"), ("2,1 ч", "экономии в неделю")]]
    slide, _ = _render(env, OutlineSlide(id="s", kind=PatternKind.stat_row, headline="Пилот подтвердил эффект", content=SlideContent(numbers=nums)), "stat_row")
    figs = _named(slide, "Figure")
    assert len(figs) == 3
    # one size for the figures of a row (a unit after a figure is set smaller on its baseline)
    sizes = {f.text_frame.paragraphs[0].runs[0].font.size.pt for f in figs}
    assert len(sizes) == 1
    size = sizes.pop()
    assert size >= 36
    for f in figs:
        text = f.text_frame.text
        lines = wrap_lines(text, "Play", size, False, f.width / EMU_PER_PT)
        assert len(lines) == 1, text
        assert NBSP in text or " " not in text  # «12 400» is one figure
        runs = f.text_frame.paragraphs[0].runs
        assert all(r.font.size.pt <= size for r in runs)


def test_statement_highlights_figures(env):
    runs = highlight_runs("Утвердить бюджет 27 млн ₽ и выделить 4 GPU-сервера до конца года", 28, "000000", "FF0000")
    accented = [r.text for r in runs if r.color == "FF0000"]
    assert any("27" in t and "₽" in t for t in accented)
    slide, _ = _render(env, OutlineSlide(id="a", kind=PatternKind.bullets, headline="Что просим", content=SlideContent(paragraphs=["Утвердить бюджет 27 млн ₽ на масштабирование и выделить 4 GPU-сервера до конца года"])), "bullets")
    st = _named(slide, "Statement")
    assert st and st[0].text_frame.paragraphs[0].runs[0].font.size >= Pt(24)


def test_russian_typesetting():
    t = typeset("В часы наплыва растёт нагрузка на GPU — резервируем 1 500 серверов за 27 млн ₽")
    assert "на" + NBSP + "GPU" in t and NBSP + "— " in t and "1" + NBSP + "500" in t and "27" + NBSP + "млн" + NBSP + "₽" in t
    assert "500" + NBSP + "серверов" in t  # a figure stays with its noun
    # a line never breaks at a no-break space
    assert all(not ln.startswith("—") for ln in wrap_lines(t, "Play", 20, False, 180))
    # a short compound keeps its hyphen (an invisible word joiner: Play has no no-break hyphen); a long one may break
    t = typeset("выделить 4 GPU-сервера для контакт-центра")
    assert "4" + NBSP + "GPU-" + WJ + "сервера" in t and "контакт-" + WJ + "центра" in t
    assert WJ not in typeset("высоконагруженная-распределённая-архитектура")
    assert text_width_pt("GPU-" + WJ + "сервера", "Play", 20) == text_width_pt("GPU-сервера", "Play", 20)


def test_figure_labels_do_not_repeat_the_heading():
    assert distinct_label("6,5", "оператор тратит в среднем 6,5 минуты", "Оператор тратит в среднем 6,5 минуты на одно обращение") == "минуты на одно обращение"
    assert distinct_label("40%", "этого времени уходит на поиск ответа", "Проблема") == "этого времени уходит на поиск ответа"


def test_display_figures_pass_the_scale_check():
    assert is_figure("6,5") and is_figure("31% → 12%") and is_figure("×4,8") and is_figure("27 млн")
    assert not is_figure("Итого") and not is_figure("оператор тратит 6,5 минуты")


def test_one_heading_size_for_the_deck(env):
    deck, manifest, ws = env
    b = DeckBuilder(deck)
    slides = [
        OutlineSlide(id="a", kind=PatternKind.bullets, headline="Коротко", content=SlideContent(bullets=["Первый тезис", "Второй тезис"])),
        OutlineSlide(id="b", kind=PatternKind.bullets, headline="Заголовок подлиннее, но всё ещё в две строки на широком слайде", content=SlideContent(bullets=["Тезис"])),
    ]
    outline = DeckOutline(title="T", slides=slides)
    ds = deck_style(b, manifest, outline)
    assert ds is deck_style(b, manifest, outline)  # decided once per deck
    sizes = []
    for s in slides:
        slide, _ = render_synth(b, LayoutSlide(outline_id=s.id, mode="synth", composition="bullets"), s, manifest, ws, outline)
        title = next(sh for sh in slide.shapes if sh.is_placeholder)
        sizes.append(title.text_frame.paragraphs[0].runs[0].font.size.pt)
    assert sizes[0] == sizes[1] == ds.head_size


def test_a_callout_stays_a_step_below_the_heading(env):
    # «Что просим» + one sentence: the ask is set larger than running text, never larger than the heading over it
    slide, _ = _render(env, OutlineSlide(id="a", kind=PatternKind.bullets, headline="Что просим", content=SlideContent(paragraphs=["Утвердить бюджет 27 млн ₽ на масштабирование и выделить 4 GPU-сервера до конца года"])), "bullets")
    head = next(sh for sh in slide.shapes if sh.is_placeholder).text_frame.paragraphs[0].runs[0].font.size.pt
    st = _named(slide, "Statement")[0].text_frame.paragraphs[0].runs[0].font.size.pt
    assert st < head


def test_the_chart_takeaway_is_the_change_never_the_last_bar_again():
    outline = DeckOutline(title="t", series=[Series(id="cc", name="Обращения", categories=["Июль", "Август", "Сентябрь"], values=[12, 31, 58], unit="тыс.")])
    spec = ChartSpec(type="column", series_ids=["cc"])
    assert _series_takeaway(spec, outline, "Обращения растут") == ("×4,8", "рост за период июль → сентябрь")
    # the heading already states the change (a multiple, a percent): no second, differently computed measure of it
    assert _series_takeaway(spec, outline, "С 12 до 58, рост в 5 раз") is None
    assert _series_takeaway(spec, outline, "Обращения выросли на 383%") is None
    assert _series_takeaway(spec, outline, "Рост в 5 раз: +46 тыс. обращений") is None  # nothing left to add
