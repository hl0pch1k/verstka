"""«Улучшить готовую презентацию» (verstka/ingest/deck_text.py): a finished deck read back as the author's «Слайд N»
text — its titles, lists, tables, charts with their data, pictures kept for the improved deck, a KPI tile read as one
line; what every slide repeats (a footer, a logo, the page number) left out."""

from __future__ import annotations

import io
import json
import re

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Pt

from verstka.ingest.deck_text import SOURCE_IMAGES, deck_text, is_source_hint, keep_pictures, source_records
from verstka.planning.brief_structure import read_structure
from verstka.planning.faithful import charts_of, expand_slides, note_table, notes_of, plan_deck, read_slide

W, H = 12192000, 6858000


def _png(color=(30, 120, 200), size=(320, 200), flat: bool = False) -> io.BytesIO:
    """A picture: a gradient (a photo has detail); `flat` — one colour, as a template draws a card's ground."""
    im = Image.new("RGB", size, color)
    if not flat:
        for x in range(size[0]):
            for y in range(0, size[1], 4):
                im.putpixel((x, y), ((color[0] + x) % 256, (color[1] + y) % 256, color[2]))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def _box(slide, x, y, w, h, text, size=18, font=None):
    tb = slide.shapes.add_textbox(Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h)))
    tf = tb.text_frame
    lines = text if isinstance(text, list) else [text]
    tf.text = lines[0]
    for ln in lines[1:]:
        tf.add_paragraph().text = ln
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            if font:
                r.font.name = font
    return tb


def _deck(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    blank = prs.slide_layouts[6]
    logo = _png((200, 30, 30), (60, 60))

    def footer(s):
        _box(s, 0.05, 0.93, 0.3, 0.05, "ООО «Ромашка» · конфиденциально", 10)
        logo.seek(0)
        s.shapes.add_picture(logo, Emu(int(W * 0.9)), Emu(int(H * 0.9)), Emu(int(W * 0.05)), Emu(int(H * 0.08)))

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Итоги квартала"
    s1.placeholders[1].text = "Отдел продаж, октябрь 2026"
    footer(s1)

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Что сделали"
    body = s2.placeholders[1].text_frame
    body.text = "Запустили онлайн-заказ"
    body.add_paragraph().text = "Открыли две новые точки"
    body.add_paragraph().text = "Обучили 40 сотрудников"
    footer(s2)

    s3 = prs.slides.add_slide(blank)
    _box(s3, 0.05, 0.05, 0.9, 0.12, "Выручка выросла на 34%", 32)
    rows = [["Месяц", "Выручка, млн ₽"], ["Июль", "1,2"], ["Август", "1,5"], ["Сентябрь", "1,9"]]
    tbl = s3.shapes.add_table(len(rows), 2, Emu(int(W * 0.05)), Emu(int(H * 0.25)), Emu(int(W * 0.5)), Emu(int(H * 0.4))).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
    footer(s3)

    s4 = prs.slides.add_slide(blank)
    _box(s4, 0.05, 0.05, 0.9, 0.12, "Доли каналов продаж", 32)
    cd = CategoryChartData()
    cd.categories = ["Онлайн", "Офлайн", "Партнёры"]
    cd.add_series("Доля, %", (55, 35, 10))
    s4.shapes.add_chart(XL_CHART_TYPE.PIE, Emu(int(W * 0.1)), Emu(int(H * 0.2)), Emu(int(W * 0.6)), Emu(int(H * 0.6)), cd)
    footer(s4)

    s5 = prs.slides.add_slide(blank)
    _box(s5, 0.05, 0.05, 0.9, 0.12, "Новая точка на Тверской", 32)
    s5.shapes.add_picture(_png(), Emu(int(W * 0.5)), Emu(int(H * 0.2)), Emu(int(W * 0.45)), Emu(int(H * 0.6)))
    # a flat panel drawn as a picture (a card's ground of the template): decoration, never a picture place
    s4.shapes.add_picture(_png((235, 238, 244), flat=True), Emu(int(W * 0.72)), Emu(int(H * 0.2)), Emu(int(W * 0.25)), Emu(int(H * 0.5)))
    _box(s5, 0.05, 0.3, 0.4, 0.1, "300 ₽", 40)
    _box(s5, 0.05, 0.41, 0.4, 0.08, "Средний чек", 16)
    _box(s5, 0.05, 0.6, 0.4, 0.2, ['{', '  "zone": "bed_1",', '  "tank_level_pct": 72', '}'], 12, font="Courier New")
    footer(s5)
    path = tmp_path / "old.pptx"
    prs.save(str(path))
    return path


def test_a_deck_reads_back_as_the_authors_slides(tmp_path):
    d = deck_text(_deck(tmp_path))
    assert d.summary() == {"title": "Итоги квартала", "slides": 5, "tables": 1, "charts": 1, "pictures": 1}
    b = d.brief
    # the cover in the brief's own cover format: «Итоги …» is no first content slide under a cover of its own
    assert b.startswith("Слайд 1. Титульный\nНазвание: «Итоги квартала».\nПодзаголовок: «Отдел продаж, октябрь 2026».")
    assert "Слайд 2. Что сделали\n\n- Запустили онлайн-заказ\n- Открыли две новые точки\n- Обучили 40 сотрудников" in b
    assert "| Месяц | Выручка, млн ₽ |\n|---|---|\n| Июль | 1,2 |" in b
    assert "[Круговая диаграмма: Доля, % — Онлайн: 55; Офлайн: 35; Партнёры: 10]" in b
    assert re.search(r"\[Место под изображение: Изображение со слайда 5 исходной презентации · [0-9a-f]{6}\]", b)
    assert "Средний чек — 300 ₽" in b  # a KPI tile is one line: the figure keeps what it measures
    assert "```\n{\n  \"zone\": \"bed_1\"," in b  # code stays code
    # what every slide repeats is the template's: never the author's text, never a picture place
    assert "Ромашка" not in b and b.count("[Место под изображение") == 1


def test_the_read_back_text_is_a_dictated_deck_with_its_chart_and_picture(tmp_path):
    d = deck_text(_deck(tmp_path))
    st = read_structure(d.brief)
    assert [sp.number for sp in st.specs] == [1, 2, 3, 4, 5] and st.specs[1].title == "Что сделали"
    assert st.title == "Итоги квартала" and st.subtitle == "Отдел продаж, октябрь 2026"
    plan = plan_deck(st.specs, None, ("structured",), title=st.title)
    assert plan.readings[0].role == "cover" and plan.readings[0].title == "Итоги квартала" and plan.title == "Итоги квартала"


def test_the_decks_pictures_go_back_to_their_places_of_the_improved_deck(tmp_path):
    d = deck_text(_deck(tmp_path))
    ws = tmp_path / "ws"
    index = keep_pictures(d, ws / "source_images")
    (hint,) = index
    assert re.fullmatch(r"Изображение со слайда 5 исходной презентации · [0-9a-f]{6}", hint) and index[hint] == "picture-01.png"
    assert is_source_hint(hint) and not is_source_hint("Вид теплицы сверху")
    assert json.loads((ws / SOURCE_IMAGES).read_text(encoding="utf-8")) == index
    pics = tmp_path / "gen" / "images"
    recs = source_records(ws, pics)
    assert len(recs) == 1 and recs[0]["status"] == "generated" and recs[0]["model"] == "original"
    assert recs[0]["hint"] == hint and (pics / recs[0]["file"]).is_file()
    assert source_records(tmp_path / "nothing", pics) == []


# ---------------------------------------------------------------------------- what a real deck holds (the review of
# 102 decks, 2026-10-07): each rule read on a deck made here, never tuned to one template


A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    return prs


def _blank(prs, title: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[5] if title else prs.slide_layouts[6])
    if title:
        s.shapes.title.text = title
    return s


def _save(prs, tmp_path, name="deck.pptx"):
    path = tmp_path / name
    prs.save(str(path))
    return path


def _section(brief: str, n: int) -> str:
    m = re.search(rf"(?ms)^Слайд {n}\..*?(?=^Слайд \d+\.|\Z)", brief)
    return m.group(0) if m else ""


def _chart(slide, kind, cats, series, fmt=None, x=0.1, y=0.2, w=0.6, h=0.6):
    cd = CategoryChartData(number_format=fmt) if fmt else CategoryChartData()
    cd.categories = cats
    for name, values in series:
        cd.add_series(name, values)
    return slide.shapes.add_chart(kind, Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h)), cd)


def test_a_chart_of_several_series_keeps_every_series_as_a_table_and_is_drawn_from_it(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги года")
    s = _blank(prs, "Выручка по городам")
    _chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ["2023", "2024", "2025"], [("Москва", (10, 12, 15)), ("СПб", (5, 6, 7)), ("Казань", (2, 3, 4))])
    b = deck_text(_save(prs, tmp_path)).brief
    sec = _section(b, 2)
    assert "[Столбчатая диаграмма]\n\n|  | Москва | СПб | Казань |\n|---|---|---|---|\n| 2023 | 10 | 5 | 2 |" in sec
    assert "| 2025 | 15 | 7 | 4 |" in sec  # every series, every category: nothing dropped by «series[:2]»
    st = read_structure(b)
    r = read_slide(2, st.specs[1].title, st.specs[1].text)
    (chart,) = charts_of(r)
    assert chart.type == "column" and chart.categories == ["2023", "2024", "2025"]
    assert [(x.name, x.values) for x in chart.series] == [("Москва", [10, 12, 15]), ("СПб", [5, 6, 7]), ("Казань", [2, 3, 4])]
    assert note_table(r) is not None  # (the table stays the chart's data: shown as a table when no chart can be drawn)
    plan = plan_deck(st.specs, None, ("structured",), title=st.title)
    sl = plan.decks["structured"][1]
    assert sl.kind.value == "chart" and len(sl.content.chart.series) == 3 and sl.content.table is None
    assert len(expand_slides(plan.decks["structured"], plan.readings)) == 2  # the table is not shown again on a slide of its own


def test_chart_values_keep_their_number_format_and_any_chart_type_is_read(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Доли каналов")
    _chart(s, XL_CHART_TYPE.PIE, ["Онлайн", "Офлайн", "Партнёры"], [("Доля", (0.45, 0.35, 0.2))], fmt="0%")
    s = _blank(prs, "Рост конверсии")
    _chart(s, XL_CHART_TYPE.BAR_CLUSTERED, ["Январь", "Февраль", "Март"], [("Конверсия", (0.012, 0.018, 0.025))], fmt="0.0%")
    s = _blank(prs, "Выручка в 3D")
    gf = _chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ["2023", "2024"], [("Выручка", (5, 7))])
    gf.chart.part._element.find(f".//{C_NS}barChart").tag = f"{C_NS}bar3DChart"  # python-pptx cannot read a 3D chart's plots
    s = _blank(prs, "Точность")
    _chart(s, XL_CHART_TYPE.LINE, ["Q1", "Q2", "Q3"], [("Доля брака", (0.0012, 0.00105, 0.0009))])
    b = deck_text(_save(prs, tmp_path)).brief
    assert "[Круговая диаграмма: Доля — Онлайн: 45%; Офлайн: 35%; Партнёры: 20%]" in b
    assert "[Горизонтальная диаграмма: Конверсия — Январь: 1,2%; Февраль: 1,8%; Март: 2,5%]" in b  # never «0,01»
    assert "[Столбчатая диаграмма: Выручка — 2023: 5; 2024: 7]" in b
    assert "Q1: 0,0012; Q2: 0,00105; Q3: 0,0009" in b  # four significant digits of a small value, never «0»


def test_a_list_stays_a_list_and_only_code_is_code(tmp_path):
    prs = _prs()
    _blank(prs, "Стратегия")
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Задачи на год"
    tf = s.placeholders[1].text_frame  # a body placeholder: its paragraphs are bulleted by the master
    tf.text = "снизить затраты на 10%;"
    for t in ["увеличить выручку;", "выйти в два новых региона."]:
        tf.add_paragraph().text = t
    s = _blank(prs)
    _box(s, 0.05, 0.05, 0.5, 0.06, "01 / Проблема", 14, font="JetBrains Mono")
    _box(s, 0.05, 0.13, 0.9, 0.12, "Клиенты уходят к конкурентам", 36)
    box = _box(s, 0.05, 0.4, 0.4, 0.3, ["снижение затрат", "рост выручки", "новые клиенты"], 18)
    for p in box.text_frame.paragraphs:  # bullets of the box's own
        etree.SubElement(p._p.get_or_add_pPr(), f"{A_NS}buChar").set("char", "•")
    _box(s, 0.5, 0.4, 0.4, 0.2, ["Продукты, упаковка", "и списания"], 16)  # a label the box wrapped: one line
    _box(s, 0.05, 0.75, 0.6, 0.2, ["def ready(zone):", "    return zone.moisture < 30"], 12, font="Consolas")
    s = _blank(prs, "Решение о поливе")
    _box(s, 0.05, 0.3, 0.6, 0.5, ["def decision(level):", "\u00a0\u00a0\u00a0\u00a0if level < 15:", "\u00a0\u00a0\u00a0\u00a0\u00a0\u00a0\u00a0\u00a0return False, \"Недостаточно воды в баке\"",
                                   "\u00a0\u00a0\u00a0\u00a0return True, \"Полив разрешён\"", "):"], 12, font="Consolas")  # messages in Russian are code's strings
    b = deck_text(_save(prs, tmp_path)).brief
    assert "- снизить затраты на 10%;\n- увеличить выручку;\n- выйти в два новых региона." in b and "```\nснизить" not in b
    sec = _section(b, 3)
    assert sec.startswith("Слайд 3. Клиенты уходят к конкурентам") and "01 / Проблема" in sec and "```\n01" not in sec
    assert "- снижение затрат\n- рост выручки\n- новые клиенты" in sec  # never «снижение затрат рост выручки …»
    assert "Продукты, упаковка и списания" in sec
    assert "```\ndef ready(zone):\n    return zone.moisture < 30\n```" in sec
    assert _section(b, 4).count("```") == 2 and "return False, \"Недостаточно воды в баке\"" in _section(b, 4)


def test_a_figure_takes_the_label_beside_it_and_never_a_label_the_layout_does_not_give_it(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Ключевые показатели")
    for i, (fig, lab) in enumerate([("148 500 ₽", "Выручка за 30 дней"), ("20 000 → 35 000 ₽", "Маркетинговый бюджет"), ("35% → 33%", "Доля расходов"),
                                    ("6 × 4 м", "Теплица"), ("≤ 3 мин", "Сборка трёх вариантов"), ("113 с", "Шаблон VK Tech")]):
        _box(s, 0.03 + (i % 3) * 0.32, 0.25 + (i // 3) * 0.3, 0.28, 0.1, fig, 36)
        _box(s, 0.03 + (i % 3) * 0.32, 0.36 + (i // 3) * 0.3, 0.28, 0.06, lab, 14)
    s = _blank(prs, "Бюджет проекта")
    _box(s, 0.05, 0.22, 0.4, 0.06, "Распределение бюджета", 18)
    for i, (lab, fig) in enumerate([("Насос и клапаны", "16 000 ₽"), ("Датчики", "14 000 ₽"), ("Монтаж", "10 000 ₽")]):
        _box(s, 0.07, 0.29 + i * 0.06, 0.2, 0.05, lab, 16)
        _box(s, 0.27, 0.29 + i * 0.06, 0.1, 0.05, fig, 16)
    s = _blank(prs, "Время")
    _box(s, 0.05, 0.3, 0.6, 0.06, "Время сборки, секунды", 16)
    _box(s, 0.05, 0.37, 0.6, 0.06, "шаблоны VK — демо-текст на двенадцать слайдов; незнакомый шаблон — прогон из видео", 11)
    _box(s, 0.05, 0.44, 0.1, 0.06, "113 с", 14)
    b = deck_text(_save(prs, tmp_path)).brief
    sec = _section(b, 2)
    for line in ("Выручка за 30 дней — 148 500 ₽", "Маркетинговый бюджет — 20 000 → 35 000 ₽", "Доля расходов — 35% → 33%", "Теплица — 6 × 4 м",
                 "Сборка трёх вариантов — ≤ 3 мин", "Шаблон VK Tech — 113 с"):
        assert line in sec, line
    sec = _section(b, 3)
    assert "Насос и клапаны — 16 000 ₽\n\nДатчики — 14 000 ₽\n\nМонтаж — 10 000 ₽" in sec  # the row's own label, never the heading
    assert "Распределение бюджета —" not in sec
    sec = _section(b, 4)
    assert "— 113 с" not in sec and sec.rstrip().endswith("113 с")  # a sentence between: no pairing the layout does not show


def test_close_boxes_keep_their_order_and_a_cards_boxes_stay_together(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Три шага")
    for k, (head, text) in enumerate([("Разбор шаблона", "палитра, шрифты, сетка"), ("Агент", "аналитик, дизайнер, критик"), ("Вёрстка", "три варианта и проверка")]):
        card = s.shapes.add_shape(5, Emu(int(W * (0.04 + k * 0.32))), Emu(int(H * 0.3)), Emu(int(W * 0.29)), Emu(int(H * 0.5)))
        card.name = f"Card {k}"
        _box(s, 0.06 + k * 0.32, 0.35 + 0.01 * k, 0.25, 0.06, head, 16)  # the headings of the cards stand 1–2 % apart
        _box(s, 0.06 + k * 0.32, 0.45, 0.25, 0.1, text, 12)
    b = deck_text(_save(prs, tmp_path)).brief
    sec = _section(b, 2)
    order = [sec.index(x) for x in ("Разбор шаблона", "палитра", "Агент", "аналитик", "Вёрстка", "три варианта")]
    assert order == sorted(order)


def test_a_capitalised_title_stays_and_a_kicker_gives_way_only_to_a_larger_heading_under_it(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "ПРОБЛЕМА")
    _box(s, 0.05, 0.3, 0.9, 0.3, ["Клиенты ждут ответа два дня.", "Половина уходит к конкурентам."], 18)
    s = _blank(prs, "ПОДХОД")
    _box(s, 0.05, 0.2, 0.9, 0.12, "Шаблон — набор правил для агента", 54)
    _box(s, 0.05, 0.5, 0.9, 0.2, "Агент собирает слайды по правилам шаблона.", 18)
    s = _blank(prs, "ДЕМО-ПРОГОН")
    _box(s, 0.7, 0.2, 0.25, 0.08, "1 мин 18 с", 60)
    _box(s, 0.7, 0.29, 0.25, 0.06, "шаблон VK Tech в видео", 12)
    s = _blank(prs)
    _box(s, 0.05, 0.05, 0.6, 0.1, "Команда", 40)
    _box(s, 0.05, 0.18, 0.8, 0.06, "Шесть человек, три города", 18)
    s = _blank(prs)
    _box(s, 0.05, 0.05, 0.1, 0.1, "1", 60)
    _box(s, 0.2, 0.06, 0.7, 0.08, "Первый шаг пилота", 28)
    b = deck_text(_save(prs, tmp_path)).brief
    assert _section(b, 2).startswith("Слайд 2. ПРОБЛЕМА\n\n")
    sec = _section(b, 3)
    assert sec.startswith("Слайд 3. Шаблон — набор правил для агента\n\nПОДХОД\n\n")  # the kicker stays, over its heading
    assert _section(b, 4).startswith("Слайд 4. ДЕМО-ПРОГОН\n\n") and "шаблон VK Tech в видео — 1 мин 18 с" in _section(b, 4)
    assert _section(b, 5).startswith("Слайд 5. Команда\n\nШесть человек, три города")
    assert _section(b, 6).startswith("Слайд 6. Первый шаг пилота")  # a figure is never the heading


def test_a_broken_shape_never_stops_the_reading(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Чернила и фигуры")
    _box(s, 0.05, 0.3, 0.6, 0.1, "Текст рядом с рукописной пометкой", 18)
    ink = etree.SubElement(s.shapes._spTree, f"{P_NS}contentPart")  # ink: python-pptx makes no shape of it
    ink.set(f"{R_NS}id", "rId99")
    tb = _box(s, 0.05, 0.5, 0.6, 0.1, "Текст в фигуре без геометрии", 18)
    tb._element.find(f"{P_NS}nvSpPr/{P_NS}cNvSpPr").attrib.pop("txBox", None)
    geom = tb._element.find(f"{P_NS}spPr/{A_NS}prstGeom")
    geom.getparent().remove(geom)
    s = _blank(prs, "Диаграмма без данных")
    gf = _chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ["A", "B"], [("S", (1, 2))])
    gf._element.find(f".//{C_NS}chart").set(f"{R_NS}id", "rId777")  # a relationship that is not there
    _box(s, 0.05, 0.85, 0.6, 0.08, "Подпись под диаграммой", 14)
    b = deck_text(_save(prs, tmp_path)).brief
    sec = _section(b, 2)
    assert "Текст рядом с рукописной пометкой" in sec and "Текст в фигуре без геометрии" in sec
    assert "Подпись под диаграммой" in _section(b, 3)


def _rgba_png(draw, size=(200, 160)) -> io.BytesIO:
    im = Image.new("RGBA", size, (0, 0, 0, 0))
    from PIL import ImageDraw

    draw(ImageDraw.Draw(im), size)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def _photo(seed: int, size=(240, 300)) -> io.BytesIO:
    """A photograph's grain: every pixel its own (a portrait, a view)."""
    import random

    rnd = random.Random(seed)
    im = Image.new("RGB", size)
    im.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(size[0] * size[1])])
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def test_every_authors_picture_gets_its_place_and_decoration_none(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    # a picture inside a picture placeholder (shape type PLACEHOLDER: read by its element, «p:pic»)
    s = prs.slides.add_slide(prs.slide_layouts[8])
    s.shapes.title.text = "Наш новый офис"
    next(ph for ph in s.placeholders if "PICTURE" in str(ph.placeholder_format.type)).insert_picture(_photo(1, (400, 300)))
    # a team: four portraits of 2.5 % of the slide each (smaller than the old 6 % floor)
    s = _blank(prs, "Команда")
    for k in range(4):
        s.shapes.add_picture(_photo(10 + k), Emu(int(W * (0.05 + 0.24 * k))), Emu(int(H * 0.3)), Emu(int(W * 0.15)), Emu(int(H * 0.17)))
    s.shapes.add_picture(_photo(2), Emu(int(W * 0.9)), Emu(int(H * 0.05)), Emu(int(W * 0.03)), Emu(int(H * 0.05)))  # an icon: 0.15 %
    # line art on transparency (a scheme's strokes) is a picture; a frame on transparency is decoration
    s = _blank(prs, "Схема")
    art = _rgba_png(lambda d, sz: [d.line((10, 10 + 20 * i, sz[0] - 10, sz[1] - 10 - 20 * i), fill=(0, 0, 0, 255), width=4) for i in range(5)])
    s.shapes.add_picture(art, Emu(int(W * 0.05)), Emu(int(H * 0.25)), Emu(int(W * 0.4)), Emu(int(H * 0.5)))
    frame = _rgba_png(lambda d, sz: d.rectangle((0, 0, sz[0] - 1, sz[1] - 1), outline=(30, 30, 200, 255), width=3))
    s.shapes.add_picture(frame, Emu(int(W * 0.5)), Emu(int(H * 0.25)), Emu(int(W * 0.4)), Emu(int(H * 0.5)))
    d = deck_text(_save(prs, tmp_path))
    b = d.brief
    assert "Изображение со слайда 2 исходной презентации" in _section(b, 2)
    sec = _section(b, 3)
    assert sec.count("[Место под изображение:") == 4 and "Изображение 4 со слайда 3" in sec
    assert _section(b, 4).count("[Место под изображение:") == 1
    assert len(d.pictures) == 6
    # more than the first picture of a slide gets its place in the improved deck: the second on a continuation slide,
    # a gallery's others described in the speaker notes
    st = read_structure(b)
    plan = plan_deck(st.specs, None, ("structured",), title=st.title)
    slides = expand_slides(plan.decks["structured"], plan.readings)
    team = [s for s, i, _ in slides if i == 2]
    hints = [s.content.photo_hint for s in team]
    assert len(team) == 2 and team[1].headline == "Команда (продолжение)" and all(h and h.startswith("Изображение") for h in hints) and len(set(hints)) == 2
    assert "Изображение 4 со слайда 3" in team[0].notes


def test_a_cropped_picture_is_kept_as_the_slide_shows_it_and_a_vector_one_only_when_it_converts(tmp_path, monkeypatch):
    import verstka.ingest.deck_text as D

    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Кадр")
    pic = s.shapes.add_picture(_photo(3, (400, 200)), Emu(int(W * 0.1)), Emu(int(H * 0.2)), Emu(int(W * 0.5)), Emu(int(H * 0.6)))
    pic.crop_left, pic.crop_right = 0.25, 0.25
    s = _blank(prs, "Схема в EMF")
    vec = s.shapes.add_picture(_photo(4, (400, 300)), Emu(int(W * 0.1)), Emu(int(H * 0.2)), Emu(int(W * 0.5)), Emu(int(H * 0.6)))
    part = vec.part.related_part(vec._element.blipFill.blip.rEmbed)
    part._blob = b"\x01\x00\x00\x00" + b"\x00" * 104  # an EMF picture (no decoder here reads it)
    from pptx.opc.packuri import PackURI

    part._partname = PackURI("/ppt/media/image99.emf")
    part._content_type = "image/x-emf"
    path = _save(prs, tmp_path)
    monkeypatch.setattr(D, "_convert_vectors", lambda blobs, timeout_s=60.0: {})
    d = deck_text(path)
    assert len(d.pictures) == 1 and "Схема в EMF" in d.brief and "[Место под изображение" not in _section(d.brief, 3)
    with Image.open(io.BytesIO(d.pictures[0][1])) as im:
        assert im.size == (200, 200)  # the crop of the frame applied: half of the width
    monkeypatch.setattr(D, "_convert_vectors", lambda blobs, timeout_s=60.0: {k: _photo(5).getvalue() for k in blobs})
    d = deck_text(path)
    assert len(d.pictures) == 2 and d.pictures[1][2] == "png" and "[Место под изображение" in _section(d.brief, 3)


def test_what_the_slides_repeat_is_left_out_but_never_a_figure_or_a_short_decks_text(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    for k in range(5):
        s = _blank(prs, f"Показатель {k + 1}")
        _box(s, 0.05, 0.93, 0.4, 0.05, "ООО «Ромашка» · конфиденциально", 10)
        _box(s, 0.4, 0.4, 0.2, 0.15, "3", 54)  # a KPI's figure at the same place on every slide
        _box(s, 0.4, 0.56, 0.2, 0.06, f"месяца до цели {k + 1}", 14)
    b = deck_text(_save(prs, tmp_path)).brief
    assert "Ромашка" not in b and b.count("— 3") == 5
    prs = _prs()
    _blank(prs, "Короткий отчёт")
    for k in range(3):
        s = _blank(prs, f"Раздел {k + 1}")
        if k < 2:
            _box(s, 0.05, 0.4, 0.6, 0.1, "Данные за 2025 год", 18)  # the author's line on two slides of a short deck
    b = deck_text(_save(prs, tmp_path, "short.pptx")).brief
    assert b.count("Данные за 2025 год") == 2


def test_speaker_notes_are_kept_as_the_slides_notes_never_as_its_text(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Регионы")
    _box(s, 0.05, 0.3, 0.6, 0.1, "Выходим в два региона.", 18)
    s.notes_slide.notes_text_frame.text = "Выходим в два региона.\nПодробно рассказать про Казань [и Пермь]: бюджет 85 000 рублей."
    b = deck_text(_save(prs, tmp_path)).brief
    sec = _section(b, 2)
    assert sec.rstrip().endswith("[Заметки докладчика: Подробно рассказать про Казань (и Пермь): бюджет 85 000 рублей.]")
    st = read_structure(b)
    spec = st.specs[1]
    assert "Заметки" not in spec.text and spec.notes.startswith("Подробно рассказать про Казань")
    assert 85000 not in [v for s in st.series for v in s.values] and not st.notes_rule  # no figure, no rule read from the notes
    plan = plan_deck(st.specs, None, ("structured",), title=st.title)
    sl = plan.decks["structured"][1]
    assert "Подробно рассказать про Казань" in sl.notes and "Казань" not in sl.model_dump_json(exclude={"notes"})
    # a brief given whole to the layout: the note in a slide's text is its notes too
    r = read_slide(2, "Регионы", "Выходим в два региона.\n\n[Заметки докладчика: Рассказать про Казань.]")
    assert [p.kind for p in r.parts] == ["para"] and notes_of(r).endswith("Рассказать про Казань.")


def _smartart(slide, items: list[tuple[int, str]]):
    """A SmartArt list: its data part (points under the document, by order) and the frame that names it."""
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI

    dgm = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    pts = ['<dgm:pt modelId="0" type="doc"><dgm:prSet/><dgm:spPr/></dgm:pt>']
    cxns = []
    parents = {-1: "0"}
    for k, (lvl, text) in enumerate(items, 1):
        pts.append(f'<dgm:pt modelId="{k}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></dgm:t></dgm:pt>')
        pts.append(f'<dgm:pt modelId="p{k}" type="pres"><dgm:prSet/><dgm:spPr/></dgm:pt>')
        parent = parents[lvl - 1]
        order = sum(1 for c in cxns if c[0] == parent)
        cxns.append((parent, str(k), order))
        parents[lvl] = str(k)
    links = "".join('<dgm:cxn modelId="c%s" srcId="%s" destId="%s" srcOrd="%d" destOrd="0"/>' % (d, s_, d, o) for s_, d, o in cxns)
    xml = (f'<dgm:dataModel xmlns:dgm="{dgm}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>{"".join(pts)}</dgm:ptLst>'
           f'<dgm:cxnLst>{links}</dgm:cxnLst></dgm:dataModel>')
    part = Part(PackURI("/ppt/diagrams/data1.xml"), "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml", slide.part.package, xml.encode("utf-8"))
    rid = slide.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData")
    frame = etree.fromstring(
        f'<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:nvGraphicFramePr><p:cNvPr id="77" name="Diagram"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        f'<p:xfrm><a:off x="{int(W * 0.1)}" y="{int(H * 0.3)}"/><a:ext cx="{int(W * 0.8)}" cy="{int(H * 0.5)}"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{dgm}"><dgm:relIds xmlns:dgm="{dgm}" r:dm="{rid}" r:lo="" r:qs="" r:cs=""/></a:graphicData></a:graphic></p:graphicFrame>'
    )
    slide.shapes._spTree.append(frame)


def test_smartart_text_is_read_as_a_list_in_its_order(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Как мы работаем")
    _smartart(s, [(0, "Заявка"), (1, "приходит в CRM"), (0, "Звонок"), (0, "Договор")])
    b = deck_text(_save(prs, tmp_path)).brief
    assert "- Заявка\n  - приходит в CRM\n- Звонок\n- Договор" in _section(b, 2)


def test_step_badges_and_a_legend_beside_its_chart_are_left_out(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Шаги")
    for k, text in enumerate(["Собрать данные", "Проверить", "Показать"]):
        _box(s, 0.05, 0.25 + 0.2 * k, 0.05, 0.08, f"0{k + 1}", 20)
        _box(s, 0.12, 0.25 + 0.2 * k, 0.6, 0.08, text, 20)
    s = _blank(prs, "Доля в выручке")
    _chart(s, XL_CHART_TYPE.PIE, ["Кофе", "Десерты", "Чай"], [("Доля", (60, 25, 15))], x=0.05, y=0.2, w=0.4, h=0.6)
    for k, (name, share) in enumerate([("Кофе", "60%"), ("Десерты", "25%"), ("Чай", "15%")]):
        _box(s, 0.5, 0.3 + 0.08 * k, 0.2, 0.06, name, 16)
        _box(s, 0.72, 0.3 + 0.08 * k, 0.08, 0.06, share, 16)
    _box(s, 0.05, 0.85, 0.9, 0.06, "Кофе даёт больше половины выручки", 16)
    b = deck_text(_save(prs, tmp_path)).brief
    sec = _section(b, 2)
    assert "01" not in sec and "02" not in sec and "Собрать данные\n\nПроверить\n\nПоказать" in sec
    sec = _section(b, 3)
    assert "[Круговая диаграмма: Доля — Кофе: 60; Десерты: 25; Чай: 15]" in sec and "Кофе даёт больше половины выручки" in sec
    assert "Кофе — 60%" not in sec and "\n\nДесерты\n" not in sec and "25%" not in sec


def test_hidden_slides_are_skipped_and_a_long_or_one_slide_deck_reads_back(tmp_path):
    prs = _prs()
    _blank(prs, "Итоги")
    s = _blank(prs, "Запасной слайд")
    _box(s, 0.05, 0.3, 0.6, 0.1, "Ответы на вопросы", 18)
    s._element.set("show", "0")
    s = _blank(prs, "Главное")
    _box(s, 0.05, 0.3, 0.6, 0.1, "Выручка выросла", 18)
    d = deck_text(_save(prs, tmp_path))
    assert "Запасной слайд" not in d.brief and "Слайд 3. Главное" in d.brief and d.slides == 2
    # «Слайд 100» and later; a deck of one slide
    text = "\n\n".join(f"Слайд {n}. Тема {n}\n\nТекст слайда {n}." for n in range(1, 106))
    assert [sp.number for sp in read_structure(text).specs][-3:] == [103, 104, 105]
    one = read_structure("Слайд 1. Итоги квартала\n\nВыручка выросла на 34%.\n")
    assert [(sp.number, sp.title) for sp in one.specs] == [(1, "Итоги квартала")]
    assert read_structure("Слайд 1 должен быть титульным, остальное придумай сам. Сделай презентацию на 5 слайдов.\n\nТема: кофе.").specs == []


def test_a_source_picture_place_takes_only_its_own_decks_picture_and_never_a_models(tmp_path):
    from pptx.util import Inches

    from verstka.images.fill import WHY_SOURCE, plan_pictures
    from verstka.images.places import CAPTION_PREFIX
    from verstka.schemas.outline import PHOTO_PLACE_NAME, DeckOutline, OutlineSlide, SlideContent

    d = deck_text(_deck(tmp_path))
    mine = tmp_path / "mine"
    keep_pictures(d, mine / "source_images")
    other_deck = tmp_path / "other.pptx"
    prs = Presentation(str(_deck(tmp_path)))
    _box(prs.slides[0], 0.5, 0.5, 0.3, 0.1, "другая версия", 12)  # another deck: other bytes, another code
    prs.save(str(other_deck))
    other = tmp_path / "other"
    keep_pictures(deck_text(other_deck), other / "source_images")
    (hint,) = [h for h, _, _ in d.pictures]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    frame = s.shapes.add_shape(5, Inches(7), Inches(1.5), Inches(5.5), Inches(4.5))
    frame.name = PHOTO_PLACE_NAME
    cap = s.shapes.add_textbox(Inches(7.5), Inches(3.2), Inches(4.5), Inches(1))
    cap.name = f"{CAPTION_PREFIX} {cap.shape_id}"
    path = tmp_path / "improved.pptx"
    prs.save(str(path))
    outline = DeckOutline(title="Итоги", slides=[OutlineSlide(id="a", kind="image_text", headline="Точка", content=SlideContent(bullets=["x"], photo_slot="Изображение", photo_hint=hint))])

    class Maker:
        calls: list = []

        def submit(self, ask, aspect):
            self.calls.append(ask)
            raise AssertionError("a model never draws a picture of the original deck")

    maker = Maker()
    (got,) = plan_pictures(path, outline, maker, reuse=source_records(other, tmp_path / "g1"), folder=tmp_path / "g1")
    assert got.future is None and got.why == WHY_SOURCE and not maker.calls  # another deck's pictures: the place stays free
    (got,) = plan_pictures(path, outline, maker, reuse=source_records(mine, tmp_path / "g2"), folder=tmp_path / "g2")
    assert got.future is not None and got.future.result().model == "original" and not maker.calls
