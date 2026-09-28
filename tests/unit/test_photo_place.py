"""A free place for the user's own photo («На этом слайде оставь свободное место под фотографию помещения. Фото я добавлю
самостоятельно»): the brief's sentences are a request, never slide text; the slide keeps its figures and leaves the
place free — the template's own photo place when a sample slide of the template has one, else a quiet frame beside the
content; the place is a neutral that reads on its ground and the audit knows it; a slide the deck lost is an error.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from verstka.schemas.common import contrast_ratio
from verstka.schemas.outline import PHOTO_PLACE_NAME

TPL = Path(__file__).resolve().parents[1] / "fixtures" / "tpl_audit" / "templates"

ASK = "На этом слайде обязательно оставь свободное место под фотографию помещения."
SELF = "Фото я добавлю самостоятельно."
KEEP = "Не заполняй это место текстом, диаграммой или сгенерированным изображением."
AREAS = "Из общей площади 180 м² тренировочный зал займет 100 м², раздевалки и душевые — 35 м², зона ресепшена и ожидания — 25 м², подсобные помещения — 20 м²."

BRIEF = f"""Тема презентации: «Фитнес-студия „Движение“». Презентация на 3 слайда.

Слайд 1. Идея проекта

Фитнес-студия «Движение» — пространство для групповых тренировок рядом с домом. Площадь студии — 180 м², максимальное количество участников одного занятия — 12 человек.

Слайд 2. Пространство студии

{AREAS} В зале предусмотрены зеркала, вентиляция и нескользящее покрытие.

{ASK} {SELF} {KEEP}

Слайд 3. Команда

Для запуска потребуются четыре тренера, два администратора и специалист по уборке. Месячный бюджет команды — 420 000 рублей.
"""


def _sat(hex_: str) -> float:
    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.0 if max(r, g, b) == 0 else (max(r, g, b) - min(r, g, b)) / max(r, g, b)


def test_the_request_for_a_photo_place_is_read_as_a_request():
    from verstka.planning.brief_structure import is_photo_instruction, photo_request

    assert photo_request(ASK) == "Фото помещения"
    assert photo_request("Оставь место для фото команды.") == "Фото команды"
    assert photo_request(SELF) == "Фото"
    assert all(is_photo_instruction(s) for s in (ASK, SELF, KEEP))
    # a photo the text talks about, or one it does not want, is no request
    for s in ("На фото видно, как вырос зал.", "Без фото, только цифры.", AREAS):
        assert photo_request(s) is None and not is_photo_instruction(s), s


def test_the_slide_spec_carries_the_photo_and_the_request_is_no_content():
    from verstka.planning.agent import _read_source
    from verstka.planning.brief_structure import read_structure

    st = read_structure(BRIEF)
    spec = next(s for s in st.specs if s.number == 2)
    assert spec.photo == "Фото помещения"
    assert st.specs[0].photo is None and st.specs[2].photo is None
    src = _read_source(f"{AREAS}\n\n{ASK} {SELF} {KEEP}")
    assert {ASK, SELF, KEEP} <= set(src.asks)
    assert not any("фото" in s.lower() or "место" in s.lower() for s in src.sentences)


@pytest.mark.parametrize(
    "ground,card,surface,text",
    [
        ("D4DECF", "EB9F41", "EB9F41", "202A79"),  # a sage ground whose cards are orange: never the accent card
        ("000000", "1A1A1A", "1A1A1A", "FFFFFF"),  # a black ground
        ("FFFFFF", "F2F2F2", "F2F2F2", "222222"),  # a card too faint to read on white
        ("EAE8F3", "000000", "000000", "000000"),  # a lavender ground with black cards
    ],
)
def test_the_photo_place_is_a_quiet_neutral_that_reads_on_its_ground(ground, card, surface, text):
    from verstka.rendering.synth import PHOTO_FILL_CONTRAST, _photo_fill

    kit = SimpleNamespace(card=SimpleNamespace(fill=card), manifest=None)
    pal = SimpleNamespace(bg=ground, surface=surface, card_fill=card, text=text)
    got = _photo_fill(kit, pal, ground)
    lo, hi = PHOTO_FILL_CONTRAST
    assert lo <= contrast_ratio(got, ground) <= hi + 0.05, got
    assert _sat(got) < 0.2 and got.upper() != card.upper() or _sat(card) < 0.2


def _boxes(slide, W: int, H: int):
    place, texts = None, []
    for sh in slide.shapes:
        box = (sh.left / W, sh.top / H, (sh.left + sh.width) / W, (sh.top + sh.height) / H)
        if sh.name == PHOTO_PLACE_NAME:
            place = box
        elif sh.has_text_frame and sh.text_frame.text.strip():
            texts.append((sh.text_frame.text, box))
    return place, texts


def _apart(a, b) -> bool:
    return a[2] <= b[0] + 0.005 or b[2] <= a[0] + 0.005 or a[3] <= b[1] + 0.005 or b[3] <= a[1] + 0.005


def test_a_template_without_a_photo_slide_gets_a_frame_beside_the_figures(tmp_path):
    """The whole path without a model: the brief's slide 2 keeps its four areas and a free frame for the photo beside
    them — no text in it, none of the request's words on the slide."""
    from verstka.pipeline.generate import generate_variants
    from verstka.planning.brief import load_brief

    p = tmp_path / "brief.md"
    p.write_text(BRIEF, encoding="utf-8")
    g = generate_variants(TPL / "lo_focus.pptx", brief=load_brief(p), strategies=["structured"], out_dir=tmp_path / "run", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, audit=True, exports=[])
    v = g.variants[0]
    i, osl = next((i, s) for i, s in enumerate(v.outline.slides) if s.spec_ref == 2)
    assert osl.content.photo_slot == "Фото помещения"
    prs = Presentation(str(tmp_path / "run" / "structured" / "deck.pptx"))
    place, texts = _boxes(prs.slides[i], prs.slide_width, prs.slide_height)
    assert place is not None and (place[2] - place[0]) >= 0.2 and (place[3] - place[1]) >= 0.3
    assert all(_apart(place, b) for _, b in texts)
    words = " ".join(t for t, _ in texts)
    for fig in ("100", "35", "25", "20"):
        assert fig in words
    assert "фото" not in words.lower() and "не заполняй" not in words.lower()
    assert not [x for x in v.audit.issues if x.slide == i + 1 and x.severity == "error"]


def _dataset_template(stem: str) -> Path | None:
    import os

    env = os.environ.get("VERSTKA_FIXTURES_DIR")
    for c in ([Path(env)] if env else []) + [Path(__file__).resolve().parents[3] / "Датасет"]:
        hit = next(iter(sorted(c.glob(f"{stem}*.pptx"))), None) if c.is_dir() else None
        if hit is not None:
            return hit
    return None


def test_a_template_with_a_photo_slide_lends_its_own_photo_place(tmp_path):
    """VK Education has content slides with a photo half: the slide is set on one of them, the photo made an empty
    place of the same size where it stood. Skipped when the VK templates (Датасет/) are not on this machine."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent

    pptx = _dataset_template("Шаблон презентации VK Education")
    if pptx is None:
        pytest.skip("VK Education template not available")
    nums = [NumberCallout(value="100 м²", label="Тренировочный зал"), NumberCallout(value="35 м²", label="Раздевалки и душевые"),
            NumberCallout(value="25 м²", label="Зона ресепшена и ожидания"), NumberCallout(value="20 м²", label="Подсобные помещения")]
    s = OutlineSlide(id="p1", kind=PatternKind.stat_row, headline="Тренировочный зал займет больше половины площади",
                     content=SlideContent(numbers=nums, photo_slot="Фото помещения"), takeaway="Тренировочный зал — основная зона студии")
    g = generate_variants(pptx, outline=DeckOutline(title="Студия", slides=[s], language="ru"), strategies=["structured"], out_dir=tmp_path / "run",
                          workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, audit=True, exports=[])
    prs = Presentation(str(tmp_path / "run" / "structured" / "deck.pptx"))
    place, texts = _boxes(prs.slides[0], prs.slide_width, prs.slide_height)
    assert place is not None and (place[2] - place[0]) >= 0.3 and (place[3] - place[1]) >= 0.5
    assert all(_apart(place, b) for _, b in texts)
    assert any("template's own photo place" in w for r in g.variants[0].render.slides for w in r.warnings)


def _ir(index: int, oid: str | None):
    from verstka.schemas.common import Bbox, BboxFrac
    from verstka.schemas.deck_ir import IRElement, IRParagraph, IRRun, IRSlide

    e = IRElement(id=f"e{index}", type="text", bbox=Bbox(x=0, y=0, w=1000, h=100), bbox_frac=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.1),
                  paragraphs=[IRParagraph(text=f"Слайд {index}", runs=[IRRun(text=f"Слайд {index}", size_pt=32, font="Arial")])], ph_type="title")
    return IRSlide(index=index, elements=[e], outline_id=oid)


def _ctx(slides, outline):
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import BboxFrac
    from verstka.schemas.deck_ir import DeckIR
    from verstka.schemas.template import SlideSize, Spacing, TemplateManifest, Tokens

    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=12192000, h=6858000), tokens=Tokens(spacing=Spacing(safe_area=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.9))), patterns=[], n_slides=1)
    return AuditContext(ir=DeckIR(source="x.pptx", slide_w=12192000, slide_h=6858000, slides=slides), manifest=man, outline=outline)


def test_a_slide_the_deck_lost_is_an_error():
    from verstka.audit.checks.integrity import slide_missing
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    o = DeckOutline(title="x", slides=[OutlineSlide(id=k, kind=PatternKind.bullets, headline=f"Слайд {k}") for k in ("a", "b", "c")])
    got = slide_missing(_ctx([_ir(1, "a"), _ir(2, "c")], o))
    assert len(got) == 1 and got[0].severity == "error" and got[0].outline_id == "b" and got[0].details["plan_slide"] == 2
    # a slide whose marker was not written is still in the deck: the count tells
    assert slide_missing(_ctx([_ir(1, "a"), _ir(2, None), _ir(3, "c")], o)) == []


def test_the_audit_leaves_the_photo_place_colour_alone():
    from verstka.audit.checks.template import color_not_in_palette
    from verstka.schemas.common import Bbox, BboxFrac
    from verstka.schemas.deck_ir import IRElement
    from verstka.schemas.template import ColorToken

    place = IRElement(id="ph", type="shape", name=PHOTO_PLACE_NAME, fill_hex="C3CCBE", bbox=Bbox(x=0, y=0, w=1000, h=1000), bbox_frac=BboxFrac(x=0.0, y=0.0, w=0.5, h=1.0))
    other = place.model_copy(update={"id": "x", "name": "Card 3"})
    ctx = _ctx([_ir(1, None)], None)
    ctx.manifest.tokens.colors = [ColorToken(hex="D4DECF"), ColorToken(hex="EB9F41"), ColorToken(hex="202A79")]
    ctx.ir.slides[0].elements.append(place)
    assert color_not_in_palette(ctx) == []
    ctx.ir.slides[0].elements.append(other)
    assert [i.element_ids for i in color_not_in_palette(ctx)] == [["x"]]


# ---------------------------------------------------------------------------- the chat: «на слайде 2 оставь место под фото»

EDIT_BRIEF = """Сервис напоминаний для команд

Слайд 1. Проблема
Сотрудники теряют задачи в потоке сообщений. 40% задач из чатов забываются. Напоминания вручную занимают 20 минут в день.

Слайд 2. Как внедряем
Запуск проходит в три этапа: пилот на одной команде, настройка сценариев, запуск на всю компанию.
"""


def _edit_outline():
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    return DeckOutline(title="Сервис напоминаний", strategy="structured", slides=[
        OutlineSlide(id="sl1", kind=PatternKind.bullets, headline="Проблема", spec_ref=1, content=SlideContent(bullets=["Сотрудники теряют задачи в потоке сообщений", "40% задач из чатов забываются"])),
        OutlineSlide(id="sl2", kind=PatternKind.bullets, headline="Как внедряем", spec_ref=2, content=SlideContent(bullets=["Пилот на одной команде", "Настройка сценариев", "Запуск на всю компанию"])),
    ], agent_log=["Аналитик: прочитал текст."])


def test_the_chat_reads_a_photo_place_request_as_an_edit_of_that_slide():
    from verstka.api.edits import looks_like_edit, parse_edit

    for msg in ("На слайде 2 оставь место под фото команды", "убери место под фото на втором слайде"):
        assert looks_like_edit(msg), msg
        req = parse_edit(msg, on_screen=1, total=2)
        assert req is not None and req.kind == "slide" and req.slides == [2]


def test_a_photo_place_is_made_and_taken_back_without_a_model():
    from verstka.planning.brief import parse_brief_text
    from verstka.planning.slide_edit import redesign_slide

    brief = parse_brief_text(EDIT_BRIEF)
    r = redesign_slide(_edit_outline(), 2, brief, None, request="На слайде 2 оставь место под фото команды")
    s = r.outline.slides[1]
    assert r.how == "rules" and s.content.photo_slot == "Фото команды"
    assert s.content.bullets == ["Пилот на одной команде", "Настройка сценариев", "Запуск на всю компанию"]
    assert "место под фото" in r.reply and r.outline.slides[0].content.photo_slot is None
    back = redesign_slide(r.outline, 2, brief, None, request="убери место под фото")
    assert back.outline.slides[1].content.photo_slot is None and "место под фото убрано" in back.reply


def test_the_designer_is_told_to_leave_the_photo_place_and_the_slide_keeps_it():
    from verstka.planning.brief import parse_brief_text
    from verstka.planning.slide_edit import redesign_slide
    from verstka.providers.mock import MockProvider
    from verstka.providers.registry import ProviderLimits, ProviderRegistry
    from verstka.skills_registry.registry import SkillsRegistry

    prompts: list[str] = []
    answer = {"kind": "process", "headline": "Запуск в три этапа", "items": [{"title": "Пилот", "text": "одна команда"}, {"title": "Настройка", "text": "сценарии"}, {"title": "Запуск", "text": "вся компания"}]}

    def script(messages):
        system, user = messages[0].content, messages[-1].content
        if "presentation designer" in system:
            prompts.append(user)
            return answer
        return {"facts": [], "series": [], "tables": []}

    p = MockProvider(script, model="qwen/qwen3.8-27b")
    providers = ProviderRegistry(roles={"llm": p, "vlm": p}, limits=ProviderLimits(max_concurrency=2, time_budget_s=60))
    r = redesign_slide(_edit_outline(), 2, parse_brief_text(EDIT_BRIEF), None, request="оставь на этом слайде место для фотографии команды",
                       skills=SkillsRegistry.load(), providers=providers)
    assert r.how == "model" and r.outline.slides[1].content.photo_slot == "Фото команды"
    assert len(prompts) == 1 and "a free place for the user's own photo (Фото команды)" in prompts[0]


# ---------------------------------------------------------------------------- the agent with a model (scripted)


def test_the_agent_keeps_the_photo_place_whatever_the_designer_answers():
    """The designer is told about the place; an answer that describes it («Здесь будет фото помещения») loses that line,
    and every variant's slide keeps the place free."""
    import re

    from verstka.planning import agent as A
    from verstka.planning.brief import parse_brief_text
    from verstka.planning.strategies import load_strategies
    from verstka.providers.mock import MockProvider
    from verstka.providers.registry import ProviderLimits, ProviderRegistry
    from verstka.skills_registry.registry import SkillsRegistry

    prompts: dict[str, str] = {}

    def script(messages):
        system, user = messages[0].content, messages[-1].content
        if "presentation designer" in system:
            heading = re.search(r"Heading of this slide: «(.+?)»", user).group(1)
            prompts[heading] = user
            if heading == "Пространство студии":
                return {"kind": "stat_row", "headline": "Тренировочный зал займет больше половины площади",
                        "numbers": [{"value": "100 м²", "label": "Тренировочный зал"}, {"value": "35 м²", "label": "Раздевалки и душевые"},
                                    {"value": "25 м²", "label": "Зона ресепшена и ожидания"}, {"value": "20 м²", "label": "Подсобные помещения"}],
                        "bullets": ["Здесь будет фото помещения"]}
            return {"kind": "bullets", "headline": heading, "bullets": ["Площадь студии — 180 м²"]}
        if "strict reviewer" in system:
            return {"issues": []}
        return {"facts": [], "series": [], "tables": []}

    p = MockProvider(script, model="qwen/qwen3.8-27b")
    reg = ProviderRegistry(roles={"llm": p, "vlm": p}, limits=ProviderLimits(max_concurrency=4, time_budget_s=120))
    res = A.run_agent(parse_brief_text(BRIEF), None, list(load_strategies().values()), skills=SkillsRegistry.load(), providers=reg)
    assert "a free place for the user's own photo (Фото помещения)" in prompts["Пространство студии"]
    assert res is not None and res.outlines
    for name, o in res.outlines.items():
        s = next(x for x in o.slides if x.spec_ref == 2)
        assert s.content.photo_slot == "Фото помещения", name
        assert not any("фото" in b.lower() for b in s.content.bullets), (name, s.content.bullets)
        shown = [n.value for n in s.content.numbers] + [c for r in (s.content.table.rows if s.content.table else []) for c in r]
        assert all(any(v in x for x in shown) for v in ("100", "35", "25", "20")), (name, shown)  # figures or a table (compact)
        assert all(x.content.photo_slot is None for x in o.slides if x.spec_ref != 2), name
