"""Подбор шрифтов под шаблон (verstka/brand/fontpick.py) и шрифты внутри PPTX (verstka/brand/fontembed.py).

Без модели и сети: шаблоны LibreOffice из фикстур (стандартные шрифты), колода VK Tech (свой Play), синтетический
шаблон python-pptx."""

from __future__ import annotations

import re
import struct
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

import verstka.brand.fontpick as fpk
from verstka.analysis.manifest import analyze_template
from verstka.brand.fontembed import FONT_DIR, FontFile, bundled_family, bundled_fonts, embed_fonts, embedded_families, used_families
from verstka.brand.fontpick import STANDARD_FONTS, choose_fonts, fonted_template, fonts_info, heading_threshold, is_standard, pick_for_template, score_pairs
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.cyrillic import embedded_coverage
from verstka.schemas.template import TemplateManifest, TypeStep

ROOT = Path(__file__).resolve().parents[2]
FIX = ROOT / "tests" / "fixtures" / "tpl_audit"
PICKED = {"Manrope", "Onest", "Unbounded", "Raleway", "Jost", "Exo 2", "IBM Plex Sans", "IBM Plex Serif", "PT Serif", "Lora",
          "Playfair Display", "Old Standard TT", "Comfortaa", "Bitter", "Nunito", "Oswald"}


def _analyze(path: Path, ws: Path) -> TemplateManifest:
    return analyze_template(path, workspace_root=ws, use_llm=False, use_vlm=False, render=False)


# ---------------------------------------------------------------------------------------------- шрифты Verstka


def test_bundled_fonts_regular_bold_cyrillic_ofl():
    fonts = bundled_fonts()
    assert PICKED <= set(fonts), PICKED - set(fonts)
    for fam in PICKED:
        faces = fonts[fam]
        assert {"regular", "bold"} <= set(faces), fam
        assert (faces["regular"].parent / "OFL.txt").exists(), fam  # лицензия рядом с каждым семейством
        for kind, path in faces.items():
            ff = FontFile(path)
            assert ff.family == fam and ff.bold == (kind == "bold") and not ff.italic
            assert ff.embeddable, fam  # OFL: fsType 0
            assert ff.unicode_ranges[0] & (1 << 9), fam  # кириллица
    total = sum(p.stat().st_size for p in FONT_DIR.rglob("*.ttf"))
    assert total < 25 * 1024 * 1024


def test_bundled_family_matches_names():
    assert bundled_family("exo 2") == "Exo 2" and bundled_family("IBMPlexSans") == "IBM Plex Sans"
    assert bundled_family("Arial") is None and bundled_family("") is None


def test_layout_measures_with_bundled_files():
    from verstka.rendering.fonts import is_installed, is_measured, text_width_pt

    for fam in ("Manrope", "Unbounded", "Old Standard TT", "Exo 2"):
        assert is_installed(fam) and is_measured(fam), fam
    s = "Выручка выросла на 26 %"
    # широкий Unbounded шире Manrope, узкий Oswald уже — меряется сам файл, а не коэффициент
    assert text_width_pt(s, "Unbounded", 24) > text_width_pt(s, "Manrope", 24) * 1.15
    assert text_width_pt(s, "Oswald", 24) < text_width_pt(s, "Manrope", 24) * 0.9


# ---------------------------------------------------------------------------------------------- встраивание


def test_eot_header_is_read_back():
    ff = FontFile(bundled_fonts()["Manrope"]["bold"])
    eot = ff.eot()
    size, data_size, version, flags = struct.unpack("<IIII", eot[:16])
    assert size == len(eot) and data_size == len(ff.data) and version == 0x00020001 and flags == 0
    assert struct.unpack("<H", eot[34:36])[0] == 0x504C
    assert eot.endswith(ff.data)
    assert struct.unpack("<I", eot[28:32])[0] == 700  # Weight


def _deck(path: Path, families: list[str]) -> Path:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, fam in enumerate(families):
        tb = s.shapes.add_textbox(Emu(300000), Emu(300000 + 600000 * i), Emu(6000000), Emu(500000))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text, r.font.name, r.font.size = "Съешь же ещё мягких булок", fam, Pt(20)
    prs.save(path)
    return path


def test_embed_fonts_package(tmp_path):
    deck = _deck(tmp_path / "d.pptx", ["Manrope", "Arial", "Unbounded"])
    assert {"Manrope", "Arial", "Unbounded"} <= used_families(deck)
    done = embed_fonts(deck)
    assert sorted(done) == ["Manrope", "Unbounded"]  # Arial — не шрифт Verstka, не встраивается
    with zipfile.ZipFile(deck) as z:
        names = z.namelist()
        pres = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()
        ct = z.read("[Content_Types].xml").decode()
    assert len([n for n in names if n.startswith("ppt/fonts/") and n.endswith(".fntdata")]) == 4
    assert 'embedTrueTypeFonts="1"' in pres
    # порядок CT_Presentation: список шрифтов после notesSz, до defaultTextStyle
    assert pres.index("<p:notesSz") < pres.index("<p:embeddedFontLst") < pres.index("<p:defaultTextStyle")
    assert re.search(r'<p:embeddedFont><p:font typeface="Manrope"[^>]*/><p:regular r:id="[^"]+"/><p:bold r:id="[^"]+"/></p:embeddedFont>', pres)
    assert rels.count("relationships/font") == 4
    assert 'Extension="fntdata" ContentType="application/x-fontdata"' in ct
    # шрифты в колоде видит и проверка кириллицы (её заголовок EOT)
    assert embedded_coverage(deck) == {"manrope": True, "unbounded": True}
    # повтор ничего не добавляет; python-pptx при сохранении колоды части шрифтов сохраняет
    assert embed_fonts(deck) == []
    prs = Presentation(str(deck))
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(tmp_path / "e.pptx")
    assert embedded_families(tmp_path / "e.pptx") == ["Manrope", "Unbounded"]
    with zipfile.ZipFile(tmp_path / "e.pptx") as z:
        assert len([n for n in z.namelist() if n.endswith(".fntdata")]) == 4


# ---------------------------------------------------------------------------------------------- правило выбора


def test_standard_fonts_list():
    for f in ("Arial", "Calibri", "Calibri Light", "Times New Roman", "Liberation Sans", "DejaVu Sans", "Carlito", "Helvetica", "Open Sans", "Noto Sans", "Aptos"):
        assert is_standard(f), f
    for f in ("Play", "Montserrat", "Lato Light", "Georgia", "Trebuchet MS", "Futura", "VK Sans Display", "Roboto", ""):
        assert not is_standard(f), f
    assert "times new roman" in STANDARD_FONTS


def test_own_font_template_is_kept(tmp_path):
    """VK Tech (Play) — свой шрифт шаблона: шрифты не меняются, копии нет."""
    m = _analyze(FIX / "vktech_short.pptx", tmp_path / "ws")
    ch = choose_fonts(m, TemplateWorkspace.open(m.template_id, tmp_path / "ws").source)
    assert not ch.apply and "Play" in ch.reason
    tid, ch2 = fonted_template(m.template_id, workspace_root=tmp_path / "ws", render=False)
    assert tid == m.template_id and not ch2.apply


def test_designer_system_font_is_kept(tmp_path):
    """Georgia и Trebuchet MS — выбор дизайнера (не шрифт по умолчанию): оставлены."""
    m = _analyze(FIX / "templates" / "synth_sidebar43.pptx", tmp_path / "ws")
    ch = choose_fonts(m, TemplateWorkspace.open(m.template_id, tmp_path / "ws").source)
    assert not ch.apply and "Georgia" in ch.reason


def test_standard_template_gets_pair_deterministically(tmp_path):
    ws = tmp_path / "ws"
    m = _analyze(FIX / "templates" / "lo_focus.pptx", ws)
    src = TemplateWorkspace.open(m.template_id, ws).source
    a = choose_fonts(m, src)
    b = choose_fonts(m, src)
    assert a.apply and a.heading in PICKED and a.body in PICKED
    assert (a.heading, a.body, a.reason) == (b.heading, b.body, b.reason)
    assert a.reason.startswith("Шрифты шаблона стандартные") and a.heading in a.reason
    # выключатель и брендбук со шрифтами — шаблон как есть
    assert not choose_fonts(m, src, enabled=False).apply
    assert not choose_fonts(m, src, brand_fonts=True).apply


def test_scores_follow_character():
    base = {k: 0.0 for k in ("dark", "colorful", "warm", "green", "cool", "purple", "earthy", "achrom", "round", "circles", "angular",
                              "diagonal", "curvy", "thin", "photo", "serif", "narrow", "dark_cool", "colorful_warm", "achrom_contrast", "circles_light")}
    top = lambda **kw: score_pairs({**base, **kw})[0][1].key  # noqa: E731
    assert top() == "modern"  # без признаков — нейтральная пара
    assert top(earthy=0.9, serif=1.0) == "vintage"
    assert top(serif=1.0, cool=0.5) == "academic"
    assert top(green=0.8, curvy=0.5) == "natural"
    assert top(round=0.9, colorful=0.6) == "friendly"
    assert top(colorful=1.0, warm=1.0, colorful_warm=1.0) == "bright"
    assert top(dark=1.0, achrom=1.0, achrom_contrast=1.0, thin=0.5) == "elegant"
    assert top(dark=0.8, cool=0.9, dark_cool=1.0, angular=0.6) == "tech"
    assert top(narrow=1.0, diagonal=0.6) == "dynamic"


def test_heading_threshold_from_scale():
    from verstka.schemas.template import Tokens

    def man(steps):
        mm = TemplateManifest.model_construct(tokens=Tokens())
        mm.tokens.typography.scale = [TypeStep(role=r, size_pt=s) for r, s in steps]
        return mm

    assert heading_threshold(man([("h1", 14), ("body", 14)])) == 20.0  # шкала пустых образцов: не весь текст — заголовок
    assert heading_threshold(man([("h1", 40), ("body", 16)])) == 34.0
    assert heading_threshold(man([("h1", 130), ("body", 24)])) == 36.0
    assert heading_threshold(man([])) == 23.8


# ---------------------------------------------------------------------------------------------- копия шаблона


def test_fonted_template_copy_analyzed_embedded_cached(tmp_path, monkeypatch):
    ws = tmp_path / "ws"
    m = _analyze(FIX / "templates" / "lo_focus.pptx", ws)
    (TemplateWorkspace.open(m.template_id, ws).dir / "title.txt").write_text("Фокус\n", encoding="utf-8")
    tid, ch = fonted_template(m.template_id, workspace_root=ws, render=False)
    assert tid != m.template_id and ch.apply
    d = ws / "templates" / tid
    pair = ch.heading if ch.heading == ch.body else f"{ch.heading} / {ch.body}"
    assert (d / "title.txt").read_text(encoding="utf-8").strip() == f"Фокус · шрифты {pair}"
    info = fonts_info(tid, ws)
    assert info["base_template_id"] == m.template_id and info["choice"]["reason"] == ch.reason
    # разбор копии видит новые шрифты, стандартных в тексте не осталось
    m2 = TemplateManifest.model_validate_json((d / "manifest.json").read_text(encoding="utf-8"))
    fams = {f.family for f in m2.tokens.typography.families}
    assert {ch.heading, ch.body} & fams and not {"Arial", "DejaVu Sans"} & {f.family for f in m2.tokens.typography.families if f.source == "text"}
    with zipfile.ZipFile(d / "source.pptx") as z:
        theme = z.read("ppt/theme/theme1.xml").decode()
    assert f'<a:latin typeface="{ch.heading}"' in theme and f'<a:latin typeface="{ch.body}"' in theme
    assert set(embedded_families(d / "source.pptx")) == {ch.heading, ch.body}
    # копия — уже с подобранными шрифтами: повторный подбор её не трогает
    assert fonted_template(tid, workspace_root=ws, render=False)[0] == tid
    assert pick_for_template(tid, workspace_root=ws).reason == ch.reason
    # кэш: без нового применения и разбора
    calls = []
    monkeypatch.setattr(fpk, "apply_fonts", lambda *a, **k: calls.append(1))
    import verstka.analysis.manifest as man

    monkeypatch.setattr(man, "analyze_template", lambda *a, **k: calls.append(2))
    assert fonted_template(m.template_id, workspace_root=ws, render=False)[0] == tid
    assert calls == []
    assert fonted_template(m.template_id, workspace_root=ws, enabled=False)[0] == m.template_id
    assert fonted_template(m.template_id, workspace_root=ws, brand_fonts=True)[0] == m.template_id


def test_html_export_carries_picked_fonts(tmp_path):
    from verstka.export.html import _font_face

    css = _font_face({"Manrope", "Arial"})
    assert "font-family:'Manrope';font-weight:700" in css and "font-family:'Play'" in css and "Arial" not in css


@pytest.mark.skipif(not Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists() and not __import__("shutil").which("soffice"), reason="нет LibreOffice")
def test_libreoffice_sets_embedded_font(tmp_path):
    """LibreOffice набирает PDF встроенным шрифтом (на macOS его нельзя указать папкой) — настоящий Unbounded."""
    import shutil
    import subprocess

    from verstka.ingest.render import pptx_to_pdf

    if not shutil.which("pdffonts"):
        pytest.skip("нет pdffonts")
    deck = _deck(tmp_path / "u.pptx", ["Unbounded"])
    embed_fonts(deck)
    pdf = pptx_to_pdf(deck, tmp_path / "pdf")
    out = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    assert "Unbounded" in out


def test_roles_placeholders_and_size(tmp_path):
    """Заголовок-заполнитель — шрифт заголовков; текст в заполнителе 32 пт — шрифт текста (не «крупный — значит
    заголовок»); надпись 44 пт вне заполнителя — шрифт заголовков; шрифт дизайнера (Georgia) не трогается."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Итоги года"
    s.placeholders[1].text = "Выручка выросла на 20 %"
    for shp, size in ((s.shapes.title, 40), (s.placeholders[1], 32)):
        r = shp.text_frame.paragraphs[0].runs[0]
        r.font.name, r.font.size = "Arial", Pt(size)
    for text, size, fam in (("+20 %", 44, "Arial"), ("подпись", 12, "Arial"), ("бренд", 44, "Georgia")):
        tb = s.shapes.add_textbox(Emu(300000), Emu(300000), Emu(3000000), Emu(800000))
        tb.text_frame.text = text
        r = tb.text_frame.paragraphs[0].runs[0]
        r.font.name, r.font.size = fam, Pt(size)
    src = tmp_path / "t.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    ch = fpk.FontChoice(apply=True, heading="Unbounded", body="Onest")
    fpk.apply_fonts(src, tmp_path / "o.pptx", m, ch)
    got = {}
    for sh in Presentation(str(tmp_path / "o.pptx")).slides[0].shapes:
        r = sh.text_frame.paragraphs[0].runs[0]
        got[sh.text_frame.text] = r.font.name
    assert got == {"Итоги года": "Unbounded", "Выручка выросла на 20 %": "Onest", "+20 %": "Unbounded", "подпись": "Onest", "бренд": "Georgia"}
