"""Rules found on the LibreOffice and Office templates of the third-party corpus (round 7): grounds read where the text
stands, colours of the template for text and charts on grounds they were not chosen for, fonts and language of the
finished file, covers of templates without a title placeholder, margins of templates whose slides are empty.

* a gradient is measured where the box is (a light top, a glow nearly clear far from its focus);
* a light brand colour on white deepens in its own hue for a chart; stock theme accents are no brand;
* the parts of a pie on a dark ground step toward white;
* symbol fonts get no stand-in; Impact is a condensed display face;
* Russian runs are tagged ru-RU and text highlights are taken off a finished deck;
* a heading's letter-spacing counts when its lines are broken;
* a first slide set in display type is the cover; a slide of empty placeholders sets the margins by them.
"""

from __future__ import annotations

import colorsys
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from verstka.schemas.common import Bbox, contrast_ratio

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _grad(stops: list[tuple[int, str, int]], lin: str = '<a:lin ang="5400000" scaled="0"/>') -> etree._Element:
    gs = "".join(f'<a:gs pos="{p}"><a:srgbClr val="{c}"><a:alpha val="{a}"/></a:srgbClr></a:gs>' for p, c, a in stops)
    return etree.fromstring(f'<a:gradFill xmlns:a="{A}"><a:gsLst>{gs}</a:gsLst>{lin}</a:gradFill>')


def test_a_gradient_is_read_where_the_box_stands():
    from verstka.rendering.charts import _grad_paint

    g = _grad([(0, "E0F4FA", 100000), (100000, "0B2A4A", 100000)])  # light at the top, navy at the bottom
    top = _grad_paint(g, {}, (0.0, 0.0, 1.0, 0.2), (1000, 1000))
    bottom = _grad_paint(g, {}, (0.0, 0.8, 1.0, 1.0), (1000, 1000))
    assert contrast_ratio(top[0], "FFFFFF") < 1.5 and contrast_ratio(bottom[0], "FFFFFF") > 8.0
    # a white glow fading from its focus at the bottom is nearly clear at the top: no ground there
    glow = _grad([(0, "FFFFFF", 90000), (75000, "FFFFFF", 20000)], '<a:path path="circle"><a:fillToRect l="50000" t="85000" r="50000" b="15000"/></a:path>')
    assert _grad_paint(glow, {}, (0.1, 0.05, 0.9, 0.2), (1000, 600))[1] < 0.5


def test_a_light_brand_colour_deepens_in_its_own_hue():
    from verstka.rendering.charts import shade_to

    got = shade_to("FFDE59", "FFFFFF", 2.6)  # a honey yellow on white
    assert got is not None and contrast_ratio(got, "FFFFFF") >= 2.6
    h0 = colorsys.rgb_to_hls(*(int("FFDE59"[i:i + 2], 16) / 255 for i in (0, 2, 4)))[0]
    h1 = colorsys.rgb_to_hls(*(int(got[i:i + 2], 16) / 255 for i in (0, 2, 4)))[0]
    assert abs(h0 - h1) < 0.02  # gold, not the theme's green


def test_stock_theme_accents_are_no_brand():
    from verstka.analysis.colors import is_stock_theme_color

    assert is_stock_theme_color("18A303") and is_stock_theme_color("4472c4")
    assert not is_stock_theme_color("0077FF")


def test_pie_parts_on_a_dark_ground_step_toward_white():
    from verstka.rendering.charts import pie_palette

    pal = pie_palette([60, 25, 15, 10, 5], "0077FF", "000000")
    assert pal[0] == "0077FF"
    cr = [contrast_ratio(c, "000000") for c in pal]
    assert all(c >= 2.0 for c in cr) and cr == sorted(cr)  # every part off the ground, lighter with its rank


def test_symbol_fonts_get_no_stand_in_and_impact_is_a_display_face():
    from verstka.rendering.cyrillic import supports_cyrillic
    from verstka.rendering.fonts import font_class

    assert supports_cyrillic("Wingdings 3") is None
    assert supports_cyrillic("Symbol") is None
    assert font_class("Impact") == "display_condensed"
    assert font_class("Gill Sans MT") == "sans"


def _deck(tmp_path, texts: list[str], highlight: bool = False):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for t in texts:
        tb = s.shapes.add_textbox(0, 0, Inches(4), Inches(1))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = t
        r.font.size = Pt(18)
        rpr = r._r.get_or_add_rPr()
        rpr.set("lang", "en-US")
        if highlight:
            hl = etree.SubElement(rpr, f"{{{A}}}highlight")
            etree.SubElement(hl, f"{{{A}}}srgbClr").set("val", "FFFFFF")
    p = tmp_path / "deck.pptx"
    prs.save(p)
    return p


def test_russian_runs_are_tagged_and_highlights_taken_off(tmp_path):
    from verstka.rendering.cyrillic import mark_russian, strip_highlights

    p = _deck(tmp_path, ["Выручка за 30 дней", "VK Tech"], highlight=True)
    assert mark_russian(p) == 1
    assert strip_highlights(p) >= 1
    with zipfile.ZipFile(p) as z:
        xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "highlight" not in xml
    runs = [(r.text, r.font._rPr.get("lang")) for sh in Presentation(p).slides[0].shapes for r in sh.text_frame.paragraphs[0].runs]
    assert ("Выручка за 30 дней", "ru-RU") in runs and ("VK Tech", "en-US") in runs


def test_letter_spacing_counts_when_a_heading_is_broken():
    from verstka.matching.scorer import display_lines

    text = "КОФЕЙНЯ «ТОЧКА КОФЕ»"
    loose = display_lines(text, "Arial", 20.0, True, 260.0)
    tracked = display_lines(text, "Arial", 20.0, True, 260.0, spc_pt=8.0)
    assert len(tracked) > len(loose)
    assert all("ТОЧКА" not in ln or "КОФЕ" in ln for ln in tracked)  # the bound name stays whole


def test_a_first_slide_set_in_display_type_is_the_cover():
    from verstka.analysis.kinds import heuristic_kind
    from verstka.analysis.shapes import ParagraphInfo, RunInfo, ShapeInfo, TextInfo
    from verstka.schemas.common import PatternKind, ShapeKind

    W, H = 12192000, 6858000

    def txt(i: int, t: str, size: float, y: float) -> ShapeInfo:
        return ShapeInfo(id=str(i), name=f"Text {i}", kind=ShapeKind.sp, bbox=Bbox(x=int(0.45 * W), y=int(y * H), w=int(0.4 * W), h=int(0.15 * H)), z=i,
                         text=TextInfo(paragraphs=[ParagraphInfo(text=t, runs=[RunInfo(text=t, size_pt=size)])]))

    shapes = [txt(1, "ELEGANT", 96, 0.15), txt(2, "PITCH", 96, 0.35), txt(3, "DECK", 96, 0.55)] + [txt(10 + i, f"Menu {i}", 14, 0.9) for i in range(6)]
    kind, _ = heuristic_kind(shapes, {}, [], 1, 20, W, H)
    assert kind == PatternKind.title


def test_empty_placeholders_set_the_margins_of_an_empty_template():
    from verstka.analysis.shapes import ShapeInfo
    from verstka.analysis.spacing import compute_spacing
    from verstka.schemas.common import ShapeKind

    W, H = 12192000, 6858000

    def ph(i: str, t: str, x: float, y: float, w: float, h: float) -> ShapeInfo:
        return ShapeInfo(id=i, name=i, kind=ShapeKind.sp, bbox=Bbox(x=int(x * W), y=int(y * H), w=int(w * W), h=int(h * H)), z=int(i), is_placeholder=True, ph_type=t)

    slide = [ph("1", "title", 0.05, 0.04, 0.9, 0.12), ph("2", "body", 0.05, 0.22, 0.9, 0.68), ph("3", "sldNum", 0.8, 0.94, 0.15, 0.05)]
    sp = compute_spacing({1: slide, 2: slide}, {}, W, H)
    assert sp.safe_area.x <= 0.06 and sp.safe_area.y <= 0.06
