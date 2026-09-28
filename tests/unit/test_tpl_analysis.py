"""Template analysis for third-party templates (TEMPLATE_FIX_PLAN, owner A): the derived type ladder of sparse
templates, chart/table text caps, English sample copy and kickers, wordmarks in the title placeholder, grounds painted
by layouts/masters, free room left by template art, cover mock-ups, stock theme accents, the safe-area foot, and font
metrics with caps/tracking."""

from __future__ import annotations

import copy
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.chrome import detect_chrome, is_kicker
from verstka.analysis.colors import ColorSample, dominant_saturated_hex, has_saturated_drawn, is_drawn, is_stock_theme_color
from verstka.analysis.components import data_text_cap
from verstka.analysis.ground import art_boxes, gradient_mean_hex, largest_free_share, painted_layers, pattern_fill_hex, slide_ground
from verstka.analysis.patterns import dedupe_patterns, mockup_boxes, mockup_on_layout
from verstka.analysis.roles import heuristic_roles, low_display_heading, wordmark_heading
from verstka.analysis.shapes import ParagraphInfo, RunInfo, ShapeInfo, SlideContext, TextInfo, extract_shapes, looks_like_placeholder, slide_background, slide_family
from verstka.analysis.spacing import extend_safe_bottom
from verstka.analysis.theme import ThemeResolver
from verstka.analysis.typography import build_type_scale, derived_size_ladder, is_sparse_scale
from verstka.ingest.package import PptxPackage
from verstka.matching.scorer import _cover_like
from verstka.rendering import fonts
from verstka.schemas.common import Bbox, BboxFrac, Family, PatternKind, ShapeKind, SlotRole
from verstka.schemas.template import ChromeElement, ClassificationTrace, Pattern, SignalVote, Slot, SlideSize, Spacing, TemplateManifest, Tokens, Typography

EMU_PT = 12700
W169, H169 = 12192000, 6858000


# ---------------------------------------------------------------------------- helpers


def _text_shape(sid: str, text: str, x: float, y: float, w: float, h: float, size: float, ph: str | None = None, W: int = W169, H: int = H169, align: str | None = None) -> ShapeInfo:
    runs = [RunInfo(text=text, size_pt=size, font="Arial")]
    ti = TextInfo(paragraphs=[ParagraphInfo(text=text, runs=runs, align=align)])
    return ShapeInfo(id=sid, name=sid, kind=ShapeKind.sp, bbox=Bbox(x=int(x * W), y=int(y * H), w=int(w * W), h=int(h * H)), z=int(sid) if sid.isdigit() else 1, is_placeholder=ph is not None, ph_type=ph, text=ti)


def _pattern(pid: str, slide: int, kind: PatternKind = PatternKind.bullets, layout: str | None = None, free: float | None = None, quality: float = 1.0) -> Pattern:
    trace = ClassificationTrace(kind=kind, heuristic=SignalVote(kind=kind, confidence=1.0))
    slots = [
        Slot(id="title_1", role=SlotRole.title, shape_id="2", bbox=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.15)),
        Slot(id="bullet_list_1", role=SlotRole.bullet_list, shape_id="3", bbox=BboxFrac(x=0.05, y=0.25, w=0.9, h=0.6)),
    ]
    return Pattern(id=pid, source_slide=slide, kind=kind, family=Family.light, slots=slots, quality=quality, classification=trace, layout_part=layout, free_share=free)


def _png(path: Path, color: tuple[int, int, int], size=(64, 36)) -> Path:
    Image.new("RGB", size, color).save(path)
    return path


# ---------------------------------------------------------------------------- T02 ladder / T09 caps


def test_sparse_scale_gets_a_ladder_relative_to_the_slide():
    h = int(6.2 * 914400)  # an 11-inch LibreOffice slide: 446 pt high, two placeholder sizes
    assert is_sparse_scale([24.0, 33.0], 24.0, h)
    ladder = derived_size_ladder([24.0, 33.0], 24.0, h)
    hpt = h / EMU_PT
    assert ladder and min(ladder) == pytest.approx(round(0.018 * hpt * 2) / 2)
    assert max(ladder) >= 0.18 * hpt  # reaches display sizes
    assert all(abs(s * 2 - round(s * 2)) < 1e-9 for s in ladder)  # 0.5 pt steps
    assert all(not abs(s - t) <= 0.04 * t for s in ladder for t in (24.0, 33.0))  # the template's own sizes stand
    assert len([s for s in ladder if s <= 0.031 * hpt]) >= 3  # something small to shrink dense text into


def test_dataset_like_scales_are_not_sparse():
    lct = [10.5, 12.0, 13.0, 14.0, 15.0, 16.0, 18.67, 20.0, 24.0, 36.0]
    assert not is_sparse_scale(lct, 14.0, 6858000)
    assert derived_size_ladder(lct, 14.0, 6858000) == []
    vk = [4.14, 6.0, 7.0, 8.12, 9.0, 9.92, 11.0, 12.0, 13.22, 14.0, 15.0, 16.0, 16.88, 18.0, 20.0, 24.0, 27.05, 29.0, 32.0, 36.0, 47.0, 48.0, 54.0, 96.0, 166.0]
    assert derived_size_ladder(vk, 9.0, 5143500) == []


def test_body_placeholder_default_too_large_makes_a_scale_sparse():
    sizes = [12.0, 13.0, 14.0, 16.0, 18.0, 20.0, 24.0, 28.0, 32.0, 36.0]
    assert not is_sparse_scale(sizes, 16.0, 6858000)
    assert is_sparse_scale(sizes, 32.0, 6858000)  # a 32 pt body on a 7.5-inch slide is not a dense-content size


def test_chart_and_table_text_capped_at_2_6_percent_of_the_slide():
    h = int(6.2 * 914400)
    sizes = [24.0, 33.0] + derived_size_ladder([24.0, 33.0], 24.0, h)
    cap = data_text_cap(sizes, h)
    assert cap <= 0.026 * h / EMU_PT and cap in sizes
    assert data_text_cap([7.0, 9.0, 12.0], 5143500) == 9.0  # VK: the 7 pt chart style stays under its cap


def test_h1_comes_from_headings_not_wordmarks_or_parked_body_text():
    slides = {}
    for i in range(1, 5):
        slides[i] = [
            _text_shape("2", f"Our target {i}", 0.07, 0.09, 0.4, 0.1, 58.5, ph="title"),
            _text_shape("3", "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor " * 2, 0.07, 0.25, 0.4, 0.5, 20.0),
        ]
    # a cover whose title placeholder holds the wordmark, and a slide with a body paragraph parked in its title
    slides[5] = [_text_shape("51", "mybrand.", 0.07, 0.1, 0.8, 0.16, 37.5, ph="title"), _text_shape("46", "MyBrand Pitch Deck", 0.09, 0.43, 0.42, 0.33, 112.5, ph="body")]
    slides[6] = [_text_shape("2", "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore " * 3, 0.07, 0.1, 0.8, 0.6, 20.0, ph="title")]
    typo = build_type_scale(slides, {}, slide_h=int(11.25 * 914400))
    assert typo.size_for("h1") == 58.5 and typo.size_for("body") == 20.0


# ---------------------------------------------------------------------------- T13 sample copy / kickers


@pytest.mark.parametrize("text", ["DEMO SLIDE", "Title Text Demo", "Slide Title Goes Here 2", "Subtitle goes here · Speaker Name",
                                  "Click to edit Master title style", "Put Your Awesome Word In Here", "IMAGE", "Your Logo", "Insert picture", "Sample text"])
def test_english_sample_copy_is_placeholder(text):
    assert looks_like_placeholder(text)


@pytest.mark.parametrize("text", ["Demo Day 2026", "The demo shows growth", "Company Presentation 2026", "Выручка выросла на 12%", "Imagine the future"])
def test_real_headings_are_not_placeholder(text):
    assert not looks_like_placeholder(text)


def test_repeated_sample_kicker_is_not_chrome_but_page_numbers_are():
    slides = {}
    for i in range(1, 6):
        slides[i] = [
            _text_shape("2", f"Heading number {i}", 0.14, 0.20, 0.34, 0.12, 60.0),
            _text_shape("3", "DEMO SLIDE", 0.14, 0.175, 0.14, 0.03, 20.0),
            _text_shape("4", "SECTION", 0.14, 0.34, 0.14, 0.03, 20.0),  # a label right under the heading: a kicker
            _text_shape("5", str(i), 0.02, 0.04, 0.04, 0.04, 25.0),
            _text_shape("6", "ACME", 0.05, 0.93, 0.2, 0.04, 10.0),
        ]
    chrome = detect_chrome(slides, W169, H169)
    texts = {c.text for c in chrome}
    assert "DEMO SLIDE" not in texts and "SECTION" not in texts
    assert "ACME" in texts
    assert any(c.signature.endswith("txt:num") for c in chrome)  # page numbers stay chrome (renumbered)


def test_is_kicker_needs_the_content_band_and_the_heading_columns():
    title = Bbox(x=int(0.14 * W169), y=int(0.2 * H169), w=int(0.34 * W169), h=int(0.12 * H169))
    assert is_kicker(_text_shape("3", "KICKER", 0.14, 0.15, 0.1, 0.03, 12.0), title, H169)
    assert not is_kicker(_text_shape("3", "LOGO", 0.8, 0.15, 0.1, 0.03, 12.0), title, H169)  # other columns
    assert not is_kicker(_text_shape("3", "TOP", 0.14, 0.02, 0.1, 0.03, 12.0), title, H169)  # above the band


# ---------------------------------------------------------------------------- T11 wordmark


def test_wordmark_in_title_placeholder_yields_the_heading_to_the_big_text():
    mark = _text_shape("51", "mybrand.", 0.07, 0.11, 0.86, 0.16, 37.5, ph="title")
    big = _text_shape("46", "MyBrand Pitch Deck", 0.09, 0.43, 0.42, 0.33, 112.5, ph="body")
    sub = _text_shape("47", "Bed & Breakfast Booking Platform", 0.13, 0.79, 0.32, 0.04, 27.0)
    shapes = [big, sub, mark]
    assert wordmark_heading(mark, shapes, H169) is big
    roles = heuristic_roles(shapes, [], set(), Typography(), W169, H169)
    assert roles["46"] == SlotRole.title and roles["51"] == SlotRole.chrome


def test_two_word_label_over_a_display_word_is_not_kept_as_chrome():
    label = _text_shape("144", "Put Your", 0.46, 0.43, 0.08, 0.04, 27.5, ph="title")
    big = _text_shape("145", "Awesome Word", 0.25, 0.45, 0.51, 0.14, 95.0)
    roles = heuristic_roles([label, big], [], set(), Typography(), W169, H169)
    assert roles["145"] == SlotRole.title and roles.get("144") != SlotRole.chrome


def test_a_real_title_placeholder_keeps_the_title_role():
    t = _text_shape("2", "Итоги квартала", 0.05, 0.05, 0.9, 0.15, 40.0, ph="title")
    kpi = _text_shape("3", "120%", 0.05, 0.4, 0.3, 0.2, 120.0)  # a big figure never takes the heading
    assert wordmark_heading(t, [t, kpi], H169) is None
    roles = heuristic_roles([t, kpi], [], set(), Typography(), W169, H169)
    assert roles["2"] == SlotRole.title


def test_low_display_heading_on_a_cover_is_the_title():
    # Marketing p1: «Marketing Report» at 185 pt from 0.49 H of a 26.67×15-inch slide, no placeholder, nothing in the
    # top 40 % — the cover's heading, not a body slot
    W, H = 24384000, 13716000
    big = _text_shape("90", "Marketing Report", 0.193, 0.488, 0.578, 0.348, 185.0, W=W, H=H)
    assert low_display_heading([big], set(), H) is big
    roles = heuristic_roles([big], [], set(), Typography(), W, H)
    assert roles["90"] == SlotRole.title
    # a note under it stays what it is
    note = _text_shape("91", "Quarterly results", 0.193, 0.84, 0.4, 0.04, 40.0, W=W, H=H)
    roles = heuristic_roles([big, note], [], set(), Typography(), W, H)
    assert roles["90"] == SlotRole.title and roles.get("91") != SlotRole.title


def test_low_display_heading_rejects_ordinary_low_text():
    W, H = 24384000, 13716000
    # not a display size (< 7 % of the slide height = 75.6 pt here)
    assert low_display_heading([_text_shape("3", "Marketing Report", 0.2, 0.5, 0.5, 0.1, 60.0, W=W, H=H)], set(), H) is None
    # a peer text of a similar size: two big texts, no single heading
    a = _text_shape("3", "Marketing Report", 0.2, 0.45, 0.5, 0.15, 185.0, W=W, H=H)
    b = _text_shape("4", "Annual Review", 0.2, 0.62, 0.5, 0.15, 150.0, W=W, H=H)
    assert low_display_heading([a, b], set(), H) is None
    # a big figure is a KPI, a text starting in the bottom 30 % is a footer, a paragraph is body copy
    assert low_display_heading([_text_shape("5", "120%", 0.2, 0.5, 0.5, 0.2, 185.0, W=W, H=H)], set(), H) is None
    assert low_display_heading([_text_shape("6", "Marketing Report", 0.2, 0.75, 0.5, 0.15, 185.0, W=W, H=H)], set(), H) is None
    assert low_display_heading([_text_shape("7", "Lorem ipsum dolor sit amet " * 4, 0.2, 0.5, 0.5, 0.3, 120.0, W=W, H=H)], set(), H) is None
    # a card member (a repeated group) is never the slide's heading
    assert low_display_heading([a], {"3"}, H) is None


# ---------------------------------------------------------------------------- T07 grounds


def test_fill_helpers_blend_patterns_and_average_gradients():
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

    class R:  # a resolver stand-in: srgb only
        def resolve_color(self, el):
            return el.get("val")

    patt = etree.fromstring(f'<a:pattFill {ns} prst="pct50"><a:fgClr><a:srgbClr val="0000FF"/></a:fgClr><a:bgClr><a:srgbClr val="FFFFFF"/></a:bgClr></a:pattFill>')
    assert pattern_fill_hex(R(), patt) == "8080FF"
    grad = etree.fromstring(f'<a:gradFill {ns}><a:gsLst><a:gs pos="0"><a:srgbClr val="FFFF00"/></a:gs><a:gs pos="100000"><a:srgbClr val="FF8000"/></a:gs></a:gsLst></a:gradFill>')
    assert gradient_mean_hex(R(), grad) == "FFC000"


def _ground_deck(tmp_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    blank = prs.slide_layouts[6]
    # slide 1: a full-bleed blue picture under a white title (MyBrand's cover)
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(str(_png(tmp_path / "blue.png", (24, 102, 245))), 0, 0, Emu(W169), Emu(H169))
    tb = s1.shapes.add_textbox(Emu(600000), Emu(2000000), Emu(8000000), Emu(1200000))
    tb.text_frame.text = "Cover title"
    tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # slide 2: a layout paints a navy full-bleed panel (Midnightblue)
    layout = prs.slide_layouts[5]
    scratch = prs.slides.add_slide(blank)
    rect = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(W169), Emu(H169))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    layout.shapes._spTree.insert(2, copy.deepcopy(rect._element))
    s2 = prs.slides.add_slide(layout)
    s2.shapes.title.text = "Heading on navy"
    # slide 3: a near-white texture picture over the white background only textures the same ground
    s3 = prs.slides.add_slide(blank)
    s3.shapes.add_picture(str(_png(tmp_path / "paper.png", (250, 250, 252))), 0, 0, Emu(W169), Emu(H169))
    # drop the scratch slide
    sld_ids = prs.slides._sldIdLst
    sld_ids.remove(list(sld_ids)[1])
    out = tmp_path / "grounds.pptx"
    prs.save(out)
    return out


def test_grounds_from_full_bleed_pictures_and_layout_panels(tmp_path):
    path = _ground_deck(tmp_path)
    pkg = PptxPackage.open(path)
    parts = pkg.slide_parts
    ctx = SlideContext(pkg, parts[0])
    g = slide_ground(pkg, parts[0], ctx, extract_shapes(pkg, parts[0], ctx))
    assert g.family == Family.dark and g.kind == "image" and g.hex.startswith("18")
    fam, bg = slide_family(pkg, parts[0], ctx, extract_shapes(pkg, parts[0], ctx))
    assert fam == Family.dark
    ctx2 = SlideContext(pkg, parts[1])
    fam2, bg2 = slide_family(pkg, parts[1], ctx2, extract_shapes(pkg, parts[1], ctx2))
    assert fam2 == Family.dark and bg2 == "2C3E50"
    assert any(la.source == "layout" and la.hex == "2C3E50" for la in painted_layers(pkg, parts[1], ctx2))
    ctx3 = SlideContext(pkg, parts[2])
    fam3, bg3 = slide_family(pkg, parts[2], ctx3, extract_shapes(pkg, parts[2], ctx3))
    assert fam3 == Family.light and bg3 == (slide_background(pkg, parts[2], ctx3)[0] or "FFFFFF")


def test_style_chain_merges_a_level_written_twice():
    from verstka.analysis.shapes import _StyleChain

    lst = etree.fromstring('<a:lstStyle xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:lvl1pPr algn="l"><a:defRPr sz="8000" b="1"/></a:lvl1pPr>'
                           '<a:lvl1pPr><a:defRPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:defRPr></a:lvl1pPr></a:lstStyle>')
    chain = _StyleChain(None, [lst], True)
    assert len(chain.def_rpr(0)) == 2  # size from the first, the white colour from the second (LibreOffice merges both)


def test_theme_resolver_reads_pattern_fills(tmp_path):
    path = _ground_deck(tmp_path)
    pkg = PptxPackage.open(path)
    r = ThemeResolver(pkg, pkg.master_parts[0])
    el = etree.fromstring('<p:spPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                          '<a:pattFill prst="pct25"><a:fgClr><a:srgbClr val="000000"/></a:fgClr><a:bgClr><a:srgbClr val="FFFFFF"/></a:bgClr></a:pattFill></p:spPr>')
    assert r.resolve_fill(el) == "BFBFBF"


# ---------------------------------------------------------------------------- T06 free room / dedupe


def test_largest_free_rectangle_share():
    area = Bbox(x=0, y=0, w=6400, h=3600)
    assert largest_free_share([], area) == 1.0
    assert largest_free_share([Bbox(x=0, y=0, w=3200, h=3600)], area) == pytest.approx(0.5, abs=0.02)
    assert largest_free_share([Bbox(x=-100, y=-100, w=7000, h=4000)], area) == 0.0


def test_dedupe_keeps_the_freer_of_two_equal_content_samples():
    cluttered = _pattern("p2", 2, layout="ppt/slideLayouts/slideLayout2.xml", free=0.4)
    clean = _pattern("p5", 5, layout="ppt/slideLayouts/slideLayout5.xml", free=0.95, quality=0.9)
    kept = dedupe_patterns([cluttered, clean])
    assert [p.id for p in kept] == ["p5"]
    # within 10 %: the usual quality order
    kept = dedupe_patterns([_pattern("p2", 2, layout="a", free=0.9), _pattern("p5", 5, layout="b", free=0.95, quality=0.9)])
    assert [p.id for p in kept] == ["p2"]
    # a divider keeps its art: no free-share preference for bookends
    kept = dedupe_patterns([_pattern("p2", 2, PatternKind.section, "a", 0.3), _pattern("p5", 5, PatternKind.section, "b", 0.9, 0.9)])
    assert [p.id for p in kept] == ["p2"]


# ---------------------------------------------------------------------------- T15 covers


def test_phone_mockup_and_empty_photo_frames_are_found():
    frame = ShapeInfo(id="5", name="phone", kind=ShapeKind.sp, bbox=Bbox(x=int(0.667 * W169), y=int(0.12 * H169), w=int(0.233 * W169), h=int(0.8 * H169)), z=5, fill_hex="15121F")
    clock = _text_shape("7", "9:41", 0.69, 0.16, 0.06, 0.04, 10.0)
    title = _text_shape("2", "Flowly App Pitch", 0.045, 0.27, 0.5, 0.27, 48.0, ph="title")
    boxes = mockup_boxes([title], W169, H169, layout_shapes=[frame, clock])
    assert len(boxes) == 1 and boxes[0].x == pytest.approx(0.667, abs=0.01)
    assert mockup_on_layout(boxes, [frame, clock], W169, H169)  # drawn by the layout: it stays on the cover
    assert not mockup_on_layout(boxes, [clock], W169, H169)
    # a broad illustration is not a device
    art = ShapeInfo(id="9", name="art", kind=ShapeKind.pic, bbox=Bbox(x=0, y=0, w=int(0.6 * W169), h=int(0.9 * H169)), z=9, image_part="ppt/media/x.png")
    assert mockup_boxes([title, art], W169, H169) == []


def _manifest(patterns: list[Pattern], h: int = 15 * 914400) -> TemplateManifest:
    return TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=int(26.67 * 914400), h=h), tokens=Tokens(), patterns=patterns)


def test_display_sized_cover_heading_is_cover_like():
    cover = _pattern("p1", 1, PatternKind.title)
    m = _manifest([cover])
    assert _cover_like(cover, m, 96.0)  # 96 pt on a 15-inch slide (1080 pt): a display size
    assert not _cover_like(cover, m, 40.0)
    first = _pattern("p3", 3, PatternKind.title).model_copy(update={"title_ph": "ctrTitle"})
    later = _pattern("p7", 7, PatternKind.title).model_copy(update={"title_ph": "ctrTitle"})
    m2 = _manifest([first, later])
    assert _cover_like(first, m2, 40.0) and not _cover_like(later, m2, 40.0)


# ---------------------------------------------------------------------------- T18 colours


def test_stock_theme_accents_and_picture_colours():
    assert is_stock_theme_color("18a303") and is_stock_theme_color("4472C4") and not is_stock_theme_color("0077FF")
    samples = [ColorSample("A23735", "fill", 1.0), ColorSample("FFFFFF", "background", 100.0)]
    assert is_drawn("A33735", samples) and not is_drawn("18A303", samples)
    assert has_saturated_drawn(samples) and not has_saturated_drawn([ColorSample("333333", "text", 1.0)])
    import numpy as np

    px = np.array([[240, 230, 210]] * 80 + [[162, 55, 53]] * 20, dtype=np.float32)
    assert dominant_saturated_hex(px) == "A23735"
    assert dominant_saturated_hex(np.array([[200, 200, 200]] * 50, dtype=np.float32)) is None


# ---------------------------------------------------------------------------- T19 safe area


def test_safe_area_foot_extends_to_content_bodies_never_shrinks():
    sp = Spacing(safe_area=BboxFrac(x=0.068, y=0.073, w=0.863, h=0.66))
    pats = [_pattern("p2", 2), _pattern("p4", 4, PatternKind.chart)]
    pats[0].slots[1].bbox = BboxFrac(x=0.07, y=0.29, w=0.86, h=0.59)  # body to 0.88
    foot = [ChromeElement(signature="n", bbox=BboxFrac(x=0.068, y=0.933, w=0.045, h=0.04), share=0.6, kind="text", sample_slide=2),
            ChromeElement(signature="edge", bbox=BboxFrac(x=0.013, y=0.822, w=0.037, h=0.178), share=0.1, kind="pic", sample_slide=1, source="background")]
    out = extend_safe_bottom(sp, pats, {}, {}, foot, W169, H169)
    assert out.safe_area.y2 == pytest.approx(0.88, abs=0.005)
    deep = Spacing(safe_area=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.88))
    assert extend_safe_bottom(deep, pats, {}, {}, foot, W169, H169).safe_area.y2 == pytest.approx(0.93, abs=0.001)


# ---------------------------------------------------------------------------- C6 / T20 fonts


def test_caps_and_tracking_widen_the_measure():
    base = fonts.text_width_pt("Показатель", "Open Sans", 20)
    caps = fonts.text_width_pt("Показатель", "Open Sans", 20, caps=True)
    tracked = fonts.text_width_pt("Показатель", "Open Sans", 20, caps=True, spc_pt=3.0)
    assert caps > base and tracked == pytest.approx(caps + 3.0 * len("Показатель"))
    assert fonts.wrap_lines("Показатель выручки", "Arial", 20, False, 120) != fonts.wrap_lines("Показатель выручки", "Arial", 20, False, 120, caps=True, spc_pt=3.0)


def test_dataset_families_keep_the_play_factor_model():
    for fam in ("Play", "Montserrat", "Montserrat Medium", "Arial", "Consolas", "Poppins Light", "Open Sans"):
        assert fonts.is_pinned(fam) and fonts.is_measured(fam)
        want = fonts._font(False).getlength("Выручка 120 000 ₽") / fonts._MEASURE_PX * 20 * fonts.width_factor(fam)
        assert fonts.text_width_pt("Выручка 120 000 ₽", fam, 20) == pytest.approx(want)


# ---------------------------------------------------------------------------- T07 bare slides / T08 theme fonts / T06 drawn boxes


def test_a_slide_without_any_background_is_white_not_its_art(tmp_path):
    """No p:bg on slide, layout or master and nothing full-bleed: the slide is white (as PowerPoint paints it), not
    the render's median — Focus's coloured triangles must not become the ground."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    p_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    for root in [prs.slide_master._element] + [lay._element for lay in prs.slide_layouts]:
        for bg in root.iter(p_ns + "bg"):
            bg.getparent().remove(bg)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    art = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, 0, 0, Emu(int(0.45 * W169)), Emu(int(0.8 * H169)))
    art.fill.solid()
    art.fill.fore_color.rgb = RGBColor(0xE8, 0x1E, 0x1E)
    path = tmp_path / "bare.pptx"
    prs.save(path)
    pkg = PptxPackage.open(path)
    part = pkg.slide_parts[0]
    ctx = SlideContext(pkg, part)
    assert slide_background(pkg, part, ctx) == (None, None)
    red = _png(tmp_path / "render.png", (232, 30, 30))  # a render dominated by the art
    g = slide_ground(pkg, part, ctx, extract_shapes(pkg, part, ctx), str(red))
    assert (g.family, g.hex, g.source) == (Family.light, "FFFFFF", "default")


def test_theme_fonts_are_listed_after_the_text_families(tmp_path):
    from verstka.analysis.manifest import analyze_template

    prs = Presentation()  # the default theme: Calibri Light / Calibri
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Emu(600000), Emu(600000 + i * 10000), Emu(8000000), Emu(1200000))
        tb.text_frame.text = f"Заголовок слайда номер {i + 1}"
        run = tb.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size = "Georgia", Pt(32)
    path = tmp_path / "georgia.pptx"
    prs.save(path)
    m = analyze_template(path, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    fams = m.tokens.typography.families
    assert fams[0].family == "Georgia" and fams[0].source == "text" and fams[0].weight > 0
    theme = [f for f in fams if f.source == "theme"]
    assert {f.family for f in theme} >= {"Calibri"} and all(f.weight == 0 for f in theme)
    assert fams.index(theme[0]) > max(i for i, f in enumerate(fams) if f.source == "text")


def test_drawn_box_follows_the_text_alignment():
    from verstka.analysis.manifest import _drawn_box

    left = _text_shape("2", "Короткий", 0.05, 0.05, 0.9, 0.1, 20.0)
    centred = _text_shape("3", "Короткий", 0.05, 0.05, 0.9, 0.1, 20.0, align="ctr")
    right = _text_shape("4", "Короткий", 0.05, 0.05, 0.9, 0.1, 20.0, align="r")
    bl, bc, br = (_drawn_box(s, W169, H169) for s in (left, centred, right))
    assert bl.w == pytest.approx(bc.w) == pytest.approx(br.w) and bl.w < 0.3
    assert bl.x == pytest.approx(0.05, abs=1e-3)
    assert bc.x + bc.w / 2 == pytest.approx(0.5, abs=1e-3)  # a centred heading inks the middle, not the left end
    assert br.x2 == pytest.approx(0.95, abs=1e-3)


@pytest.mark.skipif(not __import__("shutil").which("fc-match"), reason="fontconfig is needed to name a substitute")
def test_a_missing_font_is_a_substitute_note_not_a_reading_warning(tmp_path):
    from verstka.analysis.manifest import analyze_template

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Emu(600000), Emu(600000), Emu(8000000), Emu(1200000))
        tb.text_frame.text = f"Слайд {i + 1}: текст шрифтом, которого нет"
        run = tb.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size = "Verstka Missing Grotesk", Pt(28)
    path = tmp_path / "missing.pptx"
    prs.save(path)
    m = analyze_template(path, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    assert "Verstka Missing Grotesk" in m.font_substitutes and m.font_substitutes["Verstka Missing Grotesk"]
    assert not any("Verstka Missing Grotesk" in w for w in m.warnings)  # the narrator reads warnings as parse failures


# ---------------------------------------------------------------------------- T07 layer order, veils, freeforms / T15 mock-ups on the layout


def _bare_deck(tmp_path: Path, name: str):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W169), Emu(H169)
    return prs, prs.slides.add_slide(prs.slide_layouts[6]), tmp_path / name


def _slide_ground_of(path: Path):
    pkg = PptxPackage.open(path)
    part = pkg.slide_parts[0]
    ctx = SlideContext(pkg, part)
    return slide_ground(pkg, part, ctx, extract_shapes(pkg, part, ctx)), slide_background(pkg, part, ctx)


def test_a_full_bleed_freeform_is_art_not_the_ground(tmp_path):
    """A freeform's box is not what it paints (rays, a curve, a blot spanning the slide): never the slide's ground."""
    prs, s, path = _bare_deck(tmp_path, "freeform.pptx")
    ff = s.shapes.build_freeform(0, 0, scale=1.0)
    ff.add_line_segments([(W169, 0), (0, H169)], close=True)  # a triangle whose box is the whole slide
    tri = ff.convert_to_shape()
    tri.fill.solid()
    tri.fill.fore_color.rgb = RGBColor(0xC0, 0x10, 0x20)
    prs.save(path)
    g, (bg_hex, _) = _slide_ground_of(path)
    assert g.hex != "C01020" and g.source in ("bg", "default") and g.family == Family.light
    assert g.hex == (bg_hex or "FFFFFF")


def test_the_topmost_full_bleed_paint_is_the_ground(tmp_path):
    prs, s, path = _bare_deck(tmp_path, "stack.pptx")
    for rgb in ((0x10, 0x20, 0x50), (0xF2, 0x8C, 0x28)):  # navy under orange
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(W169), Emu(H169))
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor(*rgb)
    prs.save(path)
    g, _ = _slide_ground_of(path)
    assert g.hex == "F28C28" and g.family == Family.light and g.source == "layer:slide"


def test_a_semi_opaque_veil_is_blended_with_the_picture_under_it(tmp_path):
    prs, s, path = _bare_deck(tmp_path, "veil.pptx")
    s.shapes.add_picture(str(_png(tmp_path / "blue.png", (24, 102, 245))), 0, 0, Emu(W169), Emu(H169))
    veil = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(W169), Emu(H169))
    veil.fill.solid()
    veil.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    srgb = veil.fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    etree.SubElement(srgb, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha").set("val", "60000")
    prs.save(path)
    g, _ = _slide_ground_of(path)
    r, gg, b = (int(g.hex[k:k + 2], 16) for k in (0, 2, 4))
    # 60 % white over the blue picture: neither the picture's blue nor the veil's white
    assert (r, gg, b) == (pytest.approx(0.4 * 24 + 0.6 * 255, abs=2), pytest.approx(0.4 * 102 + 0.6 * 255, abs=2), pytest.approx(0.4 * 245 + 0.6 * 255, abs=2))
    assert g.family == Family.light


def test_a_full_bleed_layout_ground_does_not_make_a_sample_frame_layout_drawn():
    """An empty photo half of the sample (the renderer removes it) stays a sample-drawn frame although the layout's
    full-bleed ground picture covers it too; a frame the layout really draws (a phone of about its size) is
    layout-drawn."""
    half = BboxFrac(x=0.5, y=0.0, w=0.5, h=1.0)
    ground = ShapeInfo(id="1", name="bg", kind=ShapeKind.pic, bbox=Bbox(x=0, y=0, w=W169, h=H169), z=1, image_part="ppt/media/bg.png")
    assert not mockup_on_layout([half], [ground], W169, H169)
    frame = ShapeInfo(id="2", name="frame", kind=ShapeKind.sp, bbox=Bbox(x=int(0.49 * W169), y=0, w=int(0.51 * W169), h=H169), z=2, fill_hex="202020")
    assert mockup_on_layout([half], [ground, frame], W169, H169)
    placeholder = copy.copy(frame)
    placeholder.is_placeholder, placeholder.ph_type = True, "pic"
    assert not mockup_on_layout([half], [ground, placeholder], W169, H169)  # a layout's prompt, not art


def test_a_layout_drawn_mockup_costs_a_cover_more_than_a_sample_drawn_one():
    from verstka.matching.scorer import score_pattern
    from verstka.planning.strategies import get_strategy
    from verstka.schemas.outline import OutlineSlide

    slide = OutlineSlide(id="c", kind=PatternKind.title, headline="Кофейня у дома")
    base = _pattern("p1", 1, PatternKind.title).model_copy(update={"slots": [Slot(id="title_1", role=SlotRole.title, shape_id="2", bbox=BboxFrac(x=0.05, y=0.3, w=0.5, h=0.2))]})
    phone = [BboxFrac(x=0.67, y=0.12, w=0.23, h=0.8)]
    m = _manifest([base])
    strategy = get_strategy("structured")
    clean = score_pattern(slide, base, m, strategy).score
    on_sample = score_pattern(slide, base.model_copy(update={"mockup_boxes": phone}), m, strategy)
    on_layout = score_pattern(slide, base.model_copy(update={"mockup_boxes": phone, "mockup_on_layout": True}), m, strategy)
    assert on_layout.score < on_sample.score < clean
    assert any("макет устройства" in r for r in on_layout.reasons)


def test_glyphs_a_face_lacks_are_measured_in_a_fallback_not_as_notdef_boxes(monkeypatch):
    """A real face without Cyrillic (Futura Medium on macOS) or without «₽» is set by the renderer in a fallback face:
    those characters are measured as an unknown family, the characters the face has with the face itself."""
    play = fonts._font(False)
    assert fonts._has_glyph(play, "Ж") and fonts._has_glyph(play, " ") and not fonts._has_glyph(play, "")
    assert fonts._coverage_runs(play, "Abc") == [("Ab", True), ("", False), ("c", True)]
    monkeypatch.setattr(fonts, "_real_font", lambda family, bold: play)
    monkeypatch.setattr(fonts, "is_pinned", lambda family: False)
    own = play.getlength("Ab") / fonts._MEASURE_PX * 20
    fallback = play.getlength("") / fonts._MEASURE_PX * 20 * fonts._UNKNOWN_FACTOR
    assert fonts.text_width_pt("Ab", "Some Face", 20) == pytest.approx(own + fallback)


def test_the_visual_variant_never_opens_on_a_cover_whose_layout_draws_an_empty_mockup(monkeypatch):
    """The visual variant takes the template's other cover when it scores within the margin — but not a cover whose
    layout draws a phone mock-up that would stay empty (Gradient: the closing slide's ground is the better cover)."""
    import verstka.matching.matcher as matcher
    from verstka.matching.scorer import ScoreResult
    from verstka.planning.strategies import get_strategy
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    title_slot = [Slot(id="title_1", role=SlotRole.title, shape_id="2", bbox=BboxFrac(x=0.1, y=0.3, w=0.6, h=0.2))]
    best = _pattern("p5", 5, PatternKind.title, layout="ppt/slideLayouts/slideLayout4.xml").model_copy(update={"slots": title_slot})
    phone = _pattern("p1", 1, PatternKind.title, layout="ppt/slideLayouts/slideLayout2.xml").model_copy(update={"slots": title_slot, "mockup_boxes": [BboxFrac(x=0.67, y=0.12, w=0.23, h=0.8)]})
    scores = {"p5": 0.68, "p1": 0.595}
    monkeypatch.setattr(matcher, "score_pattern", lambda slide, p, *a, **k: ScoreResult(scores[p.id], [], {"text_ratio": 1.0}))
    monkeypatch.setattr(matcher, "_cover_room_ok", lambda *a: True)
    outline = DeckOutline(title="t", slides=[OutlineSlide(id="c", kind=PatternKind.title, headline="Кофейня у дома")])
    visual = get_strategy("visual")
    on_sample = _manifest([best, phone])
    assert matcher.match_outline(outline, on_sample, visual).slides[0].pattern_id == "p1"  # the other cover: variants differ
    on_layout = _manifest([best, phone.model_copy(update={"mockup_on_layout": True})])
    assert matcher.match_outline(outline, on_layout, visual).slides[0].pattern_id == "p5"



# ---------------------------------------------------------------------------- G1-19 render stand-ins for missing fonts


@pytest.fixture
def _fake_machine(monkeypatch):
    """A machine with Play, Noto Sans, DIN Condensed, Carlito and PT Serif (all measured with Play's outlines) and
    nothing else; the stand-in caches are emptied around the test."""
    have = {"play", "noto sans", "din condensed", "carlito", "pt serif"}
    play = fonts._font(False)
    for fn in (fonts.render_standin, fonts._installed_face):
        fn.cache_clear()
    monkeypatch.setattr(fonts, "_face", lambda family, bold: ("/x.ttf", 0, family.strip().lower() in have, family))
    monkeypatch.setattr(fonts, "_real_font", lambda family, bold: play)
    yield
    for fn in (fonts.render_standin, fonts._installed_face):
        fn.cache_clear()


def test_a_missing_family_gets_one_installed_face_of_its_class(_fake_machine):
    """«Open Sans» missing: LibreOffice would set digits in OpenSymbol and letters in Helvetica — the render names one
    stand-in of the family's class instead; installed families, theme refs, symbol fonts, weight-named faces and the
    families LibreOffice has a metric twin for are left alone."""
    assert fonts.render_standin("Open Sans") == "Noto Sans"
    assert fonts.render_standin("Bebas Neue") == "DIN Condensed"
    assert fonts.render_standin("Playfair Display") == "PT Serif"
    assert fonts.substitute_of("Open Sans") == "Noto Sans"  # the manifest reports what previews/PDFs show
    for fam in ("Play", "Noto Sans", "+mn-lt", "", "Wingdings", "Symbol", "Lato Light", "Open Sans SemiBold", "Montserrat-Regular", "Calibri", "Calibri Light"):
        assert fonts.render_standin(fam) is None, fam


def test_a_standin_is_never_wider_than_the_layout_measured_the_family(_fake_machine, monkeypatch):
    real = fonts._font(False)
    wide = ImageFont.truetype(str(fonts.font_path(None, True)), fonts._MEASURE_PX)  # Play Bold: wider than Play
    monkeypatch.setattr(fonts, "_real_font", lambda family, bold: wide if family == "Noto Sans" else real)
    fonts._installed_face.cache_clear()
    assert fonts.render_standin("Montserrat") == "Noto Sans"  # measured at Play × 1.2: room for the wide face
    assert fonts.render_standin("Roboto") is None  # measured at Play × 0.98: nothing narrow enough → LibreOffice's own


def test_font_classes_and_word_bounded_pins():
    assert [fonts.font_class(f) for f in ("Consolas", "Bebas Neue", "Open Sans Condensed", "Playfair Display", "Roboto Slab", "Century Gothic", "Open Sans")] == [
        "mono", "display_condensed", "condensed", "serif", "serif", "sans", "sans"]
    assert fonts.is_pinned("Play") and fonts.is_pinned("Open Sans Light") and fonts.is_pinned("Montserrat-Regular")
    assert not fonts.is_pinned("Playfair Display") and fonts.width_factor("Playfair Display") == fonts._UNKNOWN_FACTOR


def test_the_libreoffice_profile_gets_a_replacement_table(tmp_path, monkeypatch):
    import verstka.ingest.render as render
    import verstka.rendering.fonts as fonts_mod

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, fam in enumerate(("Verstka Missing Grotesk", "Play", "A & B")):
        tb = s.shapes.add_textbox(Emu(600000), Emu(600000 + i * 900000), Emu(8000000), Emu(800000))
        tb.text_frame.text = "Выручка 900 000 ₽"
        tb.text_frame.paragraphs[0].runs[0].font.name = fam
    path = tmp_path / "deck.pptx"
    prs.save(path)
    monkeypatch.setattr(fonts_mod, "render_standin", lambda fam: {"Verstka Missing Grotesk": "Noto Sans", "A & B": "PT Sans"}.get(fam))
    pairs = render.font_replacements(path)
    assert pairs == {"Verstka Missing Grotesk": "Noto Sans", "A & B": "PT Sans"}
    assert render.font_replacements(path, available={"verstka missing grotesk"}) == {"A & B": "PT Sans"}
    render._write_font_table(tmp_path / "profile", pairs)
    root = etree.parse(str(tmp_path / "profile" / "user" / "registrymodifications.xcu")).getroot()
    oor = "{http://openoffice.org/2001/registry}"
    vals = {p.get(f"{oor}name"): p.findtext("value") for p in root.iter("prop")}
    assert vals["Replacement"] == "true" and vals["Always"] == "true" and vals["OnScreenOnly"] == "false"
    got = {}
    for node in root.iter("node"):
        v = {p.get(f"{oor}name"): p.findtext("value") for p in node.iter("prop")}
        got[v["ReplaceFont"]] = v["SubstituteFont"]
    assert got == pairs


@pytest.mark.skipif(
    not (__import__("verstka.ingest.render", fromlist=["find_soffice"]).find_soffice() and __import__("shutil").which("pdffonts")),
    reason="LibreOffice and poppler are needed",
)
def test_a_missing_family_renders_in_its_standin(tmp_path):
    """End to end: a run in a family no machine has is set in the stand-in face only (no per-glyph fallback)."""
    import subprocess

    from verstka.ingest.render import pptx_to_pdf

    fonts.render_standin.cache_clear()
    sub = fonts.render_standin("Verstka Missing Grotesk")
    if not sub:
        pytest.skip("no stand-in face on this machine")
    prs = Presentation()
    tb = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Emu(600000), Emu(600000), Emu(8000000), Emu(900000))
    tb.text_frame.text = "Выручка 900 000 ₽ — 60 %"
    tb.text_frame.paragraphs[0].runs[0].font.name = "Verstka Missing Grotesk"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    pdf = pptx_to_pdf(path, tmp_path / "out")
    out = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    faces = [ln.split()[0].split("+", 1)[-1] for ln in out.splitlines()[2:] if ln.strip()]
    assert faces and all(f.replace("-", "").lower().startswith(sub.replace(" ", "").lower()) for f in faces), faces


# ---------------------------------------------------------------------------- gate 2: installed faces without Cyrillic


def test_installed_faces_with_cyrillic_keep_their_family():
    """Only a face this machine has and that lacks Cyrillic gets a Cyrillic stand-in: families with Cyrillic, missing
    families (the replacement table's job), theme references and symbol fonts never do."""
    fonts.cyrillic_standin.cache_clear()
    for fam in ("Play", "", "+mn-lt", "+mj-lt", "Wingdings", "Verstka Missing Grotesk"):
        assert fonts.cyrillic_standin(fam, False) is None and fonts.cyrillic_standin(fam, True) is None, fam
    assert fonts.is_heavy_face_name("Avenir Heavy") and fonts.is_heavy_face_name("Lato Black") and fonts.is_heavy_face_name("Open Sans SemiBold")
    assert not fonts.is_heavy_face_name("Avenir") and not fonts.is_heavy_face_name("Avenir Light") and not fonts.is_heavy_face_name("Avenir Book")


def test_a_face_without_cyrillic_gets_its_own_design_line_first(monkeypatch):
    """macOS Avenir: Book (the regular face) has no Cyrillic, Avenir Next has it → Avenir Next; a weight-named face
    («Avenir Heavy», which LibreOffice sets in Avenir Black + Helvetica) goes through its family; a named face that
    alone lacks Cyrillic («Avenir Next Heavy») → its own family (set bold by the caller); a face name fontconfig does
    not list («Futura Medium» → Futura) is checked through its family."""
    play = fonts._font(False)
    cyr = {"avenir next", "avenir next bold", "noto sans", "noto sans bold", "geo bold"}  # faces with Cyrillic here
    installed = {"avenir", "avenir heavy", "avenir next", "avenir next heavy", "avenir next condensed", "noto sans", "geo"}

    class Face:  # a face measured with Play's outlines; Cyrillic present only where listed
        def __init__(self, fam):
            self.fam = fam

        def getlength(self, text):
            return play.getlength(text)

    monkeypatch.setattr(fonts, "_face", lambda family, bold: ("/x.ttf", 0, family.strip().lower() in installed, family.strip().title()))
    monkeypatch.setattr(fonts, "_real_font", lambda family, bold: Face(family.strip().lower() + (" bold" if bold else "")))
    monkeypatch.setattr(fonts, "_has_glyph", lambda font, ch: getattr(font, "fam", "") in cyr)
    monkeypatch.setattr(fonts, "_installed_families", lambda: ("Avenir", "Avenir Next", "Avenir Next Condensed", "Noto Sans"))
    for fn in (fonts.cyrillic_standin, fonts._installed_face):
        fn.cache_clear()
    try:
        assert fonts._cyrillic_siblings("avenir") == ["Avenir Next", "Avenir Next Condensed"]  # plain width first
        assert fonts.cyrillic_standin("Avenir", False) == "Avenir Next"
        assert fonts.cyrillic_standin("Avenir Heavy", False) == "Avenir Next"
        assert fonts.cyrillic_standin("Avenir Next", False) is None  # has Cyrillic itself
        assert fonts.cyrillic_standin("Avenir Next Heavy", False) == "Avenir Next"  # only the Heavy face lacks it
        assert fonts.is_heavy_face_name("Avenir Next Heavy")
        assert fonts.cyrillic_standin("Avenir Medium", False) == "Avenir Next"  # unlisted name → its family, no Cyrillic
        assert fonts.cyrillic_standin("Avenir Next Medium", False) is None  # unlisted name → its family, has Cyrillic
        assert fonts.cyrillic_standin("Geo Bold", False) is None  # a heavy name → its family bold, which has Cyrillic
        assert fonts.cyrillic_standin("Geo", True) is None and fonts.cyrillic_standin("Geo", False) == "Noto Sans"
    finally:
        for fn in (fonts.cyrillic_standin, fonts._installed_face):
            fn.cache_clear()


def _cyr_deck(tmp_path: Path) -> Path:
    """Slide 1: an explicit «Nocyr Sans» run with Cyrillic, one with Latin only, an Arial run with Cyrillic, a heavy
    face name, a text box inheriting «Nocyr Sans» from its list style, a table cell and a chart in «Nocyr Sans»."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Emu(600000), Emu(300000), Emu(8000000), Emu(900000))
    p = tb.text_frame.paragraphs[0]
    for text, fam in (("Выручка 60 %", "Nocyr Sans"), ("Marketing", "Nocyr Sans"), ("Рост", "Arial"), ("Итоги", "Nocyr Sans Heavy")):
        r = p.add_run()
        r.text = text
        r.font.name = fam
    inh = s.shapes.add_textbox(Emu(600000), Emu(1300000), Emu(8000000), Emu(600000))
    inh.text_frame.text = "Наследует шрифт"
    body = inh.text_frame._txBody
    lst = body.find(f"{{{render_ns()}}}lstStyle")
    lst.append(etree.fromstring(f'<a:lvl1pPr xmlns:a="{render_ns()}"><a:defRPr><a:latin typeface="Nocyr Sans"/></a:defRPr></a:lvl1pPr>'))
    tbl = s.shapes.add_table(1, 1, Emu(600000), Emu(2000000), Emu(4000000), Emu(600000)).table
    tbl.cell(0, 0).text = "Ячейка"
    tbl.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name = "Nocyr Sans"
    data = CategoryChartData()
    data.categories = ["Янв", "Фев"]
    data.add_series("Выручка", (1, 2))
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(600000), Emu(2800000), Emu(4000000), Emu(2500000), data).chart
    chart.font.name = "Nocyr Sans"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def render_ns() -> str:
    return "http://schemas.openxmlformats.org/drawingml/2006/main"


def test_only_cyrillic_runs_of_a_face_without_cyrillic_change_in_the_render_copy(tmp_path, monkeypatch):
    import zipfile

    import verstka.ingest.render as render
    import verstka.rendering.fonts as fonts_mod

    path = _cyr_deck(tmp_path)
    table = {("Nocyr Sans", False): "Stand In", ("Nocyr Sans", True): "Stand In", ("Nocyr Sans Heavy", False): "Stand In"}
    monkeypatch.setattr(fonts_mod, "cyrillic_standin", lambda fam, bold=False: table.get(((fam or "").strip(), bool(bold))))
    parts = render.run_faces(path)
    slide = etree.fromstring(parts["ppt/slides/slide1.xml"])
    a = f"{{{render_ns()}}}"
    runs = {"".join(t.text or "" for t in r.iter(f"{a}t")): r for r in slide.iter(f"{a}r")}
    face = lambda r: (r.find(f"{a}rPr/{a}latin").get("typeface"), r.find(f"{a}rPr").get("b"))  # noqa: E731
    assert face(runs["Выручка 60 %"]) == ("Stand In", None)
    assert face(runs["Marketing"]) == ("Nocyr Sans", None)  # Latin chrome keeps its face
    assert face(runs["Рост"]) == ("Arial", None)
    assert face(runs["Итоги"]) == ("Stand In", "1")  # a heavy face name keeps its weight as bold
    assert face(runs["Наследует шрифт"])[0] == "Stand In"  # inherited through the list style
    assert face(runs["Ячейка"])[0] == "Stand In"  # table cell
    rpr = runs["Наследует шрифт"].find(f"{a}rPr")
    kids = [etree.QName(c).localname for c in rpr]
    assert not set(kids[kids.index("latin") + 1:]) & {"ln", "noFill", "solidFill", "gradFill", "effectLst", "highlight", "uLn", "uFill"}  # schema order
    chart_part = next(p for p in parts if p.startswith("ppt/charts/"))
    assert b'typeface="Stand In"' in parts[chart_part] and b'typeface="Nocyr Sans"' not in parts[chart_part]
    # the copy LibreOffice gets carries the rewrite; the deck itself is untouched
    copy = render.render_copy(path, tmp_path / "copy")
    with zipfile.ZipFile(copy) as z:
        assert b'typeface="Stand In"' in z.read("ppt/slides/slide1.xml")
    with zipfile.ZipFile(path) as z:
        assert b'typeface="Stand In"' not in z.read("ppt/slides/slide1.xml")
    # no such family in the deck → nothing to rewrite (the dataset decks: Play, Arial)
    monkeypatch.setattr(fonts_mod, "cyrillic_standin", lambda fam, bold=False: None)
    assert render.run_faces(path) == {}


@pytest.mark.skipif(
    not (__import__("verstka.ingest.render", fromlist=["find_soffice"]).find_soffice() and __import__("shutil").which("pdffonts")),
    reason="LibreOffice and poppler are needed",
)
def test_a_cyrillic_run_in_an_installed_face_without_cyrillic_renders_in_one_face(tmp_path):
    """End to end on a machine with such a face (macOS Avenir): the Cyrillic run is set in the stand-in only — no
    Avenir + Helvetica inside one word — while a Latin-only run keeps the family."""
    import subprocess

    from verstka.ingest.render import pptx_to_pdf

    fonts.cyrillic_standin.cache_clear()
    fam = next((f for f in ("Avenir", "Optima", "Didot") if fonts.cyrillic_standin(f, False)), None)
    if fam is None:
        pytest.skip("no installed face without Cyrillic on this machine")
    sub = fonts.cyrillic_standin(fam, False)
    prs = Presentation()
    tb = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Emu(600000), Emu(600000), Emu(8000000), Emu(900000))
    tb.text_frame.text = "Выручка Revenue 900 000 — 60 %"
    tb.text_frame.paragraphs[0].runs[0].font.name = fam
    path = tmp_path / "deck.pptx"
    prs.save(path)
    pdf = pptx_to_pdf(path, tmp_path / "out")
    out = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    faces = [ln.split()[0].split("+", 1)[-1] for ln in out.splitlines()[2:] if ln.strip()]
    assert faces and all(f.replace("-", "").lower().startswith(sub.replace(" ", "").lower()) for f in faces), faces


# ---------------------------------------------------------------------------- «₽» in faces without it


def test_a_ruble_sign_the_face_lacks_gets_a_run_of_its_own_in_the_render_copy(tmp_path, monkeypatch):
    """«₽» of a run whose face has no «₽» (Arial, Trebuchet MS, Georgia on macOS) is moved into a run of its own set in
    a face of its class that has it; the rest of the run keeps its face and properties (bold too); a face with «₽»
    (Play) and a deck without «₽» are left alone."""
    import zipfile

    import verstka.ingest.render as render
    import verstka.rendering.fonts as fonts_mod

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = s.shapes.add_textbox(Emu(600000), Emu(300000), Emu(8000000), Emu(900000)).text_frame.paragraphs[0]
    for text, fam, bold in (("Выручка 315 000 ₽ · 40%", "Norub Sans", True), ("20 ₽", "Play", False), ("Итоги ₽₽", "Nocyr Heavy", False),
                            ("7 ₽", "Missing Grotesk SemiBold", False)):
        r = p.add_run()
        r.text, r.font.name, r.font.bold = text, fam, bold
    tbl = s.shapes.add_table(1, 1, Emu(600000), Emu(2000000), Emu(4000000), Emu(600000)).table
    tbl.cell(0, 0).text = "5 ₽"
    tbl.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name = "Norub Sans"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    rub = {"Norub Sans": "Rub Sans", "Cyr Stand In": "Rub Sans", "Missing Grotesk SemiBold": "Rub Sans"}
    monkeypatch.setattr(fonts_mod, "render_standin", lambda fam: None)
    monkeypatch.setattr(fonts_mod, "cyrillic_standin", lambda fam, bold=False: "Cyr Stand In" if fam == "Nocyr Heavy" else None)
    monkeypatch.setattr(fonts_mod, "glyph_standin", lambda fam, ch="₽", bold=False: rub.get(fam))
    slide = etree.fromstring(render.run_faces(path)["ppt/slides/slide1.xml"])
    a = f"{{{render_ns()}}}"
    got = [("".join(t.text or "" for t in r.iter(f"{a}t")), r.find(f"{a}rPr/{a}latin").get("typeface"), r.find(f"{a}rPr").get("b"))
           for r in slide.iter(f"{a}r")]
    assert got == [
        ("Выручка 315 000 ", "Norub Sans", "1"), ("₽", "Rub Sans", "1"), (" · 40%", "Norub Sans", "1"),
        ("20 ₽", "Play", "0"),  # Play has «₽»
        ("Итоги ", "Cyr Stand In", "1"), ("₽₽", "Rub Sans", "1"),  # the Cyrillic stand-in first, then its «₽»; heavy → bold
        ("7 ", "Missing Grotesk SemiBold", "0"), ("₽", "Rub Sans", "0"),  # a missing weight-named face: LibreOffice's weight
        ("5 ", "Norub Sans", None), ("₽", "Rub Sans", None),  # table cell
    ], got
    copy = render.render_copy(path, tmp_path / "copy")
    with zipfile.ZipFile(copy) as z:
        assert b'typeface="Rub Sans"' in z.read("ppt/slides/slide1.xml")
    with zipfile.ZipFile(path) as z:
        assert b'typeface="Rub Sans"' not in z.read("ppt/slides/slide1.xml")  # the deck itself is untouched
    monkeypatch.setattr(fonts_mod, "glyph_standin", lambda fam, ch="₽", bold=False: None)
    monkeypatch.setattr(fonts_mod, "cyrillic_standin", lambda fam, bold=False: None)
    assert render.run_faces(path) == {}  # every face has Cyrillic and «₽» → nothing to rewrite


def test_glyph_standin_keeps_the_family_class(monkeypatch):
    """A serif without «₽» gets a serif «₽», a sans a sans one; a face with «₽», a theme reference and a symbol font
    get none; a bold run needs the stand-in's bold face to have «₽» too."""
    play = fonts._font(False)

    class Face:
        def __init__(self, fam):
            self.fam = fam

        def getlength(self, text):
            return play.getlength(text)

    rub = {"noto serif", "noto serif bold", "pt sans", "helvetica neue", "helvetica neue bold", "play"}
    installed = {"georgia", "arial", "noto serif", "pt sans", "helvetica neue", "play"}
    monkeypatch.setattr(fonts, "_face", lambda family, bold: ("/x.ttf", 0, family.strip().lower() in installed, family))
    monkeypatch.setattr(fonts, "_real_font", lambda family, bold: Face(family.strip().lower() + (" bold" if bold else "")))
    monkeypatch.setattr(fonts, "_has_glyph", lambda font, ch: getattr(font, "fam", "") in rub)
    monkeypatch.setattr(fonts, "_STANDINS", {**fonts._STANDINS, "serif": ("Noto Serif",), "sans": ("PT Sans", "Helvetica Neue")})
    for fn in (fonts.glyph_standin, fonts._installed_face):
        fn.cache_clear()
    try:
        assert fonts.glyph_standin("Georgia", "₽", False) == "Noto Serif"
        assert fonts.glyph_standin("Arial", "₽", False) == "PT Sans"
        assert fonts.glyph_standin("Arial", "₽", True) == "Helvetica Neue"  # PT Sans Bold has no «₽» here
        assert fonts.glyph_standin("Play", "₽", False) is None
        assert fonts.glyph_standin("Missing Grotesk", "₽", False) == "PT Sans"  # LibreOffice's own fallback face
        for fam in ("+mn-lt", "", "Wingdings"):
            assert fonts.glyph_standin(fam, "₽", False) is None
    finally:
        for fn in (fonts.glyph_standin, fonts._installed_face):
            fn.cache_clear()


@pytest.mark.skipif(
    not (__import__("verstka.ingest.render", fromlist=["find_soffice"]).find_soffice() and __import__("shutil").which("pdffonts")),
    reason="LibreOffice and poppler are needed",
)
def test_a_ruble_sign_in_a_face_without_it_renders_in_the_class_standin(tmp_path):
    """End to end on a machine whose Arial has no «₽» (macOS): the PDF has Arial + the stand-in, no STIX Two Math."""
    import subprocess

    from verstka.ingest.render import pptx_to_pdf

    fonts.glyph_standin.cache_clear()
    fam = next((f for f in ("Arial", "Trebuchet MS", "Verdana") if fonts._installed_face(f) and fonts.glyph_standin(f, "₽", False)), None)
    if fam is None:
        pytest.skip("every installed candidate face has «₽» here")
    sub = fonts.glyph_standin(fam, "₽", False)
    prs = Presentation()
    tb = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Emu(600000), Emu(600000), Emu(8000000), Emu(900000))
    tb.text_frame.text = "Выручка 900 000 ₽ — 60 %"
    tb.text_frame.paragraphs[0].runs[0].font.name = fam
    path = tmp_path / "deck.pptx"
    prs.save(path)
    pdf = pptx_to_pdf(path, tmp_path / "out")
    out = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    faces = {ln.split()[0].split("+", 1)[-1].split("-")[0].lower() for ln in out.splitlines()[2:] if ln.strip()}
    assert sub.replace(" ", "").lower() in faces and not any("stix" in f for f in faces), faces
