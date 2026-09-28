"""The audit compares every figure of the finished deck with the source text (checks/facts.py)."""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from verstka.audit.checks.facts import figure_not_in_brief, figure_summary
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import AuditContext

BRIEF = (
    "Ежемесячные расходы кофейни составляют 780 000 рублей: продукты — 315 000 рублей, зарплаты — 270 000 рублей, "
    "аренда — 195 000 рублей. Выручка вырастет с 900 000 до 1 138 500 рублей."
)


def _deck(tmp_path, texts, pie=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for t in texts:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, line in enumerate(t):
            tb = s.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(8), Inches(0.8))
            tb.text_frame.text = line
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        if pie is not None:
            data = CategoryChartData()
            data.categories = list(pie)
            data.add_series("Расходы", list(pie.values()))
            s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(8), Inches(2), Inches(4), Inches(4), data)
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return build_deck_ir(path)


def _ctx(ir, brief=BRIEF):
    return AuditContext(ir=ir, manifest=None, outline=None, brief_text=brief)


def test_brief_figures_and_derived_ones_pass(tmp_path):
    ir = _deck(tmp_path, [["Расходы — 780 000 ₽ в месяц", "Выручка вырастет на 26,5%", "01"]])
    ctx = _ctx(ir)
    assert figure_not_in_brief(ctx) == []
    assert ctx.figure_stats["unverified"] == 0
    assert ctx.figure_stats["derived"] >= 1  # 26,5% of 900 000 → 1 138 500


def test_an_invented_figure_is_reported(tmp_path):
    ir = _deck(tmp_path, [["Расходы — 780 000 ₽ в месяц", "Нас выбрали 12 400 клиентов"]])
    ctx = _ctx(ir)
    issues = figure_not_in_brief(ctx)
    assert len(issues) == 1 and "12 400" in issues[0].message.replace(" ", " ")
    assert ctx.figure_stats["unverified"] == 1
    assert "в нём нет" in figure_summary(ctx.figure_stats)


def test_pie_legend_shares_are_the_charts_own(tmp_path):
    pie = {"Продукты": 315000, "Зарплаты": 270000, "Аренда": 195000}
    ir = _deck(tmp_path, [["Продукты  315 000 ₽ · 40%", "Зарплаты  270 000 ₽ · 35%", "Аренда  195 000 ₽ · 25%"]], pie=pie)
    ctx = _ctx(ir)
    assert figure_not_in_brief(ctx) == []
    assert ctx.figure_stats["derived"] == 3


def test_no_brief_no_check(tmp_path):
    ir = _deck(tmp_path, [["Нас выбрали 12 400 клиентов"]])
    ctx = _ctx(ir, brief=None)
    assert figure_not_in_brief(ctx) == [] and ctx.figure_stats is None


def test_step_numbers_are_not_figures(tmp_path):
    ir = _deck(tmp_path, [["Месяц 2 — запуск комбо", "Этап 1", "3‑й мес.", "Шаг 4: обучение"]])
    ctx = _ctx(ir)
    assert figure_not_in_brief(ctx) == []


# ---------------------------------------------------------------------------- gate 3: a chart's own changes are derived

PROFIT_BRIEF = (
    "Операционная прибыль кофейни сейчас составляет 120 000 рублей в месяц. Средний чек — 300 рублей, в день 100 покупок. "
    "Аренда стоит 120 000 рублей. Зарплаты — 270 000 рублей. Продукты — 315 000 рублей. "
    "По плану на шестой месяц операционная прибыль достигнет 254 795 рублей."
)


def _column_deck(tmp_path, lines, values=(120000, 254795), cats=("Текущая", "Прогноз")):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = list(cats)
    data.add_series("Операционная прибыль", list(values))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(7), Inches(5), data)
    for i, line in enumerate(lines):
        tb = s.shapes.add_textbox(Inches(9), Inches(2 + i), Inches(3.5), Inches(0.8))
        tb.text_frame.text = line
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return build_deck_ir(path)


def test_growth_label_of_the_charts_own_values_is_derived(tmp_path):
    # the renderer's panel beside a chart of 120 000 → 254 795: «×2,1», the difference and the percent change are read
    # off the chart on the same slide (gate 3: «×2,1» was flagged figure_not_in_brief on a live coffee run)
    ir = _column_deck(tmp_path, ["×2,1", "рост: Текущая → Прогноз", "+134 795 ₽", "+112%", "в 2 раза больше"])
    ctx = _ctx(ir, brief=PROFIT_BRIEF)
    assert figure_not_in_brief(ctx) == []
    assert ctx.figure_stats["derived"] == 4 and ctx.figure_stats["unverified"] == 0


def test_a_multiple_the_chart_does_not_give_is_still_reported(tmp_path):
    ir = _column_deck(tmp_path, ["×3,5", "+99 000 ₽", "+47%"])
    ctx = _ctx(ir, brief=PROFIT_BRIEF)
    issues = figure_not_in_brief(ctx)
    assert len(issues) == 3
    assert {"«×3,5»" in i.message or "«3,5»" in i.message for i in issues} >= {True}


def test_a_decline_and_a_series_step_are_derived(tmp_path):
    # «−53%» of 254 795 → 120 000 (the renderer's (1 − ratio) of a falling series); a step between neighbours of a
    # six-month series («+18 000») is read off the chart too
    ir = _column_deck(tmp_path, ["−53%"], values=(254795, 120000))
    assert figure_not_in_brief(_ctx(ir, brief=PROFIT_BRIEF)) == []
    six = ("1-й месяц", "2-й месяц", "3-й месяц", "4-й месяц", "5-й месяц", "6-й месяц")
    ir = _column_deck(tmp_path, ["+18 000 ₽ за месяц", "+77 000 ₽"], values=(120000, 138000, 156000, 190000, 220000, 254795), cats=six)
    shown = " ".join(i.message.replace("\u00a0", " ") for i in figure_not_in_brief(_ctx(ir, brief=PROFIT_BRIEF)))
    assert "18 000" not in shown  # 120 000 → 138 000
    assert "77 000" in shown  # no two values of the series differ by it
