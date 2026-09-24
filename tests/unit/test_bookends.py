"""Covers, section dividers and closing slides: sample choice (scorer, matcher) and the clone renderer's bookend path.

Synthetic decks only (python-pptx + hand-written patterns): the checks hold for any template, not for the VK ones.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.matcher import match_outline
from verstka.matching.scorer import awkward_breaks, balanced_lines, bind_short_words, display_fit, display_lines, score_pattern, split_display_title
from verstka.planning.strategies import get_strategy
from verstka.rendering.clone import render_clone
from verstka.rendering.deck import DeckBuilder, element_bbox, slide_shape_elements
from verstka.rendering.textfill import shape_text
from verstka.schemas.common import BboxFrac, Family, PatternKind, SlotRole
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import DeckOutline, OutlineSlide
from verstka.schemas.template import Capacity, FontUsage, Pattern, SlideSize, Slot, SlotStyle, Spacing, TemplateManifest, Tokens, TypeStep, Typography

W, H = 12192000, 6858000
LONG = "Умные напоминания в VK WorkSpace: итоги пилота и план запуска"


def _manifest(patterns: list[Pattern], w: int = W, h: int = H) -> TemplateManifest:
    typo = Typography(
        families=[FontUsage(family="Play", weight=1.0)],
        scale=[TypeStep(role="display", size_pt=54), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=16), TypeStep(role="body", size_pt=12), TypeStep(role="caption", size_pt=10)],
        sizes_used=[54, 48, 47, 36, 29, 24, 20, 16, 14, 12, 10],
    )
    tokens = Tokens(typography=typo, spacing=Spacing(safe_area=BboxFrac(x=0.03, y=0.06, w=0.94, h=0.86)))
    return TemplateManifest(template_id="synthetic", source_file="s.pptx", slide_size=SlideSize(w=w, h=h), tokens=tokens, patterns=patterns, n_slides=len(patterns))


def _slot(sid: str, role: SlotRole, box: tuple[float, float, float, float], size: float, sample: str = "", chars: int = 20, lines: int = 1, align: str = "l") -> Slot:
    return Slot(id=f"{role.value}_{sid}", role=role, shape_id=sid, bbox=BboxFrac(x=box[0], y=box[1], w=box[2], h=box[3]), style=SlotStyle(font_family="Play", size_pt=size, align=align), capacity=Capacity(max_chars=chars, max_lines=lines), sample_text=sample)


# ---------------------------------------------------------------------------- typesetting helpers


def test_short_words_bind_to_the_next_word():
    assert bind_short_words("итоги пилота и план запуска") == "итоги пилота и\u00a0план запуска"
    assert bind_short_words("Умные напоминания в VK WorkSpace") == "Умные напоминания в\u00a0VK\u00a0WorkSpace"
    # a line never ends on a bound word
    for line in display_lines(bind_short_words(LONG), "Play", 40, False, 330):
        assert line.split()[-1].lower() not in ("в", "и")


def test_display_fit_grows_to_the_line_limit_before_shrinking():
    size, lines = display_fit(bind_short_words(LONG), "Play", 48, False, 560, 3)
    assert lines <= 3 and size == 48  # three lines at the sample size, not one line at a smaller one
    small, n = display_fit(bind_short_words(LONG), "Play", 48, False, 330, 3)
    assert n <= 3 and small < 48


# ---------------------------------------------------------------------------- sample choice


def _vk_like_covers() -> tuple[TemplateManifest, Pattern, Pattern]:
    """A dark title sample whose heading box holds «VK Tech» on one line (54 pt) and a speaker block, and a light
    divider with a 29 pt heading of three lines: the divider must not become the cover because of a long title."""
    cover = Pattern(id="pc", source_slide=2, kind=PatternKind.title, family=Family.dark, quality=0.85, slots=[
        _slot("1", SlotRole.title, (0.064, 0.298, 0.469, 0.156), 54, "VK Tech", chars=11),
        _slot("2", SlotRole.subtitle, (0.064, 0.459, 0.469, 0.056), 16, "Разработчик корпоративного ПО", chars=38),
        _slot("3", SlotRole.body, (0.156, 0.747, 0.469, 0.056), 16, "Имя Фамилия", chars=38),
        _slot("4", SlotRole.body, (0.156, 0.818, 0.469, 0.056), 11, "Должность", chars=56),
        _slot("5", SlotRole.body, (0.064, 0.783, 0.095, 0.056), 10, "Вставить\nфото", chars=10),
    ])
    divider = Pattern(id="pd", source_slide=3, kind=PatternKind.section, family=Family.light, slots=[
        _slot("7", SlotRole.title, (0.023, 0.171, 0.469, 0.291), 29, "Редактируемый слайд-разделитель", chars=63, lines=3),
        _slot("8", SlotRole.body, (0.023, 0.572, 0.469, 0.056), 15, "Дополнительное описание", chars=41),
    ])
    return _manifest([cover, divider], 9144000, 5143500), cover, divider


def test_a_long_title_keeps_the_title_sample_over_a_divider():
    m, cover, divider = _vk_like_covers()
    slide = OutlineSlide(id="t", kind=PatternKind.title, headline=LONG, subtitle="продуктовый комитет")
    st = get_strategy("structured")
    s_cover = score_pattern(slide, cover, m, st)
    s_div = score_pattern(slide, divider, m, st)
    assert s_cover.score > s_div.score, (s_cover.reasons, s_div.reasons)
    assert "title_fit" in s_cover.fit


def test_speaker_placeholders_are_not_clutter_on_a_cover():
    m, cover, _ = _vk_like_covers()
    bare = cover.model_copy(update={"id": "pb", "slots": cover.slots[:2]})
    slide = OutlineSlide(id="t", kind=PatternKind.title, headline="Облако для школ")
    st = get_strategy("structured")
    assert abs(score_pattern(slide, cover, m, st).score - score_pattern(slide, bare, m, st).score) < 0.02


def test_closing_slide_answers_the_cover():
    m, cover, _ = _vk_like_covers()
    title = [_slot("11", SlotRole.title, (0.063, 0.323, 0.63, 0.271), 47, "Спасибо за внимание!", chars=17)]
    dark = Pattern(id="tk_dark", source_slide=4, kind=PatternKind.thanks, family=Family.dark, slots=title, layout_part="l3")
    light = Pattern(id="tk_light", source_slide=6, kind=PatternKind.thanks, family=Family.light, slots=[s.model_copy(update={"shape_id": "12"}) for s in title], layout_part="l5")
    slide = OutlineSlide(id="z", kind=PatternKind.thanks, headline="Спасибо за внимание")
    st = get_strategy("structured")
    assert score_pattern(slide, dark, m, st, cover=cover).score > score_pattern(slide, light, m, st, cover=cover).score


def test_every_divider_of_a_deck_is_the_same_sample():
    m, cover, divider = _vk_like_covers()
    other = divider.model_copy(update={"id": "pd2", "source_slide": 5, "slots": [s.model_copy(update={"shape_id": f"2{s.shape_id}"}) for s in divider.slots]})
    m.patterns.append(other)
    heads = ["Итоги", "Результаты пилота и экономический эффект для бизнеса", "План"]
    slides = [OutlineSlide(id="t", kind=PatternKind.title, headline="Облако для школ")]
    for i, h in enumerate(heads):
        slides += [OutlineSlide(id=f"s{i}", kind=PatternKind.section, headline=h)]
    plan = match_outline(DeckOutline(title="x", slides=slides), m, get_strategy("visual"))
    ids = {p.pattern_id for p in plan.slides[1:]}
    assert len(ids) == 1 and plan.slides[0].pattern_id == "pc", [(p.mode, p.pattern_id) for p in plan.slides]


# ---------------------------------------------------------------------------- the renderer


def _tb(slide, x, y, w, h, text="", size=16.0):
    sh = slide.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    for i, line in enumerate(text.split("\n") if text else [""]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = "Play"
    return sh


def _shape(slide, kind, x, y, w, h, fill=None, line=None):
    sh = slide.shapes.add_shape(kind, Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
    else:
        sh.line.fill.background()
    return sh


def _frac(sh) -> tuple[float, float, float, float]:
    return (sh.left / W, sh.top / H, sh.width / W, sh.height / H)


def sh_id(el) -> str:
    return el.find(".//" + q("p:cNvPr")).get("id")


def _render(pptx: Path, m: TemplateManifest, p: Pattern, oslide: OutlineSlide, outline: DeckOutline, tmp: Path):
    ws = TemplateWorkspace.create(pptx, tmp / "ws")
    return render_clone(DeckBuilder(pptx), LayoutSlide(outline_id=oslide.id, mode="clone", pattern_id=p.id), oslide, p, m, ws, outline)


def _cover_deck(tmp: Path) -> tuple[Path, Pattern, dict]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    sh = {
        "title": _tb(s, 0.05, 0.30, 0.50, 0.16, "Название", size=48),
        "avatar": _shape(s, MSO_SHAPE.OVAL, 0.05, 0.70, 0.06, 0.107, fill="5A6775"),
        "speaker": _tb(s, 0.13, 0.70, 0.25, 0.107, "Имя Спикера,\nдолжность", size=16),
        "photo_box": _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.70, 0.60, 0.07, 0.12, fill="FFFFFF"),
        "photo_text": _tb(s, 0.695, 0.64, 0.08, 0.05, "Вставить\nфото", size=10),
        "frame": _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, -0.02, -0.04, 0.45, 0.25, line="FF0053"),
    }
    pptx = tmp / "cover.pptx"
    prs.save(pptx)
    slots = [
        _slot(str(sh["title"].shape_id), SlotRole.title, _frac(sh["title"]), 48, "Название", chars=12),
        _slot(str(sh["speaker"].shape_id), SlotRole.subtitle, _frac(sh["speaker"]), 16, "Имя Спикера,\nдолжность", chars=40, lines=2),
        _slot(str(sh["photo_text"].shape_id), SlotRole.body, _frac(sh["photo_text"]), 10, "Вставить\nфото", chars=10),
    ]
    return pptx, Pattern(id="pc", source_slide=1, kind=PatternKind.title, family=Family.dark, slots=slots), sh


def test_cover_drops_speaker_block_photo_box_and_fills_the_empty_frame_and_aligns_the_subtitle(tmp_path):
    pptx, p, sh = _cover_deck(tmp_path)
    m = _manifest([p])
    oslide = OutlineSlide(id="t", kind=PatternKind.title, headline=LONG, subtitle="продуктовый комитет VK WorkSpace")
    slide, warnings = _render(pptx, m, p, oslide, DeckOutline(title="x", slides=[oslide]), tmp_path)
    els = slide_shape_elements(slide)
    for gone in ("avatar", "photo_box", "photo_text"):
        assert str(sh[gone].shape_id) not in els, f"{gone} stayed on the cover"
    t_el, s_el = els[str(sh["title"].shape_id)], els[str(sh["speaker"].shape_id)]
    heading = next(sh for sh in slide.shapes if str(sh.shape_id) == str(sh_id(t_el))).text_frame.text  # breaks read as \x0b
    # a heading of two phrases is a title and its subtitle; the audience line it displaces is the kicker
    assert heading.replace("\u00a0", " ").replace("\x0b", " ").split() == "Умные напоминания в VK WorkSpace".split()
    assert shape_text(s_el).replace("\u00a0", " ") == "Итоги пилота и план запуска", "the subtitle starts with a capital"
    fb = element_bbox(els[str(sh["frame"].shape_id)])
    kicker = [e for e in els.values() if "Продуктовый комитет" in shape_text(e).replace("\u00a0", " ")]
    assert len(kicker) == 1, "the audience line stands once, as the kicker"
    kb = element_bbox(kicker[0])
    assert kb[1] >= fb[1] and kb[1] + kb[3] <= fb[1] + fb[3], "the sample's empty corner frame holds the kicker"
    tb, sb = element_bbox(t_el), element_bbox(s_el)
    assert sb[0] == tb[0], "subtitle stands on the heading's left edge"
    assert sb[1] >= tb[1] + tb[3], "subtitle stands under the heading"
    sizes = {int(r.get("sz")) for r in t_el.iter(q("a:rPr")) if r.get("sz")}
    assert sizes and min(sizes) >= 0.85 * 4800, sizes  # three lines at the sample size rather than shrinking
    breaks = t_el.findall(".//" + q("a:br"))
    assert len(breaks) <= 2, "the heading is balanced over at most three lines"


def _divider_deck(tmp: Path) -> tuple[Path, Pattern, list[str]]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = _tb(s, 0.05, 0.10, 0.40, 0.26, "Пример разделителя", size=48)
    grp = s.shapes.add_group_shape()
    dots = []
    for i in range(5):
        d = grp.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int((0.054 + i * 0.043) * W)), Emu(int(0.88 * H)), Emu(int(0.016 * W)), Emu(int(0.016 * W)))
        d.fill.solid()
        d.fill.fore_color.rgb = RGBColor.from_string("0077FF" if i == 0 else "D6ECFF")
        d.line.fill.background()
        dots.append(str(d.shape_id))
    pptx = tmp / "divider.pptx"
    prs.save(pptx)
    slots = [_slot(str(title.shape_id), SlotRole.title, _frac(title), 48, "Пример разделителя", chars=28, lines=2)]
    return pptx, Pattern(id="pd", source_slide=1, kind=PatternKind.section, family=Family.light, slots=slots), dots


def _dot_fills(slide) -> list[str]:
    out = []
    for sp in slide.shapes._spTree.iter(q("p:sp")):
        geom = sp.find(q("p:spPr") + "/" + q("a:prstGeom"))
        if geom is not None and geom.get("prst") == "ellipse":
            out.append((element_bbox(sp)[0], sp.find(q("p:spPr") + "/" + q("a:solidFill") + "/" + q("a:srgbClr")).get("val")))
    return [c for _, c in sorted(out)]


def test_divider_dots_count_the_sections_and_light_the_current_one(tmp_path):
    pptx, p, _ = _divider_deck(tmp_path)
    m = _manifest([p])
    secs = [OutlineSlide(id=f"s{i}", kind=PatternKind.section, headline=h) for i, h in enumerate(["Проблема", "Решение", "План"])]
    outline = DeckOutline(title="x", slides=[OutlineSlide(id="t", kind=PatternKind.title, headline="Т")] + secs)
    slide, _ = _render(pptx, m, p, secs[1], outline, tmp_path)
    assert _dot_fills(slide) == ["D6ECFF", "0077FF", "D6ECFF"]


def test_dots_leave_a_divider_sample_used_as_a_cover(tmp_path):
    pptx, p, _ = _divider_deck(tmp_path)
    m = _manifest([p])
    cover = OutlineSlide(id="t", kind=PatternKind.title, headline="Облако для школ")
    slide, _ = _render(pptx, m, p, cover, DeckOutline(title="x", slides=[cover]), tmp_path)
    assert _dot_fills(slide) == []
    assert not list(slide.shapes._spTree.iter(q("p:grpSp"))), "the emptied dot group goes too"


def test_a_contact_line_on_the_closing_slide_becomes_its_subtitle(tmp_path):
    from verstka.schemas.outline import SlideContent

    pptx, p, sh = _cover_deck(tmp_path)
    p = p.model_copy(update={"kind": PatternKind.thanks})
    m = _manifest([p])
    oslide = OutlineSlide(id="z", kind=PatternKind.thanks, headline="Спасибо за внимание", content=SlideContent(paragraphs=["Команда проекта", "team@example.com"]))
    slide, _ = _render(pptx, m, p, oslide, DeckOutline(title="x", slides=[oslide]), tmp_path)
    els = slide_shape_elements(slide)
    sub = shape_text(els[str(sh["speaker"].shape_id)]).replace("\u00a0", " ")
    assert "Команда проекта" in sub and "team@example.com" in sub
    assert str(sh["avatar"].shape_id) not in els
    assert bind_short_words("Команда · team@example.com") == "Команда\u00a0· team@example.com", "a separator ends its line"


# ---------------------------------------------------------------------------- typesetting a display heading


def test_a_heading_of_two_phrases_splits_into_title_and_subtitle():
    assert split_display_title(LONG) == ("Умные напоминания в VK WorkSpace", "Итоги пилота и план запуска")
    assert split_display_title("Облако для школ — итоги года и планы") == ("Облако для школ", "Итоги года и планы")
    assert split_display_title("Цель: рост") == ("Цель: рост", None), "a one-word part is no title of its own"
    assert split_display_title("ИИ-подсказчик для операторов контакт-центра") == ("ИИ-подсказчик для операторов контакт-центра", None)


def test_balanced_lines_break_after_punctuation_and_keep_adjectives_with_their_nouns():
    text = bind_short_words("Итоги пилота: рост выручки и NPS")
    lines = balanced_lines(text, "Play", 40, False, 560, 2)
    assert lines is not None and lines[0].endswith(":"), lines
    assert awkward_breaks(["Умные", "напоминания"]) == 1
    assert awkward_breaks(["Умные напоминания", "в\u00a0VK\u00a0WorkSpace"]) == 0
    # never a lone short word on the last line when a better break exists
    lines = balanced_lines(bind_short_words("Результаты пилота и экономический эффект"), "Play", 40, False, 520, 2)
    assert lines is not None and len(lines[-1].split()) > 1, lines


def test_the_audit_measures_a_paragraph_with_its_own_line_spacing():
    from verstka.audit.checks.common import text_height_needed_pt
    from verstka.schemas.common import Bbox, BboxFrac
    from verstka.schemas.deck_ir import IRElement, IRParagraph, IRRun

    def el(ls):
        return IRElement(id="1", type="text", bbox=Bbox(x=0, y=0, w=W // 2, h=H // 4), bbox_frac=BboxFrac(x=0, y=0, w=0.5, h=0.25),
                         paragraphs=[IRParagraph(text="Спасибо\nза внимание", line_spacing=ls, runs=[IRRun(text="Спасибо\nза внимание", size_pt=60)])])
    tight, lines = text_height_needed_pt(el(0.9), 1.32)
    loose, _ = text_height_needed_pt(el(None), 1.32)
    assert lines == 2 and abs(tight - 2 * 60 * 1.08) < 0.01 and abs(loose - 2 * 60 * 1.32) < 0.01


# ---------------------------------------------------------------------------- the deck's bookends as one family


def _family_template(tmp: Path) -> tuple[Path, TemplateManifest, dict[str, Pattern]]:
    """A cover (a 48 pt heading set at 90 % line spacing over a subtitle), a divider with pager dots and a closing
    slide whose sample breaks «Спасибо / за внимание»; the template's body text is 110 % (the Education case: the
    audit measured such headings with the body spacing and shrank every cover)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    t1 = _tb(s1, 0.05, 0.30, 0.55, 0.16, "Название", size=48)
    t1.text_frame.paragraphs[0].line_spacing = 0.9
    st1 = _tb(s1, 0.05, 0.50, 0.55, 0.06, "Подзаголовок презентации", size=20)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    t2 = _tb(s2, 0.05, 0.30, 0.45, 0.26, "Пример разделителя", size=48)
    t2.text_frame.paragraphs[0].line_spacing = 0.9
    grp = s2.shapes.add_group_shape()
    for i in range(4):
        d = grp.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int((0.054 + i * 0.043) * W)), Emu(int(0.88 * H)), Emu(int(0.016 * W)), Emu(int(0.016 * W)))
        d.fill.solid()
        d.fill.fore_color.rgb = RGBColor.from_string("0077FF" if i == 0 else "D6ECFF")
        d.line.fill.background()
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    t3 = _tb(s3, 0.05, 0.35, 0.6, 0.2, "Спасибо за внимание", size=48)
    t3.text_frame.paragraphs[0].line_spacing = 0.9
    pptx = tmp / "family.pptx"
    prs.save(pptx)
    pats = {
        "cover": Pattern(id="pc", source_slide=1, kind=PatternKind.title, family=Family.light, slots=[
            _slot(str(t1.shape_id), SlotRole.title, _frac(t1), 48, "Название", chars=12),
            _slot(str(st1.shape_id), SlotRole.subtitle, _frac(st1), 20, "Подзаголовок презентации", chars=40)]),
        "divider": Pattern(id="pd", source_slide=2, kind=PatternKind.section, family=Family.light, slots=[
            _slot(str(t2.shape_id), SlotRole.title, _frac(t2), 48, "Пример разделителя", chars=28, lines=2)]),
        "thanks": Pattern(id="pt", source_slide=3, kind=PatternKind.thanks, family=Family.light, slots=[
            _slot(str(t3.shape_id), SlotRole.title, _frac(t3), 48, "Спасибо\nза внимание", chars=20, lines=2)]),
    }
    m = _manifest(list(pats.values()))
    m.tokens.typography.line_spacing = 1.1
    return pptx, m, pats


def _family_deck(tmp: Path):
    from verstka.rendering.renderer import render_deck
    from verstka.schemas.layout import LayoutPlan

    pptx, m, pats = _family_template(tmp)
    heads = ["Проблема и контекст", "Решение", "Результаты пилота и экономический эффект", "План запуска"]
    slides = [OutlineSlide(id="t", kind=PatternKind.title, headline="Облако для школ: итоги года", subtitle="продуктовый комитет")]
    slides += [OutlineSlide(id=f"s{i}", kind=PatternKind.section, headline=h) for i, h in enumerate(heads)]
    slides += [OutlineSlide(id="z", kind=PatternKind.thanks, headline="Спасибо за внимание")]
    outline = DeckOutline(title="Облако для школ: итоги года", slides=slides)
    kinds = {PatternKind.title: "pc", PatternKind.section: "pd", PatternKind.thanks: "pt"}
    plan = LayoutPlan(strategy="structured", template_id="synthetic", slides=[LayoutSlide(outline_id=s.id, mode="clone", pattern_id=kinds[s.kind]) for s in slides])
    ws = TemplateWorkspace.create(pptx, tmp / "ws")
    out = tmp / "deck.pptx"
    render_deck(outline, plan, m, ws, out)
    return out, m, outline, plan


def _heading(slide) -> tuple[float, int]:
    """(size, top) of the largest text on a slide."""
    best = (0.0, 0)
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()]
        if sizes and max(sizes) > best[0]:
            best = (max(sizes), sh.top)
    return best


def test_bookends_are_one_family_cover_loudest_dividers_alike(tmp_path):
    out, m, outline, plan = _family_deck(tmp_path)
    prs = Presentation(str(out))
    heads = [_heading(s) for s in prs.slides]
    cover, dividers, closing = heads[0], heads[1:5], heads[5]
    assert all(cover[0] >= d[0] for d in dividers), heads
    assert cover[0] >= closing[0], heads
    assert len({d[0] for d in dividers}) == 1, "every divider of the deck is set at one size"
    assert len({d[1] for d in dividers}) == 1, "every divider heading starts on one line of the page"
    texts = [[sh.text_frame.text.replace("\u00a0", " ") for sh in s.shapes if sh.has_text_frame] for s in prs.slides]
    assert "02" in texts[2], "a divider carries its number in the deck"
    assert any(t.startswith("Облако для школ") for t in texts[0]) and "Итоги года" in texts[0]
    assert any("Продуктовый комитет" in t for t in texts[0]), "the audience line moves to the kicker"
    closing_title = next(sh for sh in prs.slides[5].shapes if sh.has_text_frame and sh.text_frame.text.startswith("Спасибо"))
    assert closing_title.text_frame.text.count("\x0b") == 1, "the closing keeps the sample's two lines"
    assert plan.slides[0].fit.get("title_fit", "").startswith(f"{cover[0]:g} пт"), "the plan explains the size the slide got"


def test_bookend_headings_pass_the_audit_so_autofix_leaves_them_alone(tmp_path):
    from verstka.audit.checks.layout import text_overflow
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext

    out, m, outline, _ = _family_deck(tmp_path)
    issues = text_overflow(AuditContext(ir=build_deck_ir(out, with_images=False), manifest=m, outline=outline))
    assert not issues, [i.message for i in issues]


# ---------------------------------------------------------------------------- leftovers


def test_a_picture_pointing_at_a_removed_qr_tile_leaves_with_it(tmp_path):
    from PIL import Image

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = _tb(s, 0.035, 0.235, 0.39, 0.27, "Call to action", size=66)
    tile = _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.64, 0.22, 0.29, 0.51, fill="FFFFFF")
    tile.text_frame.text = "QR-code"
    png = tmp_path / "plane.png"
    Image.new("RGB", (40, 40), (0, 119, 255)).save(png)
    plane = s.shapes.add_picture(str(png), Emu(int(0.60 * W)), Emu(int(0.70 * H)), Emu(int(0.11 * W)), Emu(int(0.14 * H)))
    far = s.shapes.add_picture(str(png), Emu(int(0.05 * W)), Emu(int(0.85 * H)), Emu(int(0.05 * W)), Emu(int(0.05 * H)))
    pptx = tmp_path / "cta.pptx"
    prs.save(pptx)
    p = Pattern(id="pd", source_slide=1, kind=PatternKind.section, family=Family.dark, slots=[
        _slot(str(title.shape_id), SlotRole.title, _frac(title), 66, "Call to action", chars=20, lines=2),
        _slot(str(tile.shape_id), SlotRole.body, _frac(tile), 14, "QR-code", chars=10)])
    oslide = OutlineSlide(id="s", kind=PatternKind.section, headline="Решение")
    slide, _ = _render(pptx, _manifest([p]), p, oslide, DeckOutline(title="x", slides=[oslide]), tmp_path)
    els = slide_shape_elements(slide)
    assert str(tile.shape_id) not in els and str(plane.shape_id) not in els, "the cursor pointing at the QR tile went with it"
    assert str(far.shape_id) in els, "a picture elsewhere on the slide stays"


def test_an_invisible_row_of_dots_leaves_the_divider(tmp_path):
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = _tb(s, 0.05, 0.10, 0.40, 0.26, "Пример разделителя", size=48)
    rows = {}
    for key, y, colors in (("seen", 0.80, ("0077FF", "BBD7F5")), ("ghost", 0.88, ("EBF3F9", "EBF3F9"))):
        rows[key] = []
        for i in range(4):
            d = _shape(s, MSO_SHAPE.OVAL, 0.054 + i * 0.043, y, 0.016, 0.016 * W / H, fill=colors[0] if i == 0 else colors[1])
            rows[key].append(str(d.shape_id))
    pptx = tmp_path / "dots.pptx"
    prs.save(pptx)
    ws = TemplateWorkspace.create(pptx, tmp_path / "ws")
    ws.slides_dir.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", (1280, 720), "#EBF3F9")  # the sample's render: the second row is painted in the ground colour
    dr = ImageDraw.Draw(im)
    for i in range(4):
        x = (0.054 + i * 0.043) * 1280
        dr.ellipse([x, 0.80 * 720, x + 0.016 * 1280, 0.80 * 720 + 0.016 * 1280], fill="#0077FF" if i == 0 else "#BBD7F5")
    im.save(ws.slide_image(1))
    p = Pattern(id="pd", source_slide=1, kind=PatternKind.section, family=Family.light, slots=[_slot(str(title.shape_id), SlotRole.title, _frac(title), 48, "Пример разделителя", chars=28, lines=2)])
    secs = [OutlineSlide(id=f"s{i}", kind=PatternKind.section, headline=h) for i, h in enumerate(["Проблема", "Решение", "План"])]
    slide, _ = render_clone(DeckBuilder(pptx), LayoutSlide(outline_id="s1", mode="clone", pattern_id="pd"), secs[1], p, _manifest([p]), ws, DeckOutline(title="x", slides=secs))
    els = slide_shape_elements(slide)
    assert not any(i in els for i in rows["ghost"]), "dots nobody can see are clutter"
    assert sum(i in els for i in rows["seen"]) == 3, "the visible pager counts the three sections"


def test_the_visual_variant_opens_on_the_other_cover_when_it_is_nearly_as_good():
    m, cover, divider = _vk_like_covers()
    cover = cover.model_copy(update={"layout_part": "l1"})
    other = cover.model_copy(update={"id": "pc2", "source_slide": 5, "layout_part": "l2",
                                     "slots": [s.model_copy(update={"shape_id": f"9{s.shape_id}", "style": s.style.model_copy(update={"size_pt": 48 if s.role == SlotRole.title else s.style.size_pt})}) for s in cover.slots]})
    m.patterns[:] = [cover, divider, other]
    outline = DeckOutline(title="x", slides=[OutlineSlide(id="t", kind=PatternKind.title, headline="Облако для школ")])
    assert match_outline(outline, m, get_strategy("structured")).slides[0].pattern_id == "pc"
    visual = match_outline(outline, m, get_strategy("visual")).slides[0]
    assert visual.pattern_id == "pc2" and any("вариант visual" in r for r in visual.reasons)


def test_a_divider_sets_its_subtitle_as_a_line_not_in_a_painted_button_and_never_repeats_its_number(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = _tb(s, 0.035, 0.235, 0.39, 0.27, "Call to action", size=66)
    button = _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.035, 0.56, 0.12, 0.06, fill="0077FF")
    button.text_frame.text = "Ссылка"
    pptx = tmp_path / "cta.pptx"
    prs.save(pptx)
    p = Pattern(id="pd", source_slide=1, kind=PatternKind.section, family=Family.dark, slots=[
        _slot(str(title.shape_id), SlotRole.title, _frac(title), 66, "Call to action", chars=20, lines=2),
        _slot(str(button.shape_id), SlotRole.subtitle, _frac(button), 14, "Ссылка", chars=10)])
    secs = [OutlineSlide(id="a", kind=PatternKind.section, headline="Проблема", subtitle="Раздел 1"),
            OutlineSlide(id="b", kind=PatternKind.section, headline="Решение", subtitle="Что показали шесть недель пилота")]
    outline = DeckOutline(title="x", slides=secs)
    texts = []
    for sec in secs:
        slide, _ = _render(pptx, _manifest([p]), p, sec, outline, tmp_path)
        els = slide_shape_elements(slide)
        texts.append([shape_text(e).replace(" ", " ") for e in els.values()])
        assert str(button.shape_id) not in els, "the painted button leaves a divider"
    assert not any("Раздел" in t for t in texts[0]) and "01" in texts[0], "«Раздел 1» under «01» says the number twice"
    assert "Что показали шесть недель пилота" in texts[1], "a real subtitle stays, as a plain line"
