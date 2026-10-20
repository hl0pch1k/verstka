"""The jury of 20.10, the pipeline's part: a picture Verstka added itself and could not get leaves no grey place (the slide
is laid out again without it), a picture the person asked for keeps its place with the reason said, «Исправить слайд»
takes an empty place away, a heading under the slide's picture is a remark, a models config without a t2i role
switches the CLI's pictures off, the manifest names the model that answered, the content checks skip what a cover
answers «no» to by design. No live model is called."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from verstka.images import providers as IP
from verstka.images.fill import read_records
from verstka.schemas.outline import PHOTO_PLACE_NAME

TPL = Path(__file__).resolve().parents[1] / "fixtures" / "tpl_audit" / "templates" / "lo_focus.pptx"

PLAIN = """Тема презентации: «Склад в сентябре». Презентация на 6 слайдов.

Слайд 1. Склад в Подольске за сентябрь

Итоги месяца для директора по логистике.

Слайд 2. Отгрузили больше

48 200 заказов — на 12% больше, чем в августе.

Слайд 3. Сборка стала быстрее

Среднее время сборки заказа — 2,4 часа против 3,1 часа в августе.

Слайд 4. Ошибок меньше цели

Доля ошибок комплектации — 0,4% при цели 0,5%.

Слайд 5. Ночная смена

Ночью собирают треть заказов, а ошибок там больше половины.

Слайд 6. Что предлагаю

Перевести ночную смену на сканеры с проверкой.
"""

ASKED = """Тема презентации: «Приложение». Презентация на 3 слайда.

Слайд 1. Запуск приложения

Что запустили в сентябре.

Слайд 2. Экран приложения

Пользователь видит баланс и последние операции на одном экране.

[Место под изображение: экран приложения с балансом]

Слайд 3. Что дальше

Добавить переводы по номеру телефона.
"""


def _places(pptx: Path) -> list[int]:
    prs = Presentation(str(pptx))
    out = []
    for k, s in enumerate(prs.slides, 1):
        for shp in s.shapes:
            if shp.name == PHOTO_PLACE_NAME:
                out.append(k)
    return out


def _build(tmp_path: Path, text: str, ip, pictures=None, strategies=("structured",), progress=None):
    from verstka.pipeline.generate import generate_variants
    from verstka.planning.brief import load_brief

    p = tmp_path / "brief.md"
    p.write_text(text, encoding="utf-8")
    return generate_variants(TPL, brief=load_brief(p), strategies=list(strategies), out_dir=tmp_path / "run", workspace_root=tmp_path / "ws",
                             use_llm=False, use_vlm=False, audit=True, exports=[], image_provider=ip, pictures=pictures, progress=progress)


def test_a_picture_verstka_added_and_could_not_get_leaves_no_grey_place(tmp_path):
    """Pictures on, the image model fails: the places Verstka added itself go, their slides are laid out again — the
    finished deck has no «Место под изображение» and the audit no photo_place_empty (R13: 40 grey places)."""
    events: list[str] = []
    text = (Path(__file__).resolve().parents[1] / "fixtures" / "plans" / "warehouse_email.md").read_text(encoding="utf-8")
    g = _build(tmp_path, text, IP.ImageChain([IP.MockImages(fail=True)]), progress=lambda m, f, **kw: events.append(m))
    v = g.variants[0]
    added = [w for w in v.warnings if w.startswith("picture added on slide")]
    assert added, "the brief should get places added by Verstka (planning/pictures.ensure_pictures)"
    assert not _places(v.out_dir / "deck.pptx")
    assert not [i for i in v.audit.issues if i.check_id == "photo_place_empty"]
    assert all(not s.content.photo_slot for s in v.outline.slides)
    assert read_records(v.out_dir) == []
    assert any("laid out again without the picture's place" in w for w in v.warnings)
    assert any("место под него убрано" in e for e in events)
    # the deck is whole: as many slides as planned, every slide with its text
    prs = Presentation(str(v.out_dir / "deck.pptx"))
    assert len(prs.slides) == len(v.outline.slides)


def test_a_picture_the_person_asked_for_keeps_its_place_and_the_reason_is_said(tmp_path):
    """Pictures off and the text asks for a picture: the place stays (the owner's rule), the timeline says why, the
    audit notes it without a fault."""
    events: list[str] = []
    g = _build(tmp_path, ASKED, IP.ImageChain([IP.MockImages()]), pictures=False, progress=lambda m, f, **kw: events.append(m))
    v = g.variants[0]
    assert _places(v.out_dir / "deck.pptx")
    recs = read_records(v.out_dir)
    assert recs and recs[0]["why"] == "генерация изображений выключена"
    assert any("которое вы просили" in e and "выключена" in e for e in events)
    pe = [i for i in v.audit.issues if i.check_id == "photo_place_empty"]
    assert pe and all(i.severity == "info" for i in pe)


def test_fix_slide_takes_the_empty_place_away(tmp_path):
    """«Исправить слайд» on photo_place_empty: the recipe — the place goes and the slide is laid out again without it."""
    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.pipeline.revise import fix_slide
    from verstka.schemas.audit import AuditReport

    g = _build(tmp_path, ASKED, IP.ImageChain([IP.MockImages()]), pictures=False)
    v = g.variants[0]
    n = _places(v.out_dir / "deck.pptx")[0]
    rep = AuditReport.model_validate_json((v.out_dir / "audit_report.json").read_text(encoding="utf-8"))
    remark = next(i for i in rep.issues if i.check_id == "photo_place_empty" and i.slide == n)
    manifest = analyze_template(TPL, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    res = fix_slide(v.out_dir, "structured", n, manifest=manifest, ws=ws, remarks=[remark])
    assert n not in _places(v.out_dir / "deck.pptx"), res
    out = json.loads((v.out_dir / "outline.json").read_text(encoding="utf-8"))
    assert not out["slides"][n - 1]["content"].get("photo_slot")
    assert not [r for r in read_records(v.out_dir) if r.get("slide") == n]


def test_a_heading_under_the_slides_picture_is_a_remark(tmp_path):
    """text_under_picture: a heading whose line runs under the slide's picture (ЛЦТ «Скринкаст до 5 минут показывает
    выбо…»); a caption set inside the picture is not one."""
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext
    from verstka.audit.checks.layout import text_under_picture

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.4), Inches(1.2), Inches(11.5), Inches(0.6))
    tb.text_frame.text = "Скринкаст до 5 минут показывает выбор визуализации по данным"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    side = s.shapes.add_textbox(Inches(0.4), Inches(3), Inches(5), Inches(1))
    side.text_frame.text = "Короткая строка"
    buf = Path(tmp_path / "pic.png")
    from PIL import Image

    Image.new("RGB", (40, 60), (90, 40, 140)).save(buf)
    pic = s.shapes.add_picture(str(buf), Inches(7.47), Emu(0), Inches(5.87), Inches(7.5))
    pic.name = "Изображение: экран"
    cap = s.shapes.add_textbox(Inches(8), Inches(6), Inches(4), Inches(0.5))
    cap.text_frame.text = "Подпись на картинке"
    path = tmp_path / "d.pptx"
    prs.save(str(path))
    ir = build_deck_ir(path, with_images=False)
    ctx = AuditContext(ir=ir, manifest=None, outline=None, ws=None, slide_images={}, render_ok=None, opens_ok=True)
    got = text_under_picture(ctx)
    assert [i.message for i in got] == ["текст «Скринкаст до 5 минут показывает выбор ви» заходит под картинку слайда"]
    assert got[0].severity == "error" and got[0].autofix.action == "rematch"


def test_cli_models_without_t2i_switch_the_pictures_off(tmp_path):
    from verstka.cli.main import _pictures_switch
    from verstka.images.providers import reset_cache

    reset_cache()
    no_t2i = tmp_path / "models.yaml"
    no_t2i.write_text("roles:\n  llm:\n    backend: mock\n", encoding="utf-8")
    assert _pictures_switch(no_t2i, False) is False
    assert _pictures_switch(no_t2i, False, True) is False  # run.yaml «pictures: true» cannot draw without a model
    assert _pictures_switch(None, False) is None  # the default profile's own switch (configs/images.yaml)
    assert _pictures_switch(None, True) is False  # no models at all
    assert _pictures_switch(None, False, False) is False


def test_the_manifest_names_the_model_that_answered():
    from verstka.pipeline.generate import answered_by
    from verstka.pipeline.run_manifest import build_run_manifest

    class _Reg:
        def describe(self):
            return {"llm": {"backend": "chain", "model": "qwen38-27b-fp8"}, "vlm": {"backend": "chain", "model": "qwen38-27b-fp8"}}

    stats = [{"ok": True, "role": "llm", "model": "qwen/qwen3.8-27b"}] * 3 + [{"ok": False, "role": "llm"}, {"ok": True, "role": "vlm", "model": "qwen38-27b-fp8"}]
    got = answered_by(stats)
    assert got == {"llm": {"qwen/qwen3.8-27b": 3}, "vlm": {"qwen38-27b-fp8": 1}}
    rm = build_run_manifest(template_id="t", template_file="f", strategy="structured", brief_hash=None, outline_hash="x", skills=None, providers=_Reg(),
                            timings={}, applied_fixes=[], audit_summary={}, slides=[], answered=got)
    assert rm["providers"]["llm"]["model"] == "qwen/qwen3.8-27b" and rm["providers"]["llm"]["configured_model"] == "qwen38-27b-fp8"
    assert rm["providers"]["vlm"]["model"] == "qwen38-27b-fp8" and "configured_model" not in rm["providers"]["vlm"]


def test_total_time_is_never_shorter_than_the_plan(tmp_path):
    g = _build(tmp_path, PLAIN, None, pictures=False)
    rm = json.loads((g.variants[0].out_dir / "run_manifest.json").read_text(encoding="utf-8"))
    t = rm["timings_s"]
    assert t["total"] >= t["plan"] and t["total"] >= t["variant"]


def test_content_checks_skip_what_a_cover_answers_no_to(tmp_path):
    """A cover has no conclusion in its heading and no content: by design — questions 1, 2, 3, 5, 10 are not asked of it."""
    from types import SimpleNamespace

    from verstka.audit.ir import build_deck_ir
    from verstka.audit.model_checks import ContentAnswer, ContentAudit, run_model_checks
    from verstka.audit.registry import AuditContext
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    prs = Presentation()
    for t in ("Склад за сентябрь", "Ошибок стало меньше"):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = t
    path = tmp_path / "d.pptx"
    prs.save(str(path))
    ir = build_deck_ir(path, with_images=False)
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="a", kind="title", headline="Склад за сентябрь"), OutlineSlide(id="b", kind="bullets", headline="Ошибок стало меньше")])
    for k, sl in enumerate(ir.slides):
        sl.outline_id = outline.slides[k].id
    img = tmp_path / "s.png"
    from PIL import Image

    Image.new("RGB", (32, 18), (255, 255, 255)).save(img)
    ctx = AuditContext(ir=ir, manifest=None, outline=outline, ws=None, slide_images={1: img, 2: img}, render_ok=True, opens_ok=True)
    no = ContentAudit(answers=[ContentAnswer(question=q, ok=False, evidence="нет") for q in (1, 5, 8)])
    skills = SimpleNamespace(run=lambda *a, **k: SimpleNamespace(parsed=no))
    providers = SimpleNamespace(has=lambda r: r == "vlm", limits=SimpleNamespace(max_concurrency=2))
    issues, _ = run_model_checks(ctx, skills, providers, use_vlm=True, use_llm=False)
    by = sorted((i.slide, i.details["question"]) for i in issues)
    assert by == [(1, 8), (2, 1), (2, 5), (2, 8)]


def test_bars_name_their_series_on_a_row_every_series_has(tmp_path):
    """Two series of bars, the second without a bar on the first row (R09: «Аренда» only in 2026): both are named on
    the first row they share — the chart has no legend, a name lost on an empty bar left 2027 unnamed."""
    from verstka.rendering.charts import add_chart
    from verstka.schemas.common import Bbox
    from verstka.schemas.outline import ChartSpec, DeckOutline, Series
    from verstka.schemas.template import ChartStyleSpec, Typography

    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    cats = ["Аренда", "Облако", "Лицензии"]
    o = DeckOutline(title="t", series=[Series(id="y26", name="2026", categories=cats, values=[36, 18, 64], unit="млн ₽"),
                                       Series(id="y27", name="2027", categories=cats, values=[0, 52, 58], unit="млн ₽")])
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = Bbox(x=600000, y=1700000, w=6700000, h=4100000)
    gf = add_chart(s, box, ChartSpec(type="bar", series_ids=["y26", "y27"]), o, ChartStyleSpec(series_colors=["2A6F97", "61A5C2"], font_size_pt=12.0), Typography(), text_hex="1A1A1A")
    cs = gf.chart._chartSpace
    assert cs.find(f".//{C}legend") is None
    for ser in cs.iter(f"{C}ser"):
        named = [int(d.find(f"{C}idx").get("val")) for d in ser.iter(f"{C}dLbl") if d.find(f"{C}showSerName") is not None and d.find(f"{C}showSerName").get("val") == "1"]
        assert named == [1], named


def test_the_compact_strip_goes_right_when_the_ground_draws_in_the_top_right():
    """The compact variant mirrors a picture's strip to the left — on a plain ground; a ground that draws an object in
    the top right (Blueprint's pencil) keeps the strip on the right, over it, so the heading never crosses it."""
    from verstka.rendering import synth
    from verstka.schemas.outline import DeckOutline

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = int(prs.slide_width), int(prs.slide_height)
    assert synth._strip_left(s, DeckOutline(title="t", strategy="compact"), W, H) is True
    assert synth._strip_left(s, DeckOutline(title="t", strategy="structured"), W, H) is False
    orig = __import__("verstka.rendering.cutouts", fromlist=["cutout_boxes"]).cutout_boxes
    import verstka.rendering.cutouts as cu

    cu.cutout_boxes = lambda slide, area=None, pad=None: [area]
    try:
        assert synth._strip_left(s, DeckOutline(title="t", strategy="compact"), W, H) is False
    finally:
        cu.cutout_boxes = orig


def test_a_long_cover_title_leaves_a_narrow_top_band_for_the_empty_area_below(tmp_path):
    """A template whose first slide is «title and content» (LibreOffice Sunset): a narrow title band at the top, an
    empty body below. A long deck title goes down into the freed area, large and centred — not 14 pt in three lines
    pressed to the top of an empty cover (QA of 20.10)."""
    from pptx.util import Emu, Pt

    from verstka.pipeline.generate import generate_variants
    from verstka.planning.brief import parse_brief_text

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(10080625), Emu(5670550)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Emu(3240000), Emu(450000), Emu(3600000), Emu(900000)
    t.text = "Заголовок"
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(33)
    b = s.placeholders[1]
    b.left, b.top, b.width, b.height = Emu(900000), Emu(1440000), Emu(8280000), Emu(3600000)
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Раздел"
    s2.placeholders[1].text = "Пункт один\nПункт два"
    tpl = tmp_path / "tpl.pptx"
    prs.save(str(tpl))
    brief = parse_brief_text("Слайд 1. Итоги запуска мобильного приложения «Почта Плюс» за первый месяц\n\nСлайд 2. Что получилось\n\nСкачали 120 тысяч раз.\n")
    g = generate_variants(tpl, brief=brief, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, audit=True, exports=[])
    v = g.variants[0]
    assert any("заголовок обложки перенесён" in w for w in v.render.slides[0].warnings)
    head = next(sh for sh in Presentation(str(v.out_dir / "deck.pptx")).slides[0].shapes if sh.has_text_frame and "Итоги запуска" in sh.text_frame.text)
    assert head.top >= Emu(1440000) and head.width > Emu(3600000)
    sizes = [r.font.size.pt for p in head.text_frame.paragraphs for r in p.runs if r.font.size]
    assert sizes and min(sizes) >= 24


def test_the_content_checks_are_on_by_default_with_the_text_model():
    """The model's content checks (the 11 questions of the ТЗ's appendix 1) run by default with the text model, in the
    API as in the CLI; asked off — off; without the text model — off."""
    from verstka.api.app import GenerateRequest, _content_checks_on

    assert _content_checks_on(GenerateRequest(template_id="t")) is True
    assert _content_checks_on(GenerateRequest(template_id="t", audit_models=False)) is False
    assert _content_checks_on(GenerateRequest(template_id="t", use_models=False)) is False
    assert _content_checks_on(GenerateRequest(template_id="t", use_models=False, audit_models=True)) is True
