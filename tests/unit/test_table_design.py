"""Designed native tables: measured geometry (measure_table) and explicit styling (add_table)."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from verstka.analysis.xmlns import q
from verstka.rendering import tables as T
from verstka.rendering.tables import add_table, cell_padding_pt, column_kinds, measure_table, typeset
from verstka.schemas.common import EMU_PER_PT, Bbox, contrast_ratio, relative_luminance
from verstka.schemas.outline import TableData
from verstka.schemas.template import TableStyleSpec, Typography, TypeStep

W, H = 12192000, 6858000
PT = EMU_PER_PT

METRICS = TableData(
    columns=["Метрика", "До пилота", "После пилота"],
    rows=[["Время обращения, мин", "6,5", "4,2"], ["Решено с первого звонка", "71%", "86%"], ["Оценка клиентов", "4,1", "4,7"]],
)
OPTIONS = TableData(
    columns=["Вариант", "Стоимость в год", "Срок внедрения", "Данные в контуре"],
    rows=[["Наш подсказчик", "9 млн ₽", "2 месяца", "да"], ["Коробочное решение", "21 млн ₽", "6 месяцев", "нет"], ["Своя разработка", "35 млн ₽", "12 месяцев", "да"]],
)
WITH_TOTAL = TableData(
    columns=["Площадка", "Операторы", "Экономия, млн ₽"],
    rows=[["Москва", "120", "18,0"], ["Казань", "180", "24,9"], ["Новосибирск", "220", "29,1"], ["Итого", "520", "72,0"]],
)
STAGES = TableData(
    columns=["Этап", "Что делаем", "Срок", "Ответственный"],
    rows=[
        ["Подготовка", "Собираем базу знаний и сценарии обращений, размечаем 5 000 диалогов", "2 недели", "Команда данных"],
        ["Пилот", "Запускаем подсказчик на 40 операторах второй линии", "1 месяц", "Контакт-центр"],
    ],
)
MATRIX = TableData(
    columns=["Возможность", "Базовый", "Про", "Корпоративный"],
    rows=[["Напоминания", "✓", "✓", "✓"], ["Умный срок", "—", "✓", "✓"], ["Эскалация", "Да", "-", "Да"], ["Дайджест", "+", "−", "+"], ["Аналитика", "—", "частично", "✓"]],
)
STYLE = TableStyleSpec(header_fill_hex="0077FF", header_text_hex="FFFFFF", body_text_hex="000000", band_fill_hex="EBF3F9", border_hex="C4C4C4", font_size_pt=12)
TYPO = Typography(scale=[TypeStep(role="h1", size_pt=36), TypeStep(role="body", size_pt=16), TypeStep(role="small", size_pt=12)], sizes_used=[12, 14, 16, 20, 36])
BOLD_TYPO = Typography(scale=[TypeStep(role="h1", size_pt=36, weight_bold_share=0.8), TypeStep(role="body", size_pt=16)], sizes_used=[12, 16, 36])


def _measure(table, width=int(W * 0.9), sizes=(20.0, 16.0, 14.0, 12.0), max_h=int(H * 0.6), fill_h=int(H * 0.36), header_bold=True, **kw):
    return measure_table(table, width, "Play", list(sizes), 1.2, max_h, fill_h, int(0.065 * H), int(0.10 * H), header_bold=header_bold, **kw)


def _slide(bg: str | None = None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if bg:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return prs, slide


def _cell_xml(cell):
    return cell._tc.tcPr


def _run(cell):
    return cell.text_frame.paragraphs[0].runs[0]


def _color(cell) -> str:
    return str(_run(cell).font.color.rgb)


def _fill(cell) -> str | None:
    sf = _cell_xml(cell).find(q("a:solidFill"))
    return sf.find(q("a:srgbClr")).get("val") if sf is not None else None


def _line(cell, side: str) -> str | None:
    ln = _cell_xml(cell).find(q("a:" + side))
    sf = ln.find(q("a:solidFill")) if ln is not None else None
    return sf.find(q("a:srgbClr")).get("val") if sf is not None else None


# ------------------------------------------------------------------------------------------------ classification


def test_column_kinds_and_typesetting():
    assert column_kinds(METRICS) == ["label", "numeric", "numeric"]
    assert column_kinds(OPTIONS) == ["label", "numeric", "numeric", "flag"]  # «9 млн ₽», «2 месяца» are quantities
    assert column_kinds(TableData(columns=["Этап", "Описание"], rows=[["1", "Вторая площадка, 300 операторов"]])) == ["label", "text"]
    assert T._is_numeric("1 500") and T._is_numeric("×2,5") and T._is_numeric("86%") and not T._is_numeric("Этап 1")
    assert typeset("Экономия, млн ₽") == "Экономия, млн ₽"
    assert typeset("Обращений в месяц") == "Обращений в месяц"
    assert typeset("1 220") == "1 220"
    rows = T.display_rows(OPTIONS)
    assert [r[3] for r in rows] == [T.CHECK, T.DASH, T.CHECK]  # boolean column → ✓ / —
    assert T.recommended_row(OPTIONS) == 0 and T.recommended_row(METRICS) is None
    assert T.is_total_row(WITH_TOTAL.rows[-1]) and not T.is_total_row(WITH_TOTAL.rows[0])


def test_numeric_detection_takes_ranges_grades_and_periods():
    for s in ("3–5 дней", "4,7 из 5", "x2", "×3", "2x", "до 5 мин", "≈ 40 %", "−18 %", "Q3 2026", "3 кв. 2026", "10–15%", "1 083 500"):
        assert T._is_numeric(s), s
    for s in ("Команда данных", "Этап 1", "ИТ и контакт-центр", "да"):
        assert not T._is_numeric(s), s
    # one range among plain numbers keeps the column numeric (right-aligned)
    t = TableData(columns=["Этап", "Срок, дней"], rows=[["Пилот", "3–5"], ["Запуск", "10"], ["Масштаб", "4,7 из 5"]])
    assert column_kinds(t) == ["label", "numeric"]


def test_boolean_columns_in_any_spelling_become_marks():
    rows = T.display_rows(MATRIX)
    assert [r[1] for r in rows] == [T.CHECK, T.DASH, T.CHECK, T.CHECK, T.DASH]  # ✓ — Да + —
    assert [r[2] for r in rows] == [T.CHECK, T.CHECK, T.DASH, T.DASH, "частично"]  # ✓ ✓ - − частично
    assert [r[3] for r in rows] == [T.CHECK] * 5  # an all-yes column is boolean too
    assert column_kinds(MATRIX) == ["label", "flag", "flag", "flag"]
    prs, slide = _slide("FFFFFF")
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.6)), MATRIX, STYLE, TYPO, accent_hex="0077FF", muted_hex="6B6B6B").table
    checks = {_color(tbl.cell(r, j)) for r in range(1, 6) for j in range(1, 4) if tbl.cell(r, j).text == T.CHECK}
    dashes = {_color(tbl.cell(r, j)) for r in range(1, 6) for j in range(1, 4) if tbl.cell(r, j).text == T.DASH}
    assert checks == {"0077FF"} and dashes == {"6B6B6B"}  # one colour per mark across every column


def test_units_move_into_the_header():
    heads = T.header_cells(OPTIONS)
    rows = T.display_rows(OPTIONS)
    assert heads[1] == "Стоимость в год, млн ₽" and [r[1] for r in rows] == ["9", "21", "35"]
    assert heads[2] == "Срок внедрения, мес." and [r[2] for r in rows] == ["2", "6", "12"]  # месяца / месяцев → мес.
    pct = TableData(columns=["Регион", "Доля"], rows=[["Москва", "34,5%"], ["Казань", "8 %"]])
    assert T.header_cells(pct)[1] == "Доля,\u00a0%" and [r[1] for r in T.display_rows(pct)] == ["34,5", "8"]
    named = TableData(columns=["Регион", "Экономия, млн"], rows=[["Москва", "18 млн ₽"], ["Казань", "24 млн ₽"]])
    assert T.header_cells(named)[1] == "Экономия, млн ₽"  # a unit the header names is replaced, not doubled
    paren = TableData(columns=["Регион", "Доля (%)"], rows=[["Москва", "34,5%"], ["Казань", "8 %"]])
    assert T.header_cells(paren)[1] == "Доля (%)" and [r[1] for r in T.display_rows(paren)] == ["34,5", "8"]  # the plan's words stay
    mixed = TableData(columns=["Показатель", "Значение"], rows=[["Время", "4,2 мин"], ["Оценка", "4,7 из 5"], ["Нагрузка", "−18 %"]])
    assert T.header_cells(mixed)[1] == "Значение" and T.display_rows(mixed)[0][1] == "4,2 мин"  # different units stay


# ------------------------------------------------------------------------------------------------ measure_table


def test_measure_widths_sum_and_no_word_breaks():
    width = int(W * 0.9)
    for table in (METRICS, OPTIONS, WITH_TOTAL, STAGES, MATRIX):
        size, cols, rows = _measure(table, width)
        assert sum(cols) == width
        assert len(rows) == len(table.rows) + 1
        px, _ = cell_padding_pt(size)
        gut = T._gutters(column_kinds(table), size)
        heads = T.header_cells(table)
        body = T.display_rows(table)
        for r, row in enumerate([heads] + body):
            for j, text in enumerate(row):
                word = T._longest_word(text, "Play", size, r == 0)
                assert word + 2 * px + gut[j] <= cols[j] / PT + 0.01, (table.columns, text, size)


def test_measure_numeric_columns_equal_and_label_capped():
    _, cols, _ = _measure(METRICS)
    assert cols[1] == cols[2] or abs(cols[1] - cols[2]) <= 2
    long_label = TableData(
        columns=["Показатель", "2024", "2025", "2026", "2027"],
        rows=[["Очень длинное название показателя, которое никак не помещается в одну строку таблицы", "1", "2", "3", "4"]],
    )
    width = int(W * 0.9)
    _, cols, _ = _measure(long_label, width, sizes=(16.0,))
    assert cols[0] <= 0.4 * width + 2 * PT or all(abs(c - cols[1]) < 2 for c in cols[2:])
    assert len({round(c / 100) for c in cols[1:4]}) == 1  # numeric columns share one width


def test_spare_width_goes_to_the_label_not_into_rivers_between_numbers():
    width = int(W * 0.9)
    fit = _measure(METRICS, width)
    size, cols, _ = fit
    _, _, nats = T._column_needs(T._prepare(METRICS), set(), width / PT, "Play", size, True)
    for j in (1, 2):  # a number column gets a little air, never more
        assert cols[j] / PT <= max(1.5 * nats[j], nats[j] + 3 * size) + 0.5
    assert cols[0] > cols[1]  # the rest goes to the label
    assert 0 < fit.content_w < width  # the natural width, for a composer that shrinks short tables


def test_measure_picks_largest_clean_size_and_steps_down_when_tall():
    size, _, _ = _measure(METRICS)
    assert size == 20.0  # a short 3×3 table reads at the lead size
    many = TableData(columns=["Город", "Операторы"], rows=[[f"Город {i}", str(100 + i)] for i in range(12)])
    size_small, _, rows = _measure(many, max_h=int(H * 0.55), fill_h=int(H * 0.3))
    assert size_small < 20.0
    assert sum(rows) <= int(H * 0.55) + 13


def test_table_sizes_stay_below_the_heading():
    # LCT: h1 20 pt, lead 20 pt → a table never competes with the heading
    sizes = T.table_sizes(20.0, 15.0, 12.0, h1=20.0, slide_h_emu=H, n_body_rows=3)
    assert max(sizes) <= 20.0 / 1.25 and sizes == [15.0, 12.0]
    assert T.table_sizes(20.0, 16.0, 12.0, h1=36.0, slide_h_emu=6858000, n_body_rows=3) == [20.0, 16.0, 12.0]
    assert T.table_sizes(20.0, 16.0, 12.0, h1=36.0, slide_h_emu=6858000, n_body_rows=7) == [16.0, 12.0]  # B above 6 rows
    assert T.table_sizes(20.0, 16.0, 12.0, h1=36.0, slide_h_emu=6858000, n_body_rows=3, compact=True) == [16.0, 12.0]
    size, _, _ = _measure(METRICS, max_size_pt=16.0)
    assert size <= 16.0


def test_measure_rows_fill_and_cap():
    size, _, rows = _measure(METRICS, fill_h=int(H * 0.36))
    assert all(r >= int(0.065 * H) - 2 for r in rows)
    assert all(r <= int(0.10 * H) + 2 for r in rows[1:])
    assert abs(sum(rows) - int(H * 0.36)) <= 4 * 2  # grown to the fill height
    assert rows[0] <= rows[1]  # the header stays compact …
    assert rows[0] >= 0.75 * rows[1] - 2  # … but reads as a band of the same rhythm
    assert len(set(rows[1:])) == 1  # body rows share one height
    # a fill target beyond what the row cap allows stops at the cap
    _, _, rows2 = _measure(METRICS, fill_h=int(H * 0.9), max_h=int(H * 0.9))
    assert all(r <= int(0.10 * H) + 2 for r in rows2[1:])


def test_measure_reports_overflow_and_density():
    fit = _measure(METRICS)
    assert not fit.overflow and not fit.too_dense
    big = TableData(columns=["Город", "Операторы"], rows=[[f"Город {i}", str(100 + i)] for i in range(15)])
    fit = _measure(big, sizes=(14.0, 12.0), max_h=int(H * 0.3), fill_h=int(H * 0.2))
    assert fit.overflow and fit.too_dense  # the composer splits the table or moves rows on
    size, cols, rows = fit  # still unpacks as before
    assert len(rows) == 16


def test_measure_prefers_one_line_short_columns_over_a_bigger_size():
    # at 20 pt the label «Коробочное решение» would wrap; a smaller clean size wins
    size, cols, _ = _measure(OPTIONS, width=int(W * 0.6), sizes=(20.0, 14.0, 12.0))
    px, _ = cell_padding_pt(size)
    labels = [r[0] for r in OPTIONS.rows]
    assert all(T._lines(t, "Play", size, False, cols[0] / PT - 2 * px) == 1 for t in labels)
    assert size < 20.0


def test_gutter_after_a_right_aligned_column():
    # «2 недели | Команда данных»: the left-aligned column after the right-aligned one starts 1 em further in
    assert column_kinds(STAGES) == ["label", "text", "numeric", "text"]
    size, cols, rows = _measure(STAGES)
    prs, slide = _slide("FFFFFF")
    tbl = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), STAGES, STYLE, TYPO, col_widths=cols, row_heights=rows, size_pt=size).table
    px, _ = cell_padding_pt(size)
    assert abs(tbl.cell(1, 3).margin_left - (px + T.GUTTER_EM * size) * PT) <= 2
    assert abs(tbl.cell(0, 3).margin_left - tbl.cell(1, 3).margin_left) <= 1  # the header follows its column
    assert abs(tbl.cell(1, 1).margin_left - px * PT) <= 2


# ------------------------------------------------------------------------------------------------ add_table


def test_add_table_styling():
    prs, slide = _slide()
    size, cols, rows = _measure(WITH_TOTAL)
    gf = add_table(slide, Bbox(x=int(W * 0.05), y=int(H * 0.25), w=sum(cols), h=sum(rows)), WITH_TOTAL, STYLE, TYPO, font_family="Play", col_widths=cols, row_heights=rows, size_pt=size, accent_hex="FF3885", band_every=0)
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    assert tblPr.find(q("a:tableStyleId")).text == T.NO_STYLE_TABLE_ID
    # header: filled, centred vertically, aligned like its column, white on the saturated fill
    for j in range(3):
        c = tbl.cell(0, j)
        tcPr = _cell_xml(c)
        assert tcPr.get("anchor") == "ctr"
        assert _fill(c) == "0077FF"
        assert _color(c) == "FFFFFF"
    assert tbl.cell(0, 1).text_frame.paragraphs[0].alignment == tbl.cell(1, 1).text_frame.paragraphs[0].alignment
    assert str(tbl.cell(1, 1).text_frame.paragraphs[0].alignment).startswith("RIGHT")
    # body hairlines in the border colour, no fill without banding
    body = _cell_xml(tbl.cell(1, 0))
    assert _line(tbl.cell(1, 0), "lnB") == "C4C4C4"
    assert body.find(q("a:noFill")) is not None
    for tag in ("a:lnL", "a:lnR", "a:lnT"):
        assert body.find(q(tag)).find(q("a:noFill")) is not None
    # schema order in tcPr: lnL, lnR, lnT, lnB, then the fill
    order = [el.tag.split("}")[1] for el in body]
    assert order[:5] == ["lnL", "lnR", "lnT", "lnB", "noFill"]
    # total row: accent rule on top (and on the bottom of the row above); regular in a regular template
    total_r = len(WITH_TOTAL.rows)
    tot = tbl.cell(total_r, 0)
    assert _run(tot).font.bold is False
    assert _line(tot, "lnT") == "FF3885"
    assert _line(tbl.cell(total_r - 1, 0), "lnB") == "FF3885"
    # every run names its typeface for latin, east-asian and complex scripts; margins scale with the size
    px, py = cell_padding_pt(size)
    for r in range(len(WITH_TOTAL.rows) + 1):
        for j in range(3):
            c = tbl.cell(r, j)
            rPr = _run(c)._r.rPr
            assert [rPr.find(q(t)).get("typeface") for t in ("a:latin", "a:ea", "a:cs")] == ["Play"] * 3
            assert abs(c.margin_left - px * PT) <= 1 and abs(c.margin_top - py * PT) <= 1
    # geometry as measured
    assert [int(c.width) for c in tbl.columns] == cols and [int(r.height) for r in tbl.rows] == rows


def test_total_row_weight_follows_the_template():
    prs, slide = _slide()
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.6)), WITH_TOTAL, STYLE, BOLD_TYPO, accent_hex="FF3885").table
    assert _run(tbl.cell(len(WITH_TOTAL.rows), 0)).font.bold is True
    assert _run(tbl.cell(0, 0)).font.bold is True


def test_add_table_flags_banding_and_recommended_row():
    prs, slide = _slide()
    size, cols, rows = _measure(OPTIONS)
    gf = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), OPTIONS, STYLE, TYPO, font_family="Play", col_widths=cols, row_heights=rows, size_pt=size, band_every=2, accent_hex="0077FF", muted_hex="8F8F8F")
    tbl = gf.table
    # boolean column: centred ✓ in accent / — muted, header centred too
    assert str(tbl.cell(0, 3).text_frame.paragraphs[0].alignment).startswith("CENTER")
    assert tbl.cell(1, 3).text == T.CHECK and tbl.cell(2, 3).text == T.DASH
    assert _color(tbl.cell(1, 3)) == "0077FF" or contrast_ratio(_color(tbl.cell(1, 3)), _fill(tbl.cell(1, 3))) >= 3.0
    assert _color(tbl.cell(2, 3)) == "8F8F8F"
    # the recommended row: an accent-hued label on an accent tint, readable at the set size
    tint = _fill(tbl.cell(1, 1))
    label = _color(tbl.cell(1, 0))
    assert T._saturation(label) >= 0.5 and contrast_ratio(label, tint) >= T._need(size, False)
    assert contrast_ratio(tint, "FFFFFF") >= 1.12 and T._saturation(tint) > 0.05
    # a short table keeps its hairlines: a forced band would give a single stripe, which reads as a highlight
    assert _cell_xml(tbl.cell(2, 0)).find(q("a:noFill")) is not None
    assert _line(tbl.cell(2, 0), "lnB") is not None
    # six body rows: banding every 2nd row replaces the inner hairlines; the last row closes the table with a rule
    six = TableData(columns=["Город", "Операторы"], rows=[[f"Город {i}", str(100 + i)] for i in range(6)])
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.6), h=int(H * 0.7)), six, STYLE, TYPO, band_every=2).table
    assert _fill(tbl.cell(2, 0)) == "EBF3F9" and _fill(tbl.cell(4, 0)) == "EBF3F9" and _fill(tbl.cell(1, 0)) is None
    assert _cell_xml(tbl.cell(1, 0)).find(q("a:lnB")).find(q("a:noFill")) is not None
    assert _line(tbl.cell(6, 0), "lnB") is not None
    assert _color(tbl.cell(1, 0)) == "000000"


def test_banding_needs_two_stripes():
    prs, slide = _slide("FFFFFF")
    three = TableData(columns=["Город", "Операторы"], rows=[["Москва", "1"], ["Казань", "2"], ["Самара", "3"]])
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.6), h=int(H * 0.5)), three, STYLE, TYPO, band_every=2).table
    assert all(_fill(tbl.cell(r, 0)) is None for r in (1, 2, 3))
    four = TableData(columns=["Город", "Операторы"], rows=[["Москва", "1"], ["Казань", "2"], ["Самара", "3"], ["Тула", "4"]])
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.6), h=int(H * 0.5)), four, STYLE, TYPO, band_every=2).table
    assert [_fill(tbl.cell(r, 0)) is not None for r in (1, 2, 3, 4)] == [False, True, False, True]


def test_add_table_backward_compatible_self_measuring():
    prs, slide = _slide()
    box = Bbox(x=int(W * 0.05), y=int(H * 0.2), w=int(W * 0.5), h=int(H * 0.6))
    gf = add_table(slide, box, METRICS, STYLE, TYPO)  # the clone/synth call: no geometry kwargs
    assert int(gf.width) == box.w
    assert int(gf.height) <= box.h + 4
    tbl = gf.table
    size = _run(tbl.cell(1, 0)).font.size.pt
    assert size >= 12.0
    assert tbl.cell(0, 0).text == "Метрика" and tbl.cell(2, 2).text == "86%"
    # no geometry kwargs and no family on the typography: still a valid table
    add_table(slide, box, TableData(columns=["A"], rows=[["x"]]), TableStyleSpec(), Typography())


def test_clone_path_sanitises_a_raw_manifest_style():
    # the manifest's first neutral as border (LCT: 000000) and a band that does not stand off the ground (VKT)
    raw = STYLE.model_copy(update={"border_hex": "000000", "band_fill_hex": "FEFFFF"})
    six = TableData(columns=["Город", "Операторы"], rows=[[f"Город {i}", str(100 + i)] for i in range(6)])
    prs, slide = _slide("FAFCFF")
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.7)), six, raw, TYPO).table
    rule = _line(tbl.cell(6, 0), "lnB")
    assert rule != "000000" and 1.12 <= contrast_ratio(rule, "FAFCFF") <= 2.2
    band = _fill(tbl.cell(2, 0))
    assert band and contrast_ratio(band, "FAFCFF") >= 1.04 and contrast_ratio("000000", band) >= 4.5
    # a dark slide read from the background: light rules, a visible band, light text
    dark = raw.model_copy(update={"body_text_hex": "FFFFFF", "band_fill_hex": None})
    prs, slide = _slide("520977")
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.7)), six, dark, TYPO).table
    rule = _line(tbl.cell(6, 0), "lnB")
    assert relative_luminance(rule) > relative_luminance("520977") and 1.2 <= contrast_ratio(rule, "520977") <= 2.0


def test_header_rule_mode_and_colours():
    prs, slide = _slide()
    plain = STYLE.model_copy(update={"header_fill_hex": None})
    gf = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.4)), METRICS, plain, TYPO, muted_hex="6B6B6B", accent_hex="0077FF")
    head = gf.table.cell(0, 0)
    assert _color(head) == "6B6B6B"
    assert _line(head, "lnB") == "0077FF"
    # white on a saturated brand fill, even when the caller asked for black
    assert T.header_text_on("0077FF", "000000") == "FFFFFF"
    assert T.header_text_on("FFD91D", "FFFFFF") == "000000"  # a pale fill keeps dark text
    spec = T.table_style_for_ground(STYLE, ground_hex="000000", text_hex="FFFFFF", accent_hex="0077FF")
    assert spec.body_text_hex == "FFFFFF" and contrast_ratio(spec.body_text_hex, spec.band_fill_hex) >= 4.5
    assert spec.header_text_hex == "FFFFFF"
    assert not T.template_bold(TYPO)
    assert T.template_bold(Typography(scale=[TypeStep(role="h1", size_pt=20, weight_bold_share=0.66)]))


def test_header_text_reads_at_its_size():
    # 11 pt regular white on #0077FF is 4.1:1: the fill is deepened a touch (no brand pairing known) …
    prs, slide = _slide("FFFFFF")
    size, cols, rows = _measure(METRICS, sizes=(11.0,))
    tbl = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), METRICS, STYLE, TYPO, col_widths=cols, row_heights=rows, size_pt=size, native=None).table
    fill = _fill(tbl.cell(0, 0))
    assert fill != "0077FF" and contrast_ratio("FFFFFF", fill) >= 4.5 and T._close(fill, "0077FF", 40)
    # … unless the template itself sets white on that blue: then the brand pairing stands at 3:1
    assert T.header_text_on("0077FF", "FFFFFF", 11.0, False, {("0077FF", "FFFFFF")}) == "FFFFFF"
    assert T.header_text_on("0077FF", "FFFFFF", 11.0, False) == "000000"
    assert T._brand_pair("FFFFFF", "0077FF", {("0077FF", "FFFFFF")})  # either way round
    # large text: 3:1 is enough, the brand fill stays exact
    size, cols, rows = _measure(METRICS, sizes=(20.0,))
    tbl = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), METRICS, STYLE, TYPO, col_widths=cols, row_heights=rows, size_pt=size, native=None).table
    assert _fill(tbl.cell(0, 0)) == "0077FF"


def test_recommended_row_tint_and_label_read_on_every_ground():
    cases = [("FFFFFF", "000000", "0077FF"), ("F4F4F6", "1C1D22", "FE095F"), ("000000", "FFFFFF", "0077FF"), ("520977", "FFFFFF", "FE095F")]
    for ground, text, accent in cases:
        prs, slide = _slide(ground)
        style = STYLE.model_copy(update={"header_fill_hex": accent, "body_text_hex": text, "band_fill_hex": None, "border_hex": None})
        size, cols, rows = _measure(OPTIONS, sizes=(14.0,))
        tbl = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), OPTIONS, style, TYPO, col_widths=cols, row_heights=rows, size_pt=size, accent_hex=accent, ground_hex=ground).table
        tint = _fill(tbl.cell(1, 1))
        assert _cell_xml(tbl.cell(1, 1)).find(q("a:solidFill")).find(q("a:srgbClr")).find(q("a:alpha")) is None  # opaque: no seams
        dark = relative_luminance(ground) < 0.18
        assert contrast_ratio(tint, ground) >= (1.3 if dark else 1.12), (ground, tint)
        label = _color(tbl.cell(1, 0))
        assert contrast_ratio(label, tint) >= T._need(size, False), (ground, label, tint)
        check = _color(tbl.cell(1, 3))
        assert contrast_ratio(check, tint) >= 3.0


def test_styles_on_dark_and_brand_grounds():
    # a brand-blue ground: never a black header; a light band carrying the ground colour, light rules
    spec = T.table_style_for_ground(STYLE, ground_hex="0077FF", text_hex="000000", accent_hex="000000", brand_pairs={("0077FF", "FFFFFF")})
    assert spec.header_fill_hex == "FFFFFF" and spec.header_text_hex == "0077FF"
    assert spec.body_text_hex == "FFFFFF"  # the template sets white on its blue
    assert relative_luminance(spec.border_hex) > relative_luminance("0077FF")
    # dark grounds: rules and bands are lighter steps of the ground, never darker seams
    for ground in ("000000", "520977"):
        style = STYLE.model_copy(update={"header_fill_hex": "FE095F", "border_hex": "231F20", "band_fill_hex": "0F0F0F"})
        spec = T.table_style_for_ground(style, ground_hex=ground, text_hex="FFFFFF", accent_hex="FE095F", divider_hex="231F20")
        assert relative_luminance(spec.border_hex) > relative_luminance(ground) and 1.2 <= contrast_ratio(spec.border_hex, ground) <= 2.0, ground
        assert relative_luminance(spec.band_fill_hex) > relative_luminance(ground) and 1.1 <= contrast_ratio(spec.band_fill_hex, ground) <= 1.6, ground
        assert spec.header_fill_hex == "FE095F"
    # a light ground: a coloured divider is not a hairline; the neutral border is
    spec = T.table_style_for_ground(STYLE.model_copy(update={"border_hex": None}), ground_hex="FFFFFF", text_hex="000000", divider_hex="A0DFE0")
    assert T._saturation(spec.border_hex) < 0.1 and 1.12 <= contrast_ratio(spec.border_hex, "FFFFFF") <= 2.2


# ------------------------------------------------------------------------------------------------ the template's own table


def _native_sample(prs, slide):
    """A sample table in the template's own style: gradient header with bold white text, dark rows with black rules
    (VK WorkSpace p14), and an accent row (Edu «Акцент»)."""
    gf = slide.shapes.add_table(5, 2, Emu(0), Emu(0), Emu(W), Emu(int(H * 0.6)))
    gf.name = "Google Shape;673;p30"
    tbl = gf.table
    grad = etree.fromstring(
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:gsLst>'
        '<a:gs pos="0"><a:srgbClr val="0077FF"/></a:gs><a:gs pos="100000"><a:srgbClr val="00AEE8"/></a:gs></a:gsLst>'
        '<a:lin ang="10800000" scaled="0"/></a:gradFill>'
    )
    for j in range(2):
        c = tbl.cell(0, j)
        c.text = "Заголовок"
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
        T._style_tcpr(c, {}, grad)
    for i in range(1, 5):
        for j in range(2):
            c = tbl.cell(i, j)
            c.text = "Текст"
            _run(c).font.color.rgb = RGBColor.from_string("FFFFFF")
            fill = ("0077FF", None) if i == 4 else ("212121", None)
            T._style_tcpr(c, {"B": ("000000", 0.75)}, fill)
    return gf


def test_native_table_style_is_read_from_the_template():
    prs, sample = _slide("000000")
    _native_sample(prs, sample)
    ctx = T.template_table_context(sample)
    nat = ctx.native
    assert nat is not None and nat.header_fill_xml is not None and T._local(nat.header_fill_xml) == "gradFill"
    assert nat.header_bold is True and nat.header_text_hex == "FFFFFF"
    assert nat.rule == ("000000", 0.75)
    assert nat.body_fill_hex == "212121"
    assert nat.accent_fill_hex == "0077FF" and nat.accent_text_hex == "FFFFFF"
    assert ("0077FF", "FFFFFF") in ctx.pairs
    # a new table in that package follows the sample: its gradient runs once across the header, bold header text
    # (it still fits), the sample's row fill and rules, the accent row for the recommended option
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("000000")
    style = STYLE.model_copy(update={"body_text_hex": "FFFFFF"})
    size, cols, rows = _measure(OPTIONS, header_bold=False, sizes=(14.0,))
    tbl = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), OPTIONS, style, TYPO, col_widths=cols, row_heights=rows, size_pt=size, header_bold=False, accent_hex="0077FF", ground_hex="000000").table
    grads = [_cell_xml(tbl.cell(0, j)).find(q("a:gradFill")) for j in range(4)]
    assert all(g is not None for g in grads)
    ends = [[gs.find(q("a:srgbClr")).get("val") for gs in g.iter(q("a:gs"))] for g in grads]
    assert all(ends[j][-1] == ends[j + 1][0] for j in range(3))  # one continuous band
    assert ends[0][0] == "00AEE8" and ends[-1][-1] == "0077FF"  # ang 180°: the last stop is on the left
    assert _run(tbl.cell(0, 0)).font.bold is True
    assert _fill(tbl.cell(2, 0)) == "212121" and _line(tbl.cell(2, 0), "lnB") == "000000"
    assert _fill(tbl.cell(1, 0)) == "0077FF" and _color(tbl.cell(1, 0)) == "FFFFFF" and _color(tbl.cell(1, 3)) == "FFFFFF"


def test_our_own_tables_are_not_read_as_samples():
    prs, slide = _slide("FFFFFF")
    add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.4)), OPTIONS, STYLE, TYPO)
    prs.part.package.__dict__.pop("_verstka_table_ctx", None)  # scan again, now with our table in the package
    assert T.template_table_context(slide).native is None


def test_slide_ground_is_read_from_the_background():
    prs, slide = _slide()
    assert T._slide_ground(slide) == ("FFFFFF", True)  # the master's bg1
    prs, slide = _slide("520977")
    assert T._slide_ground(slide) == ("520977", True)


# ------------------------------------------------------------------------------------------------ the auditor sees the table


def _audit_manifest():
    from verstka.schemas.template import ColorToken, FontUsage, SlideSize, TemplateManifest, Tokens

    colors = [ColorToken(hex="FFFFFF", roles=["background.light"]), ColorToken(hex="000000", roles=["text.primary"]), ColorToken(hex="0077FF", roles=["accent.1"])]
    typo = Typography(families=[FontUsage(family="Play", weight=1.0)], sizes_used=[12.0, 16.0, 20.0, 36.0])
    return TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(colors=colors, typography=typo), components={"table_style": {"font_size_pt": 7.0}})


def test_audit_reads_real_columns_sizes_and_contrast(tmp_path):
    from verstka.audit.checks.layout import table_cell_wrap
    from verstka.audit.checks.template import table_contrast_low
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext

    prs, slide = _slide("FFFFFF")
    size, cols, rows = _measure(MATRIX, width=int(W * 0.9), sizes=(20.0,))
    add_table(slide, Bbox(x=int(W * 0.05), y=int(H * 0.2), w=sum(cols), h=sum(rows)), MATRIX, STYLE, TYPO, col_widths=cols, row_heights=rows, size_pt=size, accent_hex="0077FF")
    # a careless table: uniform columns, a long word at 20 pt in a narrow column, grey on grey
    bad = slide.shapes.add_table(2, 4, Emu(0), Emu(int(H * 0.8)), Emu(int(W * 0.5)), Emu(int(H * 0.15))).table
    bad.cell(0, 0).text = "Корпоративный"
    for c in range(4):
        for r in range(2):
            cell = bad.cell(r, c)
            cell.text = cell.text or "Текст"
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Emu(int(20 * PT))
            run.font.color.rgb = RGBColor.from_string("9A9A9A")
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("BBBBBB")
    path = tmp_path / "t.pptx"
    prs.save(path)
    ir = build_deck_ir(path, with_images=False)
    tables = [e for e in ir.slides[0].elements if e.type == "table"]
    ours = next(e for e in tables if e.table.rows[0][0] == "Возможность")
    assert ours.table.col_widths_emu == cols  # the real grid, not width / n
    assert ours.table.cells[1][0].size_pt == size and ours.table.cells[0][0].fill_hex == "0077FF"
    ctx = AuditContext(ir=ir, manifest=_audit_manifest())
    wraps = table_cell_wrap(ctx)
    assert [i.element_ids[0] for i in wraps] == [e.id for e in tables if e is not ours]  # only the careless one
    contrast = table_contrast_low(ctx)
    assert [i.element_ids[0] for i in contrast] == [e.id for e in tables if e is not ours]
    assert contrast[0].severity == "error"  # 1.5:1


def test_a_clashing_accent_marks_the_row_instead_of_tinting_it():
    # orange mixed into forest green is olive: the row takes a lighter step of the green and an orange bar
    prs, slide = _slide("0B3D2E")
    style = STYLE.model_copy(update={"header_fill_hex": "FF7A00", "body_text_hex": "FFFFFF", "band_fill_hex": None, "border_hex": None})
    tbl = add_table(slide, Bbox(x=0, y=0, w=int(W * 0.8), h=int(H * 0.5)), OPTIONS, style, TYPO, accent_hex="FF7A00", ground_hex="0B3D2E").table
    tint = _fill(tbl.cell(1, 1))
    assert T._hue_gap(tint, "0B3D2E") < 20 and contrast_ratio(tint, "0B3D2E") >= 1.3
    assert _line(tbl.cell(1, 0), "lnL") == "FF7A00" and _line(tbl.cell(1, 1), "lnL") is None
    assert T._hue_gap(_color(tbl.cell(1, 0)), "FF7A00") < 15  # the label keeps the accent's hue


def test_the_recommended_row_is_not_a_second_header():
    # a template whose accent row has the solid header's colour (VK Education): under the header the recommended row
    # takes a tint, a solid band of the same blue would read as a second header
    prs, sample = _slide("FFFFFF")
    gf = sample.shapes.add_table(5, 2, Emu(0), Emu(0), Emu(W), Emu(int(H * 0.6)))
    gf.name = "Google Shape;674;p31"
    for i in range(5):
        for j in range(2):
            c = gf.table.cell(i, j)
            c.text = "Заголовок" if i == 0 else "Текст"
            _run(c).font.color.rgb = RGBColor.from_string("FFFFFF" if i in (0, 4) else "1A1A1A")
            T._style_tcpr(c, {"B": ("E1E1E1", 0.75)}, ("0077FF", None) if i in (0, 4) else ("FFFFFF", None))
    assert T.template_table_context(sample).native.accent_fill_hex == "0077FF"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    size, cols, rows = _measure(OPTIONS, header_bold=False, sizes=(14.0,))
    tbl = add_table(slide, Bbox(x=0, y=0, w=sum(cols), h=sum(rows)), OPTIONS, STYLE, TYPO, col_widths=cols, row_heights=rows, size_pt=size, header_bold=False, accent_hex="0077FF", ground_hex="FFFFFF").table
    hl = next(i for i in range(1, len(tbl.rows)) if tbl.cell(i, 0).text.startswith("Наш"))
    assert _fill(tbl.cell(0, 0)) == "0077FF" and _fill(tbl.cell(hl, 0)) != "0077FF"
