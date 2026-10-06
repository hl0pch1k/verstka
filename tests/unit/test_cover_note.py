"""Regression test (the live QA of 2026-10-07, Классика cover): the small print that moves under the heading's stack is measured at the stack's
width there, not at the width the art left it at the foot."""
from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from test_bookends import H, W, _frac, _manifest, _render, _shape, _slot, _tb
from verstka.rendering import clone as CL
from verstka.schemas.common import Family, PatternKind, SlotRole
from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent
from verstka.schemas.template import Pattern

GOAL = "Учебный проект о том, как датчики, программирование и анализ данных помогают выращивать растения, экономить ресурсы и сокращать ручную работу."
NOTE = "Все показатели и финансовые расчеты в презентации являются учебными допущениями"


def test_cover_small_print_moved_under_the_stack_takes_the_stacks_width(tmp_path, monkeypatch):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = _tb(s, 0.12, 0.30, 0.76, 0.14, "Основные цели", size=54)
    sub = _tb(s, 0.12, 0.50, 0.76, 0.07, "Подзаголовок образца", size=24)
    _shape(s, MSO_SHAPE.RECTANGLE, 0.0, 0.80, 0.54, 0.002, fill="FFFFFF")  # a hairline across the foot
    _shape(s, MSO_SHAPE.RECTANGLE, 0.795, 0.80, 0.205, 0.002, fill="FFFFFF")
    _shape(s, MSO_SHAPE.OVAL, 0.12, 0.90, 0.04, 0.06, fill="FFFFFF")  # a mark of the template under the rule
    _shape(s, MSO_SHAPE.RECTANGLE, 0.10, 0.22, 0.80, 0.05, fill="FFFFFF")  # a band over the heading: the stack cannot rise
    pptx = tmp_path / "cover.pptx"
    prs.save(pptx)
    slots = [_slot(str(title.shape_id), SlotRole.title, _frac(title), 54, "Основные цели", chars=20),
             _slot(str(sub.shape_id), SlotRole.subtitle, _frac(sub), 24, "Подзаголовок образца", chars=60)]
    p = Pattern(id="pc", source_slide=1, kind=PatternKind.title, family=Family.dark, slots=slots)
    m = _manifest([p])
    # the cover's photo is busy along the foot (art from a third of the width on), calm under the heading's stack
    monkeypatch.setattr(CL, "_art_edge", lambda ctx, band, x_from, hidden, designed_to=0: int(0.35 * W) if band[0] >= int(0.75 * H) else None)
    o = OutlineSlide(id="t", kind=PatternKind.title, headline="Умная теплица", subtitle="Автоматизация полива и контроль микроклимата",
                     footnote=NOTE, content=SlideContent(paragraphs=[GOAL]))
    slide, warnings = _render(pptx, m, p, o, DeckOutline(title="x", slides=[o]), tmp_path)
    shapes = {sh.name: sh for sh in slide.shapes}
    assert not (o.notes or "").strip(), (o.notes, warnings)
    assert "Цель" in shapes and "Сноска" in shapes, (sorted(shapes), warnings)
