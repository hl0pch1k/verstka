"""Composer / synth on third-party templates (TEMPLATE_FIX_PLAN, owner B): a word too wide for its box never becomes
a coordinate past the slide (T04), headings are capped and keep their inherited geometry (T01/T03/T05), sizes follow
the slide on a sparse template (T02/T09), dense blocks stay inside their area (T10), engine text never inherits the
template's capitals or tracking (T14), shape ids stay unique (T16)."""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.charts import add_chart, chart_text_capped, short_categories
from verstka.rendering.compose import Canvas, Kit, Para, Run
from verstka.rendering.deck import DeckBuilder, element_bbox
from verstka.rendering.fit import fit_size
from verstka.rendering.synth import _title_style, deck_style, render_synth
from verstka.schemas.common import EMU_PER_PT, Bbox, PatternKind
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, NumberCallout, OutlineSlide, SlideContent, SlideItem, TableData

LO_W, LO_H = 10078720, 5669280  # 11.02 × 6.2 in: a LibreOffice template
BIG_W, BIG_H = 24384000, 13716000  # 26.67 × 15 in


def _style_runs(tf, size: float, bold: bool = False, color: RGBColor = RGBColor(0x22, 0x22, 0x22)) -> None:
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Arial"


def build_sparse_deck(path: Path, w: int = LO_W, h: int = LO_H, title_pt: float = 44, body_pt: float = 32) -> Path:
    """A LibreOffice-like template: its own type scale is two sizes (a 44 pt title, a 32 pt body placeholder)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(w), Emu(h)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Company presentation"
    _style_runs(s1.shapes.title.text_frame, title_pt * 1.2, True)
    s1.placeholders[1].text = "Subtitle of the deck"
    _style_runs(s1.placeholders[1].text_frame, body_pt)
    for n in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide title {n + 1}"
        _style_runs(s.shapes.title.text_frame, title_pt, True)
        body = s.placeholders[1].text_frame
        body.text = "First point of the slide"
        body.add_paragraph().text = "Second point of the slide"
        _style_runs(body, body_pt)
    prs.save(path)
    return path


@pytest.fixture
def sparse_env(tmp_path):
    deck = build_sparse_deck(tmp_path / "sparse.pptx")
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return deck, manifest, ws


@pytest.fixture
def big_env(tmp_path):
    deck = build_sparse_deck(tmp_path / "big.pptx", BIG_W, BIG_H, title_pt=100, body_pt=25)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return deck, manifest, ws


def _render(env, slides: list[OutlineSlide], comp: str, strategy: str = "structured"):
    deck, manifest, ws = env
    b = DeckBuilder(deck)
    outline = DeckOutline(title="T", strategy=strategy, slides=slides)
    out = []
    for o in slides:
        out.append(render_synth(b, LayoutSlide(outline_id=o.id, mode="synth", composition=comp), o, manifest, ws, outline))
    return b, out


def _boxes(slide):
    for el in slide._element.cSld.find(q("p:spTree")):
        if etree.QName(el).localname in ("sp", "pic", "graphicFrame", "cxnSp", "grpSp"):
            b = element_bbox(el)
            if b:
                yield el, b


# ---------------------------------------------------------------------------------------------- T04 fit


def test_fit_size_never_returns_a_sentinel_height_for_a_word_too_wide():
    box = Bbox(x=0, y=0, w=int(1.2 * 914400), h=int(1.0 * 914400))
    res = fit_size(["рентабельность"], box, "Play", 40.0, False, [40, 32, 24])
    assert res.broken_word and not res.fits
    assert res.height_pt < 40 * 1.2 * 20  # finite: the pieces the word is broken into, never 10^6 lines
    ok = fit_size(["Выручка"], Bbox(x=0, y=0, w=4 * 914400, h=914400), "Play", 24.0, False, [24, 20])
    assert ok.fits and not ok.broken_word


def test_fit_size_ladder_only_tries_no_off_scale_ratio_sizes():
    box = Bbox(x=0, y=0, w=int(2.0 * 914400), h=int(0.3 * 914400))
    text = ["Длинный заголовок, который не помещается в одну строку"]
    free = fit_size(text, box, "Play", 30.0, False, [30, 26.5, 23.5], min_ratio=0.5)
    ladder = fit_size(text, box, "Play", 30.0, False, [30, 26.5, 23.5], min_ratio=0.5, ladder_only=True)
    assert ladder.size_pt in (30.0, 26.5, 23.5)
    assert free.size_pt <= ladder.size_pt


# ---------------------------------------------------------------------------------------------- T04 / T16 canvas


def test_canvas_clamps_boxes_to_the_slide_and_reports_it(simple_deck):
    b = DeckBuilder(simple_deck)
    slide = b.add_blank_slide(6)
    cv = Canvas(slide, font="Arial")
    W, H = b.slide_w, b.slide_h
    cv.text(Bbox(x=int(0.9 * W), y=int(0.95 * H), w=int(0.3 * W), h=int(10 ** 11)), [Para([Run("Текст", 12.0, "000000")])], name="Legend")
    el = cv.added[-1]
    x, y, w, h = element_bbox(el)
    assert 0 <= x and x + w <= W and 0 <= y and y + h <= H
    assert any("clamped" in m for m in cv.clamp_warnings())


def test_canvas_ids_stay_unique_after_a_chart(simple_deck):
    from verstka.schemas.template import ChartStyleSpec, Typography

    b = DeckBuilder(simple_deck)
    slide = b.add_blank_slide(6)
    cv = Canvas(slide, font="Arial")
    cv.text(Bbox(x=0, y=0, w=914400, h=914400), [Para([Run("До графика", 12.0, "000000")])])
    outline = DeckOutline(title="T", slides=[])
    spec = ChartSpec(type="column", categories=["A", "B", "C"], series=[InlineSeries(name="S", values=[1, 2, 3])])
    add_chart(slide, Bbox(x=0, y=914400, w=4 * 914400, h=3 * 914400), spec, outline, ChartStyleSpec(), Typography())
    cv.text(Bbox(x=0, y=0, w=914400, h=914400), [Para([Run("После графика", 12.0, "000000")])])
    ids = [el.get("id") for el in slide._element.iter(q("p:cNvPr"))]
    assert len(ids) == len(set(ids))


# ---------------------------------------------------------------------------------------------- T02 / T09 sizes


def test_sparse_template_sizes_follow_the_slide(sparse_env):
    _, manifest, _ = sparse_env
    typo = manifest.tokens.typography
    if not getattr(typo, "derived_sizes", None):
        pytest.skip("the analysis does not derive a ladder yet (contract C3)")
    kit = Kit(manifest, LO_W, LO_H, "FFFFFF")
    hpt = LO_H / EMU_PER_PT
    assert kit.body <= 0.031 * hpt + 0.5  # not the 32 pt placeholder body
    assert kit.small <= 0.025 * hpt + 0.5
    assert chart_text_capped(typo, 24.0, LO_H)


def test_chart_text_capped_only_for_large_or_sparse(simple_deck):
    from verstka.schemas.template import Typography

    assert not chart_text_capped(Typography(), 12.0, 6858000)  # 12 pt on a 540 pt slide: a dataset size
    assert chart_text_capped(Typography(), 24.0, LO_H)  # a bullet placeholder's 24 pt on an 11″ slide


def test_capped_chart_text_is_at_most_a_fortieth_of_the_slide(sparse_env):
    from verstka.schemas.template import ChartStyleSpec

    deck, manifest, _ = sparse_env
    b = DeckBuilder(deck)
    slide = b.add_blank_slide(6)
    outline = DeckOutline(title="T", slides=[])
    spec = ChartSpec(type="line", categories=[f"{i}-й месяц" for i in range(1, 7)], series=[InlineSeries(name="Выручка", values=[900, 950, 1000, 1050, 1100, 1138])])
    gf = add_chart(slide, Bbox(x=0, y=0, w=int(0.4 * LO_W), h=int(0.5 * LO_H)), spec, outline, ChartStyleSpec(font_size_pt=24.0), manifest.tokens.typography)
    sizes = [int(x.get("sz")) / 100 for x in gf.chart._chartSpace.iter() if etree.QName(x).localname in ("defRPr", "rPr") and x.get("sz")]
    assert sizes and max(sizes) <= 0.03 * LO_H / EMU_PER_PT + 0.5
    cats = [pt.text for pt in gf.chart._chartSpace.iter("{http://schemas.openxmlformats.org/drawingml/2006/chart}v")]
    assert not any("⁠" in (c or "") for c in cats)


def test_short_categories_without_joiner():
    assert "⁠" not in "".join(short_categories(["1-й месяц", "2-й месяц"], joiner=False))
    assert "⁠" in "".join(short_categories(["1-й месяц", "2-й месяц"]))  # the default keeps today's labels


# ---------------------------------------------------------------------------------------------- T03 headings


def test_display_sized_headings_are_capped(big_env):
    deck, manifest, _ = big_env
    b = DeckBuilder(deck)
    slides = [OutlineSlide(id=f"s{i}", kind=PatternKind.bullets, headline="Выручка вырастет на 26,5% за шесть месяцев работы кофейни", content=SlideContent(bullets=["Пункт"])) for i in range(3)]
    ds = deck_style(b, manifest, DeckOutline(title="T", slides=slides))
    hpt = BIG_H / EMU_PER_PT
    assert ds.head_size <= 0.062 * hpt + 0.5  # a 100 pt title on a 15″ slide is a display size, not a heading
    assert ds.head_size >= 0.04 * hpt


def test_heading_keeps_inherited_geometry_and_never_jumps_to_x0(sparse_env):
    _, manifest, _ = sparse_env
    o = OutlineSlide(id="a", kind=PatternKind.bullets, headline="Ежемесячные расходы кофейни составляют 780 000 рублей", content=SlideContent(bullets=["Продукты", "Зарплаты", "Аренда"]))
    _, [(slide, warnings)] = _render(sparse_env, [o], "bullets")
    title = next(sh for sh in slide.shapes if sh.is_placeholder)
    assert title.left is not None and title.left > 0.02 * LO_W
    assert title.top + title.height <= 0.45 * LO_H


# ---------------------------------------------------------------------------------------------- T10 dense blocks


def _dense_slides() -> list[tuple[str, OutlineSlide]]:
    steps = [SlideItem(title=f"{i}-й месяц", text=f"Запуск направления номер {i} и обучение сотрудников кофейни") for i in range(1, 7)]
    cards = [SlideItem(title=f"Сценарий развития {i}", text="Описание сценария в две-три строки текста для проверки плотного блока") for i in range(1, 7)]
    table = TableData(columns=["Показатель", "Сейчас", "Цель", "Разница"], rows=[[f"Строка {i}", f"{i * 100} ₽", f"{i * 120} ₽", f"+{i * 20} ₽"] for i in range(1, 9)])
    pie = ChartSpec(type="pie", categories=["Витрина для десертов", "Программа лояльности и учёт", "Обновление меню и фотографии", "Обучение сотрудников", "Резерв"], series=[InlineSeries(name="Вложения", values=[70000, 34000, 25000, 20000, 30000])], unit="₽")
    line = ChartSpec(type="line", categories=["Сейчас"] + [f"{i}-й месяц" for i in range(1, 7)], series=[InlineSeries(name="Выручка", values=[900, 950, 1000, 1050, 1100, 1120, 1138])], unit="тыс. ₽")
    head = "Выручка вырастет на 26,5% за шесть месяцев"
    return [
        ("process", OutlineSlide(id="p", kind=PatternKind.timeline, headline=head, content=SlideContent(items=steps, bullets=["Бюджет запуска — 180 000 ₽", "Окупаемость — 4 месяца", "Риски — низкие"]))),
        ("cards", OutlineSlide(id="c", kind=PatternKind.cards, headline=head, content=SlideContent(items=cards))),
        ("table", OutlineSlide(id="t", kind=PatternKind.table, headline=head, content=SlideContent(table=table))),
        ("chart_pair", OutlineSlide(id="cp", kind=PatternKind.chart, headline=head, content=SlideContent(chart=line, chart2=pie, bullets=["Витрина окупится за 4 месяца", "Лояльность удержит гостей"]), takeaway="Прирост выручки — 238 500 ₽ в месяц")),
        ("stat_row", OutlineSlide(id="k", kind=PatternKind.stat_row, headline=head, content=SlideContent(numbers=[NumberCallout(value=v, label=l) for v, l in [("900 000 ₽", "выручка в месяц"), ("13,3%", "операционная прибыль"), ("238 500 ₽", "прирост"), ("4 мес.", "окупаемость")]]))),
    ]


@pytest.mark.parametrize("env_name", ["sparse_env", "big_env"])
@pytest.mark.parametrize("strategy", ["structured", "visual", "compact"])
def test_dense_blocks_stay_on_the_slide(env_name, strategy, request):
    env = request.getfixturevalue(env_name)
    deck, manifest, ws = env
    W, H = (LO_W, LO_H) if env_name == "sparse_env" else (BIG_W, BIG_H)
    for comp, o in _dense_slides():
        _, [(slide, warnings)] = _render(env, [o], comp, strategy)
        for el, (x, y, w, h) in _boxes(slide):
            assert x >= -1 and y >= -1 and x + w <= W + 0.005 * W and y + h <= H + 0.005 * H, (comp, strategy, el.find(".//" + q("p:cNvPr")).get("name"), (x / W, y / H, w / W, h / H))
        assert not any("clamped" in m for m in warnings), (comp, strategy, warnings)


def test_pie_legend_stays_inside_its_box(sparse_env):
    _, slides = _render(sparse_env, [o for c, o in _dense_slides() if c == "chart_pair"], "chart_pair", "compact")
    slide, warnings = slides[0]
    legend = [(el, b) for el, b in _boxes(slide) if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith(("Legend", "Share", "Swatch"))]
    assert legend
    assert max(b[0] + b[2] for _, b in legend) <= LO_W
    assert max(b[1] + b[3] for _, b in legend) <= LO_H


# ---------------------------------------------------------------------------------------------- T14 neutral runs


def test_engine_runs_do_not_inherit_caps_or_tracking(sparse_env):
    for comp, o in _dense_slides():
        _, [(slide, _)] = _render(sparse_env, [o], comp)
        for el, _ in _boxes(slide):
            name = el.find(".//" + q("p:cNvPr")).get("name") or ""
            if el.find(".//" + q("p:ph")) is not None:
                continue  # the heading placeholder keeps the template's own heading look
            for r in el.iter(q("a:rPr")):
                if r.getparent() is not None and etree.QName(r.getparent()).localname == "r":
                    assert r.get("cap") == "none" and r.get("spc") == "0", (comp, name)


# ---------------------------------------------------------------------------------------------- T18 palette


def test_monochrome_template_accents_with_its_own_colours():
    from verstka.rendering.compose import palette_accents
    from verstka.schemas.template import ColorToken, Tokens

    mono = Tokens(colors=[
        ColorToken(hex="000000", role="background.dark", roles=["background.dark"], weight=300.0),
        ColorToken(hex="FFFFFF", role="text.primary", roles=["text.primary"], weight=6.0),
        ColorToken(hex="CCCCCC", role="neutral.1", roles=["neutral.1"], weight=1.0),
        ColorToken(hex="111111", role="surface", roles=["surface"], weight=0.1),
    ])
    acc = palette_accents(mono)
    assert acc and "0077FF" not in acc  # never a colour the template does not use
    assert set(acc) <= {"FFFFFF", "CCCCCC"}  # no ground as an accent
    branded = Tokens(colors=[ColorToken(hex="FFFFFF", role="background.light", roles=["background.light"]), ColorToken(hex="E4572E", role="accent.1", roles=["accent.1"])])
    assert palette_accents(branded) == ["E4572E"]


# ---------------------------------------------------------------------------------------------- T10 chart pairs


def test_two_bar_charts_stay_side_by_side_on_a_4_3_slide(tmp_path):
    deck = build_sparse_deck(tmp_path / "classic.pptx", 9144000, 6858000, title_pt=40, body_pt=24)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    a = ChartSpec(type="column", title="Количество покупок", categories=["Сейчас", "Цель"], series=[InlineSeries(name="Покупки", values=[100, 115])])
    b_ = ChartSpec(type="column", title="Средний чек", categories=["Сейчас", "Цель"], series=[InlineSeries(name="Чек", values=[300, 330])], unit="₽")
    o = OutlineSlide(id="pair", kind=PatternKind.chart, headline="Выручка вырастет на 238 500 ₽", content=SlideContent(chart=a, chart2=b_, bullets=["Комбо с напитками", "Программа лояльности", "Предложения для офисов"]))
    for strategy in ("structured", "visual", "compact"):
        _, [(slide, warnings)] = _render((deck, manifest, ws), [o], "chart_pair", strategy)
        charts = [b for el, b in _boxes(slide) if etree.QName(el).localname == "graphicFrame" and el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/chart}chart") is not None]
        assert len(charts) == 2, strategy
        (x1, y1, w1, h1), (x2, y2, w2, h2) = sorted(charts)
        overlap = min(y1 + h1, y2 + h2) - max(y1, y2)
        assert x2 >= x1 + w1 and overlap >= 0.8 * min(h1, h2), (strategy, charts)  # side by side, never one over the other


# ---------------------------------------------------------------------------------------------- T10 cards restructure


def _words_fit(slide) -> list[str]:
    """Words of engine text boxes wider than their box (the audit's word_break measure, 3 % tolerance)."""
    from verstka.rendering.fonts import text_width_pt

    bad = []
    for el, (x, y, w, h) in _boxes(slide):
        if el.find(".//" + q("p:ph")) is not None or etree.QName(el).localname != "sp":
            continue
        bp = el.find(".//" + q("a:bodyPr"))
        ins = sum(int(bp.get(k) or 91440) for k in ("lIns", "rIns")) if bp is not None else 2 * 91440
        usable = (w - ins) / EMU_PER_PT
        for r in el.iter(q("a:r")):
            rpr = r.find(q("a:rPr"))
            t = r.find(q("a:t"))
            if rpr is None or t is None or not rpr.get("sz"):
                continue
            lat = rpr.find(q("a:latin"))
            size = int(rpr.get("sz")) / 100
            for word in (t.text or "").split(" "):
                if word and text_width_pt(word, lat.get("typeface") if lat is not None else None, size, rpr.get("b") == "1") > usable * 1.03:
                    bad.append(word)
    return bad


def test_cards_too_narrow_for_their_words_take_two_rows(tmp_path):
    deck = build_sparse_deck(tmp_path / "classic.pptx", 9144000, 6858000, title_pt=40, body_pt=24)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    bullets = ["Партнёрства с пятью ближайшими офисами", "Программа лояльности с понятными условиями", "Продвижение в районных сообществах", "Дневные предложения с 15:00 до 18:00", "Рост маркетингового бюджета до 35 000 рублей в месяц"]
    o = OutlineSlide(id="b", kind=PatternKind.bullets, headline="Цель — увеличить число покупок со 100 до 115 в день", content=SlideContent(bullets=bullets))
    _, [(slide, warnings)] = _render((deck, manifest, ws), [o], "bullets", "visual")
    assert not _words_fit(slide), _words_fit(slide)
    texts = [b for el, b in _boxes(slide) if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith(("Card text", "Item", "Card"))]
    assert texts
    assert len({round(b[1] / 9144) for b in texts}) >= 2  # never five slivers in one row


# ---------------------------------------------------------------------------------------------- T06 art in the area


def _deck_with_corner_art(path: Path) -> tuple[Path, tuple[float, float, float, float]]:
    """A 16:9 template whose content layout draws a tree-sized shape in the lower right corner of the content band."""
    from pptx.enum.shapes import MSO_SHAPE

    path = build_sparse_deck(path, 12192000, 6858000, title_pt=36, body_pt=20)
    prs = Presentation(path)
    layout = prs.slide_layouts[1]
    tmp = prs.slides[1]
    box = (0.80, 0.62, 0.14, 0.26)
    shp = tmp.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(box[0] * 12192000)), Emu(int(box[1] * 6858000)), Emu(int(box[2] * 12192000)), Emu(int(box[3] * 6858000)))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(0x2E, 0x8B, 0x57)
    el = shp._element
    el.getparent().remove(el)
    layout.shapes._spTree.append(el)
    prs.save(path)
    return path, box


def test_content_keeps_clear_of_art_standing_in_its_area(tmp_path):
    deck, art = _deck_with_corner_art(tmp_path / "art.pptx")
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    o = [x for c, x in _dense_slides() if c == "stat_row"][0]
    _, [(slide, warnings)] = _render((deck, manifest, ws), [o], "stat_row", "structured")
    W, H = 12192000, 6858000
    ax, ay, aw, ah = art[0] * W, art[1] * H, art[2] * W, art[3] * H
    for el, (x, y, w, h) in _boxes(slide):
        if el.find(".//" + q("p:ph")) is not None:
            continue
        ox = max(0, min(x + w, ax + aw) - max(x, ax))
        oy = max(0, min(y + h, ay + ah) - max(y, ay))
        assert ox * oy <= 0.02 * w * h, (el.find(".//" + q("p:cNvPr")).get("name"), warnings)


# ---------------------------------------------------------------------------------------------- T14 heading look


def test_heading_fill_keeps_the_sample_run_look():
    from verstka.rendering.synth import _fill_keep
    from verstka.rendering.textfill import ParagraphSpec

    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="2" name="TextBox 1"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr/>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="6050"/><a:t>Title Text Demo</a:t></a:r>'
        '<a:endParaRPr lang="ru-RU" sz="10000" cap="none" spc="0" baseline="0"/></a:p></p:txBody></p:sp>'
    )
    el = etree.fromstring(xml)
    _fill_keep(el, [ParagraphSpec("Выручка за 30 дней — 900 000 ₽")], size_pt=60.5)
    runs = [r for r in el.iter(q("a:rPr"))]
    assert runs and all(r.get("cap") == "none" and r.get("spc") == "0" for r in runs)
    assert "Выручка" in "".join(t.text or "" for t in el.iter(q("a:t")))


def test_largest_empty_rectangle_of_a_cell_grid():
    from verstka.rendering.synth import _largest_empty

    n = 8
    busy = [[(i >= 5 and j >= 4) or (i == 0 and j == 0) for i in range(n)] for j in range(n)]
    got = _largest_empty(busy, Bbox(x=0, y=0, w=800, h=800), n)
    # columns 0–4 below the busy corner cell (5 × 7) beat columns 1–4 at full height (4 × 8)
    assert got is not None and (got.x, got.y, got.w, got.h) == (0, 100, 500, 700)


def test_pie_in_a_wide_low_box_keeps_its_legend_beside_the_circle(big_env):
    """A band left over a conclusion (wide, a quarter of the slide high) on a 15″ slide: the legend goes beside the
    circle at a dense size — never a dot of a pie over a legend running out of the box."""
    from verstka.rendering.compose import Composer

    deck, manifest, ws = big_env
    b = DeckBuilder(deck)
    slide = b.add_blank_slide(6)
    kit = Kit(manifest, BIG_W, BIG_H, "FFFFFF")
    pie = [o for c, o in _dense_slides() if c == "chart_pair"][0].content.chart2
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Распределение вложений", content=SlideContent(chart=pie))
    comp = Composer(slide, kit, o, DeckOutline(title="T", slides=[o]), "visual", manifest)
    box = Bbox(x=int(0.58 * BIG_W), y=int(0.55 * BIG_H), w=int(0.376 * BIG_W), h=int(0.2 * BIG_H))  # too low for the usual sizes
    n0 = len(comp.cv.tree)
    comp._pie(box, pie, center=False)
    drawn = [(el, element_bbox(el)) for el in list(comp.cv.tree)[n0:] if element_bbox(el)]
    legend = [bb for el, bb in drawn if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith(("Legend", "Share", "Swatch"))]
    assert legend
    tol = int(0.01 * BIG_H)
    assert max(x + w for x, y, w, h in legend) <= box.x + box.w + tol and max(y + h for x, y, w, h in legend) <= box.y + box.h + tol
    charts = [bb for el, bb in drawn if etree.QName(el).localname == "graphicFrame"]
    assert charts and charts[0][3] >= 0.45 * box.h  # the circle keeps most of the band's height
    assert "pie legend does not fit its box" not in comp.warnings and getattr(comp, "_pie_clean", False)


def test_heading_starts_after_a_mark_standing_over_its_left_part(simple_deck):
    """A big page number «09» kept at the left of the heading's box: the heading's letters start after its ink (not
    after its frame), and nothing moves when the mark stands elsewhere."""
    from verstka.rendering.synth import _left_clear

    b = DeckBuilder(simple_deck)
    slide = b.add_blank_slide(6)
    W, H = b.slide_w, b.slide_h
    tb = slide.shapes.add_textbox(Emu(int(0.037 * W)), Emu(int(0.04 * H)), Emu(int(0.094 * W)), Emu(int(0.08 * H)))
    tb.text_frame.text = "09"
    _style_runs(tb.text_frame, 28, True)
    head = Bbox(x=int(0.05 * W), y=int(0.04 * H), w=int(0.86 * W), h=int(0.08 * H))
    past = _left_clear(slide, head, W, H)
    assert head.x < past < int(0.037 * W) + int(0.094 * W) + int(0.015 * W)  # after the digits' ink, inside their frame
    below = Bbox(x=int(0.05 * W), y=int(0.3 * H), w=int(0.86 * W), h=int(0.08 * H))
    assert _left_clear(slide, below, W, H) == below.x  # the mark is not in this band


# ---------------------------------------------------------------------------------------------- gate 1 (writer decks)


def _picture(path: Path, photo: bool) -> Path:
    """A photograph-like picture (fine multi-coloured texture, opaque) or a flat illustration (a few solid shapes)."""
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter

    if photo:
        rng = np.random.default_rng(7)
        base = rng.integers(0, 256, size=(300, 400, 3), dtype=np.uint8)
        im = Image.fromarray(base).filter(ImageFilter.GaussianBlur(1.2))
    else:
        im = Image.new("RGB", (400, 300), (255, 255, 255))
        d = ImageDraw.Draw(im)
        d.ellipse((20, 20, 180, 180), fill=(0, 119, 255))
        d.rectangle((200, 60, 380, 260), fill=(255, 51, 153))
    im.save(path)
    return path


def _deck_with_side_picture(path: Path, pic: Path) -> Path:
    """A cover and two text slides whose layout draws a picture over its right half (the art every slide of that
    layout carries, as VK Education's «Паттерн + фото»); the title and the text at the left."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    W, H = prs.slide_width, prs.slide_height
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Обложка"
    layout = prs.slide_layouts[1]
    for ph, (y, h) in zip(layout.placeholders, [(0.08, 0.15), (0.3, 0.55)]):
        ph.left, ph.top, ph.width, ph.height = Emu(int(0.06 * W)), Emu(int(y * H)), Emu(int(0.4 * W)), Emu(int(h * H))
    # the picture goes onto the layout, with its image part
    from pptx.oxml.shapes.picture import CT_Picture

    _part, rid = layout.part.get_or_add_image_part(str(pic))
    layout.shapes._spTree.append(CT_Picture.new_pic(90, "Art 90", "", rid, int(0.52 * W), 0, int(0.48 * W), H))
    for n in range(2):
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = f"Заголовок слайда {n + 1}"
        s.placeholders[1].text = "Первый пункт слайда"
        s.placeholders[1].text_frame.add_paragraph().text = "Второй пункт слайда"
    prs.save(path)
    return path


@pytest.mark.parametrize("photo", [True, False])
def test_photo_art_never_frames_a_text_slide(tmp_path, photo):
    """G1-04: a template's sample photo (a smiling student with a laptop) is its own story — a composed text slide is
    not set beside it; a flat brand illustration stays art."""
    from verstka.rendering.synth import _choose_art, _find_art, _is_photo

    deck = _deck_with_side_picture(tmp_path / "art.pptx", _picture(tmp_path / "pic.png", photo))
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    b = DeckBuilder(deck)
    arts = [a for a in (_find_art(b, p, manifest) for p in manifest.patterns) if a is not None]
    assert arts, "the side picture is art of the template"
    assert all(a.photo is photo for a in arts)
    lay = b.source_slide(2).slide_layout.part
    pics = [r.target_part for r in lay.rels.values() if not r.is_external and r.reltype.endswith("/image")]
    assert pics and _is_photo(b, pics[0]) is photo
    o = OutlineSlide(id="s1", kind=PatternKind.bullets, headline="Япония атаковала Перл-Харбор", content=SlideContent(bullets=["7 декабря 1941 года — атака", "США вступили в войну"]))
    outline = DeckOutline(title="T", strategy="visual", slides=[o])
    ds = deck_style(b, manifest, outline)
    got = _choose_art(b, manifest, ds, "bullets", o, outline)
    assert (got is None) if photo else (got is not None)


def test_figure_split_keeps_short_units_with_the_number():
    from verstka.rendering.compose import figure_split

    assert figure_split("27 миллионов человек") == ("27", "миллионов человек")
    assert figure_split("40 стран") == ("40", "стран")
    assert figure_split("12 млн ₽ выручки") == ("12 млн ₽", "выручки")
    for v in ("300 ₽", "−69%", "1 854", "4,6 из 5", "31% → 12%", "в 2 раза"):
        assert figure_split(v) is None


@pytest.mark.parametrize("strategy", ["structured", "visual", "compact"])
def test_worded_stat_values_keep_a_figure_size_over_their_labels(simple_deck, tmp_path, strategy):
    """G1-14: «27 миллионов человек / Потери СССР» — the number is set at a figure's size (never below its label), the
    words it counts on the line under it, in the same text (the value stays whole for the audit)."""
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    nums = [NumberCallout(value="27 миллионов человек", label="Потери СССР"), NumberCallout(value="40 стран", label="Территории"), NumberCallout(value="61 государство", label="Участники")]
    o = OutlineSlide(id="s9", kind=PatternKind.stat_row, headline="СССР потерял 27 миллионов человек", content=SlideContent(numbers=nums, bullets=["Китай, Германия, Япония и Польша также понесли большие потери"]))
    _, out = _render((simple_deck, manifest, ws), [o], "stat_row", strategy)
    slide = out[0][0]
    figs, labels = [], []
    for el, _b in _boxes(slide):
        name = el.find(".//" + q("p:cNvPr")).get("name") or ""
        sizes = [int(r.get("sz")) / 100 for r in el.iter(q("a:rPr")) if r.get("sz")]
        text = "".join(t.text or "" for t in el.iter(q("a:t")))
        if name.startswith("Figure"):
            figs.append((sizes, text))
        elif name.startswith("Label"):
            labels.append(max(sizes))
    assert len(figs) == 3 and labels
    for sizes, text in figs:
        assert sizes[0] >= 1.5 * max(labels)  # the number is the hero of its tile
        assert min(sizes) >= max(labels) - 0.05  # its words never smaller than the label
    whole = ["".join(t.split()).replace("⁠", "") for _, t in figs]
    assert "27миллионовчеловек" in whole and "61государство" in whole


def test_a_year_is_never_the_hero_figure_of_a_conclusion():
    """G1-15: «Это меньше, чем в первом квартале 2024 года» has no key figure; a count stays one."""
    from verstka.rendering.compose import _aside_figure

    assert _aside_figure("Это меньше, чем в первом квартале 2024 года") is None
    assert _aside_figure("С 2019 по 2024 число выросло вдвое") is None
    assert _aside_figure("В 2025 году продано 1854 электромобиля") == "1854"
    assert _aside_figure("Продажи выросли на 12% к 2024") == "12%"


@pytest.mark.parametrize("wide", [False, True])
def test_pie_keeps_the_amounts_only_while_the_circle_keeps_its_size(sparse_env, wide):
    """G1-21 / gate 2 B1: the legend of a pie of the brief's amounts keeps «315 000 ₽ · 40%» while the circle keeps the
    pie's least size (0.22 of the slide's height, 0.4 of the box's) — in a box where the amounts would shrink the
    circle under it (the dataset's pie lost half its area so), the legend gives the shares beside a larger circle."""
    from verstka.rendering.compose import Composer

    deck, manifest, ws = sparse_env
    b = DeckBuilder(deck)
    slide = b.add_blank_slide(6)
    kit = Kit(manifest, LO_W, LO_H, "FFFFFF")
    cats = ["Продукты, упаковка и списания", "Зарплаты и связанные начисления", "Аренда", "Коммунальные услуги", "Маркетинг", "Другие операционные расходы"]
    pie = ChartSpec(type="pie", title=None, categories=cats, series=[InlineSeries(name="Расходы", values=[315000, 270000, 120000, 25000, 20000, 30000])], unit="₽")
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Структура расходов", content=SlideContent(chart=pie))
    comp = Composer(slide, kit, o, DeckOutline(title="T", slides=[o]), "compact", manifest)
    if wide:
        box = Bbox(x=int(0.05 * LO_W), y=int(0.25 * LO_H), w=int(0.6 * LO_W), h=int(0.6 * LO_H))
    else:
        box = Bbox(x=int(0.4 * LO_W), y=int(0.25 * LO_H), w=int(0.4 * LO_W), h=int(0.4 * LO_H))
    n0 = len(comp.cv.tree)
    comp._pie(box, pie, center=False)
    drawn = [(el.find(".//" + q("p:cNvPr")).get("name") or "", "".join(t.text or "" for t in el.iter(q("a:t"))), element_bbox(el)) for el in list(comp.cv.tree)[n0:]]
    shares = [t for name, t, _ in drawn if name.startswith("Share")]
    assert len(shares) == 6
    assert all("₽" in t for t in shares) if wide else not any("₽" in t for t in shares)
    circle = [bb for name, _, bb in drawn if name.startswith("Chart")]
    tol = int(0.01 * LO_H)
    assert all(bb[0] + bb[2] <= box.x + box.w + tol and bb[1] + bb[3] <= box.y + box.h + tol for _, _, bb in drawn if bb)
    assert circle and circle[0][3] >= comp._pie_min_d(box) - tol >= int(Composer.PIE_MIN_H * LO_H) - tol


def test_a_long_table_is_set_narrower_than_a_wall_over_the_slide(sparse_env, monkeypatch):
    """G1-21 / VK Tech long: an 8-row table under its heading, with a conclusion, does not cover more than 80 % of the
    safe area (the audit's fill_ratio): its rows stop growing, then it is set as wide as its columns need — at the
    same type size, at last a step smaller."""
    from verstka.rendering.compose import Composer

    rows = [["Покупки в день", "100", "115"], ["Средний чек", "300 рублей", "330 рублей"], ["Месячная выручка", "900 000 рублей", "1 138 500 рублей"], ["Переменные расходы", "315 000 рублей", "375 705 рублей"], ["Постоянные операционные расходы", "465 000 рублей", "508 000 рублей"], ["Операционная прибыль", "120 000 рублей", "254 795 рублей"], ["Операционная рентабельность", "13,3%", "22,4%"], ["Точка безубыточности", "62 покупки", "58 покупок"]]
    t = TableData(columns=["Показатель", "Сейчас", "Цель"], rows=rows)
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Операционная прибыль вырастет в 2,1 раза", takeaway="Выручка увеличивается на 26,5%, а ежемесячная операционная прибыль — примерно на 112%", content=SlideContent(table=t))
    deck, manifest, ws = sparse_env
    # a template whose safe area stops higher (a band of chrome at the foot): the table's room is the wall's size
    safe = manifest.tokens.spacing.safe_area.model_copy(update={"h": 0.8})
    tokens = manifest.tokens.model_copy(update={"spacing": manifest.tokens.spacing.model_copy(update={"safe_area": safe})})
    manifest = manifest.model_copy(update={"tokens": tokens})

    def compose():
        b = DeckBuilder(deck)
        slide = b.add_blank_slide(6)
        sx, sy, sw, sh = int(safe.x * LO_W), int(safe.y * LO_H), int(safe.w * LO_W), int(safe.h * LO_H)
        slide.shapes.add_textbox(Emu(sx), Emu(sy), Emu(sw), Emu(int(0.08 * LO_H))).text_frame.text = o.headline
        comp = Composer(slide, Kit(manifest, LO_W, LO_H, "FFFFFF"), o, DeckOutline(title="T", slides=[o], strategy="structured"), "structured", manifest)
        area = Bbox(x=sx, y=sy + int(0.1 * LO_H), w=sw, h=sh - int(0.1 * LO_H))
        n0 = len(comp.cv.tree)
        comp.compose("table", area)
        comp._notes = None  # the conclusion is drawn now: measured as a block of the slide
        frame = next(el for el, _ in _boxes(slide) if etree.QName(el).localname == "graphicFrame")
        sizes = {int(r.get("sz")) for r in frame.iter(q("a:rPr")) if r.get("sz")}
        return comp._coverage(n0, area), sizes

    with monkeypatch.context() as m:
        m.setattr(Composer, "FILL_MAX", 5.0)
        wall, sizes0 = compose()
    assert wall > 0.8  # the case: grown rows at the full width cover more than the audit allows
    got, sizes = compose()
    assert got <= Composer.FILL_MAX <= 0.8
    assert min(sizes) >= min(sizes0) * 0.75  # a step smaller at most



# ---------------------------------------------------------------------------------------------- gate 2 (B2–B6)


def _plan_slide() -> OutlineSlide:
    left = ["Витрина для десертов — 70 000 ₽", "Настройка программы лояльности и учета — 35 000 ₽", "Обновление меню, фотографии и оформление — 25 000 ₽", "Обучение сотрудников — 20 000 ₽", "Резерв — 30 000 ₽"]
    right = ["1-й месяц — учет показателей и обновление меню", "2-й месяц — запуск комбо и обучение сотрудников", "3-й месяц — запуск программы лояльности и партнерств с офисами", "4-й месяц — продвижение дневных предложений и настройка закупок", "5-й месяц — корректировка предложений по результатам продаж", "6-й месяц — оценка результатов и закрепление удачных решений"]
    items = [SlideItem(title="Разовые вложения", bullets=left), SlideItem(title="План внедрения", bullets=right)]
    return OutlineSlide(id="tc", kind=PatternKind.comparison, headline="Разовые вложения составляют 180 000 рублей", content=SlideContent(items=items), takeaway="Все разовые вложения покрываются бюджетом запуска")


@pytest.mark.parametrize("foot_room", [False, True])
def test_a_conclusion_the_block_leaves_no_room_for_stays_on_the_slide(sparse_env, foot_room):
    """B2: two long lists take the conclusion's room. The conclusion is not said aloud: it stands in the free room under
    the art the area was cut above (`foot_room`), or the block is set again above its room — never over the content."""
    from verstka.rendering.compose import Composer

    deck, manifest, ws = sparse_env
    b = DeckBuilder(deck)
    slide = b.add_blank_slide(6)
    kit = Kit(manifest, LO_W, LO_H, "FFFFFF")
    o = _plan_slide()
    comp = Composer(slide, kit, o, DeckOutline(title="T", slides=[o]), "structured", manifest)
    area = Bbox(x=int(0.05 * LO_W), y=int(0.3 * LO_H), w=int(0.9 * LO_W), h=int(0.5 * LO_H))
    room = Bbox(x=area.x, y=area.y2, w=int(0.6 * LO_W), h=int(0.15 * LO_H))
    if foot_room:
        comp.foot_room = room
    n0 = len(comp.cv.tree)
    comp.compose("two_column", area)
    assert not any("speaker notes" in w for w in comp.warnings) and "Вывод" not in (o.notes or "")
    drawn = [((el.find(".//" + q("p:cNvPr")).get("name") or ""), element_bbox(el)) for el in list(comp.cv.tree)[n0:]]
    concl = [bb for name, bb in drawn if name.startswith("Conclusion") and bb]
    others = [bb for name, bb in drawn if not name.startswith("Conclusion") and bb]
    assert concl
    tol = int(0.01 * LO_H)
    top = min(bb[1] for bb in concl)
    assert top >= max(bb[1] + bb[3] for bb in others) - tol  # under the content, never over it
    if foot_room:
        assert any("free room" in w for w in comp.warnings)
        assert all(bb[0] + bb[2] <= room.x2 + tol and bb[1] + bb[3] <= room.y2 + tol for bb in concl)
    else:
        assert any("the block set" in w or "a size smaller" in w for w in comp.warnings), comp.warnings
        assert all(bb[1] + bb[3] <= area.y2 + tol for bb in concl + others)


def test_four_short_figures_stay_the_heroes_of_their_tiles(simple_deck, tmp_path):
    """B3: «70 млн / 110 млн / 62 / 80%» over «Погибло / …» — the figures are set a clear step over their labels (a row
    of four rather than a 2×2 grid at a text size), and the visual strip does not repeat a tile's figure («80%»)."""
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    nums = [NumberCallout(value=v, label=l) for v, l in [("70 млн", "Погибло"), ("110 млн", "Мобилизовано"), ("62", "Государств"), ("80%", "Населения Земли")]]
    o = OutlineSlide(id="s7", kind=PatternKind.stat_row, headline="Война унесла жизни 70 миллионов человек", content=SlideContent(numbers=nums), takeaway="Конфликт затронул 80% населения и стал крупнейшим в истории")
    _, out = _render((simple_deck, manifest, ws), [o], "stat_row", "visual")
    slide = out[0][0]
    figs, labels, strip_figs = [], [], []
    for el, _b in _boxes(slide):
        name = el.find(".//" + q("p:cNvPr")).get("name") or ""
        sizes = [int(r.get("sz")) / 100 for r in el.iter(q("a:rPr")) if r.get("sz")]
        if name.startswith("Figure"):
            figs.append(max(sizes))
        elif name.startswith("Label"):
            labels.append(max(sizes))
        elif name.startswith("Conclusion figure"):
            strip_figs.append("".join(t.text or "" for t in el.iter(q("a:t"))))
    assert len(figs) == 4 and labels
    assert min(figs) >= 1.6 * max(labels) - 0.05
    assert not strip_figs  # «80%» is a tile already: the strip is the plain one


def test_the_hero_figure_of_a_change_is_its_target():
    """B6: «Рентабельность увеличится с 13,3% до 22,4%» shows 22,4% — where the change goes, not the old value."""
    from verstka.rendering.compose import _aside_figure

    def fig(text: str) -> str:
        return " ".join((_aside_figure(text) or "").replace("\u00a0", " ").split())

    assert fig("Рентабельность увеличится с 13,3% до 22,4%") == "22,4%"
    assert fig("Выручка вырастет от 900 000 ₽ до 1 138 500 ₽") == "1 138 500 ₽"
    assert fig("Доля расходов 35% → 33%") == "33%"
    assert fig("Операционная прибыль — 120 000 рублей") == "120 000 ₽"
    assert fig("Скидка до 15% для 300 постоянных гостей") == "15%"


def test_heading_starts_after_a_small_ornament_in_its_band(simple_deck):
    """B4: a cluster of dots of the layout (a small drawn shape, ≥ 0.3 % of the slide) standing over the heading's left
    part: the heading starts after it (Grey Elegant); large art and art outside the band move nothing."""
    from verstka.rendering.synth import _left_clear

    b = DeckBuilder(simple_deck)
    slide = b.add_blank_slide(6)
    W, H = b.slide_w, b.slide_h
    dots = slide.shapes.add_shape(1, Emu(int(-0.015 * W)), Emu(int(0.012 * H)), Emu(int(0.091 * W)), Emu(int(0.161 * H)))
    dots.fill.solid()
    dots.fill.fore_color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    head = Bbox(x=int(0.046 * W), y=int(0.124 * H), w=int(0.704 * W), h=int(0.083 * H))
    past = _left_clear(slide, head, W, H)
    assert abs(past - (int(0.076 * W) + int(0.027 * W))) <= int(0.002 * W)  # its letters 0.012 W clear of the dots
    below = Bbox(x=head.x, y=int(0.4 * H), w=head.w, h=head.h)
    assert _left_clear(slide, below, W, H) == below.x
    big = b.add_blank_slide(6)
    tri = big.shapes.add_shape(1, Emu(0), Emu(0), Emu(int(0.2 * W)), Emu(int(0.4 * H)))  # 8 % of the slide: the canvas's art
    tri.fill.solid()
    tri.fill.fore_color.rgb = RGBColor(0xEE, 0x22, 0x22)
    assert _left_clear(big, head, W, H) == head.x


def test_a_trend_chart_never_leaves_a_label_to_be_broken_inside_its_word(simple_deck):
    """B5: LibreOffice breaks a category label wider than its slot inside the word («Сейч ас», tickLblSkip ignored).
    A trend chart whose labels overrun their slots turns them (or thins them out) on every kit, and reports it
    (`verstka_turned`; `verstka_squeeze` = longest label over its slot) — a chart pair compares it between plans."""
    from verstka.schemas.template import ChartStyleSpec, Typography

    b = DeckBuilder(simple_deck)
    slide = b.add_blank_slide(6)
    W, H = b.slide_w, b.slide_h
    spec = ChartSpec(type="line", categories=["Сейчас"] + [f"{i}-й мес." for i in range(1, 7)], series=[InlineSeries(name="Выручка", values=[900, 930, 970, 1010, 1060, 1100, 1138])], unit="₽")
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="c", kind=PatternKind.chart, headline="h", content=SlideContent(chart=spec))])
    wide = add_chart(slide, Bbox(x=0, y=0, w=int(0.9 * W), h=int(0.4 * H)), spec, outline, ChartStyleSpec(), Typography(), text_hex="222222")
    narrow = add_chart(slide, Bbox(x=0, y=int(0.5 * H), w=int(0.25 * W), h=int(0.45 * H)), spec, outline, ChartStyleSpec(), Typography(), text_hex="222222")
    assert wide.verstka_squeeze < narrow.verstka_squeeze and narrow.verstka_squeeze > 1.1
    assert not wide.verstka_turned and narrow.verstka_turned

    def turned(gf) -> bool:
        ax = gf.chart._chartSpace.find(".//" + q("c:catAx"))
        bp = ax.find(q("c:txPr") + "/" + q("a:bodyPr")) if ax is not None else None
        return bp is not None and (bp.get("rot") or "0") != "0"

    cats = [c for c in narrow.chart.plots[0].categories]
    assert not turned(wide)
    assert turned(narrow) or any(c == "" for c in cats)  # turned, or every other label left blank


# ------------------------------------------------------------------ gate 3, B3-1: category labels LibreOffice never breaks

COFFEE_MONTHS = ["Сейчас"] + [f"{i}-й месяц" for i in range(1, 7)]
COFFEE_REVENUE = [900000, 930000, 970000, 1015000, 1060000, 1100000, 1138500]


def _line_chart(prs, frame: tuple[float, float], family: str, size: float):
    """A coffee revenue line chart (the short brief's s5) in a frame of `frame` = (w, h) slide shares."""
    from verstka.schemas.template import ChartStyleSpec, Typography

    W, H = prs.slide_width, prs.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    spec = ChartSpec(type="line", categories=COFFEE_MONTHS, series=[InlineSeries(name="Выручка", values=COFFEE_REVENUE)], unit="₽")
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="c", kind=PatternKind.chart, headline="h", content=SlideContent(chart=spec))])
    box = Bbox(x=int(0.03 * W), y=int(0.26 * H), w=int(frame[0] * W), h=int(frame[1] * H))
    return add_chart(slide, box, spec, outline, ChartStyleSpec(font_family=family, font_size_pt=size), Typography(primary_family=family), text_hex="222222")


def _lo_label_faults(gf) -> list[str]:
    """What LibreOffice would do to the chart's category labels (VCartesianAxis, measured in bfix6/lo): a straight
    label wider than 95 % of its tick distance wraps — a word wider than that is broken inside; a turned one is cut to
    «…» when its turned height and one average character pass the room under the axis; turned neighbours closer than
    their line height overlap."""
    from verstka.rendering.fonts import text_width_pt

    cs = gf.chart._chartSpace
    ax = cs.find(".//" + q("c:catAx"))
    rpr = ax.find(q("c:txPr") + "/" + q("a:p") + "/" + q("a:pPr") + "/" + q("a:defRPr"))
    size = int(rpr.get("sz")) / 100
    family = rpr.find(q("a:latin")).get("typeface")
    rot = int(ax.find(q("c:txPr") + "/" + q("a:bodyPr")).get("rot") or 0)
    lay = cs.find(".//" + q("c:plotArea") + "/" + q("c:layout") + "/" + q("c:manualLayout"))
    g = {k: float(lay.find(q("c:" + k)).get("val")) for k in ("x", "y", "w", "h")}
    fw, fh = gf.width / EMU_PER_PT, gf.height / EMU_PER_PT
    cats = [str(c) if c is not None else "" for c in gf.chart.plots[0].categories]
    slot = g["w"] * fw / len(cats)
    room = fh * (1 - g["y"] - g["h"])
    faults = []
    shown = [j for j, c in enumerate(cats) if c and c != "None"]
    for j in shown:
        c = cats[j]
        w = text_width_pt(c, family, size)
        if rot == 0:
            if any(text_width_pt(word, family, size) > 0.95 * slot for word in c.split()):
                faults.append(f"«{c}» broken inside a word")
            elif w > 0.95 * slot:
                faults.append(f"«{c}» wraps")
        elif 0.7071 * (w + 1.2 * size) - room > -w / len(c):
            faults.append(f"«{c}» cut to «…»")
    if rot != 0 and len(shown) > 1:
        gap = min(b - a for a, b in zip(shown, shown[1:])) * slot * 0.7071
        if gap < 1.2 * size:
            faults.append("turned labels overlap")
    return faults


@pytest.mark.skipif(__import__("sys").platform != "darwin" or not __import__("verstka.rendering.fonts", fromlist=["is_installed"]).is_installed("Play"),
                    reason="the live frame was measured on macOS with Play installed; elsewhere the labels are measured in "
                    "the face LibreOffice will draw there and may rightly be turned")
def test_coffee_months_at_half_width_on_vk_tech_are_shortened_not_broken():
    """B3-1: the live coffee deck on VK Tech (10 × 5.63 in, Play 9 pt) set its revenue chart at half width and
    LibreOffice broke «1-й мес.» into «1-й ме / с.» (33.8 pt in a 31.3 pt slot: under the old 10 % tolerance). Every
    label is measured against 95 % of its slot: «1-й месяц» → «1-й мес.» → «1 мес.», straight, all whole."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    gf = _line_chart(prs, (4031483 / 9144000, 1774287 / 5143500), "Play", 9.0)  # the live frame (20260928-034442-97788d s5)
    cats = [str(c) for c in gf.chart.plots[0].categories]
    assert cats == ["Сейчас"] + [f"{i} мес." for i in range(1, 7)]
    assert _lo_label_faults(gf) == []
    assert not gf.verstka_turned


@pytest.mark.parametrize("plot_w", [0.35, 0.29])
def test_six_month_labels_at_narrow_widths_are_shortened_turned_or_thinned_never_cut(plot_w):
    """B3-1: at 0.35 W and 0.29 W the six months no longer fit straight even shortened: they turn by 45° with room under
    the axis for their turned height (LibreOffice cuts a turned label to «1-й …» otherwise — Handdrawn at gate 3), and
    every other one is left out when turned neighbours would touch. Nothing is broken, wrapped or cut."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    gf = _line_chart(prs, (plot_w, 0.345), "Play", 9.0)
    assert _lo_label_faults(gf) == []
    cats = [str(c) for c in gf.chart.plots[0].categories if c]
    assert all(c in (["Сейчас"] + [f"{i} мес." for i in range(1, 7)]) for c in cats if c != "None")
    assert gf.verstka_turned


def test_turned_labels_get_the_room_libreoffice_measures_under_the_axis():
    """B3-1: Handdrawn (13.33 in, Verdana 14 pt) turned «1-й мес.» at 0.41 W and reserved 60.9 pt under the axis for a
    54.4 pt turned height — LibreOffice wants one more average character (6.7 pt) and cut every month to «1-й …»."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    gf = _line_chart(prs, (5029200 / 12192000, 2337491 / 6858000), "Verdana", 14.0)
    assert _lo_label_faults(gf) == []


def test_period_levels_shorten_a_time_axis_step_by_step():
    from verstka.rendering.charts import fit_categories, period_levels

    levels = period_levels(COFFEE_MONTHS)
    assert [lv[1] for lv in levels] == ["1-й месяц", "1-й мес.", "1 мес."]
    assert all(" " not in c and "⁠" not in c for lv in levels for c in lv)  # a forced wrap breaks between words
    assert period_levels(["Q1", "Q2"]) == [["Q1", "Q2"]]
    assert period_levels(["2 квартал", "Январь 2026"])[-1] == ["2 кв.", "Янв 2026"]
    # the widest label decides: «Сейчас» fits a 31.3 pt slot (29.4 ≤ 0.95 × 31.3), «1-й мес.» (33.8) does not
    fit = fit_categories(COFFEE_MONTHS, "Play", [9.0], 31.3, 60.0)
    assert fit.mode == "flat" and fit.level == 2
    # no room under the axis for turned labels: whole words on two lines when every word fits its slot
    fit = fit_categories(["Сейчас", "Через месяц", "Через год"], "Play", [9.0], 36.0, 30.0, turn_first=False)
    assert fit.mode == "wrap" and fit.lines == 2


def test_a_conclusion_is_said_aloud_once_and_never_while_the_slide_shows_it(sparse_env, monkeypatch):
    """C (gate 3): a saved plan rendered again carried «Вывод: …» from its first setting — the notes said the
    conclusion twice (short visual s4 on 12 decks) or said it while the slide showed it. The line an earlier setting
    wrote is dropped before the slide is set again; a conclusion with no room is said once."""
    from verstka.rendering.compose import Composer

    deck, manifest, ws = sparse_env
    take = "Все разовые вложения покрываются бюджетом запуска"
    said = f"Вывод: {take}."

    def setting(area_h: float, notes: str) -> OutlineSlide:
        b = DeckBuilder(deck)
        slide = b.add_blank_slide(6)
        o = OutlineSlide(id="b", kind=PatternKind.bullets, headline="Вложения", content=SlideContent(bullets=["Витрина для десертов", "Программа лояльности"]), takeaway=take, notes=notes)
        comp = Composer(slide, Kit(manifest, LO_W, LO_H, "FFFFFF"), o, DeckOutline(title="T", slides=[o]), "structured", manifest)
        comp.compose("bullets", Bbox(x=int(0.05 * LO_W), y=int(0.3 * LO_H), w=int(0.9 * LO_W), h=int(area_h * LO_H)))
        return o

    shown = setting(0.55, f"{said} Слова спикера.")
    assert said not in shown.notes and "Слова спикера." in shown.notes
    # no room anywhere (nor a lead line): the conclusion is said aloud — once
    monkeypatch.setattr(Composer, "_conclusion_spot", lambda self, n0, last=False: None)
    monkeypatch.setattr(Composer, "_rehouse_conclusion", lambda self, *a, **kw: None)
    again = setting(0.55, f"{said} Слова спикера.")
    assert again.notes.count("Вывод:") == 1 and again.notes.startswith(said)


# ------------------------------------------------------------------ gate 3, B3-2 / B3-4 / B3-5: text sized for its room


def _composer(env, o: OutlineSlide, strategy: str = "structured"):
    from verstka.rendering.compose import Composer

    deck, manifest, ws = env
    b = DeckBuilder(deck)
    slide = b.add_blank_slide(6)
    return Composer(slide, Kit(manifest, LO_W, LO_H, "FFFFFF"), o, DeckOutline(title="T", slides=[o], strategy=strategy), strategy, manifest)


def _sizes_by_name(comp, n0: int) -> dict:
    out: dict = {}
    for el in list(comp.cv.tree)[n0:]:
        name = (el.find(".//" + q("p:cNvPr")).get("name") or "").rsplit(" ", 1)[0]
        for r in el.iter(q("a:rPr")):
            if r.get("sz"):
                out.setdefault(name, set()).add(int(r.get("sz")) / 100)
    return out


def test_a_list_that_misses_its_size_by_a_hair_closes_its_rows_up_before_the_type_steps_down(sparse_env):
    """B3-2 (WWII s5): four dated lines missed the lead size by 0.02 pt and fell to the body size four points down,
    25 % of the slide left empty under them. The hairline gaps close up first; the lead size stays."""
    from verstka.rendering.compose import Composer, _emu

    texts = ["22 июня 1941 — Германия начала вторжение в СССР", "7 декабря 1941 — Япония атаковала Перл-Харбор", "1942 — Япония потерпела поражение в битве за Мидуэй", "1942 — Советский Союз начал серию побед, включая Сталинградскую битву"]
    o = OutlineSlide(id="b", kind=PatternKind.bullets, headline="К 1942 году война стала мировой", content=SlideContent(bullets=texts))
    comp = _composer(sparse_env, o)
    k = comp.kit
    w = int(0.8 * LO_W)
    text_w = int(w * 0.8)
    # the list's height at the lead size with the usual gaps (0.75 of a line over and under each hairline)
    marker_w = _emu(k.lead * 1.6)
    col_h = sum(comp.h([comp.P(t, k.lead, k.colors.text)], text_w - marker_w) for t in texts) + (len(texts) - 1) * 2 * int(max(k.lead * 0.75, 6) * EMU_PER_PT)
    area = Bbox(x=int(0.1 * LO_W), y=int(0.3 * LO_H), w=w, h=col_h - _emu(0.5))
    n0 = len(comp.cv.tree)
    comp.bullets(area, texts)
    assert min(_sizes_by_name(comp, n0)["Item"]) >= k.lead > k.body  # never the body size four points down


def test_theses_too_long_for_four_slivers_take_two_rows_and_a_readable_size(sparse_env):
    """B3-2 (WWII visual s10, Гагарин s6): four theses in a row of narrow cards were set at 12 pt (a word or two a line
    at any larger size), 60 % of each card empty. The other grid and the cards without their index numerals are
    tried: the words read at the body size or larger."""
    theses = ["Была создана Организация Объединённых Наций (ООН) для предотвращения будущих конфликтов", "Организованы международные трибуналы для осуждения лидеров стран-агрессоров", "Произошли изменения в границах государств", "Были установлены репарации"]
    o = OutlineSlide(id="c", kind=PatternKind.bullets, headline="Итоги войны", content=SlideContent(bullets=theses))
    comp = _composer(sparse_env, o, "visual")
    k = comp.kit
    area = Bbox(x=int(0.055 * LO_W), y=int(0.34 * LO_H), w=int(0.89 * LO_W), h=int(0.56 * LO_H))
    n0 = len(comp.cv.tree)
    comp.bullets(area, theses)
    sizes = _sizes_by_name(comp, n0)
    text = sizes.get("Card text") or sizes.get("Item")
    assert text and min(text) >= k.body - 0.05


def test_a_legend_name_is_written_as_the_lines_it_was_measured_in(sparse_env):
    """B3-4 (Marketing sv5): «Обновление меню и фотографии» was measured on two lines and its row sized for them; the
    renderer's narrower face set it on one and left a blank line under it. The name is written as its measured lines."""
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="h", content=SlideContent())
    comp = _composer(sparse_env, o)
    k = comp.kit
    name = "Обновление меню и фотографии"
    one = comp._set_lines(name, k.body, "222222", Emu(int(600 * EMU_PER_PT)))
    two = comp._set_lines(name, k.body, "222222", Emu(int(text_width_pt_(name, k.font, k.body) * 0.7 * EMU_PER_PT)))
    assert len(one) == 1 and one[0].text.replace(" ", " ") == name
    assert len(two) == 2 and " ".join(p.text for p in two).replace(" ", " ") == name
    assert comp.h(two, Emu(int(text_width_pt_(name, k.font, k.body) * 0.7 * EMU_PER_PT))) == comp.h([comp.P(name, k.body, "222222")], Emu(int(text_width_pt_(name, k.font, k.body) * 0.7 * EMU_PER_PT)))


def text_width_pt_(text, font, size):
    from verstka.rendering.fonts import text_width_pt

    return text_width_pt(text, font, size)


@pytest.mark.parametrize("strategy", ["structured", "compact"])
def test_a_conclusion_is_kept_on_the_slide_without_text_under_two_percent(sparse_env, strategy):
    """B3-5 (Focus / Nature long s8): to keep the conclusion on the slide the two lists shrank to 1.76 % of the slide
    height (8 pt on 6.2″). The block never goes under 2 % for it: the conclusion is set a size smaller, as the strip
    at the foot (compact), or — last — said aloud."""
    o = _plan_slide()
    comp = _composer(sparse_env, o, strategy)
    k = comp.kit
    area = Bbox(x=int(0.2 * LO_W), y=int(0.3 * LO_H), w=int(0.62 * LO_W), h=int(0.55 * LO_H))
    n0 = len(comp.cv.tree)
    comp.compose("two_column", area)
    sizes = _sizes_by_name(comp, n0)
    body = sizes.get("Column", set())
    assert body and min(body) >= 0.02 * k.hpt - 0.05
    said = "Вывод:" in (o.notes or "")
    assert ("Conclusion" in sizes) != said  # on the slide or said aloud, once


@pytest.mark.parametrize("frame_w", [0.3, 0.2])
def test_a_column_label_word_is_never_left_to_be_broken_inside(frame_w):
    """B3-1 for columns: a single word wider than 95 % of its column's slot («Нидерланды» in a narrow chart) is set
    a size smaller or the labels turn — never left for LibreOffice to break inside the word."""
    from verstka.schemas.template import ChartStyleSpec, Typography

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    spec = ChartSpec(type="column", categories=["Нидерланды", "Франция", "Германия", "Португалия"], series=[InlineSeries(name="Потери", values=[450, 600, 7000, 450])])
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="c", kind=PatternKind.chart, headline="h", content=SlideContent(chart=spec))])
    gf = add_chart(slide, Bbox(x=0, y=0, w=int(frame_w * 9144000), h=int(0.4 * 5143500)), spec, outline, ChartStyleSpec(font_family="Play", font_size_pt=9), Typography(primary_family="Play"), text_hex="222222")
    assert gf.chart._chartSpace.find(".//" + q("c:barDir")).get("val") == "col"
    assert _lo_label_faults(gf) == []


@pytest.mark.parametrize("region", [(0.72, 0.5), (0.6, 0.55), (0.5, 0.6)])
def test_a_pair_of_charts_never_flattens_its_trend_for_a_barely_larger_pie(sparse_env, region):
    """B3-3 (Sidebar43 / Focus short s5): the stacked plan set the growth line 0.80 × 0.20 H for a pie 4 % smaller. A plan
    past the first is taken only for a real gain (the pie 15 % larger or of its least size, its legend clean or keeping
    the amounts) and, stacked, only while the trend keeps a plot of 0.28 H."""
    from verstka.rendering.charts import chart_plot_bbox

    line = ChartSpec(type="line", title="Прогноз выручки", categories=COFFEE_MONTHS, series=[InlineSeries(name="Выручка", values=COFFEE_REVENUE)], unit="₽")
    pie = ChartSpec(type="pie", title="Распределение вложений", categories=["Витрина для десертов", "Программа лояльности и учет", "Обновление меню и фотографии", "Обучение сотрудников", "Резерв"], series=[InlineSeries(name="Вложения", values=[70000, 35000, 25000, 20000, 30000])], unit="₽")
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Выручка вырастет на 26,5%", content=SlideContent(chart=line, chart2=pie), takeaway="Траектория роста выручки на 26,5% за 6 месяцев")
    comp = _composer(sparse_env, o)
    comp.compose("chart_pair", Bbox(x=int(0.15 * LO_W), y=int(0.3 * LO_H), w=int(region[0] * LO_W), h=int(region[1] * LO_H)))
    frames = [sh for sh in comp.slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
    trend = next(sh for sh in frames if "LINE" in str(sh.chart.chart_type))
    other = next(sh for sh in frames if sh is not trend)
    stacked = other.top >= trend.top + trend.height - int(0.01 * LO_H)
    assert not stacked or chart_plot_bbox(trend).h >= 0.28 * LO_H


def test_the_value_axis_keeps_room_for_its_ticks_with_their_unit():
    """B3-3 follow-up: the room left of the plot was measured for «1 200 000» while the axis prints «1 200 000 ₽»; in
    a narrow pair plan LibreOffice cut every tick to «1 200 0…». The ticks are measured as printed."""
    from verstka.rendering.charts import chart_plot_bbox
    from verstka.rendering.fonts import text_width_pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(6858000)
    gf = _line_chart(prs, (0.37, 0.39), "Trebuchet MS", 14.0)
    va = gf.chart._chartSpace.find(".//" + q("c:valAx"))
    assert va.find(q("c:delete")).get("val") == "0"  # the axis is shown in this frame
    left = (chart_plot_bbox(gf).x - gf.left) / EMU_PER_PT
    assert left >= text_width_pt("1 200 000 ₽", "Trebuchet MS", 14.0) + 0.45 * 14.0 - 0.1  # the tick, with its unit, and a gap


# ------------------------------------------------------------------ gate 4 (round 5): G4-11, G4-18, G4-19, G4-20, requests

SQ_W, SQ_H = 9144000, 6858000  # 4:3


@pytest.fixture
def sparse43_env(tmp_path):
    deck = build_sparse_deck(tmp_path / "sparse43.pptx", SQ_W, SQ_H, title_pt=40, body_pt=24)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return deck, manifest, ws


def _composer_at(env, o: OutlineSlide, strategy: str, W: int, H: int):
    from verstka.rendering.compose import Composer

    deck, manifest, ws = env
    slide = DeckBuilder(deck).add_blank_slide(6)
    return Composer(slide, Kit(manifest, W, H, "FFFFFF"), o, DeckOutline(title="T", slides=[o], strategy=strategy), strategy, manifest)


def _drawn(comp, n0: int, base: str) -> list:
    """(element, bbox) of the shapes named `base N` composed after the first `n0` elements."""
    out = []
    for el in list(comp.cv.tree)[n0:]:
        name = (el.find(".//" + q("p:cNvPr")).get("name") or "").rsplit(" ", 1)[0]
        if name == base:
            out.append((el, element_bbox(el)))
    return out


def _one_line_width_pt(el) -> list[float]:
    """Per paragraph of a text shape, its runs' width on one line (pt), measured in their own face and size."""
    from verstka.rendering.fonts import text_width_pt

    widths = []
    for p in el.iter(q("a:p")):
        w = 0.0
        for r in p.iter(q("a:r")):
            rpr = r.find(q("a:rPr"))
            latin = rpr.find(q("a:latin"))
            w += text_width_pt(r.find(q("a:t")).text or "", latin.get("typeface") if latin is not None else None, int(rpr.get("sz")) / 100, rpr.get("b") == "1")
        widths.append(w)
    return widths


@pytest.mark.parametrize("strategy", ["structured", "visual", "compact"])
@pytest.mark.parametrize("region", [(0.9, 0.62), (0.55, 0.4), (0.4, 0.3), (0.3, 0.25)])
def test_a_pair_of_charts_on_a_4_3_slide_always_lands_a_plan(sparse43_env, strategy, region):
    """C's report (Classic43 / Handdrawn / LO_DNA short s5): the pair's fallback took a plan by an index past its list
    (IndexError) and the slide fell to a picture. Whatever the region — every plan clean or none — a plan is drawn."""
    line = ChartSpec(type="line", title="Прогноз выручки", categories=COFFEE_MONTHS, series=[InlineSeries(name="Выручка", values=COFFEE_REVENUE)], unit="₽")
    pie = ChartSpec(type="pie", title="Распределение вложений", categories=["Витрина для десертов", "Программа лояльности и учет", "Обновление меню и фотографии", "Обучение сотрудников", "Резерв"], series=[InlineSeries(name="Вложения", values=[70000, 35000, 25000, 20000, 30000])], unit="₽")
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Выручка вырастет на 26,5%", content=SlideContent(chart=line, chart2=pie, bullets=["Ежемесячная выручка — 900 000 рублей", "Прогноз: 1 138 500 ₽ на 6-й месяц"]), takeaway="Траектория роста выручки на 26,5% за 6 месяцев")
    comp = _composer_at(sparse43_env, o, strategy, SQ_W, SQ_H)
    comp.compose("chart_pair", Bbox(x=int(0.06 * SQ_W), y=int(0.3 * SQ_H), w=int(region[0] * SQ_W), h=int(region[1] * SQ_H)))
    frames = [sh for sh in comp.slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
    assert frames and not any("failed" in w for w in comp.warnings)


def _figures_on_one_line(comp, n0: int) -> None:
    figs = _drawn(comp, n0, "Figure")
    assert figs
    for el, (x, y, w, h) in figs:
        widths = _one_line_width_pt(el)
        assert max(widths) <= w / EMU_PER_PT + 0.5, (widths, w / EMU_PER_PT)  # one line in its box: never wrapped


@pytest.mark.parametrize("tall", [True, False])
def test_a_figure_never_wraps_in_its_tile(sparse_env, tall):
    """Focus live s3: «около 615 тыс. тонн» in a third of a narrow area was set at the tile's least size and still broke
    over its accent rule (text_overflow). Long figures stand one under the other with their labels beside them (a tall
    area), else they step down until each one reads on one line."""
    nums = [NumberCallout(value="17,9%", label="выработке электроэнергии в 2024 году"), NumberCallout(value="около 3800 тонн", label="Потребление урана в год"), NumberCallout(value="около 615 тыс. тонн", label="Разведанные запасы урана")]
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Доля атомной энергетики — 17,9%", content=SlideContent(numbers=nums))
    comp = _composer(sparse_env, o, "compact")
    k = comp.kit
    area = Bbox(x=int(0.2 * LO_W), y=int(0.35 * LO_H), w=int(0.5 * LO_W), h=int((0.5 if tall else 0.16) * LO_H))
    n0 = len(comp.cv.tree)
    comp.kpis(area, nums, [])
    _figures_on_one_line(comp, n0)
    if tall:
        assert any("one under the other" in w for w in comp.warnings)
        sizes = _sizes_by_name(comp, n0)["Figure"]
        assert max(sizes) >= k.h2  # the figures read as figures, not as 16 pt captions


def test_a_hedge_before_a_figure_is_set_at_its_units_size(sparse_env):
    """«около 615 тыс. тонн»: the hedge and the unit ride small on the figure's baseline, the number keeps the size."""
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="h", content=SlideContent())
    comp = _composer(sparse_env, o)
    k = comp.kit
    size = max(k.sizes)
    para, width = comp.figure_para("около 615 тыс. тонн", size, "111111", "777777")
    texts = [(r.text.replace(" ", " ").strip(), r.size) for r in para.runs]
    assert texts[0][0] == "около" and texts[0][1] < size
    assert texts[1] == ("615", size)
    assert texts[2][0] == "тыс. тонн" and texts[2][1] < size
    whole = comp.figure_para("более 10", size, "111111", "777777")[0]
    assert [r.text.replace(" ", " ").strip() for r in whole.runs] == ["более", "10"]


@pytest.mark.parametrize("unit,values", [(None, [42, 32, 18.6, 7.4]), ("%", [42, 32, 18.6, 7.4]), ("%", [60, 25, 15])])
def test_a_pie_of_percents_says_each_part_once(sparse_env, unit, values):
    """G4-11 (R_vk s6): «Доходы VK по направлениям, %» lost its unit and its legend read «42 · 42%», «18,6 · 19%»; the
    slices «42,0». A pie whose unit is «%», or whose parts add up to 100 ± 1, says each part once, as written."""
    from verstka.rendering.charts import pie_values_are_percents

    cats = ["Онлайн-реклама", "Онлайн-игры", "Платные сервисы", "Новые проекты"][: len(values)]
    spec = ChartSpec(type="pie", title="Доходы VK по направлениям", categories=cats, series=[InlineSeries(name="Доходы", values=values)], unit=unit)
    assert pie_values_are_percents(unit, sum(values))
    o = OutlineSlide(id="p", kind=PatternKind.chart, headline="Доходы VK", content=SlideContent(chart=spec))
    comp = _composer(sparse_env, o)
    n0 = len(comp.cv.tree)
    comp.compose("chart", Bbox(x=int(0.05 * LO_W), y=int(0.25 * LO_H), w=int(0.9 * LO_W), h=int(0.65 * LO_H)))
    shares = [etree.tostring(el, method="text", encoding="unicode").replace(" ", " ").strip() for el, _ in _drawn(comp, n0, "Share")]
    want = [f"{v:g}".replace(".", ",") + "%" for v in sorted(values, reverse=True)]
    assert shares == want and not any("·" in s for s in shares)
    gf = next(sh for sh in comp.slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart)
    assert list(gf.chart.plots[0].series[0].values) == sorted(values, reverse=True)  # the slices say the values themselves
    assert not pie_values_are_percents(None, 180000) and not pie_values_are_percents("₽", 100)


def test_captions_under_charts_step_down_before_the_conclusion_leaves_its_strip(sparse_env):
    """G4-19 (Nature short visual s4): set again above the conclusion's room, three captions under the charts stayed at
    the lead size in three columns, three lines each, fell past the room — and the hero figure «1 138 500 ₽» gave way to
    a line under the heading. They step down (body, small; three lines → two columns; never under 2 %) first."""
    o = OutlineSlide(id="v", kind=PatternKind.chart, headline="h", content=SlideContent())
    comp = _composer(sparse_env, o, "visual")
    k = comp.kit
    lines = ["Комбо с напитками и выпечкой", "Программа лояльности", "Предложения для сотрудников ближайших офисов"]
    area = Bbox(x=int(0.05 * LO_W), y=int(0.22 * LO_H), w=int(0.89 * LO_W), h=int(0.37 * LO_H))
    plain = comp._under_plan(area, lines, spread=True)
    comp._rehousing, comp._floor_2pct = True, True
    again = comp._under_plan(area, lines, spread=True)
    assert again["eh"] <= plain["eh"]
    assert again["eh"] + int(k.vgap * 1.4) <= 0.3 * area.h or again["eh"] < plain["eh"]
    sizes = {r.size for col in again["paras"] for p in col for r in p.runs}
    assert min(sizes) >= 0.02 * k.hpt - 0.05


def _theses() -> list[SlideItem]:
    return [SlideItem(title=t) for t in ("Партнёрства с пятью ближайшими офисами", "Программа лояльности с понятными условиями", "Продвижение в районных сообществах", "Дневные предложения с 15:00 до 18:00", "Рост маркетингового бюджета до 35 000 рублей в месяц")]


def test_five_theses_keep_their_numerals_where_the_room_allows(sparse_env):
    """G4-18 (dataset long visual s6: LCT, VK Education, WorkSpace lost «01…05», VK Tech kept them): the numerals are tried
    at their own size, a step over the title, hanging beside the first line — and kept while the words stay within a
    step of the size they reach without them."""
    o = OutlineSlide(id="c", kind=PatternKind.cards, headline="Цель — увеличить число покупок", content=SlideContent(items=_theses()))
    comp = _composer(sparse_env, o, "visual")
    area = Bbox(x=int(0.05 * LO_W), y=int(0.3 * LO_H), w=int(0.9 * LO_W), h=int(0.36 * LO_H))
    n0 = len(comp.cv.tree)
    comp.cards(area, _theses())
    idx = _drawn(comp, n0, "Index")
    assert [etree.tostring(el, method="text", encoding="unicode").strip() for el, _ in idx] == ["01", "02", "03", "04", "05"]
    texts = _drawn(comp, n0, "Card text")
    for (iel, ib), (tel, tb) in zip(idx, texts):
        # a numeral never stands on its words: over them, or beside their first line
        assert ib[1] + ib[3] <= tb[1] + int(0.003 * LO_H) or ib[0] + ib[2] <= tb[0] + int(0.003 * LO_W)


def test_card_words_of_a_row_start_on_one_line(sparse_env):
    """G4-18 (Focus / Sidebar s4): «Свободные часы» on two lines pushed its card's words a line below its neighbours'."""
    items = [SlideItem(title="Средний чек", text="Только 20% чеков содержат еду"), SlideItem(title="Свободные часы после обеда", text="4 из 18 мест заняты после 15:00"), SlideItem(title="Потери", text="Списания продуктов — 27 000 ₽ в месяц")]
    o = OutlineSlide(id="c", kind=PatternKind.cards, headline="h", content=SlideContent(items=items))
    comp = _composer(sparse_env, o, "structured")
    k = comp.kit
    ts = k.lead
    inner = Emu(int(text_width_pt_("Свободные часы после", k.font, ts) * 1.02 * EMU_PER_PT))
    pads = comp._title_pads(items, 3, inner, ts)
    assert pads[1] == 0 and pads[0] == pytest.approx(ts * k.line) and pads[2] == pytest.approx(ts * k.line)
    starts = []  # where each card's words start under its title (pt): the title's lines and the gap under them
    for it, pad in zip(items, pads):
        title = comp._card_paras(it, ts, k.body, k.card.colors, title_pad=pad)[0]
        starts.append(para_lines_(title, inner) * ts * k.line + title.space_after)
    assert max(starts) - min(starts) < 0.01
    # the card's measured height counts the lengthened gap (a card never shorter than its words)
    h_pad = comp._card_content_h(items[0], inner, ts, k.body, None, k.card.colors, 0, pads[0])
    assert h_pad - comp._card_content_h(items[0], inner, ts, k.body, None, k.card.colors, 0) == pytest.approx(_emu_(pads[0]), abs=_emu_(0.5))


def para_lines_(para, width_emu) -> int:
    from verstka.rendering.compose import para_lines

    return para_lines(para, width_emu / EMU_PER_PT)


def _emu_(pt: float) -> int:
    return int(pt * EMU_PER_PT)


def test_two_sentences_in_cards_read_as_statements(sparse_env):
    """W2's request (EV compact s7): two title-only cards at the body size in tall cards — 14 % of the slide in lines.
    One or two sentences in cards take the callout size under the heading, down to the lead size."""
    items = [SlideItem(title="Эксперты McKinsey включили электромобили в число революционных технологий"), SlideItem(title="В России растёт интерес к электромобилям")]
    o = OutlineSlide(id="c", kind=PatternKind.cards, headline="Электромобили — революционная технология", content=SlideContent(items=items))
    comp = _composer(sparse_env, o, "compact")
    k = comp.kit
    n0 = len(comp.cv.tree)
    comp.compose("cards", Bbox(x=int(0.05 * LO_W), y=int(0.3 * LO_H), w=int(0.9 * LO_W), h=int(0.55 * LO_H)))
    assert min(_sizes_by_name(comp, n0)["Card text"]) > k.h2 + 0.05  # past the usual cap (HEAD: the h2 size at most)


def test_six_two_line_points_are_not_squeezed_to_eleven_lines(sparse_env):
    """G4-20 (Focus long compact s8): a column of six two-line steps is twelve lines; the «≤ 11 lines» cap sent it to the
    least size with a third of the card empty."""
    o = _plan_slide()
    comp = _composer(sparse_env, o, "compact")
    k = comp.kit
    items = list(o.content.items)
    n0 = len(comp.cv.tree)
    comp.columns(Bbox(x=int(0.05 * LO_W), y=int(0.25 * LO_H), w=int(0.8 * LO_W), h=int(0.7 * LO_H)), items)
    sizes = _sizes_by_name(comp, n0)
    assert min(sizes["Column"]) >= k.body - 0.05  # HEAD: 10 pt (eleven lines) where the body size fits the card


@pytest.mark.parametrize("env_name,W,H", [("sparse43_env", SQ_W, SQ_H), ("sparse_env", LO_W, LO_H)])
def test_six_steps_with_their_budget_keep_two_percent(env_name, W, H, request):
    """G4-20 / B's open item (long visual s8 on Sidebar43, Focus, Nature, Handdrawn): six steps in a row with five
    budget lines under them were set under 2 % of the slide height. The steps stand one under the other with the
    lines beside them at ≥ 2 %."""
    env = request.getfixturevalue(env_name)
    steps = [SlideItem(title=f"{i}-й месяц", text=t) for i, t in enumerate(["Учет показателей и обновление меню", "Запуск комбо и обучение сотрудников", "Запуск программы лояльности и партнерств с офисами", "Продвижение дневных предложений и настройка закупок", "Корректировка предложений по результатам продаж", "Оценка результатов и закрепление удачных решений"], 1)]
    budget = ["Витрина для десертов — 70 000 ₽", "Настройка программы лояльности и учета — 35 000 ₽", "Обновление меню, фотографии и оформление — 25 000 ₽", "Обучение сотрудников — 20 000 ₽", "Резерв — 30 000 ₽"]
    o = OutlineSlide(id="t", kind=PatternKind.timeline, headline="Разовые вложения составляют 180 000 рублей", content=SlideContent(items=steps, bullets=budget), takeaway="Все разовые вложения покрываются бюджетом запуска")
    comp = _composer_at(env, o, "visual", W, H)
    k = comp.kit
    n0 = len(comp.cv.tree)
    comp.compose("process", Bbox(x=int(0.16 * W), y=int(0.26 * H), w=int(0.79 * W), h=int(0.63 * H)))
    sizes = _sizes_by_name(comp, n0)
    body = sizes.get("Step", set()) | sizes.get("Note", set())
    assert body and min(body) >= 0.02 * k.hpt - 0.05, sizes
    assert "Вывод" not in (o.notes or "") and "Conclusion" in sizes
