"""Template slides used as they are designed: what the analysis reads on a sample (pictures, accent cells, note fields,
icon frames), how the matcher keeps a sample's story off new content, and the gate that keeps a cloned slide from
leaving a hole where the sample had its content."""

from __future__ import annotations

import io

import numpy as np
from PIL import Image, ImageDraw

from verstka.analysis.groups import detect_repeat_groups
from verstka.analysis.pictures import _Stats, _glyphs, classify_blob
from verstka.analysis.patterns import note_field
from verstka.analysis.roles import is_empty_frame
from verstka.analysis.shapes import ParagraphInfo, RunInfo, ShapeInfo, TextInfo
from verstka.rendering.clone_check import _GRID_C, _GRID_R, largest_hole
from verstka.rendering.groups import staggered_axis
from verstka.ru import NBSP, no_orphan, typeset
from verstka.schemas.common import Bbox, BboxFrac, Family, PatternKind, ShapeKind, SlotRole
from verstka.schemas.template import Capacity, Pattern, Slot, SlotStyle

W, H = 12192000, 6858000


def _box(x, y, w, h) -> Bbox:
    return Bbox(x=int(x * W), y=int(y * H), w=int(w * W), h=int(h * H))


def _text(sid, text, size, x, y, w, h, **kw) -> ShapeInfo:
    ti = TextInfo(paragraphs=[ParagraphInfo(text=text, runs=[RunInfo(text=text, font="Play", size_pt=size, bold=False, color_hex="000000")])])
    return ShapeInfo(id=sid, name=sid, kind=ShapeKind.sp, bbox=_box(x, y, w, h), z=kw.pop("z", 5), text=ti, **kw)


def _shape(sid, x, y, w, h, fill="EBF3F9", geometry="ellipse", line=None, z=1, **kw) -> ShapeInfo:
    return ShapeInfo(id=sid, name=sid, kind=ShapeKind.sp, bbox=_box(x, y, w, h), z=z, geometry=geometry, fill_hex=fill, line_hex=line, **kw)


# ---------------------------------------------------------------------------- pictures


def _png(draw) -> np.ndarray:
    im = Image.new("RGBA", (600, 300), (0, 0, 0, 0))
    draw(ImageDraw.Draw(im))
    return np.asarray(im).astype(np.int32)


def test_a_chart_drawn_as_a_picture_shows_its_labels_a_render_does_not():
    def chart(d):
        for i, y in enumerate(range(40, 260, 44)):
            d.line([(60, y), (580, y)], fill=(200, 220, 200, 255), width=1)
            d.text((10, y - 6), str(40 - 10 * i), fill=(200, 220, 200, 255))
        for i in range(5):
            d.text((90 + i * 100, 270), f"Item {i + 1}", fill=(200, 220, 200, 255))
        d.line([(80, 200), (200, 120), (320, 160), (440, 60), (560, 70)], fill=(190, 230, 120, 255), width=3)

    def render(d):
        d.ellipse([150, 30, 450, 280], fill=(40, 120, 255, 255))
        d.ellipse([250, 90, 350, 190], fill=(200, 230, 255, 255))

    assert _glyphs(_png(chart)) >= 12
    assert _glyphs(_png(render)) <= 2


def test_glyphs_make_a_picture_the_samples_content_and_ground_needs_the_whole_slide():
    st = _Stats(coverage=0.02, colours=40, texture=0.0, sharp=0.0, hue_sd=0.0, white=0.0, glyphs=20)
    assert classify_blob(st, 0.9, 0.6) == "content"  # a transparent line chart with its axis labels
    faint = _Stats(coverage=0.02, colours=40, texture=0.0, sharp=0.0, hue_sd=0.0, white=0.0, glyphs=0)
    assert classify_blob(faint, 0.9, 0.6) == "art"  # a glow, a hairline ornament
    # a picture as large as the slide, but running off its foot, is not the slide's background
    assert classify_blob(faint, 1.0, 0.79, covers=False) == "art"
    assert classify_blob(faint, 1.0, 1.0, covers=True) == "ground"


def test_muted_photographs_are_photographs():
    toned = _Stats(coverage=1.0, colours=900, texture=0.35, sharp=0.08, hue_sd=12.0, white=0.0, lum_sd=60.0)
    assert classify_blob(toned, 0.4, 0.4) == "photo"
    portrait = _Stats(coverage=1.0, colours=460, texture=0.1, sharp=0.03, hue_sd=18.0, white=0.0, lum_sd=86.0)
    assert classify_blob(portrait, 0.44, 0.8) == "photo"
    dark_lines = _Stats(coverage=1.0, colours=444, texture=0.13, sharp=0.03, hue_sd=7.0, white=0.0, lum_sd=44.0)
    assert classify_blob(dark_lines, 1.0, 0.5) == "art"


# ---------------------------------------------------------------------------- groups, roles, note fields


def test_the_accent_step_of_a_timeline_joins_its_row():
    shapes = [_text("t", "Нумерация", 36, 0.05, 0.1, 0.9, 0.09)]
    for i, x in enumerate((0.055, 0.288, 0.521, 0.754)):
        fill = "0077FF" if i == 1 else "EBF3F9"
        shapes.append(_shape(f"c{i}", x, 0.46, 0.071, 0.127, fill=fill))
        shapes.append(_text(f"l{i}", "Здесь можно добавить пояснение", 16, x, 0.62, 0.209, 0.087))
    groups = detect_repeat_groups(shapes, W, H)
    assert len(groups) == 1 and len(groups[0].member_shape_ids) == 4
    assert any("c1" in cell for cell in groups[0].member_shape_ids)


def test_the_page_number_under_the_last_card_does_not_cost_the_card_its_column():
    shapes = [_text("t", "Заголовок", 20, 0.04, 0.06, 0.8, 0.05)]
    for i in range(5):
        y = 0.2 + i * 0.147
        shapes.append(_shape(f"f{i}", 0.488, y, 0.484, 0.115, fill="FFFFFF", geometry="roundRect", line="520977"))
        shapes.append(_text(f"b{i}", "Текст", 14, 0.506, y + 0.012, 0.448, 0.088))
    shapes.append(_text("num", "20", 12, 0.954, 0.927, 0.046, 0.053, is_placeholder=True, ph_type="sldNum"))
    groups = detect_repeat_groups(shapes, W, H)
    assert len(groups) == 1 and len(groups[0].member_shape_ids) == 5


def test_an_outline_round_an_icon_is_its_frame_not_a_text_container():
    frame = _shape("fr", 0.766, 0.204, 0.123, 0.204, fill=None, geometry="roundRect", line="FD0C50", z=9)
    icon = _shape("ic", 0.80, 0.255, 0.061, 0.101, fill="FD0C50", geometry="custom", z=3)
    assert not is_empty_frame(frame, [frame, icon], W, H)
    card = _shape("card", 0.05, 0.45, 0.27, 0.35, fill=None, geometry="roundRect", line="FD0C50", z=9)
    corner = _shape("dot", 0.06, 0.46, 0.02, 0.03, fill="FD0C50", z=3)  # a mark in the card's corner
    assert is_empty_frame(card, [card, corner], W, H)


def _slot(sid, role, x, y, w, h, text="", size=14.0, group=None) -> Slot:
    return Slot(id=sid, role=role, shape_id=sid, bbox=BboxFrac(x=x, y=y, w=w, h=h), capacity=Capacity(max_chars=100, max_lines=3), sample_text=text, style=SlotStyle(size_pt=size), group_id=group)


def test_the_note_band_at_the_foot_is_found_by_its_label():
    slots = [
        _slot("t", SlotRole.title, 0.035, 0.06, 0.64, 0.17, "Заголовок", 36),
        _slot("c1", SlotRole.card_title, 0.035, 0.27, 0.3, 0.3, "Заголовок", 18, "g1"),
        _slot("c2", SlotRole.card_title, 0.346, 0.27, 0.3, 0.3, "Заголовок", 18, "g1"),
        _slot("lab", SlotRole.body, 0.035, 0.719, 0.146, 0.058, "Примечание", 14),
        _slot("txt", SlotRole.body, 0.035, 0.807, 0.539, 0.058, "Текст", 14),
    ]
    p = Pattern(id="p15", source_slide=15, kind=PatternKind.cards, family=Family.dark, slots=slots)
    assert note_field(p, 14.0) == ("txt", "lab")
    # a speaker's name and position low on a slide is no note field
    slots2 = slots[:3] + [_slot("n", SlotRole.body, 0.15, 0.75, 0.47, 0.05, "Имя Фамилия", 16), _slot("pos", SlotRole.body, 0.15, 0.81, 0.47, 0.05, "Должность", 14)]
    assert note_field(Pattern(id="p10", source_slide=10, kind=PatternKind.cards, family=Family.dark, slots=slots2), 14.0) == (None, None)


# ---------------------------------------------------------------------------- rendering helpers


def test_a_zigzag_of_cards_is_read_along_its_line():
    top = [_box(x, 0.21, 0.207, 0.215) for x in (0.028, 0.396, 0.765)]
    bottom = [_box(x, 0.70, 0.207, 0.215) for x in (0.213, 0.581)]
    assert staggered_axis(top + bottom) == "row"
    grid = [_box(x, y, 0.4, 0.3) for y in (0.2, 0.55) for x in (0.05, 0.5)]
    assert staggered_axis(grid) is None
    assert staggered_axis([_box(0.05 + 0.3 * i, 0.3, 0.25, 0.3) for i in range(3)]) is None  # a plain row


def test_the_largest_hole_is_measured_in_the_content_area():
    grid = np.zeros((_GRID_R, _GRID_C), dtype=bool)
    area = (2, 5, _GRID_C - 2, _GRID_R - 2)
    full = largest_hole(grid, area)
    assert abs(full - (_GRID_C - 4) * (_GRID_R - 7) / (_GRID_C * _GRID_R)) < 1e-9
    grid[5:14, :] = True  # the upper part filled: the lower band stays free
    lower = largest_hole(grid, area)
    assert abs(lower - (_GRID_C - 4) * (_GRID_R - 2 - 14) / (_GRID_C * _GRID_R)) < 1e-9


def test_headings_keep_brands_compounds_and_their_last_two_words_together():
    assert f"VK{NBSP}Tech" in typeset("Пилот на шаблонах VK Tech")
    assert "Qwen3.8-⁠27B" in typeset("модель Qwen3.8-27B")
    t = no_orphan("Десятки презентаций в неделю требуют ручной вёрстки")
    assert t.endswith(f"ручной{NBSP}вёрстки")
    assert no_orphan("Два слова") == "Два слова"


# ---------------------------------------------------------------------------- the matcher keeps a sample's story off new content


def _cards_pattern(pid="pc", kind=PatternKind.bullets, **kw) -> Pattern:
    slots = [Slot(id="t", role=SlotRole.title, shape_id="t", bbox=BboxFrac(x=0.05, y=0.06, w=0.9, h=0.1), capacity=Capacity(max_chars=60, max_lines=2))]
    slots += [Slot(id=f"b{i}", role=SlotRole.body, shape_id=f"b{i}", bbox=BboxFrac(x=0.05, y=0.2, w=0.4, h=0.6), capacity=Capacity(max_chars=400, max_lines=10)) for i in range(1)]
    return Pattern(id=pid, source_slide=3, kind=kind, family=Family.light, slots=slots, **kw)


def test_a_samples_own_chart_keeps_it_off_text_but_not_off_a_chart(simple_deck, tmp_path):
    from verstka.analysis.manifest import analyze_template
    from verstka.analysis.pictures import SamplePicture
    from verstka.matching.scorer import score_pattern
    from verstka.planning.strategies import get_strategy
    from verstka.schemas.outline import ChartSpec, OutlineSlide, SlideContent

    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    p = _cards_pattern(kind=PatternKind.image_text)
    pics = [SamplePicture(shape_id="g", x=0.5, y=0.2, w=0.45, h=0.6, kind="content")]
    text = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Три шага после финала", content=SlideContent(bullets=["Инференс VK", "Картинки внутри слайдов", "Пилот на шаблонах"]))
    res = score_pattern(text, p, manifest, get_strategy("structured"), pictures=pics)
    assert res.score == 0.0 and any("не относятся к содержанию" in r for r in res.reasons)
    chart = OutlineSlide(id="c", kind=PatternKind.chart, headline="Рост", content=SlideContent(chart=ChartSpec(series_ids=["s1"])))
    assert score_pattern(chart, p, manifest, get_strategy("structured"), pictures=pics).score > 0


def test_an_empty_device_mockup_and_an_unused_note_band_weigh_on_a_content_sample(simple_deck, tmp_path):
    from verstka.analysis.manifest import analyze_template
    from verstka.matching.scorer import score_pattern
    from verstka.planning.strategies import get_strategy
    from verstka.schemas.outline import OutlineSlide, SlideContent

    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    s = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Три шага после финала", content=SlideContent(bullets=["Инференс VK", "Картинки внутри слайдов"]))
    plain = score_pattern(s, _cards_pattern(), manifest, get_strategy("structured")).score
    mock = score_pattern(s, _cards_pattern(mockup_boxes=[BboxFrac(x=0.55, y=0.2, w=0.3, h=0.6)]), manifest, get_strategy("structured"))
    assert mock.score == 0.0 and any("пустым экраном" in r for r in mock.reasons)
    noted = score_pattern(s, _cards_pattern(note_slot="b0"), manifest, get_strategy("structured"))
    assert noted.score < plain and any("поле примечания" in r for r in noted.reasons)
    s2 = s.model_copy(update={"takeaway": "Шаблон задаёт правила, Verstka их соблюдает"})
    with_take = score_pattern(s2, _cards_pattern(note_slot="b0"), manifest, get_strategy("structured"))
    assert any("встанет в поле примечания" in r for r in with_take.reasons)


# ---------------------------------------------------------------------------- reading templates: more of what samples say


def test_numbered_cards_are_not_a_drawn_table_and_keep_their_text():
    from verstka.analysis.groups import detect_drawn_table

    shapes = [_text("t", "Заголовок", 36, 0.035, 0.06, 0.64, 0.17)]
    for r, y in enumerate((0.232, 0.537)):
        for c, x in enumerate((0.036, 0.358, 0.681)):
            k = r * 3 + c
            shapes.append(_shape(f"n{k}", x, y, 0.055, 0.099, fill="212121", geometry="roundRect", z=1))
            shapes[-1].text = TextInfo(paragraphs=[ParagraphInfo(text=f"0{k + 1}", runs=[RunInfo(text=f"0{k + 1}", font="Play", size_pt=16, bold=False, color_hex="FFFFFF")])])
            shapes.append(_text(f"h{k}", "Заголовок", 20, x, y + 0.13, 0.282, 0.074))
            shapes.append(_text(f"b{k}", "Основной текст", 16, x, y + 0.189, 0.282, 0.117))  # overlaps the heading's foot
    assert detect_drawn_table(shapes, W, H) is None
    groups = detect_repeat_groups(shapes, W, H)
    assert len(groups) == 1 and len(groups[0].member_shape_ids) == 6
    assert all(len(cell) == 3 for cell in groups[0].member_shape_ids)


def test_the_accent_card_of_a_row_without_its_header_plate_stays_in_the_row():
    shapes = [_text("t", "Заголовок", 36, 0.035, 0.06, 0.64, 0.17)]
    for i, x in enumerate((0.035, 0.339)):
        shapes.append(_shape(f"c{i}", x, 0.287, 0.282, 0.502, fill="000000", geometry="roundRect", line="70C3FF", z=1))
        shapes.append(_shape(f"p{i}", x, 0.287, 0.282, 0.107, fill="00AEE8", geometry="round2SameRect", z=2))
        shapes.append(_text(f"h{i}", "Заголовок", 20, x + 0.013, 0.318, 0.257, 0.045))
        shapes.append(_text(f"b{i}", "Текст", 12, x + 0.013, 0.414, 0.257, 0.126))
    shapes.append(_shape("c2", 0.642, 0.287, 0.295, 0.502, fill="00AEE8", geometry="roundRect", line="70C3FF", z=1))
    shapes.append(_text("h2", "Заголовок", 20, 0.655, 0.32, 0.282, 0.045))
    shapes.append(_text("b2", "Текст", 12, 0.655, 0.414, 0.257, 0.126))
    groups = detect_repeat_groups(shapes, W, H)
    assert len(groups) == 1 and len(groups[0].member_shape_ids) == 3


def test_kinds_of_a_quote_an_agenda_and_a_picture_page():
    from verstka.analysis.kinds import heuristic_kind

    quote_mark = _text("q", "“", 60, 0.1, 0.2, 0.1, 0.1)
    statement = _text("s", "Write an original statement or inspiring quote.", 28, 0.1, 0.35, 0.8, 0.2)
    credit = _text("c", "— Include a credit, citation, or supporting message", 14, 0.1, 0.6, 0.8, 0.05)
    roles = {"q": SlotRole.title, "s": SlotRole.body, "c": SlotRole.body}
    assert heuristic_kind([quote_mark, statement, credit], roles, [], 18, 21, W, H)[0] == PatternKind.quote
    title = ShapeInfo(id="a", name="a", kind=ShapeKind.sp, bbox=_box(0.1, 0.2, 0.5, 0.5), z=1, text=TextInfo(paragraphs=[ParagraphInfo(text=t, has_bullet=True, runs=[RunInfo(text=t, font="Play", size_pt=32, bold=False, color_hex="FFFFFF")]) for t in ("Introduction", "About Us", "Our Projects")]))
    assert heuristic_kind([title], {"a": SlotRole.title}, [], 3, 21, W, H)[0] == PatternKind.agenda
    head = _text("h", "Заголовок в две строчки", 32, 0.02, 0.1, 0.5, 0.1)
    sub = _text("d", "Описание", 16, 0.02, 0.22, 0.4, 0.05)
    pic = ShapeInfo(id="p", name="p", kind=ShapeKind.pic, bbox=_box(0.2, 0.3, 0.6, 0.6), z=1, image_part="/ppt/media/image1.png")
    kind = heuristic_kind([head, sub, pic], {"h": SlotRole.title, "d": SlotRole.subtitle, "p": SlotRole.image}, [], 43, 54, W, H)[0]
    assert kind != PatternKind.section


def test_a_header_whose_boxes_vary_in_width_is_the_decks_furniture():
    from verstka.analysis.chrome import chrome_ids, detect_chrome

    per_slide = {}
    for i in range(1, 9):
        w = 0.08 + 0.01 * (i % 4)
        per_slide[i] = [_text(f"h{i}", "Dark Minimalist", 12, 0.5 - w / 2, 0.041, w, 0.018), _text(f"t{i}", f"Заголовок {i}", 36, 0.05, 0.2, 0.6, 0.1)]
    chrome = detect_chrome(per_slide, W, H)
    assert all(f"h{i}" in chrome_ids(per_slide[i], chrome, W, H) for i in per_slide)


# ---------------------------------------------------------------------------- charts drawn with shapes


def _bar(sid, x, y, w, h, fill="0077FF") -> ShapeInfo:
    return _shape(sid, x, y, w, h, fill=fill, geometry="rect")


def test_columns_on_a_baseline_with_their_figures_are_one_chart_area():
    """VK Tech p44: columns of different heights on a common baseline, a figure over each and a name under it — one
    data area of the sample (its native chart replaces them), the figure over the first column is no subtitle."""
    from verstka.analysis.drawn_chart import detect_drawn_chart

    shapes = [_text("t", "Заголовок в две или в одну строчку", 24, 0.02, 0.05, 0.74, 0.14)]
    for i, h in enumerate([0.58, 0.52, 0.49, 0.46, 0.40, 0.35]):
        x = 0.05 + i * 0.15
        shapes.append(_bar(f"b{i}", x, 0.81 - h, 0.08, h))
        shapes.append(_text(f"v{i}", "10", 16, x - 0.04, 0.81 - h - 0.07, 0.16, 0.06))
        shapes.append(_text(f"c{i}", "Текст", 10.5, x, 0.83, 0.08, 0.04))
    found = detect_drawn_chart(shapes, W, H)
    assert found is not None
    area, ids, _, kind, hint = found
    assert kind == "chart" and hint["direction"] == "col" and hint["label_pt"] == 10.5 and hint["value_pt"] == 16
    assert {"b0", "v0", "c5"} <= ids and "t" not in ids
    assert area.x >= 0 and area.y2 <= H  # label boxes hanging over the slide's edge stay out of it


def test_bars_back_to_back_and_a_gantt_are_charts_cards_and_stripes_are_not():
    from verstka.analysis.drawn_chart import detect_drawn_chart

    # bars from a common right edge and from a common left edge, the category names between them, a pill on top
    shapes = [_bar("hl", 0.03, 0.26, 0.33, 0.07, fill="E0F0FF"), _text("hlt", "Описание", 16, 0.12, 0.27, 0.16, 0.06)]
    for i, (wl, wr) in enumerate([(0.32, 0.34), (0.30, 0.26), (0.28, 0.25), (0.27, 0.24), (0.24, 0.28)]):
        y = 0.36 + i * 0.09
        shapes += [_bar(f"l{i}", 0.36 - wl, y, wl, 0.056), _bar(f"r{i}", 0.63, y, wr, 0.056), _text(f"m{i}", "Описание", 16, 0.42, y, 0.16, 0.06)]
    found = detect_drawn_chart(shapes, W, H)
    assert found is not None and found[4]["direction"] == "bar" and found[4]["series"] == 2
    assert {"hl", "hlt", "l0", "r4", "m2"} <= found[1]

    # a Gantt: one bar a row, shifted, over a row of months
    gantt = [_bar(f"g{i}", x, 0.28 + i * 0.11, w, 0.05) for i, (x, w) in enumerate([(0.6, 0.3), (0.5, 0.2), (0.3, 0.12), (0.4, 0.1)])]
    gantt += [_text(f"n{i}", "Текст", 12, 0.03, 0.28 + i * 0.11, 0.17, 0.04) for i in range(4)]
    gantt += [_text(f"m{i}", mon, 12, 0.19 + i * 0.13, 0.75, 0.08, 0.05) for i, mon in enumerate(["Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь"])]
    found = detect_drawn_chart(gantt, W, H)
    assert found is not None and found[4]["direction"] == "gantt"

    # a staircase of cards with headings inside is no chart, nor are equal cards on a baseline, nor bare stripes
    stairs = []
    for i, h in enumerate([0.3, 0.4, 0.5, 0.6]):
        stairs += [_bar(f"s{i}", 0.05 + i * 0.23, 0.9 - h, 0.2, h, fill="EEF4FF"), _text(f"st{i}", "Заголовок", 16, 0.06 + i * 0.23, 0.92 - h, 0.18, 0.05)]
    assert detect_drawn_chart(stairs, W, H) is None
    equal = [_bar(f"e{i}", 0.05 + i * 0.23, 0.4, 0.2, 0.4) for i in range(4)] + [_text(f"et{i}", "12", 16, 0.05 + i * 0.23, 0.3, 0.2, 0.06) for i in range(4)]
    assert detect_drawn_chart(equal, W, H) is None
    stripes = [_bar(f"x{i}", 0.6 + i * 0.05, 0.9 - h, 0.03, h) for i, h in enumerate([0.2, 0.35, 0.5, 0.3])]
    assert detect_drawn_chart(stripes, W, H) is None


# ---------------------------------------------------------------------------- measuring, gating and re-laying clones


def test_an_arrow_play_lacks_is_measured_in_the_face_the_renderer_falls_back_to():
    """Play has no «→»: measured with its «.notdef» box, «116 → 123» at 36 pt «fitted» a 172 pt box and wrapped in the
    render over the label under it."""
    from verstka.rendering.fonts import text_width_pt

    with_arrow = text_width_pt("116 → 123", "Play", 36, True)
    without = text_width_pt("116 123", "Play", 36, True)
    assert with_arrow - without >= 0.6 * 36  # the arrow takes most of an em, not a hairline
    assert with_arrow > 172


def test_live_date_and_page_fields_are_never_sample_values():
    from lxml import etree

    from verstka.analysis.shapes import has_live_field

    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    sp = etree.fromstring(f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}"><p:txBody><a:p><a:fld type="datetime1"><a:t>27.09.26</a:t></a:fld></a:p></p:txBody></p:sp>')
    assert has_live_field(sp)
    typed = etree.fromstring(f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}"><p:txBody><a:p><a:r><a:t>27.09.26</a:t></a:r></a:p></p:txBody></p:sp>')
    assert not has_live_field(typed)


def _deck_with(boxes):
    """A one-slide python-pptx deck with text boxes [(x, y, w, h, text, size)] in fractions of a 16:9 slide."""
    from pptx import Presentation
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for x, y, w, h, text, size in boxes:
        tb = slide.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
        tb.text_frame.text = text
        tb.text_frame.word_wrap = True
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
    return slide


def test_the_gate_sees_two_texts_laid_over_each_other():
    """LCT p11 with a third card: a heading that took two lines ran into the text under it."""
    from verstka.analysis.xmlns import q
    from verstka.rendering.clone_check import _texts_over_each_other

    clash = _deck_with([(0.05, 0.2, 0.25, 0.06, "На входе — шаблон и текст презентации", 20), (0.05, 0.24, 0.25, 0.3, "Verstka использует любой шаблон PowerPoint и текст", 14)])
    assert _texts_over_each_other(clash, clash._element.cSld.find(q("p:spTree")), set(), "Play")
    apart = _deck_with([(0.05, 0.2, 0.25, 0.06, "Ручная вёрстка", 20), (0.05, 0.3, 0.25, 0.3, "Подбор макета, подгонка цветов и блоков", 14)])
    assert _texts_over_each_other(apart, apart._element.cSld.find(q("p:spTree")), set(), "Play") is None


def test_a_card_drawn_as_a_picture_stretches_with_its_card_and_is_named_a_plate():
    """VK Tech p26 with three cards instead of four: the white panels drawn as pictures widen with their texts (they
    stood off them), the audit reads them as plates."""
    from pptx import Presentation
    from pptx.util import Emu

    from verstka.analysis.xmlns import q
    from verstka.rendering.deck import element_bbox
    from verstka.rendering.groups import PLATE_MARK, plates_of, transform_element

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    Image.new("RGB", (40, 100), "white").save(buf, format="PNG")
    buf.seek(0)
    pic = slide.shapes.add_picture(buf, Emu(100000), Emu(100000), Emu(2000000), Emu(5000000))
    el = pic._element
    sid = el.find(".//" + q("p:cNvPr")).get("id")
    old, new = Bbox(x=100000, y=100000, w=2000000, h=5000000), Bbox(x=100000, y=100000, w=2800000, h=5000000)
    transform_element(el, old, new, "row")
    assert element_bbox(el)[2] == 2000000  # a picture keeps its proportions
    with plates_of({sid}):
        transform_element(el, old, new, "row")
    assert element_bbox(el)[2] == 2800000 and el.find(".//" + q("p:cNvPr")).get("name").endswith(PLATE_MARK)


def test_a_title_box_over_its_text_box_is_a_stack():
    from verstka.rendering.clone import _stands_over

    s = _deck_with([(0.05, 0.3, 0.2, 0.04, "9", 12), (0.05, 0.36, 0.2, 0.4, "Демо-презентаций", 9)])
    t, b = [sh._element for sh in s.shapes]
    assert _stands_over(t, b)
    beside = _deck_with([(0.05, 0.3, 0.2, 0.04, "9", 12), (0.3, 0.3, 0.2, 0.4, "Демо-презентаций", 9)])
    t, b = [sh._element for sh in beside.shapes]
    assert not _stands_over(t, b)


# ---------------------------------------------------------------------------- figures over their labels, picture places


def _figure_box(sid="710"):
    """VK WorkSpace p17: «ххх%» at 80 pt, a line break, «данные показателя» at 16 pt in one text box."""
    runs = [RunInfo(text="ххх%", font="Inter", size_pt=80.0, bold=True, color_hex="70C3FF"), RunInfo(text="данные показателя", font="Inter", size_pt=16.0, bold=False, color_hex="70C3FF")]
    ti = TextInfo(paragraphs=[ParagraphInfo(text="ххх%\nданные показателя", runs=runs)])
    return ShapeInfo(id=sid, name=sid, kind=ShapeKind.sp, bbox=_box(0.52, 0.18, 0.39, 0.28), z=5, text=ti)


def test_a_figure_over_its_label_in_one_box_is_a_number_slot():
    from verstka.analysis.roles import figure_head

    assert figure_head(_figure_box(), 16.0) == 80.0
    plain = _text("t", "Текст описания", 16, 0.04, 0.18, 0.3, 0.1)
    assert figure_head(plain, 16.0) is None
    # a figure set at the label's size is no figure over its label
    runs = [RunInfo(text="12", font="Inter", size_pt=18.0), RunInfo(text="месяцев", font="Inter", size_pt=16.0)]
    small = ShapeInfo(id="s", name="s", kind=ShapeKind.sp, bbox=_box(0.1, 0.1, 0.2, 0.1), z=5, text=TextInfo(paragraphs=[ParagraphInfo(text="12\nмесяцев", runs=runs)]))
    assert figure_head(small, 16.0) is None


def test_a_figure_and_its_label_are_written_in_the_samples_two_styles():
    from lxml import etree

    from verstka.analysis.xmlns import q
    from verstka.rendering.clone import _compound_sizes
    from verstka.rendering.textfill import ParagraphSpec, fill_text

    P, A = "http://schemas.openxmlformats.org/presentationml/2006/main", "http://schemas.openxmlformats.org/drawingml/2006/main"
    sp = etree.fromstring(
        f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:txBody><a:bodyPr/><a:p>'
        '<a:r><a:rPr sz="8000" b="1"><a:solidFill><a:srgbClr val="70C3FF"/></a:solidFill></a:rPr><a:t>ххх%</a:t></a:r>'
        '<a:br><a:rPr sz="6600" b="1"/></a:br>'
        '<a:r><a:rPr sz="1600"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>данные показателя</a:t></a:r>'
        "</a:p></p:txBody></p:sp>"
    )
    assert _compound_sizes(sp) == (80.0, 16.0)
    fill_text(sp, [ParagraphSpec("87", bullet=False, size_pt=72.0), ParagraphSpec("Чужих шаблонов")], size_pt=16.0)
    paras = sp.findall(".//" + q("a:p"))
    fig, lab = (p.find(q("a:r") + "/" + q("a:rPr")) for p in paras)
    assert fig.get("sz") == "7200" and fig.get("b") == "1"
    assert lab.get("sz") == "1600" and lab.get("b") != "1" and lab.find(".//" + q("a:srgbClr")).get("val") == "FFFFFF"


def test_a_picture_place_drawn_as_a_shape_is_the_samples_picture_slot():
    from verstka.analysis.roles import heuristic_roles
    from verstka.schemas.template import Typography

    title = _text("t", "Заголовок", 36, 0.035, 0.06, 0.64, 0.17, is_placeholder=True, ph_type="title")
    place = _text("p", "Иллюстрация", 18, 0.39, 0.26, 0.49, 0.55, geometry="roundRect", fill_hex="0077FF")
    roles = heuristic_roles([title, place], [], set(), Typography(), W, H)
    assert roles["p"] == SlotRole.image


def test_samples_with_other_pictures_are_not_duplicates():
    """VK WorkSpace p20 (two doughnuts under «Графики») and p21 (a wide column chart under the same heading)."""
    from verstka.analysis.patterns import dedupe_patterns

    def pat(pid, n, boxes, assets):
        t = Slot(id="title_1", role=SlotRole.title, shape_id="1", bbox=BboxFrac(x=0.035, y=0.06, w=0.93, h=0.11), style=SlotStyle(size_pt=36.0), capacity=Capacity(max_chars=40, max_lines=2))
        return Pattern(id=pid, source_slide=n, kind=PatternKind.chart, family=Family.dark, slots=[t], decor_boxes=boxes, decor_assets=assets)

    doughnuts = pat("p20", 20, [BboxFrac(x=0.1, y=0.25, w=0.35, h=0.6), BboxFrac(x=0.55, y=0.25, w=0.35, h=0.6)], ["a1", "a2"])
    columns = pat("p21", 21, [BboxFrac(x=0.0, y=0.21, w=1.0, h=0.67)], ["a3"])
    twin = pat("p22", 22, [BboxFrac(x=0.1, y=0.25, w=0.35, h=0.6), BboxFrac(x=0.55, y=0.25, w=0.35, h=0.6)], ["a4", "a5"])
    kept = [p.id for p in dedupe_patterns([doughnuts, columns, twin])]
    assert kept == ["p20", "p21"]


def _bare_manifest(w: int = W, h: int = H):
    from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

    return TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=w, h=h), tokens=Tokens(), patterns=[], n_slides=1)


def test_a_footnote_keeps_its_small_print_but_running_text_does_not():
    """The slide's footnote is set at the composer's footnote size (about 2 % of the slide height — 8.12 pt on VK
    Tech's 5.6-inch slide): the gate lets it stand; a running text at that size is turned down."""
    from verstka.rendering.clone_check import content_trouble
    from verstka.schemas.outline import OutlineSlide, SlideContent

    m = _bare_manifest(9144000, 5143500)
    pat = Pattern(id="p1", source_slide=1, kind=PatternKind.bullets, family=Family.light, slots=[])
    o = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Заголовок", content=SlideContent(bullets=["Пункт"]))
    text = "Все цифры условные и используются для учебного бизнес-плана"
    small = _deck_with([(0.05, 0.9, 0.8, 0.04, text, 8.12)])
    assert "below" in (content_trouble(small, o, pat, m, []) or "")
    small.shapes[0]._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("name", "Footnote 12")
    assert content_trouble(small, o, pat, m, []) is None


def test_a_stat_slide_with_its_figures_at_a_captions_size_is_turned_down():
    from verstka.rendering.clone_check import content_trouble
    from verstka.schemas.outline import NumberCallout, OutlineSlide, SlideContent

    m = _bare_manifest()
    pat = Pattern(id="p1", source_slide=1, kind=PatternKind.stat_row, family=Family.light, slots=[])
    o = OutlineSlide(id="s", kind=PatternKind.stat_row, headline="Цифры", content=SlideContent(numbers=[NumberCallout(value="87", label="шаблонов"), NumberCallout(value="12", label="форматов")]))
    tiny = _deck_with([(0.05, 0.3, 0.2, 0.06, "87", 14), (0.3, 0.3, 0.2, 0.06, "12", 14)])
    assert "figures set at" in (content_trouble(tiny, o, pat, m, []) or "")
    big = _deck_with([(0.05, 0.3, 0.2, 0.12, "87", 48), (0.3, 0.3, 0.2, 0.12, "12", 48)])
    assert "figures set at" not in (content_trouble(big, o, pat, m, []) or "")


def test_text_hanging_under_its_own_plate_is_turned_down():
    from pptx.dml.color import RGBColor

    from verstka.rendering.clone_check import content_trouble
    from verstka.schemas.outline import OutlineSlide, SlideContent

    m = _bare_manifest()
    pat = Pattern(id="p1", source_slide=1, kind=PatternKind.cards, family=Family.light, slots=[])
    o = OutlineSlide(id="s", kind=PatternKind.cards, headline="Карточки", content=SlideContent(bullets=["Пункт"]))
    from pptx.enum.text import MSO_AUTO_SIZE

    s = _deck_with([(0.05, 0.3, 0.25, 0.05, "Основные направления работы студии", 28)])
    s.shapes[0].text_frame.auto_size = MSO_AUTO_SIZE.NONE  # a plate keeps its size: the text hangs under it
    s.shapes[0].fill.solid()
    s.shapes[0].fill.fore_color.rgb = RGBColor(0x21, 0x21, 0x21)
    assert "runs out of its plate" in (content_trouble(s, o, pat, m, []) or "")
