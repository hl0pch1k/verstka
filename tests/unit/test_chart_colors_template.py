"""Цвета графиков — только из шаблона (финальная доводка, C, критика qa5 п. 13): монохромный шаблон без акцента получает
свои цвета (чёрный и серые), а не синий VK по умолчанию; шаблон одного цвета — оттенки этого цвета, а не соседние тона
(красный и зелёный в круговой жёлто-чёрного Yellow Idea)."""

from __future__ import annotations

from types import SimpleNamespace

from pptx import Presentation
from pptx.util import Emu

from verstka.analysis.components import own_colors
from verstka.rendering.charts import _chroma, _hue_gap, add_chart, chart_palette, delta_e
from verstka.schemas.common import Bbox
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries
from verstka.schemas.template import ChartStyleSpec, ColorToken, Typography

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _tokens(colors: list[tuple[str, str, float]]):
    return SimpleNamespace(colors=[ColorToken.model_construct(hex=h, role=r, roles=[r], weight=w) for h, r, w in colors])


def test_a_monochrome_template_colours_its_charts_with_its_own_colours():
    # Piano: white ground, black ink, two greys — no accent
    piano = _tokens([("FFFFFF", "background.light", 228), ("000000", "neutral.1", 51), ("7B7B7B", "neutral.2", 4), ("CACACA", "neutral.3", 1)])
    assert own_colors(piano) == ["000000", "7B7B7B", "CACACA"]
    # Yellow Idea: one accent, text black
    yellow = _tokens([("FFFFFF", "background.light", 813), ("F2C02C", "accent.1", 274), ("000000", "text.primary", 25)])
    assert own_colors(yellow)[0] == "F2C02C"


def test_one_colour_palettes_are_its_shades_never_other_hues():
    for accent in ("F2C02C", "2F8AA5", "A23735"):
        pal = chart_palette([accent], "FFFFFF", 6, extra=[])
        assert len(set(pal)) == 6
        assert all(_hue_gap(c, accent) <= 15 for c in pal if _chroma(c) >= 12), (accent, pal)
        assert min(delta_e(a, b) for i, a in enumerate(pal) for b in pal[i + 1 :]) >= 8, (accent, pal)


def test_a_five_part_pie_of_a_yellow_template_has_no_green_or_red_slice():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(10080625), Emu(5670550)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    spec = ChartSpec(type="pie", categories=["Близость к дому", "Расписание", "Стоимость", "Тренеры", "Атмосфера"], series=[InlineSeries(name="Доля", values=[35, 25, 20, 12.5, 7.5])], unit="%")
    gf = add_chart(s, Bbox(x=500000, y=1300000, w=3000000, h=3000000), spec, DeckOutline(title="t"), ChartStyleSpec(series_colors=["F2C02C"], font_family="Play", font_size_pt=12.0), Typography(primary_family="Play"), text_hex="000000")
    fills = [e.get("val") for e in gf.chart._chartSpace.iter(f"{A}srgbClr") if e.getparent().tag == f"{A}solidFill" and e.getparent().getparent().tag == f"{C}spPr" and e.getparent().getparent().getparent().tag == f"{C}dPt"]
    assert len(fills) >= 5
    assert all(_hue_gap(c, "F2C02C") <= 20 for c in fills if _chroma(c) >= 12), fills
