"""Real mathematical formulas (verstka/rendering/equation.py): a slide's formula — LaTeX or a line of Unicode — is set
as PowerPoint sets «Вставка → Уравнение»: one mc:AlternateContent whose Choice is a native Office Math equation and
whose Fallback is the same box filled with the formula's picture (LibreOffice, Google Slides, Keynote). Each formula
carries its explanation in words; the composer lays one formula large or several with their names and explanations;
the audit, the geometry helpers, the HTML export and «Улучшить готовую» see the formula as the slide's content."""

from __future__ import annotations

import io
import re

import pytest
from lxml import etree
from PIL import Image

from verstka.rendering import equation as E

M = "{%s}" % E.M_NS


def _omml(latex: str) -> etree._Element:
    return E.omml(latex, 24, "1F2A44", "l")


# ------------------------------------------------------------------ reading a formula field


def test_a_field_of_latex_formulas_with_names_and_explanations():
    text = ("Степенная функция: $\\int x^n\\,dx = \\frac{x^{n+1}}{n+1} + C$ — показатель степени растёт на единицу; "
            "Экспонента: $\\int e^x\\,dx = e^x + C$ — экспонента равна своей первообразной; Синус: $\\int \\sin x\\,dx = -\\cos x + C$")
    fs = E.formulas_of(text)
    assert [f.label for f in fs] == ["Степенная функция", "Экспонента", "Синус"]
    assert fs[0].latex.startswith("\\int x^n") and fs[0].note == "Показатель степени растёт на единицу"
    assert fs[1].note == "Экспонента равна своей первообразной" and fs[2].note == ""
    assert fs[0].text == "∫xⁿ dx = xⁿ⁺¹/(n + 1) + C"


def test_a_line_of_unicode_is_read_into_latex():
    assert E.formulas_of("∫[a,b] f(x)dx = F(b) − F(a)")[0].latex == "\\int _{a}^{b} f(x)\\,dx = F(b) - F(a)"
    assert E.to_latex("x^(n+1)/(n+1) + C") == "\\frac{x^{n+1}}{n+1} + C"
    assert E.to_latex("S = πr²") == "S = πr^{2}"
    assert "\\sqrt{b^{2} - 4ac}" in E.to_latex("x = (−b ± √(b² − 4ac))/2a")
    f = E.formulas_of("∫ f(x) dx = F(x) + C — семейство всех первообразных функции")[0]
    assert f.latex == "\\int f(x)\\,dx = F(x) + C" and f.note == "Семейство всех первообразных функции"


def test_a_business_equation_and_broken_latex_are_no_formulas():
    assert E.formulas_of("100 покупок × 300 ₽ × 30 дней = 900 000 ₽") == []
    assert not E.is_math("100 × 300 × 30 = 900 000 рублей")
    for bad in ("\\badcmd{x}", "\\frac{a}{b", "\\frac}{"):
        assert not E.typesettable(bad), bad
    assert E.formulas_of("Плохая: $\\frac{a}{b$; Хорошая: $E = mc^2$") == [E.Formula("Хорошая", "E = mc^2", "E = mc²")]


def test_tidy_keeps_signs_and_operator_names_apart():
    assert E.tidy("n \\neq -1") == "n \\neq {-}1"
    assert E.tidy("\\dfrac{1}{2} + 2,5") == "\\frac{1}{2} + 2{,}5"
    assert E.tidy("\\lim_{\\max \\Delta x \\to 0}") == "\\lim_{\\max\\, \\Delta x \\to 0}"
    assert E.tidy("\\lim_{n \\to \\infty} x_n") == "\\lim_{n \\to \\infty} x_n"  # never «\\inf\\, ty» (live 2026-10-08)
    assert E.is_sentence("Найдём площадь под параболой y = x² на отрезке [0; 2].")
    assert not E.is_sentence("S = πr², где r — радиус")


def test_latex_reads_as_a_line_of_unicode():
    assert E.latex_text("\\int_0^2 x^2\\,dx = \\frac{8}{3}") == "∫₀²x² dx = 8/3"
    assert E.plain_math("Если $F'(x) = f(x)$, то F — первообразная") == "Если F′(x) = f(x), то F — первообразная"
    assert E.plain_math("Выручка $5 млн и $3 млн за год") == "Выручка $5 млн и $3 млн за год"  # sums of dollars stay


# ------------------------------------------------------------------ Office Math


def test_an_integral_is_an_nary_with_its_limits_and_its_operand():
    om = _omml("\\int_a^b f(x)\\,dx = F(b) - F(a)")
    nary = om.find(f".//{M}nary")
    assert nary is not None
    assert nary.find(f"{M}naryPr/{M}limLoc").get(f"{M}val") == "subSup"
    assert "".join(nary.find(f"{M}sub").itertext()) == "a" and "".join(nary.find(f"{M}sup").itertext()) == "b"
    operand = "".join(t.text for t in nary.find(f"{M}e").iter(f"{M}t"))
    assert operand.replace(" ", "") == "f(x)dx"  # the integrand and its differential, never «= F(b) …»


def test_fractions_limits_functions_fences_and_decimal_commas():
    om = _omml("\\lim_{x \\to 0} \\frac{\\sin x}{x} = 1")
    assert om.find(f".//{M}limLow") is not None and om.find(f".//{M}f/{M}num") is not None
    assert om.find(f".//{M}func/{M}fName") is not None  # «sin x» with its space, upright
    om = _omml("\\left(\\frac{a}{b}\\right)^2 + \\sqrt[3]{x} \\approx 2{,}67")
    assert om.find(f".//{M}d/{M}dPr/{M}begChr").get(f"{M}val") == "(" and om.find(f".//{M}rad/{M}deg") is not None
    assert "2,67" in [t.text for t in om.iter(f"{M}t")]
    # every run carries the slide's size and colour (DrawingML run properties, as PowerPoint writes them)
    rprs = list(om.iter("{%s}rPr" % E.A_NS))
    assert rprs and all(r.get("sz") == "2400" for r in rprs) and all(r.find(".//{%s}srgbClr" % E.A_NS).get("val") == "1F2A44" for r in rprs)


def test_the_picture_is_the_formula_at_its_size_in_its_colour():
    png, w, h = E.picture("\\int_0^2 x^2\\,dx = \\frac{8}{3}", 24, "C0392B")
    png2, w2, h2 = E.picture("\\int_0^2 x^2\\,dx = \\frac{8}{3}", 48, "C0392B")
    assert w2 == pytest.approx(2 * w, rel=0.03) and h2 == pytest.approx(2 * h, rel=0.03)
    assert (w, h) == pytest.approx(E.measure("\\int_0^2 x^2\\,dx = \\frac{8}{3}", 24))
    im = Image.open(io.BytesIO(png)).convert("RGBA")
    assert im.getpixel((0, 0))[3] == 0  # a transparent ground
    ink = [p for p in im.getdata() if p[3] > 200]
    assert ink and all(abs(p[0] - 0xC0) < 8 and abs(p[2] - 0x2B) < 8 for p in ink[:200])


def test_the_element_is_powerpoints_own_form_of_an_equation():
    f = E.formulas_of("Ньютон — Лейбниц: $\\int_a^b f(x)\\,dx = F(b) - F(a)$ — интеграл равен приращению первообразной")[0]
    el = E.equation_element(7, "Formula 7", (100, 200, 3000, 900), f, 24, "1F2A44", "rId9")
    assert etree.QName(el).localname == "AlternateContent" and E.is_equation(el)
    choice, fallback = el
    assert choice.get("Requires") == "a14" and choice.find(".//{%s}m" % E.A14_NS) is not None
    blip = fallback.find(".//{%s}blip" % E.A_NS)
    assert blip.get("{%s}embed" % E.R_NS) == "rId9"
    ids = {c.get("id") for c in el.iter("{%s}cNvPr" % E.P_NS)}
    assert ids == {"7"}
    assert E.equation_text(el) == "Ньютон — Лейбниц: ∫ₐᵇf(x) dx = F(b) − F(a)"
    assert etree.QName(E.shape_of(el)).localname == "sp" and E.shape_of(el).getparent() is fallback


# ------------------------------------------------------------------ the slide's geometry sees one shape


def test_geometry_helpers_move_both_branches(simple_deck):
    from pptx import Presentation
    from pptx.oxml import parse_xml

    from verstka.rendering.deck import element_bbox, set_element_pos, slide_shape_elements

    prs = Presentation(simple_deck)
    slide = prs.slides[0]
    f = E.formulas_of("$E = mc^2$")[0]
    el = parse_xml(etree.tostring(E.equation_element(900, "Formula 900", (100, 200, 3000, 900), f, 24, "000000", "rId1")))
    slide.shapes._spTree.append(el)
    assert element_bbox(el) == (100, 200, 3000, 900)
    set_element_pos(el, x=500, y=600)
    offs = {(o.get("x"), o.get("y")) for o in el.iter("{%s}off" % E.A_NS)}
    assert offs == {("500", "600")}
    assert slide_shape_elements(slide)["900"] is el


# ------------------------------------------------------------------ the composer


@pytest.fixture
def env(simple_deck, tmp_path):
    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace

    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _render(env, content, strategy: str = "structured"):
    from verstka.matching.compat import composition_for
    from verstka.rendering.deck import DeckBuilder
    from verstka.rendering.synth import render_synth
    from verstka.schemas.common import PatternKind
    from verstka.schemas.layout import LayoutSlide
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    deck, manifest, ws = env
    o = OutlineSlide(id="f", kind=PatternKind.bullets, headline="Таблица основных интегралов", content=content)
    comp = composition_for(o)
    slide, warnings = render_synth(DeckBuilder(deck), LayoutSlide(outline_id="f", mode="synth", composition=comp), o, manifest, ws, DeckOutline(title="T", strategy=strategy, slides=[o]))
    return slide, warnings, comp


def _equations(slide):
    tree = slide._element.cSld.find("{%s}spTree" % E.P_NS)
    return [el for el in tree if E.is_equation(el)]


FORMULAS = ("Степенная функция: $\\int x^n\\,dx = \\frac{x^{n+1}}{n+1} + C$ — показатель степени растёт на единицу; "
            "Экспонента: $\\int e^x\\,dx = e^x + C$ — экспонента равна своей первообразной; "
            "Синус: $\\int \\sin x\\,dx = -\\cos x + C$ — производная минус косинуса равна синусу")


@pytest.mark.parametrize("strategy", ["structured", "visual", "compact"])
def test_several_formulas_each_with_its_name_and_explanation(env, strategy):
    from verstka.schemas.outline import SlideContent, SlideItem

    content = SlideContent(formula=FORMULAS, paragraphs=["Основные интегралы получают, читая таблицу производных в обратную сторону."],
                           items=[SlideItem(title="Карточка", text="дизайнера")])  # cards of the designer give way to the formulas
    slide, warnings, comp = _render(env, content, strategy)
    assert comp == "formula"
    eqs = _equations(slide)
    assert len(eqs) == 3
    texts = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame).replace("\u00a0", " ")
    for word in ("Степенная функция", "Показатель степени растёт на единицу", "Экспонента равна своей первообразной", "Основные интегралы получают"):
        assert word in texts, word
    assert "$" not in texts and "\\" not in texts
    # one size for the formulas of a slide, never under the template's body size
    sizes = {r.get("sz") for el in eqs for r in el.iter("{%s}rPr" % E.A_NS) if r.getparent().tag == M + "r"}
    assert len(sizes) == 1
    # every formula's picture is a part of the file, and nothing else is
    rids = {b.get("{%s}embed" % E.R_NS) for el in eqs for b in el.iter("{%s}blip" % E.A_NS)}
    assert rids <= set(slide.part.rels) and all("image" in slide.part.rels[r].reltype for r in rids)


def test_one_formula_large_with_its_explanation_and_its_symbols(env):
    from verstka.rendering.compose import Kit
    from verstka.schemas.outline import SlideContent

    content = SlideContent(formula="$\\int f(x)\\,dx = F(x) + C$ — интегрирование восстанавливает функцию по её производной",
                           paragraphs=["Неопределённый интеграл — множество всех первообразных функции f(x)."],
                           bullets=["F(x) — первообразная функции f(x)", "C — произвольная постоянная"])
    slide, warnings, comp = _render(env, content)
    (eq,) = _equations(slide)
    texts = [sh.text_frame.text.replace("\u00a0", " ") for sh in slide.shapes if sh.has_text_frame]
    assert any("восстанавливает функцию" in t for t in texts) and any("C — произвольная постоянная" in t for t in texts)
    _deck, manifest, _ws = env
    kit = Kit(manifest, 12192000, 6858000, "FFFFFF")
    sz = int(next(r.get("sz") for r in eq.iter("{%s}rPr" % E.A_NS))) / 100
    assert sz >= kit.h2 - 0.05  # the slide's key: never smaller than a subheading


def test_an_undone_pass_takes_its_formula_picture_with_it(env):
    from verstka.rendering.compose import Canvas

    deck, manifest, ws = env
    from pptx import Presentation

    prs = Presentation(deck)
    slide = prs.slides[0]
    cv = Canvas(slide)
    f1, f2 = E.formulas_of("$E = mc^2$")[0], E.formulas_of("$a^2 + b^2 = c^2$")[0]
    from verstka.schemas.common import Bbox

    n0 = len(cv.tree)
    a = cv.equation(Bbox(x=0, y=0, w=4000000, h=800000), f1, 24, "000000")
    b = cv.equation(Bbox(x=0, y=900000, w=4000000, h=800000), f1, 24, "000000")  # the same picture, one part
    c = cv.equation(Bbox(x=0, y=1800000, w=4000000, h=800000), f2, 24, "000000")
    rid_a = next(a.iter("{%s}blip" % E.A_NS)).get("{%s}embed" % E.R_NS)
    rid_c = next(c.iter("{%s}blip" % E.A_NS)).get("{%s}embed" % E.R_NS)
    assert rid_a == next(b.iter("{%s}blip" % E.A_NS)).get("{%s}embed" % E.R_NS)

    class Shim:
        pass

    from verstka.rendering.compose import Composer

    shim = Shim()
    shim.cv, shim.slide = cv, slide
    # undo the last two: the shared picture stays (the first formula still shows it), the other goes
    Composer._undo(shim, n0 + 1)
    assert rid_a in slide.part.rels and rid_c not in slide.part.rels


# ------------------------------------------------------------------ the audit, the HTML and a deck read again


def test_the_audit_reads_a_formula_as_content_and_a_deck_read_again_keeps_it(env, tmp_path):
    from verstka.ingest.deck_text import _members

    deck, manifest, ws = env
    from verstka.schemas.outline import SlideContent

    slide, _w, _c = _render(env, SlideContent(formula="$E = mc^2$ — энергия покоя тела массы m"))
    out = tmp_path / "f.pptx"
    slide.part.package.save(str(out))
    from verstka.audit.ir import build_deck_ir

    ir = build_deck_ir(out)
    s = next(x for x in ir.slides if any(e.formula for e in x.elements))
    eq = next(e for e in s.elements if e.formula)
    assert eq.type == "picture" and eq.image_part and "E = mc²" in eq.formula
    # «Улучшить готовую»: the deck read again says «Формула: …»
    texts = [m.text_frame.text for m in _members(slide.shapes) if getattr(m, "has_text_frame", False)]
    assert any(t.startswith("Формула: ") and "mc²" in t for t in texts)


# ------------------------------------------------------------------ the model's JSON with LaTeX


def test_json_with_single_backslashes_of_latex_is_read():
    from verstka.providers.openai_compat import extract_json

    got = extract_json('{"formula": "\\int_a^b f(x)\\,dx = \\sqrt{2} \\cdot \\pi", "q": "a\\"b"}', {"formula"})
    assert got == {"formula": "\\int_a^b f(x)\\,dx = \\sqrt{2} \\cdot \\pi", "q": 'a"b'}
