"""qa4 audit checks: a score of 100 means what it says — the logo inside the slide, a figure over its label, the type
floor, a legend clear of the axis title, the data deck's summary checked against its data (not the one-line brief)."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import AuditContext
from verstka.ingest.workspace import TemplateWorkspace
from verstka.schemas.common import PatternKind
from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent

W, H = 12192000, 6858000
ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    return manifest, TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")


def _deck(path: Path, build) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for fn in build:
        fn(prs.slides.add_slide(prs.slide_layouts[6]))
    prs.save(str(path))
    return path


def _text(slide, x, y, w, h, text, size):
    tb = slide.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return tb


def test_a_logo_cut_by_the_slide_edge_is_an_error(env, tmp_path):
    from verstka.audit.checks.template import logo_off_slide

    manifest, ws = env
    buf = io.BytesIO()
    Image.new("RGBA", (400, 100), (0, 119, 255, 255)).save(buf, "PNG")

    def logo(x):
        def put(s):
            _text(s, 0.1, 0.4, 0.6, 0.1, "Содержимое слайда", 18)
            pic = s.shapes.add_picture(io.BytesIO(buf.getvalue()), Emu(int(x * W)), Emu(int(0.9 * H)), Emu(int(0.12 * W)), Emu(int(0.04 * H)))
            pic._element.nvPicPr.cNvPr.set("name", "Логотип VK Cloud")
            pic._element.nvPicPr.cNvPr.set("descr", "Логотип VK Cloud")
        return put

    ir = build_deck_ir(_deck(tmp_path / "d.pptx", [logo(0.82), logo(0.92)]))
    got = logo_off_slide(AuditContext(ir=ir, manifest=manifest, ws=ws))
    assert [i.slide for i in got] == [2] and got[0].severity == "error"


def test_figures_set_at_their_labels_size_and_small_text_are_found(env, tmp_path):
    from verstka.audit.checks.template import body_below_template, figure_label_ratio, text_too_small

    manifest, ws = env

    def kpi(s):
        _text(s, 0.05, 0.05, 0.9, 0.1, "Выручка растёт быстрее числа клиентов", 28)
        _text(s, 0.05, 0.3, 0.3, 0.08, "1,24 млрд ₽", 20)  # Marketing's 43 over 30.5 pt, as a ratio
        _text(s, 0.05, 0.4, 0.3, 0.1, "Выручка, +23% к III кв. 2025", 16)
        _text(s, 0.05, 0.6, 0.8, 0.1, "Партнёры получают скидку на обучение сотрудников и доступ к витринам", 9)
        _text(s, 0.05, 0.8, 0.8, 0.05, "Мелкая приписка", 6)

    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Выручка растёт быстрее числа клиентов",
                     content=SlideContent(numbers=[NumberCallout(value="1,24 млрд ₽", label="Выручка, +23% к III кв. 2025")]))
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="t", kind=PatternKind.title, headline="Т"), o, OutlineSlide(id="z", kind=PatternKind.thanks, headline="С")])
    ir = build_deck_ir(_deck(tmp_path / "d.pptx", [lambda s: None, kpi, lambda s: None]))
    for s, osl in zip(ir.slides, outline.slides):
        s.outline_id = osl.id
    ctx = AuditContext(ir=ir, manifest=manifest, outline=outline, ws=ws)
    ratio = figure_label_ratio(ctx)
    assert ratio and ratio[0].slide == 2 and "1,6" in ratio[0].message.replace(".", ",")
    assert body_below_template(ctx) and body_below_template(ctx)[0].details["size"] == 9
    tiny = text_too_small(ctx)
    assert tiny and tiny[0].details["size"] == 6 and tiny[0].severity == "error"


def _scatter_deck(path: Path, names: list[str], width: float) -> Path:
    """A scatter whose X axis title sits right over a bottom legend (the chart writer's manual layout of qa4 c04)."""
    from lxml import etree

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    data = XyChartData()
    for k, nm in enumerate(names):
        ser = data.add_series(nm)
        for i in range(5):
            ser.add_data_point(i + k, 2 * i + k)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Emu(int(0.05 * W)), Emu(int(0.3 * H)), Emu(int(width * W)), Emu(int(0.45 * H)), data)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11.5)
    C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    cs = ch._chartSpace
    ax = cs.findall(f".//{C}valAx")[0]
    ax.find(C + "axPos").set("val", "b")
    title = etree.fromstring(
        '<c:title xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<c:tx><c:rich><a:bodyPr/><a:p><a:r><a:rPr sz="1150"/><a:t>Подготовка, ч/нед.</a:t></a:r></a:p></c:rich></c:tx>'
        '<c:layout><c:manualLayout><c:xMode val="edge"/><c:yMode val="edge"/><c:x val="0.42"/><c:y val="0.84"/></c:manualLayout></c:layout>'
        '<c:overlay val="0"/></c:title>')
    ax.insert(list(ax).index(ax.find(C + "axPos")) + 1, title)
    prs.save(str(path))
    return path


def test_a_bottom_legend_that_wraps_onto_the_axis_title_is_found(env, tmp_path):
    from verstka.audit.checks.layout import legend_over_axis_title

    manifest, ws = env
    names = ["Технический", "Естественнонаучный", "Гуманитарный"]
    narrow = build_deck_ir(_scatter_deck(tmp_path / "n.pptx", names, 0.4))
    got = legend_over_axis_title(AuditContext(ir=narrow, manifest=manifest, ws=ws))
    assert got and "2 строки" in got[0].message
    wide = build_deck_ir(_scatter_deck(tmp_path / "w.pptx", names, 0.9))
    assert legend_over_axis_title(AuditContext(ir=wide, manifest=manifest, ws=ws)) == []


def test_a_data_decks_summary_redone_with_larger_figures_is_checked_against_its_data(simple_deck, tmp_path):
    """qa4 live checks (NEW-E1): «Исправить слайд» with «цифры крупнее» set the sales summary as figures — «20% — больше
    120 мин · 3 — больше 10 часов» among them, the data story's own facts — and `figure_not_in_brief` checked it against
    the one-line brief: 100 → 97. A data deck's summary carries the data's figures: they are checked against them."""
    from verstka.audit.runner import run_audit
    from verstka.cli.main import _read_data
    from verstka.data.story import build_data_story
    from verstka.pipeline.generate import generate_variants
    from verstka.planning.brief import load_brief

    brief_path = tmp_path / "brief.md"
    brief_path.write_text("Итоги продаж VK Cloud за 9 месяцев 2026 года для правления. Покажи рост выручки, сравни регионы, "
                          "что даёт выручку и где узкие места. Данные — в приложенных файлах.", encoding="utf-8")
    story = build_data_story(load_brief(brief_path), _read_data([str(ROOT / "examples" / "data-sales-2026")]), use_llm=False)
    outline = story.outline
    summary = next(s for s in outline.slides if s.data is None and s.content.items and s.kind == PatternKind.cards)
    assert any("120" in (it.text or "") for it in summary.content.items)  # the bottleneck card, as the live run had it
    # the summary as the model redid it (od-report, «qa4 live checks»): the figures lead, the conclusions are their labels
    nums = [NumberCallout(value="+23% против +18%", label="Выручка и клиенты"), NumberCallout(value="41%", label="Москва даёт 41% выручки, Казань растёт быстрее всех"),
            NumberCallout(value="46%", label="Прямые продажи приносят почти половину выручки"), NumberCallout(value="20% — больше 120 мин", label="3 — больше 10 часов")]
    summary.kind = PatternKind.stat_row
    summary.content = SlideContent(numbers=nums)
    res = generate_variants(simple_deck, outline=outline, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws",
                            use_llm=False, use_vlm=False, render_images=False, audit=False, autofix=False, exports=[], pictures=False, eyes=False)
    v = res.variants[0]
    vdir = tmp_path / "out" / "structured"
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    rep = run_audit(vdir / "deck.pptx", manifest, v.outline, ws, render=False, use_vlm=False, use_llm=False, brief_text=brief_path.read_text(encoding="utf-8"))
    flagged = [i.message for i in rep.issues if i.check_id == "figure_not_in_brief"]
    assert not flagged, flagged
