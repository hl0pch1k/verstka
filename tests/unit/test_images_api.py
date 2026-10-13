"""«Заменить изображение» through the API: a build with pictures (a fake text-to-image model, no live call), a new
picture on the person's wish — only that slide changes, the previous version is kept —, «верни как было» brings the
old picture back, «Убрать изображение» leaves the free place; 503 without a model, 422 on a slide without a place,
409 while the variant is changing."""

from __future__ import annotations

import time
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from test_api import _wait, client  # noqa: F401  (shared fixture)
from test_images import GREENHOUSE

from verstka.analysis.xmlns import q
from verstka.images import providers as IP

TPL = Path(__file__).resolve().parents[1] / "fixtures" / "tpl_audit" / "templates" / "lo_focus.pptx"


def _blip_of(deck: Path, n: int) -> bytes:
    prs = Presentation(str(deck))
    slide = prs.slides[n - 1]
    pic = next(el for el in slide._element.cSld.find(q("p:spTree")) if etree.QName(el).localname == "pic")
    rid = pic.find(".//" + q("a:blip")).get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    return slide.part.related_part(rid).blob


@pytest.fixture
def built(client, monkeypatch):  # noqa: F811
    c, mod = client
    mock = IP.MockImages()
    chain = IP.ImageChain([mock])
    monkeypatch.setattr(mod, "image_provider", lambda: chain)
    with open(TPL, "rb") as f:
        r = c.post("/api/templates", files={"file": ("lo_focus.pptx", f, "application/octet-stream")}, data={"use_models": "false"})
    tid = _wait(c, r.json()["job_id"])["result"]["template_id"]
    r = c.post("/api/generations", json={"template_id": tid, "brief": GREENHOUSE, "strategies": ["structured"], "use_models": False, "exports": [], "pictures": True})
    assert r.status_code == 200, r.text
    gid = r.json()["generation_id"]
    job = _wait(c, r.json()["job_id"], timeout=240)
    assert job["status"] == "done", job.get("error")
    g = c.get(f"/api/generations/{gid}").json()
    v = g["variants"][0]
    return c, mod, mock, gid, v


def test_the_variant_payload_lists_its_pictures(built):
    c, mod, mock, gid, v = built
    pics = v["pictures"]
    assert sorted(p["status"] for p in pics) == ["free", "generated"]
    gen = next(p for p in pics if p["status"] == "generated")
    assert gen["model_label"] == "Mock T2I" and "kept" not in gen and gen["hint"].startswith("Экран панели управления")
    assert c.get("/api/health").json()["images"]["configured"] is True


def test_each_place_says_where_it_is_on_the_slide(built):
    """The interface frames a picture on the slide («Изображение на слайде»): every place's box in slide fractions
    (`bbox`), from its EMU box and the deck's slide size; the EMU box itself stays inside."""
    from verstka.images.fill import read_records

    c, mod, mock, gid, v = built
    vdir = mod.store.generation_dir(gid) / "structured"
    prs = Presentation(str(vdir / "deck.pptx"))
    sw, sh = prs.slide_width, prs.slide_height
    boxes = {r["id"]: r["box"] for r in read_records(vdir)}
    assert v["pictures"]
    for p in v["pictures"]:
        x, y, w, h = boxes[p["id"]]
        assert p["bbox"] == pytest.approx({"x": x / sw, "y": y / sh, "w": w / sw, "h": h / sh}, abs=1e-4)
        b = p["bbox"]
        assert 0 <= b["x"] < 1 and 0 <= b["y"] < 1 and 0 < b["w"] <= 1 and 0 < b["h"] <= 1
        assert "box" not in p


def test_a_new_picture_on_a_wish_and_back_again(built):
    c, mod, mock, gid, v = built
    n = next(p for p in v["pictures"] if p["status"] == "generated")["slide"]
    vdir = mod.store.generation_dir(gid) / "structured"
    before = _blip_of(vdir / "deck.pptx", n)
    r = c.post(f"/api/generations/{gid}/structured/slides/{n}/image", json={"wish": "вид сверху, больше зелени"})
    assert r.status_code == 200, r.text
    job = _wait(c, r.json()["job_id"])
    assert job["status"] == "done", job.get("error")
    res = job["result"]
    assert res["applied"] and res["kind"] == "picture" and res["version"] == 1 and "больше зелени" in res["reply"]
    assert "вид сверху, больше зелени" in mock.calls[-1]["prompt"]
    after = _blip_of(vdir / "deck.pptx", n)
    assert after != before
    g = c.get(f"/api/generations/{gid}").json()["variants"][0]
    assert next(p for p in g["pictures"] if p["slide"] == n)["wish"] == "вид сверху, больше зелени"
    assert g["edits"][-1]["kind"] == "picture"
    # «верни как было»: the old picture, the old record
    r = c.post(f"/api/generations/{gid}/structured/edits", json={"message": "верни как было", "slide": n})
    job = _wait(c, r.json()["job_id"])
    assert job["status"] == "done" and job["result"]["changed"], job
    assert _blip_of(vdir / "deck.pptx", n) == before
    g = c.get(f"/api/generations/{gid}").json()["variants"][0]
    assert next(p for p in g["pictures"] if p["slide"] == n)["wish"] is None


def test_the_picture_taken_out_leaves_the_free_place(built):
    c, mod, mock, gid, v = built
    n = next(p for p in v["pictures"] if p["status"] == "generated")["slide"]
    r = c.post(f"/api/generations/{gid}/structured/slides/{n}/image", json={"action": "remove"})
    job = _wait(c, r.json()["job_id"])
    assert job["status"] == "done" and job["result"]["applied"], job
    vdir = mod.store.generation_dir(gid) / "structured"
    prs = Presentation(str(vdir / "deck.pptx"))
    names = [sh.name for sh in prs.slides[n - 1].shapes]
    assert "Место для фото" in names and any(x.startswith("Photo caption") for x in names)
    g = c.get(f"/api/generations/{gid}").json()["variants"][0]
    assert next(p for p in g["pictures"] if p["slide"] == n)["status"] == "removed"


def test_errors_say_why(built, monkeypatch):
    c, mod, mock, gid, v = built
    n = next(p for p in v["pictures"] if p["status"] == "generated")["slide"]
    assert c.post(f"/api/generations/{gid}/structured/slides/1/image", json={}).status_code == 422  # the cover has no place
    assert c.post(f"/api/generations/{gid}/structured/slides/99/image", json={}).status_code == 422
    with mod._fix_lock:
        mod._fix_active.add((gid, "structured"))
    try:
        assert c.post(f"/api/generations/{gid}/structured/slides/{n}/image", json={}).status_code == 409
    finally:
        with mod._fix_lock:
            mod._fix_active.discard((gid, "structured"))
    monkeypatch.setattr(mod, "image_provider", lambda: None)
    r = c.post(f"/api/generations/{gid}/structured/slides/{n}/image", json={})
    # a person reads only that pictures are off; what to set is in the server's log
    assert r.status_code == 503 and r.json()["detail"] == "Модель изображений не подключена"


def test_the_helper_chat_changes_a_picture_on_request(built):
    c, mod, mock, gid, v = built
    n = next(p for p in v["pictures"] if p["status"] == "generated")["slide"]
    r = c.post("/api/chat", json={"session_id": "pic", "message": f"замени картинку на слайде {n} на вид сверху", "generation_id": gid, "strategy": "structured"}).json()
    assert r["intent"] == "picture" and r["actions"] and r["actions"][0]["type"] == "edit", r
    job = _wait(c, r["actions"][0]["job_id"])
    assert job["status"] == "done" and job["result"]["changed"] and job["result"]["slide"] == n, job
    assert "вид сверху" in mock.calls[-1]["prompt"]
    # a slide without a picture or a place: the chat says so instead of starting a job
    r = c.post("/api/chat", json={"session_id": "pic", "message": "нарисуй другое изображение на слайде 1", "generation_id": gid, "strategy": "structured"}).json()
    assert not r["actions"] and "нет изображения" in r["reply"]


@pytest.mark.parametrize(
    "error, states, words",
    [
        ("Z-Image: no credits (402) on the balance (Недостаточно средств)", ["no_credits"], "модель изображений недоступна"),
        ("Z-Image: access refused (401/403): HTTP 401, check the key", [], "модель изображений недоступна"),
        # Polza's wording names both windows; the chain told the daily cap by its Retry-After and holds the model
        ("Z-Image: Достигнут минутный (дневной) лимит запросов к модели", ["quota"], "после полуночи по московскому времени"),
        ("Z-Image: the request was refused by moderation (403)", ["ready"], "отказалась рисовать"),
        ("Z-Image: 502 server error, congested upstream", ["ready"], "чуть позже"),
    ],
)
def test_a_failed_picture_change_says_when_trying_again_helps(error, states, words):
    from verstka.api.app import _picture_failure_reply

    reply = _picture_failure_reply(error, states)
    assert words in reply and reply.startswith("Не получилось") and "слайд остался как был" in reply
