"""The finished file and what the audit reads in it (round of 2026-10-04):

* pictures are stored in the .pptx as they are, never deflated again at every save; the XML parts stay deflated;
* the copy LibreOffice renders drops an alpha channel that is opaque everywhere (the same pixels, a third of the time);
* a picture of the slide's own under a text is that text's ground (code on a black panel is read as on black);
* a text set with a tighter letter-spacing is measured with it (one line, as the renderer set it);
* a cover's subtitle steps down the scale to stand above its small print.
"""

from __future__ import annotations

import io
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches, Pt


def _png(color: tuple, size=(1200, 600), mode="RGBA") -> bytes:
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, format="PNG")
    return buf.getvalue()


def _deck_with_picture(tmp_path, png: bytes, text: str = "", color: str = "F5F5F5", spc: int | None = None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic_path = tmp_path / "ground.png"
    pic_path.write_bytes(png)
    slide.shapes.add_picture(str(pic_path), 0, Inches(2), prs.slide_width, Inches(5.5))
    if text:
        tb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(6), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(28)
        from pptx.dml.color import RGBColor

        run.font.color.rgb = RGBColor.from_string(color)
        if spc is not None:
            run._r.get_or_add_rPr().set("spc", str(spc))
    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    return out


def test_pictures_are_stored_as_they_are_and_xml_stays_deflated(tmp_path):
    import verstka  # noqa: F401 - installs the package writer

    deck = _deck_with_picture(tmp_path, _png((10, 120, 200, 255)), "Текст")
    with zipfile.ZipFile(deck) as z:
        kinds = {i.filename: i.compress_type for i in z.infolist()}
    media = [n for n in kinds if n.startswith("ppt/media/")]
    assert media and all(kinds[n] == zipfile.ZIP_STORED for n in media)
    assert kinds["ppt/slides/slide1.xml"] == zipfile.ZIP_DEFLATED
    from verstka.pkgzip import rewrite

    rewrite(deck, {"ppt/slides/slide1.xml": zipfile.ZipFile(deck).read("ppt/slides/slide1.xml")})
    assert len(Presentation(str(deck)).slides) == 1  # the rewritten package opens


def test_the_render_copy_drops_an_alpha_channel_opaque_everywhere(tmp_path):
    from verstka.ingest.render import _without_opaque_alpha

    opaque = _png((10, 20, 30, 255), size=(1600, 900))
    got = _without_opaque_alpha(opaque)
    assert got is not None and Image.open(io.BytesIO(got)).mode == "RGB"
    assert list(Image.open(io.BytesIO(got)).getdata())[:3] == [(10, 20, 30)] * 3
    # a picture that is see-through somewhere keeps its alpha; a small one is not worth the work
    see_through = _png((10, 20, 30, 128), size=(1600, 900))
    assert _without_opaque_alpha(see_through) is None
    assert _without_opaque_alpha(_png((10, 20, 30, 255), size=(40, 40))) is None


def test_a_picture_of_the_slide_under_a_text_is_its_ground(tmp_path):
    from verstka.audit.checks.common import ground_of, text_elements
    from verstka.audit.ir import build_deck_ir

    deck = _deck_with_picture(tmp_path, _png((0, 0, 0, 255), mode="RGBA"), '"timestamp": "2026-06-15T10:30"', color="F5F5F5")
    ir = build_deck_ir(deck)
    s = ir.slides[0]
    e = next(x for x in text_elements(s) if "timestamp" in x.text)
    color, kind = ground_of(s, e, s.background_hex)
    assert color == "000000" and kind == "picture"  # light code on the black panel: no «contrast» to fix by blackening it


def test_text_is_measured_at_the_letter_spacing_it_is_set_with(tmp_path):
    from verstka.audit.checks.common import text_elements, text_height_needed_pt
    from verstka.audit.ir import build_deck_ir
    from verstka.rendering.fonts import text_width_pt

    text = "Почему обычной теплице нужен постоянный контроль"
    width_in = (text_width_pt(text, None, 28.0, False) - 0.6 * len(text)) / 72 + 0.25  # wide enough only when tightened
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1), Emu(int(width_in * 914400)), Inches(0.6))
    tb.text_frame.word_wrap = True
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(28)
    run._r.get_or_add_rPr().set("spc", "-140")
    deck = tmp_path / "spc.pptx"
    prs.save(str(deck))
    e = next(x for x in text_elements(build_deck_ir(deck).slides[0]) if x.text.startswith("Почему"))
    assert e.spc_pt == -1.4
    assert text_height_needed_pt(e, 1.2)[1] == 1


def test_a_cover_subtitle_steps_down_to_stand_above_its_small_print():
    from verstka.rendering.synth import _sub_height, _sub_size_above

    sub = "Автоматизация полива и контроль микроклимата"
    w = int(Inches(8))
    scale = [12.0, 16.0, 20.0, 24.0, 28.0, 36.0, 54.0]
    y, foot = int(Inches(5.6)), int(Inches(7.0))  # two lines at a display size do not fit above the small print
    size = _sub_size_above(sub, None, 54.0, 12.0, scale, w, y, foot)
    assert size < 54.0 and y + _sub_height(sub, None, size, w) <= foot
    # with room enough it keeps the template's own size; never below half of it (the caller says the note aloud then)
    assert _sub_size_above(sub, None, 54.0, 12.0, scale, w, y, int(Inches(9.5))) == 54.0
    assert _sub_size_above(sub, None, 54.0, 12.0, scale, w, y, y + 1) >= 27.0


def test_the_code_face_is_never_refonted_into_the_templates(tmp_path):
    """A code slide sets its lines in a monospace face whatever the template's: the audit's «font not in the template»
    never takes it for a slip (its fix set the code of LCT in Montserrat — indents no longer lined up)."""
    import json
    from pathlib import Path

    import pytest

    from verstka.audit.checks.template import font_not_in_template
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext
    from verstka.schemas.template import TemplateManifest

    lct = Path(__file__).resolve().parents[2] / "workspace" / "templates" / "1604aa9146b82cac" / "manifest.json"
    if not lct.exists():
        pytest.skip("the LCT template is not in the workspace")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for face, text in (("Consolas", "def irrigation_decision(soil_moisture):"), ("Comic Sans MS", "Чужой шрифт")):
        tb = slide.shapes.add_textbox(Inches(1), Inches(1 + 2 * (face != "Consolas")), Inches(6), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.name = face
    deck = tmp_path / "code.pptx"
    prs.save(str(deck))
    manifest = TemplateManifest.model_validate(json.loads(lct.read_text(encoding="utf-8")))
    issues = font_not_in_template(AuditContext(ir=build_deck_ir(deck), manifest=manifest))
    said = " ".join(i.message for i in issues)
    assert "Comic Sans" in said and "Consolas" not in said


def test_the_finished_deck_has_no_machine_markers_and_the_audit_still_maps_its_slides(tmp_path):
    from verstka.audit.ir import build_deck_ir
    from verstka.rendering.renderer import strip_outline_marks

    prs = Presentation()
    for i in (1, 2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tf = slide.notes_slide.notes_text_frame
        tf.text = ("Заметка автора.\n" if i == 1 else "") + f"[verstka:sl{i}]"
    deck = tmp_path / "marks.pptx"
    prs.save(str(deck))
    assert build_deck_ir(deck).slides[0].outline_id == "sl1"
    assert strip_outline_marks(deck) == 2
    prs2 = Presentation(str(deck))
    notes = [s.notes_slide.notes_text_frame.text for s in prs2.slides]
    assert notes == ["Заметка автора.", ""]
    assert all(s.outline_id is None for s in build_deck_ir(deck).slides)
