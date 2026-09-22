"""Regression tests for the analysis review findings (R1, AUD-2, R3, R4c, F5).

Synthetic decks are built with python-pptx; the VK-template cases are integration tests
that are skipped when the dataset directory is absent.
"""

from __future__ import annotations

import copy
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.chrome import chrome_ids, detect_chrome
from verstka.analysis.groups import detect_repeat_groups
from verstka.analysis.kinds import heuristic_kind
from verstka.analysis.roles import heuristic_roles
from verstka.analysis.shapes import SlideContext, extract_shapes
from verstka.analysis.xmlns import NS, find, q
from verstka.audit.ir import build_deck_ir
from verstka.ingest.package import PptxPackage
from verstka.schemas.common import PatternKind, ShapeKind, SlotRole
from verstka.schemas.template import RepeatGroup, TemplateManifest, TypeStep, Typography

W, H = 12192000, 6858000

TYPO = Typography(scale=[TypeStep(role="h1", size_pt=32), TypeStep(role="body", size_pt=16), TypeStep(role="caption", size_pt=10)])


# ---------------------------------------------------------------------------- helpers


def _blank_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(W)
    prs.slide_height = Emu(H)
    return prs


def _shapes(pptx: Path, n: int = 1):
    pkg = PptxPackage.open(pptx)
    part = pkg.slide_parts[n - 1]
    return pkg, extract_shapes(pkg, part, SlideContext(pkg, part))


def _shared_ids(g: RepeatGroup) -> set[str]:
    """Shape ids that occur in more than one cell of the group."""
    seen: dict[str, int] = {}
    dup: set[str] = set()
    for ci, cell in enumerate(g.member_shape_ids):
        for sid in cell:
            if sid in seen and seen[sid] != ci:
                dup.add(sid)
            seen.setdefault(sid, ci)
    return dup


def _analyze_deck(pptx: Path) -> dict[int, tuple[list, list[RepeatGroup], dict[str, SlotRole]]]:
    pkg = PptxPackage.open(pptx)
    per_slide = {i: extract_shapes(pkg, part, SlideContext(pkg, part)) for i, part in enumerate(pkg.slide_parts, 1)}
    chrome = detect_chrome(per_slide, *pkg.slide_size)
    ids = {i: chrome_ids(per_slide[i], chrome, *pkg.slide_size) for i in per_slide}
    out = {}
    for i, shapes in per_slide.items():
        groups = detect_repeat_groups(shapes, *pkg.slide_size, chrome_ids=ids[i])
        roles = heuristic_roles(shapes, groups, ids[i], TYPO, *pkg.slide_size)
        out[i] = (shapes, groups, roles)
    return out


def _vk(fixtures_dir, needle: str) -> Path:
    if fixtures_dir is None:
        pytest.skip("VK templates not available")
    p = next((p for p in fixtures_dir.glob("*.pptx") if needle.lower() in p.name.lower()), None)
    if p is None:
        pytest.skip(f"template {needle} not found")
    return p


# ---------------------------------------------------------------------------- R1: satellites stolen from a neighbouring card


def _four_cards_deck(out: Path) -> Path:
    """Four identical cards in a row whose gap (0.03·W) is smaller than the satellite distance (0.06·W):
    card N's title/body boxes lie within satellite reach of card N-1."""
    prs = _blank_prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Emu(int(W * 0.05)), Emu(int(H * 0.08)), Emu(int(W * 0.8)), Emu(int(H * 0.12)))
    t.text_frame.text = "Four cards"
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    for i in range(4):
        x, y = int(W * (0.05 + i * 0.23)), int(H * 0.30)
        w, h = int(W * 0.20), int(H * 0.45)
        rect = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(0xED, 0xF3, 0xFC)
        rect.line.fill.background()
        rect.name = f"Card {i + 1}"
        title = s.shapes.add_textbox(Emu(x + int(w * 0.05)), Emu(y + int(h * 0.08)), Emu(int(w * 0.9)), Emu(int(h * 0.18)))
        title.text_frame.text = f"Card {i + 1}"
        r = title.text_frame.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.size = Pt(20)
        body = s.shapes.add_textbox(Emu(x + int(w * 0.05)), Emu(y + int(h * 0.35)), Emu(int(w * 0.9)), Emu(int(h * 0.5)))
        body.text_frame.word_wrap = True
        body.text_frame.text = f"Body of card {i + 1}"
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(out)
    return out


def test_r1_neighbour_card_text_is_not_a_satellite(tmp_path):
    pkg, shapes = _shapes(_four_cards_deck(tmp_path / "cards4.pptx"))
    groups = detect_repeat_groups(shapes, W, H)
    assert len(groups) == 1
    g = groups[0]
    assert g.axis == "row"
    assert len(g.member_shape_ids) == 4, g.member_shape_ids
    assert all(len(cell) == 3 for cell in g.member_shape_ids), g.member_shape_ids
    assert not _shared_ids(g), g.member_shape_ids
    assert g.max_n >= 4
    roles = heuristic_roles(shapes, groups, set(), TYPO, W, H)
    assert list(roles.values()).count(SlotRole.card_title) == 4
    assert list(roles.values()).count(SlotRole.card_body) == 4
    kind, _ = heuristic_kind(shapes, roles, groups, 2, 10, W, H)
    assert kind == PatternKind.cards


def _card(s, x: float, y: float, w: float, h: float, title: str, fill=(0xED, 0xF3, 0xFC)):
    rect = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h)))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(*fill)
    rect.line.fill.background()
    t = s.shapes.add_textbox(Emu(int(W * (x + 0.01))), Emu(int(H * (y + 0.03))), Emu(int(W * (w - 0.02))), Emu(int(H * 0.06)))
    t.text_frame.text = title
    r = t.text_frame.paragraphs[0].runs[0]
    r.font.bold = True
    r.font.size = Pt(20)
    b = s.shapes.add_textbox(Emu(int(W * (x + 0.01))), Emu(int(H * (y + 0.12))), Emu(int(W * (w - 0.02))), Emu(int(H * (h - 0.15))))
    b.text_frame.text = f"Body of {title}"
    b.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    return rect


def test_r1_stray_empty_text_box_does_not_break_the_cell_composition(tmp_path):
    """A page-number box under the last card only (VK Tech slide 16) must not drop that card from the group."""
    prs = _blank_prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(3):
        _card(s, 0.05 + i * 0.31, 0.30, 0.28, 0.45, f"Card {i + 1}")
    s.shapes.add_textbox(Emu(int(W * 0.90)), Emu(int(H * 0.78)), Emu(int(W * 0.05)), Emu(int(H * 0.04)))  # empty, invisible
    prs.save(tmp_path / "stray.pptx")
    pkg, shapes = _shapes(tmp_path / "stray.pptx")
    groups = detect_repeat_groups(shapes, W, H)
    assert len(groups) == 1 and len(groups[0].member_shape_ids) == 3, [g.member_shape_ids for g in groups]
    assert not _shared_ids(groups[0])


def test_r1_text_inside_a_wider_foreign_card_is_not_a_satellite(tmp_path):
    """Four equal cards plus a wide card whose title box lies right under card 2 (VK Tech slide 18)."""
    prs = _blank_prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(3):
        _card(s, 0.03 + i * 0.32, 0.24, 0.30, 0.30, f"Card {i + 1}")
    # second row starts 0.015·H below the first: the title boxes of the wide card (and of card 4) are within
    # the 0.05·H "below" satellite distance of the cards above them
    _card(s, 0.03, 0.555, 0.30, 0.30, "Card 4")
    _card(s, 0.35, 0.555, 0.62, 0.30, "Wide card")  # different width → not in the cluster
    prs.save(tmp_path / "wide.pptx")
    pkg, shapes = _shapes(tmp_path / "wide.pptx")
    groups = detect_repeat_groups(shapes, W, H)
    g = max(groups, key=lambda g: len(g.member_shape_ids))
    assert len(g.member_shape_ids) == 4 and g.axis == "grid", g.member_shape_ids
    assert all(len(cell) == 3 for cell in g.member_shape_ids), g.member_shape_ids
    wide_title = next(sh for sh in shapes if sh.plain_text == "Wide card")
    assert all(wide_title.id not in cell for cell in g.member_shape_ids)


def test_r1_label_grazing_a_decorative_picture_stays_with_its_icon(tmp_path):
    """Icon + label rows where the last label overlaps the bounding box of a big illustration (WorkSpace slide 13)."""
    prs = _blank_prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    img = tmp_path / "illu.png"
    Image.new("RGB", (400, 400), (20, 40, 80)).save(img)
    # the illustration's bounding box covers ~45% of the third label, whose centre stays outside it
    s.shapes.add_picture(str(img), Emu(int(W * 0.62)), Emu(0), height=Emu(int(H * 0.52)))
    for i in range(3):
        x = 0.03 + i * 0.225
        chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(W * x)), Emu(int(H * 0.32)), Emu(int(W * 0.05)), Emu(int(H * 0.09)))
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor(0x00, 0xAE, 0xE8)
        lbl = s.shapes.add_textbox(Emu(int(W * (x + 0.06))), Emu(int(H * 0.32)), Emu(int(W * 0.15)), Emu(int(H * 0.084)))
        lbl.text_frame.text = "Текст"
        lbl.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(tmp_path / "graze.pptx")
    pkg, shapes = _shapes(tmp_path / "graze.pptx")
    groups = detect_repeat_groups(shapes, W, H)
    g = max(groups, key=lambda g: len(g.member_shape_ids))
    assert len(g.member_shape_ids) == 3 and all(len(cell) == 2 for cell in g.member_shape_ids), g.member_shape_ids


@pytest.mark.integration
@pytest.mark.parametrize("needle", ["VK Tech", "WorkSpace", "Education"])
def test_r1_vk_templates_no_shape_in_two_cells(fixtures_dir, needle):
    res = _analyze_deck(_vk(fixtures_dir, needle))
    bad = [(i, g.id, sorted(_shared_ids(g))) for i, (_, groups, _) in res.items() for g in groups if _shared_ids(g)]
    assert bad == [], bad


@pytest.mark.integration
def test_r1_vk_tech_slide_14_and_20(fixtures_dir):
    res = _analyze_deck(_vk(fixtures_dir, "VK Tech"))
    _, groups14, roles14 = res[14]
    g = max(groups14, key=lambda g: len(g.member_shape_ids))
    assert g.axis == "row" and len(g.member_shape_ids) == 4, g.member_shape_ids
    assert g.member_shape_ids[0] == ["665", "666", "667", "668", "669"]
    # the card title is the visible «Безопасность» box; the emptied 14 pt box above it is a leftover, not a slot
    assert roles14["673"] == SlotRole.card_title and roles14["674"] == SlotRole.card_body
    assert roles14["671"] not in (SlotRole.card_title, SlotRole.card_body)
    _, groups20, _ = res[20]
    g = max(groups20, key=lambda g: len(g.member_shape_ids))
    assert g.axis == "column" and len(g.member_shape_ids) == 4, g.member_shape_ids
    # slide 16: eight equal cards plus a stray page-number box; slide 18: four equal cards next to a wide one
    _, groups16, _ = res[16]
    g = max(groups16, key=lambda g: len(g.member_shape_ids))
    assert g.axis == "grid" and len(g.member_shape_ids) == 8, g.member_shape_ids
    _, groups18, _ = res[18]
    g = max(groups18, key=lambda g: len(g.member_shape_ids))
    assert g.axis == "grid" and len(g.member_shape_ids) == 4, g.member_shape_ids


@pytest.mark.integration
def test_r1_workspace_slide_13_icon_row_keeps_three_cells(fixtures_dir):
    res = _analyze_deck(_vk(fixtures_dir, "WorkSpace"))
    _, groups, _ = res[13]
    icon_rows = [g for g in groups if any("649" in cell for cell in g.member_shape_ids)]
    assert icon_rows and len(icon_rows[0].member_shape_ids) == 3, [g.member_shape_ids for g in groups]


# ---------------------------------------------------------------------------- AUD-2: fill alpha


def _alpha_deck(out: Path) -> Path:
    prs = _blank_prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, alpha in enumerate([29804, None, 50000]):
        rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * (0.05 + i * 0.3))), Emu(int(H * 0.3)), Emu(int(W * 0.25)), Emu(int(H * 0.4)))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(0x00, 0x77, 0xFF)
        rect.name = f"rect{i}"
        rect.text_frame.text = f"text {i}"
        if alpha is not None:
            clr = find(rect._element, "p:spPr/a:solidFill/a:srgbClr")
            etree.SubElement(clr, q("a:alpha")).set("val", str(alpha))
    # scheme colour with alpha
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.05)), Emu(int(H * 0.75)), Emu(int(W * 0.25)), Emu(int(H * 0.15)))
    rect.name = "scheme"
    sf = find(rect._element, "p:spPr/a:solidFill")
    if sf is None:
        spPr = find(rect._element, "p:spPr")
        sf = etree.SubElement(spPr, q("a:solidFill"))
        # solidFill must precede a:ln in spPr
        spPr.remove(sf)
        spPr.insert(1, sf)
    for child in list(sf):
        sf.remove(child)
    etree.SubElement(etree.SubElement(sf, q("a:schemeClr"), val="accent1"), q("a:alpha")).set("val", "40000")
    prs.save(out)
    return out


def test_aud2_fill_alpha_extracted_and_passed_to_ir(tmp_path):
    p = _alpha_deck(tmp_path / "alpha.pptx")
    pkg, shapes = _shapes(p)
    by_name = {s.name: s for s in shapes}
    assert by_name["rect0"].fill_hex == "0077FF"
    assert by_name["rect0"].fill_alpha == pytest.approx(0.29804)
    assert by_name["rect1"].fill_alpha == 1.0
    assert by_name["rect2"].fill_alpha == pytest.approx(0.5)
    assert by_name["scheme"].fill_hex and by_name["scheme"].fill_alpha == pytest.approx(0.4)
    ir = build_deck_ir(p)
    els = {e.name: e for e in ir.slides[0].elements}
    assert els["rect0"].fill_alpha == pytest.approx(0.29804)
    assert els["rect1"].fill_alpha == 1.0
    assert els["scheme"].fill_alpha == pytest.approx(0.4)


@pytest.mark.integration
def test_aud2_workspace_cards_are_translucent(fixtures_dir):
    pkg, shapes = _shapes(_vk(fixtures_dir, "WorkSpace"), 3)
    s = next(s for s in shapes if s.id == "413")
    assert s.fill_hex == "0077FF" and s.fill_alpha == pytest.approx(0.05098)


# ---------------------------------------------------------------------------- R3: table frames with a dummy xfrm


def _table_deck(out: Path) -> Path:
    prs = _blank_prs()
    # slide 1: dummy xfrm (0.246×0.437 of the slide) but grid 0.9·W × 0.6·H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    gf = s.shapes.add_table(3, 4, Emu(int(W * 0.05)), Emu(int(H * 0.2)), Emu(int(W * 0.9)), Emu(int(H * 0.6)))
    ext = find(gf._element, "p:xfrm/a:ext")
    ext.set("cx", str(int(W * 0.246)))
    ext.set("cy", str(int(H * 0.437)))
    # slide 2: xfrm larger than the grid → xfrm wins
    s = prs.slides.add_slide(prs.slide_layouts[6])
    gf = s.shapes.add_table(2, 2, Emu(int(W * 0.1)), Emu(int(H * 0.2)), Emu(int(W * 0.4)), Emu(int(H * 0.3)))
    ext = find(gf._element, "p:xfrm/a:ext")
    ext.set("cx", str(int(W * 0.8)))
    ext.set("cy", str(int(H * 0.5)))
    # slide 3: grid wider than the slide remainder → clamped to the slide edge
    s = prs.slides.add_slide(prs.slide_layouts[6])
    gf = s.shapes.add_table(2, 2, Emu(int(W * 0.5)), Emu(int(H * 0.6)), Emu(int(W * 0.9)), Emu(int(H * 0.6)))
    ext = find(gf._element, "p:xfrm/a:ext")
    ext.set("cx", str(int(W * 0.1)))
    ext.set("cy", str(int(H * 0.1)))
    prs.save(out)
    return out


def test_r3_table_frame_size_from_grid(tmp_path):
    p = _table_deck(tmp_path / "tables.pptx")
    pkg, shapes = _shapes(p, 1)
    t = next(s for s in shapes if s.frame_kind == "table")
    assert t.table_dims == (3, 4)
    assert t.bbox.x == int(W * 0.05) and t.bbox.y == int(H * 0.2)
    assert t.bbox.w == pytest.approx(W * 0.9, abs=W * 0.002)
    assert t.bbox.h == pytest.approx(H * 0.6, abs=H * 0.002)
    pkg, shapes = _shapes(p, 2)
    t = next(s for s in shapes if s.frame_kind == "table")
    assert t.bbox.w == int(W * 0.8) and t.bbox.h == int(H * 0.5)
    pkg, shapes = _shapes(p, 3)
    t = next(s for s in shapes if s.frame_kind == "table")
    assert t.bbox.x2 <= W and t.bbox.y2 <= H
    assert t.bbox.w == pytest.approx(W * 0.5, abs=W * 0.002) and t.bbox.h == pytest.approx(H * 0.4, abs=H * 0.002)
    ir = build_deck_ir(p)
    e = next(e for e in ir.slides[0].elements if e.type == "table")
    assert e.bbox_frac.w == pytest.approx(0.9, abs=0.002) and e.bbox_frac.h == pytest.approx(0.6, abs=0.002)


@pytest.mark.integration
def test_r3_workspace_slide_14_table_is_full_width(fixtures_dir):
    pkg, shapes = _shapes(_vk(fixtures_dir, "WorkSpace"), 14)
    sw, sh = pkg.slide_size
    t = next(s for s in shapes if s.id == "673")
    assert t.frame_kind == "table"
    assert t.bbox.w / sw == pytest.approx(1.0, abs=0.01)
    assert t.bbox.h / sh == pytest.approx(0.789, abs=0.01)
    assert t.bbox.x2 <= sw and t.bbox.y2 <= sh


# ---------------------------------------------------------------------------- R4c: empty subtitle placeholder under the title


def _title_slide_deck(out: Path) -> Path:
    prs = _blank_prs()
    # slide 1: title + empty body placeholder directly under it (a designer's subtitle slot)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    title = s.shapes.title
    title.left, title.top, title.width, title.height = Emu(int(W * 0.06)), Emu(int(H * 0.40)), Emu(int(W * 0.8)), Emu(int(H * 0.16))
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Emu(int(W * 0.06)), Emu(int(H * 0.60)), Emu(int(W * 0.8)), Emu(int(H * 0.08))
    # slide 2: title + tall empty body placeholder (a real content slot, stays bullet_list)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Content"
    body = s.placeholders[1]
    body.top, body.height = Emu(int(H * 0.28)), Emu(int(H * 0.55))
    # slide 3: title + six equal empty body placeholders stacked in a column (an agenda list, not a subtitle)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Agenda"
    title = s.shapes.title
    title.top, title.height = Emu(int(H * 0.08)), Emu(int(H * 0.12))
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Emu(int(W * 0.06)), Emu(int(H * 0.25)), Emu(int(W * 0.6)), Emu(int(H * 0.07))
    tree = find(s._element, "p:cSld/p:spTree")
    for k in range(1, 6):
        el = copy.deepcopy(body._element)
        nv = find(el, "p:nvSpPr/p:cNvPr")
        nv.set("id", str(100 + k))
        nv.set("name", f"Body {k}")
        find(el, "p:spPr/a:xfrm/a:off").set("y", str(int(H * (0.25 + k * 0.11))))
        tree.append(el)
    prs.save(out)
    return out


def test_r4c_empty_body_placeholder_under_title_is_subtitle(tmp_path):
    p = _title_slide_deck(tmp_path / "title.pptx")
    pkg, shapes = _shapes(p, 1)
    groups = detect_repeat_groups(shapes, W, H)
    roles = heuristic_roles(shapes, groups, set(), TYPO, W, H)
    ph = next(s for s in shapes if s.is_placeholder and s.ph_type == "body")
    assert not ph.has_text
    assert roles[ph.id] == SlotRole.subtitle
    assert SlotRole.title in roles.values()
    kind, _ = heuristic_kind(shapes, roles, groups, 1, 10, W, H)
    assert kind == PatternKind.title

    pkg, shapes = _shapes(p, 2)
    groups = detect_repeat_groups(shapes, W, H)
    roles = heuristic_roles(shapes, groups, set(), TYPO, W, H)
    ph = next(s for s in shapes if s.is_placeholder and s.ph_type == "body")
    assert roles[ph.id] == SlotRole.bullet_list

    pkg, shapes = _shapes(p, 3)
    groups = detect_repeat_groups(shapes, W, H)
    roles = heuristic_roles(shapes, groups, set(), TYPO, W, H)
    phs = [s for s in shapes if s.is_placeholder and s.ph_type == "body"]
    assert len(phs) == 6
    assert SlotRole.subtitle not in roles.values()


@pytest.mark.integration
def test_r4c_vk_templates_subtitle_placeholders(fixtures_dir):
    edu = _analyze_deck(_vk(fixtures_dir, "Education"))
    for n, sid in [(1, "236"), (2, "242"), (53, "1084")]:
        assert edu[n][2][sid] == SlotRole.subtitle, (n, sid, edu[n][2][sid])
    tech = _analyze_deck(_vk(fixtures_dir, "VK Tech"))
    assert SlotRole.subtitle not in tech[7][2].values()
    assert SlotRole.subtitle not in tech[8][2].values()


# ---------------------------------------------------------------------------- F5: tables are not text boxes for the audit


def _table_text_deck(out: Path) -> Path:
    prs = _blank_prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Emu(int(W * 0.05)), Emu(int(H * 0.05)), Emu(int(W * 0.6)), Emu(int(H * 0.1)))
    tb.text_frame.text = "Table slide"
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.7)), Emu(int(H * 0.05)), Emu(int(W * 0.2)), Emu(int(H * 0.1)))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0xED, 0xF3, 0xFC)
    t = s.shapes.add_table(5, 4, Emu(int(W * 0.05)), Emu(int(H * 0.2)), Emu(int(W * 0.9)), Emu(int(H * 0.6))).table
    for r in range(5):
        for c in range(4):
            t.cell(r, c).text = f"r{r}c{c}"
    prs.save(out)
    return out


def test_f5_slide_texts_excludes_tables(tmp_path):
    ir = build_deck_ir(_table_text_deck(tmp_path / "table_text.pptx"))
    sl = ir.slides[0]
    types = [e.type for e in sl.elements]
    assert types.count("text") == 1 and types.count("table") == 1 and types.count("shape") == 1
    assert [e.type for e in sl.texts] == ["text"]
    table = next(e for e in sl.elements if e.type == "table")
    assert table.table is not None and len(table.table.rows) == 5 and all(len(r) == 4 for r in table.table.rows)
    assert table.has_text  # cell text stays available for consumers that want it
    assert {e.type for e in sl.all_text} == {"text", "table"}


# ---------------------------------------------------------------------------- manifest cache invalidation


def test_analysis_version_bumped():
    assert int(TemplateManifest.model_fields["analysis_version"].default) >= 8
