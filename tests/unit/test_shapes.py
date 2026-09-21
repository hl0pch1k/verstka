from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches

from verstka.analysis.shapes import SlideContext, extract_shapes, looks_like_placeholder, slide_family
from verstka.ingest.package import PptxPackage
from verstka.schemas.common import Family, ShapeKind


def _shapes(pptx: Path, n: int):
    pkg = PptxPackage.open(pptx)
    part = pkg.slide_parts[n - 1]
    ctx = SlideContext(pkg, part)
    return pkg, ctx, extract_shapes(pkg, part, ctx)


def test_title_placeholder_inherits_size_and_bbox(simple_deck):
    pkg, ctx, shapes = _shapes(simple_deck, 1)
    title = next(s for s in shapes if s.ph_type == "ctrTitle")
    assert title.is_placeholder and title.plain_text == "Quarterly results"
    assert title.text.max_size_pt >= 40  # inherited from master titleStyle
    assert title.bbox.w > 0 and title.bbox.h > 0  # inherited from layout
    assert title.text.dominant_font  # theme font resolved
    sub = next(s for s in shapes if s.ph_type == "subTitle")
    assert sub.plain_text == "Q3 2026 overview" and sub.text.has_bullets is False
    fam, bg = slide_family(pkg, pkg.slide_parts[0], ctx, shapes)
    assert fam == Family.light


def test_cards_fill_geometry_and_runs(simple_deck):
    pkg, ctx, shapes = _shapes(simple_deck, 2)
    cards = [s for s in shapes if s.fill_hex == "EDF3FC"]
    assert len(cards) == 3 and all(c.geometry == "roundRect" for c in cards)
    assert all(c.corner_radius and c.corner_radius > 0 for c in cards)
    titles = [s for s in shapes if s.plain_text.startswith("Pillar")]
    assert len(titles) == 3
    assert titles[0].text.bold_share == 1.0 and titles[0].text.max_size_pt == 20 and titles[0].text.dominant_color == "0077FF"
    footer = next(s for s in shapes if s.plain_text == "ACME")
    assert footer.text.max_size_pt == 10


def test_picture_and_table_frame(simple_deck):
    pkg, ctx, shapes = _shapes(simple_deck, 3)
    pic = next(s for s in shapes if s.kind == ShapeKind.pic)
    assert pic.image_part and pic.image_part.startswith("ppt/media/") and pic.image_ext == "png"


def test_group_transform(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    grp = slide.shapes.add_group_shape()
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(1), Inches(1))
    # move the group by +1in without touching child offsets
    xfrm = grp._element.grpSpPr.xfrm
    xfrm.off.x = Inches(2)
    xfrm.off.y = Inches(2)
    p = tmp_path / "grp.pptx"
    prs.save(p)
    pkg, ctx, shapes = _shapes(p, 1)
    rects = [s for s in shapes if s.geometry == "rect"]
    assert len(rects) == 2 and rects[0].group_path
    assert rects[0].bbox.x == Inches(2) and rects[0].bbox.y == Inches(2)
    assert rects[1].bbox.x == Inches(4)


def test_placeholder_regex():
    assert looks_like_placeholder("Заголовок в две или в одну строчку")
    assert looks_like_placeholder("Lorem ipsum dolor sit amet")
    assert looks_like_placeholder("Имя Фамилия")
    assert not looks_like_placeholder("Выручка выросла на 12% за квартал")
