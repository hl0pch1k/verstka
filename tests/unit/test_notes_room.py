"""Regression tests (the live QA of 2026-10-07): the author's planned lines stay on the slide when a denser setting holds them."""
from __future__ import annotations

import pytest
from lxml import etree

from verstka.analysis.manifest import analyze_template
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.deck import DeckBuilder
from verstka.rendering.synth import render_synth
from verstka.schemas.common import EMU_PER_PT, PatternKind
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, TableData
from verstka.schemas.template import TypeStep


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _render(env, o: OutlineSlide, comp: str, strategy: str):
    deck, manifest, ws = env
    outline = DeckOutline(title="T", strategy=strategy, slides=[o])
    return render_synth(DeckBuilder(deck), LayoutSlide(outline_id=o.id, mode="synth", composition=comp), o, manifest, ws, outline)


def _texts(slide) -> str:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            out.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(out).replace(" ", " ").replace("⁠", "")


def _table_size(slide) -> float:
    for sh in slide.shapes:
        if sh.has_table:
            szs = [int(r.get("sz")) / 100 for r in sh._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr") if r.get("sz")]
            return min(szs)
    raise AssertionError("no table")


ROWS = [["Покупки в день", "100", "115"], ["Средний чек", "300 рублей", "330 рублей"], ["Месячная выручка", "900 000 рублей", "1 138 500 рублей"],
        ["Переменные расходы", "315 000 рублей", "375 705 рублей"], ["Постоянные операционные расходы", "465 000 рублей", "508 000 рублей"],
        ["Операционная прибыль", "120 000 рублей", "254 795 рублей"], ["Операционная рентабельность", "13,3%", "22,4%"]]
LINES = ["Рост постоянных расходов: +43 000 ₽", "Персонал: 20 000 ₽", "Маркетинг: 15 000 ₽", "Лояльность и учет: 5 000 ₽", "Коммунальные услуги: 3 000 ₽"]


def _small_under_floor(manifest, hpt: float) -> float:
    """The template's small step set at 2.1 % of the slide height: the Kit snaps «small» to it (it allows 1.9 %), the
    table must be allowed to use it too."""
    small = round(0.021 * hpt * 2) / 2
    typo = manifest.tokens.typography
    near = lambda x: small < x <= 0.022 * hpt + 0.6  # noqa: E731 — no step of the template between it and 2.2 % H
    typo.scale = [s for s in typo.scale if s.role != "small"] + [TypeStep(role="small", size_pt=small)]
    typo.sizes_used = sorted({x for x in typo.sizes_used if not near(x)} | {small})
    typo.derived_sizes = [x for x in (typo.derived_sizes or []) if not near(x)]
    return small


@pytest.mark.parametrize("strategy", ["structured", "compact", "visual"])
def test_a_table_steps_down_to_the_templates_small_size_before_its_lines_leave(env, strategy):
    deck, manifest, ws = env
    hpt = manifest.slide_size.h / EMU_PER_PT
    small = _small_under_floor(manifest, hpt)
    assert small < 0.022 * hpt
    from verstka.rendering.compose import Kit
    assert Kit(manifest, manifest.slide_size.w, manifest.slide_size.h, "FFFFFF").small == small  # the Kit takes it
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Операционная прибыль вырастет до 254 795 ₽ к шестому месяцу",
                     takeaway="Выручка увеличивается на 26,5%, а прибыль — примерно на 112%",
                     content=SlideContent(table=TableData(columns=["Показатель", "Сейчас", "Цель"], rows=ROWS), bullets=list(LINES)))
    slide, warnings = _render(env, o, "table", strategy)
    text = _texts(slide)
    missing = [t for t in LINES if t.replace("рублей", "₽") not in text and t not in text]
    assert not missing, (missing, warnings)
    assert "• " not in (o.notes or ""), o.notes
    assert "Вывод:" not in (o.notes or ""), o.notes


@pytest.mark.parametrize("strategy", ["structured", "visual"])
def test_number_tiles_beside_a_photo_place_set_denser_before_the_conclusion_leaves(env, strategy):
    nums = [NumberCallout(value="25,4 °C", label="Температура воздуха"), NumberCallout(value="67%", label="Влажность воздуха"),
            NumberCallout(value="72%", label="Заполненность бака"), NumberCallout(value="40 сек", label="Последнее обновление")]
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Панель показывает текущие показания и историю за 24 часа",
                     takeaway="Данные обновляются каждые 40 секунд",
                     content=SlideContent(numbers=nums, bullets=["Графики за сутки рядом с текущими значениями", "Предупреждения объясняют причину и требуемое действие", "Отметки включения полива и вентиляции на графике"], photo_slot="Фото", photo_hint="Экран панели"))
    slide, warnings = _render(env, o, "stat_row", strategy)
    names = [sh.name for sh in slide.shapes]
    assert any(n.startswith("Conclusion") for n in names), (warnings, o.notes)
    assert "Вывод:" not in (o.notes or ""), o.notes


def test_the_audit_reads_the_typesetters_minus_and_ranges():
    """«-31,25%» is set «−31,25%» (a true minus) and «5-6 тыс.» «5⁠–⁠6 тыс.»: the planned line is on the slide."""
    from verstka.audit.checks.integrity import _norm_text, _present
    from verstka.ru import typeset

    for line in ("Эксперимент: -31,25% воды, -12,5 ч ручной работы в месяц", "Рост 5-6 тыс. гостей в месяц", "-5 °C ночью"):
        assert _present(line, _norm_text(typeset(line))), (line, typeset(line))


# ---------------------------------------------------------------------------------------------- the review of 3.0 (R4)

ROWS10 = [["Покупки в день", "100", "115", "+15%"], ["Средний чек", "300 рублей", "330 рублей", "+10%"], ["Доля чеков с едой", "20%", "30%", "+10 п. п."],
          ["Месячная выручка", "900 000 рублей", "1 138 500 рублей", "+26,5%"], ["Переменные расходы", "315 000 рублей", "375 705 рублей", "+19,3%"],
          ["Постоянные расходы", "465 000 рублей", "508 000 рублей", "+9,2%"], ["Списания продуктов", "27 000 рублей", "18 000 рублей", "−33%"],
          ["Операционная прибыль", "120 000 рублей", "254 795 рублей", "+112%"], ["Операционная рентабельность", "13,3%", "22,4%", "+9,1 п. п."],
          ["Возвращающиеся гости", "25%", "35%", "+10 п. п."]]


def _two_tables():
    big = OutlineSlide(id="big", kind=PatternKind.table, headline="Десять показателей кофейни: сейчас и через полгода",
                       takeaway="Прибыль растёт быстрее выручки: постоянные расходы растут медленнее, чем продажи",
                       content=SlideContent(table=TableData(columns=["Показатель", "Сейчас", "Цель", "Изменение"], rows=ROWS10),
                                            bullets=["Расчёт на 30 дней в месяце", "Цены поставщиков без изменений", "Налоги и кредиты не учтены"]))
    small = OutlineSlide(id="sm", kind=PatternKind.table, headline="Три сценария запуска", content=SlideContent(table=TableData(
        columns=["Сценарий", "Покупки в день", "Прибыль в месяц"], rows=[["Осторожный", "105", "160 000 рублей"], ["Базовый", "115", "254 795 рублей"], ["Смелый", "125", "330 000 рублей"]])))
    return big, small


def _render_in(env, deck_slides: list[OutlineSlide], o: OutlineSlide, strategy: str):
    deck, manifest, ws = env
    outline = DeckOutline(title="T", strategy=strategy, slides=deck_slides)
    return render_synth(DeckBuilder(deck), LayoutSlide(outline_id=o.id, mode="synth", composition="table"), o, manifest, ws, outline)


@pytest.mark.parametrize("strategy", ["structured", "compact", "visual"])
def test_a_table_of_the_deck_is_never_set_at_another_tables_dense_size(env, strategy):
    """MyBrand compact s3: a three-row table on an almost empty slide stepped from 22.5 to 17.5 pt — the deck's one table
    size was measured with the steps under the body size a ten-row table of another slide took. The deck's size is
    measured on the regular ladder (the small size only from 2.2 % of the slide height); the steps are that slide's."""
    deck, manifest, ws = env
    _small_under_floor(manifest, manifest.slide_size.h / EMU_PER_PT)
    from verstka.rendering.compose import Kit

    kit = Kit(manifest, manifest.slide_size.w, manifest.slide_size.h, "FFFFFF")
    big, small = _two_tables()
    slide, _ = _render_in(env, [big, small], small.model_copy(deep=True), strategy)
    assert _table_size(slide) >= kit.body - 0.05, (_table_size(slide), kit.body)


@pytest.mark.parametrize("strategy", ["structured", "compact", "visual"])
def test_the_lines_under_a_table_are_never_larger_than_the_table(env, strategy):
    """MyBrand s9: the table stepped from 22.5 to 17.5 pt to keep its conclusion, the lines under it stayed at 22.5 pt —
    the hierarchy inverted. The lines under a table are never set larger than the table — nor, outside the dense pass,
    under the template's body size (running text keeps the body size: the lines follow a table in small print down to
    it, not below); a dense setting stays only where it brings the lines onto the slide (VK Tech compact s2: a ten-row
    table at 8 pt for lines that left anyway)."""
    deck, manifest, ws = env
    _small_under_floor(manifest, manifest.slide_size.h / EMU_PER_PT)
    from verstka.rendering.compose import Kit

    kit = Kit(manifest, manifest.slide_size.w, manifest.slide_size.h, "FFFFFF")
    big, small = _two_tables()
    o = big.model_copy(deep=True)
    slide, warnings = _render_in(env, [big, small], o, strategy)
    size = _table_size(slide)
    lines = [sh for sh in slide.shapes if sh.has_text_frame and sh.name.startswith("Note")]
    sizes = {r.font.size.pt for sh in lines for p in sh.text_frame.paragraphs for r in p.runs if r.font.size is not None and r.text.strip()}
    dense = any("dense" in w for w in warnings)
    assert not sizes or max(sizes) <= max(size, 0.0 if dense else kit.text_floor) + 0.05, (sizes, size, warnings)
    assert not sizes or dense or min(sizes) >= kit.text_floor - 0.05, (sizes, warnings)
    aloud = "• " in (o.notes or "")
    assert not aloud or size >= kit.small - 0.05, (size, warnings)  # said aloud anyway: the table keeps its size


def test_a_table_on_a_sample_never_lies_over_the_slides_lines(tmp_path):
    """Кофейня s2: a list sample (one text box under the heading) took the slide's three lines and its conclusion, and
    the ten-row table was drawn over them — the gate reads text over text, not a table over text. A native object that
    would lie over the slide's own text fails on such a sample: the renderer takes another one or composes the slide."""
    from pptx import Presentation as _P
    from pptx.util import Emu as _E, Pt as _Pt

    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import render_clone

    W, H = 12192000, 6858000
    prs = _P()
    prs.slide_width, prs.slide_height = _E(W), _E(H)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Название"
    s1.placeholders[1].text = "Подзаголовок"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Заголовок слайда"
    box = s2.shapes.add_textbox(_E(int(W * 0.05)), _E(int(H * 0.25)), _E(int(W * 0.9)), _E(int(H * 0.65)))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Первый тезис списка"
    for t in ("Второй тезис списка", "Третий тезис списка"):
        box.text_frame.add_paragraph().text = t
    for p in box.text_frame.paragraphs:
        p.runs[0].font.size = _Pt(18)
    prs.save(tmp_path / "list.pptx")
    manifest = analyze_template(tmp_path / "list.pptx", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    pattern = next(p for p in manifest.patterns if p.source_slide == 2)
    big, _ = _two_tables()
    o = big.model_copy(update={"takeaway": None})
    with pytest.raises(RuntimeError, match="lie over the slide's text"):
        render_clone(DeckBuilder(tmp_path / "list.pptx"), LayoutSlide(outline_id="big", mode="clone", pattern_id=pattern.id), o, pattern, manifest, ws, DeckOutline(title="T", slides=[o]))


def test_the_slides_lines_never_take_a_label_inside_the_samples_chart(tmp_path):
    """VK Education p45: the sample's chart is a picture with a caption printed inside it («Диаграмма про Петербург»).
    The slide's line went into that caption and the native chart, drawn in the picture's place, covered it — the
    sample was turned down («would lie over the slide's text»). The caption is the sample chart's own: the line is
    set beside the chart, clear of it."""
    from PIL import Image, ImageDraw
    from pptx import Presentation as _P
    from pptx.util import Emu as _E, Pt as _Pt

    from verstka.rendering.clone import render_clone
    from verstka.rendering.clone_check import ink_box
    from verstka.rendering.deck import element_bbox
    from verstka.schemas.common import Bbox
    from verstka.schemas.outline import ChartSpec, Series

    W, H = 12192000, 6858000
    img = Image.new("RGB", (900, 700), (245, 245, 250))
    d = ImageDraw.Draw(img)
    for i, h in enumerate((180, 260, 340, 420, 520)):
        d.rectangle((80 + i * 160, 620 - h, 180 + i * 160, 620), fill=(40 + 40 * i, 90, 220 - 30 * i))
        d.text((95 + i * 160, 640), f"Item {i + 1}", fill=(60, 60, 60))
    for j, y in enumerate(range(100, 620, 104)):
        d.text((20, y), str(50 - 10 * j), fill=(60, 60, 60))
    img.save(tmp_path / "chart.png")
    prs = _P()
    prs.slide_width, prs.slide_height = _E(W), _E(H)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Название"
    s1.placeholders[1].text = "Подзаголовок"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Диаграмма продаж по месяцам"
    s2.shapes.add_picture(str(tmp_path / "chart.png"), _E(int(W * 0.05)), _E(int(H * 0.26)), _E(int(W * 0.60)), _E(int(H * 0.64)))
    cap = s2.shapes.add_textbox(_E(int(W * 0.20)), _E(int(H * 0.45)), _E(int(W * 0.16)), _E(int(H * 0.15)))
    cap.text_frame.word_wrap = True
    cap.text_frame.text = "Подпись внутри диаграммы"
    cap.text_frame.paragraphs[0].runs[0].font.size = _Pt(14)
    prs.save(tmp_path / "labelled.pptx")
    manifest = analyze_template(tmp_path / "labelled.pptx", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    pattern = next(p for p in manifest.patterns if p.source_slide == 2)
    line = "Органический рост без промо: сотрудники подключали коллег сами"
    o = OutlineSlide(id="c", kind=PatternKind.chart, headline="Аудитория росла в 10 раз за пять месяцев",
                     content=SlideContent(chart=ChartSpec(series_ids=["s1"]), paragraphs=[line]))
    outline = DeckOutline(title="T", slides=[o], series=[Series(id="s1", name="Пользователи", categories=["Май", "Июнь", "Июль"], values=[1200, 3400, 6100])])
    slide, warns = render_clone(DeckBuilder(tmp_path / "labelled.pptx"), LayoutSlide(outline_id="c", mode="clone", pattern_id=pattern.id), o, pattern, manifest, ws, outline)
    assert "Органический рост" in _texts(slide), warns
    frames = [Bbox(x=b[0], y=b[1], w=b[2], h=b[3]) for b in (element_bbox(f) for f in slide._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}graphicFrame")) if b]
    assert frames
    fam = manifest.tokens.typography.primary_family
    held = [sh for sh in slide.shapes if sh.has_text_frame and "Органический" in sh.text_frame.text]
    under = [sh.name for sh in held for f in frames if (ib := ink_box(sh._element, fam)) is not None and f.intersection(ib) > 0.1 * ib.area]
    assert not under, (under, warns)
