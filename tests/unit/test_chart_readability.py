"""Порог читаемости подписей графиков и таблиц (финальная доводка, направление C, 17.10): подписи категорий, значений,
осей и легенды не мельче 2,4 % высоты слайда или ¾ основного кегля шаблона (не больше 2,6 % высоты); что не помещается
на пороге — прореживается (ось времени, точки линии, интервалы гистограммы), переносится на 2–3 строки или режется «…»
у слова, но не уменьшается и не поворачивается на 45°. Таблица — не мельче того же порога."""

from __future__ import annotations

import re

import pytest
from pptx import Presentation
from pptx.util import Emu

from verstka.rendering.charts import add_chart, fit_categories, readable_floor, readable_pt, template_body_pt
from verstka.rendering.tables import default_sizes, measure_table, readable_sizes, readable_table_pt, table_need_h
from verstka.schemas.common import Bbox
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, TableData
from verstka.schemas.template import ChartStyleSpec, TypeStep, Typography

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
OUTLINE = DeckOutline(title="t")


def _typo(body: float) -> Typography:
    return Typography(primary_family="Play", scale=[TypeStep(role="h1", size_pt=round(body * 2)), TypeStep(role="body", size_pt=body), TypeStep(role="small", size_pt=round(body * 0.8))])


def _slide(w_emu: int, h_emu: int):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(w_emu), Emu(h_emu)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _sizes(gf) -> list[float]:
    """Every size the chart's own text is set at (its axes, legend, labels, title)."""
    xml = gf.chart.part.blob.decode("utf-8", "ignore")
    return [int(x) / 100 for x in re.findall(r'\bsz="(\d+)"', xml)]


def _in_cat_axis(el) -> bool:
    while el is not None:
        if el.tag == f"{C}catAx":
            return True
        el = el.getparent()
    return False


def _rots(gf) -> set[str]:
    xml = gf.chart.part.blob.decode("utf-8", "ignore")
    return set(re.findall(r'<a:bodyPr[^>]*\brot="(-?\d+)"', xml))


# -------------------------------------------------------------------------------------------------- the threshold


def test_readable_threshold_is_relative_to_the_slide_height_and_the_template_body():
    # 2.4 % of the height is the floor: 13 pt on a 540 pt slide, 10 pt on VK Tech's 405 pt (its body is 9 pt)
    assert readable_pt(540.0, 14.0) == 13.0
    assert readable_pt(405.0, 9.0) == 10.0
    # a larger body raises it, up to 2.6 % of the height (a 24 pt bullet body is no reason for 18 pt labels)
    assert readable_pt(446.0, 15.0) == 11.5  # 0.75 × 15 = 11.25 (under 2.6 % × 446 = 11.6), above 2.4 % × 446 = 10.7
    assert readable_pt(446.0, 24.0) == 12.0  # capped: 11.6, up to half a point
    assert readable_pt(446.0, None) == 11.0
    assert template_body_pt(_typo(15.0)) == 15.0 and template_body_pt(Typography()) is None


def test_readable_floor_reads_the_slide_and_the_typography():
    prs, s = _slide(9144000, 5143500)
    assert readable_floor(s, _typo(9.0)) == 10.0
    prs, s = _slide(12192000, 6858000)
    assert readable_floor(s, _typo(18.0)) == 13.5  # 0.75 × 18 = 13.5 ≤ 2.6 % × 540 = 14.04


# -------------------------------------------------------------------------------------------------- charts


def _bars(n: int, words: int = 1) -> ChartSpec:
    names = ["Москва", "Санкт-Петербург", "Екатеринбург", "Казань", "Новосибирск", "Краснодар", "Ростов-на-Дону", "Нижний Новгород", "Самара", "Уфа", "Пермь", "Воронеж"]
    return ChartSpec(type="bar", categories=names[:n], series=[InlineSeries(name="Выручка", values=[1433 - 100 * i for i in range(n)])], unit="млн ₽")


@pytest.mark.parametrize("h_pt", [100, 143, 200, 300])
@pytest.mark.parametrize("n", [4, 8, 12])
def test_bar_labels_never_go_under_the_threshold_in_any_frame(h_pt, n):
    """A tight frame of VK Tech (720 × 405 pt, 9 pt chart style) used to set 8 regions at 9 pt and 12 at 7.5 pt: now every
    label is at the threshold or above; a frame too short says so (`verstka_rows_fit`, `verstka_min_height`)."""
    prs, s = _slide(9144000, 5143500)
    gf = add_chart(s, Bbox(x=500000, y=1300000, w=8000000, h=int(h_pt * 12700)), _bars(n), OUTLINE, ChartStyleSpec(series_colors=["0077FF"], font_family="Play", font_size_pt=9.0), _typo(9.0), text_hex="000000")
    cs = gf.chart._chartSpace
    names = int(cs.find(f".//{C}catAx/{C}txPr/{A}p/{A}pPr/{A}defRPr").get("sz")) / 100
    others = [int(e.get("sz")) / 100 for e in cs.iter(f"{A}defRPr") if e.get("sz") and not _in_cat_axis(e)]
    assert min(others) >= 10.0 - 1e-6, (h_pt, n, sorted(set(others)))  # values, ticks, titles: the threshold
    assert names >= 8.5 - 1e-6  # the row names may give ROW_NAMES_GIVE (0.85) of it while frames are cut for 2 % labels
    assert gf.verstka_font_pt >= 10.0
    if not gf.verstka_rows_fit:
        assert gf.verstka_min_height > h_pt


def test_values_stay_on_the_bars_at_the_label_size_before_they_leave_them():
    """VK Tech's sample frame for 8 regions (675 × 143 pt): the values keep their place on the bars (at the labels' size)
    instead of going to a value axis."""
    prs, s = _slide(9144000, 5143500)
    gf = add_chart(s, Bbox(x=500000, y=1300000, w=int(675 * 12700), h=int(143 * 12700)), _bars(8), OUTLINE, ChartStyleSpec(series_colors=["0077FF"], font_family="Play", font_size_pt=9.0), _typo(9.0), text_hex="000000")
    assert gf.chart.plots[0].has_data_labels


@pytest.mark.parametrize("w_share", [0.25, 0.35, 0.5, 0.9])
@pytest.mark.parametrize("kind", ["line", "column", "histogram"])
def test_category_labels_are_never_turned_by_45_degrees(kind, w_share):
    """No label is turned: a line's points and a histogram's bins are thinned out (tilted a degree), other columns wrap."""
    prs, s = _slide(9144000, 5143500)
    if kind == "histogram":
        cats = [f"{a}–{a + 20}" for a in range(0, 200, 20)] + ["больше 200"]
    elif kind == "line":
        cats = ["Сейчас"] + [f"{i}-й месяц" for i in range(1, 7)]
    else:
        cats = ["Прямые продажи", "Партнёры", "Маркетплейс", "Онлайн", "Розница", "Опт"]
    spec = ChartSpec(type=kind, categories=cats, series=[InlineSeries(name="Значение", values=[100 + 7 * i for i in range(len(cats))])], unit="шт.")
    gf = add_chart(s, Bbox(x=300000, y=1300000, w=int(9144000 * w_share), h=int(160 * 12700)), spec, OUTLINE, ChartStyleSpec(series_colors=["0077FF"], font_family="Play", font_size_pt=9.0), _typo(9.0), text_hex="000000")
    assert _rots(gf) <= {"0", "-60000", "-5400000"}  # (straight, a 1° tilt of thinned labels, an axis title along its axis)
    ax = gf.chart._chartSpace.find(f".//{C}catAx/{C}txPr/{A}p/{A}pPr/{A}defRPr")
    bar_dir = gf.chart._chartSpace.find(f".//{C}barDir")
    rows = bar_dir is not None and bar_dir.get("val") == "bar"  # (long names: horizontal bars, whose row names may give)
    assert int(ax.get("sz")) / 100 >= (8.5 if rows else 10.0)


def test_fit_categories_never_turns_names_and_cuts_a_word_too_wide_last():
    names = ["Автоматизация", "Маршрутизация", "Обучение", "Контроль", "Прочее", "Другое"]
    got = fit_categories(names, "Play", [10.0], 30.0, 60.0)
    assert got.mode == "wrap" and not got.fits  # (a word wider than its slot: cut to «…», the frame is too narrow)
    assert all(t.endswith("…") or t in names for t in got.labels)
    got = fit_categories(["Прямые продажи", "Партнёры", "Маркетплейс"], "Play", [10.0], 70.0, 60.0)
    assert got.mode in ("flat", "wrap") and got.fits


# -------------------------------------------------------------------------------------------------- tables


def _table(rows: int = 10) -> TableData:
    return TableData(columns=["Клиент", "Отрасль", "Регион", "Выручка, млн ₽", "Доля, %"],
                     rows=[[f"ООО «Клиент {i}»", "Промышленность", "Москва", f"{200 - 10 * i:.1f}", f"{6 - 0.3 * i:.1f}"] for i in range(rows)])


def test_table_threshold_is_the_chart_labels_one():
    assert readable_table_pt(5143500, _typo(9.0)) == 10.0
    assert readable_table_pt(6858000, _typo(14.0)) == 13.0
    assert readable_sizes([15.0, 11.0, 9.0], 10.0) == [15.0, 11.0]
    assert readable_sizes([9.0, 8.0], 10.0) == [10.0]


def test_measure_table_never_sets_a_size_under_its_min_and_reports_overflow():
    t = _table(10)
    fit = measure_table(t, 675 * 12700, "Play", [11.0, 9.0], 1.2, max_h_emu=137 * 12700, fill_h_emu=0, min_row_h_emu=0, max_row_h_emu=0, min_size_pt=10.0)
    assert fit.size_pt >= 10.0 and fit.overflow  # (137 pt holds no ten rows at a readable size: the composer must give room)
    need = table_need_h(t, 675 * 12700, "Play", 11.0)
    fit = measure_table(t, 675 * 12700, "Play", [11.0, 9.0], 1.2, max_h_emu=need + 12700, fill_h_emu=0, min_row_h_emu=0, max_row_h_emu=0, min_size_pt=10.0)
    assert fit.size_pt == 11.0 and not fit.overflow


def test_default_table_sizes_stop_at_the_threshold():
    sizes = default_sizes(_typo(9.0), 5143500, 7.0)
    assert min(sizes) >= 10.0


# -------------------------------------------------------------------------------------------------- audit


def test_the_audit_reads_a_chart_label_size_and_flags_one_under_the_threshold(tmp_path):
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext, run_checks
    from verstka.schemas.template import TemplateManifest

    prs, s = _slide(9144000, 5143500)
    gf = add_chart(s, Bbox(x=500000, y=1300000, w=8000000, h=3000000), _bars(4), OUTLINE, ChartStyleSpec(series_colors=["0077FF"], font_family="Play", font_size_pt=9.0), _typo(9.0), text_hex="000000")
    path = tmp_path / "c.pptx"
    prs.save(path)
    ir = build_deck_ir(path, with_images=False)
    ir.slides[0].outline_id = "c"  # (a slide the engine set: the template's own sample charts are not checked)
    chart = next(e.chart for e in ir.slides[0].elements if e.chart is not None)
    assert chart.label_pt is not None and chart.label_pt >= 10.0
    ctx = AuditContext(ir=ir, manifest=TemplateManifest.model_construct(template_id="t"))
    assert not run_checks(ctx, only={"chart_text_small"})[0]
    small = ir.model_copy(deep=True)
    next(e.chart for e in small.slides[0].elements if e.chart is not None).label_pt = 7.0
    got, _ = run_checks(AuditContext(ir=small, manifest=TemplateManifest.model_construct(template_id="t")), only={"chart_text_small"})
    assert len(got) == 1 and got[0].severity == "warn"
