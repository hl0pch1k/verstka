"""Unseen templates, reliability round (2026-10-17): a content slide never on the cover's layout when the template has no
content sample; a tilted heading measured by its letters (the chart under it, the conclusion clear of it); a chart's
conclusion drops its strip before the plot is flattened; a table announced by a line ending in a colon stands under
that line; cards led by figures that no grid holds become rows; the audit takes a template's own heading face, a
conclusion said shorter, a heading on its layout's ribbon and a line over a blot under its descenders as they are."""

from __future__ import annotations

import math

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.compose import Composer, Kit
from verstka.rendering.deck import DeckBuilder, element_bbox
from verstka.schemas.common import Bbox, PatternKind
from verstka.schemas.data import SlideData
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, OutlineSlide, SlideContent, SlideItem, TableData

W43, H43 = 9144000, 6858000


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _composer(env, o: OutlineSlide, strategy: str = "structured", W: int | None = None, H: int | None = None) -> Composer:
    deck, manifest, _ = env
    W = W or manifest.slide_size.w
    H = H or manifest.slide_size.h
    prs = Presentation(str(deck))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return Composer(slide, Kit(manifest, W, H, "FFFFFF"), o, DeckOutline(title="T", strategy=strategy, slides=[o]), strategy, manifest)


def _named(comp: Composer, n0: int, base: str) -> list[tuple]:
    out = []
    for el in list(comp.cv.tree)[n0:]:
        nv = el.find(".//" + q("p:cNvPr"))
        if nv is not None and (nv.get("name") or "").rsplit(" ", 1)[0] == base:
            out.append((el, element_bbox(el)))
    return out


# ---- the canvas of a template whose only content-like sample reads as a divider ---------------------------------------


def _cover_and_empty_content(path, tmp_path):
    """A template like LibreOffice's: slide 1 is the cover (its title mid-slide over art), slide 2 an empty «title,
    content» page — no text on it, so the analysis reads it as a divider."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W43), Emu(H43)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Company presentation"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    for ph in s2.placeholders:
        ph.text_frame.text = ""
    prs.save(str(path))
    return path


def test_a_content_slide_takes_the_cleanest_layout_when_no_sample_is_a_content_page(tmp_path):
    from verstka.rendering.synth import _layout_for

    deck = _cover_and_empty_content(tmp_path / "t.pptx", tmp_path)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    # (the empty page read as a divider, as the analysis reads a LibreOffice template's)
    manifest = manifest.model_copy(update={"patterns": [p if p.source_slide == 1 else p.model_copy(update={"kind": PatternKind.section}) for p in manifest.patterns]})
    assert {p.kind for p in manifest.patterns} <= {PatternKind.title, PatternKind.section}
    b = DeckBuilder(deck)
    fam = manifest.patterns[0].family
    layout, _ = _layout_for(b, manifest, fam, PatternKind.chart)
    title = next(ph for ph in layout.placeholders if "TITLE" in str(ph.placeholder_format.type))
    assert int(title.top) <= 0.2 * H43, "a content slide's heading stands at the top, not mid-slide on the cover's layout"
    cover_layout = b.prs.slides[0].slide_layout
    assert layout.part.partname != cover_layout.part.partname


# ---- a tilted heading ---------------------------------------------------------------------------------------------------


def _tilted_title(slide, W: int, H: int, rot_deg: float = -8.0, anchor: str = "ctr"):
    sp = slide.shapes.add_textbox(Emu(int(0.03 * W)), Emu(int(0.12 * H)), Emu(int(0.8 * W)), Emu(int(0.17 * H)))
    sp.text_frame.text = "Выручка — на 23% больше, чем в сентябре 2025"
    for r in sp.text_frame.paragraphs[0].runs:
        r.font.size = Pt(26)
    sp.rotation = rot_deg % 360
    sp.text_frame._txBody.find(q("a:bodyPr")).set("anchor", anchor)
    return sp


def test_a_tilted_heading_ends_where_its_letters_end_not_its_box(env):
    from verstka.rendering.synth import _keep_turned_heading

    deck, manifest, _ = env
    prs = Presentation(str(deck))
    W, H = prs.slide_width, prs.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sp = _tilted_title(slide, W, H)
    box = (int(sp.left), int(sp.top), int(sp.width), int(sp.height))
    low = _keep_turned_heading(sp, manifest, W, H, [])
    # the turned box's own lowest corner (narrowed to its letters: the left end is where the box turns lowest)
    t = math.radians(-8.0)
    x, y, w, h = int(sp.left), int(sp.top), int(sp.width), int(sp.height)
    cx, cy = x + w / 2, y + h / 2
    box_low = max(cy + (px - cx) * math.sin(t) + (py - cy) * math.cos(t) for px in (x, x + w) for py in (y, y + h))
    assert low is not None and low < box_low - 0.02 * H, (low / H, box_low / H, box)


def test_the_room_under_a_tilted_heading_starts_under_its_turned_letters():
    from verstka.rendering.clone_notes import turned_box

    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sp = _tilted_title(slide, W, H)
    b = element_bbox(sp._element)
    letters = Bbox(x=b[0], y=b[1] + b[3] // 2 - int(0.03 * H), w=b[2], h=int(0.06 * H))
    turned = turned_box(sp._element, letters)
    # the left end of a heading tilted up to the right stands lower than its unturned letters
    assert turned.y2 > letters.y2 + 0.04 * H
    sp.rotation = 0
    assert turned_box(sp._element, letters) == letters


# ---- a chart's conclusion drops its strip before the plot is flattened -------------------------------------------------


def test_a_data_charts_conclusion_drops_its_strip_before_the_plot_is_flattened(env):
    months = [f"м{i}" for i in range(1, 22)]
    spec = ChartSpec(type="line", categories=months, unit="млн ₽", series=[InlineSeries(name="Выручка", values=[300 + 6 * i for i in range(21)])])
    take = "Сентябрь 2026 к сентябрю 2025: 423 млн ₽ против 344 млн ₽"
    o = OutlineSlide(id="d", kind=PatternKind.chart, headline="Выручка — на 23% больше", content=SlideContent(chart=spec), takeaway=take,
                     footnote="Источник: Выручка по месяцам.csv", data=SlideData(dataset_id="d", file="Выручка по месяцам.csv"))
    c = _composer(env, o)
    k = c.kit
    k.card.fill = "E8E8E8"  # the template's cards have a fill: the conclusion would take a strip
    c._comp = "chart_text"
    tall = Bbox(x=int(0.04 * k.W), y=int(0.2 * k.H), w=int(0.9 * k.W), h=int(0.7 * k.H))
    c._reserve_notes(tall)
    assert c._notes[1]["strip"], "with room enough, the conclusion keeps its strip"
    short = Bbox(x=int(0.04 * k.W), y=int(0.4 * k.H), w=int(0.9 * k.W), h=int(0.36 * k.H))
    body = c._reserve_notes(short)
    plan = c._notes[1]
    assert not plan["strip"], "a plot under its need: the strip's padding goes first"
    assert take.rstrip(".") in plan["para"].text.replace("\xa0", " "), "the same words — not said shorter"
    assert body.h > short.h - plan["h"] - int(k.vgap * 3) - c.h([c.P(o.footnote, c._footnote_size(), k.colors.muted)], short.w)


# ---- a table under the line that announces it ---------------------------------------------------------------------------


def test_a_line_ending_in_a_colon_stands_over_the_table_it_announces(env):
    t = TableData(columns=["Месяц", "Май", "Июнь", "Июль", "Август", "Сентябрь"], rows=[["Активные пользователи, чел.", "1200", "3400", "6100", "9800", "12400"]])
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Результаты пилота",
                     content=SlideContent(bullets=["В пилоте участвовали 12 400 сотрудников из 37 компаний.", "Динамика активных пользователей пилота:"], table=t))
    c = _composer(env, o)
    k = c.kit
    n0 = len(c.cv.tree)
    c.compose("table", Bbox(x=int(0.05 * k.W), y=int(0.22 * k.H), w=int(0.9 * k.W), h=int(0.65 * k.H)))
    intro = _named(c, n0, "Intro")
    frames = [b for el, b in ((el, element_bbox(el)) for el in list(c.cv.tree)[n0:]) if etree.QName(el).localname == "graphicFrame"]
    assert intro and frames
    text = "".join(x.text or "" for x in intro[0][0].iter(q("a:t")))
    assert "Динамика активных пользователей пилота:" in text and "12" in text
    assert intro[0][1][1] + intro[0][1][3] <= frames[0][1], "the announcing line stands over the table"


# ---- cards led by figures that no grid holds ----------------------------------------------------------------------------


def test_figure_cards_no_grid_holds_become_rows_with_whole_words(tmp_path):
    from tests.unit.test_tpl_compose import build_sparse_deck

    deck = build_sparse_deck(tmp_path / "sq.pptx", W43, H43, title_pt=40, body_pt=24)
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    items = [SlideItem(number="+23%", title="Выручка и клиенты", text="Выручка растёт быстрее числа клиентов: +23% против +18%"),
             SlideItem(number="41%", title="Выручка по регионам", text="Москва даёт 41% выручки, Казань растёт быстрее всех"),
             SlideItem(number="46%", title="Каналы продаж", text="Прямые продажи приносят почти половину выручки — 46%"),
             SlideItem(number="20%", title="Узкое место: время ответа поддержки", text="20% обращений — больше 120 мин, 3 — больше 10 часов")]
    o = OutlineSlide(id="s", kind=PatternKind.cards, headline="Итоги", content=SlideContent(items=items), footnote="Источник: Итоги квартала.txt")
    slide = DeckBuilder(deck).add_blank_slide(6)
    c = Composer(slide, Kit(manifest, W43, H43, "FFFFFF"), o, DeckOutline(title="T", slides=[o], strategy="structured"), "structured", manifest)
    # a column between art at both sides of a 4:3 slide: two thirds of its width, a third of its height
    area = Bbox(x=int(0.115 * W43), y=int(0.52 * H43), w=int(0.67 * W43), h=int(0.36 * H43))
    n0 = len(c.cv.tree)
    c.compose("cards", area)
    assert not _named(c, n0, "Card"), "no slivers of cards"
    assert len(_named(c, n0, "Figure")) == 4 and len(_named(c, n0, "Label")) == 4
    for el, b in [(el, element_bbox(el)) for el in list(c.cv.tree)[n0:]]:
        assert b[1] + b[3] <= area.y2 + int(0.005 * H43), "nothing runs past the area"
    for el, _ in _named(c, n0, "Label"):
        sizes = {int(r.get("sz")) / 100 for r in el.iter(q("a:rPr")) if r.get("sz")}
        assert min(sizes) >= 11.0 - 0.05


# ---- the audit ------------------------------------------------------------------------------------------------------------


def test_a_conclusion_said_shorter_with_the_rest_in_the_notes_is_on_the_slide():
    from verstka.audit.checks.integrity import _alnum, _said_shorter

    take = "Медиана — 54 минуты. 3 обращения — больше 10 часов: это выбросы, на графике они в столбце «больше 180»."
    have = "80% обращений получают ответ за 2 часа Медиана — 54 минуты Источник: Время ответа поддержки.csv"
    notes_ok = _alnum("Значения далеко от остальных. 3 обращения — больше 10 часов: это выбросы, на графике они в столбце «больше 180».")
    assert _said_shorter(take, have, _alnum(have), notes_ok)
    assert not _said_shorter(take, have, _alnum(have), _alnum("Половина значений меньше 54 мин.")), "the rest lost: an issue"
    assert not _said_shorter(take, "80% обращений", _alnum("80% обращений"), notes_ok), "the first sentence not on the slide: an issue"


def test_a_line_over_a_blot_under_its_descenders_is_not_on_it():
    from verstka.audit.checks.layout import letter_band

    row = Bbox(x=0, y=1000, w=5000, h=1200)
    band = letter_band(row)
    assert band.y > row.y and band.y2 < row.y2
    blot = Bbox(x=100, y=row.y2 - 150, w=400, h=900)  # starts in the line's descender room
    assert blot.intersection(band) == 0


def test_a_heading_inside_its_layouts_heading_box_is_home():
    from verstka.audit.checks.layout import _title_home

    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    home = _title_home(slide, W, H)
    title = slide.slide_layout.placeholders[0]
    assert home is not None and home.x < int(title.left) and home.x2 > int(title.left) + int(title.width)


def test_the_face_a_templates_own_heading_slot_is_set_in_is_the_templates(tmp_path):
    from types import SimpleNamespace

    from verstka.audit.checks.template import _allowed_fonts

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    # the sample's heading is an empty placeholder whose paragraph end names the face (no text of the template is set
    # in it, so the template's statistics never count it)
    p0 = s.shapes.title.text_frame.paragraphs[0]._p
    end = etree.SubElement(p0, q("a:endParaRPr"))
    end.set("lang", "ru-RU")
    end.set("sz", "4000")
    etree.SubElement(end, q("a:latin")).set("typeface", "Georgia")
    s.placeholders[1].text_frame.text = "Текст слайда"
    prs.save(str(tmp_path / "g.pptx"))
    manifest = analyze_template(tmp_path / "g.pptx", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    assert any(sl.style.font_family == "Georgia" for p in manifest.patterns for sl in p.slots)
    assert "georgia" not in {f.family.lower() for f in manifest.tokens.typography.families if f.weight > 0}
    ctx = SimpleNamespace(manifest=manifest, ir=SimpleNamespace(embedded_fonts=[], theme_fonts=[]), ws=None)
    assert "georgia" in _allowed_fonts(ctx)


def test_a_watermark_under_the_letters_is_a_mark_a_blot_under_the_line_is_not(tmp_path):
    from tests.unit.test_composer_qa4 import H as QH, W as QW, _ground_deck, _png
    from verstka.rendering.picture_lines import mark_share

    def paper(d):
        d.rectangle((0, 0, 1920, 1080), fill=(236, 220, 186, 255))
        d.ellipse((1000, 420, 1300, 720), outline=(214, 196, 160, 255), width=40)  # a pale watermark ring
        d.ellipse((120, 905, 190, 975), fill=(20, 20, 40, 255))  # an ink blot just under a line at the foot

    prs = _ground_deck(tmp_path / "paper.pptx", _png(paper))
    s = prs.slides[0]

    def box(x0, y0, x1, y1):
        return Bbox(x=int(x0 / 1920 * QW), y=int(y0 / 1080 * QH), w=int((x1 - x0) / 1920 * QW), h=int((y1 - y0) / 1080 * QH))

    over_ring = mark_share(s, box(960, 520, 1500, 620), box(960, 540, 1500, 600), sheet_share=0.55)
    over_blot = mark_share(s, box(100, 850, 700, 925), box(100, 862, 700, 900), sheet_share=0.55)
    assert over_ring is not None and over_ring >= 0.04
    assert over_blot is not None and over_blot < 0.04


def test_a_line_grazing_the_fine_row_of_a_blot_is_not_on_it():
    from verstka.rendering.layers import CELLS_H, CELLS_W, FINE_H, FINE_W, busy_for_line

    col, row = 1, 11
    std = [0.0] * (CELLS_W * CELLS_H)
    fine = [0.0] * (FINE_W * FINE_H)
    std[row * CELLS_W + col] = 74.0
    fine[(row * 2 + 1) * FINE_W + col * 2] = 91.0  # the blot in the cell's lower fine row
    top = (row * 2) / FINE_H
    assert not busy_for_line(std, fine, col, row, top + 0.004, top + 1 / FINE_H + 0.004), "the line only grazes the blot's row"
    assert busy_for_line(std, fine, col, row, top + 1 / FINE_H, top + 2 / FINE_H - 0.001), "a line on the blot's row stands on it"


def test_a_table_past_its_area_keeps_its_source_line_under_it(env):
    rows = [[f"Клиент {i}", "Промышленность", "Москва", f"{200 - i * 9},4", f"{6 - i * 0.3:.1f}".replace(".", ",")] for i in range(10)]
    t = TableData(columns=["Клиент", "Отрасль", "Регион", "Выручка, млн ₽", "Доля, %"], rows=rows)
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Топ-10 клиентов", content=SlideContent(table=t), footnote="Источник: Клиенты.xlsx")
    c = _composer(env, o)
    k = c.kit
    n0 = len(c.cv.tree)
    c.compose("table", Bbox(x=int(0.1 * k.W), y=int(0.5 * k.H), w=int(0.7 * k.W), h=int(0.18 * k.H)))
    frames = [element_bbox(el) for el in list(c.cv.tree)[n0:] if etree.QName(el).localname == "graphicFrame"]
    foot = _named(c, n0, "Footnote")
    assert frames and foot
    if frames[0][1] + frames[0][3] > int(0.68 * k.H):  # the table ran past its area
        assert foot[0][1][1] >= frames[0][1] + frames[0][3], "the source line is never printed over the table's last rows"


def test_a_category_label_cut_with_an_ellipsis_is_still_its_category():
    from verstka.audit.checks.data import _same_labels

    want = ["Прикладная информатика", "Электроэнергетика и электротехника", "Лингвистика"]
    assert _same_labels(want, ["Прикладная информатика", "Электроэнергетика и…", "Лингвистика"])
    assert not _same_labels(want, ["Лингвистика", "Электроэнергетика и…", "Прикладная информатика"])


def test_a_sample_whose_headline_runs_off_its_band_is_refused(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    from verstka.rendering.clone_check import heading_off_its_ground

    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    layout = prs.slide_layouts[5]  # title only
    tmp = prs.slides.add_slide(layout)
    band = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(int(0.03 * H)), Emu(int(0.45 * W)), Emu(int(0.2 * H)))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x0D, 0x84, 0xA1)
    band.line.fill.background()
    el = band._element
    el.getparent().remove(el)
    layout.shapes._spTree.insert(2, el)  # the band is the layout's, under the heading
    title = tmp.shapes.title
    title.left, title.top, title.width, title.height = Emu(int(0.03 * W)), Emu(int(0.05 * H)), Emu(int(0.4 * W)), Emu(int(0.16 * H))
    title.text_frame.text = "Итоги"
    for r in title.text_frame.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(28)
    prs.save(str(tmp_path / "band.pptx"))
    manifest = analyze_template(tmp_path / "band.pptx", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    pattern = next(p for p in manifest.patterns if any(s.role.value == "title" for s in p.slots))
    deck = Presentation(str(tmp_path / "band.pptx"))
    slide = deck.slides[pattern.source_slide - 1]
    t = slide.shapes.title
    assert heading_off_its_ground(slide, pattern, manifest) is None, "the sample's own short heading stands on its band"
    t.width = Emu(int(0.9 * W))
    t.text_frame.paragraphs[0].runs[0].text = "Топ-10 клиентов дают 38% выручки компании за девять месяцев"
    assert heading_off_its_ground(slide, pattern, manifest), "a white headline running past the band onto the white page"


def test_what_changed_most_beside_a_table_stays_inside_the_area(env):
    t = TableData(columns=["Показатель", "Первый набор", "Второй набор"],
                  rows=[["Студентов, поддержка и сопровождение преподавателей", "640", "1 100"], ["Дошли до конца программы обучения", "71%", "76%"],
                        ["Выпускников с подтверждённым трудоустройством", "38", "0"]])
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Результаты первого набора", content=SlideContent(table=t))
    c = _composer(env, o, strategy="visual")
    k = c.kit
    area = Bbox(x=int(0.06 * k.W), y=int(0.23 * k.H), w=int(0.88 * k.W), h=int(0.3 * k.H))
    n0 = len(c.cv.tree)
    c.compose("table", area)
    from verstka.rendering.fonts import wrap_lines

    for el, b in _named(c, n0, "Takeaway"):
        assert b[1] + b[3] <= area.y2 + 1, "the column of changes never runs past the area"
        need = 0.0
        for para in el.iter(q("a:p")):
            text = "".join(t.text or "" for t in para.iter(q("a:t")))
            rpr = next((r for r in para.iter(q("a:rPr")) if r.get("sz")), None)
            if not text.strip() or rpr is None:
                continue
            size = int(rpr.get("sz")) / 100
            need += len(wrap_lines(text, k.font, size, rpr.get("b") == "1", b[2] / 12700)) * size * 1.15
        assert need * 12700 <= b[3] * 1.02, "the column's lines fit its box (not cut to the area's height)"
