"""Owner C of TEMPLATE_FIX_PLAN («clone+audit»): deck helpers (C1), template layers (C2), textfill (C5, T11, T14),
and the audit's honesty (T08, T17): layout/master grounds and art are seen, and real defects still fire.

Fixture decks in tests/fixtures/tpl_audit are cut from the baseline replay of the third-party corpus (one or two
slides each, pictures shrunk) with the manifests they were rendered with (gzipped)."""

from __future__ import annotations

import copy
import gzip
import io
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from verstka.analysis.xmlns import q

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "tpl_audit"


# ------------------------------------------------------------------------------------------------ helpers


def _blank_template(tmp_path: Path) -> Path:
    path = tmp_path / "tpl.pptx"
    Presentation().save(str(path))
    return path


def _xfrm(el) -> tuple[int, int, int, int] | None:
    x = el.find(q("p:spPr") + "/" + q("a:xfrm"))
    if x is None or x.find(q("a:off")) is None or x.find(q("a:ext")) is None:
        return None
    off, ext = x.find(q("a:off")), x.find(q("a:ext"))
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def _layout_ph(layout, ph_type: str):
    for shp in layout.placeholders:
        ph = shp._element.find(".//" + q("p:ph"))
        if (ph.get("type") or "body") == ph_type:
            return shp
    raise AssertionError(ph_type)


def _audit(deck: str, manifest: str):
    from verstka.audit.runner import run_audit
    from verstka.schemas.template import TemplateManifest

    man = TemplateManifest.model_validate_json(gzip.open(FIX / manifest, "rt", encoding="utf-8").read())
    return run_audit(FIX / deck, man, None, render=False, use_llm=False, use_vlm=False)


def _serious(report, check_id: str | None = None):
    return [i for i in report.issues if i.severity in ("error", "warn") and (check_id is None or i.check_id == check_id)]


# ------------------------------------------------------------------------------------------------ C1: deck helpers


def test_add_layout_slide_writes_inherited_geometry(tmp_path):
    from verstka.rendering.deck import DeckBuilder

    b = DeckBuilder(_blank_template(tmp_path))
    layout = b.prs.slide_layouts[1]  # title and content: its placeholders inherit from the master
    ref = Presentation(str(tmp_path / "tpl.pptx")).slides.add_slide(Presentation(str(tmp_path / "tpl.pptx")).slide_layouts[1])
    want = (ref.shapes.title.left, ref.shapes.title.top, ref.shapes.title.width, ref.shapes.title.height)
    assert _xfrm(ref.shapes.title._element) is None  # python-pptx leaves the geometry inherited
    slide = b.add_layout_slide(layout)
    assert slide in b.created
    assert _xfrm(slide.shapes.title._element) == want
    assert all(_xfrm(s._element) is not None for s in slide.placeholders)


def test_materialize_xfrm_completes_a_partial_write(tmp_path):
    from verstka.rendering.deck import materialize_xfrm

    prs = Presentation(str(_blank_template(tmp_path)))
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    lx, ly, lw, lh = title.left, title.top, title.width, title.height
    title.width = Emu(3_000_000)  # python-pptx writes a:ext cx=… cy=0 and no a:off
    assert _xfrm(title._element) is None
    assert materialize_xfrm(title) is True
    assert _xfrm(title._element) == (lx, ly, 3_000_000, lh)
    assert materialize_xfrm(title) is False  # complete now: a no-op


def test_materialize_placeholders_counts_and_keeps_explicit_boxes(tmp_path):
    from verstka.rendering.deck import materialize_placeholders

    prs = Presentation(str(_blank_template(tmp_path)))
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.left, slide.shapes.title.top, slide.shapes.title.width, slide.shapes.title.height = 10, 20, 30, 40
    assert materialize_placeholders(slide) == 1  # the body only: the title already has its own box
    assert _xfrm(slide.shapes.title._element) == (10, 20, 30, 40)


# ------------------------------------------------------------------------------------------------ C2: layers


def _layout_shape(prs, layout, kind, x, y, w, h, rgb) -> None:
    """Add a filled shape to a layout, under its placeholders (python-pptx only adds shapes to slides: the shape is
    drawn on a scratch slide and moved)."""
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    shp = scratch.shapes.add_shape(kind, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(*rgb)
    el = shp._element
    el.getparent().remove(el)
    layout.shapes._spTree.insert(2, el)
    lst = prs.slides._sldIdLst
    sid = lst[-1]
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)


def _banded(tmp_path: Path, band_rgb=(0x1F, 0x2A, 0x44)):
    """A slide on a layout that paints a band under the title and an ellipse (art) on the right."""
    prs = Presentation(str(_blank_template(tmp_path)))
    W, H = int(prs.slide_width), int(prs.slide_height)
    layout = prs.slide_layouts[5]  # title only
    _layout_shape(prs, layout, MSO_SHAPE.OVAL, int(0.7 * W), int(0.35 * H), int(0.25 * W), int(0.5 * H), (0xE0, 0x40, 0x40))
    _layout_shape(prs, layout, MSO_SHAPE.RECTANGLE, 0, 0, W, int(0.2 * H), band_rgb)
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.left, slide.shapes.title.top = int(0.05 * W), int(0.04 * H)
    slide.shapes.title.width, slide.shapes.title.height = int(0.8 * W), int(0.12 * H)
    return prs, slide, W, H


def test_drawn_layers_heading_band_and_free_rect(tmp_path):
    from verstka.rendering.layers import art_boxes, drawn_layers, free_rect, heading_band
    from verstka.schemas.common import Bbox

    prs, slide, W, H = _banded(tmp_path)
    layers = drawn_layers(slide)
    assert [l.source for l in layers if not l.placeholder][:2] == ["layout", "layout"]
    tb = slide.shapes.title
    band = heading_band(slide, Bbox(x=tb.left, y=tb.top, w=tb.width, h=tb.height))
    assert band is not None and band.source == "layout" and band.box.h == int(0.2 * H)
    arts = art_boxes(slide)
    assert any(b.x >= int(0.7 * W) - 1 for b in arts)  # the ellipse is art
    area = Bbox(x=int(0.05 * W), y=int(0.25 * H), w=int(0.9 * W), h=int(0.7 * H))
    free = free_rect(slide, area)
    assert free.x2 <= int(0.7 * W) and free.w >= int(0.55 * W)  # content keeps off the ellipse
    # no art in the way: the area itself
    left = Bbox(x=int(0.05 * W), y=int(0.25 * H), w=int(0.5 * W), h=int(0.7 * H))
    assert free_rect(slide, left) == left


def test_is_ground_full_bleed_calm_layer(tmp_path):
    from verstka.rendering.layers import art_boxes, drawn_layers, is_ground

    prs = Presentation(str(_blank_template(tmp_path)))
    W, H = int(prs.slide_width), int(prs.slide_height)
    layout = prs.slide_layouts[5]
    _layout_shape(prs, layout, MSO_SHAPE.RECTANGLE, 0, 0, W, H, (0x10, 0x20, 0x30))
    slide = prs.slides.add_slide(layout)
    lay = next(l for l in drawn_layers(slide) if l.source == "layout")
    assert lay.full_bleed and is_ground(lay)
    assert art_boxes(slide) == []


def test_gradient_position_linear_and_radial():
    from verstka.rendering.layers import gradient_position

    lin = etree.fromstring('<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:lin ang="0" scaled="1"/></a:gradFill>')
    assert gradient_position(lin, 0.0, 0.5, 100, 50) == pytest.approx(0.0)
    assert gradient_position(lin, 1.0, 0.5, 100, 50) == pytest.approx(1.0)
    down = etree.fromstring('<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:lin ang="5400000" scaled="1"/></a:gradFill>')
    assert gradient_position(down, 0.3, 0.25, 100, 50) == pytest.approx(0.25)
    # a radial gradient from the bottom centre (LibreOffice Sunset): 0 at the focus, 1 at the top edge
    rad = etree.fromstring('<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:path path="circle"><a:fillToRect l="50000" t="100000" r="50000" b="0"/></a:path></a:gradFill>')
    assert gradient_position(rad, 0.5, 1.0, 100, 50) == pytest.approx(0.0)
    assert gradient_position(rad, 0.5, 0.0, 100, 50) == pytest.approx(1.0)
    assert gradient_position(rad, 0.5, 0.5, 100, 50) == pytest.approx(0.5)


def test_ground_under_reads_the_layout_band_unless_a_slide_picture_hides_it(tmp_path):
    from PIL import Image

    from verstka.rendering.layers import ground_under
    from verstka.schemas.common import Bbox

    prs, slide, W, H = _banded(tmp_path)
    box = Bbox(x=int(0.1 * W), y=int(0.05 * H), w=int(0.3 * W), h=int(0.1 * H))
    got = ground_under(slide, box)
    assert got is not None and got[0] == "1F2A44" and got[1].source == "layout"
    assert ground_under(slide, Bbox(x=int(0.1 * W), y=int(0.5 * H), w=int(0.3 * W), h=int(0.1 * H))) is None
    buf = io.BytesIO()
    Image.new("RGB", (32, 16), (250, 250, 250)).save(buf, "PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, 0, 0, W, int(0.3 * H))
    assert ground_under(slide, box) is None


# ------------------------------------------------------------------------------------------------ C5: textfill


def _text_sp(ppr: str = "", body: str = "") -> etree._Element:
    return etree.fromstring(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="2" name="t"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/>'
        f'<p:txBody><a:bodyPr {body}/><a:lstStyle/><a:p>{ppr}<a:r><a:rPr lang="ru-RU" sz="2000"/><a:t>old</a:t></a:r></a:p></p:txBody></p:sp>'
    )


def test_fill_text_reset_indent_and_neutral_runs():
    from verstka.rendering.textfill import ParagraphSpec, fill_text

    el = _text_sp('<a:pPr marL="1270000" indent="-300000"/>')
    assert fill_text(el, [ParagraphSpec("Заголовок", bullet=False)], size_pt=24, reset_indent=True, neutral_runs=True)
    ppr = el.find(".//" + q("a:pPr"))
    assert ppr.get("marL") == "0" and ppr.get("indent") == "0"
    rpr = el.find(".//" + q("a:rPr"))
    assert (rpr.get("cap"), rpr.get("spc"), rpr.get("baseline")) == ("none", "0", "0")
    # defaults keep today's behaviour
    el2 = _text_sp('<a:pPr marL="1270000" indent="-300000"/>')
    fill_text(el2, [ParagraphSpec("Заголовок", bullet=False)], size_pt=24)
    assert el2.find(".//" + q("a:pPr")).get("marL") == "1270000"
    assert el2.find(".//" + q("a:rPr")).get("spc") is None


def test_effective_insets_and_inherited_caps_spc_walk_the_placeholder_chain(tmp_path):
    from verstka.rendering.textfill import effective_insets, inherited_caps_spc

    prs = Presentation(str(_blank_template(tmp_path)))
    layout = prs.slide_layouts[5]
    ltitle = _layout_ph(layout, "title")._element
    body_pr = ltitle.find(q("p:txBody") + "/" + q("a:bodyPr"))
    body_pr.set("lIns", "194000")
    body_pr.set("tIns", "12700")
    master = layout.slide_master._element
    lvl1 = master.find(q("p:txStyles") + "/" + q("p:titleStyle") + "/" + q("a:lvl1pPr"))
    defrpr = lvl1.find(q("a:defRPr"))
    defrpr.set("cap", "all")
    defrpr.set("spc", "500")
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    assert effective_insets(title) == (194000, 12700, 91440, 45720)
    assert inherited_caps_spc(title) == (True, 5.0)
    # the slide's own run wins over the chain
    title.text = "x"
    title.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("spc", "0")
    assert inherited_caps_spc(title) == (True, 0.0)


# ------------------------------------------------------------------------------------------------ T08: audit helpers


def test_same_family_and_template_chrome():
    from verstka.audit.checks.common import is_template_chrome
    from verstka.audit.checks.template import same_family
    from verstka.schemas.common import Bbox, BboxFrac
    from verstka.schemas.deck_ir import IRElement

    assert same_family("Bebas", "Bebas Neue") and same_family("Montserrat-Regular", "montserrat")
    assert not same_family("Arial", "Arimo") and not same_family("PT", "PT Sans")
    num = IRElement(id="4", type="text", bbox=Bbox(x=0, y=0, w=10, h=10), bbox_frac=BboxFrac(x=0.9, y=0.9, w=0.05, h=0.05), ph_type="sldNum")
    assert is_template_chrome(num)
    assert not is_template_chrome(num.model_copy(update={"ph_type": "body"}))


def test_region_paint_reads_a_gradient_where_the_text_is():
    from verstka.audit.checks.common import region_paint
    from verstka.audit.ir import gradient_cells
    from verstka.schemas.common import Bbox, BboxFrac
    from verstka.schemas.deck_ir import IRElement

    grad = etree.fromstring(
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:gsLst>'
        '<a:gs pos="0"><a:srgbClr val="FFFF00"/></a:gs><a:gs pos="100000"><a:srgbClr val="FF8000"/></a:gs></a:gsLst>'
        '<a:path path="circle"><a:fillToRect l="50000" t="100000" r="50000" b="0"/></a:path></a:gradFill>'
    )

    class _Res:
        def resolve_color(self, el):
            return el.get("val")

    cells = gradient_cells(_Res(), grad, 1000, 600)
    panel = IRElement(id="master:5", type="shape", bbox=Bbox(x=0, y=0, w=1000, h=600), bbox_frac=BboxFrac(x=0, y=0, w=1, h=1), fill_hex="FFC000", source="master", paint_kind="gradient", cells=cells)
    top_hex, top_a = region_paint(panel, Bbox(x=300, y=10, w=400, h=60))
    low_hex, _ = region_paint(panel, Bbox(x=400, y=540, w=200, h=50))
    assert top_a == pytest.approx(1.0)
    assert top_hex.startswith("FF") and int(top_hex[2:4], 16) < 0x95  # orange at the top
    assert int(low_hex[2:4], 16) > 0xE0  # yellow at the focus


# ------------------------------------------------------------------------------------------------ T17: honesty


def test_audit_still_flags_a_heading_line_spilling_out_of_its_band():
    """Midnightblue long structured slide 3: a two-line heading in a one-line navy band — the second line white on
    white. Its box lies mostly inside the band, so the band is its ground; its lines are not."""
    r = _audit("midnight_long_s3.pptx", "manifest_3ddedfc06fd7ec9f.json.gz")
    spill = [i for i in _serious(r, "contrast_low") if (i.details or {}).get("ground") == "band_spill"]
    assert spill and spill[0].severity == "error"


def test_audit_still_flags_content_over_template_art_and_off_slide():
    """Candy long visual slide 3: legend texts over the layout's art, a chart past the slide's edge."""
    r = _audit("candy_long_visual_s3.pptx", "manifest_fec48b04f7ff51ea.json.gz")
    assert any(i.severity == "error" for i in _serious(r, "content_over_art"))
    assert _serious(r, "out_of_bounds")


def test_audit_flags_content_on_the_focus_triangles():
    r = _audit("focus_long_s4_s6.pptx", "manifest_368fc0232d0871b3.json.gz")
    over = _serious(r, "content_over_art")
    assert {i.slide for i in over} == {1, 2}


def test_audit_flags_a_heading_at_the_slide_edge_through_the_frame():
    """Sunset short slide 4: the heading written at x=0, over the frame and across the layout's rule."""
    r = _audit("sunset_short_s4.pptx", "manifest_7d37bee880693301.json.gz")
    assert _serious(r, "margin_violation")
    assert any((i.details or {}).get("rule") for i in _serious(r, "content_over_art"))
    # white labels on the light middle of the gradient panel are unreadable there (measured where they stand)
    assert any(i.severity == "error" and (i.details or {}).get("ground") == "gradient" for i in _serious(r, "contrast_low"))


def test_audit_flags_a_two_line_heading_out_of_the_blue_curve_band():
    r = _audit("bluecurve_long_s2.pptx", "manifest_bb8044c91669bb1f.json.gz")
    assert any(i.severity == "error" for i in _serious(r, "contrast_low"))


def test_audit_leaves_the_template_own_layers_alone():
    """Must NOT flag: VK Tech short structured (every slide), the Blue_Curve heading inside its band, the MyBrand
    cover title on its layout's gradient picture."""
    vk = _audit("vktech_short.pptx", "manifest_cbbe3aa6a21d23ce.json.gz")
    assert _serious(vk) == [] and vk.summary.score == 100
    bc = _audit("bluecurve_inband_s2_s3.pptx", "manifest_bluecurve_inband.json.gz")
    assert _serious(bc, "contrast_low") == [] and _serious(bc, "content_over_art") == []
    mb = _audit("mybrand_short_s1.pptx", "manifest_a4d0f0fa139a4368.json.gz")
    assert _serious(mb, "contrast_low") == [] and _serious(mb, "fill_ratio") == []


# ------------------------------------------------------------------------------------------------ T16 / T15: autofix


def _deck_with_card(tmp_path: Path):
    prs = Presentation(str(_blank_template(tmp_path)))
    W, H = int(prs.slide_width), int(prs.slide_height)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(0.6 * W), int(0.8 * H), int(0.3 * W), int(0.3 * H))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    card.name = "Card 7"
    txt = slide.shapes.add_textbox(int(0.62 * W), int(0.85 * H), int(0.26 * W), int(0.1 * H))
    txt.text_frame.text = "Текст карточки"
    txt.name = "Card title 8"
    return prs, slide, card, txt, W, H


def test_move_block_moves_a_card_with_its_text(tmp_path):
    from verstka.audit.autofix import _move_block
    from verstka.schemas.common import Bbox

    prs, slide, card, txt, W, H = _deck_with_card(tmp_path)
    safe = Bbox(x=int(0.05 * W), y=int(0.05 * H), w=int(0.9 * W), h=int(0.9 * H))
    plan = _move_block(txt._element, slide, safe, W, H, into_safe=False)
    assert plan is not None and len(plan) == 2
    moved = {id(el): box for el, box in plan}
    dy_card = moved[id(card._element)][1] - card.top
    dy_text = moved[id(txt._element)][1] - txt.top
    assert dy_card == dy_text < 0 and moved[id(card._element)][1] + card.height <= H


def test_move_block_refuses_a_move_that_creates_an_overlap(tmp_path):
    from verstka.audit.autofix import _move_block
    from verstka.schemas.common import Bbox

    prs, slide, card, txt, W, H = _deck_with_card(tmp_path)
    other = slide.shapes.add_textbox(int(0.6 * W), int(0.62 * H), int(0.3 * W), int(0.12 * H))
    other.text_frame.text = "Другой блок"
    safe = Bbox(x=int(0.05 * W), y=int(0.05 * H), w=int(0.9 * W), h=int(0.9 * H))
    assert _move_block(txt._element, slide, safe, W, H, into_safe=False) is None


def test_rematch_never_gives_a_cover_an_empty_photo_frame():
    from verstka.audit.autofix import _rematch
    from verstka.schemas.common import BboxFrac, PatternKind, SlotRole
    from verstka.schemas.layout import LayoutPlan
    from verstka.schemas.outline import DeckOutline
    from verstka.schemas.template import TemplateManifest

    fx = Path(__file__).resolve().parents[1] / "fixtures"
    outline = DeckOutline.model_validate_json((fx / "outline_demo.json").read_text())
    cover = next(s for s in outline.slides if s.kind == PatternKind.title)
    cover.content.image_hint = None
    man = TemplateManifest.model_validate_json(gzip.open(FIX / "manifest_a4d0f0fa139a4368.json.gz", "rt", encoding="utf-8").read())
    covers = [p for p in man.patterns if p.kind == PatternKind.title]
    assert covers
    mock = copy.deepcopy(covers[0])
    mock.id, mock.mockup_boxes = "pmock", [BboxFrac(x=0.6, y=0.1, w=0.2, h=0.8)]
    plain = copy.deepcopy(covers[0])
    plain.id, plain.mockup_boxes = "pplain", []
    plain.slots = [s for s in plain.slots if s.role != SlotRole.image]
    man.patterns += [mock, plain]
    plan = LayoutPlan.model_validate({"template_id": man.template_id, "strategy": "structured", "slides": [{"outline_id": cover.id, "mode": "clone", "pattern_id": "pcurrent", "alternatives": [["pmock", 0.9], ["pplain", 0.5]]}]})
    assert _rematch(plan, outline, cover.id, man) == "rematch → pplain"


def test_shades_of_palette_colours_are_the_template_own():
    from verstka.audit.checks.template import is_shade_of

    assert is_shade_of("16A085", ["2C3E50", "1ABC9C"])  # «accent 1, darker 25 %»
    assert is_shade_of("BA8202", ["E8A202"])
    assert not is_shade_of("47798F", ["77CAEE", "009BDD"])  # a greyed blue is not a theme variant
    assert not is_shade_of("0077FF", ["000000", "CCCCCC"])
    assert not is_shade_of("777777", ["000000", "FFFFFF"])  # greys are not shades of black


def test_text_on_a_template_triangle_stands_on_the_triangle():
    from verstka.audit.checks.common import ground_of
    from verstka.schemas.common import Bbox, BboxFrac, Family
    from verstka.schemas.deck_ir import IRElement, IRParagraph, IRRun, IRSlide

    W = H = 1000
    tri = IRElement(id="layout:3", type="shape", bbox=Bbox(x=0, y=0, w=1000, h=1000), bbox_frac=BboxFrac(x=0, y=0, w=1, h=1),
                    fill_hex="F10D0C", geometry="triangle", source="layout", paint_kind="solid",
                    outline=[[(0, 0), (1000, 0), (0, 1000)]])
    txt = lambda x, y: IRElement(id="9", type="text", bbox=Bbox(x=x, y=y, w=200, h=100), bbox_frac=BboxFrac(x=x / W, y=y / H, w=0.2, h=0.1), z=5,
                                 paragraphs=[IRParagraph(text="Заголовок", runs=[IRRun(text="Заголовок", size_pt=32, color_hex="FFFFFF")])])
    slide = IRSlide(index=1, family=Family.light, background_hex="FFFFFF", template_elements=[tri])
    on = txt(100, 100)  # wholly inside the triangle's upper-left half
    slide.elements = [on]
    assert ground_of(slide, on, "FFFFFF") == ("F10D0C", "template")
    off = txt(600, 600)  # inside the triangle's box, outside what it paints
    slide.elements = [off]
    assert ground_of(slide, off, "FFFFFF") == (None, None)


def test_snap_up_lands_on_the_scale():
    from verstka.rendering.clone import _snap_up

    assert _snap_up(14.58, [12.5, 17.5, 20.0]) == 17.5
    assert _snap_up(17.4, [12.5, 17.5, 20.0]) == 17.5
    assert _snap_up(30.0, [12.5, 17.5, 20.0]) == 30.0


def test_cover_lines_calm_ground_and_speaker_notes(tmp_path):
    """T15 (clone covers): a line box that a layout triangle only partly paints is not calm (the tip of a triangle in
    it already counts); a box wholly on the triangle or wholly off it is; a line with no room goes to the speaker notes
    once."""
    from types import SimpleNamespace

    from verstka.rendering.clone import _calm_at, _note_to_speaker_notes, _straddles
    from verstka.schemas.common import Bbox

    prs = Presentation(str(_blank_template(tmp_path)))
    W, H = int(prs.slide_width), int(prs.slide_height)
    layout = prs.slide_layouts[6]
    _layout_shape(prs, layout, MSO_SHAPE.ISOSCELES_TRIANGLE, 0, 0, int(0.5 * W), int(0.6 * H), (0xF1, 0x0D, 0x0C))
    slide = prs.slides.add_slide(layout)
    ctx = SimpleNamespace(slide=slide)
    on = Bbox(x=int(0.2 * W), y=int(0.4 * H), w=int(0.1 * W), h=int(0.1 * H))  # inside the triangle near its base
    off = Bbox(x=int(0.6 * W), y=int(0.1 * H), w=int(0.3 * W), h=int(0.1 * H))  # white ground
    edge = Bbox(x=int(0.3 * W), y=int(0.25 * H), w=int(0.25 * W), h=int(0.1 * H))  # across the right edge
    assert _calm_at(ctx, on) and _calm_at(ctx, off)
    assert not _calm_at(ctx, edge)
    box = slide.shapes.add_textbox(edge.x, edge.y, edge.w, edge.h)
    box.text_frame.text = "Подзаголовок"
    assert _straddles(ctx, box._element)
    calm = slide.shapes.add_textbox(off.x, off.y, off.w, off.h)
    calm.text_frame.text = "Дата"
    assert not _straddles(ctx, calm._element)
    osl = SimpleNamespace(notes="Главная цель")
    _note_to_speaker_notes(osl, "Все цифры  условные")
    _note_to_speaker_notes(osl, "Все цифры условные")
    assert osl.notes == "Главная цель\nВсе цифры условные"


def test_ir_measures_a_heading_its_layout_sets_in_capitals(tmp_path):
    """A title whose layout placeholder sets cap="all" shows in capitals: the IR says so and the audit measures the
    upper-cased text (a heading measured in lower case fitted one line, rendered on two, over the subtitle)."""
    from verstka.audit.checks.common import text_height_needed_pt
    from verstka.audit.ir import build_deck_ir

    prs = Presentation(str(_blank_template(tmp_path)))
    layout = prs.slide_layouts[5]  # title only
    ph = _layout_ph(layout, "title")
    txb = ph._element.find(q("p:txBody"))
    lst = txb.find(q("a:lstStyle"))
    if lst is None:
        lst = etree.SubElement(txb, q("a:lstStyle"))
        txb.remove(lst)
        txb.insert(1, lst)
    lvl = etree.SubElement(lst, q("a:lvl1pPr"))
    etree.SubElement(lvl, q("a:defRPr")).set("cap", "all")
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Больше прибыли с каждой чашки"
    plain = prs.slides.add_slide(prs.slide_layouts[6])
    tb = plain.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height // 4)
    tb.text_frame.text = "Больше прибыли с каждой чашки"
    path = tmp_path / "caps.pptx"
    prs.save(str(path))
    ir = build_deck_ir(path, with_images=False)
    title = next(e for e in ir.slides[0].elements if e.type == "text")
    assert title.caps is True
    assert next(e for e in ir.slides[1].elements if e.type == "text").caps is False
    # the same heading, measured in a box where its lower-case letters just fit one line
    from verstka.rendering.fonts import text_width_pt

    lower_w = text_width_pt(title.text, None, 40.0)
    for p in title.paragraphs:
        for r in p.runs:
            r.size_pt, r.font = 40.0, None
    title.insets_emu = (0, 0, 0, 0)
    title.bbox = title.bbox.model_copy(update={"w": int(lower_w * 1.03 * 12700)})
    _, n_caps = text_height_needed_pt(title)
    title.caps = False
    _, n_lower = text_height_needed_pt(title)
    assert n_lower == 1 and n_caps == 2


# ------------------------------------------------------------------------------------------------ gate round 1 (C)

_GOAL = "Цель: увеличить ежемесячную операционную прибыль со 120 000 до 255 000 рублей"


def _cover_ctx(tmp_path: Path):
    """A 13.33″ cover: heading, subtitle and a goal box wrapped on two lines in a 0.64 W column, a logo at the foot."""
    from types import SimpleNamespace

    from verstka.schemas.common import BboxFrac

    prs = Presentation(str(_blank_template(tmp_path)))
    W, H = 12192000, 6858000
    prs.slide_width, prs.slide_height = W, H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def box(x, y, w, h, text, size):
        tb = slide.shapes.add_textbox(int(x * W), int(y * H), int(w * W), int(h * H))
        tb.text_frame.text = text
        from pptx.util import Pt

        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        return tb._element

    t_el = box(0.035, 0.363, 0.642, 0.231, "Больше прибыли с каждой чашки", 54)
    s_el = box(0.035, 0.639, 0.642, 0.054, "План развития кофейни «Точка кофе» на 6 месяцев", 20)
    g_el = box(0.035, 0.707, 0.642, 0.079, _GOAL, 16)
    safe = BboxFrac(x=0.035, y=0.06, w=0.93, h=0.88)
    ctx = SimpleNamespace(
        W=W, H=H, slide=slide, warnings=[], filled=set(), id_of=lambda el: None,
        manifest=SimpleNamespace(tokens=SimpleNamespace(spacing=SimpleNamespace(safe_area=safe))),
        ws=SimpleNamespace(slide_image=lambda n: tmp_path / "missing.png"), pattern=SimpleNamespace(source_slide=1),
    )
    return ctx, t_el, s_el, g_el


def test_cover_small_print_stays_on_the_slide_under_a_one_line_goal(tmp_path):
    """G1-01: no foot room for the brief's small print («Все исходные данные и прогнозы условные» on the first slide):
    the goal that wrapped needlessly takes one line in the width free at its own height and the note goes under it;
    where something stands beside the goal, the goal is said aloud and the note takes its place. Never the note first."""
    from types import SimpleNamespace

    from verstka.rendering.clone import _note_under_goal
    from verstka.schemas.common import Bbox

    ctx, t_el, s_el, g_el = _cover_ctx(tmp_path)
    W, H = ctx.W, ctx.H
    logo = Bbox(x=int(0.033 * W), y=int(0.88 * H), w=int(0.22 * W), h=int(0.06 * H))
    note = (int(0.035 * W), int(0.138 * W), int(0.059 * H), 12.0)
    osl = SimpleNamespace(notes="")
    top, kept = _note_under_goal(ctx, osl, g_el, _GOAL, 16.0, "Play", "l", t_el, s_el, [], [logo], note, logo.y)
    gb = _xfrm(kept)
    assert kept is g_el and top is not None
    assert gb[3] < 0.8 * int(0.079 * H)  # one line (it had two)
    assert gb[1] + gb[3] <= top and top + note[2] <= logo.y  # the note under the goal, clear of the logo
    assert osl.notes == ""  # nothing said aloud

    # a picture beside the goal: no one-line goal; the goal leaves the stack, the note stands under the subtitle. The
    # goal waits in ctx.goal_to_place: the cover puts it on a calm panel of the layout once its lines stand, else it
    # is said aloud (round 4.1, C3-1)
    ctx, t_el, s_el, g_el = _cover_ctx(tmp_path)
    pic = Bbox(x=int(0.5 * W), y=int(0.70 * H), w=int(0.45 * W), h=int(0.1 * H))
    top, kept = _note_under_goal(ctx, osl, g_el, _GOAL, 16.0, "Play", "l", t_el, s_el, [], [logo, pic], note, logo.y)
    sb = _xfrm(s_el)
    assert kept is None and top is not None and top >= sb[1] + sb[3]
    assert top + note[2] <= logo.y
    assert "255 000" in (ctx.goal_to_place or "") and g_el.getparent() is None and osl.notes == ""

    # no room even there: nothing moves, the caller sends the note to the notes
    ctx, t_el, s_el, g_el = _cover_ctx(tmp_path)
    low = Bbox(x=0, y=int(0.70 * H), w=W, h=int(0.3 * H))
    top, kept = _note_under_goal(ctx, SimpleNamespace(notes=""), g_el, _GOAL, 16.0, "Play", "l", t_el, s_el, [], [low], note, low.y)
    assert top is None and kept is g_el and g_el.getparent() is not None


def _dataset_template(stem: str) -> Path | None:
    import os

    env = os.environ.get("VERSTKA_FIXTURES_DIR")
    for c in ([Path(env)] if env else []) + [Path(__file__).resolve().parents[3] / "Датасет"]:
        hit = next(iter(sorted(c.glob(f"{stem}*.pptx"))), None) if c.is_dir() else None
        if hit is not None:
            return hit
    return None


def test_workspace_long_cover_keeps_the_brief_small_print_on_the_slide(tmp_path):
    """G1-01 on the dataset: the long coffee cover on VK WorkSpace carries «Все исходные данные и прогнозы условные» on
    slide 1 (the brief asks for it in small print on the first slide), clear of the goal, the logo and the art — not in
    the speaker notes. Skipped when the VK templates (Датасет/) are not on this machine."""
    pptx = _dataset_template("VK_WorkSpace")
    if pptx is None:
        pytest.skip("VK WorkSpace template not available")
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    cover = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Больше прибыли с каждой чашки", subtitle="План развития кофейни “Точка кофе” на 6 месяцев",
                         content=SlideContent(paragraphs=[_GOAL]), footnote="Все исходные данные и прогнозы условные",
                         notes="Главная цель: увеличить ежемесячную операционную прибыль со 120 000 до 255 000 рублей.")
    outline = DeckOutline(title="Больше прибыли с каждой чашки", slides=[cover], language="ru")
    generate_variants(pptx, outline=outline, strategies=["structured"], out_dir=tmp_path / "run", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, audit=False, exports=[])
    prs = Presentation(str(tmp_path / "run" / "structured" / "deck.pptx"))
    s = prs.slides[0]
    texts = {sh.name: sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()}
    note = next((sh for sh in texts.values() if "условные" in sh.text_frame.text), None)
    assert note is not None, [sh.text_frame.text for sh in texts.values()]
    assert "условные" not in (s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    nb = (note.left, note.top, note.left + note.width, note.top + note.height)
    assert nb[3] <= prs.slide_height
    for sh in s.shapes:
        if sh.shape_id == note.shape_id or not (sh.width and sh.height):
            continue
        ob = (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
        inter_w = min(nb[2], ob[2]) - max(nb[0], ob[0])
        inter_h = min(nb[3], ob[3]) - max(nb[1], ob[1])
        assert not (inter_w > 0 and inter_h > 0), f"the note overlaps «{sh.name}»"


def _panels_deck(tmp_path: Path, text_top: float, second=(0.07, 0.37, 0.77, 0.53)):
    """A cover of two overlapping panels (the Marketing cover: a grey panel over the top, a dark one from the middle
    down, neither holding the other) and a white heading box starting at `text_top`."""
    from pptx.util import Pt

    prs = Presentation(str(_blank_template(tmp_path)))
    W, H = int(prs.slide_width), int(prs.slide_height)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for (x, y, w, h), rgb in (((0.06, 0.0, 0.94, 0.86), (0x7C, 0x7D, 0x87)), (second, (0x45, 0x47, 0x54))):
        sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x * W), int(y * H), int(w * W), int(h * H))
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor(*rgb)
        sp.line.fill.background()
    tb = s.shapes.add_textbox(int(0.19 * W), int(text_top * H), int(0.58 * W), int(0.3 * H))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = "Больше прибыли с каждой чашки, больше гостей и больше выручки"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    path = tmp_path / "panels.pptx"
    prs.save(str(path))
    return path


def test_text_outside_card_ignores_a_second_panel_under_the_start_and_still_flags_a_card_below(tmp_path):
    """G1-02: a heading that starts on two overlapping panels (neither holds the other) does not «run out of its card»
    into the one that also holds its first line; one that starts on the upper panel only and hangs below it into the
    lower panel still does (the column of a two-column slide whose list hangs into the conclusion strip)."""
    from verstka.audit.checks.layout import text_outside_card
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import BboxFrac
    from verstka.schemas.template import SlideSize, Spacing, TemplateManifest, Tokens

    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=9144000, h=6858000), tokens=Tokens(spacing=Spacing(safe_area=BboxFrac(x=0.04, y=0.04, w=0.92, h=0.92))), patterns=[], n_slides=1)
    # starts on both panels (y 0.5): the lower one is its ground too — no issue
    ir = build_deck_ir(_panels_deck(tmp_path, 0.5), with_images=False)
    assert not text_outside_card(AuditContext(ir=ir, manifest=man))
    # starts on the upper panel only and hangs below its foot (0.86) into a strip under it: flagged
    strip = (0.07, 0.87, 0.77, 0.12)
    ir = build_deck_ir(_panels_deck(tmp_path, 0.66, second=strip), with_images=False)
    got = text_outside_card(AuditContext(ir=ir, manifest=man))
    assert got and got[0].check_id == "text_outside_card"


def test_figures_keep_their_thousands_and_unit_on_display_lines():
    """G1-16: «120 000» never breaks as «120 / 000», «255 000 рублей» never leaves «рублей» alone on a line."""
    from verstka.matching.scorer import bind_figures, bind_short_words

    got = bind_short_words(_GOAL)
    assert "120 000" in got and "255 000 рублей" in got
    assert bind_figures("выручка 42 % и 18,6 млрд ₽, 1 854 машины") == "выручка 42 % и 18,6 млрд ₽, 1 854 машины"
    assert bind_figures("в 2019 2020 годах") == "в 2019 2020 годах"  # two years are not one figure


def _covers_manifest():
    """VK Tech-like: a centred cover whose cube stands right under a one-line subtitle (p1, scored first) and a
    left-aligned cover with room under its heading (p2)."""
    from verstka.schemas.common import BboxFrac, Family, PatternKind, SlotRole
    from verstka.schemas.template import Capacity, FontUsage, Pattern, SlideSize, Slot, SlotStyle, Spacing, TemplateManifest, Tokens, TypeStep, Typography

    def slot(sid, role, box, size, sample, chars=30, align="l"):
        return Slot(id=f"{role.value}_{sid}", role=role, shape_id=sid, bbox=BboxFrac(x=box[0], y=box[1], w=box[2], h=box[3]), style=SlotStyle(font_family="Play", size_pt=size, align=align), capacity=Capacity(max_chars=chars, max_lines=1), sample_text=sample)

    centred = Pattern(id="p1", source_slide=1, kind=PatternKind.title, family=Family.dark, quality=1.0, layout_part="l1", slots=[
        slot("1", SlotRole.title, (0.271, 0.266, 0.469, 0.156), 48, "VK Tech", 11, "ctr"),
        slot("2", SlotRole.subtitle, (0.271, 0.428, 0.469, 0.056), 16, "Разработчик корпоративного ПО", 38, "ctr"),
    ])
    left = Pattern(id="p2", source_slide=2, kind=PatternKind.title, family=Family.dark, quality=0.9, layout_part="l2", slots=[
        slot("3", SlotRole.title, (0.064, 0.298, 0.469, 0.156), 44, "VK Tech", 11),
        slot("4", SlotRole.subtitle, (0.064, 0.459, 0.469, 0.056), 16, "Разработчик корпоративного ПО", 38),
        slot("5", SlotRole.body, (0.156, 0.747, 0.469, 0.056), 16, "Имя Фамилия", 38),
        slot("6", SlotRole.body, (0.156, 0.818, 0.469, 0.056), 11, "Должность", 56),
    ])
    typo = Typography(families=[FontUsage(family="Play", weight=1.0)], scale=[TypeStep(role="display", size_pt=54), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=16), TypeStep(role="body", size_pt=12), TypeStep(role="caption", size_pt=10)], sizes_used=[54, 48, 36, 24, 16, 14, 12, 10])
    tokens = Tokens(typography=typo, spacing=Spacing(safe_area=BboxFrac(x=0.03, y=0.06, w=0.94, h=0.86)))
    return TemplateManifest(template_id="vkl", source_file="s.pptx", slide_size=SlideSize(w=9144000, h=5143500), tokens=tokens, patterns=[centred, left], n_slides=2)


def test_a_cover_whose_stack_does_not_fit_its_sample_gives_way_to_one_that_holds_it():
    """G1-20: «История VK» with a 50-character subtitle does not take the centred cover whose room ends at a one-line
    subtitle (the cube under it) when the left-aligned cover within the margin holds the stack; a short subtitle keeps
    the centred cover."""
    from verstka.matching.matcher import _cover_room_ok, match_outline
    from verstka.planning.strategies import get_strategy
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    m = _covers_manifest()
    long_sub = OutlineSlide(id="t", kind=PatternKind.title, headline="История VK", subtitle="От почтового сервиса до технологической корпорации")
    short_sub = OutlineSlide(id="t", kind=PatternKind.title, headline="История VK", subtitle="Хронология компании")
    by_id = {p.id: p for p in m.patterns}
    assert not _cover_room_ok(long_sub, by_id["p1"], m) and _cover_room_ok(long_sub, by_id["p2"], m)
    for st in ("structured", "compact"):
        got = match_outline(DeckOutline(title="x", slides=[long_sub]), m, get_strategy(st)).slides[0]
        assert got.pattern_id == "p2" and any("весь блок текста" in r for r in got.reasons), (st, got.reasons, got.alternatives)
        assert dict(got.alternatives)["p1"] > got.score  # the centred cover scored first and gave way
    if _cover_room_ok(short_sub, by_id["p1"], m):
        first = match_outline(DeckOutline(title="x", slides=[short_sub]), m, get_strategy("structured")).slides[0]
        assert first.pattern_id == "p1" or not any("весь блок текста" in r for r in first.reasons)


def test_a_rerender_starts_from_the_notes_the_render_did_not_add():
    """A cover whose first render said its goal aloud (no room on the clone sample) and whose rematch draws it on the
    slide must not repeat it in the speaker notes: autofix takes the said-aloud cover lines off before re-rendering."""
    from verstka.audit.autofix import _unsay
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    cover = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Больше прибыли", content=SlideContent(paragraphs=[_GOAL]), footnote="Все цифры условные",
                         notes="Главная цель: рост прибыли.\n" + _GOAL + "\nВсе цифры  условные")
    other = OutlineSlide(id="sl2", kind=PatternKind.bullets, headline="План", notes="Все цифры условные")
    o = DeckOutline(title="x", slides=[cover, other])
    assert _unsay(o) == 2
    assert o.slides[0].notes == "Главная цель: рост прибыли." and o.slides[1].notes == "Все цифры условные"


def test_cover_small_print_takes_the_template_dark_text_on_a_light_ground(monkeypatch):
    """G1-18: small print under an orange heading on white (4.2:1) is set in the template's dark text colour at the
    small size, not 14 pt bold orange; where the line colour reads, or on a dark ground, it keeps its colour."""
    from types import SimpleNamespace

    import verstka.rendering.clone as C
    from verstka.schemas.common import Bbox
    from verstka.schemas.template import ColorToken, Tokens

    toks = Tokens(colors=[ColorToken(hex="FFFFFF", role="background.light", roles=["background.light"]), ColorToken(hex="009BDD", role="accent.2", roles=["accent.2"]), ColorToken(hex="DD4100", role="accent.3", roles=["accent.3"])])
    ctx = SimpleNamespace(manifest=SimpleNamespace(tokens=toks))
    box = Bbox(x=0, y=0, w=100, h=10)
    monkeypatch.setattr(C, "_ground_hex", lambda ctx, b: "FFFFFF")
    assert C._small_print_color(ctx, box, "DD4100") == "000000"  # no text.primary: the darkest colour that reads
    toks.colors.append(ColorToken(hex="333333", role="text.primary", roles=["text.primary"]))
    assert C._small_print_color(ctx, box, "DD4100") == "333333"
    assert C._small_print_color(ctx, box, "222222") is None  # it reads: kept
    monkeypatch.setattr(C, "_ground_hex", lambda ctx, b: "0077FF")
    assert C._small_print_color(ctx, box, "FFFFFF") is None  # a brand ground: the 14 pt bold rule stays


# ------------------------------------------------------------------------------------------------ gate 2 (C1–C5)

TPL = FIX / "templates"
_LONG_NOTE = "Все исходные данные и прогнозы условные"


def _render_cover(tmp_path: Path, template: str, long: bool = True):
    """Slide 1 of the coffee brief (long: heading, subtitle, goal, small print; short: a two-phrase heading and small
    print) rendered on a corpus template of tests/fixtures/tpl_audit/templates. Returns (presentation, slide 1)."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    if long:
        cover = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Больше прибыли с каждой чашки", subtitle="План развития кофейни “Точка кофе” на 6 месяцев",
                             content=SlideContent(paragraphs=[_GOAL]), footnote=_LONG_NOTE, notes="Главная цель: увеличить ежемесячную операционную прибыль со 120 000 до 255 000 рублей.")
    else:
        cover = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Кофейня «Точка кофе»: план увеличения прибыли за 6 месяцев", footnote="Все цифры условные")
    outline = DeckOutline(title=cover.headline, slides=[cover], language="ru")
    out = tmp_path / f"run_{template}_{'long' if long else 'short'}"
    generate_variants(TPL / f"{template}.pptx", outline=outline, strategies=["structured"], out_dir=out, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, audit=False, exports=[])
    prs = Presentation(str(out / "structured" / "deck.pptx"))
    return prs, prs.slides[0]


def _shape_with(slide, text: str):
    return next((sh for sh in slide.shapes if sh.has_text_frame and text in sh.text_frame.text.replace("\xa0", " ")), None)


def _notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def test_grey_elegant_long_cover_keeps_the_brief_disclaimer_on_slide_one(tmp_path):
    """C1: on LO Grey Elegant the long cover's goal fills the column under the heading; the small print the brief asks
    for on the first slide goes into the column beside the stack (under the date, right of the rule) — not into the
    speaker notes — inside the slide and clear of every other text."""
    prs, s = _render_cover(tmp_path, "lo_grey_elegant")
    note = _shape_with(s, "условные")
    assert note is not None, [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "условные" not in _notes(s)
    assert _shape_with(s, "Цель:") is not None  # the goal stays too
    nb = (note.left, note.top, note.left + note.width, note.top + note.height)
    assert nb[0] >= 0 and nb[2] <= prs.slide_width and nb[3] <= prs.slide_height
    for sh in s.shapes:
        if sh.shape_id == note.shape_id or not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        ob = (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
        assert not (min(nb[2], ob[2]) > max(nb[0], ob[0]) and min(nb[3], ob[3]) > max(nb[1], ob[1])), f"the note overlaps «{sh.text_frame.text[:30]}»"


def test_grey_elegant_short_cover_keeps_its_small_print_on_the_slide(tmp_path):
    """C1 (short brief): «Все цифры условные» is on slide 1 of LO Grey Elegant, not in the notes."""
    _, s = _render_cover(tmp_path, "lo_grey_elegant", long=False)
    assert _shape_with(s, "условные") is not None and "условные" not in _notes(s)


def test_focus_cover_heading_lines_stay_on_their_triangle(tmp_path):
    """C2: on LO Focus the cover heading stands on the red triangle; each of its lines (a slanted edge moves with the
    height) keeps its letters on the triangle — the audit's line-by-line ground check finds nothing off it."""
    from verstka.audit.checks.common import text_elements
    from verstka.audit.checks.template import _line_off_ground
    from verstka.audit.ir import build_deck_ir

    _render_cover(tmp_path, "lo_focus")
    ir = build_deck_ir(tmp_path / "run_lo_focus_long" / "structured" / "deck.pptx")
    sl = ir.slides[0]
    title = next(e for e in text_elements(sl) if e.ph_type in ("title", "ctrTitle"))
    assert "чашки" in title.text
    assert title.dominant_size and title.dominant_size >= 24  # not shrunk to body size to fit
    assert _line_off_ground(sl, title, sl.background_hex or "FFFFFF", title.dominant_color or "FFFFFF", 3.0) is None


def test_audit_flags_a_heading_line_off_its_triangle():
    """C2/C4: the gate-2 Focus cover (28 pt heading, its second line centred wider than the red triangle at its own
    height — «с» on the white gap) is flagged by contrast_low's line-by-line ground check; its goal only in the speaker
    notes is flagged by content_in_notes."""
    from verstka.audit.runner import run_audit
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent
    from verstka.schemas.template import TemplateManifest

    man = TemplateManifest.model_validate_json(gzip.open(FIX / "manifest_368fc0232d0871b3.json.gz", "rt", encoding="utf-8").read())
    cover = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Больше прибыли с каждой чашки", subtitle="План развития кофейни “Точка кофе” на 6 месяцев",
                         content=SlideContent(paragraphs=[_GOAL]), footnote=_LONG_NOTE)
    r = run_audit(FIX / "focus_long_cover_g2.pptx", man, DeckOutline(title="x", slides=[cover]), render=False, use_llm=False, use_vlm=False)
    off = [i for i in _serious(r, "contrast_low") if (i.details or {}).get("ground") == "line_off_shape"]
    assert off and "чашки" in off[0].message
    notes = _serious(r, "content_in_notes")
    assert len(notes) == 1 and notes[0].details["line"] == "goal" and notes[0].details["in_notes"]


def test_small_print_size_stays_small_on_big_captions_and_big_slides():
    """C3: the cover small print is at most the scale step under the subtitle, 0.7 of it and 2.6 % of the slide
    height, never under 1.8 % of it: a template whose caption is 20 pt (Sidebar 4:3, subtitle 20 pt) and a 26.67″ slide
    (subtitle 40 pt) no longer set it at the subtitle's size."""
    from verstka.rendering.clone import _small_print_size
    from verstka.schemas.common import EMU_PER_PT

    h43 = 6858000
    size = _small_print_size(20.0, h43, [10, 12, 14, 16, 20, 24, 28, 36], [10, 12, 14, 16, 20, 24, 28, 36], 20.0)
    assert size <= 0.7 * 20.0 and size <= 0.026 * h43 / EMU_PER_PT and size >= 0.018 * h43 / EMU_PER_PT - 0.5
    h_huge = int(15 * 914400)
    big = _small_print_size(40.0, h_huge, [24, 28, 32, 40, 48, 88], [24, 28, 32, 40, 48, 88], 28.0)
    assert big <= 0.7 * 40.0 and big <= 0.026 * h_huge / EMU_PER_PT + 0.01
    assert _small_print_size(28.0, 6858000, [12, 14, 16, 18, 28], [12, 14, 16, 18, 28], 12.0) <= 18.0  # the old rule's size


@pytest.mark.parametrize("template", ["synth_sidebar43", "synth_huge2667"])
def test_cover_small_print_is_set_small_under_the_stack(tmp_path, template):
    """C3: on Sidebar 4:3 (20 pt caption) and Huge 26.67″ covers the small print is at most 0.7 of the subtitle and
    2.6 % of the slide height, and stands under the heading's stack (not free-floating mid-slide)."""
    from verstka.schemas.common import EMU_PER_PT

    prs, s = _render_cover(tmp_path, template)
    note = _shape_with(s, "условные")
    sub = _shape_with(s, "План развития")
    assert note is not None and sub is not None
    size = lambda sh: max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0.0)
    assert size(note) <= 0.7 * size(sub) + 0.01
    assert size(note) <= 0.026 * prs.slide_height / EMU_PER_PT + 0.51
    assert note.top >= sub.top + sub.height  # under the stack


def test_vivid_long_cover_keeps_its_goal_between_the_subtitle_and_the_small_print(tmp_path):
    """C5: the LO Vivid long cover (two-line heading over the blue band) keeps the goal on the slide, between the
    subtitle and the small print, instead of saying it aloud."""
    prs, s = _render_cover(tmp_path, "lo_vivid")
    goal, sub, note = _shape_with(s, "Цель:"), _shape_with(s, "План развития"), _shape_with(s, "условные")
    assert goal is not None and sub is not None and note is not None
    assert "Цель:" not in _notes(s)
    tol = int(0.01 * prs.slide_height)
    assert sub.top + sub.height <= goal.top + tol and goal.top + goal.height <= note.top + tol


# ------------------------------------------------------------------------------------------------ gate 2 (C4) checks


def _ir_slide(index: int, texts: list[tuple[str, tuple[float, float, float, float], float]], notes: str = "", oid: str | None = None, extra=()):
    """An IR slide of text boxes (text, frac box, size) on a 13.33×7.5″ slide."""
    from verstka.schemas.common import Bbox, BboxFrac
    from verstka.schemas.deck_ir import IRElement, IRParagraph, IRRun, IRSlide

    W, H = 12192000, 6858000
    els = []
    for i, (t, (x, y, w, h), size) in enumerate(texts):
        els.append(IRElement(id=f"e{index}_{i}", type="text", bbox=Bbox(x=int(x * W), y=int(y * H), w=int(w * W), h=int(h * H)), bbox_frac=BboxFrac(x=x, y=y, w=w, h=h),
                             paragraphs=[IRParagraph(text=t, runs=[IRRun(text=t, size_pt=size, font="Arial")])], ph_type="title" if i == 0 else None, z=i))
    els.extend(extra)
    return IRSlide(index=index, elements=els, notes=notes, outline_id=oid)


def _ctx(slides, outline=None):
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import BboxFrac
    from verstka.schemas.deck_ir import DeckIR
    from verstka.schemas.template import SlideSize, Spacing, TemplateManifest, Tokens

    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=12192000, h=6858000), tokens=Tokens(spacing=Spacing(safe_area=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.9))), patterns=[], n_slides=1)
    return AuditContext(ir=DeckIR(source="x.pptx", slide_w=12192000, slide_h=6858000, slides=slides), manifest=man, outline=outline)


def test_content_in_notes_flags_a_conclusion_said_only_aloud_and_accepts_its_callout_form():
    from verstka.audit.checks.integrity import content_in_notes
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    o = DeckOutline(title="x", slides=[
        OutlineSlide(id="a", kind=PatternKind.bullets, headline="Расходы", content=SlideContent(bullets=["Аренда 120 000 ₽"]), takeaway="Все разовые вложения покрываются бюджетом запуска"),
        OutlineSlide(id="b", kind=PatternKind.chart, headline="Прибыль", takeaway="Прогноз: 254 795 ₽ в месяц", footnote="Налоги не учитываются"),
        OutlineSlide(id="c", kind=PatternKind.bullets, headline="План", takeaway="Рост на 26,5% за 6 месяцев"),
    ])
    s1 = _ir_slide(1, [("Расходы", (0.05, 0.05, 0.9, 0.1), 32), ("Аренда 120 000 ₽", (0.05, 0.3, 0.9, 0.1), 18)], notes="Вывод: Все разовые вложения покрываются бюджетом запуска. [verstka:a]", oid="a")
    s2 = _ir_slide(2, [("Прибыль", (0.05, 0.05, 0.9, 0.1), 32), ("254 795 ₽", (0.6, 0.3, 0.3, 0.1), 40), ("Прогноз", (0.6, 0.4, 0.3, 0.05), 14), ("Налоги не учитываются", (0.05, 0.9, 0.5, 0.04), 10)], oid="b")
    s3 = _ir_slide(3, [("План", (0.05, 0.05, 0.9, 0.1), 32)], oid="c")
    got = content_in_notes(_ctx([s1, s2, s3], o))
    assert [(i.slide, i.details["line"], i.details["in_notes"]) for i in got] == [(1, "takeaway", True), (3, "takeaway", False)]
    assert all(i.severity == "warn" for i in got)


def test_timeline_order_reads_dates_and_flags_a_list_out_of_order():
    from verstka.audit.checks.integrity import date_key, timeline_order
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent, SlideItem

    assert date_key("22 июня 1941") == (1941, 6, 22)
    assert date_key("Апрель — июнь 1940") == (1940, 4, None)
    assert date_key("1931–1937") == (1931, None, None)
    assert date_key("III кв. 2025") == (2025, 7, None)
    assert date_key("Мая 1945") == (1945, 5, None) and date_key("Март 1945") == (1945, 3, None)
    assert date_key("2800 студентов") is None and date_key("1500 ₽") is None and date_key("1-й месяц") is None
    items = lambda *ts: SlideContent(items=[SlideItem(title=t) for t in ts])
    o = DeckOutline(title="x", slides=[
        OutlineSlide(id="a", kind=PatternKind.timeline, headline="Война на Тихом океане", content=items("22 июня 1941", "1931–1937", "Июль 1937", "7 декабря 1941")),
        OutlineSlide(id="b", kind=PatternKind.timeline, headline="Итоги", content=items("1 сентября 1939", "1941 год", "Июнь 1941", "8 мая 1945")),
        OutlineSlide(id="c", kind=PatternKind.timeline, headline="Назад", content=items("2026", "2025", "2024")),
        OutlineSlide(id="d", kind=PatternKind.process, headline="Этапы", content=items("1-й месяц", "3-й месяц", "2-й месяц")),
    ])
    slides = [_ir_slide(i + 1, [("x", (0.05, 0.05, 0.9, 0.1), 32)], oid=oid) for i, oid in enumerate("abcd")]
    got = timeline_order(_ctx(slides, o))
    assert [i.slide for i in got] == [1]
    assert "«1931–1937» после «22 июня 1941»" in got[0].message


def test_fill_ratio_flags_a_one_sentence_slide_even_beside_template_art():
    """C4: a content slide whose text under the heading is one short sentence is nearly empty (warn), whatever a
    picture of the template fills beside it; a hero figure, a chart slide and a cover are not."""
    from verstka.audit.checks.density import fill_ratio
    from verstka.schemas.common import Bbox, BboxFrac, PatternKind
    from verstka.schemas.deck_ir import IRElement
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    W, H = 12192000, 6858000
    art = IRElement(id="pic", type="picture", bbox=Bbox(x=int(0.04 * W), y=int(0.12 * H), w=int(0.55 * W), h=int(0.85 * H)), bbox_frac=BboxFrac(x=0.04, y=0.12, w=0.55, h=0.85), z=9)
    o = DeckOutline(title="x", slides=[
        OutlineSlide(id="a", kind=PatternKind.title, headline="История"),
        OutlineSlide(id="b", kind=PatternKind.bullets, headline="VK начала как почтовый сервис", content=SlideContent(paragraphs=["Основана в 1998 году"])),
        OutlineSlide(id="c", kind=PatternKind.big_number, headline="Охват", content=SlideContent(paragraphs=["2800"])),
        OutlineSlide(id="d", kind=PatternKind.bullets, headline="Итог", content=SlideContent(paragraphs=["Основана в 1998 году"])),
    ])
    slides = [
        _ir_slide(1, [("История", (0.05, 0.3, 0.5, 0.2), 54)], oid="a"),
        _ir_slide(2, [("VK начала как почтовый сервис", (0.63, 0.06, 0.3, 0.2), 28), ("Основана в 1998 году", (0.66, 0.34, 0.31, 0.06), 16)], oid="b", extra=[art]),
        _ir_slide(3, [("Охват", (0.05, 0.06, 0.5, 0.1), 28), ("2800", (0.05, 0.2, 0.8, 0.3), 120)], oid="c"),
        _ir_slide(4, [("Итог", (0.05, 0.06, 0.5, 0.1), 28), ("Основана в 1998 году", (0.05, 0.3, 0.5, 0.06), 16)], oid="d"),
    ]
    got = fill_ratio(_ctx(slides, o))
    sparse = [i for i in got if (i.details or {}).get("text_only")]
    assert [i.slide for i in sparse] == [2, 4] and all(i.severity == "warn" for i in sparse)  # the last slide too: it is no closing slide


# ------------------------------------------------------------------------------------------------ round 4.1 (after gate 3)


def test_content_in_notes_accepts_the_visual_aside_of_a_two_figure_conclusion():
    """The visual variant sets «За 30 дней выручка составляет 900 000 рублей» beside its pie as «900 000 ₽» over «За
    30 дней выручка» (the line has two figures, so kpi_callout declines it): the conclusion is on the slide."""
    from verstka.audit.checks.integrity import content_in_notes
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    o = DeckOutline(title="x", slides=[OutlineSlide(id="a", kind=PatternKind.chart, headline="Кофе приносит 60% выручки", takeaway="За 30 дней выручка составляет 900 000 рублей")])
    s = _ir_slide(1, [("Кофе приносит 60% выручки", (0.05, 0.05, 0.9, 0.1), 32), ("900 000 ₽", (0.7, 0.3, 0.25, 0.1), 40), ("За 30 дней выручка", (0.7, 0.4, 0.25, 0.05), 14)], oid="a")
    assert content_in_notes(_ctx([s], o)) == []
    bare = _ir_slide(1, [("Кофе приносит 60% выручки", (0.05, 0.05, 0.9, 0.1), 32), ("900 000 ₽", (0.7, 0.3, 0.25, 0.1), 40)], oid="a")
    assert len(content_in_notes(_ctx([bare], o))) == 1  # the figure alone, without its words, is not the conclusion


def test_fix_slide_sends_the_new_content_checks_to_the_designer():
    """«Исправить слайд»: a conclusion said aloud, dates out of order and a slide that is one picture are the slide
    designer's to fix (CONTENT_CHECKS), each with a hint; the note names the whole line that left the slide."""
    from verstka.api import remarks as RM
    from verstka.audit.registry import all_checks, load_builtin_checks
    from verstka.schemas.audit import Issue

    for cid in ("content_in_notes", "timeline_order", "slide_is_picture"):
        assert cid in RM.CONTENT_CHECKS and RM.FIX_HINT.get(cid)
    assert RM.FIX_HINT.get("content_over_art")
    load_builtin_checks()
    known = {spec.id for spec, _ in all_checks()}
    assert RM.CONTENT_CHECKS - {"slide_content", "deck_coherence"} <= known  # no stale id
    long = "За 30 дней выручка составляет 900 000 рублей при среднем чеке 300 рублей и 100 покупках в день"
    i = Issue(id="content_in_notes-2-6", slide=2, check_id="content_in_notes", severity="warn", kind="deterministic",
              message=f"Вывод «{long[:50]}…» есть только в заметках докладчика: на слайде не нашлось места",
              details={"line": "takeaway", "text": long, "in_notes": True})
    note = RM.remarks_note([i])
    assert f"вывод «{long}» есть только в заметках докладчика" in note and "Как исправить: верни эту строку на слайд" in note


def _cover_audit(tmp_path: Path, template: str, long: bool = True):
    """Render slide 1 of the coffee brief on a fixture template and audit it with the analysed manifest (no render)."""
    import glob

    from verstka.audit.runner import run_audit
    from verstka.schemas.template import TemplateManifest

    prs, s = _render_cover(tmp_path, template, long=long)
    run = tmp_path / f"run_{template}_{'long' if long else 'short'}" / "structured"
    man = TemplateManifest.model_validate_json(Path(glob.glob(str(tmp_path / "ws" / "templates" / "*" / "manifest.json"))[0]).read_text())
    from verstka.schemas.outline import DeckOutline

    outline = DeckOutline.model_validate_json((run / "outline.json").read_text())
    report = run_audit(run / "deck.pptx", man, outline, render=False, use_llm=False, use_vlm=False)
    return prs, s, report


def test_focus_long_cover_keeps_its_goal_on_a_calm_panel(tmp_path):
    """C3-1: the LO Focus long cover's goal found no calm place under the heading or beside the subtitle (the white
    centre between the triangles); it no longer goes to the speaker notes: it stands on a calm painted panel of the
    layout, every line on that one panel at its own height, in a colour that reads there as small text, clear of every
    other text — and the audit finds nothing on slide 1."""
    from verstka.rendering.layers import ground_under
    from verstka.schemas.common import Bbox, EMU_PER_PT, contrast_ratio

    prs, s, report = _cover_audit(tmp_path, "lo_focus")
    goal = _shape_with(s, "Цель:")
    assert goal is not None, [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "Цель:" not in _notes(s).split("[verstka")[0].replace("Главная цель:", "")
    size = max(r.font.size.pt for p in goal.text_frame.paragraphs for r in p.runs if r.font.size)
    assert size >= 0.018 * prs.slide_height / EMU_PER_PT - 0.05
    color = str(goal.text_frame.paragraphs[0].runs[0].font.color.rgb)
    lines = goal.text_frame.text.replace("\x0b", "\n").split("\n")
    assert 1 <= len(lines) <= 4
    pitch = int(1.2 * size * EMU_PER_PT)
    top = goal.top + goal.text_frame.margin_top
    layer = None
    for i in range(len(lines)):
        row = Bbox(x=goal.left + goal.text_frame.margin_left, y=top + i * pitch, w=max(goal.width - goal.text_frame.margin_left - goal.text_frame.margin_right, 1), h=pitch)
        got = ground_under(s, Bbox(x=row.x + row.w // 3, y=row.y, w=row.w // 3, h=row.h))
        assert got is not None
        layer = layer or got[1]
        assert got[1].el is layer.el  # one panel under every line
        assert contrast_ratio(color, got[0]) >= 4.5
    for sh in s.shapes:
        if sh.shape_id == goal.shape_id or not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        a = (goal.left, goal.top, goal.left + goal.width, goal.top + goal.height)
        b = (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
        assert not (min(a[2], b[2]) > max(a[0], b[0]) and min(a[3], b[3]) > max(a[1], b[1])), f"the goal overlaps «{sh.text_frame.text[:30]}»"
    bad = [i for i in report.issues if i.slide == 1 and i.severity in ("error", "warn")]
    assert bad == [], [(i.check_id, i.message) for i in bad]


def test_panel_run_follows_a_triangle_edge():
    """layers.panel_run: the run of the LO Focus red triangle over a band narrows as the band goes down to its tip."""
    from verstka.rendering.layers import drawn_layers, panel_run

    prs = Presentation(str(TPL / "lo_focus.pptx"))
    s = prs.slides[0]
    H = prs.slide_height
    red = next(l for l in drawn_layers(s) if l.fill_hex == "F10D0C")
    hi = panel_run(red, int(0.05 * H), int(0.08 * H))
    lo = panel_run(red, int(0.35 * H), int(0.38 * H))
    assert hi and lo and (lo[0][1] - lo[0][0]) < (hi[0][1] - hi[0][0])
    assert panel_run(red, int(0.6 * H), int(0.65 * H)) == []  # past its tip


def test_candy_cover_subtitle_keeps_to_the_calm_part_of_the_photo(tmp_path):
    """C3-2: on LO Candy the cover subtitle ran from the blurred table onto the candies on one line. It is now set on
    lines that each keep to the photo's calm stretch at their own height (two lines, at most a step smaller), and the
    audit — which sees a line running onto the busy part of a full-slide photo — finds nothing on it."""
    from verstka.schemas.common import EMU_PER_PT

    prs, s, report = _cover_audit(tmp_path, "lo_candy", long=False)
    sub = _shape_with(s, "План увеличения")
    assert sub is not None
    lines = sub.text_frame.text.replace("\x0b", "\n").split("\n")
    assert len(lines) == 2 and " ".join(lines).replace("\xa0", " ") == "План увеличения прибыли за 6 месяцев"
    art = [i for i in report.issues if i.slide == 1 and i.check_id == "content_over_art"]
    assert art == [], [i.message for i in art]
    assert sub.left + sub.width <= 0.6 * prs.slide_width  # off the candies


def test_audit_sees_a_cover_line_running_onto_the_busy_part_of_a_photo():
    """C3-2 (audit): the gate-3 LO Candy short cover — the subtitle on one line from the calm table onto the candies —
    is flagged by content_over_art line by line (warn), though the photo is the slide's ground, not art."""
    from verstka.audit.checks.layout import content_over_art
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext

    ir = build_deck_ir(FIX / "candy_short_cover_g3.pptx")
    got = content_over_art(AuditContext(ir=ir, manifest=None, outline=None))
    assert len(got) == 1 and got[0].severity == "warn" and got[0].details.get("photo")
    assert "План увеличения" in got[0].message


def test_audit_sees_a_line_running_onto_a_photo_the_slide_itself_holds(tmp_path):
    """The same line-by-line photo check on a cover whose photo is a picture of the slide itself (many user templates
    put it there, not in the layout); a line standing on a card of its own over the photo is not on the photo."""
    from pptx.dml.color import RGBColor as _RGB
    from pptx.util import Pt

    from verstka.audit.checks.layout import content_over_art
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext

    src = Presentation(str(TPL / "lo_candy.pptx"))
    photo = max((p for p in src.part.package.iter_parts() if p.content_type.startswith("image/")), key=lambda p: len(p.blob))
    img = tmp_path / ("photo.png" if photo.content_type.endswith("png") else "photo.jpg")
    img.write_bytes(photo.blob)
    for card in (False, True):
        prs = Presentation(str(_blank_template(tmp_path)))
        W, H = prs.slide_width, prs.slide_height
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(img), 0, 0, W, H)
        tb = s.shapes.add_textbox(int(0.08 * W), int(0.40 * H), int(0.82 * W), int(0.09 * H))
        tb.text_frame.text = "План увеличения прибыли за 6 месяцев"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        if card:
            tb.fill.solid()
            tb.fill.fore_color.rgb = _RGB(0xFF, 0xFF, 0xFF)
        path = tmp_path / f"slide_photo_{card}.pptx"
        prs.save(path)
        got = content_over_art(AuditContext(ir=build_deck_ir(path), manifest=None, outline=None))
        photo_hits = [i for i in got if (i.details or {}).get("photo")]
        assert (len(photo_hits) == 0) if card else (len(photo_hits) == 1), [i.message for i in got]


def test_audit_tells_a_line_over_an_ink_blot_from_a_line_just_above_it(tmp_path):
    """The photo check reads where inside a busy coarse cell the busy part lies (the fine grid): on LO Vintage's paper a
    list line standing just above the ink blots at the foot shares their coarse cell but not their pixels — no remark
    (a near false positive of the full-corpus replay, round 4.1); the same line set over the blots is flagged."""
    from pptx.util import Pt

    from verstka.audit.checks.layout import content_over_art
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext

    got = {}
    # 0.782: the line's letters (12 pt) lie in the coarse cell row of the blot's top (0.786–0.857 of the height) but
    # above the blot itself (its top at 0.82); 0.86: on the blots
    for y in (0.782, 0.86):
        prs = Presentation(str(_blank_template(tmp_path)))
        W, H = prs.slide_width, prs.slide_height
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(FIX / "vintage_paper.jpg"), 0, 0, W, H)
        tb = s.shapes.add_textbox(int(0.03 * W), int(y * H), int(0.4 * W), int(0.04 * H))
        tb.text_frame.text = "Обновление меню — 25 000 ₽"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        path = tmp_path / f"vintage_{y}.pptx"
        prs.save(path)
        ir = build_deck_ir(path)
        got[y] = [i for i in content_over_art(AuditContext(ir=ir, manifest=None, outline=None)) if (i.details or {}).get("photo")]
    assert got[0.782] == [] and len(got[0.86]) == 1
