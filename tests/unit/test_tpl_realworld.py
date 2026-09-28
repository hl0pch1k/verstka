"""Real-world templates (free Google Slides / Canva / SlidesCarnival decks, Office themes): the rules found on them.

* a palette PNG's transparency lives in its palette (tRNS): Office «Ion»'s glows are ≤ 14 % opaque and never art;
* a cover whose title says «Consulting» is no comparison, and a template's first slide is its cover;
* credits, «How to use this template» and resource pages are service slides, never a layout; a watermark domain in
  a header is not;
* a header's or a footer's sample values («JOHN DOE», «NEW YORK», «2023», the template site) are not the deck's;
* a deep burgundy or navy used for headings is the brand accent;
* a pie whose accent is kin to the slide's ground steps toward white or black;
* a heading column past the middle of an empty slide returns to the page grid; a column never mirrors itself shut;
* a family without Cyrillic is measured and set in a stand-in of its class.
"""

from __future__ import annotations

import io
import struct
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from verstka.schemas.common import Bbox


# ---------------------------------------------------------------------------------------------- transparency


def _glow_png(alpha: int) -> bytes:
    """A palette PNG whose one colour (white) is `alpha` opaque through the tRNS chunk — the way Office themes store
    their glows."""
    im = Image.new("P", (64, 64), 0)
    im.putpalette([255, 255, 255] + [0, 0, 0] * 255)
    buf = io.BytesIO()
    im.save(buf, format="PNG", transparency=bytes([alpha]))
    return buf.getvalue()


def test_a_faint_palette_png_is_no_art(tmp_path):
    from verstka.rendering.layers import art_layers

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    png = tmp_path / "glow.png"
    png.write_bytes(_glow_png(30))
    s.shapes.add_picture(str(png), 0, Emu(int(prs.slide_height * 0.4)), Emu(int(prs.slide_width * 0.33)), Emu(int(prs.slide_height * 0.6)))
    assert Image.open(png).mode == "P"
    assert art_layers(s) == []  # at most 12 % opaque everywhere: the ground shows through
    solid = tmp_path / "solid.png"
    Image.new("RGB", (64, 64), (200, 30, 30)).save(solid)
    s.shapes.add_picture(str(solid), Emu(int(prs.slide_width * 0.6)), 0, Emu(int(prs.slide_width * 0.3)), Emu(int(prs.slide_height * 0.3)))
    assert len(art_layers(s)) == 1


# ---------------------------------------------------------------------------------------------- kinds and service slides


def test_consulting_is_no_comparison():
    from verstka.analysis.kinds import _COMPARE_RE

    assert not _COMPARE_RE.search("McKinsey Consulting")
    assert not _COMPARE_RE.search("Consolidated report")
    for t in ("Pros & Cons", "Плюсы и минусы", "До и после", "SaaS vs on-premise", "Тарифы"):
        assert _COMPARE_RE.search(t), t


def _shapes_of(texts: list[tuple[str, float]], W: int = 12192000, H: int = 6858000):
    from verstka.analysis.shapes import ParagraphInfo, ShapeInfo, TextInfo
    from verstka.schemas.common import ShapeKind

    return [
        ShapeInfo(id=str(i + 1), name=f"Text {i + 1}", kind=ShapeKind.sp, bbox=Bbox(x=int(0.1 * W), y=int(y * H), w=int(0.8 * W), h=int(0.08 * H)), z=i,
                  text=TextInfo(paragraphs=[ParagraphInfo(text=t)]))
        for i, (t, y) in enumerate(texts)
    ]


@pytest.mark.parametrize(
    "texts,service",
    [
        ([("Credits", 0.1), ("This presentation template is free for everyone to use thanks to the following:", 0.3), ("SlidesCarnival for the presentation template", 0.5)], True),
        ([("How to Use This Presentation", 0.05), ("Click on the “Google Slides” button below this presentation preview", 0.3)], True),
        ([("Resource Page", 0.05), ("Use these design resources in your Canva Presentation", 0.2)], True),
        ([("Our Company", 0.1), ("Who we are? Briefly elaborate on what you want to discuss.", 0.4), ("SLIDESCARNIVAL.COM", 0.01)], False),
        ([("Выручка за 30 дней", 0.1), ("Кофе — 60 % выручки", 0.4)], False),
    ],
)
def test_service_slides_of_free_templates(texts, service):
    from verstka.analysis.patterns import reference_reason

    shapes = _shapes_of(texts)
    got = reference_reason(shapes, {})
    assert (got is not None and "служебный" in got) is service, got


def test_header_and_footer_sample_values():
    from verstka.analysis.shapes import looks_like_sample_value

    for t in ("JOHN DOE", "NEW YORK", "2023", "SLIDESCARNIVAL.COM", "BRAND NAME", "www.reallygreatsite.com", "+123-456-7890",
              "hello@reallygreatsite.com", "Back to overview page", "Date: dd/mm/yyyy", "Borcelle", "July 2024", "Q3 2023"):
        assert looks_like_sample_value(t), t
    for t in ("VK Tech", "Сентябрь 2026", "2023 год", "Отчёт 2023", "12", "Олимпиада 2026", "Далее по плану"):
        assert not looks_like_sample_value(t), t


# ---------------------------------------------------------------------------------------------- colours


def test_a_deep_burgundy_heading_colour_is_the_accent():
    from verstka.analysis.colors import assign_color_roles
    from verstka.schemas.template import ColorToken

    def tok(hx, **cw):
        return ColorToken(hex=hx, weight=sum(cw.values()), context_weight=cw, contexts={k: 1 for k in cw})

    tokens = [tok("D5B58F", background=13.5), tok("000000", text=80.0), tok("5F1A1F", text=6.0, fill=4.0), tok("1A1A2E", text=2.0)]
    roles = {t.hex: t.role for t in assign_color_roles(tokens)}
    assert roles["5F1A1F"] == "accent.1"  # the parchment template's burgundy headings and panels
    assert not (roles.get("1A1A2E") or "").startswith("accent")  # a tinted black stays a neutral


def test_pie_slices_on_a_ground_of_the_same_kin_step_toward_white():
    from verstka.rendering.charts import delta_e, pie_palette

    pal = pie_palette([60, 25, 15], "FF5A5F", "D6006F")  # a coral accent on a magenta slide
    assert all(delta_e(pal[i], pal[i + 1]) >= 20 for i in range(len(pal) - 1))
    blue = pie_palette([60, 25, 15], "0077FF", "FFFFFF")  # the usual case keeps its tints toward the ground
    assert blue[0] == "0077FF" and all(delta_e(blue[i], blue[i + 1]) >= 20 for i in range(2))


# ---------------------------------------------------------------------------------------------- fonts


def _eot(cyr: bool, subset: bool = False) -> bytes:
    head = bytearray(80)
    struct.pack_into("<IIII", head, 0, 80, 0, 0x00020001, (1 if subset else 0) | 4)
    struct.pack_into("<H", head, 34, 0x504C)
    struct.pack_into("<I", head, 36, (1 << 9) if cyr else 1)
    struct.pack_into("<I", head, 52, (1 << 2) if cyr else 1)
    return bytes(head)


def test_embedded_fonts_tell_their_cyrillic():
    from verstka.rendering.cyrillic import _eot_cyrillic

    assert _eot_cyrillic(_eot(True)) is True
    assert _eot_cyrillic(_eot(False)) is False
    assert _eot_cyrillic(_eot(True, subset=True)) is False  # a subset keeps the template's own characters only
    assert _eot_cyrillic(b"junk") is None


def test_stand_ins_keep_the_class_and_prefer_the_templates_own(tmp_path):
    from verstka.rendering.cyrillic import cyrillic_substitutes, measured_family, deck_fonts

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for fam in ("Cardo", "DM Sans", "Anton", "Inter"):
        tb = s.shapes.add_textbox(0, 0, Inches(3), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "Sample"
        run.font.name = fam
    p = tmp_path / "t.pptx"
    prs.save(p)
    subs = cyrillic_substitutes(p)
    assert "Inter" not in subs
    assert subs["Cardo"] == "Georgia"  # a serif stays a serif
    assert subs["DM Sans"] == "Inter"  # the template's own sans with Cyrillic
    assert subs["Anton"] == "Impact"  # a heavy condensed display face
    with deck_fonts(subs):
        assert measured_family("DM Sans") == "Inter"
    assert measured_family("DM Sans") == "DM Sans"


def test_the_finished_deck_sets_russian_text_in_the_stand_in(tmp_path):
    from verstka.rendering.cyrillic import apply_to_pptx

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for text in ("Выручка за 30 дней", "SLIDESCARNIVAL"):
        tb = s.shapes.add_textbox(0, 0, Inches(3), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.name = "DM Sans"
    p = tmp_path / "deck.pptx"
    prs.save(p)
    assert apply_to_pptx(p, {"DM Sans": "Calibri"}) >= 1
    runs = [(r.text, r.font.name) for sh in Presentation(p).slides[0].shapes for r in sh.text_frame.paragraphs[0].runs]
    assert ("Выручка за 30 дней", "Calibri") in runs
    assert ("SLIDESCARNIVAL", "DM Sans") in runs  # Latin text the template keeps keeps its face
