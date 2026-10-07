"""Text measured as it is set, and composed content kept clear of the template's rules (round of 2026-10-04, QA on the
eight templates of the dev workspace):

* a heading over its text in one box is measured at its own step and weight, with the sample's space between the
  paragraphs (a WorkSpace card's last line hung under its plate);
* a box that sets no line spacing anywhere is single-spaced, whatever spacing the template's other texts use;
* the audit reads the spacing a paragraph inherits and the space between paragraphs;
* a rule of the layout through a composed area leaves the area its wider side; one on its edge pushes the edge in;
* a heading inherits its vertical anchor from the master through the layout;
* a mock-up slide shown alone keeps its first sentence as its lede; cards keep what a composed card holds;
* a cover of the author's text gets no invented date;
* the fixes of a finished deck (no markers in its notes) find their slides by the outline's order.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from verstka.analysis.xmlns import q
from verstka.schemas.common import EMU_PER_PT, Bbox

_CARD = """<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
  xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:nvSpPr><p:cNvPr id="5" name="Card"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="431800" y="1841500"/><a:ext cx="3636000" cy="2022000"/></a:xfrm>
    <a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="232323"/></a:solidFill></p:spPr>
  <p:txBody><a:bodyPr anchor="t" lIns="108000" tIns="72000" rIns="91425" bIns="45700" wrap="square"><a:noAutofit/></a:bodyPr><a:lstStyle/>
    <a:p><a:pPr><a:spcBef><a:spcPts val="0"/></a:spcBef></a:pPr><a:r><a:rPr lang="ru-RU" sz="2000" b="1"><a:latin typeface="Play"/></a:rPr><a:t>Первый этап развития</a:t></a:r></a:p>
    <a:p><a:pPr><a:spcBef><a:spcPts val="600"/></a:spcBef></a:pPr><a:r><a:rPr lang="ru-RU" sz="1800"><a:latin typeface="Play"/></a:rPr><a:t>{body}</a:t></a:r></a:p>
  </p:txBody></p:sp>"""

_LONG = ("Повышение надежности: проверка датчиков, контроль фактического расхода воды, резервное питание контроллера "
         "и удобный журнал событий.")


def test_a_heading_over_its_text_is_measured_as_set():
    from verstka.rendering.fit import fit_size

    texts = ["Первый этап развития", _LONG]
    box = Bbox(x=0, y=0, w=int(286.3 * EMU_PER_PT), h=int(159.2 * EMU_PER_PT))
    ins = (108000, 72000, 91425, 45700)
    plain = fit_size(texts, box, "Play", 18.0, False, [], insets_emu=ins, line_spacing=1.2, min_ratio=0.6)
    as_set = fit_size(texts, box, "Play", 18.0, False, [], insets_emu=ins, line_spacing=1.2, para_spacing_pt=6.0, min_ratio=0.6, head=(20 / 18, True))
    assert as_set.height_pt > plain.height_pt + 6.0 - 0.01 or as_set.size_pt < plain.size_pt
    # set at the size it was fitted, the heading and the text keep inside the card's inner height
    inner = (box.h - ins[1] - ins[3]) / EMU_PER_PT
    assert as_set.fits and as_set.height_pt <= inner


def test_ink_box_counts_the_space_between_paragraphs():
    from verstka.rendering.clone_check import ink_box

    sp = etree.fromstring(_CARD.format(body=_LONG))
    with_gap = ink_box(sp, "Play")
    sp2 = etree.fromstring(_CARD.format(body=_LONG).replace('<a:spcPts val="600"/>', '<a:spcPts val="0"/>'))
    without = ink_box(sp2, "Play")
    assert with_gap is not None and without is not None
    assert abs((with_gap.h - without.h) / EMU_PER_PT - 6.0) < 0.05


def _composed_rule_slide(tmp_path):
    """A slide on a layout that draws a hairline down the middle of the slide (the line moved from a scratch slide
    into the layout's own shape tree: python-pptx draws no shapes on layouts)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height
    scratch = prs.slides.add_slide(layout)
    line = scratch.shapes.add_connector(1, W // 2, int(0.06 * H), W // 2, int(0.94 * H))
    line.line.width = Pt(1)
    layout.shapes._spTree.append(line._element)
    slide = prs.slides.add_slide(layout)
    return prs, slide, W, H


def test_a_rule_through_the_area_leaves_its_wider_side(tmp_path):
    from verstka.rendering.layers import clear_of_rules

    prs, slide, W, H = _composed_rule_slide(tmp_path)
    area = Bbox(x=int(0.44 * W), y=int(0.3 * H), w=int(0.45 * W), h=int(0.6 * H))
    got = clear_of_rules(slide, area)
    assert got.x > W // 2 and got.x2 == area.x2 and got.h == area.h


def test_a_rule_on_the_edge_of_the_area_pushes_the_edge_in(tmp_path):
    from verstka.rendering.layers import clear_of_rules

    prs, slide, W, H = _composed_rule_slide(tmp_path)
    area = Bbox(x=int(0.05 * W), y=int(0.15 * H), w=int(0.455 * W), h=int(0.7 * H))  # its right edge a hair past the rule
    got = clear_of_rules(slide, area)
    assert got.x == area.x and got.x2 < W // 2 - int(0.01 * W)


def test_a_heading_inherits_its_anchor_from_the_master():
    from verstka.rendering.synth import _title_anchor

    prs = Presentation()
    master = prs.slide_master
    for ph in master.placeholders:
        if "TITLE" in str(ph.placeholder_format.type):
            ph.text_frame._txBody.find(q("a:bodyPr")).set("anchor", "b")
    layout = prs.slide_layouts[5]  # «Title Only»: its own title sets no anchor
    for ph in layout.placeholders:
        if "TITLE" in str(ph.placeholder_format.type):
            bp = ph.text_frame._txBody.find(q("a:bodyPr"))
            if bp is not None and bp.get("anchor"):
                del bp.attrib["anchor"]
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    assert _title_anchor(title) == "b"


def test_a_mockup_shown_alone_keeps_its_first_sentence():
    from verstka.planning import faithful as F

    r = F.read_slide(8, "Панель наблюдения", "Панель управления объединяет текущие показания. Пользователь быстро видит, какие зоны требуют внимания.\n\n"
                     "[Место под изображение: экран панели управления теплицей]")
    room = F.Room()
    room.kinds.add("mockup")
    room.side_text["mockup"] = 0  # the template's screen stands alone under its heading
    forms = F.forms_of(r, room, "visual")
    f = forms[0]
    assert f.kind == "mockup" and not f.slide.content.bullets
    assert f.slide.subtitle and f.slide.subtitle.startswith("Панель управления объединяет текущие показания")


def test_cards_keep_what_a_composed_card_holds():
    from verstka.planning import faithful as F

    paras = ["Условия внутри теплицы меняются в течение дня. Солнечным утром температура может быстро расти, после "
             "проветривания — снижаться, а влажность грунта зависит от полива.",
             "Ручной полив часто выполняется по расписанию. Такой подход прост, но не всегда соответствует потребностям растений.",
             "В учебном проекте уход занимал в среднем 45 минут в день. В это время входили осмотр растений и полив.",
             "Автоматизация позволяет перейти к непрерывному наблюдению. Человек продолжает отвечать за растения."]
    r = F.read_slide(3, "Почему теплице нужен контроль", "\n\n".join(paras))
    room = F.Room()
    room.kinds.add("cards")
    room.cells["cards"] = [(4, 4, 90)]  # the template's card samples hold ninety characters a cell
    form = next(f for f in F.forms_of(r, room, "structured") if f.kind == "cards")
    texts = [i.text for i in form.slide.content.items]
    assert len(texts) == 4 and all(t.count(". ") >= 1 for t in texts), texts  # two sentences each, not the first alone


def test_a_cover_of_the_authors_text_gets_no_invented_date():
    from types import SimpleNamespace

    from verstka.rendering.clone import _bookend_texts, _deck_date
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    for verbatim in (True, False):
        cover = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Умная теплица", subtitle="Автоматизация полива", verbatim=verbatim)
        ctx = SimpleNamespace(outline=DeckOutline(title="Умная теплица", slides=[cover], language="ru"))
        footer = _bookend_texts(ctx, cover)[3]
        assert footer == (None if verbatim else _deck_date("ru"))
    # a cover whose own lines date it (an improved deck's «Октябрь 2026») is never dated twice
    for line in ("Октябрь 2026", "Отчёт за окт. 2026", "10.2026", "October 2026"):
        dated = OutlineSlide(id="sl1", kind=PatternKind.title, headline="Умная теплица", subtitle=line)
        ctx = SimpleNamespace(outline=DeckOutline(title="Умная теплица", slides=[dated], language="ru"))
        assert _bookend_texts(ctx, dated)[3] is None, line


def test_the_fixes_of_a_finished_deck_find_their_slides_by_order(tmp_path):
    from verstka.audit.autofix import _ir_slides
    from verstka.schemas.audit import AuditReport
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    prs = Presentation()
    for _ in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Текст"
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))
    outline = DeckOutline(title="Т", slides=[OutlineSlide(id="a", kind=PatternKind.bullets, headline="Один"), OutlineSlide(id="b", kind=PatternKind.bullets, headline="Два")])
    report = AuditReport(deck=str(deck), template_id="t", strategy="structured")
    assert [s.outline_id for s in _ir_slides(report, outline)] == ["a", "b"]
    assert [s.outline_id for s in _ir_slides(report)] == [None, None]


def test_the_audit_reads_inherited_spacing_and_paragraph_gaps(tmp_path):
    from verstka.audit.checks.common import text_elements, text_height_needed_pt
    from verstka.audit.ir import build_deck_ir

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(3))
    tf = tb.text_frame
    tf.text = "Первый"
    p2 = tf.add_paragraph()
    p2.text = "Второй"
    for p in tf.paragraphs:
        p.runs[0].font.size = Pt(20)
    lst = tf._txBody.find(q("a:lstStyle"))
    lvl = etree.SubElement(lst, q("a:lvl1pPr"))
    ln = etree.SubElement(lvl, q("a:lnSpc"))
    etree.SubElement(ln, q("a:spcPct")).set("val", "150000")
    sb = etree.SubElement(lvl, q("a:spcBef"))
    etree.SubElement(sb, q("a:spcPts")).set("val", "1000")
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))
    ir = build_deck_ir(deck, with_images=False)
    e = next(x for x in text_elements(ir.slides[0]) if "Первый" in x.text)
    assert all(abs((p.line_spacing or 0) - 1.5) < 1e-6 for p in e.paragraphs)
    need, lines = text_height_needed_pt(e)
    assert lines == 2 and abs(need - (2 * 20 * 1.2 * 1.5 + 10)) < 0.01


def test_a_cover_is_dated_by_a_month_and_a_year_only():
    """The month stems of 3.0 dated «Маркетинг 2026», «Маркетплейсов 2026», «Декларация 2025» and «Версия 1.2025»: those
    covers lost their date. A month is one of its word forms or a dotted abbreviation, a numeric month has two digits
    and is no version number — and a dated line the two-phrase heading moved over the title (the kicker) dates it too."""
    from types import SimpleNamespace

    from verstka.rendering.clone import _bookend_texts, _dated, _deck_date
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    dated = ["Октябрь 2026", "октября 2026", "в октябре 2026", "Отчёт за окт. 2026", "Сент. 2026", "май 2026", "Мая 2026", "марта 2026",
             "10.2026", "03/2026", "01.10.2026", "October 2026", "Oct 2026", "15 января 2026"]
    undated = ["Маркетинг 2026", "Отчёт для маркетплейсов 2026", "Декларация 2025", "Версия 1.2025", "Версия 10.2025", "v10.2025", "Маяк 2026",
               "Майские праздники 2026", "Августовский отчёт 2026", "1.2025", "Стратегия 2026", "2025/2026", "Decision 2025"]
    assert [t for t in dated if not _dated(t)] == []
    assert [t for t in undated if _dated(t)] == []

    def cover(headline: str, subtitle: str):
        o = OutlineSlide(id="sl1", kind=PatternKind.title, headline=headline, subtitle=subtitle)
        return _bookend_texts(SimpleNamespace(outline=DeckOutline(title=headline, slides=[o], language="ru")), o)

    for line in ("Маркетинг 2026", "Отчёт для маркетплейсов 2026", "Декларация 2025", "Версия 1.2025"):
        assert cover("Умная теплица", line)[3] == _deck_date("ru"), line
    head, sub, kicker, footer = cover("Умная теплица: итоги пилота", "Октябрь 2026")
    assert (head, kicker, footer) == ("Умная теплица", "Октябрь 2026", None), (head, sub, kicker, footer)


def test_running_text_is_set_on_the_templates_scale_or_not_in_the_sample():
    """VK Education s10: two columns at 9.9 pt under the template's smallest 12 pt (the audit's size_not_in_scale): a box
    whose text fits only that far under the scale is no room for it — the gate takes a roomier sample or the composer.
    Within the audit's tolerance (a size reads as the template's) and for headings nothing changes."""
    from types import SimpleNamespace

    from verstka.rendering.clone import SCALE_TOL, _SlideCtx
    from verstka.schemas.common import BboxFrac, SlotRole
    from verstka.schemas.template import Capacity, Slot, SlotStyle

    ctx = SimpleNamespace(grow_scale=[12.0, 14.0, 16.0, 20.0])

    def slot(role):
        return Slot(id="s", role=role, shape_id="s", bbox=BboxFrac(x=0.1, y=0.1, w=0.3, h=0.2), style=SlotStyle(size_pt=16.0), capacity=Capacity(max_chars=100, max_lines=4))

    assert _SlideCtx.off_scale(ctx, slot(SlotRole.card_body), 9.9)
    assert _SlideCtx.off_scale(ctx, slot(SlotRole.bullet_list), 12.0 - SCALE_TOL - 0.1)
    assert not _SlideCtx.off_scale(ctx, slot(SlotRole.body), 12.0 - SCALE_TOL + 0.05)
    assert not _SlideCtx.off_scale(ctx, slot(SlotRole.title), 9.9)
    assert not _SlideCtx.off_scale(ctx, None, 9.9)


def _ruled_template(out):
    """A grid sample as Классика draws it: three rows, a heading column (24 pt) beside a text column (18 pt), the rules
    between the rows drawn by the slide's layout (the gate reads the rules of the slide itself only)."""
    from pptx.enum.shapes import MSO_CONNECTOR

    W, H = 12192000, 6858000
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Название"
    s1.placeholders[1].text = "Подзаголовок"
    layout = prs.slide_layouts[5]
    s2 = prs.slides.add_slide(layout)
    s2.shapes.title.text = "Заголовок слайда"
    rows = (0.19, 0.45, 0.72)
    for y, label in zip(rows, ("Первая строка", "Вторая строка", "Третья строка")):
        head = s2.shapes.add_textbox(Emu(int(W * 0.055)), Emu(int(H * y)), Emu(int(W * 0.238)), Emu(int(H * 0.175)))
        head.text_frame.word_wrap = True
        head.text_frame.text = label
        head.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        body = s2.shapes.add_textbox(Emu(int(W * 0.338)), Emu(int(H * y)), Emu(int(W * 0.623)), Emu(int(H * 0.175)))
        body.text_frame.word_wrap = True
        body.text_frame.text = "Текст строки сетки на две строки"
        body.text_frame.add_paragraph().text = "Вторая строка текста"
        for p in body.text_frame.paragraphs:
            p.runs[0].font.size = Pt(18)
    s3 = prs.slides.add_slide(prs.slide_layouts[1])  # (the template's small steps: a note, a caption, a body)
    s3.shapes.title.text = "Текст"
    for i, size in enumerate((12, 14, 16)):
        tb = s3.shapes.add_textbox(Emu(int(W * 0.1)), Emu(int(H * (0.3 + 0.15 * i))), Emu(int(W * 0.6)), Emu(int(H * 0.1)))
        tb.text_frame.text = "Пояснение к слайду"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    tree = layout.shapes._spTree
    for y in (0.40, 0.67):
        ln = s2.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(W * 0.03)), Emu(int(H * y)), Emu(int(W * 0.97)), Emu(int(H * y)))
        ln.line.width = Pt(1)
        tree.append(ln._element)  # (moved to the layout: drawn under every slide of it)
    prs.save(out)
    return out


def test_a_text_never_grows_past_a_rule_of_the_templates_grid(tmp_path):
    """Классика s10: the bodies of a ruled grid grew from 94 to 129 and 168 pt over the hairline under their row — the
    seventh line struck through by it (a rule has no area: the room under a box never stopped at it, and the gate reads
    the slide's own rules only). A box grows down to the rule under it, never past it: the text steps down the scale
    inside, or the box says it overflows and the gate turns the sample down."""
    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import render_clone
    from verstka.rendering.deck import DeckBuilder
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.layout import LayoutSlide
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent, SlideItem

    tpl = _ruled_template(tmp_path / "ruled.pptx")
    manifest = analyze_template(tpl, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    pattern = next(p for p in manifest.patterns if p.source_slide == 2)
    now = ["100 покупок в день", "Средний чек 300 рублей", "20% чеков с едой", "Простой после 15:00", "Списания 27 000 рублей", "25% гостей возвращаются", "Прибыль 120 000 рублей"]
    later = ["115 покупок в день", "Средний чек 330 рублей", "30% чеков с едой", "Дневные предложения", "Списания 18 000 рублей", "35% гостей возвращаются", "Прибыль 254 795 рублей"]
    s = OutlineSlide(id="s2", kind=K.comparison, headline="Сейчас и через полгода", content=SlideContent(columns=[SlideItem(title="Сейчас", bullets=now), SlideItem(title="Через полгода", bullets=later)]))
    outline = DeckOutline(title="Т", slides=[s])
    slide, warnings = render_clone(DeckBuilder(tpl), LayoutSlide(outline_id="s2", mode="clone", pattern_id=pattern.id), s, pattern, manifest, ws, outline)
    rules = sorted(sh.top for sh in slide.slide_layout.shapes if sh.height is not None and sh.height <= 0.006 * 6858000 and sh.width >= 0.3 * 12192000)
    assert len(rules) == 2
    grown = []
    for sh in slide.shapes:
        if not sh.has_text_frame or not any(t in sh.text_frame.text.replace("\u00a0", " ") for t in now + later):
            continue
        under = [y for y in rules if y > sh.top + 0.01 * 6858000]
        if under:
            grown.append(sh.top + sh.height)
            assert sh.top + sh.height <= under[0], (sh.text_frame.text[:30], sh.top / 12700, sh.height / 12700, under[0] / 12700)
    assert grown, [sh.text_frame.text[:30] for sh in slide.shapes if sh.has_text_frame]


def test_the_audit_sees_a_hairline_of_the_layout_through_a_text(tmp_path):
    """A connector or a line of the layout has no area: the template layers of the audit's IR left it out, and Классика's
    seventh line struck through by its grid passed every check. The audit reads the layout's hairlines from the deck and
    flags the text they cross, line by line (a rule through a block's last line lies in the block's lower margin)."""
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext, load_builtin_checks, run_checks
    from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

    W, H = 12192000, 6858000
    tpl = _ruled_template(tmp_path / "ruled.pptx")
    prs = Presentation(str(tpl))
    slide = prs.slides[1]
    boxes = [sh for sh in slide.shapes if sh.has_text_frame and "Текст строки" in sh.text_frame.text]
    first = min(boxes, key=lambda sh: sh.top)
    # grown over the hairline under its row (at 40 % of the height), a line of it right across the hairline
    first.top, first.height = Emu(int(first.top) + int(11 * 12700)), Emu(int(H * 0.30))
    tf = first.text_frame
    for i in range(5):
        tf.add_paragraph().text = f"Ещё одна строка текста номер {i + 1}"
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(18)
    for sh in boxes:
        if sh is not first:
            sh.text_frame.text = "Короткий текст"
    prs.save(str(tmp_path / "deck.pptx"))
    load_builtin_checks()
    ir = build_deck_ir(tmp_path / "deck.pptx", with_images=False)
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(), patterns=[], n_slides=3)
    got, _ = run_checks(AuditContext(ir=ir, manifest=man), only={"content_over_art"})
    hits = [i for i in got if i.slide == 2 and (i.details or {}).get("hairline")]
    assert len(hits) == 1 and "Текст строки" in hits[0].message and hits[0].severity == "warn", [(i.slide, i.message, i.details) for i in got]
    # the rows that keep inside their cells are not flagged
    assert not [i for i in got if "Короткий текст" in i.message]
