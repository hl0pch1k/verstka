"""Any file a person calls a template reaches the analysis as a .pptx (verstka/ingest/convert.py): PowerPoint templates
(.potx), macro-enabled decks (.pptm), slide shows (.ppsx), themes (.thmx) are repackaged; legacy .ppt and OpenDocument
.odp go through LibreOffice; PDF, Keynote, Word and the like are refused with what to upload instead. A template
with (almost) no slides of its own gets sample slides from its layouts (verstka/ingest/samples.py)."""

from __future__ import annotations

import io
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from verstka.ingest.convert import CT_PRESENTATION, ConvertError, check_upload, sniff, to_pptx
from verstka.ingest.samples import MIN_SLIDES, ensure_samples, layout_kind, plan_samples

from test_api import _wait, client  # noqa: F401  (shared fixture)

CT_TEMPLATE = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
CT_SHOW = "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml"
CT_MACRO = "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"


def _deck(tmp: Path, n_slides: int = 2) -> Path:
    prs = Presentation()
    for i in range(n_slides):
        s = prs.slides.add_slide(prs.slide_layouts[1 if i else 0])
        s.shapes.title.text = f"Слайд {i + 1}"
    p = tmp / "deck.pptx"
    prs.save(p)
    return p


def _retyped(src: Path, dst: Path, ct: str, macro: bool = False) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            d = zin.read(n)
            if n == "[Content_Types].xml":
                s = re.sub(r'(PartName="/ppt/presentation.xml" ContentType=")[^"]+', lambda m: m.group(1) + ct, d.decode())
                if macro:
                    s = s.replace("</Types>", '<Override PartName="/ppt/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
                d = s.encode()
            if n == "ppt/_rels/presentation.xml.rels" and macro:
                d = d.decode().replace("</Relationships>", '<Relationship Id="rIdV" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>').encode()
            zout.writestr(n, d)
        if macro:
            zout.writestr("ppt/vbaProject.bin", b"\xd0\xcf\x11\xe0macro")
    return dst


def _main_ct(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        ct = z.read("[Content_Types].xml").decode()
    return re.search(r'PartName="/ppt/presentation.xml" ContentType="([^"]+)"', ct).group(1)


@pytest.mark.parametrize("ext,ct,macro", [(".potx", CT_TEMPLATE, False), (".ppsx", CT_SHOW, False), (".pptm", CT_MACRO, True)])
def test_office_variants_are_repackaged_as_a_pptx(tmp_path, ext, ct, macro):
    src = _retyped(_deck(tmp_path), tmp_path / f"Бренд{ext}", ct, macro)
    assert sniff(src) == "pptx" and check_upload(src) == "pptx"
    out = to_pptx(src)
    assert out.name == "Бренд.pptx" and out.exists()
    assert _main_ct(out) == CT_PRESENTATION
    with zipfile.ZipFile(out) as z:
        assert not any(n.startswith("ppt/vbaProject") for n in z.namelist())  # macros never kept
        assert "vbaProject" not in z.read("ppt/_rels/presentation.xml.rels").decode()
    assert len(Presentation(out).slides) == 2


def test_a_plain_pptx_is_analysed_as_it_is(tmp_path):
    src = _deck(tmp_path)
    assert to_pptx(src) == src


def _thmx(tmp: Path) -> Path:
    """A theme package the way PowerPoint writes one: the presentation, master, layouts and theme under theme/."""
    src = tmp / "base.pptx"
    Presentation().save(src)
    dst = tmp / "Тема.thmx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for n in zin.namelist():
            d = zin.read(n)
            if n.startswith("ppt/") and not n.startswith(("ppt/presProps", "ppt/viewProps", "ppt/tableStyles")) and "printerSettings" not in n:
                zout.writestr("theme/" + n[4:], d)
            elif n == "[Content_Types].xml":
                s = d.decode().replace('PartName="/ppt/', 'PartName="/theme/')
                zout.writestr(n, s)
            elif n == "_rels/.rels":
                zout.writestr(n, d.decode().replace('Target="ppt/presentation.xml"', 'Target="theme/theme/themeManager.xml"'))
        zout.writestr("theme/theme/themeManager.xml", "<a:themeManager xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\"/>")
    return dst


def test_a_theme_becomes_a_presentation_with_its_layouts(tmp_path):
    src = _thmx(tmp_path)
    assert sniff(src) == "thmx"
    out = to_pptx(src)
    prs = Presentation(out)
    assert len(prs.slides) == 0 and len(prs.slide_layouts) >= 9


@pytest.mark.parametrize(
    "name,payload,words",
    [
        ("brochure.pdf", b"%PDF-1.4\n%x\n", "PDF"),
        ("photo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64, "картинка"),
        ("junk.pptx", b"x" * 256, "не похож"),
    ],
)
def test_what_is_not_a_presentation_is_refused_with_what_to_upload(tmp_path, name, payload, words):
    p = tmp_path / name
    p.write_bytes(payload)
    with pytest.raises(ConvertError) as e:
        check_upload(p)
    assert words in str(e.value)


def test_keynote_word_and_plain_zips_say_what_to_do(tmp_path):
    key = tmp_path / "deck.key"
    with zipfile.ZipFile(key, "w") as z:
        z.writestr("Index/Document.iwa", b"x")
    doc = tmp_path / "notes.docx"
    with zipfile.ZipFile(doc, "w") as z:
        z.writestr("word/document.xml", b"<w/>")
    arc = tmp_path / "pack.zip"
    with zipfile.ZipFile(arc, "w") as z:
        z.writestr("a.txt", b"x")
    for p, words in ((key, "Keynote"), (doc, "Word"), (arc, "архив")):
        with pytest.raises(ConvertError) as e:
            check_upload(p)
        assert words in str(e.value)


@pytest.mark.skipif(shutil.which("soffice") is None and not Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists(), reason="LibreOffice not installed")
def test_legacy_ppt_goes_through_libreoffice(tmp_path):
    import subprocess

    from verstka.ingest.render import find_soffice

    src = _deck(tmp_path)
    subprocess.run([find_soffice(), f"-env:UserInstallation={(tmp_path / 'lo').as_uri()}", "--headless", "--convert-to", "ppt", "--outdir", str(tmp_path), str(src)], capture_output=True, timeout=180)
    ppt = tmp_path / "deck.ppt"
    assert ppt.exists()
    named = tmp_path / "Старый шаблон.ppt"
    ppt.rename(named)
    assert sniff(named) == "legacy"
    out = to_pptx(named)
    assert out.name == "Старый шаблон.pptx" and len(Presentation(out).slides) == 2


# ---------------------------------------------------------------------------------------- samples from layouts


def test_standard_layouts_are_read_by_kind():
    prs = Presentation()
    kinds = {l.name: layout_kind(l) for l in prs.slide_layouts}
    assert kinds["Title Slide"] == "cover"
    assert kinds["Title and Content"] == "content"
    assert kinds["Section Header"] == "section"
    assert kinds["Two Content"] == "two"
    assert kinds["Comparison"] == "comparison"
    assert kinds["Title Only"] == "title_only"
    assert kinds["Content with Caption"] == "caption"
    assert kinds["Blank"] is None
    assert kinds["Title and Vertical Text"] is None and kinds["Vertical Title and Text"] is None


def test_a_template_without_slides_gets_samples_from_its_layouts(tmp_path):
    src = tmp_path / "source.pptx"
    Presentation().save(src)
    made = ensure_samples(src)
    prs = Presentation(src)
    assert made == len(prs.slides) >= 6
    assert (tmp_path / "original.pptx").exists() and len(Presentation(tmp_path / "original.pptx").slides) == 0
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text]
    assert "Название презентации" in texts and "Заголовок слайда" in texts
    # a new analysis starts from the original again: the same samples, not twice as many
    assert ensure_samples(src) == made and len(Presentation(src).slides) == made


def test_a_template_with_a_cover_only_gets_the_other_kinds(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Обложка"
    src = tmp_path / "source.pptx"
    prs.save(src)
    kinds = [layout_kind(l) for l in plan_samples(Presentation(src))]
    assert "cover" not in kinds and kinds[:3] == ["content", "section", "two"]
    ensure_samples(src)
    assert len(Presentation(src).slides) == 1 + len(kinds)


def test_a_template_with_enough_slides_is_left_as_it_is(tmp_path):
    prs = Presentation()
    for _ in range(MIN_SLIDES):
        prs.slides.add_slide(prs.slide_layouts[1])
    src = tmp_path / "source.pptx"
    prs.save(src)
    assert ensure_samples(src) == 0 and not (tmp_path / "original.pptx").exists()


def test_the_analysis_of_a_bare_template_finds_its_layouts(tmp_path):
    from verstka.analysis.manifest import analyze_template

    src = tmp_path / "Тема.pptx"
    Presentation().save(src)
    m = analyze_template(src, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    kinds = {p.kind.value for p in m.patterns}
    assert m.n_slides >= 6 and {"title", "bullets", "section"} <= kinds
    assert m.source_file == "Тема.pptx"


# ---------------------------------------------------------------------------------------- the upload


def test_a_potx_upload_is_analysed_under_its_own_name(client, tmp_path):
    c, mod = client
    potx = _retyped(_deck(tmp_path, 1), tmp_path / "Фирменный.potx", CT_TEMPLATE)
    with open(potx, "rb") as f:
        r = c.post("/api/templates", files={"file": ("Фирменный.potx", f, "application/octet-stream")}, data={"use_models": "false"})
    assert r.status_code == 200, r.text
    job = _wait(c, r.json()["job_id"])
    assert job["status"] == "done", job
    card = next(t for t in c.get("/api/templates").json() if t["template_id"] == job["result"]["template_id"])
    assert card["source_file"] == "Фирменный.potx"
    assert job["result"]["n_patterns"] >= 4  # the one-slide template got samples from its layouts


def test_an_upload_that_is_not_a_presentation_says_what_to_upload(client):
    c, mod = client
    r = c.post("/api/templates", files={"file": ("report.pdf", io.BytesIO(b"%PDF-1.4\n%x\n"), "application/pdf")}, data={"use_models": "false"})
    assert r.status_code == 400 and "PowerPoint" in r.json()["detail"]
    r = c.post("/api/templates", files={"file": ("empty.pptx", io.BytesIO(b""), "application/octet-stream")}, data={"use_models": "false"})
    assert r.status_code == 400 and "пустой" in r.json()["detail"]


def test_a_password_protected_file_says_so(tmp_path):
    p = tmp_path / "secret.pptx"
    p.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 512 + "EncryptionInfo".encode("utf-16-le") + b"\0" * 64)
    assert sniff(p) == "encrypted"
    with pytest.raises(ConvertError) as e:
        check_upload(p)
    assert "паролем" in str(e.value)
