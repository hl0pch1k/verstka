"""The clone path after the swarm review of 2026-10-05: a sample's say over its master stays with its clones, a word
bound by no-break spaces is measured as one word, a clone that cannot take the slide is one more sample tried."""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches


def test_a_clone_keeps_its_samples_hidden_master_graphics(tmp_path):
    from verstka.rendering.deck import DeckBuilder

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Образец"
    s._element.set("showMasterSp", "0")
    src = tmp_path / "tpl.pptx"
    prs.save(str(src))
    b = DeckBuilder(src)
    new = b.clone_slide(1)
    assert new._element.get("showMasterSp") == "0"


W, H = 12192000, 6858000
PT = 12700


def _slide(boxes):
    """A one-slide deck with text boxes [(x, y, w, h, text, size)]: x, y in fractions of a 16:9 slide, w, h in EMU."""
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for x, y, w, h, text, size in boxes:
        tb = slide.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w)), Emu(int(h)))
        tb.text_frame.text = text
        tb.text_frame.word_wrap = True
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
    return slide


def _gate(slide, native=()):
    from verstka.rendering.clone_check import content_trouble
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import OutlineSlide, SlideContent
    from verstka.schemas.template import Family, Pattern, SlideSize, TemplateManifest, Tokens

    if native:
        slide.__dict__["_verstka_native"] = set(native)
    m = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(), patterns=[], n_slides=1)
    pat = Pattern(id="p1", source_slide=1, kind=PatternKind.bullets, family=Family.light, slots=[])
    o = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Заголовок", content=SlideContent(bullets=["Пункт"]))
    return content_trouble(slide, o, pat, m, []) or ""


def _between(size: float) -> float:
    """A line width (pt) «автоматизации» fits and «до автоматизации» bound by a no-break space does not."""
    from verstka.rendering.fonts import text_width_pt

    one = text_width_pt("автоматизации", None, size, False)
    pair = text_width_pt("до автоматизации", None, size, False)
    assert pair > one * 1.1
    return (one + pair) / 2


def test_words_bound_by_a_no_break_space_are_measured_as_one_word():
    """MyBrand p12: «… до автоматизации» at 30 pt fitted a 250 pt box by the check's own split and rendered as
    «до автоматизаци / и» — the renderer never breaks at a no-break space, so the pair is one word to it."""
    from verstka.rendering.fit import _measure

    width = _between(30.0)
    lines, broken = _measure(["уход за теплицей в день до автоматизации"], None, 30.0, False, width)
    assert broken
    assert not _measure(["уход за теплицей в день до автоматизации"], None, 30.0, False, width)[1]
    # the gate reads it so too
    box = (width + 14.4) * PT + 2
    assert "breaks mid-word" in _gate(_slide([(0.05, 0.3, box, 1.5 * 72 * PT, "Уход за теплицей в день до автоматизации", 30)]))
    assert "breaks mid-word" not in _gate(_slide([(0.05, 0.3, box, 1.5 * 72 * PT, "Уход за теплицей в день до автоматизации", 30)]))


def test_a_paragraph_margin_narrows_its_lines_in_the_gate():
    from verstka.analysis.xmlns import q
    from verstka.rendering.fonts import text_width_pt

    one = text_width_pt("автоматизации", None, 30.0, False)
    box = (one * 1.05 + 14.4) * PT
    s = _slide([(0.05, 0.3, box, 2 * 72 * PT, "Этап автоматизации", 30)])
    assert "breaks mid-word" not in _gate(s)
    s = _slide([(0.05, 0.3, box, 2 * 72 * PT, "Этап автоматизации", 30)])
    p = s.shapes[0]._element.find(".//" + q("a:p"))
    ppr = p.find(q("a:pPr"))
    if ppr is None:
        from lxml import etree

        ppr = etree.SubElement(p, q("a:pPr"))
        p.insert(0, ppr)
    ppr.set("marL", str(int(0.2 * one * PT)))
    assert "breaks mid-word" in _gate(s)


def test_the_design_systems_own_texts_are_checked_for_broken_words():
    """Coffee chart forced onto VK Education p50: the side text drawn with the pie broke as «рентабельно / сть» and
    the gate let it pass — its sizes are the design system's, its words are not exempt."""
    width = _between(20.0)
    box = (width + 14.4) * PT + 2
    s = _slide([(0.05, 0.3, box, 1.5 * 72 * PT, "Рост до автоматизации", 20)])
    sid = s.shapes[0]._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").get("id")
    assert "breaks mid-word" in _gate(s, native=[sid])


def test_a_text_in_a_scaled_group_is_measured_on_the_slide():
    """VK Tech's team sample keeps its cells in a group drawn at 0.376 of its size: a role box 148 pt wide in the
    group is 56 pt on the slide. Its text is fitted — and gated — at the 56 pt."""
    from lxml import etree

    from verstka.analysis.xmlns import q
    from verstka.rendering.clone import _text_frame

    s = _slide([(0.05, 0.3, 148 * PT, 40 * PT, "Руководитель", 14)])
    tree = s._element.cSld.find(q("p:spTree"))
    sp = s.shapes[0]._element
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    grp = etree.fromstring(
        f'<p:grpSp xmlns:p="{P}" xmlns:a="{A}"><p:nvGrpSpPr><p:cNvPr id="90" name="team"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        f'<p:grpSpPr><a:xfrm><a:off x="600000" y="2000000"/><a:ext cx="{int(0.376 * 2000000)}" cy="{int(0.376 * 1000000)}"/>'
        '<a:chOff x="0" y="0"/><a:chExt cx="2000000" cy="1000000"/></a:xfrm></p:grpSpPr></p:grpSp>'
    )
    tree.append(grp)
    grp.append(sp)
    sp.find(q("p:spPr") + "/" + q("a:xfrm") + "/" + q("a:off")).set("x", "0")
    sp.find(q("p:spPr") + "/" + q("a:xfrm") + "/" + q("a:off")).set("y", "0")
    b = (0, 0, int(148 * PT), int(40 * PT))
    f = _text_frame(sp, b)
    assert abs(f.w - 0.376 * 148 * PT) < PT and abs(f.h - 0.376 * 40 * PT) < PT
    # «Руководитель» at 14 pt is wider than the 56 pt the box has on the slide
    assert "breaks mid-word" in _gate(s)


def test_a_short_last_row_of_widened_cells_stands_centred():
    """VK WorkSpace p12, five items in its 2×4 grid: rows of three wider cards, the two of the last row centred under
    them — not a hole at the bottom right."""
    from verstka.rendering.groups import _rebalance_grid, cell_bbox
    from verstka.schemas.common import Bbox

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    w_s, h_s = prs.slide_width, prs.slide_height
    boxes, cells = [], []
    for r in range(2):
        for c in range(4):
            x, y, w, h = int(0.035 * w_s + c * 0.235 * w_s), int(0.3 * h_s + r * 0.3 * h_s), int(0.2 * w_s), int(0.25 * h_s)
            boxes.append(Bbox(x=x, y=y, w=w, h=h))
            if len(cells) < 5:
                cells.append([s.shapes.add_shape(1, x, y, w, h)._element])
    _rebalance_grid(cells, [[] for _ in cells], boxes, w_s, h_s)
    got = [cell_bbox(c) for c in cells]
    assert len({b.y for b in got[:3]}) == 1 and len({b.y for b in got[3:]}) == 1 and got[3].y > got[0].y
    x0, x1 = boxes[0].x, boxes[3].x + boxes[3].w
    mid = (got[3].x + got[4].x + got[4].w) / 2
    assert abs(mid - (x0 + x1) / 2) < 0.01 * w_s
    assert abs(got[0].x - x0) < 0.005 * w_s  # the full row keeps the grid's left edge


def test_the_texts_of_one_row_of_cards_start_at_one_height(tmp_path):
    """VK WorkSpace p6, three cards: the text under the one-line heading «Датчики» started 46 pt under the texts of
    its neighbours, whose headings took two lines — each card had stacked its text under its own heading."""
    from test_review_clone import _manifest, _new_prs, _pattern, _slot, _tb

    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import _SlideCtx, _stack_cells
    from verstka.rendering.deck import DeckBuilder, element_bbox
    from verstka.schemas.common import PatternKind, SlotRole
    from verstka.schemas.outline import DeckOutline

    prs, blank = _new_prs()
    s = prs.slides.add_slide(blank)
    heads = ["Датчики", "Проверка показаний", "Правила полива"]
    shapes = []
    for k, h in enumerate(heads):
        x = 0.05 + k * 0.31
        t = _tb(s, x, 0.30, 0.25, 0.15, h, size=24, bold=True)
        b = _tb(s, x, 0.47, 0.25, 0.30, "Система сравнивает показания и включает полив по расписанию.", size=16)
        shapes.append((t, b))
    pptx = tmp_path / "cards.pptx"
    prs.save(pptx)
    slots = [(_slot(t, SlotRole.card_title, size=24, bold=True), _slot(b, SlotRole.card_body, size=16)) for t, b in shapes]
    pattern = _pattern("p1", 1, PatternKind.cards, [x for pair in slots for x in pair])
    builder = DeckBuilder(pptx)
    slide = builder.clone_slide(1)
    ctx = _SlideCtx(builder, slide, pattern, _manifest([pattern]), TemplateWorkspace.create(pptx, tmp_path / "ws"), DeckOutline(title="T", slides=[]))
    stacks = [(ctx.els[ts.shape_id], ts, [(ctx.els[bs.shape_id], bs)]) for ts, bs in slots]
    assert _stack_cells(ctx, stacks)
    tops = [element_bbox(ctx.els[bs.shape_id])[1] for _, bs in slots]
    assert max(tops) - min(tops) <= 12700, tops


def test_the_labels_under_compound_figures_share_one_size():
    from lxml import etree

    from verstka.analysis.xmlns import q
    from verstka.rendering.clone import _harmonize_compound_labels

    P, A = "http://schemas.openxmlformats.org/presentationml/2006/main", "http://schemas.openxmlformats.org/drawingml/2006/main"

    def box(lab_sz: int):
        return etree.fromstring(
            f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:txBody><a:bodyPr/>'
            f'<a:p><a:r><a:rPr sz="5400" b="1"/><a:t>87%</a:t></a:r></a:p>'
            f'<a:p><a:r><a:rPr sz="{lab_sz}"/><a:t>подпись</a:t></a:r><a:endParaRPr sz="{lab_sz}"/></a:p></p:txBody></p:sp>'
        )

    els = [box(3000), box(3300), box(2900), box(3300)]
    _harmonize_compound_labels(els)
    for e in els:
        fig, lab = e.findall(".//" + q("a:p"))
        assert fig.find(".//" + q("a:rPr")).get("sz") == "5400"
        assert {r.get("sz") for r in lab.iter(q("a:rPr"))} == {"2900"} and lab.find(q("a:endParaRPr")).get("sz") == "2900"


def test_an_icon_goes_with_the_tile_that_framed_it(tmp_path):
    """VK WorkSpace p4 on a slide of the author's text: the sample's pictograms went, their black tiles with a blue
    outline stayed as four empty frames beside the texts. The card holding the text stays."""
    from PIL import Image
    from test_review_clone import _manifest, _new_prs, _pattern, _slot, _tb

    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import _drop_icon, _SlideCtx
    from verstka.rendering.deck import DeckBuilder
    from verstka.schemas.common import PatternKind, SlotRole
    from verstka.schemas.outline import DeckOutline

    prs, blank = _new_prs()
    s = prs.slides.add_slide(blank)
    card = _tb(s, 0.05, 0.30, 0.40, 0.25, "", fill="F2F2F2")
    tile = _tb(s, 0.07, 0.33, 0.06, 0.1, "", fill="000000")
    png = tmp_path / "i.png"
    Image.new("RGBA", (64, 64), (0, 119, 255, 255)).save(png)
    from pptx.util import Emu

    icon = s.shapes.add_picture(str(png), Emu(int(0.08 * 12192000)), Emu(int(0.345 * 6858000)), Emu(int(0.04 * 12192000)), Emu(int(0.07 * 6858000)))
    text = _tb(s, 0.15, 0.33, 0.28, 0.2, "Резервное питание датчиков", size=16)
    pptx = tmp_path / "icon.pptx"
    prs.save(pptx)
    pattern = _pattern("p1", 1, PatternKind.cards, [_slot(text, SlotRole.card_body, size=16)])
    builder = DeckBuilder(pptx)
    slide = builder.clone_slide(1)
    ctx = _SlideCtx(builder, slide, pattern, _manifest([pattern]), TemplateWorkspace.create(pptx, tmp_path / "ws"), DeckOutline(title="T", slides=[]))
    ids = {n: str(sh.shape_id) for n, sh in (("card", card), ("tile", tile), ("icon", icon), ("text", text))}
    _drop_icon(ctx, ctx.els[ids["icon"]])
    left = {str(sh.shape_id) for sh in slide.shapes}
    assert ids["icon"] not in left and ids["tile"] not in left
    assert ids["card"] in left and ids["text"] in left


def test_a_chip_beside_its_line_rides_with_that_line():
    """VK WorkSpace p12: a row of four «chip + text» cells — each chip stands left of its text, outside the cell's
    band along the row. An emptied cell must take its chip with it: four arrow chips floated over nothing."""
    from test_review_clone import _group, _new_prs, _tb

    from verstka.rendering.groups import cell_riders

    prs, blank = _new_prs()
    s = prs.slides.add_slide(blank)
    texts, chips = [], []
    for k in range(4):
        x = 0.085 + k * 0.23
        chips.append(_tb(s, x - 0.045, 0.27, 0.035, 0.06, "", fill="0077FF"))
        texts.append(_tb(s, x, 0.26, 0.16, 0.08, f"Текст {k + 1}", size=14))
    g = _group("g2", [[t] for t in texts], "row", 4)
    cells = [[t._element] for t in texts]
    riders = cell_riders(s, g, cells)
    for k in range(4):
        assert [str(r.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").get("id")) for r in riders[k]] == [str(chips[k].shape_id)]


def test_the_rows_of_a_card_grid_are_anchored_alike(tmp_path):
    """VK Education p8 with five items: the top row's texts (6/5/5 lines) were set to start at the card tops, the
    bottom row's (4/4) stayed centred — two rows of one grid, two alignments."""
    from test_review_clone import _manifest, _new_prs, _pattern, _slot, _tb

    from verstka.analysis.xmlns import q
    from verstka.ingest.workspace import TemplateWorkspace
    from verstka.rendering.clone import _SlideCtx, _top_align_rows
    from verstka.rendering.deck import DeckBuilder
    from verstka.schemas.common import PatternKind, SlotRole
    from verstka.schemas.outline import DeckOutline

    prs, blank = _new_prs()
    s = prs.slides.add_slide(blank)
    long = "Система сравнивает показания датчиков и включает полив по расписанию, когда почва сохнет быстрее обычного."
    short = "Полив по расписанию и по показаниям датчиков."
    texts = [long, short, short, short, short]
    boxes = []
    for k, t in enumerate(texts):
        row, col = divmod(k, 3)
        tb = _tb(s, 0.05 + col * 0.31, 0.25 + row * 0.35, 0.28, 0.3, t, size=16)
        tb.text_frame._txBody.find(q("a:bodyPr")).set("anchor", "ctr")
        boxes.append(tb)
    pptx = tmp_path / "grid.pptx"
    prs.save(pptx)
    pattern = _pattern("p1", 1, PatternKind.cards, [_slot(b, SlotRole.card_body, size=16) for b in boxes])
    builder = DeckBuilder(pptx)
    slide = builder.clone_slide(1)
    ctx = _SlideCtx(builder, slide, pattern, _manifest([pattern]), TemplateWorkspace.create(pptx, tmp_path / "ws"), DeckOutline(title="T", slides=[]))
    els = [ctx.els[str(b.shape_id)] for b in boxes]
    _top_align_rows(ctx, els)
    anchors = [e.find(q("p:txBody") + "/" + q("a:bodyPr")).get("anchor") for e in els]
    assert anchors == ["t"] * 5, anchors


def test_a_clone_table_widens_up_to_the_change_column_beside_it():
    """The visual variant draws a before/after table at 64 % of the area and its change column («+6%») beside it; the
    clone then widened the table to the sample's whole area — under the column."""
    from types import SimpleNamespace

    from pptx.util import Emu

    from verstka.analysis.xmlns import q
    from verstka.rendering.clone import _table_to_width
    from verstka.rendering.deck import element_bbox
    from verstka.schemas.common import Bbox

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tree = s._element.cSld.find(q("p:spTree"))
    before = set(tree)
    box = Bbox(x=int(0.05 * W), y=int(0.25 * H), w=int(0.9 * W), h=int(0.6 * H))
    s.shapes.add_table(3, 3, Emu(box.x), Emu(box.y), Emu(int(0.4 * W)), Emu(int(0.3 * H)))
    side = s.shapes.add_textbox(Emu(int(0.62 * W)), Emu(box.y), Emu(int(0.3 * W)), Emu(int(0.2 * H)))
    side.text_frame.text = "+6%"
    _table_to_width(SimpleNamespace(slide=s, W=W, H=H), before, box)
    gf = next(el for el in tree if el.tag.endswith("graphicFrame"))
    b = element_bbox(gf)
    assert b[0] + b[2] <= int(0.62 * W) - int(0.01 * W)
    assert b[2] > int(0.5 * W)  # widened as far as the column lets it (0.4 → 0.55 of the slide)


def test_a_placeholder_anchored_by_its_layout_says_so_on_the_slide(tmp_path):
    """Office template, compact slide 9: the layout anchors the heading at the foot of its tall box; the gate read
    the text at the box's top and the conclusion strip was drawn over the heading."""
    from verstka.analysis.xmlns import q
    from verstka.rendering.clone_check import ink_box
    from verstka.rendering.deck import DeckBuilder

    prs = Presentation()
    layout = prs.slide_layouts[5]  # «Title Only»
    lay_title = next(ph for ph in layout.placeholders if ph.placeholder_format.type is not None and "TITLE" in str(ph.placeholder_format.type))
    lay_title._element.find(".//" + q("a:bodyPr")).set("anchor", "b")
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Финансовый результат"
    src = tmp_path / "tpl.pptx"
    prs.save(str(src))
    b = DeckBuilder(src)
    new = b.clone_slide(1)
    t = new.shapes.title._element
    assert t.find(q("p:txBody") + "/" + q("a:bodyPr")).get("anchor") == "b"
    from verstka.rendering.deck import element_bbox

    box = element_bbox(t)
    ib = ink_box(t, None)
    assert ib.y + ib.h > box[1] + box[3] - 0.3 * box[3]  # the ink stands at the foot of the box


def test_an_autofixed_cover_never_takes_the_cover_another_variant_opens_on():
    """WorkSpace greenhouse: structured's cover p1 was turned down by the audit and rematched to the next sample of its
    list, p24 — the visual variant's own cover: all three variants opened on one slide. The plan then describes p24."""
    from verstka.audit.autofix import _rematch
    from verstka.matching.matcher import headline_key
    from verstka.schemas.common import PatternKind
    from verstka.schemas.layout import LayoutPlan
    from verstka.schemas.outline import DeckOutline, OutlineSlide

    cover = OutlineSlide(id="t", kind=PatternKind.title, headline="Умная теплица")
    outline = DeckOutline(title="Умная теплица", slides=[cover])
    plan = LayoutPlan.model_validate({"template_id": "t", "strategy": "structured", "slides": [{"outline_id": "t", "mode": "clone", "pattern_id": "p1", "score": 1.2, "reasons": ["обложка шаблона: +0.10"], "alternatives": [["p24", 0.94], ["p3", 0.7]]}]})
    assert _rematch(plan, outline, "t", None, others={headline_key(cover): {"p24"}}) == "rematch → p3"
    ps = plan.for_outline("t")
    assert ps.score == 0.7 and not any("обложка шаблона" in r for r in ps.reasons) and all(a != "p3" for a, _ in ps.alternatives)


def test_twin_samples_of_a_template_never_stand_on_neighbouring_slides():
    """VK Education p26/p27 — one row of coloured shapes and captions drawn twice: slides 4 and 5 showed the same
    design back to back, the penalties compared sample ids only."""
    from verstka.matching.matcher import _assign_content, layout_signature
    from verstka.matching.scorer import ScoreResult
    from verstka.planning.strategies import load_strategies
    from verstka.schemas.common import BboxFrac, Family, PatternKind, SlotRole
    from verstka.schemas.template import Pattern, Slot, SlotStyle

    def pat(pid: str, y: float) -> Pattern:
        t = Slot(id="t", role=SlotRole.title, shape_id="1", bbox=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.1), style=SlotStyle())
        b = Slot(id="b", role=SlotRole.body, shape_id="2", bbox=BboxFrac(x=0.05, y=y, w=0.9, h=0.5), style=SlotStyle())
        return Pattern(id=pid, source_slide=int(pid[1:]), kind=PatternKind.cards, family=Family.light, slots=[t, b], layout_part="slideLayout2.xml")

    a, b, c = pat("p26", 0.3), pat("p27", 0.3), pat("p9", 0.45)
    assert layout_signature(a) == layout_signature(b) != layout_signature(c)
    cands = [(ScoreResult(0.9, [], {}), a), (ScoreResult(0.89, [], {}), b), (ScoreResult(0.8, [], {}), c)]
    entries = [{"index": 0, "key": "x", "cands": cands}, {"index": 1, "key": "y", "cands": cands}]
    st = load_strategies()["structured"]
    got = _assign_content(entries, st, None, {p.id: layout_signature(p) for p in (a, b, c)})
    assert {got[0][2].id, got[1][2].id} in ({"p26", "p9"}, {"p27", "p9"}), (got[0][2].id, got[1][2].id)


def _bare(kind, slots, groups=None, note=None):
    from verstka.schemas.common import Family
    from verstka.schemas.template import Pattern

    return Pattern(id="p1", source_slide=1, kind=kind, family=Family.light, slots=slots, repeat_groups=groups or [], note_slot=note)


def _manifest_of(patterns):
    from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

    return TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(), patterns=patterns, n_slides=len(patterns))


def test_the_scorer_sees_what_the_gate_will_turn_down():
    """A formula is the composer's on every template (the gate turns each sample down) — no sample is planned for it;
    four items on three tiles drawn in the template are merged by the render — a fallback at most; a subtitle with
    no room under the heading is a miss the gate would find."""
    from verstka.matching.matcher import match_outline
    from verstka.matching.scorer import _gate_room, score_pattern
    from verstka.planning.strategies import load_strategies
    from verstka.schemas.common import BboxFrac, PatternKind, SlotRole
    from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem
    from verstka.schemas.template import Capacity, RepeatGroup, Slot, SlotStyle

    st = load_strategies()["structured"]
    title = Slot(id="t", role=SlotRole.title, shape_id="1", bbox=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.12), style=SlotStyle(size_pt=32), capacity=Capacity(max_chars=60, max_lines=2))
    cells = [Slot(id=f"c{k}", role=SlotRole.card_title, shape_id=str(10 + k), bbox=BboxFrac(x=0.05 + 0.31 * k, y=0.3, w=0.28, h=0.4), style=SlotStyle(size_pt=18), capacity=Capacity(max_chars=80, max_lines=3), group_id="g") for k in range(3)]
    g = RepeatGroup(id="g", member_shape_ids=[[c.shape_id] for c in cells], min_n=3, max_n=3, axis="row", cell_bbox=BboxFrac(x=0.05, y=0.3, w=0.28, h=0.4), fixed=True)
    tiles = _bare(PatternKind.cards, [title, *cells], [g])
    four = OutlineSlide(id="s", kind=PatternKind.cards, headline="Почему нужен контроль", content=SlideContent(items=[SlideItem(title=f"Причина {k}") for k in range(4)]))
    assert score_pattern(four, tiles, _manifest_of([tiles]), st).score < st.synth_threshold
    # a formula: composed whatever the template holds
    nums = _bare(PatternKind.stat_row, [title, *[c.model_copy(update={"role": SlotRole.number}) for c in cells]], [g.model_copy(update={"fixed": False})])
    f = OutlineSlide(id="f", kind=PatternKind.big_number, headline="Расчет", content=SlideContent(formula="100 × 300 × 30 = 900 000 ₽", numbers=[NumberCallout(value="900 000 ₽", label="Выручка")]))
    plan = match_outline(DeckOutline(title="T", slides=[f]), _manifest_of([nums]), st)
    assert plan.for_outline("f").mode == "synth"
    # the subtitle: the cards start right under the heading
    tight = _bare(PatternKind.cards, [title, *[c.model_copy(update={"bbox": BboxFrac(x=c.bbox.x, y=0.18, w=c.bbox.w, h=0.5)}) for c in cells]], [g])
    sub = four.model_copy(update={"subtitle": "Цель — поднять средний чек с 300 до 330 рублей за полгода работы кофейни"})
    assert any("подзаголовку" in r for r in _gate_room(sub, tight, _manifest_of([tight])))
    assert not any("подзаголовку" in r for r in _gate_room(sub, tiles, _manifest_of([tiles])))


def test_a_line_left_to_the_notes_carries_its_fix():
    """Классика: the cover's goal and footnote were only in the speaker notes and the issue had no autofix — the build
    never tried the template's other cover, which holds both."""
    from test_tpl_clone_audit import _ctx, _ir_slide

    from verstka.audit.checks.integrity import content_in_notes
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    o = DeckOutline(title="x", slides=[OutlineSlide(id="c", kind=PatternKind.bullets, headline="План", content=SlideContent(bullets=["Шаг"]), takeaway="Рост на 26,5% за 6 месяцев")])
    got = content_in_notes(_ctx([_ir_slide(1, [("План", (0.05, 0.05, 0.9, 0.1), 32)], oid="c")], o))
    assert got and got[0].autofix is not None and got[0].autofix.action == "rematch" and got[0].autofix.params.get("outline_id") == "c"


def test_an_author_list_item_the_plan_lost_is_reported():
    """Eight one-line steps planned as six bullets: the slide matched its plan, the audit scored 97 — the two steps
    were only in the notes. The audit reads the brief's list against the slides of that section."""
    from test_tpl_clone_audit import _ctx, _ir_slide

    from verstka.audit.checks.integrity import _item_present, author_item_missing
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    brief = "Слайд 1. Переезд\n\nСлайд 2. Восемь шагов переезда\n— Найти помещение;\n— Подписать договор аренды;\n— Настроить пропуска;\n\nСлайд 3. Итоги\n"
    o = DeckOutline(title="x", slides=[OutlineSlide(id="s2", kind=PatternKind.bullets, headline="Восемь шагов переезда", spec_ref=2, verbatim=True, content=SlideContent(bullets=["Найти помещение", "Подписать договор аренды"]))])
    ctx = _ctx([_ir_slide(1, [("Восемь шагов переезда", (0.05, 0.05, 0.9, 0.1), 32), ("Найти помещение", (0.05, 0.3, 0.9, 0.05), 18), ("Подписать договор аренды", (0.05, 0.4, 0.9, 0.05), 18)], oid="s2")], o)
    ctx.brief_text = brief
    got = author_item_missing(ctx)
    assert len(got) == 1 and got[0].details["missing"] == ["Настроить пропуска"] and got[0].autofix is not None
    # a figure set large over its label is the line «100 покупок в день»
    assert _item_present("100 покупок в день", "100|покупоквдень") and not _item_present("3 000 покупок в месяц", "100|покупоквдень")
