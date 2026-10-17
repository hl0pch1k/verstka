"""Колода из семи файлов примера на трёх шаблонах VK и на незнакомом (LibreOffice «Forestbird», при разработке не
использовался), без модели: история строится, вариант визуализации датасета один во всех трёх вариантах и на всех
шаблонах (и в плане, и в самих файлах .pptx), `OutlineSlide.data` заполнено, аудит без замечаний к данным, а вся
работа — история, вёрстка трёх вариантов и аудит на четырёх шаблонах — укладывается в 60 секунд.

Шаблоны VK — из папки «Датасет» (VERSTKA_FIXTURES_DIR или рядом с репозиторием); их разбор берётся из рабочей папки
владельца, если он там свежий (иначе шаблон разбирается заново — это время в 60 секунд не входит). Второй набор данных
(`examples/data-science-2026/`) проходит тот же путь на незнакомом шаблоне без падений."""

from __future__ import annotations

import json
import os
import shutil
import time
from pathlib import Path

import pytest

from verstka.audit.ir import build_deck_ir
from verstka.audit.runner import run_audit
from verstka.data.readers import read_file
from verstka.data.story import build_data_story
from verstka.ingest.workspace import TemplateWorkspace, file_sha256
from verstka.pipeline.generate import generate_variants
from verstka.schemas.outline import Brief
from verstka.schemas.template import TemplateManifest

ROOT = Path(__file__).resolve().parents[2]
SALES = ROOT / "examples" / "data-sales-2026"
SCIENCE = ROOT / "examples" / "data-science-2026"
UNSEEN = ROOT / "tests" / "fixtures" / "data_story" / "lo_forestbird.pptx"
BRIEF = ("Итоги продаж VK Cloud за 9 месяцев 2026 года для правления. Покажи рост выручки, сравни регионы, что даёт "
         "выручку и где узкие места. Данные — в приложенных файлах.")
SCIENCE_BRIEF = ("Итоги университета за 2026 год для учёного совета. Покажи, как подготовка влияет на баллы ЕГЭ, как устроен "
                 "приём по формам обучения, как растёт наука и на что ушёл бюджет НИОКР. Данные — в приложенных файлах.")
VK_NAMES = ("vk tech", "workspace", "education")
DATA_CHECKS = {"headline_no_conclusion", "headline_figure_not_in_data", "data_text_figure_not_in_data", "chart_values_mismatch", "chart_no_units",
               "chart_no_source", "chart_overloaded", "table_raw_dump", "chart_from_untypable", "data_split_inconsistent"}
DATA_EXT = {".csv", ".xlsx", ".json", ".txt"}


def _vk_templates() -> list[Path]:
    env = os.environ.get("VERSTKA_FIXTURES_DIR")
    dirs = ([Path(env)] if env else []) + [p / "Датасет" for p in ROOT.parents]
    for d in dirs:
        if d.is_dir():
            found = [next((p for p in d.glob("*.pptx") if n in p.name.lower()), None) for n in VK_NAMES]
            if all(found):
                return found  # type: ignore[return-value]
    return []


def _cached(template: Path, root: Path) -> bool:
    """The owner's analysis of the template, when it is current: copied into the test's workspace."""
    tid = file_sha256(template)[:16]
    for ws in [p / "workspace" / "templates" / tid for p in ROOT.parents]:
        m = ws / "manifest.json"
        if m.is_file():
            try:
                version = json.loads(m.read_text(encoding="utf-8")).get("analysis_version")
            except (OSError, ValueError):
                continue
            if version == TemplateManifest.model_fields["analysis_version"].default:
                shutil.copytree(ws, root / "templates" / tid, dirs_exist_ok=True)
                return True
    return False


def _tables(folder: Path):
    return [t for p in sorted(folder.iterdir()) if p.suffix.lower() in DATA_EXT and p.is_file() for t in read_file(p)]


@pytest.fixture(scope="module")
def templates(tmp_path_factory):
    from verstka.analysis.manifest import analyze_template

    root = tmp_path_factory.mktemp("ws")
    found = _vk_templates() + [UNSEEN]
    for t in found:
        if not _cached(t, root):
            analyze_template(t, workspace_root=root, use_llm=False, use_vlm=False)  # (not timed: a template is analysed once)
    return root, found


def _build(templates, tmp_path_factory, folder: Path, brief: Brief):
    root, found = templates
    tables = _tables(folder)
    t0 = time.monotonic()
    out = {}
    for t in found:
        story = build_data_story(brief, tables)
        res = generate_variants(t, outline=story.outline, out_dir=tmp_path_factory.mktemp("out"), workspace_root=root,
                                use_llm=False, use_vlm=False, audit=False, autofix=False, exports=[], render_images=False, pictures=False)
        ws = TemplateWorkspace.open(res.template_id, root)
        audits = {v.strategy: run_audit(v.out_dir / "deck.pptx", res.manifest, v.outline, ws, render=False, use_vlm=False, use_llm=False, strategy=v.strategy) for v in res.variants}
        out[t.name] = (story, res, audits)
    return out, time.monotonic() - t0, len(found)


@pytest.fixture(scope="module")
def built(templates, tmp_path_factory):
    return _build(templates, tmp_path_factory, SALES, Brief(text=BRIEF, slide_count=9))


@pytest.fixture(scope="module")
def science(templates, tmp_path_factory):
    return _build(templates, tmp_path_factory, SCIENCE, Brief(text=SCIENCE_BRIEF))


def test_the_sample_deck_builds_on_every_template_in_a_minute(built):
    out, seconds, n = built
    if n < 4:
        pytest.skip("the VK templates are not available (VERSTKA_FIXTURES_DIR): only the unseen one was built")
    assert seconds <= 60, f"{seconds:.1f} s for 4 templates × 3 variants"


def test_a_dataset_keeps_its_chart_in_every_variant_and_template(built):
    out, _, _ = built
    planned: dict[str, set] = {}
    drawn: dict[str, set] = {}
    for name, (story, res, audits) in out.items():
        assert len(res.variants) == 3
        for v in res.variants:
            data = [s for s in v.outline.slides if s.data is not None]
            assert len(data) == 7, (name, v.strategy)
            ir = build_deck_ir(v.out_dir / "deck.pptx")
            by_id = {s.outline_id: s for s in ir.slides}
            for k, s in enumerate(v.outline.slides):
                if s.data is None:
                    continue
                c = s.content
                planned.setdefault(s.data.file, set()).add(c.chart.type if c.chart else "table" if c.table else "cards")
                ir_slide = by_id.get(s.id) or ir.slides[k]
                charts = [(e.chart.type, e.chart.grouping, e.chart.gap_width == 0) for e in ir_slide.elements if e.chart]
                tables = [e for e in ir_slide.elements if e.table]
                drawn.setdefault(s.data.file, set()).add(tuple(charts) if charts else ("table",) if tables else ("cards",))
    assert all(len(v) == 1 for v in planned.values()), planned
    assert all(len(v) == 1 for v in drawn.values()), drawn
    assert planned["Выручка по месяцам.csv"] == {"line"} and planned["Регионы.xlsx"] == {"bar"}
    assert planned["Каналы продаж.json"] == {"doughnut"} and planned["Время ответа поддержки.csv"] == {"histogram"}
    assert planned["Клиенты.xlsx"] == {"table"} and planned["Итоги квартала.txt"] == {"cards"}


def test_the_variants_differ_in_layout_not_in_the_chart(built):
    out, _, _ = built
    from verstka.pipeline.revise import slide_fingerprints

    for name, (story, res, audits) in out.items():
        prints = {v.strategy: tuple(slide_fingerprints(v.out_dir / "deck.pptx")) for v in res.variants}
        assert len(set(prints.values())) == 3, name  # three different decks of one story


def test_the_slides_carry_their_data_for_why_this(built):
    out, _, _ = built
    for story, res, _ in out.values():
        for v in res.variants:
            for s in v.outline.slides:
                if s.data is None:
                    continue
                assert s.data.options and s.data.task is not None and s.data.headline_facts and s.data.profile_line
                assert s.footnote and s.footnote.startswith("Источник:")


def test_the_audit_finds_nothing_wrong_with_the_data(built):
    out, _, _ = built
    for name, (_, _, audits) in out.items():
        for strategy, rep in audits.items():
            bad = [(i.slide, i.check_id, i.message) for i in rep.issues if i.check_id in DATA_CHECKS and i.severity in ("error", "warn")]
            assert bad == [], (name, strategy, bad)


def test_the_second_data_set_builds_on_the_unseen_template(tmp_path):
    story = build_data_story(Brief(text=SCIENCE_BRIEF), _tables(SCIENCE))
    res = generate_variants(UNSEEN, outline=story.outline, out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", strategies=["structured"],
                            use_llm=False, use_vlm=False, audit=False, autofix=False, exports=[], render_images=False, pictures=False)
    v = res.variants[0]
    assert (v.out_dir / "deck.pptx").is_file()
    assert len([s for s in v.outline.slides if s.data is not None]) >= 5


def _decor_over_frames(pptx: Path, outline) -> list[tuple]:
    """(slide, file, where the decor comes from, share of the frame it covers) for every chart or table frame of a data
    slide that the template's decor — the art its layout and master draw (`layers.art_layers`: grounds, rules and the
    slide's own shapes aside) and the pictures kept on the slide — stands on."""
    from lxml import etree
    from pptx import Presentation

    from verstka.rendering.deck import element_bbox
    from verstka.rendering.layers import art_layers, opaque_box
    from verstka.schemas.common import Bbox

    prs = Presentation(str(pptx))
    out = []
    for k, (slide, osl) in enumerate(zip(prs.slides, outline.slides), 1):
        if osl.data is None:
            continue
        tree = slide._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spTree")
        frames = [Bbox(x=b[0], y=b[1], w=b[2], h=b[3]) for el in tree if etree.QName(el).localname == "graphicFrame" and (b := element_bbox(el)) and b[2] > 0]
        for lay in art_layers(slide):
            if lay.kind in ("graphicFrame", "cxnSp") or (lay.source == "slide" and not lay.picture) or lay.cover >= 0.85:
                continue
            box = lay.box
            if lay.kind == "pic":
                part = slide.part if lay.source == "slide" else slide.slide_layout.part if lay.source == "layout" else slide.slide_layout.slide_master.part
                try:
                    box = opaque_box(part, lay.el, box)
                except Exception:  # noqa: BLE001
                    pass
            for f in frames:
                share = f.intersection(box) / max(1, f.area)
                if share > 0.002:
                    out.append((k, osl.data.file, lay.source, round(share, 3)))
    return out


def test_no_template_decor_stands_on_a_chart_or_table(built):
    out, _, _ = built
    for name, (_, res, _) in out.items():
        for v in res.variants:
            assert _decor_over_frames(v.out_dir / "deck.pptx", v.outline) == [], (name, v.strategy)


@pytest.mark.parametrize("which", ["built", "science"])
def test_both_example_sets_pass_the_data_checks_on_every_template(which, request):
    """The deterministic story and the deterministic audit never disagree: no data error on either example set."""
    out, _, _ = request.getfixturevalue(which)
    for name, (story, res, audits) in out.items():
        for strategy, rep in audits.items():
            errors = [(i.slide, i.check_id, i.message) for i in rep.issues if i.check_id in DATA_CHECKS and i.severity == "error"]
            assert errors == [], (which, name, strategy, errors)


def test_the_closing_slides_read_as_a_deliberate_end(templates, tmp_path):
    """qa3 leftovers: a data deck's «Что дальше» is numbered steps — a sample that numbers its steps, or the composer's
    — never three short lines in a ruled list or in tall cards (ЛЦТ p15, VK Tech p24/p25/p28), at the template's body
    size or larger; the summary's cards are led by their figures, each at least 1.5× the cards' text (ЛЦТ p17 set them
    at 24 pt over 18 pt)."""
    import re

    from pptx import Presentation

    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import OutlineSlide, SlideContent, SlideItem

    root, found = templates
    story = build_data_story(Brief(text=BRIEF, slide_count=9), _tables(SALES))
    o = story.outline
    steps = ["Сократить хвост обращений дольше двух часов", "Развивать партнёрский канал в Казани и Краснодаре", "Удержать темп роста выручки в IV квартале"]
    at = next(k + 1 for k, s in enumerate(o.slides) if s.id == "summary")
    o.slides.insert(at, OutlineSlide(id="next-steps", kind=PatternKind.process, headline="Что дальше", content=SlideContent(items=[SlideItem(title=x) for x in steps])))
    cards = next(s for s in o.slides if s.id == "summary").content.items

    def norm(t: str) -> str:
        return re.sub(r"\s+", "", t or "")

    def shapes(group):
        for x in group:
            if x.shape_type == 6:  # (a group: its members — a sample's steps are often grouped)
                yield from shapes(x.shapes)
            else:
                yield x

    head = norm(next(s for s in o.slides if s.id == "summary").headline).lower()

    def sizes(slide, texts: list[str], exact: bool) -> list[float]:
        keys = [norm(t) for t in texts if norm(t)]
        out = []
        for shape in shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                t = norm(p.text)
                if t.lower() == head:
                    continue  # (the summary's headline says its two strongest conclusions — qa4 NEW-S1)
                if t and any(t == x if exact else t[:16] == x[:16] for x in keys):
                    out += [r.font.size.pt for r in p.runs if r.font.size]
        return out

    for t in found:
        res = generate_variants(t, outline=o, out_dir=tmp_path / t.stem, workspace_root=root, use_llm=False, use_vlm=False, audit=False,
                                autofix=False, exports=[], render_images=False, pictures=False)
        body = res.manifest.tokens.typography.size_for("body", 14.0)
        for v in res.variants:
            plan = json.loads((v.out_dir / "layout_plan.json").read_text(encoding="utf-8"))["slides"]
            deck = Presentation(v.out_dir / "deck.pptx")
            k = next(i for i, s in enumerate(plan) if s["outline_id"] == "next-steps")
            if plan[k]["mode"] == "clone":
                pat = next(p for p in res.manifest.patterns if p.id == plan[k]["pattern_id"])
                assert any(x.role.value == "number" for x in pat.slots), (t.name, v.strategy, pat.id)
            got = sizes(deck.slides[k], steps, exact=False)
            assert got and min(got) >= body - 0.5, (t.name, v.strategy, got, body)
            j = next(i for i, s in enumerate(plan) if s["outline_id"] == "summary")
            figs = sizes(deck.slides[j], [i.number for i in cards], exact=True)
            words = sizes(deck.slides[j], [x for i in cards for x in (i.title, i.text) if len(norm(x)) >= 4], exact=False)
            assert figs and words and min(figs) >= 1.5 * max(words) - 0.5, (t.name, v.strategy, figs, words)
