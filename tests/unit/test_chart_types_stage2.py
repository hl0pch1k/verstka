"""Вторая ступень, движок диаграмм (verstka.rendering.charts): гистограмма, рассеяние, накопительные и сгруппированные
столбцы/полосы, линия с разрывами, подписи осей, строка источника, управление выделением — родные объекты PowerPoint в
стиле шаблона; чтение пропусков аудитом (audit/ir.py, audit/checks/writing.py)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from verstka.audit.checks.writing import _chart_rows
from verstka.audit.ir import build_deck_ir
from verstka.rendering.charts import (
    NO_HIGHLIGHT,
    _chroma,
    add_chart,
    chart_data_ok,
    delta_e,
    eased,
    effective_chart_type,
    prefer_bar,
    resolve_series,
    source_said,
    source_text,
    xy_series,
)
from verstka.schemas.common import Bbox, contrast_ratio
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, OutlineSlide, SlideContent, XYSeries
from verstka.schemas.template import ChartStyleSpec, TypeStep, Typography

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
W, H = 12192000, 6858000
BOX = Bbox(x=int(W * 0.06), y=int(H * 0.22), w=int(W * 0.88), h=int(H * 0.68))
# VK Tech-like template: its chart colours and its face
PALETTE = ["0077FF", "00D3E6", "005EFF", "FF6C6C", "0563C1"]
STYLE = ChartStyleSpec(series_colors=PALETTE, font_family="Play", font_size_pt=7.0)
TYPO = Typography(scale=[TypeStep(role="display", size_pt=47), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=12), TypeStep(role="body", size_pt=9), TypeStep(role="small", size_pt=7)])
TEXT = "000000"
OUTLINE = DeckOutline(title="t")
BINS = ["0–30", "30–60", "60–90", "90–120", "120–180", "180–240", "240–360", "360–600", "> 600"]
COUNTS = [148, 121, 63, 78, 52, 22, 15, 10, 3]
QS = ["IV кв. 2024", "I кв. 2025", "II кв. 2025", "III кв. 2025", "IV кв. 2025", "I кв. 2026", "II кв. 2026", "III кв. 2026"]
NPS = [41, 43, None, 46, 48, 49, None, 52]
REG = ["Москва", "Санкт-Петербург", "Казань", "Новосибирск", "Екатеринбург"]


def _slide(bg: str = "FFFFFF"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return prs, s


def _gf(spec: ChartSpec, bg: str = "FFFFFF"):
    prs, s = _slide(bg)
    gf = add_chart(s, BOX, spec, OUTLINE, STYLE, TYPO, text_hex=TEXT if bg == "FFFFFF" else "FFFFFF")
    return prs, s, gf


def _cs(spec: ChartSpec, bg: str = "FFFFFF"):
    return _gf(spec, bg)[2].chart._chartSpace


def _hist(**kw) -> ChartSpec:
    return ChartSpec(type="histogram", categories=BINS, series=[InlineSeries(name="Обращения", values=COUNTS)], **kw)


def _stacked(kind: str = "stacked_column", **kw) -> ChartSpec:
    return ChartSpec(type=kind, categories=["I кв.", "II кв.", "III кв.", "IV кв."], series=[
        InlineSeries(name="Прямые продажи", values=[120, 134, 142, 151]), InlineSeries(name="Партнёры", values=[71, 80, 86, 92]),
        InlineSeries(name="Маркетплейс", values=[40, 46, 55, 63]), InlineSeries(name="Прочее", values=[22, 21, 25, 27])], unit="млн ₽", **kw)


def _scatter(**kw) -> ChartSpec:
    xs = [120, 180, 220, 260, 310, 350, 400, 450, 520, 590]
    ys = [110, 160, 230, 240, 300, 330, 390, 420, 500, 560]
    return ChartSpec(type="scatter", xy=[XYSeries(name="Клиенты", x=xs, y=ys)], **kw)


def _title_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(f"{A}t"))


def _ser_fill(ser) -> str:
    return ser.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val")


# ---------------------------------------------------------------------------------------------- types


def test_new_types_map_to_native_chart_kinds():
    assert effective_chart_type(_hist(), OUTLINE) == "histogram"
    assert effective_chart_type(_scatter(), OUTLINE) == "scatter"
    assert effective_chart_type(_stacked(), OUTLINE) == "stacked_column"
    assert effective_chart_type(_stacked("stacked_bar"), OUTLINE) == "stacked_bar"
    grouped = ChartSpec(type="grouped_column", categories=["I", "II"], series=[InlineSeries(name="2025", values=[1, 2]), InlineSeries(name="2026", values=[2, 3])])
    assert effective_chart_type(grouped, OUTLINE) == "column"
    assert effective_chart_type(grouped.model_copy(update={"type": "grouped_bar"}), OUTLINE) == "bar"
    # long category labels turn stacked and grouped columns horizontal, as plain columns; a histogram's bins never turn
    long_cats = ["Автоматизация обращений", "Маршрутизация звонков"]
    stacked_long = ChartSpec(type="stacked_column", categories=long_cats, series=[InlineSeries(name="a", values=[1, 2]), InlineSeries(name="b", values=[2, 1])])
    assert effective_chart_type(stacked_long, OUTLINE) == "stacked_bar"
    assert effective_chart_type(grouped.model_copy(update={"categories": ["Автоматизация обращений клиентов", "Маршрутизация звонков операторам"]}), OUTLINE) == "bar"
    many_bins = ChartSpec(type="histogram", categories=[f"{i * 10}–{i * 10 + 10}" for i in range(12)], series=[InlineSeries(name="n", values=list(range(1, 13)))])
    assert not prefer_bar(many_bins, OUTLINE) and effective_chart_type(many_bins, OUTLINE) == "histogram"


def test_histogram_bins_touch_with_a_hairline_of_the_ground():
    cs = _cs(_hist())
    bc = cs.find(f".//{C}barChart")
    assert bc.find(f"{C}barDir").get("val") == "col"
    assert bc.find(f"{C}grouping").get("val") == "clustered"
    assert bc.find(f"{C}gapWidth").get("val") == "0"
    assert bc.find(f"{C}overlap").get("val") == "0"
    ser = bc.find(f"{C}ser")
    assert _ser_fill(ser) == "0077FF"  # the template's accent
    seam = ser.find(f"{C}spPr/{A}ln")
    assert seam.find(f"{A}solidFill/{A}srgbClr").get("val") == "FFFFFF" and int(seam.get("w")) == int(0.75 * 12700)
    # the counts sit over the bins
    assert ser.find(f"{C}dLbls/{C}dLblPos").get("val") == "outEnd"


def test_histogram_highlight_steps_the_other_bins_back():
    cs = _cs(_hist(highlight_index=0))
    ser = cs.find(f".//{C}barChart/{C}ser")
    base = _ser_fill(ser)
    dpt = {int(d.find(f"{C}idx").get("val")): d.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val") for d in ser.findall(f"{C}dPt")}
    assert dpt[0] == "0077FF" and base != "0077FF"
    assert _chroma(base) > 12 and contrast_ratio(base, "FFFFFF") >= 1.4  # a tint of the accent, never grey


def test_stacked_bars_are_parts_of_one_whole():
    for kind, bar_dir in (("stacked_column", "col"), ("stacked_bar", "bar")):
        cs = _cs(_stacked(kind))
        bc = cs.find(f".//{C}barChart")
        assert bc.find(f"{C}barDir").get("val") == bar_dir
        assert bc.find(f"{C}grouping").get("val") == "stacked"
        assert bc.find(f"{C}overlap").get("val") == "100"
        assert 40 <= int(bc.find(f"{C}gapWidth").get("val")) <= 180
        sers = bc.findall(f"{C}ser")
        fills = [_ser_fill(s) for s in sers]
        # the template's categorical colours, the accent first, every two clearly apart; «Прочее» quiet (no hue)
        assert fills[0] == "0077FF" and set(fills[:3]) <= set(PALETTE + ["0077FF"]) | set(fills[:3])
        assert all(delta_e(a, b) >= 20 for i, a in enumerate(fills) for b in fills[i + 1 :])
        assert _chroma(fills[3]) < 12
        # a hairline of the ground between the segments, each value inside its segment
        assert all(s.find(f"{C}spPr/{A}ln/{A}solidFill/{A}srgbClr").get("val") == "FFFFFF" for s in sers)
        assert all(s.find(f"{C}dLbls/{C}dLblPos").get("val") == "ctr" for s in sers)
        # parts need their legend
        assert cs.find(f".//{C}legend") is not None


def test_stacked_series_never_take_the_story_paint():
    # «2025 / 2026» stacked are two parts, not a past and a lead: two template colours, not a tint of the accent
    spec = ChartSpec(type="stacked_column", categories=["I", "II", "III"], series=[InlineSeries(name="2025", values=[3, 4, 5]), InlineSeries(name="2026", values=[4, 5, 6])])
    fills = [_ser_fill(s) for s in _cs(spec).iter(f"{C}ser")]
    assert fills[0] == "0077FF" and fills[1] in PALETTE[1:]


def test_grouped_bars_highlight_a_category_and_keep_the_series_order():
    spec = ChartSpec(type="grouped_bar", categories=REG, series=[InlineSeries(name="2025", values=[480, 260, 150, 120, 98]), InlineSeries(name="2026", values=[560, 300, 207, 131, 110])], unit="млн ₽", highlight_index=2)
    cs = _cs(spec)
    bc = cs.find(f".//{C}barChart")
    assert bc.find(f"{C}grouping").get("val") == "clustered" and bc.find(f"{C}barDir").get("val") == "bar"
    sers = bc.findall(f"{C}ser")
    for ser in sers:
        base = _ser_fill(ser)
        dpt = {int(d.find(f"{C}idx").get("val")): d.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val") for d in ser.findall(f"{C}dPt")}
        assert 2 not in dpt and set(dpt) == {0, 1, 3, 4}
        assert all(v == eased(base, "FFFFFF") for v in dpt.values())
        assert all(contrast_ratio(v, "FFFFFF") < contrast_ratio(base, "FFFFFF") for v in dpt.values())  # lighter, never darker
    # the stepped-back group keeps its order of strength: the lead stays stronger than its past
    lead, past = _ser_fill(sers[1]), _ser_fill(sers[0])
    assert contrast_ratio(eased(lead, "FFFFFF"), "FFFFFF") > contrast_ratio(eased(past, "FFFFFF"), "FFFFFF")
    # the first row still names its series («2025: 480»)
    first = [d for d in sers[0].iter(f"{C}dLbl") if d.find(f"{C}idx").get("val") == "0"][0]
    assert first.find(f"{C}showSerName").get("val") == "1"


def test_planner_bars_of_several_series_ignore_a_highlight_as_before():
    spec = ChartSpec(type="column", categories=["Q1", "Q2", "Q3"], series=[InlineSeries(name="Web", values=[12, 15, 19]), InlineSeries(name="Mobile", values=[8, 11, 16])], highlight_index=2)
    assert not list(_cs(spec).iter(f"{C}dPt"))


# ---------------------------------------------------------------------------------------------- gaps and highlight


def test_line_with_gaps_breaks_and_ends_on_its_last_value():
    spec = ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True)
    cs = _cs(spec)
    assert cs.find(f"{C}chart/{C}dispBlanksAs").get("val") == "gap"
    cache = cs.find(f".//{C}lineChart/{C}ser/{C}val/{C}numRef/{C}numCache")
    assert cache.find(f"{C}ptCount").get("val") == "8"
    assert sorted(int(p.get("idx")) for p in cache.findall(f"{C}pt")) == [0, 1, 3, 4, 5, 7]  # a gap, never a zero
    # a missing tail: the end (the highlight) is the last point with a value
    tail = spec.model_copy(update={"series": [InlineSeries(name="NPS", values=[41, 43, None, 46, 48, 49, 52, None])]})
    ser = _cs(tail).find(f".//{C}lineChart/{C}ser")
    big = [d for d in ser.findall(f"{C}dPt") if d.find(f"{C}marker/{C}size") is not None and d.find(f"{C}marker/{C}size").get("val") == "11"]
    assert [int(d.find(f"{C}idx").get("val")) for d in big] == [6]


def test_every_kind_writes_blanks_as_gaps():
    for spec in (_hist(), _stacked(), ChartSpec(type="column", categories=["a", "b"], series=[InlineSeries(name="x", values=[1, None])])):
        assert _cs(spec).find(f"{C}chart/{C}dispBlanksAs").get("val") == "gap"


def test_no_highlight_leaves_the_line_without_an_accented_point():
    spec = ChartSpec(type="line", categories=QS[:5], series=[InlineSeries(name="NPS", values=[41, 43, 45, 46, 48])], highlight_index=NO_HIGHLIGHT)
    ser = _cs(spec).find(f".//{C}lineChart/{C}ser")
    assert not ser.findall(f"{C}dPt")
    default = _cs(spec.model_copy(update={"highlight_index": None})).find(f".//{C}lineChart/{C}ser")
    assert [int(d.find(f"{C}idx").get("val")) for d in default.findall(f"{C}dPt")] == [4]  # the last point by default


# ---------------------------------------------------------------------------------------------- axis titles


def _axis_title(cs, axis: str):
    t = cs.find(f".//{C}plotArea/{C}{axis}/{C}title")
    return t


def test_axis_titles_in_the_template_face_and_muted_colour():
    cs = _cs(_hist(x_title="Минут до ответа", y_title="Число обращений"))
    for axis, text, rot in (("catAx", "Минут до ответа", "0"), ("valAx", "Число обращений", "-5400000")):
        t = _axis_title(cs, axis)
        assert t is not None and _title_text(t) == text
        r = t.find(f".//{A}r/{A}rPr")
        assert r.find(f"{A}latin").get("typeface") == "Play"
        color = r.find(f"{A}solidFill/{A}srgbClr").get("val")
        assert color != TEXT and contrast_ratio(color, "FFFFFF") >= 4.5  # the muted text colour, still readable
        assert r.get("b") == "0"
        assert t.find(f"{C}tx/{C}rich/{A}bodyPr").get("rot") == rot
        assert t.find(f"{C}overlay").get("val") == "0"
        assert t.find(f"{C}layout/{C}manualLayout") is not None
    # the value axis's ticks are replaced by the labels: the axis stays only to carry its title
    va = cs.find(f".//{C}valAx")
    assert va.find(f"{C}delete").get("val") == "0" and va.find(f"{C}tickLblPos").get("val") == "none"
    assert va.find(f"{C}majorGridlines") is None


def test_value_axis_title_carries_the_word_unit_once():
    spec = ChartSpec(type="line", categories=[f"м{i}" for i in range(12)], series=[InlineSeries(name="Выручка", values=[float(300 + 5 * i) for i in range(12)])], unit="млн ₽", y_title="Выручка")
    cs = _cs(spec)
    assert _title_text(_axis_title(cs, "valAx")) == "Выручка, млн ₽"
    assert cs.find(f"{C}chart/{C}title") is None  # no unit caption above: the axis says it
    # a title that names its unit already is kept as written
    cs = _cs(spec.model_copy(update={"y_title": "Выручка, млн ₽"}))
    assert _title_text(_axis_title(cs, "valAx")) == "Выручка, млн ₽"
    # horizontal bars: the categories' title along the left axis, the values' under the bars
    bars = ChartSpec(type="bar", categories=REG, series=[InlineSeries(name="Выручка", values=[508, 214, 131, 98, 87])], unit="млн ₽", x_title="Регион", y_title="Выручка")
    cs = _cs(bars)
    assert cs.find(f".//{C}catAx/{C}title/{C}tx/{C}rich/{A}bodyPr").get("rot") == "-5400000"
    assert cs.find(f".//{C}valAx/{C}title/{C}tx/{C}rich/{A}bodyPr").get("rot") == "0"


def test_axis_titles_take_room_so_the_plot_moves_in():
    plain = _cs(_hist())
    titled = _cs(_hist(x_title="Минут до ответа", y_title="Число обращений"))

    def ml(cs):
        el = cs.find(f".//{C}plotArea/{C}layout/{C}manualLayout")
        return {k: float(el.find(f"{C}{k}").get("val")) for k in ("x", "y", "w", "h")}

    a, b = ml(plain), ml(titled)
    assert b["x"] > a["x"] and b["y"] + b["h"] < a["y"] + a["h"]
    # the left title sits left of the plot, the bottom one under it
    yl = titled.find(f".//{C}valAx/{C}title/{C}layout/{C}manualLayout")
    xl = titled.find(f".//{C}catAx/{C}title/{C}layout/{C}manualLayout")
    assert float(yl.find(f"{C}x").get("val")) < b["x"]
    assert float(xl.find(f"{C}y").get("val")) > b["y"] + b["h"]


def test_pies_take_no_axis_titles():
    spec = ChartSpec(type="pie", categories=["a", "b"], series=[InlineSeries(name="s", values=[60, 40])], unit="%", x_title="x", y_title="y")
    cs = _cs(spec)
    assert not list(cs.iter(f"{C}title")) or all(_title_text(t) not in ("x", "y") for t in cs.iter(f"{C}title"))


# ---------------------------------------------------------------------------------------------- scatter


def test_scatter_is_native_xy_points_with_quiet_axes():
    assert chart_data_ok(_scatter(), OUTLINE) and resolve_series(_scatter(), OUTLINE) == []
    assert not chart_data_ok(ChartSpec(type="scatter"), OUTLINE)
    assert xy_series(ChartSpec(type="scatter", xy=[XYSeries(name="a", x=[1, 2, 3], y=[4, 5])])) == [("a", [1.0, 2.0], [4.0, 5.0])]
    cs = _cs(_scatter(x_title="Средний чек, тыс. ₽", y_title="Выручка на клиента", unit="тыс. ₽"))
    sc = cs.find(f".//{C}scatterChart")
    assert sc is not None
    ser = sc.find(f"{C}ser")
    xs = [float(p.find(f"{C}v").text) for p in ser.findall(f"{C}xVal//{C}pt")]
    ys = [float(p.find(f"{C}v").text) for p in ser.findall(f"{C}yVal//{C}pt")]
    assert xs[:3] == [120.0, 180.0, 220.0] and ys[:3] == [110.0, 160.0, 230.0]
    # points only: the series has no line, its markers are ringed circles in the accent
    assert ser.find(f"{C}spPr/{A}ln/{A}noFill") is not None
    mk = ser.find(f"{C}marker")
    assert mk.find(f"{C}symbol").get("val") == "circle"
    assert mk.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val") == "0077FF"
    # two value axes: X under a hairline baseline, Y with quiet gridlines; both titled
    axes = cs.findall(f".//{C}plotArea/{C}valAx")
    assert len(axes) == 2
    by_pos = {a.find(f"{C}axPos").get("val"): a for a in axes}
    assert by_pos["l"].find(f"{C}majorGridlines") is not None and by_pos["b"].find(f"{C}majorGridlines") is None
    assert _title_text(by_pos["b"].find(f"{C}title")) == "Средний чек, тыс. ₽"
    assert _title_text(by_pos["l"].find(f"{C}title")) == "Выручка на клиента, тыс. ₽"
    # every axis id pairs with the other (LibreOffice reads them as signed 32-bit)
    ids = {a.find(f"{C}axId").get("val") for a in axes}
    assert all(a.find(f"{C}crossAx").get("val") in ids for a in axes) and all(int(i) < 2**31 for i in ids)


def test_scatter_highlight_is_one_larger_point_with_its_value():
    cs = _cs(_scatter(highlight_index=3))
    ser = cs.find(f".//{C}scatterChart/{C}ser")
    dpt = ser.find(f"{C}dPt")
    assert dpt.find(f"{C}idx").get("val") == "3"
    assert dpt.find(f"{C}marker/{C}spPr/{A}solidFill/{A}srgbClr").get("val") == "0077FF"
    assert ser.find(f"{C}marker/{C}spPr/{A}solidFill/{A}srgbClr").get("val") != "0077FF"
    labels = ser.findall(f"{C}dLbls/{C}dLbl")
    assert [d.find(f"{C}idx").get("val") for d in labels] == ["3"]
    assert ser.find(f"{C}dLbls/{C}showVal").get("val") == "0"
    # (the series' label list sits before its X values: schema order)
    names = [c.tag.split("}")[1] for c in ser]
    assert names.index("dLbls") < names.index("xVal")


def test_scatter_of_several_series_uses_the_palette_and_a_legend():
    spec = ChartSpec(type="scatter", xy=[XYSeries(name="Москва", x=[1, 2, 3], y=[3, 4, 5]), XYSeries(name="Регионы", x=[1, 2, 3], y=[1, 2, 2])])
    cs = _cs(spec)
    fills = [s.find(f"{C}marker/{C}spPr/{A}solidFill/{A}srgbClr").get("val") for s in cs.iter(f"{C}ser")]
    assert fills[0] == "0077FF" and delta_e(fills[0], fills[1]) >= 20
    assert cs.find(f".//{C}legend") is not None


# ---------------------------------------------------------------------------------------------- source line


def test_source_line_under_the_chart_inside_its_frame():
    assert source_text("Выручка по месяцам.csv") == "Источник: Выручка по месяцам.csv"
    assert source_text("Источник: Росстат") == "Источник: Росстат"
    assert source_text("  ") is None
    prs, s, gf = _gf(_hist(source="Время ответа поддержки.csv"))
    src = gf.verstka_source
    assert src is not None and src.name == "Chart source"
    assert src.text_frame.text == "Источник: Время ответа поддержки.csv"
    # under the chart, both inside the caller's frame
    assert gf.top == BOX.y and gf.top + gf.height <= src.top and src.top + src.height <= BOX.y + BOX.h + 1
    run = src.text_frame.paragraphs[0].runs[0]
    # small print: the template's face, never larger than the chart's category labels
    cat_size = int(gf.chart._chartSpace.find(f".//{C}catAx/{C}txPr/{A}p/{A}pPr/{A}defRPr").get("sz")) / 100
    assert run.font.name == "Play" and run.font.size.pt <= cat_size
    color = str(run.font.color.rgb)
    assert color != TEXT and contrast_ratio(color, "FFFFFF") >= 4.5
    # no source, no line; the chart keeps the whole frame
    prs, s, gf = _gf(_hist())
    assert gf.verstka_source is None and gf.height == BOX.h
    assert not [sh for sh in s.shapes if sh.name == "Chart source"]


def test_source_said_by_the_slide_is_not_repeated():
    assert source_said("Выручка по месяцам.csv", "Источник: Выручка по месяцам.csv")
    assert source_said("Источник: Регионы.xlsx", "Данные: регионы.xlsx, выгрузка CRM")
    assert not source_said("Регионы.xlsx", "Все цифры условные")
    assert not source_said("Регионы.xlsx", None)
    assert source_said(None, None)  # nothing to print


def test_source_line_on_a_dark_ground_stays_readable():
    prs, s, gf = _gf(_hist(source="Время ответа поддержки.csv"), bg="0B1220")
    color = str(gf.verstka_source.text_frame.paragraphs[0].runs[0].font.color.rgb)
    assert contrast_ratio(color, "0B1220") >= 4.5


# ---------------------------------------------------------------------------------------------- audit reads gaps


def _save(prs, tmp_path: Path, name: str) -> Path:
    p = tmp_path / name
    prs.save(str(p))
    return p


def _ir_chart(path: Path):
    ir = build_deck_ir(path)
    return next(e.chart for e in ir.slides[0].elements if e.type == "chart")


def test_audit_reads_a_gap_in_place(tmp_path):
    spec = ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True, y_title="NPS", x_title="Квартал")
    prs, s, gf = _gf(spec)
    ch = _ir_chart(_save(prs, tmp_path, "gaps.pptx"))
    assert ch.categories == QS
    assert ch.series[0].values == [41.0, 43.0, None, 46.0, 48.0, 49.0, None, 52.0]  # every value against its own quarter
    assert ch.title is None  # an axis title is not the chart's title
    assert not ch.has_value_axis  # (the axis kept for its title shows no values)


def test_audit_reads_a_scatter_and_a_histogram(tmp_path):
    prs, s, gf = _gf(_scatter(y_title="Выручка на клиента", x_title="Средний чек"))
    ch = _ir_chart(_save(prs, tmp_path, "xy.pptx"))
    assert ch.type == "scatter" and ch.series[0].values[:3] == [110.0, 160.0, 230.0]
    assert ch.categories[:3] == ["120", "180", "220"] and ch.has_value_axis
    prs, s, gf = _gf(_hist())
    ch = _ir_chart(_save(prs, tmp_path, "hist.pptx"))
    assert ch.type == "column" and ch.categories == BINS and ch.series[0].values == [float(v) for v in COUNTS]


def test_chart_rows_keep_values_with_their_categories():
    class Ctx:
        outline = None

    osl = OutlineSlide(id="s1", kind="chart", headline="NPS", content=SlideContent(chart=ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)])))
    rows = _chart_rows(osl, Ctx())
    assert rows == [[(c, v, 1.0) for c, v in zip(QS, NPS) if v is not None]]
    assert ("III кв. 2025", 46, 1.0) in rows[0]  # not shifted onto «II кв. 2025»


# ---------------------------------------------------------------------------------------------- the slide composer


def _compose(env, oslide: OutlineSlide):
    from verstka.rendering.deck import DeckBuilder
    from verstka.rendering.synth import render_synth
    from verstka.schemas.layout import LayoutSlide

    deck, manifest, ws = env
    b = DeckBuilder(deck)
    outline = DeckOutline(title="T", strategy="structured", slides=[oslide])
    slide, _ = render_synth(b, LayoutSlide(outline_id=oslide.id, mode="synth", composition="chart"), oslide, manifest, ws, outline)
    return slide


@pytest.fixture
def env(simple_deck, tmp_path):
    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace

    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def test_composer_prints_the_source_once(env):
    chart = ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True, source="NPS по кварталам.csv", y_title="NPS")
    with_foot = OutlineSlide(id="s1", kind="chart", headline="NPS вырос до 52", footnote="Источник: NPS по кварталам.csv", content=SlideContent(chart=chart))
    slide = _compose(env, with_foot)
    texts = [" ".join(sh.text_frame.text.replace("\u00a0", " ").split()) for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
    assert sum("NPS по кварталам.csv" in t for t in texts) == 1  # the slide's footnote
    assert not [sh for sh in slide.shapes if sh.name == "Chart source"]
    # no footnote: the chart prints its own source line
    slide = _compose(env, with_foot.model_copy(update={"footnote": None}))
    assert [sh.text_frame.text for sh in slide.shapes if sh.name == "Chart source"] == ["Источник: NPS по кварталам.csv"]
    # the value axis's title names the measure: no unit caption repeats it over the chart
    assert not [sh for sh in slide.shapes if sh.name in ("Unit", "Chart title")]


# ---------------------------------------------------------------------------------------------- LibreOffice preview


def test_libreoffice_renders_every_new_kind(tmp_path):
    from verstka.ingest.render import find_soffice, render_slides

    if not find_soffice():
        pytest.skip("LibreOffice is not installed")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    specs = [
        _hist(x_title="Минут до ответа", y_title="Число обращений", source="Время ответа поддержки.csv"),
        _stacked(y_title="Выручка"), _stacked("stacked_bar"),
        ChartSpec(type="grouped_column", categories=["I", "II", "III"], series=[InlineSeries(name="2025", values=[3, 4, 5]), InlineSeries(name="2026", values=[4, 5, None])], highlight_index=1),
        ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True, y_title="NPS"),
        _scatter(x_title="Средний чек", y_title="Выручка", highlight_index=2),
    ]
    for spec in specs:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_chart(s, BOX, spec, OUTLINE, STYLE, TYPO, text_hex=TEXT)
    pngs = render_slides(_save(prs, tmp_path, "kinds.pptx"), tmp_path / "png", dpi=40, keep_pdf=False)
    assert len(pngs) == len(specs) and all(p.stat().st_size > 0 for p in pngs)
