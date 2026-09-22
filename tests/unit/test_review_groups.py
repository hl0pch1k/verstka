"""Review fixes: repeat-group riders/z-order/column axis (R7), synth layout+title (R2), card radius (EXTRA-4), scorer gates (R6)."""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.scorer import score_pattern
from verstka.planning.strategies import get_strategy
from verstka.rendering.deck import DeckBuilder, element_bbox, slide_shape_elements
from verstka.rendering.fit import fit_size
from verstka.rendering.groups import adjust_group, cell_bbox
from verstka.rendering.synth import _layout_for, card_adj, render_synth
from verstka.rendering.textfill import shape_text
from verstka.schemas.common import EMU_PER_PT, Bbox, BboxFrac, Family, PatternKind, SlotRole
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem
from verstka.schemas.template import CardSpec, Capacity, Pattern, RepeatGroup, Slot

W, H = 12192000, 6858000


def _vk(fixtures_dir, needle: str) -> Path:
    if fixtures_dir is None:
        pytest.skip("VK templates not available")
    p = next((p for p in fixtures_dir.glob("*.pptx") if needle.lower() in p.name.lower()), None)
    if p is None:
        pytest.skip(f"template {needle} not found")
    return p


# ---- synthetic decks --------------------------------------------------------------------------------


def _text(slide, x, y, w, h, text, size=14, name=None):
    tb = slide.shapes.add_textbox(Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    if name:
        tb.name = name
    return tb


def _shape(slide, x, y, w, h, kind=MSO_SHAPE.OVAL, name=None):
    sh = slide.shapes.add_shape(kind, Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H)))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x00, 0x77, 0xFF)
    if name:
        sh.name = name
    return sh


def _sid(shape) -> str:
    return str(shape.shape_id)


def _row_deck(path: Path):
    """Title + 3 text cells in a row + one icon chip above each cell + a footer (chrome)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = _text(s, 0.05, 0.08, 0.9, 0.12, "Заголовок слайда", size=32, name="Title")
    cells, chips = [], []
    for i in range(3):
        x = 0.05 + i * 0.31
        chips.append(_shape(s, x, 0.36, 0.07, 0.12, name=f"Chip {i + 1}"))
        cells.append(_text(s, x, 0.55, 0.24, 0.2, f"Тезис {i + 1}", name=f"Cell {i + 1}"))
    footer = _text(s, 0.05, 0.92, 0.1, 0.05, "ACME", size=10, name="Footer")
    prs.save(path)
    group = RepeatGroup(id="g1", member_shape_ids=[[_sid(c)] for c in cells], max_n=3, axis="row", gap=0.07, cell_bbox=BboxFrac(x=0.05, y=0.55, w=0.24, h=0.2))
    return group, [_sid(c) for c in cells], [_sid(c) for c in chips], _sid(title), _sid(footer)


def _column_deck(path: Path):
    """3 text cells stacked in a column, each with a small marker to its left."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cells, marks = [], []
    for i in range(3):
        y = 0.2 + i * 0.25
        marks.append(_shape(s, 0.05, y, 0.04, 0.07, name=f"Mark {i + 1}"))
        cells.append(_text(s, 0.12, y, 0.6, 0.18, f"Пункт {i + 1}", name=f"Cell {i + 1}"))
    prs.save(path)
    group = RepeatGroup(id="g1", member_shape_ids=[[_sid(c)] for c in cells], max_n=3, axis="column", gap=0.07, cell_bbox=BboxFrac(x=0.12, y=0.2, w=0.6, h=0.18))
    return group, [_sid(c) for c in cells], [_sid(m) for m in marks]


def _card_pair_deck(path: Path):
    """2 cards (background rect + text) in a row with room for a third."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cells = []
    for i in range(2):
        x = 0.05 + i * 0.3
        bg = _shape(s, x, 0.4, 0.26, 0.3, kind=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Card {i + 1}")
        tx = _text(s, x + 0.02, 0.45, 0.22, 0.2, f"Карточка {i + 1}", name=f"Card text {i + 1}")
        cells.append([_sid(bg), _sid(tx)])
    prs.save(path)
    return RepeatGroup(id="g1", member_shape_ids=cells, max_n=3, axis="row", gap=0.04, cell_bbox=BboxFrac(x=0.05, y=0.4, w=0.26, h=0.3))


def _by_name(slide) -> dict[str, etree._Element]:
    return {nv.get("name"): nv.getparent().getparent() for nv in slide._element.cSld.find(q("p:spTree")).iter(q("p:cNvPr"))}


# ---- R7: riders, z-order, column axis ---------------------------------------------------------------


def test_adjust_group_moves_and_removes_riders(tmp_path):
    group, cell_ids, chip_ids, title_id, footer_id = _row_deck(tmp_path / "row.pptx")
    b = DeckBuilder(tmp_path / "row.pptx")
    slide = b.clone_slide(1)
    before = {k: element_bbox(v) for k, v in _by_name(slide).items()}
    cells, _ = adjust_group(slide, group, 2, W, H, b.next_shape_id(slide))
    assert len(cells) == 2
    after = _by_name(slide)
    # the third chip left with its cell; the second chip travelled with the redistributed second cell
    assert "Chip 3" not in after and "Cell 3" not in after
    assert element_bbox(after["Chip 1"])[0] == before["Chip 1"][0]
    assert element_bbox(after["Chip 2"])[0] - before["Chip 2"][0] == element_bbox(after["Cell 2"])[0] - before["Cell 2"][0] > 0
    assert element_bbox(after["Chip 2"])[0] == element_bbox(after["Cell 2"])[0]
    # wide title and far-away footer are not riders
    assert element_bbox(after["Title"]) == before["Title"] and element_bbox(after["Footer"]) == before["Footer"]


def test_adjust_group_protected_ids_are_never_riders(tmp_path):
    group, cell_ids, chip_ids, title_id, footer_id = _row_deck(tmp_path / "row.pptx")
    b = DeckBuilder(tmp_path / "row.pptx")
    slide = b.clone_slide(1)
    before = {k: element_bbox(v) for k, v in _by_name(slide).items()}
    adjust_group(slide, group, 2, W, H, b.next_shape_id(slide), protected_ids=set(chip_ids) | {title_id, footer_id})
    after = _by_name(slide)
    assert "Chip 3" in after and all(element_bbox(after[f"Chip {i}"]) == before[f"Chip {i}"] for i in (1, 2, 3))


def test_adjust_group_duplicate_keeps_z_order(tmp_path):
    group = _card_pair_deck(tmp_path / "cards.pptx")
    b = DeckBuilder(tmp_path / "cards.pptx")
    slide = b.clone_slide(1)
    cells, _ = adjust_group(slide, group, 3, W, H, b.next_shape_id(slide))
    assert len(cells) == 3
    tree = slide._element.cSld.find(q("p:spTree"))
    order = list(tree)
    z = [[order.index(e) for e in c] for c in cells]
    # inside every cell the background stays below its text, and the copy sits after the source cell
    assert all(zc == sorted(zc) for zc in z), z
    assert min(z[2]) > max(z[1])
    assert shape_text(cells[2][1]).startswith("Карточка")


def test_adjust_group_column_axis_redistributes_with_riders(tmp_path):
    group, cell_ids, mark_ids = _column_deck(tmp_path / "col.pptx")
    b = DeckBuilder(tmp_path / "col.pptx")
    slide = b.clone_slide(1)
    before = {k: element_bbox(v) for k, v in _by_name(slide).items()}
    cells, _ = adjust_group(slide, group, 2, W, H, b.next_shape_id(slide))
    assert len(cells) == 2
    after = _by_name(slide)
    assert "Cell 3" not in after and "Mark 3" not in after
    # second cell spread down towards the end of the original span, x untouched, marker follows
    c2_before, c2_after = before["Cell 2"], element_bbox(after["Cell 2"])
    assert c2_after[1] > c2_before[1] and c2_after[0] == c2_before[0]
    span_end = int(0.2 * H) + 2 * (int(0.18 * H) + int(0.07 * H)) + int(0.18 * H)
    assert abs(c2_after[1] + c2_after[3] - span_end) <= 2
    assert element_bbox(after["Mark 2"])[1] - before["Mark 2"][1] == c2_after[1] - c2_before[1]


# ---- R2: synth layout choice and title height ---------------------------------------------------------


def _pattern(pid, kind, family, layout_part, slots=(), groups=()):
    return Pattern(id=pid, source_slide=1, kind=kind, family=family, layout_part=layout_part, slots=list(slots), repeat_groups=list(groups))


def _slot(sid, role, cap, group_id=None, bbox=None):
    return Slot(id=sid, role=role, shape_id=sid, bbox=bbox or BboxFrac(x=0.1, y=0.3, w=0.3, h=0.2), capacity=Capacity(max_chars=cap, max_lines=3), group_id=group_id)


def test_layout_for_prefers_family_layout_over_last(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    b = DeckBuilder(simple_deck)
    last = str(b.prs.slide_layouts[-1].part.partname).lstrip("/")
    light_part = "ppt/slideLayouts/slideLayout6.xml"
    patterns = [
        _pattern("p1", PatternKind.bullets, Family.light, light_part),
        _pattern("p2", PatternKind.bullets, Family.light, light_part),
        _pattern("p3", PatternKind.cards, Family.dark, last),
    ]
    m = manifest.model_copy(update={"patterns": patterns})
    # no light stat_row sample: fall back to the layout most used by light samples, never the last layout
    layout, fam = _layout_for(b, m, Family.light, PatternKind.stat_row)
    assert str(layout.part.partname).lstrip("/") == light_part and fam == Family.light
    # only dark samples exist: use their layout and report the real family so the palette matches the background
    m2 = manifest.model_copy(update={"patterns": [patterns[2]]})
    layout2, fam2 = _layout_for(b, m2, Family.light, PatternKind.bullets)
    assert str(layout2.part.partname).lstrip("/") == last and fam2 == Family.dark
    # nothing resolvable: a title-only-like layout, still not blindly the last one
    m3 = manifest.model_copy(update={"patterns": [_pattern("p9", PatternKind.bullets, Family.light, "ppt/slideLayouts/slideLayout99.xml")]})
    layout3, _ = _layout_for(b, m3, Family.light, PatternKind.bullets)
    assert layout3.name == "Title Only"


def test_synth_title_is_fitted_and_content_starts_below(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    b = DeckBuilder(simple_deck)
    headline = "Очень длинный заголовок, который занимает несколько строк и раньше перекрывался карточками и буллетами под ним"
    oslide = OutlineSlide(id="x", kind=PatternKind.bullets, headline=headline, content=SlideContent(bullets=["Первый тезис", "Второй тезис", "Третий тезис"]))
    outline = DeckOutline(title="T", slides=[oslide])
    slide, warnings = render_synth(b, LayoutSlide(outline_id="x", mode="synth", composition="bullets"), oslide, manifest, ws, outline)
    title = next(sh for sh in slide.shapes if sh.is_placeholder)
    body = next(sh for sh in slide.shapes if not sh.is_placeholder and sh.has_text_frame and "Первый" in sh.text_frame.text)
    sz = title._element.find(".//" + q("a:rPr")).get("sz")
    assert sz is not None, "fitted size must be written into the title placeholder"
    size_pt = int(sz) / 100
    box = Bbox(x=title.left, y=title.top, w=title.width, h=title.height)
    res = fit_size([headline], box, manifest.tokens.typography.primary_family, size_pt, False, [], line_spacing=manifest.tokens.typography.line_spacing)
    assert res.lines >= 2 and res.fits, "headline must fit inside the (possibly enlarged) placeholder"
    assert body.top >= title.top + title.height


@pytest.mark.integration
def test_vk_tech_synth_layout_is_light_from_second_master(fixtures_dir, tmp_path):
    tech = _vk(fixtures_dir, "VK Tech")
    manifest = analyze_template(tech, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    b = DeckBuilder(tech)
    assert len(b.layouts()) == sum(len(m.slide_layouts) for m in b.prs.slide_masters) > len(b.prs.slide_layouts)
    light_parts = {p.layout_part for p in manifest.patterns if p.family == Family.light}
    layout, fam = _layout_for(b, manifest, Family.light, PatternKind.bullets)
    assert "Спасибо" not in layout.name and fam == Family.light
    assert str(layout.part.partname).lstrip("/") in light_parts
    # content kinds without a sample of their own land on the layout of the 36 «Свободный дизайн» samples (second master)
    layout2, fam2 = _layout_for(b, manifest, Family.light, PatternKind.stat_row)
    assert str(layout2.part.partname).lstrip("/") == "ppt/slideLayouts/slideLayout11.xml" and fam2 == Family.light


# ---- EXTRA-4: absolute card radius, capped gutter -------------------------------------------------------


def test_card_adj_is_absolute_radius_capped():
    inch = 914400
    assert card_adj(int(0.12 * inch), 3 * inch, 2 * inch) == pytest.approx(0.06)
    assert card_adj(int(0.5 * inch), inch, 4 * inch) == 0.25
    assert card_adj(0, inch, inch) == 0.0


def test_synth_cards_are_not_capsules_and_gutter_is_capped(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    # VK Education-like tokens: pill chips give typical_radius 0.5, cards have no radius of their own; huge gutter
    tokens = manifest.tokens.model_copy(deep=True)
    tokens.shapes.typical_radius = 0.5
    tokens.spacing.gutter = 0.0824
    comps = manifest.components.model_copy(deep=True)
    comps.card = CardSpec(fill_hex="EDF3FC", radius=None, width_frac=0.19, height_frac=0.48)
    m = manifest.model_copy(update={"tokens": tokens, "components": comps})
    b = DeckBuilder(simple_deck)
    items = [SlideItem(title=f"Сценарий {i}", text="Короткое описание") for i in range(1, 4)]
    oslide = OutlineSlide(id="c", kind=PatternKind.cards, headline="Три сценария", content=SlideContent(items=items))
    slide, _ = render_synth(b, LayoutSlide(outline_id="c", mode="synth", composition="cards"), oslide, m, ws, DeckOutline(title="T", slides=[oslide]))
    rects = [sh for sh in slide.shapes if sh.shape_type is not None and "AUTO_SHAPE" in str(sh.shape_type) and sh.width > W * 0.15]
    assert len(rects) == 3
    for r in rects:
        assert r.adjustments[0] <= 0.25
        assert r.adjustments[0] * min(r.width, r.height) <= 0.13 * 914400  # ≈ 0.12 inch corner
    xs = sorted(r.left for r in rects)
    gap = xs[1] - (xs[0] + rects[0].width)
    assert gap <= 0.03 * W + 1


# ---- R6: hard capacity gates in the scorer ---------------------------------------------------------------


def _bullets_slide():
    return OutlineSlide(id="b", kind=PatternKind.bullets, headline="Сотрудники теряют задачи", content=SlideContent(bullets=["Задачи живут в четырёх инструментах", "Статус уточняют вручную в чате", "Дедлайны срываются в трети случаев"]))


def _stat_slide(n=3):
    nums = [NumberCallout(value=v, label=l) for v, l in [("12 400", "участников пилота"), ("+34%", "задач завершены в срок"), ("2,1 ч", "экономии в неделю")][:n]]
    return OutlineSlide(id="s", kind=PatternKind.stat_row, headline="Пилот подтвердил эффект", content=SlideContent(numbers=nums))


def test_scorer_zero_content_capacity_is_hard_gate(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    visual = get_strategy("visual")
    icons = [_slot("t", SlotRole.title, 41), _slot("i1", SlotRole.icon, 0, "g1"), _slot("i2", SlotRole.icon, 0, "g1")]
    grp = RepeatGroup(id="g1", member_shape_ids=[["i1"], ["i2"]], max_n=2, axis="row", cell_bbox=BboxFrac(x=0.1, y=0.3, w=0.3, h=0.3))
    freeform = _pattern("p38", PatternKind.freeform, Family.light, None, icons, [grp])
    res = score_pattern(_bullets_slide(), freeform, manifest, visual)
    assert res.score == 0.0 and any("нет ни одного слота под содержимое" in r for r in res.reasons)
    # a chart slide is exempt: the chart goes into the big image area
    chart = OutlineSlide(id="ch", kind=PatternKind.chart, headline="Рост", content=SlideContent(chart={"type": "column", "series_ids": ["s1"]}))
    image = _pattern("p40", PatternKind.image_text, Family.light, None, [_slot("t", SlotRole.title, 41), _slot("img", SlotRole.image, 0, bbox=BboxFrac(x=0.1, y=0.3, w=0.6, h=0.5))])
    assert score_pattern(chart, image, manifest, visual).score > 0


def test_scorer_tiny_content_capacity_penalises_kind(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    structured = get_strategy("structured")
    small = _pattern("ps", PatternKind.bullets, Family.light, None, [_slot("t", SlotRole.title, 41), _slot("b", SlotRole.bullet_list, 20)])
    big = _pattern("pb", PatternKind.bullets, Family.light, None, [_slot("t", SlotRole.title, 41), _slot("b", SlotRole.bullet_list, 200)])
    r_small = score_pattern(_bullets_slide(), small, manifest, structured)
    r_big = score_pattern(_bullets_slide(), big, manifest, structured)
    assert r_small.score < r_big.score - 0.2
    assert any("ёмкости слотов" in r or "ёмкость слотов" in r for r in r_small.reasons)


def test_scorer_number_holders_gate_and_weight_only_on_kind(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    visual = get_strategy("visual")
    # Education p18-like: 2-cell number column + one standalone number, labels aside
    slots = [_slot("t", SlotRole.title, 45), _slot("n1", SlotRole.number, 8, "g1"), _slot("n2", SlotRole.number, 8, "g1"), _slot("n3", SlotRole.number, 3), _slot("l1", SlotRole.number_label, 126, "g2"), _slot("l2", SlotRole.number_label, 126, "g2"), _slot("l3", SlotRole.number_label, 102)]
    g1 = RepeatGroup(id="g1", member_shape_ids=[["n1"], ["n2"]], max_n=2, axis="column", cell_bbox=BboxFrac(x=0.1, y=0.3, w=0.3, h=0.2))
    g2 = RepeatGroup(id="g2", member_shape_ids=[["l1"], ["l2"]], max_n=2, axis="column", cell_bbox=BboxFrac(x=0.4, y=0.3, w=0.3, h=0.2))
    two = _pattern("p18", PatternKind.stat_row, Family.light, None, slots, [g1, g2])
    res = score_pattern(_stat_slide(3), two, manifest, visual)
    assert any("числа не поместятся" in r for r in res.reasons)
    assert res.score < visual.synth_threshold, res
    # two numbers fit the same pattern and score well above the threshold
    ok = score_pattern(_stat_slide(2), two, manifest, visual)
    assert ok.score >= visual.synth_threshold and not any("не поместятся" in r for r in ok.reasons)
    # ×1.3 strategy weight only scales the kind term: a pattern with failed capacity stays below 1.0
    slots3 = [_slot("t", SlotRole.title, 45)] + [_slot(f"n{i}", SlotRole.number, 8, "g1") for i in range(3)] + [_slot(f"l{i}", SlotRole.number_label, 60, "g1") for i in range(3)]
    g3 = RepeatGroup(id="g1", member_shape_ids=[[f"n{i}", f"l{i}"] for i in range(3)], max_n=3, axis="row", cell_bbox=BboxFrac(x=0.1, y=0.3, w=0.25, h=0.3))
    three = _pattern("p45", PatternKind.stat_row, Family.light, None, slots3, [g3])
    full = score_pattern(_stat_slide(3), three, manifest, visual)
    assert full.score > 0.8
    five = OutlineSlide(id="s5", kind=PatternKind.stat_row, headline="Пять чисел", content=SlideContent(numbers=[NumberCallout(value=str(i), label="метрика") for i in range(5)]))
    lost = score_pattern(five, three, manifest, visual)
    assert lost.score < 1.0 and lost.score < visual.synth_threshold


@pytest.mark.integration
def test_vk_tech_bullets_do_not_land_on_icon_only_pattern(fixtures_dir, tmp_path):
    from verstka.matching.matcher import match_outline

    manifest = analyze_template(_vk(fixtures_dir, "VK Tech"), workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    outline = DeckOutline.model_validate_json((Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json").read_text(encoding="utf-8"))
    plan = match_outline(outline, manifest, get_strategy("visual"))
    by_id = {s.outline_id: s for s in plan.slides}
    sl3 = by_id["sl3"]
    if sl3.mode == "clone":
        p = next(p for p in manifest.patterns if p.id == sl3.pattern_id)
        assert any(s.role in (SlotRole.body, SlotRole.bullet_list, SlotRole.card_body, SlotRole.card_title) for s in p.slots), sl3
    sl4 = by_id["sl4"]
    if sl4.mode == "clone":
        p = next(p for p in manifest.patterns if p.id == sl4.pattern_id)
        assert not any("не поместятся" in r for r in sl4.reasons)
