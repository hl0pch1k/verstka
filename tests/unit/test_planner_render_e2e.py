from pathlib import Path

import pytest
from pptx import Presentation

from verstka.analysis.manifest import analyze_template
from verstka.analysis.shapes import looks_like_placeholder
from verstka.ingest.render import find_pdftoppm, find_soffice
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.matcher import match_outline
from verstka.pipeline.generate import generate_variants
from verstka.planning.brief import parse_brief_text
from verstka.planning.condense import trim_words
from verstka.planning.facts import basic_facts
from verstka.planning.outline import basic_outline, plan_outline, validate_outline
from verstka.planning.strategies import get_strategy
from verstka.providers.registry import ProviderRegistry
from verstka.rendering.renderer import render_deck
from verstka.schemas.common import PatternKind
from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent
from verstka.skills_registry.registry import SkillsRegistry

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json"
BRIEF = """---
audience: продуктовый комитет
тип: фича
slides: 8
---
# Умные напоминания

В пилоте участвовали 12 400 сотрудников. Доля завершённых в срок задач выросла на 34%. Экономия 2,1 часа в неделю на человека.

## Что сделали
Напоминание из чата превращает сообщение в задачу одним касанием. Умный срок предлагает дату по контексту. Эскалация уходит руководителю автоматически.

| Месяц | Май | Июнь | Июль |
|---|---|---|---|
| Активные пользователи | 1200 | 3400 | 6100 |
"""


def _deck_texts(pptx: Path) -> list[str]:
    prs = Presentation(str(pptx))
    out = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return out


def test_basic_facts_and_outline():
    brief = parse_brief_text(BRIEF)
    facts = basic_facts(brief.text)
    values = {f.value for f in facts.facts}
    assert "12 400" in values and "34" in values
    assert len(facts.tables) == 1 and len(facts.series) == 1 and facts.series[0].values == [1200.0, 3400.0, 6100.0]
    outline = basic_outline(brief, facts, get_strategy("structured"), 8)
    kinds = [s.kind for s in outline.slides]
    assert kinds[0] == PatternKind.title and kinds[-1] == PatternKind.thanks
    # a one-row numeric table is a time series: it becomes a chart, and the same figures are not repeated as a table
    assert PatternKind.stat_row in kinds and PatternKind.chart in kinds and PatternKind.table not in kinds
    chart = next(s for s in outline.slides if s.kind == PatternKind.chart)
    assert "с 1 200 до 6 100" in chart.headline  # the heading states what the series shows
    stat = next(s for s in outline.slides if s.kind == PatternKind.stat_row)
    assert [n.value for n in stat.content.numbers][:3] == ["12 400", "+34%", "2,1 ч"]
    assert len(outline.slides) <= 9
    assert trim_words("один два три четыре пять шесть", 3) == "один два три"


def test_validate_outline_density():
    s = OutlineSlide(id="x", kind=PatternKind.bullets, headline="Заголовок " * 5, content=SlideContent(bullets=[("слово " * 25).strip()] * 9))
    o = DeckOutline(title="T", slides=[s])
    v = validate_outline(o, None, 12)
    body = next(sl for sl in v.slides if sl.id == "x")
    assert len(body.content.bullets) == 6 and all(len(b.split()) <= 15 for b in body.content.bullets)
    assert "Не вошло" in body.notes
    assert v.slides[0].kind == PatternKind.title and v.slides[-1].kind == PatternKind.thanks


def test_plan_outline_with_mock_llm(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    demo = DeckOutline.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    planned = {"title": demo.title, "subtitle": demo.subtitle, "slides": [s.model_dump() for s in demo.slides]}
    providers = ProviderRegistry.mock(
        {
            "Return the JSON plan only": planned,
            "List the issues": {"issues": []},
            "Extract facts, series and tables": {"facts": [f.model_dump() for f in demo.facts], "series": [s.model_dump() for s in demo.series], "tables": []},
        }
    )
    skills = SkillsRegistry.load()
    brief = parse_brief_text(BRIEF)
    from verstka.planning.facts import extract_facts

    facts, _ = extract_facts(brief, skills, providers)
    assert len(facts.facts) == 6
    outline, warnings = plan_outline(brief, manifest, get_strategy("visual"), facts, skills, providers, target=12)
    assert len(outline.slides) in (12, 13) and outline.strategy == "visual" and not any("failed" in w for w in warnings)


def test_render_demo_outline_on_simple_deck(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    outline = DeckOutline.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    plan = match_outline(outline, manifest, get_strategy("structured"))
    result = render_deck(outline, plan, manifest, ws, tmp_path / "deck.pptx")
    prs = Presentation(str(result.pptx_path))
    assert len(prs.slides) == 12
    texts = [t.replace("\u00a0", " ") for t in _deck_texts(result.pptx_path)]  # no-break spaces are typesetting
    assert not any(looks_like_placeholder(t) for t in texts if t.strip()), [t for t in texts if looks_like_placeholder(t)]
    assert any("Умные напоминания" in t for t in texts) and any("Спасибо" in t for t in texts)
    # the cards slide carries the item titles
    assert any("Напоминание из чата" in t for t in texts) and any("Дайджест" in t for t in texts)
    # native chart and table exist somewhere in the deck
    has_chart = any(sh.has_chart for s in prs.slides for sh in s.shapes)
    has_table = any(sh.has_table for s in prs.slides for sh in s.shapes)
    assert has_chart and has_table
    modes = {s.mode for s in result.slides}
    assert "clone" in modes and "synth" in modes
    if find_soffice() and find_pdftoppm():
        from verstka.ingest.render import render_slides

        imgs = render_slides(result.pptx_path, tmp_path / "png", dpi=40)
        assert len(imgs) == 12


def test_generate_variants_offline(simple_deck, tmp_path):
    brief = parse_brief_text(BRIEF)
    res = generate_variants(simple_deck, brief=brief, strategies=["structured", "compact"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False)
    assert len(res.variants) == 2
    for v in res.variants:
        assert (v.out_dir / "deck.pptx").exists() and (v.out_dir / "outline.json").exists() and (v.out_dir / "layout_plan.json").exists()
        assert len(Presentation(str(v.out_dir / "deck.pptx")).slides) == len(v.outline.slides)
    assert len(res.variants[1].outline.slides) <= len(res.variants[0].outline.slides)


def test_variants_are_planned_side_by_side_with_a_model(simple_deck, tmp_path, monkeypatch):
    """With a model each plan is a chain of slow calls; the strategies are independent, so the deck waits for the
    slowest chain rather than the sum (the 5-minute budget). Each plan gets its own copy of the facts."""
    import threading
    import time as _time

    import verstka.pipeline.generate as gen

    analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    live, peak, seen_facts, lock = [0], [0], [], threading.Lock()

    def slow_plan(brief, manifest, strategy, facts, skills=None, providers=None, target=None):
        with lock:
            live[0] += 1
            peak[0] = max(peak[0], live[0])
            seen_facts.append(id(facts))
        _time.sleep(0.4)
        with lock:
            live[0] -= 1
        return basic_outline(brief, facts, strategy, target), []

    monkeypatch.setattr(gen, "plan_outline", slow_plan)
    providers = ProviderRegistry.mock({})
    t = _time.time()
    res = generate_variants(simple_deck, brief=parse_brief_text(BRIEF), out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", providers=providers, skills=SkillsRegistry.load(), use_vlm=False, audit=False, autofix=False, exports=[], render_images=False)
    assert [v.strategy for v in res.variants] == ["structured", "visual", "compact"]
    assert peak[0] == 3 and len(set(seen_facts)) == 3
    assert all(v.timings["plan"] >= 0.4 for v in res.variants) and _time.time() - t < 3 * 0.4 + 5


def test_a_spent_model_budget_still_delivers_the_decks(simple_deck, tmp_path):
    analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    providers = ProviderRegistry.mock({"*": {"facts": [], "series": [], "tables": []}})
    providers.limits.time_budget_s = 0
    res = generate_variants(simple_deck, brief=parse_brief_text(BRIEF), strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", providers=providers, skills=SkillsRegistry.load(), use_vlm=False, audit=False, autofix=False, exports=[], render_images=False)
    v = res.variants[0]
    assert (v.out_dir / "deck.pptx").exists() and len(v.outline.slides) >= 5
    assert any("time budget" in w for w in v.warnings) and providers.roles["llm"].calls == []
