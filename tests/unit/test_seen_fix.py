"""«Исправить слайд» по замечанию «Взгляда дизайнера» (audit/seen_fix.py, pipeline/revise.py): the remark names no
element, the rules fix it in place by measuring — a plate where no colour reads on the grounds the slide shows, a text
box cut short of the picture it runs under, no letter-spacing, a larger size — on a cover too; what cannot be fixed
in place says why; the model looks again and a remark it still sees stays (owner's remark 18.10, item 8)."""

from __future__ import annotations

import json
import time
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from verstka.analysis.xmlns import q
from verstka.audit import seen_fix as SF
from verstka.schemas.audit import Issue

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json"


@pytest.fixture
def manifest(simple_deck, tmp_path):
    from verstka.analysis.manifest import analyze_template

    return analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)


def _deck_with_text(src: Path, out: Path, *, color=(0x36, 0x39, 0x3B), spc: int = 0, wide_under_picture: bool = False) -> tuple[Path, str]:
    """Slide 2 of the simple deck with a heading-sized text box at the left (dark grey on the white slide); with
    `wide_under_picture` the box runs under a picture on the right. Returns (deck, the box's shape id)."""
    prs = Presentation(str(src))
    s = prs.slides[1]
    W, H = prs.slide_width, prs.slide_height
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith(("Pillar", "Description")):
            sh._element.getparent().remove(sh._element)
        elif sh.name.startswith("Card"):
            sh._element.getparent().remove(sh._element)
    w = int(W * (0.85 if wide_under_picture else 0.5))
    box = s.shapes.add_textbox(Emu(int(W * 0.06)), Emu(int(H * 0.40)), Emu(w), Emu(int(H * 0.2)))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Финансовый бренд года"
    r = box.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = RGBColor(*color)
    if spc:
        r._r.get_or_add_rPr().set("spc", str(spc))
    if wide_under_picture:
        pic = out.parent / "pic.png"
        Image.new("RGB", (40, 60), (40, 60, 120)).save(pic)
        s.shapes.add_picture(str(pic), Emu(int(W * 0.62)), Emu(int(H * 0.2)), Emu(int(W * 0.3)), Emu(int(H * 0.6)))
    prs.save(str(out))
    return out, str(box.shape_id)


def _picture(deck: Path, out: Path, n: int, dark_under: str | None = None) -> Path:
    """A picture of slide n as a renderer may show it: white, with a black block under shape `dark_under`."""
    prs = Presentation(str(deck))
    W, H = prs.slide_width, prs.slide_height
    im = Image.new("RGB", (400, int(400 * H / W)), (255, 255, 255))
    if dark_under:
        sh = next(s for s in prs.slides[n - 1].shapes if str(s.shape_id) == dark_under)
        d = ImageDraw.Draw(im)
        k = 400 / W
        d.rectangle([int(sh.left * k) - 4, int(sh.top * k) - 4, int((sh.left + sh.width) * k) + 4, int((sh.top + sh.height) * k) + 4], fill=(0, 0, 0))
        d.text((int(sh.left * k) + 4, int(sh.top * k) + 4), "Финансовый", fill=(54, 57, 59))
    im.save(out)
    return out


def _shape(deck: Path, n: int, sid: str):
    return next(s for s in Presentation(str(deck)).slides[n - 1].shapes if str(s.shape_id) == sid)


def _remark(code: str, slide: int, k: int = 0, where: str = "заголовок") -> Issue:
    from verstka.audit.eyes import CODES, FIX_HINTS

    return Issue(id=f"design_review-{slide}-{k}", slide=slide, check_id="design_review", severity="warn", kind="model",
                 message=f"{CODES[code].capitalize()} ({where}): тёмно-серый текст на чёрном", details={"code": code, "severity": "major"},
                 suggestion=FIX_HINTS[code])


# ---------------------------------------------------------------------------- measuring


def test_image_ground_is_the_box_majority_not_the_letters():
    im = Image.new("RGB", (100, 50), (0, 0, 0))
    ImageDraw.Draw(im).rectangle([10, 10, 40, 20], fill=(90, 90, 90))  # the letters: a minority
    got = SF.image_ground(im, (0, 0, 1000, 500), 1000, 500)
    assert got is not None and got[0] == "000000" and got[1] > 0.8
    assert SF.image_ground(im, (2000, 0, 10, 10), 1000, 500) is None  # off the picture


def test_best_color_reads_on_every_ground():
    best, worst = SF.best_color(["36393B", "FFFFFF", "4A5EE6"], ["000000"])
    assert best == "FFFFFF" and worst > 15
    # black in one render, white in another: no colour reads well on both
    _, worst = SF.best_color(["36393B", "FFFFFF", "000000"], ["000000", "FFFFFF"])
    assert worst < 4.6


def test_code_of_reads_only_the_review_remarks():
    assert SF.code_of(_remark("low_contrast", 1)) == "low_contrast"
    assert SF.code_of(Issue(id="x", slide=1, check_id="overlap", severity="warn", kind="deterministic", message="m")) is None


# ---------------------------------------------------------------------------- the fixes


def test_dark_text_on_a_block_only_the_render_shows_gets_a_plate(simple_deck, manifest, tmp_path):
    """The owner's cover: dark-grey text that PowerPoint shows on white and the previews on a black block — no colour
    reads on both, so the text gets a plate of the block's black and white letters: alike everywhere."""
    deck, sid = _deck_with_text(simple_deck, tmp_path / "d.pptx")
    pic = _picture(deck, tmp_path / "p.png", 2, dark_under=sid)
    res = SF.fix_seen(deck, 2, ["low_contrast"], manifest, pic)
    assert res.cured == {"low_contrast"} and not res.left
    assert [r["kind"] for r in res.records] == ["plate"] and res.records[0]["element_id"] == sid
    sh = _shape(deck, 2, sid)
    fill = sh._element.find(q("p:spPr")).find(q("a:solidFill"))
    assert fill is not None and fill.find(q("a:srgbClr")).get("val") == "000000"
    colors = {c.get("val") for c in sh._element.find(q("p:txBody")).iter(q("a:srgbClr"))}
    assert colors == {"FFFFFF"}


def test_the_preview_the_person_saw_counts_too(simple_deck, manifest, tmp_path):
    """The renderer drew the block black for the preview and white now: the preview the person saw is a ground too."""
    deck, sid = _deck_with_text(simple_deck, tmp_path / "d.pptx")
    now = _picture(deck, tmp_path / "now.png", 2)
    seen = _picture(deck, tmp_path / "seen.png", 2, dark_under=sid)
    res = SF.fix_seen(deck, 2, ["low_contrast"], manifest, now, [seen])
    assert [r["kind"] for r in res.records] == ["plate"]


def test_text_that_reads_everywhere_is_left_alone(simple_deck, manifest, tmp_path):
    deck, _ = _deck_with_text(simple_deck, tmp_path / "d.pptx")
    before = deck.read_bytes()
    res = SF.fix_seen(deck, 2, ["low_contrast"], manifest, _picture(deck, tmp_path / "p.png", 2))
    # nothing changed is nothing cured: the slide is laid out again for the remark (never «остался прежним»)
    assert res.cured == set() and res.untouched == {"low_contrast"} and res.records == [] and deck.read_bytes() == before


def test_text_running_under_a_picture_is_cut_short(simple_deck, manifest, tmp_path):
    deck, sid = _deck_with_text(simple_deck, tmp_path / "d.pptx", wide_under_picture=True)
    w0 = _shape(deck, 2, sid).width
    res = SF.fix_seen(deck, 2, ["overlap"], manifest)
    assert res.cured == {"overlap"} and res.records and "отодвинут" in res.records[0]["result"].lower()
    sh = _shape(deck, 2, sid)
    pic = next(s for s in Presentation(str(deck)).slides[1].shapes if s.shape_type == 13)
    assert sh.width < w0 and sh.left + sh.width < pic.left


def test_spaced_capitals_that_break_words_lose_their_spacing(simple_deck, manifest, tmp_path):
    deck, sid = _deck_with_text(simple_deck, tmp_path / "d.pptx", spc=4000)  # 40 pt between letters
    prs = Presentation(str(deck))
    sh = next(s for s in prs.slides[1].shapes if str(s.shape_id) == sid)
    sh.width = Emu(int(prs.slide_width * 0.22))
    prs.save(str(deck))
    res = SF.fix_seen(deck, 2, ["word_broken"], manifest)
    assert res.cured == {"word_broken"}, res
    rpr = _shape(deck, 2, sid)._element.find(".//" + q("a:rPr"))
    assert rpr.get("spc") == "0"


def test_a_shrunk_heading_takes_a_size_of_the_template_scale(simple_deck, manifest, tmp_path):
    deck, sid = _deck_with_text(simple_deck, tmp_path / "d.pptx")
    prs = Presentation(str(deck))
    sh = next(s for s in prs.slides[1].shapes if str(s.shape_id) == sid)
    sh.width = Emu(int(prs.slide_width * 0.17))
    prs.save(str(deck))
    res = SF.fix_seen(deck, 2, ["word_broken"], manifest)
    assert res.cured == {"word_broken"} and res.records, res
    size = int(_shape(deck, 2, sid)._element.find(".//" + q("a:rPr")).get("sz")) / 100
    typo = manifest.tokens.typography
    scale = {float(x) for x in (typo.sizes_used or [st.size_pt for st in typo.scale])}
    assert size < 28 and size in scale, (size, sorted(scale))


def test_small_text_grows_to_the_body_size(simple_deck, manifest, tmp_path):
    from verstka.audit.autofix import _text_floors

    deck, sid = _deck_with_text(simple_deck, tmp_path / "d.pptx")
    prs = Presentation(str(deck))
    sh = next(s for s in prs.slides[1].shapes if str(s.shape_id) == sid)
    sh.text_frame.paragraphs[0].runs[0].font.size = Pt(7)
    prs.save(str(deck))
    res = SF.fix_seen(deck, 2, ["too_small"], manifest)
    body, least = _text_floors(manifest)
    assert res.cured == {"too_small"} and [r["kind"] for r in res.records] == ["grow_text"]
    size = int(_shape(deck, 2, sid)._element.find(".//" + q("a:rPr")).get("sz")) / 100
    assert size >= min(max(body, least), 28) - 0.5


def test_what_cannot_be_fixed_in_place_says_why(simple_deck, manifest, tmp_path):
    deck, _ = _deck_with_text(simple_deck, tmp_path / "d.pptx")
    res = SF.fix_seen(deck, 2, ["unbalanced", "low_contrast"], manifest, None)
    assert not res.cured and set(res.left) == {"unbalanced", "low_contrast"}
    assert "пустоты" in res.left["unbalanced"] and "картинку" in res.left["low_contrast"]
    why = res.why_for(True, "обложки")
    assert why[0].isupper() and "другие образцы шаблона для обложки выходят хуже" in why and "другой вариант" in why


# ---------------------------------------------------------------------------- «Исправить слайд» end to end (offline)


def _wait(c, job_id, timeout=300):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/jobs/{job_id}").json()
        if j["status"] in ("done", "failed"):
            return j
        time.sleep(0.2)
    raise TimeoutError(job_id)


@pytest.fixture
def built(tmp_path, monkeypatch, simple_deck):
    """A variant built offline on the simple template, previews drawn without LibreOffice."""
    import importlib

    from fastapi.testclient import TestClient

    import verstka.api.app as app_module
    import verstka.pipeline.generate as G
    from verstka.pipeline import revise as R

    monkeypatch.setenv("VERSTKA_WORKSPACE", str(tmp_path / "ws"))
    importlib.reload(app_module)

    def fake(vdir, manifest, title, exports, *, images=True, dpi=110):
        out = []
        (Path(vdir) / "slides").mkdir(exist_ok=True)
        for k, fp in enumerate(R.slide_fingerprints(Path(vdir) / "deck.pptx"), 1):
            path = Path(vdir) / "slides" / f"slide-{k:03d}.jpg"
            Image.new("RGB", (64, 36), tuple(int(fp[i:i + 2], 16) for i in (0, 2, 4))).save(path)
            out.append(path)
        return {}, out, [], None

    monkeypatch.setattr(G, "render_outputs", fake)
    c = TestClient(app_module.app)
    with open(simple_deck, "rb") as f:
        r = c.post("/api/templates", files={"file": ("simple.pptx", f, "application/octet-stream")}, data={"use_models": "false"})
    tid = _wait(c, r.json()["job_id"])["result"]["template_id"]
    outline = json.loads(FIXTURE.read_text(encoding="utf-8"))
    r = c.post("/api/generations", json={"template_id": tid, "outline": outline, "strategies": ["structured"], "use_models": False, "exports": []})
    gid = r.json()["generation_id"]
    assert _wait(c, r.json()["job_id"])["status"] == "done"
    return c, app_module, tid, gid, app_module.store.generation_dir(gid) / "structured"


def _with_remark(vdir: Path, remark: Issue) -> None:
    from verstka.schemas.audit import AuditReport

    rep = AuditReport.model_validate_json((vdir / "audit_report.json").read_text(encoding="utf-8"))
    rep.issues = [i for i in rep.issues if i.slide != remark.slide or not (i.severity in ("error", "warn") or i.kind == "model")] + [remark]
    rep.recompute()
    (vdir / "audit_report.json").write_text(rep.model_dump_json(indent=2), encoding="utf-8")


def _title_id(vdir: Path) -> str:
    s = Presentation(str(vdir / "deck.pptx")).slides[0]
    return str(next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()).shape_id)


def test_cover_contrast_remark_is_fixed_by_the_button(built, monkeypatch, tmp_path):
    """The owner's case: «Исправить слайд» on the cover with the review's «тёмно-серый текст на чёрном» — the cover
    is never laid out again, yet the slide changes (a plate under the heading) and the remark is gone."""
    c, mod, tid, gid, vdir = built
    sid = _title_id(vdir)
    _with_remark(vdir, _remark("low_contrast", 1))
    monkeypatch.setattr(SF, "slide_picture", lambda deck, n, out_dir, timeout_s=120.0: _picture(Path(deck), Path(tmp_path) / f"shot-{time.time_ns()}.png", n, dark_under=sid))
    r = c.post(f"/api/generations/{gid}/structured/slides/1/fix", json={"issue_ids": ["design_review-1-0"]})
    job = _wait(c, r.json()["job_id"])
    assert job["status"] == "done", job["error"]
    res = job["result"]
    assert res["applied"] is True and res["fixed"] == ["design_review-1-0"], res["reply"]
    assert "подложка" in res["reply"] and res["why"] is None
    audit = json.loads((vdir / "audit_report.json").read_text(encoding="utf-8"))
    assert not any(i["check_id"] == "design_review" and i["slide"] == 1 for i in audit["issues"])
    fill = _shape(vdir / "deck.pptx", 1, sid)._element.find(q("p:spPr")).find(q("a:solidFill"))
    assert fill is not None


def test_a_remark_the_model_still_sees_stays_and_says_so(built, monkeypatch, tmp_path):
    from verstka.pipeline import revise as R

    c, mod, tid, gid, vdir = built
    sid = _title_id(vdir)
    remark = _remark("low_contrast", 1)
    _with_remark(vdir, remark)
    before = (vdir / "deck.pptx").read_bytes()
    monkeypatch.setattr(SF, "slide_picture", lambda deck, n, out_dir, timeout_s=120.0: _picture(Path(deck), Path(tmp_path) / f"shot-{time.time_ns()}.png", n, dark_under=sid))
    monkeypatch.setattr(SF, "relook", lambda *a, **kw: ["low_contrast"])
    res = R.fix_slide(vdir, "structured", 1, manifest=mod.store.manifest(tid), ws=mod.store.workspace(tid), remarks=[remark], use_models=True, exports=[])
    assert not res.applied and res.fixed == []
    assert res.why.startswith("Дизайнер всё ещё видит: текст плохо читается на фоне")
    assert (vdir / "deck.pptx").read_bytes() == before


def test_a_remark_that_cannot_be_fixed_on_a_cover_says_why(built, monkeypatch):
    import verstka.audit.eyes as E
    from verstka.pipeline import revise as R

    monkeypatch.setattr(E, "relayout", lambda *a, **kw: "skip")  # no other sample for this cover

    c, mod, tid, gid, vdir = built
    remark = _remark("unbalanced", 1, where="весь слайд")
    _with_remark(vdir, remark)
    res = R.fix_slide(vdir, "structured", 1, manifest=mod.store.manifest(tid), ws=mod.store.workspace(tid), remarks=[remark], use_models=False, exports=[])
    assert not res.applied
    assert res.why.startswith("Пустоты и перекос на месте не исправить") and "другие образцы шаблона для обложки выходят хуже" in res.why


def test_letters_below_a_code_box_widen_it_to_its_panel(simple_deck, manifest, tmp_path):
    """The owner's code slide: the renderer wraps a code font wider than measured and the last line falls out under
    the black panel — the picture shows it, the box takes the whole panel."""
    prs = Presentation(str(simple_deck))
    s = prs.slides[1]
    W, H = prs.slide_width, prs.slide_height
    for sh in list(s.shapes):  # the slide holds the panel and its code only
        sh._element.getparent().remove(sh._element)
    panel = s.shapes.add_shape(1, Emu(int(W * 0.45)), Emu(int(H * 0.2)), Emu(int(W * 0.5)), Emu(int(H * 0.5)))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0, 0, 0)
    box = s.shapes.add_textbox(Emu(int(W * 0.47)), Emu(int(H * 0.22)), Emu(int(W * 0.3)), Emu(int(H * 0.46)))
    box.text_frame.word_wrap = True
    box.text_frame.text = "\n".join(f"line {k}" for k in range(8))
    deck = tmp_path / "code.pptx"
    prs.save(str(deck))
    sid = str(box.shape_id)
    im = Image.new("RGB", (400, int(400 * H / W)), (250, 250, 250))
    d = ImageDraw.Draw(im)
    k = 400 / W
    d.rectangle([int(panel.left * k), int(panel.top * k), int((panel.left + panel.width) * k), int((panel.top + panel.height) * k)], fill=(0, 0, 0))
    y = int((box.top + box.height) * k) + 18  # a line of letters under the panel
    d.rectangle([int(box.left * k) + 10, y, int(box.left * k) + 90, y + 4], fill=(120, 120, 120))
    pic = tmp_path / "shot.png"
    im.save(pic)
    from verstka.audit.ir import build_deck_ir

    ir = build_deck_ir(deck, with_images=False)
    e = next(x for x in ir.slides[1].elements if str(x.id) == sid)
    assert SF.ink_below(Image.open(pic), e, ir.slides[1], ir.slide_w, ir.slide_h) > 0.02
    recs, touched, still = SF.fix_spill(deck, 2, manifest, pic)
    assert touched == {sid} and "растянут на всю карточку" in recs[0]["result"]
    sh = _shape(deck, 2, sid)
    assert sh.left + sh.width > box.left + box.width and sh.left + sh.width <= panel.left + panel.width


def test_a_remark_with_nothing_to_fix_in_place_lays_the_slide_out_again(built, monkeypatch):
    """«Исправить слайд» never answers «остался прежним» while a remark is on the slide: the rules found nothing to
    change and there is no model for the designer: the slide is laid out on another sample of the template (or
    composed), only that slide changes, and the reply says what was done."""
    from verstka.pipeline import revise as R

    c, mod, tid, gid, vdir = built
    remark = _remark("unbalanced", 3, where="весь слайд")
    _with_remark(vdir, remark)
    before = R.slide_fingerprints(vdir / "deck.pptx")
    res = R.fix_slide(vdir, "structured", 3, manifest=mod.store.manifest(tid), ws=mod.store.workspace(tid), remarks=[remark], use_models=False, exports=[])
    assert res.applied and res.how == "relayout" and res.fixed == [remark.id], res.reply
    assert "собран" in res.reply
    after = R.slide_fingerprints(vdir / "deck.pptx")
    assert after[2] != before[2] and R.changed_slides(before, after, skip={3}) == []


def test_nothing_left_to_change_says_the_remarks_are_gone(built, monkeypatch):
    """The slide cannot be changed (no other layout) and the model, looking again, sees nothing: the remark goes and
    the reply says «замечаний больше нет» — not «остался прежним, добавьте пожелание»."""
    import verstka.audit.eyes as E
    from verstka.pipeline import revise as R

    c, mod, tid, gid, vdir = built
    remark = _remark("unbalanced", 1, where="весь слайд")
    _with_remark(vdir, remark)
    before = (vdir / "deck.pptx").read_bytes()
    monkeypatch.setattr(E, "relayout", lambda *a, **kw: "skip")
    monkeypatch.setattr(SF, "slide_picture", lambda deck, n, out_dir, timeout_s=120.0: _picture(Path(deck), Path(out_dir).parent / f"shot-{time.time_ns()}.png", n))
    monkeypatch.setattr(SF, "relook", lambda *a, **kw: [])
    res = R.fix_slide(vdir, "structured", 1, manifest=mod.store.manifest(tid), ws=mod.store.workspace(tid), remarks=[remark], use_models=True, exports=[])
    assert not res.applied and res.why.startswith("Замечаний больше нет") and res.fixed == [remark.id]
    assert res.reply.startswith("Слайд 1: замечаний больше нет")
    assert (vdir / "deck.pptx").read_bytes() == before
    audit = json.loads((vdir / "audit_report.json").read_text(encoding="utf-8"))
    assert not any(i["id"] == remark.id for i in audit["issues"])
