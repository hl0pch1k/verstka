from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.audit.autofix import autofix_loop, plan_fixes
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import all_checks
from verstka.audit.runner import run_audit
from verstka.ingest.workspace import TemplateWorkspace
from verstka.matching.matcher import match_outline
from verstka.planning.strategies import get_strategy
from verstka.rendering.renderer import render_deck
from verstka.schemas.outline import DeckOutline

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json"
W, H = 12192000, 6858000


def _defective_deck(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    # slide 1: clean title
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s1.shapes.add_textbox(Emu(int(W * 0.1)), Emu(int(H * 0.3)), Emu(int(W * 0.8)), Emu(int(H * 0.2)))
    tb.text_frame.text = "Чистый заголовок"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    # slide 2: overflow + overlap + foreign font + placeholder + off-slide
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    a = s2.shapes.add_textbox(Emu(int(W * 0.1)), Emu(int(H * 0.2)), Emu(int(W * 0.3)), Emu(int(H * 0.08)))
    a.text_frame.word_wrap = True
    a.text_frame.text = "Очень длинный текст, который никак не поместится в такую маленькую рамку, потому что слов слишком много и они продолжаются и продолжаются без конца"
    a.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    b = s2.shapes.add_textbox(Emu(int(W * 0.15)), Emu(int(H * 0.22)), Emu(int(W * 0.3)), Emu(int(H * 0.08)))
    b.text_frame.text = "Наложенный блок"
    b.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    b.text_frame.paragraphs[0].runs[0].font.name = "Comic Sans MS"
    c = s2.shapes.add_textbox(Emu(int(W * 0.95)), Emu(int(H * 0.6)), Emu(int(W * 0.2)), Emu(int(H * 0.1)))
    c.text_frame.text = "Lorem ipsum dolor"
    # slide 3: eight long bullets in one box
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    d = s3.shapes.add_textbox(Emu(int(W * 0.1)), Emu(int(H * 0.15)), Emu(int(W * 0.8)), Emu(int(H * 0.7)))
    tf = d.text_frame
    tf.text = "Заголовок списка"
    for i in range(8):
        p = tf.add_paragraph()
        p.text = ("слово " * 18).strip()
        p.level = 1
    # slide 4: only a title
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    e = s4.shapes.add_textbox(Emu(int(W * 0.1)), Emu(int(H * 0.1)), Emu(int(W * 0.8)), Emu(int(H * 0.12)))
    e.text_frame.text = "Одинокий заголовок"
    # slide 5: clean content
    s5 = prs.slides.add_slide(prs.slide_layouts[6])
    f = s5.shapes.add_textbox(Emu(int(W * 0.1)), Emu(int(H * 0.1)), Emu(int(W * 0.8)), Emu(int(H * 0.12)))
    f.text_frame.text = "Финал"
    prs.save(str(path))
    return path


def test_registry_and_defective_deck(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ids = {spec.id for spec, _ in all_checks()}
    for expected in ("out_of_bounds", "overlap", "text_overflow", "font_not_in_template", "placeholder_text", "empty_slide", "too_many_bullets", "bullet_too_long", "fill_ratio", "contrast_low", "chart_missing_labels", "duplicate_slides"):
        assert expected in ids
    deck = _defective_deck(tmp_path / "bad.pptx")
    ir = build_deck_ir(deck)
    assert ir.n_slides == 5 and ir.slides[1].texts
    report = run_audit(deck, manifest, render=False)
    by_check = {}
    for i in report.issues:
        by_check.setdefault(i.check_id, set()).add(i.slide)
    assert 2 in by_check["text_overflow"]
    assert 2 in by_check["overlap"]
    assert 2 in by_check["out_of_bounds"] or 2 in by_check.get("text_clipped", set())
    assert 2 in by_check["placeholder_text"]
    assert 3 in by_check["too_many_bullets"] and 3 in by_check["bullet_too_long"]
    assert 4 in by_check["empty_slide"]
    assert 1 not in by_check.get("text_overflow", set()) and 1 not in by_check.get("overlap", set())
    assert report.summary.errors >= 4 and report.summary.score < 100
    # bullets in a box longer than 15 words must be autofixable
    assert any(i.autofix and i.autofix.action == "condense_text" for i in report.issues if i.check_id == "bullet_too_long")


def test_generated_deck_audits_and_autofix(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    outline = DeckOutline.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    plan = match_outline(outline, manifest, get_strategy("structured"))
    out = tmp_path / "deck.pptx"
    render_deck(outline, plan, manifest, ws, out)
    ir = build_deck_ir(out)
    assert all(s.outline_id for s in ir.slides), [s.outline_id for s in ir.slides]
    chart = next(e for s in ir.slides for e in s.elements if e.type == "chart")
    assert chart.chart and chart.chart.series and chart.chart.series[0].values[-1] == 12400 and chart.chart.has_data_labels
    table = next(e for s in ir.slides for e in s.elements if e.type == "table")
    assert table.table and table.table.rows[0][0] == "Сценарий"
    report = run_audit(out, manifest, outline, ws, render=False, strategy="structured")
    assert not [i for i in report.issues if i.check_id == "placeholder_text"]
    fixes = plan_fixes(report, plan)
    final, plan2, outline2, _ = autofix_loop(out, report, outline, plan, manifest, ws, max_iterations=2)
    assert final.iterations <= 2
    assert final.summary.errors <= report.summary.errors
