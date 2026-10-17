"""Вторая ступень, HTML-версия графиков (verstka.export.svg_charts) и аудит: гистограмма, рассеяние, накопительные и
сгруппированные столбцы/полосы, разрывы линии, подписи осей, строка источника — разметкой SVG, как в PPTX; IR с
пропусками (schemas/deck_ir.py, audit/ir.py) и проверка author_item_missing на графике с пропусками."""

from __future__ import annotations

import re

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from verstka.audit.ir import build_deck_ir
from verstka.export.svg_charts import chart_svg
from verstka.rendering.charts import SEAM_PT, add_chart, label_on
from verstka.schemas.common import contrast_ratio
from verstka.schemas.common import Bbox
from verstka.schemas.deck_ir import IRAxisScale, IRChart, IRChartSeries, IRPointLabel
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, OutlineSlide, SlideContent, XYSeries
from verstka.schemas.template import ChartStyleSpec, TypeStep, Typography

W, H = 12192000, 6858000
BOX = Bbox(x=int(W * 0.06), y=int(H * 0.22), w=int(W * 0.88), h=int(H * 0.68))
PALETTE = ["0077FF", "00D3E6", "005EFF", "FF6C6C", "0563C1"]
STYLE = ChartStyleSpec(series_colors=PALETTE, font_family="Play", font_size_pt=7.0)
TYPO = Typography(scale=[TypeStep(role="display", size_pt=47), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=12), TypeStep(role="body", size_pt=9), TypeStep(role="small", size_pt=7)])
OUTLINE = DeckOutline(title="t")
BINS = ["0–30", "30–60", "60–90", "90–120", "120–180", "180–240"]
COUNTS = [148, 121, 63, 78, 52, 22]
QS = ["IV кв. 2024", "I кв. 2025", "II кв. 2025", "III кв. 2025", "IV кв. 2025", "I кв. 2026", "II кв. 2026", "III кв. 2026"]
NPS = [41, 43, None, 46, 48, 49, None, 52]
STACK = [InlineSeries(name="Прямые продажи", values=[120, 134, 142, 151]), InlineSeries(name="Партнёры", values=[71, 80, 86, 92]), InlineSeries(name="Прочее", values=[22, 21, 25, 27])]
REG = ["Москва", "Санкт-Петербург", "Казань"]
XS = [120, 180, 220, 260, 310, 350, 400, 450, 520, 590]
YS = [110, 160, 230, 240, 300, 330, 390, 420, 500, 560]

RECT_RE = re.compile(r'<rect x="([\d.\-]+)" y="([\d.\-]+)" width="([\d.\-]+)" height="([\d.\-]+)" fill="#([0-9A-F]{6})"/>')
CIRCLE_RE = re.compile(r'<circle cx="([\d.\-]+)" cy="([\d.\-]+)" r="([\d.\-]+)" fill="#([0-9A-F]{6})"([^/]*)/>')


def _ir(spec: ChartSpec, tmp_path, name: str = "c.pptx", style: ChartStyleSpec = STYLE) -> IRChart:
    """The chart the PPTX holds, read back by the audit's IR (what the HTML export draws)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    add_chart(s, BOX, spec, OUTLINE, style, TYPO, text_hex="000000")
    path = tmp_path / name
    prs.save(str(path))
    ir = build_deck_ir(path, with_images=False)
    return next(e.chart for e in ir.slides[0].elements if e.type == "chart")


def _svg(chart: IRChart) -> str:
    return chart_svg(chart, BOX.w / 12700, BOX.h / 12700, PALETTE, text_hex="000000", font="Play", font_px=9.0)


def _rects(svg: str) -> list[tuple[float, float, float, float, str]]:
    return [(float(x), float(y), float(w), float(h), c) for x, y, w, h, c in RECT_RE.findall(svg) if float(h) > 2 and float(w) > 2 and float(w) != float(h)]


def _texts(svg: str) -> list[str]:
    return [re.sub(r"<[^>]+>", "", t) for t in re.findall(r"<text[^>]*>.*?</text>", svg)]


# ---------------------------------------------------------------------------------------------- the IR


def test_ir_series_keep_gaps_as_none():
    s = IRChartSeries(name="NPS", values=[41, None, 46])  # the schema takes a gap (no setting after validation)
    assert s.values == [41.0, None, 46.0]
    assert IRChartSeries.model_validate(s.model_dump()).values == [41.0, None, 46.0]


def test_ir_reads_grouping_gaps_titles_and_scales(tmp_path):
    hist = _ir(ChartSpec(type="histogram", categories=BINS, series=[InlineSeries(name="Обращения", values=COUNTS)], x_title="Минут до ответа", y_title="Число обращений"), tmp_path, "h.pptx")
    assert hist.type == "column" and hist.gap_width == 0 and hist.grouping == "clustered"
    assert hist.x_title == "Минут до ответа" and hist.y_title == "Число обращений" and hist.title is None
    stacked = _ir(ChartSpec(type="stacked_bar", categories=["I кв.", "II кв.", "III кв.", "IV кв."], series=STACK, unit="млн ₽"), tmp_path, "s.pptx")
    assert stacked.type == "bar" and stacked.grouping == "stacked" and stacked.overlap == 100 and stacked.gap_width is not None
    line = _ir(ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True, y_title="NPS"), tmp_path, "l.pptx")
    assert line.series[0].values == [41.0, 43.0, None, 46.0, 48.0, 49.0, None, 52.0]
    assert line.series[0].point_sizes == {7: 11.0} and line.series[0].marker_size == 8.0  # the end point, larger
    xy = _ir(ChartSpec(type="scatter", xy=[XYSeries(name="Клиенты", x=XS, y=YS)], x_title="Средний чек", y_title="Выручка", highlight_index=2), tmp_path, "x.pptx")
    ser = xy.series[0]
    assert xy.type == "scatter" and ser.x_values[:3] == [120.0, 180.0, 220.0] and ser.values[:3] == [110.0, 160.0, 230.0]
    assert xy.x_title == "Средний чек" and xy.y_title == "Выручка"
    assert xy.x_scale is not None and xy.x_scale.max >= 590 and xy.x_scale.step and xy.value_scale.max >= 560
    assert ser.color and ser.marker_size and 0 < (ser.marker_alpha or 0) < 1 and ser.marker_line == "FFFFFF"
    assert ser.point_sizes[2] > ser.marker_size and ser.point_colors[2] == "0077FF"
    assert xy.axis_color and xy.rule_color  # the X axis's own colours, as a category axis's


def test_author_item_missing_reads_a_chart_with_gaps():
    """A list of the author drawn as a line with a missing quarter: the check reads the chart's figures and skips the
    gap, never falls over on it (the registry would silently drop the check for the whole deck)."""
    from test_tpl_clone_audit import _ctx, _ir_slide

    from verstka.audit.checks.integrity import author_item_missing
    from verstka.schemas.common import Bbox as B
    from verstka.schemas.common import BboxFrac, PatternKind
    from verstka.schemas.deck_ir import IRElement

    chart = IRElement(id="c1", type="chart", bbox=B(x=0, y=0, w=100, h=100), bbox_frac=BboxFrac(x=0.1, y=0.3, w=0.8, h=0.6),
                      chart=IRChart(type="line", categories=["I кв.", "II кв.", "III кв."], series=[IRChartSeries(name="NPS", values=[41.5, None, 46])]))
    brief = "Слайд 1. Обложка\n\nСлайд 2. NPS по кварталам\n— 41,5\n— 46\n— 58\n\nСлайд 3. Итоги\n"
    o = DeckOutline(title="x", slides=[OutlineSlide(id="s2", kind=PatternKind.chart, headline="NPS по кварталам", spec_ref=2, verbatim=True, content=SlideContent())])
    ctx = _ctx([_ir_slide(1, [("NPS по кварталам", (0.05, 0.05, 0.9, 0.1), 32)], oid="s2", extra=[chart])], o)
    ctx.brief_text = brief
    got = author_item_missing(ctx)
    assert len(got) == 1 and got[0].details["missing"] == ["58"]  # 41,5 and 46 are on the chart; 58 is not


# ---------------------------------------------------------------------------------------------- the SVG


def test_histogram_bins_touch_parted_by_a_hairline_of_the_ground(tmp_path):
    ch = _ir(ChartSpec(type="histogram", categories=BINS, series=[InlineSeries(name="Обращения", values=COUNTS)], highlight_index=0, x_title="Минут до ответа"), tmp_path)
    svg = _svg(ch)
    bars = sorted(_rects(svg))
    assert len(bars) == len(BINS)
    slot = bars[1][0] - bars[0][0]
    for a, b in zip(bars, bars[1:]):
        assert b[0] - a[0] == pytest.approx(slot, abs=0.2)  # every bin one slot on
        assert b[0] - (a[0] + a[2]) == pytest.approx(SEAM_PT, abs=0.15)  # a hairline apart, not a column gap
    assert bars[0][4] == "0077FF" and {b[4] for b in bars[1:]} == {ch.series[0].color}  # the highlighted bin in the accent
    assert "148" in _texts(svg) and "Минут до ответа" in _texts(svg)


def test_stacked_columns_stack_their_segments_with_values_inside(tmp_path):
    ch = _ir(ChartSpec(type="stacked_column", categories=["I кв.", "II кв.", "III кв.", "IV кв."], series=STACK, unit="млн ₽", y_title="Выручка"), tmp_path)
    svg = _svg(ch)
    bars = _rects(svg)
    assert len(bars) == 12
    first = sorted((b for b in bars if abs(b[0] - min(x for x, *_ in bars)) < 0.5), key=lambda b: -b[1])  # bottom first
    assert [b[4] for b in first] == [s.color for s in ch.series]  # Прямые at the base, «Прочее» quiet on top
    for lower, upper in zip(first, first[1:]):
        assert lower[1] - (upper[1] + upper[3]) == pytest.approx(SEAM_PT, abs=0.15)  # one on another, a hairline apart
    heights = [b[3] + SEAM_PT for b in first]
    assert heights[0] / heights[1] == pytest.approx(120 / 71, rel=0.03)
    # each value inside its own segment, in the colour the chart gives it (white on the accent)
    m = re.search(r'<text x="([\d.]+)" y="([\d.]+)"[^>]*fill="#([0-9A-F]{6})"[^>]*>120</text>', svg)
    assert m and first[0][1] < float(m.group(2)) < first[0][1] + first[0][3] and m.group(3) == "FFFFFF"
    assert all(n in _texts(svg) for n in ("Прямые продажи", "Партнёры", "Прочее"))  # the legend
    assert re.search(r'transform="rotate\(-90[^"]*"[^>]*>Выручка, млн ₽</text>', svg)  # the value axis's title, turned


def _ink_of(svg: str, label: str) -> str:
    """The fill of the SVG <text> that prints `label`."""
    m = re.search(rf'<text[^>]*fill="#([0-9A-F]{{6}})"[^>]*>{re.escape(label)}</text>', svg)
    assert m, label
    return m.group(1)


def test_values_on_segments_follow_the_pptx_rule_white_on_blue_and_pink(tmp_path):
    # one rule for the PPTX and the HTML: the template's white on its accents from 3:1 (0077FF 4.1:1, FF3985 3.3:1) —
    # and the same ink in both for the same fill; a light template colour (8BBFFA) no longer paints a series as it is: a
    # series stands off its ground at 3:1 (a deeper shade), so its values read white as well
    style = ChartStyleSpec(series_colors=["0077FF", "FF3985", "8BBFFA"], font_family="Play", font_size_pt=7.0)
    cats = ["I кв.", "II кв.", "III кв."]
    vals = {"Один": [51, 61, 71], "Два": [32, 42, 52], "Три": [23, 33, 43]}
    series = [InlineSeries(name=n, values=v) for n, v in vals.items()]
    for kind in ("stacked_column", "stacked_bar"):
        ch = _ir(ChartSpec(type=kind, categories=cats, series=series), tmp_path, f"{kind}.pptx", style)
        assert [s.color for s in ch.series][:2] == ["0077FF", "FF3985"] and contrast_ratio(ch.series[2].color, "FFFFFF") >= 3.0
        svg = _svg(ch)
        for s in ch.series:
            for v in vals[s.name]:
                ink = _ink_of(svg, str(v))
                assert ink == s.label_color == label_on(s.color, "000000", "FFFFFF"), (kind, s.color, ink, s.label_color)  # PPTX == SVG
        assert _ink_of(svg, "51") == "FFFFFF" and _ink_of(svg, "32") == "FFFFFF"  # the brand blue and the pink keep white


def test_values_on_a_dark_deck_take_the_ground_on_a_light_fill_and_stay_white_on_the_accent(tmp_path):
    style = ChartStyleSpec(series_colors=["0077FF", "8BBFFA"], font_family="Play", font_size_pt=7.0)
    ch = _ir(ChartSpec(type="stacked_column", categories=["I", "II"], series=[InlineSeries(name="Один", values=[51, 61]), InlineSeries(name="Два", values=[23, 33])]), tmp_path, "d.pptx", style)
    svg = chart_svg(ch, BOX.w / 12700, BOX.h / 12700, PALETTE, text_hex="FFFFFF", font="Play", font_px=9.0, ground="0B1220")
    assert _ink_of(svg, "51") == "FFFFFF"
    assert label_on("8BBFFA", "FFFFFF", "0B1220") == "0B1220"  # a light fill: white text does not read on it, the ground does
    assert _ink_of(svg, "23") == label_on(ch.series[1].color, "FFFFFF", "0B1220")  # (the series' own colour, as the PPTX paints it)


def test_pie_and_doughnut_slice_values_follow_the_same_rule(tmp_path):
    style = ChartStyleSpec(series_colors=["0077FF", "FF3985", "8BBFFA"], font_family="Play", font_size_pt=7.0)
    for kind in ("pie", "doughnut"):
        ch = _ir(ChartSpec(type=kind, categories=["a", "b", "c"], series=[InlineSeries(name="Доли", values=[50, 30, 20])], unit="%"), tmp_path, f"{kind}.pptx", style)
        s = ch.series[0]
        svg = _svg(ch)
        inks = re.findall(r'<text[^>]*fill="#([0-9A-F]{6})"[^>]*>(\d+%?)</text>', svg)
        assert len(inks) == 3 and len(s.point_colors) == 3
        for j, (ink, _) in enumerate(inks):
            fill = s.point_colors[j]
            assert ink == s.point_labels[j].color == label_on(fill, "000000", "FFFFFF"), (kind, j, fill)  # PPTX == SVG
        assert inks[0][0] == "FFFFFF"  # the first slice is the accent 0077FF


def test_stacked_bars_run_left_to_right_top_down(tmp_path):
    ch = _ir(ChartSpec(type="stacked_bar", categories=REG, series=[InlineSeries(name="Новые", values=[210, 120, 95]), InlineSeries(name="Постоянные", values=[480, 260, 150])], unit="млн ₽", x_title="Регион", y_title="Выручка"), tmp_path)
    svg = _svg(ch)
    bars = _rects(svg)
    moscow = sorted((b for b in bars if abs(b[1] - min(y for _, y, *_ in bars)) < 0.5), key=lambda b: b[0])
    assert len(moscow) == 2 and moscow[1][0] - (moscow[0][0] + moscow[0][2]) == pytest.approx(SEAM_PT, abs=0.15)
    assert moscow[1][2] / moscow[0][2] == pytest.approx(480 / 210, rel=0.03)
    # horizontal bars: the categories' title on the left (turned), the values' title under the plot
    assert re.search(r'rotate\(-90[^"]*"[^>]*>Регион</text>', svg) and re.search(r'<text x="[\d.]+" y="[\d.]+"[^>]*>Выручка, млн ₽</text>', svg)


def test_grouped_bars_side_by_side_with_the_highlighted_category(tmp_path):
    spec = ChartSpec(type="grouped_bar", categories=REG, series=[InlineSeries(name="2025", values=[480, 260, 150]), InlineSeries(name="2026", values=[560, 300, 207])], unit="млн ₽", highlight_index=2)
    ch = _ir(spec, tmp_path)
    svg = _svg(ch)
    bars = sorted(_rects(svg), key=lambda b: b[1])
    assert len(bars) == 6
    a, b = bars[0], bars[1]  # Москва: 2025 above 2026, just apart (overlap −8 % of a bar)
    assert b[1] - (a[1] + a[3]) == pytest.approx(0.08 * a[3], abs=0.3)
    assert "2025: 480" in _texts(svg)  # the series named on the first row instead of a legend
    kazan = bars[4:]
    assert {x[4] for x in kazan} == {s.color for s in ch.series} and not {x[4] for x in bars[:4]} & {s.color for s in ch.series}
    col = _ir(spec.model_copy(update={"type": "grouped_column", "categories": ["I", "II", "III"]}), tmp_path, "gc.pptx")
    cols = _rects(_svg(col))
    assert len(cols) == 6 and len({round(c[2], 1) for c in cols}) == 1


def test_line_breaks_at_a_gap_and_ends_on_its_last_value(tmp_path):
    ch = _ir(ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True, y_title="NPS", x_title="Квартал"), tmp_path)
    svg = _svg(ch)
    lines = re.findall(r'<polyline points="([^"]+)"', svg)
    assert [len(p.split()) for p in lines] == [2, 3]  # 41–43 | 46–48–49; the lone 52 is a point, never joined across
    dots = CIRCLE_RE.findall(svg)
    assert len(dots) == 6 and float(dots[-1][2]) > float(dots[0][2]) and dots[-1][3] == "0077FF"  # the end point larger
    assert all("stroke=" in d[4] for d in dots)  # ringed in the ground's colour
    assert "52" in _texts(svg) and "NPS" in _texts(svg) and "Квартал" in _texts(svg)
    # a missing tail: the line ends (and is labelled) on its last point with a value
    tail = IRChart(type="line", categories=["a", "b", "c", "d"], series=[IRChartSeries(name="s", values=[3, 4, 5, None], point_labels={2: IRPointLabel(position="r")})], has_data_labels=False)
    svg = chart_svg(tail, 600, 300, PALETTE)
    assert len(re.findall(r"<polyline", svg)) == 1 and "5" in _texts(svg)
    area = IRChart(type="area", categories=list("abcde"), series=[IRChartSeries(name="s", values=[3, 4, None, 5, 6])])
    assert len(re.findall(r'<path d="M', chart_svg(area, 600, 300, PALETTE))) == 2  # two areas, not one over a zero


def test_scatter_is_translucent_points_on_two_round_axes(tmp_path):
    ch = _ir(ChartSpec(type="scatter", xy=[XYSeries(name="Клиенты", x=XS, y=YS)], unit="тыс. ₽", x_title="Средний чек", y_title="Выручка на клиента", highlight_index=2), tmp_path)
    svg = _svg(ch)
    assert "<polyline" not in svg and not _rects(svg)  # points only
    dots = CIRCLE_RE.findall(svg)
    assert len(dots) == len(XS)
    plain = [d for d in dots if 'fill-opacity="' in d[4]]
    assert len(plain) == len(XS) - 1 and all(f'stroke="#{ch.series[0].marker_line}"' in d[4] for d in dots)
    last = dots[-1]  # the highlighted point: larger, in the accent, drawn over the others
    assert last[3] == "0077FF" and float(last[2]) > float(plain[0][2]) and "fill-opacity" not in last[4]
    texts = _texts(svg)
    xs = ch.x_scale
    ticks = [round(xs.min + i * xs.step) for i in range(int((xs.max - xs.min) / xs.step) + 1)]
    assert all(str(t) in texts for t in ticks if t < 1000)  # the X ticks the PPTX wrote
    assert "230" in texts and "Средний чек" in texts and any(t.startswith("Выручка на клиента") for t in texts)
    # the point is placed by its X: the larger X, the further right
    cx = [float(d[0]) for d in plain]
    assert cx == sorted(cx)


def test_scatter_of_several_series_has_a_legend():
    ch = IRChart(type="scatter", has_legend=True, series=[IRChartSeries(name="A", values=[1, 2], x_values=[1, 2], color="0077FF", marker_size=7), IRChartSeries(name="B", values=[3, 1], x_values=[1.5, 3], color="FF6C6C", marker_size=7)],
                 value_scale=IRAxisScale(min=0, max=4, step=1, format="#,##0"), x_scale=IRAxisScale(min=0, max=4, step=1, format="#,##0"))
    svg = chart_svg(ch, 600, 300, PALETTE)
    assert len(CIRCLE_RE.findall(svg)) == 4 and "A" in _texts(svg) and "B" in _texts(svg)
    assert {"0", "1", "2", "3", "4"} <= set(_texts(svg))


def test_axis_titles_take_their_room_and_stay_inside_the_frame():
    base = IRChart(type="column", categories=["a", "b"], series=[IRChartSeries(name="s", values=[1, 2])], has_data_labels=True, has_value_axis=False)
    plain = _rects(chart_svg(base, 600, 300, PALETTE))
    titled = chart_svg(base.model_copy(update={"x_title": "Квартал", "y_title": "Выручка, млн ₽"}), 600, 300, PALETTE)
    assert '<g transform="translate(' in titled  # the plot moves in, right of the turned title
    t_bars = _rects(titled)
    assert max(b[1] + b[3] for b in t_bars) < max(b[1] + b[3] for b in plain)  # and up, over the bottom title
    m = re.search(r'<text x="([\d.]+)" y="([\d.]+)" transform="rotate\(-90[^>]*>Выручка, млн ₽</text>', titled)
    assert m and 0 < float(m.group(1)) < 20 and 0 < float(m.group(2)) < 300
    # a pie takes none
    pie = IRChart(type="pie", categories=["a", "b"], series=[IRChartSeries(name="s", values=[1, 2])], x_title="x", y_title="y")
    assert "rotate(-90" not in chart_svg(pie, 400, 300, PALETTE)


# ---------------------------------------------------------------------------------------------- end to end


def test_html_export_draws_every_new_kind(simple_deck, tmp_path):
    """A deck of every new chart kind built without models, exported to HTML: each chart is SVG markup (no picture of
    it), the source line is the slide's own text, once."""
    from verstka.pipeline.generate import generate_variants

    slides = [
        OutlineSlide(id="s1", kind="title", headline="Продажи в 2026 году", subtitle="Итоги девяти месяцев"),
        OutlineSlide(id="s2", kind="chart", headline="80% обращений закрываются за 2 часа",
                     content=SlideContent(chart=ChartSpec(type="histogram", categories=BINS, series=[InlineSeries(name="Обращения", values=COUNTS)], x_title="Минут до ответа", y_title="Число обращений", source="Время ответа поддержки.csv"))),
        OutlineSlide(id="s3", kind="chart", headline="NPS вырос до 52", content=SlideContent(chart=ChartSpec(type="line", categories=QS, series=[InlineSeries(name="NPS", values=NPS)], gaps=True, y_title="NPS"))),
        OutlineSlide(id="s4", kind="chart", headline="Прямые продажи растут", content=SlideContent(chart=ChartSpec(type="stacked_column", categories=["I кв.", "II кв.", "III кв.", "IV кв."], series=STACK, unit="млн ₽", y_title="Выручка"))),
        OutlineSlide(id="s5", kind="chart", headline="Выручка по регионам", content=SlideContent(chart=ChartSpec(type="grouped_bar", categories=REG, series=[InlineSeries(name="2025", values=[480, 260, 150]), InlineSeries(name="2026", values=[560, 300, 207])], unit="млн ₽", highlight_index=2))),
        OutlineSlide(id="s6", kind="chart", headline="Чем выше чек, тем больше выручка", content=SlideContent(chart=ChartSpec(type="scatter", xy=[XYSeries(name="Клиенты", x=XS, y=YS)], unit="тыс. ₽", x_title="Средний чек", y_title="Выручка на клиента"))),
    ]
    outline = DeckOutline(title="Продажи в 2026 году", strategy="structured", slides=slides)
    res = generate_variants(simple_deck, outline=outline, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False,
                            render_images=False, audit=True, autofix=False, exports=["html"], agent=False, writer=False, pictures=False, eyes=False)
    v = res.variants[0]
    page = (v.out_dir / "deck.html").read_text(encoding="utf-8")
    sections = page.split('<section class="slide"')[1:]
    assert len(sections) == len(slides)
    svgs = [re.search(r"<svg.*?</svg>", s, re.S) for s in sections[1:]]
    assert all(svgs) and "<img" not in page
    hist, line, stacked, grouped, xy = (m.group(0) for m in svgs)
    assert "Минут до ответа" in hist and len(_rects(hist)) == len(BINS)
    assert len(re.findall(r"<polyline", line)) == 2
    assert len(_rects(stacked)) == 12 and "Прочее" in stacked
    assert "2025: 480" in grouped
    assert len(CIRCLE_RE.findall(xy)) == len(XS) and "Средний чек" in xy
    assert page.count("Время ответа поддержки.csv") == 1  # the source line under the chart, the slide's own text
    assert v.audit is not None and not [i for i in v.audit.issues if i.check_id in ("author_item_missing",) or "chart" in i.check_id]
