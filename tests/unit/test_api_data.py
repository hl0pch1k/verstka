"""Вторая ступень, API (описание решения §6): файлы с данными (POST/GET/DELETE /api/datasets), брендбук (задача разбора
с шагами, расхождения с шаблоном, логотипы), примеры «данные продаж» и «брендбук VK Cloud», поля сборки
(`dataset_ids`, `brandbook_id`, `brand_priority`) и брендированные копии шаблонов в библиотеке.

Без ключей и без модели: временная рабочая папка, сеть закрыта (tests/conftest.py)."""

from __future__ import annotations

import gzip
import io
import json
import shutil
import threading
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import verstka.api.app as A
from verstka.api.store import Store
from verstka.schemas.brand import BrandConflict

ROOT = Path(__file__).resolve().parents[2]
SALES = ROOT / "examples" / "data-sales-2026"
BRAND = ROOT / "examples" / "brandbook-vk-cloud"
FIX = ROOT / "tests" / "fixtures" / "tpl_audit"
VK_TID = "cbbe3aa6a21d23ce"  # разбор шаблона VK Tech в tests/fixtures/tpl_audit

REVENUE = (
    "Месяц;Выручка, млн ₽\n"
    + "\n".join(f"2025-{m:02d};{100 + m * 7}" for m in range(1, 13))
    + "\n"
).encode("utf-8")


@pytest.fixture
def api(tmp_path, monkeypatch):
    store = Store(tmp_path)
    monkeypatch.setattr(A, "store", store)
    for var in ("VERSTKA_MAX_DATA_MB", "VERSTKA_MAX_BRAND_MB", "VERSTKA_MAX_XLSX_MB", "VERSTKA_MAX_TEXT_MB", "VERSTKA_MAX_UPLOAD_MB"):
        monkeypatch.delenv(var, raising=False)
    return TestClient(A.app), store, tmp_path


def _wait(jid: str, timeout: float = 60) -> A.Job:
    t0 = time.time()
    while time.time() - t0 < timeout:
        job = A.runner.get(jid)
        # the runner sets the status before its last event: wait for that event (its message is the job's last word)
        if job is not None and job.events and job.events[-1].get("status") in ("done", "failed"):
            return job
        time.sleep(0.05)
    raise TimeoutError(jid)


def _dirs(root: Path, name: str) -> list[Path]:
    return sorted(p for p in (root / name).iterdir()) if (root / name).is_dir() else []


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================================================ файлы с данными


def test_a_csv_becomes_a_card_and_is_kept(api):
    c, store, root = api
    r = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))])
    assert r.status_code == 200, r.text
    (card,) = r.json()
    assert set(card) == {"id", "file", "kind", "rows", "cols", "nature", "nature_label", "line", "issues", "sheet"}
    assert card["file"] == "Выручка.csv" and card["kind"] == "csv" and card["rows"] == 12 and card["cols"] == 2
    assert card["nature"] == "time_series" and card["nature_label"] == "Временной ряд" and card["sheet"] is None
    assert card["line"].startswith("Временной ряд · 12 строк")
    d = root / "datasets" / card["id"]
    assert sorted(p.name for p in d.iterdir()) == ["plan.json", "profile.json", "source.csv", "summary.json", "table.json"]
    assert (d / "source.csv").read_bytes() == REVENUE
    assert json.loads((d / "summary.json").read_text(encoding="utf-8")) == card
    # the table and its profile are known by the dataset's id (SlideData.dataset_id, the card and the request agree)
    assert json.loads((d / "table.json").read_text(encoding="utf-8"))["id"] == card["id"]
    assert json.loads((d / "profile.json").read_text(encoding="utf-8"))["table_id"] == card["id"]


def test_a_book_gives_a_card_per_sheet(api):
    c, _, root = api
    book = _xlsx({
        "Регионы": [["Регион", "Выручка"], ["Москва", 410], ["Казань", 120], ["Самара", 90]],
        "Каналы": [["Канал", "Доля, %"], ["Прямые", 46], ["Партнёры", 28], ["Маркетплейс", 17], ["Онлайн", 9]],
    })
    r = c.post("/api/datasets", files=[("files", ("Продажи.xlsx", book, "application/octet-stream"))])
    assert r.status_code == 200, r.text
    cards = r.json()
    assert [x["sheet"] for x in cards] == ["Регионы", "Каналы"]
    assert {x["kind"] for x in cards} == {"xlsx"} and len({x["id"] for x in cards}) == 2
    for x in cards:
        assert (root / "datasets" / x["id"] / "source.xlsx").read_bytes() == book


def test_the_seven_sample_files_read_in_one_request(api):
    c, _, _ = api
    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in sorted(SALES.iterdir()) if p.name != "SCENARIO.md"]
    r = c.post("/api/datasets", files=files)
    assert r.status_code == 200, r.text
    cards = {x["file"]: x for x in r.json()}
    assert len(cards) == 7
    assert cards["Выручка по месяцам.csv"]["nature"] == "time_series"
    assert any("пропуск" in i for i in cards["NPS по кварталам.csv"]["issues"])
    assert any("выброс" in i for i in cards["Время ответа поддержки.csv"]["issues"])


def test_get_says_the_card_the_profile_and_a_preliminary_choice(api):
    c, _, _ = api
    did = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))]).json()[0]["id"]
    got = c.get(f"/api/datasets/{did}").json()
    assert got["id"] == did and got["file"] == "Выручка.csv" and got["nature"] == "time_series"
    assert got["profile"]["table_id"] == did and got["profile"]["time_axis"]["points"] == 12
    assert got["plan"]["task"]["task"] == "dynamics" and got["plan"]["options"][0]["kind"] == "line"


def test_get_stands_without_the_preliminary_choice(api, monkeypatch):
    import verstka.data.plan as PL

    monkeypatch.setattr(PL, "plan_dataset", lambda p, t: 1 / 0)
    c, _, _ = api
    did = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))]).json()[0]["id"]
    got = c.get(f"/api/datasets/{did}").json()
    assert got["id"] == did and got["profile"]["table_id"] == did and got["plan"] is None


def test_delete_removes_the_files(api):
    c, _, root = api
    did = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))]).json()[0]["id"]
    assert c.delete(f"/api/datasets/{did}").json() == {"deleted": True}
    assert not (root / "datasets" / did).exists()
    assert c.get(f"/api/datasets/{did}").status_code == 404
    assert c.delete(f"/api/datasets/{did}").json() == {"deleted": False}


@pytest.mark.parametrize("bad", ["..", "..%2F..%2Fruns", "a.b", "-x", "x" * 65, "%2Fetc%2Fpasswd"])
def test_dataset_ids_are_validated(api, bad):
    c, _, _ = api
    assert c.get(f"/api/datasets/{bad}").status_code in (404, 405)
    assert c.delete(f"/api/datasets/{bad}").status_code in (404, 405)


@pytest.mark.parametrize("name,data,words", [
    ("Пусто.csv", b"", "Файл пустой"),
    ("Отчёт.pdf", b"%PDF-1.4 ...", "CSV, XLSX, JSON"),
    ("Скан.png", b"\x89PNG\r\n", "CSV, XLSX, JSON"),
    ("Выгрузка.dat", bytes(range(256)) * 8, "двоичные данные"),
])
def test_unreadable_files_are_said_in_words(api, name, data, words):
    c, _, root = api
    r = c.post("/api/datasets", files=[("files", (name, data, "application/octet-stream"))])
    assert r.status_code == 400 and words in r.json()["detail"]
    assert _dirs(root, "datasets") == []


def test_a_reader_slip_is_a_human_400(api, monkeypatch):
    import verstka.data.readers as R

    monkeypatch.setattr(R, "read_bytes", lambda name, data: (_ for _ in ()).throw(KeyError("xl/workbook.xml")))
    c, _, _ = api
    r = c.post("/api/datasets", files=[("files", ("Книга.xlsx", b"PK\x03\x04broken", "application/octet-stream"))])
    assert r.status_code == 400 and "не повреждён" in r.json()["detail"]


def test_a_file_that_reads_but_is_not_for_a_chart_is_a_card(api):
    c, _, _ = api
    text = "Имя;Комментарий\nИван;Хороший клиент\nПётр;Плохой клиент\nАнна;Новый клиент\n".encode()
    r = c.post("/api/datasets", files=[("files", ("Отзывы.csv", text, "text/csv"))])
    assert r.status_code == 200, r.text
    (card,) = r.json()
    assert card["nature"] == "untypable" and card["issues"][0].startswith("Не для графика: ")


def test_a_profile_slip_still_gives_a_card(api, monkeypatch):
    import verstka.data.profile as P

    monkeypatch.setattr(P, "profile_table", lambda t: 1 / 0)
    c, _, _ = api
    r = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))])
    assert r.status_code == 200, r.text
    (card,) = r.json()
    assert card["nature"] == "untypable" and card["rows"] == 12 and card["issues"] == ["Не для графика: не удалось разобрать колонки"]


def test_size_and_count_limits(api, monkeypatch):
    c, _, root = api
    monkeypatch.setenv("VERSTKA_MAX_DATA_MB", "0.0001")  # ~100 байт
    r = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))])
    assert r.status_code == 413 and "МБ" in r.json()["detail"] and "«Выручка.csv»" in r.json()["detail"]
    monkeypatch.delenv("VERSTKA_MAX_DATA_MB")
    many = [("files", (f"f{i}.csv", REVENUE, "text/csv")) for i in range(A._MAX_DATA_FILES + 1)]
    r = c.post("/api/datasets", files=many)
    assert r.status_code == 413 and f"до {A._MAX_DATA_FILES} файлов" in r.json()["detail"]
    monkeypatch.setattr(A, "_MAX_DATA_TABLES", 1)
    book = _xlsx({"a": [["x", "y"], ["a", 1], ["b", 2]], "b": [["x", "y"], ["c", 3], ["d", 4]]})
    r = c.post("/api/datasets", files=[("files", ("Книга.xlsx", book, "application/octet-stream"))])
    assert r.status_code == 413 and "таблиц" in r.json()["detail"]
    assert _dirs(root, "datasets") == []


def test_a_bad_file_among_several_leaves_nothing_behind(api):
    c, _, root = api
    r = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv")), ("files", ("Пусто.csv", b"  ", "text/csv"))])
    assert r.status_code == 400 and r.json()["detail"] == "«Пусто.csv»: файл пустой"
    assert _dirs(root, "datasets") == []


def test_no_file_is_said_in_words(api):
    c, _, _ = api
    r = c.post("/api/datasets", data={"x": "1"})
    assert r.status_code == 400 and "Выберите файл" in r.json()["detail"]


@pytest.mark.parametrize("raw,shown", [
    ("../../evil.csv", "evil.csv"),
    ("..\\..\\win.csv", "win.csv"),
    ("/etc/passwd.csv", "passwd.csv"),
    ("Отчёт: продажи?.csv", "Отчёт: продажи?.csv"),  # a data file's name is never a path: kept as people wrote it
    ("a<b>|*.csv", "a<b>|*.csv"),
    ("..", "данные"),
])
def test_file_names_are_safe(api, raw, shown):
    c, _, root = api
    r = c.post("/api/datasets", files=[("files", (raw, REVENUE, "text/csv"))])
    assert r.status_code == 200, r.text
    (card,) = r.json()
    assert card["file"] == shown
    d = root / "datasets" / card["id"]
    assert [p.name for p in d.iterdir() if p.name.startswith("source")] == ["source.csv"] or shown == "данные"
    assert c.get(f"/api/datasets/{card['id']}").json()["plan"]["source"].endswith(shown)
    assert not (root.parent / "evil.csv").exists() and not (root / "evil.csv").exists()


def test_safe_file_name_keeps_the_extension_of_a_long_name():
    name = A.safe_file_name("Очень " * 60 + ".xlsx", "x")
    assert len(name) <= 120 and name.endswith(".xlsx")
    assert A.safe_file_name("a\x00b\tc\x7f.csv", "x") == "a b c .csv"
    assert A.safe_file_name(".hidden.csv", "x") == "hidden.csv" and A.safe_file_name("...", "x") == "x"
    assert A.safe_file_name("Отчёт: продажи?.pdf", "x") == "Отчёт продажи .pdf"  # a brandbook's name becomes a path
    assert A.data_file_name("C:\\Users\\Отчёт: продажи?.csv") == "Отчёт: продажи?.csv"
    assert A.data_file_name("a\x00b\tc.csv") == "a b c.csv" and A.data_file_name("..") == "данные"
    assert len(A.data_file_name("Очень " * 60 + ".csv")) <= 120


# ============================================================================================ брендбук


def _upload_brand(c, **kw):
    r = c.post("/api/brandbooks", **kw)
    assert r.status_code == 200, r.text
    body = r.json()
    assert set(body) == {"job_id", "brandbook_id"}
    return body["brandbook_id"], _wait(body["job_id"])


def test_tokens_brandbook_steps_and_result(api):
    c, store, root = api
    bid, job = _upload_brand(c, files={"file": ("Брендбук VK Cloud.json", (BRAND / "tokens.json").read_bytes(), "application/json")})
    assert job.status == "done", job.error
    steps = [e for e in job.events if e.get("step")]
    assert [e["step"] for e in steps] == ["open", "palette", "fonts", "logos", "rules", "compare"]
    assert steps[1]["message"] == "Палитра: 6 цветов" and len(steps[1]["brand"]["colors"]) == 6
    assert [f["family"] for f in steps[2]["brand"]["fonts"]] == ["VK Sans Display", "VK Sans Text"]
    assert steps[-1]["message"] == "Сверка с шаблоном: шаблон не выбран"
    progress = [e["progress"] for e in steps]
    assert progress == sorted(progress) and 0 < progress[0] and progress[-1] < 1
    assert job.result["brand"]["id"] == bid and job.result["brand"]["file"] == "Брендбук VK Cloud.json"
    assert job.result["conflicts"] == []
    got = c.get(f"/api/brandbooks/{bid}").json()
    assert got["brand"] == job.result["brand"] and got["conflicts"] == []
    assert not (root / "brandbooks" / bid / ".parsing").exists()
    assert store.read_brandbook(bid).id == bid


def test_pdf_brandbook_logos_are_served_and_nothing_else(api):
    c, _, root = api
    bid, job = _upload_brand(c, files={"file": ("Брендбук VK Cloud.pdf", (BRAND / "guideline.pdf").read_bytes(), "application/pdf")})
    assert job.status == "done", job.error
    brand = job.result["brand"]
    assert brand["source"] == "pdf" and brand["pages"] == 12 and brand["file"] == "Брендбук VK Cloud.pdf"
    logos = [lg["file"] for lg in brand["logos"]]
    assert logos and all(f.startswith("logos/") for f in logos)
    r = c.get(f"/api/brandbooks/{bid}/files/{logos[0]}")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    for bad in ("brand.json", "source/Брендбук%20VK%20Cloud.pdf", "logos/..%2Fbrand.json", "..%2F..%2Fdatasets", "logos/../brand.json", "%2Fetc%2Fpasswd"):
        assert c.get(f"/api/brandbooks/{bid}/files/{bad}").status_code == 404, bad


def _folder_parts(prefix: str = "brandbook-vk-cloud") -> list:
    parts = [("files", (f"{prefix}/tokens.json", (BRAND / "tokens.json").read_bytes(), "application/json"))]
    for p in sorted((BRAND / "logos").glob("*.png")):
        parts.append(("files", (f"{prefix}/logos/{p.name}", p.read_bytes(), "image/png")))
    parts.append(("files", (f"{prefix}/.DS_Store", b"\x00\x01", "application/octet-stream")))
    return parts


def test_a_folder_brandbook(api):
    c, _, root = api
    bid, job = _upload_brand(c, files=_folder_parts())
    assert job.status == "done", job.error
    brand = job.result["brand"]
    assert brand["source"] == "folder" and brand["file"] == "brandbook-vk-cloud"
    assert {lg["variant"] for lg in brand["logos"]} == {"primary", "white", "mark", "mono"}
    assert not list((root / "brandbooks" / bid).rglob(".DS_Store"))
    assert c.get(f"/api/brandbooks/{bid}/files/{brand['logos'][0]['file']}").status_code == 200


def test_folder_paths_cannot_leave_the_brandbook(api):
    c, _, root = api
    parts = _folder_parts() + [("files", ("brandbook-vk-cloud/../../../evil.json", b'{"a": 1}', "application/json"))]
    r = c.post("/api/brandbooks", files=parts)
    assert r.status_code == 400 and "путём" in r.json()["detail"]
    assert _dirs(root, "brandbooks") == []
    assert not list(root.parent.rglob("evil.json"))


def test_folder_path_rules():
    assert str(A.folder_path("Бренд/logos/main.svg")) == "Бренд/logos/main.svg"
    assert str(A.folder_path("/abs//x\\y.png")) == "abs/x/y.png"
    assert A.folder_path("Бренд/.git/config") is None and A.folder_path("__MACOSX/a.png") is None
    assert A.folder_path("/".join(["d"] * 9) + "/a.png") is None
    with pytest.raises(A.HTTPException):
        A.folder_path("a/../b.png")


def _zip(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


def test_a_zip_brandbook_and_its_limits(api, monkeypatch):
    c, _, root = api
    content = {"VK Cloud/tokens.json": (BRAND / "tokens.json").read_bytes(), "VK Cloud/logos/primary.png": (BRAND / "logos" / "primary.png").read_bytes()}
    bid, job = _upload_brand(c, files={"file": ("Брендбук.zip", _zip(content), "application/zip")})
    assert job.status == "done", job.error
    assert job.result["brand"]["logos"] and job.result["brand"]["colors"]
    monkeypatch.setattr(A, "_MAX_ZIP_ENTRIES", 1)
    r = c.post("/api/brandbooks", files={"file": ("Брендбук.zip", _zip(content), "application/zip")})
    assert r.status_code == 413 and "больше 1 файлов" in r.json()["detail"]
    monkeypatch.setattr(A, "_MAX_ZIP_ENTRIES", 1000)
    monkeypatch.setattr(A, "_MAX_ZIP_UNPACKED", 1000)
    bomb = _zip({"big.json": b" " * 100_000})
    r = c.post("/api/brandbooks", files={"file": ("Бомба.zip", bomb, "application/zip")})
    assert r.status_code == 413 and "распакованном" in r.json()["detail"]
    r = c.post("/api/brandbooks", files={"file": ("Битый.zip", b"PK\x03\x04 not a zip", "application/zip")})
    assert r.status_code == 400 and "повреждён" in r.json()["detail"]
    assert [p.name for p in _dirs(root, "brandbooks")] == [bid]


def test_brandbook_size_limit(api, monkeypatch):
    c, _, root = api
    monkeypatch.setenv("VERSTKA_MAX_BRAND_MB", "0.001")
    r = c.post("/api/brandbooks", files={"file": ("Брендбук.pdf", (BRAND / "guideline.pdf").read_bytes(), "application/pdf")})
    assert r.status_code == 413 and "МБ" in r.json()["detail"]
    r = c.post("/api/brandbooks", files=_folder_parts())
    assert r.status_code == 413 and "Папка брендбука больше" in r.json()["detail"]
    monkeypatch.delenv("VERSTKA_MAX_BRAND_MB")
    monkeypatch.setattr(A, "_MAX_BRAND_FOLDER_FILES", 2)
    r = c.post("/api/brandbooks", files=_folder_parts())
    assert r.status_code == 413 and "больше 2 файлов" in r.json()["detail"]
    assert _dirs(root, "brandbooks") == []


@pytest.mark.parametrize("kw,status,words", [
    ({"data": {"template_id": "x"}}, 400, "Выберите брендбук"),
    ({"files": {"file": ("Брендбук.docx", b"x", "application/octet-stream")}}, 400, "не подходит"),
    ({"files": {"file": ("Брендбук.fig", b"x", "application/octet-stream")}}, 400, "Figma"),
    ({"files": {"file": ("Брендбук.pdf", b"", "application/pdf")}}, 400, "пустой"),
    ({"data": {"figma_url": "https://example.com/brand"}}, 400, "не похожа на файл Figma"),
    ({"files": {"file": ("a.json", b"{}", "application/json")}, "data": {"figma_url": "https://www.figma.com/design/AbC123/x"}}, 400, "что-то одно"),
])
def test_brandbook_requests_said_in_words(api, kw, status, words):
    c, _, root = api
    r = c.post("/api/brandbooks", **kw)
    assert r.status_code == status and words in r.json()["detail"], r.text
    assert _dirs(root, "brandbooks") == []


def test_a_brandbook_without_a_brand_fails_its_job_and_leaves_nothing(api):
    c, _, root = api
    bid, job = _upload_brand(c, files={"file": ("Пусто.json", b'{"a": 1}', "application/json")})
    assert job.status == "failed" and job.message.startswith("Не получилось: ")
    assert not (root / "brandbooks" / bid).exists()
    assert c.get(f"/api/brandbooks/{bid}").status_code == 404


def test_a_reader_slip_fails_the_job_in_words(api, monkeypatch):
    import verstka.brand.read as BR

    monkeypatch.setattr(BR, "read_brandbook", lambda *a, **k: 1 / 0)
    c, _, _ = api
    _, job = _upload_brand(c, files={"file": ("Брендбук.pdf", b"%PDF-1.4", "application/pdf")})
    assert job.status == "failed" and "проверьте файл" in job.message


def test_no_vision_model_without_a_key(api, monkeypatch):
    import verstka.brand.read as BR

    seen = {}
    real = BR.read_brandbook

    def spy(source, **kw):
        seen.update(kw)
        return real(source, **{**kw, "use_vlm": False, "providers": None})

    monkeypatch.setattr(BR, "read_brandbook", spy)
    c, _, _ = api
    _, job = _upload_brand(c, files={"file": ("Брендбук.pdf", (BRAND / "guideline.pdf").read_bytes(), "application/pdf")})
    assert job.status == "done", job.error
    assert seen["use_vlm"] is False and seen["providers"] is None and seen["skills"] is None


def test_a_figma_link_is_read_by_the_reader(api, monkeypatch):
    import verstka.brand.read as BR

    monkeypatch.setattr(BR, "read_brandbook", lambda src, **k: _figma_spec(src))
    c, _, _ = api
    bid, job = _upload_brand(c, data={"figma_url": "https://www.figma.com/design/AbC123XyZ/VK-Cloud?node-id=1-2"})
    assert job.status == "done", job.error
    assert job.result["brand"]["source"] == "figma" and job.result["brand"]["id"] == bid


def _figma_spec(src):
    from verstka.schemas.brand import BrandColor, BrandSpec

    assert isinstance(src, str) and "figma.com" in src
    return BrandSpec(name="VK Cloud", source="figma", file=src, colors=[BrandColor(hex="#0077FF", role="primary")])


def test_parsing_then_deleted(api, monkeypatch):
    """GET while the brandbook is read → 409; DELETE during the reading → the folder goes and the job fails."""
    import verstka.brand.read as BR

    gate = threading.Event()
    real = BR.read_brandbook

    def slow(source, **kw):
        gate.wait(10)
        return real(source, **kw)

    monkeypatch.setattr(BR, "read_brandbook", slow)
    c, _, root = api
    r = c.post("/api/brandbooks", files={"file": ("tokens.json", (BRAND / "tokens.json").read_bytes(), "application/json")})
    bid, jid = r.json()["brandbook_id"], r.json()["job_id"]
    g = c.get(f"/api/brandbooks/{bid}")
    assert g.status_code == 409 and "разбирается" in g.json()["detail"]
    assert c.delete(f"/api/brandbooks/{bid}").json() == {"deleted": True}
    gate.set()
    job = _wait(jid)
    assert job.status == "failed" and "убрали" in job.message
    assert not (root / "brandbooks" / bid).exists()


def test_delete_brandbook(api):
    c, _, root = api
    bid, job = _upload_brand(c, files={"file": ("tokens.json", (BRAND / "tokens.json").read_bytes(), "application/json")})
    assert c.delete(f"/api/brandbooks/{bid}").json() == {"deleted": True}
    assert not (root / "brandbooks" / bid).exists()
    assert c.get(f"/api/brandbooks/{bid}").status_code == 404
    assert c.delete(f"/api/brandbooks/{bid}").json() == {"deleted": False}
    for bad in ("..", "a.b", "..%2Fdatasets"):
        assert c.get(f"/api/brandbooks/{bad}").status_code in (404, 405)
        assert c.delete(f"/api/brandbooks/{bad}").status_code in (404, 405)


def test_conflicts_with_the_chosen_template(api, monkeypatch):
    c, store, _ = api
    sentinel = object()
    monkeypatch.setattr(store, "manifest", lambda tid: sentinel if tid == "tpl0000000000001" else None)
    calls = []

    def fake(manifest, spec, brand_dir, priority="brand"):
        calls.append((manifest, spec.id, brand_dir))
        return [BrandConflict(what="Шрифт заголовков", template="Play", brand="VK Sans Display", result="VK Sans Display, замена — Play", kind="font")]

    monkeypatch.setattr(A, "_brand_conflicts", fake)
    bid, job = _upload_brand(c, files={"file": ("tokens.json", (BRAND / "tokens.json").read_bytes(), "application/json")}, data={"template_id": "tpl0000000000001"})
    assert job.status == "done", job.error
    assert [x["what"] for x in job.result["conflicts"]] == ["Шрифт заголовков"]
    compare = [e for e in job.events if e.get("step") == "compare"][0]
    assert compare["message"] == "Сверка с шаблоном: 1 расхождение"
    assert calls[0][0] is sentinel and calls[0][1] == bid and calls[0][2] == store.brandbook_dir(bid)
    assert c.get(f"/api/brandbooks/{bid}?template_id=tpl0000000000001").json()["conflicts"][0]["result"] == "VK Sans Display, замена — Play"
    assert c.get(f"/api/brandbooks/{bid}").json()["conflicts"] == []
    assert c.get(f"/api/brandbooks/{bid}?template_id=ffffffffffffffff").json()["conflicts"] == []
    # a slip of the comparison: the brandbook stands, no conflicts
    monkeypatch.setattr(A, "_brand_conflicts", lambda *a: 1 / 0)
    assert c.get(f"/api/brandbooks/{bid}?template_id=tpl0000000000001").json()["conflicts"] == []


def test_real_conflicts_vk_tech_and_the_sample_brandbook(api, monkeypatch):
    """SCENARIO.md: шрифт заголовков Play → VK Sans Display, второй акцент #FF3885 → #FF3985 (verstka.brand.conflicts)."""
    import verstka.brand.apply as ap

    monkeypatch.setattr(ap, "font_available", lambda f: f == "Play")
    c, _, root = api
    d = root / "templates" / VK_TID
    d.mkdir(parents=True)
    (d / "manifest.json").write_bytes(gzip.open(FIX / f"manifest_{VK_TID}.json.gz").read())
    shutil.copyfile(FIX / "vktech_short.pptx", d / "source.pptx")
    bid, job = _upload_brand(c, files={"file": ("Брендбук VK Cloud.pdf", (BRAND / "guideline.pdf").read_bytes(), "application/pdf")}, data={"template_id": VK_TID})
    assert job.status == "done", job.error
    got = {x["what"]: x for x in job.result["conflicts"]}
    assert got["Шрифт заголовков"]["brand"] == "VK Sans Display"
    assert any(x["template"] == "#FF3885" and x["brand"] == "#FF3985" for x in got.values())
    again = c.get(f"/api/brandbooks/{bid}?template_id={VK_TID}").json()
    assert again["conflicts"] == job.result["conflicts"]


# ============================================================================================ примеры


def test_the_sales_example(api):
    c, _, _ = api
    ex = c.get("/api/examples/data-sales-2026").json()
    assert ex["name"] == "data-sales-2026" and ex["slides"] == 9
    assert ex["brief"].startswith("Итоги продаж VK Cloud за 9 месяцев 2026 года для правления.") and "«" not in ex["brief"]
    assert [f["name"] for f in ex["files"]] == [
        "Выручка по месяцам.csv", "Регионы.xlsx", "Каналы продаж.json", "Время ответа поддержки.csv", "Клиенты.xlsx", "NPS по кварталам.csv", "Итоги квартала.txt",
    ]
    for f in ex["files"]:
        r = c.get(f["url"])
        assert r.status_code == 200 and r.content == (SALES / f["name"]).read_bytes()
    assert ex["brandbook"]["name"] == "Брендбук VK Cloud.pdf"
    assert c.get(ex["brandbook"]["url"]).content == (BRAND / "guideline.pdf").read_bytes()
    for bad in ("SCENARIO.md", "..%2FSCENARIO.md", "..%2F..%2Fpyproject.toml", "%2Fetc%2Fpasswd"):
        assert c.get(f"/api/examples/data-sales-2026/files/{bad}").status_code == 404, bad


def test_the_brandbook_example_and_unknown_ones(api):
    c, _, _ = api
    ex = c.get("/api/examples/brandbook-vk-cloud").json()
    assert ex["files"] == [] and ex["brandbook"]["name"] == "Брендбук VK Cloud.pdf"
    assert c.get(ex["brandbook"]["url"]).status_code == 200
    assert c.get("/api/examples/brandbook-vk-cloud/files/tokens.json").status_code == 404
    assert c.get("/api/examples/nope").status_code == 404
    assert c.get("/api/examples/a.b").status_code == 404


# ============================================================================================ сборка


@pytest.fixture
def gen(api, monkeypatch):
    c, store, root = api
    monkeypatch.setattr(store, "manifest", lambda tid: object() if tid == "0123456789abcdef" else None)
    submitted = []

    def submit(kind, fn):
        submitted.append(fn)
        return A.Job(id=f"job{len(submitted):05d}", kind=kind)

    monkeypatch.setattr(A.runner, "submit", submit)
    return c, store, root, submitted


def test_the_build_validates_and_keeps_the_data_fields(gen, monkeypatch):
    c, store, root, submitted = gen
    did = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))]).json()[0]["id"]
    monkeypatch.setattr(A.runner, "get", lambda jid: None)
    bid = "bb0000000001"
    (root / "brandbooks" / bid).mkdir()
    from verstka.schemas.brand import BrandColor, BrandSpec

    store.write_brandbook(bid, BrandSpec(name="VK Cloud", source="tokens", colors=[BrandColor(hex="#0077FF", role="primary")]))
    body = {"template_id": "0123456789abcdef", "brief": "Итоги продаж", "dataset_ids": [did, did], "brandbook_id": bid, "brand_priority": "template"}
    r = c.post("/api/generations", json=body)
    assert r.status_code == 200, r.text
    meta = store.read_generation_meta(r.json()["generation_id"])
    assert meta["dataset_ids"] == [did] and meta["brandbook_id"] == bid and meta["brand_priority"] == "template"
    listed = {g["id"]: g for g in c.get("/api/generations").json()}[r.json()["generation_id"]]
    assert listed["dataset_ids"] == [did] and listed["brandbook_id"] == bid and listed["brand_priority"] == "template"
    # the generation receives them read (Task 10 builds the deck of them)
    got = {}
    monkeypatch.setattr(A, "_run_generation", lambda gid, gdir, req, job, data=None: got.update(req=req, data=data) or {})
    submitted[-1](A.Job(id="j1", kind="generate"))
    data = got["data"]
    assert data.dataset_ids == [did] and data.tables[0].id == did and data.profiles[0].nature == "time_series"
    # (the brandbook's logos were copied into the build's own folder before the job started: Task 10, review minor 10)
    gdir = store.generation_dir(r.json()["generation_id"])
    assert data.brand.id == bid and data.brand_dir == gdir / "brand" and data.brand_priority == "template"
    assert data.describe() == {"datasets": [{"id": did, "file": "Выручка.csv", "sheet": None, "kind": "csv"}],
                               "brandbook": {"id": bid, "name": "VK Cloud", "file": None, "source": "tokens"}}


@pytest.mark.parametrize("patch,status,words", [
    ({"dataset_ids": ["000000000000"]}, 404, "файлов с данными больше нет"),
    ({"dataset_ids": ["../../runs"]}, 404, "файлов с данными больше нет"),
    ({"brandbook_id": "000000000000"}, 404, "Брендбука больше нет"),
    ({"brandbook_id": "../templates"}, 404, "Брендбука больше нет"),
    ({"brand_priority": "both"}, 422, None),
])
def test_the_build_refuses_missing_inputs(gen, patch, status, words):
    c, store, root, submitted = gen
    r = c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "Итоги", **patch})
    assert r.status_code == status, r.text
    if words:
        assert words in r.json()["detail"]
    assert submitted == [] and _dirs(root, "runs") == []


def test_the_build_waits_for_a_brandbook_being_read(gen, monkeypatch):
    c, store, root, submitted = gen
    bid = "bb0000000002"
    (root / "brandbooks" / bid).mkdir()
    (root / "brandbooks" / bid / ".parsing").write_text("live00000001", encoding="utf-8")
    monkeypatch.setattr(A.runner, "get", lambda jid: A.Job(id=jid, kind="brandbook", status="running") if jid == "live00000001" else None)
    r = c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "Итоги", "brandbook_id": bid})
    assert r.status_code == 409 and "разбирается" in r.json()["detail"]
    # the job is gone (the server restarted): the brandbook is not there
    monkeypatch.setattr(A.runner, "get", lambda jid: None)
    r = c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "Итоги", "brandbook_id": bid})
    assert r.status_code == 404


def test_too_many_datasets(gen, monkeypatch):
    c, _, _, _ = gen
    monkeypatch.setattr(A, "_MAX_GENERATION_DATASETS", 2)
    r = c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "Итоги", "dataset_ids": ["a00000000001", "a00000000002", "a00000000003"]})
    assert r.status_code == 422 and "до 2" in r.json()["detail"]


def test_a_build_without_data_is_as_before(gen, monkeypatch):
    c, store, _, submitted = gen
    r = c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "Итоги"})
    assert r.status_code == 200
    meta = store.read_generation_meta(r.json()["generation_id"])
    assert meta["dataset_ids"] == [] and meta["brandbook_id"] is None and meta["brand_priority"] == "brand"
    got = {}
    monkeypatch.setattr(A, "_run_generation", lambda gid, gdir, req, job, data=None: got.update(data=data) or {})
    submitted[-1](A.Job(id="j2", kind="generate"))
    assert got["data"].empty and got["data"].describe() == {}


# ============================================================================================ брендированные копии шаблонов


def _template(root: Path, tid: str, base: str | None = None) -> Path:
    d = root / "templates" / tid
    d.mkdir(parents=True)
    (d / "source.pptx").write_bytes(b"x")
    (d / "manifest.json").write_text(json.dumps({"template_id": tid, "source_file": "VK Tech.pptx", "patterns": []}), encoding="utf-8")
    if base:
        (d / "brand.json").write_text(json.dumps({"base_template_id": base, "brand_id": "bb0000000001", "priority": "brand"}), encoding="utf-8")
    return d


def test_branded_copies_are_not_in_the_library(api):
    c, store, root = api
    _template(root, "aaaaaaaaaaaaaaaa")
    _template(root, "bbbbbbbbbbbbbbbb", base="aaaaaaaaaaaaaaaa")
    assert [t["template_id"] for t in c.get("/api/templates").json()] == ["aaaaaaaaaaaaaaaa"]
    assert store.branded_copies("aaaaaaaaaaaaaaaa") == ["bbbbbbbbbbbbbbbb"]
    assert store.branded_copies("../x") == []


def test_deleting_a_template_takes_its_branded_copies(api):
    c, store, root = api
    _template(root, "aaaaaaaaaaaaaaaa")
    _template(root, "bbbbbbbbbbbbbbbb", base="aaaaaaaaaaaaaaaa")
    _template(root, "cccccccccccccccc", base="aaaaaaaaaaaaaaaa")
    _template(root, "dddddddddddddddd")
    _template(root, "eeeeeeeeeeeeeeee", base="dddddddddddddddd")
    assert c.delete("/api/templates/aaaaaaaaaaaaaaaa").json() == {"deleted": True}
    assert sorted(p.name for p in (root / "templates").iterdir()) == ["dddddddddddddddd", "eeeeeeeeeeeeeeee"]


def test_a_deck_built_on_a_branded_copy_holds_its_template(api, monkeypatch):
    c, store, root = api
    _template(root, "aaaaaaaaaaaaaaaa")
    _template(root, "bbbbbbbbbbbbbbbb", base="aaaaaaaaaaaaaaaa")
    gid, _ = store.new_generation_dir()
    job = A.Job(id="live0002", kind="generate", status="running")
    monkeypatch.setitem(A.runner.jobs, job.id, job)
    store.write_generation_meta(gid, {"id": gid, "template_id": "bbbbbbbbbbbbbbbb", "status": "running", "job_id": job.id})
    r = c.delete("/api/templates/aaaaaaaaaaaaaaaa")
    assert r.status_code == 409 and "собирается" in r.json()["detail"]
    assert (root / "templates" / "bbbbbbbbbbbbbbbb").exists()


# ============================================================================================ ошибки по-русски, размеры, кэш


def test_folder_file_and_folder_with_one_name_is_a_calm_400(api):
    c, _, root = api
    png = (BRAND / "logos" / "primary.png").read_bytes()
    cases = [
        [("files", ("Бренд/a", b"x", "text/plain")), ("files", ("Бренд/a/b.png", png, "image/png"))],  # a file, then a folder
        [("files", ("Бренд/a/b.png", png, "image/png")), ("files", ("Бренд/a", b"x", "text/plain"))],  # a folder, then a file
        [("files", ("Бренд/a?/b.png", png, "image/png")), ("files", ("Бренд/a*", b"x", "text/plain"))],  # names that clean alike
    ]
    for parts in cases:
        r = c.post("/api/brandbooks", files=parts)
        assert r.status_code == 400 and "называются одинаково" in r.json()["detail"], r.text
    assert _dirs(root, "brandbooks") == []


@pytest.mark.parametrize("call,detail", [
    (lambda c: c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "x", "brand_priority": "both"}), "brand_priority: нужно «brand» или «template»"),
    (lambda c: c.post("/api/generations", json={"template_id": "0123456789abcdef", "brief": "x", "slides": "z"}), "slides: нужно целое число"),
    (lambda c: c.post("/api/generations", json={"brief": "x"}), "template_id: обязательное поле"),
    (lambda c: c.post("/api/generations", content=b"{bad", headers={"content-type": "application/json"}), "запрос — не JSON"),
    (lambda c: c.post("/api/brandbooks", data={"use_models": "abc", "figma_url": "x"}), "use_models: нужно «да» или «нет» (true / false)"),
    (lambda c: c.post("/api/datasets", data={"files": "x"}), "files.0: нужен файл"),
])
def test_validation_errors_are_said_in_russian(api, call, detail):
    c, _, _ = api
    r = call(c)
    assert r.status_code == 422 and r.json() == {"detail": detail}


def test_an_outline_that_does_not_read_is_said_in_russian():
    from pydantic import ValidationError

    from verstka.schemas.outline import DeckOutline

    with pytest.raises(ValidationError) as e:
        DeckOutline.model_validate({"title": "x", "slides": [{"kind": "nope"}]})
    msg = A._validation_message(e.value)
    assert "slides.0.kind: нужно «title», «section»" in msg and "slides.0.headline: обязательное поле" in msg and "Input should" not in msg


def test_framework_errors_are_said_in_russian(api):
    c, _, _ = api
    parts = [("files", (f"f{i}.csv", b"a", "text/csv")) for i in range(1001)]
    r = c.post("/api/datasets", files=parts)
    assert r.status_code == 400 and r.json()["detail"].startswith("Слишком много файлов в одном запросе — до 1000")
    assert c.get("/api/nope").json()["detail"] == "Не найдено"
    assert c.put("/api/health").json()["detail"] == "Такой запрос здесь не принимается"
    # the API's own errors are as they were
    assert c.get("/api/datasets/000000000000").json() == {"detail": "Файл с данными не найден — приложите его снова"}


def test_an_oversize_body_is_refused_before_it_is_read(api, monkeypatch):
    """Content-Length over the route's limit → 413 at once: the multipart parser never runs (nothing goes to disk)."""
    from starlette.formparsers import MultiPartParser

    async def never(self):
        raise AssertionError("the body was parsed")

    monkeypatch.setattr(MultiPartParser, "parse", never)
    c, _, root = api
    monkeypatch.setenv("VERSTKA_MAX_DATA_MB", "0.01")  # 10 files × 10 КБ + 1 МБ of headers
    body = b"x" * (A._SLACK + 200_000)
    r = c.post("/api/datasets", files=[("files", ("big.csv", body, "text/csv"))])
    assert r.status_code == 413 and r.json()["detail"].startswith("Файлы с данными больше") and "по одному" in r.json()["detail"]
    monkeypatch.setenv("VERSTKA_MAX_BRAND_MB", "0.01")
    r = c.post("/api/brandbooks", files={"file": ("b.pdf", body, "application/pdf")})
    assert r.status_code == 413 and r.json()["detail"].startswith("Брендбук больше 0,01 МБ")
    monkeypatch.setenv("VERSTKA_MAX_UPLOAD_MB", "0.01")
    r = c.post("/api/templates", files={"file": ("t.pptx", body, "application/octet-stream")})
    assert r.status_code == 413 and "МБ" in r.json()["detail"]
    assert _dirs(root, "datasets") == [] and _dirs(root, "brandbooks") == []
    # a body within the limits goes on: the parser is reached (here it breaks — said in Russian, as the framework's words are)
    monkeypatch.delenv("VERSTKA_MAX_DATA_MB")
    r = c.post("/api/datasets", files=[("files", ("small.csv", b"a;b\n1;2\n", "text/csv"))])
    assert r.status_code == 400 and r.json()["detail"] == "Запрос не прочитан — отправьте его ещё раз"


def test_a_huge_table_keeps_its_first_rows_and_says_so(api, monkeypatch):
    c, _, root = api
    monkeypatch.setattr(A, "_MAX_DATA_ROWS", 5000)
    rows = "Месяц;Выручка\n" + "\n".join(f"2025-{1 + i % 12:02d};{i}" for i in range(12000)) + "\n"
    r = c.post("/api/datasets", files=[("files", ("Много.csv", rows.encode(), "text/csv"))])
    assert r.status_code == 200, r.text
    (card,) = r.json()
    assert card["rows"] == 5000 and card["issues"][0] == "Прочитаны первые 5 000 строк из 12 000"
    assert "5 000 строк" in card["line"]
    table = json.loads((root / "datasets" / card["id"] / "table.json").read_text(encoding="utf-8"))
    assert len(table["rows"]) == 5000
    assert c.get(f"/api/datasets/{card['id']}").json()["issues"][0].startswith("Прочитаны первые")
    # a wide table: the cap is by cells
    monkeypatch.setattr(A, "_MAX_DATA_ROWS", 200_000)
    monkeypatch.setattr(A, "_MAX_DATA_CELLS", 4000)
    wide = ";".join(f"К{j}" for j in range(4)) + "\n" + "\n".join(";".join(str(i * j) for j in range(4)) for i in range(3000)) + "\n"
    (card,) = c.post("/api/datasets", files=[("files", ("Широкая.csv", wide.encode(), "text/csv"))]).json()
    assert card["rows"] == 1000 and card["issues"][0] == "Прочитаны первые 1 000 строк из 3 000"


def test_cap_rows_keeps_totals_and_row_units_in_step():
    t = A.Table(id="t", file="x.csv", kind="csv", columns=["a", "b"], rows=[[str(i), i] for i in range(3000)], totals=[5, 2500])
    capped, cut = A._cap_rows(t)
    assert cut is None and capped is t
    old = A._MAX_DATA_ROWS
    try:
        A._MAX_DATA_ROWS = 1000
        capped, cut = A._cap_rows(t)
    finally:
        A._MAX_DATA_ROWS = old
    assert len(capped.rows) == 1000 and capped.totals == [5] and cut == "Прочитаны первые 1 000 строк из 3 000"


def test_a_book_of_excel_has_its_own_limit(api, monkeypatch):
    c, _, root = api
    book = _xlsx({"Лист": [["Регион", "Выручка"]] + [[f"Регион {i}", i] for i in range(300)]})
    monkeypatch.setenv("VERSTKA_MAX_XLSX_MB", str(len(book) / 2 / 1024 / 1024))
    r = c.post("/api/datasets", files=[("files", ("Регионы.xlsx", book, "application/octet-stream"))])
    assert r.status_code == 413 and r.json()["detail"].startswith("Книга Excel «Регионы.xlsx» больше") and "CSV" in r.json()["detail"]
    # a CSV of the same size is fine, and so is a «.xlsx» that is really text
    text = ("Регион;Выручка\n" + "\n".join(f"Регион {i};{i}" for i in range(2000)) + "\n").encode()
    assert len(text) > len(book) / 2
    assert c.post("/api/datasets", files=[("files", ("Регионы.csv", text, "text/csv"))]).status_code == 200
    assert c.post("/api/datasets", files=[("files", ("Выгрузка.xlsx", text, "application/octet-stream"))]).status_code == 200
    # a text with a table has its own limit too
    monkeypatch.setenv("VERSTKA_MAX_TEXT_MB", str(len(text) / 2 / 1024 / 1024))
    r = c.post("/api/datasets", files=[("files", ("Итоги.txt", text, "text/plain"))])
    assert r.status_code == 413 and r.json()["detail"].startswith("Текст «Итоги.txt» больше") and "CSV" in r.json()["detail"]
    assert c.post("/api/datasets", files=[("files", ("Итоги.csv", text, "text/csv"))]).status_code == 200


def test_the_plan_is_computed_once(api, monkeypatch):
    import verstka.data.plan as PL

    c, store, root = api
    did = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))]).json()[0]["id"]
    first = c.get(f"/api/datasets/{did}").json()
    monkeypatch.setattr(PL, "plan_dataset", lambda *a, **k: (_ for _ in ()).throw(AssertionError("recomputed")))
    monkeypatch.setattr(store, "read_dataset", lambda *a: (_ for _ in ()).throw(AssertionError("the table was read")))
    again = c.get(f"/api/datasets/{did}").json()
    assert again == first and again["plan"]["task"]["task"] == "dynamics"


def test_a_dataset_kept_before_the_cache_gets_it_once(api):
    c, store, root = api
    did = c.post("/api/datasets", files=[("files", ("Выручка.csv", REVENUE, "text/csv"))]).json()[0]["id"]
    d = root / "datasets" / did
    (d / "plan.json").unlink()
    (d / "summary.json").unlink()
    got = c.get(f"/api/datasets/{did}").json()
    assert got["plan"]["options"][0]["kind"] == "line" and got["file"] == "Выручка.csv"
    assert (d / "plan.json").is_file() and (d / "summary.json").is_file()


def test_interrupted_brandbooks_are_swept_at_start(api, monkeypatch):
    import os

    import verstka.providers.registry as REG

    c, store, root = api
    old, fresh, done = (root / "brandbooks" / n for n in ("old000000001", "new000000001", "fin000000001"))
    for d in (old, fresh, done):
        (d / "source").mkdir(parents=True)
        (d / ".parsing").write_text("gone0000", encoding="utf-8")
    (done / "brand.json").write_text("{}", encoding="utf-8")
    past = time.time() - A._BRAND_SWEEP_AGE - 60
    os.utime(old / ".parsing", (past, past))
    os.utime(done / ".parsing", (past, past))
    asked = []
    monkeypatch.setattr(REG, "default_models_path", lambda: asked.append(1) or Path("configs/models.yaml"))
    with TestClient(A.app):  # the lifespan: the models profile is logged, interrupted brandbooks go
        pass
    assert asked and not old.exists() and fresh.exists() and done.exists()


def test_a_brandbook_removed_at_the_last_moment_is_not_done(api, monkeypatch):
    c, store, root = api
    monkeypatch.setattr(store, "write_brandbook", lambda bid, spec: False)
    bid, job = _upload_brand(c, files={"file": ("tokens.json", (BRAND / "tokens.json").read_bytes(), "application/json")})
    assert job.status == "failed" and job.message == f"Не получилось: {A._BRAND_REMOVED}"
    assert not (root / "brandbooks" / bid).exists()


WEB = ROOT / "web"
_FOLD_TS = """
import { readFileSync } from "node:fs";
import { foldBrandEvent, BRAND_RUN0, brandRows, brandTiles } from %s;
const events = JSON.parse(readFileSync(process.argv[2], "utf8"));
let run = BRAND_RUN0;
const trail = [];
for (const ev of events) { run = foldBrandEvent(run, ev); trail.push(run.done); }
console.log(JSON.stringify({ trail, run, rows: brandRows(run, %s, "VK Tech"), tiles: brandTiles(run) }));
"""


def _client_fold(events: list[dict], src: str, tmp_path: Path) -> dict:
    """The job's real events through the web client's own foldBrandEvent (web/src/ui/create/dataRules.ts, bundled by the
    web's esbuild and run by node) — what the analysis card shows step by step."""
    import subprocess

    node = shutil.which("node")
    if not node or not (WEB / "node_modules" / "esbuild").is_dir():
        pytest.skip("node and web/node_modules are needed to run the client's fold")
    entry, bundle, data = tmp_path / "fold.ts", tmp_path / "fold.mjs", tmp_path / "events.json"
    entry.write_text(_FOLD_TS % (json.dumps(str(WEB / "src" / "ui" / "create" / "dataRules.ts")), json.dumps(src)), encoding="utf-8")
    build = f"require('esbuild').buildSync({{entryPoints: [{json.dumps(str(entry))}], outfile: {json.dumps(str(bundle))}, bundle: true, platform: 'node', format: 'esm', logLevel: 'error'}})"
    done = subprocess.run([node, "-e", build], cwd=WEB, capture_output=True, text=True, timeout=120)
    assert done.returncode == 0, done.stderr
    data.write_text(json.dumps(events, ensure_ascii=False), encoding="utf-8")
    done = subprocess.run([node, str(bundle), str(data)], capture_output=True, text=True, timeout=60)
    assert done.returncode == 0, done.stderr
    return json.loads(done.stdout)


def _steps_said(events: list[dict]) -> list[int]:
    """Before each event, how many steps the client may count (the latest step said, by its index, + 1)."""
    order = ["open", "palette", "fonts", "logos", "rules", "compare"]
    out, done = [], 0
    for e in events:
        if e.get("step") in order:
            done = max(done, order.index(e["step"]) + 1)
        out.append(done)
    return out


def _pdf_with_the_model(c, monkeypatch, tmp_path):
    """A PDF read by a (mock) vision model through the API: the steps come as the reader finishes them."""
    from verstka.providers.registry import ProviderRegistry

    at_model: list[list] = []
    reading = {"name": "VK Cloud", "colors": [{"hex": "#2787F5", "role": "chart"}], "rules": [{"text": "Фото — только в полный слайд", "kind": "other"}]}

    def answer(messages):
        at_model.append([(e.get("step"), e["message"]) for e in A.runner.get(jids[-1]).events])
        return reading

    registry = ProviderRegistry.mock({"brand guideline": answer})
    monkeypatch.setattr(A, "providers", lambda: registry)
    jids: list[str] = []
    orig = A.runner.submit
    monkeypatch.setattr(A.runner, "submit", lambda kind, fn: (lambda job: (jids.append(job.id), job)[1])(orig(kind, fn)))
    bid, job = _upload_brand(c, files={"file": ("Брендбук VK Cloud.pdf", (BRAND / "guideline.pdf").read_bytes(), "application/pdf")})
    assert job.status == "done", job.error
    return job, at_model


def test_pdf_steps_come_as_the_reader_reads_and_fold_on_the_client(api, monkeypatch, tmp_path):
    c, _, _ = api
    job, at_model = _pdf_with_the_model(c, monkeypatch, tmp_path)
    events = job.events
    # all five steps reached the job before the model answered; the page progress and the model's line came under «rules»
    assert [k for k, _ in at_model[0] if k] == ["open", "palette", "fonts", "logos", "rules"]
    assert at_model[0][-1][1] == "Модель читает 6 страниц — это до двух минут"
    pages = [e for e in events if e["message"].startswith("Готовлю страницы для модели — ")]
    assert [e["message"].rsplit(" ", 3)[-3:] for e in pages] == [[str(k), "из", "6"] for k in range(1, 7)]
    assert all("step" not in e for e in pages)
    steps = [e["step"] for e in events if e.get("step")]
    assert steps == ["open", "palette", "fonts", "logos", "rules", "palette", "rules", "compare"]  # the model refined two
    first = next(e for e in events if e.get("step") == "open")
    assert first["message"] == "Открыт файл: 12 страниц" and first["brand"]["pages"] == 12 and first["brand"]["id"] == job.result["brand"]["id"]
    # the progress never runs ahead of the steps said: the client counts floor(progress × 6) steps for an event without one
    for e, said in zip(events, _steps_said(events)):
        if not e.get("step") and e["status"] == "running":
            assert int(e["progress"] * 6) <= max(said, 0), e
    progress = [e["progress"] for e in events]
    assert progress == sorted(progress)
    folded = _client_fold(events, "pdf", tmp_path)
    assert folded["trail"] == sorted(folded["trail"]) and folded["run"]["done"] == 6
    assert all(d <= said for d, said in zip(folded["trail"], _steps_said(events)) if said < 6)
    facts = [r["b"] for r in folded["rows"]]  # the step's result after the client's label (the labels are the client's own)
    assert facts[:2] == ["12 страниц", "6 цветов"] and facts[4] == str(len(job.result["brand"]["rules"])) and facts[-1] == "шаблон не выбран"
    assert all(r["state"] == "done" for r in folded["rows"])
    tiles = [t["meta"] for t in folded["tiles"]]
    assert all(t["on"] for t in folded["tiles"]) and tiles[0] == "2 шрифта" and tiles[1] == "6 цветов" and tiles[2].startswith("4")
    roles = {c["hex"]: c["role"] for c in folded["run"]["brand"]["colors"]}
    assert roles["#2787F5"] == "chart"  # the refined palette reached the tiles


@pytest.mark.parametrize("name,src,first", [("tokens.json", "json", "цвета, шрифты, отступы"), ("archive", "dir", "логотипы и примеры слайдов")])
def test_fast_sources_say_their_steps_at_the_end_and_fold_on_the_client(api, tmp_path, name, src, first):
    c, _, _ = api
    if name == "archive":
        content = {"VK Cloud/tokens.json": (BRAND / "tokens.json").read_bytes(), "VK Cloud/logos/primary.png": (BRAND / "logos" / "primary.png").read_bytes()}
        upload = {"file": ("Брендбук.zip", _zip(content), "application/zip")}
    else:
        upload = {"file": (name, (BRAND / name).read_bytes(), "application/json")}
    _, job = _upload_brand(c, files=upload)
    assert job.status == "done", job.error
    steps = [e["step"] for e in job.events if e.get("step")]
    # an archive is said open before its reading («2 файла»), then the reader's own words at its end
    assert steps == (["open"] if name == "archive" else []) + ["open", "palette", "fonts", "logos", "rules", "compare"]
    if name == "archive":
        assert job.events[[e.get("step") for e in job.events].index("open")]["message"] == "Открыт архив: 2 файла"
    folded = _client_fold(job.events, src, tmp_path)
    assert folded["run"]["done"] == 6 and folded["trail"] == sorted(folded["trail"])
    assert folded["rows"][0]["b"] == first  # the reader's result of «open» (the label is the client's own)


def test_steps_a_reader_did_not_say_are_said_after_it(api, monkeypatch):
    """A reader that does not call on_step (an older one, a fake): its five steps are said from the brandbook read."""
    import verstka.brand.read as BR

    monkeypatch.setattr(BR, "read_brandbook", lambda src, **k: BR.read_tokens(src))
    c, _, _ = api
    _, job = _upload_brand(c, files={"file": ("tokens.json", (BRAND / "tokens.json").read_bytes(), "application/json")})
    assert job.status == "done", job.error
    assert [e["step"] for e in job.events if e.get("step")] == ["open", "palette", "fonts", "logos", "rules", "compare"]


def test_brand_step_keys_and_progress_bounds():
    assert [A.brand_step_key(w) for w in ("открыт файл", "прочитаны токены", "открыта папка", "открыт архив", "открыт файл Figma", "палитра", "шрифты", "логотипы", "правила", "страница", "модель")] == \
        ["open"] * 5 + ["palette", "fonts", "logos", "rules", None, None]
    job = A.Job(id="j", kind="brandbook", status="running")
    ev = A.BrandStepEvents(job, "b1")
    ev("страница", "6 из 6")  # before any step: under the first one
    ev("открыта папка", "логотипы и примеры слайдов", {"source": "folder"})
    ev("страница", "3 из 6")
    ev("модель", "читает 6 страниц")
    said = _steps_said(job.events)
    assert all(int(e["progress"] * 6) <= s for e, s in zip(job.events, said) if not e.get("step"))
    assert job.events[1]["brand"] == {"id": "b1", "source": "folder"}


def test_brand_opened_for_each_source(tmp_path):
    z = tmp_path / "b.zip"
    z.write_bytes(_zip({"VK/tokens.json": b"{}", "VK/logos/a.png": b"x", "VK/": b""}))
    assert A._brand_opened(z)[0] == "Открыт архив: 2 файла"
    d = tmp_path / "Бренд"
    (d / "logos").mkdir(parents=True)
    (d / "logos" / "a.png").write_bytes(b"x")
    assert A._brand_opened(d) == ("Открыта папка: 1 файл", {"source": "folder", "file": "Бренд"})
    assert A._brand_opened("https://www.figma.com/design/AbC/x") is None
    assert A._brand_opened(tmp_path / "t.json") is None
    pdf = tmp_path / "g.pdf"  # a PDF's reader says «open» itself, as soon as its text is read
    shutil.copyfile(BRAND / "guideline.pdf", pdf)
    assert A._brand_opened(pdf) is None


def test_the_t2i_hint_names_polza_when_vk_has_no_image_model(monkeypatch):
    import verstka.images.providers as IP

    monkeypatch.setattr(IP, "_CACHE", {})
    monkeypatch.setenv("VERSTKA_MODELS", "configs/models.vk.yaml")
    monkeypatch.delenv("VK_T2I_MODEL", raising=False)
    hint = A._t2i_hint()
    assert "POLZA_API_KEY" in hint and "VK_T2I_MODEL" in hint
    monkeypatch.setattr(IP, "_CACHE", {})
    monkeypatch.setenv("VERSTKA_MODELS", "configs/models.yaml")
    assert A._t2i_hint() == "POLZA_API_KEY in .env"
