from pathlib import Path

from lxml import etree
from pptx import Presentation

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.render import find_pdftoppm, find_soffice
from verstka.rendering.charts import add_chart, number_format
from verstka.rendering.deck import DeckBuilder, element_bbox, slide_shape_elements
from verstka.rendering.fit import fit_size
from verstka.rendering.fonts import measure_text_lines, text_width_pt, wrap_lines
from verstka.rendering.groups import adjust_group, cell_bbox
from verstka.rendering.tables import add_table
from verstka.rendering.textfill import ParagraphSpec, fill_text, shape_text
from verstka.schemas.common import Bbox, PatternKind
from verstka.schemas.outline import ChartSpec, DeckOutline, TableData

import pytest

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json"
needs_office = pytest.mark.skipif(find_soffice() is None or find_pdftoppm() is None, reason="LibreOffice/poppler not installed")


def test_font_metrics_and_wrap():
    assert text_width_pt("Hello", "Play", 18) > text_width_pt("Hi", "Play", 18) > 0
    lines = wrap_lines("Каждый третий дедлайн срывается из-за забытой договорённости в чате", "Play", 18, False, 150)
    assert len(lines) >= 3 and all(text_width_pt(l, "Play", 18) <= 150 + 1e-6 for l in lines)
    assert measure_text_lines("short", "Play", 18, False, 300) == 1


def test_fit_size_steps_down_scale():
    box = Bbox(x=0, y=0, w=914400 * 3, h=int(914400 * 0.5))  # 3in × 0.5in
    long = ["Каждый третий дедлайн срывается из-за забытой договорённости в чате, а менеджеры тратят часы на контроль"]
    r = fit_size(long, box, "Play", 18.0, scale_sizes=[18, 14, 12, 10.5, 9])
    assert r.size_pt < 18.0
    short = fit_size(["Короткий тезис"], box, "Play", 18.0, scale_sizes=[18, 14])
    assert short.fits and short.size_pt == 18.0


def test_deck_builder_clone_and_delete(simple_deck, tmp_path):
    b = DeckBuilder(simple_deck)
    s2 = b.clone_slide(2)
    s3 = b.clone_slide(3)
    b.clone_slide(2)
    b.delete_original_slides()
    out = b.save(tmp_path / "out.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    # picture relationship resolved on the cloned slide 3
    pic = next(sh for sh in prs.slides[1].shapes if sh.shape_type is not None and sh.shape_type == 13)
    assert pic.image.content_type == "image/png"
    # every r:embed on cloned slides points to an existing relationship
    for slide in prs.slides:
        rids = set(slide.part.rels.keys())
        for el in slide._element.iter():
            v = el.get(q("r:embed"))
            if v is not None:
                assert v in rids
    assert len([sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.startswith("Pillar")]) == 3


def test_fill_text_keeps_style_and_bullets(simple_deck):
    b = DeckBuilder(simple_deck)
    slide = b.clone_slide(2)
    els = slide_shape_elements(slide)
    body = next(e for e in els.values() if shape_text(e).startswith("Description"))
    rPr_before = body.find(".//" + q("a:rPr"))
    assert rPr_before.get("sz") == "1400"
    ok = fill_text(body, [ParagraphSpec("Первый", bullet=True), ParagraphSpec("Второй", bullet=True), ParagraphSpec("Третий пункт", bullet=True, bold=True)])
    assert ok
    ps = body.findall(".//" + q("a:p"))
    assert len(ps) == 3 and shape_text(body) == "Первый\nВторой\nТретий пункт"
    assert all(p.find(q("a:pPr")).find(q("a:buChar")) is not None for p in ps)
    assert all(r.get("sz") == "1400" for r in body.iter(q("a:rPr")))
    assert ps[2].find(".//" + q("a:rPr")).get("b") == "1"
    fill_text(body, [ParagraphSpec("Один абзац")], size_pt=11)
    assert body.find(".//" + q("a:rPr")).get("sz") == "1100"


def test_adjust_group_shrink_and_grow(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    cards = next(p for p in manifest.patterns if p.kind == PatternKind.cards)
    group = cards.repeat_groups[0]
    W, H = manifest.slide_size.w, manifest.slide_size.h
    # shrink 3 → 2 with even redistribution
    b = DeckBuilder(simple_deck)
    slide = b.clone_slide(cards.source_slide)
    cells, nid = adjust_group(slide, group, 2, W, H, b.next_shape_id(slide))
    assert len(cells) == 2
    boxes = [cell_bbox(c) for c in cells]
    assert boxes[1].x > boxes[0].x + boxes[0].w
    assert len([e for e in slide_shape_elements(slide).values() if shape_text(e).startswith("Pillar")]) == 2
    # grow 3 → 5 when the group allows it (max_n forced)
    group_big = group.model_copy(update={"max_n": 5})
    b2 = DeckBuilder(simple_deck)
    slide2 = b2.clone_slide(cards.source_slide)
    cells2, nid2 = adjust_group(slide2, group_big, 5, W, H, b2.next_shape_id(slide2))
    # three 26%-wide cards already span the slide: no duplicate fits, growth is capped by the slide bounds
    assert len(cells2) == 3
    # shrink the cards first so that duplicates fit, then grow
    b4 = DeckBuilder(simple_deck)
    slide4 = b4.clone_slide(cards.source_slide)
    for e in slide_shape_elements(slide4).values():
        box = element_bbox(e)
        if box and box[2] > W * 0.2:
            from verstka.rendering.deck import set_element_pos
            set_element_pos(e, x=int(box[0] * 0.6), w=int(box[2] * 0.6))
    small = group.model_copy(update={"max_n": 5, "gap": 0.02})
    cells4, _ = adjust_group(slide4, small, 5, W, H, b4.next_shape_id(slide4))
    assert len(cells4) >= 4
    xs = [cell_bbox(c).x for c in cells4]
    assert xs == sorted(xs) and len(set(xs)) == len(xs)
    slide2 = slide4
    ids = [nv.get("id") for nv in slide2._element.iter(q("p:cNvPr"))]
    assert len(ids) == len(set(ids))
    # clamp to max_n
    b3 = DeckBuilder(simple_deck)
    slide3 = b3.clone_slide(cards.source_slide)
    cells3, _ = adjust_group(slide3, group, 7, W, H, b3.next_shape_id(slide3))
    assert len(cells3) == max(group.max_n, 3)


def test_chart_and_table_native(simple_deck, tmp_path):
    outline = DeckOutline.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    b = DeckBuilder(simple_deck)
    slide = b.add_blank_slide(6)
    W, H = b.slide_w, b.slide_h
    add_chart(slide, Bbox(x=int(W * 0.05), y=int(H * 0.2), w=int(W * 0.5), h=int(H * 0.6)), ChartSpec(type="column", series_ids=["s1"], unit="чел.", highlight_index=4), outline, manifest.components.chart_style, manifest.tokens.typography, text_hex="000000")
    add_table(slide, Bbox(x=int(W * 0.58), y=int(H * 0.2), w=int(W * 0.38), h=int(H * 0.5)), TableData(columns=["Сценарий", "Базовый", "Про"], rows=[["Напоминание", "да", "да"], ["Эскалация", "нет", "12%"]]), manifest.components.table_style, manifest.tokens.typography)
    b.delete_original_slides()
    out = b.save(tmp_path / "ct.pptx")
    prs = Presentation(str(out))
    shapes = list(prs.slides[0].shapes)
    charts = [s for s in shapes if s.has_chart]
    tables = [s for s in shapes if s.has_table]
    assert len(charts) == 1 and len(tables) == 1
    assert charts[0].chart.plots[0].series[0].values[-1] == 12400
    assert tables[0].table.cell(0, 0).text == "Сценарий" and tables[0].table.cell(2, 2).text == "12%"
    assert number_format("%") == '0"%"'
    if find_soffice() and find_pdftoppm():
        from verstka.ingest.render import render_slides

        imgs = render_slides(out, tmp_path / "png", dpi=50)
        assert len(imgs) == 1
