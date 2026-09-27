"""«Исправить слайд» (REMARKS_SPEC §2): the remarks' words (api/remarks.py), one slide spliced into a deck with every
other slide checked by its fingerprint (pipeline/revise.py), the slide designer's fix mode with a scripted fake model
(planning/slide_edit.py), and the API job end to end offline — only the fixed slide changes, «верни как было» puts the
deck back exactly."""

from __future__ import annotations

import json
import threading
import time
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

from verstka.api import remarks as RM
from verstka.pipeline import revise as R
from verstka.schemas.audit import FixAction, Issue
from verstka.schemas.common import BboxFrac

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "outline_demo.json"


def _issue(id_: str, check: str, severity: str = "error", slide: int = 8, kind: str = "deterministic", message: str = "", bboxes=None, **kw) -> Issue:
    return Issue(id=id_, slide=slide, check_id=check, severity=severity, kind=kind, message=message or check, bboxes=[BboxFrac(**b) for b in (bboxes or [])], **kw)


# ---------------------------------------------------------------------------- remarks.py


def test_where_reads_the_place():
    assert RM.where([{"x": 0.022, "y": 0.967, "w": 0.306, "h": 0.0415}]) == "внизу слева"
    assert RM.where([{"x": 0.4, "y": 0.4, "w": 0.2, "h": 0.2}]) == "в центре"
    assert RM.where([{"x": 0.7, "y": 0.05, "w": 0.2, "h": 0.1}]) == "вверху справа"
    assert RM.where([{"x": 0.0, "y": 0.0, "w": 0.95, "h": 0.9}]) == "на большей части слайда"
    assert RM.where([]) == "" and RM.where([{"x": 1.2, "y": 0.5, "w": 0.1, "h": 0.1}]) == ""


def test_remarks_note_numbers_hints_and_wishes():
    titles = RM.check_titles()
    assert "(VLM)" not in titles["slide_content"] and titles["contrast_low"].endswith("4,5:1")
    missing = _issue("content_missing-8-13", "content_missing", message="6 из 12 текстов плана нет на слайде: «1-й месяц», «2-й месяц»…",
                     details={"missing": [f"{k}-й месяц" for k in range(1, 7)], "wanted": 12})
    overflow = _issue("text_overflow-8-2", "text_overflow", severity="warn", message="«Выручка растёт» не поместился", bboxes=[{"x": 0.05, "y": 0.8, "w": 0.3, "h": 0.1}])
    model = _issue("slide_content-8-1", "slide_content", severity="info", kind="model", message="вывод не следует из цифр", suggestion="сформулируй вывод по цифрам")
    note = RM.remarks_note([model, overflow, missing], titles)
    lines = note.splitlines()
    assert lines[0].startswith("- The quality check of the rendered slide found these problems")
    # severity first: the error, the warning, the model's remark
    assert lines[1].startswith("  1. Часть запланированного контента пропала со слайда — 6 из 12 текстов плана нет на слайде: «1-й месяц»")
    assert "«6-й месяц»" in lines[1] and "Как исправить: верни пропавшие тексты" in lines[1]
    assert lines[2].startswith("  2. Текст не поместился в свою рамку — «Выручка растёт» не поместился (внизу слева). Как исправить: сократи текст")
    assert lines[3].startswith("  3. Содержание слайда — вывод не следует из цифр.") and "Как исправить: сформулируй вывод по цифрам." in lines[3]
    assert "The person also asks" not in note
    with_wish = RM.remarks_note([missing], titles, wishes="покажи этапами")
    assert with_wish.splitlines()[-1].startswith("- The person also asks: «покажи этапами». Do exactly what they ask")
    assert RM.remarks_note([], titles, wishes="таблицей").startswith("- The person also asks: «таблицей»")


def test_critic_line_plurals_and_cap():
    one = [_issue("a", "content_missing")]
    assert RM.critic_line(8, one) == "Критик: слайд 8 — 1 замечание: часть запланированного контента пропала со слайда."
    two = [_issue("a", "text_clipped", slide=1), _issue("b", "margin_violation", severity="warn", slide=1)]
    assert RM.critic_line(1, two) == "Критик: слайд 1 — 2 замечания: текст обрезан краем слайда; контент заходит в поля у краёв."
    five = [_issue(str(k), c) for k, c in enumerate(["overlap", "text_overflow", "word_break", "too_many_bullets", "overlap"])]
    line = RM.critic_line(3, five)
    assert line.startswith("Критик: слайд 3 — 5 замечаний: два блока наложились друг на друга; ") and line.endswith(" и ещё 1.")


def test_match_fixed_counts_per_check_and_marks_new():
    requested = [_issue("o1", "overlap", bboxes=[{"x": 0.1, "y": 0.1, "w": 0.1, "h": 0.1}]), _issue("o2", "overlap", bboxes=[{"x": 0.1, "y": 0.5, "w": 0.1, "h": 0.1}]),
                 _issue("m1", "content_missing")]
    after = [_issue("o9", "overlap"), _issue("b1", "bullet_too_long", severity="warn"), _issue("g1", "grid_alignment", severity="info")]
    fixed, remaining = RM.match_fixed(requested, after)
    assert fixed == ["o1", "m1"]  # one overlap of two is gone: the first in number order counts as fixed
    assert [(r["check_id"], r["new"]) for r in remaining] == [("overlap", False), ("bullet_too_long", True)]  # info notes are not remarks


def test_match_fixed_never_calls_a_remark_that_stayed_fixed():
    """Two alike-titled remarks of one check, one of them still there by its element: that one stays (with the id the
    person saw), the other is fixed — not «the first in number order»."""
    top = _issue("mv-1-3", "margin_violation", severity="warn", slide=1, message="«Заголовок» заходит в поле", element_ids=["s1"], bboxes=[{"x": 0.1, "y": 0.1, "w": 0.3, "h": 0.1}])
    foot = _issue("mv-1-4", "margin_violation", severity="warn", slide=1, message="«Сноска» заходит в поле", element_ids=["s9"], bboxes=[{"x": 0.1, "y": 0.9, "w": 0.3, "h": 0.05}])
    still = _issue("mv-1-2", "margin_violation", severity="warn", slide=1, message="«Заголовок» заходит в поле", element_ids=["s1"], bboxes=[{"x": 0.1, "y": 0.1, "w": 0.3, "h": 0.1}])
    fixed, remaining = RM.match_fixed([top, foot], [still])
    assert fixed == ["mv-1-4"] and [r["id"] for r in remaining] == ["mv-1-3"] and not remaining[0]["new"]


def test_carry_ids_keeps_the_ids_of_the_remarks_that_stayed():
    """The audit numbers its issues through the deck: after a fix of slide 1 every later id shifts by one. Carried over,
    the remarks that stayed keep their ids on every slide, a remark that changed its words keeps its id, and a new
    remark never takes the id of a fixed one."""
    old = [
        _issue("text_clipped-1-1", "text_clipped", slide=1, message="Сноска обрезана", element_ids=["s9"]),
        _issue("margin_violation-1-2", "margin_violation", severity="warn", slide=1, message="«Сноска» в поле", element_ids=["s9"]),
        _issue("margin_violation-1-3", "margin_violation", severity="warn", slide=1, message="«Итог» в поле", element_ids=["s4"]),
        _issue("grid_alignment-2-4", "grid_alignment", severity="info", slide=2, message="край не по сетке", element_ids=["s2"]),
        _issue("content_missing-8-5", "content_missing", slide=8, message="6 из 12 текстов плана нет на слайде"),
        _issue("slide_content-3-9", "slide_content", severity="info", slide=3, kind="model", message="вывод не следует из цифр"),
    ]
    new = [
        _issue("margin_violation-1-1", "margin_violation", severity="warn", slide=1, message="«Итог» в поле", element_ids=["s4"]),
        _issue("overlap-1-2", "overlap", slide=1, message="«Сноска» легла на «Итог»", element_ids=["s4", "s9"]),
        _issue("grid_alignment-2-3", "grid_alignment", severity="info", slide=2, message="край не по сетке", element_ids=["s2"]),
        _issue("content_missing-8-4", "content_missing", slide=8, message="3 из 12 текстов плана нет на слайде"),
        _issue("margin_violation-1-5", "margin_violation", severity="warn", slide=1, message="«Сноска» в поле", element_ids=["s9"]),
    ]
    got = RM.carry_ids(old, new)
    assert [i.id for i in got] == ["margin_violation-1-3", got[1].id, "grid_alignment-2-4", "content_missing-8-5", "margin_violation-1-2"]
    assert got[1].id not in {i.id for i in old} and got[1].id.startswith("overlap-1-")  # new: an id no old remark had
    assert len({i.id for i in got}) == len(got) and [i.message for i in got] == [i.message for i in new]
    # the fix's own view of slide 1: the clipped footnote is fixed, the margin remarks stayed, the overlap is new
    fixed, remaining = RM.match_fixed(old[:3], [i for i in got if i.slide == 1])
    assert fixed == ["text_clipped-1-1"]
    assert {r["id"] for r in remaining} == {"margin_violation-1-2", "margin_violation-1-3", got[1].id}
    assert [r["new"] for r in remaining if r["check_id"] == "overlap"] == [True]
    # a report: per-slide index follows the carried ids
    from verstka.schemas.audit import AuditReport

    a, b = AuditReport(deck="a", template_id="t", issues=old), AuditReport(deck="b", template_id="t", issues=new).recompute()
    RM.carry_report_ids(a, b)
    assert b.per_slide[1] == ["margin_violation-1-3", got[1].id, "margin_violation-1-2"] and b.per_slide[8] == ["content_missing-8-5"]


def test_preview_line_says_the_work_goes_on():
    req = [_issue("a", "content_missing"), _issue("b", "text_clipped")]
    assert RM.preview_line(["a", "b"], req, []) == "Проверка: замечаний на слайде не осталось — готовлю превью слайда."
    left = [{"id": "b", "check_id": "text_clipped", "severity": "error", "message": "", "new": False}]
    assert RM.preview_line(["a"], req, left) == "Проверка: исправлено 1 из 2 — готовлю превью слайда."
    assert RM.preview_line([], [], left) == "Проверка: слайд переделан — готовлю превью слайда."


def test_job_runner_speaks_russian():
    """The runner's own messages reach the interface: «Начинаю», «Готово», «Не получилось: {the reason}»."""
    from verstka.api.jobs import JobRunner

    runner = JobRunner()

    def boom(job):
        raise RuntimeError("Модель не ответила\nподробности")

    def mute(job):
        raise ValueError()

    def settled(job):
        t0 = time.time()
        while not (job.events and job.events[-1]["status"] in ("done", "failed")):
            assert time.time() - t0 < 10
            time.sleep(0.02)
        return job

    ok, bad, empty = (settled(runner.submit("t", fn)) for fn in (lambda job: {"ok": True}, boom, mute))
    assert [e["message"] for e in ok.events] == ["Начинаю", "Готово"]
    assert bad.events[-1]["message"] == "Не получилось: Модель не ответила" and bad.error.startswith("Модель не ответила\n")
    assert empty.events[-1]["message"] == "Не получилось: ValueError" and empty.error.startswith("ValueError\n")


def test_fix_reply_cases():
    req = [_issue("a", "content_missing")]
    assert RM.fix_reply(8, ["a"], req, [], 77, 87, [], True, None, [], "было — пункты, стало — этапы (6 шагов)") == \
        "Слайд 8 исправлен: было — пункты, стало — этапы (6 шагов). Замечаний на нём не осталось. Оценка 77 → 87."
    assert RM.fix_reply(8, ["a"], req, [], 77, 77, [], True, None, [], "") == "Слайд 8 исправлен — замечаний на нём не осталось."
    left = [{"id": "x", "check_id": "bullet_too_long", "severity": "warn", "message": "", "new": True}]
    assert RM.fix_reply(8, [], req * 2, left, 77, 84, ["пожелания не учтены: модель не ответила"], True, None, [4, 5]) == \
        "Слайд 8: исправлено 0 из 2, осталось: пункт списка длиннее 15 слов. Оценка 77 → 84. Пожелания не учтены: модель не ответила. Изменились и слайды 4, 5."
    assert RM.fix_reply(8, [], req, [], 77, 77, [], False, "Замечания остались — попробуйте добавить пожелание", []) == \
        "Не получилось исправить слайд 8: замечания остались — попробуйте добавить пожелание."


# ---------------------------------------------------------------------------- splice + fingerprints


def _chart_data(cats, vals) -> CategoryChartData:
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Ряд", vals)
    return cd


def test_splice_takes_one_slide_and_keeps_the_rest(simple_deck: Path, tmp_path: Path):
    # the base deck: slide 2 has a chart of its own that the donor drops
    base = tmp_path / "base.pptx"
    prs = Presentation(str(simple_deck))
    prs.slides[1].shapes.add_chart(XL_CHART_TYPE.PIE, Emu(100000), Emu(100000), Emu(2000000), Emu(1500000), _chart_data(["Старое", "Было"], (1, 2)))
    prs.save(str(base))
    # the donor: the same deck with slide 2 changed — new text, a new picture, a new native chart, the old chart gone
    donor = tmp_path / "donor.pptx"
    dp = Presentation(str(base))
    s2 = dp.slides[1]
    s2.shapes.title.text = "Три опоры — новый заголовок"
    old_chart = next(sh for sh in s2.shapes if sh.has_chart)
    old_chart._element.getparent().remove(old_chart._element)  # its relationship stays behind, like a replaced sample chart
    img = tmp_path / "new.png"
    Image.new("RGB", (120, 80), (220, 40, 40)).save(img)
    s2.shapes.add_picture(str(img), Emu(500000), Emu(3000000), width=Emu(1200000))
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(6000000), Emu(2500000), Emu(4000000), Emu(3000000), _chart_data(["Q1", "Q2", "Q3"], (10, 20, 30)))
    dp.save(str(donor))

    out = tmp_path / "out.pptx"
    R.splice_slide(base, donor, 2, out)
    assert R.changed_slides(base, out, skip={2}) == []
    fo, fd, fb = R.slide_fingerprints(out), R.slide_fingerprints(donor), R.slide_fingerprints(base)
    assert fo[1] == fd[1] and fo[1] != fb[1]
    # it reopens; the chart has its workbook; the picture came along
    re = Presentation(str(out))
    assert re.slides[1].shapes.title.text == "Три опоры — новый заголовок"
    charts = [sh for sh in re.slides[1].shapes if sh.has_chart]
    assert len(charts) == 1 and list(charts[0].chart.plots[0].categories) == ["Q1", "Q2", "Q3"]
    wb = charts[0].chart.part.chart_workbook.xlsx_part
    assert wb is not None and wb.blob[:2] == b"PK"
    assert any(sh.shape_type == 13 for sh in re.slides[1].shapes)  # MSO_SHAPE_TYPE.PICTURE
    # the old slide's chart (and its workbook) are not written any more
    names = zipfile.ZipFile(out).namelist()
    assert sum(1 for n in names if n.startswith("ppt/charts/chart")) == 1
    assert sum(1 for n in names if n.startswith("ppt/embeddings/")) == 1


def test_fingerprints_ignore_the_workbook_timestamp(tmp_path: Path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    for path in (a, b):
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_chart(XL_CHART_TYPE.PIE, Emu(0), Emu(0), Emu(3000000), Emu(2000000), _chart_data(["А", "Б"], (1, 2)))
        prs.save(str(path))
        time.sleep(1.1)  # the embedded workbook's creation time differs
    assert R.slide_fingerprints(a) == R.slide_fingerprints(b)
    assert R.changed_slides(a, [*R.slide_fingerprints(b), "extra"]) == [2]


# ---------------------------------------------------------------------------- the designer's fix mode (fake model)


BRIEF = """Сервис напоминаний для команд

Слайд 1. Проблема
Сотрудники теряют задачи в потоке сообщений. 40% задач из чатов забываются. Напоминания вручную занимают 20 минут в день.

Слайд 2. Как внедряем
Запуск проходит в три этапа: пилот на одной команде, настройка сценариев, запуск на всю компанию.
"""


def _fake_designer(answer: dict):
    from verstka.providers.mock import MockProvider
    from verstka.providers.registry import ProviderLimits, ProviderRegistry

    prompts: list[str] = []

    def script(messages):
        system, user = messages[0].content, messages[-1].content
        if "presentation designer" in system:
            prompts.append(user)
            return answer
        return {"facts": [], "series": [], "tables": []}

    p = MockProvider(script, model="qwen/qwen3.8-27b")
    return ProviderRegistry(roles={"llm": p, "vlm": p}, limits=ProviderLimits(max_concurrency=2, time_budget_s=60)), prompts


def _two_slide_outline():
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent

    return DeckOutline(title="Сервис напоминаний", strategy="structured", slides=[
        OutlineSlide(id="sl1", kind=PatternKind.bullets, headline="Проблема", spec_ref=1, content=SlideContent(bullets=["Сотрудники теряют задачи в потоке сообщений", "40% задач из чатов забываются"])),
        OutlineSlide(id="sl2", kind=PatternKind.bullets, headline="Как внедряем", spec_ref=2, content=SlideContent(bullets=["Пилот на одной команде", "Настройка сценариев", "Запуск на всю компанию"])),
    ], agent_log=["Аналитик: прочитал текст."])


def test_redesign_slide_fix_mode_with_a_fake_designer():
    from verstka.planning.brief import parse_brief_text
    from verstka.planning.slide_edit import redesign_slide
    from verstka.skills_registry.registry import SkillsRegistry

    answer = {"kind": "process", "headline": "Запуск в три этапа", "items": [{"title": "Пилот", "text": "одна команда"}, {"title": "Настройка", "text": "сценарии"}, {"title": "Запуск", "text": "вся компания"}], "rationale": "Этапы читаются по порядку."}
    providers, prompts = _fake_designer(answer)
    events: list[dict] = []
    remark = _issue("content_missing-2-1", "content_missing", slide=2, message="2 из 4 текстов плана нет на слайде: «Пилот…»", details={"missing": ["Пилот на одной команде", "Настройка сценариев"]})
    note = RM.remarks_note([remark], wishes="покажи этапами")
    r = redesign_slide(
        _two_slide_outline(), 2, parse_brief_text(BRIEF), None, request="покажи этапами", note=note,
        start_line="переделываю по замечаниям и пожеланию «покажи этапами»", rationale="Исправлены замечания: часть запланированного контента пропала со слайда · по вашей просьбе «покажи этапами»",
        log_line="Исправление: слайд 2 — {what}.", skills=SkillsRegistry.load(), providers=providers, progress=events.append,
    )
    assert r.how == "model" and r.what.startswith("было — список (3 пункта), стало — ")
    s = r.outline.slides[1]
    assert s.id == "sl2" and s.kind.value in ("process", "timeline") and [it.title for it in s.content.items] == ["Пилот", "Настройка", "Запуск"]
    assert s.rationale.startswith("Исправлены замечания: часть запланированного контента пропала со слайда · по вашей просьбе «покажи этапами».")
    assert r.outline.agent_log[-1] == f"Исправление: слайд 2 — {r.what}."
    # the designer saw the remarks and the wishes; the timeline shows its start and its result under «Дизайнер»
    assert len(prompts) == 1 and "The quality check of the rendered slide found these problems" in prompts[0] and "«покажи этапами»" in prompts[0]
    msgs = [e["message"] for e in events if e.get("step") == "designer"]
    assert msgs[0] == "Переделываю по замечаниям и пожеланию «покажи этапами»." and len(msgs) == 2


def test_redesign_slide_without_a_model():
    from verstka.planning.brief import parse_brief_text
    from verstka.planning.slide_edit import redesign_slide, revise_slide

    brief = parse_brief_text(BRIEF)
    with pytest.raises(ValueError, match="model_unavailable"):
        redesign_slide(_two_slide_outline(), 2, brief, None, note="- fix it")
    # the chat's edit keeps its words
    with pytest.raises(ValueError, match="без модели могу только сменить форму"):
        revise_slide(_two_slide_outline(), 2, "перепиши заголовок", brief, None)


# ---------------------------------------------------------------------------- the API, offline


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("VERSTKA_WORKSPACE", str(tmp_path / "ws"))
    import importlib

    from fastapi.testclient import TestClient

    import verstka.api.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app), app_module


def _wait(c, job_id, timeout=300):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/jobs/{job_id}").json()
        if j["status"] in ("done", "failed"):
            return j
        time.sleep(0.2)
    raise TimeoutError(job_id)


def _fake_render_outputs(monkeypatch):
    """No LibreOffice here: previews drawn from each slide's fingerprint (an equal slide, an equal picture)."""
    import verstka.pipeline.generate as G

    def fake(vdir, manifest, title, exports, *, images=True, dpi=110):
        fps = R.slide_fingerprints(Path(vdir) / "deck.pptx")
        out = []
        if images:
            (Path(vdir) / "slides").mkdir(exist_ok=True)
            for k, fp in enumerate(fps, 1):
                path = Path(vdir) / "slides" / f"slide-{k:03d}.jpg"
                Image.new("RGB", (64, 36), tuple(int(fp[i:i + 2], 16) for i in (0, 2, 4))).save(path)
                out.append(path)
        return {}, out, [], None

    monkeypatch.setattr(G, "render_outputs", fake)


def _push_out(pptx: Path, index: int) -> None:
    """Slide `index`'s widest text box pushed past the right edge of the slide: a real out-of-bounds remark."""
    prs = Presentation(str(pptx))
    s = prs.slides[index - 1]
    box = max((sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.width), key=lambda sh: sh.width)
    box.left = int(prs.slide_width - box.width * 0.6)
    prs.save(str(pptx))


def test_slide_fix_api_only_this_slide_and_exact_undo(client, simple_deck, monkeypatch):
    from verstka.audit.runner import run_audit
    from verstka.ingest.render import find_soffice
    from verstka.schemas.outline import DeckOutline

    c, mod = client
    if find_soffice() is None:
        _fake_render_outputs(monkeypatch)
    with open(simple_deck, "rb") as f:
        r = c.post("/api/templates", files={"file": ("simple.pptx", f, "application/octet-stream")}, data={"use_models": "false"})
    tid = _wait(c, r.json()["job_id"])["result"]["template_id"]
    outline = json.loads(FIXTURE.read_text(encoding="utf-8"))
    r = c.post("/api/generations", json={"template_id": tid, "outline": outline, "strategies": ["structured"], "use_models": False, "exports": []})
    gid = r.json()["generation_id"]
    assert _wait(c, r.json()["job_id"])["status"] == "done"
    vdir = mod.store.generation_dir(gid) / "structured"

    # slide 3 (a list) gets a real remark: its text box pushed out of the slide, and the audit written again
    k = 3
    _push_out(vdir / "deck.pptx", k)
    manifest, ws = mod.store.manifest(tid), mod.store.workspace(tid)
    o = DeckOutline.model_validate_json((vdir / "outline.json").read_text(encoding="utf-8"))
    rep = run_audit(vdir / "deck.pptx", manifest, o, ws, render=False, strategy="structured")
    (vdir / "audit_report.json").write_text(rep.model_dump_json(indent=2), encoding="utf-8")
    asked = [i for i in rep.issues if i.slide == k and RM.is_stage(i)]
    pushed = {i.id for i in asked if i.check_id in ("out_of_bounds", "text_clipped", "margin_violation")}
    assert pushed and any(i.autofix and i.autofix.action == "move_inside" for i in asked), [(i.check_id, i.autofix) for i in asked]
    # every slide of the synthetic template also has a remark the rules cannot fix (its 10 pt footer is not in the scale)
    unfixable = {i.check_id for i in asked if i.id not in pushed}
    before = R.slide_fingerprints(vdir / "deck.pptx")
    before_images = {p.name: p.read_bytes() for p in (vdir / "slides").glob("slide-*.jpg")}

    # validation: a slide that is not there, a slide without remarks, remarks of another slide
    base = f"/api/generations/{gid}/structured/slides"
    r = c.post(f"{base}/99/fix", json={})
    assert r.status_code == 422 and r.json()["detail"] == "В этом варианте 12 слайдов — слайда 99 нет"
    clean = 5
    cleaned = rep.model_copy(deep=True)
    cleaned.issues = [i for i in cleaned.issues if i.slide != clean or not RM.is_stage(i)]
    (vdir / "audit_report.json").write_text(cleaned.model_dump_json(indent=2), encoding="utf-8")
    r = c.post(f"{base}/{clean}/fix", json={})
    assert r.status_code == 422 and r.json()["detail"] == f"На слайде {clean} нет замечаний"
    (vdir / "audit_report.json").write_text(rep.model_dump_json(indent=2), encoding="utf-8")
    r = c.post(f"{base}/{k}/fix", json={"issue_ids": ["nope-1-1"]})
    assert r.status_code == 422 and r.json()["detail"] == f"На слайде {k} нет таких замечаний"
    assert c.post(f"/api/generations/{gid}/nope/slides/{k}/fix", json={}).json()["detail"] == "Вариант не найден"

    # the fix, held until a second request has been refused
    gate = threading.Event()
    real = R.fix_slide

    def held(*a, **kw):
        gate.wait(30)
        return real(*a, **kw)

    monkeypatch.setattr(R, "fix_slide", held)
    r = c.post(f"{base}/{k}/fix", json={"wishes": "  "})
    assert r.status_code == 200
    job_id = r.json()["job_id"]
    busy = c.post(f"{base}/{k}/fix", json={})
    assert busy.status_code == 409 and busy.json()["detail"] == "Этот вариант уже меняется — дождитесь, пока агент закончит"
    gate.set()
    job = _wait(c, job_id)
    assert job["status"] == "done", job["error"]
    res = job["result"]
    assert res["kind"] == "fix" and res["applied"] is True and res["changed"] is True and res["slide"] == k
    assert res["how"] == "autofix" and res["version"] == 1 and res["at"]
    assert set(res["requested"]) == {i.id for i in asked} and set(res["fixed"]) == pushed
    assert {x["check_id"] for x in res["remaining"]} == unfixable and not any(x["new"] for x in res["remaining"])
    assert res["changed_other_slides"] is False and res["other_slides"] == []
    # the reply names what is left, or — when the template layers leave no unfixable remark on this slide — says so
    if unfixable:
        assert res["reply"].startswith(f"Слайд {k}: исправлено {len(pushed)} из {len(asked)}, осталось: ")
    else:
        assert res["reply"].startswith(f"Слайд {k} исправлен") and "не осталось" in res["reply"]
    assert res["new_score"] > res["score_before"]
    assert res["notes"] == [] and res["why"] is None and res["errors"] == 0
    steps = [e["step"] for e in job.get("agent", [])]
    assert steps[:1] == ["critic"] and {"designer", "compile", "check"} <= set(steps)
    assert all(e["slide"] == k and e["variant"] == "structured" for e in job["agent"])
    # while the previews render, the check step says so: the timeline never reads as finished before the job is
    last = job["agent"][-1]
    verdict = f"Исправлено {len(pushed)} из {len(asked)}" if unfixable else "Замечаний на слайде не осталось"
    assert last["step"] == "check" and last["message"] == f"{verdict} — готовлю превью слайда."
    checks = [e["message"] for e in job["agent"] if e["step"] == "check"]
    assert last["progress"] >= 0.85 and (checks[-2].startswith(f"Исправлено {len(pushed)} из {len(asked)}, осталось:") if unfixable else len(checks) >= 1)
    # the remarks keep their ids through the fix: the ones left on slide k are the ones the person asked about, and
    # every other slide's remarks keep theirs although the audit numbers them through the deck
    fixed_audit = json.loads((vdir / "audit_report.json").read_text(encoding="utf-8"))
    stage_after = {i["id"] for i in fixed_audit["issues"] if i["slide"] == k and RM.is_stage(Issue.model_validate(i))}
    assert stage_after == {x["id"] for x in res["remaining"]} and stage_after <= {i.id for i in asked} - pushed
    others = lambda issues: sorted((i["slide"], i["check_id"], i["id"]) for i in issues if i["slide"] not in (0, k))  # noqa: E731
    assert others(fixed_audit["issues"]) == others(json.loads(rep.model_dump_json())["issues"])

    # only slide k changed; the other previews are the very same files
    after = R.slide_fingerprints(vdir / "deck.pptx")
    assert R.changed_slides(before, after, skip={k}) == [] and after[k - 1] != before[k - 1]
    after_images = {p.name: p.read_bytes() for p in (vdir / "slides").glob("slide-*.jpg")}
    assert all(after_images[nm] == data for nm, data in before_images.items() if nm != f"slide-{k:03d}.jpg")
    g = c.get(f"/api/generations/{gid}").json()
    v = g["variants"][0]
    assert v["edits"][-1]["kind"] == "fix" and v["edits"][-1]["slides"] == [k] and v["edits"][-1]["request"] == f"Исправь замечания на слайде {k}"
    assert v["edits"][-1]["version"] == 1 and v["edits"][-1]["how"] == "autofix" and v["edits"][-1]["fixed"] == res["fixed"]
    assert g["summary"]["structured"]["score"] == res["new_score"]
    audit = json.loads((vdir / "audit_report.json").read_text(encoding="utf-8"))
    assert ".structured.fix" not in json.dumps(audit) and audit["deck"] == str(vdir / "deck.pptx")
    assert any(f.get("action") == "slide_fix" and f.get("slide") == k for f in audit["applied_fixes"])
    assert not (vdir.parent / ".structured.fix").exists()

    # «верни как было» restores the deck exactly (no re-render), and the previews with it
    r = c.post(f"/api/generations/{gid}/structured/edits", json={"message": "верни как было", "slide": k})
    job = _wait(c, r.json()["job_id"])
    assert job["status"] == "done", job["error"]
    assert job["result"]["reply"].startswith(f"Вернул как было до правки «Исправь замечания на слайде {k}»")
    assert R.slide_fingerprints(vdir / "deck.pptx") == before
    assert {p.name: p.read_bytes() for p in (vdir / "slides").glob("slide-*.jpg")} == before_images
    assert json.loads((vdir / "audit_report.json").read_text(encoding="utf-8"))["summary"]["score"] == rep.summary.score
    assert R.version_numbers(vdir) == []


def test_slide_fix_designer_path_keeps_the_other_slides(client, simple_deck, monkeypatch):
    """The designer's path (a fake model redesigns slide 3): the scratch deck is rendered whole, but only slide 3 is
    spliced in — every other slide keeps its fingerprint."""
    from verstka.ingest.render import find_soffice
    from verstka.skills_registry.registry import SkillsRegistry

    c, mod = client
    if find_soffice() is None:
        _fake_render_outputs(monkeypatch)
    with open(simple_deck, "rb") as f:
        r = c.post("/api/templates", files={"file": ("simple.pptx", f, "application/octet-stream")}, data={"use_models": "false"})
    tid = _wait(c, r.json()["job_id"])["result"]["template_id"]
    outline = json.loads(FIXTURE.read_text(encoding="utf-8"))
    r = c.post("/api/generations", json={"template_id": tid, "outline": outline, "strategies": ["structured"], "use_models": False, "exports": []})
    gid = r.json()["generation_id"]
    assert _wait(c, r.json()["job_id"])["status"] == "done"
    vdir = mod.store.generation_dir(gid) / "structured"
    k = 3
    old = outline["slides"][k - 1]
    answer = {"kind": "cards", "headline": old["headline"], "items": [{"title": b.split(" ")[0], "text": b} for b in old["content"]["bullets"][:3]], "rationale": "Карточки читаются быстрее."}
    providers, prompts = _fake_designer(answer)
    remark = Issue(id=f"fill_ratio-{k}-1", slide=k, check_id="fill_ratio", severity="warn", kind="deterministic", message="слайд заполнен слишком плотно", outline_id=old["id"],
                   autofix=FixAction(action="none"))
    before = R.slide_fingerprints(vdir / "deck.pptx")
    events: list = []
    res = R.fix_slide(
        vdir, "structured", k, manifest=mod.store.manifest(tid), ws=mod.store.workspace(tid), remarks=[remark], wishes="покажи карточками",
        use_models=True, skills=SkillsRegistry.load(), providers=providers, exports=[], progress=lambda *a, **kw: events.append(a[0] if a else kw),
    )
    assert len(prompts) == 1 and "«покажи карточками»" in prompts[0] and "Слайд слишком пустой или слишком плотный" in prompts[0]
    assert res.how == "model" and res.applied, res.reply
    assert res.other == [] and R.changed_slides(before, vdir / "deck.pptx", skip={k}) == []
    assert R.slide_fingerprints(vdir / "deck.pptx")[k - 1] != before[k - 1]
    o = json.loads((vdir / "outline.json").read_text(encoding="utf-8"))
    # the outline says what the deck shows (the synthetic template has no cards: the render lays them out otherwise)
    assert o["slides"][k - 1]["kind"] != "bullets" and o["slides"][k - 1]["rationale"].startswith("Исправлены замечания: слайд слишком пустой или слишком плотный · по вашей просьбе «покажи карточками».")
    assert [s["id"] for s in o["slides"]] == [s["id"] for s in outline["slides"]] and o["agent_log"][-1].startswith(f"Исправление: слайд {k} — было — ")
    agent = [e for e in events if isinstance(e, dict) and e.get("type") == "agent"]
    assert [e["step"] for e in agent][:2] == ["critic", "designer"] and agent[1]["message"] == "Переделываю по замечаниям и пожеланию «покажи карточками»."
    assert R.read_edits(vdir)[-1]["request"] == f"Исправь слайд {k}: покажи карточками"


def test_slide_fix_designer_tries_another_form_when_content_is_still_lost(client, simple_deck, monkeypatch):
    """The designer's first version keeps a list that cannot show the slide's steps (content_missing again): it is
    told so and tries another form once; the better version is kept."""
    from verstka.ingest.render import find_soffice
    from verstka.providers.mock import MockProvider
    from verstka.providers.registry import ProviderLimits, ProviderRegistry
    from verstka.schemas.audit import AuditReport
    from verstka.skills_registry.registry import SkillsRegistry

    c, mod = client
    if find_soffice() is None:
        _fake_render_outputs(monkeypatch)
    with open(simple_deck, "rb") as f:
        r = c.post("/api/templates", files={"file": ("simple.pptx", f, "application/octet-stream")}, data={"use_models": "false"})
    tid = _wait(c, r.json()["job_id"])["result"]["template_id"]
    outline = json.loads(FIXTURE.read_text(encoding="utf-8"))
    r = c.post("/api/generations", json={"template_id": tid, "outline": outline, "strategies": ["structured"], "use_models": False, "exports": []})
    gid = r.json()["generation_id"]
    assert _wait(c, r.json()["job_id"])["status"] == "done"
    vdir = mod.store.generation_dir(gid) / "structured"
    k = 3
    old = outline["slides"][k - 1]
    steps = [{"title": f"Этап {i}", "text": b} for i, b in enumerate(old["content"]["bullets"], 1)]
    answers = [{"kind": "bullets", "headline": old["headline"], "bullets": old["content"]["bullets"][:2], "items": steps}, {"kind": "process", "headline": old["headline"], "items": steps}]
    prompts: list[str] = []

    def script(messages):
        if "presentation designer" in messages[0].content:
            prompts.append(messages[-1].content)
            return answers[min(len(prompts), 2) - 1]
        return {"facts": [], "series": [], "tables": []}

    p = MockProvider(script, model="qwen/qwen3.8-27b")
    providers = ProviderRegistry(roles={"llm": p, "vlm": p}, limits=ProviderLimits(max_concurrency=2, time_budget_s=60))
    remark = Issue(id=f"content_missing-{k}-1", slide=k, check_id="content_missing", severity="error", kind="deterministic", message="2 из 5 текстов плана нет на слайде",
                   outline_id=old["id"], autofix=FixAction(action="rematch", params={"outline_id": old["id"]}))
    rep = AuditReport.model_validate_json((vdir / "audit_report.json").read_text(encoding="utf-8"))
    rep.issues.append(remark)
    rep.recompute()
    (vdir / "audit_report.json").write_text(rep.model_dump_json(indent=2), encoding="utf-8")
    before = R.slide_fingerprints(vdir / "deck.pptx")
    events: list = []
    res = R.fix_slide(
        vdir, "structured", k, manifest=mod.store.manifest(tid), ws=mod.store.workspace(tid), remarks=[remark], use_models=True,
        skills=SkillsRegistry.load(), providers=providers, exports=[], progress=lambda *a, **kw: events.append(a[0] if a else kw),
    )
    assert len(prompts) == 2 and "Your previous version of this slide used the kind «bullets»" in prompts[1]
    assert res.applied and res.how == "model" and res.fixed == [remark.id] and res.what.startswith("было — список (3 пункта), стало — ")
    assert R.changed_slides(before, vdir / "deck.pptx", skip={k}) == []
    lines = [e["message"] for e in events if isinstance(e, dict) and e.get("step") == "designer"]
    assert "Пробую другую форму — прежняя не вместила всё." in lines
    shares = [e["progress"] for e in events if isinstance(e, dict) and e.get("type") == "agent"]
    assert shares == sorted(shares)  # the progress never goes back


def test_restore_version_puts_the_kept_files_back(simple_deck: Path, tmp_path: Path, monkeypatch):
    """A version keeps the deck, the plan, the audit, the previews and the exports; restoring it copies them back
    (nothing rendered). An older version without previews gets them rendered again from its deck."""
    import verstka.pipeline.generate as G
    from verstka.schemas.audit import AuditReport

    vdir = tmp_path / "structured"
    (vdir / "slides").mkdir(parents=True)
    import shutil

    shutil.copy2(simple_deck, vdir / "deck.pptx")
    (vdir / "outline.json").write_text('{"title": "x", "slides": []}', encoding="utf-8")
    (vdir / "layout_plan.json").write_text('{"strategy": "structured", "template_id": "t"}', encoding="utf-8")
    (vdir / "audit_report.json").write_text(AuditReport(deck="elsewhere/deck.pptx", template_id="t").model_dump_json(), encoding="utf-8")
    (vdir / "deck.pdf").write_bytes(b"%PDF old")
    for k in (1, 2, 3):
        Image.new("RGB", (32, 18), (k * 40, 0, 0)).save(vdir / "slides" / f"slide-{k:03d}.jpg")
    old_files = {p.relative_to(vdir).as_posix(): p.read_bytes() for p in vdir.rglob("*") if p.is_file()}
    calls: list[bool] = []

    def fake(vd, manifest, title, exports, *, images=True, dpi=110):
        calls.append(images)
        imgs = []
        for k in (1, 2, 3):
            path = Path(vd) / "slides" / f"slide-{k:03d}.jpg"
            path.parent.mkdir(exist_ok=True)
            Image.new("RGB", (32, 18), (0, 0, 255)).save(path)
            imgs.append(path)
        return {}, imgs, [], True

    monkeypatch.setattr(G, "render_outputs", fake)
    n = R.snapshot(vdir)
    assert (vdir / "versions" / f"v{n}" / "slides" / "slide-002.jpg").exists() and (vdir / "versions" / f"v{n}" / "deck.pdf").exists()
    # the variant changes
    prs = Presentation(str(vdir / "deck.pptx"))
    prs.slides[0].shapes.title.text = "Другое"
    prs.save(str(vdir / "deck.pptx"))
    (vdir / "deck.pdf").write_bytes(b"%PDF new")
    Image.new("RGB", (32, 18), (0, 255, 0)).save(vdir / "slides" / "slide-001.jpg")
    report = R.restore_version(vdir, n, None, "x", ["pdf"])
    now = {p.relative_to(vdir).as_posix(): p.read_bytes() for p in vdir.rglob("*") if p.is_file() and "versions" not in p.parts}
    assert now == old_files | {"audit_report.json": now["audit_report.json"]} and calls == []
    assert report.deck == str(vdir / "deck.pptx") and report.slide_images[2] == str(vdir / "slides" / "slide-002.jpg")
    # an older version: no previews, no pdf kept — they are rendered again from the kept deck
    shutil.rmtree(vdir / "versions" / f"v{n}" / "slides")
    (vdir / "versions" / f"v{n}" / "deck.pdf").unlink()
    R.restore_version(vdir, n, None, "x", ["pdf"])
    assert calls == [True] and R.slide_fingerprints(vdir / "deck.pptx") == R.slide_fingerprints(simple_deck)


def test_critic_notes_of_redone_slides_are_dropped():
    """After a fix (or a chat redesign) of a slide, the build's critic notes about its old version leave «Почему так»;
    an undone edit brings them back."""
    from verstka.api.app import _redone_slides

    edits = [
        {"version": 1, "kind": "fix", "slides": [1]}, {"version": None, "undo_of": 1, "kind": "undo", "slides": []},
        {"version": 1, "kind": "fix", "slides": [8]}, {"version": 2, "kind": "slide", "slides": [3]},
        {"version": None, "undo_of": 2, "kind": "undo", "slides": []}, {"version": 2, "kind": "swap", "slides": [4, 5]},
    ]
    assert _redone_slides(edits) == {8}
    assert _redone_slides([]) == set()
