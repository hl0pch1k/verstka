"""Design on any template (2026-10-17, direction D): a lone thesis is the slide's message, not a list of one row; a
short block stands at the optical middle of its area, never in the top third over an empty foot (the audit says so:
`bottom_empty`); a template whose scale says nothing of its headings lets running text grow to the deck's heading; the
lines of one list on a photo ground stand on one plate; the free place for the author's photo is the deck's own (a
tint of its accent, a dashed edge, a pictogram) and the audit tells it from a grey stub (`photo_place_empty`); a
table's text keeps the readable threshold; a chart grows down to the conclusion set at the foot."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.analysis.xmlns import q
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.compose import Canvas, Composer, Kit, Para, Run
from verstka.rendering.deck import element_bbox
from verstka.schemas.common import Bbox, PatternKind, contrast_ratio
from verstka.schemas.outline import PHOTO_PLACE_NAME, DeckOutline, OutlineSlide, SlideContent, TableData


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _composer(env, o: OutlineSlide, strategy: str = "structured"):
    deck, manifest, _ = env
    prs = Presentation(str(deck))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = manifest.slide_size.w, manifest.slide_size.h
    kit = Kit(manifest, W, H, "FFFFFF")
    kit.head_size = 28.0
    return Composer(slide, kit, o, DeckOutline(title="T", strategy=strategy, slides=[o]), strategy, manifest), W, H


def _names(c: Composer, n0: int) -> list[str]:
    return [(el.find(".//" + q("p:cNvPr")).get("name") or "") for el in list(c.cv.tree)[n0:]]


THESIS = ("Функция «Умные напоминания» превращает сообщение в задачу одним касанием, предлагает срок по контексту "
          "переписки и присылает утренний дайджест.")


def test_a_lone_thesis_is_the_slides_message_not_a_list_of_one_row(env):
    o = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Что сделали", content=SlideContent(bullets=[THESIS]))
    c, W, H = _composer(env, o)
    n0 = len(c.cv.tree)
    c.compose("bullets", Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.88 * W), h=int(0.68 * H)))
    names = _names(c, n0)
    assert not any(n.startswith("Marker") for n in names), names
    assert any(n.startswith("Statement") for n in names), names
    # (on the slide's fallbacks too: a process slide whose steps are one line)
    o2 = OutlineSlide(id="p", kind=PatternKind.process, headline="План запуска", content=SlideContent(bullets=[THESIS]))
    c2, _, _ = _composer(env, o2)
    n0 = len(c2.cv.tree)
    c2.compose("process", Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.88 * W), h=int(0.68 * H)))
    assert not any(n.startswith("Marker") for n in _names(c2, n0))


def test_a_short_block_stands_at_the_optical_middle_of_its_area(env):
    o = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Контекст", content=SlideContent(bullets=["Сотрудники теряют задачи в чатах", "Ручные напоминания ставят 18%"]))
    c, W, H = _composer(env, o)
    area = Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.88 * W), h=int(0.7 * H))
    n0 = len(c.cv.tree)
    c.compose("bullets", area)
    boxes = [element_bbox(el) for el in list(c.cv.tree)[n0:]]
    top = min(b[1] for b in boxes if b)
    bottom = max(b[1] + b[3] for b in boxes if b)
    above, below = top - area.y, area.y2 - bottom
    assert above > 0.05 * H, "a short list is not left in the top third"
    assert below >= above and below - above <= 0.25 * area.h, (above / H, below / H)


def test_a_full_block_keeps_the_content_line_under_the_heading(env):
    rows = [f"Пункт {i}: строка списка, которая занимает почти всю ширину колонки и переносится на вторую строку" for i in range(6)]
    o = OutlineSlide(id="s", kind=PatternKind.bullets, headline="Много", content=SlideContent(bullets=rows))
    c, W, H = _composer(env, o)
    area = Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.88 * W), h=int(0.6 * H))
    n0 = len(c.cv.tree)
    c.compose("bullets", area)
    assert min(element_bbox(el)[1] for el in list(c.cv.tree)[n0:] if element_bbox(el)) <= area.y + int(0.01 * H)
    assert not any("optical middle" in w for w in c.warnings)


# ---------------------------------------------------------------------------------------------- the audit: bottom_empty


def _text_slide(path: Path, body_y: float, lines: int = 2) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Emu(int(0.06 * W)), Emu(int(0.06 * H)), Emu(int(0.88 * W)), Emu(int(0.1 * H)))
    t.text_frame.text = "Заголовок слайда"
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    b = s.shapes.add_textbox(Emu(int(0.06 * W)), Emu(int(body_y * H)), Emu(int(0.88 * W)), Emu(int(0.05 * H * lines)))
    b.text_frame.word_wrap = True
    b.text_frame.text = "\n".join(f"Строка текста номер {i}" for i in range(lines))
    for p in b.text_frame.paragraphs:
        p.runs[0].font.size = Pt(18)
    prs.save(str(path))
    return path


def _audit_ctx(deck: Path, kind: PatternKind = PatternKind.bullets):
    from verstka.audit.ir import build_deck_ir
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import BboxFrac
    from verstka.schemas.template import SlideSize, Spacing, TemplateManifest, Tokens

    ir = build_deck_ir(deck)
    o = DeckOutline(title="T", slides=[OutlineSlide(id="a", kind=PatternKind.title, headline="T"), OutlineSlide(id="b", kind=kind, headline="Заголовок слайда")])
    ir.slides[0].outline_id = "b"
    ir.slides[0].index = 2  # (a content slide, not the cover)
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=ir.slide_w, h=ir.slide_h), tokens=Tokens(spacing=Spacing(safe_area=BboxFrac(x=0.05, y=0.05, w=0.9, h=0.88))), patterns=[], n_slides=1)
    return AuditContext(ir=ir, manifest=man, outline=o)


def test_the_audit_sees_a_block_in_the_top_third_over_an_empty_foot(tmp_path):
    from verstka.audit.checks.layout import bottom_empty

    got = bottom_empty(_audit_ctx(_text_slide(tmp_path / "top.pptx", 0.2)))
    assert len(got) == 1 and got[0].check_id == "bottom_empty" and got[0].details["empty_below"] > 0.42
    # the same lines at the optical middle, a full slide, and a slide whose air is meant (one large figure): silent
    assert bottom_empty(_audit_ctx(_text_slide(tmp_path / "mid.pptx", 0.42))) == []
    assert bottom_empty(_audit_ctx(_text_slide(tmp_path / "full.pptx", 0.2, lines=11))) == []
    assert bottom_empty(_audit_ctx(_text_slide(tmp_path / "big.pptx", 0.2), PatternKind.big_number)) == []


# ------------------------------------------------------------------------------------------------- clone: growth cap


def test_running_text_grows_to_the_decks_heading_on_a_template_whose_scale_has_none():
    from verstka.rendering.clone import _short_cap, _text_cap

    typo = SimpleNamespace(size_for=lambda role, default=0.0: 14.0 if role in ("body", "h1") else default)
    slide = SimpleNamespace(__dict__={})
    ctx = SimpleNamespace(typo=typo, slide=slide)
    assert _text_cap(ctx) == 14.0  # (no heading known: the scale's)
    slide.__dict__["_verstka_head_size"] = 25.0
    assert _text_cap(ctx) >= 19.0
    assert 20.0 <= _short_cap(ctx) <= 25.0 * 0.85
    # only a roomy box inside the slide lets a short text grow that far (a narrow side box at the edge keeps the cap)
    from verstka.rendering.clone import _roomy
    from verstka.schemas.common import BboxFrac

    assert _roomy(SimpleNamespace(bbox=BboxFrac(x=0.05, y=0.25, w=0.9, h=0.6)))
    assert not _roomy(SimpleNamespace(bbox=BboxFrac(x=0.72, y=0.2, w=0.3, h=0.6)))
    # a template with a real heading size keeps its own bound
    typo2 = SimpleNamespace(size_for=lambda role, default=0.0: {"body": 9.0, "h1": 24.0}.get(role, default))
    assert _text_cap(SimpleNamespace(typo=typo2, slide=SimpleNamespace(__dict__={"_verstka_head_size": 40.0}))) == 16.0


# ------------------------------------------------------------------------------------------------- one plate per list


def _photo_ground(path: Path) -> Path:
    """A layout: dark purple over the slide, an opaque white glow over the right part of the middle band."""
    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    layout = prs.slide_layouts[6]
    scratch = prs.slides.add_slide(layout)
    ground = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    glow = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(0.45 * W)), Emu(int(0.52 * H)), Emu(int(0.55 * W)), Emu(int(0.12 * H)))
    for sh, rgb in ((ground, "560043"), (glow, "FFFFFF")):
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(rgb)
        sh.line.fill.background()
    for sh in (glow, ground):
        el = sh._element
        el.getparent().remove(el)
        layout.shapes._spTree.insert(2, el)
    prs.save(str(path))
    return path


def test_the_rows_of_a_list_on_a_photo_stand_on_one_plate(tmp_path):
    from verstka.rendering.synth import _ink_pass

    deck = _photo_ground(tmp_path / "g.pptx")
    manifest = analyze_template(deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    prs = Presentation(str(deck))
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cv = Canvas(s)
    rows = []
    for i, y in enumerate((0.3, 0.42, 0.56)):  # the third row crosses the dark ground and the white glow
        cv.ellipse(Bbox(x=int(0.1 * W), y=int((y + 0.015) * H), w=int(0.01 * W), h=int(0.01 * W)), "FFFFFF", name="Marker")
        rows.append(cv.text(Bbox(x=int(0.12 * W), y=int(y * H), w=int(0.8 * W), h=int(0.06 * H)), [Para([Run(f"Пункт списка номер {i + 1}, длинная строка через весь слайд", 16.0, "FFFFFF", False, "Arial")])], name="Item"))
    pal = SimpleNamespace(text="FFFFFF", text2="FFFFFF", bg="560043")
    _ink_pass(s, list(cv.added), pal, manifest, [])
    plates = [el for el in s._element.iter(q("p:sp")) if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Plate")]
    assert len(plates) == 1, "one plate under the whole list, not under its third row"
    pb = element_bbox(plates[0])
    for r in rows:
        b = element_bbox(r)
        assert pb[1] <= b[1] and pb[1] + pb[3] >= b[1] + int(0.03 * H)
    fill = plates[0].find(".//" + q("a:solidFill") + "/" + q("a:srgbClr")).get("val")
    for r in rows:
        assert all(contrast_ratio(c.get("val"), fill) >= 4.5 for c in r.iter(q("a:srgbClr")))
    tree = plates[0].getparent()
    assert all(tree.index(plates[0]) < tree.index(el) for el in cv.added), "the plate lies under the markers and the rows"


# ---------------------------------------------------------------------------------------------------- the photo place


def _sat(h: str) -> float:
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.0 if max(r, g, b) == 0 else (max(r, g, b) - min(r, g, b)) / max(r, g, b)


@pytest.mark.parametrize("ground,accent", [("FFFFFF", "FF7A00"), ("FFF7D6", "F2C200"), ("101820", "2E9BFF")])
def test_the_photo_place_is_a_tint_of_the_accent_with_a_dashed_edge(ground, accent):
    from verstka.rendering.synth import PHOTO_FILL_CONTRAST, _photo_fill, _photo_outline

    kit = SimpleNamespace(card=SimpleNamespace(fill=None), manifest=None, colors=SimpleNamespace(accent=accent))
    pal = SimpleNamespace(bg=ground, surface=None, card_fill=None, text="000000", accent=accent)
    got = _photo_fill(kit, pal, ground)
    lo, hi = PHOTO_FILL_CONTRAST
    assert lo <= contrast_ratio(got, ground) <= hi + 0.05, got
    assert _sat(got) >= 0.05, "never a grey stub where the template has an accent"
    edge = _photo_outline(kit, pal, got)
    assert contrast_ratio(edge, got) > 1.15


def test_the_audit_tells_a_designed_photo_place_from_a_grey_stub():
    from verstka.audit.checks.integrity import photo_place_empty
    from verstka.audit.registry import AuditContext
    from verstka.schemas.common import BboxFrac
    from verstka.schemas.deck_ir import DeckIR, IRElement, IRSlide
    from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

    def ctx(fill: str, line, own: bool):
        place = IRElement(id="ph", type="shape", name=PHOTO_PLACE_NAME, fill_hex=fill, line_hex=line, bbox=Bbox(x=0, y=0, w=1000, h=1000), bbox_frac=BboxFrac(x=0.5, y=0.2, w=0.4, h=0.6))
        o = DeckOutline(title="T", slides=[OutlineSlide(id="a", kind=PatternKind.bullets, headline="Пространство", content=SlideContent(bullets=["Зал: 100 м²"], photo_slot="Фото помещения", photo_own=own))])
        man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=12192000, h=6858000), tokens=Tokens(), patterns=[], n_slides=1)
        return AuditContext(ir=DeckIR(source="x.pptx", slide_w=12192000, slide_h=6858000, slides=[IRSlide(index=1, elements=[place], outline_id="a")]), manifest=man, outline=o)

    designed = photo_place_empty(ctx("FFE9D6", "FFB27A", True))
    assert [i.severity for i in designed] == ["info"]
    stub = photo_place_empty(ctx("EEEEEE", None, True))
    assert [i.severity for i in stub] == ["warn"]
    # a place the person asked a picture for (described in the text, «добавь картинку») and the picture did not come:
    # it stays as asked, a note with its reason, not a fault (the owner's rule of 20.10); a place Verstka added itself
    # never reaches a finished deck empty (pipeline/generate.relay_without_pictures) — the audit still warns of one
    missing = photo_place_empty(ctx("FFE9D6", "FFB27A", False))
    assert [i.severity for i in missing] == ["info"] and "картинки нет" in missing[0].message


# ---------------------------------------------------------------------------------------------------------- tables


def test_a_tables_text_keeps_the_readable_threshold(env):
    tables = pytest.importorskip("verstka.rendering.tables")
    if not hasattr(tables, "readable_table_pt"):
        pytest.skip("no readable threshold of tables in this tree")
    t = TableData(columns=["Регион", "Выручка", "Доля"], rows=[[f"Регион {i}", f"{100 + i * 7}", f"{i + 3}%"] for i in range(10)])
    o = OutlineSlide(id="t", kind=PatternKind.table, headline="Регионы", content=SlideContent(table=t))
    c, W, H = _composer(env, o)
    n0 = len(c.cv.tree)
    c.compose("table", Bbox(x=int(0.06 * W), y=int(0.2 * H), w=int(0.88 * W), h=int(0.62 * H)))
    sizes = [int(r.get("sz")) / 100 for el in list(c.cv.tree)[n0:] for r in el.iter(q("a:rPr")) if r.get("sz") and el.find(".//" + q("a:tbl")) is not None]
    assert sizes and min(sizes) >= c._table_readable() - 0.05, (sizes, c._table_readable())


# -------------------------------------------------------------------------------------------- clone: chart to the strip


def test_a_chart_grows_down_to_the_conclusion_at_the_foot(tmp_path):
    from verstka.rendering.clone_notes import _chart_down_to

    prs = Presentation()
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["a", "b", "c"]
    cd.add_series("s", (1, 2, 3))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(int(0.05 * W)), Emu(int(0.2 * H)), Emu(int(0.9 * W)), Emu(int(0.35 * H)), cd)
    el = gf._element
    bottom = int(0.55 * H)
    _chart_down_to(s, bottom, int(0.7 * H), int(0.05 * W), int(0.95 * W), H)
    b = element_bbox(el)
    assert abs(b[1] + b[3] - int(0.7 * H)) <= 2 and b[1] == int(0.2 * H)
    # never more than half again its height, and a chart that is not the content's foot stays
    gf2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(int(0.05 * W)), Emu(int(0.2 * H)), Emu(int(0.9 * W)), Emu(int(0.2 * H)), cd)
    _chart_down_to(s, int(0.4 * H), int(0.9 * H), int(0.05 * W), int(0.95 * W), H)
    assert element_bbox(gf2._element)[3] <= int(0.2 * H * 1.5) + 2


def test_a_change_to_nothing_is_said_as_the_change_not_as_minus_hundred_percent():
    from verstka.rendering.compose import table_deltas

    t = TableData(columns=["Статья", "Сейчас", "В VK Cloud"], rows=[["Поддержка", "38", "0"], ["Облако", "1", "27"], ["Команда", "14", "8"]])
    got = {lab: v for v, lab, _span in table_deltas(t)}
    assert got["Поддержка"] == "−38" and got["Облако"] == "+26" and got["Команда"] == "−43%"
    assert not any("100%" in v for v in got.values())


def test_an_empty_place_verstka_added_itself_is_a_warning():
    from verstka.audit.checks.integrity import photo_place_empty
    from verstka.audit.registry import AuditContext
    from verstka.planning.pictures import AUTO_LABEL
    from verstka.schemas.common import BboxFrac
    from verstka.schemas.deck_ir import DeckIR, IRElement, IRSlide
    from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

    place = IRElement(id="ph", type="shape", name=PHOTO_PLACE_NAME, fill_hex="FFE9D6", line_hex="FFB27A", bbox=Bbox(x=0, y=0, w=1000, h=1000), bbox_frac=BboxFrac(x=0.5, y=0.2, w=0.4, h=0.6))
    o = DeckOutline(title="T", slides=[OutlineSlide(id="a", kind=PatternKind.bullets, headline="Склад", content=SlideContent(bullets=["Зал: 100 м²"], photo_slot=AUTO_LABEL, photo_hint="Склад"))])
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=12192000, h=6858000), tokens=Tokens(), patterns=[], n_slides=1)
    got = photo_place_empty(AuditContext(ir=DeckIR(source="x.pptx", slide_w=12192000, slide_h=6858000, slides=[IRSlide(index=1, elements=[place], outline_id="a")]), manifest=man, outline=o))
    assert [i.severity for i in got] == ["warn"] and "картинки нет" in got[0].message


# ------------------------------------------------------------------------------- unseen templates: readable, filled


def test_a_statement_under_a_label_heading_is_as_large_as_the_label(env):
    """«Предложение» only names the slide: the statement under it is set as large as that heading, not a step under it
    (a short ask at body size over an empty slide — unseen templates, 20.10); under a heading that says the conclusion
    the statement stays a step below."""
    from verstka.schemas.common import PatternKind

    o = OutlineSlide(id="a", kind=PatternKind.bullets, headline="Предложение", content=SlideContent(bullets=["Перенести хранилище в облако"]))
    c, W, H = _composer(env, o)
    assert c._under_heading() >= 0.95 * c.kit.head_size - 0.05
    o2 = OutlineSlide(id="b", kind=PatternKind.bullets, headline="Перенос в облако окупится за семь месяцев", content=SlideContent(bullets=["Перенести хранилище"]))
    c2, _, _ = _composer(env, o2)
    assert c2._under_heading() <= 0.9 * c2.kit.head_size + 0.05


def test_a_row_of_short_steps_keeps_its_lines_readable(env):
    """Four short steps («1 000 обращений») with long lines: the row of steps stays (a column of four names over empty
    gaps read worse — Vintage, 20.10) and the lines under it read at 2.6 % of the slide height at least (10 pt before)."""
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import SlideItem

    items = [SlideItem(title=t) for t in ("1 000 обращений", "500 записей", "350 посещений", "175 покупателей")]
    lines = ["На первые три месяца продвижения выделяется 240 000 рублей.",
             "План на этот период — получить 1 000 обращений, записать 500 человек на пробное занятие и продать 175 абонементов.",
             "Из них 90 000 рублей направят на рекламу в социальных сетях, 60 000 рублей — на районные сообщества, 40 000 — на партнёрства.",
             "Средняя стоимость привлечения покупателя составит около 1 371 рубля."]
    o = OutlineSlide(id="p", kind=PatternKind.process, headline="Привлечение клиентов", content=SlideContent(items=items, bullets=lines))
    c, W, H = _composer(env, o)
    n0 = len(c.cv.tree)
    c.compose("process", Bbox(x=int(0.06 * W), y=int(0.22 * H), w=int(0.88 * W), h=int(0.62 * H)))
    names = _names(c, n0)
    sizes = [int(r.get("sz")) / 100 for el in list(c.cv.tree)[n0:] if (el.find(".//" + q("p:cNvPr")).get("name") or "").startswith("Note") for r in el.iter(q("a:rPr")) if r.get("sz")]
    assert any(n.startswith("Note") for n in names) and sizes and min(sizes) >= 0.026 * H / 12700 - 0.05, (names, sizes)


def test_a_step_beside_its_year_keeps_the_names_capital(env):
    """A timeline set as steps beside the slide's lines: «1684 — Лейбниц ввёл…», never «лейбниц» (R27, 20.10)."""
    from verstka.schemas.common import PatternKind
    from verstka.schemas.outline import SlideItem

    items = [SlideItem(title="1684", text="Лейбниц ввёл обозначение производной"), SlideItem(title="1736", text="Вышел «Метод флюксий» Ньютона")]
    lines = ["Оба подхода сошлись в одном понятии производной, которым пользуются до сих пор в анализе и физике.",
             "Обозначения Лейбница прижились в учебниках, точки Ньютона — в механике."]
    o = OutlineSlide(id="t", kind=PatternKind.timeline, headline="История производной", content=SlideContent(items=items, bullets=lines))
    c, W, H = _composer(env, o)
    assert c._steps_beside(Bbox(x=int(0.06 * W), y=int(0.22 * H), w=int(0.88 * W), h=int(0.62 * H)), items, lines)
    text = " ".join("".join(t.text or "" for t in el.iter(q("a:t"))) for el in c.cv.tree)
    assert "Лейбниц" in text and "лейбниц" not in text
