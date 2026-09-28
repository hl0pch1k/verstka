"""Gate 4 (G4-13, G4-21): the audit reads the words a written deck shows — lines cut out of their sentences, clock times
and project numbers as key figures, charts of two quantities, repeated headings, lines that open with «Также» or
«они», and a topic deck whose text was never written. Cases are the gate-4 decks' own lines (writer replays)."""

from __future__ import annotations

from verstka.audit.registry import AuditContext, run_checks
from verstka.schemas.common import PatternKind
from verstka.schemas.deck_ir import DeckIR, IRSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, NumberCallout, OutlineSlide, SlideContent, SlideItem
from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

MARK = "Текст написан агентом Verstka по статье «Вторая мировая война» из Википедии (лицензия CC BY-SA). Проверьте факты перед выступлением."
NEW = {"line_fragment", "figure_is_time", "orphan_opener", "duplicate_heading", "chart_mixed_units", "writer_failed"}


def _deck(slides: list[OutlineSlide], written: bool = True, writer=None, brief: str | None = None) -> AuditContext:
    """A deck of the plan's slides: a cover first (its notes carry the writer's attribution when `written`)."""
    cover = OutlineSlide(id="cover", kind=PatternKind.title, headline="Тема", notes=MARK if written else "")
    o = DeckOutline(title="Тема", slides=[cover, *slides])
    ir = DeckIR(source="x.pptx", slide_w=12192000, slide_h=6858000, slides=[IRSlide(index=i + 1, outline_id=s.id, notes=s.notes) for i, s in enumerate(o.slides)])
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=12192000, h=6858000), tokens=Tokens(), patterns=[], n_slides=1)
    return AuditContext(ir=ir, manifest=man, outline=o, writer=writer, brief_text=brief)


def _issues(ctx: AuditContext, check: str):
    got, _ = run_checks(ctx, only={check})
    return [i for i in got if i.severity != "info"]


def _bullets(sid: str, head: str, bullets: list[str], notes: str) -> OutlineSlide:
    return OutlineSlide(id=sid, kind=PatternKind.bullets, headline=head, notes=notes, content=SlideContent(bullets=bullets))


# ------------------------------------------------------------------------------------------------ line_fragment


def test_line_fragment_flags_a_line_cut_out_of_its_sentence():
    """WWII s7: «6 июня 1944 года союзные силы США» — the sentence's start without its predicate (the Normandy landing
    lost); Гагарин A1 s3/s4: a line that ends on «только», a line stopped before the sentence's «, но …»."""
    ww2 = _bullets("s7", "Переломные события (продолжение)", [
        "6 июня 1944 года союзные силы США",
        "В августе 1945 года СССР вступил в войну против Японии",
    ], "6 июня 1944 года союзные силы США, Великобритании и Канады после двух месяцев отвлекающих манёвров провели "
       "крупнейшую десантную операцию в истории и высадились в Нормандии. В августе 1945 года СССР вступил в войну против Японии.")
    gag3 = _bullets("s3", "Начало", ["Выключение двигателя произошло только"],
                    "Выключение двигателя произошло только после срабатывания дублирующего механизма, но корабль уже поднялся на орбиту.")
    gag4 = _bullets("s4", "Ход событий", ["В конце полёта ТДУ конструктора Исаева проработала успешно"],
                    "В конце полёта ТДУ конструктора Исаева проработала успешно, но отключилась на секунду раньше, в результате "
                    "чего автоматика выдала запрет на штатное разделение отсеков.")
    got = _issues(_deck([ww2, gag3, gag4]), "line_fragment")
    assert [i.slide for i in got] == [2, 3, 4]
    assert "6 июня 1944 года союзные силы США" in got[0].message and "нет сказуемого" in got[0].message
    assert "кончается на «только»" in got[1].message
    assert "но отключилась на секунду раньше" in got[2].message and got[2].details["lines"][0]["text"].startswith("В конце полёта")


def test_line_fragment_leaves_whole_sentences_list_items_and_timeline_entries_alone():
    """A whole sentence, an item of the sentence's enumeration («Агрессивная политика нацистской Германии» of «Основными
    причинами стали …»), a clause cut before a plain «, а …», a timeline entry under its date and a faithful deck (no
    writer) are proper lines."""
    s2 = _bullets("s2", "Война началась в 1939 году", [
        "Агрессивная политика нацистской Германии",
        "Рост милитаризма в Японии",
        "Какао-бобы растирали в пасту с маисом и острым перцем",
        "В августе 1945 года СССР вступил в войну против Японии",
    ], "Основными причинами стали агрессивная политика нацистской Германии, ограничения Версальского договора и рост "
       "милитаризма в Японии. Какао-бобы растирали в пасту с маисом и острым перцем, а напиток взбивали до получения пены. "
       "В августе 1945 года СССР вступил в войну против Японии.")
    tl = OutlineSlide(id="s3", kind=PatternKind.timeline, headline="Mail.ru провела IPO", notes="В 2010 году компания провела IPO на Лондонской фондовой бирже под названием Mail.ru Group.",
                      content=SlideContent(items=[SlideItem(title="2010", text="IPO на Лондонской фондовой бирже под названием Mail.ru Group"),
                                                  SlideItem(title="6 и 9 августа 1945", text="США бомбардировали Хиросиму и Нагасаки")]))
    card = OutlineSlide(id="s5", kind=PatternKind.cards, headline="Причины войны", notes="Агрессивная политика нацистской Германии привела к аншлюсу Австрии и захвату Чехословакии.",
                        content=SlideContent(items=[SlideItem(title="Агрессивная политика нацистской Германии", text="Аншлюс Австрии, захват Чехословакии")]))
    assert _issues(_deck([s2, tl, card]), "line_fragment") == []  # a card's noun-phrase title is the card's heading
    cut = _bullets("s4", "Нормандия", ["6 июня 1944 года союзные силы США"], "6 июня 1944 года союзные силы США высадились в Нормандии.")
    assert _issues(_deck([cut], written=False), "line_fragment") == []  # faithful mode: the user's own words


def test_line_fragment_reads_a_circumstance_torn_off_its_clause():
    s = _bullets("s3", "12 апреля 1961 года стартовал «Восток-1»", ["С Юрием Гагариным на борту"],
                 "12 апреля 1961 года с космодрома Байконур стартовал корабль «Восток-1» с Юрием Гагариным на борту.")
    got = _issues(_deck([s]), "line_fragment")
    assert len(got) == 1 and "обстоятельство" in got[0].message


# ------------------------------------------------------------------------------------------------ figure_is_time


def test_figure_is_time_flags_a_clock_time_and_a_project_number_shown_as_quantities():
    """Гагарин L1 s4: «10 ч» of «в 10 часов 53 минуты» (a clock time read as ten hours); energy s5: «22220» of «ледоколы
    проекта 22220». «106 минут» (a duration) and «10:53» shown as a time are figures."""
    gag = OutlineSlide(id="s4", kind=PatternKind.stat_row, headline="Ход событий", notes="Корабль сделал один оборот вокруг Земли, и в 10 часов 53 минуты посадка произошла в районе деревни Смеловка. Полёт длился 106 минут.",
                       content=SlideContent(numbers=[NumberCallout(value="10 ч", label="Корабль сделал один оборот вокруг Земли"), NumberCallout(value="106 минут", label="Полёт длился")]))
    energy = OutlineSlide(id="s5", kind=PatternKind.stat_row, headline="Строительство АЭС", notes="Также развивается транспортная ядерная энергетика, включая ледоколы проекта 22220. Россия строит более 10 атомных энергоблоков.",
                          content=SlideContent(numbers=[NumberCallout(value="22220", label="проекта ледоколы"), NumberCallout(value="более 10", label="атомных энергоблоков строит Россия")]))
    clock = OutlineSlide(id="s6", kind=PatternKind.big_number, headline="Посадка", notes="Посадка произошла в 10:53 по московскому времени.",
                         content=SlideContent(numbers=[NumberCallout(value="10:53", label="посадка по московскому времени")]))
    got = _issues(_deck([gag, energy, clock]), "figure_is_time")
    assert [(i.slide, i.details["value"], i.details["kind"]) for i in got] == [(2, "10 ч", "time"), (3, "22220", "code")]
    assert "10 часов 53 минуты" in got[0].message and "проекта 22220" in got[1].message


# ------------------------------------------------------------------------------------------------ orphan_opener


def test_orphan_opener_flags_lines_that_point_outside_themselves():
    vk5 = _bullets("s5", "Продукты и сервисы", ["VK владеет социальными сетями", "Также она владеет поисковой системой «Поиск Mail»"], "")
    vk3 = _bullets("s3", "Первым продуктом стал почтовый сервис Mail.ru", ["Вдохновлённые покупкой Hotmail Microsoft, они предложили создать публичный почтовый сервис"], "")
    ev3 = OutlineSlide(id="ev3", kind=PatternKind.big_number, headline="В том же году было продано 17,8 тысячи электромобилей", content=SlideContent(numbers=[NumberCallout(value="17,8 тыс", label="электромобилей продано в 2024 году")]))
    ww9 = OutlineSlide(id="ww9", kind=PatternKind.stat_row, headline="Итоги и потери", content=SlideContent(numbers=[NumberCallout(value="62", label="государства в ней"), NumberCallout(value="более 70 млн", label="человек погибло")]))
    choc = _bullets("ch6", "Итоги и последствия", ["Большим гурманом и любителем шоколада был и её первый министр Никита Панин"], "")
    got = _issues(_deck([vk5, vk3, ev3, ww9, choc]), "orphan_opener")
    assert [i.slide for i in got] == [2, 3, 4, 5, 6]
    msgs = [i.message for i in got]
    assert "«Также»" in msgs[0] and "«они»" in msgs[1] and "«В том же году»" in msgs[2] and "«ней»" in msgs[3] and "«её»" in msgs[4]


def test_orphan_opener_keeps_lines_with_their_own_names():
    ok = OutlineSlide(id="s2", kind=PatternKind.cards, headline="Полёт Гагарина", content=SlideContent(
        bullets=["Гагарин и его экипаж готовились два года", "Когда VK купила Mail.ru, она стала крупнейшим холдингом", "2012 — компания запустила мессенджер"],
        items=[SlideItem(title="Mail.ru Group", text="Она владеет «Одноклассниками»")]))
    assert _issues(_deck([ok]), "orphan_opener") == []
    faithful = _bullets("s3", "План", ["Также запускаем доставку"], "")
    assert _issues(_deck([faithful], written=False), "orphan_opener") == []


# ------------------------------------------------------------------------------------------------ duplicate_heading and chart_mixed_units


def test_duplicate_heading_flags_the_second_slide_with_the_same_heading():
    """VK s8 repeated s2 «VK была основана в 1998 году»; «… (продолжение)» and the closing slide are other headings."""
    slides = [
        _bullets("s2", "VK была основана в 1998 году", ["a b c"], ""),
        _bullets("s3", "Ход событий", ["a b c"], ""),
        _bullets("s4", "Ход событий (продолжение)", ["a b c"], ""),
        _bullets("s8", "VK была основана в 1998 году.", ["a b c"], ""),
        OutlineSlide(id="end", kind=PatternKind.thanks, headline="Тема"),
    ]
    got = _issues(_deck(slides, written=False), "duplicate_heading")
    assert [(i.slide, i.details["other"]) for i in got] == [(5, 2)]


def test_chart_mixed_units_flags_a_chart_of_two_quantities():
    """Energy s6: a column chart of 36 power units against 54 countries; a pie of shares in % and a chart of one counted
    thing are one quantity."""
    mixed = OutlineSlide(id="s6", kind=PatternKind.chart, headline="Россия — четвёртая по мощности атомной генерации",
                         notes="В стране эксплуатируются 36 энергоблоков общей мощностью около 28,6 ГВт. Россия экспортирует ядерные технологии и топливо в 54 страны.",
                         content=SlideContent(chart=ChartSpec(type="column", categories=["энергоблоков эксплуатируются в стране", "россия экспортирует ядерные технологии"], series=[InlineSeries(name="x", values=[36, 54])])))
    pie = OutlineSlide(id="s7", kind=PatternKind.chart, headline="Доходы VK по направлениям",
                       notes="Основные источники дохода: онлайн-реклама (42 %), онлайн-игры (32 %) и дополнительные платные сервисы (18,6 %).",
                       content=SlideContent(chart=ChartSpec(type="pie", categories=["Онлайн-реклама", "Онлайн-игры", "Дополнительные платные сервисы"], series=[InlineSeries(name="%", values=[42, 32, 18.6])])))
    evs = OutlineSlide(id="s8", kind=PatternKind.chart, headline="Электромобили в 2024 году",
                       notes="В 2024 году в России зарегистрировали 59,6 тыс. электромобилей. Продали 17,8 тысячи электромобилей.",
                       content=SlideContent(chart=ChartSpec(type="column", categories=["Зарегистрировали", "Продали"], series=[InlineSeries(name="x", values=[59.6, 17.8])])))
    got = _issues(_deck([mixed, pie, evs]), "chart_mixed_units")
    assert [i.slide for i in got] == [2]
    assert "36 энергоблоков" in got[0].message and "54 страны" in got[0].message


# ------------------------------------------------------------------------------------------------ writer_failed (G4-21)


def _skeleton() -> list[OutlineSlide]:
    return [
        OutlineSlide(id="ag", kind=PatternKind.agenda, headline="О чём поговорим", content=SlideContent(items=[SlideItem(title=h) for h in ("Контекст", "Главное", "Детали")])),
        *[OutlineSlide(id=f"sec{k}", kind=PatternKind.section, headline=h, notes=f"Добавьте сюда тезисы и цифры раздела «{h}»") for k, h in enumerate(("Контекст", "Главное", "Детали"))],
        OutlineSlide(id="end", kind=PatternKind.thanks, headline="Спасибо за внимание"),
    ]


def test_writer_failed_makes_an_unwritten_topic_deck_an_error():
    """ws/runs/20260928-061318-ff267f: the model was down, the topic deck was a skeleton of dividers and scored 100."""
    from verstka.schemas.audit import AuditReport

    got = _issues(_deck(_skeleton(), written=False, writer={"status": "failed", "reason": "all model links failed"}), "writer_failed")
    assert len(got) == 1 and got[0].severity == "error" and got[0].slide == 0
    assert "модель недоступна" in got[0].message and got[0].details == {"status": "failed", "dividers": 3}
    rep = AuditReport(deck="x.pptx", template_id="t", issues=got).recompute()
    assert rep.summary.score < 100
    assert "отказалась" in _issues(_deck(_skeleton(), written=False, writer="refused"), "writer_failed")[0].message
    # a written deck, a deck with content slides and a brief that was not a topic (no writer) are not flagged
    assert _issues(_deck(_skeleton(), written=False, writer={"status": "written"}), "writer_failed") == []
    assert _issues(_deck(_skeleton(), written=False, writer=None), "writer_failed") == []
    real = [_bullets("s2", "Факты", ["12 апреля 1961 года Гагарин полетел в космос"], ""), *_skeleton()]
    assert _issues(_deck(real, written=False, writer={"status": "failed"}), "writer_failed") == []


def test_run_audit_and_autofix_pass_the_writer_record_through(tmp_path):
    """The generation hands the writer's record to run_audit (and autofix_loop): a skeleton deck of a failed writer
    gets the error from the real entry point."""
    import inspect

    from pptx import Presentation

    from verstka.audit.autofix import autofix_loop
    from verstka.audit.runner import run_audit

    assert "writer" in inspect.signature(autofix_loop).parameters
    prs = Presentation()
    for _ in range(6):
        prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "deck.pptx"
    prs.save(path)
    cover = OutlineSlide(id="cover", kind=PatternKind.title, headline="Полёт Гагарина")
    outline = DeckOutline(title="Полёт Гагарина", slides=[cover, *_skeleton()])
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=prs.slide_width, h=prs.slide_height), tokens=Tokens(), patterns=[], n_slides=1)
    rep = run_audit(path, man, outline, None, use_vlm=False, use_llm=False, render=False, writer={"status": "failed"})
    assert any(i.check_id == "writer_failed" and i.severity == "error" for i in rep.issues)
    rep2 = run_audit(path, man, outline, None, use_vlm=False, use_llm=False, render=False)
    assert not any(i.check_id == "writer_failed" for i in rep2.issues)


# ------------------------------------------------------------------------------------------------ remarks


def test_the_slide_designer_gets_the_writing_checks_with_hints():
    from verstka.api import remarks as RM
    from verstka.schemas.audit import Issue

    for cid in ("line_fragment", "figure_is_time", "chart_mixed_units", "duplicate_heading", "orphan_opener"):
        assert cid in RM.CONTENT_CHECKS and RM.FIX_HINT.get(cid)
    lines = [{"text": "6 июня 1944 года союзные силы США", "role": "bullet", "why": "нет сказуемого"},
             {"text": "Выключение двигателя произошло только", "role": "bullet", "why": "кончается на «только»"}]
    i = Issue(id="line_fragment-7-1", slide=7, check_id="line_fragment", severity="warn", kind="deterministic",
              message="строка «6 июня 1944 года союзные силы США» — обрывок предложения: нет сказуемого (и ещё 1)", details={"lines": lines})
    note = RM.remarks_note([i])
    assert "«Выключение двигателя произошло только» (кончается на «только»)" in note and "Как исправить: возьми из текста слайда предложение целиком" in note
