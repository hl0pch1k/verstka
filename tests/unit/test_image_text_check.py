"""No writing in a slide picture (images/text_check.py): the vision model looks at every picture before it goes on a
slide; a picture with letters or digits on it is drawn again from a stricter prompt, and a place whose every picture
had writing stays free and says why. The image and vision models are scripts here."""

from __future__ import annotations

from pptx import Presentation

from tests.unit.test_images import _deck_with_places
from verstka.images import providers as IP
from verstka.images.fill import apply_pictures, plan_pictures
from verstka.images.prompt import ImageAsk, build_prompt
from verstka.images.service import ImageMaker
from verstka.images.text_check import WHY_TEXT, TextCheck
from verstka.providers.mock import MockProvider
from verstka.providers.registry import ProviderRegistry
from verstka.skills_registry.registry import SkillsRegistry

SKILLS = SkillsRegistry.load()


def _eyes(dirty: int, calls: list):
    """A vision model that sees writing on the first `dirty` pictures it is shown."""

    def answer(messages):
        calls.append(len(messages[-1].images))
        return {"text": "some", "what": "digits on a screen"} if len(calls) <= dirty else {"text": "none", "what": ""}

    p = MockProvider(answer, model="qwen/qwen3.8-27b")
    return ProviderRegistry(roles={"vlm": p})


def test_a_picture_with_writing_is_drawn_again_from_a_stricter_prompt(tmp_path):
    mock = IP.MockImages()
    looks: list = []
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "pics", deck_title="Умная теплица", skills=SKILLS, providers=_eyes(1, looks))
    made = maker.make(ImageAsk(hint="Экран панели управления теплицей с показаниями датчиков", headline="Панель", deck="Умная теплица"), "4:3")
    maker.close()
    assert len(mock.calls) == 2 and looks == [1, 1]  # each picture looked at once, alone
    assert made.attempts == 2 and made.text_checked
    assert not mock.calls[0]["prompt"].startswith("A wordless image") and mock.calls[1]["prompt"].startswith("A wordless image")
    assert "Every surface is blank" in mock.calls[1]["prompt"]
    assert len(list((tmp_path / "pics").glob("*.jpg"))) == 1  # the picture with writing is not kept


def test_a_place_whose_every_picture_has_writing_stays_free_and_says_why(tmp_path):
    path, outline = _deck_with_places(tmp_path)
    mock = IP.MockImages()
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "pics", deck_title=outline.title, skills=SKILLS, providers=_eyes(99, []))
    res = apply_pictures(path, plan_pictures(path, outline, maker), wait_s=10)
    maker.close()
    assert len(mock.calls) == 3 and res.changed == []
    assert res.records[0]["status"] == "free" and res.records[0]["why"] == WHY_TEXT
    # the third attempt draws a flat illustration, never a screen
    assert "Clean modern flat vector illustration made of simple geometric shapes" in mock.calls[2]["prompt"]
    assert not [sh for sh in Presentation(str(path)).slides[0].shapes if sh.shape_type == 13]


def test_without_a_vision_model_the_first_picture_is_taken(tmp_path):
    mock = IP.MockImages()
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "pics", deck_title="d")
    made = maker.make(ImageAsk(hint="Вид теплицы сверху", headline="h", deck="d"), "4:3")
    maker.close()
    assert len(mock.calls) == 1 and made.attempts == 1 and not made.text_checked


def test_the_strict_prompts_keep_the_subject_and_the_rules():
    ask = ImageAsk(hint="Экран приложения", headline="h", deck="d", subject_en="A laptop showing a dashboard with three tiles")
    p0, s0 = build_prompt(ask, None)
    p1, s1 = build_prompt(ask, None, strict=1)
    p2, s2 = build_prompt(ask, None, strict=2)
    assert s0 == s1 == "screen" and s2 == "illustration"
    assert all("A laptop showing a dashboard" in p for p in (p0, p1, p2))
    assert "pictograms and progress bars instead of any text" in p0 and "screen shows only soft coloured rounded shapes" in p1
    assert all(len(p) <= 1000 for p in (p0, p1, p2))


def test_the_vision_models_answer_is_read_leniently():
    assert TextCheck.model_validate({"text": "нет"}).clean and TextCheck.model_validate({"text": False}).clean
    assert not TextCheck.model_validate({"text": "yes", "what": "  a sign  "}).clean
    assert TextCheck.model_validate({"text": True}).text == "some" and TextCheck.model_validate({"what": None}).what == ""


def _grader(fits: list[int], calls: list, seen_ctx: list):
    """A vision model that sees no writing and grades the pictures it is shown by `fits`, in turn."""

    def answer(messages):
        calls.append(1)
        seen_ctx.append(messages[-1].content if hasattr(messages[-1], "content") else "")
        k = min(len(calls), len(fits)) - 1
        return {"text": "none", "what": "", "fits": fits[k], "flaws": "subject small" if fits[k] < 7 else ""}

    return ProviderRegistry(roles={"vlm": MockProvider(answer, model="qwen/qwen3.8-27b")})


def test_a_picture_that_fits_its_slide_badly_is_drawn_again_and_the_best_is_kept(tmp_path):
    """Pictures to the ideal (20.10): the vision model also grades how well a picture fits its slide (0–10); one under
    quality.min_fit is drawn again from the same prompt (no stricter rules: it had no writing), and the best clean
    picture is kept — the others are not left in the folder."""
    mock = IP.MockImages()
    calls: list = []
    ctx: list = []
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "pics", deck_title="Пекарня", skills=SKILLS, providers=_grader([4, 6, 5], calls, ctx))
    made = maker.make(ImageAsk(hint="Витрина пекарни с хлебом", headline="40 магазинов", deck="Пекарня"), "1:1")
    maker.close()
    assert len(mock.calls) == 3 and made.attempts == 3 and made.fits == 6
    assert not any(c["prompt"].startswith("A wordless image") for c in mock.calls)
    assert len(list((tmp_path / "pics").glob("*.jpg"))) == 1
    assert any("40 магазинов" in str(c) for c in ctx)  # the vision model is told the slide and the subject
    mock2, calls2 = IP.MockImages(), []
    maker = ImageMaker(IP.ImageChain([mock2]), None, tmp_path / "p2", deck_title="Пекарня", skills=SKILLS, providers=_grader([9], calls2, []))
    assert maker.make(ImageAsk(hint="Витрина", headline="h", deck="d"), "1:1").fits == 9 and len(mock2.calls) == 1
    maker.close()


def test_the_grade_is_read_leniently():
    assert TextCheck.model_validate({"text": "none", "fits": "7/10"}).fits == 7
    assert TextCheck.model_validate({"text": "none", "fits": 12}).fits == 10 and TextCheck.model_validate({"text": "none", "fits": "?"}).fits == 10
    assert TextCheck.model_validate({"text": "none"}).fits == 10  # an answer without a grade never rejects a picture
