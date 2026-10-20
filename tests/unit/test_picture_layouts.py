"""A generated picture takes a form of the slide, not a box in it (the owner, 18.10, п. 9: «картинки всегда
прямоугольник, не подходят ни по месту, ни по форме»): a side column bleeds to the slide's edge, a deck of circles gets
a round picture, a deck of rounded cards a rounded one, a place over most of the slide is its background with an
underlay under the text; the picture is cropped to its frame by what it shows; it is asked in its frame's composition
and only where the slide has room for it."""

from __future__ import annotations

import io

import numpy as np
import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from verstka.analysis.xmlns import q
from verstka.images import palette as P
from verstka.images.places import CAPTION_PREFIX, places_in, put_picture, take_picture_out
from verstka.images.prompt import ImageAsk, build_prompt, template_style
from verstka.rendering import picture_layouts as PL
from verstka.schemas.outline import PHOTO_PLACE_NAME, DeckOutline, OutlineSlide, SlideContent

W, H = Inches(13.333), Inches(7.5)


def _deck(place=(7.2, 1.6, 5.4, 5.0), title_w=6.0, extra=None, circles=0, rounded=0, geom=MSO_SHAPE.RECTANGLE, number=False):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(title_w), Inches(1))
    t.text_frame.text = "Головной офис"
    s.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(5.8), Inches(3)).text_frame.text = "Короткий текст рядом с картинкой."
    if number:
        n = s.shapes.add_textbox(Inches(12.6), Inches(7.0), Inches(0.5), Inches(0.4))
        n.text_frame.text = "6"
    x, y, w, h = place
    frame = s.shapes.add_shape(geom, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.name = PHOTO_PLACE_NAME
    cap = s.shapes.add_textbox(Inches(x + 0.3), Inches(y + 1), Inches(w - 0.6), Inches(1))
    cap.text_frame.text = "Место под изображение"
    cap.name = f"{CAPTION_PREFIX} {cap.shape_id}"
    for k in range(circles):
        s2 = prs.slides.add_slide(prs.slide_layouts[6])
        s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1 + k), Inches(1), Inches(3), Inches(3))
    for k in range(rounded):
        s2 = prs.slides.add_slide(prs.slide_layouts[6])
        s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1 + k * 0.1), Inches(4), Inches(2))
    if extra:
        extra(s)
    return prs, s


def _frame(prs, s):
    pl = places_in(prs, None)[0]
    return pl, PL.frames_of(prs, [pl])[(1, pl.shape_id)]


def _jpeg(w=1200, h=900) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (90, 120, 200)).save(buf, "JPEG")
    return buf.getvalue()


def test_a_side_column_bleeds_to_the_edge_and_bottom_but_never_over_the_title():
    prs, s = _deck(title_w=12.0)  # a title spanning the slide: the picture starts under it
    pl, fr = _frame(prs, s)
    assert fr.kind == "bleed" and fr.edges == ("right", "bottom")
    assert fr.box[0] == pl.box[0] and fr.box[0] + fr.box[2] == W and fr.box[1] == pl.box[1] and fr.box[1] + fr.box[3] == H
    # a short title leaves the column free up to the top: the picture is the right half of the slide
    prs, s = _deck(title_w=5.0)
    pl, fr = _frame(prs, s)
    assert fr.kind == "bleed" and set(fr.edges) == {"right", "top", "bottom"} and fr.box[1] == 0 and fr.composition == "tall"


def test_the_title_box_counts_by_its_words_not_its_width():
    def sized(s):
        t = s.shapes[0]
        t.width = Inches(12.4)
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size = Pt(28)
    prs, s = _deck(title_w=6.0, extra=sized)
    _, fr = _frame(prs, s)
    assert "top" in fr.edges  # «Головной офис» at 28 pt ends far left of the picture's column


def test_no_bleed_over_text_or_a_logo_between_the_place_and_the_edge():
    def logo(s):
        s.shapes.add_picture(io.BytesIO(_jpeg(200, 200)), Inches(12.85), Inches(3), Inches(0.4), Inches(0.4))
    prs, s = _deck(place=(6.6, 1.6, 5.4, 5.0), extra=logo)
    _, fr = _frame(prs, s)
    assert fr.kind == "plain" and fr.box == places_in(prs, None)[0].box


def test_a_page_number_the_bleed_covers_is_put_over_the_picture():
    prs, s = _deck(number=True)
    pl, fr = _frame(prs, s)
    assert fr.kind == "bleed" and fr.raise_ids
    put_picture(s, pl, _jpeg(), descr="x", frame=fr)
    tree = s._element.cSld.find(q("p:spTree"))
    ids = [PL._id(e) for e in tree]
    assert ids.index(fr.raise_ids[0]) > ids.index(pl.shape_id)


def test_a_deck_of_circles_gets_a_round_picture_and_rounded_cards_a_rounded_one():
    prs, s = _deck(place=(4, 1.5, 5, 4.5), circles=3)  # a place in the middle of the slide: no bleed
    pl, fr = _frame(prs, s)
    assert fr.kind == "circle" and fr.geom == "ellipse" and fr.box[2] == fr.box[3] == min(pl.box[2], pl.box[3])
    assert fr.composition == "circle"
    prs, s = _deck(place=(4, 1.5, 5, 4.5), rounded=4)
    _, fr = _frame(prs, s)
    assert fr.kind == "rounded" and fr.geom == "roundRect" and 0 < fr.adj <= 50000
    put_picture(s, places_in(prs, None)[0], _jpeg(), descr="x", frame=fr)
    pic = next(e for e in s._element.cSld.find(q("p:spTree")) if etree.QName(e).localname == "pic")
    g = pic.find(".//" + q("a:prstGeom"))
    assert g.get("prst") == "roundRect" and g.find(q("a:avLst") + "/" + q("a:gd")) is not None


def test_a_form_the_template_drew_stays():
    prs, s = _deck(place=(4, 1.5, 5, 4.5), geom=MSO_SHAPE.HEXAGON, circles=3)
    _, fr = _frame(prs, s)
    assert fr.kind == "template" and fr.geom is None


def test_a_place_over_most_of_the_slide_is_its_background_with_an_underlay_under_the_text():
    prs, s = _deck(place=(0.3, 0.3, 12.7, 6.9))
    pl, fr = _frame(prs, s)
    assert fr.kind == "background" and fr.box == (0, 0, W, H) and fr.to_back and fr.scrim is not None
    assert fr.scrim_color == "FFFFFF"  # the slide's ground, not the free place's fill
    kept = put_picture(s, pl, _jpeg(), descr="x", frame=fr)
    tree = s._element.cSld.find(q("p:spTree"))
    kids = [e for e in tree if etree.QName(e).localname in ("sp", "pic")]
    assert etree.QName(kids[0]).localname == "pic" and PL._name(kids[1]) == "Подложка под текст"
    alpha = kids[1].find(".//" + q("a:alpha"))
    assert alpha is not None and 50000 <= int(alpha.get("val")) < 100000
    # «Убрать изображение»: the place comes back, the underlay goes
    assert take_picture_out(s, pl.shape_id, kept)
    assert not any(PL._name(e) == "Подложка под текст" for e in tree)
    assert any(PL._name(e) == PHOTO_PLACE_NAME for e in tree)


def test_a_sample_photo_the_place_already_covers_does_not_stop_the_bleed():
    def under(s):
        pic = s.shapes.add_picture(io.BytesIO(_jpeg()), Inches(7.2), Inches(1.6), Inches(5.4), Inches(5.0))
        tree = s._element.cSld.find(q("p:spTree"))
        tree.remove(pic._element)
        tree.insert(2, pic._element)  # under the place
    prs, s = _deck(extra=under)
    _, fr = _frame(prs, s)
    assert fr.kind == "bleed"


def test_the_crop_keeps_the_subject_not_the_middle():
    img = Image.new("RGB", (1600, 800), (200, 205, 210))
    px = np.asarray(img).copy()
    px[300:600, 100:400] = (20, 40, 160)  # the subject at the left
    img = Image.fromarray(px)
    l, t, r, b = P.focus_crop(img, 1.0)
    assert (r - l) == 800 and l <= 100 and r >= 400
    out = P.fit_to(img, 1, 1, focus=True)
    a = np.asarray(out)
    assert (a[..., 2] > 140).mean() > 0.05  # the blue subject is in the crop
    assert P.fit_to(img, 1, 1, focus=False).size == (800, 800)


def test_the_picture_is_asked_in_its_frames_composition_and_the_templates_style():
    pal = P.Palette(dark="0B1F44", light="F4F6FA", accents=["FF6A00"])
    tall, _ = build_prompt(ImageAsk(hint="Стеклянный фасад офиса", frame="tall"), pal)
    assert "Vertical portrait composition" in tall
    circle, _ = build_prompt(ImageAsk(hint="Стеклянный фасад офиса", frame="circle"), pal)
    assert "circular crop" in circle
    plain, _ = build_prompt(ImageAsk(hint="Стеклянный фасад офиса"), pal)
    assert "composition" not in plain.lower() or "uncluttered composition" in plain
    # the author's words leave the style open: the template's own pictures set it
    _, style = build_prompt(ImageAsk(hint="Процесс обучения", template_style="illustration"), pal)
    assert style == "illustration"
    _, style = build_prompt(ImageAsk(hint="Фото команды", template_style="illustration"), pal)
    assert style == "photo"

    class A:
        def __init__(self, kind):
            self.kind, self.width, self.height = kind, 1000, 800

    class M:
        assets = [A("illustration")] * 10 + [A("photo")]

    assert template_style(M()) == "illustration"
    M.assets = [A("photo")] * 3 + [A("illustration")] * 4
    assert template_style(M()) == "photo"
    M.assets = []
    assert template_style(M()) == ""


def test_a_picture_only_beside_a_short_text():
    from verstka.planning.pictures import room_for_picture

    assert room_for_picture("Короткий лид.", ["раз", "два", "три"])
    assert not room_for_picture("Лид.", ["1", "2", "3", "4", "5"])
    assert not room_for_picture("Лид.", ["пункт"], ["1 — a", "2 — b", "3 — c"])
    assert not room_for_picture("Лид.", other=True)
    assert not room_for_picture("x" * 500)


def test_the_writers_note_keeps_its_description():
    """The writer's note glued to a list was read as the list's line: the place «Фото» lost what it shows and the
    picture was drawn from the headline (the owner's run of 17.10: a glass trophy)."""
    from verstka.planning.brief_structure import read_structure
    from verstka.planning.faithful import read_parts
    from verstka.planning.topic_writer import TopicSlide, render_slide

    lines = render_slide(6, TopicSlide(title="Бренд", lead="Самый узнаваемый бренд.", points=["Первое место."], picture="Стеклянный фасад офиса банка"), "Бренд")
    text = "Слайд 1. Титульный\nНазвание: «Банк».\n" + "\n".join(lines) + "\n"
    spec = next(sp for sp in read_structure(text).specs if sp.photo)
    img = [p for p in read_parts(spec.text) if p.kind == "image"]
    assert img and img[0].text.startswith("Стеклянный фасад")


def test_a_rendered_place_gets_its_frame_through_plan_and_apply(tmp_path):
    from verstka.images import providers as IP
    from verstka.images.fill import apply_pictures, plan_pictures
    from verstka.images.service import ImageMaker

    prs, s = _deck(title_w=5.0)
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    outline = DeckOutline(title="Банк", slides=[OutlineSlide(id="a", kind="image_text", headline="Головной офис",
                                                             content=SlideContent(bullets=["x"], photo_slot="Изображение", photo_hint="Стеклянный фасад офиса"))])
    mock = IP.MockImages()
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "cache", deck_title="Банк")
    pend = plan_pictures(path, outline, maker)
    assert pend[0].frame.kind == "bleed" and pend[0].aspect == "3:4" and "Vertical portrait" in mock_prompt(pend)
    res = apply_pictures(path, pend, wait_s=10)
    maker.close()
    rec = res.records[0]
    assert rec["status"] == "generated" and rec["frame"]["kind"] == "bleed" and rec["box"] == list(pend[0].frame.box)


def mock_prompt(pend) -> str:
    return pend[0].future.result(timeout=10).prompt


def test_a_note_the_designer_copied_into_the_slide_leaves_its_text():
    from verstka.planning.pictures import strip_picture_notes

    o = DeckOutline(title="Банк", slides=[OutlineSlide(id="s8", kind="bullets", headline="Головной офис", content=SlideContent(
        bullets=["Офис на улице Вавилова", "[Место под изображение: стеклянный фасад офиса]"], photo_slot="Фото", photo_hint="Стеклянный фасад офиса"))])
    warns: list = []
    strip_picture_notes(o, warns)
    assert o.slides[0].content.bullets == ["Офис на улице Вавилова"] and o.slides[0].content.photo_hint and warns


def test_a_content_slides_picture_is_a_full_height_strip_away_from_the_layouts_logo():
    """The owner (18.10, after 88ef1e9): «картинка — от самого верха до самого низа, а с другой стороны текст». The
    strip takes ~44 % of the slide at the right edge — the left one for the compact variant — unless the layout draws
    a logo there; the text's columns stay beside it."""
    prs, s = _deck()
    sw = int(PL.STRIP_SHARE * W)
    assert PL.picture_strip(s, W, H) == (W - sw, 0, sw, H)
    assert PL.picture_strip(s, W, H, left=True) == (0, 0, sw, H)
    # a logo the layout draws at the bottom left: the compact variant's strip goes right
    tmp = s.shapes.add_picture(io.BytesIO(_jpeg(200, 200)), Inches(0.4), Inches(6.8), Inches(1), Inches(0.4))
    el = tmp._element
    el.getparent().remove(el)
    s.slide_layout._element.cSld.find(q("p:spTree")).append(el)
    assert PL.picture_strip(s, W, H, left=True) == (W - sw, 0, sw, H)
    x, w = PL.beside_strip((W - sw, 0, sw, H), Inches(0.5), W - Inches(1), W)
    assert x == Inches(0.5) and x + w <= W - sw - int(0.03 * W)
    x, w = PL.beside_strip((0, 0, sw, H), Inches(0.5), W - Inches(1), W)
    assert x >= sw + int(0.03 * W) and x + w == W - Inches(0.5)


@pytest.fixture
def env_deck(simple_deck, tmp_path):
    from verstka.analysis.manifest import analyze_template
    from verstka.ingest.workspace import TemplateWorkspace

    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    return simple_deck, manifest, TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")


def _picture_deck(heads: dict[int, str], n: int = 10, kinds: dict[int, str] | None = None):
    from verstka.schemas.common import PatternKind as K

    sl = [OutlineSlide(id="t", kind=K.title, headline="Склад в Подольске")]
    for k in range(2, n + 1):
        kind = (kinds or {}).get(k, "bullets")
        c = SlideContent(bullets=["Короткий пункт.", "Ещё пункт."]) if kind != "table" else SlideContent(table={"columns": ["a", "b"], "rows": [["1", "2"]]})
        sl.append(OutlineSlide(id=f"s{k}", kind=K(kind), headline=heads.get(k, f"Итоги квартала {k}"), content=c, spec_ref=k))
    return DeckOutline(title="Склад", slides=sl)


def test_pictures_go_only_where_the_slide_shows_a_thing_one_can_see():
    """The jury of 20.10 (and VK's experts): a cloud for «облачные сервисы», an arrow «рост» on a slide about rising
    costs, an envelope, a notebook — pictures that say nothing. A picture is added only on a short slide whose heading
    names a thing one can see (a place, a device, a process at work), the same slide in every variant, described by
    that thing and the heading; a deck of money and plans keeps none (never forced to two)."""
    from verstka.planning.pictures import AUTO_LABEL, concrete_subject, ensure_pictures

    assert concrete_subject("Перевод на сканеры и 6 кладовщиков стоят 2,9 млн ₽ в год") == ("сканер", "сканеры")
    assert concrete_subject("Ошибки комплектации 0,8% превышают цель")[0] == "склад"
    for abstract in ("Расходы на облачные сервисы выросли на 12%", "Выручка растёт быстрее плана", "Два правила для сложных случаев данных",
                     "Глубина 1642 метров", "Q3: квартал закрыт с прибылью", "Мониторинг и машинное обучение", "Письмо клиентам и отчёт"):
        assert concrete_subject(abstract) is None, abstract
    heads = {3: "Перевод на сканеры окупится за год", 7: "Новый склад в Казани принимает 3 000 заказов"}
    a, b = _picture_deck(heads), _picture_deck(heads)
    said = ensure_pictures([a, b])
    got = [(i, s.content.photo_hint) for i, s in enumerate(a.slides) if s.content.photo_slot]
    assert [i for i, _ in got] == [2, 6] and len(said) == 2
    assert got[0][1] == "Сканер — Перевод на сканеры окупится за год" and got[1][1].startswith("Склад — ")
    assert [(i, s.content.photo_hint) for i, s in enumerate(b.slides) if s.content.photo_slot] == got
    assert all(a.slides[i].content.photo_slot == AUTO_LABEL and not a.slides[i].content.photo_own for i, _ in got)
    assert ensure_pictures([a]) == []  # enough already
    # money and plans only: no picture, and the run says why
    plain = _picture_deck({})
    said = ensure_pictures([plain])
    assert not any(s.content.photo_slot for s in plain.slides) and said and said[0].startswith("no picture added")


def test_one_thing_is_pictured_once_and_never_beside_steps_or_three_cards():
    """One arrow on two slides (the jury's R05 s2 and s10): a thing the deck already shows is not pictured again (the
    author's own picture of it included); steps, a timeline, two columns or three cards need the slide's width (R04 s15
    «генерац/ии», R07 s7 «Канареечн/ая») and get no picture."""
    from verstka.planning.pictures import ensure_pictures
    from verstka.schemas.outline import SlideItem

    heads = {2: "Склад в Казани", 4: "Склад в Подольске", 6: "Смартфон кладовщика", 8: "Экран приложения курьера", 10: "Сканеры на складе"}
    o = _picture_deck(heads, n=12, kinds={10: "process"})
    ensure_pictures([o])
    things = [s.headline for s in o.slides if s.content.photo_slot]
    assert things == ["Склад в Казани", "Смартфон кладовщика"]  # one warehouse, one phone; the process slide never
    o = _picture_deck({3: "Сканеры на складе"}, n=6)
    o.slides[2].content.items = [SlideItem(title=t, text="коротко") for t in ("Тест", "Трафик", "Учения")]
    o.slides[2].content.bullets = []
    assert not ensure_pictures([o])[0].startswith("picture added") and not o.slides[2].content.photo_slot
    o = _picture_deck({3: "Сканеры на складе", 5: "Новые сканеры в зоне приёмки"}, n=8)
    o.slides[2].content.photo_slot, o.slides[2].content.photo_hint = "Изображение", "сканер в руке кладовщика"
    ensure_pictures([o])
    assert not o.slides[4].content.photo_slot  # the author already shows a scanner


def test_an_added_picture_leaves_a_slide_whose_words_it_would_break(env_deck):
    """The jury of 20.10: the picture's strip narrowed the column and a long word broke («генерац/ии»). The renderer
    lays such a slide out again without the added picture (an author's picture stays); a heading under its label stops
    before the strip."""
    from verstka.planning.pictures import AUTO_LABEL
    from verstka.rendering.deck import DeckBuilder
    from verstka.rendering.synth import render_synth
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.layout import LayoutSlide
    from verstka.schemas.outline import SlideItem

    deck, manifest, ws = env_deck

    def render(oslide):
        outline = DeckOutline(title="T", strategy="structured", slides=[oslide])
        b = DeckBuilder(deck)
        slide, warns = render_synth(b, LayoutSlide(outline_id=oslide.id, mode="synth", composition="cards"), oslide, manifest, ws, outline)
        return b, slide, warns

    long = "Сверхвысокопроизводительная"
    items = [SlideItem(title=f"{long} {k}", text=f"{long}") for k in range(4)]
    o = OutlineSlide(id="c", kind=K.cards, headline="Склад и сканеры", content=SlideContent(items=items, photo_slot=AUTO_LABEL, photo_hint="Склад — Склад и сканеры"))
    b, slide, warns = render(o)
    assert any("the added picture left out" in w for w in warns)
    assert o.content.photo_slot is None and len(b.created) == 1
    assert not [sh for sh in slide.shapes if sh.name == PHOTO_PLACE_NAME]
    # a short slide keeps its picture, its text clear of the strip
    o = OutlineSlide(id="d", kind=K.bullets, headline="Склад в Казани", content=SlideContent(bullets=["Две смены.", "Сорок человек."], photo_slot=AUTO_LABEL, photo_hint="Склад — Склад в Казани"))
    b, slide, warns = render(o)
    places = [sh for sh in slide.shapes if sh.name == PHOTO_PLACE_NAME]
    assert places and not any("left out" in w for w in warns)
    assert PL.crowded_beside(slide, (places[0].left, places[0].top, places[0].width, places[0].height), b.slide_w, b.slide_h) is None


def test_crowded_beside_sees_a_broken_word_and_a_heading_under_the_strip():
    prs, s = _deck()
    sw = int(PL.STRIP_SHARE * W)
    strip = (W - sw, 0, sw, H)
    # «_deck»'s place stands where the strip would: ours, never counted
    assert PL.crowded_beside(s, strip, W, H) is None
    t = s.shapes.add_textbox(Inches(0.6), Inches(5.6), Inches(1.2), Inches(1))
    t.text_frame.text = "Канареечная"
    t.text_frame.word_wrap = True
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    assert "Канареечная" in PL.crowded_beside(s, strip, W, H)
    t._element.getparent().remove(t._element)
    h = s.shapes.add_textbox(Inches(0.6), Inches(0.2), Inches(12), Inches(0.6))
    h.text_frame.word_wrap = True
    h.text_frame.text = "Два правила для сложных случаев данных и ещё несколько слов для длины"
    h.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    assert "under the picture" in PL.crowded_beside(s, strip, W, H)


def test_a_place_verstka_chose_is_never_drawn_from_its_heading_alone(tmp_path):
    from verstka.images import providers as IP
    from verstka.images.prompt import ImageAsk as Ask
    from verstka.images.service import ImageMaker
    from verstka.planning.pictures import AUTO_LABEL

    mock = IP.MockImages()
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path, deck_title="Франция")
    try:
        maker.make(Ask(hint="Капетинги и формирование королевства", label=AUTO_LABEL), "3:4")
        raised = None
    except Exception as e:  # noqa: BLE001
        raised = e
    maker.close()
    assert raised is not None and "описание" in getattr(raised, "why", "") and not mock.calls
