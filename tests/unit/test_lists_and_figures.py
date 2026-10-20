"""Regression tests (the live QA of 2026-10-07, Кофейня and VK WorkSpace): columns with their own lists never run together in one box of a sample; the figures of a card slide always have a place; the autofix never cuts a list; a chart in thousands is one quantity."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

W, H = 12192000, 6858000


def _box(slide, x, y, w, h, text, size, bold=False):
    tb = slide.shapes.add_textbox(Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h)))
    tb.text_frame.word_wrap = True
    lines = text.split("\n")
    tb.text_frame.text = lines[0]
    for ln in lines[1:]:
        tb.text_frame.add_paragraph().text = ln
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
    return tb


def _place(ph, x, y, w, h, size, bold=False):
    ph.left, ph.top, ph.width, ph.height = Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h))
    for p in ph.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold


def _two_column_template(out: Path) -> Path:
    """A «Comparison» sample as the coffee template has it (its slide 4): the heading at the left, a heading
    placeholder over a list placeholder in each of two columns, the left list with its sample lines, the right one
    empty and its heading of two lines — no repeat group, four standalone text slots."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Название презентации"
    s1.placeholders[1].text = "Подзаголовок"
    s2 = prs.slides.add_slide(prs.slide_layouts[4])  # Comparison
    s2.shapes.title.text = "Заголовок слайда"
    s2.placeholders[1].text = "Вариант Б"
    s2.placeholders[2].text_frame.text = "Первый тезис слева"
    s2.placeholders[2].text_frame.add_paragraph().text = "Второй тезис с деталью"
    s2.placeholders[3].text_frame.text = "Тезис справа"
    s2.placeholders[3].text_frame.add_paragraph().text = "Ещё один факт"
    _place(s2.shapes.title, 0.008, 0.179, 0.27, 0.383, 30)
    _place(s2.placeholders[1], 0.337, 0.18, 0.317, 0.078, 14, bold=True)
    _place(s2.placeholders[2], 0.337, 0.256, 0.317, 0.715, 13)
    _place(s2.placeholders[3], 0.666, 0.18, 0.317, 0.078, 14, bold=True)
    _place(s2.placeholders[4], 0.667, 0.256, 0.317, 0.715, 13)
    prs.save(out)
    return out


def test_column_stacks_pairs_each_heading_with_the_list_under_it():
    from verstka.matching.scorer import column_stacks
    from verstka.schemas.common import BboxFrac, SlotRole
    from verstka.schemas.template import Capacity, Slot, SlotStyle

    def sl(i, role, x, y, w, h, lines):
        return Slot(id=i, role=role, shape_id=i, bbox=BboxFrac(x=x, y=y, w=w, h=h), style=SlotStyle(), capacity=Capacity(max_chars=80 * lines, max_lines=lines))

    right_head, left_head = sl("6", SlotRole.body, 0.666, 0.18, 0.317, 0.078, 2), sl("4", SlotRole.body, 0.337, 0.18, 0.317, 0.078, 2)
    left, right = sl("5", SlotRole.bullet_list, 0.337, 0.256, 0.317, 0.715, 24), sl("7", SlotRole.bullet_list, 0.667, 0.256, 0.317, 0.715, 24)
    got = column_stacks([right_head, left_head, left, right])
    assert [(h.shape_id, b.shape_id) for h, b in got] == [("4", "5"), ("6", "7")]
    # a list beside a note set lower and narrower (coffee p6) is no pair of columns
    note = sl("8", SlotRole.bullet_list, 0.008, 0.51, 0.319, 0.47, 12)
    big = sl("9", SlotRole.bullet_list, 0.425, 0.19, 0.524, 0.78, 26)
    assert column_stacks([note, big]) == []


def test_columns_with_lists_never_run_together_in_one_box_of_a_sample(tmp_path):
    """Coffee visual/compact s10: two columns (6 + 3 lines) on the coffee template's two-column sample were set as one
    list of nine in its right box, the left column empty (too_many_bullets); three columns of one line each in the
    same box covered 14% of the slide (fill_ratio). Each column takes the sample's column, else the slide is composed."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent, SlideItem

    tpl = _two_column_template(tmp_path / "tpl.pptx")
    two = OutlineSlide(id="s2", kind=K.two_column, headline="Пять показателей и три меры", content=SlideContent(columns=[
        SlideItem(title="Показатели", bullets=["Покупки в день", "Средний чек", "Доля чеков с едой", "Доля переменных расходов", "Операционная прибыль"]),
        SlideItem(title="Меры", bullets=["Тестировать акции небольшими запусками", "Сравнивать поставщиков", "Корректировать графики смен"]),
    ]))
    three = OutlineSlide(id="s3", kind=K.two_column, headline="Три барьера: чек, трафик, потери", content=SlideContent(columns=[
        SlideItem(title="Низкий средний чек", bullets=["Только 20% чеков содержат еду"]),
        SlideItem(title="Простой в дневные часы", bullets=["После 15:00 занято 4 из 18 мест"]),
        SlideItem(title="Потери продуктов", bullets=["Списания — 27 000 ₽ в месяц"]),
    ]))
    o = DeckOutline(title="Т", slides=[OutlineSlide(id="s1", kind=K.title, headline="Т"), two, three])
    res = generate_variants(tpl, outline=o, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render_images=False, audit=True, autofix=False)
    v = res.variants[0]
    prs = Presentation(str(v.out_dir / "deck.pptx"))
    for idx, cols in ((1, two.content.columns), (2, three.content.columns)):
        boxes = [sh.text_frame.text for sh in prs.slides[idx].shapes if sh.has_text_frame]
        # no text box carries the lines of two columns
        for b in boxes:
            assert sum(1 for col in cols if any(line in b for line in col.bullets)) <= 1, (idx, b)
    assert not [i for i in v.audit.issues if i.check_id in ("too_many_bullets", "fill_ratio") and i.severity != "info"]


def test_the_autofix_never_cuts_a_list():
    """Coffee s10: «condense_text» cut a column of eight lines to six — two of the author's risks left the slide and
    the fix reported «skip»."""
    from verstka.audit.autofix import _condense
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.outline import DeckOutline, OutlineSlide, SlideContent, SlideItem

    lines = ["Покупки в день", "Средний чек", "Доля чеков с едой", "Доля переменных расходов", "Операционная прибыль", "Слабый отклик на предложения", "Рост закупочных цен", "Перегрузка сотрудников"]
    s = OutlineSlide(id="s", kind=K.two_column, headline="Х", content=SlideContent(bullets=list(lines), columns=[SlideItem(title="А", bullets=list(lines)), SlideItem(title="Б", bullets=["x y"])]))
    o = DeckOutline(title="T", slides=[s])
    _condense(o, "s", None, None)
    assert o.slides[0].content.columns[0].bullets == lines
    assert o.slides[0].content.bullets == lines


def test_a_form_switch_never_leaves_a_slides_figures_without_a_place():
    """Coffee visual s6: the compiler's variety pass made a row of figures (100 → 115, 148 500 ₽) with its lines into
    cards — the cards drew no figures and «Дневные предложения с 15:00 до 18:00» was cut at «15:»."""
    from verstka.planning.compile import _convert, _item_of
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.outline import NumberCallout, SlideContent

    c = SlideContent(bullets=["Партнерства с пятью ближайшими офисами", "Дневные предложения с 15:00 до 18:00"], paragraphs=["Бюджет вырастет с 20 000 до 35 000 ₽ в месяц"],
                     numbers=[NumberCallout(value="100 → 115", label="Покупок в день"), NumberCallout(value="148 500 ₽", label="Выручка за 30 дней")])
    assert _convert(c, K.stat_row, K.cards) is None
    it = _item_of("Дневные предложения с 15:00 до 18:00")
    assert (it.title, it.text) == ("Дневные предложения с 15:00 до 18:00", "")
    it = _item_of("Витрина для десертов — 70 000 рублей")
    assert (it.title, it.text) == ("Витрина для десертов", "70 000 рублей")


def test_an_item_slides_figures_are_drawn_by_the_composer(tmp_path, simple_deck):
    """Whatever form the plan gives (an edited plan, a reshape): cards with key figures show the figures."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem

    s = OutlineSlide(id="s2", kind=K.cards, headline="Четыре канала дадут 15 покупок в день", content=SlideContent(
        items=[SlideItem(title=t) for t in ("Партнерства с офисами", "Программа лояльности", "Районные сообщества", "Дневные предложения")],
        numbers=[NumberCallout(value="100 → 115", label="Покупок в день"), NumberCallout(value="148 500 ₽", label="Выручка за 30 дней")]))
    o = DeckOutline(title="Т", slides=[OutlineSlide(id="s1", kind=K.title, headline="Т"), s])
    res = generate_variants(simple_deck, outline=o, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render_images=False, audit=True, autofix=False)
    v = res.variants[0]
    assert not [i for i in v.audit.issues if i.check_id == "content_missing"]


def test_chart_mixed_units_reads_a_chart_in_thousands_as_thousands():
    """Workspace s8: a bar chart of one-time investments in «тыс. ₽» (70, 35, 25, 20, 30) was read as «35% выручки»,
    «25%», «20% чеков», «30 рабочих дней» — numbers of other sentences that share no word with the bars."""
    from verstka.audit.registry import AuditContext, run_checks
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.deck_ir import DeckIR, IRSlide
    from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, OutlineSlide, SlideContent
    from verstka.schemas.template import SlideSize, TemplateManifest, Tokens

    brief = Path(__file__).resolve().parents[2] / "tests/fixtures/briefs/coffee_long.md"
    cats = ["Витрина для десертов", "Настройка программы лояльности и учета", "Обновление меню, фотографии и оформление", "Обучение сотрудников", "Резерв"]
    s = OutlineSlide(id="s8", kind=K.chart, headline="Бюджет запуска — 180 000 ₽", content=SlideContent(chart=ChartSpec(type="bar", unit="тыс. ₽", categories=cats, series=[InlineSeries(name="Сумма", values=[70, 35, 25, 20, 30])])))
    o = DeckOutline(title="T", slides=[OutlineSlide(id="c", kind=K.title, headline="T"), s])
    ir = DeckIR(source="x.pptx", slide_w=W, slide_h=H, slides=[IRSlide(index=i + 1, outline_id=x.id, notes=x.notes) for i, x in enumerate(o.slides)])
    man = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(), patterns=[], n_slides=1)
    got, _ = run_checks(AuditContext(ir=ir, manifest=man, outline=o, brief_text=brief.read_text(encoding="utf-8")), only={"chart_mixed_units"})
    assert [i for i in got if i.severity != "info"] == []


# ---------------------------------------------------------------------------------------------- the review of 3.0 (R1, R2)


def _cards_and_list_template(out: Path) -> Path:
    """A sample of three cards in a row (a heading over a text in each) with a free list box under them."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Название презентации"
    s1.placeholders[1].text = "Подзаголовок"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Заголовок слайда"
    for i in range(3):
        x, y, w, h = 0.06 + i * 0.30, 0.25, 0.26, 0.36
        rect = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(W * x)), Emu(int(H * y)), Emu(int(W * w)), Emu(int(H * h)))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(0xED, 0xF3, 0xFC)
        rect.line.fill.background()
        _box(s2, x + 0.02, y + 0.03, w - 0.04, 0.08, f"Пункт {i + 1}", 20, bold=True)
        _box(s2, x + 0.02, y + 0.13, w - 0.04, 0.20, "Описание пункта в две строки текста", 14)
    _box(s2, 0.06, 0.66, 0.86, 0.26, "Первый вывод под карточками\nВторой вывод под карточками\nТретий вывод", 18)
    prs.save(out)
    return out


def _sizes(shape) -> set[float]:
    return {r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size is not None and r.text.strip()}


def test_a_freeform_slide_keeps_its_items_beside_its_figures(tmp_path, simple_deck):
    """Stress s11 (15 of 24 decks): a freeform slide with items and key figures kept only «+112% — Рост прибыли» — the
    guard for figures beside an item form's items took it before it was set as a list, and its items left the slide."""
    from verstka.pipeline.generate import generate_variants
    from verstka.rendering.fallbacks import prepare_slide
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.layout import LayoutSlide
    from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem

    s = OutlineSlide(id="s2", kind=K.freeform, headline="Ключевые решения и цифры", content=SlideContent(
        items=[SlideItem(title="Комбо", text="390 рублей за капучино и круассан"), SlideItem(title="Витрина", text="70 000 ₽ разовых вложений"),
               SlideItem(title="Лояльность", text="каждый шестой напиток в подарок")],
        numbers=[NumberCallout(value="70 000 ₽", label="Витрина"), NumberCallout(value="+112%", label="Рост прибыли")]))
    o = DeckOutline(title="Т", slides=[OutlineSlide(id="s1", kind=K.title, headline="Т"), s])
    out, ps, notes = prepare_slide(s, LayoutSlide(outline_id="s2", mode="clone", pattern_id="p2"), o)
    # every item is a line of the list; the figure an item gives («Витрина — 70 000 ₽ …») is not said twice
    assert out.kind == K.bullets and ps.mode == "synth" and ps.composition == "bullets", (out.kind, ps)
    assert out.content.bullets == ["+112% — Рост прибыли", "Комбо — 390 рублей за капучино и круассан", "Витрина — 70 000 ₽ разовых вложений",
                                   "Лояльность — каждый шестой напиток в подарок"], out.content.bullets
    assert not out.content.items and not out.content.numbers
    res = generate_variants(simple_deck, outline=o, strategies=["structured"], out_dir=tmp_path / "out", workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render_images=False, audit=True, autofix=False)
    v = res.variants[0]
    assert not [i for i in v.audit.issues if i.check_id == "content_missing"], [i.details for i in v.audit.issues if i.check_id == "content_missing"]
    shown = " ".join(sh.text_frame.text for sh in Presentation(str(v.out_dir / "deck.pptx")).slides[1].shapes if sh.has_text_frame).replace(" ", " ").replace("⁠", "")
    for t in ("Комбо", "Витрина", "Лояльность", "Рост прибыли"):
        assert t in shown, (t, shown)


def test_a_figure_a_card_gives_is_compared_as_a_quantity():
    """Кофейня s7: «70 000 ₽ — Витрина для десертов» stood as a line beside the card «Витрина для десертов — 70 000
    рублей»: the figure was looked up as characters. Quantities are compared — ₽ is рублей, a no-break space a space,
    «млн» a million, «процентов» %, a minus any dash — with the words of the figure's label; another fact that happens
    to have the same number keeps its line."""
    from verstka.rendering.fallbacks import figure_given, figures_beside_items
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem

    given = [
        (NumberCallout(value="70 000 ₽", label="Витрина для десертов"), "Витрина для десертов — 70 000 рублей"),
        (NumberCallout(value="70 000 ₽", label="Витрина"), "Витрина — 70 000 руб."),
        (NumberCallout(value="1,2 млн ₽", label="Выручка"), "Выручка — 1 200 000 рублей в месяц"),
        (NumberCallout(value="35%", label="Возвращающиеся гости"), "Возвращающихся гостей — 35 процентов"),
        (NumberCallout(value="−33%", label="Списания продуктов"), "Списания продуктов: -33 %"),
    ]
    for n, line in given:
        assert figure_given(n, [line]), (n.value, line)
    assert not figure_given(NumberCallout(value="30 000 ₽", label="Резерв"), ["Обучение сотрудников — 30 000 рублей"])
    assert not figure_given(NumberCallout(value="180 000 ₽", label="Все вложения"), ["Витрина для десертов — 70 000 рублей"])
    assert not figure_given(NumberCallout(value="70 000 %", label="Витрина"), ["Витрина — 70 000 рублей"])
    s = OutlineSlide(id="s7", kind=K.cards, headline="Вложения в запуск — 180 000 рублей", content=SlideContent(
        items=[SlideItem(title="Витрина для десертов", text="70 000 рублей"), SlideItem(title="Программа лояльности и учёт", text="35 000 рублей"),
               SlideItem(title="Меню и оформление", text="25 000 рублей"), SlideItem(title="Обучение сотрудников", text="20 000 рублей")],
        numbers=[NumberCallout(value="70 000 ₽", label="Витрина для десертов"), NumberCallout(value="180 000 ₽", label="Все вложения"), NumberCallout(value="30 000 ₽", label="Резерв")]))
    out, lines = figures_beside_items(s, DeckOutline(title="T", slides=[s]))
    assert lines == ["180 000 ₽ — Все вложения", "30 000 ₽ — Резерв"] and out.content.bullets == lines and not out.content.numbers
    assert out.kind == K.cards and len(out.content.items) == 4


def test_lines_beside_items_never_take_an_items_place_nor_a_larger_size(tmp_path):
    """Кофейня structured s7: on the template's two-column sample three figure lines took the main column at 21.5 pt and
    the four cards ran together in a 14 pt list beside them. The lines beside the items are secondary: never in a place
    the sample gives the items (its columns), never larger than the items' own text — the sample without another box
    for them leaves the slide to the composer, which sets them under the cards."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.common import PatternKind as K
    from verstka.schemas.outline import DeckOutline, NumberCallout, OutlineSlide, SlideContent, SlideItem

    cards = OutlineSlide(id="s2", kind=K.cards, headline="Вложения в запуск — 180 000 рублей", content=SlideContent(
        items=[SlideItem(title="Витрина для десертов", text="70 000 рублей"), SlideItem(title="Программа лояльности и учёт", text="35 000 рублей"),
               SlideItem(title="Меню и оформление", text="25 000 рублей"), SlideItem(title="Обучение сотрудников", text="20 000 рублей")],
        numbers=[NumberCallout(value="70 000 ₽", label="Витрина для десертов"), NumberCallout(value="180 000 ₽", label="Все вложения"), NumberCallout(value="30 000 ₽", label="Резерв")]))
    three = OutlineSlide(id="s3", kind=K.cards, headline="Три шага запуска", content=SlideContent(
        items=[SlideItem(title="Меню", text="Новые позиции и цены"), SlideItem(title="Команда", text="Обучение продажам"), SlideItem(title="Учёт", text="Списания каждый день")],
        bullets=["Бюджет запуска — 180 000 рублей", "Срок — шесть месяцев"]))
    for name, tpl, slide in (("two", _two_column_template(tmp_path / "two.pptx"), cards), ("cards", _cards_and_list_template(tmp_path / "cards.pptx"), three)):
        o = DeckOutline(title="Т", slides=[OutlineSlide(id="s1", kind=K.title, headline="Т"), slide])
        res = generate_variants(tpl, outline=o, strategies=["structured"], out_dir=tmp_path / f"out_{name}", workspace_root=tmp_path / f"ws_{name}", use_llm=False, use_vlm=False, render_images=False, audit=True, autofix=False)
        v = res.variants[0]
        assert not [i for i in v.audit.issues if i.check_id == "content_missing"], (name, [i.details for i in v.audit.issues if i.check_id == "content_missing"])
        prs = Presentation(str(v.out_dir / "deck.pptx"))
        boxes = [sh for sh in prs.slides[1].shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        norm = lambda t: t.replace(" ", " ").replace("⁠", "")  # noqa: E731
        lines = [sh for sh in boxes if any(b.split(" — ")[0] in norm(sh.text_frame.text) for b in v.outline.slides[1].content.bullets)]
        items = [sh for sh in boxes if any(it.text in norm(sh.text_frame.text) for it in slide.content.items) and sh not in lines]
        assert lines and items, (name, [norm(sh.text_frame.text) for sh in boxes])
        assert max(max(_sizes(sh)) for sh in lines) <= min(min(_sizes(sh)) for sh in items) + 0.05, (name, [(norm(sh.text_frame.text)[:30], _sizes(sh)) for sh in boxes])
        if name == "two":
            # the sample's two columns are the cards' places: the figure lines take neither — the slide is composed
            assert v.plan.for_outline("s2").mode == "synth", v.plan.for_outline("s2")


def test_samples_are_scored_on_the_figures_set_as_lines():
    """The scorer reads an item form's key figures as the renderer sets them (lines under the items): a sample is not
    marked down for number boxes the slide will not use, and one that keeps no box for the lines beside the items
    (they never take an item's place) does not win over the composer."""
    from verstka.matching.scorer import lines_beside_items, room_for_lines, score_pattern
    from verstka.planning.strategies import get_strategy
    from verstka.rendering.fallbacks import figures_beside_items
    from verstka.schemas.common import BboxFrac, Family, PatternKind as K, SlotRole
    from verstka.schemas.outline import NumberCallout, OutlineSlide, SlideContent, SlideItem
    from verstka.schemas.template import Capacity, Pattern, Slot, SlotStyle, SlideSize, TemplateManifest, Tokens

    def sl(i, role, x, y, w, h, lines, size=14.0):
        return Slot(id=i, role=role, shape_id=i, bbox=BboxFrac(x=x, y=y, w=w, h=h), style=SlotStyle(size_pt=size), capacity=Capacity(max_chars=60 * lines, max_lines=lines))

    title = sl("2", SlotRole.title, 0.008, 0.179, 0.27, 0.383, 5, 30.0)
    two = Pattern(id="p4", kind=K.two_column, family=Family.light, source_slide=4, slots=[title, sl("4", SlotRole.body, 0.337, 0.18, 0.317, 0.078, 2), sl("5", SlotRole.bullet_list, 0.337, 0.256, 0.317, 0.715, 24),
                                                                       sl("6", SlotRole.body, 0.666, 0.18, 0.317, 0.078, 2), sl("7", SlotRole.bullet_list, 0.667, 0.256, 0.317, 0.715, 24)])
    s = OutlineSlide(id="s7", kind=K.cards, headline="Вложения в запуск", content=SlideContent(
        items=[SlideItem(title=f"Статья {i}", text=f"{i}0 000 рублей") for i in range(1, 5)],
        numbers=[NumberCallout(value="180 000 ₽", label="Все вложения"), NumberCallout(value="30 000 ₽", label="Резерв")]))
    shown = figures_beside_items(s, None)[0]
    assert lines_beside_items(shown) and not room_for_lines(two, None)
    m = TemplateManifest(template_id="t", source_file="t.pptx", slide_size=SlideSize(w=W, h=H), tokens=Tokens(), patterns=[two], n_slides=1)
    st = get_strategy("structured")
    got, want = score_pattern(s, two, m, st), score_pattern(shown, two, m, st)
    assert (got.score, got.reasons) == (want.score, want.reasons)
    assert got.score < st.synth_threshold and "строкам рядом с пунктами нет места в образце" in got.reasons, got
