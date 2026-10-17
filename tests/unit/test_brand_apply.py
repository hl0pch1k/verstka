"""Брендбук → брендированная копия шаблона (verstka/brand/apply.py) и расхождения (verstka/brand/conflicts.py).

Без модели и сети: синтетический шаблон python-pptx, настоящий шаблон из фикстур LibreOffice, разбор VK Tech из фикстур
и пример брендбука `examples/brandbook-vk-cloud`."""

from __future__ import annotations

import gzip
import json
import random
import re
import zipfile
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

import verstka.brand.apply as ap
import verstka.brand.logos as lg
import verstka.brand.palette as pal
from verstka.schemas.common import contrast_ratio
from lxml import etree
from verstka.analysis.manifest import analyze_template
from verstka.brand.apply import apply_brand, brand_fingerprint, branded_info, branded_template
from verstka.brand.palette import ColorPlan, plan_colors, theme_info, tone
from verstka.brand.conflicts import brand_conflicts
from verstka.ingest.workspace import TemplateWorkspace
from verstka.schemas.brand import BrandColor, BrandFont, BrandLogo, BrandSpec
from verstka.schemas.template import TemplateManifest

ROOT = Path(__file__).resolve().parents[2]
FIX = ROOT / "tests" / "fixtures" / "tpl_audit"
BRAND_EXAMPLE = ROOT / "examples" / "brandbook-vk-cloud"
W, H = 12192000, 6858000

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ---------------------------------------------------------------------------------------------- материалы


def _logo_png(path: Path, badge, ink, size=(400, 100)) -> Path:
    """Логотип: знак-плашка и «надпись» полосой на прозрачном фоне."""
    im = Image.new("RGBA", size, (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    s = size[1]
    d.rounded_rectangle([0, 0, s - 1, s - 1], radius=s // 5, fill=badge)
    d.rectangle([int(s * 1.25), s // 3, size[0] - 8, 2 * s // 3], fill=ink)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def _photo(path: Path) -> Path:
    rnd = random.Random(7)
    im = Image.new("RGB", (320, 240))
    im.putdata([(rnd.randrange(256), rnd.randrange(80), rnd.randrange(80)) for _ in range(320 * 240)])
    im.save(path, quality=90)
    return path


def _brand_dir(tmp: Path) -> Path:
    d = tmp / "brandbook"
    _logo_png(d / "logos" / "primary.png", (0, 119, 255, 255), (11, 18, 32, 255))
    _logo_png(d / "logos" / "white.png", (255, 255, 255, 255), (255, 255, 255, 255))
    im = Image.new("RGBA", (128, 128), (0, 0, 0, 0))
    ImageDraw.Draw(im).rounded_rectangle([0, 0, 127, 127], radius=24, fill=(0, 119, 255, 255))
    im.save(d / "logos" / "mark.png")
    return d


def _brand(**kw) -> BrandSpec:
    data = dict(
        name="Облако",
        source="tokens",
        colors=[
            BrandColor(hex="#0077FF", role="primary", name="Синий"),
            BrandColor(hex="#2787F5", role="secondary"),
            BrandColor(hex="#0B1220", role="text"),
            BrandColor(hex="#FFFFFF", role="background"),
            BrandColor(hex="#FF3985", role="accent"),
            BrandColor(hex="#00C2FF", role="chart"),
        ],
        fonts=[
            BrandFont(role="heading", family="Brand Display", fallback="Play", weight=700),
            BrandFont(role="body", family="Brand Text", fallback="Play", weight=400),
        ],
        logos=[
            BrandLogo(name="Основной", file="logos/primary.png", variant="primary", ground="light"),
            BrandLogo(name="Белый", file="logos/white.png", variant="white", ground="dark"),
            BrandLogo(name="Знак", file="logos/mark.png", variant="mark", ground="any"),
        ],
        logo_place="bottom_right",
        margins_px=48,
    )
    data.update(kw)
    return BrandSpec(**data)


def _text(slide, text, x, y, w, h, font, size, color, bold=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.text_frame.text = text
    r = tb.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold = font, Pt(size), bold
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def _card(slide, x, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(int(H * 0.35)), Emu(int(W * 0.25)), Emu(int(H * 0.3)))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(color)
    sp.line.fill.background()
    return sp


def build_template(tmp: Path, *, logo: bool = True) -> Path:
    """Шаблон: светлый фон F7F7F7, текст 1F2937, акценты 1F6FEB и E8590C, светлая плашка DCE8FD, заголовки Georgia,
    текст Verdana, фото на втором слайде, диаграмма на третьем, логотип «знак + надпись» внизу слева."""
    logo_path = _logo_png(tmp / "src" / "acme-logo.png", (31, 111, 235, 255), (31, 41, 55, 255))
    photo = _photo(tmp / "src" / "photo.jpg")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for i in range(4):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("F7F7F7")
        _text(s, f"Итоги квартала {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
        _text(s, "Выручка выросла, клиентов стало больше", int(W * 0.05), int(H * 0.22), int(W * 0.8), int(H * 0.08), "Verdana", 14, "1F2937")
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.05)), Emu(int(H * 0.2)), Emu(int(W * 0.3)), Emu(int(H * 0.01)))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("1F6FEB")  # фирменная полоса шаблона: главный акцент
        bar.line.fill.background()
        if i == 1:
            s.shapes.add_picture(str(photo), Emu(int(W * 0.05)), Emu(int(H * 0.33)), Emu(int(W * 0.5)), Emu(int(H * 0.5))).name = "Фото"
        elif i == 2:
            cd = CategoryChartData()
            cd.categories = ["янв", "фев", "мар"]
            cd.add_series("Выручка", (1, 2, 3))
            gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(int(W * 0.05)), Emu(int(H * 0.33)), Emu(int(W * 0.6)), Emu(int(H * 0.5)), cd)
            ser = gf.chart.plots[0].series[0]
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGBColor.from_string("1F6FEB")
            gf.chart.font.name = "Verdana"
        else:
            _card(s, int(W * 0.05), "1F6FEB")
            _card(s, int(W * 0.36), "E8590C" if i == 0 else "1F6FEB")
            _card(s, int(W * 0.67), "DCE8FD")
            _text(s, "+23%", int(W * 0.06), int(H * 0.4), int(W * 0.2), int(H * 0.12), "Georgia", 40, "FFFFFF")
        if logo:
            pic = s.shapes.add_picture(str(logo_path), Emu(int(W * 0.04)), Emu(int(H * 0.9)), height=Emu(int(H * 0.05)))
            pic.name = "Логотип ACME"
    out = tmp / "src" / "acme.pptx"
    prs.save(out)
    return out


@pytest.fixture
def fonts_here(monkeypatch):
    """На «машине» есть шрифты бренда — в копию пишутся они, а не замена."""
    monkeypatch.setattr(ap, "font_available", lambda f: f in {"Brand Display", "Brand Text", "Play"})


@pytest.fixture
def fonts_missing(monkeypatch):
    """Шрифтов бренда нет — пишется замена из брендбука (Play)."""
    monkeypatch.setattr(ap, "font_available", lambda f: f == "Play")


def _analyze(path: Path, ws: Path) -> TemplateManifest:
    return analyze_template(path, workspace_root=ws, use_llm=False, use_vlm=False, render=False)


def _parts(pptx: Path, pattern: str) -> dict[str, str]:
    with zipfile.ZipFile(pptx) as z:
        return {n: z.read(n).decode("utf-8") for n in z.namelist() if re.match(pattern, n)}


def _all(pptx: Path, pattern: str) -> str:
    return "\n".join(_parts(pptx, pattern).values())


def _hexes(xml: str) -> set[str]:
    return set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml))


@pytest.fixture
def synthetic(tmp_path):
    src = build_template(tmp_path)
    m = _analyze(src, tmp_path / "ws")
    return src, m, _brand_dir(tmp_path)


# ---------------------------------------------------------------------------------------------- цвета


def test_color_math_tones():
    assert tone("0077FF") == "chromatic"
    assert tone("D6ECFF") == "tint"
    assert tone("000000") == tone("0B1220") == "dark"
    assert tone("FAFCFF") == "light"
    assert tone("8F8F8F") == "grey"


def test_synthetic_roles_as_expected(synthetic):
    """Предпосылка остальных тестов: разбор видит роли так, как задумано в шаблоне."""
    _, m, _ = synthetic
    roles = {t.hex: t.roles for t in m.tokens.colors}
    assert "background.light" in roles["F7F7F7"]
    assert "text.primary" in roles["1F2937"]
    assert any(r.startswith("accent.") for r in roles["1F6FEB"])
    assert any(r.startswith("accent.") for r in roles["E8590C"])


def test_explicit_colors_mapped_by_role(synthetic, fonts_here, tmp_path):
    src, m, bdir = synthetic
    out = tmp_path / "branded.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    slides = _all(out, r"ppt/slides/slide\d+\.xml$")
    hx = _hexes(slides)
    # главный акцент → основной цвет бренда, текст → текст бренда, фон → фон бренда
    assert {"0077FF", "0B1220", "FFFFFF"} <= hx
    assert not {"1F6FEB", "1F2937", "F7F7F7", "E8590C"} & hx
    assert res.colors["#1F6FEB"] == "#0077FF"
    assert res.colors["#1F2937"] == "#0B1220"
    assert res.colors["#F7F7F7"] == "#FFFFFF"
    # второй акцент (оранжевый): близкого тона у бренда нет (розовый «акцент» — только для близкого выделения) — по
    # порядку ролей, дополнительный цвет бренда
    assert res.colors["#E8590C"] == "#2787F5"
    # светлая плашка — тот же светлый оттенок, только в тоне бренда
    tint = res.colors.get("#DCE8FD", "#DCE8FD").lstrip("#")
    assert tone(tint) in ("tint", "light") and pal.hls(tint)[1] > 0.85
    # белые цифры на синей плашке остаются белыми: контраст не ломается
    assert "FFFFFF" in hx
    # диаграмма: заливка серии тоже в цветах бренда
    charts = _all(out, r"ppt/charts/chart\d+\.xml$")
    assert "0077FF" in _hexes(charts) and "1F6FEB" not in _hexes(charts)


def test_theme_colors_and_fonts_rewritten(synthetic, fonts_here, tmp_path):
    src, m, bdir = synthetic
    out = tmp_path / "branded.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    before = theme_info(etree.fromstring(_parts(src, r"ppt/theme/theme1\.xml$")["ppt/theme/theme1.xml"].encode()))
    after = theme_info(etree.fromstring(_parts(out, r"ppt/theme/theme1\.xml$")["ppt/theme/theme1.xml"].encode()))
    assert before.stock  # стандартная тема python-pptx: её акценты занимает бренд по порядку
    assert after.slots["accent1"] == "0077FF"
    assert after.slots["accent2"] == "C0504D"  # красный слот — цвет ошибки и падения: в брендбуке пары нет, остаётся
    assert after.slots["accent3"] == "2787F5"
    assert after.slots["dk1"] == "0B1220"
    assert after.major == "Brand Display" and after.minor == "Brand Text"


def test_fonts_rewritten_by_role(synthetic, fonts_here, tmp_path):
    src, m, bdir = synthetic
    out = tmp_path / "branded.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    slides = _all(out, r"ppt/slides/slide\d+\.xml$")
    assert "Georgia" not in slides and "Verdana" not in slides
    # заголовок (36 пт) — шрифтом заголовков, текст 14 пт — шрифтом текста; начертание — как у шаблона (полужирный
    # не навязывается: подбор кегля образцов шаблона не меняется)
    title = re.search(r'<a:rPr[^>]*sz="3600"[^>]*>.*?</a:rPr>', slides, re.S).group(0)
    assert 'typeface="Brand Display"' in title and 'b="1"' not in title
    body = re.search(r'<a:rPr[^>]*sz="1400"[^>]*>.*?</a:rPr>', slides, re.S).group(0)
    assert 'typeface="Brand Text"' in body and 'b="1"' not in body
    # мастера, макеты и диаграммы
    masters = _all(out, r"ppt/(slideMasters|slideLayouts)/.*\.xml$")
    assert re.search(r'typeface="Brand (Display|Text)"', masters) or "+mj-lt" in masters
    assert 'typeface="Brand Text"' in _all(out, r"ppt/charts/chart\d+\.xml$")
    assert res.fonts == {"heading": "Brand Display", "body": "Brand Text"}


def test_missing_brand_font_falls_back(synthetic, fonts_missing, tmp_path):
    """Шрифта бренда нет на машине — в копию пишется его замена из брендбука, как велит брендбук."""
    src, m, bdir = synthetic
    out = tmp_path / "branded.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    slides = _all(out, r"ppt/slides/slide\d+\.xml$")
    assert 'typeface="Play"' in slides and "Brand Display" not in slides
    assert res.fonts == {"heading": "Play", "body": "Play"}


def test_pictures_untouched(synthetic, fonts_here, tmp_path):
    src, m, bdir = synthetic
    out = tmp_path / "branded.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        photos = [n for n in a.namelist() if n.startswith("ppt/media/") and n.endswith((".jpg", ".jpeg"))]
        assert photos
        for n in photos:
            assert a.read(n) == b.read(n)  # пиксели фото те же
        s2a, s2b = a.read("ppt/slides/slide2.xml").decode(), b.read("ppt/slides/slide2.xml").decode()
    photo_pic = lambda x: re.search(r'<p:pic>(?:(?!</p:pic>).)*name="Фото".*?</p:pic>', x, re.S).group(0)  # noqa: E731
    pa, pb = photo_pic(s2a), photo_pic(s2b)
    assert re.search(r"<a:blip [^>]*>", pa).group(0) == re.search(r"<a:blip [^>]*>", pb).group(0)
    assert re.search(r"<a:xfrm>.*?</a:xfrm>", pa, re.S).group(0) == re.search(r"<a:xfrm>.*?</a:xfrm>", pb, re.S).group(0)


def test_blip_recolor_kept():
    """Цвета внутри картинки (перекраска, дуотон) — часть картинки: не трогаются."""
    xml = (
        f'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A[1:-1]}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:cSld><p:spTree>'
        '<p:pic><p:blipFill><a:blip r:embed="rId2"><a:duotone><a:srgbClr val="1F6FEB"/><a:srgbClr val="FFFFFF"/></a:duotone></a:blip></p:blipFill>'
        '<p:spPr><a:ln><a:solidFill><a:srgbClr val="1F6FEB"/></a:solidFill></a:ln></p:spPr></p:pic></p:spTree></p:cSld></p:sld>'
    )
    root = etree.fromstring(xml.encode())
    plan = ColorPlan(table={"1F6FEB": "0077FF"})
    from collections import Counter

    assert ap._recolor(root, plan, Counter())
    vals = [e.get("val") for e in root.iter(f"{A}srgbClr")]
    assert vals == ["1F6FEB", "FFFFFF", "0077FF"]


# ---------------------------------------------------------------------------------------------- логотипы


def test_logos_replaced_with_brand_logo(synthetic, fonts_here, tmp_path):
    """Логотип шаблона в том же углу, что велит брендбук (слева внизу): логотип бренда встаёт на его место."""
    src, m, bdir = synthetic
    out = tmp_path / "branded.pptx"
    res = apply_brand(src, out, _brand(logo_place="bottom_left"), manifest=m, brand_dir=bdir)
    assert res.logos_replaced == 4 and res.logos_added == 0
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        brand_media = [n for n in names if n.startswith("ppt/media/brand-logo-")]
        assert brand_media and all("primary" in n for n in brand_media)  # светлый фон → основной логотип
        assert z.read(brand_media[0]) == (bdir / "logos" / "primary.png").read_bytes()
        for i in range(1, 5):
            xml = z.read(f"ppt/slides/slide{i}.xml").decode()
            rels = z.read(f"ppt/slides/_rels/slide{i}.xml.rels").decode()
            logo = re.search(r'<p:pic>(?:(?!</p:pic>).)*name="Логотип Облако".*?</p:pic>', xml, re.S)
            assert logo, f"slide {i}: логотип бренда"
            rid = re.search(r'r:embed="([^"]+)"', logo.group(0)).group(1)
            assert re.search(rf'Id="{rid}"[^>]*Target="\.\./media/brand-logo-primary', rels) or re.search(
                rf'Target="\.\./media/brand-logo-primary[^"]*"[^>]*Id="{rid}"', rels
            )
            assert "Логотип ACME" not in xml
            # пропорции логотипа бренда и та же высота рамки
            ext = re.search(r'<a:ext cx="(\d+)" cy="(\d+)"', logo.group(0))
            cx, cy = int(ext.group(1)), int(ext.group(2))
            assert abs(cx / cy - 392 / 100) < 0.15
            assert abs(cy - H * 0.05) < H * 0.01
    # фото и диаграмма не задеты
    assert "Фото" in _parts(out, r"ppt/slides/slide2\.xml$")["ppt/slides/slide2.xml"]


def test_dark_ground_gets_white_logo(tmp_path, fonts_here):
    """Белая надпись логотипа шаблона — тёмный фон: ставится белый логотип бренда."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    white = _logo_png(tmp_path / "src" / "logo-white.png", (255, 255, 255, 255), (255, 255, 255, 255))
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("101828")
        _text(s, "Итоги", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "FFFFFF")
        s.shapes.add_picture(str(white), Emu(int(W * 0.8)), Emu(int(H * 0.05)), height=Emu(int(H * 0.06))).name = "Logo"
    src = tmp_path / "src" / "dark.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "dark-branded.pptx"
    res = apply_brand(src, out, _brand(logo_place="top_right"), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_replaced == 3
    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("ppt/media/brand-logo-")]
    assert media and all("white" in n for n in media)
    # тёмный фон шаблона брендбук со светлым фоном не переворачивает: та же светлота, оттенок — нейтральных бренда
    grounds = [h for h in _hexes(_all(out, r"ppt/slides/slide\d+\.xml$")) if h != "FFFFFF"]
    assert grounds and all(abs(pal.lightness(h) - pal.lightness("101828")) < 1.5 for h in grounds)


def test_no_template_logo_brand_rule_adds_logo(tmp_path, fonts_here):
    """Логотипа в шаблоне нет, в брендбуке есть правило места — логотип бренда встаёт на макеты содержательных слайдов
    в правый нижний угол с полями брендбука; без правила — ничего не добавляется."""
    src = build_template(tmp_path, logo=False)
    m = _analyze(src, tmp_path / "ws")
    bdir = _brand_dir(tmp_path)
    out = tmp_path / "with-logo.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    assert res.logos_added >= 1
    layouts = _parts(out, r"ppt/slideLayouts/slideLayout\d+\.xml$")
    hit = [x for x in layouts.values() if "Логотип Облако" in x]
    assert hit
    off = re.search(r'name="Логотип Облако".*?<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"', hit[0], re.S)
    x, y, cx, cy = (int(v) for v in off.groups())
    margin = 48 / 1920 * W
    assert abs(W - (x + cx) - margin) < W * 0.002 and abs(H - (y + cy) - margin) < H * 0.01
    out2 = tmp_path / "no-rule.pptx"
    res2 = apply_brand(src, out2, _brand(logo_place=None), manifest=m, brand_dir=bdir)
    assert res2.logos_added == 0
    assert "Логотип Облако" not in _all(out2, r"ppt/slideLayouts/.*\.xml$")


def test_added_logo_takes_one_fallback_corner(tmp_path, fonts_here):
    """Угол из брендбука занят содержимым у всех образцов — логотип встаёт в одно запасное место для всего шаблона
    (тот же край, другая сторона), ни с чем не сталкиваясь, и это сказано словами."""
    src = build_template(tmp_path, logo=False)
    prs = Presentation(str(src))
    for s in prs.slides:
        card = _card(s, int(W * 0.5), "1F6FEB")  # вся правая половина низа занята
        card.top, card.width, card.height = Emu(int(H * 0.62)), Emu(int(W * 0.48)), Emu(int(H * 0.36))
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "out.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1 and res.logo_covered == res.logo_samples
    assert res.logo_place == "fallback" and any(res.logo_where in w for w in res.warnings)
    assert any(k in res.logo_where for k in ("левом нижнем углу", "левой половине", "левым нижним углом")), res.logo_where
    boxes = _logo_boxes(out)
    assert {round(x / W, 2) for _, (x, _y, _cx, _cy) in boxes} and all(x < W * 0.5 for _, (x, _y, _cx, _cy) in boxes)
    assert _ir_collisions(out) == []


def test_added_logo_goes_over_background_pictures(tmp_path, fonts_here):
    """Образцы шаблона закрывают угол фоновой картинкой во весь слайд (логотип макета скрылся бы под ней) — логотип
    встаёт на сами образцы поверх фона."""
    src = build_template(tmp_path, logo=False)
    bg = tmp_path / "src" / "bg.png"
    Image.new("RGB", (160, 90), (240, 240, 240)).save(bg)
    prs = Presentation(str(src))
    for s in prs.slides:
        pic = s.shapes.add_picture(str(bg), 0, 0, Emu(W), Emu(H))
        s.shapes._spTree.remove(pic._element)
        s.shapes._spTree.insert(2, pic._element)  # фон — под всем
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "out.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1
    slides = _parts(out, r"ppt/slides/slide\d+\.xml$")
    assert any("Логотип Облако" in x for x in slides.values())
    # на макет образцов логотип не лёг (он оказался бы под фоновой картинкой слайдов)
    assert "Логотип Облако" not in _parts(out, r"ppt/slideLayouts/slideLayout7\.xml$")["ppt/slideLayouts/slideLayout7.xml"]


# ---------------------------------------------------------------------------------------------- приоритет


def test_priority_template_only_fills_gaps(synthetic, fonts_here, tmp_path):
    """«Шаблон важнее»: свои цвета, шрифты и логотип шаблона остаются; бренд занимает только пустые места —
    стандартную тему офисного пакета и шрифты по умолчанию."""
    src, m, bdir = synthetic
    out = tmp_path / "tpl.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, priority="template", brand_dir=bdir)
    slides = _all(out, r"ppt/slides/slide\d+\.xml$")
    hx = _hexes(slides)
    assert {"1F6FEB", "E8590C", "1F2937", "F7F7F7"} <= hx
    assert "Georgia" in slides and "Verdana" in slides
    assert 'b="1"' not in re.search(r'<a:rPr[^>]*sz="3600"[^>]*>', slides).group(0)
    assert res.logos_replaced == 0 and "Логотип ACME" in slides
    theme = theme_info(etree.fromstring(_parts(out, r"ppt/theme/theme1\.xml$")["ppt/theme/theme1.xml"].encode()))
    assert theme.slots["accent1"] == "0077FF"  # стандартная тема — пустое место
    assert theme.major == "Brand Display"  # Calibri темы — шрифт по умолчанию, не свой у шаблона
    assert theme.slots["dk1"] == "000000"


def test_priority_brand_vs_template_on_real_template(tmp_path, fonts_here):
    """Настоящий шаблон LibreOffice (lo_vivid): брендбук перекрашивает акценты и текст, меняет шрифт; при приоритете
    шаблона свои цвета и Noto Sans остаются."""
    src = FIX / "templates" / "lo_vivid.pptx"
    ws = tmp_path / "ws"
    m = _analyze(src, ws)
    bdir = _brand_dir(tmp_path)
    brand = _brand()
    chrom = {"#0077FF", "#2787F5", "#00C2FF", "#FF3985"}
    out = tmp_path / "vivid-brand.pptx"
    res = apply_brand(src, out, brand, manifest=m, brand_dir=bdir)
    # текст (тёмный бирюзовый) → текст бренда, голубой акцент → цвет бренда того же тона
    assert res.colors["#04617B"] == "#0B1220"
    assert res.colors["#009EDA"] in chrom
    vis = _all(out, r"ppt/(slides|slideLayouts|slideMasters)/.*\.xml$")
    assert not {"04617B", "009EDA"} & _hexes(vis)
    assert 'typeface="Noto Sans"' not in vis and re.search(r'typeface="Brand (Display|Text)"', vis)
    theme = theme_info(etree.fromstring(_parts(out, r"ppt/theme/theme1\.xml$")["ppt/theme/theme1.xml"].encode()))
    assert theme.slots["accent1"] == "0077FF" and theme.major == "Brand Display"
    # разбор копии: акценты — цвета бренда или их оттенки (заливки шаблона с модификаторами яркости)
    bid = branded_template(m.template_id, brand, "brand", workspace_root=ws, brand_dir=bdir, render=False)
    copy = TemplateManifest.model_validate_json((ws / "templates" / bid / "manifest.json").read_text(encoding="utf-8"))
    acc = [t.hex for t in copy.tokens.colors if (t.role or "").startswith("accent.")]
    assert acc and {"0077FF", "2787F5"} <= set(acc)
    assert all(min(pal.hue_dist(pal.hue(a), pal.hue(b[1:])) for b in chrom) <= 25 for a in acc)
    assert not {"1A82BF", "009EDA", "04617B"} & {t.hex for t in copy.tokens.colors}
    assert {f.family for f in copy.tokens.typography.families} & {"Brand Display", "Brand Text"}
    # шаблон важнее: свои цвета и Noto Sans остаются
    out2 = tmp_path / "vivid-template.pptx"
    apply_brand(src, out2, brand, manifest=m, priority="template", brand_dir=bdir)
    vis2 = _all(out2, r"ppt/(slides|slideLayouts|slideMasters)/.*\.xml$")
    assert {"04617B", "009EDA"} <= _hexes(vis2) and 'typeface="Noto Sans"' in vis2


def test_contrast_kept_light_on_dark(tmp_path, fonts_here):
    """Брендбук со светлым фоном и тёмным текстом на тёмном шаблоне: светлый текст и тёмный фон шаблона остаются."""
    src = FIX / "templates" / "lo_vivid.pptx"
    m = _analyze(src, tmp_path / "ws")
    plan = plan_colors(m, _brand(colors=[BrandColor(hex="#0077FF", role="primary"), BrandColor(hex="#FFFFFF", role="text"), BrandColor(hex="#0B1220", role="background")]), "brand")
    # текст шаблона тёмный, текст бренда светлый — не меняется; фон шаблона светлый, фон бренда тёмный — не меняется
    text = next(t.hex for t in m.tokens.colors if "text.primary" in t.roles)
    bg = next(t.hex for t in m.tokens.colors if "background.light" in t.roles)
    assert plan.map(text) == text and plan.map(bg) == bg
    assert len(plan.warnings) == 2


# ---------------------------------------------------------------------------------------------- копия в хранилище и кэш


def test_branded_template_stored_analyzed_and_cached(synthetic, fonts_here, tmp_path, monkeypatch):
    src, _, bdir = synthetic
    ws = tmp_path / "ws"
    base = TemplateWorkspace.create(src, ws)
    (base.dir / "title.txt").write_text("ACME\n", encoding="utf-8")
    brand = _brand()
    bid = branded_template(base.template_id, brand, "brand", workspace_root=ws, brand_dir=bdir, render=False)
    assert bid != base.template_id
    d = ws / "templates" / bid
    assert (d / "source.pptx").exists() and (d / "manifest.json").exists()
    assert (d / "title.txt").read_text(encoding="utf-8").strip() == "ACME · брендбук Облако"
    m = TemplateManifest.model_validate_json((d / "manifest.json").read_text(encoding="utf-8"))
    assert m.template_id == bid
    # разбор копии видит бренд: палитра и шрифты — бренда, весь конвейер дальше работает с ними
    hexes = {t.hex for t in m.tokens.colors}
    assert "0077FF" in hexes and "1F6FEB" not in hexes
    fams = {f.family for f in m.tokens.typography.families}
    assert {"Brand Display", "Brand Text"} & fams and not {"Georgia", "Verdana"} & fams
    info = branded_info(bid, ws)
    assert info["base_template_id"] == base.template_id and info["priority"] == "brand" and info["brand_name"] == "Облако"
    # логотип ACME слева внизу, брендбук велит справа внизу — логотип бренда там, и это сказано
    assert info["result"]["logo_note"] == "в шаблоне слева внизу, по брендбуку — справа внизу; поставлен справа внизу"
    assert info["result"]["logo_covered"] == info["result"]["logo_samples"] >= 1

    # кэш: тот же брендбук и приоритет — без нового применения и разбора
    calls = []
    monkeypatch.setattr(ap, "apply_brand", lambda *a, **k: calls.append(1))
    import verstka.analysis.manifest as man

    monkeypatch.setattr(man, "analyze_template", lambda *a, **k: calls.append(2))
    assert branded_template(base.template_id, brand.model_copy(update={"id": "другой-id", "warnings": ["x"]}), "brand", workspace_root=ws, brand_dir=bdir) == bid
    assert calls == []
    # другой приоритет и другой брендбук — другие копии
    assert ap.branded_id(base.template_id, brand_fingerprint(brand, bdir), "template") != bid
    other = brand.model_copy(update={"colors": [BrandColor(hex="#00AA55", role="primary")]})
    assert brand_fingerprint(other, bdir) != brand_fingerprint(brand, bdir)
    # логотип поменялся в файле — другой отпечаток
    fp = brand_fingerprint(brand, bdir)
    _logo_png(bdir / "logos" / "primary.png", (255, 0, 0, 255), (0, 0, 0, 255))
    assert brand_fingerprint(brand, bdir) != fp


def test_branded_template_unchanged_copy_reuses_analysis(tmp_path, fonts_here, monkeypatch):
    """Брендбук ничего не меняет (шаблон важнее, своё всё) — копия получает разбор шаблона без нового разбора."""
    src = FIX / "templates" / "lo_vivid.pptx"
    ws = tmp_path / "ws"
    m = _analyze(src, ws)
    import verstka.analysis.manifest as man

    called = []
    monkeypatch.setattr(man, "analyze_template", lambda *a, **k: called.append(1))
    # шрифты шаблона свои (Noto Sans), тема — стандартная, но у бренда нет хроматических цветов и логотипов
    brand = _brand(colors=[BrandColor(hex="#222222", role="text")], fonts=[], logos=[], logo_place=None)
    monkeypatch.setattr(ap, "_retheme", lambda *a, **k: False)
    bid = branded_template(m.template_id, brand, "template", workspace_root=ws, render=False)
    assert called == []
    m2 = TemplateManifest.model_validate_json((ws / "templates" / bid / "manifest.json").read_text(encoding="utf-8"))
    assert m2.template_id == bid and len(m2.patterns) == len(m.patterns)
    assert (ws / "templates" / bid / "source.pptx").read_bytes() == TemplateWorkspace.open(m.template_id, ws).source.read_bytes()


def test_unknown_template_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        branded_template("0123456789abcdef", _brand(), workspace_root=tmp_path)


# ---------------------------------------------------------------------------------------------- расхождения


def _vk_manifest() -> TemplateManifest:
    data = json.loads(gzip.open(FIX / "manifest_cbbe3aa6a21d23ce.json.gz").read())
    return TemplateManifest.model_validate(data)


def by_logo_result(rows) -> str:
    return next(c.result for c in rows if c.what == "Логотип")


def test_conflicts_vk_tech_and_vk_cloud_as_in_scenario(monkeypatch):
    """SCENARIO.md: шрифт заголовков Play → VK Sans Display (замена — Play), второй акцент #FF3885 → #FF3985."""
    from verstka.brand.read import read_brandbook

    monkeypatch.setattr(ap, "font_available", lambda f: f == "Play")
    brand = read_brandbook(BRAND_EXAMPLE)
    m = _vk_manifest()
    src = FIX / "vktech_short.pptx"  # колода на шаблоне VK Tech: его тема и макеты
    got = brand_conflicts(m, brand, source=src, brand_dir=BRAND_EXAMPLE)
    by = {c.what: c for c in got}
    # ровно четыре строки, как принял контроллер
    assert [(c.what, c.template, c.brand, c.result) for c in got] == [
        ("Шрифт заголовков", "Play", "VK Sans Display", "VK Sans Display, замена — Play"),
        ("Шрифт текста", "Play", "VK Sans Text", "VK Sans Text, замена — Play"),
        ("Второй акцент", "#FF3885", "#FF3985", "#FF3985"),
        ("Логотип", "логотип шаблона", "логотип VK Cloud", by_logo_result(got)),
    ]
    head = by["Шрифт заголовков"]
    assert (head.template, head.brand, head.result, head.kind) == ("Play", "VK Sans Display", "VK Sans Display, замена — Play", "font")
    acc = by["Второй акцент"]
    assert (acc.template, acc.brand, acc.result, acc.kind) == ("#FF3885", "#FF3985", "#FF3985", "color")
    # основной акцент совпадает (#0077FF), чёрный и #0B1220, #FAFCFF и белый — для человека одно и то же
    assert "Основной акцент" not in by and "Цвет текста" not in by and "Фон" not in by
    # логотип VK Tech слева внизу, брендбук велит справа внизу: в расхождении сказано, где он встанет
    assert by["Логотип"].result.startswith("логотип VK Cloud: в шаблоне слева внизу, по брендбуку — справа внизу; ")
    # шаблон важнее — в презентации остаётся шаблон
    tpl = {c.what: c for c in brand_conflicts(m, brand, priority="template", source=src, brand_dir=BRAND_EXAMPLE)}
    assert tpl["Шрифт заголовков"].result == "Play, как в шаблоне"
    assert tpl["Второй акцент"].result == "#FF3885, как в шаблоне"


def test_conflicts_brand_font_installed(monkeypatch):
    from verstka.brand.read import read_brandbook

    monkeypatch.setattr(ap, "font_available", lambda f: True)
    got = {c.what: c for c in brand_conflicts(_vk_manifest(), read_brandbook(BRAND_EXAMPLE / "tokens.json"), source=FIX / "vktech_short.pptx")}
    assert got["Шрифт заголовков"].result == "VK Sans Display"
    # у токенов нет файлов логотипов: логотип шаблона остаётся, и это сказано
    assert (got["Логотип"].brand, got["Логотип"].result) == ("в брендбуке нет файлов логотипа", "логотип шаблона оставлен")


def test_conflicts_without_template_file():
    """Файла шаблона нет — расхождения по разбору (тема неизвестна: акценты — из разбора)."""
    got = brand_conflicts(_vk_manifest(), _brand(), workspace_root="/nonexistent")
    assert any(c.what == "Шрифт заголовков" for c in got)
    assert all(c.what and c.template and c.brand and c.result for c in got)


# ============================================================================================== правки после ревью
# «брендбук важнее шаблона — на любом шаблоне»: красный, коричневый, оливковый шаблоны; контраст; место логотипа


def _zip_put(path: Path, parts: dict[str, bytes]) -> None:
    import io as _io

    buf = _io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, parts.pop(item.filename, None) or zin.read(item))
        for name, data in parts.items():
            zout.writestr(name, data)
    path.write_bytes(buf.getvalue())


def _set_theme(path: Path, slots: dict[str, str]) -> None:
    with zipfile.ZipFile(path) as z:
        xml = z.read("ppt/theme/theme1.xml").decode("utf-8")
    for slot, hexv in slots.items():
        xml = re.sub(rf"<a:{slot}>.*?</a:{slot}>", f'<a:{slot}><a:srgbClr val="{hexv}"/></a:{slot}>', xml, flags=re.S)
    _zip_put(path, {"ppt/theme/theme1.xml": xml.encode("utf-8")})


def build_custom(tmp: Path, name: str, *, bg: str, text: str, cards: list[tuple[str, str]], theme: dict[str, str] | None = None,
                 n: int = 4, footer: bool = False, bg_picture: str | None = None, logos: list[tuple[str, float]] | None = None,
                 card_size: float = 14) -> Path:
    """Шаблон из `n` слайдов: фон `bg`, заголовок цвета `text`, карточки (заливка, цвет текста на ней), по желанию —
    текст колонтитула в правом нижнем углу, фоновая картинка во весь слайд, логотипы (имя, x) внизу."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    logo_png = _logo_png(tmp / "src" / f"{name}-logo.png", (31, 111, 235, 255), (31, 41, 55, 255))
    pic_png = None
    if bg_picture:
        pic_png = tmp / "src" / f"{name}-bg.png"
        Image.new("RGB", (160, 90), tuple(int(bg_picture[i:i + 2], 16) for i in (0, 2, 4))).save(pic_png)
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        if pic_png:
            pic = s.shapes.add_picture(str(pic_png), 0, 0, Emu(W), Emu(H))
            s.shapes._spTree.remove(pic._element)
            s.shapes._spTree.insert(2, pic._element)
        _text(s, f"Итоги квартала {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, text)
        for k, (fill, ink) in enumerate(cards):
            sp = _card(s, int(W * (0.05 + 0.31 * k)), fill)
            tf = sp.text_frame
            tf.text = "Выручка выросла на 23 % к сентябрю"
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.name = Pt(card_size), "Verdana"
            r.font.color.rgb = RGBColor.from_string(ink)
        if footer:
            _text(s, f"стр. {i + 1}", int(W * 0.84), int(H * 0.9), int(W * 0.12), int(H * 0.06), "Verdana", 12, text)
        for lname, lx in logos or []:
            if lname.endswith("only-first") and i:
                continue
            pic = s.shapes.add_picture(str(logo_png), Emu(int(W * lx)), Emu(int(H * 0.9)), height=Emu(int(H * 0.05)))
            pic.name = lname
    out = tmp / "src" / f"{name}.pptx"
    prs.save(out)
    if theme:
        _set_theme(out, theme)
    return out


def _vk_cloud():
    from verstka.brand.read import read_brandbook

    return read_brandbook(BRAND_EXAMPLE)


def _hexes_of(pptx: Path) -> set[str]:
    return _hexes(_all(pptx, r"ppt/(slides|slideLayouts|slideMasters|theme)/.*\.xml$"))


def test_red_theme_accent_like_lct_becomes_brand_primary(tmp_path, fonts_missing):
    """Красный акцент темы (как #FF0053 у ЛЦТ), разбором отмеченный «падением», — фирменный цвет шаблона: его заменяет
    основной цвет бренда; цвета диаграмм копии начинаются с основного цвета, «акцент» бренда — не первым."""
    src = build_custom(tmp_path, "red", bg="F3F4F5", text="1C1D22", cards=[("FF0053", "FFFFFF"), ("FF0053", "FFFFFF")],
                       theme={"accent1": "FF0053", "accent2": "FFD6E3", "accent3": "FC3777", "accent4": "8A83D1", "accent5": "520977",
                              "accent6": "2D1451", "hlink": "FF0053"}, card_size=24)
    ws = tmp_path / "ws"
    m = _analyze(src, ws)
    for t in m.tokens.colors:
        if t.hex == "FF0053":
            t.semantic = "negative"  # как у ЛЦТ: разбор считает красный «падением»
    TemplateWorkspace.open(m.template_id, ws).manifest_path.write_text(m.model_dump_json(), encoding="utf-8")
    brand = _vk_cloud()
    plan = ap.brand_plan(src, m, brand).colors
    assert plan.map("FF0053") == "0077FF" and "FF0053" not in plan.keep
    assert plan.map("FC3777") == "FF3985"  # близкий по тону розовый акцент шаблона — «акцент» бренда
    out = tmp_path / "red-brand.pptx"
    apply_brand(src, out, brand, manifest=m, brand_dir=BRAND_EXAMPLE)
    assert "FF0053" not in _hexes_of(out)
    bid = branded_template(m.template_id, brand, "brand", workspace_root=ws, brand_dir=BRAND_EXAMPLE, render=False)
    copy = TemplateManifest.model_validate_json((ws / "templates" / bid / "manifest.json").read_text(encoding="utf-8"))
    series = copy.components.chart_style.series_colors
    assert series[0] == "0077FF" and series.index("FF3985") > series.index("2787F5") if "FF3985" in series else series[0] == "0077FF"
    rows = {c.what: c for c in brand_conflicts(m, brand, source=src, brand_dir=BRAND_EXAMPLE)}
    assert rows["Основной акцент"].template == "#FF0053" and rows["Основной акцент"].result == "#0077FF"


def test_dark_olive_like_classic_gets_brand_shade_not_crimson(tmp_path, fonts_missing):
    """Тёмно-оливковый фон (как #59512B у «Классики») и коричневые — не в розовый «акцент» бренда (#840033): тёмный
    оттенок основного цвета бренда той же светлоты; коричневые нейтральные — холодные той же светлоты."""
    src = build_custom(tmp_path, "olive", bg="59512B", text="FFFFFF", cards=[("474134", "FFFFFF"), ("625342", "FFFFFF")],
                       theme={"accent1": "CEC5A7", "accent2": "59512B", "accent3": "827D57", "accent4": "B6A592", "accent5": "625342",
                              "accent6": "474134", "dk2": "474134"})
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "olive-brand.pptx"
    apply_brand(src, out, _vk_cloud(), manifest=m, brand_dir=BRAND_EXAMPLE)
    hexes = _hexes_of(out)
    assert not {"59512B", "474134", "625342", "840033", "FF3985"} & hexes
    pinkish = [h for h in hexes if pal.chroma(h) > 0.15 and pal.hue_dist(pal.hue(h), 337) <= 40]
    assert not pinkish
    plan = ap.brand_plan(src, m, _vk_cloud()).colors
    olive = plan.map("59512B")
    assert pal.hue_dist(pal.hue(olive), 212) <= 15 and abs(pal.lightness(olive) - pal.lightness("59512B")) < 2
    for brown in ("474134", "625342"):
        cool = plan.map(brown)
        assert abs(pal.lightness(cool) - pal.lightness(brown)) < 2 and pal.lab(cool)[2] < 0  # холодный (синеватый)


def test_brown_like_coffee_theme_accent_and_dark_ground_mapped(tmp_path, fonts_missing):
    """Коричневый шаблон (как «Кофейня»): акцент темы #913F37 («падение» по разбору) и бордовый фон обложки #6D2F29 —
    цвета бренда; основной акцент — в расхождениях; чёрный не сдвигается в #00000D."""
    src = build_custom(tmp_path, "coffee", bg="FFFFFF", text="151513", cards=[("913F37", "FFFFFF"), ("000000", "FFFFFF")],
                       theme={"accent1": "913F37", "accent2": "72726A", "accent3": "9B9B8F", "accent4": "64514F", "accent5": "9E655B",
                              "accent6": "A1A088", "dk1": "151513", "dk2": "322924"})
    prs = Presentation(str(src))
    cover = prs.slides[0]
    cover.background.fill.fore_color.rgb = RGBColor.from_string("6D2F29")
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    for t in m.tokens.colors:
        if t.hex == "913F37":
            t.semantic = "negative"
    brand = _vk_cloud()
    plan = ap.brand_plan(src, m, brand).colors
    assert pal.hue_dist(pal.hue(plan.map("913F37")), 212) <= 15
    assert pal.hue_dist(pal.hue(plan.map("6D2F29")), 212) <= 15
    assert plan.map("000000") in ("000000", "0B1220")
    out = tmp_path / "coffee-brand.pptx"
    apply_brand(src, out, brand, manifest=m, brand_dir=BRAND_EXAMPLE)
    assert not {"913F37", "6D2F29", "00000D"} & _hexes_of(out)
    rows = {c.what: c for c in brand_conflicts(m, brand, source=src, brand_dir=BRAND_EXAMPLE)}
    assert rows["Основной акцент"].template == "#913F37"


def test_kept_main_accent_is_a_conflict(tmp_path, fonts_missing):
    """Брендбук без фирменных цветов: акцент шаблона остаётся — и это видно в расхождениях, с причиной."""
    src = build_custom(tmp_path, "keep", bg="FFFFFF", text="1F2937", cards=[("E8590C", "FFFFFF")], theme={"accent1": "E8590C", "accent2": "C2410C", "accent3": "9A3412", "accent4": "FDBA74", "accent5": "7C2D12", "accent6": "FED7AA"})
    m = _analyze(src, tmp_path / "ws")
    brand = _brand(colors=[BrandColor(hex="#0B1220", role="text"), BrandColor(hex="#FFFFFF", role="background")])
    rows = {c.what: c for c in brand_conflicts(m, brand, source=src, brand_dir=_brand_dir(tmp_path))}
    assert rows["Основной акцент"].result == "#E8590C, как в шаблоне: в брендбуке нет фирменных цветов"


def test_white_text_on_purple_card_stays_readable(tmp_path, fonts_missing):
    """Белый текст на фиолетовой карточке (как #753A93 у ЛЦТ, 7,55 : 1) остаётся ≥ 4,5 : 1: цвет бренда берётся со
    светлотой карточки шаблона (#0077FF с белым — только 4,1 : 1)."""
    src = build_custom(tmp_path, "purple", bg="F3F4F5", text="1C1D22", cards=[("753A93", "FFFFFF"), ("753A93", "FFFFFF")],
                       theme={"accent1": "753A93", "accent2": "976BAE", "accent3": "520977", "accent4": "8A83D1", "accent5": "2D1451", "accent6": "DCCEE4"})
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "purple-brand.pptx"
    res = apply_brand(src, out, _vk_cloud(), manifest=m, brand_dir=BRAND_EXAMPLE)
    card = res.colors["#753A93"].lstrip("#")
    assert pal.hue_dist(pal.hue(card), pal.hue("0077FF")) <= 8  # тёмный оттенок основного, не бирюза
    assert contrast_ratio("FFFFFF", card) >= 4.5


def test_link_on_light_card_stays_readable(tmp_path, fonts_missing):
    """Ссылка #0563C1 мелким текстом на #EEF2F7 (5,25 : 1, как «Перейти» у VK Tech) не падает ниже 4,5 : 1."""
    src = build_custom(tmp_path, "link", bg="FFFFFF", text="000000", cards=[("EEF2F7", "0563C1"), ("0077FF", "FFFFFF")],
                       theme={"accent1": "0077FF", "accent2": "FFFFFF", "accent3": "FF3885", "accent4": "202020", "accent5": "F7FAFE",
                              "accent6": "D6ECFF", "hlink": "0563C1"}, card_size=12)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "link-brand.pptx"
    res = apply_brand(src, out, _vk_cloud(), manifest=m, brand_dir=BRAND_EXAMPLE)
    link = res.colors.get("#0563C1", "#0563C1").lstrip("#")
    ground = res.colors.get("#EEF2F7", "#EEF2F7").lstrip("#")
    assert contrast_ratio(link, ground) >= 4.5


def _slide_logo_media(pptx: Path) -> list[str]:
    """Картинки логотипов бренда на самих слайдах (не на макетах)."""
    out = []
    with zipfile.ZipFile(pptx) as z:
        for n in z.namelist():
            if re.match(r"ppt/slides/slide\d+\.xml$", n):
                xml = z.read(n).decode("utf-8")
                rels = z.read(n.replace("slides/", "slides/_rels/") + ".rels").decode("utf-8")
                for m in re.finditer(r'name="Логотип[^"]*".*?r:embed="([^"]+)"', xml, re.S):
                    t = re.search(rf'Id="{m.group(1)}"[^>]*Target="([^"]+)"', rels) or re.search(rf'Target="([^"]+)"[^>]*Id="{m.group(1)}"', rels)
                    out.append(t.group(1))
    return out


def _logo_boxes(pptx: Path) -> list[tuple[str, tuple[int, int, int, int]]]:
    out = []
    for name, xml in _parts(pptx, r"ppt/(slides|slideLayouts)/.*\.xml$").items():
        for m in re.finditer(r'name="Логотип [^"]*".*?<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"', xml, re.S):
            out.append((name, tuple(int(v) for v in m.groups())))
    return out


def test_added_logo_avoids_footer_text(tmp_path, fonts_here):
    """Номер страницы в правом нижнем углу (там, где логотипу место по брендбуку): логотип встаёт рядом, не на текст,
    или не ставится — со спокойным объяснением."""
    src = build_custom(tmp_path, "foot", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF"), ("DCE8FD", "1F2937"), ("1F6FEB", "FFFFFF")], footer=True)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "foot-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_covered >= 1 or any("нет свободного места" in w for w in res.warnings)
    assert _ir_collisions(out) == []  # в том числе с номером страницы
    from verstka.audit.ir import build_deck_ir

    for s in build_deck_ir(out).slides:
        for e in list(s.elements) + list(s.template_elements):
            if (e.name or "").startswith("Логотип"):
                assert e.bbox.y > H * 0.75 and e.bbox.x > W * 0.5  # у того же края и угла, рядом с номером


def test_added_logo_variant_by_real_ground(tmp_path, fonts_here):
    """Образцы светлого шаблона, но в углу — тёмная фоновая картинка: логотип встаёт поверх неё белым вариантом."""
    src = build_custom(tmp_path, "darkpic", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")], bg_picture="101828")
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "darkpic-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1
    media = _slide_logo_media(out)
    assert media and all("white" in n for n in media)


def test_replaced_logo_keeps_one_side_and_readable_size(tmp_path, fonts_here):
    """Два логотипа шаблона (свой слева на всех слайдах, партнёр справа на части) — логотип бренда везде слева, а не
    то слева, то справа; маленькая рамка — логотип не мельче читаемого."""
    src = build_custom(tmp_path, "two", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")],
                       logos=[("Logo ACME", 0.04), ("Logo Partner only-first", 0.85)])
    prs = Presentation(str(src))
    for s in prs.slides:
        for sh in s.shapes:
            if sh.name.startswith("Logo ACME"):
                sh.height = Emu(int(H * 0.015))
                sh.width = Emu(int(H * 0.06))
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "two-brand.pptx"
    res = apply_brand(src, out, _brand(logo_place="bottom_left"), manifest=m, brand_dir=_brand_dir(tmp_path))
    boxes = _logo_boxes(out)
    assert len(boxes) == 4 and res.logos_removed == 1
    assert all(x < W * 0.5 for _, (x, _y, _cx, _cy) in boxes)
    assert all(cy >= H * lg.MIN_WORDMARK_H * 0.99 for _, (_x, _y, _cx, cy) in boxes)


def test_theme_object_defaults_and_table_styles_rewritten(synthetic, fonts_here, tmp_path):
    src, m, bdir = synthetic
    with zipfile.ZipFile(src) as z:
        theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
    defaults = (
        '<a:objectDefaults><a:spDef><a:spPr><a:solidFill><a:srgbClr val="1F6FEB"/></a:solidFill></a:spPr><a:bodyPr/>'
        '<a:lstStyle><a:defPPr><a:defRPr><a:latin typeface="Georgia"/></a:defRPr></a:defPPr></a:lstStyle></a:spDef></a:objectDefaults>'
    )
    theme = re.sub(r"<a:objectDefaults\s*/>|<a:objectDefaults>.*?</a:objectDefaults>", defaults, theme, flags=re.S)
    if "<a:objectDefaults>" not in theme:
        theme = theme.replace("</a:themeElements>", "</a:themeElements>" + defaults)
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><a:tblStyleLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'def="{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"><a:tblStyle styleId="{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}" styleName="Brand">'
        '<a:wholeTbl><a:tcTxStyle><a:font><a:latin typeface="Georgia"/></a:font><a:srgbClr val="1F2937"/></a:tcTxStyle>'
        '<a:tcStyle><a:fill><a:solidFill><a:srgbClr val="1F6FEB"/></a:solidFill></a:fill></a:tcStyle></a:wholeTbl>'
        '<a:firstRow><a:tcStyle><a:fill><a:solidFill><a:srgbClr val="5AC831"/></a:solidFill></a:fill></a:tcStyle></a:firstRow></a:tblStyle></a:tblStyleLst>'
    )
    _zip_put(src, {"ppt/theme/theme1.xml": theme.encode("utf-8"), "ppt/tableStyles.xml": styles.encode("utf-8")})
    out = tmp_path / "defaults.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    th = _parts(out, r"ppt/theme/theme1\.xml$")["ppt/theme/theme1.xml"]
    od = re.search(r"<a:objectDefaults>.*?</a:objectDefaults>", th, re.S).group(0)
    assert 'val="0077FF"' in od and 'typeface="Brand Text"' in od
    ts = _parts(out, r"ppt/tableStyles\.xml$")["ppt/tableStyles.xml"]
    assert 'val="0077FF"' in ts and 'val="0B1220"' in ts and 'typeface="Brand Text"' in ts
    # цвет стиля таблиц вне палитры шаблона (зелёная шапка) — тоже цвет бренда, своей светлоты
    assert "5AC831" not in ts
    green = re.search(r'<a:firstRow>.*?srgbClr val="([0-9A-F]{6})"', ts, re.S).group(1)
    assert abs(pal.lightness(green) - pal.lightness("5AC831")) < 2 and min(pal.hue_dist(pal.hue(green), pal.hue(b)) for b in ("0077FF", "2787F5", "00C2FF")) <= 8


def test_image_probe_cache_keyed_by_hash(synthetic):
    src, m, _ = synthetic
    lg.find_template_logos(ap.Package(src), m)
    assert lg._PROBES and all(isinstance(k, str) and len(k) == 64 for k in lg._PROBES)
    assert not any(isinstance(x, (bytes, bytearray)) for v in lg._PROBES.values() if v for x in v)


def test_logo_path_confined_to_brand_dir(tmp_path):
    bdir = _brand_dir(tmp_path)
    ok = lg.logo_path(BrandLogo(name="Основной", file="logos/primary.png"), bdir)
    assert ok is not None and ok.is_relative_to(bdir.resolve())
    for bad in ("/etc/hosts", "../secret.png", "logos/../../secret.png", "C:/secret.png"):
        assert lg.logo_path(BrandLogo(name="x", file=bad), bdir) is None
    assert lg.logo_path(BrandLogo(name="x", file="logos/primary.png"), None) is None


def test_conflicts_module_uses_public_api_only():
    import ast

    tree = ast.parse((ROOT / "verstka" / "brand" / "conflicts.py").read_text(encoding="utf-8"))
    private = [a.name for n in ast.walk(tree) if isinstance(n, ast.ImportFrom) and (n.module or "").startswith("verstka.brand") for a in n.names if a.name.startswith("_")]
    assert private == []


# ============================================================================================== правки после ревью, круг 2
# Оракул столкновений — независимый от `verstka.brand.logos`: IR аудита (рамки заполнителей из макета и мастера, слои
# макета и мастера), как `collide_ir.py` ревьюера.


def _xml_lines(pptx: Path, part: str) -> list[tuple[float, float, float, float]]:
    """Линии нулевой толщины и соединители части (IR аудита их не видит): рамки (x, y, w, h) с допуском."""
    with zipfile.ZipFile(pptx) as z:
        try:
            xml = z.read(part).decode("utf-8")
        except KeyError:
            return []
    out = []
    for m in re.finditer(r"<p:(?:cxnSp|sp)>(.*?)</p:(?:cxnSp|sp)>", xml, re.S):
        body = m.group(1)
        if "<p:ph" in body or "grpSp" in body:
            continue
        x = re.search(r'<a:off x="(-?\d+)" y="(-?\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/>', body)
        if not x or "<a:ln" not in body:
            continue
        bx, by, bw, bh = (int(v) for v in x.groups())
        if bw == 0 or bh == 0 or 'prst="line"' in body or m.group(0).startswith("<p:cxnSp"):
            t = W * 0.004
            out.append((bx - t, by - t, bw + 2 * t, bh + 2 * t))
    return out


def _ir_collisions(pptx: Path) -> list[tuple]:
    from verstka.audit.ir import build_deck_ir

    ir = build_deck_ir(pptx, with_images=True)
    sw, sh = ir.slide_w, ir.slide_h
    hits = []
    with zipfile.ZipFile(pptx) as z:
        names = set(z.namelist())
        master_of = {}
        for n in names:
            if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n):
                rel = z.read(n.replace("slideLayouts/", "slideLayouts/_rels/") + ".rels").decode("utf-8")
                mm = re.search(r'Target="\.\./slideMasters/(slideMaster\d+\.xml)"', rel)
                master_of[n] = "ppt/slideMasters/" + mm.group(1) if mm else None
    for s in ir.slides:
        if s.layout_part:
            for lx, ly, lw, lh in _xml_lines(pptx, s.layout_part) + (_xml_lines(pptx, master_of.get(s.layout_part)) if master_of.get(s.layout_part) else []):
                for e in list(s.elements) + list(s.template_elements):
                    if (e.name or "").startswith("Логотип") and min(e.bbox.x + e.bbox.w, lx + lw) - max(e.bbox.x, lx) > 0 and min(e.bbox.y + e.bbox.h, ly + lh) - max(e.bbox.y, ly) > 0:
                        hits.append((s.index, "line", "layout/master", "", ""))
        allel = list(s.elements) + list(s.template_elements)
        for lg_el in [e for e in allel if (e.name or "").startswith("Логотип")]:
            la = lg_el.bbox.w * lg_el.bbox.h or 1
            for e in allel:
                if e is lg_el or (e.name or "").startswith("Логотип"):
                    continue
                ix = max(0, min(lg_el.bbox.x + lg_el.bbox.w, e.bbox.x + e.bbox.w) - max(lg_el.bbox.x, e.bbox.x))
                iy = max(0, min(lg_el.bbox.y + lg_el.bbox.h, e.bbox.y + e.bbox.h) - max(lg_el.bbox.y, e.bbox.y))
                if ix * iy <= 0.03 * la:
                    continue
                if e.type == "shape" and e.line_hex and not e.fill_hex and e.geometry != "custom":
                    # рамка-контур (и во весь слайд): столкновение — только с её краями
                    t = 0.004 * sw
                    x, y, w, h = e.bbox.x, e.bbox.y, e.bbox.w, e.bbox.h
                    lx, ly, lw, lh = lg_el.bbox.x, lg_el.bbox.y, lg_el.bbox.w, lg_el.bbox.h
                    crosses = any(
                        min(lx + lw, ex + ew) - max(lx, ex) > 0 and min(ly + lh, ey + eh) - max(ly, ey) > 0
                        for ex, ey, ew, eh in ((x - t, y - t, w + 2 * t, 2 * t), (x - t, y + h - t, w + 2 * t, 2 * t), (x - t, y - t, 2 * t, h + 2 * t), (x + w - t, y - t, 2 * t, h + 2 * t))
                    )
                    if crosses:
                        hits.append((s.index, "frame-edge", e.source, e.name, ""))
                    continue
                if e.bbox.w * e.bbox.h >= 0.85 * sw * sh:
                    continue  # фон слайда
                if e.type == "shape" and not e.fill_hex and not e.line_hex:
                    continue  # пустая рамка
                hits.append((s.index, e.type, e.source, e.name, e.text[:20]))
    return hits


def _layout_with_low_body(prs) -> object:
    """Макет «Заголовок и объект» с рамкой текста до самого низа и правого края: заполнители слайдов её наследуют."""
    lay = prs.slide_layouts[1]
    body = lay.placeholders[1]
    body.left, body.top, body.width, body.height = Emu(int(W * 0.05)), Emu(int(H * 0.25)), Emu(int(W * 0.92)), Emu(int(H * 0.72))
    return lay


def test_inherited_placeholder_frames_are_obstacles(tmp_path, fonts_here):
    """Заполнитель слайда без своей рамки (рамка — у макета) занимает низ слайда: логотип не встаёт на него — ни по
    проверке IR, ни на макете, куда ляжет содержимое собранных слайдов."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    lay = _layout_with_low_body(prs)
    for i in range(3):
        s = prs.slides.add_slide(lay)
        s.shapes.title.text = f"Итоги квартала {i + 1}"
        s.placeholders[1].text = "Выручка выросла\nКлиентов стало больше\nСредний чек вырос"
        assert s.placeholders[1]._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm") is None
    src = tmp_path / "inherit.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "inherit-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1 and _ir_collisions(out) == []
    # рамка текста отступила от логотипа (охранное поле), но не больше чем на 30 %: содержимое собранных слайдов
    # туда не зайдёт
    from verstka.audit.ir import build_deck_ir

    for s in build_deck_ir(out).slides:
        body = next(e for e in s.elements if e.is_placeholder and e.ph_type in ("body", "obj"))
        assert body.bbox.h >= 0.7 * H * 0.72 - 1 and body.bbox.w >= 0.75 * W * 0.92 - 1


def test_stroke_only_arrow_and_number_circle_are_obstacles(tmp_path, fonts_here):
    """Стрелка-контур без заливки (произвольная фигура) и кружок с номером в правом нижнем углу — препятствия."""
    src = build_custom(tmp_path, "arrow", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF"), ("DCE8FD", "1F2937"), ("1F6FEB", "FFFFFF")])
    prs = Presentation(str(src))
    for s in prs.slides:
        ff = s.shapes.build_freeform(int(W * 0.70), int(H * 0.93), scale=1.0)
        ff.add_line_segments([(int(W * 0.97), int(H * 0.93)), (int(W * 0.95), int(H * 0.91))], close=False)
        arrow = ff.convert_to_shape()
        arrow.fill.background()
        arrow.line.color.rgb = RGBColor.from_string("1F2937")
        circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(W * 0.94)), Emu(int(H * 0.86)), Emu(int(H * 0.06)), Emu(int(H * 0.06)))
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor.from_string("1F6FEB")
        circle.text_frame.text = "5"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "arrow-brand.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert _ir_collisions(out) == []


def _set_slide_bg(slide, xml_inner: str) -> None:
    from lxml import etree as _et

    csld = slide._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld")
    old = csld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if old is not None:
        csld.remove(old)
    bg = _et.fromstring(
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:bgPr>{xml_inner}<a:effectLst/></p:bgPr></p:bg>'
    )
    csld.insert(0, bg)


def test_photo_background_gets_white_logo(tmp_path, fonts_here):
    """Фон слайда — тёмное фото через `p:bg` (как у ЛЦТ): логотип бренда — белый, по настоящим точкам под ним."""
    photo = tmp_path / "src" / "purple.png"
    photo.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (320, 180), (60, 20, 100)).save(photo)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _, rid = s.part.get_or_add_image_part(str(photo))
        _set_slide_bg(s, f'<a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>')
        _text(s, f"Итоги {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "FFFFFF")
        _text(s, "Выручка выросла на 23 %", int(W * 0.05), int(H * 0.3), int(W * 0.5), int(H * 0.1), "Verdana", 16, "FFFFFF")
    src = tmp_path / "photo.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "photo-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1
    media = _slide_logo_media(out) + [n for n in re.findall(r'Target="\.\./media/(brand-logo-[^"]+)"', _all(out, r"ppt/slideLayouts/_rels/slideLayout7\.xml\.rels$"))]
    assert media and all("white" in n for n in media)


def test_light_tint_of_theme_colour_gets_dark_logo(tmp_path, fonts_here):
    """Фон — цвет темы с модификаторами (accent6, lumMod 20 % + lumOff 80 % — почти белый, как у «Классики»): логотип
    бренда — основной (тёмный), а не белый на светлом."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _set_slide_bg(s, '<a:solidFill><a:schemeClr val="accent6"><a:lumMod val="20000"/><a:lumOff val="80000"/></a:schemeClr></a:solidFill>')
        _text(s, f"Итоги {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
        _text(s, "Выручка выросла на 23 %", int(W * 0.05), int(H * 0.3), int(W * 0.5), int(H * 0.1), "Verdana", 16, "1F2937")
    src = tmp_path / "tint.pptx"
    prs.save(src)
    _set_theme(src, {"accent6": "1F4E79"})
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "tint-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1
    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("ppt/media/brand-logo-")]
    assert media and all("primary" in n for n in media)


def test_dark_cards_and_links_stay_in_primary_family(tmp_path, fonts_missing):
    """Тёмные фиолетовые карточки и ссылка без близкого тона у бренда — тёмные оттенки основного цвета, одна семья,
    а не бирюза из затемнённого голубого (ЛЦТ, слайд 3)."""
    src = build_custom(tmp_path, "violet", bg="F3F4F5", text="1C1D22", cards=[("753A93", "FFFFFF"), ("520977", "FFFFFF")],
                       theme={"accent1": "FF0053", "accent2": "FFD6E3", "accent3": "753A93", "accent4": "520977", "accent5": "8A83D1",
                              "accent6": "2D1451", "hlink": "0000EE"})
    m = _analyze(src, tmp_path / "ws")
    plan = ap.brand_plan(src, m, _vk_cloud()).colors
    for h in ("753A93", "520977", "0000EE"):
        got = plan.map(h)
        assert pal.hue_dist(pal.hue(got), pal.hue("0077FF")) <= 8, (h, got)


def test_font_swap_keeps_one_line_in_its_frame(tmp_path, fonts_missing):
    """Большая цифра, которая помещалась в свою рамку одной строкой, после смены шрифта тоже помещается: кегль
    уменьшается в пределах 20 %, а не переносится («16M» у MyBrand)."""
    from verstka.rendering.fonts import text_width_pt

    old = text_width_pt("118M", "Trebuchet MS", 82.5)
    new = text_width_pt("118M", "Play", 82.5)
    if new <= old * 1.02:
        pytest.skip("замена не шире шрифта шаблона на этой машине")
    inner = (old + new) / 2
    src = build_custom(tmp_path, "num", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")])
    prs = Presentation(str(src))
    for s in prs.slides:
        tb = _text(s, "118M", int(W * 0.5), int(H * 0.3), int(inner * 12700 + 2 * 91440), int(H * 0.3), "Trebuchet MS", 82.5, "1F2937")
        tb.name = "Число"
        tb.text_frame.word_wrap = True
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "num-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.refit >= 1
    xml = _parts(out, r"ppt/slides/slide1\.xml$")["ppt/slides/slide1.xml"]
    rpr = re.search(r'name="Число".*?<a:rPr[^>]*sz="(\d+)"[^>]*>.*?typeface="([^"]+)"', xml, re.S)
    size, family = int(rpr.group(1)) / 100, rpr.group(2)
    assert text_width_pt("118M", family, size) <= inner + 0.5


def test_second_brand_conflicts_call_uses_caches(tmp_path, fonts_missing, monkeypatch):
    """Второй вызов расхождений на том же файле не перечитывает пакет (тема, логотипы, пары «текст — фон» — из кэша)."""
    src = build_custom(tmp_path, "cache", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")], logos=[("Logo ACME", 0.04)])
    m = _analyze(src, tmp_path / "ws")
    brand_conflicts(m, _vk_cloud(), source=src, brand_dir=BRAND_EXAMPLE)
    made = []
    real = ap.Package

    class Counting(real):
        def __init__(self, *a, **k):
            made.append(1)
            super().__init__(*a, **k)

    monkeypatch.setattr(lg, "Package", Counting)
    monkeypatch.setattr(pal, "Package", Counting)
    import time

    t = time.time()
    brand_conflicts(m, _vk_cloud(), source=src, brand_dir=BRAND_EXAMPLE)
    assert made == [] and time.time() - t < 0.5


def test_near_identical_kept_background_makes_no_warning(synthetic, fonts_here, monkeypatch):
    """Фон шаблона оставлен ради контраста, но он для человека тот же, что фон бренда (ΔE ≤ 15), — предупреждения нет."""
    src, m, _ = synthetic
    monkeypatch.setattr(pal, "contrast_ratio", lambda a, b: 2.0)
    plan = plan_colors(m, _brand(), "brand", pal.template_theme(src))
    assert "F7F7F7" in plan.keep and plan.kept_reason["F7F7F7"] == "contrast"
    assert not any("F7F7F7" in w for w in plan.warnings)


def test_logo_does_not_cross_a_frame_edge(tmp_path, fonts_here):
    """Рамка-контур во весь слайд (как у «Классики»): логотип не встаёт на её край — внутри рамки или вне её."""
    src = build_custom(tmp_path, "frame", bg="F8F4EC", text="1F2937", cards=[("1F6FEB", "FFFFFF")])
    prs = Presentation(str(src))
    for s in prs.slides:
        fr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.015)), Emu(int(H * 0.02)), Emu(int(W * 0.97)), Emu(int(H * 0.93)))
        fr.fill.background()
        fr.line.color.rgb = RGBColor.from_string("1F2937")
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "frame-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1
    assert _ir_collisions(out) == []


def test_logo_does_not_sit_on_a_layout_line(tmp_path, fonts_here):
    """Линия нулевой толщины на макете поперёк низа слайда (линейка «Классики»; IR её не видит) — логотип не на ней."""
    src = build_custom(tmp_path, "rule", bg="F8F4EC", text="1F2937", cards=[("1F6FEB", "FFFFFF")])
    with zipfile.ZipFile(src) as z:
        lay = z.read("ppt/slideLayouts/slideLayout7.xml").decode("utf-8")
    line = (
        '<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="90" name="Линейка"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr><p:spPr><a:xfrm>'
        f'<a:off x="0" y="{int(H * 0.93)}"/><a:ext cx="{W}" cy="0"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
        '<a:ln w="12700"><a:solidFill><a:srgbClr val="1F2937"/></a:solidFill></a:ln></p:spPr></p:cxnSp>'
    )
    lay = lay.replace("</p:spTree>", line + "</p:spTree>", 1)
    _zip_put(src, {"ppt/slideLayouts/slideLayout7.xml": lay.encode("utf-8")})
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "rule-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_added >= 1
    assert _ir_collisions(out) == []


# ============================================================================================== правки после ревью, круг 3


_OVERRIDE_DARK = (
    '<p:clrMapOvr><a:overrideClrMapping bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1" accent2="accent2" '
    'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/></p:clrMapOvr>'
)


def _dark_layout(pptx: Path, part: str) -> None:
    """Макет с фоном bg2 и своей картой цветов bg2 → dk2 (как «Title with Picture 6» у «Кофейни»): тёмный при светлой
    теме."""
    with zipfile.ZipFile(pptx) as z:
        xml = z.read(part).decode("utf-8")
    assert "<a:masterClrMapping/>" in xml
    xml = xml.replace("<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>", _OVERRIDE_DARK)
    bg = '<p:bg><p:bgPr><a:solidFill><a:schemeClr val="bg2"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'
    xml = re.sub(r"(<p:cSld[^>]*>)", lambda mm: mm.group(1) + bg, xml, count=1)
    _zip_put(pptx, {part: xml.encode("utf-8")})


def _part_logo_media(pptx: Path, part: str) -> list[str]:
    """Картинки логотипа бренда на части (слайд или макет)."""
    with zipfile.ZipFile(pptx) as z:
        xml = z.read(part).decode("utf-8")
        rels = z.read(re.sub(r"([^/]+)$", r"_rels/\1.rels", part)).decode("utf-8")
    out = []
    for mm in re.finditer(r'name="Логотип[^"]*".*?r:embed="([^"]+)"', xml, re.S):
        t = re.search(rf'Id="{mm.group(1)}"[^>]*Target="([^"]+)"', rels) or re.search(rf'Target="([^"]+)"[^>]*Id="{mm.group(1)}"', rels)
        out.append(t.group(1))
    return out


def test_dark_colour_map_override_gets_white_logo(tmp_path, fonts_here):
    """Фон макета — bg2, а его карта цветов bg2 → dk2: макет тёмный при светлой теме. Логотип на нём — белый: и на
    макете образцов, и на пустом макете; на обычном светлом макете — основной."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[5])  # «Только заголовок» — станет тёмным
        s.shapes.title.text = f"Итоги квартала {i + 1}"
        _text(s, "Выручка выросла на 23 % к сентябрю", int(W * 0.05), int(H * 0.35), int(W * 0.5), int(H * 0.1), "Verdana", 16, "FFFFFF")
        _card(s, int(W * 0.05), "1F6FEB")
    src = tmp_path / "ovr.pptx"
    prs.save(src)
    _dark_layout(src, "ppt/slideLayouts/slideLayout6.xml")  # макет образцов
    _dark_layout(src, "ppt/slideLayouts/slideLayout7.xml")  # пустой макет без образцов
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "ovr-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_samples >= 1 and res.logo_covered == res.logo_samples
    on_samples = _part_logo_media(out, "ppt/slideLayouts/slideLayout6.xml") + _slide_logo_media(out)
    assert on_samples and all("white" in n for n in on_samples), on_samples
    blank = _part_logo_media(out, "ppt/slideLayouts/slideLayout7.xml")
    assert blank and all("white" in n for n in blank), blank
    light = _part_logo_media(out, "ppt/slideLayouts/slideLayout2.xml")
    assert all("primary" in n for n in light), light


def test_colour_modifiers_as_powerpoint_applies_them():
    """Модификаторы цвета темы: lumMod 20 % + lumOff 80 % — почти белый того же тона, shade — темнее, tint — светлее;
    `_fill_hex` берёт цвет слота темы и применяет их."""
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

    def el(inner: str) -> etree._Element:
        return etree.fromstring(f"<a:schemeClr {ns} val=\"accent6\">{inner}</a:schemeClr>")

    light = lg._with_mods("1F4E79", el('<a:lumMod val="20000"/><a:lumOff val="80000"/>'))
    assert pal.lightness(light) > 85 and pal.hue_dist(pal.hue(light), pal.hue("1F4E79")) <= 10
    assert lg._with_mods("FFFFFF", el('<a:shade val="50000"/>')) == "808080"
    assert lg._with_mods("000000", el('<a:tint val="25000"/>')) == "BFBFBF"
    assert lg._with_mods("FFFFFF", el('<a:lumMod val="50000"/>')) == "808080"
    fill = etree.fromstring(f'<a:solidFill {ns}><a:schemeClr val="accent6"><a:lumMod val="20000"/><a:lumOff val="80000"/></a:schemeClr></a:solidFill>')
    assert lg._fill_hex(fill, {"accent6": "1F4E79"}) == light


def test_unused_layout_with_photo_background_gets_white_logo(tmp_path, fonts_here):
    """Пустой макет (без образцов) с тёмным фото фоном `p:bg`: логотип на нём — белый, по точкам фото под ним; на
    светлых образцах — основной."""
    photo = tmp_path / "src" / "night.png"
    photo.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (320, 180), (24, 20, 48)).save(photo)
    src = build_custom(tmp_path, "lay", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")])
    prs = Presentation(str(src))
    lay = prs.slide_layouts[1]  # «Заголовок и объект» — образцов на нём нет
    _, rid = lay.part.get_or_add_image_part(str(photo))
    _set_slide_bg(lay, f'<a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "lay-brand.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    dark = _part_logo_media(out, "ppt/slideLayouts/slideLayout2.xml")
    assert dark and all("white" in n for n in dark), dark
    light = _slide_logo_media(out) + _part_logo_media(out, "ppt/slideLayouts/slideLayout7.xml")
    assert light and all("primary" in n for n in light), light


def test_logo_keeps_clear_space_from_a_card(tmp_path, fonts_here):
    """Карточка кончается в миллиметре над углом логотипа: логотип не встаёт вплотную — между ними охранное поле не
    меньше 1 % высоты слайда (иначе — другое место или чуть меньший логотип)."""
    src = build_custom(tmp_path, "gap", bg="FFFFFF", text="1F2937", cards=[("DCE8FD", "1F2937")])
    margin, h = 48 / 1920 * W, 0.045 * H
    top = H - margin - h
    card = (int(W * 0.5), int(H * 0.55), int(W * 0.48), int(top - 0.005 * H - H * 0.55))
    prs = Presentation(str(src))
    for s in prs.slides:
        sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Emu(v) for v in card))
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string("1F6FEB")
        sp.line.fill.background()
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "gap-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    boxes = _logo_boxes(out)
    assert res.logos_added >= 1 and boxes
    cx, cy, cw, ch = card
    for _, (x, y, w, hh) in boxes:
        dx = max(0, max(x, cx) - min(x + w, cx + cw))
        dy = max(0, max(y, cy) - min(y + hh, cy + ch))
        assert (dx * dx + dy * dy) ** 0.5 >= 0.0095 * H, (x / W, y / H, dx / H, dy / H)
    assert _ir_collisions(out) == []


def test_dark_tone_reachable_only_by_darkening_becomes_primary_shade(tmp_path, fonts_missing):
    """Тёмно-зелёный акцент шаблона (#168253, lo_focus): ближайший тон бренда — светлый голубой, и он получился бы
    только затемнением ниже L* 50 (петрольный). Вместо него — оттенок основного: одна семья."""
    src = build_custom(tmp_path, "focus", bg="FFFFFF", text="1F2937", cards=[("E8A202", "1F2937"), ("168253", "FFFFFF")],
                       theme={"accent1": "E8A202", "accent2": "168253", "accent3": "F10D0C", "accent4": "EA7500"})
    m = _analyze(src, tmp_path / "ws")
    plan = ap.brand_plan(src, m, _brand()).colors
    got = plan.map("168253")
    assert pal.hue_dist(pal.hue(got), pal.hue("0077FF")) <= 8, got


def test_link_colours_follow_brand_primary(tmp_path, fonts_missing):
    """Ссылки — основной цвет бренда: явный розовый цвет текста со ссылкой (ЛЦТ, слайд 30) — оттенок основного своей
    светлоты (контраст тот же); офисный #0563C1 в теме заметок — основной."""
    src = build_custom(tmp_path, "links", bg="2D1451", text="FFFFFF", cards=[("FF0053", "FFFFFF"), ("F92571", "FFFFFF")],
                       theme={"accent1": "FF0053", "accent2": "F92571", "accent3": "8A83D1"})
    prs = Presentation(str(src))
    for s in prs.slides:
        tb = _text(s, "https://www.flaticon.com/", int(W * 0.05), int(H * 0.8), int(W * 0.5), int(H * 0.06), "Verdana", 14, "F92571")
        tb.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://www.flaticon.com/"
    prs.slides[0].notes_slide.notes_text_frame.text = "Заметка"
    prs.save(src)
    with zipfile.ZipFile(src) as z:
        notes_rels = z.read("ppt/notesMasters/_rels/notesMaster1.xml.rels").decode("utf-8")
        notes_theme = "ppt/theme/" + re.search(r'Target="\.\./theme/(theme\d+\.xml)"', notes_rels).group(1)
        xml = z.read(notes_theme).decode("utf-8")
    xml = re.sub(r"<a:hlink>.*?</a:hlink>", '<a:hlink><a:srgbClr val="0563C1"/></a:hlink>', xml, flags=re.S)
    _zip_put(src, {notes_theme: xml.encode("utf-8")})
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "links-brand.pptx"
    apply_brand(src, out, _vk_cloud(), manifest=m, brand_dir=BRAND_EXAMPLE)
    slide = _parts(out, r"ppt/slides/slide1\.xml$")["ppt/slides/slide1.xml"]
    rpr = re.search(r"<a:rPr[^>]*>(?:(?!</a:rPr>).)*?hlinkClick(?:(?!</a:rPr>).)*?</a:rPr>|<a:rPr[^>]*>(?:(?!</a:rPr>).)*?</a:rPr>(?=<a:t>https)", slide, re.S)
    link = re.search(r'srgbClr val="([0-9A-F]{6})"', rpr.group(0)).group(1)
    assert pal.hue_dist(pal.hue(link), pal.hue("0077FF")) <= 8, link
    assert abs(pal.lightness(link) - pal.lightness("F92571")) <= 3  # контраст с фоном тот же
    theme = _parts(out, re.escape(notes_theme) + "$")[notes_theme]
    hl = re.search(r'<a:hlink><a:srgbClr val="([0-9A-F]{6})"/>', theme).group(1)
    assert hl == "0077FF", hl


def test_ir_cache_keeps_two_templates(tmp_path):
    """IR большого шаблона весит десятки мегабайт: в долгоживущем сервере кэш держит только два последних файла."""
    files = [build_custom(tmp_path, f"ir{i}", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")], n=1) for i in range(3)]
    pal._IRS.clear()
    for f in files:
        assert pal.source_ir(f) is not None
    assert len(pal._IRS) <= pal.IR_CACHE_SIZE == 2
    assert pal.source_ir(files[-1]) is pal.source_ir(files[-1])


def test_logo_goes_on_layouts_next_to_empty_footers(tmp_path, fonts_here):
    """Колонтитулы мастера — дата, подпись, номер — тянутся через весь низ слайда (4 : 3, как у synth_sidebar43). На
    собранных слайдах они пусты: логотип встаёт на макет рядом с ними, не заходя туда, где встанет номер слайда, когда
    его включат; новый слайд на этом макете получает логотип без столкновений."""
    prs = Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Итоги квартала {i + 1}"
        s.placeholders[1].text = "Выручка выросла\nКлиентов стало больше"
    src = tmp_path / "foot43.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "foot43-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_layouts >= 1 and res.logo_covered == res.logo_samples
    boxes = [b for part, b in _logo_boxes(out) if part == "ppt/slideLayouts/slideLayout2.xml"]
    assert boxes
    sw = 9144000
    master = _parts(out, r"ppt/slideMasters/slideMaster1\.xml$")["ppt/slideMasters/slideMaster1.xml"]
    num = re.search(r'<p:ph type="sldNum".*?<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/>', master, re.S)
    nx, ny, ncx, ncy = (int(v) for v in num.groups())
    tw = 3 * 0.6 * 12 * 12700 + 91440 * 2
    number = (nx + ncx - tw, ny, tw, ncy)  # где встанет номер слайда, когда его включат
    for x, y, cx, cy in boxes:
        assert y + cy > ny  # в полосе колонтитулов, рядом с ними, а не над ней
        assert min(x + cx, number[0] + number[2]) - max(x, number[0]) <= 0 or min(y + cy, ny + ncy) - max(y, ny) <= 0
        assert x + cx <= sw
    prs2 = Presentation(str(out))
    prs2.slides.add_slide(prs2.slide_layouts[1])
    fresh = tmp_path / "foot43-fresh.pptx"
    prs2.save(fresh)
    assert [h for h in _ir_collisions(fresh) if h[0] == 4] == []


def test_layout_logo_matches_the_layout_own_ground(tmp_path, fonts_here):
    """Макет тёмный (фото-фон ЛЦТ), а все его образцы — со своим светлым фоном: основной логотип на макете был бы
    тёмным на тёмном у нового слайда, собранного прямо по макету. Логотип встаёт на сами образцы (основной), а на
    макете — только подходящий его собственному фону."""
    src = build_custom(tmp_path, "darklay", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")])
    prs = Presentation(str(src))
    _set_slide_bg(prs.slide_layouts[6], '<a:solidFill><a:srgbClr val="1E1B4B"/></a:solidFill>')
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "darklay-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_covered == res.logo_samples >= 1
    on_layout = _part_logo_media(out, "ppt/slideLayouts/slideLayout7.xml")
    assert all("white" in n for n in on_layout), on_layout
    on_slides = _slide_logo_media(out)
    assert on_slides and all("primary" in n for n in on_slides), on_slides


# ============================================================================================== круг 4: логотип в колоде


def test_added_logo_carries_the_brand_logo_mark(tmp_path, fonts_here):
    """Добавленный логотип помечен (`p:cNvPr/@title`), подпись для чтения с экрана — «Логотип …»; заменённый логотип
    шаблона — нет (он и так часть оформления шаблона)."""
    from verstka.brand.package import BRAND_LOGO_MARK

    src = build_custom(tmp_path, "mark", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")])
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "mark-brand.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    tags = re.findall(r'<p:cNvPr [^>]*name="Логотип[^"]*"[^>]*>', _all(out, r"ppt/(slides|slideLayouts)/[^/]*\.xml$"))
    assert tags and all(f'title="{BRAND_LOGO_MARK}"' in t and 'descr="Логотип Облако"' in t for t in tags)
    rep = tmp_path / "rep-brand.pptx"
    src2 = build_custom(tmp_path, "rep", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")], logos=[("Logo ACME", 0.04)])
    apply_brand(src2, rep, _brand(logo_place="bottom_left"), manifest=_analyze(src2, tmp_path / "ws2"), brand_dir=_brand_dir(tmp_path))
    assert BRAND_LOGO_MARK not in _all(rep, r"ppt/slides/[^/]*\.xml$")


def _marked_logo_slide(tmp_path):
    """Слайд колоды python-pptx с логотипом бренда (меткой) в правом нижнем углу и текстом сверху."""
    from verstka.brand.package import BRAND_LOGO_MARK

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, "Итоги квартала", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
    png = _logo_png(tmp_path / "logo.png", (0, 119, 255, 255), (11, 18, 32, 255))
    pic = s.shapes.add_picture(str(png), Emu(int(W * 0.86)), Emu(int(H * 0.9)), Emu(int(W * 0.1)), Emu(int(H * 0.045)))
    pic.name = "Логотип Облако"
    pic._element.nvPicPr.cNvPr.set("title", BRAND_LOGO_MARK)
    return prs, s


def test_brand_logo_is_set_aside_and_put_back_in_place(tmp_path):
    """Сборка колоды откладывает логотип бренда склонированного образца и возвращает его на то же место поверх
    слайда, если новое содержимое туда не зашло; зашло (рамка текста до самого угла) — логотип не возвращается, его
    картинка уходит из связей слайда, а в предупреждениях сказано почему."""
    from verstka.rendering.brand_logo import DROPPED, hold_brand_logos, put_back_brand_logos

    prs, s = _marked_logo_slide(tmp_path)
    box = (s.shapes[-1].left, s.shapes[-1].top, s.shapes[-1].width, s.shapes[-1].height)
    assert hold_brand_logos(s) == 1 and not any(sh.name.startswith("Логотип") for sh in s.shapes)
    _text(s, "Выручка выросла", int(W * 0.05), int(H * 0.3), int(W * 0.5), int(H * 0.1), "Verdana", 16, "1F2937")
    warns: list[str] = []
    put_back_brand_logos(s, warns)
    logo = [sh for sh in s.shapes if sh.name.startswith("Логотип")]
    assert len(logo) == 1 and (logo[0].left, logo[0].top, logo[0].width, logo[0].height) == box and warns == []
    assert s.shapes[-1].name.startswith("Логотип")  # поверх всего
    ids = [sh.shape_id for sh in s.shapes]
    assert len(ids) == len(set(ids))

    prs, s = _marked_logo_slide(tmp_path)
    hold_brand_logos(s)
    _text(s, "Вывод слайда во всю ширину низа", int(W * 0.05), int(H * 0.86), int(W * 0.92), int(H * 0.1), "Verdana", 16, "1F2937")
    rels_before = len(s.part.rels)
    warns = []
    put_back_brand_logos(s, warns)
    assert not any(sh.name.startswith("Логотип") for sh in s.shapes) and warns == [DROPPED]
    assert len(s.part.rels) == rels_before - 1


def test_brand_off_slides_are_untouched(tmp_path):
    """Обычная колода (без метки логотипа бренда): отложить нечего, вернуть нечего — слайд не меняется ни в одном
    байте."""
    from verstka.rendering.brand_logo import hold_brand_logos, put_back_brand_logos

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    png = _logo_png(tmp_path / "plain.png", (0, 119, 255, 255), (11, 18, 32, 255))
    s.shapes.add_picture(str(png), Emu(100), Emu(100), Emu(1000000), Emu(250000)).name = "Логотип шаблона"
    before = etree.tostring(s._element)
    assert hold_brand_logos(s) == 0
    assert put_back_brand_logos(s, []) == (s, [])
    assert etree.tostring(s._element) == before and "_verstka_brand_logos" not in s.__dict__


def _deck(src: Path, out: Path) -> tuple[object, Path]:
    """Колода без модели по демо-плану, один вариант (structured)."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.outline import DeckOutline

    d = json.loads((ROOT / "tests" / "fixtures" / "outline_demo.json").read_text(encoding="utf-8"))
    d.pop("_source", None)
    got = generate_variants(src, outline=DeckOutline.model_validate(d), strategies=["structured"], out_dir=out, workspace_root=out / "ws",
                            use_llm=False, use_vlm=False, render_images=False, audit=False, autofix=False, exports=[], pictures=False, eyes=False)
    return got.variants[0], next(out.rglob("deck.pptx"))


def _no_logo_hooks(monkeypatch) -> None:
    """Сборка как до логотипа бренда на склонированных слайдах: ничего не откладывается и не возвращается."""
    import verstka.rendering.clone as cl
    import verstka.rendering.synth as sy

    for mod in (cl, sy):
        monkeypatch.setattr(mod, "hold_brand_logos", lambda slide, *a, **k: 0)
        monkeypatch.setattr(mod, "put_back_brand_logos", lambda slide, warns, *a, **k: (slide, warns))


def test_generated_deck_keeps_brand_logo_on_cloned_slides(tmp_path, fonts_here, monkeypatch):
    """Колода, собранная без модели на брендированной копии lo_vivid (логотип — на самих образцах): склонированные
    образцы и слайды на их холсте несут логотип бренда, без столкновений с содержимым; раньше сборка убирала его как
    картинку содержимого."""
    src = FIX / "templates" / "lo_vivid.pptx"
    copy = tmp_path / "vivid-brand.pptx"
    res = apply_brand(src, copy, _brand(), manifest=_analyze(src, tmp_path / "ws"), brand_dir=_brand_dir(tmp_path))
    assert res.logo_covered >= 1 and _slide_logo_media(copy)
    v, deck = _deck(copy, tmp_path / "out")
    assert any(s.mode == "clone" for s in v.render.slides)
    kept = _slide_logo_media(deck)
    assert len(kept) >= len(v.render.slides) // 2, kept
    assert _ir_collisions(deck) == []
    _no_logo_hooks(monkeypatch)
    _, deck0 = _deck(copy, tmp_path / "out0")
    assert len(_slide_logo_media(deck0)) < len(kept)


def test_ordinary_deck_is_the_same_with_and_without_logo_hooks(tmp_path, monkeypatch):
    """Обычная колода (шаблон без брендбука): каждая часть файла та же, что у сборки без отложенного логотипа."""
    src = FIX / "templates" / "lo_vivid.pptx"
    _, deck = _deck(src, tmp_path / "out")
    _no_logo_hooks(monkeypatch)
    _, deck0 = _deck(src, tmp_path / "out0")
    assert _parts_of(deck) == _parts_of(deck0)


def _parts_of(pptx: Path) -> dict[str, bytes]:
    """Части файла; у книги Excel диаграммы — её части без времени создания (оно своё у каждой сборки)."""
    import io as _io

    out = {}
    with zipfile.ZipFile(pptx) as z:
        for n in z.namelist():
            data = z.read(n)
            if n.endswith(".xlsx"):
                with zipfile.ZipFile(_io.BytesIO(data)) as x:
                    for m in x.namelist():
                        if m != "docProps/core.xml":
                            out[f"{n}!{m}"] = x.read(m)
            else:
                out[n] = data
    return out


# ============================================================================================== круг 5: живой прогон


def _wordmark_template(tmp: Path, name: str = "Acme_Pitch_Deck", text: str = "acme.", n: int = 4) -> Path:
    """Шаблон с логотипом без картинки, как у MyBrand: кружок-знак (группа фигур) и надпись «acme.» в заголовке-
    заполнителе с отступом слева — в левом верхнем углу обложки и нескольких слайдов."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[5] if i else prs.slide_layouts[0])
        if i == 0:
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = RGBColor.from_string("1F4FE0")
        grp = s.shapes.add_group_shape()
        c = grp.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(W * 0.05)), Emu(int(H * 0.05)), Emu(int(H * 0.08)), Emu(int(H * 0.08)))
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "1F4FE0")
        c.line.fill.background()
        bar = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.05 + H * 0.02)), Emu(int(H * 0.085)), Emu(int(H * 0.04)), Emu(int(H * 0.01)))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("1F2937")
        tb = _text(s, text, int(W * 0.05 + H * 0.1), int(H * 0.05), int(W * 0.3), int(H * 0.08), "Verdana", 28, "FFFFFF" if i == 0 else "1F2937")
        tb.text_frame.word_wrap = False
        if i == 0:
            s.shapes.title.text = "Итоги года"
        else:
            s.shapes.title.text = f"Раздел {i}"
            s.shapes.title.top = Emu(int(H * 0.2))
            _text(s, "Выручка выросла на 23 % к сентябрю", int(W * 0.05), int(H * 0.4), int(W * 0.5), int(H * 0.1), "Verdana", 16, "1F2937")
    out = tmp / "src" / f"{name}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def test_text_and_shape_wordmark_is_found_and_replaced(tmp_path, fonts_here):
    """Логотип шаблона из надписи и знака-фигур («mybrand.» у MyBrand): находится по имени шаблона в имени файла,
    вместе со знаком; в копии его нет — на его месте логотип бренда (на тёмной обложке — белый)."""
    src = _wordmark_template(tmp_path)
    m = _analyze(src, tmp_path / "ws")
    pkg = ap.Package(src)
    found = lg.find_text_logos(pkg, m, ("Calibri", "Calibri"), src)
    assert len(found) == 4 and all(lk.text == "acme." and len(lk.els) == 2 and lk.by_name for lk in found)
    out = tmp_path / "acme-brand.pptx"
    res = apply_brand(src, out, _brand(logo_place="top_left"), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_replaced == 4
    slides = _all(out, r"ppt/slides/slide\d+\.xml$")
    assert ">acme.<" not in slides and "<a:t>acme.</a:t>" not in slides
    assert "white" in _part_logo_media(out, "ppt/slides/slide1.xml")[0]
    assert all("primary" in _part_logo_media(out, f"ppt/slides/slide{i}.xml")[0] for i in (2, 3, 4))
    for _, (x, y, cx, cy) in _logo_boxes(out):
        assert x < W * 0.3 and y < H * 0.2  # там же, где был логотип шаблона


def test_short_texts_that_are_not_logos_stay(tmp_path):
    """Не логотип: адрес сайта, дата, число, надпись без знака и не имя шаблона, надпись посреди слайда."""
    src = _wordmark_template(tmp_path, name="Report_Deck", text="nordline.example")
    assert lg.find_text_logos(ap.Package(src), _analyze(src, tmp_path / "ws"), ("Calibri", "Calibri"), src) == []
    src2 = _wordmark_template(tmp_path, name="Report_Deck2", text="2026")
    assert lg.find_text_logos(ap.Package(src2), _analyze(src2, tmp_path / "ws2"), ("Calibri", "Calibri"), src2) == []
    # строчная надпись со знаком на трёх слайдах — логотип, даже если имя шаблона другое
    src3 = _wordmark_template(tmp_path, name="Report_Deck3", text="nordic.")
    assert len(lg.find_text_logos(ap.Package(src3), _analyze(src3, tmp_path / "ws3"), ("Calibri", "Calibri"), src3)) == 4


def test_template_logo_moves_to_the_brandbook_corner(synthetic, fonts_here, tmp_path):
    """Брендбук велит справа внизу, логотип шаблона слева внизу: логотип бренда встаёт справа внизу, логотипа шаблона
    нет, столкновений нет, и сказано, что сделано."""
    src, m, bdir = synthetic
    out = tmp_path / "moved.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    assert res.logo_note == "в шаблоне слева внизу, по брендбуку — справа внизу; поставлен справа внизу"
    assert "Логотип: в шаблоне слева внизу, по брендбуку — справа внизу; поставлен справа внизу." in res.warnings
    got = Presentation(str(out))
    kinds = {p.source_slide: getattr(p.kind, "value", p.kind) for p in m.patterns}
    for i, s in enumerate(got.slides, 1):
        logos = [sh for owner in (s, s.slide_layout) for sh in owner.shapes if sh.name.startswith("Логотип")]
        assert len(logos) <= 1, i  # ни на одном слайде нет двух логотипов
        if kinds.get(i) not in ("title", "section", "thanks"):
            assert logos and logos[0].left > W * 0.5 and logos[0].top > H * 0.7, i
    assert "Логотип ACME" not in _all(out, r"ppt/(slides|slideLayouts)/[^/]*\.xml$")
    assert res.logo_covered == res.logo_samples and _ir_collisions(out) == []
    rows = {c.what: c for c in brand_conflicts(m, _brand(), source=src, brand_dir=bdir)}
    assert rows["Логотип"].result.endswith("поставлен справа внизу")


def test_template_logo_stays_where_the_brandbook_corner_is_taken(tmp_path, fonts_here):
    """Угол из брендбука занят карточкой на всех слайдах: логотип бренда — на месте логотипа шаблона (слева внизу), и
    это сказано; на одном слайде из трёх угол свободен — там логотип переезжает, остальные — как в шаблоне."""
    src = build_custom(tmp_path, "busy", bg="FFFFFF", text="1F2937", cards=[("DCE8FD", "1F2937")], logos=[("Logo ACME", 0.04)])
    prs = Presentation(str(src))
    for s in prs.slides:
        sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.55)), Emu(int(H * 0.7)), Emu(int(W * 0.43)), Emu(int(H * 0.28)))
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string("1F6FEB")
        sp.line.fill.background()
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "busy-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_note == "в шаблоне слева внизу, по брендбуку — справа внизу; оставлен слева внизу: угол из брендбука занят"
    boxes = [b for part, b in _logo_boxes(out) if "/slides/" in part]
    assert boxes and all(x < W * 0.5 for (x, _y, _cx, _cy) in boxes) and res.logos_replaced >= 1


def test_layout_logo_moves_and_blocked_slides_keep_a_copy_of_the_layout(tmp_path, fonts_here):
    """Логотип шаблона на макете (как у VK Tech). На двух слайдах из трёх угол брендбука свободен, на одном — занят
    карточкой: макет получает логотип в углу брендбука, а слайд с карточкой — копию макета с логотипом в углу
    шаблона. Файл открывается, копия — в списке макетов мастера, столкновений нет."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    lay = prs.slide_layouts[6]
    png = _logo_png(tmp_path / "src" / "acme.png", (31, 111, 235, 255), (31, 41, 55, 255))
    tmp_slide = prs.slides.add_slide(lay)
    pic = tmp_slide.shapes.add_picture(str(png), Emu(int(W * 0.04)), Emu(int(H * 0.9)), height=Emu(int(H * 0.05)))
    pic.name = "Logo ACME"
    _, rid = lay.part.get_or_add_image_part(str(png))
    el = pic._element
    el.blipFill.blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    lay.shapes._spTree.append(el)
    sld_ids = prs.slides._sldIdLst
    prs.part.drop_rel(sld_ids[0].rId)
    sld_ids.remove(sld_ids[0])
    for i in range(3):
        s = prs.slides.add_slide(lay)
        _text(s, f"Итоги квартала {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
        _card(s, int(W * 0.05), "DCE8FD")
        if i == 2:
            sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(W * 0.55)), Emu(int(H * 0.7)), Emu(int(W * 0.43)), Emu(int(H * 0.28)))
            sp.fill.solid()
            sp.fill.fore_color.rgb = RGBColor.from_string("1F6FEB")
    src = tmp_path / "src" / "laylogo.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "laylogo-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_note and re.search(r"поставлен справа внизу, а на 1 из \d слайд\w* шаблона, где этот угол занят, — слева внизу", res.logo_note), res.logo_note
    got = Presentation(str(out))
    names = [l.name for mm in got.slide_masters for l in mm.slide_layouts]
    assert any("логотип как в шаблоне" in n for n in names)
    sides = []
    for s in got.slides:
        logos = [sh for owner in (s, s.slide_layout) for sh in owner.shapes if sh.name.startswith("Логотип")]
        assert len(logos) == 1
        sides.append("right" if logos[0].left > W * 0.5 else "left")
    assert sides == ["right", "right", "left"]
    assert _ir_collisions(out) == []


def test_logo_only_on_covers_content_slides_get_the_brand_logo(tmp_path, fonts_here):
    """Логотип шаблона только на обложке (как у VK Education): обложка — логотип бренда на его месте, содержательные
    слайды — логотип бренда в углу из брендбука."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    png = _logo_png(tmp_path / "src" / "edu.png", (31, 111, 235, 255), (31, 41, 55, 255))
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Образовательная программа"
    cover.placeholders[1].text = "Итоги года"
    cover.shapes.add_picture(str(png), Emu(int(W * 0.04)), Emu(int(H * 0.05)), height=Emu(int(H * 0.06))).name = "Logo EDU"
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, f"Итоги квартала {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
        for k in range(3):
            _card(s, int(W * (0.05 + 0.31 * k)), "DCE8FD")
    src = tmp_path / "src" / "edu.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "edu-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_replaced == 1 and res.logos_added >= 1
    assert res.logo_samples >= 1 and res.logo_covered == res.logo_samples
    cover_logo = [b for part, b in _logo_boxes(out) if part == "ppt/slides/slide1.xml"]
    assert cover_logo and cover_logo[0][0] < W * 0.3 and cover_logo[0][1] < H * 0.2


def _dataset(name: str):
    import os

    env = os.environ.get("VERSTKA_FIXTURES_DIR")
    for d in ([Path(env)] if env else []) + [p / "Датасет" for p in ROOT.parents]:
        if d.is_dir():
            hit = next((p for p in d.glob("*.pptx") if name in p.name.lower()), None)
            if hit is not None:
                return hit
    return None


def test_vk_education_content_slides_get_the_logo(tmp_path, fonts_missing):
    """VK Education: логотип шаблона — только на обложках и финальных; содержательные слайды получают логотип VK Cloud
    (раньше — ни на одном)."""
    src = _dataset("education")
    if src is None:
        pytest.skip("шаблон VK Education не найден (VERSTKA_FIXTURES_DIR)")
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "edu-brand.pptx"
    res = apply_brand(src, out, _vk_cloud(), manifest=m, brand_dir=BRAND_EXAMPLE)
    assert res.logo_samples >= 10 and res.logo_covered >= 0.8 * res.logo_samples
    assert res.logos_added >= 10 and _ir_collisions(out) == []


def test_brandbook_without_logo_files_says_so(tmp_path, fonts_here):
    """Брендбук из tokens.json без файлов логотипа: одно спокойное предупреждение — в отчёте копии, в `brand.json`
    (и чтения брендбука, и применения) и в `branded_warnings`, которые печатает CLI и пишет run_manifest."""
    from verstka.brand.read import NO_LOGO_FILES

    src = build_custom(tmp_path, "nofiles", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")], logos=[("Logo ACME", 0.04)])
    res = apply_brand(src, tmp_path / "a.pptx", _brand(logos=[]), manifest=_analyze(src, tmp_path / "ws"), brand_dir=_brand_dir(tmp_path))
    assert NO_LOGO_FILES in res.warnings
    ws = tmp_path / "ws3"
    base = TemplateWorkspace.create(src, ws)
    brand = _brand(logos=[], warnings=[NO_LOGO_FILES])
    bid = branded_template(base.template_id, brand, "brand", workspace_root=ws, brand_dir=_brand_dir(tmp_path), render=False)
    info = branded_info(bid, ws)
    assert NO_LOGO_FILES in info["brand_warnings"] and NO_LOGO_FILES in info["result"]["warnings"]
    assert ap.branded_warnings(bid, ws).count(NO_LOGO_FILES) == 1
    assert ap.branded_warnings(base.template_id, ws) == []


def test_status_red_is_kept_unless_the_brand_has_a_danger_colour(tmp_path, fonts_missing):
    """Красный «ошибки и падения» не становится синим бренда: слот стандартной темы и красный акцент шаблона остаются
    красными; если в брендбуке есть цвет ошибки — становятся им. Розово-малиновый главный акцент (ЛЦТ) — по-прежнему
    цвет бренда."""
    src = build_custom(tmp_path, "red", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF"), ("C9211E", "FFFFFF")])
    m = _analyze(src, tmp_path / "ws")
    plan = ap.brand_plan(src, m, _brand()).colors
    assert plan.map("C0504D") == "C0504D" and plan.map("C9211E") == "C9211E"
    assert plan.map_strict("C9211E") == "C9211E"
    danger = _brand(colors=_brand().colors + [BrandColor(hex="#E5243B", role="danger")])
    plan2 = ap.brand_plan(src, m, danger).colors
    assert plan2.map("C0504D") == "E5243B" and plan2.map("C9211E") == "E5243B"
    assert pal.status_red("C9211E") and not pal.status_red("FF0053") and not pal.status_red("FF3985")
    lct = build_custom(tmp_path, "lctred", bg="F3F4F5", text="1C1D22", cards=[("FF0053", "FFFFFF")], theme={"accent1": "FF0053", "accent2": "FC3777"})
    assert pal.hue_dist(pal.hue(ap.brand_plan(lct, _analyze(lct, tmp_path / "ws3"), _brand()).colors.map("FF0053")), pal.hue("0077FF")) <= 8


# ============================================================================================== круг 6: живой прогон 2


def test_cover_without_any_logo_gets_the_brand_logo(tmp_path, fonts_here):
    """У шаблона нет своего логотипа (Forestbird): логотип бренда получают и содержательные слайды, и обложка — на
    том же месте, если оно свободно и там."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Итоги года"
    cover.placeholders[1].text = "Для правления"
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, f"Итоги квартала {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
        _card(s, int(W * 0.05), "DCE8FD")
    src = tmp_path / "src" / "nologo.pptx"
    src.parent.mkdir(parents=True, exist_ok=True)
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    assert any(getattr(p.kind, "value", p.kind) == "title" for p in m.patterns)
    out = tmp_path / "nologo-brand.pptx"
    res = apply_brand(src, out, _brand(), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logo_covers >= 1
    got = Presentation(str(out))
    first = got.slides[0]
    logos = [sh for owner in (first, first.slide_layout) for sh in owner.shapes if sh.name.startswith("Логотип")]
    assert len(logos) == 1 and _ir_collisions(out) == []


def test_missing_logo_files_reach_the_brandbook_layer(tmp_path, monkeypatch):
    """tokens.json без файлов логотипа: предупреждение — в самом брендбуке (`BrandSpec.warnings`) и в расхождениях, и
    `GET /api/brandbooks/{id}` отдаёт оба — слой «Брендбук» их показывает."""
    import time

    from fastapi.testclient import TestClient

    import verstka.api.app as A
    from verstka.api.store import Store
    from verstka.brand.read import NO_LOGO_FILES, read_brandbook

    assert NO_LOGO_FILES in read_brandbook(BRAND_EXAMPLE / "tokens.json").warnings
    store = Store(tmp_path)
    monkeypatch.setattr(A, "store", store)
    monkeypatch.setattr(ap, "font_available", lambda f: f == "Play")
    tid = "cbbe3aa6a21d23ce"
    d = tmp_path / "templates" / tid
    d.mkdir(parents=True)
    (d / "manifest.json").write_bytes(gzip.open(FIX / f"manifest_{tid}.json.gz").read())
    (d / "source.pptx").write_bytes((FIX / "vktech_short.pptx").read_bytes())
    c = TestClient(A.app)
    r = c.post("/api/brandbooks", files={"file": ("Брендбук.json", (BRAND_EXAMPLE / "tokens.json").read_bytes(), "application/json")}, data={"template_id": tid})
    assert r.status_code == 200, r.text
    bid, jid = r.json()["brandbook_id"], r.json()["job_id"]
    t0 = time.time()
    while time.time() - t0 < 60:
        job = A.runner.get(jid)
        if job is not None and job.events and job.events[-1].get("status") in ("done", "failed"):
            break
        time.sleep(0.05)
    assert job.status == "done", job.error
    got = c.get(f"/api/brandbooks/{bid}?template_id={tid}").json()
    assert NO_LOGO_FILES in got["brand"]["warnings"]
    row = next(x for x in got["conflicts"] if x["what"] == "Логотип")
    assert row["brand"] == "в брендбуке нет файлов логотипа" and row["result"] == "логотип шаблона оставлен"


def test_background_pictures_with_partner_marks_stay_as_they_are(tmp_path, fonts_here):
    """Знаки партнёров шаблона (у ЛЦТ — организаторы, впечатаны в фон) принадлежат содержанию: копия под брендбук
    фоновые картинки не трогает ни при каком приоритете."""
    bg = tmp_path / "src" / "bg.png"
    bg.parent.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", (1600, 900), (70, 25, 110))
    d = ImageDraw.Draw(im)
    for k in range(4):
        d.rectangle([900 + k * 170, 40, 1030 + k * 170, 90], outline=(255, 255, 255), width=4)
    im.save(bg)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _, rid = s.part.get_or_add_image_part(str(bg))
        _set_slide_bg(s, f'<a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>')
        _text(s, f"Итоги {i + 1}", int(W * 0.05), int(H * 0.2), int(W * 0.6), int(H * 0.14), "Georgia", 36, "FFFFFF")
    src = tmp_path / "src" / "partners.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    before = {n: b for n, b in ((n, zipfile.ZipFile(src).read(n)) for n in zipfile.ZipFile(src).namelist()) if n.startswith("ppt/media/")}
    for prio in ("brand", "template"):
        out = tmp_path / f"p-{prio}.pptx"
        apply_brand(src, out, _brand(), manifest=m, priority=prio, brand_dir=_brand_dir(tmp_path))
        with zipfile.ZipFile(out) as z:
            assert all(z.read(n) == b for n, b in before.items())


def test_title_caps_follow_only_a_brandbook_rule(synthetic, fonts_here, tmp_path):
    """Капс заголовков — оформление шаблона (ЛЦТ): остаётся, пока брендбук о нём молчит; правило «заголовки — без
    капса» его снимает."""
    from verstka.schemas.brand import BrandRule

    src, m, bdir = synthetic
    with zipfile.ZipFile(src) as z:
        master = z.read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
    master = re.sub(r"(<p:titleStyle><a:lvl1pPr[^>]*>(?:(?!</a:lvl1pPr>).)*?<a:defRPr)", lambda mm: mm.group(1) + ' cap="all"', master, count=1, flags=re.S)
    assert 'cap="all"' in master
    _zip_put(src, {"ppt/slideMasters/slideMaster1.xml": master.encode("utf-8")})
    assert ap.caps_rule(_vk_cloud()) is None
    out = tmp_path / "caps.pptx"
    apply_brand(src, out, _brand(), manifest=m, brand_dir=bdir)
    assert 'cap="all"' in _parts(out, r"ppt/slideMasters/slideMaster1\.xml$")["ppt/slideMasters/slideMaster1.xml"]
    nocaps = _brand(rules=[BrandRule(text="Заголовки — без капса, как в предложении", kind="type")])
    assert ap.caps_rule(nocaps) is False
    out2 = tmp_path / "nocaps.pptx"
    res = apply_brand(src, out2, nocaps, manifest=m, brand_dir=bdir)
    assert res.caps_changed >= 1 and 'cap="all"' not in _parts(out2, r"ppt/slideMasters/slideMaster1\.xml$")["ppt/slideMasters/slideMaster1.xml"]


# ============================================================================================== круг 7: живой прогон 3


def _group_lockup_xml(rid_mark: str, rid_word: str, box: tuple[int, int, int, int], child_w: int, name: str = "Logo ACME") -> str:
    """Группа-логотип (знак и надпись) с рамкой `box`; ширина детей в координатах группы — `child_w`."""
    x, y, w, h = box
    return (
        f'<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:nvGrpSpPr><p:cNvPr id="900" name="{name} group"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        f'<p:grpSpPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/><a:chOff x="0" y="0"/><a:chExt cx="{child_w}" cy="{h}"/></a:xfrm></p:grpSpPr>'
        f'<p:pic><p:nvPicPr><p:cNvPr id="901" name="{name} mark"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid_mark}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{h}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        f'<p:pic><p:nvPicPr><p:cNvPr id="902" name="{name} word"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid_word}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{int(h * 1.2)}" y="{int(h * 0.3)}"/><a:ext cx="{child_w - int(h * 1.2)}" cy="{int(h * 0.4)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic></p:grpSp>'
    )


def _abs_pics(holder) -> list[tuple[str, tuple[float, float, float, float]]]:
    """Картинки части с их рамками на слайде (через преобразования групп)."""
    out = []
    for pic in holder._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}pic"):
        b = lg.xfrm_box(pic.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr/" + A + "xfrm"))
        groups = [g for g in pic.iterancestors() if g.tag == "{http://schemas.openxmlformats.org/presentationml/2006/main}grpSp"]
        out.append((pic.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").get("name"), lg.to_abs(b, groups)))
    return out


def test_wider_brand_logo_in_a_template_group_refits_the_group(tmp_path, fonts_here):
    """Логотип шаблона — группа (знак и надпись), логотип бренда шире: после замены рамка группы — там, где видно
    логотип бренда, и он внутри слайда (VK Tech: картинка 1 676 934 EMU в группе шириной 1 232 613)."""
    (tmp_path / "src").mkdir(parents=True, exist_ok=True)
    mark = tmp_path / "src" / "mark.png"
    Image.new("RGBA", (100, 100), (31, 111, 235, 255)).save(mark)
    word = _logo_png(tmp_path / "src" / "word.png", (31, 41, 55, 255), (31, 41, 55, 255), size=(300, 100))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    box = (int(W * 0.03), int(H * 0.9), int(W * 0.08), int(H * 0.05))
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, f"Итоги квартала {i + 1}", int(W * 0.05), int(H * 0.06), int(W * 0.8), int(H * 0.14), "Georgia", 36, "1F2937")
        _, r1 = s.part.get_or_add_image_part(str(mark))
        _, r2 = s.part.get_or_add_image_part(str(word))
        s.shapes._spTree.append(parse_xml(_group_lockup_xml(r1, r2, box, box[2])))
    src = tmp_path / "src" / "grp.pptx"
    prs.save(src)
    m = _analyze(src, tmp_path / "ws")
    out = tmp_path / "grp-brand.pptx"
    res = apply_brand(src, out, _brand(logo_place="bottom_left"), manifest=m, brand_dir=_brand_dir(tmp_path))
    assert res.logos_replaced == 3
    got = Presentation(str(out))
    for s in got.slides:
        grp = next(sh for sh in s.shapes if sh.shape_type is not None and sh._element.tag.endswith("}grpSp"))
        gx, gy, gw, gh = grp.left, grp.top, grp.width, grp.height
        pics = [b for name, b in _abs_pics(s) if name.startswith("Логотип")]
        assert pics
        for x, y, w, h in pics:
            assert gx - 2 <= x and x + w <= gx + gw + 2 and gy - 2 <= y and y + h <= gy + gh + 2
            assert 0 <= x and x + w <= W


def test_one_corner_keeps_a_grouped_logo_inside_the_slide(tmp_path):
    """Сборка ставит логотипы колоды в один угол (`one_corner`). Логотип макета — группа, чья рамка уже логотипа бренда
    внутри (копия старой версии): её мерят по детям и ставят так, что логотип не выходит за край («VK Clouc»)."""
    from verstka.brand.package import BRAND_LOGO_MARK
    from verstka.rendering.brand_logo import one_corner

    png = _logo_png(tmp_path / "logo.png", (0, 119, 255, 255), (11, 18, 32, 255))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    lay = prs.slide_layouts[5]
    _, r1 = lay.part.get_or_add_image_part(str(png))
    narrow = int(W * 0.06)  # рамка группы и chExt — уже картинки внутри
    grp = parse_xml(_group_lockup_xml(r1, r1, (int(W * 0.03), int(H * 0.9), narrow, int(H * 0.045)), narrow))
    wide = grp.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}pic")[1]
    wide.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("name", "Логотип Облако")
    wide.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("descr", "Логотип Облако")
    wide.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr/" + A + "xfrm/" + A + "ext").set("cx", str(int(W * 0.12)))
    lay.shapes._spTree.append(grp)
    for i in range(4):
        s = prs.slides.add_slide(prs.slide_layouts[6] if i < 3 else lay)
        if i < 3:
            pic = s.shapes.add_picture(str(png), Emu(int(W * 0.85)), Emu(int(H * 0.9)), Emu(int(W * 0.12)), Emu(int(H * 0.045)))
            pic.name = "Логотип Облако"
            pic._element.nvPicPr.cNvPr.set("title", BRAND_LOGO_MARK)
            pic._element.nvPicPr.cNvPr.set("descr", "Логотип Облако")
    assert one_corner(prs) >= 1
    for name, (x, y, w, h) in _abs_pics(prs.slides[3].slide_layout):
        if name.startswith("Логотип"):
            assert x > W * 0.5 and x + w <= W * 0.99 and y + h <= H * 0.99, (x / W, (x + w) / W)


def test_missing_logo_warning_in_cli_output_and_run_manifest(tmp_path, fonts_here, monkeypatch):
    """Брендбук без файлов логотипа (tokens.json): `verstka generate --brand` печатает предупреждение, и оно же — в
    `run_manifest.warnings` колоды на брендированной копии."""
    import verstka.cli.main as cli
    from verstka.brand.read import NO_LOGO_FILES

    src = FIX / "templates" / "lo_vivid.pptx"
    printed: list[str] = []
    monkeypatch.setattr(cli.console, "print", lambda *a, **k: printed.append(" ".join(str(x) for x in a)))

    class _Status:
        def update(self, *a, **k):
            pass

    ws = tmp_path / "ws"
    copy = cli._branded(src, str(BRAND_EXAMPLE / "tokens.json"), "brand", ws, None, None, use_vlm=False, status=_Status())
    assert any(NO_LOGO_FILES in line for line in printed), printed
    _, deck = _deck_in(copy, tmp_path / "out", ws)
    rm = json.loads((deck.parent / "run_manifest.json").read_text(encoding="utf-8"))
    assert rm["warnings"].count(NO_LOGO_FILES) == 1


def _deck_in(src: Path, out: Path, ws: Path):
    """Колода без модели по демо-плану (structured) в рабочей папке `ws` (там же брендированная копия)."""
    from verstka.pipeline.generate import generate_variants
    from verstka.schemas.outline import DeckOutline

    d = json.loads((ROOT / "tests" / "fixtures" / "outline_demo.json").read_text(encoding="utf-8"))
    d.pop("_source", None)
    got = generate_variants(src, outline=DeckOutline.model_validate(d), strategies=["structured"], out_dir=out, workspace_root=ws,
                            use_llm=False, use_vlm=False, render_images=False, audit=False, autofix=False, exports=[], pictures=False, eyes=False)
    return got.variants[0], next(out.rglob("deck.pptx"))


# ============================================================================================== финальное ревью


def test_grey_template_text_keeps_its_tone(tmp_path, fonts_missing):
    """Основной текст шаблона серый (#666666 у lo_focus), текст бренда почти чёрный: серый остаётся серым (своя
    светлота, оттенок бренда) и это сказано, иначе подписи обложки темнеют и вёрстка обложки ломается (README C2).
    Тёмный цветной текст шаблона (бирюзовый lo_vivid) по-прежнему становится текстом бренда."""
    src = build_custom(tmp_path, "grey", bg="FFFFFF", text="666666", cards=[("E8A202", "666666")])
    m = _analyze(src, tmp_path / "ws")
    plan = ap.brand_plan(src, m, _vk_cloud()).colors
    assert plan.text_pair and plan.text_pair[0] == "666666"
    got = plan.map("666666")
    assert abs(pal.lightness(got) - pal.lightness("666666")) < 3
    assert any("#666666" in w and "серый" in w for w in plan.warnings)
    src2 = build_custom(tmp_path, "teal", bg="FFFFFF", text="04617B", cards=[("E8A202", "04617B")])
    plan2 = ap.brand_plan(src2, _analyze(src2, tmp_path / "ws2"), _vk_cloud()).colors
    assert plan2.map("04617B") == "0B1220"


def test_template_priority_conflicts_say_what_the_deck_gets(synthetic, fonts_missing, tmp_path):
    """«Шаблон важнее»: итог каждой строки расхождений — то, что запишет копия с этим приоритетом: слоты стандартной
    темы и стандартный шрифт занимает бренд, остальное — как в шаблоне."""
    src, m, bdir = synthetic
    brand = _brand()
    out = tmp_path / "tpl.pptx"
    res = apply_brand(src, out, brand, manifest=m, priority="template", brand_dir=bdir)
    rows = brand_conflicts(m, brand, priority="template", source=src, brand_dir=bdir)
    assert rows
    for c in rows:
        if c.kind == "color" and c.template.startswith("#"):
            written = res.colors.get(c.template)
            if written:
                assert c.result.startswith(written), (c.what, c.result, written)
            else:
                assert c.result.startswith(f"{c.template}, как в шаблоне"), (c.what, c.result)
        if c.kind == "font":
            assert ("как в шаблоне" in c.result) == (c.template not in res.font_swaps), (c.what, c.result, res.font_swaps)
        if c.kind == "logo":
            assert "логотип Облако" not in c.result or "как в шаблоне" in c.result


def test_brand_caches_survive_parallel_builds(tmp_path):
    """Кэши брендбука общие для потоков сервера: параллельное чтение и вытеснение не роняют сборку (финальное ревью
    I7)."""
    import threading

    files = [build_custom(tmp_path, f"c{i}", bg="FFFFFF", text="1F2937", cards=[("1F6FEB", "FFFFFF")], n=1) for i in range(6)]
    pngs = [_logo_png(tmp_path / f"l{i}.png", (i * 40, 119, 255, 255), (11, 18, 32, 255)).read_bytes() for i in range(6)]
    errors: list[BaseException] = []

    def work(k: int) -> None:
        try:
            for j in range(12):
                f = files[(k + j) % len(files)]
                pal.template_theme(f)
                pal.source_ir(f)
                lg.image_probe(pngs[(k + j) % len(pngs)])
        except BaseException as e:  # noqa: BLE001
            errors.append(e)

    threads = [threading.Thread(target=work, args=(k,)) for k in range(6)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert errors == [] and len(pal._IRS) <= pal.IR_CACHE_SIZE


def test_api_returns_conflicts_for_the_requested_priority(tmp_path, monkeypatch):
    """`GET /api/brandbooks/{id}?template_id=…&brand_priority=template` считает расхождения с приоритетом шаблона:
    `result` — то, что получит колода («…, как в шаблоне»), а не итог «брендбук важнее»."""
    import time

    from fastapi.testclient import TestClient

    import verstka.api.app as A
    from verstka.api.store import Store

    store = Store(tmp_path)
    monkeypatch.setattr(A, "store", store)
    monkeypatch.setattr(ap, "font_available", lambda f: f == "Play")
    tid = "cbbe3aa6a21d23ce"
    d = tmp_path / "templates" / tid
    d.mkdir(parents=True)
    (d / "manifest.json").write_bytes(gzip.open(FIX / f"manifest_{tid}.json.gz").read())
    (d / "source.pptx").write_bytes((FIX / "vktech_short.pptx").read_bytes())
    c = TestClient(A.app)
    r = c.post("/api/brandbooks", files={"file": ("Брендбук.json", (BRAND_EXAMPLE / "tokens.json").read_bytes(), "application/json")})
    bid, jid = r.json()["brandbook_id"], r.json()["job_id"]
    t0 = time.time()
    while time.time() - t0 < 60:
        job = A.runner.get(jid)
        if job is not None and job.events and job.events[-1].get("status") in ("done", "failed"):
            break
        time.sleep(0.05)
    brand_rows = {x["what"]: x["result"] for x in c.get(f"/api/brandbooks/{bid}?template_id={tid}").json()["conflicts"]}
    tpl_rows = {x["what"]: x["result"] for x in c.get(f"/api/brandbooks/{bid}?template_id={tid}&brand_priority=template").json()["conflicts"]}
    assert brand_rows["Шрифт заголовков"] == "VK Sans Display, замена — Play"
    assert tpl_rows["Шрифт заголовков"] == "Play, как в шаблоне"
    assert tpl_rows["Второй акцент"] == "#FF3885, как в шаблоне"
