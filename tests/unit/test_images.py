"""Pictures inside slides (verstka/images): the t2i backends read their hosts' answers, a chain skips links without a
key and links that cannot answer, the picture is graded to the template's palette and cropped to its place, the places
the author described get pictures, the author's own photo never does, and the free place comes back on request.
No live model is called: the hosts are replaced by fakes."""

from __future__ import annotations

import base64
import io
import json
import time

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from verstka.analysis.xmlns import q
from verstka.images import palette as P
from verstka.images import providers as IP
from verstka.images.fill import apply_pictures, plan_pictures, read_records, write_records
from verstka.images.places import CAPTION_PREFIX, places_in, take_picture_out
from verstka.images.prompt import ImageAsk, build_prompt, style_of
from verstka.images.service import ImageMaker
from verstka.schemas.outline import PHOTO_PLACE_NAME, DeckOutline, OutlineSlide, SlideContent


def _png(w=64, h=48, color=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


class _Resp:
    def __init__(self, status=200, obj=None, content=b"", headers=None):
        self.status_code = status
        self._obj = obj
        self.content = content or (json.dumps(obj).encode() if obj is not None else b"")
        self.text = self.content.decode("utf-8", "replace") if obj is not None else ""
        self.headers = headers or {}

    def json(self):
        return self._obj


def _polza(monkeypatch, *, create, statuses=(), picture=None, gets=None):
    """Polza's Media API faked: POST /media answers `create`, each GET /media/{id} the next of `statuses`, a GET of the
    picture's CDN link answers `picture`."""
    import httpx

    seen = {"posts": [], "gets": [] if gets is None else gets}
    pending = list(statuses)

    def fake_post(url, json=None, headers=None, timeout=None):
        seen["posts"].append({"url": url, "body": json, "auth": headers.get("Authorization")})
        return create if isinstance(create, _Resp) else _Resp(obj=create)

    def fake_get(url, headers=None, timeout=None, follow_redirects=None):
        seen["gets"].append(url)
        if "/media/" in url:
            nxt = pending.pop(0)
            return nxt if isinstance(nxt, _Resp) else _Resp(obj=nxt)
        r = _Resp(content=picture or _png())
        r.raise_for_status = lambda: None
        return r

    monkeypatch.setattr(httpx, "post", fake_post)
    monkeypatch.setattr(httpx, "get", fake_get)
    return seen


def test_polza_media_sends_the_prompt_and_aspect_and_reads_a_finished_picture(monkeypatch):
    done = {"id": "gen_1", "status": "completed", "data": [{"url": "https://s3.polza.ai/f/1.png"}], "usage": {"output_units": 1, "cost_rub": 1.4, "cost": 1.4}}
    seen = _polza(monkeypatch, create=done)
    p = IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "pza_test", label="Z-Image")
    r = p.generate("a greenhouse", aspect="16:9", seed=7)
    post = seen["posts"][0]
    assert post["url"] == "https://polza.ai/api/v1/media" and post["auth"] == "Bearer pza_test"
    # Z-Image takes no seed: the input is the prompt and one of Polza's aspect ratios
    assert post["body"] == {"model": "tongyi-mai/z-image", "input": {"prompt": "a greenhouse", "aspect_ratio": "16:9"}}
    assert seen["gets"] == ["https://s3.polza.ai/f/1.png"]  # no polling: the answer was ready
    assert r.data[:8] == b"\x89PNG\r\n\x1a\n" and r.cost_rub == 1.4 and r.label == "Z-Image" and r.seed == 7


def test_polza_media_polls_a_pending_task_until_the_picture_is_ready(monkeypatch):
    import time as _time

    monkeypatch.setattr(_time, "sleep", lambda s: None)
    seen = _polza(monkeypatch, create={"id": "gen_2", "status": "pending"}, statuses=[
        {"id": "gen_2", "status": "processing"},
        {"id": "gen_2", "status": "completed", "data": {"url": "https://s3.polza.ai/f/2.png"}, "usage": {"cost_rub": 1.4}},
    ])
    r = IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "pza_test").generate("x", aspect="4:3")
    assert seen["gets"] == ["https://polza.ai/api/v1/media/gen_2", "https://polza.ai/api/v1/media/gen_2", "https://s3.polza.ai/f/2.png"]
    assert r.cost_rub == 1.4  # the status's data may be one object as well as a list
    # another endpoint's usage.cost is in its own money: only cost_rub counts there
    from verstka.images.providers import _usage_rub
    assert _usage_rub({"usage": {"cost": 0.04}}) == 0.04 and _usage_rub({"usage": {"cost": 0.04}}, polza=False) is None


def test_polza_media_reports_a_failed_task_and_a_task_that_never_ends(monkeypatch):
    import time as _time

    monkeypatch.setattr(_time, "sleep", lambda s: None)
    _polza(monkeypatch, create={"id": "g", "status": "pending"}, statuses=[{"id": "g", "status": "failed", "error": {"message": "NSFW"}}])
    with pytest.raises(IP.ImageError, match="failed .NSFW"):
        IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "k").generate("x")
    _polza(monkeypatch, create={"id": "g", "status": "pending"}, statuses=[{"id": "g", "status": "processing"}] * 50)
    clock = iter(range(0, 10_000, 30))
    monkeypatch.setattr(_time, "monotonic", lambda: float(next(clock)))
    with pytest.raises(IP.ImageError, match="not ready in time"):
        IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "k", timeout_s=90).generate("x")


def test_polza_media_never_polls_past_the_generations_deadline(monkeypatch):
    import time as _time

    now, slept = [1000.0], []
    monkeypatch.setattr(_time, "monotonic", lambda: now[0])

    def nap(s):
        slept.append(s)
        now[0] += s

    monkeypatch.setattr(_time, "sleep", nap)
    _polza(monkeypatch, create={"id": "g", "status": "pending"}, statuses=[{"id": "g", "status": "processing"}] * 50)
    with pytest.raises(IP.ImageError, match="not ready in time|no time left"):
        IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "k", timeout_s=90).generate("x", deadline=1010.0)
    assert slept and now[0] <= 1010.0  # the link's 90 s would allow more; the generation's deadline does not


def test_the_pictures_download_fits_the_generations_deadline(monkeypatch):
    import time as _time

    import httpx

    monkeypatch.setattr(_time, "monotonic", lambda: 1000.0)
    done = {"id": "g", "status": "completed", "data": [{"url": "https://s3.polza.ai/f/1.png"}], "usage": {"cost_rub": 1.4}}
    monkeypatch.setattr(httpx, "post", lambda url, json=None, headers=None, timeout=None: _Resp(obj=done))
    timeouts: list = []

    def cdn(url, headers=None, timeout=None, follow_redirects=None):
        timeouts.append(timeout)
        r = _Resp(content=_png())
        r.raise_for_status = lambda: None
        return r

    monkeypatch.setattr(httpx, "get", cdn)
    r = IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "k", timeout_s=90).generate("x", deadline=1012.0)
    assert r.cost_rub == 1.4 and timeouts == [12.0]  # the CDN's own 30 s would run past the deadline


def test_polza_media_keeps_the_prompt_within_polzas_limit_and_takes_the_nearest_aspect(monkeypatch):
    done = {"id": "g", "status": "completed", "data": [{"url": "https://s3.polza.ai/f/3.png"}]}
    seen = _polza(monkeypatch, create=done)
    IP.PolzaMedia("tongyi-mai/z-image", "https://polza.ai/api/v1", "k").generate("y" * 1500, aspect="21:9")
    body = seen["posts"][0]["body"]["input"]
    assert len(body["prompt"]) == IP.PROMPT_MAX_CHARS and body["aspect_ratio"] == "16:9"


def test_picture_of_reads_urls_and_chat_answers():
    data = _png()
    url = "data:image/png;base64," + base64.b64encode(data).decode()
    assert IP.picture_of({"data": [{"url": url}]}) == data
    assert IP.picture_of({"data": {"url": url}}) == data  # Polza's media status documents one object
    assert IP.picture_of({"choices": [{"message": {"images": [{"type": "image_url", "image_url": {"url": url}}]}}]}) == data
    assert IP.picture_of({"data": [{"url": "https://x/y.png"}]}, fetch=lambda u: data) == data
    assert IP.picture_of({"data": []}) is None


def test_a_chain_skips_links_without_a_key_and_holds_a_link_without_money(monkeypatch):
    import httpx

    calls = []

    def fake_post(url, json=None, headers=None, timeout=None):
        calls.append(json["model"])
        if json["model"] == "broke":
            return _Resp(status=402, obj={"error": {"code": "INSUFFICIENT_BALANCE", "message": "Недостаточно средств на балансе"}})
        return _Resp(obj={"id": "g", "status": "completed", "data": [{"url": "data:image/png;base64," + base64.b64encode(_png()).decode()}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    chain = IP.build_image_provider({"backend": "polza_media", "chain": [
        {"model": "nokey", "api_key": ""}, {"model": "broke", "api_key": "k"}, {"model": "ok", "api_key": "k"}]})
    assert [d["state"] for d in chain.describe()] == ["off", "ready", "ready"]
    r = chain.generate("x", aspect="1:1")
    assert r.model == "ok" and calls == ["broke", "ok"]
    chain.generate("y", aspect="1:1")
    assert calls == ["broke", "ok", "ok"]  # the link without money is skipped for a while
    assert chain.describe()[1]["state"] == "no_credits"


def _seq_post(monkeypatch, answers):
    """POST /media answers the next of `answers` each time (a _Resp, or a media object); returns the bodies sent."""
    import httpx

    sent, queue = [], list(answers)

    def fake_post(url, json=None, headers=None, timeout=None):
        sent.append(json)
        nxt = queue.pop(0) if len(queue) > 1 else queue[0]
        return nxt if isinstance(nxt, _Resp) else _Resp(obj=nxt)

    monkeypatch.setattr(httpx, "post", fake_post)
    return sent


_DONE = {"id": "g", "status": "completed", "data": [{"url": "data:image/png;base64," + base64.b64encode(_png()).decode()}], "usage": {"cost_rub": 1.4}}
_MINUTE = _Resp(status=429, obj={"error": {"code": "TOO_MANY_REQUESTS", "message": "Достигнут минутный лимит запросов к модели. Повторите позже."}}, headers={"retry-after": "40"})
_DAY = _Resp(status=429, obj={"error": {"code": "TOO_MANY_REQUESTS", "message": "Достигнут дневной лимит запросов к модели. Повторите позже."}}, headers={"retry-after": "7200"})
_HOSTS = _Resp(status=429, obj={"error": {"code": "TOO_MANY_REQUESTS", "message": "Провайдер модели ограничил частоту запросов."}})
_RESERVE = _Resp(status=402, obj={"error": {"code": "INSUFFICIENT_BALANCE", "message": "Недостаточно средств: на балансе 1.20 ₽, зарезервировано 4.20 ₽"}})


@pytest.fixture
def naps(monkeypatch):
    """time.sleep of the picture chain recorded instead of slept."""
    import time as _time

    slept: list[float] = []
    monkeypatch.setattr(_time, "sleep", lambda s: slept.append(round(s, 1)))
    return slept


def _one(model="tongyi-mai/z-image"):
    return IP.build_image_provider({"backend": "polza_media", "model": model, "api_key": "k"})


def test_the_only_picture_model_rides_out_polzas_minute_cap(monkeypatch, naps):
    sent = _seq_post(monkeypatch, [_MINUTE, _DONE])
    chain = _one()
    assert chain.generate("x").cost_rub == 1.4 and len(sent) == 2
    assert naps == [40.0] and chain.describe()[0]["state"] == "ready"  # waited the window out, never held


def test_the_hosts_429_gets_a_pause_and_two_in_a_row_hold_the_model(monkeypatch, naps):
    sent = _seq_post(monkeypatch, [_HOSTS, _DONE])
    chain = _one()
    chain.generate("x")
    assert len(sent) == 2 and naps == [2.0] and chain.describe()[0]["state"] == "ready"
    _seq_post(monkeypatch, [_HOSTS])
    with pytest.raises(IP.ImageError):
        chain.generate("y")
    assert chain.describe()[0]["state"] == "congested"  # two in a row: a short hold


def test_a_reserve_402_waits_for_the_parallel_pictures_and_never_holds(monkeypatch, naps):
    sent = _seq_post(monkeypatch, [_RESERVE, _RESERVE, _DONE])
    chain = _one()
    assert chain.generate("x").cost_rub == 1.4 and len(sent) == 3
    assert naps == [8.0, 8.0] and chain.describe()[0]["state"] == "ready"


def test_polzas_daily_cap_holds_the_picture_model_until_its_reset(monkeypatch, naps):
    _seq_post(monkeypatch, [_DAY])
    chain = _one()
    with pytest.raises(IP.ImageError):
        chain.generate("x")
    assert chain.describe()[0]["state"] == "quota" and naps == []
    assert 7190 < chain.holds[0].until - time.time() <= 7200  # its Retry-After: the cap resets then
    with pytest.raises(IP.ImageError, match="skipped"):
        chain.generate("y")  # no request until the cap resets


def test_a_wait_never_runs_past_the_generations_deadline(monkeypatch, naps):
    import time as _time

    sent = _seq_post(monkeypatch, [_MINUTE, _DONE])
    chain = _one()
    with pytest.raises(IP.ImageError):
        chain.generate("x", deadline=_time.monotonic() + 30)  # 40 s window + a picture's time do not fit
    assert len(sent) == 1 and naps == [] and chain.describe()[0]["state"] == "rate"


def test_a_short_hold_of_the_only_link_is_waited_out_by_the_next_picture(monkeypatch, naps):
    import time as _time

    chain = _one()
    chain._hold(0, "rate", 20.0, "a minute window")
    sent = _seq_post(monkeypatch, [_DONE])
    assert chain.generate("x").cost_rub == 1.4 and len(sent) == 1
    assert len(naps) == 1 and 15 < naps[0] <= 20
    chain._hold(0, "rate", 20.0, "a minute window")
    with pytest.raises(IP.ImageError, match="skipped"):
        chain.generate("y", deadline=_time.monotonic() + 25)  # no room to wait: skipped


def test_status_polling_survives_a_404_race_and_a_5xx(monkeypatch, naps):
    statuses = [_Resp(status=404, obj={"error": {"code": "NOT_FOUND", "message": "Not found"}}), _Resp(status=502, obj={"error": {"code": "BAD_GATEWAY", "message": "x"}}),
                {"id": "g9", "status": "completed", "data": [{"url": "https://s3.polza.ai/f/9.png"}], "usage": {"cost_rub": 1.4}}]
    seen = _polza(monkeypatch, create={"id": "g9", "status": "pending"}, statuses=statuses)
    chain = _one()
    assert chain.generate("x").cost_rub == 1.4
    assert len(seen["posts"]) == 1  # one task, paid once: a failed status request is asked again, not a new task
    assert chain.describe()[0]["state"] == "ready"


def test_a_failed_task_of_the_hosts_passing_trouble_is_asked_again(monkeypatch, naps):
    sent = _seq_post(monkeypatch, [{"id": "a", "status": "failed", "error": {"code": "BAD_GATEWAY", "message": "upstream"}}, _DONE])
    assert _one().generate("x").cost_rub == 1.4 and len(sent) == 2


def test_a_picture_the_cdn_does_not_give_is_an_image_error_without_a_hold(monkeypatch, naps):
    import httpx

    seen = _polza(monkeypatch, create={"id": "g", "status": "completed", "data": [{"url": "https://s3.polza.ai/f/x.png"}]})

    def fake_get(url, headers=None, timeout=None, follow_redirects=None):
        seen["gets"].append(url)
        return httpx.Response(404, request=httpx.Request("GET", url))

    monkeypatch.setattr(httpx, "get", fake_get)
    chain = _one()
    with pytest.raises(IP.ImageError, match="could not be downloaded"):
        chain.generate("x")
    assert chain.describe()[0]["state"] == "ready"  # this picture is lost, the next one is asked as usual


def test_a_chain_without_any_key_is_off():
    chain = IP.build_image_provider({"backend": "polza_media", "chain": [{"model": "a", "api_key": ""}]})
    assert chain.off
    with pytest.raises(IP.ImageError):
        chain.generate("x")


def test_the_default_config_draws_with_z_image_on_polza(monkeypatch):
    from pathlib import Path

    monkeypatch.setenv("POLZA_API_KEY", "pza_test")
    IP.reset_cache()
    try:
        chain = IP.load_image_provider(Path(__file__).resolve().parents[2] / "configs" / "models.yaml")
        assert [(lk.name, lk.model, lk.label, lk.off) for lk in chain.links] == [("polza_media", "tongyi-mai/z-image", "Z-Image", False)]
        assert chain.links[0].base_url == "https://polza.ai/api/v1"
    finally:
        IP.reset_cache()
    monkeypatch.delenv("POLZA_API_KEY")
    try:
        assert IP.load_image_provider(Path(__file__).resolve().parents[2] / "configs" / "models.yaml").off
    finally:
        IP.reset_cache()


def test_the_images_module_follows_the_same_profile_resolver(monkeypatch, tmp_path):
    """load_image_provider() without a path asks registry.default_models_path(), the one function that picks
    $VERSTKA_MODELS, the VK profile (both VK variables) or configs/models.yaml."""
    from verstka.providers import registry as R

    mine = tmp_path / "mine.yaml"
    mine.write_text("roles:\n  t2i: {backend: mock, model: m, label: Mine}\nlimits: {}\n", encoding="utf-8")
    seen: list[str] = []
    real = R.default_models_path
    monkeypatch.setattr(R, "default_models_path", lambda: (seen.append("asked"), mine)[1])
    IP.reset_cache()
    try:
        assert [lk.label for lk in IP.load_image_provider().links] == ["Mine"] and seen == ["asked"]
    finally:
        IP.reset_cache()
    monkeypatch.setattr(R, "default_models_path", real)  # the real resolver again, driven by the environment alone
    monkeypatch.setenv("POLZA_API_KEY", "pza_test")
    monkeypatch.delenv("VERSTKA_MODELS", raising=False)
    monkeypatch.setenv("VK_INFERENCE_API_KEY", "vk-test")
    monkeypatch.setenv("VK_PROJECT_ID", "PRJ-test")
    IP.reset_cache()
    try:
        vk_chain = IP.load_image_provider()
        assert vk_chain is IP.load_image_provider(R.default_models_path())  # the cache is keyed by that very file
        assert R.default_models_path().name == "models.vk.yaml"
        monkeypatch.setenv("VERSTKA_MODELS", "configs/models.yaml")
        assert IP.load_image_provider() is not vk_chain and R.default_models_path().name == "models.yaml"
    finally:
        IP.reset_cache()


def test_the_vk_profile_draws_with_z_image_on_polza(monkeypatch):
    """VK inference has no text-to-image model: the VK profile's t2i role is the same Z-Image on Polza as the default
    profile's; without POLZA_API_KEY it is off; VK_T2I_MODEL puts VK's own model first, Z-Image answering behind it."""
    from pathlib import Path

    vk = Path(__file__).resolve().parents[2] / "configs" / "models.vk.yaml"
    default = Path(__file__).resolve().parents[2] / "configs" / "models.yaml"
    monkeypatch.setenv("VK_INFERENCE_API_KEY", "vk-test")
    monkeypatch.setenv("VK_PROJECT_ID", "PRJ-test")
    monkeypatch.setenv("POLZA_API_KEY", "pza_test")
    IP.reset_cache()
    try:
        chain = IP.load_image_provider(vk)
        live = [lk for lk in chain.links if not lk.off]
        assert len(live) == 1 and not chain.off
        z = live[0]
        assert (z.name, z.model, z.label, z.base_url, z.api_key, z.timeout_s) == ("polza_media", "tongyi-mai/z-image", "Z-Image", "https://polza.ai/api/v1", "pza_test", 90.0)
        ref = IP.load_image_provider(default).links[0]  # the very link of the default profile
        assert (ref.name, ref.model, ref.label, ref.base_url, ref.timeout_s) == (z.name, z.model, z.label, z.base_url, z.timeout_s)
        assert [d["state"] for d in chain.describe()] == ["off", "ready"]  # VK_T2I_MODEL unset: VK's own link is never asked
        with pytest.raises(IP.ImageError, match="no API key or model"):
            chain.links[0]._post("https://example.invalid/x", {}, None)
    finally:
        IP.reset_cache()
    # no Polza key: no pictures, as before (the slides keep their free places)
    monkeypatch.delenv("POLZA_API_KEY")
    try:
        assert IP.load_image_provider(vk).off
    finally:
        IP.reset_cache()
    # VK offers its own model: its link is asked first, on the same endpoint with the project header; Z-Image is behind it
    monkeypatch.setenv("POLZA_API_KEY", "pza_test")
    monkeypatch.setenv("VK_T2I_MODEL", "open-t2i-model")
    try:
        first, second = IP.load_image_provider(vk).links
        assert (first.name, first.model, first.off) == ("openai_images", "open-t2i-model", False)
        assert first.base_url.endswith("/inference-providers-router/00000000-0000-0000-0000-000000000001/v1") and first.api_key == "vk-test"
        assert first.headers == {"Project-ID": "PRJ-test"} and (second.name, second.model, second.off) == ("polza_media", "tongyi-mai/z-image", False)
    finally:
        IP.reset_cache()


def test_nearest_aspect():
    assert IP.nearest_aspect(1600, 900) == "16:9"
    assert IP.nearest_aspect(900, 1200) == "3:4"
    assert IP.nearest_aspect(1000, 1000) == "1:1"
    assert IP.nearest_aspect(2100, 900) == "16:9"  # Polza's Z-Image has no 21:9: the widest it takes


def test_grading_brings_a_picture_to_the_palette_and_fit_keeps_proportions():
    pal = P.Palette(dark="0B1F44", light="F4F6FA", accents=["FF6A00"])
    img = Image.new("RGB", (300, 200), (40, 200, 60))  # a green the palette does not have
    before = P.palette_match(img, pal)
    after = P.palette_match(P.grade(img, pal, 0.8), pal)
    assert after > before
    assert P.grade(img, pal, 0.0).getpixel((0, 0)) == (40, 200, 60)
    out = P.fit_to(Image.new("RGB", (1600, 900)), 3, 4, max_side=800)
    assert abs(out.size[0] / out.size[1] - 0.75) < 0.01 and max(out.size) <= 800


def test_colour_names_read_as_words():
    assert P.color_name("FFFFFF") == "white"
    assert "orange" in P.color_name("FF6A00")
    assert "blue" in P.color_name("0B1F44")


def test_the_prompt_carries_subject_style_palette_and_no_text():
    pal = P.Palette(dark="0B1F44", light="F4F6FA", accents=["FF6A00"])
    prompt, style = build_prompt(ImageAsk(hint="Экран панели управления теплицей", headline="Панель наблюдения", deck="Умная теплица"), pal)
    assert style == "screen"
    assert "Экран панели управления теплицей" in prompt and "orange" in prompt and "no text of any kind" in prompt
    # Z-Image draws what it is given as text: no headline, no deck title, no quotes, no colour codes in the prompt
    assert "Панель наблюдения" not in prompt and "Умная теплица" not in prompt and "«" not in prompt and "#" not in prompt
    p2, _ = build_prompt(ImageAsk(hint="Вид теплицы сверху", headline="Теплица", deck="Т", wish="больше зелени, утро"), pal)
    assert p2.startswith("больше зелени, утро") and "Вид теплицы сверху" in p2
    assert style_of("схема полива", {"styles": [{"match": "схем", "style": "illustration"}], "default_style": "photo"}) == "illustration"


def _deck_with_places(tmp_path):
    """Two slides, each with a free place and its caption, as the renderer leaves them."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(6), Inches(1)).text_frame.text = f"Слайд {i + 1}"
        frame = s.shapes.add_shape(5, Inches(7), Inches(1.5), Inches(5.5), Inches(4.5))  # a rounded rectangle
        frame.name = PHOTO_PLACE_NAME
        cap = s.shapes.add_textbox(Inches(7.5), Inches(3.2), Inches(4.5), Inches(1))
        cap.text_frame.text = "Место под изображение"
        cap.name = f"{CAPTION_PREFIX} {cap.shape_id}"
        over = s.shapes.add_shape(1, Inches(6.8), Inches(1.3), Inches(0.5), Inches(0.5))  # something drawn over the place
        over.name = "Over"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    outline = DeckOutline(title="Умная теплица", slides=[
        OutlineSlide(id="a", kind="image_text", headline="Панель наблюдения", content=SlideContent(bullets=["x"], photo_slot="Изображение", photo_hint="Экран панели управления")),
        OutlineSlide(id="b", kind="image_text", headline="Помещение", content=SlideContent(bullets=["y"], photo_slot="Фото помещения", photo_own=True)),
    ])
    return path, outline


def test_places_are_found_with_their_captions_and_owners(tmp_path):
    path, outline = _deck_with_places(tmp_path)
    places = places_in(Presentation(str(path)), outline)
    assert [(p.slide, p.own, len(p.caption_ids)) for p in places] == [(1, False, 1), (2, True, 1)]
    assert places[0].hint == "Экран панели управления" and places[0].headline == "Панель наблюдения"


def test_the_described_place_gets_a_picture_and_the_own_photo_place_stays_free(tmp_path):
    path, outline = _deck_with_places(tmp_path)
    mock = IP.MockImages()
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "cache", deck_title=outline.title)
    pending = plan_pictures(path, outline, maker)
    res = apply_pictures(path, pending, wait_s=10)
    write_records(tmp_path, res.records)
    maker.close()
    assert res.changed == [1] and len(mock.calls) == 1
    assert "Экран панели управления" in mock.calls[0]["prompt"]
    prs = Presentation(str(path))
    tree1 = prs.slides[0]._element.cSld.find(q("p:spTree"))
    names = [el.find(".//" + q("p:cNvPr")).get("name") for el in tree1 if el.find(".//" + q("p:cNvPr")) is not None]
    pics = [el for el in tree1 if etree.QName(el).localname == "pic"]
    assert len(pics) == 1 and not any(n.startswith(CAPTION_PREFIX) for n in names) and PHOTO_PLACE_NAME not in names
    # the picture keeps the place's position under what was drawn over it; the place in the right column runs out
    # to the slide's right edge and bottom (rendering/picture_layouts.py: a bleed is square-cornered)
    assert names.index("Over") > [n for n in names].index(pics[0].find(".//" + q("p:cNvPr")).get("name"))
    assert pics[0].find(".//" + q("a:prstGeom")).get("prst") == "rect"
    off, ext = pics[0].find(".//" + q("a:off")), pics[0].find(".//" + q("a:ext"))
    assert int(off.get("x")) + int(ext.get("cx")) == prs.slide_width and int(off.get("y")) + int(ext.get("cy")) == prs.slide_height
    assert "сгенерировано" in pics[0].find(".//" + q("p:cNvPr")).get("descr")
    tree2 = prs.slides[1]._element.cSld.find(q("p:spTree"))
    assert not [el for el in tree2 if etree.QName(el).localname == "pic"]
    recs = read_records(tmp_path)
    assert [(r["slide"], r["status"]) for r in recs] == [(1, "generated"), (2, "free")]
    assert recs[1]["why"].startswith("место под ваше фото")
    # «Убрать изображение»: the free place and its caption come back
    rec = recs[0]
    assert take_picture_out(prs.slides[0], rec["shape_id"], rec["kept"])
    names = [el.find(".//" + q("p:cNvPr")).get("name") for el in tree1 if el.find(".//" + q("p:cNvPr")) is not None]
    assert PHOTO_PLACE_NAME in names and any(n.startswith(CAPTION_PREFIX) for n in names)


def test_without_a_model_the_places_stay_free(tmp_path):
    path, outline = _deck_with_places(tmp_path)
    pending = plan_pictures(path, outline, None)
    res = apply_pictures(path, pending)
    assert res.changed == [] and {r["status"] for r in res.records} == {"free"}


def test_the_variants_share_a_picture_and_a_failing_model_leaves_the_place(tmp_path):
    path, outline = _deck_with_places(tmp_path)
    mock = IP.MockImages()
    maker = ImageMaker(IP.ImageChain([mock]), None, tmp_path / "cache", deck_title=outline.title)
    f1 = maker.submit(ImageAsk(hint="h", headline="x", deck="d"), "4:3")
    f2 = maker.submit(ImageAsk(hint="h", headline="x", deck="d"), "4:3")
    assert f1 is f2
    f1.result(timeout=5)
    assert len(mock.calls) == 1
    maker.close()
    bad = ImageMaker(IP.ImageChain([IP.MockImages(fail=True)]), None, tmp_path / "cache2", deck_title=outline.title)
    res = apply_pictures(path, plan_pictures(path, outline, bad), wait_s=5)
    bad.close()
    assert res.changed == [] and res.records[0]["why"] == "модель изображений не ответила"


# ---------------------------------------------------------------------------- the whole path (no live model)

GREENHOUSE = """Тема презентации: «Умная теплица». Презентация на 3 слайда.

Слайд 1. Умная теплица

Автоматический полив и контроль климата для учебной теплицы.

Слайд 2. Панель наблюдения

Панель управления объединяет текущие показания датчиков. Пользователь быстро видит, какие зоны требуют внимания.

[Место под изображение: экран панели управления теплицей с показаниями датчиков]

Слайд 3. Помещение

Теплица площадью 40 м² стоит во дворе школы. Внутри — шесть грядок и бак для воды на 200 литров.

На этом слайде обязательно оставь свободное место под фотографию помещения. Фото я добавлю самостоятельно. Не заполняй это место текстом, диаграммой или сгенерированным изображением.
"""


def _pics_and_places(slide):
    tree = slide._element.cSld.find(q("p:spTree"))
    names = [el.find(".//" + q("p:cNvPr")).get("name") for el in tree if el.find(".//" + q("p:cNvPr")) is not None]
    return [n for n in names if n.startswith("Изображение")], [n for n in names if n == PHOTO_PLACE_NAME]


def test_a_build_puts_pictures_in_the_described_places_only(tmp_path):
    from pathlib import Path as _P

    from verstka.pipeline.generate import generate_variants
    from verstka.planning.brief import load_brief

    tpl = _P(__file__).resolve().parents[1] / "fixtures" / "tpl_audit" / "templates" / "lo_focus.pptx"
    p = tmp_path / "brief.md"
    p.write_text(GREENHOUSE, encoding="utf-8")
    mock = IP.MockImages()
    g = generate_variants(tpl, brief=load_brief(p), strategies=["structured", "visual"], out_dir=tmp_path / "run", workspace_root=tmp_path / "ws",
                          use_llm=False, use_vlm=False, audit=True, exports=[], image_provider=IP.ImageChain([mock]))
    assert mock.calls and all("экран панели управления теплицей" in c["prompt"].lower() for c in mock.calls)
    assert len(mock.calls) == 1  # the variants share the place's picture
    for v in g.variants:
        prs = Presentation(str(v.out_dir / "deck.pptx"))
        by_spec = {s.spec_ref: i for i, s in enumerate(v.outline.slides)}
        pics2, places2 = _pics_and_places(prs.slides[by_spec[2]])
        pics3, places3 = _pics_and_places(prs.slides[by_spec[3]])
        assert len(pics2) == 1 and not places2, v.strategy
        assert not pics3 and len(places3) == 1, v.strategy  # the author's own photo: the place stays free
        recs = read_records(v.out_dir)
        assert sorted(r["status"] for r in recs) == ["free", "generated"]
        assert next(r for r in recs if r["status"] == "free")["own"] is True
        rm = json.loads((v.out_dir / "run_manifest.json").read_text(encoding="utf-8"))
        assert rm["pictures"]["generated"] == 1 and rm["pictures"]["own"] == 1
    assert any((tmp_path / "run" / "images").glob("*.jpg"))


def test_a_build_with_pictures_off_keeps_every_place_free(tmp_path):
    from pathlib import Path as _P

    from verstka.pipeline.generate import generate_variants
    from verstka.planning.brief import load_brief

    tpl = _P(__file__).resolve().parents[1] / "fixtures" / "tpl_audit" / "templates" / "lo_focus.pptx"
    p = tmp_path / "brief.md"
    p.write_text(GREENHOUSE, encoding="utf-8")
    mock = IP.MockImages()
    g = generate_variants(tpl, brief=load_brief(p), strategies=["structured"], out_dir=tmp_path / "run", workspace_root=tmp_path / "ws",
                          use_llm=False, use_vlm=False, audit=True, exports=[], image_provider=IP.ImageChain([mock]), pictures=False)
    assert not mock.calls
    recs = read_records(g.variants[0].out_dir)
    assert {r["status"] for r in recs} == {"free"} and any(r["why"] == "генерация изображений выключена" for r in recs)


def test_a_caption_hanging_out_of_its_trimmed_place_goes_with_it(tmp_path):
    """The renderer may trim a frame after captioning it (aligned to the content's foot): the caption then hangs below
    the frame — it is still the place's, and goes when the picture comes (VK Tech, slide «Панель наблюдения»)."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    frame = s.shapes.add_shape(1, Inches(5), Inches(1.5), Inches(3.6), Inches(1.6))
    frame.name = PHOTO_PLACE_NAME
    cap = s.shapes.add_textbox(Inches(5.3), Inches(2.85), Inches(3), Inches(0.55))  # its centre below the frame
    cap.name = f"{CAPTION_PREFIX} {cap.shape_id}"
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="a", kind="image_text", headline="H", content=SlideContent(photo_slot="Изображение", photo_hint="экран"))])
    (place,) = places_in(prs, outline)
    assert place.caption_ids == [str(cap.shape_id)]


def test_a_bracketed_note_of_several_sentences_describes_a_picture_not_the_authors_photo():
    """«[Место под изображение: вид сверху. На изображении — три грядки. Подписи показывают датчики.]» is cut into
    sentences; none of them is the author's request for their own photo — the place gets a generated picture (the
    live run of 2026-10-06 left both described places of the greenhouse brief free). The author's own words still
    keep a place for their photo."""
    from verstka.planning.brief_structure import read_structure

    brief = (
        "Слайд 1. Проект «Теплица»\n\nПлан на год\n\n"
        "Слайд 2. Наша команда\n\nКоманда из пяти человек работает над проектом второй год. Фото команды я добавлю "
        "самостоятельно, не заполняй это место сгенерированным изображением.\n\n"
        "Слайд 3. Схема\n\nСистема состоит из датчиков и контроллера.\n\n"
        "[Место под изображение: вид теплицы сверху. На изображении — три грядки, бак и насос. Подписи показывают "
        "датчики.]\n\n"
        "Слайд 4. Спасибо!\n"
    )
    specs = read_structure(brief).specs
    assert [(s.photo is not None, s.photo_own) for s in specs] == [(False, False), (True, True), (True, False), (False, False)]
