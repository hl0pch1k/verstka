"""Native chart design rules (verstka.rendering.charts): template palette, muted bars, labels, layout, audit."""

import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from verstka.audit.checks.integrity import chart_missing_labels
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import AuditContext
from verstka.export.svg_charts import chart_svg, format_value
from verstka.rendering.charts import (
    _chroma,
    _is_data_color,
    _lightness,
    _rule_color,
    _seg_in_box,
    add_chart,
    chart_palette,
    delta_e,
    doughnut_hole_bbox,
    effective_chart_type,
    label_format,
    muted_tint,
    number_format,
    place_point_labels,
    prefer_bar,
    series_roles,
    short_unit,
    slide_ground,
    unit_caption,
)
from verstka.schemas.common import Bbox, contrast_ratio, relative_luminance
from verstka.schemas.outline import ChartSpec, DeckOutline, Series
from verstka.schemas.template import ChartStyleSpec, TemplateManifest, TypeStep, Typography

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
W, H = 12192000, 6858000
BOX = Bbox(x=int(W * 0.05), y=int(H * 0.25), w=int(W * 0.55), h=int(H * 0.6))
# LCT-like manifest colours: the first two are the same pink (the audit's «grey or near-identical series» case)
STYLE = ChartStyleSpec(series_colors=["FE095F", "FF0053", "753A93", "976BAE"], font_family="Montserrat", font_size_pt=12.0)
BOLD = Typography(scale=[TypeStep(role="display", size_pt=36, weight_bold_share=1.0), TypeStep(role="h1", size_pt=20, weight_bold_share=0.66), TypeStep(role="body", size_pt=14), TypeStep(role="small", size_pt=12)])
MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн"]
OUTLINE = DeckOutline(
    title="t",
    series=[
        Series(id="cc", name="Обращений через подсказчик, тыс.", categories=["Июль", "Август", "Сентябрь"], values=[12, 31, 58], unit="тыс."),
        Series(id="ws", name="Активные пользователи", categories=["Май", "Июнь", "Июль", "Август", "Сентябрь"], values=[1200, 3400, 6100, 9800, 12400], unit="чел."),
        Series(id="long", name="Экономия", categories=["Автоматизация обращений", "Подсказки", "Маршрутизация"], values=[42, 0, 12.5], unit="млн ₽ в год"),
        Series(id="a", name="2024", categories=["Q1", "Q2", "Q3"], values=[30, 42, 51], unit="%"),
        Series(id="b", name="2025", categories=["Q1", "Q2", "Q3"], values=[38, 49, 63], unit="%"),
        Series(id="c", name="План", categories=["Q1", "Q2", "Q3"], values=[40, 50, 60], unit="%"),
        Series(id="web", name="Web", categories=["Q1", "Q2", "Q3"], values=[12, 15, 19], unit="тыс."),
        Series(id="mob", name="Mobile", categories=["Q1", "Q2", "Q3"], values=[8, 11, 16], unit="тыс."),
        Series(id="was", name="Было", categories=["Время ответа оператора", "Повторные обращения"], values=[48, 22], unit="мин"),
        Series(id="now", name="Стало", categories=["Время ответа оператора", "Повторные обращения"], values=[31, 14], unit="мин"),
        Series(id="ln", name="NPS", categories=["Янв", "Фев", "Мар", "Апр"], values=[31.5, 34.2, 33.8, 38.1], unit="п.п."),
        Series(id="rise", name="NPS", categories=MONTHS, values=[31.5, 34.2, 33.8, 38.1, 41.0, 44.6], unit="п.п."),
        Series(id="m12", name="Выручка", categories=["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"], values=[110, 118, 125, 131, 128, 140, 152, 149, 163, 171, 180, 196], unit="млн ₽"),
        Series(id="pie", name="Каналы", categories=["Чат", "Телефон", "Почта"], values=[460, 310, 230], unit="тыс."),
        Series(id="dn", name="Бюджет", categories=["Разработка", "Инфраструктура", "Маркетинг", "Поддержка", "Прочее"], values=[420, 260, 150, 110, 60], unit="млн ₽"),
        Series(id="tiny", name="Доли", categories=["Основной канал", "Резерв", "Тест"], values=[92.4, 7.0, 0.6], unit="%"),
        Series(id="zero", name="Сбои", categories=["Q1", "Q2", "Q3"], values=[3, 0, 2], unit="шт."),
        Series(id="yrs", name="Клиенты", categories=[str(y) for y in range(2016, 2026)], values=list(range(1, 11)), unit="тыс."),
        Series(id="mon", name="Выручка", categories=["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август"], values=list(range(1, 9)), unit="млн ₽"),
    ],
)


def _slide(bg: str = "FFFFFF", w: int = W, h: int = H):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(w), Emu(h)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return prs, s


def _fill(el) -> str:
    return el.find(f"{C}spPr/{A}solidFill/{A}srgbClr").get("val")


def _chart(spec, bg="FFFFFF", typo=None, **kw):
    prs, s = _slide(bg)
    gf = add_chart(s, BOX, spec, OUTLINE, STYLE, typo or Typography(), text_hex=kw.pop("text_hex", "1A1A1A"), **kw)
    return gf.chart._chartSpace


def _point_labels(ser) -> dict:
    return {int(d.find(f"{C}idx").get("val")): d for d in ser.iter(f"{C}dLbl")}


def _ser_fmt(ser) -> str:
    return ser.find(f"{C}dLbls/{C}numFmt").get("formatCode")


# ---------------------------------------------------------------------------------------------- palette


def test_palette_is_distinct_visible_and_led_by_the_accent():
    for ground in ("F4F5F6", "702D8D", "000000"):
        pal = chart_palette(["FE095F", "FF0053", "753A93", "976BAE", "FFD6E3"], ground, 3, extra=["8A83D1", "520977"])
        assert pal[0] == "FE095F"
        assert "FF0053" not in pal  # a near-duplicate of the accent never becomes a second series
        assert all(contrast_ratio(c, ground) >= 1.4 for c in pal)
        assert all(delta_e(a, b) >= 20 for a, b in zip(pal, pal[1:]))
    # more series than the template has colours: tints/shades fill in, never a failure
    pal = chart_palette(["0077FF"], "FFFFFF", 4)
    assert len(pal) == 4 and all(contrast_ratio(c, "FFFFFF") >= 1.4 for c in pal)
    assert all(delta_e(a, b) >= 15 for a, b in zip(pal, pal[1:]))


def test_five_colour_palettes_are_pairwise_distinct_and_never_text_or_background():
    """The four templates' own inputs (theme accents incl. lt2/dk2, manifest colours) on their grounds: five slices
    stay ≥ ΔE 20 apart pairwise, and a near-black or near-white grey never becomes a slice."""
    cases = [
        ("FE095F", ["FFD6E3", "FF0053", "FC3777", "8A83D1", "520977", "2D1451", "753A93", "976BAE", "F92571"], ("F4F5F6", "702D8D")),
        ("0077FF", ["FFFFFF", "FF3885", "202020", "F7FAFE", "D6ECFF", "00D3E6", "005EFF", "FF6C6C", "0563C1"], ("FAFCFF",)),
        ("0077FF", ["FF3885", "7CEDF8", "EBF3F9", "D8FAFD", "FFEBF3", "3782BA", "C3A3E2", "8DBBDD"], ("FFFFFF",)),
        ("0077FF", ["00E9FF", "AAFBFF", "EDF3FC", "7C8A9A", "202020", "00AEE8", "EE5959", "6DBCFF"], ("000000",)),
    ]
    for accent, extra, grounds in cases:
        for ground in grounds:
            pal = chart_palette([accent], ground, 5, extra=extra)
            assert len(set(pal)) == 5, (ground, pal)
            assert all(delta_e(a, b) >= 20 for i, a in enumerate(pal) for b in pal[i + 1 :]), (ground, pal)
            assert all(_is_data_color(c) for c in pal), (ground, pal)
            assert "202020" not in pal and "EDF3FC" not in pal
    assert not _is_data_color("202020") and not _is_data_color("EDF3FC") and _is_data_color("808080")
    assert _lightness("FFFFFF") > 92


def test_muted_bars_are_a_tint_of_the_accent_not_grey():
    for accent, ground in (("0077FF", "FFFFFF"), ("0077FF", "000000"), ("FE095F", "F4F5F6"), ("FE095F", "702D8D"), ("0077FF", "FAFCFF")):
        m = muted_tint(accent, ground)
        assert contrast_ratio(m, ground) >= 1.4, (accent, ground, m)
        assert delta_e(m, accent) >= 20, (accent, ground, m)
        assert _chroma(m) > 12, "muted bars keep the brand hue"
        if relative_luminance(ground) > 0.5:
            assert relative_luminance(m) > relative_luminance(accent)
        else:
            assert relative_luminance(m) < relative_luminance(accent)
    # on a dark ground the stepped-back bars keep body instead of sinking into it (was 1.67:1 on black)
    assert contrast_ratio(muted_tint("0077FF", "000000"), "000000") >= 2.0
    assert contrast_ratio(muted_tint("FE095F", "702D8D"), "702D8D") > 1.45


def test_baseline_rule_is_a_quiet_line_on_the_texts_side_of_the_ground():
    # a divider barely darker than a black ground is invisible; black on a purple ground is not a hairline
    for neutral, text, ground in (("212121", "FFFFFF", "000000"), ("000000", "FFFFFF", "702D8D")):
        r = _rule_color(neutral, text, ground)
        assert r != neutral and contrast_ratio(r, ground) >= 1.45
        assert relative_luminance(r) > relative_luminance(ground)
    assert _rule_color("E1E1E1", "1A1A1A", "FFFFFF") == "E1E1E1"  # the template's own divider when it fits


# ---------------------------------------------------------------------------------------------- units and formats


def test_unit_once_as_a_short_suffix_and_zero_hidden():
    assert short_unit("тыс.") == "тыс."
    assert short_unit("млн ₽ в год") == "млн ₽"
    assert short_unit("тыс. обращений") == "тыс."
    fmt = label_format("тыс.", [12, 31, 58])
    assert fmt == '#,##0" тыс.";"−"#,##0" тыс.";'  # true minus; empty zero section: a zero prints nothing
    assert '0.0"%"' in label_format("%", [4.5, 12])
    assert "в год" not in label_format("млн ₽ в год", [42, 12.5]) and "₽" in label_format("млн ₽ в год", [42])
    assert unit_caption(ChartSpec(type="column", series_ids=["long"]), OUTLINE) == "млн ₽ в год"
    # a word unit, even a short one, is written once above the chart: «12 / 31 / 58 тыс.» reads as two scales
    assert unit_caption(ChartSpec(type="column", series_ids=["cc"]), OUTLINE) == "тыс."
    assert unit_caption(ChartSpec(type="pie", series_ids=["cc"]), OUTLINE) is None
    assert number_format("%") == '0"%"'  # the old helper keeps its contract


def test_word_unit_is_written_once_above_and_every_label_is_a_plain_number():
    cs = _chart(ChartSpec(type="column", series_ids=["ws"]))
    ser = cs.find(f".//{C}barChart/{C}ser")
    assert "чел." not in _ser_fmt(ser)
    with_unit = [i for i, d in _point_labels(ser).items() if d.find(f"{C}numFmt") is not None and "чел." in d.find(f"{C}numFmt").get("formatCode")]
    assert with_unit == []  # no label carries the unit alone
    assert "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t")) == "чел."
    # the chart-type default keeps the unit (editors, audit)
    assert "чел." in cs.find(f".//{C}barChart/{C}dLbls/{C}numFmt").get("formatCode")
    # a one-glyph unit rides on every label
    cs = _chart(ChartSpec(type="column", series_ids=["a", "b"]))
    assert all('"%"' in _ser_fmt(s) for s in cs.iter(f"{C}ser"))


def test_long_unit_goes_to_a_caption_or_the_title_or_stays_with_the_caller():
    cs = _chart(ChartSpec(type="column", series_ids=["long"]))
    assert "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t")) == "млн ₽ в год"
    ser = cs.find(f".//{C}barChart/{C}ser")
    assert "₽" not in _ser_fmt(ser) and all("₽" not in (d.find(f"{C}numFmt").get("formatCode") if d.find(f"{C}numFmt") is not None else "") for d in _point_labels(ser).values())
    # a real title takes the unit after a comma
    cs = _chart(ChartSpec(type="column", series_ids=["cc"], title="Обращения"))
    assert "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t")) == "Обращения, тыс."
    # several series with a word unit: a caption, plain labels
    cs = _chart(ChartSpec(type="column", series_ids=["web", "mob"]))
    assert "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t")) == "тыс."
    # the caller already wrote the unit right above the chart: no second caption
    prs, s = _slide()
    tb = s.shapes.add_textbox(Emu(BOX.x), Emu(BOX.y - 400000), Emu(BOX.w), Emu(300000))
    tb.text_frame.text = "млн ₽ в год"
    gf = add_chart(s, BOX, ChartSpec(type="column", series_ids=["long"]), OUTLINE, STYLE, Typography(), text_hex="1A1A1A")
    cs = gf.chart._chartSpace
    assert cs.find(f".//{C}title") is None or not "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t"))
    assert "₽" not in _ser_fmt(cs.find(f".//{C}ser"))


def test_prefer_bar_for_long_labels_or_many_categories():
    assert prefer_bar(ChartSpec(type="column", series_ids=["long"]), OUTLINE)
    assert prefer_bar(ChartSpec(type="column", series_ids=["mon"]), OUTLINE)  # 8 month names
    assert not prefer_bar(ChartSpec(type="column", series_ids=["yrs"]), OUTLINE)  # 10 short years still fit columns
    assert not prefer_bar(ChartSpec(type="column", series_ids=["cc"]), OUTLINE)
    assert not prefer_bar(ChartSpec(type="line", series_ids=["long"]), OUTLINE)
    assert effective_chart_type(ChartSpec(type="column", series_ids=["long"]), OUTLINE) == "bar"


def test_series_roles_from_names():
    assert series_roles(["Было", "Стало"]) == ["past", "lead"]
    assert series_roles(["2024", "2025", "План"]) == ["past", "lead", "plan"]
    assert series_roles(["Факт", "План"]) == ["lead", "plan"]
    assert series_roles(["Выручка", "Цель"]) == ["lead", "plan"]
    assert series_roles(["До внедрения", "После внедрения"]) == ["past", "lead"]
    assert series_roles(["Web", "Mobile"]) is None
    assert series_roles(["Москва"]) is None


# ---------------------------------------------------------------------------------------------- charts


def test_single_series_column_is_designed():
    cs = _chart(ChartSpec(type="column", series_ids=["cc"], unit="тыс.", highlight_index=2), neutral_hex="E1E1E1")
    bar = cs.find(f".//{C}barChart")
    # no value axis, no gridlines; category labels drawn, horizontal, readable
    va, ca = cs.find(f".//{C}valAx"), cs.find(f".//{C}catAx")
    assert va.find(f"{C}delete").get("val") in ("1", "true")
    assert va.find(f"{C}majorGridlines") is None
    assert ca.find(f"{C}delete").get("val") in ("0", "false")
    assert ca.find(f"{C}tickLblSkip").get("val") == "1"
    assert ca.find(f"{C}txPr/{A}bodyPr").get("rot") == "0"
    assert int(ca.find(f".//{A}defRPr").get("sz")) >= 1050
    assert 60 <= int(bar.find(f"{C}gapWidth").get("val")) <= 80
    # highlight in accent.1, the others its tint toward the ground
    ser = bar.find(f"{C}ser")
    assert _fill(ser) == muted_tint("FE095F", "FFFFFF")
    dpt = ser.find(f"{C}dPt")
    assert dpt.find(f"{C}idx").get("val") == "2" and _fill(dpt) == "FE095F"
    # the highlight label is set in the accent (a template without bold steps it up a size instead); the word unit
    # is written once, above the chart, and every label is a plain number of that one scale
    plot_lbls = bar.find(f"{C}dLbls")
    assert "тыс." in plot_lbls.find(f"{C}numFmt").get("formatCode") and plot_lbls.find(f"{C}showVal").get("val") == "1"
    hl = _point_labels(ser)[2]
    assert "тыс." not in hl.find(f"{C}numFmt").get("formatCode")
    assert "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t")) == "тыс."
    hl_rpr = hl.find(f".//{A}defRPr")
    assert hl_rpr.get("b") == "0" and hl_rpr.find(f"{A}solidFill/{A}srgbClr").get("val") == "FE095F"
    assert int(hl_rpr.get("sz")) > int(ser.find(f"{C}dLbls/{C}txPr//{A}defRPr").get("sz"))
    # transparent chart, plot area laid out by hand, axis ids LibreOffice can read (signed 32-bit)
    assert cs.find(f"{C}spPr/{A}noFill") is not None and cs.find(f".//{C}plotArea/{C}spPr/{A}noFill") is not None
    ml = cs.find(f".//{C}plotArea/{C}layout/{C}manualLayout")
    assert ml is not None and ml.find(f"{C}layoutTarget").get("val") == "inner"
    assert all(0 <= int(e.get("val")) < 2**31 for e in cs.iter(f"{C}axId", f"{C}crossAx"))


def test_bold_emphasis_only_where_the_template_speaks_in_bold():
    cs = _chart(ChartSpec(type="column", series_ids=["cc"], highlight_index=2), typo=BOLD)
    ser = cs.find(f".//{C}barChart/{C}ser")
    assert _point_labels(ser)[2].find(f".//{A}defRPr").get("b") == "1"
    cs = _chart(ChartSpec(type="line", series_ids=["ln"]), typo=Typography())
    assert all(d.find(f".//{A}defRPr").get("b") != "1" for d in cs.iter(f"{C}dLbl") if d.find(f".//{A}defRPr") is not None)


def test_sizes_come_from_the_template_scale():
    """A 720×405 pt slide with a 7/9/12/24/47 scale and a 7 pt chart style: labels at 9 (on the scale, the body
    size), values a step up on the scale — not an off-scale 10.5."""
    typo = Typography(scale=[TypeStep(role="display", size_pt=47), TypeStep(role="h1", size_pt=24), TypeStep(role="h2", size_pt=12), TypeStep(role="body", size_pt=9), TypeStep(role="small", size_pt=7)])
    prs, s = _slide("FFFFFF", 9144000, 5143500)
    box = Bbox(x=500000, y=1300000, w=5000000, h=3200000)
    gf = add_chart(s, box, ChartSpec(type="column", series_ids=["cc"], highlight_index=2), OUTLINE, ChartStyleSpec(series_colors=["0077FF"], font_size_pt=7.0), typo, text_hex="000000")
    cs = gf.chart._chartSpace
    assert cs.find(f".//{C}catAx//{A}defRPr").get("sz") == "900"
    assert cs.find(f".//{C}ser/{C}dLbls/{C}txPr//{A}defRPr").get("sz") == "1200"


def test_long_labels_turn_horizontal_top_down_with_a_real_zero_said():
    cs = _chart(ChartSpec(type="column", series_ids=["long"], highlight_index=0))
    assert cs.find(f".//{C}barDir").get("val") == "bar"
    assert cs.find(f".//{C}catAx/{C}scaling/{C}orientation").get("val") == "maxMin"
    ser = cs.find(f".//{C}ser")
    zero = _point_labels(ser)[1]
    # a bar row with a real zero says «0» at the axis in the muted colour instead of looking like missing data
    assert zero.find(f"{C}delete") is None and zero.find(f"{C}numFmt").get("formatCode") == "0"
    fmt = cs.find(f".//{C}barChart/{C}dLbls/{C}numFmt").get("formatCode")
    assert "₽" in fmt and "в год" not in fmt
    # columns still hide a zero
    cs = _chart(ChartSpec(type="column", series_ids=["zero"]))
    assert cs.find(f".//{C}barDir").get("val") == "col"
    assert _point_labels(cs.find(f".//{C}ser"))[1].find(f"{C}delete").get("val") == "1"


def test_story_series_paint_the_latest_in_the_accent():
    cs = _chart(ChartSpec(type="column", series_ids=["a", "b", "c"]), bg="F4F5F6")
    sers = list(cs.iter(f"{C}ser"))
    assert _fill(sers[1]) == "FE095F"  # 2025, not 2024, takes the brand colour
    assert _fill(sers[0]) == muted_tint("FE095F", "F4F5F6")  # the earlier year steps back
    plan = sers[2].find(f"{C}spPr")
    assert plan.find(f"{A}noFill") is not None and plan.find(f"{A}ln/{A}solidFill/{A}srgbClr").get("val") == "FE095F"
    assert cs.find(f".//{C}legend/{C}legendPos").get("val") == "b"
    # names without a story keep distinct categorical colours
    cs = _chart(ChartSpec(type="column", series_ids=["web", "mob"]), bg="F4F5F6")
    cols = [_fill(s) for s in cs.iter(f"{C}ser")]
    assert cols[0] == "FE095F" and delta_e(cols[0], cols[1]) >= 20


def test_horizontal_series_are_named_on_the_first_row_not_in_a_legend():
    cs = _chart(ChartSpec(type="bar", series_ids=["was", "now"]))
    assert cs.find(f".//{C}legend") is None
    for ser in cs.iter(f"{C}ser"):
        first = _point_labels(ser)[0]
        assert first.find(f"{C}showSerName").get("val") == "1" and first.find(f"{C}separator").text == ": "
    assert _fill(list(cs.iter(f"{C}ser"))[1]) == "FE095F"  # Стало in the accent


def test_line_is_an_accent_line_with_markers_and_labels_off_the_line():
    cs = _chart(ChartSpec(type="line", series_ids=["ln"]))
    ser = cs.find(f".//{C}lineChart/{C}ser")
    assert ser.find(f"{C}spPr/{A}ln").get("w") == str(int(2.25 * 12700))
    assert ser.find(f"{C}marker/{C}symbol").get("val") == "circle"
    assert cs.find(f".//{C}lineChart/{C}dLbls/{C}dLblPos").get("val") == "t"
    assert "0.0" in cs.find(f".//{C}lineChart/{C}dLbls/{C}numFmt").get("formatCode")  # data with decimals keeps them
    # a steady rise: the labels of the rising points move off the line (nudged or re-sided), leader lines off
    cs = _chart(ChartSpec(type="line", series_ids=["rise"]))
    ser = cs.find(f".//{C}lineChart/{C}ser")
    lbls = _point_labels(ser)
    moved = [i for i, d in lbls.items() if d.find(f"{C}layout/{C}manualLayout") is not None or d.find(f"{C}dLblPos").get("val") != "t"]
    assert len(moved) >= 2
    assert ser.find(f"{C}dLbls/{C}showLeaderLines").get("val") == "0"
    # no value axis over a raised minimum: no baseline pretending to be zero
    assert cs.find(f".//{C}catAx/{C}spPr/{A}ln/{A}noFill") is not None


def test_label_placement_keeps_labels_off_the_segments():
    pts = [(50.0 + 60 * i, 300.0 - 40 * i) for i in range(6)]  # a steady rise
    sizes = [(40.0, 14.0)] * 6
    placed = place_point_labels(pts, sizes, (0, 0, 500, 400), 6.0)
    segs = list(zip(pts, pts[1:]))
    for (x, y), (pos, dx), (w, h) in zip(pts, placed, sizes):
        if pos == "t":
            box = (x + dx - w / 2, y - 6 - h, x + dx + w / 2, y - 6)
        elif pos == "b":
            box = (x + dx - w / 2, y + 6, x + dx + w / 2, y + 6 + h)
        elif pos == "r":
            box = (x + 6, y - h / 2, x + 6 + w, y + h / 2)
        else:
            box = (x - 6 - w, y - h / 2, x - 6, y + h / 2)
        assert sum(_seg_in_box(a, b, box) for a, b in segs) < 1.0, (x, y, pos, dx)


def test_several_lines_are_named_at_their_ends_on_a_round_axis():
    cs = _chart(ChartSpec(type="line", series_ids=["web", "mob"]))
    assert cs.find(f".//{C}legend") is None
    for ser in cs.iter(f"{C}ser"):
        end = _point_labels(ser)[2]
        assert end.find(f"{C}showSerName").get("val") == "1" and end.find(f"{C}dLblPos").get("val") == "r"
    va = cs.find(f".//{C}valAx")
    step = float(va.find(f"{C}majorUnit").get("val"))
    lo, hi = float(va.find(f"{C}scaling/{C}min").get("val")), float(va.find(f"{C}scaling/{C}max").get("val"))
    assert 3 <= round((hi - lo) / step) <= 7 and abs(lo / step - round(lo / step)) < 1e-6 and abs(hi / step - round(hi / step)) < 1e-6
    # 12 months: a round step, not LibreOffice's automatic 70, 90 … 210
    cs = _chart(ChartSpec(type="line", series_ids=["m12"]))
    va = cs.find(f".//{C}valAx")
    step = float(va.find(f"{C}majorUnit").get("val"))
    assert step in (10, 20, 25, 50) and float(va.find(f"{C}scaling/{C}max").get("val")) >= 196


def test_area_is_a_soft_fill_under_an_edge_line():
    cs = _chart(ChartSpec(type="area", series_ids=["ln"]))
    area = cs.find(f".//{C}areaChart")
    sp = area.find(f"{C}ser/{C}spPr")
    assert sp.find(f"{A}ln/{A}noFill") is not None  # no outline around the polygon
    assert sp.find(f"{A}solidFill/{A}srgbClr/{A}alpha") is not None
    edge = area.getnext()
    assert edge.tag == f"{C}lineChart"
    eser = edge.find(f"{C}ser")
    assert eser.find(f"{C}spPr/{A}ln").get("w") == str(int(2.25 * 12700))
    assert [a.get("val") for a in edge.findall(f"{C}axId")] == [a.get("val") for a in area.findall(f"{C}axId")]
    end = _point_labels(eser)[3]
    assert end.find(f"{C}dLblPos").get("val") == "r" and "п.п." not in end.find(f"{C}numFmt").get("formatCode")
    assert "п.п." in "".join(t.text for t in cs.find(f".//{C}title").iter(f"{A}t"))  # the unit once, above


def test_pie_shows_percentages_without_leader_lines():
    cs = _chart(ChartSpec(type="doughnut", series_ids=["pie"]))
    dl = cs.find(f".//{C}doughnutChart/{C}dLbls")
    assert '"%"' in dl.find(f"{C}numFmt").get("formatCode")
    assert dl.find(f"{C}showLeaderLines").get("val") == "0"
    shares = [float(v.text) for v in cs.find(f".//{C}ser/{C}val").iter(f"{C}v")]
    assert abs(sum(shares) - 100) < 0.5
    fills = [_fill(p) for p in cs.iter(f"{C}dPt")]
    assert all(delta_e(a, b) >= 20 for a, b in zip(fills, fills[1:]))


def test_pie_labels_that_do_not_fit_their_slice_go_outside_and_other_is_quiet():
    cs = _chart(ChartSpec(type="pie", series_ids=["tiny"]))
    lbls = _point_labels(cs.find(f".//{C}pieChart/{C}ser"))
    assert lbls[2].find(f"{C}dLblPos").get("val") == "outEnd"
    assert lbls[0].find(f"{C}dLblPos").get("val") in ("ctr", "inEnd")
    cs = _chart(ChartSpec(type="doughnut", series_ids=["dn"]))
    fills = [_fill(p) for p in cs.iter(f"{C}dPt")]
    assert _chroma(fills[4]) < 12  # «Прочее» without a hue of its own
    assert all(_chroma(f) >= 12 for f in fills[:4])


def test_doughnut_label_wider_than_the_ring_at_its_angle_is_left_to_the_legend():
    # «110 000 ₽» at nine o'clock lies across the ring: set there it runs over the edge into the total in the hole
    team = Series(id="team", name="Бюджет команды", categories=["Оплата работы тренеров", "Администраторов", "Уборки", "Резерв на замены"], values=[260000, 110000, 30000, 20000], unit="₽")
    outline = OUTLINE.model_copy(update={"series": [*OUTLINE.series, team]})
    prs, s = _slide()
    box = Bbox(x=int(W * 0.05), y=int(H * 0.25), w=int(W * 0.3), h=int(H * 0.55))
    gf = add_chart(s, box, ChartSpec(type="doughnut", series_ids=["team"]), outline, STYLE, Typography(), text_hex="1A1A1A", legend=False, amounts=True)
    lbls = _point_labels(gf.chart._chartSpace.find(f".//{C}doughnutChart/{C}ser"))
    assert lbls[1].find(f"{C}delete") is not None
    # short shares keep their labels wherever they sit, three o'clock included
    cs = _chart(ChartSpec(type="doughnut", series_ids=["pie"]))
    assert all(d.find(f"{C}delete") is None for d in _point_labels(cs.find(f".//{C}doughnutChart/{C}ser")).values())


def test_ground_is_read_from_the_slide_and_drives_the_muted_colour():
    prs, s = _slide("112233")
    assert slide_ground(s, BOX) == "112233"
    card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(BOX.x - 100000), Emu(BOX.y - 100000), Emu(BOX.w + 200000), Emu(BOX.h + 200000))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    assert slide_ground(s, BOX) == "FFFFFF"  # the card under the chart, not the slide
    prs2, s2 = _slide("112233")
    gf = add_chart(s2, BOX, ChartSpec(type="column", series_ids=["cc"], highlight_index=2), OUTLINE, STYLE, Typography(), text_hex="FFFFFF")
    assert _fill(gf.chart._chartSpace.find(f".//{C}ser")) == muted_tint("FE095F", "112233")
    # an explicit ground wins, and on an accent-coloured ground the text colour leads (white bars on brand blue)
    prs3, s3 = _slide("FFFFFF")
    gf3 = add_chart(s3, BOX, ChartSpec(type="column", series_ids=["cc"], highlight_index=2), OUTLINE, ChartStyleSpec(series_colors=["0077FF"]), Typography(), text_hex="FFFFFF", ground_hex="0077FF")
    dpt = gf3.chart._chartSpace.find(f".//{C}dPt")
    assert _fill(dpt) == "FFFFFF"


def test_every_chart_run_names_latin_ea_and_cs_typefaces():
    for spec in (ChartSpec(type="column", series_ids=["cc"], title="Обращения"), ChartSpec(type="doughnut", series_ids=["pie"])):
        cs = _chart(spec)
        latins = list(cs.iter(f"{A}latin"))
        assert latins
        for latin in latins:
            ea, csf = latin.getnext(), latin.getnext().getnext()
            assert ea.tag == f"{A}ea" and csf.tag == f"{A}cs" and ea.get("typeface") == latin.get("typeface")


def test_doughnut_hole_box_is_centred_on_the_plot():
    prs, s = _slide()
    gf = add_chart(s, BOX, ChartSpec(type="doughnut", series_ids=["pie"]), OUTLINE, STYLE, Typography(), text_hex="1A1A1A")
    hole = doughnut_hole_bbox(gf)
    assert hole is not None and BOX.x < hole.x < BOX.x + BOX.w and BOX.y < hole.y < BOX.y + BOX.h
    assert math.isclose(hole.w, hole.h, rel_tol=0.01)


def test_generated_charts_pass_the_label_audit(tmp_path):
    prs, s = _slide("FFFFFF")
    specs = [
        ChartSpec(type="column", series_ids=["cc"], highlight_index=2),
        ChartSpec(type="column", series_ids=["long"]),
        ChartSpec(type="column", series_ids=["a", "b", "c"]),
        ChartSpec(type="bar", series_ids=["was", "now"]),
        ChartSpec(type="line", series_ids=["ln"]),
        ChartSpec(type="line", series_ids=["a", "b"]),
        ChartSpec(type="area", series_ids=["ln"]),
        ChartSpec(type="pie", series_ids=["pie"]),
        ChartSpec(type="doughnut", series_ids=["pie"]),
    ]
    for spec in specs:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        add_chart(sl, BOX, spec, OUTLINE, STYLE, Typography(), text_hex="1A1A1A")
    path = tmp_path / "charts.pptx"
    prs.save(str(path))
    ir = build_deck_ir(path, with_images=False)
    charts = [e for sl in ir.slides for e in sl.elements if e.type == "chart"]
    assert len(charts) == len(specs)
    issues = chart_missing_labels(AuditContext(ir=ir, manifest=TemplateManifest.model_construct()))
    assert issues == [], [i.message for i in issues]


# ---------------------------------------------------------------------------------------------- the web version


def test_svg_export_draws_the_same_design(tmp_path):
    prs, s = _slide("FFFFFF")
    add_chart(s, BOX, ChartSpec(type="column", series_ids=["ws"], highlight_index=4), OUTLINE, STYLE, Typography(), text_hex="1A1A1A")
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    add_chart(sl, BOX, ChartSpec(type="line", series_ids=["web", "mob"]), OUTLINE, STYLE, Typography(), text_hex="1A1A1A")
    path = tmp_path / "c.pptx"
    prs.save(str(path))
    ir = build_deck_ir(path, with_images=False)
    col, line = [e.chart for sl in ir.slides for e in sl.elements if e.type == "chart"]
    assert col.series[0].point_colors[4] == "FE095F" and col.series[0].color == muted_tint("FE095F", "FFFFFF")
    svg = chart_svg(col, 600, 300, ["0077FF"])
    assert 'fill="#FE095F"' in svg and f'fill="#{muted_tint("FE095F", "FFFFFF")}"' in svg  # highlight kept
    assert ">12 400<" in svg and svg.count("чел.") == 1  # the unit once (the caption above), Russian grouping
    assert ">1 200<" in svg and "rx=" not in svg  # plain labels, square bars
    assert line.series_labels and "Mobile: 16" in chart_svg(line, 600, 300, ["0077FF"])
    assert format_value(12.5, '0.0"%"') == "12,5%" and format_value(-3, '#,##0;"−"#,##0;') == "−3" and format_value(0, "0;0;") == ""
