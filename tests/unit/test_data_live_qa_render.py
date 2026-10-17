"""Живой прогон второй ступени (live QA), вёрстка: годы не выделяются акцентом, «чел. ,» не бывает, вывод слайда с
данными говорит первым предложением, а не уходит в заметки; рисунок на подложке-картинке (карандаш и ластик Blueprint)
находится по пикселям и не ложится под заголовок и диаграмму; подсказки плейсхолдеров макета не видны в LibreOffice;
рисунок образца о своём (карта США) не встаёт рядом с данными колоды."""

from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu

from verstka.rendering.compose import _first_sentence, highlight_runs, label_without_figure
from verstka.rendering.cutouts import cutout_boxes, sheet_box
from verstka.schemas.common import Bbox

ROOT = Path(__file__).resolve().parents[2]


# ── B3 / L6: typography of the composed words ──────────────────────────────────────────────────────────────────────
def test_a_year_is_not_set_in_the_accent():
    runs = highlight_runs("Сентябрь 2026 к сентябрю 2025: +23%, 423 млн ₽", 12, "000000", "FF0000")
    accent = [r.text.strip() for r in runs if r.color == "FF0000"]
    assert accent == ["+23%", "423 млн ₽"]


def test_a_figure_cut_out_of_brackets_leaves_no_space_before_the_comma():
    got = label_without_figure("Принято на 1 курс — 2 425 чел. (+6%), отчислено студентов — 412 чел. (−8%). Сравнение с 2025.", "+6%")
    assert "чел., отчислено" in got and " ," not in got


def test_the_first_sentence_of_a_conclusion():
    assert _first_sentence("Остальные 230 клиентов — 2 167 млн ₽, 62%. Крупнейший — АО «Северный капитал».") == "Остальные 230 клиентов — 2 167 млн ₽, 62%"
    assert _first_sentence("Выручка выросла на 23%") == ""
    assert _first_sentence("Рост 5 млн. Это хорошо.") == ""  # «млн.» ends no sentence


# ── L1: the objects a picture ground draws on its sheet ───────────────────────────────────────────────────────────
def _ground(shade: bool = False) -> bytes:
    """A blue grid ground with a white sheet; a yellow «pencil» and a pale «eraser» lying on its top right corner, or
    (`shade`) only a soft grey shade in its top left corner."""
    im = Image.new("RGB", (1600, 900), (47, 136, 167))
    d = ImageDraw.Draw(im)
    d.rectangle([60, 50, 1540, 860], fill=(255, 255, 255))
    if shade:
        for k in range(300):
            g = 255 - int(40 * (1 - k / 300))
            d.line([(60 + k, 50), (60, 50 + k)], fill=(g, g, g), width=2)
    else:
        d.polygon([(1150, 190), (1180, 160), (1500, 0), (1530, 30)], fill=(240, 200, 60), outline=(180, 140, 20))
        d.rounded_rectangle([1380, 120, 1520, 230], radius=20, fill=(250, 222, 210), outline=(225, 170, 150), width=4)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _slide_on(png: bytes):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
    tree = slide.shapes._spTree
    tree.remove(pic._element)
    tree.insert(2, pic._element)  # under everything: the ground
    return prs, slide


def test_the_pencil_and_the_eraser_are_found_and_a_shade_is_not():
    prs, slide = _slide_on(_ground())
    W, H = int(prs.slide_width), int(prs.slide_height)
    area = Bbox(x=int(0.05 * W), y=int(0.03 * H), w=int(0.9 * W), h=int(0.9 * H))
    cuts = cutout_boxes(slide, area, pad=0)
    assert cuts and all(c.x >= 0.65 * W and c.y2 <= 0.3 * H for c in cuts)
    assert any(c.x <= 0.75 * W for c in cuts) and any(c.x2 >= 0.93 * W for c in cuts)  # the pencil and the eraser
    sheet = sheet_box(slide, area)
    assert sheet is not None and abs(sheet.x - 0.0375 * W) < 0.01 * W and abs(sheet.y - 0.056 * H) < 0.015 * H
    prs2, slide2 = _slide_on(_ground(shade=True))
    assert cutout_boxes(slide2, area, pad=0) == []


def test_a_plain_slide_has_no_cutouts():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    assert cutout_boxes(slide) == [] and sheet_box(slide) is None


# ── L5: the layouts' prompts are not drawn ────────────────────────────────────────────────────────────────────────
def test_the_layout_prompts_are_emptied_but_not_the_fields():
    from verstka.rendering.renderer import _quiet_layout_prompts

    prs = Presentation()
    layout = prs.slide_layouts[1]
    for ph in layout.placeholders:
        if ph.has_text_frame and not ph.text_frame.text:
            ph.text_frame.text = "Альтернативные решения"
    prs.slides.add_slide(layout)
    fields = {ph.placeholder_format.idx: ph.text_frame.text for ph in layout.placeholders
              if (ph._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}ph").get("type") or "body") in ("dt", "ftr", "sldNum")}
    n = _quiet_layout_prompts(SimpleNamespace(prs=prs))
    assert n >= 2
    for ph in layout.placeholders:
        kind = ph._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}ph").get("type") or "body"
        if kind in ("dt", "ftr", "sldNum"):
            assert ph.text_frame.text == fields[ph.placeholder_format.idx]
        else:
            assert ph.text_frame.text.strip() == ""


# ── B4: a sample's drawing of its own subject ─────────────────────────────────────────────────────────────────────
def test_a_drawing_of_many_outlines_is_a_subject_and_a_few_shapes_are_decor():
    from verstka.rendering.clone import _subject_drawing

    def ctx_with(n_parts: int):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shapes = []
        for k in range(n_parts):
            fb = slide.shapes.build_freeform(Emu(900000 + (k % 8) * 500000), Emu(1200000 + (k // 8) * 500000))
            fb.add_line_segments([(Emu(1300000 + (k % 8) * 500000), Emu(1300000 + (k // 8) * 500000)), (Emu(1000000 + (k % 8) * 500000), Emu(1600000 + (k // 8) * 500000))])
            shapes.append(fb.convert_to_shape())
        slide.shapes.add_group_shape(shapes)
        W, H = int(prs.slide_width), int(prs.slide_height)
        return SimpleNamespace(slide=slide, W=W, H=H, pattern=SimpleNamespace(chrome_shape_ids=set()), id_of=lambda el: None)

    assert _subject_drawing(ctx_with(40)) is not None
    assert _subject_drawing(ctx_with(5)) is None
