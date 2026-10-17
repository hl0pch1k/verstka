"""qa4 composer round: the typographic floor of every slide (a key figure 1.6 times its label, running text never
under the template's body size), labels on chips, lines and marks drawn inside the template's ground picture, the
chart shown as a table of ten rows, and the audit checks that make a score of 100 mean what it says."""

from __future__ import annotations

import io

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from verstka.analysis.manifest import analyze_template
from verstka.audit.ir import build_deck_ir
from verstka.audit.registry import AuditContext
from verstka.ingest.workspace import TemplateWorkspace
from verstka.rendering.clone_check import FIG_RATIO, SMALL_SHARE, body_floor, closing_trouble, figure_pairs, is_index_number, small_floor
from verstka.rendering.compose import Kit, bare_label, distinct_label, lead_figures, says_again, unrepeat_item
from verstka.rendering.deck import DeckBuilder
from verstka.rendering.picture_lines import cell_marks, ground_lines, picture_lines
from verstka.rendering.synth import _kind_label, render_synth
from verstka.schemas.common import Bbox, PatternKind
from verstka.schemas.layout import LayoutSlide
from verstka.schemas.outline import ChartSpec, DeckOutline, InlineSeries, NumberCallout, OutlineSlide, SlideContent, SlideItem

W, H = 12192000, 6858000


@pytest.fixture
def env(simple_deck, tmp_path):
    manifest = analyze_template(simple_deck, workspace_root=tmp_path / "ws", use_llm=False, use_vlm=False, render=False)
    ws = TemplateWorkspace.open(manifest.template_id, tmp_path / "ws")
    return simple_deck, manifest, ws


def _render(env, oslide: OutlineSlide, comp: str, strategy: str = "structured"):
    deck, manifest, ws = env
    b = DeckBuilder(deck)
    outline = DeckOutline(title="T", strategy=strategy, slides=[oslide])
    return render_synth(b, LayoutSlide(outline_id=oslide.id, mode="synth", composition=comp), oslide, manifest, ws, outline)


def _sizes(sh) -> list[float]:
    return [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()]


def _digit_size(sh) -> float:
    return max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and any(ch.isdigit() for ch in r.text)), default=0.0)


# ---- the floor ------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("strategy", ["structured", "visual", "compact"])
def test_every_figure_of_a_kpi_row_is_at_least_1_6_times_its_label(env, strategy):
    nums = [NumberCallout(value="1,24 млрд ₽", label="Выручка, +23% к III кв. 2025"), NumberCallout(value="3 480", label="Клиенты, +18% к III кв. 2025"),
            NumberCallout(value="356 тыс. ₽", label="Средний чек, +4% к III кв. 2025"), NumberCallout(value="52", label="NPS, +6 п. к III кв. 2025")]
    o = OutlineSlide(id="k", kind=PatternKind.stat_row, headline="Выручка растёт быстрее числа клиентов: +23% против +18%", content=SlideContent(numbers=nums))
    slide, _ = _render(env, o, "stat_row", strategy)
    figs = [sh for sh in slide.shapes if sh.name.startswith("Figure")]
    labs = [sh for sh in slide.shapes if sh.name.startswith("Label")]
    assert len(figs) == len(labs) == 4
    _, manifest, _ = env
    for f, lab in zip(figs, labs):
        assert _digit_size(f) >= FIG_RATIO * max(_sizes(lab)) - 0.3, (f.text_frame.text, _digit_size(f), max(_sizes(lab)))
        assert min(_sizes(lab)) >= body_floor(manifest) - 0.05


def test_before_and_after_lead_with_their_figures(env):
    cols = [SlideItem(title="Было", text="31% срывов дедлайнов"), SlideItem(title="Стало", text="12% срывов дедлайнов")]
    o = OutlineSlide(id="c", kind=PatternKind.comparison, headline="Срывы дедлайнов снизились с 31% до 12%", content=SlideContent(columns=cols))
    slide, _ = _render(env, o, "comparison")
    by_text = {}
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    by_text[t.replace(" ", " ")] = max(r.font.size.pt for r in p.runs if r.font.size)
    assert by_text["31%"] >= FIG_RATIO * by_text["Было"] and by_text["12%"] >= FIG_RATIO * by_text["Стало"], by_text
    pairs = lead_figures(cols)
    assert [(x.number, x.text) for x in pairs] == [("31%", "срывов дедлайнов"), ("12%", "срывов дедлайнов")]
    assert lead_figures([SlideItem(title="До 1 декабря", text="Подписание соглашений"), SlideItem(title="Февраль", text="Набор")]) is None


def test_the_floors_follow_the_template(env):
    _, manifest, _ = env
    m = manifest.model_copy(deep=True)
    m.tokens.typography.scale = []
    hpt = H / 12700
    assert body_floor(m) == min(14.0, round(0.031 * hpt * 2) / 2)
    assert small_floor(m) == pytest.approx(SMALL_SHARE * hpt)
    k = Kit(m, W, H, "FFFFFF")
    assert k.small_text >= k.text_floor and k.text_floor <= k.body


def test_a_label_is_never_a_bare_per_unit():
    head = "Контроль статусов отнимает до 5 часов в неделю"
    lab = "Менеджеры тратят в неделю на контроль статусов"
    assert distinct_label("до 5 часов", lab, head) == lab
    assert bare_label("в неделю") and bare_label("на сотрудника") and not bare_label("минуты на одно обращение")


def test_a_card_that_says_its_title_again_keeps_its_words_alone():
    it = unrepeat_item(SlideItem(title="Утренний дайджест", text="присылает утренний дайджест по задачам команды"))
    assert it.title == "" and it.text.startswith("Присылает")
    keep = SlideItem(title="Практика", text="Каждый модуль заканчивается практическим проектом")
    assert unrepeat_item(keep) == keep  # a one-word label is a label
    assert not says_again("Срок по контексту", "Система предлагает дату")


# ---- chips ----------------------------------------------------------------------------------------------------------

def test_a_chip_names_the_kind_of_slide_truly():
    steps = OutlineSlide(id="n", kind=PatternKind.two_column, headline="Следующие шаги: соглашения до 1 декабря", content=SlideContent(columns=[SlideItem(title="До 1 декабря", text="a"), SlideItem(title="Февраль", text="b")]))
    assert _kind_label(steps) == "План"
    pair = OutlineSlide(id="p", kind=PatternKind.two_column, headline="Срывы дедлайнов снизились", content=SlideContent(columns=[SlideItem(title="Было", text="31%"), SlideItem(title="Стало", text="12%")]))
    assert _kind_label(pair) == "Сравнение"
    other = OutlineSlide(id="o", kind=PatternKind.two_column, headline="Практика и наставники", content=SlideContent(columns=[SlideItem(title="Практика", text="a"), SlideItem(title="Наставники", text="b")]))
    assert _kind_label(other) == ""
    assert _kind_label(OutlineSlide(id="s", kind=PatternKind.process, headline="Четыре действия")) == "Шаги"


def _chip_deck(path, label: str):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(0.04 * W)), Emu(int(0.06 * H)), Emu(int(0.3 * W)), Emu(int(0.06 * H)))
    chip.fill.solid()
    chip.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x66)
    tb = s.shapes.add_textbox(Emu(int(0.05 * W)), Emu(int(0.06 * H)), Emu(int(0.28 * W)), Emu(int(0.06 * H)))
    tb.text_frame.text = label
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    body = s.shapes.add_textbox(Emu(int(0.05 * W)), Emu(int(0.3 * H)), Emu(int(0.8 * W)), Emu(int(0.3 * H)))
    body.text_frame.text = "Текст слайда с содержимым и выводом"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    prs.save(str(path))
    return path


def test_the_audit_sees_a_headline_on_a_chip(env, tmp_path):
    from verstka.audit.checks.template import headline_in_chip

    _, manifest, ws = env
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="a", kind=PatternKind.cards, headline="Четыре главных вывода")])
    bad = build_deck_ir(_chip_deck(tmp_path / "bad.pptx", "ЧЕТЫРЕ ГЛАВНЫХ ВЫВОДА"))
    bad.slides[0].outline_id = "a"
    got = headline_in_chip(AuditContext(ir=bad, manifest=manifest, outline=outline, ws=ws))
    assert got and got[0].check_id == "headline_in_chip"
    ok = build_deck_ir(_chip_deck(tmp_path / "ok.pptx", "ГЛАВНОЕ"))
    ok.slides[0].outline_id = "a"
    assert headline_in_chip(AuditContext(ir=ok, manifest=manifest, outline=outline, ws=ws)) == []


# ---- the ground picture: lines and marks -------------------------------------------------------------------------------

def _png(draw_fn, size=(1920, 1080)) -> bytes:
    im = Image.new("RGBA", size, (247, 248, 253, 255))
    draw_fn(ImageDraw.Draw(im))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def test_a_hairline_inside_a_ground_picture_is_found():
    def panel(d):
        d.rounded_rectangle((60, 46, 1854, 942), radius=24, fill=(255, 255, 255, 255), outline=(225, 229, 238, 255))
        d.line((958, 47, 958, 941), fill=(225, 229, 238, 255), width=2)

    lines = picture_lines(_png(panel))
    mid = [ln for ln in lines if ln[0] and abs(ln[1] - 0.499) < 0.01]
    assert mid and mid[0][3] - mid[0][2] > 0.7  # the divider down the panel

    def gradient(d):
        for x in range(1920):
            d.line((x, 0, x, 1080), fill=(200 + x // 40, 210, 240, 255))

    assert picture_lines(_png(gradient)) == []


def _ground_deck(path, png: bytes):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(png), 0, 0, Emu(W), Emu(H))
    prs.save(str(path))
    return Presentation(str(path))


def test_the_marks_a_ground_draws_in_cards_are_found(tmp_path):
    def cards(d):
        for i, x in enumerate((60, 980)):
            d.rounded_rectangle((x, 200, x + 880, 560), radius=20, fill=(255, 255, 255, 255))
            d.text((x + 600, 380), "0", fill=(210, 228, 252, 255))
            d.ellipse((x + 620, 400, x + 800, 540), outline=(210, 228, 252, 255), width=30)

    prs = _ground_deck(tmp_path / "g.pptx", _png(cards))
    s = prs.slides[0]
    cells = [Bbox(x=int(60 / 1920 * W), y=int(200 / 1080 * H), w=int(880 / 1920 * W), h=int(360 / 1080 * H)), Bbox(x=int(980 / 1920 * W), y=int(200 / 1080 * H), w=int(880 / 1920 * W), h=int(360 / 1080 * H))]
    marks = cell_marks(s, cells)
    assert all(m is not None for m in marks)
    assert marks[0].x > cells[0].x + 0.5 * cells[0].w  # the numeral at the card's right

    plain = _ground_deck(tmp_path / "p.pptx", _png(lambda d: d.rounded_rectangle((60, 200, 940, 560), radius=20, fill=(255, 255, 255, 255))))
    assert cell_marks(plain.slides[0], cells[:1]) == [None]
    assert ground_lines(plain.slides[0]) == [] or all(not ln.crosses(cells[0], inset=0.12) for ln in ground_lines(plain.slides[0]))


# ---- figures and labels on a cloned slide -------------------------------------------------------------------------------

def test_an_items_index_is_no_key_figure():
    assert is_index_number("1", 0) and is_index_number("02", 1) and is_index_number("2025", 3)
    assert not is_index_number("+23%", 0) and not is_index_number("41%", 1)
    o = OutlineSlide(id="x", kind=PatternKind.process, headline="Шаги", content=SlideContent(items=[SlideItem(title="Старт", text="a", number="1"), SlideItem(title="Итоги квартала", text="b", number="+23%")]))
    assert [p[0] for p in figure_pairs(o)] == ["+23%"]


def test_the_summary_headline_is_no_card_text():
    from types import SimpleNamespace

    from lxml import etree

    def para(text: str, size: int) -> str:
        return f'<a:p><a:r><a:rPr sz="{size * 100}"/><a:t>{text}</a:t></a:r></a:p>'

    ns = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    head = "Выручка растёт быстрее клиентов: +23% против +18%"
    xml = (f'<p:sld {ns}><p:cSld><p:spTree>'
           f'<p:sp><p:txBody>{para(head, 24)}</p:txBody></p:sp>'
           f'<p:sp><p:txBody>{para("+23%", 36)}{para("Выручка растёт быстрее клиентов", 18)}</p:txBody></p:sp>'
           f'</p:spTree></p:cSld></p:sld>')
    slide = SimpleNamespace(_element=etree.fromstring(xml), __dict__={})
    slide.__dict__["_verstka_closing"] = [("+23%", "Выручка растёт быстрее клиентов", "")]
    assert closing_trouble(slide, 14.0, head) is None
    assert closing_trouble(slide, 14.0) is not None  # without the headline its 24 pt reads as a card's text


# ---- a chart shown as a table -----------------------------------------------------------------------------------------

def test_a_chart_shown_as_a_table_keeps_ten_rows_with_totals():
    """qa4 NEW-H1: a stacked bar of 13 rows that did not fit became a raw 13×4 table at 5–6 pt with no totals — now the
    largest categories, the rest as «Прочие», and «Всего» with each series' share (the headline's «53%»)."""
    from types import SimpleNamespace

    from verstka.rendering.compose import Composer as _C

    cats = [f"Направление {i}" for i in range(1, 14)]
    spec = ChartSpec(type="stacked_bar", categories=cats, x_title="Направление", unit="чел.",
                     series=[InlineSeries(name="Бюджет", values=[60 + i for i in range(13)]), InlineSeries(name="Платное", values=[100 + 3 * i for i in range(13)]),
                             InlineSeries(name="Целевое", values=[20] * 13)])
    outline = DeckOutline(title="T", slides=[OutlineSlide(id="t", kind=PatternKind.chart, headline="Больше половины приёма — платное", content=SlideContent(chart=spec))])
    tbl = _C._chart_as_table(SimpleNamespace(outline=outline, TABLE_ROWS=10), spec)
    assert len(tbl.rows) <= 10
    assert tbl.rows[-1][0] == "Всего" and tbl.rows[-2][0].startswith("Прочие")
    paid = sum(100 + 3 * i for i in range(13))
    total = sum(60 + i for i in range(13)) + paid + 20 * 13
    assert f"({round(100 * paid / total)}%)" in tbl.rows[-1][2]  # each series' share of the whole
    assert tbl.columns[-1] == "Всего" and len(tbl.columns) <= 5
    # every category is counted: the rows shown and «Прочие» add up to the whole
    shown = sum(float(r[-1].replace("\u00a0", "").replace(" ", "").replace(",", ".")) for r in tbl.rows[:-1])
    assert shown == pytest.approx(total)
